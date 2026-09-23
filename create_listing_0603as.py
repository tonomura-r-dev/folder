import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUTPUT_PATH = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー\架電リスト_リスティング_2026-06-03as.xlsx"
HEADER = ["企業名", "LP URL", "広告種別", "法人番号", "商材", "検索KW", "業界"]
COL_WIDTHS = [32, 56, 16, 16, 22, 32, 22]
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="メイリオ", size=10, bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
ROW_FILLS = [PatternFill("solid", fgColor="F2F7FC"), PatternFill("solid", fgColor="FFFFFF")]
ROW_FONT = Font(name="メイリオ", size=9)
ROW_ALIGN = Alignment(vertical="center", wrap_text=False)
URL_FONT = Font(name="メイリオ", size=9, color="0563C1", underline="single")

DATA = [
    # ── ハイクラス転職（LP） ──
    ("株式会社リクルート（リクルートダイレクトスカウト）", "https://careercarver.jp/doc/contents/pages/cs/ad/lp001_01.html", "リスティング広告", "", "ハイクラス転職・スカウト・年収UP",         "ハイクラス 転職 スカウト, 年収 800万以上 転職, キャリアカーバー 評判", "転職エージェント"),
    # ── 名刺管理SaaS（LP） ──
    ("Sansan株式会社",   "https://jp.sansan.com/lp/workstyle/",              "リスティング広告", "", "クラウド名刺管理・営業DX・法人向け",         "名刺管理 クラウド 比較, Sansan 料金 評判, 営業 DX ツール おすすめ",     "BtoB SaaS"),
    # ── UIデザイン支援（無料相談LP） ──
    ("株式会社グッドパッチ",   "https://design-partnership.goodpatch.com/faq",    "リスティング広告", "", "UI/UXデザイン支援・プロダクト改善相談",       "UIデザイン 外注 会社, UX 設計 依頼, プロダクト デザイン 改善 相談",    "デザイン支援"),
]


def make_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "架電リスト"

    for col_idx, header in enumerate(HEADER, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
    ws.row_dimensions[1].height = 20

    for row_idx, row_data in enumerate(DATA, 2):
        fill = ROW_FILLS[(row_idx - 2) % 2]
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.fill = fill
            cell.alignment = ROW_ALIGN
            if col_idx == 2:
                cell.hyperlink = value
                cell.font = URL_FONT
            else:
                cell.font = ROW_FONT
        ws.row_dimensions[row_idx].height = 16

    for col_idx, width in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADER))}1"

    wb.save(OUTPUT_PATH)
    print(f"Saved: {OUTPUT_PATH}")
    print(f"Rows: {len(DATA)}")


if __name__ == "__main__":
    make_xlsx()
