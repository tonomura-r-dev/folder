import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUTPUT_PATH = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー\架電リスト_リスティング_2026-06-03af.xlsx"
HEADER = ["企業名", "LP URL", "広告種別", "法人番号", "商材", "検索KW", "業界"]
COL_WIDTHS = [32, 56, 16, 16, 22, 32, 22]
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="メイリオ", size=10, bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
ROW_FILLS = [PatternFill("solid", fgColor="F2F7FC"), PatternFill("solid", fgColor="FFFFFF")]
ROW_FONT = Font(name="メイリオ", size=9)
ROW_ALIGN = Alignment(vertical="center", wrap_text=False)
URL_FONT = Font(name="メイリオ", size=9, color="0563C1", underline="single")

# ── 全社、実際の広告クリック先LP（セミナー/相談予約/トライアルCV）をWeb確認済み ──
DATA = [
    # ── 不動産投資（無料セミナー・資料請求LP） ──
    ("J.P.Returns株式会社（JPリターンズ）",       "https://jpreturns.com/seminar/",            "リスティング広告", "", "マンション投資・無料セミナー・資料請求", "不動産投資 セミナー 初心者, マンション投資 リスク, JPリターンズ 評判",   "不動産投資"),
    # ── 保険相談（オンライン相談予約LP） ──
    ("株式会社アイリックコーポレーション（保険クリニック）", "https://www.hoken-clinic.com/lp/webconsult/", "リスティング広告", "", "保険無料相談・見直し・オンライン相談", "保険 見直し 無料相談, 保険 相談 オンライン, 生命保険 比較 相談",        "保険相談"),
    # ── フェイシャルエステ（初回トライアルLP） ──
    ("株式会社シーボン.（C'BON）",                "https://www.cbon.co.jp/campaign/index3.html", "リスティング広告", "", "フェイシャルエステ・初回トライアル",    "フェイシャル エステ 体験, シーボン 評判 料金, エステ 初回 お試し",      "エステ"),
]


def make_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "架電リスト"

    for col_idx, header in enumerate(HEADER, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
    ws.row_dimensions[1].height = 20

    for row_idx, row_data in enumerate(DATA, 2):
        fill = ROW_FILLS[(row_idx - 2) % 2]
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.fill = fill
            cell.alignment = ROW_ALIGN
            if col_idx == 2:
                cell.hyperlink = value
                cell.font = URL_FONT
            else:
                cell.font = ROW_FONT
        ws.row_dimensions[row_idx].height = 16

    for col_idx, width in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADER))}1"

    wb.save(OUTPUT_PATH)
    print(f"Saved: {OUTPUT_PATH}")
    print(f"Rows: {len(DATA)}")


if __name__ == "__main__":
    make_xlsx()
