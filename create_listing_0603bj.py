import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUTPUT_PATH = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー\架電リスト_リスティング_2026-06-03bj.xlsx"
HEADER = ["企業名", "LP URL", "広告種別", "法人番号", "商材", "検索KW", "業界"]
COL_WIDTHS = [32, 56, 16, 16, 22, 32, 22]
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="メイリオ", size=10, bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
ROW_FILLS = [PatternFill("solid", fgColor="F2F7FC"), PatternFill("solid", fgColor="FFFFFF")]
ROW_FONT = Font(name="メイリオ", size=9)
ROW_ALIGN = Alignment(vertical="center", wrap_text=False)
URL_FONT = Font(name="メイリオ", size=9, color="0563C1", underline="single")

DATA = [
    # ── クラウド監視SaaS ──
    ("Datadog Japan株式会社",        "https://www.datadoghq.com/free-datadog-trial/",         "リスティング広告", "", "Datadog・APM・ログ分析・14日間無料",          "Datadog 料金 評判, クラウド 監視 SaaS, APM ログ 分析 ツール",       "DevOps SaaS"),
    # ── オブザーバビリティSaaS ──
    ("New Relic株式会社",            "https://newrelic.com/jp/lp/start-new-relic-for-free-goog-da-jp", "リスティング広告", "", "New Relic・オブザーバビリティ・無料100GB", "New Relic 料金 評判, オブザーバビリティ SaaS, システム 監視 無料",   "DevOps SaaS"),
    # ── コード管理・CI/CD ──
    ("GitHub Inc.",                  "https://github.com/pricing",                             "リスティング広告", "", "GitHub・コード管理・Actions・Copilot AI",      "GitHub 料金 評判, CI/CD ツール 比較, コード 管理 SaaS",              "DevOps SaaS"),
    # ── エラー監視SaaS ──
    ("Sentry Inc.",                  "https://sentry.io/lp/ja/era-monitoringu/",               "リスティング広告", "", "Sentry・エラー監視・APM・デバッグ支援",       "エラー 監視 SaaS 比較, Sentry 料金 評判, バグ 追跡 ツール",          "DevOps SaaS"),
    # ── CI/CD SaaS ──
    ("CircleCI, LLC",                "https://circleci.com/pricing/",                          "リスティング広告", "", "CircleCI・CI/CD・継続的デリバリー・無料枠",   "CircleCI 料金 評判, CI CD ツール 比較, 自動テスト デプロイ",         "DevOps SaaS"),
    # ── IAM SaaS ──
    ("Okta Japan株式会社",           "https://www.okta.com/ja-jp/pricing/",                    "リスティング広告", "", "Okta・IAM・SSO・アクセス管理",                "Okta 料金 評判, IAM SaaS 比較, シングルサインオン 企業",              "セキュリティSaaS"),
    # ── コミュニケーションAPI ──
    ("Twilio Inc.",                  "https://www.twilio.com/en-us/pricing",                   "リスティング広告", "", "Twilio・SMS・電話API・コミュニケーション",     "Twilio 料金 評判, SMS API 比較, 通知 API SaaS",                     "コミュニケーションAPI"),
    # ── 決済API ──
    ("Stripe Japan株式会社",         "https://stripe.com/pricing",                             "リスティング広告", "", "Stripe・オンライン決済・API・クレジットカード", "Stripe 料金 評判, 決済 API SaaS, オンライン 決済 比較",              "フィンテックSaaS"),
    # ── データウェアハウスSaaS ──
    ("Snowflake Inc.",               "https://signup.snowflake.com/?_l=ja",                    "リスティング広告", "", "Snowflake・AIデータクラウド・無料トライアル", "Snowflake 料金 評判, データウェアハウス SaaS, クラウド 分析",        "データSaaS"),
    # ── SIEM/SOC SaaS ──
    ("Splunk Japan合同会社",         "https://www.splunk.com/en_us/products/pricing.html",     "リスティング広告", "", "Splunk・SIEM・ログ分析・セキュリティ",         "Splunk 料金 評判, SIEM SaaS 比較, セキュリティ ログ 分析",          "セキュリティSaaS"),
    # ── ヘッドレスCMS ──
    ("Contentful GmbH",             "https://www.contentful.com/pricing/",                    "リスティング広告", "", "Contentful・ヘッドレスCMS・APIファースト",     "ヘッドレス CMS 比較, Contentful 料金 評判, コンテンツ 管理 API",    "CMS SaaS"),
    # ── API開発プラットフォーム ──
    ("Postman Inc.",                 "https://www.postman.com/jp/",                            "リスティング広告", "", "Postman・API開発・テスト・コラボレーション",  "Postman 料金 評判, API テスト ツール, API 開発 SaaS",               "DevOps SaaS"),
    # ── フロントエンドSaaS ──
    ("Vercel Inc.",                  "https://vercel.com/pricing",                             "リスティング広告", "", "Vercel・フロントエンドデプロイ・エッジ",       "Vercel 料金 評判, デプロイ SaaS 比較, フロントエンド ホスティング", "クラウドSaaS"),
    # ── インシデント管理SaaS ──
    ("PagerDuty株式会社",            "https://www.pagerduty.com/sign-up/",                     "リスティング広告", "", "PagerDuty・インシデント管理・14日間無料",      "PagerDuty 料金 評判, インシデント 管理 SaaS, 障害 対応 ツール",     "DevOps SaaS"),
    # ── ビジネスチャット ──
    ("Lark Technologies株式会社",    "https://www.larksuite.com/paid/chat-alternative-jp",     "リスティング広告", "", "Lark・チャット+ドキュメント+会議統合",        "Lark ビジネスチャット 評判, コラボ SaaS 無料 比較, Lark 料金",       "コラボSaaS"),
    # ── ドキュメント管理 ──
    ("Notion Labs Japan合同会社",    "https://www.notion.so/ja-jp/pricing",                    "リスティング広告", "", "Notion・ドキュメント・DB・プロジェクト管理",  "Notion 料金 評判, ドキュメント 管理 SaaS, チーム wiki ツール",       "コラボSaaS"),
    # ── デザインツール ──
    ("Canva Japan株式会社",          "https://www.canva.com/ja_jp/pricing/",                   "リスティング広告", "", "Canva・グラフィックデザイン・無料プランあり", "Canva 料金 評判, デザイン ツール 無料, LP 作成 ツール",               "デザインSaaS"),
    # ── フォームSaaS ──
    ("Typeform S.L.",                "https://www.typeform.com/pricing",                       "リスティング広告", "", "Typeform・AIフォーム・3.5倍のレスポンス",     "フォーム 作成 SaaS 比較, Typeform 料金 評判, 調査 申し込み",         "フォームSaaS"),
    # ── CRM営業管理 ──
    ("Pipedrive株式会社",            "https://www.pipedrive.com/en/pricing",                   "リスティング広告", "", "Pipedrive・CRM・営業パイプライン・14日無料",   "Pipedrive 料金 評判, CRM 営業 ツール, セールス 管理 SaaS",          "CRM SaaS"),
    # ── プロダクト分析SaaS ──
    ("Amplitude Inc.",               "https://amplitude.com/ja-jp/get-started",                "リスティング広告", "", "Amplitude・AIデジタル分析・45,000社導入",      "Amplitude 料金 評判, プロダクト 分析 SaaS, デジタル アナリティクス", "プロダクト分析SaaS"),
]


def make_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "架電リスト"

    for col_idx, header in enumerate(HEADER, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
    ws.row_dimensions[1].height = 20

    for row_idx, row_data in enumerate(DATA, 2):
        fill = ROW_FILLS[(row_idx - 2) % 2]
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.fill = fill
            cell.alignment = ROW_ALIGN
            if col_idx == 2:
                cell.hyperlink = value
                cell.font = URL_FONT
            else:
                cell.font = ROW_FONT
        ws.row_dimensions[row_idx].height = 16

    for col_idx, width in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADER))}1"

    wb.save(OUTPUT_PATH)
    print(f"Saved: {OUTPUT_PATH}")
    print(f"Rows: {len(DATA)}")


if __name__ == "__main__":
    make_xlsx()
