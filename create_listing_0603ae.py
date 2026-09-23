import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUTPUT_PATH = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー\架電リスト_リスティング_2026-06-03ae.xlsx"
HEADER = ["企業名", "LP URL", "広告種別", "法人番号", "商材", "検索KW", "業界"]
COL_WIDTHS = [32, 56, 16, 16, 22, 32, 22]
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="メイリオ", size=10, bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
ROW_FILLS = [PatternFill("solid", fgColor="F2F7FC"), PatternFill("solid", fgColor="FFFFFF")]
ROW_FONT = Font(name="メイリオ", size=9)
ROW_ALIGN = Alignment(vertical="center", wrap_text=False)
URL_FONT = Font(name="メイリオ", size=9, color="0563C1", underline="single")

# ── 全社、実際の広告クリック先LP（定期/トライアル/商品購入CVページ）をWeb確認済み ──
DATA = [
    # ── マットレス（120日間フリートライアルLP） ──
    ("株式会社Morght（NELLマットレス）",    "https://nell.life/experience/trial",      "リスティング広告", "", "マットレス・120日トライアル",     "マットレス おすすめ 腰痛, NELL マットレス 評判, マットレス 通販 比較",        "寝具"),
    # ── 家庭用脱毛器（公式販売LP） ──
    ("株式会社エムロック（ケノン）",         "https://xn--rckyc9e.com/",                "リスティング広告", "", "家庭用脱毛器ケノン・光美容器", "脱毛器 家庭用 おすすめ, ケノン 口コミ 効果, 光 美容器 ランキング",          "美容家電"),
    # ── まつ毛美容液（公式商品LP） ──
    ("水橋保寿堂製薬株式会社（エマーキット）", "https://mizu-ho.com/product/emaked",     "リスティング広告", "", "まつ毛美容液 EMAKED",          "まつ毛 美容液 おすすめ, エマーキット 効果 口コミ, まつ毛 伸ばす 方法",      "化粧品"),
    # ── 健康食品（すっぽんサプリ定期LP） ──
    ("株式会社ていねい通販（すっぽん小町）", "https://www.teinei.co.jp/item/komachi/",  "リスティング広告", "", "すっぽんコラーゲンサプリ・定期", "すっぽん小町 評判 効果, コラーゲン サプリ おすすめ, すっぽん サプリ 比較", "健康食品"),
    # ── 韓国スキンケア（毛穴ケア公式EC） ──
    ("APR Japan合同会社（MEDICUBE）",       "https://themedicube.jp/",                 "リスティング広告", "", "毛穴ケアスキンケア・トナーパッド", "メディキューブ 毛穴 効果, medicube 評判, 毛穴 ケア 韓国 コスメ",          "スキンケア"),
    # ── 育毛剤（超トク定期コースLP） ──
    ("株式会社ユーピーエス（フィンジア）",   "https://finjia.jp/dis1/",                 "リスティング広告", "", "スカルプエッセンス・育毛・定期",  "育毛剤 おすすめ メンズ, フィンジア 効果 口コミ, スカルプ ケア 比較",       "育毛"),
]


def make_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "架電リスト"

    for col_idx, header in enumerate(HEADER, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
    ws.row_dimensions[1].height = 20

    for row_idx, row_data in enumerate(DATA, 2):
        fill = ROW_FILLS[(row_idx - 2) % 2]
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.fill = fill
            cell.alignment = ROW_ALIGN
            if col_idx == 2:
                cell.hyperlink = value
                cell.font = URL_FONT
            else:
                cell.font = ROW_FONT
        ws.row_dimensions[row_idx].height = 16

    for col_idx, width in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADER))}1"

    wb.save(OUTPUT_PATH)
    print(f"Saved: {OUTPUT_PATH}")
    print(f"Rows: {len(DATA)}")


if __name__ == "__main__":
    make_xlsx()
