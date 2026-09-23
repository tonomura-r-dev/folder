import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUTPUT_PATH = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー\架電リスト_リスティング_2026-05-29k.xlsx"
HEADER = ["企業名", "LP URL", "広告種別", "電話番号", "商材", "検索KW", "業界"]
COL_WIDTHS = [32, 50, 16, 16, 20, 32, 22]
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="メイリオ", size=10, bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
ROW_FILLS = [PatternFill("solid", fgColor="F2F7FC"), PatternFill("solid", fgColor="FFFFFF")]
ROW_FONT = Font(name="メイリオ", size=9)
ROW_ALIGN = Alignment(vertical="center", wrap_text=False)
URL_FONT = Font(name="メイリオ", size=9, color="0563C1", underline="single")

DATA = [
    ("株式会社MJリサーチ",        "https://mj-research.co.jp/lp/af002.html",                 "リスティング広告", "", "浮気調査・探偵サービス",         "浮気調査 費用, 探偵事務所 おすすめ, 浮気 証拠 集め方",             "探偵・調査"),
    ("株式会社クロル",             "https://curolu.co.jp/lp/uut/",                             "リスティング広告", "", "浮気調査・探偵サービス",         "総合探偵社 口コミ, 浮気 バレない 調査, 探偵 無料相談",             "探偵・調査"),
    ("さくら幸子探偵事務所",       "https://www.sakurasachiko.jp/lp/uwaki/",                   "リスティング広告", "", "浮気調査・探偵サービス",         "浮気調査 女性探偵, 信頼できる 探偵事務所, 不倫 証拠 探偵",         "探偵・調査"),
    ("東京女性探偵社",             "https://tokyo.o-praca.jp/lp/",                             "リスティング広告", "", "浮気調査・探偵サービス",         "女性探偵 浮気調査, 東京 探偵社 女性スタッフ, 不倫調査 安心",       "探偵・調査"),
    ("株式会社MR",                 "https://www.tantei-mr.co.jp/lp/premarital/",               "リスティング広告", "", "探偵・身辺調査（結婚前調査）",   "結婚前 身辺調査 費用, 婚前調査 探偵 依頼, 相手の過去 調査 探偵",   "探偵・調査"),
    ("株式会社生活総合サービス",   "https://www.teinei.co.jp/lp/bocodeco/article/001/aa/",    "リスティング広告", "", "青汁・健康食品（定期通販）",     "青汁 通販 おすすめ, 飲みやすい 青汁 初回, 健康食品 定期便 申込",   "健康食品・青汁D2C"),
    ("株式会社日本第一製薬",       "https://jp-no1.co.jp/lp/jset029",                         "リスティング広告", "", "ダイエットサプリ（通販）",       "ダイエットサプリ 効果 口コミ, 脂肪燃焼 サプリ, 痩せる サプリ 通販", "健康食品・サプリD2C"),
    ("純藍株式会社",               "http://junai-inc.co.jp/lp/01/sp/index.html",              "リスティング広告", "", "藍の青汁（健康食品）",           "藍 青汁 無農薬 通販, 国産 青汁 定期便, 健康食品 藍草 初回",         "健康食品・青汁D2C"),
    ("株式会社ルックルック",       "https://looklook.co.jp/lp/kaimin/",                        "リスティング広告", "", "睡眠サプリ（通販）",             "睡眠サプリ 効果 口コミ, 不眠 改善 サプリ, グリシン テアニン 睡眠", "健康食品・サプリD2C"),
    ("株式会社エポーラー",         "https://epauler.co.jp/lp/wt/cp_001/pc/",                  "リスティング広告", "", "プラセンタ美容液（通販）",       "プラセンタ 美容液 効果, 老化防止 サプリ 女性, 肌ケア プラセンタ",  "美容・健康食品D2C"),
    ("株式会社ミックコスモ",       "https://www.miccosmo.co.jp/lp/whitelabelplus/",            "リスティング広告", "", "プラセンタ化粧品・サプリ（通販）", "プラセンタ 化粧品 おすすめ, エイジングケア 通販, 美白 プラセンタ",  "美容・健康食品D2C"),
    ("株式会社健美舎",             "https://www.kenbishya.co.jp/lp/kenkainosusume/",           "リスティング広告", "", "グルコサミン・関節サプリ（通販）", "グルコサミン サプリ 効果, 膝 関節 サプリメント, 軟骨 サポート 通販", "健康食品・サプリD2C"),
    ("株式会社ココロハ",           "https://cocoloha.co.jp/lp/toroli/",                        "リスティング広告", "", "ローヤルゼリーサプリ（通販）",   "ローヤルゼリー サプリ 効果, 女性ホルモン 更年期 サプリ, 疲労回復", "健康食品・サプリD2C"),
    ("PharmaX株式会社",            "https://yojo.co.jp/lp/",                                   "リスティング広告", "", "パーソナル漢方（定期通販）",     "漢方 オンライン 相談, パーソナル漢方 体質改善, 漢方薬 通販 初回",   "健康食品・漢方D2C"),
    ("株式会社me",                 "https://maroa.co.jp/lp/001/",                              "リスティング広告", "", "酵素ペースト（美容健康食品）",   "酵素 ペースト 美容, 腸活 発酵食品 通販, デトックス 酵素 定期",     "健康食品・サプリD2C"),
    ("株式会社ラメリア・ジャパン", "https://www.lamellia-japan.co.jp/lp/pre/news601_all/",     "リスティング広告", "", "エイジングケアサプリ（通販）",   "エイジングケア サプリ 50代, シミ しわ 予防 サプリ, アンチエイジング", "美容・健康食品D2C"),
    ("株式会社ジネコ",             "https://jineko.co.jp/lp/pregna_man/",                      "リスティング広告", "", "妊活サプリ・男性用（通販）",     "妊活 サプリ 男性 亜鉛, 精子 改善 サプリ, 男性不妊 サプリメント",   "健康食品・妊活D2C"),
    ("株式会社JUNOa",              "https://junoa.co.jp/lp/cpl/ec/smp/",                       "リスティング広告", "", "妊活サプリ（定期通販）",         "妊活 サプリメント おすすめ, 妊活 腸活 サプリ, 葉酸 妊活 栄養",     "健康食品・妊活D2C"),
    ("natural tech株式会社",       "https://brands.naturaltech.jp/mitas/lp/",                  "リスティング広告", "", "葉酸サプリ・mitas（定期通販）", "葉酸サプリ 妊活 おすすめ, 妊婦 葉酸 必要量, 温活 葉酸 サプリ",     "健康食品・妊活D2C"),
    ("株式会社feileB",             "https://lepeelorganics.jp/lp?u=navi-ninkatsumen",          "リスティング広告", "", "オーガニックサプリ（通販）",     "オーガニックサプリ 無添加, フェリチン 鉄分 サプリ, 無農薬 サプリ", "健康食品・サプリD2C"),
]


def make_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "架電リスト"

    # Header
    for col_idx, header in enumerate(HEADER, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
    ws.row_dimensions[1].height = 20

    # Data
    for row_idx, row_data in enumerate(DATA, 2):
        fill = ROW_FILLS[(row_idx - 2) % 2]
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.fill = fill
            cell.alignment = ROW_ALIGN
            if col_idx == 2:  # LP URL column
                cell.hyperlink = value
                cell.font = URL_FONT
            else:
                cell.font = ROW_FONT
        ws.row_dimensions[row_idx].height = 16

    # Column widths
    for col_idx, width in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # Freeze pane & auto filter
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADER))}1"

    wb.save(OUTPUT_PATH)
    print(f"Saved: {OUTPUT_PATH}")
    print(f"Rows: {len(DATA)}")


if __name__ == "__main__":
    make_xlsx()
