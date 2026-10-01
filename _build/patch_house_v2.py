# -*- coding: utf-8 -*-
"""ハウス食品グループ（Come on House）LINE公式アカウント運用のご提案：殿村さんの元資料（42枚）を直す（2026-10-01）
A 別業種の文言を削除・置換（S14の表＝不動産売却 → Come on House 向けに作り直し／S29・S31＝買取の画面のページは削除）
  S17の6か月後をSIM（S18）の数字に／サンクスLINE誘導 初期15万・月額5万／S2の(?)・アカウント表記／S36の誤字・社内語／古い注記
B 文末を言い切りに／字体をメイリオに／出典のない数字を削除（S6・S7・S9）／「SIM決定版」→出典の書き方／S19・S20の数字を圧縮しない
C 重複（S10＝S24と同じ図）を削除／アジェンダS12の章名を中身に合わせる／費用対効果（S17・S18）を「費用感」の章へ移動
  「公式LINE」→「LINE公式アカウント」
  python3 _build/patch_house_v2.py <元資料.pptx>
"""
import sys
from pathlib import Path
from pptx import Presentation
from pptx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "20261001_ハウス食品グループ株式会社御中_LINE公式アカウント運用のご提案.pptx"
prs = Presentation(sys.argv[1])
S = list(prs.slides)
log = []


def paras(s):
    yield from s.shapes._spTree.iter(qn("a:p"))


def rep(no, old, new, whole=False):
    """no=元資料のページ番号。1つのrunの中なら書式を残して置き換え、runをまたぐときは段落ごと書き換える"""
    hit = 0
    for p in paras(S[no - 1]):
        ts = list(p.iter(qn("a:t")))
        full = "".join(t.text or "" for t in ts)
        if whole and full.strip() != old:
            continue
        if old not in full:
            continue
        one = next((t for t in ts if t.text and old in t.text), None)
        if one is not None:
            one.text = one.text.replace(old, new)
        else:
            ts[0].text = full.replace(old, new)
            for t in ts[1:]:
                t.text = ""
        hit += 1
    assert hit, (no, old)
    log.append((no, old, new))


def drop_shape(no, cond):
    n = 0
    for sh in list(S[no - 1].shapes):
        if cond(sh):
            sh._element.getparent().remove(sh._element); n += 1
    assert n, (no, "drop")


def cell_rows(no, rows):
    tbls = [sh.table for sh in S[no - 1].shapes if sh.has_table]
    flat = [r for t in tbls for r in list(t.rows)[1:]]
    assert len(flat) == len(rows)
    for r, vals in zip(flat, rows):
        for c, v in zip(r.cells, vals):
            ts = list(c.text_frame._txBody.iter(qn("a:t")))
            ps = c.text_frame._txBody.findall(qn("a:p"))
            for extra in ps[1:]:
                c.text_frame._txBody.remove(extra)
            ts = list(ps[0].iter(qn("a:t")))
            ts[0].text = v
            for t in ts[1:]:
                t.text = ""


# ---------------- A ----------------
rep(2, "10％前後(?)", "10％前後")
rep(2, "LINE公式アカウントの保有はない。", "LINE公式アカウントは広告用のみ（運用中のアカウントなし）")
drop_shape(2, lambda sh: sh.has_text_frame and "現状広告用アカウントとして存在" in sh.text_frame.text)
cell_rows(14, [
    ["①接触", "献立・料理のヒントを探している", "1回読んで終わり、再訪のきっかけがない", "コラム閲覧時のポップアップで「新着レシピをLINEで受け取る」"],
    ["②サイト離脱", "会員登録まではしない", "メールアドレスの入力が面倒で離脱", "ワンタップのLINE追加を、会員登録の代わりの入口に"],
    ["③育成", "気になる企画・商品はある", "メルマガが開かれず、届かない", "コンテンツ更新通知（月4回）で再訪を促す"],
    ["④リード獲得", "会員登録・イベント応募", "登録直後の熱量が続かない", "サンクスLINE誘導で、登録直後に友だち化"],
    ["⑤リード有効化", "好みや家族構成が分からない", "一斉配信で自分事化されない", "3問のアンケートで好み・世帯を把握し、出し分け配信"],
    ["⑥再育成", "登録後に休眠化", "キャンペーン後に再訪のきっかけがない", "総選挙の結果発表・季節の企画配信で呼び戻す"],
    ["⑦マネタイズ", "商品を買うきっかけが少ない", "サイトの閲覧が店頭の購買につながらない", "対象商品のレシピ提案・クーポン・マストバイキャンペーン"],
    ["⑧ファン化/紹介", "好きな商品・企画が定着", "参加が一部の会員に限られる", "投稿・投票企画への参加、口コミ・シェアの依頼"]])
for old, new in [("13,000前後", "12,983人"), ("30％", "26.9％"), ("9,000人程度", "9,487人"), ("1,000程度", "3,795件"), ("4,000件", "7,779件")]:
    rep(17, old, new)
rep(30, "・購入完了", "・会員登録完了", whole=True)
rep(30, "・セミナー予約", "・イベント応募", whole=True)
rep(30, "・登録完了", "・キャンペーン応募", whole=True)
rep(30, "初期：10万円", "初期：15万円")
rep(30, "月額3万円～", "月額5万円")
drop_shape(30, lambda sh: sh.has_text_frame and "暫定措置として導入済アカウント" in sh.text_frame.text)
rep(35, "転職ユーザーのカスタマージャーニー", "会員ユーザーのカスタマージャーニー")
rep(36, "LINEOAの運用業務・CR制作業務を一部外出ししてご依頼依頼も。", "LINE公式アカウントの運用業務・クリエイティブ制作業務の一部のみのご依頼も可能")
for r in next(sh for sh in S[35].shapes if sh.has_table).table.rows:
    cs = r.cells
    if "サンクスLINE誘導ツール" in cs[4].text:
        t5 = list(cs[5].text_frame._txBody.iter(qn("a:t")))
        t5[0].text = "¥150,000"
        from copy import deepcopy
        b6 = cs[6].text_frame._txBody
        for p in b6.findall(qn("a:p")):
            b6.remove(p)
        b6.append(deepcopy(cs[5].text_frame._txBody.findall(qn("a:p"))[0]))
        list(b6.iter(qn("a:t")))[0].text = "¥50,000"

# ---------------- B ----------------
END = [(2, "をもとにした想定数値となります。", "をもとにした想定数値"),
       (2, "今回ご提案の内容は記載させていただいております。", "今回のご提案内容を記載"),
       (3, "今回ご提案の内容は記載させていただいております。", "今回のご提案内容を記載"),
       (11, "導入支援が可能です。", "導入支援が可能"),
       (25, "としての活用等も可能です", "としての活用等も可能"),
       (27, "★を記載しております。", "★を記載"),
       (28, "実施が可能です。", "実施可能"),
       (32, "行うことが出来る機能です。", "行うことができる機能"), (33, "行うことが出来る機能です。", "行うことができる機能"),
       (32, "APIを使用した機能です。", "APIを使用した機能"), (33, "APIを使用した機能です。", "APIを使用した機能"),
       (32, "利用可能機能が異なります。", "利用可能な機能が異なる"), (33, "利用可能機能が異なります。", "利用可能な機能が異なる"),
       (35, "月次対応内に含んでおります。", "月次対応内に含む"),
       (36, "都度発注でのご依頼も各単価で受けさせていただきます。", "都度発注でのご依頼も各単価で対応可能"),
       (41, "前後する可能性がございます。", "前後する可能性あり")]
for no, o, n in END:
    rep(no, o, n)
# 出典のない数字を削除
rep(6, "開封率に関してはメールマガジンの6倍を誇ると言われています。", "メールマガジンより開封されやすく、すぐに読まれる")
for o, n in [("10％", "低い"), ("60％", "高い"), ("1時間", "時間がかかる"), ("15分", "すぐに読まれる")]:
    rep(6, o, n, whole=True)
rep(7, "「今すぐ客」以外の約99％の層を", "「今すぐ客」以外の大多数の層を")
drop_shape(7, lambda sh: sh.has_text_frame and sh.text_frame.text.strip() in ("1-2％", "9-9.5％", "80％"))
rep(7, "集客施策におけるCVRは高くても2%程度。", "")
rep(9, "商材平均で最大３%のユーザーは有効化/マネタイズ可能 ※生涯CVR とされている。", "流入元ごとに、LINEへの誘導と配信施策を組み合わせる。")
drop_shape(9, lambda sh: sh.has_text_frame and (sh.text_frame.text.strip() in ("３％", "1.5％", "2.5％", "2％") or "CVするユーザー割合" in sh.text_frame.text))
rep(9, "→ 商談・申込への誘導", "→ 会員登録・企画参加への誘導")
rep(9, "→ 商談・再購入誘導", "→ 企画参加・商品購入へ")
rep(9, "ホットリード:即スタッフ対応", "反応の高い方へ企画・キャンペーン案内")
# S19・S20：数字を圧縮しない（SIM＝S18の数字）・社内語・言い切り
for o, n in [("初月に約7,100人", "初月に7,113人"), ("ブロック率25〜27%と低位。", "ブロック率は6ヶ月目でも26.9%と低位。"),
             ("（月 +約1,200人）。6ヶ月で友だち約13,000人。", "。6ヶ月で友だち12,983人。"), ("約 13,000 人", "12,983 人"),
             ("初月 約7,100人", "初月 7,113人"), ("ブロック率25〜27%と低く、", "ブロック率は6ヶ月目でも26.9%で、"),
             ("6ヶ月後 約3,800件/月", "6ヶ月後 3,795件/月"),
             ("到達（配信可）約9,500人 × 反応率40%（SIM）。6ヶ月累計 約17,800件。", "到達（配信可）9,487人 × 反応率40%。6ヶ月累計 17,802件。"),
             ("※数値はSIM決定版より。ブロック率・反応率は想定値。費用：初期15万円＋月額5万円（投稿代行）。", "※出典：弊社シミュレーション。ブロック率・反応率は想定値"),
             ("継続接点に変えます。", "継続接点に変える")]:
    rep(19, o, n)
for o, n in [("（現行資料の課題）", ""), ("習慣化します。", "習慣化する"), ("ブロック率も25〜27%と低位。", "ブロック率も6ヶ月目で26.9%と低位。"),
             ("6ヶ月で約13,000人。", "6ヶ月で12,983人。"), ("※SIM決定版より。", "※出典：弊社シミュレーション。")]:
    rep(20, o, n)

# ---------------- C ----------------
for no in (19, 20, 21, 22):
    for p in paras(S[no - 1]):
        for t in p.iter(qn("a:t")):
            if t.text and "公式LINE" in t.text:
                t.text = t.text.replace("公式LINE", "LINE公式アカウント")
        full = "".join(x.text or "" for x in p.iter(qn("a:t")))
        if "公式LINE" in full.replace("LINE公式", ""):        # runが分かれている見出し（追加施策｜公式LINE）
            ts = list(p.iter(qn("a:t")))
            ts[0].text = full.replace("｜公式LINE", "｜LINE公式アカウント")
            for x in ts[1:]:
                x.text = ""
rep(12, "現状ステータスの整理", "想定動線")

# 字体をメイリオに統一
AFTER = {qn(x) for x in ("a:sym", "a:hlinkClick", "a:hlinkMouseOver", "a:rtl", "a:extLst")}
for s in prs.slides:
    for tag in ("a:rPr", "a:endParaRPr", "a:defRPr"):
        for rpr in s.shapes._spTree.iter(qn(tag)):
            fs = []
            for ft in ("a:latin", "a:ea", "a:cs"):
                e = rpr.find(qn(ft))
                if e is None:
                    e = rpr.makeelement(qn(ft), {})
                else:
                    rpr.remove(e)
                e.set("typeface", "メイリオ")
                for k in ("panose", "pitchFamily", "charset"):
                    e.attrib.pop(k, None)
                fs.append(e)
            nxt = next((c for c in rpr if c.tag in AFTER), None)
            for e in fs:
                (nxt.addprevious(e) if nxt is not None else rpr.append(e))

# 並べ替え：S10・S29・S31を削除、S17・S18を「費用感」（S34）の直後へ
lst = prs.slides._sldIdLst
ids = list(lst)
drop = {10, 29, 31}
order = [n for n in range(1, 43) if n not in drop and n not in (17, 18)]
k = order.index(34) + 1
order[k:k] = [17, 18]
for n in drop:
    prs.part.drop_rel(ids[n - 1].rId)
for el in list(lst):
    lst.remove(el)
for n in order:
    lst.append(ids[n - 1])
prs.save(str(OUT))
print("saved:", OUT.name, len(order), "枚 / 置き換え", len(log), "件")
