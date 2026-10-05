# -*- coding: utf-8 -*-
"""サンクスLINE誘導 4枚資料の重複整理（2026-10-05 殿村さん指示「同義テキストの乱立を禁止」）。
ベースは殿村さんがPCで保存した版。各スライドに固有の役割だけを持たせる。
  1 何か＝仕組みと想定友だち獲得数だけ（数字は1回だけ見せる）
  2 なぜ重要か＝CV後を売上につなげる必要性だけ（数字・LTVの話は出さない）
  3 初回売上までどう使うか＝来店・面談・購入までの施策＋実績
  4 その後どうLTVを伸ばすか＝再来店・再購入の施策＋実績
実績はLINEヤフー公式か弊社実績だけ。業界名だけ書く（社名・出所は書かない）。％で統一。
  python3 _build/patch_thanks_line_4p_dedupe.py <殿村さん保存版.pptx>
"""
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

sys.path.insert(0, str(Path(__file__).resolve().parent))
from pptx_parts import *  # noqa

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "20261005_サンクスLINE誘導のご提案.pptx"
PALE = "E4E8F6"
prs = Presentation(sys.argv[1])
SW = prs.slide_width / 360000
L, R = 2.2, SW - 2.2
CW = R - L
S = list(prs.slides)


def shp(s, name):
    return next(x for x in s.shapes if x.name == name)


def lead(s, text):
    set_paras(shp(s, "TextBox 2"), [(text, 0)])


def flow(s, x0, x1, y, h, steps, size=13, ov=0.35):
    n = len(steps)
    w = (x1 - x0 + (n - 1) * ov) / n
    for i, (ls, f, tc) in enumerate(steps):
        x = x0 + i * (w - ov)
        shape(s, MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON, x, y, w, h, None, fill=f, adj=0.25)
        pad = 0.3 if i == 0 else 0.7
        label(s, x + pad, y, w - pad - 0.6, h, [(t, size, True, tc, 0) for t in ls], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def results(s, y, items, head="実績"):
    """items = [(業界, 施策, 指標, 数字, 比較条件)]。業界名 → 施策 → 大きな数字 の順。"""
    label(s, L, y, 10.0, 0.8, [(head, 13, True, NAVY, 0)], anchor=MSO_ANCHOR.MIDDLE)
    cw = CW / len(items)
    for i, (ind, how, metric, num, cond) in enumerate(items):
        x = L + i * cw
        if i:
            vline(s, x, y + 1.0, y + 5.4, LGRAY, 0.75)
        size = 34 if len(num) <= 7 else 26
        label(s, x + 0.4, y + 0.9, cw - 0.8, 4.6, [(ind, 14, True, NAVY, 2), (how, 10.5, False, GRAY, 6), (metric, 11.5, False, INK, 0),
                                                  (num, size, True, NAVY, 0), (cond, 9.5, False, GRAY, 0)], align=PP_ALIGN.CENTER)


# ================= 1 何か：数字は1回だけ =================
s = S[0]
for sh in list(s.shapes):
    if (sh.has_text_frame and sh.text_frame.text.startswith("CV：100件　×")) or (sh.shape_type == 9 and sh.top > 13 * 360000):
        sh._element.getparent().remove(sh._element)
from pptx.util import Cm, Pt  # noqa: E402
for sh in s.shapes:  # 数字の段を少し下げて余白をそろえ、「CV：100件」が矢印に重ならないよう大きさを合わせる
    if sh.top is not None and 8.5 * 360000 < sh.top < 13 * 360000:
        sh.top = sh.top + Cm(1.3)
    if sh.name in ("TextBox 18", "TextBox 19", "TextBox 21"):
        for p_ in sh.text_frame.paragraphs:
            for r in p_.runs:
                if r.font.size and r.font.size.pt >= 40:
                    r.font.size = Pt(42)
                elif r.font.size and r.font.size.pt >= 18:
                    r.font.size = Pt(18)
    if sh.name == "TextBox 18":
        sh.left = Cm(1.6)

# ================= 2 なぜ重要か：CV後を売上につなげる必要性だけ =================
s = S[1]
keep_header_only(s)
lead(s, "CV後に接点がなければ、来店・購入の前に離脱が起きる")
LW = 4.6
rows = [("サンクスLINE", "なし", GRAY, 5.4,
         [(["CV"], LGRAY, INK), (["その後の", "接点なし"], LGRAY, INK), (["来店・購入", "の前に離脱"], LGRAY, INK), (["売上に", "ならない"], "F2F2F2", GRAY)]),
        ("サンクスLINE", "あり", GREEN_TX, 10.4,
         [(["CV"], PALE, NAVY), (["LINEで", "フォロー"], GREEN, WHITE), (["来店・購入", "・成約"], NAVY, WHITE), (["売上に", "なる"], NAVY, WHITE)])]
for a, b, c, y, steps in rows:
    label(s, L, y, LW, 3.0, [(a, 12, True, GRAY, 0), (b, 18, True, c, 0)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    flow(s, L + LW + 0.3, R, y, 3.0, steps, size=14)
hline(s, L, R, 9.15, LGRAY, 0.75)

# ================= 3 初回売上まで：施策＋実績 =================
s = S[2]
for sh in list(s.shapes):
    if sh.top is not None and sh.top > 8.5 * 360000:
        sh._element.getparent().remove(sh._element)
results(s, 8.9, [("美容クリニック", "予約確認・前日のお知らせをLINEで配信", "予約後の来院率", "40〜50%改善", ""),
                 ("皮膚科クリニック", "友だち追加後のステップ配信", "LINE経由の予約数", "約20%増", "ステップ配信の実施前との比較"),
                 ("就職支援", "LINEのチャットで面談へ誘導", "LINEで問い合わせた方のうち", "55%", "が面談を予約")])

# ================= 4 その後のLTV：施策＋実績 =================
s = S[3]
keep_header_only(s)
lead(s, "初回で終わらせず、継続フォローで再来店・再購入へ")
stairs = [("初回来店・", "購入・成約"), ("LINEで", "継続接点"), ("フォロー", "配信"), ("再来店・", "再購入"), ("アップセル／", "クロスセル"), ("LTV", "向上")]
n = len(stairs)
gap = 0.3
gw = (CW - gap * (n - 1)) / n
base = 9.4
for i, (a, b) in enumerate(stairs):
    h = 1.8 + i * 0.45
    x = L + i * (gw + gap)
    f, tc = (GREEN, WHITE) if i == n - 1 else ((NAVY, WHITE) if i >= 3 else (PALE, NAVY))
    rect(s, x, base - h, gw, h, f)
    label(s, x, base - h, gw, h, [(a, 12, True, tc, 0), (b, 12, True, tc, 0)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
label(s, L, 9.75, 3.2, 0.8, [("施策の例", 11.5, True, NAVY, 0)], anchor=MSO_ANCHOR.MIDDLE)
items = ["来店後フォロー", "関連サービス案内", "クーポン配信", "セグメント配信"]
cw = (CW - 3.4 - 0.3 * 3) / 4
for i, t in enumerate(items):
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, L + 3.4 + i * (cw + 0.3), 9.75, cw, 0.8, [(t, 11, True, NAVY, 0)], fill=CARD, adj=0.5,
          margins=(0.1, 0, 0.1, 0))
results(s, 10.9, [("美容室", "予約後のフォロー", "次回予約客数", "20%増", "約200名 → 約240名"),
                  ("居酒屋", "来店翌日のアンケートと、再来店前のクーポン", "リピーターの売上割合", "7.6% → 12.9%", "2025年1月と3月の比較"),
                  ("焼肉店", "来店翌日に、アンケートとクーポンを自動配信", "リピーター率", "19.4% → 40%超", "2024年1月と2025年1月の比較")])
label(s, L, 17.2, CW, 0.6, [("※サンクスLINE単体の効果ではなく、LINEでつながった後の施策の例（美容室・居酒屋・焼肉店はLINEミニアプリを併用）", 9, False, GRAY, 0)])

prs.save(str(OUT))
print("saved:", OUT.name, len(prs.slides), "枚")
