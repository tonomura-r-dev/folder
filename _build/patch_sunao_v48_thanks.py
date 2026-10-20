# -*- coding: utf-8 -*-
"""すなお ver4.7 → ver4.8：サンクスLINE誘導のページを、Appendixの「サンクスLINE｜機能・費用」（動画のQRつき）に差し替え、費用を足す（2026-10-20 殿村さん）。

    python -m _build.patch_sunao_v48_thanks <ver4.7.pptx> <Appendix ver1.2.pptx> <業界別ver3.7.pptx> [出力.pptx]

  殿村さん「サンクスLINEのページはマジでコレが見やすい」「費用感乗せれてないからなぁ」。
  ・FMTの「新規友だちの追加（サンクスLINE誘導）」2枚 → Appendix ver1.2 の1枚目（機能・費用・動画のQR）1枚に
  ・費用の札（初期：15万円／月額3万円〜＋設定地点追加）は、業界別ver3.7の2枚目にある札を写して、黒い「サンクスLINE誘導」の右に並べる
    （その分、「フォーム送信後のLINE誘導を自動化するツール」と注記は1段下げる）
  ・このあと patch_sunao_v4_colors を当てて、赤い文字・強調の赤枠・黄色い文字を外す（紺・黒・白・灰色。緑はLINEの部分だけ）
  ほかのページは触らない（出力を省くと上書き）。
"""
import copy
import sys

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches

from _build.build_sunao_v1 import all_meiryo
from _build.patch_tanaka_v5 import copy_slide


def title_of(s):
    for sh in s.shapes:
        if sh.has_text_frame and sh.top < Inches(0.6) and sh.text_frame.text.strip():
            return sh.text_frame.text.strip()
    return ""


def by_id(s, sid):
    return next(sh for sh in s.shapes if sh.shape_id == sid)


def paste(dst, src_shape):
    el = copy.deepcopy(src_shape._element)
    nid = max(sh.shape_id for sh in dst.shapes) + 1
    el.find(".//" + qn("p:cNvPr")).set("id", str(nid))
    dst.shapes._spTree.append(el)
    return by_id(dst, nid)


def move(sh, x=None, y=None, w=None, h=None):
    for attr, v in (("left", x), ("top", y), ("width", w), ("height", h)):
        if v is not None:
            setattr(sh, attr, Inches(v))


def main(src, appendix_path, gyokai_path, out):
    prs = Presentation(src)
    app = Presentation(appendix_path).slides[0]
    gyo = Presentation(gyokai_path).slides[1]
    assert title_of(app) == "サンクスLINE｜機能・費用" and title_of(gyo) == "サンクスLINE｜機能・費用"
    olds = [s for s in prs.slides if title_of(s) == "新規友だちの追加（サンクスLINE誘導）"]
    assert len(olds) == 2, len(olds)

    new = copy_slide(prs, app)
    # 費用の札を黒い「サンクスLINE誘導」の右に並べる
    init, mon = paste(new, by_id(gyo, 26)), paste(new, by_id(gyo, 28))
    move(init, x=4.3, y=0.92, w=1.45, h=0.53)
    move(mon, x=5.85, y=0.92, w=1.95, h=0.53)
    # 説明と注記を1段下げる（WEBサイト・動作イメージの見出しにかからない幅で）
    move(by_id(new, 9), x=4.3, y=1.5, w=5.2, h=0.3)
    note = by_id(new, 14)
    move(note, x=4.3, y=1.8, w=2.35, h=0.22)
    note.text_frame.paragraphs[0].runs[0].text = "※既存APIツールとの併用も可能"
    for r in note.text_frame.paragraphs[0].runs[1:]:
        r._r.getparent().remove(r._r)
    t = by_id(new, 22)                                          # タイトルは前後のページ（新規友だちの追加（…））にそろえる
    t.text_frame.paragraphs[0].runs[0].text = "新規友だちの追加（サンクスLINE誘導）"
    for r in t.text_frame.paragraphs[0].runs[1:]:
        r._r.getparent().remove(r._r)
    all_meiryo([new])

    lst = prs.slides._sldIdLst
    ids = {prs.part.related_part(e.rId): e for e in lst}
    pos = list(lst).index(ids[olds[0].part])
    el = ids[new.part]
    lst.remove(el)
    lst.insert(pos, el)
    for s in olds:
        e = ids[s.part]
        prs.part.drop_rel(e.rId)
        lst.remove(e)
    prs.save(out)
    print("saved:", out, len(prs.slides), "枚")


if __name__ == "__main__":
    main(*sys.argv[1:4], sys.argv[4] if len(sys.argv) > 4 else sys.argv[1])
