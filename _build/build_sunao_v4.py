# -*- coding: utf-8 -*-
"""すなおコンサルティング ver4：提案FMT（`_templates/DYM_LINEOA_提案FMT_202607（最新の枠組み）.pptx`・63枚）に、ver3の中身を当てはめる。

    python -m _build.build_sunao_v4 <ver3.pptx> <SIM ver1.2.xlsx> [出力.pptx]

  2026-10-19 殿村さん「FMTはこれに戻します」「一回そのFMTに合わせて埋めれるところを全部埋めてほしい」。
  ・ページは外さない（これから保険営業向けの数字を入れていくため）。
  ・すなおの中身がまだ無いページ（シーズナリティ・前後検索・Day0〜15・休眠リード・長期ニーズ・事例）は、
    不動産の中身だけ消して枠を残す（数字は「--」、画像は差し込み枠）。通知メッセージのページはFMTのまま（不動産の中身なし）。
  ・ver3の中身の当てはめ先：サマリー・目標→要件定義／保険営業の方の環境→業界STUDY／他社のLINE運用→業界STUDY＞訴求軸／
    現状と課題→想定動線・現状把握／本質課題→要件定義（補足ポイント）／気をつける点→LINEOA施策（15日）のリード／
    配信の設計・友だちの入口→全体設計・トラック分岐／概算お見積り→想定の費用対効果。
    FMTに入る場所が無い3枚（料金・成果報酬型プラン・なぜLINEか）は、ver3のページをそのまま写して同じ章に入れる。
  ・SIMの数字はSIM ver1.2（LP UU 仮に3,000）。成果報酬プランの固定費は、ツール費＋アカウント費（コンサル費は成果報酬に置き換わる）。
  ・FMTの直し：サンクスLINE誘導は初期15万円・月額3万円〜（社内向けの価格改定の注記は外す）／都度発注のツール費／
    スケジュールの「約1か月」を外す／サービス内容の「転職ユーザー」→「保険営業の方」。
"""
import copy
import io
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import openpyxl
from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from _build.arauchi_v4.lib import add_box, add_table, move
from _build.arauchi_v10.build import clone, runs
from _build.build_sunao_v1 import all_meiryo
from _build.patch_tanaka_v5 import copy_slide

ROOT = Path(__file__).resolve().parent.parent
FMT = ROOT / "_templates" / "DYM_LINEOA_提案FMT_202607（最新の枠組み）.pptx"
OUT = ROOT / "202610_株式会社すなおコンサルティング御中_LINE公式アカウント運用のご提案_ver4.pptx"
NAVY = "1F2A5C"            # FMTの紺
LBLUE = "E5E8EE"           # FMTの月の箱の薄いグレー
SUB = "5A6378"
GRAY = "7F7F7F"
TEXT = "333333"
RED = {"FF0000", "C00000", "C8102E"}
EMPTY = None


# ---------------------------------------------------------------- 部品
def find(s, sid):
    """グループの中まで探す"""
    def walk(shapes):
        for sh in shapes:
            if sh.shape_id == sid:
                return sh
            if sh.shape_type == 6:
                r = walk(sh.shapes)
                if r is not None:
                    return r
        return None
    sh = walk(s.shapes)
    assert sh is not None, sid
    return sh


def drop(s, *ids):
    for sid in ids:
        e = find(s, sid)._element
        e.getparent().remove(e)


def _kind(rpr):
    """文字の書式の種類：n 通常／b 太字／r 赤／br 太字の赤"""
    if rpr is None:
        return "n"
    bold = rpr.get("b") == "1"
    clr = rpr.find(".//" + qn("a:srgbClr"))
    red = clr is not None and clr.get("val") in RED
    return ("b" if bold else "") + ("r" if red else "") or "n"


def write(sh, paras, ptemp=None, size=None):
    """段落ごとに書き直す。paras の1要素＝1段落：文字列（型にした段落の最初の文字の書式）／
    [(文字, 種類), ...]（種類＝n 通常・b 太字・r 赤・br 太字の赤。元の図形にある同じ種類の文字の書式を使う）／None（空行）。
    型にする段落は ptemp[i]（無ければ i 番目。足りなければ最後の段落）"""
    body = sh.text_frame._txBody
    orig = [copy.deepcopy(p) for p in body.findall(qn("a:p"))]
    sty = {}
    for p in orig:
        for r in p.findall(qn("a:r")):
            sty.setdefault(_kind(r.find(qn("a:rPr"))), r.find(qn("a:rPr")))
    for p in body.findall(qn("a:p")):
        body.remove(p)
    for i, item in enumerate(paras):
        k = ptemp[i] if ptemp else i
        t = orig[min(k, len(orig) - 1)]
        p = copy.deepcopy(t)
        first = p.find(qn("a:r"))
        first_rpr = copy.deepcopy(first.find(qn("a:rPr"))) if first is not None else None
        for tag in ("a:r", "a:br", "a:fld"):
            for e in p.findall(qn(tag)):
                p.remove(e)
        end = p.find(qn("a:endParaRPr"))
        segs = [] if item is None else ([(item, None)] if isinstance(item, str) else item)
        for text, key in segs:
            rpr = sty.get(key) if key else first_rpr
            if rpr is None:
                rpr = first_rpr if first_rpr is not None else next(iter(sty.values()), None)
            r = etree.Element(qn("a:r"))
            if rpr is not None:
                r.append(copy.deepcopy(rpr))
                if size:
                    r[0].set("sz", str(int(size * 100)))
            etree.SubElement(r, qn("a:t")).text = text
            if end is not None:
                end.addprevious(r)
            else:
                p.append(r)
        body.append(p)


def put(s, sid, *lines, size=None):
    write(find(s, sid), list(lines), size=size)


def struct(sh, items):
    """「見出し（太字）＋箇条書き＋空行」の形の文字の箱を書き直す。items＝[("h", 文字) / ("t", 文字) / ("e",)]"""
    ps = sh.text_frame._txBody.findall(qn("a:p"))
    head = next(i for i, p in enumerate(ps) if p.find(qn("a:r")) is not None
                and _kind(p.find(qn("a:r")).find(qn("a:rPr"))).startswith("b"))
    bodyi = next(i for i, p in enumerate(ps) if i != head and p.find(qn("a:r")) is not None
                 and not _kind(p.find(qn("a:r")).find(qn("a:rPr"))).startswith("b"))
    emp = next((i for i, p in enumerate(ps) if p.find(qn("a:r")) is None), head)
    tmpl = {"h": head, "t": bodyi, "e": emp}
    write(sh, [None if it[0] == "e" else it[1] for it in items], ptemp=[tmpl[it[0]] for it in items])


def _first_rpr(ps):
    for p in ps:
        r = p.find(qn("a:r"))
        if r is not None and r.find(qn("a:rPr")) is not None:
            return copy.deepcopy(r.find(qn("a:rPr")))
    return None


def set_cell(cell, text, like=None, force=False):
    """表のセルの文字を差し替える。i行目は元のi番目の段落の書式（足りなければ最後）。改行は段落を分ける。
    元のセルが空で書式が無いときは、like のセルの書式を使う"""
    tx = cell._tc.find(qn("a:txBody"))
    ps = tx.findall(qn("a:p"))
    tmpls = [copy.deepcopy(p) for p in ps]
    base = _first_rpr(ps)
    if like is not None and (base is None or force):
        base = _first_rpr(like._tc.find(qn("a:txBody")).findall(qn("a:p")))
        if force:
            tmpls = [copy.deepcopy(like._tc.find(qn("a:txBody")).find(qn("a:p")))]
    for p in ps:
        tx.remove(p)
    for i, line in enumerate(str(text).split("\n")):
        tmpl = tmpls[min(i, len(tmpls) - 1)]
        rpr = _first_rpr([tmpl])
        if rpr is None:
            rpr = base
        p = copy.deepcopy(tmpl)
        for tag in ("a:r", "a:br", "a:fld"):
            for e in p.findall(qn(tag)):
                p.remove(e)
        r = etree.Element(qn("a:r"))
        if rpr is not None:
            r.append(copy.deepcopy(rpr))
        etree.SubElement(r, qn("a:t")).text = line
        end = p.find(qn("a:endParaRPr"))
        if end is not None:
            end.addprevious(r)
        else:
            p.append(r)
        tx.append(p)


def slot(s, x, y, w, h, text):
    """まだ中身が無い所の破線の差し込み枠"""
    add_box(s, x, y, w, h, [runs((text, dict(size=12, bold=True, color=GRAY)))], line=GRAY, dash="dash",
            align="ctr", anchor="ctr", name=f"差し込み {text}")


def geo(sh):
    return [Emu(v).inches for v in (sh.left, sh.top, sh.width, sh.height)]


def swap_for_slot(s, sid, text):
    x, y, w, h = geo(find(s, sid))
    drop(s, sid)
    slot(s, x, y, w, h, text)


def fill_color(sh, fill, text_color=None):
    sp = sh._element.spPr
    for tag in ("a:solidFill", "a:noFill", "a:gradFill"):
        for e in sp.findall(qn(tag)):
            sp.remove(e)
    f = etree.Element(qn("a:solidFill"))
    etree.SubElement(f, qn("a:srgbClr"), val=fill)
    prst = sp.find(qn("a:prstGeom"))
    (prst.addnext(f) if prst is not None else sp.append(f))
    if text_color:
        for r in sh._element.iter(qn("a:rPr")):
            for e in list(r):
                if e.tag in (qn("a:solidFill"),):
                    r.remove(e)
            sf = etree.Element(qn("a:solidFill"))
            etree.SubElement(sf, qn("a:srgbClr"), val=text_color)
            r.insert(0, sf)


# ---------------------------------------------------------------- SIM
def load_sim(path):
    """SIMをLibreOfficeで計算し直したコピーから読む（コピーは読むだけ）。row(r) → [D, E..J]"""
    work = Path(tempfile.mkdtemp())
    try:
        shutil.copy(path, work / "sim.xlsx")
        subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", str(work / "c"),
                        str(work / "sim.xlsx")], check=True, capture_output=True, timeout=180)
        ws = openpyxl.load_workbook(work / "c" / "sim.xlsx", data_only=True).worksheets[0]
        rows = {r: [ws[f"{c}{r}"].value for c in "DEFGHIJ"] for r in list(range(12, 42)) + list(range(44, 52))
                + list(range(65, 74))}
        rows["name1"], rows["name2"] = ws["C8"].value, ws["C9"].value
        rows["tanka1"], rows["tanka2"] = ws["D8"].value or 0, ws["D9"].value or 0
        return rows
    finally:
        shutil.rmtree(work)


def n(v):
    return int(v + 0.5) if isinstance(v, (int, float)) else v


def yen(v):
    return f"¥{n(v):,}"


def pct(v, d=1):
    return f"{v * 100:.{d}f}%"


# ---------------------------------------------------------------- 現状ステータスの整理
def p_cover(s):
    put(s, 7, "株式会社すなおコンサルティング 御中")
    sub = find(s, 6)
    write(sub, ["広告で集めた方をLINEでナーチャリングをし、", "セミナー参加・入会へ"])
    move(sub, y=2.55, h=0.72)


def p_doshen(s):
    """想定動線（カスタマージャーニーの図＋ユーザー動線の帯＋サービス特性）"""
    put(s, 32, "セミナー予約")
    put(s, 24, "セミナー参加")
    put(s, 36, "サロン入会")
    put(s, 9, "広告からサロン入会まで8フェーズの動線を可視化し、各フェーズで発生する",
        "離脱要因とLINE施策の打ち手を整理する。")
    write(find(s, 10), [
        "・またサービス特性として、",
        None,
        [("【対象は保険営業の方】", "b"), ("： 一般の方向けではなく、保険営業の方の営業力向上を目的としたセミナー・サロン", "n")],
        [("【セミナーは無料・オンライン】", "b"), ("： セミナー参加の段階では費用の不安は小さく、金額の判断はサロン入会の段階", "n")],
        [("【イベントは年4回】", "b"), ("： イベントとイベントの間の時期は、集めた方との接点が切れやすい", "n")],
        None,
        "・貴社サービスとして",
        [("【広告で、1回のイベントで81件のセミナー予約を集めている】", "b")],
        [("【セミナーの参加後に、サロンの入会を案内できる】", "b")],
        None,
        [("という強みを持っており、", "n"), ("集めた方と、次のイベント・入会まで接点を持ち続けること", "br"),
         ("はメリットにつながる想定。", "n")],
    ])


ROWS_8F = [
    ("営業の成果を上げたいが、まだ情報収集の段階", "日程が合わないと、申し込まずに離れる", "離脱防止ポップアップ（次のセミナーの案内をLINEで受け取れる）"),
    ("今回は申し込まない", "次のセミナーまで、案内を届ける手段が無い", "ポップアップからLINE友だち化"),
    ("どんなセミナーか、まだよく分からない", "中身が見えず、予約に進まない", "見どころ動画・参加特典をステップ配信"),
    ("セミナーを予約した", "当日までに日が空き、参加を忘れる", "サンクスLINE誘導で友だち化・前日／当日のお知らせ"),
    ("セミナーに参加し、関心が高まっている", "参加後、入会を迷ったまま時間が経つ", "イベント直後にサロンの申込みフォームを案内"),
    ("参加したが未入会／予約したが不参加", "次のイベントまで接点が無い", "平常期の配信（月2回）で次のイベントへ"),
    ("サロンの入会を検討", "入会後の中身・成果が見えない", "入会した方の声・サロンの中身を配信"),
    ("サロンに入会", "--", "次のイベントの案内・紹介"),
]


def p_doshen_table(s):
    for tid, rows in ((18, ROWS_8F[:4]), (20, ROWS_8F[4:])):
        tbl = find(s, tid).table
        for r, vals in enumerate(rows, start=1):
            for c, v in enumerate(vals, start=1):
                set_cell(tbl.cell(r, c), v)


def p_genjo(s):
    """現状把握：LPのステータス（広告実績）／LINEのステータス"""
    for sid, t in ((65, "LP UU"), (71, "広告費用"), (78, "CV数")):
        put(s, sid, t)
    for sid, t in ((106, "3,000"), (108, "約2,760"), (110, "約211万円"), (118, "81件"), (120, "セミナー予約"),
                   (122, "なし"), (116, "--"), (114, "--"), (112, "広告")):
        put(s, sid, t)
    put(s, 31, "LP（セミナー申込みページ）ステータス")
    put(s, 67, "※1回のイベントの数値。LP UUは", "　広告のクリック数からの想定数値。")
    write(find(s, 43), [[("広告でセミナー予約を獲得（1回 81件・CPA 26,043円）。", "n")],
                        [("申し込まずに離れた方を、次のイベントまで案内する手段が無い。", "r")]])
    put(s, 53, "参加して入会していない方へ、", "入会まで案内を続ける手段が無い状況。")
    drop(s, 63)                                              # 「※現状広告用アカウントとして存在」（不動産の案件の注記）
    swap_for_slot(s, 10, "LPの画面")


def p_season(s):
    box = geo(find(s, 7))                                    # 不動産のGoogleトレンドの画像
    drop(s, 7, 6, 10, 12, 15)                                # 画像と、繁忙期の囲み・「やや繁忙期」
    slot(s, *box, "Googleトレンドの画像（保険営業の方の検索の季節変動）")


def p_zengo(s, pic_id):
    """前後検索の画像のページ：画像と、画像の上の言葉の囲み（小さな枠）を外し、差し込み枠にする"""
    box = geo(find(s, pic_id))
    drop(s, pic_id)
    for sh in list(s.shapes):
        if sh.shape_type != 1 or sh.has_text_frame and sh.text_frame.text.strip():
            continue
        x, y, w, h = geo(sh)
        if h <= 0.25 and w < 1.3 and x < 9.2 and 1.9 < y < 6.6:
            sh._element.getparent().remove(sh._element)
    slot(s, *box, "前後検索の画像")
    put(s, 22, "キーワード")


def p_zengo_sum(s):
    for sid in (7, 11, 17):
        put(s, sid, "--")
    drop(s, 19)


def p_study(s):
    """業界STUDY：保険営業の方の環境（ver3の市場環境のページの数字）"""
    put(s, 22, "営業職員数")
    write(find(s, 24), ["約24万人で横ばい（2020〜2024年度）", "生命保険の営業職員（各年度末）", None])
    put(s, 26, "保険の加入経路")
    write(find(s, 28), ["加入経路のうち「生命保険会社の営業職員」が56.7%（2024年調査）", None])
    put(s, 30, "お客様との接触回数")
    write(find(s, 32), ["コロナ前と比べ、接触回数が「減少した」営業職員が55.5%（2023年調査）", None])
    write(find(s, 34), ["毎年ほぼ同じ規模の方が、営業を続けている", "一度つながった方に、次のイベントも届けられる", None])
    write(find(s, 36), ["保険の加入は、今も営業職員が中心", "営業力を上げたいニーズは続く", None])
    write(find(s, 38), ["お客様と会える回数が減っている", "成果を出す営業のやり方が求められている", None])
    drop(s, 9, 16)
    # 営業職員数の棒グラフ（図形で作る）
    data = (("2020年度", 248601), ("2021年度", 242947), ("2022年度", 247081), ("2023年度", 240903), ("2024年度", 241507))
    add_box(s, 0.7, 3.9, 9.4, 0.3, [runs(("生命保険の営業職員数（各年度末・名）", dict(size=12, bold=True, color=NAVY)))],
            anchor="ctr", ins=(0, 0, 0, 0), name="グラフの見出し")
    base, top, bw, gap = 6.15, 4.55, 1.1, 0.75
    x0 = 0.7 + (9.4 - (5 * bw + 4 * gap)) / 2
    for k, (yr, v) in enumerate(data):
        hgt = (base - top) * v / 250000
        x = x0 + k * (bw + gap)
        add_box(s, x, base - hgt, bw, hgt, [runs(("", dict(size=10)))], fill=NAVY, name=f"棒 {yr}")
        add_box(s, x - 0.2, base - hgt - 0.3, bw + 0.4, 0.28, [runs((f"{v:,}", dict(size=12, bold=True, color=NAVY)))],
                align="ctr", anchor="b", ins=(0, 0, 0, 0), name=f"値 {yr}")
        add_box(s, x - 0.2, base + 0.03, bw + 0.4, 0.25, [runs((yr, dict(size=11, color=TEXT)))], align="ctr",
                anchor="t", ins=(0, 0, 0, 0), name=f"年度 {yr}")
    note = find(s, 8)
    write(note, ["※営業職員数：生命保険協会「2025年版 生命保険の動向」／加入経路：生命保険文化センター「2024（令和6）年度 生命保険に関する全国実態調査」",
                 "※接触回数：矢野経済研究所「生命保険の営業職員アンケート調査」（2023年4〜6月、生命保険の営業職員400名）"])
    move(note, x=0.42, y=6.5, w=9.95, h=0.42)
    drop(s, 10)


def p_study_axis(s, ver3):
    """業界STUDY＞訴求軸：他社のLINE運用（ver3）"""
    put(s, 7, "野沢塾")
    write(find(s, 12), ["保険営業の方向けの講座（月2回）", "サイトのボタンから友だち追加", "あいさつで特典の動画を届ける"])
    put(s, 14, "R&C保険営業ラボ")
    write(find(s, 16), ["保険営業の方向けの無料の学習サービス（2026年9月開始）",
                        "登録時の5問のアンケートに合わせて、資料・動画を段階的に配信", None])
    put(s, 18, "貴社の訴求軸")
    write(find(s, 20), ["イベントの予約・参加に合わせた案内", "参加後、サロンの入会へつなげる", None])
    # 左の画像＝ver3の「野沢塾のあいさつ」の画面。右は他社の運用のまとめ
    x, y, w, h = geo(find(s, 4))
    drop(s, 4, 6)
    pic = next(sh for sh in ver3.shapes if sh.shape_type == 13)
    blob = pic.image.blob
    ratio = pic.width / pic.height
    ph = h - 0.8
    pw = min(w, ph * ratio)
    ph = pw / ratio
    s.shapes.add_picture(io.BytesIO(blob), Inches(x + (w - pw) / 2), Inches(y), Inches(pw), Inches(ph))
    add_box(s, x, y + ph + 0.05, w, 0.28, [runs(("野沢塾のあいさつ（LINE公式アカウント）", dict(size=10, color=GRAY)))],
            align="ctr", anchor="t", ins=(0, 0, 0, 0), name="画像の説明")
    add_box(s, 5.25, 2.6, 4.9, 1.3,
            [runs(("保険営業の方向けの他社も、LINEで特典を渡し、", dict(size=14, bold=True, color=NAVY))),
             runs(("友だちに学びを届けている", dict(size=14, bold=True, color=NAVY)))],
            fill=LBLUE, align="ctr", anchor="ctr", line_spacing=120, name="他社のまとめ")
    n22, n24 = find(s, 22), find(s, 24)
    write(n22, ["※野沢塾：LINE公式アカウント（2026年10月19日確認）・公式サイト（2026年10月18日確認）より。画面の名前とパスワードは伏せている"])
    write(n24, ["※R&C保険営業ラボ：R&C株式会社のプレスリリースより。招待制のため、画面は未掲載"])
    move(n22, x=0.42, y=6.5, w=9.95, h=0.24)
    move(n24, x=0.42, y=6.74, w=9.95, h=0.24)


# ---------------------------------------------------------------- 施策の詳細（配信の効率化）
def p_yoken(s, active6):
    put(s, 27, "広告で集めた方をLINEにため、次のイベント参加数・入会数を伸ばし、LTVを最大化する", size=16)
    put(s, 46, f"6か月目に追える人数{active6}人（ブロックされずに届く友だち）を目指す。",
        "イベント期はセミナー予約・参加、イベントの後はサロン入会へつなげる。")
    write(find(s, 32), [
        "　・ポップアップツールとサンクスLINE誘導で、友だちを集める",
        "　　　LPを離れる方と、セミナーを予約した方を、次のイベントまで追える友だちに。",
        "　・興味あり・予約・参加の段階ごとに、案内を変える",
        "　　　見どころ動画・参加特典 → 前日・当日のお知らせ → サロンの申込みフォーム。",
        "　・平常期の配信で、次のイベントまでつながりを保つ",
        "　　　イベントの無い時期は月2回に絞り、ブロックを抑えながら追える人数を積み上げる。",
        "＋α施策として…。",
        [("　・過去のセミナー予約者・参加者へのご案内", "b"), ("　※通知メッセージ", "br")],
        "　　　過去の予約者の情報（電話番号）を保有している場合、リストを活用した配信が可能。",
        [("　・サロンの申込み完了ページからのLINE誘導", "b"), ("　※サンクスLINE誘導", "br")],
    ], ptemp=[0, 1, 2, 3, 4, 5, 6, 7, 8, 9])


def p_yoken2(s):
    put(s, 8, "本質課題：イベントが終わると、集めた方との接点が切れる")
    for sid, head, body in ((12, "01 集める", "離れる方・予約した方を、友だちにする手段が無い。"),
                            (16, "02 追う", "次のイベントまで、情報を届ける手段が無い。"),
                            (20, "03 つなげる", "参加した方に、入会を案内する手段が無い。")):
        write(find(s, sid), [head, body])
    for sid, t in ((22, "レバー①友だち獲得"), (28, "レバー②予約・参加の向上"), (35, "レバー③入会の向上")):
        put(s, sid, t)
    write(find(s, 24), ["ポップアップツールで、LPを離れる方を友だちに", "サンクスLINE誘導で、予約した方を友だちに", None])
    write(find(s, 31), ["見どころ動画・参加特典で、予約へ", "前日・当日のお知らせで、当日の参加へ", None])
    write(find(s, 37), ["イベント直後に、サロンの申込みフォームを案内", "平常期の配信で、次のイベントまでつながる", None])


def p_hiyou(s, sim):
    tot = sim[73]
    ev = [k for k, p in enumerate(sim[12][1:]) if p == "イベント期"][0]
    nv = [k for k, p in enumerate(sim[12][1:]) if p != "イベント期"][0]
    put(s, 8, f"{n(tot[0]):,}円")
    write(find(s, 10), [f"イベント期 {n(tot[1 + ev]):,}円〜", f"平常期 {n(tot[1 + nv]):,}円〜"], size=16)
    init = (("プロフ設定", "有"), ("挨拶メッセ", "有"), ("Step配信", "有（5本）"), ("リッチメニュー", "有"),
            ("ツール導入", "有（2種）"), ("タグ・GA設定", "有"), ("アカウント開設", "有"))
    for (a, b), (k, v) in zip(((23, 27), (31, 35), (39, 43), (48, 52), (56, 60), (64, 68), (72, 76)), init):
        put(s, a, k)
        put(s, b, v)
    mon = (("イベント期の配信", f"月{n(sim[30][1 + ev])}件"), ("平常期の配信", f"月{n(sim[30][1 + nv])}件"),
           ("Step修正", "随時"), ("メニュー修正", "随時"), ("ツール保守", "有"), ("個別チャット", "対応なし ※貴社対応"))
    for (a, b), (k, v) in zip(((25, 29), (33, 37), (41, 46), (50, 54), (58, 62), (66, 70)), mon):
        put(s, a, k)
        put(s, b, v)
    after = (("友だち数", f"{n(sim[28][-1]):,}人"), ("ブロック率", pct(sim[27][-1], 0)),
             ("追える人数", f"{n(sim[29][-1]):,}人"),
             (sim["name1"], f"{sum(n(x) for x in sim[36][1:])}件（6か月）"),
             (sim["name2"], f"{sum(n(x) for x in sim[37][1:])}件（6か月）"))
    for (a, b), (k, v) in zip(((80, 9), (82, 12), (84, 14), (86, 18), (22, 26)), after):
        put(s, a, k)
        put(s, b, v)


def p_hiyou_detail(s, sim, lead_src):
    lead = clone(s, find(lead_src, 19))
    write(lead, ["月ごとの友だち数・追える人数・セミナー予約・サロン入会と、費用の見込み。"])
    put(s, 6, "想定の費用対効果（月ごと）")
    head = dict(fill=NAVY, color="FFFFFF", bold=True, size=12)
    lab = dict(fill="F2F2F2", color=NAVY, bold=True, size=12, align="l")
    val = dict(fill="FFFFFF", color=TEXT, size=12)

    def row(label, init, vals):
        return [dict(text=label, **lab), dict(text=init, **val)] + [dict(text=v, **val) for v in vals]

    m = range(1, 7)
    rows = [[dict(text="項目", **head), dict(text="初期", **head)] + [dict(text=f"{k}か月目", **head) for k in m],
            row("時期", "--", [sim[12][k] for k in m]),
            row("LP UU", "--", [f"{n(sim[14][k]):,}" for k in m]),
            row("友だちの増加", "--", [f"{n(sim[26][k]):,}人" for k in m]),
            row("友だち数", "--", [f"{n(sim[28][k]):,}人" for k in m]),
            row("ブロック率", "--", [pct(sim[27][k]) for k in m]),
            row("追える人数", "--", [f"{n(sim[29][k]):,}人" for k in m]),
            row("クリック数", "--", [f"{n(sim[35][k]):,}" for k in m]),
            row(sim["name1"], "--", [f"{n(sim[36][k])}件" for k in m]),
            row(sim["name2"], "--", [f"{n(sim[37][k])}件" for k in m]),
            row("費用（税抜）", f"{n(sim[73][0]):,}円", [f"{n(sim[73][k]):,}円" for k in m])]
    add_table(s, 0.5, 2.2, [1.55, 1.1] + [1.18] * 6, [0.36] * len(rows), rows, name="月ごとの見込み")
    add_box(s, 0.5, 2.2 + 0.36 * len(rows) + 0.1, 9.9, 0.45,
            [runs(("※LP UUは、広告のクリック数から置いた想定数値（イベント期＝1回のイベント／平常期＝少量の広告）。"
                   "友だちの増加＝ポップアップツール＋サンクスLINE誘導。追える人数＝ブロックされずに届く友だちの数", dict(size=9, color=GRAY)))],
            anchor="t", ins=(0, 0, 0, 0), name="注記")


def p_kangae15(s):
    write(find(s, 3), ["LINE友だち追加～15日を境に大きくアクティブ率が下がる（ブロック率が上がる）",
                       "為、段階に合わせて配信を変え、イベントの無い時期は回数を絞る。"])


def p_zentai(s):
    put(s, 99, "予約完了")
    put(s, 41, "予約完了")
    put(s, 103, "参加")
    put(s, 101, "サロン入会")
    put(s, 62, "・前日のお知らせ")
    write(find(s, 58), ["・見どころ動画", "・参加特典", "・ステップ配信"])
    write(find(s, 108), ["平常期の配信"])
    write(find(s, 10), [[("●イベント直後に、サロンの", "n")], [("申込みフォームを案内", "n")],
                        [("（申込みはいつでも可能）", "n")], None,
                        [("※", "r"), ("個別チャット対応は効率化のみで基本貴社対応", "r")]])
    write(find(s, 106), ["AD", None, None, None, None, None])


def p_track(s):
    put(s, 6, "イベント期の層（予約した方）")
    put(s, 9, "平常期の層（予約していない方）")
    struct(find(s, 7), [("h", "特徴"), ("t", "セミナーを予約済み・参加の意思がある"), ("t", "当日までに日が空く"), ("e",),
                        ("h", "メリット"), ("t", "日程ごとの前日・当日のお知らせで、当日の参加につなげられる"), ("e",),
                        ("h", "デメリット"), ("t", "当日に参加しないと、そのまま離れやすい")])
    struct(find(s, 11), [("h", "特徴"), ("t", "LPを見たが、今回は申し込んでいない"), ("t", "次のイベントまで間がある"), ("e",),
                         ("h", "メリット"), ("t", "平常期の配信で、次のイベントまでつながりを保てる"),
                         ("t", "次のイベントの案内を、広告より先に届けられる"), ("e",),
                         ("h", "デメリット"), ("t", "予約・入会まで時間がかかる")])
    struct(find(s, 13), [("h", "結論・推奨"),
                         ("t", "友だちの入口（ポップアップツール／サンクスLINE誘導）で自動で分け、予約した方は当日の参加へ、"),
                         ("t", "予約していない方は平常期の配信で次のイベントへつなげる。")])


def p_day15(s, overlay=False):
    put(s, 12, "検索直後の疑問に即応し、15日以内でセミナー予約へ誘導する")
    for sid in (9, 16, 26):
        write(find(s, sid), [find(s, sid).text_frame.paragraphs[0].text, "--"])
    put(s, 5, "3問タップで検討段階を分岐・タグ付け", "（--）")
    for sid in (10, 18, 28, 38, 11, 20, 30, 39, 41):
        put(s, sid, "--")
    for sid in (13, 22, 32, 40):
        put(s, sid, "--")
    for sid in (43, 45, 47, 49, 51):
        put(s, sid, "[検索データ]", "--")
    if overlay:
        put(s, 27, "【その他セグメント配信の活用に関して】", "Day0のアンケートの回答に合わせて、配信の中身を出し分ける。")


def p_kyumin(s):
    put(s, 4, "休眠リード開拓と参加後フォローアップ")
    put(s, 17, "--件", "保有リスト")


def p_kyumin_season(s):
    put(s, 4, "休眠リード開拓と参加後フォローアップ")
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip().endswith("月") and len(sh.text_frame.text.strip()) <= 3:
            fill_color(sh, LBLUE, SUB)
    put(s, 43, "--")


def p_choki(s):
    put(s, 4, "D. 参加後フォローアップ")
    put(s, 8, "参加直後", "--")
    put(s, 11, "参加後30〜60日", "--")


# ---------------------------------------------------------------- 新規友だち獲得・費用感・サポート体制
def p_thanks(s, second=False):
    put(s, 26, "初期：15万円")
    drop(s, 10)                                              # 社内向けの「2026年7月〜の価格改定」の注記
    if second:
        put(s, 62, "セミナー予約")
        put(s, 1025, "サロン申込み")
        put(s, 1045, "予約内容の登録", "↓", "LINEに遷移", "↓", "前日・当日のお知らせが開始")
        put(s, 1047, "申込み内容の登録", "↓", "LINEに遷移", "↓", "入会後のご案内が開始")
        a, b = geo(find(s, 1028)), geo(find(s, 1031))
        drop(s, 1028, 1031)
        slot(s, a[0], a[1], b[0] + b[2] - a[0], max(a[1] + a[3], b[1] + b[3]) - a[1], "LPの予約フォーム")
        x, _, w, _ = geo(find(s, 1025))
        _, y, _, h = geo(find(s, 1049))
        drop(s, 1049)
        slot(s, x, y, w, h, "サロンの申込みページ")


def p_seika_sim(s, sim):
    write(find(s, 5), ["固定費用を最低限に抑えた成果報酬プラン。", "友だち追加はポップアップ・サンクスLINE誘導で実施。"])
    tbl = find(s, 10).table
    m = range(1, 7)
    flags = {5: 18, 6: 19, 4: 17, 7: 20, 8: 21}
    vals = {1: 14, 2: 15, 3: 16, 9: 22, 10: 23, 11: 24, 12: 25, 13: 26, 15: 28, 16: 29, 17: 30, 18: 31,
            19: 32, 20: 33, 21: 34, 22: 35, 23: 36, 24: 37, 25: 38, 26: 39, 27: 40}
    keep = copy.deepcopy(tbl.cell(5, 4)._tc)                 # 「実施予定」の書式（5行目）。空のセルにも同じ書式を使う
    for k in m:
        tbl.cell(5, 3 + k)._tc.getparent().replace(tbl.cell(5, 3 + k)._tc, copy.deepcopy(keep))
    for r in range(1, 28):
        if r in flags:
            for k in m:
                set_cell(tbl.cell(r, 3 + k), "実施予定" if sim[flags[r]][k] else "", like=tbl.cell(5, 3 + k), force=True)
            continue
        if r == 14:
            set_cell(tbl.cell(r, 3), pct(sim[27][0]))
            for k in m:
                set_cell(tbl.cell(r, 3 + k), pct(sim[27][k]))
            continue
        if r in vals:
            row = sim[vals[r]]
            set_cell(tbl.cell(r, 3), row[0] if isinstance(row[0], str) else f"{n(row[0] or 0):,}")
            for k in m:
                set_cell(tbl.cell(r, 3 + k), f"{n(row[k]):,}")
    set_cell(tbl.cell(10, 1), "ポップアップツール")
    set_cell(tbl.cell(5, 1), "ポップアップツール")
    set_cell(tbl.cell(23, 1), sim["name1"])
    set_cell(tbl.cell(24, 1), sim["name2"])
    set_cell(tbl.cell(25, 1), "-")
    set_cell(tbl.cell(26, 1), "-")


def seika_rows(sim):
    """成果報酬プランの費用：固定費＝ツール費＋アカウント費（コンサル費は成果報酬に置き換わる）＋成果件数×成果単価"""
    m = range(1, 7)
    fixed = [sum((sim[r][k] or 0) for r in (65, 66, 67, 68, 70, 71, 72)) for k in m]
    reward = [n(sim[36][k]) * sim["tanka1"] for k in m]     # 成果報酬はセミナー予約だけ（サロン入会は計測できない。2026-10-20 佐村さん）
    total = [f + r for f, r in zip(fixed, reward)]
    cvs = [n(sim[36][k]) + n(sim[37][k]) for k in m]
    clicks = [sim[35][k] for k in m]
    return dict(init=sim[73][0], fixed=fixed, reward=reward, total=total, cvs=cvs, clicks=clicks,
                getu=[sim[46][k] for k in m], cvr=[sim[41][k] for k in m])


def p_seika_cost(s, sim):
    write(find(s, 5), ["固定費用を最低限に抑えた成果報酬プラン。", "友だち追加はポップアップ・サンクスLINE誘導で実施。"])
    d = seika_rows(sim)
    set_cell(find(s, 12).table.cell(0, 0), yen(d["init"]))
    lo, hi = min(d["total"]), max(d["total"])
    set_cell(find(s, 14).table.cell(0, 0), f"{yen(lo)} – {n(hi):,}\n※セミナー予約 {n(sim['tanka1']):,}円／件で設定")
    t = find(s, 2).table
    for k in range(6):
        c = 4 + k
        set_cell(t.cell(1, c), "¥0")
        set_cell(t.cell(2, c), yen(d["fixed"][k]))
        set_cell(t.cell(3, c), yen(d["getu"][k]))
        set_cell(t.cell(4, c), yen(d["total"][k] / d["clicks"][k]) if d["clicks"][k] else "--")
        set_cell(t.cell(5, c), pct(d["cvr"][k], 2))
        set_cell(t.cell(6, c), yen(d["total"][k] / d["cvs"][k]) if d["cvs"][k] else "--")
        set_cell(t.cell(7, c), yen(d["reward"][k]))
        set_cell(t.cell(8, c), yen(d["total"][k]))
    set_cell(t.cell(2, 3), yen(d["init"]))
    set_cell(t.cell(8, 3), yen(d["init"]))


def p_service(s):
    put(s, 30, "カスタマージャーニー（8フェーズ）")


def p_tsudo(s):
    t = find(s, 14).table
    for r in range(len(t.rows)):
        if t.cell(r, 4).text.startswith("サンクスLINE誘導ツール"):
            set_cell(t.cell(r, 5), "¥150,000")
            set_cell(t.cell(r, 6), "¥30,000〜", like=t.cell(r, 5))
        if t.cell(r, 4).text.startswith("離脱防止ツール"):
            set_cell(t.cell(r, 6), "¥30,000", like=t.cell(r, 5), force=True)


def p_schedule(s):
    put(s, 19, "サイト・LINEへの実装は、初期設計が整い次第、順次進行")
    drop(s, 21, 22)                                          # 「約1か月」の札と点線の矢印


def p_jirei(s):
    for sid in (3, 8, 14):
        put(s, sid, "--")
    for sid in (10, 12, 16, 18, 20):
        put(s, sid, "--")


# ---------------------------------------------------------------- 組み立て
def main(ver3_path, sim_path, out):
    sim = load_sim(sim_path)
    prs = Presentation(FMT)
    sl = list(prs.slides)
    ver3 = Presentation(ver3_path)
    v3 = list(ver3.slides)
    assert find(v3[5], 2).text_frame.text == "なぜLINEか"
    assert find(v3[12], 2).text_frame.text == "料金"
    assert find(v3[14], 9).text_frame.text == "成果報酬型プラン"
    assert find(v3[7], 2).text_frame.text == "他社のLINE運用"
    S = lambda i: sl[i - 1]                                    # noqa: E731（FMTの何枚目か）

    p_cover(S(1))
    p_doshen(S(3))
    p_doshen_table(S(4))
    p_genjo(S(5))
    p_season(S(6))
    p_zengo(S(7), 4)
    p_zengo(S(8), 9)
    p_zengo_sum(S(9))
    p_study(S(10))
    p_study_axis(S(11), v3[7])
    p_yoken(S(13), n(sim[29][-1]))
    p_yoken2(S(14))
    p_hiyou(S(15), sim)
    p_hiyou_detail(S(16), sim, S(4))
    p_kangae15(S(18))
    p_zentai(S(19))
    for i in (20, 21, 24):
        p_track(S(i))
    p_day15(S(22), overlay=True)
    p_day15(S(23))
    p_kyumin(S(25))
    p_kyumin_season(S(26))
    p_choki(S(27))
    p_thanks(S(32))
    p_thanks(S(33), second=True)
    p_seika_sim(S(37), sim)
    p_seika_cost(S(38), sim)
    p_service(S(39))
    p_tsudo(S(40))
    p_schedule(S(45))
    p_jirei(S(62))

    # FMTに入る場所が無いver3の3枚を写す（料金・成果報酬型プラン・なぜLINEか）
    v_seika = copy_slide(prs, v3[14])
    v_ryokin = copy_slide(prs, v3[12])
    v_naze = copy_slide(prs, v3[5])
    order = sl[:38] + [v_seika, sl[38], v_ryokin] + sl[39:46] + [v_naze] + sl[46:]
    lst = prs.slides._sldIdLst
    by_part = {prs.part.related_part(e.rId): e for e in lst}
    for e in list(lst):
        lst.remove(e)
    for s in order:
        lst.append(by_part[s.part])
    all_meiryo(order)
    prs.save(out)
    print("saved:", Path(out).name, len(prs.slides), "枚")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2], sys.argv[3] if len(sys.argv) > 3 else OUT)
