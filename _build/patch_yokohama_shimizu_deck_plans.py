# -*- coding: utf-8 -*-
"""横浜シミズ：数字・費用のページを、タナカホームの3ページの形に作り直す（2026-10-20 殿村さん「毎回FMTが変わるのはなんなの？」「修正してください」「15万ですすめて」）。

    python -m _build.patch_yokohama_shimizu_deck_plans <横浜シミズ資料.pptx> <タナカホーム資料.pptx> <出力.pptx>

  タナカホーム（202610_タナカホーム株式会社御中_LINE公式アカウント運用のご提案.pptx）の13〜15ページ
  「概算お見積り」「料金プラン：3つのプランの費用」「運用内容：プランごとの対応範囲」をそのままコピーし、数字と言葉だけ入れ替える。
  横浜シミズの旧13（月ごとの表）・旧14（1プランの費用表）は外し、この3ページを同じ場所に入れる（16枚→17枚）。
  プラン（タナカホームと同じく、初期費用は3プランとも同じ。月額は運用の範囲で変わる）：
    ① 初期設計・構築のみ  初期295,000円〜／月額65,000円〜（ツール・アカウントだけ）
    ② 運用代行・効率改善  初期295,000円〜／月額215,000円〜（運用150,000円〜）＝SIM ver2.0（殿村さん確定）
    ③ 運用代行フルサポート 初期295,000円〜／月額315,000円〜（運用250,000円〜＝仮の案）
    勤務前の案内の送信・個別対応はDYMでは持たない（殿村さん「送信代行はいらない」「個別対応はこっちで賄わない」）
    ※最初の版は10/19のAIの金額（①初期795,000・③初期695,000）を混ぜて、構築だけのプランが一番高くなっていた（殿村さん「むちゃくちゃ」）
    ツール・LINE関連：初期 離脱防止15,000＋サンクスLINE誘導150,000＋アカウント開設30,000／月額 30,000＋30,000〜＋利用料5,000〜
"""
import copy
import sys

from pptx import Presentation
from pptx.oxml.ns import qn

from _build.patch_tanaka_v5 import copy_slide

TANAKA_PAGES = (13, 14, 15)
PLAN_HEAD = ["項目", "プラン①\n初期設計・構築のみ", "プラン②\n運用代行・効率改善", "プラン③\n運用代行フルサポート"]


def set_text(tf, text):
    """1段落目・1つ目のrunの書式で文字を入れる。改行は段落を分ける"""
    ps = tf._txBody.findall(qn("a:p"))
    proto = copy.deepcopy(ps[0])
    for p in ps:
        tf._txBody.remove(p)
    for line in text.split("\n"):
        p = copy.deepcopy(proto)
        runs = p.findall(qn("a:r"))
        for r in runs[1:]:
            p.remove(r)
        runs[0].find(qn("a:t")).text = line
        tf._txBody.append(p)


def by_id(s, sid):
    return next(sh for sh in s.shapes if sh.shape_id == sid)


def fill_estimate(s):
    assert by_id(s, 339).text_frame.text.startswith("概算お見積り")
    texts = {
        339: "概算お見積り：月215,000円で、6か月目に新規応募 月6件・再応募 月3件",
        381: "プラン②｜友だちの入口を増やし、勤務の連絡と次回イベントの募集で、新規応募・再応募を増やす",
        343: "¥295,000〜",
        344: "¥215,000〜",
        # 初期の実施内容
        365: "離脱防止", 366: "導入",
        367: "サンクス\nLINE誘導", 368: "導入",
        369: "サイトの\nLINEボタン", 370: "設置",
        371: "ステップ配信", 372: "4通を構築",
        373: "勤務前の案内", 374: "型をつくる",
        # 月次の実施内容
        349: "企画配信", 350: "月4回",
        351: "ステップ配信", 352: "4通の改善",
        353: "セグメント配信", 354: "設定・改善",
        355: "勤務前の案内", 356: "型の改善",
        357: "月次レポート", 358: "1回",
        359: "友だちの入口", 360: "運用",
        # 6か月後（SIM ver2.0 の6か月目）
        361: "友だち数", 375: "984人",
        362: "新規応募", 376: "6件/月",
        364: "再応募", 378: "3件/月",
    }
    for sid, t in texts.items():
        set_text(by_id(s, sid).text_frame, t)


def fill_table(tbl, rows):
    for r, vals in zip(tbl.rows, rows):
        for c, v in zip(r.cells, vals):
            set_text(c.text_frame, v)


def fill_price(s):
    assert by_id(s, 635).text_frame.text.startswith("料金プラン")
    fill_table(by_id(s, 639).table, [
        PLAN_HEAD,
        ["初期費用", "295,000円〜", "295,000円〜", "295,000円〜"],
        ["月額運用費", "なし", "150,000円〜", "250,000円〜"],
        ["月額ツール・LINE関連費", "65,000円〜", "65,000円〜", "65,000円〜"],
        ["月額費用", "65,000円〜", "215,000円〜", "315,000円〜"],
    ])
    notes = [sh for sh in s.shapes if sh.name == "注記"]
    assert len(notes) == 4
    for sh, t in zip(notes, [
        "※初期費用＝運用設計・構築 100,000円〜＋離脱防止 15,000円＋サンクスLINE誘導 150,000円＋LINE公式アカウント開設 30,000円",
        "※月額ツール・LINE関連費＝離脱防止 30,000円＋サンクスLINE誘導 30,000円〜＋LINE公式アカウント利用料 5,000円〜（ライトプラン想定）",
        "※契約期間は6か月〜。概算お見積りのページはプラン②の場合",
        "※施策内容・配信本数・連携範囲等により、費用は個別にお見積もりいたします。",
    ]):
        set_text(sh.text_frame, t)


def fill_scope(s):
    assert by_id(s, 2).text_frame.text.startswith("運用内容")
    shape = by_id(s, 6)
    tbl = shape.table
    trs = tbl._tbl.findall(qn("a:tr"))
    assert "LINE広告" in "".join(t.text or "" for t in trs[-1].iter(qn("a:t")))
    shape.height -= int(trs[-1].get("h"))
    tbl._tbl.remove(trs[-1])                                  # LINE広告の行は今回の提案に無い
    fill_table(tbl, [
        PLAN_HEAD,
        ["離脱防止・\nサンクスLINE誘導", "導入・運用", "導入・運用", "導入・運用"],
        ["サイトのLINEボタン", "設置", "設置", "設置"],
        ["ステップ配信", "14日間・4通\n（構築のみ）", "14日間・4通\n（構築・改善）", "14日間・4通\n（構築・改善）"],
        ["企画配信\n（次回イベントの募集）", "なし（貴社で運用）", "月4回\n（設計・作成）", "月4回\n（設計・作成・画像制作）"],
        ["勤務前の案内", "型の作成のみ", "型の作成・改善\n（送信は貴社）", "型の作成・改善\n（送信は貴社）"],
        ["結果の分析・改善\n（月次レポート）", "なし", "対応", "対応"],
    ])


def main(src, tanaka, out):
    prs = Presentation(src)
    assert len(prs.slides) == 16
    old13, old14 = prs.slides[12], prs.slides[13]
    assert old13.shapes[0].text_frame.text.startswith("概算お見積り｜")
    assert old14.shapes[0].text_frame.text.startswith("運用プラン・費用｜")
    tk = Presentation(tanaka)
    new = [copy_slide(prs, tk.slides[n - 1]) for n in TANAKA_PAGES]
    fill_estimate(new[0])
    fill_price(new[1])
    fill_scope(new[2])

    lst = prs.slides._sldIdLst
    ids = list(lst)                                           # 0〜15 が元の16枚、16〜18 が足した3枚
    for el in ids[16:]:
        lst.remove(el)
    for el in (ids[12], ids[13]):                             # 旧13・旧14を外す
        prs.part.drop_rel(el.rId)
        lst.remove(el)
    anchor = ids[11]
    for el in reversed(ids[16:]):
        anchor.addnext(el)
    assert len(prs.slides) == 17
    prs.save(out)
    print("saved:", out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2], sys.argv[3])
