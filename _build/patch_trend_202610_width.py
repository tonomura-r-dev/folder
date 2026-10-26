"""
10月号 P.3〜P.7：カードの横幅を文字量に合わせて縮め、左右中央に置く（2026-10-26 殿村さん指示）。
同じページのカードは、いちばん長い行に合わせて同じ幅にそろえる。
使い方: python3 _build/patch_trend_202610_width.py <pptx>
"""
import sys
from pptx import Presentation
from pptx.util import Inches, Emu
from pptx.enum.shapes import MSO_SHAPE

F = sys.argv[1]
prs = Presentation(F)
SW = prs.slide_width
MAX_W = Inches(9.64)
PAD = Inches(0.45)  # PowerPointのメイリオは確認用フォントより幅が出るので余裕を持たせる


def text_w(s, pt):
    return Inches(sum((0.55 if ord(c) < 0x2000 else 1.0) for c in s) * pt / 72)


def is_card(x):
    return x.shape_type == 1 and getattr(x, "auto_shape_type", None) == MSO_SHAPE.ROUNDED_RECTANGLE \
        and x.top > Inches(3.0) and x.has_text_frame


for slide in list(prs.slides)[2:7]:
    cards = [x for x in slide.shapes if is_card(x)]
    if not cards:
        continue
    need = 0
    for c in cards:
        tf = c.text_frame
        for p in tf.paragraphs:
            if not p.runs:
                continue
            pt = p.runs[0].font.size.pt if p.runs[0].font.size else 13
            for seg in p.text.split("\x0b"):
                need = max(need, text_w(seg, pt) + tf.margin_left + tf.margin_right + PAD)
    w = min(int(need), MAX_W)
    for c in cards:
        c.width, c.left = w, int((SW - w) / 2)
    print(slide.slide_id, round(Emu(w).inches, 2))

prs.save(F)
print("saved", F)
