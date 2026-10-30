# -*- coding: utf-8 -*-
"""ジム・フィットネス業界 ver5（47枚）を、10/29の提案FMT（MAX版・53枚）の章立てにそろえる（2026-10-30 殿村さん「FMTにそろえるだけ」）。

    python -m _build.build_fitness_fmt1029 <ver5.pptx> <FMT最新.pptx> <FMT具体施策枠のみ.pptx> <塾 新FMT版.pptx> <出力.pptx>

  形は塾の新FMT版（_build/build_juku_fmt1029.py）と同じ。
  ・01 施策背景：FMTの弊社紹介〜ポジショニング・LINEOA公式基礎・実例紹介＋ジムの背景（検索需要と店舗数・業界のLINEの実態）
  ・02 LINE施策全体図：FMTのファネル・施策パターン・全体図・施策一覧＋ジムの全体設計・8フェーズ・前後検索・シーズナリティ・年間の配信設計
  ・03 施策詳細：FMTの具体施策のイメージ8枚（ジムの中身で埋める）＋カーブスの事例・ステップ配信・リッチメニュー・あいさつ・企画・シミュレーション
  ・04 料金形態／05 サポート体制／06 次回ご提案に向けて：FMTのページ（旧「ご提供プラン」は外す）
  ・FMTの新規友だちの追加〜提供施策・裏表紙のあとに Appendix（旧の汎用ページ：背景・考え方・施策展開図・改善モデル・具体施策18）
  FMTから入れたページは黄色を白に、料金表のサンクスLINE誘導は初期15万円・月額3万円〜、離脱防止は初期1.5万円・月額3万円。
  ジムのページの中の「P◯」は新しい並びに振り直す。
"""
import sys

from pptx import Presentation

import _build.build_juku_fmt1029 as J0
from _build.build_juku_fmt1029 import fill_gutai, renumber
from _build.patch_tanaka_v5 import copy_slide

GUTAI = [  # 札, [入口, LINE, 配信, ゴール], (目的, 対象, 中身, 見る数字)  右の文は「|」で改行
    ("LINEの友だちへ",
     [("サイト・広告・店頭・SNS", "離脱防止ポップアップ／店頭POP・チラシのQR／|サンクスLINE誘導"),
      ("友だち追加", "無料体験と30秒診断を受け取る"), ("あいさつメッセージ", "30秒 体づくりタイプ診断"), "友だちとしてつながり続ける"],
     ("入会を決めていない方とも、|LINEでつながる",
      "ジムを探している方|（「続けられるか不安」な方も含む）",
      "離脱防止ポップアップ／|店頭POP・チラシのQR／|体験予約の完了画面からのLINE誘導",
      "友だち追加数／入口ごとの追加数")),
    ("露出の強化",
     ["友だち全体", ("企画配信", "1月・4〜7月の需要の波に合わせる"),
      "期限つきの無料体験・入会特典を毎月案内", "体験予約"],
     ("需要の波に合わせて、|体験予約のピークをつくる",
      "友だち全体|（閑散期に追加した方も含む）",
      "1月の入会金無料キャンペーン／|特典を毎月替え、締切前に告知／|会員の声の動画",
      "配信到達数・開封率／|企画配信からの体験予約数")),
    ("CVの最大化",
     ["友だち追加から14日以内の方", ("診断4回", "体づくり・性格・ダイエット・続けられるか"),
      ("14日間の配信", "無料体験のご案内はDay0・6・14"), "体験予約（CV①）"],
     ("反応が高い追加後14日の間に、|体験予約へ進んでいただく",
      "友だち追加から14日以内の方",
      "答えるだけの診断を4回／|1分料金シミュレーター／|不安Q&A（勧誘・解約・混雑）",
      "体験予約数／診断の回答率")),
    ("CPOの改善（CV質）",
     ["体験を予約した方", "予約の確認・通知メッセージ",
      ("体験の前日までの配信", "持ち物・アクセス・当日の流れ"), "来店（CV②）・入会"],
     ("体験を予約した後、当日来られる方を|増やし、入会につなげる",
      "体験を予約した方",
      "予約確定・前日のお知らせ／|「見学だけOK・入会を迫りません」／|日程の変更はトークで受け付け",
      "来店率／キャンセル率／|来店から入会の割合")),
    ("休眠顧客の有効化",
     ["体験だけで止まった方・休眠会員", "通知メッセージ（電話番号でLINEに届く）",
      ("1月・新生活の時期の再案内", "入会キャンペーン・期限つきの無料体験"), "体験予約・再入会"],
     ("止まっている方に、需要が|高まる時期にもう一度ご案内する",
      "体験に来たが入会しなかった方／|休会・退会した方の保有リスト",
      "通知メッセージで再アプローチ／|1月・4月の前に再案内／|前回の体験に触れた文面",
      "再案内からの体験予約数／|再入会数")),
    ("マネタイズ最大化",
     ["診断に答えた方", ("診断タグ＋行動タグ", "目的・経験・頻度・性格タイプ／予約の有無"),
      "同じタイプの会員の声・1回あたりの料金", "入会"],
     ("一人ひとりに合う内容を届け、|入会まで進めていただく",
      "診断に答えた方／|体験に来て入会を迷っている方",
      "回答と行動で配信を出し分け／|「週2回なら1回あたり」の料金／|入会金無料・事務手数料0円の案内",
      "入会数／入会率／ブロック率")),
    ("リピート施策",
     ["入会した方", "入会タグでリッチメニューを会員用に",
      ("入会後14日の配信", "持ち物・始め方・効果が出る時期"), "継続・PT／物販の申込"],
     ("入会直後の離脱を防ぎ、|続けていただく",
      "入会した方",
      "入会後14日の配信（4通）／|フォームチェック・体組成測定の予約／|来店が2週間途切れた方へのフォロー",
      "PT・物販の申込数／|継続率（会員管理システムで比較）")),
    ("コミュニティ運営",
     ["会員", ("会員のタグ付け", "リッチメニュー③に会員向けの入口"),
      "会員の声・イベント・お友だち紹介の案内", "継続・紹介（CV③）"],
     ("会員とのつながりを強くし、|紹介につなげる",
      "在籍中の会員",
      "会員の声・店舗のイベントの案内／|LINEの転送でお友だちに紹介／|紹介の特典の案内（中身は貴社で決める）",
      "紹介からの友だち追加数／|紹介からの体験予約数")),
]


def main(gym_path, fmt_path, fmt_empty_path, juku_path, out):
    J0.TEXT_FIX = [("貴社保有リスト をターゲット", "貴社保有リストをターゲット"),
                   ("転職ユーザーのカスタマージャーニー", "入会を検討する方のカスタマージャーニー")]
    prs = Presentation(gym_path)
    fmt, emp, juku = Presentation(fmt_path), Presentation(fmt_empty_path), Presentation(juku_path)
    assert len(prs.slides) == 47 and len(fmt.slides) == 53 and len(emp.slides) == 53
    G = {i + 1: s for i, s in enumerate(prs.slides)}

    def add(src):
        s = copy_slide(prs, src)
        J0.fmt_fix(s)
        return s

    F = {n: add(fmt.slides[n - 1]) for n in list(range(2, 7)) + list(range(8, 54)) if not 18 <= n <= 25}
    E = {n: add(emp.slides[n - 1]) for n in range(18, 26)}
    for n, page in zip(range(18, 26), GUTAI):
        fill_gutai(E[n], page)
    apx = copy_slide(prs, juku.slides[74])           # 「アジェンダ／Appendix（参考資料）」
    assert "Appendix" in "".join(sh.text_frame.text for sh in apx.shapes if sh.has_text_frame)

    order = ([G[1]] + [F[n] for n in range(2, 7)] + [G[15], G[9], F[8], F[9], F[10]]
             + [F[11], F[12], G[13], G[14], G[23], F[13], F[14], F[15], F[16], G[19], G[24], G[20], G[16], G[17], G[18]]
             + [F[17]] + [E[n] for n in range(18, 26)]
             + [G[n] for n in (10, 11, 22, 7, 25, 26, 27, 28, 29, 30, 31, 32, 38)]
             + [F[n] for n in range(29, 54)]
             + [apx] + [G[n] for n in (3, 5, 4, 6, 8, 33, 34, 35, 36, 37, 42, 43, 44, 45, 46)])
    assert len(order) == 89, len(order)
    mapping = {}
    for i, s in enumerate(order, 1):
        for n, gs in G.items():
            if gs is s:
                mapping[n] = i
    for n, s in G.items():
        if n in mapping:
            renumber(s, mapping)

    lst = prs.slides._sldIdLst
    ids = {id(s._element): sid for s, sid in zip(prs.slides, list(lst))}
    keep = [ids[id(s._element)] for s in order]
    keep_set = set(id(x) for x in keep)
    for sid in list(lst):
        lst.remove(sid)
        if id(sid) not in keep_set:
            prs.part.drop_rel(sid.rId)
    for sid in keep:
        lst.append(sid)
    assert len(prs.slides) == 89
    prs.save(out)
    print("saved:", out)


if __name__ == "__main__":
    main(*sys.argv[1:6])
