# -*- coding: utf-8 -*-
"""提案FMT（MAX版／最強資料）P41〜P44「提案に向けて明確化したいポイント①〜④」の表を、殿村さんが直したP43の形にそろえる（2026-10-29）。

    python -m _build.patch_fmt_max_tables_p43 <殿村さんPC保存版.pptx> <出力.pptx>

  P43：表の左端0.92・幅8.92インチ（中央）、見出し行 y1.95 高さ0.42、行は y2.41〜、下の枠も 0.92／8.92。
  P41・P42・P44の表の左端・幅・見出し行を同じにする。文字・色・行の数は変えない。
"""
import copy
import sys

from pptx import Presentation
from pptx.oxml.ns import qn

E = 914400
L, W, HEAD_Y, HEAD_H, ROW_Y = 0.92, 8.92, 1.95, 0.42, 2.41


def table_cells(s):
    return [sh for sh in s.shapes if sh.name.startswith("Rectangle") and 1.85 * E < sh.top < 6.45 * E]


def fit(s, cols, row_h, row_step, foot=None):
    """cols = [(左端, 幅)…]。セルを列（今の左端の順）と行（今の上端の順）で振り直す"""
    cells = table_cells(s)
    xs = sorted({sh.left for sh in cells})
    ys = sorted({sh.top for sh in cells})
    assert len(xs) == len(cols), (len(xs), len(cols))
    for sh in cells:
        c, r = xs.index(sh.left), ys.index(sh.top)
        sh.left, sh.width = int(cols[c][0] * E), int(cols[c][1] * E)
        if r == 0:
            sh.top, sh.height = int(HEAD_Y * E), int(HEAD_H * E)
        else:
            sh.top, sh.height = int((ROW_Y + (r - 1) * row_step) * E), int(row_h * E)
    if foot:
        name, y, h = foot
        f = next(sh for sh in s.shapes if sh.name == name)
        f.left, f.width = int(L * E), int(W * E)
        if y is not None:
            f.top, f.height = int(y * E), int(h * E)


def br_after(sh, word):
    """word の直後で改行（a:br）。ランを切って同じ書式の a:br を入れる"""
    for p in sh.text_frame.paragraphs:
        for r in p.runs:
            if word in r.text:
                head, tail = r.text.split(word, 1)
                r.text = head + word
                new = copy.deepcopy(r._r)
                new.find(qn("a:t")).text = tail
                br = copy.deepcopy(r._r)
                br.tag = qn("a:br")
                br.remove(br.find(qn("a:t")))
                r._r.addnext(br)
                br.addnext(new)
                return
    raise SystemExit("見つからない: " + word)


def main(src, out):
    prs = Presentation(src)
    sl = prs.slides
    assert "④" in sl[43].shapes[0].text_frame.text
    # P43 の列：項目 0.92/1.87・確認内容 2.82/4.37・記入 7.22/2.62
    # 言葉の途中で折り返さない幅にする（P41の施策の例・P42の項目は広め）。下の枠は2行のときは行の切れ目で改行
    fit(sl[40], [(0.92, 1.37), (2.32, 4.15), (6.50, 0.90), (7.43, 2.41)], 0.96, 1.0, ("Rectangle 23", 6.45, 0.66))  # P41
    fit(sl[41], [(0.92, 2.27), (3.22, 4.37), (7.62, 2.22)], 0.59, 0.62, ("Rectangle 27", None, None))              # P42
    fit(sl[43], [(0.92, 1.87), (2.82, 7.02)], 0.62, 0.665, ("Rectangle 24", 6.45, 0.66))                            # P44
    br_after(next(sh for sh in sl[40].shapes if sh.name == "Rectangle 23"), "優先度は低い。")
    br_after(next(sh for sh in sl[43].shapes if sh.name == "Rectangle 24"), "配信にかかわる")
    prs.save(out)
    print("saved:", out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
