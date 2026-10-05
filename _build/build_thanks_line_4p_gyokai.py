# -*- coding: utf-8 -*-
"""サンクスLINE誘導 4枚資料・業界別版（2026-10-05 殿村さん指示）。
既存の4枚（20261005_サンクスLINE誘導のご提案.pptx）を土台に、別ファイルで作る。
  1 機能性＝サンクスLINE誘導とは（5ステップの流れ＋想定 100件→60人→30人）
  2 重要性＝CVで接点が終わる／LINEで続く の比較だけ
  3 利用シーン＝人材・不動産投資・ホビー・ハウスメーカー（2×2・CV→LINE→利用シーン）
  4 実績＝LINEヤフー公式事例 旅行（観光バス）・クリニック（皮膚科）。出典・対象・比較条件・期間を書く
事例の出典（2026-10-05 本文で確認）：
  琴平バス https://www.lycbiz.com/jp/case-study/line-official-account/kotobus/
  アクネクリニック https://www.lycbiz.com/jp/case-study/line-official-account/acne-clinic/
  python3 _build/build_thanks_line_4p_gyokai.py <20261005_サンクスLINE誘導のご提案.pptx>
"""
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

sys.path.insert(0, str(Path(__file__).resolve().parent))
from pptx_parts import *  # noqa

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "20261005_サンクスLINE誘導のご提案_業界別.pptx"
PALE = "E4E8F6"
prs = Presentation(sys.argv[1])
SW = prs.slide_width / 360000
L, R = 2.2, SW - 2.2
CW = R - L
MID = (L + R) / 2
S = list(prs.slides)


def shp(s, name):
    return next(x for x in s.shapes if x.name == name)


def head(s, title, lead):
    set_paras(shp(s, "TextBox 1"), [(title, 0)])
    set_paras(shp(s, "TextBox 2"), [(lead, 0)])


def flow(s, x0, y, h, w, steps, size=13, ov=0.35):
    """steps = [(行のリスト, 塗り, 文字色)]。幅 w の矢羽根を ov ずつ重ねて並べる。"""
    for i, (ls, f, tc) in enumerate(steps):
        x = x0 + i * (w - ov)
        shape(s, MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON, x, y, w, h, None, fill=f, adj=0.25)
        pad = 0.3 if i == 0 else 0.7
        label(s, x + pad, y, w - pad - 0.6, h, [(t, size, True, tc, 0) for t in ls], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def step_w(x0, x1, n, ov=0.35):
    return (x1 - x0 + (n - 1) * ov) / n


# ================= 1 機能性：何の機能か＋想定値 =================
s = S[0]
for sh in list(s.shapes):
    if sh.name in [f"Pentagon 5"] + [f"Chevron {k}" for k in (7, 9, 11, 13, 15)] + [f"TextBox {k}" for k in range(6, 18)]:
        sh._element.getparent().remove(sh._element)
set_paras(shp(s, "TextBox 2"), [("CVの完了画面から、LINEの友だち追加へ案内する機能", 0)])
w = step_w(L, R, 5)
flow(s, L, 4.9, 2.6, w, [(["サイト・LP", "来訪"], LGRAY, INK), (["CV"], LGRAY, INK), (["完了画面"], LGRAY, INK),
                         (["サンクス", "LINE誘導"], NAVY, WHITE), (["LINE", "友だち追加"], GREEN, WHITE)], size=14)
label(s, L + (w - 0.35) - 1.0, 7.6, w + 2.0, 0.9, [("予約・申込・問い合わせ・会員登録など", 9.5, False, GRAY, 0)], align=PP_ALIGN.CENTER)
set_paras(shp(s, "TextBox 29"), [("※誘導可能率・友だち追加率は想定値", 0)])

# ================= 2 重要性：CVで接点が終わる／続く =================
s = S[1]
keep_header_only(s)
set_paras(shp(s, "TextBox 2"), [("CVで接点が終わると、面談・商談・購入・継続利用につながらない", 0)])
LW = 4.6
x0 = L + LW + 0.3
w = step_w(x0, R, 4)
label(s, L, 5.4, LW, 3.0, [("サンクスLINE", 12, True, GRAY, 0), ("なし", 18, True, GRAY, 0)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
flow(s, x0, 5.4, 3.0, w, [(["サイト・LP"], LGRAY, INK), (["CV"], LGRAY, INK), (["接点終了"], "F2F2F2", GRAY)], size=15)
hline(s, L, R, 9.15, LGRAY, 0.75)
label(s, L, 10.4, LW, 3.0, [("サンクスLINE", 12, True, GRAY, 0), ("あり", 18, True, GREEN_TX, 0)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
flow(s, x0, 10.4, 3.0, w, [(["サイト・LP"], PALE, NAVY), (["CV"], PALE, NAVY), (["サンクス", "LINE誘導"], NAVY, WHITE),
                           (["LINE上で", "継続接点"], GREEN, WHITE)], size=15)
label(s, x0 + 3 * (w - 0.35), 13.6, w, 0.9, [("その後のご案内・やり取りができる", 10, False, GRAY, 0)], align=PP_ALIGN.CENTER)

# ================= 3 利用シーン：ターゲット4業界 =================
s = S[2]
keep_header_only(s)
head(s, "ターゲット業界別｜サンクスLINE活用イメージ", "CVの後にLINEで送る内容は、業界ごとに変わる")
inds = [("人材", "採用・転職・求人サービス", ["応募・", "会員登録"], ["面談日時の案内", "求人紹介", "選考のご案内"]),
        ("不動産投資", "資料請求・個別相談・セミナー", ["資料請求・", "相談予約"], ["面談のリマインド", "セミナー案内", "物件情報"]),
        ("ホビー（toC）", "商品購入・会員登録", ["購入・", "会員登録"], ["新商品・再入荷", "限定商品の案内", "キャンペーン・クーポン"]),
        ("ハウスメーカー", "資料請求・展示場の来場予約", ["資料請求・", "来場予約"], ["来場前日のリマインド", "施工事例", "見学会・個別相談"])]
TOP, ROWH = 4.5, 6.55
vline(s, MID, TOP + 0.3, TOP + 2 * ROWH - 0.3, LGRAY, 0.75)
hline(s, L, R, TOP + ROWH, LGRAY, 0.75)
for k, (name, sub, cv, scenes) in enumerate(inds):
    qx = L + (k % 2) * (CW / 2) + (0.6 if k % 2 else 0)
    qy = TOP + (k // 2) * ROWH + 0.45
    qw = CW / 2 - 0.6
    # 業界名（番号の丸＋名前＋CV例）
    shape(s, MSO_SHAPE.OVAL, qx, qy, 0.95, 0.95, [(str(k + 1), 12, True, WHITE, 0)], fill=NAVY, margins=(0, 0, 0, 0))
    label(s, qx + 1.2, qy - 0.1, 6.5, 1.15, [(name, 16, True, NAVY, 0)], anchor=MSO_ANCHOR.MIDDLE)
    label(s, qx + 1.2 + len(name) * 0.62 + 0.5, qy, 8.0, 0.95, [(sub, 10, False, GRAY, 0)], anchor=MSO_ANCHOR.MIDDLE)
    # CV → LINE → 利用シーン
    fy, fh = qy + 1.55, 3.6
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, qx, fy + 0.5, 3.6, fh - 1.0, [(t, 12, True, NAVY, 0) for t in cv], fill=PALE, adj=0.15,
          margins=(0.1, 0, 0.1, 0))
    arrow_line(s, qx + 3.75, fy + fh / 2, qx + 4.75, fy + fh / 2, NAVY, 1.75)
    shape(s, MSO_SHAPE.OVAL, qx + 4.9, fy + fh / 2 - 1.15, 2.3, 2.3, [("LINE", 13, True, WHITE, 0)], fill=GREEN, margins=(0, 0, 0, 0))
    arrow_line(s, qx + 7.35, fy + fh / 2, qx + 8.35, fy + fh / 2, NAVY, 1.75)
    ch = 1.0
    for j, t in enumerate(scenes):
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, qx + 8.5, fy + 0.15 + j * (ch + 0.15), qw - 8.5, ch, [(t, 11.5, True, NAVY, 0)],
              fill=WHITE, line=NAVY, lw=0.75, adj=0.5, margins=(0.1, 0, 0.1, 0))
label(s, L, 17.55, CW, 0.5, [("※利用シーンは想定例", 9, False, GRAY, 0)])

# ================= 4 実績：LINEヤフー公式事例 =================
s = S[3]
keep_header_only(s)
head(s, "LINEヤフー活用事例｜旅行・クリニック", "予約後・友だち追加後のLINE配信が、友だち化・予約の増加につながった")
cases = [("旅行（観光バス）", [(["高速バス", "の予約"], PALE, NAVY), (["乗車前日の", "案内をLINEで"], GREEN, WHITE), (["LINE", "友だち追加"], NAVY, WHITE)],
          "予約番号・便名・乗車場所を、乗車前日の夕方に配信（LINE通知メッセージ）",
          "案内を受け取った方の友だち追加率", "約70%",
          ["導入から約2か月で、4,398人が友だち追加（2021年10月〜）", "予約忘れ・乗車前の問い合わせも減少"],
          "出典：LINEヤフー for Business 導入事例（琴平バス／2022年6月公開・取材先調べ）"),
         ("クリニック（皮膚科）", [(["LINE", "友だち追加"], PALE, NAVY), (["ステップ配信", "全5通"], GREEN, WHITE), (["予約"], NAVY, WHITE)],
          "ニキビ・治療法の解説 → 患者の声 → 無料カウンセリングのクーポン、の順に自動配信",
          "LINE経由の予約数", "約20%増",
          ["ステップ配信の実施前との比較", "友だち追加広告による友だち数の増加も含む"],
          "出典：LINEヤフー for Business 導入事例（アクネクリニック／2022年4月公開・取材先調べ）")]
vline(s, MID, 4.8, 16.4, LGRAY, 0.75)
for k, (name, steps, how, metric, num, notes, src) in enumerate(cases):
    cx = L + k * (CW / 2) + (0.6 if k else 0)
    cw = CW / 2 - 0.6
    label(s, cx, 4.5, cw, 1.0, [(name, 17, True, NAVY, 0)], anchor=MSO_ANCHOR.MIDDLE)
    w = step_w(cx, cx + cw, 3)
    flow(s, cx, 5.75, 2.3, w, steps, size=12)
    label(s, cx, 8.2, cw, 0.8, [(how, 10, False, GRAY, 0)], align=PP_ALIGN.CENTER)
    label(s, cx, 9.45, cw, 0.8, [(metric, 13, True, INK, 0)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    label(s, cx, 10.2, cw, 2.6, [(num, 48, True, NAVY, 0)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    label(s, cx, 13.0, cw, 1.6, [(t, 11, False, INK, 2) for t in notes], align=PP_ALIGN.CENTER)
    label(s, cx, 15.1, cw, 0.9, [(src, 8.5, False, GRAY, 0)], align=PP_ALIGN.CENTER)
label(s, L, 17.2, CW, 0.6, [("※どちらもサンクスLINE誘導そのものではなく、LINEでつながった後の活用事例", 9, False, GRAY, 0)])

prs.save(str(OUT))
print("saved:", OUT.name, len(prs.slides), "枚")
