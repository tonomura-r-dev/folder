# -*- coding: utf-8 -*-
"""ワクトリ v4：14ページ「他社のLINE活用のまとめ」の言い過ぎを直す（2026-10-28 殿村さん）。
毎日配信を確認できたのは4社のうち2社（ダイブ・リゾバ.com）だけなので、「他社は毎日」と言い切らない。

    python -m _build.patch_wakutori_v4_p14 <v4.pptx> [出力.pptx]
"""
import sys

from pptx import Presentation

from _build.patch_wakutori_v2_shisaku import by_id, set_lines

TEXTS = {7: "他社は、追加の直後から登録・相談へつないでいる。リゾートバイトの集客で、LINEは必須の接点になっている。"}
CELL = ("2. 追加の後も、毎日配信する", "2. 追加の後も、毎日配信する会社もある")


def main(src, out):
    prs = Presentation(src)
    s = prs.slides[13]
    for sid, t in TEXTS.items():
        set_lines(by_id(s, sid), [t])
    sh = by_id(s, 63)
    assert sh.text_frame.text == CELL[0]
    set_lines(sh, [CELL[1]])
    prs.save(out)
    print("saved:", out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else sys.argv[1])
