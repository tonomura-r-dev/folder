# -*- coding: utf-8 -*-
"""すなお ver5.1 → ver5.2：長期ニーズトラック「D. 参加後フォローアップ」の下に「ほかの配信案」を足す（2026-10-20 殿村さん「ほか案って感じで出しておいて」「他の案」）。見出しは「他の案」。

    python -m _build.patch_sunao_v52_hokaan <ver5.1.pptx> [出力.pptx]

  LPの「よくある質問」と「講師紹介」から2つ。
  ・よくある質問：次のセミナーの案内に添える（地方からも参加できる・録画はない・初心者でも分かる・参加は無料）
  ・講師紹介：入会を迷っている方へ（契約件数全国1位で表彰・週2日の営業で毎年TOT）
  足した分、見出し〜注記をまとめて上げる（見出しの下から下端までの空きの上1/3に寄せ直す）。ほかのページは触らない（出力を省くと上書き）。
"""
import sys

from pptx import Presentation
from pptx.util import Emu, Inches

from _build.arauchi_v4.lib import add_box
from _build.arauchi_v10.build import runs
from _build.build_sunao_v4 import GRAY, NAVY, TEXT
from _build.patch_sunao_pc44_to_v49 import slide

HOKA = [
    ("よくある質問", "次のセミナーの案内に添える",
     "地方からも参加できる・録画はない・初心者でも分かる・参加は無料"),
    ("講師紹介", "入会を迷っている方へ",
     "谷岡さんの経歴：契約件数全国1位で表彰・週2日の営業で毎年TOT"),
]
ROW_H, ROW_GAP = 0.56, 0.08
BAND = "F5F7FA"


def main(src, out):
    prs = Presentation(src)
    s = slide(prs, "長期ニーズトラック：シーズナリティ連動の中長期設計")
    head = next(sh for sh in s.shapes if sh.shape_id == 4)
    note = next(sh for sh in s.shapes if sh.name == "注記")
    assert not any(sh.name == "ほかの配信案" for sh in s.shapes), "ver5.2 済み"

    old_top = Emu(head.top).inches
    old_end = Emu(note.top).inches                                  # 注記の上＝配信内容の下＋0.2
    block = 0.2 + 0.32 + ROW_H * 2 + ROW_GAP                         # 余白＋見出し＋2行
    group = (old_end - old_top) + block + 0.25
    top = 1.1 + (7.05 - 1.1 - group) / 3
    dy = Inches(top - old_top)
    for sh in s.shapes:
        if sh.shape_id in (3, 5, 7):                                 # タイトル・リード・空の本文枠はそのまま
            continue
        sh.top = sh.top + dy

    y = old_end + (top - old_top) - 0.2 + 0.2                        # 配信内容の下＋0.2
    add_box(s, 0.5, y, 9.9, 0.32, [runs(("他の案", dict(size=13, bold=True, color=NAVY)))],
            anchor="ctr", ins=(0.1, 0, 0, 0), name="ほかの配信案")
    y += 0.32
    for i, (label, use, body) in enumerate(HOKA):
        add_box(s, 0.5, y, 2.1, ROW_H, [runs((label, dict(size=13, bold=True, color=NAVY)))],
                fill=BAND, anchor="ctr", ins=(0.2, 0.02, 0.05, 0.02), name="ほかの配信案 項目")
        add_box(s, 2.6, y, 7.8, ROW_H,
                [runs((use, dict(size=12, bold=True, color=TEXT))), runs((body, dict(size=12, color=TEXT)))],
                fill=BAND, anchor="ctr", ins=(0.05, 0.02, 0.1, 0.02), name="ほかの配信案 内容")
        y += ROW_H + ROW_GAP
    note.top = Inches(y - ROW_GAP + 0.2)
    r = note.text_frame.paragraphs[0].runs[0]
    assert r.text.startswith("※配信内容は、貴社LPの")
    r.text = "※配信内容は、貴社LPの「受講生の声」「講師メッセージ」「よくある質問」「講師紹介」をもとにした例"
    prs.save(out)
    print("saved:", out, len(prs.slides), "枚")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else sys.argv[1])
