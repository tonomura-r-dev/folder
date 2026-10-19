# -*- coding: utf-8 -*-
"""すなおコンサルティング：本質課題のページのメッセージを差し替える（2026-10-19 殿村さん「別の言い方ないかな」→ 案2）。

    python -m _build.patch_sunao_honshitsu_msg <今の資料.pptx> [出力.pptx]

  「ためて追い続ける場所が無い。この場所を、LINE公式アカウントでつくる」
  →「イベントが終わると、集めた方との接点が途切れる。／途切れない接点を、LINE公式アカウントでつくる」（2行）。
  2行なので、メッセージの枠は new_page の2行のときと同じ位置・高さにする。ほかのページは触らない（出力を省くと上書き）。
"""
import sys

from pptx import Presentation

from _build.arauchi_v4.lib import move, set_paras
from _build.build_sunao_v2 import HONSHITSU_MSG

OLD = "ためて追い続ける場所が無い。この場所を、LINE公式アカウントでつくる"


def main(src, out):
    prs = Presentation(src)
    hits = [sh for s in prs.slides for sh in s.shapes if sh.has_text_frame and sh.text_frame.text == OLD]
    assert len(hits) == 1, f"差し替え元が {len(hits)} か所"
    ms = hits[0]
    set_paras(ms, HONSHITSU_MSG.split("\n"))
    move(ms, y=0.81, h=0.64)
    prs.save(out)
    print("saved:", out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else sys.argv[1])
