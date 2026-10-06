# -*- coding: utf-8 -*-
"""サンクスLINE誘導・業界別版に「全体座組」の1枚を先頭へ追加（2026-10-06 殿村さん指示）。
既存4枚は変更しない。AD提案資料からの橋渡しページ：
  広告で反響獲得（オレンジ）→ CVを獲得 →｜→ サンクスLINE誘導（緑）→ LINEでナーチャリング（緑）→ 売上・LTV最大化（濃紺）
上に「AD領域／LINE領域」の帯、下に結論1行だけ。想定数値・業界例・事例は入れない（後続ページの役割）。
  python3 _build/patch_thanks_line_gyokai_v12_zakumi.py <業界別ver1.1.pptx> [出力]
"""
import copy
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

sys.path.insert(0, str(Path(__file__).resolve().parent))
from pptx_parts import *  # noqa

ROOT = Path(__file__).resolve().parent.parent
OUT = Path(sys.argv[2]) if len(sys.argv) > 2 else ROOT / "20261006_サンクスLINE誘導のご提案_業界別ver1.2.pptx"
ORANGE, ORANGE_BG = "F59B21", "FEF0DE"  # listing.pdf のオレンジ（塗り・線）
ORANGE_TX = "C46A00"  # 白・薄いオレンジの上の文字は濃いオレンジ（殿村さん指示：読みやすく）
LINE_BG = "C9D8EE"

prs = Presentation(sys.argv[1])
ref = prs.slides[0]
assert ref.shapes[0].text_frame.text.startswith("CV後の接点をLINEに残す"), "1枚目が「サンクスLINE誘導とは」ではない"
s = prs.slides.add_slide(ref.slide_layout)
for sh in list(s.shapes):
    sh._element.getparent().remove(sh._element)
for nm in ("TextBox 1", "TextBox 2", "Connector 3"):
    s.shapes._spTree.append(copy.deepcopy(next(x for x in ref.shapes if x.name == nm)._element))
by = {x.name: x for x in s.shapes}
set_paras(by["TextBox 1"], [("広告で反響を獲得し、LINEで売上・LTVまで最大化", 0)])
set_paras(by["TextBox 2"], [("集客からCV後のフォローまで、一気通貫で顧客接点を設計", 0)])
# 先頭へ移動
lst = prs.slides._sldIdLst
el = lst[-1]
lst.remove(el)
lst.insert(0, el)

Y0, Y1 = 6.8, 14.5          # 図の上端・下端
MY = (Y0 + Y1) / 2          # 矢印の高さ
HH = 1.7                    # 見出しの高さ
X1, W1 = 2.2, 5.8           # 広告で反響獲得
X2, W2 = 8.8, 4.4           # CVを獲得
X3, W3 = 14.0, 3.6          # サンクスLINE誘導
X4, W4 = 18.4, 6.6          # LINEでナーチャリング
X5, W5 = 25.8, 5.9          # 売上・LTV最大化
BX = (X2 + W2 + X3) / 2     # AD／LINEの境目

# 領域の帯
rect(s, X1, 5.15, X2 + W2 - X1, 0.14, ORANGE)
label(s, X1, 4.45, 6.0, 0.7, [("AD領域", 13, True, ORANGE_TX, 0)], anchor=MSO_ANCHOR.MIDDLE)
rect(s, X3, 5.15, X5 + W5 - X3, 0.14, GREEN)
label(s, X3, 4.45, 6.0, 0.7, [("LINE領域", 13, True, GREEN_TX, 0)], anchor=MSO_ANCHOR.MIDDLE)
vline(s, BX, 4.5, Y1 + 0.2, LGRAY, 1.25, dash=True)


def column(x, w, head, head_fill, head_color, body_fill, line):
    rect(s, x, Y0, w, HH, head_fill, [(head, 15, True, head_color, 0)])
    shape(s, MSO_SHAPE.RECTANGLE, x, Y0 + HH, w, Y1 - Y0 - HH, None, fill=body_fill, line=line, lw=1.0)


# 1 広告で反響獲得（オレンジ）
column(X1, W1, "広告で反響獲得", ORANGE, WHITE, WHITE, ORANGE)
label(s, X1, Y0 + HH + 0.5, W1, 3.4, [(t, 13, False, INK, 6) for t in ("検索広告", "SNS広告", "その他WEB広告")], align=PP_ALIGN.CENTER)
hline(s, X1 + 0.6, X1 + W1 - 0.6, Y1 - 1.6, LGRAY, 0.75)
label(s, X1, Y1 - 1.5, W1, 1.3, [("サイト・LPへ集客", 13, True, ORANGE_TX, 0)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

# 2 CVを獲得
column(X2, W2, "CVを獲得", ORANGE_BG, ORANGE_TX, WHITE, ORANGE)
label(s, X2, Y0 + HH + 0.5, W2, 4.6, [(t, 12, False, INK, 5) for t in ("申込", "問い合わせ", "購入", "資料請求 など")], align=PP_ALIGN.CENTER)

# 3 サンクスLINE誘導（LINE領域の入口）
shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, X3, MY - 1.9, W3, 3.8, [("サンクス", 15, True, WHITE, 0), ("LINE誘導", 15, True, WHITE, 0)],
      fill=GREEN, adj=0.2, margins=(0.1, 0, 0.1, 0))

# 4 LINEでナーチャリング（トーク画面風）
column(X4, W4, "LINEでナーチャリング", GREEN_BG, GREEN_TX, LINE_BG, None)
by_ = Y0 + HH + 0.45
for i, t in enumerate(("リマインド", "情報提供", "キャンペーン案内", "個別案内")):
    y = by_ + i * 1.32
    shape(s, MSO_SHAPE.OVAL, X4 + 0.35, y + 0.12, 0.7, 0.7, None, fill=GREEN)
    b = shape(s, MSO_SHAPE.ROUNDED_RECTANGULAR_CALLOUT, X4 + 1.35, y, W4 - 1.75, 0.95, [(t, 12, True, INK, 0)], fill=WHITE,
              margins=(0.1, 0, 0.1, 0))
    b.adjustments[0], b.adjustments[1] = -0.56, -0.15

# 5 売上・LTV最大化（濃紺）
column(X5, W5, "売上・LTV最大化", NAVY, WHITE, NAVY, None)
label(s, X5, Y0 + HH + 0.45, W5, 5.0, [(t, 14, True, WHITE, 7) for t in ("面談・商談", "購入", "再購入", "継続利用")],
      align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
hline(s, X5 + 0.6, X5 + W5 - 0.6, Y0 + HH + 0.05, "5A6390", 0.75)

# 矢印
for xa, xb in ((X1 + W1, X2), (X2 + W2, X3), (X3 + W3, X4), (X4 + W4, X5)):
    arrow_line(s, xa + 0.12, MY, xb - 0.12, MY, NAVY, 2.0)

# 結論（1行だけ）
label(s, 2.2, 15.5, 29.5, 1.4, [("CVをゴールにせず、広告で獲得したユーザーをLINEで育成し、その後の売上までつなげる。", 17, True, NAVY, 0)],
      align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

prs.save(str(OUT))
print("saved:", OUT.name, len(prs.slides), "枚")
