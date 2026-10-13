# -*- coding: utf-8 -*-
"""v4 の検品：v3 との差分ページ、禁止語、数字（SIM）、S24の表、出典のリンク。

    python -m _build.arauchi_v4.checks <v3.pptx> <v4.pptx> <SIM.xlsx>
"""
import re
import sys
import zipfile

import openpyxl
from lxml import etree
from pptx import Presentation

from _build.arauchi_v4.build import c14n, slide_order
from _build.arauchi_v4.pages import HEAD, ROWS

CHANGED = {3, 5, 6, 7, 9, 11, 12, 13, 24}
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def page_texts(slide):
    out = []
    for el in slide._element.iter(A + "t"):
        out.append(el.text or "")
    return "".join(out)


def all_text(slide):
    paras = []
    for p in slide._element.iter(A + "p"):
        paras.append("".join((t.text or "") for t in p.iter(A + "t")))
    return "\n".join(paras)


def main(v3, v4, sim):
    ok = True
    def check(cond, msg):
        nonlocal ok
        print(("OK  " if cond else "NG  ") + msg)
        ok = ok and cond

    z3, z4 = zipfile.ZipFile(v3), zipfile.ZipFile(v4)
    o3, o4 = slide_order(z3), slide_order(z4)
    check(len(o3) == len(o4) == 46, f"枚数 v3={len(o3)} v4={len(o4)}（46枚）")
    # v3 と v4 でスライドXMLが違うページ
    diff = {i + 1 for i, (a, b) in enumerate(zip(o3, o4)) if c14n(z3.read(a)) != c14n(z4.read(b))}
    check(diff == CHANGED, f"v3と中身が違うページ＝{sorted(diff)}（期待 {sorted(CHANGED)}）")
    # 画像・ノート・レイアウトの参照先が同じ
    same_rel = True
    for a, b in zip(o3, o4):
        ra = a.replace("slides/", "slides/_rels/") + ".rels"
        rb = b.replace("slides/", "slides/_rels/") + ".rels"
        sa = {(r.get("Type").split("/")[-1], r.get("Target")) for r in etree.fromstring(z3.read(ra)) if not r.get("Type").endswith("/hyperlink")}
        sb = {(r.get("Type").split("/")[-1], r.get("Target")) for r in etree.fromstring(z4.read(rb)) if not r.get("Type").endswith("/hyperlink")}
        same_rel &= sa == sb
    check(same_rel, "画像・ノート・レイアウトの参照先は v3 と同じ")

    prs = Presentation(v4)
    texts = {i + 1: all_text(s) for i, s in enumerate(prs.slides)}
    alltxt = "\n".join(texts[i] for i in range(1, 47))
    for bad in ("松竹梅", "スコープ", "許容CPA", "♾", "∞"):
        check(bad not in alltxt, f"『{bad}』が資料に出てこない")
    # 『ヒアリング』『CPA』は汎用ページ（S32・S34）に元からある。直したページ（1〜17・24）に無いことを見る
    mine = "\n".join(texts[i] for i in list(range(1, 18)) + [24])
    for bad in ("ヒアリング", "CPA"):
        check(bad not in mine, f"『{bad}』が直したページ（1〜17・24）に無い（汎用ページS32・S34の元からの記載は変更なし）")
    check("予算" not in alltxt, "『予算』が資料に出てこない（予算のヒアリング項目を作っていない）")

    # --- SIM の値
    wb = openpyxl.load_workbook(sim, data_only=True)
    ws = wb[wb.sheetnames[0]]
    cv1_m1, cv2_m1, tot_m1 = ws["E36"].value, ws["E37"].value, ws["E40"].value
    cv1_m6, cv2_m6, tot_m6, sum_all = ws["J36"].value, ws["J37"].value, ws["J40"].value, ws["K40"].value
    friends6 = ws["J28"].value
    new1, new6 = ws["E26"].value, ws["J26"].value
    cvr1, cvr6 = ws["E41"].value, ws["J41"].value
    init, month = ws["D73"].value, ws["E73"].value
    print(f"SIM: ①{cv1_m1}→{cv1_m6} ②{cv2_m1}→{cv2_m6} 合計{tot_m1}→{tot_m6} 累計{sum_all} 友だち{friends6:.0f} 初期{init} 月額{month}")
    print(f"     物件を送る方の割合 {cv1_m1/new1:.1%}→{cv1_m6/new6:.1%} ／ CVR {cvr1:.2%}→{cvr6:.2%}")
    check((cv1_m6, cv2_m6, tot_m6, sum_all) == (13, 10, 23, 90), "SIM：6か月目 ①13・②10・合計23／6か月累計90")
    check(f"月{tot_m1}件→6か月目 月{tot_m6}件" in texts[12], f"S12：1か月目 月{tot_m1}件→6か月目 月{tot_m6}件")
    t13 = texts[13]
    for s in (f"約{friends6:,.0f}人", f"{cv1_m6}件/月", f"{cv2_m6}件/月", f"{tot_m6}件/月", f"6か月累計{sum_all}件",
              "¥215,000", "¥135,000", "初期21.5万円・月額13.5万円", "月次運用範囲内で適宜調整"):
        check(s in t13, f"S13：『{s}』")
    check("18件" not in t13 and "77件" not in t13 and "12件/月" not in t13 and "6件/月" not in t13, "S13：旧数字（18件・77件・12件/月・6件/月）が残っていない")
    check("18件" not in texts[12] and "月8件" not in texts[12], "S12：旧数字が残っていない")
    check((init, month) == (215000, 135000), "SIM：初期215,000・月額135,000")
    t9 = texts[9]
    r1, r2 = f"{round(cv1_m1/new1*100)}→{round(cv1_m6/new6*100)}%", f"{cvr1*100:.1f}→{cvr6*100:.1f}%"
    check(r1 in t9 and r2 in t9, f"S9：{r1}（①÷新しい友だち）／{r2}（②÷クリック）")
    check("25→52" not in t9 and "1.2%" not in t9, "S9：旧数字が残っていない")

    # --- S24
    shp = [s for s in prs.slides[23].shapes if s.has_table][0]
    tbl = shp.table
    got = [[c.text for c in r.cells] for r in tbl.rows]
    check(got[0] == HEAD and [tuple(r) for r in got[1:]] == ROWS, "S24：表の中身が指定どおり")
    spec = [("初期費用", "21.5万円", "21.5万円", "21.5万円"), ("月額コンサル費", "10万円", "5万円", "なし"),
            ("月額ツール・LINE費", "3.5万円", "3.5万円", "別途実費"), ("月額合計", "13.5万円", "8.5万円", "運用契約なし"),
            ("企画配信", "月4〜5本", "月1〜2本", "なし"), ("ステップ配信", "10通", "なし", "なし"),
            ("離脱防止", "導入・運用", "導入・運用", "初期導入支援"), ("継続改善", "対応", "限定的", "なし"),
            ("API連携・開発", "必要に応じ別途見積", "なし", "なし")]
    check([tuple(r) for r in got[1:]] == spec, "S24：ユーザー指定の表と一字一句一致")
    for kw in ("プラン①", "プラン②", "プラン③"):
        check(kw in texts[24], f"S24：『{kw}』表記")

    # --- 出典（S5・S6）：URLを文字で出さない・リンクは付ける
    for n in (5, 6):
        t = texts[n]
        check("http" not in t and "確認）" not in t and "確認日" not in t, f"S{n}：出典にURL・確認日の文字が無い")
        z = prs.slides[n - 1]
        links = [r.target_ref for r in z.part.rels.values() if r.reltype.endswith("/hyperlink")]
        print(f"     S{n} のリンク: {links}")
        check(len(links) >= 3, f"S{n}：出典名にハイパーリンクが付いている（{len(links)}件）")
    # --- 100円の注意書き
    check("100円賃貸が対応できる物件に限ります" in texts[5], "S5：100円賃貸の注意書きが残っている")
    # --- S7
    t7 = texts[7]
    check("導入しません" not in t7 and "今回は導入しない" not in t7, "S7：『外部ツールは導入しない』の断定が無い")
    check("別途お見積り" in t7 and "API連携・開発" in t7, "S7：API連携・開発は別途お見積り")
    # --- S3/S11
    check("最大の改善ポイントは、LINE友だち追加後に物件の問い合わせへつなげる導線の整備" in texts[3] and "⑤" in texts[3], "S3：⑤を最大の改善ポイントとして明示")
    check("最重要課題" in texts[11] and "⑤リード有効化" in texts[11], "S11：最重要課題＝⑤リード有効化")
    check("家賃負担" not in alltxt, "『家賃負担の軽減』の訴求が残っていない")
    print("\n総合:", "全部OK" if ok else "NGあり")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main(*sys.argv[1:4]))
