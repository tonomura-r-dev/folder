# -*- coding: utf-8 -*-
"""あらうち提案資料 v4（2026-10-13 殿村さん指示）の部品。v3 の common.py に、図形・表・リンクの追加を足したもの。"""
import copy

from lxml import etree
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Inches

from _build.arauchi_v3.common import *  # noqa: F401,F403
from _build.arauchi_v3.common import ids, txbody, set_paras, set_size_all, delete, move  # noqa: F401

FONT = "メイリオ"
PANOSE = dict(panose="020B0604030504040204", pitchFamily="50", charset="-128")
NAVY = "1F285A"
RED = "C00000"
TEXT = "333333"


def emu(v):
    return int(Inches(v))


def make_rpr(size, bold=False, color=TEXT, link_rid=None):
    rpr = etree.Element(qn("a:rPr"), lang="ja-JP", altLang="en-US", sz=str(int(round(size * 100))),
                        b="1" if bold else "0", dirty="0")
    fill = etree.SubElement(rpr, qn("a:solidFill"))
    etree.SubElement(fill, qn("a:srgbClr"), val=color)
    etree.SubElement(rpr, qn("a:latin"), typeface=FONT, **PANOSE)
    etree.SubElement(rpr, qn("a:ea"), typeface=FONT, **PANOSE)
    if link_rid:
        etree.SubElement(rpr, qn("a:hlinkClick"), {qn("r:id"): link_rid})
    return rpr


def make_para(runs, size, bold=False, color=TEXT, align="l", before=None, after=None, line=None, slide=None):
    """runs: 文字列、または [(文字, {bold, color, size, link}), ...]"""
    p = etree.Element(qn("a:p"))
    ppr = etree.SubElement(p, qn("a:pPr"), algn=align)
    if line:
        ln = etree.SubElement(ppr, qn("a:lnSpc"))
        etree.SubElement(ln, qn("a:spcPct"), val=str(int(line * 1000)))
    if before is not None:
        sb = etree.SubElement(ppr, qn("a:spcBef"))
        etree.SubElement(sb, qn("a:spcPts"), val=str(int(before * 100)))
    if after is not None:
        sa = etree.SubElement(ppr, qn("a:spcAft"))
        etree.SubElement(sa, qn("a:spcPts"), val=str(int(after * 100)))
    etree.SubElement(ppr, qn("a:buNone"))
    if isinstance(runs, str):
        runs = [(runs, {})]
    for text, opt in runs:
        rid = None
        if opt.get("link"):
            rid = slide.part.relate_to(opt["link"], RT.HYPERLINK, is_external=True)
        if text == "" and len(runs) == 1:
            continue
        for k, part in enumerate(text.split("\n")):
            if k:
                br = etree.SubElement(p, qn("a:br"))
                br.append(make_rpr(opt.get("size", size), opt.get("bold", bold), opt.get("color", color)))
            r = etree.SubElement(p, qn("a:r"))
            r.append(make_rpr(opt.get("size", size), opt.get("bold", bold), opt.get("color", color), rid))
            t = etree.SubElement(r, qn("a:t"))
            t.text = part
    end = etree.SubElement(p, qn("a:endParaRPr"), lang="ja-JP", altLang="en-US", sz=str(int(round(size * 100))), dirty="0")
    return p


def add_box(slide, x, y, w, h, paras, size=12, bold=False, color=TEXT, fill=None, line=None, line_w=0.75,
            dash=None, align="l", anchor="ctr", ins=(0.08, 0.04, 0.08, 0.04), shape="rect", before=None, after=None,
            line_spacing=None, name=None):
    """図形（四角）を追加して返す。paras は make_para の runs の並び（段落ごと）"""
    from pptx.enum.shapes import MSO_SHAPE
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if shape == "round" else MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    el = shp._element
    st = el.find(qn("p:style"))
    if st is not None:
        el.remove(st)
    if name:
        shp.name = name
    sp_pr = el.spPr
    for tag in ("a:solidFill", "a:noFill", "a:ln"):
        for e in sp_pr.findall(qn(tag)):
            sp_pr.remove(e)
    if fill:
        f = etree.SubElement(sp_pr, qn("a:solidFill"))
        etree.SubElement(f, qn("a:srgbClr"), val=fill)
    else:
        etree.SubElement(sp_pr, qn("a:noFill"))
    ln = etree.SubElement(sp_pr, qn("a:ln"))
    if line:
        ln.set("w", str(int(line_w * 12700)))
        f = etree.SubElement(ln, qn("a:solidFill"))
        etree.SubElement(f, qn("a:srgbClr"), val=line)
        if dash:
            etree.SubElement(ln, qn("a:prstDash"), val=dash)
    else:
        etree.SubElement(ln, qn("a:noFill"))
    tb = el.find(qn("p:txBody"))
    bp = tb.find(qn("a:bodyPr"))
    for k in list(bp.attrib):
        del bp.attrib[k]
    for c in list(bp):
        bp.remove(c)
    bp.set("wrap", "square")
    bp.set("lIns", str(emu(ins[0])))
    bp.set("tIns", str(emu(ins[1])))
    bp.set("rIns", str(emu(ins[2])))
    bp.set("bIns", str(emu(ins[3])))
    bp.set("rtlCol", "0")
    bp.set("anchor", anchor)
    etree.SubElement(bp, qn("a:noAutofit"))
    for p in tb.findall(qn("a:p")):
        tb.remove(p)
    for runs in paras:
        tb.append(make_para(runs, size, bold, color, align, before, after, line_spacing, slide))
    return shp


# ---------------------------------------------------------------- 表
def _ln(tag, w_pt, color):
    ln = etree.Element(qn(tag), w=str(int(w_pt * 12700)), cap="flat", cmpd="sng", algn="ctr")
    f = etree.SubElement(ln, qn("a:solidFill"))
    etree.SubElement(f, qn("a:srgbClr"), val=color)
    etree.SubElement(ln, qn("a:prstDash"), val="solid")
    etree.SubElement(ln, qn("a:round"))
    etree.SubElement(ln, qn("a:headEnd"), type="none", w="med", len="med")
    etree.SubElement(ln, qn("a:tailEnd"), type="none", w="med", len="med")
    return ln


def add_table(slide, x, y, col_w, row_h, rows, name="表"):
    """rows[r][c] = dict(text, size, bold, color, fill, align, borders={l,r,t,b: (pt, color)}, ins=(l,t,r,b))
    列幅 col_w・行の高さ row_h はインチ。PowerPointの表と同じ作り（a:tbl／セルごとに書式を明示）。"""
    nid = max(sh.shape_id for sh in slide.shapes) + 1
    gf = etree.SubElement(slide.shapes._spTree, qn("p:graphicFrame"))
    nv = etree.SubElement(gf, qn("p:nvGraphicFramePr"))
    etree.SubElement(nv, qn("p:cNvPr"), id=str(nid), name=f"{name} {nid - 1}")
    cnv = etree.SubElement(nv, qn("p:cNvGraphicFramePr"))
    etree.SubElement(cnv, qn("a:graphicFrameLocks"), noGrp="1")
    etree.SubElement(nv, qn("p:nvPr"))
    xf = etree.SubElement(gf, qn("p:xfrm"))
    etree.SubElement(xf, qn("a:off"), x=str(emu(x)), y=str(emu(y)))
    etree.SubElement(xf, qn("a:ext"), cx=str(sum(emu(w) for w in col_w)), cy=str(sum(emu(h) for h in row_h)))
    g = etree.SubElement(gf, qn("a:graphic"))
    gd = etree.SubElement(g, qn("a:graphicData"), uri="http://schemas.openxmlformats.org/drawingml/2006/table")
    tbl = etree.SubElement(gd, qn("a:tbl"))
    etree.SubElement(tbl, qn("a:tblPr"))
    grid = etree.SubElement(tbl, qn("a:tblGrid"))
    for w in col_w:
        etree.SubElement(grid, qn("a:gridCol"), w=str(emu(w)))
    for r, cells in enumerate(rows):
        tr = etree.SubElement(tbl, qn("a:tr"), h=str(emu(row_h[r])))
        for c in cells:
            tc = etree.SubElement(tr, qn("a:tc"))
            for attr in ("gridSpan", "rowSpan"):            # 結合セル（先頭のセルに span、結合された側に hMerge / vMerge）
                if c.get(attr):
                    tc.set(attr, str(c[attr]))
            for attr in ("hMerge", "vMerge"):
                if c.get(attr):
                    tc.set(attr, "1")
            tx = etree.SubElement(tc, qn("a:txBody"))
            etree.SubElement(tx, qn("a:bodyPr"))
            etree.SubElement(tx, qn("a:lstStyle"))
            lines = c["text"] if isinstance(c["text"], list) else [c["text"]]
            for ln_text in lines:
                tx.append(make_para(ln_text, c.get("size", 13), c.get("bold", False), c.get("color", TEXT),
                                    c.get("align", "ctr"), slide=slide))
            ins = c.get("ins", (0.08, 0.03, 0.08, 0.03))
            tcpr = etree.SubElement(tc, qn("a:tcPr"), marL=str(emu(ins[0])), marR=str(emu(ins[2])),
                                    marT=str(emu(ins[1])), marB=str(emu(ins[3])), anchor="ctr")
            b = c.get("borders", {})
            for key, tag in (("l", "a:lnL"), ("r", "a:lnR"), ("t", "a:lnT"), ("b", "a:lnB")):
                w_pt, col = b.get(key, (0.75, "BFC8DC"))
                tcpr.append(_ln(tag, w_pt, col))
            f = etree.SubElement(tcpr, qn("a:solidFill"))
            etree.SubElement(f, qn("a:srgbClr"), val=c.get("fill", "FFFFFF"))
    return gf
