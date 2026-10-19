# -*- coding: utf-8 -*-
"""すなおコンサルティング：本編の左上の紺のラベルを、全ページ同じ大きさにそろえる（2026-10-19 殿村さん「大きさも統一してほしい」）。

    python -m _build.patch_sunao_label_size <今の資料.pptx> [出力.pptx]

  ラベル＝幅1.97・高さ0.55インチ（大半のページと同じ）。右のメッセージは、ラベルの右0.2インチから始める。
  広かったのは「LINE公式アカウント」「気をつける点と対策」「成果報酬型プラン」の3つ。Appendixは触らない（出力を省くと上書き）。
"""
import sys

from pptx import Presentation
from pptx.util import Emu

from _build.build_sunao_v2 import fix_label


def label_and_message(s):
    """左上のラベル（左端・上から0.7〜1.1インチ）と、その右のメッセージ"""
    lb = [sh for sh in s.shapes if sh.has_text_frame and Emu(sh.left).inches < 0.05 and 0.7 < Emu(sh.top).inches < 1.1]
    if len(lb) != 1:
        return None, None
    ms = [sh for sh in s.shapes if sh.has_text_frame and sh is not lb[0]
          and 0.75 < Emu(sh.top).inches < 1.0 and 1.5 < Emu(sh.left).inches < 3.5]
    return lb[0], (ms[0] if len(ms) == 1 else None)


def main(src, out):
    prs = Presentation(src)
    slides = list(prs.slides)
    back = next(i for i, s in enumerate(slides) if any(sh.has_text_frame and "Appendix" == sh.text_frame.text.strip()
                                                       for sh in s.shapes))
    n = 0
    for s in slides[:back]:                                 # Appendixの扉より前（本編）だけ
        lb, ms = label_and_message(s)
        if lb is None:
            continue
        assert ms is not None, lb.text_frame.text
        fix_label(lb, ms)
        n += 1
    prs.save(out)
    print("saved:", out, f"（ラベル {n} 枚）")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else sys.argv[1])
