# -*- coding: utf-8 -*-
"""サンクスLINE：殿村さんがChatGPTで生成した画像を差し込む（2026-10-01）。上書き保存。
- 「LINEは、申込みの「後」につながる」のページ：四角と矢印の図を _images/thanks_line_journey.png に差し替え（下の一文は残す）
- 「実績と費用」のページ：実績カード・費用を画像 _images/thanks_line_jisseki.png 1枚に置き換える（締めの一文と出典はパワポの文字で残す）
  python3 _build/patch_thanks_line_images.py <pptx> [<pptx> ...]
"""
import sys
from pathlib import Path
from pptx import Presentation
from pptx.util import Cm

ROOT = Path(__file__).resolve().parent.parent
IMG = ROOT / "_images"
for path in sys.argv[1:]:
    prs = Presentation(path)
    SW = prs.slide_width / 360000
    for s in prs.slides:
        title = s.shapes[0].text_frame.text if s.shapes[0].has_text_frame else ""
        if title.startswith("LINEは、申込みの「後」"):
            for sh in list(s.shapes):
                if sh.name in ("Rounded Rectangle 4", "Right Arrow 5", "Rounded Rectangle 6", "Right Arrow 7",
                               "Right Arrow 9", "Rounded Rectangle 14", "Rounded Rectangle 15",
                               "Rounded Rectangle 16", "Rounded Rectangle 17"):
                    sh._element.getparent().remove(sh._element)
            W = 21.0
            H = W * 941 / 1672
            pic = s.shapes.add_picture(str(IMG / "thanks_line_journey.png"), Cm((SW - W) / 2), Cm(4.15), Cm(W), Cm(H))
            pic.name = "広告とLINEの役割"
            band = next(x for x in s.shapes if x.name == "Rounded Rectangle 13")
            band.left, band.top, band.width = Cm(1.46), Cm(4.15 + H + 0.3), Cm(SW - 2.92)
        if title.startswith("実績と費用"):
            # 実績カード・費用の図形を、画像（_images/thanks_line_jisseki.png）1枚に置き換える
            keep_text = ("実績と費用", "弊社の運用実績", "広告のご提案とあわせて", "出典")
            for sh in list(s.shapes):
                txt = sh.text_frame.text if sh.has_text_frame else ""
                if sh.name == "Connector 3" or any(txt.startswith(k) for k in keep_text):
                    continue
                sh._element.getparent().remove(sh._element)
            for r in s.shapes[0].text_frame.paragraphs[0].runs[1:]:
                r.text = ""
            s.shapes[0].text_frame.paragraphs[0].runs[0].text = "実績と費用"
            W = 20.0
            H = W * 941 / 1672
            pic = s.shapes.add_picture(str(IMG / "thanks_line_jisseki.png"), Cm((SW - W) / 2), Cm(4.15), Cm(W), Cm(H))
            pic.name = "実績と費用"
            for sh in s.shapes:
                txt = sh.text_frame.text if sh.has_text_frame else ""
                if txt.startswith("広告のご提案とあわせて"):
                    sh.top, sh.height = Cm(4.15 + H + 0.25), Cm(1.2)
                if txt.startswith("出典"):
                    sh.top = Cm(4.15 + H + 1.5)
    prs.save(path)
    print("saved:", path)
