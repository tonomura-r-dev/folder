"""
LINE広告 トレンドレポート 2026年10月 ビルドスクリプト。

LINEOA版トレンドレポートと同じトンマナ（trend_report_engine.py）で作り、
分類タグだけ広告向けに差し替える。
内容は2026-10-26に殿村さんと決めた3項目。LINEヤフー公式の発表で裏付けが取れた事実だけを載せる
（配信面の振り分け先・縦型動画の比重・データ統合ターゲティング・自動入札の一本化・
オーディエンスの有効期限は裏付けが取れなかったため省いた）。

使い方: python3 _build/build_line_ads_trend_202610.py
"""

from pathlib import Path
from pptx import Presentation
from pptx.util import Pt, Inches
import trend_report_engine as engine

ROOT = Path(__file__).resolve().parent.parent
OUT = str(ROOT / "LINE広告トレンドレポート_2026年10月.pptx")
COVER = "2026年10月　LINE広告トレンドレポート"

# FMTのタグ名 → 広告版のタグ名（ハイライトはFMTの名前で指定し、最後に表示名を差し替える）
TAG_RENAME = {
    "LINE公式アカウント": ["LINE広告"],
    "開発\nLINE API": ["LINEヤフー", "広告"],
    "オプション\n商材": ["配信面"],
}

TOPICS = [
    dict(
        bracket="LINE広告",
        headline="旧「LINE広告」の配信停止（10/26）",
        industry="LINE広告を運用中の広告主",
        category="LINE公式アカウント",
        description_lines=[
            "2026年10月26日（月）に、旧「LINE広告」の新規作成・編集・広告配信が停止された。",
            "日本向けの広告はすべて「未配信」になり、停止後の管理画面は閲覧のみ。",
            "2027年3月末頃に提供終了となり、それ以降はレポートも確認できなくなる。",
        ],
        render="flow",
        steps=[
            ("2026年6月30日", "新規アカウントの作成が停止"),
            ("2026年10月26日", "新規作成・編集・広告配信が停止"),
            ("停止後〜", "管理画面は閲覧のみ（レポート・明細は確認可）"),
            ("2027年3月末頃", "提供終了（レポートも確認不可）"),
        ],
        note_headline="レポートは提供終了までに保存",
        note_desc="移行先のLINEヤフー広告には過去の配信実績が引き継がれないため、比較に使うレポートは2027年3月末頃の提供終了までに手元に残しておく。",
        source="https://www.lycbiz.com/jp/news/line-ads/20260925/",
    ),
    dict(
        bracket="LINEヤフー広告",
        headline="ディスプレイ広告への移行で必要な対応",
        industry="LINE広告を運用中の広告主",
        category="開発\nLINE API",
        description_lines=[
            "配信を続けるには「LINEヤフー広告 ディスプレイ広告」への切り替えが必要。",
            "アカウントや配信設定は自動では移らず、移行ツールか手動で移す。",
            "タグ・コンバージョン設定と自動入札の学習データは引き継がれない。",
        ],
        render="grid_notes",
        items=[
            ("移行ツールで移せるもの", "キャンペーン・広告グループ・広告の設定、画像・動画、オーディエンス。"),
            ("移せないもの", "LINE Tagとコンバージョン設定、過去の配信実績・レポート、自動入札の学習データ。"),
            ("移行後にやること", "タグとコンバージョンを設定し直す。学習データがない状態から始まるため、移行直後は学習期間を見込む。"),
            ("移行ツールの対象", "日本向け配信・運用型・審査が承認済みのアカウント。予約型や審査待ち・否認のものは対象外。"),
        ],
        source="https://www.lycbiz.com/jp/column/line-ads/service-information/migration-guide/",
    ),
    dict(
        bracket="LINE VOOM",
        headline="サービス終了に伴う広告枠の終了",
        industry="LINE広告を運用中の広告主",
        category="オプション\n商材",
        description_lines=[
            "2026年9月30日（水）11時に、LINE VOOMのサービスが終了した。",
            "あわせて、LINE VOOM面への広告配信も終了している。",
            "LINEアプリのVOOMタブは、ショッピングタブに切り替わる。",
        ],
        render="before_after",
        as_is="LINE VOOM面に広告を配信できる",
        to_be_lines=["VOOM面への広告配信は終了", "9/30 11時にサービス終了。投稿・閲覧もできない"],
        merits=[
            ("VOOMタブはショッピングタブへ",
             "LINEアプリ下部のVOOMタブは、ショッピングタブに切り替わる。"),
            ("VOOM面向けの設定は見直し",
             "VOOM面を前提にした配信設定・クリエイティブは、残りの配信面に合わせて見直す。"),
        ],
        source="https://www.lycbiz.com/jp/news/ly-ads/20260730/",
    ),
]

for t in TOPICS:
    t["title"] = f"【{t['bracket']}】{t['headline']}"


def rename_tags(path):
    prs = Presentation(path)
    for slide in list(prs.slides)[2:]:
        for sh in slide.shapes:
            if not sh.has_text_frame:
                continue
            key = "\n".join(p.text for p in sh.text_frame.paragraphs).replace("\x0b", "\n")
            for old, new in TAG_RENAME.items():
                if key.replace("\n", "") == old.replace("\n", ""):
                    engine.set_multiline_text(sh, new)
                    if len("".join(new)) > 6:  # 長いタグ名は枠に収まるよう小さく
                        for para in sh.text_frame.paragraphs:
                            for r in para.runs:
                                r.font.size = Pt(9)
    # 表紙タイトルは「LINEOA」より長いので、中央揃えのまま枠を左右に広げて1行に収める
    title = next(sh for sh in prs.slides[0].shapes if sh.has_text_frame and "トレンドレポート" in sh.text_frame.text)
    title.left, title.width = Inches(0.45), Inches(9.92)
    engine.force_font(prs)
    prs.save(path)


if __name__ == "__main__":
    engine.build_standalone(OUT, COVER, TOPICS)
    rename_tags(OUT)
