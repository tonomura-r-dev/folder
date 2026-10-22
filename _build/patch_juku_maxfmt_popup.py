# -*- coding: utf-8 -*-
"""塾の新FMT版：「離脱防止」を「ポップアップバナー」に言い換える（2026-10-22 殿村さん）。

    python -m _build.patch_juku_maxfmt_popup <入力.pptx> <出力.pptx>

  「離脱防止ポップアップ」「離脱防止バナー」「ポップアップ / 離脱防止」などの組み合わせは、
  「ポップアップバナーポップアップ」のように重ならないよう、長い言い方から先に置き換える。
  本編・Appendixとも全ページ。文字だけ変える（図形・画像はそのまま）。
"""
import copy
import sys

from pptx import Presentation
from pptx.oxml.ns import qn

NEW = "ポップアップバナー"
PAIRS = [  # 長い言い方から順に
    ("ポップアップ / 離脱防止", NEW),
    ("ポップアップ(離脱防止 など)", NEW),
    ("LP離脱防止の設置", "LPの" + NEW + "の設置"),
    ("離脱防止ポップアップ", NEW),
    ("離脱防止バナー", NEW),
    ("離脱防止", NEW),
]


def runs_of(s, name, first):
    """図形 name の中で、最初の文字が first の段落の a:r の一覧"""
    sh = next(x for x in s.shapes if x.name == name and first in "".join(t.text or "" for t in x._element.iter(qn("a:t"))))
    for p in sh._element.iter(qn("a:p")):
        rs = p.findall(qn("a:r"))
        if rs and (rs[0].find(qn("a:t")).text or "").startswith(first):
            return p, rs
    raise KeyError((name, first))


def keep_first(p, rs, text):
    rs[0].find(qn("a:t")).text = text
    for r in rs[1:]:
        p.remove(r)


def two_lines(s, name, first, line1, line2):
    """「離脱防止」＋「ポップアップ」の2段の箱を「ポップアップ」＋「バナー」の2段に（同じ書式で）"""
    p1, rs1 = runs_of(s, name, first)
    keep_first(p1, rs1, line1)
    p2 = p1.getnext()
    p2.getparent().replace(p2, copy.deepcopy(p1))
    p2 = p1.getnext()
    p2.findall(qn("a:r"))[0].find(qn("a:t")).text = line2


def main(src, out):
    prs = Presentation(src)
    slides = list(prs.slides)
    # 組み合わせの言い方が段落の途中で a:r に分かれているところ（先にページを指定して直す）
    for no in (44, 45):
        keep_first(*runs_of(slides[no - 1], "Google Shape;752;p28", "新規友だちの追加"), "新規友だちの追加（" + NEW + "）")
        keep_first(*runs_of(slides[no - 1], "グループ化 2", "ポップアップ"), NEW)
    two_lines(slides[16], "Rounded Rectangle 40", "離脱防止", "ポップアップ", "バナー")      # P17 導線の箱
    two_lines(slides[62], "Rectangle 36", "00 離脱防止", "00 ポップアップ", "バナー")        # Appendix 設計図① の表
    p, rs = runs_of(slides[42], "Google Shape;178;p7", "離脱防止バナー")                     # P43 カードの見出し（20pt だと幅に収まらないので16pt）
    rs[0].find(qn("a:rPr")).set("sz", "1600")
    count = 0
    for no, s in enumerate(prs.slides, 1):
        for t in s._element.iter(qn("a:t")):
            if t.text and "離脱防止" in t.text:
                for old, new in PAIRS:
                    if old in t.text:
                        count += t.text.count(old)
                        t.text = t.text.replace(old, new)
        for p in s._element.iter(qn("a:p")):
            text = "".join(t.text or "" for t in p.iter(qn("a:t")))
            assert "離脱防止" not in text, (no, text)
            assert NEW + "ポップアップ" not in text and NEW + "バナー" not in text and "ポップアップ (" not in text, (no, text)
            assert "ポップアップ / " not in text and "ポップアップ(" not in text, (no, text)
    prs.save(out)
    print("saved:", out, count, "か所")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
