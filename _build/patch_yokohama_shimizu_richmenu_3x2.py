# -*- coding: utf-8 -*-
"""横浜シミズ：「配信イメージ」のリッチメニューを 縦3×横2 → 横3×縦2 に組み直す（2026-10-21 殿村さん「リッチメニューは横3×縦2やろ」）。

    python -m _build.patch_yokohama_shimizu_richmenu_3x2 <資料.pptx> [出力.pptx]

  6つの枠（仕事を探す／勤務の案内／希望条件を変える／よくある質問／お問い合わせ／横浜シミズとは）を、
  灰色の枠の中に横3×縦2で並べ直す。並びは左上から 仕事を探す→勤務の案内→希望条件を変える／よくある質問→お問い合わせ→横浜シミズとは。
  文字の大きさ・色・言葉は変えない。枠が細くなる分、長い言葉だけ改行を入れる。
  灰色の枠の空きは、上1/3・下2/3 に分ける（CLAUDE.md 2026-10-19）。ほかのページは触らない。出力を省くと上書き。
"""
import copy
import sys

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches

ORDER = [33, 34, 35, 36, 37, 38]          # 左上から右へ、1段目→2段目
BREAKS = {                                # 枠が細くなるので、ここだけ改行を入れる
    33: ["仕事を探す", "募集中の\nイベント"],
    34: ["勤務の案内", "集合・持ち物"],
    35: ["希望条件を\n変える", "エリア・曜日"],
    36: ["よくある\n質問", "給与・服装など"],
    37: ["お問い\n合わせ", "チャットで相談"],
    38: ["横浜シミズ\nとは", "会社・現場\nの紹介"],
}
MARGIN, GAP, TILE_H = 0.06, 0.05, 1.0
INS = 18288                               # 枠の内側の左右の余白 0.02インチ


def by_id(s, sid):
    return next(sh for sh in s.shapes if sh.shape_id == sid)


def set_para(p, text):
    """段落の最初のrunの書式で text を入れる。改行は a:br で入れる"""
    runs = p.findall(qn("a:r"))
    for r in runs[1:]:
        p.remove(r)
    for br in p.findall(qn("a:br")):
        p.remove(br)
    r0 = runs[0]
    lines = text.split("\n")
    r0.find(qn("a:t")).text = lines[0]
    prev = r0
    for line in lines[1:]:
        br = copy.deepcopy(r0)
        br.tag = qn("a:br")
        br.remove(br.find(qn("a:t")))
        prev.addnext(br)
        r = copy.deepcopy(r0)
        r.find(qn("a:t")).text = line
        br.addnext(r)
        prev = r


def main(src, out):
    prs = Presentation(src)
    s = prs.slides[11]
    assert by_id(s, 31).text_frame.text == "リッチメニュー（構成案）"
    panel = by_id(s, 32)
    assert [by_id(s, sid).text_frame.text.split("\n")[0] for sid in ORDER] == \
        ["仕事を探す", "勤務の案内", "希望条件を変える", "よくある質問", "お問い合わせ", "横浜シミズとは"]
    w = (panel.width - Inches(MARGIN * 2) - Inches(GAP * 2)) / 3
    grid_h = Inches(TILE_H * 2 + GAP)
    top = panel.top + (panel.height - grid_h) / 3
    for i, sid in enumerate(ORDER):
        sh = by_id(s, sid)
        col, row = i % 3, i // 3
        sh.left = int(panel.left + Inches(MARGIN) + col * (w + Inches(GAP)))
        sh.top = int(top + row * Inches(TILE_H + GAP))
        sh.width, sh.height = int(w), Inches(TILE_H)
        bp = sh._element.find(qn("p:txBody")).find(qn("a:bodyPr"))
        bp.set("lIns", str(INS))
        bp.set("rIns", str(INS))
        for p, text in zip(sh._element.find(qn("p:txBody")).findall(qn("a:p")), BREAKS[sid]):
            set_para(p, text)
    prs.save(out)
    print("saved:", out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else sys.argv[1])
