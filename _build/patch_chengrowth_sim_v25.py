# ver2.4 → ver2.5：サンクスLINEをSIMから外す（CVが付かないのにCPAに費用が乗るため）。資料ではオプションとして別に見せる
import sys
from openpyxl import load_workbook
from openpyxl.workbook.properties import CalcProperties
SRC,OUT=sys.argv[1],sys.argv[2]
wb=load_workbook(SRC); ws=wb['SIM3_LINE登録2%']
for c in 'EFGHIJ':
    ws[f'{c}19']=False
    ws[f'{c}67']=0
ws['D67']=0
ws['N67']='サンクスLINE（初期10万・月5万）はオプション。CVが付かないためSIMの費用に入れない'
ws['N24']='完了画面（サンクスLINE）：オプションのためSIMには入れない'
wb.calculation=CalcProperties(fullCalcOnLoad=True)
wb.save(OUT)
