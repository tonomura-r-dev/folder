# -*- coding: utf-8 -*-
"""すなお：殿村さんのPC保存版 ver4.4 に、ver4.5〜4.8 の直しだけを当てて ver4.9 にする（2026-10-20 殿村さん「おけ」）。

    python -m _build.patch_sunao_pc44_to_v49 <PC保存のver4.4.pptx> <Appendix ver1.2.pptx> <業界別ver3.7.pptx> [出力.pptx]

  作り直しはしない（殿村さんがPCで直した概算お見積り＝月次レポート・定例会の行を残すため）。当てるのは次の4つだけ。
  1. 運用内容：プラン②の「結果の分析・改善（月次レポート・定例会）」を「対応」に（ver4.5）
  2. 現状把握：LPの画面（殿村さんのスクショ）を差し込み枠に貼り、LINEのステータスを全部「なし」に（ver4.6）
  3. シーズナリティ：Googleトレンド「保険営業」のグラフを差し込み枠に貼り、リードの文を差し替え（ver4.6）
  4. 新規友だちの追加（サンクスLINE誘導）：2枚 → Appendixのページ（動画のQRつき）1枚＋費用の札（ver4.8）。
     色の直し（赤い注記→紺・強調の赤枠を外す）はこの1枚だけに当てる（ほかのページは殿村さんの版のまま）
"""
import sys

from pptx import Presentation
from pptx.util import Emu, Inches

from _build.build_sunao_v4 import LP_IMG, TREND_IMG, place_pic
from _build.patch_sunao_v4_colors import fix_slide, fix_unreadable
from _build import patch_sunao_v48_thanks as v48


def title_of(s):
    for sh in s.shapes:
        if sh.has_text_frame and sh.top < Inches(0.6) and sh.text_frame.text.strip():
            return sh.text_frame.text.strip()
    return ""


def slide(prs, title):
    hit = [s for s in prs.slides if title_of(s) == title]
    assert len(hit) == 1, (title, len(hit))
    return hit[0]


def set_text(sh, text):
    """最初の段落・最初の文字列の書式のまま、文字を差し替える（段落は1つにする）"""
    tf = sh.text_frame
    for p in tf.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    runs = tf.paragraphs[0].runs
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def swap_slot(s, label, img):
    """差し込み枠（名前が「差し込み 〇〇」）を外して、同じ場所に画像を収める"""
    sh = next(x for x in s.shapes if x.name.startswith("差し込み") and x.text_frame.text.strip().startswith(label))
    box = [Emu(v).inches for v in (sh.left, sh.top, sh.width, sh.height)]
    sh._element.getparent().remove(sh._element)
    place_pic(s, img, *box)


def main(src, appendix, gyokai, out):
    prs = Presentation(src)

    # 1. 運用内容：プラン②も月次レポート・定例会に対応
    s = slide(prs, "運用内容：プランごとの対応範囲")
    t = next(sh for sh in s.shapes if sh.has_table).table
    r = next(i for i in range(len(t.rows)) if t.cell(i, 0).text.startswith("結果の分析・改善"))
    if t.cell(r, 2).text != "対応":                           # 殿村さんのPC保存版では、すでに「対応」に直してあった
        t.cell(r, 2).text_frame.paragraphs[0].runs[0].text = "対応"

    # 2. 現状把握：LPの画面とLINEのステータス
    s = slide(prs, "現状把握")
    for sid in (171, 173, 175, 179, 181, 183, 185, 187, 189):
        sh = next(x for x in s.shapes if x.shape_id == sid)
        assert sh.text_frame.text.strip() == "--", (sid, sh.text_frame.text)
        set_text(sh, "なし")
    swap_slot(s, "LPの画面", LP_IMG)

    # 3. シーズナリティ：Googleトレンド「保険営業」
    s = slide(prs, "ニーズ調査（シーズナリティ）")
    swap_slot(s, "Googleトレンドの画像", TREND_IMG)
    set_text(next(x for x in s.shapes if x.shape_id == 5),
             "「保険営業」の検索は1年中ほぼ同じ水準。イベントの無い月も、平常期の配信でつながりを保つ。")

    prs.save(out)

    # 4. サンクスLINE誘導：Appendixのページ＋費用の札（ver4.8と同じ手順）。色の直しは差し替えた1枚だけ
    v48.main(out, appendix, gyokai, out)
    prs = Presentation(out)
    new = slide(prs, "新規友だちの追加（サンクスLINE誘導）")
    print("colors:", fix_slide(new), fix_unreadable(new))
    prs.save(out)
    print("saved:", out, len(prs.slides), "枚")


if __name__ == "__main__":
    main(*sys.argv[1:4], sys.argv[4] if len(sys.argv) > 4 else sys.argv[1])
