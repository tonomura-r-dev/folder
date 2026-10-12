# -*- coding: utf-8 -*-
"""文字の大きさを一律で上げる（2026-10-12 殿村さん「先に文字の大きさを整えて」）。

    python -m _build.arauchi_v3.quick_bump <入力.pptx> <出力.pptx> --slides 8,9,18,19,...

14pt未満の文字を「+2pt、ただし最低11pt、上限14pt」にする（指定したスライドだけ）。
グループの中・表のセルも対象。大きさの指定が無い（継承の）文字は数だけ数えて触らない。
"""
import argparse

from pptx import Presentation
from pptx.oxml.ns import qn


def new_size(cur):
    if cur >= 14:
        return cur
    return min(max(cur + 2, 11), 14)


def txbodies(shapes):
    for sh in shapes:
        if sh.shape_type == 6:
            yield from txbodies(sh.shapes)
            continue
        if sh.has_text_frame:
            yield sh.text_frame._txBody
        if getattr(sh, "has_table", False) and sh.has_table:
            for row in sh.table.rows:
                for c in row.cells:
                    yield c.text_frame._txBody


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("src")
    ap.add_argument("out")
    ap.add_argument("--slides", required=True)
    a = ap.parse_args()
    nums = [int(x) for x in a.slides.split(",") if x.strip()]
    prs = Presentation(a.src)
    inherit = 0
    for n in nums:
        s = prs.slides[n - 1]
        for tb in txbodies(s.shapes):
            for rpr in list(tb.iter(qn("a:rPr"))) + list(tb.iter(qn("a:endParaRPr"))):
                sz = rpr.get("sz")
                if sz is None:
                    if rpr.tag == qn("a:rPr"):
                        inherit += 1
                    continue
                rpr.set("sz", str(int(round(new_size(int(sz) / 100) * 100))))
    prs.save(a.out)
    print("saved", a.out, "inherited runs untouched:", inherit)


if __name__ == "__main__":
    main()
