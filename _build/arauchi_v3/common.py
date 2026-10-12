# -*- coding: utf-8 -*-
"""あらうち提案資料 v3（2026-10-12 殿村さん指示「文字を大きく・S1・S15・S16/S17統合」）の共通部品。

各スライドの直しは sNN.py（NN＝元の資料のスライド番号）に apply(ctx) として書く。
ctx.prs＝Presentation、ctx.orig[n]＝元のn枚目のスライド、ctx.dup_after(slide)＝スライドを複製して直後に置く。
"""
import copy

from lxml import etree
from pptx.oxml.ns import qn
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.util import Emu, Inches, Pt

R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

# ---------------------------------------------------------------- 文字サイズの基準（全スライド共通）
FLOOR_BODY = 11.0     # 本文・表の中・フロー図の文字の最低
FLOOR_NOTE = 10.0     # 注記・出典・URLの最低
KEEP_FROM = 14.0      # これ以上の大きさの文字は、無理に大きくしなくてよい


def ids(slide):
    return {sh.shape_id: sh for sh in slide.shapes}


def iter_shapes(shapes):
    """グループの中まで含めて図形をたどる"""
    for sh in shapes:
        yield sh
        if sh.shape_type == 6:
            yield from iter_shapes(sh.shapes)


def txbody(shape):
    return shape.text_frame._txBody


def all_rpr(shape):
    tb = txbody(shape)
    return list(tb.iter(qn("a:rPr"))) + list(tb.iter(qn("a:endParaRPr")))


def set_size_all(shape, pt):
    """図形の中の文字をすべて pt にする"""
    for rpr in all_rpr(shape):
        rpr.set("sz", str(int(round(pt * 100))))


def bump_sizes(shape, mapping=None, add=None, floor=None, cap=None):
    """図形の中の文字の大きさを、元の大きさに応じて上げる。
    mapping={元pt: 新pt}、add＝一律に足すpt、floor＝最低、cap＝上限（KEEP_FROM以上の文字は触らない）"""
    for rpr in all_rpr(shape):
        sz = rpr.get("sz")
        if sz is None:
            continue
        cur = int(sz) / 100
        if cur >= KEEP_FROM and mapping is None:
            continue
        new = cur
        if mapping and cur in mapping:
            new = mapping[cur]
        elif add:
            new = cur + add
        if floor:
            new = max(new, floor)
        if cap:
            new = min(new, max(cap, cur))
        rpr.set("sz", str(int(round(new * 100))))


def set_color(shape, para_idx, hex_):
    """段落 para_idx の全 run の文字色を hex_ にする"""
    p_el = txbody(shape).findall(qn("a:p"))[para_idx]
    for rpr in [r.find(qn("a:rPr")) for r in p_el.findall(qn("a:r"))]:
        for f in rpr.findall(qn("a:solidFill")):
            rpr.remove(f)
        fill = etree.Element(qn("a:solidFill"))
        etree.SubElement(fill, qn("a:srgbClr"), val=hex_)
        ln = rpr.find(qn("a:ln"))
        if ln is not None:
            ln.addnext(fill)
        else:
            rpr.insert(0, fill)


def set_insets(shape, l=None, t=None, r=None, b=None):
    """文字の余白（インチ）。None は変えない"""
    bp = txbody(shape).find(qn("a:bodyPr"))
    for k, v in (("lIns", l), ("tIns", t), ("rIns", r), ("bIns", b)):
        if v is not None:
            bp.set(k, str(int(Inches(v))))


def set_anchor(shape, anchor):
    """上下の位置：t / ctr / b"""
    txbody(shape).find(qn("a:bodyPr")).set("anchor", anchor)


def set_autofit_none(shape):
    """「図形を文字に合わせる」等を外す（大きさを固定する）"""
    bp = txbody(shape).find(qn("a:bodyPr"))
    for tag in ("a:spAutoFit", "a:normAutofit", "a:noAutofit"):
        for el in bp.findall(qn(tag)):
            bp.remove(el)
    etree.SubElement(bp, qn("a:noAutofit"))


def _ppr(p_el):
    ppr = p_el.find(qn("a:pPr"))
    if ppr is None:
        ppr = etree.Element(qn("a:pPr"))
        p_el.insert(0, ppr)
    return ppr


def set_line_spacing(shape, pct):
    """全段落の行間を pct（%）にする"""
    for p_el in txbody(shape).findall(qn("a:p")):
        ppr = _ppr(p_el)
        for old in ppr.findall(qn("a:lnSpc")):
            ppr.remove(old)
        ln = etree.Element(qn("a:lnSpc"))
        etree.SubElement(ln, qn("a:spcPct"), val=str(int(pct * 1000)))
        ppr.insert(0, ln)


def set_spacing(shape, before=None, after=None):
    """全段落の段落前・段落後の間隔（pt）"""
    for p_el in txbody(shape).findall(qn("a:p")):
        ppr = _ppr(p_el)
        for tag, v in (("a:spcBef", before), ("a:spcAft", after)):
            if v is None:
                continue
            for old in ppr.findall(qn(tag)):
                ppr.remove(old)
            el = etree.Element(qn(tag))
            etree.SubElement(el, qn("a:spcPts"), val=str(int(v * 100)))
            ln = ppr.find(qn("a:lnSpc"))
            if tag == "a:spcBef":
                (ln.addnext(el) if ln is not None else ppr.insert(0, el))
            else:
                sb = ppr.find(qn("a:spcBef"))
                anchor = sb if sb is not None else ln
                (anchor.addnext(el) if anchor is not None else ppr.insert(0, el))


def make_run(template_rpr, text):
    r = etree.Element(qn("a:r"))
    rpr = copy.deepcopy(template_rpr) if template_rpr is not None else etree.Element(qn("a:rPr"), lang="ja-JP")
    rpr.tag = qn("a:rPr")
    r.append(rpr)
    t = etree.SubElement(r, qn("a:t"))
    t.text = text
    return r


def set_paras(shape, paras, size=None, plain=False, tmpl_rpr=None, align=None):
    """文字を差し替える。paras＝段落ごとの [(文字, 太字 or None), ...]（文字だけの段落は "文字" でも可）。
    各段落の最初の run の書式を土台にする。段落が足りなければ最後の段落を複製する。"""
    tb = txbody(shape)
    ps = tb.findall(qn("a:p"))
    norm = [[(p, None)] if isinstance(p, str) else p for p in paras]
    first_rpr = tb.find(".//" + qn("a:rPr"))
    if first_rpr is None:
        first_rpr = tmpl_rpr
    while len(ps) < len(norm):
        ps[-1].addnext(copy.deepcopy(ps[-1]))
        ps = tb.findall(qn("a:p"))
    for p_el, runs in zip(ps, norm):
        rs = p_el.findall(qn("a:r"))
        base_rpr = rs[0].find(qn("a:rPr")) if rs else first_rpr
        base_rpr = copy.deepcopy(base_rpr) if base_rpr is not None else None
        for child in list(p_el):
            if child.tag in (qn("a:r"), qn("a:br"), qn("a:fld")):
                p_el.remove(child)
        end = p_el.find(qn("a:endParaRPr"))
        for text, bold in runs:
            r = make_run(base_rpr, text)
            rpr = r.find(qn("a:rPr"))
            if bold is not None:
                rpr.set("b", "1" if bold else "0")
            if plain:
                for f in rpr.findall(qn("a:solidFill")):
                    rpr.remove(f)
            (end.addprevious(r) if end is not None else p_el.append(r))
        if align is not None:
            _ppr(p_el).set("algn", align)
    for p_el in ps[len(norm):]:
        tb.remove(p_el)
    if size is not None:
        set_size_all(shape, size)


def delete(shape):
    el = shape._element
    el.getparent().remove(el)


def clone_into(dst_slide, src_shape, dy=0.0):
    """図形を複製して dst_slide の最前面に置く（idは振り直す）"""
    el = copy.deepcopy(src_shape._element)
    for ext in el.iter(qn("a:extLst")):     # creationId の重複を避ける
        ext.getparent().remove(ext)
        break
    dst_slide.shapes._spTree.append(el)
    nid = max(sh.shape_id for sh in dst_slide.shapes) + 1
    el.find(".//" + qn("p:cNvPr")).set("id", str(nid))
    sh = [s for s in dst_slide.shapes if s.shape_id == nid][-1]
    if dy:
        sh.top = sh.top + Inches(dy)
    return sh


def bring_to_front(slide, shape):
    slide.shapes._spTree.append(shape._element)


def move(shape, x=None, y=None, w=None, h=None):
    """インチで位置・大きさを指定（None は変えない）"""
    if x is not None:
        shape.left = Inches(x)
    if y is not None:
        shape.top = Inches(y)
    if w is not None:
        shape.width = Inches(w)
    if h is not None:
        shape.height = Inches(h)


# ---------------------------------------------------------------- スライドの複製・削除
class Ctx:
    def __init__(self, prs):
        self.prs = prs
        self.orig = {i + 1: s for i, s in enumerate(prs.slides)}
        self.orig_id = {i + 1: s.slide_id for i, s in enumerate(prs.slides)}   # 消したスライドは slide_id を引けなくなるので先に控える
        self.id_to_no = {v: k for k, v in self.orig_id.items()}
        self.parent = {}          # 複製したスライド → 元のスライド番号
        self.to_delete = []

    def index_of(self, slide):
        for i, s in enumerate(self.prs.slides):
            if s.slide_id == slide.slide_id:
                return i
        raise ValueError("slide not found")

    def dup_after(self, slide):
        """slide を丸ごと複製し、slide のすぐ後ろ（すでに複製があればその後ろ）に置いて返す。
        画像などの参照も引き継ぐ。ノートは複製しない。"""
        prs = self.prs
        new = prs.slides.add_slide(slide.slide_layout)
        tree = new.shapes._spTree
        for el in list(tree):
            if etree.QName(el).localname not in ("nvGrpSpPr", "grpSpPr"):
                tree.remove(el)
        for el in slide.shapes._spTree:
            if etree.QName(el).localname not in ("nvGrpSpPr", "grpSpPr"):
                tree.append(copy.deepcopy(el))
        # 背景
        src_bg = slide._element.cSld.find(qn("p:bg"))
        if src_bg is not None:
            new._element.cSld.insert(0, copy.deepcopy(src_bg))
        rid_map = {}
        for rid, rel in slide.part.rels.items():
            if rel.reltype in (RT.SLIDE_LAYOUT, RT.NOTES_SLIDE):
                continue
            if rel.is_external:
                rid_map[rid] = new.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
            else:
                rid_map[rid] = new.part.relate_to(rel.target_part, rel.reltype)
        for el in new._element.iter():
            for attr, val in list(el.attrib.items()):
                if attr.startswith("{%s}" % R_NS) and val in rid_map:
                    el.set(attr, rid_map[val])
        # 並び：元のスライド（とその複製）の直後
        root_no = self.parent.get(slide.slide_id, None)
        if root_no is None:
            root_no = self.id_to_no[slide.slide_id]
        lst = prs.slides._sldIdLst
        new_el = lst[-1]
        lst.remove(new_el)
        family = [i for i, s in enumerate(prs.slides) if s.slide_id == slide.slide_id or self.parent.get(s.slide_id) == root_no]
        lst.insert(max(family) + 1, new_el)
        self.parent[new.slide_id] = root_no
        return new

    def delete_slide(self, slide):
        prs = self.prs
        lst = prs.slides._sldIdLst
        for sld_id in list(lst):
            if prs.part.related_part(sld_id.rId) is slide.part:
                prs.part.drop_rel(sld_id.rId)
                lst.remove(sld_id)
                return
        raise ValueError("slide not found")

    def mapping(self):
        """元のスライド番号 → 出来上がりのスライド番号（複数あり）"""
        out = {}
        for i, s in enumerate(self.prs.slides, 1):
            no = self.parent.get(s.slide_id)
            if no is None:
                no = self.id_to_no[s.slide_id]
            out.setdefault(no, []).append(i)
        return out
