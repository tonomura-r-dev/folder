# -*- coding: utf-8 -*-
"""v3 の機械チェック：情報が消えていないか・文字が大きくなったか・はみ出していないか。

    python -m _build.arauchi_v3.checks <元の資料.pptx> <出力.pptx> [--orig 2,3,4]

出力.pptx.map.json（build.py が書く）を使って、元のスライドと出来上がりのスライドを対応させる。
--orig を付けると、その元番号のスライドだけ調べる。
"""
import argparse
import json
import re
import sys

from pptx import Presentation
from pptx.oxml.ns import qn

FLOOR_BODY = 11.0
FLOOR_NOTE = 10.0
KEEP_FROM = 14.0
RESTRUCTURED = {1, 15, 16, 17}   # 指示で作り直したスライド（文言の照合は目視で）


def shapes_deep(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type == 6:
            yield from shapes_deep(sh.shapes)


def text_frames(slide):
    for sh in shapes_deep(slide.shapes):
        if sh.has_text_frame:
            yield sh, sh.text_frame
        if getattr(sh, "has_table", False) and sh.has_table:
            for row in sh.table.rows:
                for c in row.cells:
                    yield sh, c.text_frame


def paras(slide):
    out = []
    for _, tf in text_frames(slide):
        for p in tf.paragraphs:
            t = re.sub(r"\s+", "", "".join(r.text for r in p.runs))
            if t:
                out.append(t)
    return out


def run_sizes(slide):
    """(shape_id, 文字, 大きさpt or None)"""
    out = []
    for sh, tf in text_frames(slide):
        for p in tf.paragraphs:
            for r in p.runs:
                if r.text.strip():
                    out.append((sh.shape_id, r.text, r.font.size.pt if r.font.size else None))
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("src")
    ap.add_argument("out")
    ap.add_argument("--orig", default="")
    a = ap.parse_args()
    only = {int(x) for x in a.orig.split(",") if x.strip()} if a.orig else None
    A, B = Presentation(a.src), Presentation(a.out)
    m = {int(k): v for k, v in json.loads(open(a.out + ".map.json", encoding="utf-8").read()).items()}
    W, H = B.slide_width, B.slide_height
    n_err = 0
    for n in range(1, 48):
        if only is not None and n not in only:
            continue
        a_sl = A.slides[n - 1]
        new_pages = m.get(n, [])
        b_sls = [B.slides[i - 1] for i in new_pages]
        head = f"元S{n} → 出来上がり {new_pages if new_pages else '（削除）'}"
        msgs = []
        # 1) 情報が消えていないか
        if b_sls:
            joined = "".join("".join(paras(s)) for s in b_sls)
            missing = [p for p in paras(a_sl) if p not in joined]
            if missing:
                tag = "要目視（作り直したスライド）" if n in RESTRUCTURED else "★消えた文言"
                for p in missing:
                    msgs.append(f"  {tag}: {p[:70]}")
                if n not in RESTRUCTURED:
                    n_err += len(missing)
        # 2) 文字の大きさ
        a_min = {}
        for sid, t, sz in run_sizes(a_sl):
            if sz is not None:
                a_min[sid] = min(a_min.get(sid, 99), sz)
        for page, s in zip(new_pages, b_sls):
            b_min = {}
            for sid, t, sz in run_sizes(s):
                if sz is None:
                    continue
                b_min[sid] = min(b_min.get(sid, 99), sz)
                if sz < FLOOR_NOTE:
                    msgs.append(f"  ★p{page} [{sid}] {sz}pt（10pt未満）: {t[:30]}")
                    n_err += 1
                elif sz < FLOOR_BODY:
                    msgs.append(f"  p{page} [{sid}] {sz}pt（11pt未満＝注記ならOK）: {t[:30]}")
            for sid, mn in b_min.items():
                if sid in a_min and a_min[sid] < KEEP_FROM and mn <= a_min[sid]:
                    msgs.append(f"  ★p{page} [{sid}] 大きくなっていない {a_min[sid]}→{mn}pt")
                    n_err += 1
            # 3) スライドの外にはみ出した図形
            for sh in shapes_deep(s.shapes):
                if sh.left is None or sh.width is None:
                    continue
                if sh.left < -10000 or sh.top < -10000 or sh.left + sh.width > W + 10000 or sh.top + sh.height > H + 10000:
                    msgs.append(f"  ★p{page} [{sh.shape_id}] スライドの外にはみ出し")
                    n_err += 1
        print(head)
        print("\n".join(msgs) if msgs else "  OK")
    print(f"\n★の数: {n_err}")
    sys.exit(1 if n_err else 0)


if __name__ == "__main__":
    main()
