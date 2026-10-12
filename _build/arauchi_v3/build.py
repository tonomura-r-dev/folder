# -*- coding: utf-8 -*-
"""あらうち提案資料 v3 をつくる（2026-10-12 殿村さん指示）。

    python -m _build.arauchi_v3.build <元の資料.pptx> <出力.pptx> [--only 1,15,17]

- 元の資料＝殿村さんが10/12に送り直した47枚（S1〜S17修正版をPCで保存したもの）
- 各スライドの直しは sNN.py の apply(ctx)（NN＝元のスライド番号）。--only を付けると、その番号の直しだけ当てる
- 最後に元のS16（章の扉「施策の詳細（新規友だち獲得）」）を消す（S17と統合したため）
- 出力の横に .map.json（元の番号 → 出来上がりの番号）を書く
"""
import argparse
import importlib
import json
import sys
from pathlib import Path

from pptx import Presentation

from _build.arauchi_v3.common import Ctx, ids, set_paras

HERE = Path(__file__).resolve().parent


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("src")
    ap.add_argument("out")
    ap.add_argument("--only", default="")
    a = ap.parse_args()
    only = {int(x) for x in a.only.split(",") if x.strip()} if a.only else None

    prs = Presentation(a.src)
    assert len(prs.slides) == 47, "元の資料は47枚のはず"
    ctx = Ctx(prs)
    for n in range(1, 48):
        if only is not None and n not in only:
            continue
        if not (HERE / f"s{n:02d}.py").exists():
            continue
        mod = importlib.import_module(f"_build.arauchi_v3.s{n:02d}")
        mod.apply(ctx)
    # 元のS16（章の扉）を消す：S17と1枚に統合した（--only でS17を当てたときだけ）
    if only is None or 17 in only:
        ctx.delete_slide(ctx.orig[16])
    # S11の「出典は6ページ」を、元のS6が最終的に何ページ目になったかに合わせる
    m = ctx.mapping()
    s6_page = m[6][0]
    for sh in ctx.orig[11].shapes:
        if sh.has_text_frame and "（出典は6ページ）" in sh.text_frame.text and s6_page != 6:
            for r in sh._element.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}t"):
                if r.text and "（出典は6ページ）" in r.text:
                    r.text = r.text.replace("（出典は6ページ）", f"（出典は{s6_page}ページ）")
    prs.save(a.out)
    Path(a.out + ".map.json").write_text(json.dumps({str(k): v for k, v in sorted(m.items())}, ensure_ascii=False))
    print("saved:", a.out, "slides:", len(prs.slides))


if __name__ == "__main__":
    main()
