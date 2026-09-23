# -*- coding: utf-8 -*-
"""Slide 22から成果報酬プランの記載を削除し、9月開始の数値に更新"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
TARGET = ROOT / "20260804_株式会社ビイサイドプランニング御中_公式LINEのご提案v05.pptx"

NAVY = "1F285A"; INK = "333333"; MUT = "808080"; WHITE = "FFFFFF"

prs = Presentation(str(TARGET))
slides = list(prs.slides)


def set_font(run, size, bold=None, color=INK, name="メイリオ"):
    f = run.font; f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    f.name = name
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        e = rPr.find(qn(tag))
        if e is None:
            e = rPr.makeelement(qn(tag), {}); rPr.append(e)
        e.set("typeface", name)
    f.color.rgb = RGBColor.from_string(color)


ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
ANCH = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}


def reset_tf(tf):
    for para in tf.paragraphs[1:]:
        para._p.getparent().remove(para._p)
    for r in list(tf.paragraphs[0].runs):
        r._r.getparent().remove(r._r)


def put_text(tf, paras, anchor="t", ml=0.06, mr=0.06, mt=0.03, mb=0.03):
    reset_tf(tf); tf.word_wrap = True
    tf.margin_left = Inches(ml); tf.margin_right = Inches(mr)
    tf.margin_top = Inches(mt); tf.margin_bottom = Inches(mb)
    tf.vertical_anchor = ANCH[anchor]
    first = True
    for p in paras:
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        para.alignment = ALIGN[p.get("align", "l")]
        if p.get("sa") is not None:
            para.space_after = Pt(p["sa"])
        if p.get("ls") is not None:
            para.line_spacing = p["ls"]
        for (t, sz, b, c) in p["runs"]:
            r = para.add_run(); r.text = t
            set_font(r, sz, b, c)


def find(slide, needle):
    for sh in slide.shapes:
        if sh.has_text_frame and needle in sh.text_frame.text:
            return sh
    return None


# ---------- Slide 22 ----------
s = slides[22]  # 新スライド挿入で1つ後ろへ

# ① 成果報酬プランの記載を削除
opt = find(s, "成果報酬プランの場合")
if opt is not None:
    opt._element.getparent().remove(opt._element)
    print("removed: 成果報酬プラン行")

# ② 9月開始の数値へ更新（CV / CPA）
for sh in s.shapes:
    if not sh.has_text_frame:
        continue
    t = sh.text_frame.text.strip()
    if t.startswith("135件"):
        put_text(sh.text_frame, [{"runs": [("146件（半年累計）", 9.5, None, INK)], "align": "c"}],
                 anchor="m", ml=0.03, mr=0.03, mt=0.0, mb=0.0)
        print("updated: CV 135件 -> 146件")
    elif t.startswith("13,287"):
        put_text(sh.text_frame, [{"runs": [("12,286円", 9.5, None, INK)], "align": "c"}],
                 anchor="m", ml=0.03, mr=0.03, mt=0.0, mb=0.0)
        print("updated: CPA 13,287 -> 12,286")

# ③ 帯の獲得単価を更新
ans = find(s, "LINEで増える来場者の獲得単価")
if ans is not None:
    put_text(ans.text_frame,
             [{"runs": [("LINEで増える来場者の獲得単価　", 12, True, WHITE),
                        ("約13,500円", 19, True, WHITE),
                        ("　（現状の来場単価 16,000〜18,000円 → 16〜25%減）", 11.5, None, WHITE)],
               "align": "c"}],
             anchor="m", ml=0.0, mr=0.0, mt=0.0, mb=0.0)
    print("updated: 獲得単価 13,700 -> 13,500")

# ④ 注記の算出根拠を更新
note = find(s, "CV＝フェア事前申込")
if note is not None:
    put_text(note.text_frame,
             [{"runs": [("※CV＝フェア事前申込＋来場引き上げの合計", 9, None, MUT)], "sa": 2},
              {"runs": [("※獲得単価＝半年のLINE費用÷増分来場者数", 9, None, MUT)]}],
             anchor="m", ml=0.0, mr=0.0)
    print("updated: 注記")

prs.save(str(TARGET))
print("saved:", TARGET)
