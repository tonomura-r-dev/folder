# -*- coding: utf-8 -*-
"""あらうち提案資料 v6（2026-10-13 殿村さん指示「成果報酬プラン（プラン④）の追加」）。

元＝v5（殿村さんがアップしたファイルのつなぎを直した版）。直す・足すのは3ページだけ：
  p24 料金プランの比較（プラン①〜④）  … 元の「3プラン比較」を4プランに。料金の行だけにして、運用内容の行は次のページへ
  p25 運用内容の比較（プラン①〜③）    … 元の運用内容の行（企画配信・ステップ配信…）。プラン④は26ページへ案内
  p26 プラン④｜成果報酬型            … 有効CVの数え方・料金・参考試算・注記
ほかのページは変更しない（13ページの費用対効果はプラン①の試算のまま）。
"""
import copy

from lxml import etree

from _build.arauchi_v4.lib import (NAVY, RED, TEXT, add_box, add_table, delete, ids, make_para, qn, set_paras)

GREEN = "00B050"          # 資料で使っている緑（S7・S13）
GREEN_TINT = "E6F4EA"
NAVY_TINT = "E8ECF6"
PANEL = "F4F7FF"
GRID = "BFC8DC"
GREY_TXT = "444444"


# ---------------------------------------------------------------- 部品
def find_shape(slide, name=None, text=None):
    for sh in slide.shapes:
        if name is not None and sh.name != name:
            continue
        if text is not None and not (sh.has_text_frame and text in sh.text_frame.text):
            continue
        return sh
    raise KeyError((name, text))


def clone_shape(dst, src, new_id):
    el = copy.deepcopy(src._element)
    for ext in list(el.iter(qn("a:extLst"))):
        ext.getparent().remove(ext)
    el.find(".//" + qn("p:cNvPr")).set("id", str(new_id))
    dst.shapes._spTree.append(el)
    return [s for s in dst.shapes if s.shape_id == new_id][-1]


def new_page(prs, like, index):
    """like と同じレイアウトの空のスライドを作り、index 番目（0始まり）に置く"""
    s = prs.slides.add_slide(like.slide_layout)
    tree = s.shapes._spTree
    for el in list(tree):
        if etree.QName(el).localname not in ("nvGrpSpPr", "grpSpPr"):
            tree.remove(el)
    lst = prs.slides._sldIdLst
    el = lst[-1]
    lst.remove(el)
    lst.insert(index, el)
    return s


def header(slide, kit, tab, headline, tab_fill=None):
    """見出し（タイトル・「あらうち様向けプラン」・左の紺のタブ・要旨）をそろえる"""
    nid = 2
    out = {}
    for key in ("title", "label", "tab", "headline"):
        out[key] = clone_shape(slide, kit[key], nid)
        nid += 1
    set_paras(out["tab"], [tab])
    set_paras(out["headline"], [headline])
    if tab_fill:
        sp = out["tab"]._element.spPr
        for tag in ("a:solidFill",):
            for e in sp.findall(qn(tag)):
                sp.remove(e)
        f = etree.Element(qn("a:solidFill"))
        etree.SubElement(f, qn("a:srgbClr"), val=tab_fill)
        sp.find(qn("a:prstGeom")).addnext(f)
    return out


def hdr_cell(paras, fill, size=13, bold=True, **kw):
    d = dict(text=paras, size=size, bold=bold, color="FFFFFF", fill=fill, align="ctr",
             borders={"l": (0.75, "FFFFFF"), "r": (0.75, "FFFFFF"), "t": (0.75, "FFFFFF"), "b": (0.75, "FFFFFF")})
    d.update(kw)
    return d


def body_cell(text, fill="FFFFFF", size=13, bold=False, color=TEXT, **kw):
    d = dict(text=text, size=size, bold=bold, color=color, fill=fill, align="ctr",
             borders={"l": (0.75, GRID), "r": (0.75, GRID), "t": (0.75, GRID), "b": (0.75, GRID)})
    d.update(kw)
    return d


def plan_head(num, name, fill):
    return hdr_cell([[(num, {"size": 14, "bold": True})], [(name, {"size": 12, "bold": False})]], fill)


def merged_away(fill, **kw):
    d = body_cell("", fill=fill)
    d.update(kw)
    return d


COLS = [1.7, 2.075, 2.075, 2.075, 2.075]
MONTHLY = ["13.5万円", "8.5万円", "ツール利用料\n＋手数料", "8,000円×\n有効CV件数"]


# ---------------------------------------------------------------- p24 料金プランの比較（プラン①〜④）
def p24(s):
    """既存の見出し4つは残し、表と注記を作り直す（運用内容の行は p25 へ移す）"""
    for sh in list(s.shapes):
        if sh.has_text_frame and "ご希望の運用範囲" in sh.text_frame.text:
            set_paras(sh, ["料金の払い方が異なる4つのプランをご用意しました。"])
        elif sh.has_text_frame and sh.text_frame.text.strip() == "プラン比較":
            set_paras(sh, ["料金プランの比較"])
        elif sh.shape_type == 19 or sh.name == "注記":     # 表・注記は作り直す
            delete(sh)
    grp_l = hdr_cell("コンサル型（毎月固定）", NAVY, size=12, gridSpan=3)
    grp_r = hdr_cell("成果報酬型（成果連動）", GREEN, size=12)
    rows = [
        [hdr_cell("項目", NAVY, size=14, rowSpan=2), grp_l, merged_away(NAVY, hMerge=True), merged_away(NAVY, hMerge=True), grp_r],
        [merged_away(NAVY, vMerge=True),
         plan_head("プラン①", "総合運用", NAVY), plan_head("プラン②", "獲得重視", NAVY), plan_head("プラン③", "初期導入", NAVY),
         plan_head("プラン④", "成果報酬型", GREEN)],
    ]

    def row(label, vals, fills=None, size=13, bold=False, label_size=13):
        out = [body_cell(label, fill=NAVY_TINT, size=label_size, bold=True, color=NAVY)]
        for i, v in enumerate(vals):
            out.append(body_cell(v, fill=(GREEN_TINT if i == 3 else "FFFFFF"), size=size, bold=bold))
        return out

    rows.append(row("初期費用", ["21.5万円"] * 4))
    rows.append(row("月額コンサル費", ["10万円", "5万円", "なし", "なし"]))
    rows.append(row("月額ツール・\nLINE費", ["3.5万円", "3.5万円", "別途実費", "確認中"]))
    rows.append(row("成果単価", ["なし", "なし", "なし", "8,000円／有効CV"], bold=False))
    rows[-1][4]["bold"] = True
    rows[-1][4]["color"] = "007A3D"
    last = [body_cell("月額費用", fill=NAVY, size=13, bold=True, color="FFFFFF")]
    for i, v in enumerate(MONTHLY):
        last.append(body_cell(v, fill=("D5ECDC" if i == 3 else NAVY_TINT), size=14, bold=True,
                              color=("007A3D" if i == 3 else NAVY)))
    rows.append(last)
    # 区切り：月額費用の行の上は紺の太い線
    for c in rows[-1]:
        c["borders"] = dict(c["borders"], t=(1.5, NAVY))
    heights = [0.38, 0.64, 0.44, 0.44, 0.64, 0.44, 0.68]
    add_table(s, 0.42, 1.48, COLS, heights, rows, name="料金プラン比較表")
    y = 1.48 + sum(heights) + 0.1
    # 払い方の違い（一目で）
    add_box(s, 0.42, y, 5.3, 0.42, [[("コンサル型（プラン①〜③）｜毎月の固定費でご提供", {})]], size=12.5, bold=True, color=NAVY,
            fill=NAVY_TINT, line=NAVY, line_w=1.0, align="ctr", anchor="ctr", ins=(0.05, 0.02, 0.05, 0.02), name="コンサル型の説明")
    add_box(s, 5.84, y, 4.58, 0.42, [[("成果報酬型（プラン④）｜有効CVの件数で変動", {})]], size=12.5, bold=True, color="007A3D",
            fill=GREEN_TINT, line=GREEN, line_w=1.0, align="ctr", anchor="ctr", ins=(0.05, 0.02, 0.05, 0.02), name="成果報酬型の説明")
    notes = [
        "※価格はあらうち様向けの個別提案価格です（DYMの標準価格ではありません）。初期費用21.5万円は4プラン共通です。",
        "※コンサル型（プラン①〜③）と成果報酬型（プラン④）は別契約です。固定のコンサル費と成果報酬を併せて請求することはありません。",
        "※プラン③は月額コンサル契約がありません。ツールを継続して利用する場合の利用料は別途実費です。",
        "※プラン④は成果保証ではなく、実際の有効CV件数に応じた請求です（詳細は26ページ）。専用フォームの構築に必要な技術要件・ツール費用は確認中のため、確定次第ご案内します。",
    ]
    add_box(s, 0.42, y + 0.5, 10.0, 7.1 - (y + 0.5), [[(t, {})] for t in notes], size=11, color=GREY_TXT, align="l", anchor="t",
            ins=(0.02, 0.02, 0.02, 0.02), after=2, name="注記")


# ---------------------------------------------------------------- p25 運用内容の比較
def p25(prs, kit, s_like):
    s = new_page(prs, s_like, 24)
    header(s, kit, "運用内容の比較", "コンサル型の運用内容を比較します（プラン④は26ページ）。")
    rows = [[hdr_cell("項目", NAVY, size=14),
             plan_head("プラン①", "総合運用", NAVY), plan_head("プラン②", "獲得重視", NAVY), plan_head("プラン③", "初期導入", NAVY),
             plan_head("プラン④", "成果報酬型", GREEN)]]
    last = [body_cell("月額費用", fill=NAVY, size=13, bold=True, color="FFFFFF")]
    for i, v in enumerate(MONTHLY):
        last.append(body_cell(v, fill=("D5ECDC" if i == 3 else NAVY_TINT), size=14, bold=True, color=("007A3D" if i == 3 else NAVY)))
    rows.append(last)
    ops = [("企画配信", ["月4〜5本", "月1〜2本", "なし"]),
           ("ステップ配信", ["10通", "なし", "なし"]),
           ("離脱防止", ["導入・運用", "導入・運用", "初期導入支援"]),
           ("継続改善", ["対応", "限定的", "なし"]),
           ("API連携・開発", ["必要に応じ別途見積", "なし", "なし"])]
    pointer = [[("専用フォームで\n有効CVを計測し、", {})], [("件数に応じて\n成果報酬をご請求", {})], [("（26ページ）", {"size": 12})]]
    for i, (label, vals) in enumerate(ops):
        r = [body_cell(label, fill=NAVY_TINT, size=13, bold=True, color=NAVY)]
        r += [body_cell(v, size=13) for v in vals]
        if i == 0:
            r.append(body_cell(pointer, fill=GREEN_TINT, size=13, bold=True, color="007A3D", rowSpan=len(ops)))
        else:
            r.append(merged_away(GREEN_TINT, vMerge=True))
        rows.append(r)
    heights = [0.62, 0.66] + [0.5] * len(ops)
    add_table(s, 0.42, 1.48, COLS, heights, rows, name="運用内容比較表")
    y = 1.48 + sum(heights) + 0.2
    notes = [
        "※API連携・開発は基本料金に含まれません。必要に応じて、プラン①で別途お見積りします。",
        "※13ページの試算（6か月目に月23件）は、プラン①｜総合運用の基本施策を対象としたものです。プラン②・③の成果は含みません。",
        "※プラン④の参考試算（月間の有効CV 23件（仮））は、26ページに記載しています。",
    ]
    add_box(s, 0.42, y, 10.0, 7.1 - y, [[(t, {})] for t in notes], size=11, color=GREY_TXT, align="l", anchor="t",
            ins=(0.02, 0.02, 0.02, 0.02), after=2, name="注記")
    return s


# ---------------------------------------------------------------- p26 プラン④｜成果報酬型
def p26(prs, kit, s_like):
    s = new_page(prs, s_like, 25)
    header(s, kit, "プラン④｜成果報酬型", "固定のコンサル費はなく、有効CVの件数に応じてご請求します。", tab_fill=GREEN)

    # ---- 有効CVの数え方（4つの流れ）
    add_box(s, 0.42, 1.46, 5.0, 0.3, [[("有効CVの数え方", {})]], size=13, bold=True, color=NAVY, align="l", anchor="ctr",
            ins=(0.02, 0.0, 0.02, 0.0), name="見出し：有効CVの数え方")
    steps = [
        ("1", "友だち", "新規・既存の友だち\nどちらも成果の対象です"),
        ("2", "専用フォーム", "LINE上の専用フォームで\n物件情報を送信（CV地点）"),
        ("3", "有効CVの集計", "有効CVとして集計\n重複・テスト送信などの\n無効CVは除外"),
        ("4", "成果報酬", "有効CV件数 × 8,000円を\n月額成果報酬として\nご請求します"),
    ]
    x0, w, gap, y0, h = 0.42, 2.28, 0.28, 1.8, 1.5
    for i, (n, title, body) in enumerate(steps):
        x = x0 + i * (w + gap)
        fill = GREEN if i == 3 else NAVY
        add_box(s, x, y0, w, 0.4, [[(f"STEP {n}　{title}", {})]], size=13, bold=True, color="FFFFFF", fill=fill, align="ctr", anchor="ctr",
                ins=(0.05, 0.02, 0.05, 0.02), name=f"流れ{n}の見出し")
        add_box(s, x, y0 + 0.4, w, h - 0.4, [[(t, {})] for t in body.split("\n")], size=12, color=TEXT, fill=PANEL, line="D9D9D9", line_w=0.75,
                align="ctr", anchor="ctr", ins=(0.06, 0.04, 0.06, 0.04), name=f"流れ{n}の内容")
        if i < 3:
            from pptx.enum.shapes import MSO_SHAPE
            ar = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, _in(x + w + 0.03), _in(y0 + h / 2 - 0.12), _in(gap - 0.06), _in(0.24))
            el = ar._element
            st = el.find(qn("p:style"))
            if st is not None:
                el.remove(st)
            sp = el.spPr
            for tag in ("a:solidFill", "a:noFill", "a:ln"):
                for e in sp.findall(qn(tag)):
                    sp.remove(e)
            f = etree.SubElement(sp, qn("a:solidFill"))
            etree.SubElement(f, qn("a:srgbClr"), val="7F7F7F")
            ln = etree.SubElement(sp, qn("a:ln"))
            etree.SubElement(ln, qn("a:noFill"))
            ar.name = f"矢印{i + 1}"
            tb = el.find(qn("p:txBody"))
            if tb is not None:
                el.remove(tb)

    # ---- 料金（左）
    y1 = 3.5
    add_box(s, 0.42, y1, 4.85, 0.36, [[("料金（プラン④）", {})]], size=13, bold=True, color="FFFFFF", fill=GREEN, align="ctr", anchor="ctr",
            ins=(0.05, 0.02, 0.05, 0.02), name="見出し：料金")
    price_rows = []
    for label, val, vc in (("初期費用", "215,000円", TEXT), ("月額固定コンサル費", "なし", TEXT),
                           ("成果単価", "8,000円／有効CV", TEXT), ("月額成果報酬", "有効CV件数 × 8,000円", "007A3D")):
        price_rows.append([body_cell(label, fill=NAVY_TINT, size=13, bold=True, color=NAVY, align="l", ins=(0.12, 0.03, 0.08, 0.03)),
                           body_cell(val, size=14, bold=True, color=vc, align="l", ins=(0.12, 0.03, 0.08, 0.03))])
    add_table(s, 0.42, y1 + 0.36, [2.15, 2.7], [0.44] * 4, price_rows, name="料金表")

    # ---- 参考試算（右）
    add_box(s, 5.5, y1, 4.92, 0.36, [[("参考試算（月間）※成果保証ではありません", {})]], size=13, bold=True, color="FFFFFF", fill=GREEN, align="ctr",
            anchor="ctr", ins=(0.05, 0.02, 0.05, 0.02), name="見出し：参考試算")
    calc = [
        [("月間の有効CV　23件（仮）", {"size": 13, "bold": True})],
        [("成果単価　8,000円／有効CV", {"size": 13, "bold": True})],
        [("月間成果報酬　", {"size": 14, "bold": True}), ("184,000円（試算）", {"size": 22, "bold": True, "color": RED})],
        [("13ページの試算（6か月目）の23件を参考値にしています。\n実際の有効CV件数に応じたご請求です。", {"size": 11, "color": GREY_TXT})],
    ]
    add_box(s, 5.5, y1 + 0.36, 4.92, 1.76, calc, size=13, color=TEXT, fill=PANEL, line="D9D9D9", line_w=0.75, align="l", anchor="ctr",
            ins=(0.15, 0.06, 0.12, 0.06), after=3, name="参考試算")

    # ---- 注記
    notes = [
        "※成果の対象は、新規・既存の友だちの両方です。CV地点は、LINE上の専用フォームからの物件情報の送信です。",
        "※専用フォームの設置・計測により、有効CV数を把握します。重複・テスト送信などの無効CVは成果対象から除外します。",
        "※コンサル型（プラン①〜③）と成果報酬型（プラン④）は別契約です。固定のコンサル費と成果報酬を併せて請求することはありません。",
        "※専用フォームの構築に必要な技術要件・ツール費用は確認中です。確定次第ご案内します。",
    ]
    yn = y1 + 0.36 + 1.76 + 0.1
    add_box(s, 0.42, yn, 10.0, 7.1 - yn, [[(t, {})] for t in notes], size=11, color=GREY_TXT, align="l", anchor="t",
            ins=(0.02, 0.02, 0.02, 0.02), after=2, name="注記")
    return s


def _in(v):
    from pptx.util import Inches
    return Inches(v)


def apply(prs):
    s24 = prs.slides[23]
    kit = {"title": find_shape(s24, text="弊社からご提供できる運用プラン"),
           "tab": find_shape(s24, text="プラン比較"),
           "headline": find_shape(s24, text="ご希望の運用範囲"),
           "label": find_shape(s24, text="あらうち様向けプラン")}
    # 新しいページの見出し用に、元の見出しを控える（p24 を書き換える前にコピーしておく）
    kit_copy = {}
    for k, sh in kit.items():
        kit_copy[k] = sh
    p25s = p25(prs, kit_copy, s24)
    p26s = p26(prs, kit_copy, s24)
    p24(s24)
    return s24, p25s, p26s
