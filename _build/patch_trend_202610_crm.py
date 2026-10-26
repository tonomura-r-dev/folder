"""
10月号 CRMオプションを1枚 → 4枚（P.11〜P.14）に組み直す（2026-10-26 殿村さん指示）。
P.11 今回の導入事例／P.12 おすすめの対象・よくある悩み・料金／P.13 導入で変わること・4つの機能／P.14 業種別の活用例
FAQ・今後の予定・利用条件・期待できる効果は載せない。分類タグ＝オプション商材、逆三角形＝伸長、目次はP.11-14の1行。
使い方: python3 _build/patch_trend_202610_crm.py <pptx>
"""
import sys
import copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.oxml.ns import qn
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
import trend_report_engine as engine

F = sys.argv[1]
prs = Presentation(F)
BOTTOM = 7.05
L, W = 0.6, 9.64


def text(sh):
    return sh.text_frame.text.strip() if sh.has_text_frame else ""


s3 = list(prs.slides)[2]
CARD_SRC = next(x for x in s3.shapes if x.top > Inches(3.0) and text(x).startswith("変更点"))
crm_idx = next(i for i, s in enumerate(prs.slides) if any(text(x) == "【CRMオプション】導入事例の追加" for x in s.shapes))
assert crm_idx == len(prs.slides) - 1
for _ in range(3):
    engine.duplicate_slide(prs, crm_idx)
pages = list(prs.slides)[crm_idx:crm_idx + 4]


def clear_body(s):
    head = next(x for x in s.shapes if text(x) == "■ 詳細")
    for x in [x for x in s.shapes if x.top >= head.top + head.height - Inches(0.02) and x.top < Inches(7.0)]:
        x._element.getparent().remove(x._element)
    return Emu(head.top + head.height).inches + 0.08


def new_id(s):
    return str(max(int(e.get("id")) for e in s._element.iter(qn("p:cNvPr"))) + 1)


def card(s, x, y, w, h, heading, lines, body_pt=13, head_pt=18, accent=None):
    el = copy.deepcopy(CARD_SRC._element)
    s.shapes._spTree.append(el)
    cn = el.find(".//" + qn("p:cNvPr"))
    cn.set("id", new_id(s))
    c = next(sh for sh in s.shapes if sh._element is el)
    tf = c.text_frame
    tmpl_h, tmpl_b = copy.deepcopy(tf.paragraphs[0]._p), copy.deepcopy(tf.paragraphs[1]._p)
    for p in list(tf.paragraphs):
        p._p.getparent().remove(p._p)
    items = [(tmpl_h, heading, head_pt, None)]
    if accent:
        items.append((tmpl_h, accent, head_pt + 4, engine.ACCENT))
    items += [(tmpl_b, t, body_pt, None) for t in lines]
    for tm, t, pt, color in items:
        pe = copy.deepcopy(tm)
        for br in pe.findall(qn("a:br")):
            pe.remove(br)
        for r in pe.findall(qn("a:r"))[1:]:
            pe.remove(r)
        tf._txBody.append(pe)
    for p, (tm, t, pt, color) in zip(tf.paragraphs, items):
        p.runs[0].text = t
        p.runs[0].font.size = Pt(pt)
        if color is not None:
            p.runs[0].font.color.rgb = color
            p.runs[0].font.bold = True
            p.space_before = Pt(0)
    tf._txBody.find(qn("a:bodyPr")).set("anchor", "t")
    c.left, c.top, c.width, c.height = Inches(x), Inches(y), Inches(w), Inches(h)
    return c


def put(cell, val, size, bold, color, fill, align):
    cell.fill.solid()
    cell.fill.fore_color.rgb = fill
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = cell.margin_right = Inches(0.12)
    cell.margin_top = cell.margin_bottom = Inches(0.03)
    p = cell.text_frame.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = val
    r.font.size, r.font.bold, r.font.color.rgb = Pt(size), bold, color


def table(s, x, y, colw, rowh, rows, header=True, size=13):
    gf = s.shapes.add_table(len(rows), len(colw), Inches(x), Inches(y), Inches(sum(colw)), Inches(sum(rowh)))
    t = gf.table
    for j, w in enumerate(colw):
        t.columns[j].width = Inches(w)
    for i, h in enumerate(rowh):
        t.rows[i].height = Inches(h)
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            if header and i == 0:
                put(t.cell(i, j), v, size + 1, True, engine.WHITE, engine.NAVY, PP_ALIGN.CENTER)
            elif j == 0:
                put(t.cell(i, j), v, size, True, engine.NAVY, engine.LIGHT_BLUE, PP_ALIGN.CENTER)
            else:
                put(t.cell(i, j), v, size, False, engine.BODY_GRAY, engine.WHITE, PP_ALIGN.LEFT)


def header(s, title, summary):
    engine.set_single_run_text(next(x for x in s.shapes if text(x).startswith("【CRMオプション】")), title)
    d = next(x for x in s.shapes if x.has_text_frame and text(x).startswith("CRMオプションの導入事例") and x.top < Inches(2.9))
    engine.set_multiline_text(d, summary)
    engine.set_tag_highlight(s, "オプション\n商材")
    engine.set_funnel_highlight(s, ["伸長"])


def mid(area_top, total):
    return area_top + (BOTTOM - area_top - total) / 2


cw, gap = (W - 0.2) / 2, 0.2

# ---------- P.11 導入事例 ----------
s = pages[0]
header(s, "【CRMオプション】導入事例の追加", [
    "CRMオプションの媒体資料（Ver2.0・2026年9月更新）に、タオル美術館様の導入事例が2件追加された。",
    "フォームで取った情報をもとに、クーポンの配信と商品の出し分けを行った事例。",
    "フォームの回答数は約2.0倍、配信のクリック率は約2.8倍になった。",
])
at = clear_body(s)
h = 2.85
y = mid(at, h)
card(s, L, y, cw, h, "① 誕生月クーポンでフォームの回答を促進", [
    "・課題：リサーチ機能でアンケートを取っていたが、回答数が伸びない",
    "・解決策：誕生月を聞くフォームを置き、回答した月ごとにクーポンを配信。店頭チラシ・リッチメニュー・あいさつメッセージで回答を促した",
    "・結果：リサーチ機能のときと比べて、ひと月あたりの回答数が約2.0倍",
], body_pt=12, head_pt=15, accent="回答数 約2.0倍")
card(s, L + cw + gap, y, cw, h, "② キャラクターコラボ商品の出し分け", [
    "・課題：コラボ商品を一斉配信しても、開封・クリックにつながらない",
    "・解決策：フォームの回答で好きなキャラクターのタグを付け、コラボ商品の配信を出し分け",
    "・結果：一斉配信と比べて、クリック率が約2.8倍",
], body_pt=12, head_pt=15, accent="クリック率 約2.8倍")

# ---------- P.12 おすすめの対象・よくある悩み・料金 ----------
s = pages[1]
header(s, "【CRMオプション】おすすめの対象とよくある悩み", [
    "CRMオプションは、LINE公式アカウントの中で手軽に完結する顧客管理（CRM）の機能。",
    "フォームで顧客のニーズを集め、顧客ごとにデータを整理して、顧客ごとの配信に使える。",
    "月額5,000円（税別）から始められる。",
])
at = clear_body(s)
h1, h2 = 1.85, 1.45
y = mid(at, h1 + 0.2 + h2)
card(s, L, y, cw, h1, "おすすめの対象", [
    "・配信しているが、成果が伸び悩んでいる",
    "・顧客管理ができていない",
    "・チャット対応が中心で、配信をうまく使えていない",
], body_pt=12, head_pt=16)
card(s, L + cw + gap, y, cw, h1, "料金", [
    "・月額5,000円（税別）",
    "・チャットProオプションを含む",
    "・LINE公式アカウントのプラン代は別途かかる",
    "・初期サポート：20,000円（税別）・1回のみ",
], body_pt=12, head_pt=16)
card(s, L, y + h1 + 0.2, W, h2, "よくある悩み", [
    "・誰に送るかが曖昧：興味のある人・ない人が混ざった配信になっている",
    "・顧客情報が整理されていない：チャット・タグなどに情報が散らばっている",
    "・運用が担当者まかせ：経験や感覚に頼っていて、情報が配信に使える形で貯まっていない",
], body_pt=12, head_pt=16)

# ---------- P.13 導入で変わること・4つの機能 ----------
s = pages[2]
header(s, "【CRMオプション】導入で変わることと4つの機能", [
    "CRMオプションを入れると、一斉配信が中心の運用から、顧客ごとに最適な配信に変えられる。",
    "フォーム・ユーザーリスト・ユーザーリスト配信・オートメーションの4つの機能を、",
    "LINE公式アカウントの管理画面ひとつで使える。",
])
at = clear_body(s)
rowh = [0.34, 0.32, 0.32, 0.32]
th, ch = sum(rowh), 1.75
y = mid(at, th + 0.2 + ch)
table(s, L, y, [1.9, 3.87, 3.87], rowh, [
    ("導入で変わること", "Before", "After"),
    ("配信の方法", "一斉配信が中心", "顧客ごとに最適な配信"),
    ("顧客の理解", "あいまいな分類", "実データで分類"),
    ("データ管理", "手作業", "自動で整理"),
], size=13)
card(s, L, y + th + 0.2, W, ch, "4つの機能", [
    "・フォーム：質問への回答に応じて、自動でタグを付ける・外す",
    "・ユーザーリスト：顧客ごとのニーズ・属性・最後の行動がひと目で分かる",
    "・ユーザーリスト配信：タグ・属性で配信先を絞り込める",
    "・オートメーション：フォームの回答・タグの付与・誕生日をきっかけに自動配信",
], body_pt=13, head_pt=16)

# ---------- P.14 業種別の活用例 ----------
s = pages[3]
header(s, "【CRMオプション】業種別の活用例", [
    "フォームで聞いた内容をタグにして、その後の配信に使う例が、業種別に紹介されている。",
    "教育・習い事、ガソリンスタンド、アパレル・雑貨、スーパー・小売の4業種。",
    "フォームの出し方（QRコード・メッセージ配信・あいさつメッセージ）も業種ごとに違う。",
])
at = clear_body(s)
rowh = [0.36] + [0.62] * 4
y = mid(at, sum(rowh))
table(s, L, y, [1.95, 3.85, 3.84], rowh, [
    ("業種", "フォームで聞くこと", "その後の配信"),
    ("教育・習い事", "体験授業の後に、悩み・伸ばしたい分野を聞く", "同じ悩みの生徒の成功事例を自動配信"),
    ("ガソリンスタンド", "給油時にQRコードで、1か月の給油の頻度を聞く", "回答の翌日に、お礼のクーポンを自動配信"),
    ("アパレル・雑貨", "来店客に、気になったアイテムを聞く", "回答の翌日に、商品の紹介とクーポンを自動配信"),
    ("スーパー・小売", "あいさつメッセージで、好みのカテゴリーを聞く", "後日、好みに合ったお得情報を配信"),
], size=12)

# ---------- 目次 ----------
toc = list(prs.slides)[1]
page_sh = next(x for x in toc.shapes if text(x).startswith("P."))
lines = ["P.11-14" if p.text == "P.11" else p.text for p in page_sh.text_frame.paragraphs]
engine.set_multiline_text(page_sh, lines)

engine.force_font(prs)
prs.save(F)
print("saved", F, len(prs.slides))
