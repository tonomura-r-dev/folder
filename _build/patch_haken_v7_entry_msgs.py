# -*- coding: utf-8 -*-
"""人材業界（派遣・バイト・単発）：入口①〜⑤の14日ステップ5枚を、見本どおり「メッセージの見出し＋中身」で書く（2026-10-22 殿村さん）。

    python -m _build.patch_haken_v7_entry_msgs <殿村さんのPC保存版ver3.pptx> <出力.pptx>

  v6（Dayカード＋検索データ）からの変更：
  ・カードの見出しを、LINEで送るメッセージの見出しにする（見本の「転職で損しない3つのポイント」の書き方）
  ・ねらいの1行をやめて、メッセージに入れる中身を3行で書く
  ・リードの左に紺の札「ステップ配信」を置き、リードを大きくする（見本の「LINEOA施策」の札。殿村さん指定で「ステップ配信」）
  ・下の流れのまとめは入れない（殿村さん）
  入口・登録する方・ゴール・役割・検索データ・ボタンは v6 の PAGES をそのまま使う。
"""
import sys

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

from _build.patch_haken_v5_entry_pages import by_name, c, set_one
from _build.patch_haken_v6_entry_cards import (BAND_GAP, BLACK, COLS, CW, DAYS, GAP, GRAY, INK,
                                               LGRAY, NAVY, NOTE, PAGES, PANEL, W, WHITE, X0)
from _build.patch_tanaka_v5 import copy_slide
from _build.pptx_parts import rect, rich, shape

MSGS = [  # 入口ごと × Day0・2・6・10・14 の（メッセージの見出し, 中身3行）
    [("ご応募ありがとうございます！", ["3問の求人診断で希望を確認", "面接・職場見学までの流れ", "日程の確認・変更はLINEで"]),
     ("面接・見学の前に確認したい3つのこと", ["持ち物（本人確認書類など）", "服装の目安", "会場までの行き方・連絡先"]),
     ("ほかにも、条件に合う求人があります", ["診断の回答に合う求人3件", "時給・エリア・曜日を並べて比較", "気になる求人はそのまま応募"]),
     ("派遣で働く前に知っておきたいこと", ["3年ルールのしくみ", "契約期間と更新の流れ", "よくある質問"]),
     ("今週の新着と、条件の見直し", ["条件に合う新着のまとめ", "「条件を変える」で希望を更新", "次の応募のご案内"])],
    [("30秒で、合うお仕事が分かります", ["3問の求人診断", "回答に合う求人3件", "気になる求人はそのまま応募"]),
     ("派遣会社を選ぶときの3つのポイント", ["求人の数とエリア", "登録から就業までのサポート", "お給料の受け取り方"]),
     ("派遣のしくみ、3分で分かります", ["3年ルールのしくみ", "契約期間と更新の流れ", "この回は求人をお送りしない"]),
     ("応募のあとは、こう進みます", ["応募から就業までの流れ", "持ち物・服装の目安", "日程の変更はLINEで"]),
     ("今週の新着をまとめました", ["条件に合う新着のまとめ", "気になる求人はそのまま応募", "条件の変更もこちらから"])],
    [("このLINEでできること", ["3問の求人診断", "条件に合う求人のお知らせ", "面接・見学の日程の確認"]),
     ("登録から働き始めるまで、サポートします", ["登録から就業までの流れ", "担当者のサポート内容", "よくある質問"]),
     ("派遣で働く前に知っておきたいこと", ["3年ルールのしくみ", "契約期間と更新の流れ", "よくある質問"]),
     ("応募のあとは、こう進みます", ["応募から就業までの流れ", "持ち物・服装の目安", "日程の変更はLINEで"]),
     ("いま人気の求人をまとめました", ["人気の求人のまとめ", "気になる求人はそのまま応募", "条件の変更もこちらから"])],
    [("掲載エリアのお仕事をまとめました", ["3問の求人診断（エリアは掲載地域）", "掲載エリアの求人3件", "気になる求人はそのまま応募"]),
     ("近くの新着が入りました", ["掲載エリアの新着", "時給・曜日を並べて比較", "条件の変更もこちらから"]),
     ("扶養内・掛け持ちで働くときの目安", ["扶養内で働くときの目安", "掛け持ちのときの確定申告", "合う働き方の選び方"]),
     ("応募のあとは、こう進みます", ["応募から就業までの流れ", "持ち物・服装の目安", "日程の変更はLINEで"]),
     ("今週の新着をまとめました", ["条件に合う新着のまとめ", "気になる求人はそのまま応募", "条件の変更もこちらから"])],
    [("次のお仕事、いつ頃から探しますか？", ["次に探す時期を伺う", "希望条件の確認（3問）", "時期に合わせてご案内"]),
     ("探し始めの方に、先にお届けします", ["時期が近い方に合う求人3件", "前回の希望条件をもとに", "気になる求人はそのまま応募"]),
     ("契約を更新しないときの、次の探し方", ["契約満了までの流れ", "次のお仕事の探し方", "よくある質問"]),
     ("前回と近いお仕事の新着です", ["前回と近い職種の新着", "経験を生かせる求人", "条件の変更もこちらから"]),
     ("今週の新着と、条件の見直し", ["条件に合う新着のまとめ", "「条件を変える」で希望を更新", "再応募のご案内"])],
]

BAND_Y, BAND_H = 1.38, 0.56
CARD_Y, CARD_H = 2.06, 4.6
HEAD_H = 0.62
MSG_Y, MSG_H = 0.7, 0.82                        # メッセージの見出し（カード上端から）
BODY_Y, BODY_H = 1.54, 1.5                      # 中身3行
DATA_Y, DATA_H = 3.08, 0.86                     # 検索データの灰色の枠
BTN_H, BTN_B = 0.4, 0.1
PAD = 0.1
TAG = (0.34, 0.74, 1.55, 0.5)                   # 紺の札「ステップ配信」
LEAD_X = 2.02


def build(s, page, msgs):
    title, lead, entry, who, goal, steps = page
    keep = {"TextBox 1", "TextBox 2", "TextBox 41"}
    for sh in list(s.shapes):
        if sh.name not in keep:
            sh._element.getparent().remove(sh._element)
    set_one(by_name(s, "TextBox 1"), title)
    set_one(by_name(s, "TextBox 41"), NOTE, size=10)
    lead_box = by_name(s, "TextBox 2")
    set_one(lead_box, lead, size=14)
    r = lead_box.text_frame.paragraphs[0].runs[0]
    r.font.bold = True
    lead_box.left, lead_box.width = int(c(LEAD_X) / 2.54 * 914400), int(c(10.5 - LEAD_X) / 2.54 * 914400)
    lead_box.top, lead_box.height = int(TAG[1] * 914400), int(TAG[3] * 914400)
    lead_box.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    x, y, w, h = TAG
    rect(s, c(x), c(y), c(w), c(h), NAVY, [("ステップ配信", 14, False, WHITE, 0)])

    bw = (W - BAND_GAP * 2) / 3
    for k, (lab, val) in enumerate([("入口", entry), ("登録する方", who), ("ゴール（CV）", goal)]):
        bx = X0 + (bw + BAND_GAP) * k
        rect(s, c(bx), c(BAND_Y), c(bw), c(BAND_H), WHITE, line=NAVY, lw=1.0)
        rich(s, c(bx + 0.05), c(BAND_Y), c(bw - 0.1), c(BAND_H),
             [[(lab, 10, True, NAVY)], [(val, 12, True, INK)]], align=PP_ALIGN.CENTER)

    for k, ((role, _msg, _aim, data, btn), (head_txt, body)) in enumerate(zip(steps, msgs)):
        cx = X0 + (CW + GAP) * k
        rect(s, c(cx), c(CARD_Y), c(CW), c(CARD_H), WHITE, line=LGRAY, lw=0.75)
        rect(s, c(cx), c(CARD_Y), c(CW), c(HEAD_H), BLACK if k == COLS - 1 else NAVY)
        rich(s, c(cx), c(CARD_Y), c(CW), c(HEAD_H),
             [[(DAYS[k], 16, True, WHITE)], [(role, 10.5, False, WHITE)]], align=PP_ALIGN.CENTER)
        rich(s, c(cx + PAD), c(CARD_Y + MSG_Y), c(CW - PAD * 2), c(MSG_H), [[(head_txt, 13, True, NAVY)]],
             anchor=MSO_ANCHOR.TOP)
        rich(s, c(cx + PAD), c(CARD_Y + BODY_Y), c(CW - PAD * 2), c(BODY_H),
             [[("・" + t, 10.5, False, INK)] for t in body], anchor=MSO_ANCHOR.TOP)
        rect(s, c(cx + PAD), c(CARD_Y + DATA_Y), c(CW - PAD * 2), c(DATA_H), PANEL, line=LGRAY, lw=0.5)
        rich(s, c(cx + PAD + 0.04), c(CARD_Y + DATA_Y + 0.04), c(CW - PAD * 2 - 0.08), c(DATA_H - 0.08),
             [[("［検索データ］", 9, True, GRAY)], [(data, 10, False, INK)]], anchor=MSO_ANCHOR.TOP)
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, c(cx + PAD), c(CARD_Y + CARD_H - BTN_B - BTN_H),
              c(CW - PAD * 2), c(BTN_H), [(btn, 11.5, True, NAVY, 0)], fill=WHITE, line=NAVY, lw=1.25,
              margins=(0.05, 0.02, 0.05, 0.02))


def main(src, out):
    prs = Presentation(src)
    slides = list(prs.slides)
    assert len(slides) == 35
    base = slides[24]
    assert by_name(base, "TextBox 1").text_frame.text == "［第4章］配信設計｜入口ごとの14日ステップ"
    lst = prs.slides._sldIdLst
    anchor = list(lst)[24]
    for page, msgs in zip(PAGES, MSGS):
        new = copy_slide(prs, base)
        e = list(lst)[-1]
        lst.remove(e)
        anchor.addnext(e)
        anchor = e
        build(new, page, msgs)
    assert len(prs.slides) == 40
    prs.save(out)
    print("saved:", out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
