"""
10月号（正）の本文エリア（「■ 詳細」より下のカード・表など）を左右の余白が同じになるよう中央に寄せ、
カードの文字を大きくする（2026-10-21 殿村さん指摘：右端まで寄っていて見づらい）。
使い方: python3 _build/patch_trend_202610_center.py <pptx>
"""
import sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.shapes import MSO_SHAPE

F = sys.argv[1]
prs = Presentation(F)
W = prs.slide_width
FOOTER_TOP = Inches(7.0)

for slide in list(prs.slides)[2:]:
    head = next(sh for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip() == "■ 詳細")
    body_top = head.top + head.height - Inches(0.02)
    body = [sh for sh in slide.shapes if sh is not head and sh.top >= body_top and sh.top < FOOTER_TOP]
    left = min(sh.left for sh in body)
    right = max(sh.left + sh.width for sh in body)
    new_left, new_right = left, W - left
    k = (new_right - new_left) / (right - left)
    for sh in body:
        sh.left = int(new_left + (sh.left - left) * k)
        sh.width = int(sh.width * k)

    for sh in body:
        if sh.shape_type != 1 or not sh.has_text_frame:
            continue
        if sh.auto_shape_type != MSO_SHAPE.ROUNDED_RECTANGLE:
            continue
        paras = sh.text_frame.paragraphs
        if len(paras) != 2 or not paras[0].runs or paras[0].runs[0].font.size != Pt(17):
            continue
        small = sh.height < Inches(1.0)
        for r in paras[0].runs:
            r.font.size = Pt(18 if small else 19)
        for r in paras[1].runs:
            r.font.size = Pt(11 if small else 13)

prs.save(F)
print("saved", F)

# 表は枠（graphicFrame）を縮めても列幅が残るので、列幅も枠に合わせる
prs = Presentation(F)
for slide in list(prs.slides)[2:]:
    for sh in slide.shapes:
        if getattr(sh, "has_table", False):
            cols = sh.table.columns
            tot = sum(c.width for c in cols)
            if tot != sh.width:
                for c in cols:
                    c.width = int(c.width * sh.width / tot)
prs.save(F)
