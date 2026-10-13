# -*- coding: utf-8 -*-
"""展示会 LINE｜前提予算の図解（1枚・4:3）。2026-10-13 「簡潔に、図と線で示してほしい」を受けて作成。
上：今のMeta（店舗ごと・そのまま）／中：LINE広告の枠を全店舗共通で1つ（線で各エリアへ）／下：予約 → サンクスLINE → 来場。
  python3 _build/build_tenjikai_budget_zu.py <Appendix ver1.2.pptx> [出力.pptx]
"""
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Cm

sys.path.insert(0, str(Path(__file__).resolve().parent))
from pptx_parts import (GRAY, GREEN, INK, LGRAY, NAVY, WHITE, arrow_line, find, hline,  # noqa: E402
                        keep_header_only, rich, set_paras, shape, vline)

ROOT = Path(__file__).resolve().parent.parent
SRC = Path(sys.argv[1])
OUT = Path(sys.argv[2]) if len(sys.argv) > 2 else ROOT / "20261013_展示会LINE_前提予算の図解.pptx"

TITLE = "ご提案｜LINE広告の予算枠を、全店舗共通で1つ設ける"
LEAD = "店舗ごとに積まず、1つの枠から各エリアに配信し、申し込んだ方は来場まで追う"
X0, X1 = 1.79, 25.76
W = X1 - X0
GREEN_BG, GREEN_TX = "E8F8EE", "0B7A3B"


def P(t, size, bold=False, color=INK, sa=0):
    return (t, size, bold, color, sa)


def box(s, x, y, w, h, paras, fill=WHITE, line=NAVY, lw=1.0, align=PP_ALIGN.CENTER):
    return shape(s, MSO_SHAPE.RECTANGLE, x, y, w, h, paras, fill=fill, line=line, lw=lw, align=align,
                 anchor=MSO_ANCHOR.MIDDLE, margins=(0.1, 0.05, 0.1, 0.05))


def label(s, x, y, w, txt, color=GRAY, size=11.5, align=PP_ALIGN.LEFT):
    rich(s, x, y, w, 0.6, [[(txt, size, True, color)]], align=align, anchor=MSO_ANCHOR.MIDDLE)


prs = Presentation(str(SRC))
lst = prs.slides._sldIdLst
keep = [i for i, s in enumerate(prs.slides)
        if {"TextBox 1", "TextBox 2", "Connector 3"} <= {sh.name for sh in s.shapes}][:1]
for i, el in reversed(list(enumerate(list(lst)))):
    if i not in keep:
        prs.part.drop_rel(el.rId)
        lst.remove(el)
s = prs.slides[0]
keep_header_only(s)
t1, t2 = find(s, "TextBox 1"), find(s, "TextBox 2")
t1.left, t1.top, t1.width, t1.height = Cm(1.52), Cm(0.38), Cm(18.1), Cm(0.94)
t2.left, t2.top, t2.width, t2.height = Cm(1.79), Cm(2.27), Cm(25.31), Cm(0.96)
set_paras(t1, [(TITLE, 0)])
set_paras(t2, [(LEAD, 0)])
for t in (t1, t2):
    for r in t.text_frame._txBody.iter(qn("a:rPr")):
        r.set("sz", "1600")

# ── 上段：今のMeta（店舗ごと・そのまま）
stores = ["銀座店", "三鷹", "高松", "日向", "ほか"]
n, g = len(stores), 0.45
LW = 4.0                       # 左の見出し列の幅
sw = (W - LW - g * (n - 1)) / n
sx0 = X0 + LW
y_meta = 4.6
label(s, X0, y_meta + 0.35, LW - 0.2, "今のMeta", NAVY, 13)
label(s, X0, y_meta + 0.95, LW - 0.2, "店舗ごと・そのまま", GRAY, 10.5)
cx = []
for i, nm in enumerate(stores):
    x = sx0 + i * (sw + g)
    box(s, x, y_meta, sw, 1.8, [P(nm, 14, True, NAVY), P("Meta広告", 11, False, GRAY)], line=NAVY)
    cx.append(x + sw / 2)

# ── 中段：LINE広告の枠（全店舗共通で1つ）→ 線で各エリアへ
y_bus = y_meta + 2.6          # 各エリアへ伸ばす横線
y_line = y_bus + 1.1           # 枠の上端
for c in cx:
    vline(s, c, y_meta + 1.8, y_bus, GREEN, 2.0, dash=True)
hline(s, cx[0], cx[-1], y_bus, GREEN, 2.5)
mid = (cx[0] + cx[-1]) / 2
vline(s, mid, y_bus, y_line, GREEN, 2.5)
bx, bw = sx0, W - LW
box(s, bx, y_line, bw, 2.4, [P("LINE広告の予算枠（全店舗共通で1つ）", 14, True, WHITE, 4),
                              P("広告費 月30万円 ＋ 運用手数料20%（6万円）＝ 月36万円", 15, True, WHITE)],
    fill=GREEN, line=GREEN)
label(s, X0, y_line + 0.6, LW - 0.2, "追加の予算", GREEN_TX, 13)
label(s, X0, y_line + 1.2, LW - 0.2, "（ご提案）", GREEN_TX, 10.5)

# ── 下段：予約 → サンクスLINE → 来場
y_flow = y_line + 2.4 + 1.3
arrow_line(s, mid, y_line + 2.4 + 0.08, mid, y_flow - 0.08, NAVY, w=2.0)
steps = [("予約", WHITE, NAVY), ("サンクスLINE\n完了画面からLINEへ", WHITE, NAVY),
         ("前日・当日の\nお知らせ", WHITE, NAVY), ("来場", NAVY, WHITE)]
m, fg = len(steps), 0.8
fw = (bw - fg * (m - 1)) / m
for i, (t, fill, tc) in enumerate(steps):
    x = bx + i * (fw + fg)
    lines = t.split("\n")
    paras = [P(lines[0], 13.5, True, tc)] + [P(l, 10.5, False, tc if fill == NAVY else GRAY) for l in lines[1:]]
    box(s, x, y_flow, fw, 1.9, paras, fill=fill, line=NAVY)
    if i < m - 1:
        arrow_line(s, x + fw + 0.1, y_flow + 0.95, x + fw + fg - 0.1, y_flow + 0.95, NAVY, w=2.0)
label(s, X0, y_flow + 0.65, LW - 0.2, "来場まで追う", NAVY, 13)

# 下の一文
rich(s, bx, y_flow + 2.4, bw, 1.3,
     [[("店舗ごとに30万円を足すのではなく、1店舗分（30万円）で全店舗のエリアに配信", 13.5, True, NAVY)],
      [("サンクスLINE：初期15万円・月額3万円〜", 11.5, False, GRAY)]],
     align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE)

# 字体をメイリオに統一（latin → ea → cs の順）
AFTER = {qn(x) for x in ("a:sym", "a:hlinkClick", "a:hlinkMouseOver", "a:rtl", "a:extLst")}
for sl in prs.slides:
    for r in sl.shapes._spTree.iter(qn("a:r")):
        if r.find(qn("a:rPr")) is None:
            r.insert(0, r.makeelement(qn("a:rPr"), {"lang": "ja-JP"}))
    for tag in ("a:rPr", "a:endParaRPr", "a:defRPr"):
        for rpr in sl.shapes._spTree.iter(qn(tag)):
            for ft in ("a:latin", "a:ea", "a:cs"):
                for e in rpr.findall(qn(ft)):
                    rpr.remove(e)
            nxt = next((c for c in rpr if c.tag in AFTER), None)
            for ft in ("a:latin", "a:ea", "a:cs"):
                e = rpr.makeelement(qn(ft), {"typeface": "メイリオ"})
                if nxt is not None:
                    nxt.addprevious(e)
                else:
                    rpr.append(e)

prs.save(str(OUT))
print("saved:", OUT.name)
