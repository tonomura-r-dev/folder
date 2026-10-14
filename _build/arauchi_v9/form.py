# -*- coding: utf-8 -*-
"""あらうち提案資料 v9 の直し（2026-10-14 殿村さん指示）：CVを「新規作成の1つのフォーム」に一本化する。

    python -m _build.arauchi_v9.form <直す前のv9.pptx> <出力.pptx>

  直す前のv9 ＝ commit 506784f のv9（殿村さんがPowerPointで保存したv8(2)に、build.py でS15・S16を直したもの）。
  殿村さんの指示：「v9を修正して出してほしい」。版は上げない（v10は「ちゃんと修正する」ときだけ）。出すのはv9だけ。

決まっていること（殿村さん）
  ・CVは1つのフォームの送信だけ（新規作成。チャットでURLを送る形のままでは計測できないため）。新規CV・リピートCVの区別はしない
  ・「フォームで計測をする」（S4。「チャットで受けているので計測できない」とは書かない）／現状の言い方は「厚くできる」（S2・S4）
  ・ツールは離脱防止のみ（Lステップ・Lメッセージは入れない）。標準機能で運用
  ・件数は殿村さんのSIM（2026-10-14 14:14保存）：CV＝5,7,8,10,12,13／累計55件。**数値感はまだ直し中なので、ここ（CV）だけ差し替えれば全ページに反映される**
  ・既存の友だち（1,241人）経由の問い合わせは試算に入れず「上積み」（SIMにはまだ入れていない）
  ・CVRは触らない（殿村さん：あらうちはこれ以上言われても面倒なので一旦スルー）
  ・フォームの作成・設置の費用の扱いは未確定 → 資料には書かない（初期設計項目にも入れない）

変えるページ（v9の番号）
  S2 現状把握／S3 想定動線／S4 LINEアカウントの現状／S8 想定動線／S9 想定セグメント×成果／S11 要件定義（課題）／
  S12 要件定義（配信目標・棒グラフを1系列に）／S13 想定の費用対効果／S16 運用内容の比較（注記の件数）／S17 実施施策
"""
import sys
import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn

from _build.arauchi_v4.build import c14n, normalize_decl, rel_set, slide_order
from _build.arauchi_v4.lib import TEXT, add_box, delete, move, set_paras
from _build.arauchi_v8.pages import label, replace_in_paragraphs, shape_by_id

# ---------------------------------------------------------------- 数字の正本（殿村さんのSIM・2026-10-14 14:14保存。CV＝①の「物件問い合わせ」だけ）
CV = [5, 7, 8, 10, 12, 13]
TOTAL = sum(CV)
assert TOTAL == 55 and CV[0] == 5 and CV[-1] == 13

BAR = "1F497D"          # S12の見出しの紺
RED_TXT = "FF3300"      # S12の目標の赤
CHANGED = [2, 3, 4, 8, 9, 11, 12, 13, 16, 17]


def one(s, old, new, n=1):
    got = replace_in_paragraphs(s, old, new)
    assert got == n, (old, got)


def keep_one_run(shape, para_idx, run_idx, text):
    """段落の run_idx 番目の書式のまま全文を入れ、ほかの run を消す"""
    p = shape.text_frame._txBody.findall(qn("a:p"))[para_idx]
    runs = p.findall(qn("a:r"))
    runs[run_idx].find(qn("a:t")).text = text
    for k, r in enumerate(runs):
        if k != run_idx:
            p.remove(r)


# ---------------------------------------------------------------- S2 現状把握：「改善の余地」→「厚くできる」
def p02(s):
    one(s, "改善の余地があると想定", "標準機能のまま厚くできると想定")
    one(s, "改善の余地あり。", "標準機能で厚くできる余地あり。")


# ---------------------------------------------------------------- S3 想定動線：フォームで送信
def p03(s):
    one(s, "ポータルサイトで物件検索", "ポータルで物件検索")
    one(s, "→ 物件情報の送信（CV）の流れ", "→ フォームで物件を送信（CV）の流れ")


# ---------------------------------------------------------------- S4 LINEアカウントの現状：「フォームで計測をする」
def p04(s):
    hd = shape_by_id(s, 137)
    # 2行目：太字の run だけを残して全文を差し替える（「一方で、…がなく、…未設置です」→ 厚くする＋フォームで計測）
    keep_one_run(hd, 1, 1, "送るメリット・安心材料やボタン、リッチメニューを足して厚くし、フォームで計測をします。")


# ---------------------------------------------------------------- S8 想定動線
def p08(s):
    one(s, "物件情報の送信（CV）につなげます", "フォーム送信（CV）につなげます")
    set_paras(shape_by_id(s, 185), ["フォームで送信", "（CV）"])


# ---------------------------------------------------------------- S9 想定セグメント×成果
def p09(s):
    one(s, "物件の送り方を案内", "フォームの送り方を案内")
    one(s, "→ 物件URLの送信", "→ フォームで物件を送信")
    one(s, "→ 物件の送信（CV）", "→ フォーム送信（CV）")
    one(s, "→ 物件の送信", "→ フォームで物件を送信", n=2)
    # 右の列：6か月目の問い合わせ（新規友だち経由のみ。既存友だち経由は上積み＝試算に入れない）
    set_paras(shape_by_id(s, 272), [f"約{CV[-1]}件", "新規友だち経由", "（フォーム送信）"])
    set_paras(shape_by_id(s, 273), ["上積み", "既存友だち経由", "（試算には含みません）"])
    hd = shape_by_id(s, 293)
    set_paras(hd, [
        "新しい友だちには物件の送り方を、今の友だちには配信とリッチメニューで問い合わせのきっかけを届けます。",
        "今の友だちは見込み顧客として活用し、上積みの問い合わせや将来の紹介につなげます（試算に含みません）。",
        "※問い合わせ＝フォームでの物件情報の送信（CV）の件数。新規友だち経由の月間の試算",
    ])


# ---------------------------------------------------------------- S11 課題
def p11(s):
    one(s, "「物件を送る」まで案内します", "「フォームで物件を送る」まで案内します")


# ---------------------------------------------------------------- S12 要件定義：配信目標・棒グラフ（1系列）
def p12(s):
    set_paras(shape_by_id(s, 325), [f"1か月目 月{CV[0]}件 → 6か月目 月{CV[-1]}件（試算）"])
    set_paras(shape_by_id(s, 331), ["フォームでの送り方に加え、営業電話なし・確認だけでも可などの送る理由と、フォームへのボタンを追加"])
    groups = [sh for sh in s.shapes if sh.shape_type == 6 and sh.name.startswith("グラフ（問い合わせ件数")]
    assert len(groups) == 1
    delete(groups[0])
    x0, pitch, bar_w, base_y, scale = 4.18, 0.70, 0.44, 2.52, 0.55 / max(CV)         # 最大の月＝0.55in（これまでの33件の棒と同じ高さ）
    before = {sh.shape_id for sh in s.shapes}
    add_box(s, x0, base_y, pitch * 6, 0.012, [[("", {})]], fill="A6A6A6", line=None, name="グラフ_基準線")
    for i, v in enumerate(CV):
        cx = x0 + pitch * i + pitch / 2
        h = v * scale
        add_box(s, cx - bar_w / 2, base_y - h, bar_w, h, [[("", {})]], fill=BAR, line=None, name=f"グラフ_問い合わせ_{i + 1}か月目")
        edge = i in (0, len(CV) - 1)
        label(s, cx - pitch / 2, base_y - h - 0.32, pitch, 0.30, f"{v}件" if edge else str(v), 14 if edge else 12,
              bold=True, color=RED_TXT if edge else TEXT, name=f"グラフ_件数_{i + 1}か月目")
        label(s, cx - pitch / 2, base_y + 0.03, pitch, 0.21, f"{i + 1}か月目", 10, color="595959", name=f"グラフ_月_{i + 1}")
    add_box(s, 8.58, 1.76, 0.14, 0.14, [[("", {})]], fill=BAR, line=None, name="グラフ_凡例")
    label(s, 8.78, 1.70, 1.45, 0.48, ["問い合わせ", "（フォーム送信）"], 11, color=TEXT, align="l", anchor="t", name="グラフ_凡例文字")
    label(s, 8.58, 2.46, 1.65, 0.47, ["6か月累計", f"{TOTAL}件（試算）"], 11, bold=True, color=TEXT, align="l", anchor="t", name="グラフ_累計")
    grp = s.shapes.add_group_shape(shapes=[sh for sh in s.shapes if sh.shape_id not in before])
    grp.name = "グラフ（問い合わせ件数・月間の試算）"


# ---------------------------------------------------------------- S13 想定の費用対効果
def p13(s):
    one(s, "月33件を見込みます", f"月{CV[-1]}件を見込みます")
    set_paras(shape_by_id(s, 362), ["問い合わせ"])
    set_paras(shape_by_id(s, 376), [f"{CV[-1]}件/月"])
    for sid in (363, 377, 379):                                     # 問い合わせ②・その件数・「（6か月累計126件）」
        delete(shape_by_id(s, sid))
    lab, val = shape_by_id(s, 364), shape_by_id(s, 378)             # 「合計 33件/月」→「6か月累計 55件」（②の段へ詰める）
    set_paras(lab, ["6か月累計"])
    set_paras(val, [f"{TOTAL}件"])
    move(lab, y=3.18)
    move(val, y=3.18)
    box = shape_by_id(s, 382)                                       # 試算の前提：既存の友だち経由は上積み（試算外）を足す。箱は空いた分だけ上へ
    set_paras(box, ["＜試算の前提＞",
                    "プラン①｜総合運用の基本施策を対象とした想定値です。",
                    "API連携・追加開発、紹介による追加成果は含みません。",
                    "既存の友だち経由は上積み（試算外）。",
                    "成果を保証するものではありません。"])
    move(box, y=4.12, h=2.10)
    one(s, "問い合わせ①＝友だち追加直後の物件の問い合わせ、②＝ステップ配信・企画配信などからの問い合わせ",
        "問い合わせ＝フォームでの物件情報の送信（CV）の件数")


# ---------------------------------------------------------------- S16 運用内容の比較（注記の件数）
def p16(s):
    one(s, "6か月目に月33件", f"6か月目に月{CV[-1]}件")


# ---------------------------------------------------------------- S17 実施施策
def p17(s):
    one(s, "物件URLの送信を後押しします", "フォームでの物件URLの送信を後押しします")
    set_paras(shape_by_id(s, 395), ["フォームで", "物件URLの送信"])
    set_paras(shape_by_id(s, 423), ["フォームでの送り方"])
    for sid in (425, 443):
        set_paras(shape_by_id(s, sid), ["フォームで", "物件の送信"])


APPLY = {2: p02, 3: p03, 4: p04, 8: p08, 9: p09, 11: p11, 12: p12, 13: p13, 16: p16, 17: p17}
assert sorted(APPLY) == CHANGED


# ---------------------------------------------------------------- 差し込み
def repack(base_path, built_path, out_path):
    """直したスライドのXMLだけを、base（PowerPointで保存した元のzip）に差し込む。ほかのパーツは1バイトも変えない"""
    zb, zn = zipfile.ZipFile(base_path), zipfile.ZipFile(built_path)
    base_order, built_order = slide_order(zb), slide_order(zn)
    assert len(base_order) == len(built_order) == 47, (len(base_order), len(built_order))
    replaced = {}
    for i, (b_name, n_name) in enumerate(zip(base_order, built_order), 1):
        if c14n(zb.read(b_name)) != c14n(zn.read(n_name)):
            replaced[b_name] = normalize_decl(zn.read(n_name))
        b_rels = b_name.replace("slides/", "slides/_rels/") + ".rels"
        n_rels = n_name.replace("slides/", "slides/_rels/") + ".rels"
        if rel_set(zb.read(b_rels)) != rel_set(zn.read(n_rels)):
            raise SystemExit(f"想定外のrelsの違い: {b_rels}")
    changed_nos = sorted(i for i, n in enumerate(base_order, 1) if n in replaced)
    assert changed_nos == CHANGED, f"変わったスライド {changed_nos} != {CHANGED}"
    with zipfile.ZipFile(out_path, "w") as zo:
        for info in zb.infolist():
            data = replaced.pop(info.filename) if info.filename in replaced else zb.read(info.filename)
            zo.writestr(info, data, compress_type=info.compress_type)
    assert not replaced
    return changed_nos


def main(base, out):
    tmp = out + ".tmp.pptx"
    prs = Presentation(base)
    for no, fn in APPLY.items():
        fn(prs.slides[no - 1])
    prs.save(tmp)
    changed = repack(base, tmp, out)
    Path(tmp).unlink()
    print("saved:", out)
    print("差し替えたスライド（v9の番号）:", changed)


if __name__ == "__main__":
    main(*sys.argv[1:3])
