# -*- coding: utf-8 -*-
"""あらうち提案資料 v9 の情報量を減らす（2026-10-15 殿村さん指示）。

    python -m _build.arauchi_v9.trim <殿村さんがPowerPointで直したv9.pptx> <出力.pptx>

  ベース＝殿村さんがPowerPointで直したv9（2026-10-14 17時に受領・commit f5853fd のv9）。
  殿村さん：「汎用部分以外のところの情報量が少し多い。削れるところ削っていきましょう」
  → 削る候補13項目を出し、「2を左下だけ残して、それ以外は修正して」。版は上げない（v9のまま・v9だけ出す）。

  1  S2   表の「--」4行（投稿数・ステップ・アンケート・想定コスト）を削除。広告の4行は「広告：なし」の1行に
  2  S3   「⑤最重要課題」の印は左下（ユーザー動線の⑤タブ）だけ残す：「◀ 最重要課題（仮説）」を削除、⑤の赤い箱は①〜④と同じ色に
  3  S5＋S6 2枚を1枚に（6項目→4項目）。外す＝S5の③物件数で選ぶ・S6の②オンライン契約。S6はページごと削除
  4  S7   結論を1行に（API連携はS16（新14）の表に、離脱防止の注記はS8・S12・S19にある）
  5  S8   下段「参考：CV以降の個別対応」の6ブロックを削除（定期対応項目は残す）
  6  S9   下の4行（問い合わせた方・受付時間外・迷っている方・他社で内見済み）を削除
  7  S10  ページごと削除（汎用の考え方のページ）
  8  S11  02の「電話33.4%…」と、結論の「最重要課題は…（⑤リード有効化）です。」（上の赤い見出しと同じ）を削除
  9  S12  「手法」の説明を各1行に
  10 S13  試算の前提を「プラン①の基本施策での想定値です。成果を保証するものではありません。」に
  11 S16  「月額費用」の行と、注記2つ（API・プラン③）を削除。残る注記の「13ページ」→「11ページ」（S6・S10を抜いたため）
  12 S17  ステップ配信の「ターゲット層」「セグメント方法」の列を削除（見出しに「新しい友だちへ追加後に自動配信」）
  13 S18  下の「役割分担」の枠を削除

  ページ番号の参照は S16 の「13ページ」だけ（ほかに「◯ページ」の参照は無い）。
"""
import copy
import re
import sys
import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches

from _build.arauchi_v4.build import c14n, normalize_decl, rel_set, slide_order
from _build.arauchi_v4.lib import delete, move, set_paras
from _build.arauchi_v8.pages import replace_in_paragraphs, shape_by_id

N_BASE = 47
DROP = [6, 10]                                   # ページごと削除（v9の番号）
CHANGED = [2, 3, 5, 7, 8, 9, 11, 12, 13, 16, 17, 18]
NEW_PAGE_OF_S13 = 13 - sum(1 for d in DROP if d < 13)
assert NEW_PAGE_OF_S13 == 11


def one(s, old, new, n=1):
    got = replace_in_paragraphs(s, old, new)
    assert got == n, (old, got)


def inch(v):
    return Emu(v).inches


def scale_center(sh, y0, k, hk=None):
    """y0 を基準に縦方向へ k 倍に広げる（中心の位置を k 倍・高さを hk 倍。hk 省略時は k）"""
    hk = k if hk is None else hk
    top, h = inch(sh.top), inch(sh.height)
    c = y0 + (top + h / 2 - y0) * k
    nh = h * hk
    sh.top, sh.height = Inches(c - nh / 2), Inches(nh)


def paras(sh):
    return sh.text_frame._txBody.findall(qn("a:p"))


def drop_paras(sh, idxs):
    ps = paras(sh)
    for i in sorted(idxs, reverse=True):
        ps[i].getparent().remove(ps[i])


def set_sizes(sh, old, new):
    for tag in ("a:rPr", "a:endParaRPr"):
        for r in sh._element.iter(qn(tag)):
            if r.get("sz") == str(old):
                r.set("sz", str(new))


def next_id(slide):
    return max(int(e.get("id")) for e in slide._element.iter(qn("p:cNvPr"))) + 1


def clone_shape(dst, src_shape):
    el = copy.deepcopy(src_shape._element)
    el.find(".//" + qn("p:cNvPr")).set("id", str(next_id(dst)))
    dst.shapes._spTree.append(el)
    return shape_by_id(dst, int(el.find(".//" + qn("p:cNvPr")).get("id")))


# ---------------------------------------------------------------- 1 S2 現状把握：表の行を減らす
def p02(s):
    # WEBサイトステータス：広告4行 → 「広告：なし」の1行
    set_paras(shape_by_id(s, 44), ["広告"])
    set_paras(shape_by_id(s, 64), ["なし"])
    for sid in (45, 65, 52, 72, 46, 66, 59, 80, 60, 81, 61, 82, 58, 79):
        delete(shape_by_id(s, sid))
    left = [(42, 63), (44, 64), (49, 69), (50, 70), (51, 71), (48, 68), (47, 67)]      # サイトUU／広告／月間CV数／CV地点／LINE動線／リード保有数／他SNS
    right = [(53, 74), (54, 75), (55, 76), (62, 83), (57, 78), (56, 77), (73, 84)]      # 友だち数／アクティブ数／ブロック率／プロフ動線／ツール／リッチメニュー／あいさつ
    y0, pitch, h = 2.10, 0.46, 0.40
    for rows in (left, right):
        for i, pair in enumerate(rows):
            for sid in pair:
                move(shape_by_id(s, sid), y=y0 + pitch * i, h=h)
    move(shape_by_id(s, 43), y=y0 + pitch * 6 + h + 0.10)                               # ※いずれも貴社ヒヤリング…


# ---------------------------------------------------------------- 2 S3：⑤最重要課題の印は左下だけ
def p03(s):
    delete(shape_by_id(s, 122))                          # ◀ 最重要課題（仮説）
    five, four = shape_by_id(s, 93), shape_by_id(s, 100)  # ⑤の箱を④と同じ色・文字に
    sp5, sp4 = five._element.spPr, four._element.spPr
    sp5.replace(sp5.find(qn("a:solidFill")), copy.deepcopy(sp4.find(qn("a:solidFill"))))
    for r in five._element.iter(qn("a:rPr")):
        r.set("sz", "1200")
        f = r.find(qn("a:solidFill"))
        f.getparent().replace(f, copy.deepcopy(four._element.find(".//" + qn("a:rPr")).find(qn("a:solidFill"))))


# ---------------------------------------------------------------- 3 S5＋S6：市場分析を1枚（4項目）に
CARD3_BLUE = "5B7DB1"


def p05(s, s6):
    # 見出し：1行目はS5のまま、2行目はS6の2行目（太字）
    lead5, lead6 = shape_by_id(s, 142), shape_by_id(s6, 157)
    p5, p6 = paras(lead5), paras(lead6)
    assert len(p5) == len(p6) == 2
    p5[1].getparent().replace(p5[1], copy.deepcopy(p6[1]))
    # S5の③（物件数で選ぶ）を外し、S6の①（文字でのやり取り）・③（URLを送るだけの競合）を足す
    for sid in (149, 150, 151):
        delete(shape_by_id(s, sid))
    c3 = [clone_shape(s, shape_by_id(s6, sid)) for sid in (158, 159, 160)]
    c4 = [clone_shape(s, shape_by_id(s6, sid)) for sid in (164, 165, 166)]
    for sh in c3:                                        # 03は青を一段明るく（01紺→02青→03明るい青→04灰）
        for c in sh._element.iter(qn("a:srgbClr")):
            if c.get("val") == "1F2A5C":
                c.set("val", CARD3_BLUE)
    set_paras(c3[0], ["03"])
    set_paras(c4[0], ["04"])
    set_paras(shape_by_id(s, 147), ["ポータルサイトで探す"])         # 物件はポータルサイトで探す（4列で1行に収める）
    set_paras(c3[1], ["文字でのやり取りを希望"])                       # 電話より、文字でのやり取り
    set_paras(c4[1], ["URLを送るだけの競合"])                          # URLを送るだけの競合も
    cards = [[shape_by_id(s, i) for i in (143, 144, 145)], [shape_by_id(s, i) for i in (146, 147, 148)], c3, c4]
    w, gap = 2.385, 0.15
    for i, (num, title, body) in enumerate(cards):
        x = 0.42 + (w + gap) * i
        move(num, x=x, y=1.62, w=w, h=0.56)
        move(title, x=x, y=2.18, w=w, h=0.50)
        move(body, x=x, y=2.68, w=w, h=3.92)
        for sh, ins in ((title, 0.08), (body, 0.12)):
            bp = sh.text_frame._txBody.find(qn("a:bodyPr"))
            bp.set("lIns", str(Inches(ins)))
            bp.set("rIns", str(Inches(ins)))


# ---------------------------------------------------------------- 4 S7：結論を1行に
def p07(s):
    c = shape_by_id(s, 176)
    ps = paras(c)
    assert "".join(t.text for t in ps[1].iter(qn("a:t"))).startswith("基本施策は")
    drop_paras(c, range(2, len(ps)))
    move(c, h=0.85)


# ---------------------------------------------------------------- 5 S8：下段の「参考：CV以降の個別対応」を削除
def p08(s):
    for sid in [194, 217, 218] + list(range(220, 231)):
        delete(shape_by_id(s, sid))
    # 上段（対新規ユーザー領域）を縦に広げて、空いた分を使う
    y0, k = 1.92, (5.40 - 1.92) / 2.66
    top_ids = [180, 184, 185, 186, 187, 188, 233] + list(range(195, 217))
    for sid in top_ids:
        scale_center(shape_by_id(s, sid), y0, k)
    # 下段：見出しを「定期対応項目」に、5項目を3列に並べる
    lab = shape_by_id(s, 183)
    set_paras(lab, ["<定期対応項目（一部抽出）>"])
    move(lab, y=5.48)
    move(shape_by_id(s, 181), y=5.84, h=1.16)
    w, gap = 2.82, 0.25
    x0 = 0.50 + (10.05 - (3 * w + 2 * gap)) / 2
    for i, sid in enumerate((189, 190, 192, 191, 193)):
        move(shape_by_id(s, sid), x=x0 + (w + gap) * (i % 3), y=6.01 + 0.46 * (i // 3), w=w, h=0.36)


# ---------------------------------------------------------------- 6 S9：下の4行を削除し、上の2段を広げる
def p09(s):
    for sid in [246, 247, 249, 250, 252, 253, 254, 255] + list(range(256, 265)) + [274, 275, 276, 277] + list(range(281, 293)):
        delete(shape_by_id(s, sid))
    y0, k = 2.18, 2.30
    for sid in (239, 240, 241, 242, 244, 245, 248, 251):          # 段（背景・矢印・対象・LINE誘導）はそのまま k 倍
        scale_center(shape_by_id(s, sid), y0, k)
    for sid in (265, 266, 268, 269, 278, 279):                    # 施策の箱は高さを控えめに
        scale_center(shape_by_id(s, sid), y0, k, hk=1.6)
    for sid in (267, 270, 280, 272):                              # 矢印の線・合計の文字は位置だけ
        scale_center(shape_by_id(s, sid), y0, k, hk=1.0)
    # 施策の箱を少し広げて、文字を11→12ptに
    for a, arrow, b in ((265, 267, 266), (268, 270, 269), (278, 280, 279)):
        move(shape_by_id(s, a), x=3.80, w=2.05)
        move(shape_by_id(s, arrow), x=5.85, w=0.30)
        move(shape_by_id(s, b), x=6.15, w=2.15)
    for sid in (242, 248, 265, 266, 268, 269, 278, 279):         # LINE誘導の2つ（幅1.56in）は11ptのまま
        set_sizes(shape_by_id(s, sid), 1100, 1200)


# ---------------------------------------------------------------- 8 S11
def p11(s):
    body = shape_by_id(s, 311)
    ps = paras(body)
    assert "33.4" in "".join(t.text for t in ps[1].iter(qn("a:t")))
    drop_paras(body, [1])
    one(s, "最重要課題は、友だち追加後の問い合わせ獲得（⑤リード有効化）です。", "")


# ---------------------------------------------------------------- 9 S12：手法の説明を1行に
def p12(s):
    set_paras(shape_by_id(s, 330), ["サイトを離れようとした方に、LINEでの問い合わせを案内"])
    set_paras(shape_by_id(s, 331), ["送る理由（営業電話なし等）とフォームへのボタンを追加"])
    set_paras(shape_by_id(s, 332), ["「物件を送る」を常に表示。よくある質問は自動で回答"])
    set_paras(shape_by_id(s, 333), ["14日間・全10通と月4回程度の配信で、送信を後押し"])


# ---------------------------------------------------------------- 10 S13：試算の前提
def p13(s):
    b = shape_by_id(s, 382)
    ps = paras(b)
    texts = ["".join(t.text for t in p.iter(qn("a:t"))) for p in ps]
    assert texts[2].startswith("API連携") and texts[3].startswith("成果を保証"), texts
    one(s, "プラン①｜総合運用の基本施策を対象とした想定値です。", "プラン①の基本施策での想定値です。")
    drop_paras(b, [2])
    move(b, h=1.20)


# ---------------------------------------------------------------- 11 S16：月額費用の行と注記2つ
def p16(s):
    gf = shape_by_id(s, 6)
    tbl = gf.table._tbl
    tr = tbl.tr_lst[1]
    assert "月額費用" in "".join(t.text for t in tr.iter(qn("a:t")))
    tbl.remove(tr)
    for r in tbl.tr_lst:                                          # 空いた分、行を1.2倍に
        r.set("h", str(int(int(r.get("h")) * 1.2)))
    gf.height = sum(int(r.get("h")) for r in tbl.tr_lst)
    note = shape_by_id(s, 7)
    ps = paras(note)
    texts = ["".join(t.text for t in p.iter(qn("a:t"))) for p in ps]
    assert texts[0].startswith("※API") and "13ページ" in texts[1] and texts[2].startswith("※") and "プラン③" in texts[2], texts
    drop_paras(note, [0, 2])
    one(s, "13ページ", f"{NEW_PAGE_OF_S13}ページ")
    move(note, y=inch(gf.top + gf.height) + 0.14, h=0.40)


# ---------------------------------------------------------------- 12 S17：ステップ配信の2列を削除
def p17(s):
    rows = [(388, 389, 390, 391, 392, 393, 394, 395, 396, 397),
            (398, 399, 400, 401, 402, 403, 404, 405, 406, 407),
            (408, 409, 410, 411, 412, 413, 414, 415, 416, 417)]
    for day, tgt, a1, seg, a2, cont, a3, aim, a4, cta in rows:
        for sid in (tgt, seg, a2):
            delete(shape_by_id(s, sid))
        move(shape_by_id(s, day), x=0.42, w=1.90)                 # 企画配信の列にそろえる
        move(shape_by_id(s, a1), x=2.32)
        move(shape_by_id(s, cont), x=2.58, w=2.98)
        move(shape_by_id(s, a3), x=5.56)
        move(shape_by_id(s, aim), x=5.82, w=2.72)
        move(shape_by_id(s, a4), x=8.54)
        move(shape_by_id(s, cta), x=8.80, w=1.60)
    for sid in (455, 456):
        delete(shape_by_id(s, sid))
    move(shape_by_id(s, 457), x=2.58, w=2.98)
    move(shape_by_id(s, 458), x=5.82, w=2.72)
    one(s, "ステップ配信・全10通", "ステップ配信・全10通｜新しい友だちへ追加後に自動配信")
    move(shape_by_id(s, 387), w=9.98)


# ---------------------------------------------------------------- 13 S18：役割分担の枠を削除
def p18(s):
    for sid in (471, 472, 473, 474, 475, 478):
        delete(shape_by_id(s, sid))
    bottom = 6.60
    for sid in (466, 468, 470):                                   # カードを下へ伸ばし、文字を14→15ptに
        sh = shape_by_id(s, sid)
        move(sh, h=bottom - inch(sh.top))
        set_sizes(sh, 1400, 1500)
    mid = (2.12 + bottom) / 2
    for sid in (476, 477):
        sh = shape_by_id(s, sid)
        move(sh, y=mid - inch(sh.height) / 2)


def apply(prs):
    sl = prs.slides
    p02(sl[1])
    p03(sl[2])
    p05(sl[4], sl[5])
    p07(sl[6])
    p08(sl[7])
    p09(sl[8])
    p11(sl[10])
    p12(sl[11])
    p13(sl[12])
    p16(sl[15])
    p17(sl[16])
    p18(sl[17])


# ---------------------------------------------------------------- 差し込み（S6・S10はパーツごと外す）
def repack(base_path, built_path, out_path):
    zb, zn = zipfile.ZipFile(base_path), zipfile.ZipFile(built_path)
    base_order, built_order = slide_order(zb), slide_order(zn)
    assert len(base_order) == len(built_order) == N_BASE, (len(base_order), len(built_order))
    drops = [base_order[d - 1] for d in DROP]
    gone = set()
    for drop in drops:
        drop_rels = drop.replace("slides/", "slides/_rels/") + ".rels"
        gone |= {drop, drop_rels}
        for r in etree.fromstring(zb.read(drop_rels)):
            assert r.get("Type").endswith("/slideLayout"), r.get("Type")      # ノート・画像・リンクなし
    for name in zb.namelist():
        if name.endswith(".rels") and name not in gone and name != "ppt/_rels/presentation.xml.rels":
            data = zb.read(name).decode("utf-8")
            for drop in drops:
                assert 'Target="../slides/%s"' % drop.split("/")[-1] not in data, name
                assert 'Target="%s"' % drop.split("/")[-1] not in data, name

    prels = zb.read("ppt/_rels/presentation.xml.rels").decode("utf-8")
    pres = zb.read("ppt/presentation.xml").decode("utf-8")
    ct = zb.read("[Content_Types].xml").decode("utf-8")
    for drop in drops:
        m = re.search(r'<Relationship [^>]*Target="slides/%s"[^>]*/>' % re.escape(drop.split("/")[-1]), prels)
        rid = re.search(r'Id="(rId\d+)"', m.group(0)).group(1)
        prels = prels.replace(m.group(0), "")
        m2 = re.search(r'<p:sldId [^>]*r:id="%s"[^>]*/>' % rid, pres)
        pres = pres.replace(m2.group(0), "")
        ct, n = re.subn(r'<Override [^>]*PartName="/%s"[^>]*/>' % re.escape(drop), "", ct)
        assert n == 1, drop

    app = zb.read("docProps/app.xml").decode("utf-8")
    nb, nn = N_BASE, N_BASE - len(DROP)
    assert f"<Slides>{nb}</Slides>" in app
    app = app.replace(f"<Slides>{nb}</Slides>", f"<Slides>{nn}</Slides>")
    head = f"<vt:lpstr>スライド タイトル</vt:lpstr></vt:variant><vt:variant><vt:i4>{nb}</vt:i4>"
    assert head in app
    app = app.replace(head, head.replace(f"<vt:i4>{nb}</vt:i4>", f"<vt:i4>{nn}</vt:i4>"))
    m = re.search(r'<vt:vector size="(\d+)" baseType="lpstr">', app)
    size = int(m.group(1))
    app = app.replace(m.group(0), f'<vt:vector size="{size - len(DROP)}" baseType="lpstr">')
    for _ in DROP:
        assert "<vt:lpstr>PowerPoint プレゼンテーション</vt:lpstr>" in app
        app = app.replace("<vt:lpstr>PowerPoint プレゼンテーション</vt:lpstr>", "", 1)

    replaced, changed_nos = {}, []
    for i, (b_name, n_name) in enumerate(zip(base_order, built_order), 1):
        if i in DROP:
            continue
        if c14n(zb.read(b_name)) != c14n(zn.read(n_name)):
            replaced[b_name] = normalize_decl(zn.read(n_name))
            changed_nos.append(i)
        b_rels = b_name.replace("slides/", "slides/_rels/") + ".rels"
        n_rels = n_name.replace("slides/", "slides/_rels/") + ".rels"
        if rel_set(zb.read(b_rels)) != rel_set(zn.read(n_rels)):
            raise SystemExit(f"想定外のrelsの違い: {b_rels}")
    assert changed_nos == CHANGED, f"変わったスライド {changed_nos} != {CHANGED}"

    with zipfile.ZipFile(out_path, "w") as zo:
        for info in zb.infolist():
            name = info.filename
            if name in gone:
                continue
            if name == "ppt/presentation.xml":
                data = pres.encode("utf-8")
            elif name == "ppt/_rels/presentation.xml.rels":
                data = prels.encode("utf-8")
            elif name == "[Content_Types].xml":
                data = ct.encode("utf-8")
            elif name == "docProps/app.xml":
                data = app.encode("utf-8")
            elif name in replaced:
                data = replaced.pop(name)
            else:
                data = zb.read(name)
            zo.writestr(info, data, compress_type=info.compress_type)
    assert not replaced
    return changed_nos, sorted(gone)


def main(base, out):
    tmp = out + ".tmp.pptx"
    prs = Presentation(base)
    apply(prs)
    prs.save(tmp)
    changed, gone = repack(base, tmp, out)
    Path(tmp).unlink()
    print("saved:", out)
    print("差し替えたスライド（v9の番号）:", changed)
    print("外したパーツ:", gone)


if __name__ == "__main__":
    main(*sys.argv[1:3])
