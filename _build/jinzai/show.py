# -*- coding: utf-8 -*-
"""人材（派遣・バイト・単発）：非表示のうち6枚を表示して流れに入れ、残りの非表示5枚を外す（2026-10-20 殿村さん）。

    python -m _build.jinzai.show <殿村さん保存版（2026-10-20）.pptx> <出力.pptx>

  殿村さん：「30,31,32,33,34,39は入れたいかも」（38万円→28万円は殿村さんが直し済み）
  → 30・31（全体設計）は第2章の全体設計（6）のあと、34（他社分析）は求職者の心理（9）のあと、
     39・32（施策展開図）は第3章の先頭（タウンワークの型のあと）、33（アクティブ率）は第4章の先頭に置く。
  → 35（他社分析②）・36（ブレスト）・37（転職の診断）・38（具体施策）・40（改善モデル①）はパーツごと外す。
  → 見出しに［第N章］を付け、本文のページ参照（P10・P11・P16）を新しい番号に付け替える。
     39の他業界の言葉（学習・再購入・クロスセル）は人材の言葉に。
  ほかのページは1バイトも変えない（殿村さんの保存版のzipに、直したスライドのXMLだけ差し込む）。
"""
import re
import sys
import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation

import _build.juku_fmt.build as jb
from _build.arauchi_v4.build import slide_order
from _build.arauchi_v8.pages import replace_in_paragraphs

N_BASE = 40
SHOW = [30, 31, 32, 33, 34, 39]
DROP = [35, 36, 37, 38, 40]
ORDER = [1, 2, 3, 4, 5, 6, 30, 31, 7, 8, 9, 34, 10, 11, 12, 39, 32, 13, 14, 15, 16, 33,
         17, 18, 19, 20, 21, 22, 23, 24, 25, 26, 27, 28, 29]
NEW_POS = {u: i for i, u in enumerate(ORDER, 1)}
PKG_RELS = "{http://schemas.openxmlformats.org/package/2006/relationships}"


def title_of(z, name):
    return "".join(re.findall(r"<a:t>([^<]*)</a:t>", z.read(name).decode("utf-8"))[:3])


def rels_of(name):
    d, f = name.rsplit("/", 1)
    return f"{d}/_rels/{f}.rels"


def resolve(base, target):
    parts = base.split("/")[:-1] + target.split("/")
    out = []
    for p in parts:
        if p == "..":
            out.pop()
        elif p != ".":
            out.append(p)
    return "/".join(out)


# ---------------------------------------------------------------- 1) 並べ替え・表示・外す（zipのまま）
def assemble(base, out):
    z = zipfile.ZipFile(base)
    order = slide_order(z)
    assert len(order) == N_BASE
    for n, t in [(30, "全体設計｜LINEで対策できる領域"), (34, "他社分析①｜主要エージェントのLINE運用"),
                 (35, "他社分析②｜運用の型は3つに分かれる"), (36, "具体企画のブレスト"), (40, "改善モデル①")]:
        assert title_of(z, order[n - 1]).startswith(t), (n, title_of(z, order[n - 1]))
    names = set(z.namelist())
    replace = {}
    for n in SHOW:
        xml = z.read(order[n - 1]).decode("utf-8")
        new, c = re.subn(r'(<p:sld\b[^>]*?) show="0"', r"\1", xml, count=1)
        assert c == 1, n
        replace[order[n - 1]] = new

    # 外すスライド：スライド・rels・ノート（＋ノートのrels）。画像などは、ほかから参照が無ければ外す
    gone, maybe = set(), set()
    for n in DROP:
        s = order[n - 1]
        assert 'show="0"' in z.read(s).decode("utf-8")[:2000], n
        gone |= {s, rels_of(s)}
        for r in etree.fromstring(z.read(rels_of(s))):
            if r.get("TargetMode") == "External":
                continue
            t = resolve(s, r.get("Target"))
            kind = r.get("Type").split("/")[-1]
            if kind == "notesSlide":
                gone |= {t, rels_of(t)}
            elif kind != "slideLayout":
                maybe.add(t)
    kept_refs = set()
    for name in names - gone:
        if name.endswith(".rels"):
            src = name.replace("_rels/", "").removesuffix(".rels")
            for r in etree.fromstring(z.read(name)):
                if r.get("TargetMode") != "External":
                    kept_refs.add(resolve(src, r.get("Target")) if not name.startswith("_rels/") else r.get("Target"))
    gone |= {m for m in maybe if m not in kept_refs}
    for name in names - gone:                                  # 外したパーツへの参照が残っていないこと
        if name.endswith(".rels") and name != "ppt/_rels/presentation.xml.rels":
            src = name.replace("_rels/", "").removesuffix(".rels")
            for r in etree.fromstring(z.read(name)):
                if r.get("TargetMode") != "External":
                    assert resolve(src, r.get("Target")) not in gone, (name, r.get("Target"))

    prels = z.read("ppt/_rels/presentation.xml.rels").decode("utf-8")
    pres = z.read("ppt/presentation.xml").decode("utf-8")
    ct = z.read("[Content_Types].xml").decode("utf-8")

    def rid_of_part(name):
        m = re.search(r'<Relationship [^>]*Target="slides/%s"[^>]*/>' % re.escape(name.split("/")[-1]), prels)
        return m, re.search(r'Id="(rId\d+)"', m.group(0)).group(1)

    m = re.search(r"<p:sldIdLst>(.*?)</p:sldIdLst>", pres, re.S)
    entries = {re.search(r'r:id="(rId\d+)"', e).group(1): e for e in re.findall(r"<p:sldId [^>]*/>", m.group(1))}
    pres = pres.replace(m.group(0), "<p:sldIdLst>" + "".join(entries[rid_of_part(order[u - 1])[1]] for u in ORDER)
                        + "</p:sldIdLst>")
    for n in DROP:
        mm, _ = rid_of_part(order[n - 1])
        prels = prels.replace(mm.group(0), "")
    for g in gone:
        ct, c = re.subn(r'<Override [^>]*PartName="/%s"[^>]*/>' % re.escape(g), "", ct)
        assert c <= 1

    app = z.read("docProps/app.xml").decode("utf-8")
    nn = len(ORDER)
    assert f"<Slides>{N_BASE}</Slides>" in app
    app = app.replace(f"<Slides>{N_BASE}</Slides>", f"<Slides>{nn}</Slides>")
    head = f"<vt:lpstr>スライド タイトル</vt:lpstr></vt:variant><vt:variant><vt:i4>{N_BASE}</vt:i4>"
    assert head in app
    app = app.replace(head, head.replace(f"<vt:i4>{N_BASE}</vt:i4>", f"<vt:i4>{nn}</vt:i4>"))
    mv = re.search(r'<vt:vector size="(\d+)" baseType="lpstr">', app)
    app = app.replace(mv.group(0), f'<vt:vector size="{int(mv.group(1)) - len(DROP)}" baseType="lpstr">')
    for _ in DROP:
        assert "<vt:lpstr>PowerPoint プレゼンテーション</vt:lpstr>" in app
        app = app.replace("<vt:lpstr>PowerPoint プレゼンテーション</vt:lpstr>", "", 1)
    hidden = sum(1 for u in ORDER if u not in SHOW and 'show="0"' in z.read(order[u - 1]).decode("utf-8")[:2000])
    app, c = re.subn(r"<HiddenSlides>\d+</HiddenSlides>", f"<HiddenSlides>{hidden}</HiddenSlides>", app)
    assert c == 1
    replace.update({"ppt/presentation.xml": pres, "ppt/_rels/presentation.xml.rels": prels,
                    "[Content_Types].xml": ct, "docProps/app.xml": app})

    with zipfile.ZipFile(out, "w") as zo:
        for info in z.infolist():
            if info.filename in gone:
                continue
            data = replace.get(info.filename)
            zo.writestr(info, data.encode("utf-8") if data is not None else z.read(info.filename),
                        compress_type=info.compress_type)
    return sorted(gone)


# ---------------------------------------------------------------- 2) 文字
def p(u):
    return f"P{NEW_POS[u]}"


def edit(prs):
    at = {u: prs.slides[i - 1] for u, i in NEW_POS.items()}
    R = replace_in_paragraphs
    for u, old, new in [
            (30, "全体設計｜LINEで対策できる領域", "［第2章］全体設計｜LINEで対策できる領域"),
            (31, "全体設計｜施策の考え方", "［第2章］全体設計｜施策の考え方"),
            (34, "他社分析①｜主要エージェント", "［第2章］他社分析｜主要エージェント"),
            (32, "構築｜施策展開図（初期・月次）", "［第3章］構築｜施策展開図（初期・月次）"),
            (33, "構築｜友だち追加後15日で", "［第4章］配信設計｜追加後15日で"),
            # ページ参照
            (12, "リッチメニュー（P16）", f"リッチメニュー（{p(16)}）"),
            (12, "実際の画面はP11。", f"実際の画面は{p(11)}。"),
            (16, "（2019年公開・P10）", f"（2019年公開・{p(10)}）"),
            (25, "（事例はP10〜P11）", f"（事例は{p(10)}〜{p(11)}）"),
            # 39：他業界の言葉
            (39, "アンケ回答例→学習コンテンツ", "アンケート回答例→お役立ち情報"),
            (39, "学習が進行する導線設計", "応募に進む導線設計"),
            (39, "◯再購入・クロスセル", "◯再応募・次の仕事の案内"),
            (39, "再購入時期、併売商品を設定", "満了の時期・近い職種を設定"),
            (39, "LINEで査定", "LINEで条件確認"),
            # アジェンダ：第2章に他社分析
            (2, "データと参考事例（タウンワーク）", "データ・他社分析・タウンワーク")]:
        assert R(at[u], old, new) == 1, (u, old)
    t = [sh for sh in at[39].shapes if sh.has_text_frame and sh.text_frame.text == "施策展開図"]
    assert len(t) == 1
    assert R(at[39], "施策展開図", "［第3章］構築｜施策展開図") >= 1


def main(base, out):
    mid, tmp = out + ".mid.pptx", out + ".tmp.pptx"
    gone = assemble(base, mid)
    prs = Presentation(mid)
    edit(prs)
    prs.save(tmp)
    jb.ORDER = ORDER                      # repack は並びの長さだけ見る
    jb.repack(mid, tmp, out)
    Path(mid).unlink()
    Path(tmp).unlink()
    shown = sum(1 for s in Presentation(out).slides if s._element.get("show") != "0")
    print("saved:", out, "／", len(ORDER), "枚（表示", shown, "枚）／外したパーツ:", len(gone))
    for g in gone:
        print("  -", g)


if __name__ == "__main__":
    main(*sys.argv[1:3])
