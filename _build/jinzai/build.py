# -*- coding: utf-8 -*-
"""人材（派遣・バイト・単発）業界：各版の良いところをつなぎ合わせる＋タウンワークを運用の参考事例に（2026-10-20 殿村さん）。

    python -m _build.jinzai.build <章立て版（37枚・殿村さん保存版 2026-09-15 15時）> <出力.pptx>

  殿村さん：「どれが一番良い資料？それに、タウンワークのLINE導入事例を入れれたら」
           →「各資料の良いところを持ち寄って良い資料につなぎ合わせて」
  受け取った5つ（_data/jinzai/）：
    章立て版（9/15 15時・37枚中20枚表示）＝一番新しい。文言も一番正確（通知メッセージは審査制、仮説の明示、就業率が前提）
    9/14版（37枚中28枚表示）／28枚のPDF（9/15 11時）／27枚のPDF（9/15 14時）＝章立て前の版。文言は章立て版のほうが新しい
    施策ビジュアル（①〜④）＝件数が「○件」のままの下書き。本資料には入れない（別のPPTXのまま）
  → 章立て版を土台に、9/14版で表示していた「データの裏付け」のページを表示に戻す：
       検討背景（3）・前後検索（8）・希望条件5問（11）・7日ステップの全体像（13）・求職者の心理（25）・効果測定の設計（30）
     汎用のページ（26〜29）・転職の他社分析（31・32）・ブレストなど（33〜37）は非表示のまま。
  → タウンワーク（LINEヤフー公式の導入事例・2019年1月公開）を、ジム（カーブス）・塾（明光義塾）と同じ3枚で入れる：
       第2章の最後＝運用の参考事例｜タウンワーク／タウンワークの実際の画面（2026年9月 弊社確認のスクショ4枚）
       第3章の先頭＝タウンワークの型を分解すると、3つの仕組み
     「LINEOA実績・信頼性」（21）の「同業種の導入事例：未取得」を埋め、リッチメニュー（12）に参照を足す。
  → アジェンダの章の中身を新しい並びに合わせ、社内向けのメモ（要確認・差込枠・未取得・提案までに確認）を消す。
  ほかは元のスライドのまま（直したスライドのXMLと、足した3枚・画像だけを差し込む）。
"""
import copy
import random
import re
import sys
import zipfile
from io import BytesIO
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.oxml.ns import qn

import _build.juku_fmt.build as jb
from _build.arauchi_v4.build import slide_order
from _build.arauchi_v4.lib import add_box
from _build.arauchi_v8.pages import replace_in_paragraphs, shape_by_id
from _build.gym_curves.build import add_pic, keep_only, place, round_box, set_one
from _build.gym_fmt.build import fix_fonts
from _build.juku_fmt.build import CT_SLIDE, RT_SLIDE, rels_of
from _build.juku_fmt.plan import set_lead

ROOT = Path(__file__).resolve().parent.parent.parent
SHOTS = ROOT / "_data/参考資料/タウンワークLINE_20260928"
N_BASE = 37
TPL = 21                                     # 型にするページ（見出し・説明・下の帯・注記がある）
SHOW = [3, 8, 11, 13, 25, 30]                 # 表示に戻すページ
ORDER = ([1, 2, 3, 4, 5, 6, 7, 8, 25, ("N", 1), ("N", 2), ("N", 3), 9, 10, 11, 12,
          13, 14, 15, 16, 17, 18, 19, 20, 21, 30, 22, 23, 24] +
         [26, 27, 28, 29, 31, 32, 33, 34, 35, 36, 37])
NEW_POS = {key: i for i, key in enumerate(ORDER, 1)}
RT_IMAGE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"

NAVY, NAVY2, INK, GREY = "002060", "1E2761", "333333", "595959"
PALE, BORDER = "EEF2F9", "D9D9D9"
X0, W = 0.47, 9.89
BAND_Y = 5.63

CASE_TITLE = "バイト応募件数約80％増！リッチメニューを生かしたタウンワークのLINE活用事例"
SCREEN_SHOTS = [
    ("1　あいさつ", "1_あいさつ.png", ["追加の直後に「気になるテーマを", "タップ」と時間帯別のカード"]),
    ("2　リッチメニュー", "2_リッチメニュー.png", ["時間帯・学生歓迎・単発など", "よく探す条件がワンタップ"]),
    ("3　移動先のサイト", "3_遷移先サイト.png", ["条件を入れた状態の", "検索結果ページ"]),
    ("4　会員連携の画面", "4_連携リンク先_リクルートIDログイン.png", ["合う求人を受け取るための", "会員連携（リクルートID）"]),
]
SHOT_ASPECT = 1170 / 2532


def shot_bytes(name):
    im = Image.open(SHOTS / name).convert("RGB")
    im = im.resize((585, 1266), Image.LANCZOS)
    buf = BytesIO()
    im.save(buf, "JPEG", quality=90)
    return buf.getvalue()


def title_of(z, name):
    m = re.search(r"<a:t>([^<]*)</a:t>", z.read(name).decode("utf-8"))
    return m.group(1) if m else ""


# ---------------------------------------------------------------- 1) 組み立て（zipのまま）
def assemble(base, out):
    z = zipfile.ZipFile(base)
    order = slide_order(z)
    assert len(order) == N_BASE
    for n, t in [(TPL, "［第5章］LINEOA実績・信頼性"), (3, "［第1章］検討背景｜応募は「決断」ではなく「とりあえず」"),
                 (8, "［第2章］ニーズ調査｜前後検索"), (13, "［第4章］配信設計｜7日ステップの全体像"),
                 (25, "求職者の心理｜不安は、応募したあとに出てくる"), (30, "成果｜効果測定の設計"),
                 (12, "［第3章］構築｜リッチメニュー（3タブ × 6枠）")]:
        assert title_of(z, order[n - 1]) == t, (n, title_of(z, order[n - 1]))
    names = set(z.namelist())
    ct = z.read("[Content_Types].xml").decode("utf-8")
    prels = z.read("ppt/_rels/presentation.xml.rels").decode("utf-8")
    pres = z.read("ppt/presentation.xml").decode("utf-8")
    assert '<Default Extension="jpeg"' in ct

    def rid_of_part(name):
        m = re.search(r'<Relationship [^>]*Target="slides/%s"[^>]*/>' % re.escape(name.split("/")[-1]), prels)
        return re.search(r'Id="(rId\d+)"', m.group(0)).group(1)

    rid_of = {n: rid_of_part(order[n - 1]) for n in range(1, N_BASE + 1)}
    next_rid = max(int(x) for x in re.findall(r'Id="rId(\d+)"', prels)) + 1
    next_slide = max(int(m.group(1)) for n in names for m in [re.match(r"ppt/slides/slide(\d+)\.xml$", n)] if m) + 1
    replace = {}

    # 表示に戻す（<p:sld show="0"> の show を外す）
    for n in SHOW:
        xml = z.read(order[n - 1]).decode("utf-8")
        new, c = re.subn(r'(<p:sld\b[^>]*?) show="0"', r"\1", xml, count=1)
        assert c == 1, n
        replace[order[n - 1]] = new

    tpl = z.read(order[TPL - 1]).decode("utf-8")
    tpl_rels = etree.fromstring(z.read(rels_of(order[TPL - 1])))
    for r in list(tpl_rels):
        if r.get("Type").endswith("/notesSlide"):
            tpl_rels.remove(r)
    assert [r.get("Type").split("/")[-1] for r in tpl_rels] == ["slideLayout"]
    used = set(re.findall(r'creationId[^>]*val="(\d+)"', "".join(z.read(o).decode("utf-8") for o in order)))
    rnd = random.Random(202610201)
    add = {}
    for k in (1, 2, 3):
        dst = f"ppt/slides/slide{next_slide}.xml"
        next_slide += 1
        cid = str(rnd.randrange(10 ** 9, 4 * 10 ** 9))
        assert cid not in used
        xml, c = re.subn(r'(<p14:creationId [^>]*val=")\d+(")', rf"\g<1>{cid}\2", tpl, count=1)
        assert c <= 1                               # 21枚目は python-pptx で作られたページで creationId が無い
        add[dst] = xml.encode("utf-8")
        rels = copy.deepcopy(tpl_rels)
        if k == 2:                                  # 実際の画面：スクショ4枚（パーツ名は英数字だけ）
            for j, (_, fn, _) in enumerate(SCREEN_SHOTS, 1):
                media = f"ppt/media/townwork_{j}.jpeg"
                assert media not in names
                add[media] = shot_bytes(fn)
                etree.SubElement(rels, "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship",
                                 Id=f"rId{100 + j}", Type=RT_IMAGE, Target=f"../media/townwork_{j}.jpeg")
        add[rels_of(dst)] = etree.tostring(rels, xml_declaration=True, encoding="UTF-8", standalone=True)
        ct = ct.replace("</Types>", f'<Override PartName="/{dst}" ContentType="{CT_SLIDE}"/></Types>')
        rid = f"rId{next_rid}"
        next_rid += 1
        prels = prels.replace("</Relationships>", f'<Relationship Id="{rid}" Type="{RT_SLIDE}" Target="slides/{dst.split("/")[-1]}"/></Relationships>')
        rid_of[("N", k)] = rid

    m = re.search(r"<p:sldIdLst>(.*?)</p:sldIdLst>", pres, re.S)
    entries = {re.search(r'r:id="(rId\d+)"', e).group(1): e for e in re.findall(r"<p:sldId [^>]*/>", m.group(1))}
    next_id = max(int(x) for x in re.findall(r'<p:sldId id="(\d+)"', m.group(1))) + 1
    lst = []
    for key in ORDER:
        if rid_of[key] in entries:
            lst.append(entries[rid_of[key]])
        else:
            lst.append(f'<p:sldId id="{next_id}" r:id="{rid_of[key]}"/>')
            next_id += 1
    pres = pres.replace(m.group(0), "<p:sldIdLst>" + "".join(lst) + "</p:sldIdLst>")

    app = z.read("docProps/app.xml").decode("utf-8")
    n_new = len(ORDER)
    m_ttl = re.search(r'<vt:vector size="(\d+)" baseType="lpstr">(.*?)</vt:vector>', app, re.S)
    titles = re.findall(r"<vt:lpstr>.*?</vt:lpstr>", m_ttl.group(2))
    assert f"<Slides>{N_BASE}</Slides>" in app
    app = app.replace(f"<Slides>{N_BASE}</Slides>", f"<Slides>{n_new}</Slides>")
    app, c = re.subn(r"(<vt:lpstr>スライド タイトル</vt:lpstr></vt:variant><vt:variant><vt:i4>)\d+(</vt:i4>)", rf"\g<1>{n_new}\2", app)
    assert c == 1
    extra = "<vt:lpstr>PowerPoint プレゼンテーション</vt:lpstr>" * 3
    hidden = sum(1 for n in range(1, N_BASE + 1) if n not in SHOW and 'show="0"' in z.read(order[n - 1]).decode("utf-8")[:2000])
    app, c = re.subn(r"<HiddenSlides>\d+</HiddenSlides>", f"<HiddenSlides>{hidden}</HiddenSlides>", app)
    assert c == 1
    app = app.replace(m_ttl.group(0), f'<vt:vector size="{len(titles) + 3}" baseType="lpstr">' + m_ttl.group(2) + extra + "</vt:vector>")
    replace.update({"ppt/presentation.xml": pres, "ppt/_rels/presentation.xml.rels": prels,
                    "[Content_Types].xml": ct, "docProps/app.xml": app})

    with zipfile.ZipFile(out, "w") as zo:
        for info in z.infolist():
            data = replace.get(info.filename)
            zo.writestr(info, data.encode("utf-8") if data is not None else z.read(info.filename),
                        compress_type=info.compress_type)
        for name, data in add.items():
            zo.writestr(name, data, compress_type=zipfile.ZIP_STORED if name.endswith(".jpeg") else zipfile.ZIP_DEFLATED)


# ---------------------------------------------------------------- 2) 中身
def chip(s, x, y, w, h, text, size=11, fill=NAVY2):
    return round_box(s, x, y, w, h, [[(text, {})]], adj=20000, size=size, bold=True, color="FFFFFF", fill=fill,
                     align="ctr", ins=(0.04, 0, 0.04, 0))


def base_page(s, title, lead, band=None, foot=None):
    keep_only(s, {2, 3, 18, 19, 20})
    set_one(s, 2, title)
    set_lead(s, lead)
    place(shape_by_id(s, 3), h=0.64)
    el = shape_by_id(s, 19)._element                      # 帯の文字は作り直す（元の文字は書式が引き継げない）
    el.getparent().remove(el)
    if band is None:
        el = shape_by_id(s, 18)._element
        el.getparent().remove(el)
    else:
        place(shape_by_id(s, 18), y=BAND_Y, h=0.41)      # 帯の位置はほかのページ（5.63in）にそろえる
        add_box(s, X0, BAND_Y, W, 0.41, [[(band, {})]], size=12.5, bold=True, color="FFFFFF", align="ctr",
                ins=(0.05, 0, 0.05, 0))
    set_one(s, 20, foot)


def p(n):
    return f"P{NEW_POS[n]}"


def page_case(s):
    """運用の参考事例｜タウンワーク"""
    base_page(s, "［第2章］運用の参考事例｜タウンワーク",
              ["アルバイト求人のタウンワークは、LINEを「求人を探す入口」にしました（LINEヤフー公式の導入事例）。",
               "リッチメニューに人気の検索条件を並べ、ワンタップで探せるようにしたことで、応募が増えています。"],
              band="求人を「届ける」だけでなく、「探す入口」をLINEに置いたことが応募につながった。",
              foot=f"出典｜LINEヤフー for Business 導入事例「{CASE_TITLE}」（2019年1月公開・数値は取材時点）")
    kw, gap = 3.15, (W - 3 * 3.15) / 2
    for i, (lab, big, sub) in enumerate([
            ("リッチメニューのタップ数", "＋20%", "リッチメニューを充実させる前と比べて"),
            ("リッチメニュー経由の応募数", "＋79%", "同じく、充実させる前と比べて"),
            ("LINE経由の応募の内訳", "リッチメニューが最多", "配信（プッシュ）からの応募を上回る")]):
        round_box(s, X0 + i * (kw + gap), 1.55, kw, 1.05,
                  [[(lab, {"size": 10.5, "bold": True, "color": NAVY})],
                   [(big, {"size": 20 if i < 2 else 15, "bold": True, "color": NAVY})],
                   [(sub, {"size": 9, "color": GREY})]], adj=6000, fill=PALE, line=BORDER, align="ctr",
                  line_spacing=100)
    add_box(s, X0, 2.72, W, 0.28, [[("何をしたか", {})]], size=11, bold=True, color=NAVY, ins=(0, 0, 0, 0))
    steps = [("1　リッチメニューに条件を並べる", ["よく探されるバイトの条件を", "ボタンにして、ワンタップで検索"]),
             ("2　トークで条件を選ぶ", ["「トーク検索」で条件を選ぶと、", "合う求人がカード形式で届く"]),
             ("3　そのまま求人ページへ", ["条件を入れた状態の検索結果に", "移動し、サイトで応募する"])]
    for i, (head, lines) in enumerate(steps):
        x = X0 + i * (kw + gap)
        chip(s, x, 3.02, kw, 0.3, head, size=10.5)
        add_box(s, x, 3.32, kw, 0.78, [[(t, {})] for t in lines], size=10.5, color=INK, fill="FFFFFF", line=BORDER,
                align="ctr", anchor="ctr", ins=(0.06, 0.03, 0.06, 0.03), line_spacing=105)
    add_box(s, X0, 4.27, W, 0.62,
            [[("一番タップされたのは、一番上に置いた「人気のバイト検索条件」。", {"bold": True, "color": NAVY})],
             [("リッチメニューは、トーク画面を開くと必ず目に入る場所。そこに「探す」を置いています。", {})]],
            size=10.5, color=INK, fill="FFFFFF", line=NAVY, ins=(0.14, 0.04, 0.14, 0.04), line_spacing=105)
    add_box(s, X0, 5.0, W, 0.6,
            [[("※ LINE公式アカウントの開設は2017年4月、リッチメニューの強化は2018年2月から（Messaging APIを利用）。", {})],
             [("※ 今後は、利用者の状況や行動に応じて、リッチメニューを出し分けることを構想していると紹介されています。", {})]],
            size=9, color=GREY, ins=(0, 0, 0, 0), line_spacing=100)


def page_screens(s):
    """タウンワークの実際の画面（スクショ4枚）"""
    base_page(s, "［第2章］タウンワークの実際の画面",
              ["友だち追加の直後に「気になるテーマをタップ」と案内し、リッチメニューにはよく探す条件が並びます。",
               "条件を選ぶと、条件を入れた状態の検索結果ページへ移動し、応募はサイトで行う作りです。"],
              foot="画面：タウンワーク LINE公式アカウント（2026年9月 弊社確認）")
    colw = 2.3
    gap = (W - 4 * colw) / 3
    ih = 4.15
    iw = ih * SHOT_ASPECT
    for j, (head, _, cap) in enumerate(SCREEN_SHOTS):
        x = X0 + j * (colw + gap)
        chip(s, x, 1.55, colw, 0.3, head, size=10.5)
        add_pic(s, f"rId{101 + j}", f"タウンワーク_{head[2:]}", x + (colw - iw) / 2, 1.93, iw, ih)
        add_box(s, x - 0.1, 6.14, colw + 0.2, 0.55, [[(t, {})] for t in cap], size=9.5, color=INK, align="ctr", anchor="t",
                ins=(0, 0, 0, 0), line_spacing=100)


def page_pattern(s):
    """タウンワークの型を分解すると、3つの仕組み"""
    base_page(s, "［第3章］タウンワークの型を分解すると、3つの仕組み",
              ["公式事例（2019年）と実際の画面（2026年9月）から読み取れる仕組みは3つです。",
               f"この型を、このあとのあいさつ（{p(10)}）・希望条件（{p(11)}）・リッチメニュー（{p(12)}）に取り入れます。"],
              band="LINEの中で全部を完結させない。「次に探すときの入口」をLINEに置く。",
              foot=f"※ 出典：LINEヤフー for Business 導入事例（タウンワーク・2019年1月公開）・タウンワーク LINE公式アカウント"
                   f"（2026年9月 弊社確認）。実際の画面は{p(('N', 2))}。")
    kw, gap = 3.15, (W - 3 * 3.15) / 2
    cols = [
        ("① リッチメニューを「探す入口」に",
         ["時間帯（午前中・お昼から・夕方〜夜・深夜・夜勤）や「学生歓迎」「単発・1日OK」など、よく探す条件をワンタップに。",
          "一番上に、一番人気の条件を置く。"],
         ["タップ数 ＋20%", "リッチメニュー経由の応募 ＋79%"]),
        ("② 追加の直後に、探し方を案内",
         ["あいさつで「気になるテーマをタップ」と案内し、時間帯別のカードを出す。",
          "通知が多いと感じたら、オフにする方法も先に伝えている。"],
         ["あいさつ＋時間帯別のカード", "（2026年9月 弊社確認）"]),
        ("③ 応募はサイトで、LINEは入口",
         ["条件を選ぶと、条件を入れた状態の検索結果ページへ移動。応募はサイトで行う。",
          "合う求人を受け取るために、会員とLINEを連携してもらう。"],
         ["移動先＝条件入りの検索結果", "会員連携＝リクルートID"]),
    ]
    for i, (head, body, ex) in enumerate(cols):
        x = X0 + i * (kw + gap)
        chip(s, x, 1.55, kw, 0.34, head, size=11)
        round_box(s, x, 1.99, kw, 1.95, [[(t, {})] for t in body], adj=5000, size=11.5, color=INK, fill=PALE,
                  line=BORDER, anchor="t", ins=(0.12, 0.1, 0.1, 0.06), after=6, line_spacing=110)
        round_box(s, x, 4.06, kw, 1.0,
                  [[("実際の例", {"size": 10, "bold": True, "color": NAVY})]] + [[(t, {})] for t in ex],
                  adj=5000, size=10.5, color=INK, fill="FFFFFF", line=BORDER, anchor="t", ins=(0.1, 0.08, 0.08, 0.04),
                  line_spacing=105)


def set_runs(sh, texts):
    """段落ごとに、最初のrunの文字だけ差し替える（書式はそのまま）。None はそのまま"""
    ps = sh.text_frame.paragraphs
    assert len(ps) == len(texts), (sh.shape_id, len(ps))
    for para, t in zip(ps, texts):
        if t is None:
            continue
        runs = para.runs
        runs[0].text = t
        for r in runs[1:]:
            r._r.getparent().remove(r._r)


def fix_agenda(s):
    by = {sh.shape_id: sh for sh in s.shapes}
    set_runs(by[15], ["主戦場は契約満了からの再応募。データと参考事例（タウンワーク）"])
    set_runs(by[20], ["タウンワークの型・友だち追加の動線・あいさつ・リッチメニュー"])
    set_runs(by[30], ["運用サポート体制・LINEOA実績・効果測定の設計"])
    assert replace_in_paragraphs(s, "本丸は「第4章 育成・効率化」", "本丸は「第4章 配信設計」") == 1
    place(by[41], y=6.28)                                  # 2行の注記がフッターに重なっていた


def fix_results(s):
    """LINEOA実績・信頼性：同業種の導入事例＝タウンワーク"""
    by = {sh.shape_id: sh for sh in s.shapes}
    set_runs(by[14], ["タウンワーク｜リッチメニュー経由の応募"])
    by[14].text_frame._txBody.find(qn("a:bodyPr")).set("wrap", "none")
    set_runs(by[15], ["＋79%"])
    set_runs(by[17], [f"数字はLINEヤフー公式の値です（事例は{p(('N', 1))}〜{p(('N', 3))}）"])
    el = by[19]._element                                    # 帯の文字に書式が無く、紺の帯に黒い文字になっていた → 白で作り直す
    el.getparent().remove(el)
    add_box(s, X0, by[18].top / 914400, W, by[18].height / 914400, [[("詳細は次回の商談でご確認いただけます", {})]],
            size=12.5, bold=True, color="FFFFFF", align="ctr", ins=(0.05, 0, 0.05, 0))
    set_runs(by[20], ["出典｜LINE国内MAU・開封率：LINEヤフー公式値／導入事例：LINEヤフー for Business 導入事例"
                      "（タウンワーク・2019年1月公開）"])


def fix_others(at):
    assert replace_in_paragraphs(at[25], "求職者の心理｜", "［第2章］求職者の心理｜") == 1
    assert replace_in_paragraphs(at[30], "成果｜効果測定の設計", "［第5章］成果｜効果測定の設計") == 1
    assert replace_in_paragraphs(at[9], "（用途は要確認）", "（利用できる用途は事前に確認）") == 1
    assert replace_in_paragraphs(at[11], "（★派遣を基準に設計）", "（派遣を基準に設計）") == 1
    assert replace_in_paragraphs(at[12], "タブ構成とボタン名は運用しながら調整します",
                                 f"タブ構成とボタン名は運用しながら調整します。タウンワークもリッチメニューに人気の検索条件を並べ、"
                                 f"リッチメニュー経由の応募が＋79%（公式事例・{p(('N', 1))}）") == 1


def edit(prs):
    sl = prs.slides
    at = {key: sl[i - 1] for key, i in NEW_POS.items()}
    fix_agenda(at[2])
    fix_results(at[TPL])
    fix_others(at)
    page_case(at[("N", 1)])
    page_screens(at[("N", 2)])
    page_pattern(at[("N", 3)])
    return sum(fix_fonts(at[k]) for k in (("N", 1), ("N", 2), ("N", 3), 2, TPL, 9, 11, 12, 25, 30))


def main(base, out):
    mid, tmp = out + ".mid.pptx", out + ".tmp.pptx"
    assemble(base, mid)
    prs = Presentation(mid)
    fixed = edit(prs)
    prs.save(tmp)
    jb.ORDER = ORDER                      # repack は並びの長さだけ見る
    jb.repack(mid, tmp, out)
    Path(mid).unlink()
    Path(tmp).unlink()
    shown = sum(1 for s in Presentation(out).slides if s._element.get("show") != "0")
    print("saved:", out, "／", len(ORDER), "枚（表示", shown, "枚）／メイリオにそろえた文字:", fixed)


if __name__ == "__main__":
    main(*sys.argv[1:3])
