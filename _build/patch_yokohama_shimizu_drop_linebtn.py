# -*- coding: utf-8 -*-
"""横浜シミズ：殿村さんのPC保存版の「概算お見積り」から、初期の実施内容「サイトのLINEボタン｜設置」の行を外す（2026-10-21）。

    python -m _build.patch_yokohama_shimizu_drop_linebtn <資料.pptx> [出力.pptx]

  殿村さん「そもそも、サイトのLINEボタンをこっちが触ること自体変でしょ」。運用内容の表の行は殿村さんがPCで外し済み。
  下の2行（ステップ配信・勤務前の案内）を1行ずつ上げる。ほかのページは触らない。出力を省くと上書き。
"""
import sys

from pptx import Presentation


def by_id(s, sid):
    return next(sh for sh in s.shapes if sh.shape_id == sid)


def main(src, out):
    prs = Presentation(src)
    s = prs.slides[12]
    assert by_id(s, 339).text_frame.text.startswith("概算お見積り")
    assert by_id(s, 369).text_frame.text == "サイトの\nLINEボタン"
    tops = {sid: by_id(s, sid).top for sid in (369, 371)}
    for sid in (369, 370):
        sh = by_id(s, sid)
        sh._element.getparent().remove(sh._element)
    for sids, top in (((371, 372), tops[369]), ((373, 374), tops[371])):
        for sid in sids:
            by_id(s, sid).top = top
    prs.save(out)
    print("saved:", out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else sys.argv[1])
