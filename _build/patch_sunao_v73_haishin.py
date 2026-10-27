# -*- coding: utf-8 -*-
"""すなお ver7.2 → ver7.3（案）：「配信の内容」の章に、前後検索・診断・15日のステップ・16日以降を入れる（2026-10-27 殿村さん）。

    python -m _build.patch_sunao_v73_haishin <ver7.2.pptx> <ver7.3.pptx>

  殿村さんのメモ：前後検索でアクティブ率が高い → 2日おきに配信／0〜15日と16日以降の配信を組む／
  ステップ配信だけでなく、市場データ（前後検索）の分析から、ユーザー・業界に合わせて設計したことを見せる。
  前後検索は「投資セミナー」（色分け済みの画像）。診断＝タグ付けのため（どんな方かを分かりやすくする）。
  並び：P32（重要な考え方）→ 前後検索 → 診断 → ステップ配信8通（P34と差し替え）→ 16日以降 → P33 → P35。
  P33（出し分け）を外すかは殿村さんの確認待ちなので、この版では残して16日以降の後ろへ。ほかのページは変えない。
"""
import sys

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches

from _build.patch_tanaka_v5 import copy_slide
from _build.pptx_parts import drop_slide, rect, rich, shape

NAVY, INK, SUB, LINE, LIGHT, PALE, WHITE = "1F2A5C", "333333", "777777", "D0D5E0", "EEF1F8", "F5F7FA", "FFFFFF"
ZENGO = {"A": "3467B2", "B": "8E4EC6", "C": "00897B"}       # 前後検索の画像の色分け（make_sunao_zengo_color と同じ）
IMG = "_images/sunao_zengo_toushi_seminar_color.png"
KEEP = {"Google Shape;156;p7", "正方形/長方形 2", "正方形/長方形 3", "直線コネクタ 1", "テキスト ボックス 8"}
IN = 2.54


def i(v):
    return v * IN


def set_text(sh, text):
    p = sh.text_frame.paragraphs[0]._p
    rs = p.findall(qn("a:r"))
    for r in rs[1:]:
        p.remove(r)
    rs[0].find(qn("a:t")).text = text
    for extra in sh.text_frame.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)


def new_page(prs, base, title, lead):
    s = copy_slide(prs, base)
    for sh in list(s.shapes):
        if sh.name not in KEEP:
            sh._element.getparent().remove(sh._element)
    by = {sh.name: sh for sh in s.shapes}
    set_text(by["Google Shape;156;p7"], title)
    by["Google Shape;156;p7"].width = Inches(7.6)
    set_text(by["テキスト ボックス 8"], lead)
    return s


def note(s, y, lines):
    rich(s, i(0.46), i(y), i(9.9), i(0.2 * len(lines) + 0.1), [[(t, 9, False, SUB)] for t in lines], anchor=MSO_ANCHOR.TOP)


def insight(s, y, text, h=0.78):
    rect(s, i(0.46), i(y), i(9.9), i(h), PALE, [("示唆", 13, True, NAVY, 2), (text, 13, False, INK, 0)],
         align=PP_ALIGN.LEFT, margins=(0.3, 0.12, 0.3, 0.12))


def table(s, x, y, widths, head, rows, rh=0.7, hh=0.4, size=12):
    xs = [x]
    for w in widths[:-1]:
        xs.append(xs[-1] + w + 0.03)
    for cx, w, t in zip(xs, widths, head):
        rect(s, i(cx), i(y), i(w), i(hh), NAVY, [(t, 13, True, WHITE, 0)], align=PP_ALIGN.CENTER)
    for r, row in enumerate(rows):
        ry = y + hh + 0.03 + (rh + 0.03) * r
        for k, (cx, w, t) in enumerate(zip(xs, widths, row)):
            if t == "^":                                   # 上のセルとつなげる
                continue
            span = 1
            while r + span < len(rows) and rows[r + span][k] == "^":
                span += 1
            rect(s, i(cx), i(ry), i(w), i(rh * span + 0.03 * (span - 1)), LIGHT if k == 0 else WHITE,
                 [(line, size, k == 0, NAVY if k == 0 else INK, 0) for line in t.split("\n")],
                 align=PP_ALIGN.LEFT, line=LINE, lw=0.75, margins=(0.25, 0.05, 0.15, 0.05))
    return y + hh + 0.03 + (rh + 0.03) * len(rows)


# ── 1. 前後検索「投資セミナー」 ──────────────────────────────────
def page_zengo(prs, base):
    s = new_page(prs, base, "前後検索「投資セミナー」：比べる・確かめる・始めるが15日続く",
                 "セミナーを探した方は、その日にほかのセミナーと比べて不安を確かめ、1〜4日後に始め方を調べる")
    w = 6.25
    s.shapes.add_picture(IMG, Inches(0.46), Inches(1.78), Inches(w), Inches(w * 1167 / 2000))
    items = [("A", "比べる（0日前後）", "ほかのセミナー・イベント", "投資セミナー 東京\n資産運用セミナー・マネーセミナー"),
             ("B", "確かめる（0日前後）", "不安・安全性", "投資詐欺・金融庁"),
             ("C", "始める（1〜4日後）", "自分で始める方法", "口座開設・ポートフォリオ\netfとは・nisa つみたて投資枠")]
    for k, (cat, head, what, kw) in enumerate(items):
        y = 1.78 + 1.24 * k
        shape(s, MSO_SHAPE.OVAL, i(6.95), i(y + 0.08), i(0.18), i(0.18), fill=ZENGO[cat])
        rich(s, i(7.2), i(y), i(3.2), i(1.15),
             [[(head, 13, True, NAVY)], [(what, 12, False, INK)]] + [[(t, 11, False, SUB)] for t in kw.split("\n")],
             anchor=MSO_ANCHOR.TOP)
    insight(s, 5.6, "友だち追加から15日間は関心が続くため、2日おきに「比べる → 確かめる → 始める」の順で届ける")
    note(s, 6.5, ["出典｜LINEヤフーの前後検索データ（「投資セミナー」）"])
    return s


# ── 2. 診断 ──────────────────────────────────────────────────
def page_shindan(prs, base):
    s = new_page(prs, base, "診断：最初の2問で、どんな方かが分かるようにする",
                 "答えでタグを付け、その後の配信をその方に合わせる")
    end = table(s, 0.46, 2.1, [1.6, 3.95, 1.45, 2.81], ["質問", "選択肢", "付けるタグ", "配信での使い方"],
                [["1問目\nあてはまる方", "保険営業の方／一般の方", "保険営業\n一般", "一般の方には、一般の方向けの\nセミナーを案内"],
                 ["2問目\nいちばんの関心", "保険営業の方：成約を上げたい／\n紹介を増やしたい／働く時間を減らしたい", "成約・紹介・\n時間",
                  "受講生の声を、同じ悩み\nだった方の声にする"],
                 ["^", "一般の方：資産形成を始めたい／\n保険を見直したい", "資産形成・\n見直し", "一般の方向けの案内の\n中身を変える"]],
                rh=0.78)
    insight(s, end + 0.3, "一度の診断で、その方に合う受講生の声とセミナーを届けられる。聞くのは友だち追加直後の1回だけにする")
    note(s, end + 1.25, ["※答えはボタンで選ぶ形。押したボタンで、その後の配信を分ける"])
    return s


# ── 3. ステップ配信（8通） ──────────────────────────────────────
STEPS = [
    ("0日目", "あいさつ・参加特典", "参加特典をご案内し、\n2問の診断で関心を確認"),
    ("2日目｜比べる", "ほかのセミナーとの違い", "両想い営業・ライフプラン\nブロックを動画で紹介"),
    ("4日目｜確かめる", "安心して参加できる理由", "講師の経歴と、参加は\n無料であることをお伝え"),
    ("6日目｜始める", "セミナーで分かること", "面談の進め方など、\n参加して持ち帰れること"),
    ("8日目", "受講生の声", "診断の答えに合わせて、\n同じ悩みだった方の声"),
    ("10日目", "よくある質問", "地方から参加できる・\n初心者でも分かる"),
    ("12日目", "次回セミナーのご案内", "開催日程と\n予約ページをご案内"),
    ("14日目", "予約締切のご案内", "締切が近いことを\nお伝えし、予約を後押し"),
]


def page_step(prs, base):
    s = new_page(prs, base, "ステップ配信（8通・案）：関心が続く15日間は、2日おきに届ける",
                 "前後検索の順番（比べる → 確かめる → 始める）に合わせて、セミナーの予約まで案内する")
    w, gap = 2.37, 0.14
    for k, (day, head, body) in enumerate(STEPS):
        x = 0.46 + (w + gap) * (k % 4)
        y = 1.85 + 2.0 * (k // 4)
        rect(s, i(x), i(y), i(w), i(0.4), NAVY, [(day, 13, True, WHITE, 0)], align=PP_ALIGN.CENTER)
        rect(s, i(x), i(y + 0.44), i(w), i(0.46), LIGHT, [(head, 13, True, NAVY, 0)], align=PP_ALIGN.CENTER)
        rect(s, i(x), i(y + 0.94), i(w), i(0.86), PALE, [(t, 12, False, INK, 0) for t in body.split("\n")],
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, margins=(0.2, 0.08, 0.15, 0.05))
    note(s, 5.95, ["※募集期間外は、次回の案内開始時に届ける。入口（離脱防止ポップアップ／通常の友だち追加）で、0日目のあいさつの言葉を変える",
                   "※一般の方を選んだ方には、一般の方向けの案内を中心に配信",
                   "※配信内容は、貴社LPの「生産性4倍セミナーの特徴」「受講生の声」「よくある質問」「講師紹介」をもとにした例"])
    return s


# ── 4. 16日以降 ────────────────────────────────────────────────
def page_after(prs, base):
    s = new_page(prs, base, "16日以降の配信：開催の有無で、回数と中身を変える",
                 "15日を過ぎたら回数を絞ってつながりを保ち、開催月は予約・参加まで後押しする")
    lanes = [
        ("開催のない月", "月1〜2回", 2.1, [("講師のコラム", "営業のヒントを届ける"), ("受講生の声", "成果が出た方の話"),
                                       ("次回の予告", "次のセミナーの時期")]),
        ("開催月", "月4〜5回", 3.85, [("開催のご案内", "日程と予約ページ"), ("見どころ", "動画・受講生の声"),
                                   ("予約締切", "締切をお知らせ"), ("前日・当日", "予約した方へ\nリマインド"),
                                   ("参加後", "サロンのご案内\n（次のページ）")]),
    ]
    for name, freq, y, boxes in lanes:
        h = 1.4
        rect(s, i(0.46), i(y), i(1.75), i(h), NAVY, [(name, 13, True, WHITE, 2), (freq, 12, False, WHITE, 0)], align=PP_ALIGN.CENTER)
        n = len(boxes)
        aw = 0.24
        bw = (10.36 - 2.36 - aw * (n - 1)) / n
        for k, (head, body) in enumerate(boxes):
            x = 2.36 + (bw + aw) * k
            rect(s, i(x), i(y), i(bw), i(h), WHITE, [(head, 13, True, NAVY, 3)] + [(t, 12, False, INK, 0) for t in body.split("\n")],
                 align=PP_ALIGN.CENTER, line=LINE, lw=0.75, margins=(0.08, 0.05, 0.08, 0.05))
            if k < n - 1:
                rich(s, i(x + bw), i(y), i(aw), i(h), [[("▶", 11, False, NAVY)]], align=PP_ALIGN.CENTER)
    note(s, 5.45, ["※前日・当日のリマインドと参加後のフォローは、予約・参加した方を特定できる範囲で実施（予約システム等との連携は、可否を確認のうえ設計）"])
    return s


def main(src, out):
    prs = Presentation(src)
    slides = list(prs.slides)
    assert len(slides) == 61
    p32, p33, p34, p35 = slides[31], slides[32], slides[33], slides[34]
    assert p34.shapes[0].text_frame.text.startswith("友だち追加直後のステップ配信")
    new = [page_zengo(prs, p34), page_shindan(prs, p34), page_step(prs, p34), page_after(prs, p34)]
    lst = prs.slides._sldIdLst
    el = {int(e.get("id")): e for e in lst}
    order = [el[x.slide_id] for x in [p32] + new + [p33]]
    for e in order[1:]:
        lst.remove(e)
    anchor = order[0]
    for e in order[1:]:
        anchor.addnext(e)
        anchor = e
    drop_slide(prs, p34)
    texts = ["".join(sh.text_frame.text for sh in sl.shapes if sh.has_text_frame)[:20] for sl in list(prs.slides)[31:38]]
    print(texts)
    assert len(prs.slides) == 64
    prs.save(out)
    print("saved:", out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
