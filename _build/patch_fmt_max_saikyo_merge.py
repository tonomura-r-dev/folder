# -*- coding: utf-8 -*-
"""MAX版FMTと最強資料を同じ中身にそろえる（2026-10-29 殿村さん）。

    python -m _build.patch_fmt_max_saikyo_merge <MAX版.pptx> <最強資料.pptx> <出力.pptx>

  土台はMAX版（リード・アジェンダ・サンクスLINEの図などが新しい方）。最強資料から次を移す。
  ・P21 具体施策のイメージ：札の「露出の強化」（P22と重複）→「LINEの友だちへ」
  ・P41〜P43 提案に向けて明確化したいポイント①〜③：最強資料の記入欄の列（度合い・目標／数値・内容）を足す
  ・P45（④その他の古い1行版。P44の表と重複）を外す
  名前は「ポップアップ」にそろえる（MAX版は元からポップアップ）。
"""
import copy
import sys

from pptx import Presentation

E = 914400


def cells(s):
    return [sh for sh in s.shapes if sh.name.startswith("Rectangle") and sh.top > 1.85 * E and sh.top < 6.4 * E]


def add_col(ms, ss):
    mc, sc = cells(ms), cells(ss)
    xcol = max(sh.left for sh in sc)
    extra = sorted([sh for sh in sc if sh.left == xcol], key=lambda x: x.top)
    dx = min(sh.left for sh in sc) - min(sh.left for sh in mc)
    for sh in mc:
        sh.left += dx
    assert max(sh.left + sh.width for sh in mc) <= xcol + E * 0.05, "列が重なる"
    first = min(sh.left for sh in mc)
    rows = sorted([sh for sh in mc if sh.left == first], key=lambda x: x.top)
    assert len(rows) == len(extra), (len(rows), len(extra))
    anchor = rows[-1]._element
    for r, sh in zip(rows, extra):
        el = copy.deepcopy(sh._element)
        anchor.addnext(el)
        anchor = el
    # 位置を行に合わせる
    added = [sh for sh in ms.shapes if sh.left == xcol]
    for r, sh in zip(rows, sorted(added, key=lambda x: x.top)):
        sh.top, sh.height = r.top, r.height
    ids = [int(x.get("id")) for x in ms._element.iter() if x.tag.endswith("}cNvPr")]
    nxt = max(ids) + 1
    for sh in added:
        sh._element.nvSpPr.cNvPr.set("id", str(nxt))
        nxt += 1


def main(max_path, sk_path, out):
    m, s = Presentation(max_path), Presentation(sk_path)
    assert len(m.slides) == 54 and len(s.slides) == 54
    tag = next(sh for sh in m.slides[20].shapes if sh.has_text_frame and sh.text_frame.text == "露出の強化")
    r = tag.text_frame.paragraphs[0].runs
    r[0].text = "LINEの友だちへ"
    for x in r[1:]:
        x.text = ""
    for mn, sn in [(41, 42), (42, 43), (43, 44)]:
        add_col(m.slides[mn - 1], s.slides[sn - 1])
    p45 = m.slides[44]
    assert "④" in p45.shapes[0].text_frame.text or any("④" in sh.text_frame.text for sh in p45.shapes if sh.has_text_frame)
    lst = m.slides._sldIdLst
    sid = list(lst)[44]
    m.part.drop_rel(sid.rId)
    lst.remove(sid)
    assert len(m.slides) == 53
    m.save(out)
    print("saved:", out)


if __name__ == "__main__":
    main(*sys.argv[1:4])
