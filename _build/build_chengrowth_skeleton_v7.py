#!/usr/bin/env python3
"""
チェングロウス アジェンダv7 の簡易スケルトン資料を作る。
中身は入れず、章・タイトル・要旨だけを箱として並べる（構成確認用）。
出力: 20260928_株式会社チェングロウス御中_LINE公式アカウント運用のご提案_構成案.pptx
"""
import sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

NAVY = RGBColor(0x1F, 0x28, 0x5A)
TITLE_NAVY = RGBColor(0x00, 0x20, 0x60)
GRAY = RGBColor(0x33, 0x33, 0x33)
LIGHT_BG = RGBColor(0xF4, 0xF7, 0xFF)
BORDER = RGBColor(0xD9, 0xD9, 0xD9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "メイリオ"

SW, SH = Inches(10.83), Inches(7.5)

# (chapter, title, summary, note) — noteは差し込み待ち等の注記のみ
SLIDES = [
    ("0", "表紙", "自動車求人Navi LINE公式アカウント運用のご提案", ""),
    ("0", "アジェンダ", "業界の課題→貴社の現状→全体設計→施策→効果・費用の順に進め、最後に弊社の運用のご案内（汎用）をまとめる", ""),

    ("1", "整備士は超売り手市場", "求人倍率5.46倍。事業者の約6割が人手不足を感じている", ""),
    ("1", "担い手が増えない", "整備学校の入学者は約4割減、平均年齢は47.2歳。有資格者の取り合いは続く", ""),
    ("1", "職種ごとの採用難易度", "営業は12.55倍で整備士以上に不足、受付は0.96倍で人が集まる。全職種を対象にしつつ、主役は整備士", ""),
    ("1", "採用側の悩み", "「応募が少ない」52.4%、「条件に合う人が来ない」42.9%。数と質の両方に手を打つ必要がある", ""),

    ("2", "3つの経路を併用して比べている", "エージェント35%・総合34%・特化33%。特化サイトに来る人の84%は有資格者", ""),
    ("2", "比べる軸", "転職理由の1位は給与。軸は地域・給与・休日・資格との適合・資格取得支援・相談のしやすさ", ""),
    ("2", "検索は一年中ある", "山は3月、谷は12月だが差は約1.4倍", ""),
    ("2", "前後検索①「カーディーラー」", "働くことへの関心が出るが、車の購入の検索に埋もれて広告では狙いにくい", ""),
    ("2", "前後検索②「自動車整備士」", "検索前後の悩みから、LINEで届ける内容と順番を決める", "差し込み待ち：LINEヤフー前後検索「自動車整備士」未取得"),
    ("2", "他社のLINE活用", "総合大手はLINEで大規模に求職者を抱えている。整備士特化で本格的に使っている所はまだ少ない", ""),
    ("2", "まとめ", "広告を増やすより「来た人を逃さない受け皿」が効く", ""),

    ("3", "貴社の強み", "オートバックスグループの信頼、転職支援、資格取得支援。選ばれる材料はすでにそろっている", ""),
    ("3", "今の受け皿", "入口は「応募」か「転職支援の申込」の2つだけ。比べている途中の人が残る場所がない", ""),
    ("3", "今の数字", "整備士の応募単価は6万円（目標は2〜2.5万円）。登録はほぼ発生しておらず、登録と応募を分けて見られていない", ""),
    ("3", "できること・できないこと", "市場の倍率や広告単価は変えられない。来た人を残す・登録を軽くする・応募まで追うことはできる", ""),

    ("4", "最適解は、実績のある型をお手本にすること", "お手本はタウンワーク。リッチメニューで条件を選んですぐ求人を探せるようにし、応募数が79%増えた（LINEヤフー公式）", ""),
    ("4", "貴社向けに変えるところ", "①LINEだけで登録が完了する（タウンワークはリクルートIDが必要）②条件は整備士向け③LINEを通らずに応募した人も通知メッセージでLINEに集める", ""),
    ("4", "業界課題と施策の対応表", "課題一つひとつに、だから何をするかを対応させる", ""),
    ("4", "8フェーズ", "取りこぼしが起きている「サイト離脱」と「登録から応募まで」にLINEを置く", ""),
    ("4", "CVの置き方", "CV①は「LINEで登録」で完了（LINEアカウントが会員証になる）、CV②は応募。職種・資格・エリアは登録後にLINEで聞く", ""),
    ("4", "規模の近い参考", "特化型の求人サイト（保育box）では、LINEの応募率がメルマガの5倍（1.25%と0.25%）", ""),
    ("4", "サイトリニューアルの要件3点", "①「LINEで登録・ログイン」②会員名簿にLINEの番号を保存し、条件で選んで送れるようにする ③条件別の検索結果ページ", ""),
    ("4", "LINE連携で必要なこと", "開発はDYMで対応可能（都度見積もり）。貴社には取得する項目と運用ルールを決めていただく", ""),
    ("4", "施策の全体像", "広告で呼び込み、サイトで受け止め、LINEで残して登録・応募・面談まで運ぶ", ""),

    ("5", "友だち追加の動線一覧", "「離脱防止」「LINEで登録ボタン」「完了画面」の3本", ""),
    ("5", "動線00 離脱防止ポップアップ", "帰ろうとした人に「新着求人をLINEで受け取る」を出し、応募しなかった人も残す", ""),
    ("5", "動線01「LINEで登録」ボタン", "求人詳細と登録ページに置く。登録と同時に友だち追加の案内も出るので、CV①と友だちが一度に増える", ""),
    ("5", "Profile+で入力を減らす", "氏名・電話番号・生年月日などが自動で入り、確認して押すだけで登録が終わる（利用には申請と契約が必要）", ""),
    ("5", "Profile+と通知メッセージの組み合わせ", "LINEから来た人はProfile+で登録してLINEで連絡する。LINEを通らずに応募した人には通知メッセージで「応募受付完了」を届け、その場で友だち追加してもらう", ""),
    ("5", "動線02 完了画面", "応募後の連絡をLINEに寄せ、不採用や辞退のときにも別の求人を案内できる状態を残す", ""),
    ("5", "あいさつメッセージ", "全員共通。1通目で「LINEで登録」、2通目で「すぐ転職したい？」を聞く", ""),
    ("5", "リッチメニュー", "職種・地域・条件をワンタップで検索できるようにし、未経験・資格取得支援の入口も置く", ""),
    ("5", "「すぐ転職したい？」の分岐", "すぐの人は求人から応募へ。「まだ」の人はタグで絞り、興味のある内容だけ届けて配信数を抑える", ""),
    ("5", "登録から応募までのステップ配信", "迷う期間をステップ配信で追う", ""),
    ("5", "条件別の新着求人配信", "会員名簿の条件（職種・資格・エリア）で選んで、合う新着求人だけを届ける", ""),
    ("5", "年間の企画配信と再育成", "3月の山に向けて前倒しで配信し、応募しなかった人はクリックの履歴で絞って再び声をかける", ""),
    ("5", "通知メッセージ", "「応募受付完了」「会員登録完了」の公式のひな形を使う（販促は不可）。面談日程の連絡に使えるかは申込時に確認", ""),
    ("5", "よくある質問の自動応答", "資格・受験料・勤務地などは自動応答で返し、担当者の手間を増やさない", ""),

    ("6", "効果測定の設計", "CV①とCV②を分けて測り、登録から応募への率・面談予約・ブロック率を月ごとに見る", ""),
    ("6", "成果シミュレーション", "広告SIMの数字に合わせて、友だち数・登録・応募の見込みを示す", "差し込み待ち：DYM広告チームの広告SIM（Meta30万・リス70万）確認中"),
    ("6", "費用対効果の考え方", "投資額（初期20万・月20万＋都度見積もり）に対して、何ヶ月で回収できるかを示す", "数字はSIM確定後"),
    ("6", "スケジュール", "11〜12月に予算決定 → Profile+・通知メッセージの申請 → 4月にサイトと同時に開設 → 翌3月の山までに友だちを貯める", ""),

    ("7", "LINEヤフー公式の実績", "求人・人材（保育box・トライト・ランスタッドなど）に加え、LINEで会員登録が増えた事例（登録と同時の連携率 約9割／新規会員登録 約3.5倍）", ""),
    ("7", "次に確認していただきたいこと", "取得する項目・通知メッセージの運用・面談の体制が決まれば、構築に入れる", ""),

    ("8", "【汎用】なぜLINEか①", "ユーザーとの連絡にLINE公式アカウントを使う企業が増えている", ""),
    ("8", "【汎用】なぜLINEか②", "若い層を中心に、電話やメールを避ける動きがある", ""),
    ("8", "【汎用】LINE施策における重要な考え方", "新規の友だちを増やす動線と、既存の友だちへの配信を、両方整える", ""),
    ("8", "【汎用】施策展開図", "「友だちの新規追加」と「友だちの育成・応募への着地」の2つの動線で組む", ""),
    ("8", "【汎用】改善モデル①　友だち・リードを増やす", "広告・サイトの集客地点に「LINEで登録」を加え、応募まで行かない人もリストに残す", ""),
    ("8", "【汎用】改善モデル②　面談に引き上げる", "友だちになった人を、まず話を聞くだけの面談へつなぐ", ""),
    ("8", "【汎用】弊社の運用プラン①　サービス内容", "予算型コンサル（初期20万・月20万）で、構築から配信・効果測定・改善までを行う", ""),
    ("8", "【汎用】弊社の運用プラン②　都度発注", "LINE連携の開発・Profile+・通知メッセージ・離脱防止ツールなどは、必要なときに都度お見積もり", ""),
    ("8", "【汎用】弊社体制図", "貴社の窓口と、弊社の担当（運用・配信・分析・開発）の役割分担を示す", ""),
    ("0", "裏表紙", "株式会社DYM", ""),
]


def set_font(run, size, bold=False, color=GRAY):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    for tag in ("latin", "ea", "cs"):
        el = rPr.makeelement(f"{{http://schemas.openxmlformats.org/drawingml/2006/main}}{tag}", {"typeface": FONT})
        rPr.append(el)


def add_textbox(slide, l, t, w, h, text, size, bold=False, color=GRAY, align=PP_ALIGN.LEFT, anchor=None):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    if anchor:
        tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_font(run, size, bold, color)
    return box


def add_rect(slide, l, t, w, h, fill, line=None):
    from pptx.enum.shapes import MSO_SHAPE
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line:
        shp.line.color.rgb = line
        shp.line.width = Pt(0.75)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def build():
    prs = Presentation()
    prs.slide_width = SW
    prs.slide_height = SH
    blank = prs.slide_layouts[6]

    total = len(SLIDES)
    for i, (chap, title, summary, note) in enumerate(SLIDES, 1):
        slide = prs.slides.add_slide(blank)
        # 背景
        add_rect(slide, 0, 0, SW, SH, WHITE)

        if title in ("表紙", "裏表紙"):
            add_rect(slide, 0, 0, SW, SH, NAVY)
            add_textbox(slide, Inches(0.8), Inches(3.0), Inches(9.2), Inches(1.0),
                        summary, 30, True, WHITE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
            add_textbox(slide, Inches(0.8), Inches(4.0), Inches(9.2), Inches(0.5),
                        "構成案（簡易スケルトン）" if title == "表紙" else "",
                        14, False, RGBColor(0xCC, 0xCC, 0xCC), PP_ALIGN.CENTER)
            continue

        # 章タグ
        add_rect(slide, Inches(0.5), Inches(0.35), Inches(1.4), Inches(0.4), NAVY)
        add_textbox(slide, Inches(0.5), Inches(0.35), Inches(1.4), Inches(0.4),
                    f"{chap}章" if chap != "0" else "導入", 12, True, WHITE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)

        # ページ番号
        add_textbox(slide, Inches(9.6), Inches(0.35), Inches(0.9), Inches(0.4),
                    f"{i}/{total}", 11, False, GRAY, PP_ALIGN.RIGHT)

        # タイトル
        add_textbox(slide, Inches(0.5), Inches(0.9), Inches(9.8), Inches(0.7),
                    title, 20, True, TITLE_NAVY)

        # 区切り線
        add_rect(slide, Inches(0.5), Inches(1.65), Inches(9.8), Pt(1.5), NAVY)

        # 本文カード（要旨）
        add_rect(slide, Inches(0.5), Inches(2.0), Inches(9.8), Inches(3.2), LIGHT_BG, BORDER)
        add_textbox(slide, Inches(0.85), Inches(2.25), Inches(9.1), Inches(2.7),
                    summary, 16, False, GRAY, PP_ALIGN.LEFT, MSO_ANCHOR.TOP)

        # 差し込み待ち等の注記
        if note:
            add_rect(slide, Inches(0.5), Inches(5.4), Inches(9.8), Inches(0.55),
                     RGBColor(0xFF, 0xF3, 0xCD), RGBColor(0xE0, 0xB0, 0x00))
            add_textbox(slide, Inches(0.7), Inches(5.4), Inches(9.4), Inches(0.55),
                        f"⚠ {note}", 12, True, RGBColor(0x99, 0x66, 0x00), PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)

    out = "20260928_株式会社チェングロウス御中_LINE公式アカウント運用のご提案_構成案.pptx"
    prs.save(out)
    print(f"saved: {out} ({total} slides)")


if __name__ == "__main__":
    build()
