# -*- coding: utf-8 -*-
"""サンクスLINE誘導・業界別 ver3.1 → ver3.2（2026-10-07 殿村さん指示「二枚目は500%この画像」）。
2枚目「サンクスLINE｜機能・費用」を、DYM共通資料31枚目の見た目（殿村さん提示の画像）に忠実に組み直す。他のページは触らない。
  左上：黒帯「サンクスLINE誘導」＋紺「初期：15万円」「月額3万円〜/+設定地点追加」／右上：「フォーム送信後のLINE誘導を自動化するツール」＋赤い注記
  左：WEBサイト（フォーム入力 ▶ サンクス画面へ）／中央：吹き出し「CV完了後 約2〜3秒でLINEへ自動遷移」→太い矢印→緑「DYM開発ツール」
  右：動作イメージ（LINE未追加の場合＝友だち登録ページへ遷移する。／LINE追加済みの場合＝トーク画面へ遷移する。赤枠）
  下：薄緑の帯「フォーム入力内容をLINEトーク入力画面へ引き継ぎ可能」＋緑の点線矢印
このページだけリード行とグレーの区切り線を外す（画像にないため）。
  python3 _build/patch_thanks_line_gyokai_v32_s2.py <業界別ver3.1.pptx> [出力]
"""
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Cm, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

sys.path.insert(0, str(Path(__file__).resolve().parent))
from pptx_parts import *  # noqa

ROOT = Path(__file__).resolve().parent.parent
IMG = ROOT / "_images"
OUT = Path(sys.argv[2]) if len(sys.argv) > 2 else ROOT / "20261007_サンクスLINE誘導のご提案_業界別ver3.2.pptx"
BLACK, RED, YELLOW = "000000", "C00000", "FFD966"
LBLUE, DGRAY, MGRAY = "C9D8EE", "595959", "BFBFBF"
LIME, LIME_BG = "92D050", "E2F7E1"
OFF = "F2F2F2"

prs = Presentation(sys.argv[1])
s = prs.slides[1]
assert s.shapes[0].text_frame.text.startswith("サンクスLINE｜機能・費用"), "2枚目が機能・費用ではない"
for sh in list(s.shapes):
    if sh.name != "TextBox 1":
        sh._element.getparent().remove(sh._element)


def pic(name, x, y, h):
    im_w, im_h = {"form": (351, 479), "thanks": (252, 448), "add": (296, 526), "talk": (296, 526)}[name]
    w = h * im_w / im_h
    p = s.shapes.add_picture(str(IMG / f"thanks_tool_{name}.png"), Cm(x), Cm(y), Cm(w), Cm(h))
    p.line.color.rgb = rgb(LGRAY)
    p.line.width = Pt(0.5)
    return w


# ---- 左上：名称と費用 ----
shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 2.2, 2.4, 9.4, 1.3, [("サンクスLINE誘導", 14, True, WHITE, 0)], fill=BLACK, adj=0.2)
shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 2.2, 3.9, 4.5, 1.5, [("初期：15万円", 14, True, WHITE, 0)], fill=NAVY, adj=0.15)
shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 7.1, 3.9, 4.5, 1.5, [("月額3万円〜", 13, True, WHITE, 0), ("+設定地点追加", 10, True, YELLOW, 0)], fill=NAVY,
      adj=0.15, margins=(0.1, 0, 0.1, 0))
# ---- 右上：一言と注記 ----
label(s, 13.0, 2.4, 18.7, 1.4, [("フォーム送信後のLINE誘導を自動化するツール", 18, False, INK, 0)], anchor=MSO_ANCHOR.MIDDLE)
label(s, 13.0, 3.85, 18.7, 0.7, [("※既存APIツールとの併用についても、問題なく稼働が可能。", 9.5, False, RED, 0)], anchor=MSO_ANCHOR.MIDDLE)

# ---- 左：WEBサイト（フォーム ▶ サンクス画面）----
rect(s, 4.0, 5.6, 6.0, 0.75, OFF, [("WEBサイト", 11, True, GRAY, 0)])
PH, PY = 7.0, 6.8
w_form = PH * 351 / 479
w_th = PH * 252 / 448
x_form, x_th = 2.6, 9.0
shape(s, MSO_SHAPE.RECTANGLE, 2.2, 6.45, x_th + w_th + 0.4 - 2.2, 8.7, None, fill=WHITE, line=LBLUE, lw=1.0)
pic("form", x_form, PY, PH)
rect(s, x_form, PY + PH + 0.2, w_form, 0.8, MGRAY, [("フォーム入力", 11, True, WHITE, 0)])
shape(s, MSO_SHAPE.CHEVRON, 8.0, PY + PH / 2 - 0.5, 0.75, 1.0, None, fill=NAVY, adj=0.5)
pic("thanks", x_th, PY, PH)
label(s, x_th + 0.2, PY + 0.42, w_th - 0.4, 0.55, [("WEBサイト名", 8, True, GRAY, 0)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
rect(s, x_th, PY + PH + 0.2, w_th, 0.8, DGRAY, [("サンクス画面へ", 11, True, WHITE, 0)])

# ---- 中央：吹き出し → 太い矢印 → DYM開発ツール ----
co = shape(s, MSO_SHAPE.RECTANGULAR_CALLOUT, 14.0, 7.3, 5.2, 2.6,
           [("CV完了後", 12, False, INK, 0), ("約2〜3秒で", 12, False, INK, 0), ("LINEへ自動遷移", 12, False, INK, 0)],
           fill=WHITE, line=LGRAY, lw=1.0, align=PP_ALIGN.LEFT, margins=(0.4, 0.1, 0.2, 0.1))
co.adjustments[0], co.adjustments[1] = 0.0, 0.85
shape(s, MSO_SHAPE.RIGHT_ARROW, 13.8, 10.6, 5.6, 1.9, None, fill=NAVY, adj=0.5)
rect(s, 19.7, 8.0, 2.6, 6.5, LIME, [(t, 12, True, INK, 0) for t in ("DYM", "開発", "ツール")])

# ---- 右：動作イメージ ----
RX, RW = 23.4, 8.3
rect(s, RX, 5.6, RW, 0.75, NAVY, [("動作イメージ", 11, True, WHITE, 0)])
cases = [(6.55, 4.3, "LINE未追加の場合", ["友だち登録ページへ", "遷移する。"], "add", 4.0),
         (11.05, 4.5, "LINE追加済みの場合", ["トーク画面へ遷移する。"], "talk", 4.2)]
for fy, fh, nm, txt, img, ih in cases:
    shape(s, MSO_SHAPE.RECTANGLE, RX, fy, RW, fh, None, fill=WHITE, line=LGRAY, lw=1.0)
    rect(s, RX + 0.2, fy + 0.2, 4.6, 0.65, OFF, [(nm, 10, True, INK, 0)])
    label(s, RX + 0.2, fy + 1.0, 4.6, fh - 1.2, [(t, 11, False, INK, 0) for t in txt], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    iw = ih * 296 / 526
    px = RX + RW - 0.2 - iw
    pic(img, px, fy + 0.15, ih)
    if img == "add":
        label(s, px + 0.1, fy + 0.15 + ih * 0.245, iw - 0.2, ih * 0.21, [("ロゴ", 8, True, GRAY, 0)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        label(s, px, fy + 0.15 + ih * 0.475, iw, 0.4, [("アカウント名", 7, True, INK, 0)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    else:
        tb = label(s, px + iw * 0.22, fy + 0.15 + ih * 0.045, iw * 0.42, 0.35, [("アカウント名", 5.5, True, INK, 0)], align=PP_ALIGN.CENTER,
                   anchor=MSO_ANCHOR.MIDDLE)
        tb.text_frame.word_wrap = False
        # 吹き出し（入力内容の引き継ぎ）を赤枠で強調
        shape(s, MSO_SHAPE.RECTANGLE, px + iw * 0.27, fy + 0.15 + ih * 0.26, iw * 0.58, ih * 0.265, None, fill=None, line=RED, lw=1.5)
        talk_box = (px + iw * 0.27, fy + 0.15 + ih * 0.26 + ih * 0.265)
# DYM開発ツール → 2つの動作（緑の点線）
arrow_line(s, 22.3, 9.6, RX - 0.1, 7.3, GREEN, 1.5, dash=True)
arrow_line(s, 22.3, 12.9, RX - 0.1, 12.0, GREEN, 1.5, dash=True)

# ---- 下：入力内容の引き継ぎ ----
shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 7.0, 15.6, 15.5, 2.0,
      [("フォーム入力内容を", 14, False, INK, 0), ("LINEトーク入力画面へ引き継ぎ可能", 14, False, INK, 0)], fill=LIME_BG, line=LIME, lw=1.0, adj=0.3)
arrow_line(s, 7.0, 16.2, x_form + 0.6, PY + PH + 1.05, GREEN, 1.5, dash=True)
arrow_line(s, 22.5, 16.3, talk_box[0] + 0.3, talk_box[1] + 0.05, GREEN, 1.5, dash=True)

prs.save(str(OUT))
print("saved:", OUT.name, len(prs.slides), "枚")
