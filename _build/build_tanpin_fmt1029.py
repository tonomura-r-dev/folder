# -*- coding: utf-8 -*-
"""単品通販業界（殿村さんPC保存版・36枚）を、10/29の提案FMT（MAX版・53枚）の章立てにそろえる（2026-10-30 殿村さん「FMTに合わせて修正して」）。

    python -m _build.build_tanpin_fmt1029 <PC保存版.pptx> <FMT最新.pptx> <FMT具体施策枠のみ.pptx> <塾 新FMT版.pptx> <出力.pptx>

  形はジムの新FMT版（_build/build_fitness_fmt1029.py）と同じ。
  ・01 施策背景：FMTの弊社紹介〜ポジショニング＋市場全体①〜③・単品通販の課題①〜⑤＋FMTのLINEOA公式基礎・実例紹介＋LINEOA実績
  ・02 LINE施策全体図：FMTのファネル＋LINEでやるべきこと①②・全体設計①〜③＋FMTの施策パターン・全体図・施策一覧
  ・03 施策詳細：FMTの具体施策のイメージ8枚（単品通販の中身で埋める）＋構築①〜⑤・配信設計①〜⑥・改善①②・成果・締め①②
  ・04〜06：FMTのページ。旧の費用プラン・運用スケジュール・サポート体制は裏表紙の後ろのAppendixへ
  ・旧の資料アジェンダと旧裏表紙は外す。ページの中の「S◯」は新しい並びの「P◯」に振り直す。
"""
import re
import sys

from pptx import Presentation

import _build.build_juku_fmt1029 as J0
from _build.build_juku_fmt1029 import fill_gutai, renumber
from _build.patch_tanaka_v5 import copy_slide

GUTAI = [  # 札, [入口, LINE, 配信, ゴール], (目的, 対象, 中身, 見る数字)  右の文は「|」で改行
    ("LINEの友だちへ",
     [("購入者・既存顧客", "同梱物QR／購入完了画面／マイページ"), ("友だち追加", "「解約・休止もLINEでできます」"),
      ("あいさつメッセージ", "初期アンケート（使用頻度）"), "友だちとしてつながり続ける"],
     ("購入した方を、確実に|LINEの友だちにする",
      "初回購入した方／|定期で購入している方",
      "同梱物QR（主役）／|購入完了画面・マイページ／|通知メッセージで既存顧客へ",
      "友だち追加数／|購入者のうちの追加率")),
    ("露出の強化",
     ["友だち全体", ("企画配信", "季節・新商品・限定セット"),
      "年間の企画投稿カレンダーに沿って案内", "まとめ買い・関連商品の購入"],
     ("季節や新商品の情報を届け、|接触の量を増やす",
      "友だち全体|（定期・単品で購入の方を含む）",
      "年間の企画投稿カレンダー／|新商品・限定セットの案内／|売り込みにならない情報の配信",
      "配信到達数・開封率／|企画配信からの購入数")),
    ("CVの最大化",
     ["初回購入した方", ("使用頻度のタグ", "内容量÷1日の使用量で、切れる頃を逆算"),
      ("購入後の配信", "使用フォロー → 切れる頃のお知らせ"), "2回目の購入（F2）・定期引上"],
     ("初回で離脱する方を減らし、|2回目の購入につなげる",
      "初回購入した方",
      "効果実感までの目安を伝える配信／|「そろそろ切れる頃です」の案内／|F2の後に定期・まとめ買いの案内",
      "F2転換率／定期引上率")),
    ("CPOの改善（CV質）",
     ["定期で購入している方", ("LINE ID連携", "定期の回数・次回発送日が分かる"),
      ("次回発送の数日前の配信", "「今回はスキップできます」"), "定期の継続"],
     ("定期を続けやすくし、|お客様1人あたりの成果を高める",
      "定期で購入している方",
      "次回発送前のスキップの案内／|お届け間隔・数量の変更／|節目（3回目・6回目）の継続特典",
      "月次解約率／継続回数")),
    ("休眠顧客の有効化",
     ["友だちでない既存顧客・解約した方", "通知メッセージ（電話番号でLINEに届く）",
      ("発送・決済エラー・再入荷のお知らせ", "通知の中に友だち追加の導線"), "友だち追加・購入の再開"],
     ("LINEで届いていない既存顧客に、|もう一度届ける",
      "購入したが友だちでない方／|解約した方・決済エラーで|止まっている方",
      "通知の中に友だち追加の導線／|決済エラーの放置を防ぐ／|同梱物QRと合わせて到達率を上げる",
      "通知からの友だち追加数／|購入の再開数")),
    ("マネタイズ最大化",
     ["解約・休止を相談した方", ("解約アンケート", "理由を1タップで"),
      ("理由に応じた選択肢", "スキップ・周期変更・プラン変更・休止"), "継続・休止（解約も止めない）"],
     ("解約の手前で選択肢を出し、|続けられる方には続けていただく",
      "リッチメニュー「解約・休止の|ご相談」を押した方",
      "余っている→スキップ・周期変更／|高い→プラン・数量の変更／|一時的な事情→休止",
      "解約の手前で片づいた件数／|月次解約率")),
    ("リピート施策",
     ["定期・単品で購入している方", ("購入内容・回数のタグ", "LINE ID連携で自動で付く"),
      ("節目・季節の配信", "継続特典・まとめ買い・関連商品"), "継続・関連商品の購入"],
     ("購入の回数と内容に合わせ、|次の購入につなげる",
      "定期・単品で購入している方",
      "3回目・6回目の継続特典／|使い方の見直し・成分の解説／|関連商品・まとめ買いの案内",
      "継続回数／関連商品の購入数")),
    ("コミュニティ運営",
     ["長く続けている方", ("継続回数のタグ", "リッチメニュー「使い方」に会員向けの入口"),
      "使い方のコツ・よくある質問", "継続・紹介"],
     ("長く続けている方とのつながりを|強くし、紹介につなげる",
      "定期を長く続けている方",
      "長く続けている方への限定案内／|使い方のコツ・よくある質問／|LINEの転送でお友だちに紹介",
      "継続回数／|紹介からの友だち追加数")),
]


def renumber_s(s, mapping):
    """単品通販のページの中の「S◯」「S◯-◯」を新しい番号「P◯」「P◯〜P◯」に"""
    for t in s._element.iter("{%s}t" % J0.A_NS):
        if t.text and re.search(r"S\d\d", t.text):
            def rep(m):
                a, b = int(m.group(1)), m.group(2)
                assert a in mapping, ("消えたページを指している", a, t.text)
                if b:
                    assert int(b) in mapping, ("消えたページを指している", b, t.text)
                    return "P%d〜P%d" % (mapping[a], mapping[int(b)])
                return "P%d" % mapping[a]
            t.text = re.sub(r"S(\d\d)(?:-(\d\d))?", rep, t.text)


def main(src_path, fmt_path, fmt_empty_path, juku_path, out):
    J0.TEXT_FIX = [("貴社保有リスト をターゲット", "貴社保有リストをターゲット"),
                   ("転職ユーザーのカスタマージャーニー", "購入者のカスタマージャーニー")]
    prs = Presentation(src_path)
    fmt, emp, juku = Presentation(fmt_path), Presentation(fmt_empty_path), Presentation(juku_path)
    assert len(prs.slides) == 36 and len(fmt.slides) == 53 and len(emp.slides) == 53
    G = {i + 1: s for i, s in enumerate(prs.slides)}

    def add(src):
        s = copy_slide(prs, src)
        J0.fmt_fix(s)
        return s

    F = {n: add(fmt.slides[n - 1]) for n in list(range(2, 7)) + list(range(8, 54)) if not 18 <= n <= 25}
    E = {n: add(emp.slides[n - 1]) for n in range(18, 26)}
    for n, page in zip(range(18, 26), GUTAI):
        fill_gutai(E[n], page)
    apx = copy_slide(prs, juku.slides[74])           # 「アジェンダ／Appendix（参考資料）」
    assert "Appendix" in "".join(sh.text_frame.text for sh in apx.shapes if sh.has_text_frame)

    order = ([G[1]] + [F[n] for n in range(2, 7)] + [G[n] for n in range(3, 11)] + [F[8], F[9], F[10], G[35]]
             + [F[11], F[12]] + [G[n] for n in range(11, 16)] + [F[13], F[14], F[15], F[16]]
             + [F[17]] + [E[n] for n in range(18, 26)] + [G[n] for n in range(16, 30)] + [G[33], G[34]]
             + [F[n] for n in range(29, 54)]
             + [apx, G[30], G[31], G[32]])
    assert len(order) == 83, len(order)
    mapping = {}
    for i, s in enumerate(order, 1):
        for n, gs in G.items():
            if gs is s:
                mapping[n] = i
    for n, s in G.items():
        if n in mapping:
            renumber_s(s, mapping)

    lst = prs.slides._sldIdLst
    ids = {id(s._element): sid for s, sid in zip(prs.slides, list(lst))}
    keep = [ids[id(s._element)] for s in order]
    keep_set = set(id(x) for x in keep)
    for sid in list(lst):
        lst.remove(sid)
        if id(sid) not in keep_set:
            prs.part.drop_rel(sid.rId)
    for sid in keep:
        lst.append(sid)
    assert len(prs.slides) == 83
    prs.save(out)
    print("saved:", out)


if __name__ == "__main__":
    main(*sys.argv[1:6])
