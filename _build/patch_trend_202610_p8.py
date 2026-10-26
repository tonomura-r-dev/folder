"""
10月号 P.8（ユニット単位の統計情報）を直す（2026-10-26 殿村さん指示）。
- AS-IS ▶ TO-BE と注意カード2枚をやめて、「変更点」カード1枚＋下に注意の※1文
- カードはP.3の「変更点」カードを複製して使う。幅は一番長い行に合わせ、左右・上下中央
- 分類タグ＝開発 LINE API、逆三角形＝伸長
使い方: python3 _build/patch_trend_202610_p8.py <pptx>
"""
import sys
import copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.oxml.ns import qn
import trend_report_engine as engine

F = sys.argv[1]
prs = Presentation(F)
slides = list(prs.slides)
s3, s8 = slides[2], slides[7]
assert "ユニット単位の統計情報" in s8.shapes.title.text if s8.shapes.title else True


def text(sh):
    return sh.text_frame.text.strip() if sh.has_text_frame else ""


assert any("ユニット単位の統計情報" in text(x) for x in s8.shapes)
head = next(x for x in s8.shapes if text(x) == "■ 詳細")
area_top = head.top + head.height + Inches(0.08)
for x in [x for x in s8.shapes if x.top >= head.top + head.height - Inches(0.02) and x.top < Inches(7.0)]:
    x._element.getparent().remove(x._element)

# P.3 の「変更点」カードを複製
src = next(x for x in s3.shapes if x.top > Inches(3.0) and text(x).startswith("変更点"))
el = copy.deepcopy(src._element)
s8.shapes._spTree.append(el)
cn = el.find(".//" + qn("p:cNvPr"))
cn.set("id", str(max(int(e.get("id")) for e in s8._element.iter(qn("p:cNvPr"))) + 1))
card = next(x for x in s8.shapes if x._element is el)

BODY = [
    "・ユニットごとの統計情報を、LINE通知メッセージでもAPIで取得できるようになった。",
    "・ユニットごとに、開封数・URLタップ数などを確認できる。",
]
p0, p1, p2 = card.text_frame.paragraphs
for p, t in ((p1, BODY[0]), (p2, BODY[1])):
    p.runs[0].text = t
    for r in list(p.runs[1:]):
        r._r.getparent().remove(r._r)
    for br in p._p.findall(qn("a:br")):
        br.getparent().remove(br)

NOTE = ["※ ユニット名は送るときに付ける（1通に1つ。後から付けたり変えたりはできない。月1,000種類まで）。",
        "※ 20人未満の数字は表示されない。送信数は取れないため、開封率の分母は自社で持っておく。"]


def text_w(s, pt):
    return sum((0.55 if ord(c) < 0x2000 else 1.0) for c in s) * pt / 72


tf = card.text_frame
w = max(text_w(t, 14) for t in BODY) + Emu(tf.margin_left + tf.margin_right).inches + 0.45
w = min(w, 9.64)
h = 0.32 + 20 * 1.3 / 72 + 2 * (14 * 1.45 + 6) / 72
note_h = 0.62
total = h + 0.15 + note_h
top = Emu(area_top).inches + (7.05 - Emu(area_top).inches - total) / 2
SW = Emu(prs.slide_width).inches
card.left, card.top, card.width, card.height = Inches((SW - w) / 2), Inches(top), Inches(w), Inches(h)

nb = s8.shapes.add_textbox(Inches((SW - w) / 2), Inches(top + h + 0.15), Inches(w), Inches(note_h))
nb.text_frame.word_wrap = True
for i, t in enumerate(NOTE):
    p = nb.text_frame.paragraphs[0] if i == 0 else nb.text_frame.add_paragraph()
    r = p.add_run()
    r.text = t
    r.font.size, r.font.color.rgb = Pt(12), engine.BODY_GRAY

engine.set_tag_highlight(s8, "開発\nLINE API")
engine.set_funnel_highlight(s8, ["伸長"])
engine.force_font(prs)
prs.save(F)
print("saved", F, round(w, 2), round(h, 2))
