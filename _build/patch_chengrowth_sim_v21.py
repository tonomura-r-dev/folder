# SIM ver2.0 に「SIM3_LINE登録2%」（通知メッセ無・LINEで登録=UU×2.0%）と比較シートを追加して ver2.1 で保存
import sys
from openpyxl import load_workbook
from openpyxl.workbook.properties import CalcProperties
SRC,OUT=sys.argv[1],sys.argv[2]
wb=load_workbook(SRC)
s3=wb.copy_worksheet(wb['SIM2_通知メッセ無']); s3.title='SIM3_LINE登録2%'
c3=wb.copy_worksheet(wb['SIM2_比較']); c3.title='SIM3_比較'
for row in c3.iter_rows():
    for c in row:
        if isinstance(c.value,str) and "'SIM2_通知メッセ無'!" in c.value:
            c.value=c.value.replace("'SIM2_通知メッセ無'!","'SIM3_LINE登録2%'!")
s3['A4']='株式会社チェングロウス御中_LINEOA施策提案【SIM】（通知メッセージ無・LINEで登録2.0%）'
ramp=[0.62,0.81,0.93,1.00,1.03,1.05]
cv=[1,2,3,4,5,6]   # 配信経由の応募（仮。殿村さんが修正）
for i,c in enumerate('EFGHIJ'):
    s3[f'{c}26']=f'=SUM({c}22:{c}25)+{c}14*0.02*{ramp[i]}'
    s3[f'{c}36']=cv[i]
s3['N26']='＋サイト内「LINEで登録」：UU×2.0%×立ち上げ係数（Profile+で入力ほぼゼロになる前提）'
order=['SIM1_通知メッセ有','SIM1_比較','SIM2_通知メッセ無','SIM2_比較','SIM3_LINE登録2%','SIM3_比較','SIM考え方']
wb._sheets=[wb[n] for n in order]
wb.calculation=CalcProperties(fullCalcOnLoad=True)
wb.save(OUT)
