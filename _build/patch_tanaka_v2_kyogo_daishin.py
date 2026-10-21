# -*- coding: utf-8 -*-
"""タナカホーム v2：競合の2ページを「大進建設がメイン・一条工務店とセキスイハイムはサブ」に組み直す（2026-10-21 殿村さん）。

    python -m _build.patch_tanaka_v2_kyogo_daishin <殿村さんのPC保存版v2.pptx> <出力.pptx>

  10/19の「大進建設」のLINEは群馬の同名の別会社だった。青森の大進建設（daishink.co.jp）のLINEを、
  サイトのポップアップから殿村さんが追加して撮った画面（_data/tanaka/競合LINE_20261021/）に差し替える。
  画面に出る殿村さんの名前は灰色で伏せる（_images/tanaka_line_daishin_*.png）。
  5ページ「競合のLINE運用」：左＝大進建設（入口と特典／あいさつ／キャンペーンとリッチメニュー）を大きく、
    右＝ほかの大手（一条工務店の画面2枚＋1行、セキスイハイム1行）を小さく。追加後の配信の枠は外す（2社とも届いていない）。
  6ページ「競合の取り組みに対して」：表をやめ、真ん中に項目（きっかけ・入口・あいさつ・リッチメニュー・どこまで追うか）、
    左に大進建設、右にタナカホーム（弊社の運用でつくる形）を並べる。下の「ポイントは…」の枠は殿村さんの版のまま。
  ほかの29枚は触らない。
"""
import copy
import io
import sys
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches

from _build.patch_yokohama_shimizu_deck_plans import set_text

ROOT = Path(__file__).resolve().parent.parent
SRC_DIR = ROOT / "_data/tanaka/競合LINE_20261021"
AISATSU = ROOT / "_images/tanaka_line_daishin_aisatsu.png"
MENU = ROOT / "_images/tanaka_line_daishin_menu.png"
NAME_BOXES = [(168, 405, 538, 475), (612, 1700, 778, 1762), (170, 1762, 392, 1822)]   # 「殿村亮太@DYM」の位置（1170×2532）
NAVY, GRAY, INK = "1F285A", "7F7F7F", "333333"

ROWS = [   # 項目, 大進建設, タナカホーム（弊社の運用でつくる形）
    ("友だちになるきっかけ", "LINE登録で最大4万円分のギフトカード",
     "来場特典に加え、\n施工事例・ルームツアーなどの情報"),
    ("友だちの入口", "サイトのポップアップ",
     "サイトのポップアップ＋\nサイトを閉じるとき・資料請求の完了直後"),
    ("あいさつ", "お名前入りのあいさつと、キャンペーンの案内。\n「まずは資料請求から」",
     "予約の案内に加え、\n平屋か2階建てか・予算の簡易診断へ"),
    ("リッチメニュー", "キャンペーン・資料請求・モデルハウス・\nイベント・土地情報",
     "今の予約・ホームページ・SNSに、\n施工事例・ルームツアー・資金相談を足す"),
    ("どこまで追うか", "特典で、資料請求・予約来場から、\n来場後の事前審査まで",
     "弊社が運用を代行して、\n資料請求の前から契約まで追う設計にする"),
]


def mask_images():
    im = Image.open(SRC_DIR / "大進建設_あいさつ.png").convert("RGB")
    d = ImageDraw.Draw(im)
    for b in NAME_BOXES:
        d.rounded_rectangle(b, radius=10, fill=(200, 200, 200))
    im.save(AISATSU)
    Image.open(SRC_DIR / "大進建設_キャンペーン_リッチメニュー.png").convert("RGB").save(MENU)


def by_id(s, sid):
    return next(sh for sh in s.shapes if sh.shape_id == sid)


def place(sh, x, y, w, h):
    sh.left, sh.top, sh.width, sh.height = Inches(x), Inches(y), Inches(w), Inches(h)


def remove(s, *sids):
    for sid in sids:
        el = by_id(s, sid)._element
        el.getparent().remove(el)


def clone(s, tmpl, name, text, x, y, w, h):
    el = copy.deepcopy(tmpl if not hasattr(tmpl, "_element") else tmpl._element)
    s.shapes._spTree.append(el)
    sh = s.shapes[-1]
    sh._element.find(qn("p:nvSpPr")).find(qn("p:cNvPr")).set("id", str(max(x.shape_id for x in s.shapes) + 1))
    sh.name = name
    place(sh, x, y, w, h)
    set_text(sh.text_frame, text)
    return sh


def set_image(s, pic, path, drop=True):
    blip = pic._element.find(".//" + qn("a:blip"))
    old = blip.get(qn("r:embed"))
    _, rid = s.part.get_or_add_image_part(io.BytesIO(path.read_bytes()))
    blip.set(qn("r:embed"), rid)
    if drop and old != rid:                                   # 写した画像の枠は、元の画像を使い続けるので外さない
        s.part.drop_rel(old)


def fill(sh, color):
    sp = sh._element.find(qn("p:spPr"))
    sf = sp.find(qn("a:solidFill"))
    sf.find(qn("a:srgbClr")).set("val", color)


def run_props(sh, color=None, bold=None, size=None):
    for r in sh._element.iter(qn("a:rPr")):
        if size:
            r.set("sz", str(size))
        if bold is not None:
            r.set("b", "1" if bold else "0")
        if color:
            sf = r.find(qn("a:solidFill"))
            sf.find(qn("a:srgbClr")).set("val", color)


def page5(s):
    assert by_id(s, 2).text_frame.text.startswith("競合のLINE運用")
    set_text(by_id(s, 2).text_frame, "競合のLINE運用：大進建設は特典で、資料請求から事前審査まで後押しする")
    set_text(by_id(s, 6).text_frame, "大進建設を中心に、一条工務店・セキスイハイムの友だち追加の入口とあいさつを、実際の画面で見る")

    # 左：大進建設（3列）
    place(by_id(s, 9), 0.45, 1.62, 6.1, 0.36)
    cols = [(0.45, 1.75), (2.35, 2.0), (4.5, 2.05)]
    (ax, aw), (bx, bw), (cx, cw) = cols
    place(by_id(s, 11), ax, 2.05, aw, 0.46)
    place(by_id(s, 8), ax + (aw - 1.75) / 2, 2.56, 1.75, 1.75)
    place(by_id(s, 12), ax, 4.4, aw, 0.6)
    set_text(by_id(s, 12).text_frame, "LINE登録で\n最大4万円分")
    place(by_id(s, 13), bx, 2.05, bw, 0.46)
    pic_b = by_id(s, 14)
    set_image(s, pic_b, AISATSU)
    pic_b.name = "大進建設のあいさつ"
    place(pic_b, bx + (bw - 1.45) / 2, 2.56, 1.45, 3.14)
    place(by_id(s, 15), bx, 5.76, bw, 0.78)
    set_text(by_id(s, 15).text_frame, "お名前入りのあいさつと、キャンペーンの案内")
    head_c = by_id(s, 16)
    place(head_c, cx, 2.05, cw, 0.46)
    set_text(head_c.text_frame, "リッチメニュー")
    head_c.name = "見出し リッチメニュー"
    pic_c = copy.deepcopy(pic_b._element)                     # あいさつの画像の枠を写して、中身を差し替える
    s.shapes._spTree.append(pic_c)
    pic_c = s.shapes[-1]
    pic_c._element.find(qn("p:nvPicPr")).find(qn("p:cNvPr")).set("id", str(max(x.shape_id for x in s.shapes) + 1))
    pic_c.name = "大進建設のキャンペーンとリッチメニュー"
    set_image(s, pic_c, MENU, drop=False)
    place(pic_c, cx + (cw - 1.45) / 2, 2.56, 1.45, 3.14)
    clone(s, by_id(s, 15), "画面の説明", "資料請求・モデルハウス・\nイベント・土地情報", cx - 0.05, 5.76, cw + 0.1, 0.78)
    remove(s, 17, 24, 25)

    # 右：ほかの大手（参考）
    sub = by_id(s, 10)
    place(sub, 6.8, 1.62, 3.58, 0.36)
    set_text(sub.text_frame, "ほかの大手（参考）")
    sub.name = "会社 ほかの大手"
    fill(sub, GRAY)
    place(by_id(s, 26), 6.65, 1.62, 0.01, 4.95)
    lab1 = by_id(s, 18)
    place(lab1, 6.8, 2.05, 3.58, 0.38)
    set_text(lab1.text_frame, "一条工務店")
    lab1.name = "見出し 一条工務店"
    place(by_id(s, 19), 6.85, 2.5, 0.95, 2.06)
    place(by_id(s, 22), 7.9, 2.5, 0.95, 2.06)
    place(by_id(s, 20), 8.95, 2.5, 1.43, 2.06)
    set_text(by_id(s, 20).text_frame, "1分のアンケートへ案内。メニューに性能・暮らし・実例")
    lab2 = by_id(s, 21)
    place(lab2, 6.8, 4.75, 3.58, 0.38)
    set_text(lab2.text_frame, "セキスイハイム")
    lab2.name = "見出し セキスイハイム"
    place(by_id(s, 23), 6.8, 5.2, 3.58, 0.6)
    set_text(by_id(s, 23).text_frame, "LINEで、カタログ無料診断・無料間取り作成を案内")
    set_text(by_id(s, 27).text_frame,
             "※各社のLINE公式アカウントより（大進建設 2026年10月21日、一条工務店 2026年10月19日、セキスイハイム 2026年10月16日確認）。"
             "サイトのポップアップは大進建設の公式サイトより（2026年10月16日確認）")


def page6(s, tmpl_head, tmpl_cap):
    assert by_id(s, 2).text_frame.text.startswith("競合の取り組みに対して")
    set_text(by_id(s, 6).text_frame, "特典で事前審査まで後押しする大進建設に対し、資料請求の前から契約まで追う設計で応える")
    remove(s, 9, 10, 11)                                      # 表と差し込みの枠を外す

    lx, lw, mx, mw, rx, rw = 0.55, 3.72, 4.4, 2.0, 6.53, 3.75
    h = clone(s, tmpl_head, "会社 大進建設", "大進建設", lx, 1.72, lw, 0.36)
    fill(h, GRAY)
    clone(s, tmpl_head, "会社 タナカホーム", "タナカホーム（弊社の運用でつくる形）", rx, 1.72, rw, 0.36)
    y, row_h, gap = 2.2, 0.6, 0.1
    for i, (item, daishin, tanaka) in enumerate(ROWS):
        last = i == len(ROWS) - 1
        chip = clone(s, tmpl_head, f"項目 {item}", item, mx, y + 0.08, mw, row_h - 0.16)
        if last:
            run_props(chip, size=1200)
        else:                                                 # 白地に紺の枠・紺の文字
            sp = chip._element.find(qn("p:spPr"))
            fill(chip, "FFFFFF")
            ln = sp.find(qn("a:ln"))
            for c in list(ln):
                ln.remove(c)
            ln.set("w", "12700")
            sf = ln.makeelement(qn("a:solidFill"), {})
            clr = sf.makeelement(qn("a:srgbClr"), {"val": NAVY})
            sf.append(clr)
            ln.append(sf)
            run_props(chip, color=NAVY, size=1200)
        left = clone(s, tmpl_cap, "大進建設", daishin, lx, y, lw, row_h)
        left.text_frame._txBody.find(qn("a:bodyPr")).set("anchor", "ctr")
        right = clone(s, tmpl_cap, "タナカホーム", tanaka, rx, y, rw, row_h)
        right.text_frame._txBody.find(qn("a:bodyPr")).set("anchor", "ctr")
        run_props(right, color=NAVY, bold=True)
        y += row_h + gap
    set_text(by_id(s, 12).text_frame,
             "※各社のLINE公式アカウントより（大進建設 2026年10月21日、一条工務店 2026年10月19日確認）。"
             "入口・ポップアップは各社の公式サイトより（2026年10月16日確認）")


def main(src, out):
    mask_images()
    prs = Presentation(src)
    assert len(prs.slides) == 31
    s5, s6 = prs.slides[4], prs.slides[5]
    tmpl_head = copy.deepcopy(by_id(s5, 9)._element)          # 紺の帯（会社名）と、画面の説明（12pt）を型にする
    tmpl_cap = copy.deepcopy(by_id(s5, 12)._element)
    page5(s5)
    page6(s6, tmpl_head, tmpl_cap)
    prs.save(out)
    print("saved:", out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
