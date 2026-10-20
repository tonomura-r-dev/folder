# -*- coding: utf-8 -*-
"""横浜シミズ：資料（殿村さん版・20261020・16枚）のS13・S14に、SIM ver2.0（殿村さん確定）の数字を入れる（2026-10-20）。

    python -m _build.patch_yokohama_shimizu_deck_sim <資料.pptx> <出力.pptx>

  S13（6か月の見込み）：表をSIMの行に組み替える（友だち追加・累計・ブロック率・配信できる人数・クリック数・新規応募・再応募・合計）。
    タイトルの「試算」を外して「概算お見積り｜〜」に（CLAUDE.md）。下の3つの枠と注記もSIMの前提に合わせる。
  S14（運用プラン・費用）：月額 165,000円〜／初期 295,000円〜 に。離脱防止ポップアップ・LINE公式アカウント開設の行を足す。
    「月6本」→「企画配信 月4回・ステップ配信 4通」。
  ほかの14枚は触らない。数字は小数を出さない（四捨五入した整数）。リードの2行目の色など、元の書式はそのまま使う。
"""
import copy
import sys

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches

MONTHS = 6
# SIM ver2.0（殿村さん確定）を四捨五入した整数
FRIENDS_ADD = [164] * 6
FRIENDS = [164, 328, 492, 656, 820, 984]
BLOCK = ["32%", "33%", "34%", "34%", "35%", "36%"]
ACTIVE = [112, 221, 327, 430, 531, 628]
CLICKS = [107, 211, 313, 330, 407, 481]
NEW_CV = [2, 2, 4, 5, 5, 6]
RE_CV = [0, 0, 1, 2, 3, 3]
ALL_CV = [a + b for a, b in zip(NEW_CV, RE_CV)]


def by_id(s, sid):
    return next(sh for sh in s.shapes if sh.shape_id == sid)


def set_runs_text(p, text):
    """段落の最初のrunの書式で text にする（ほかのrunは消す）"""
    runs = p._p.findall(qn("a:r"))
    for r in runs[1:]:
        p._p.remove(r)
    runs[0].find(qn("a:t")).text = text


def set_paras(tf, lines):
    """1段落目（見出し）と2段落目（本文）の書式を型にして、行を差し替える"""
    ps = tf.paragraphs
    head, body = ps[0]._p, ps[1]._p
    txBody = head.getparent()
    for p in list(txBody.findall(qn("a:p"))):
        txBody.remove(p)
    for i, text in enumerate(lines):
        p = copy.deepcopy(head if i == 0 else body)
        txBody.append(p)
        runs = p.findall(qn("a:r"))
        for r in runs[1:]:
            p.remove(r)
        runs[0].find(qn("a:t")).text = text


def set_cell(cell, text):
    ps = cell.text_frame.paragraphs
    for p in ps[1:]:
        p._p.getparent().remove(p._p)
    set_runs_text(ps[0], text)


def rebuild_rows(tbl_shape, templates, data, heights):
    """templates[i] の行（a:tr）を型にして data の行を作り直す"""
    tbl = tbl_shape.table._tbl
    trs = tbl.findall(qn("a:tr"))
    proto = [copy.deepcopy(trs[i]) for i in templates]
    for tr in trs[1:]:
        tbl.remove(tr)
    for tr_proto, row, h in zip(proto, data, heights):
        tr = copy.deepcopy(tr_proto)
        tr.set("h", str(int(Inches(h))))
        tbl.append(tr)
        for tc, text in zip(tr.findall(qn("a:tc")), row):
            p = tc.find(qn("a:txBody")).findall(qn("a:p"))
            for extra in p[1:]:
                extra.getparent().remove(extra)
            runs = p[0].findall(qn("a:r"))
            for r in runs[1:]:
                p[0].remove(r)
            runs[0].find(qn("a:t")).text = text


def fmt(n):
    return f"{n:,}"


def slide13(s):
    title = by_id(s, 2)
    assert title.text_frame.text.startswith("目標・6か月の試算（SIM）")
    set_runs_text(title.text_frame.paragraphs[0], "概算お見積り｜友だち獲得から新規応募・再応募まで")
    lead = by_id(s, 3).text_frame.paragraphs
    set_runs_text(lead[0], "サイトの訪問者数（UU）月8,947をもとに、友だち獲得から新規応募・再応募までの6か月の見込みを算出。")
    set_runs_text(lead[1], "月の応募数と各率は仮に置いた数字。実数をいただければ、置き換えて算出し直す。")

    tshape = by_id(s, 4)
    grid = tshape.table._tbl.find(qn("a:tblGrid")).findall(qn("a:gridCol"))
    for gc, w in zip(grid, [1.62, 2.83] + [0.74] * 6 + [1.0]):
        gc.set("w", str(int(Inches(w))))
    rows = [
        ["友だち追加数", "サイトのLINEボタン 45＋離脱防止 91＋サンクスLINE誘導 28", *map(fmt, FRIENDS_ADD), fmt(FRIENDS[-1])],
        ["累計の友だち数", "前月 ＋ 友だち追加", *map(fmt, FRIENDS), "—"],
        ["ブロック率", "32%から毎月少しずつ上昇", *BLOCK, "—"],
        ["配信できる友だち数", "累計 ×（1 − ブロック率）", *map(fmt, ACTIVE), "—"],
        ["配信のクリック数", "企画配信 月4回＋ステップ配信 4通", *map(fmt, CLICKS), fmt(sum(CLICKS))],
        ["新規応募", "応募前に友だちになった方の応募", *map(fmt, NEW_CV), fmt(sum(NEW_CV))],
        ["再応募（勤務希望）", "一度勤務した方の、次回イベントへの応募", *map(fmt, RE_CV), fmt(sum(RE_CV))],
        ["応募の合計", "新規応募 ＋ 再応募", *map(fmt, ALL_CV), fmt(sum(ALL_CV))],
    ]
    # 型：1〜5行目は元の「新規応募数」の行（白）、6〜8行目は元の「再応募」の行（薄い紺の網かけ）
    rebuild_rows(tshape, [1, 1, 1, 1, 1, 5, 5, 5], rows, [0.46, 0.32, 0.32, 0.32, 0.32, 0.32, 0.32, 0.32])
    tshape.table.rows[0].height = Inches(0.34)
    tshape.height = sum(r.height for r in tshape.table.rows)

    y, h = 4.68, 1.36
    for sid in (5, 6, 7):
        sh = by_id(s, sid)
        sh.top, sh.height = Inches(y), Inches(h)
    set_paras(by_id(s, 5).text_frame, [
        "いただいている数値",
        "・サイトUU 8,947（月間）",
        "※月の応募数は、仮にUUの1%（89件）で計算",
    ])
    set_paras(by_id(s, 6).text_frame, [
        "仮定（弊社で設定）",
        "・サンクスLINE誘導：応募×遷移80%×追加40%",
        "・LINEボタン・離脱防止：弊社の標準の係数",
        "・応募：配信からのクリックに対する応募",
        "※採用サイトの完了画面に設置できる前提",
    ])
    set_paras(by_id(s, 7).text_frame, [
        "含めていないもの",
        "・今いる登録スタッフ（LINE公式アカウントは新規開設のため、0人から）",
        "・求人媒体経由の応募",
        "・勤務連絡の効率化（時間の削減）",
    ])
    set_runs_text(by_id(s, 8).text_frame.paragraphs[0],
                  "※見込みであり、成果を保証するものではない。1か月目＝運用開始の月（構築の完了後）。"
                  "広告経由と自然流入は区別していない。人数・件数は四捨五入。")


def slide14(s):
    assert by_id(s, 2).text_frame.text.startswith("運用プラン・費用｜")
    lead = by_id(s, 3).text_frame.paragraphs
    t = lead[0].runs[0].text
    assert "13.5万円〜" in t
    set_runs_text(lead[0], t.replace("13.5万円〜", "165,000円〜"))

    m = by_id(s, 21)
    for tshape in (m, by_id(s, 22)):
        g = tshape.table._tbl.find(qn("a:tblGrid")).findall(qn("a:gridCol"))
        g[0].set("w", str(int(Inches(3.55))))
        g[1].set("w", str(int(Inches(1.6))))
    rebuild_rows(m, [1, 2, 2, 3, 4], [
        ["DYMの運用費（運用代行・効率改善プラン／企画配信 月4回・ステップ配信 4通）", "100,000円"],
        ["離脱防止ポップアップ（ツール費）", "30,000円"],
        ["サンクスLINE誘導（ツール費）", "30,000円〜"],
        ["LINE公式アカウント利用料（ライトプラン想定）", "5,000円〜"],
        ["合計（予算：月額20万円以内）", "165,000円〜"],
    ], [0.46, 0.36, 0.36, 0.36, 0.42])
    m.table.rows[0].height = Inches(0.34)
    m.height = sum(r.height for r in m.table.rows)

    f = by_id(s, 22)
    f.top = Inches(3.95)
    rebuild_rows(f, [1, 2, 2, 2, 3], [
        ["LINEの初期構築（あいさつ・リッチメニュー・タグ設計・勤務前の案内の型）", "100,000円〜"],
        ["離脱防止ポップアップの導入", "15,000円"],
        ["サンクスLINE誘導の導入", "150,000円"],
        ["LINE公式アカウントの開設", "30,000円"],
        ["合計", "295,000円〜"],
    ], [0.46, 0.36, 0.36, 0.36, 0.4])
    f.table.rows[0].height = Inches(0.34)
    f.height = sum(r.height for r in f.table.rows)

    set_paras(by_id(s, 24).text_frame, [
        "DYM",
        "・配信の設計・作成（企画配信 月4回・ステップ配信 4通）／タグ・セグメント配信の設定",
        "・離脱防止ポップアップ・サンクスLINE誘導の導入・設定／勤務前の案内の型づくり",
        "・月次レポート・改善のご提案",
    ])
    p = by_id(s, 25).text_frame.paragraphs[3]
    assert p.runs[0].text == "・個別のチャット対応・応募フォームへの設置のご協力"
    set_runs_text(p, "・個別のチャット対応・サイトと応募フォームへの設置のご協力")
    note = by_id(s, 20).text_frame.paragraphs[0]
    assert "月6本の配信とは別に" in note.runs[0].text
    set_runs_text(note, note.runs[0].text.replace("月6本の配信とは別に", "企画配信・ステップ配信とは別に"))


def main(src, out):
    prs = Presentation(src)
    slides = list(prs.slides)
    assert len(slides) == 16
    slide13(slides[12])
    slide14(slides[13])
    prs.save(out)
    print("saved:", out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
