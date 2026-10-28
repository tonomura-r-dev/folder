# -*- coding: utf-8 -*-
"""ワクトリ v2：料金の3ページ（36〜38）に SIM ver1.0 の数字を入れ、成果報酬型（プラン③）を外す（2026-10-28 殿村さん「成果報酬消そう」）。

    python -m _build.patch_wakutori_v2_ryokin <v2.pptx> [出力.pptx]

  数字は `株式会社ワクトリ御中_LINEOA施策提案【SIM】ver1.0.xlsx`（殿村さんPC保存版＋SIM②を同じ形に）のまま。
    プラン①（SIM①・月3投稿）：初期 365,000円（運用設計 200,000＋ポップアップ 15,000＋サンクスLINE誘導 150,000）／
      月額 360,000円（コンサル 300,000＋ポップアップ 30,000＋サンクスLINE誘導 30,000）／6か月目 友だち 87,495人・応募 月234件
    プラン②（SIM②・月5投稿）：初期 265,000円（運用設計 250,000＋ポップアップ 15,000）／月額 380,000円（コンサル 350,000＋ポップアップ 30,000）
  LINE公式アカウント利用料（アカウント費・通数費）は別。ENTRY と応募は同じ登録フォームなので、概算お見積りの成果は「応募」の1つ。
  出力を省くと上書き。
"""
import copy
import sys

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches

from _build.patch_wakutori_v2_shisaku import by_id, remove, set_lines

EST = {
    339: ["概算お見積り：月36万円で、6か月目に友だち 87,495人・応募 月234件"],
    381: ["プラン①（月3投稿）で、友だちの入口を増やし、検討の順番に沿った配信で応募を増やす"],
    343: ["¥365,000〜"], 344: ["¥360,000〜"],
    364: ["応募"], 378: ["234件/月"], 375: ["87,495人"],
}
HEAD = ["項目", "プラン①（月3投稿）\nコンサル基本", "プラン②（月5投稿）\nコンサル基本"]
PLAN = [HEAD,
        ["初期費用", "365,000円〜", "265,000円〜"],
        ["月額コンサルティング費", "300,000円〜", "350,000円〜"],
        ["月額ツール・LINE関連費", "60,000円〜", "30,000円〜"],
        ["月額費用", "360,000円〜", "380,000円〜"]]
PLAN_NOTES = {
    640: "※初期費用＝運用設計・構築（プラン① 200,000円〜／プラン② 250,000円〜）＋ポップアップ 15,000円。プラン①はサンクスLINE誘導 150,000円を含む",
    641: "※月額ツール・LINE関連費＝ポップアップ 30,000円。プラン①はサンクスLINE誘導 30,000円〜を含む（プラン②はオプション）。LINE公式アカウント利用料は別",
}
SCOPE = [HEAD,
         ["ポップアップ", "導入・運用", "導入・運用"],
         ["適正バイト診断・\nステップ配信", "設計・構築・改善", "設計・構築・改善"],
         ["企画配信", "月3回程度", "月5回程度"],
         ["開始時期・エリア別の配信", "対応", "対応"],
         ["結果の分析・改善\n（月次レポート・定例会）", "対応", "対応"],
         ["サンクスLINE誘導", "導入・運用", "オプション"]]


def drop_last_col(tbl_shape, widths):
    tbl = tbl_shape._element.find(".//" + qn("a:tbl"))
    grid = tbl.find(qn("a:tblGrid"))
    grid.remove(grid.findall(qn("a:gridCol"))[-1])
    for tr in tbl.findall(qn("a:tr")):
        tr.remove(tr.findall(qn("a:tc"))[-1])
    for gc, w in zip(grid.findall(qn("a:gridCol")), widths):
        gc.set("w", str(int(Inches(w))))


def fill_table(tbl_shape, body):
    t = tbl_shape.table
    for r, row in enumerate(body):
        for c, text in enumerate(row):
            txb = t.cell(r, c)._tc.find(qn("a:txBody"))
            ps = txb.findall(qn("a:p"))
            tmpl = copy.deepcopy(ps[0])
            for x in tmpl.findall(qn("a:r"))[1:] + tmpl.findall(qn("a:br")):
                tmpl.remove(x)
            for p in ps:
                txb.remove(p)
            for line in text.split("\n"):
                p = copy.deepcopy(tmpl)
                p.find(qn("a:r")).find(qn("a:t")).text = line
                txb.append(p)


def main(src, out):
    prs = Presentation(src)
    est, plan, scope = (prs.slides[k] for k in (35, 36, 37))
    assert by_id(est, 339).text_frame.text.startswith("概算お見積り")

    for sid, lines in EST.items():
        set_lines(by_id(est, sid), lines)
    ent_y = by_id(est, 362).top                                # ENTRY の行を外し、応募を上へ
    remove(est, 362)
    remove(est, 376)
    by_id(est, 364).top = ent_y
    by_id(est, 378).top = ent_y
    tree = est.shapes._spTree                                  # 初期の実施内容にサンクスLINE誘導を1行足す（ポップアップの行を写す）
    nid = max(int(e.get("id")) for e in tree.iter(qn("p:cNvPr"))) + 1
    for k, (sid, text) in enumerate(((365, "サンクスLINE誘導"), (366, "導入"))):
        el = copy.deepcopy(by_id(est, sid)._element)
        el.find(".//" + qn("p:cNvPr")).set("id", str(nid + k))
        by_id(est, 382 if sid == 365 else 383)._element.addnext(el)
        sh = next(s for s in est.shapes if s.shape_id == nid + k)
        sh.top = by_id(est, 359).top
        set_lines(sh, [text])

    tp = next(sh for sh in plan.shapes if sh.has_table)
    drop_last_col(tp, [2.5, 3.75, 3.75])
    fill_table(tp, PLAN)
    set_lines(by_id(plan, 635), ["料金プラン：2つのプランの費用"])            # 成果報酬型を外したので2つ
    set_lines(by_id(plan, 637), ["運用範囲に合わせて、2つのプランをご用意"])
    for sid, text in PLAN_NOTES.items():
        set_lines(by_id(plan, sid), [text])

    ts = next(sh for sh in scope.shapes if sh.has_table)
    drop_last_col(ts, [2.5, 3.75, 3.75])
    fill_table(ts, SCOPE)
    set_lines(by_id(scope, 5), ["2つのプランの運用内容の比較"])
    set_lines(by_id(scope, 7), ["※プラン②のサンクスLINE誘導はオプション（初期150,000円・月額30,000円〜）"])
    prs.save(out)
    print("saved:", out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else sys.argv[1])
