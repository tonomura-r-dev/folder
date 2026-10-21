# -*- coding: utf-8 -*-
"""人材業界（派遣・バイト・単発）：リッチメニューのページ（21枚目）を組み直す（2026-10-21 殿村さん）。

    python -m _build.patch_haken_richmenu <入力.pptx> <出力.pptx>

  ボタンは横3×縦2の6つ：仕事情報／マイページ／Web登録・来社予約（上の段＝探す・進む）、
  よくある質問／社員インタビュー／お問い合わせ（下の段＝不安を減らす）。
  「通知を止める」「今は働いている」は外し、「今すぐ探す」「条件に合う新着」は「仕事情報」にまとめた（殿村さん指示）。
  メニューは画像ではなく図形で組む（ボタン名だけ。アイコンは入れない）。右は段ごとの役割を2つ（ボタン名は図にあるので書かない）。帯はリードと同じなので外した。
  ほかのページは変えない。
"""
import sys

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

from _build.pptx_parts import GREEN, LGRAY, WHITE, find, rect, rich, set_para, shape

NAVY = "1E2761"
SUB = "64748B"
LEAD = "求人を探す入口と、応募前の不安を減らすボタンを並べる。"
NOTE = "※ 画面はイメージ。ボタン名は運用しながら調整する"
BUTTONS = [["仕事情報", "マイページ", "Web登録・\n来社予約"],
           ["よくある質問", "社員インタビュー", "お問い合わせ"]]
GROUPS = [("探す・進む", ["条件に合う求人を探し、", "応募・登録はサイトで行う"]),
          ("不安を減らす", ["職場や働き方の疑問に、", "応募の前にお答えする"])]
REMOVE = {"Picture 26", "Rectangle 23", "TextBox 24"} | {f"Rectangle {n}" for n in (5, 8, 11, 14, 17, 20)} \
    | {f"TextBox {n}" for n in range(6, 23)}

# 配置（cm）
FX, FY, FW, FH = 1.27, 4.57, 15.37, 10.03   # メニュー全体（リッチメニューに近い横長）
PAD = 0.13
RX, RW = 17.4, 8.9                           # 右の説明


def main(src, out):
    prs = Presentation(src)
    s = list(prs.slides)[20]
    assert find(s, "TextBox 1").text_frame.text.endswith("リッチメニュー")
    for sh in list(s.shapes):
        if sh.name in REMOVE:
            sh._element.getparent().remove(sh._element)
    t2 = find(s, "TextBox 2")
    set_para(t2, 0, LEAD)
    while len(t2.text_frame.paragraphs) > 1:
        p = t2.text_frame.paragraphs[1]._p
        p.getparent().remove(p)
    set_para(find(s, "TextBox 25"), 0, NOTE)

    rect(s, FX, FY, FW, FH, "F2F2F2")
    tw = (FW - PAD * 4) / 3
    th = (FH - PAD * 3) / 2
    for r, row in enumerate(BUTTONS):
        for c, name in enumerate(row):
            first = (r, c) == (0, 0)
            shape(s, MSO_SHAPE.RECTANGLE, FX + PAD + c * (tw + PAD), FY + PAD + r * (th + PAD), tw, th,
                  [(t, 14, True, WHITE if first else NAVY, 0) for t in name.split("\n")],
                  fill=GREEN if first else WHITE, line=None if first else LGRAY, lw=1.0)
    for r, (head, desc) in enumerate(GROUPS):
        cy = FY + PAD + r * (th + PAD) + th / 2
        h = 2.1
        rect(s, RX, cy - h / 2, 0.12, h, NAVY)
        rich(s, RX + 0.3, cy - h / 2, RW - 0.3, h,
             [[(head, 14, True, NAVY)]] + [[(d, 12, False, SUB)] for d in desc],
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE)
    prs.save(out)
    print("saved:", out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
