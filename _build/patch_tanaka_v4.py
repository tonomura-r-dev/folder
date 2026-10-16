# -*- coding: utf-8 -*-
"""タナカホーム v4：殿村さんのPC保存版v2（7枚）を、2026-10-16に固めた骨組みに合わせて直す。

    python -m _build.patch_tanaka_v4 <PC保存版v2.pptx> <出力.pptx>

  1 表紙：タイトル＝殿村さん確定「集めた見込み客を逃さずナーチャリングをし、来場・契約へ変えるLINE仕組みの構築」（2行）。
          同じ意味のサブタイトルは外す
  2 現状と打ち手：表に「LINE公式アカウント」の行を足す（来場特典で友だちを集め、来場予約のご案内が中心）
  3 課題：来場を決める前の検討中の方と、LINEでつながれていない（前後検索の動き＋今のLINEの範囲）
  4 解決策（新規）：資料請求の時点から、来場・契約までを追う（今のLINE／弊社のご提案の比較）
  5 解決策（殿村さんのタイトルのまま）：配信を検討の順番に並べ直す
  旧5 現場体制：外す（殿村さん「不要」）
  6 今後の展開：型の③を「改善」に、締めを「つながる範囲を広げる」に
"""
import copy
import sys

from pptx import Presentation

from _build.arauchi_v4.lib import NAVY, RED, TEXT, add_box, add_table, set_paras, move
from _build.arauchi_v10.build import TINT, LINE_C, by_id, clone, meiryo, runs
from _build.build_tanaka_v1 import GRAY, head_cell, cell, label_cell, note, band
from _build.build_tanaka_v2 import heading, flow, X0, W

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
HDR = (2, 3, 4, 5, 6)   # PC保存版の本文ページ共通：タイトル・リード帯2つ・区切り線・リード


def drop(s, *ids):
    for sid in ids:
        e = by_id(s, sid)._element
        e.getparent().remove(e)


def set_lines(sh, lines):
    """書式を残したまま段落の文字を差し替える（最初の段落・最初の文字列を型にする）"""
    body = sh._element.find(f".//{{http://schemas.openxmlformats.org/presentationml/2006/main}}txBody")
    ps = body.findall(f"{A}p")
    tmpl = ps[0]
    for p in ps:
        body.remove(p)
    for ln in lines:
        p = copy.deepcopy(tmpl)
        rs = p.findall(f"{A}r")
        for r in rs[1:]:
            p.remove(r)
        rs[0].find(f"{A}t").text = ln
        body.append(p)


def set_title_lead(s, title, lead):
    set_lines(by_id(s, HDR[0]), [title])
    set_lines(by_id(s, HDR[4]), [lead])


def new_page(prs, src, title, lead):
    s = prs.slides.add_slide(src.slide_layout)
    for sh in list(s.shapes):
        sh._element.getparent().remove(sh._element)
    for sid in HDR:
        clone(s, by_id(src, sid))
    hdr = [sh for sh in s.shapes]
    set_lines(hdr[0], [title])
    set_lines(hdr[4], [lead])
    return s


# ---------------------------------------------------------------- 1 表紙
def p_cover(s):
    t = by_id(s, 26)
    set_lines(t, ["集めた見込み客を逃さずナーチャリングをし、", "来場・契約へ変えるLINE仕組みの構築"])
    bp = t._element.find(f".//{A}bodyPr")
    bp.set("wrap", "square")
    bp.set("anchor", "b")
    for c in list(bp):
        bp.remove(c)
    bp.append(bp.makeelement(f"{A}noAutofit", {}))
    move(t, x=0.6, y=2.5, w=9.63, h=1.3)
    drop(s, 27)


# ---------------------------------------------------------------- 2 現状と打ち手
def p_genjo(s):
    drop(s, 8)
    lw = 4.75
    rows = [[label_cell("来場予約\n（4展示場）", 11), cell("目標 月50〜60件に対し、現状 月35〜40件", 11)],
            [label_cell("来場予約のCPA", 11), cell("Web広告で約30,000円\n（ポータル・特典費を含めると注文住宅 約50,000円）", 10.5)],
            [label_cell("青森店", 11), cell("予約のキャンセルが2〜3割\n八戸・おいらせより認知が浅い", 11)],
            [label_cell("自社サイトの\n資料請求", 11), cell("月一桁", 11)],
            [label_cell("検索広告のCPA\n（9月）", 11), cell("14,796円 → 47,500円（前年同月比）", 11)],
            [label_cell("LINE公式\nアカウント", 11), cell("来場特典のキャンペーンで友だちを集め、\n見学会・来場予約のご案内を中心に配信", 10.5)]]
    add_table(s, X0, 2.3, [1.55, lw - 1.55], [0.52] * 6, rows, name="現在地")


# ---------------------------------------------------------------- 3 課題
def p_kadai(s):
    set_title_lead(s, "課題：来場を決める前の検討中の方と、LINEでつながれていない",
                   "検討者は約2週間かけて比較を続ける。この期間の接点が、来場を左右する")
    drop(s, *range(7, 20))
    flow(s, [("流入\n広告・SEO・ポータル", NAVY, "FFFFFF"), ("資料請求", NAVY, "FFFFFF"),
             ("来場予約・来場", NAVY, "FFFFFF"), ("商談・契約", NAVY, "FFFFFF")], 1.85, 0.72, size=11)
    texts = [("検討の動き", "当日に大手ハウスメーカーを一斉に比べ、その後約2週間、費用・住宅ローン、性能・間取りを調べ続ける"),
             ("資料請求の後", "今の友だち登録のご案内は来場特典とセット。来場を決める前の方には、友だちになる理由がない"),
             ("来場予約の後", "LINEで見学会・来場予約をご案内できている。青森店は予約のキャンセルが2〜3割"),
             ("来場の後", "競合（一条工務店・大進建設）と比べながら、契約まで検討が続く")]
    n, aw = 4, 0.32
    bw = (W - aw * (n - 1)) / n
    for k, (h1, t) in enumerate(texts):
        x = X0 + k * (bw + aw)
        add_box(s, x, 2.7, bw, 1.6, [runs((h1, dict(size=10.5, bold=True, color=NAVY))),
                                     runs((t, dict(size=10.5, color=TEXT)))],
                line=LINE_C, anchor="t", line_spacing=120, ins=(0.1, 0.08, 0.1, 0.06), name=f"状態{k + 1}")
    w2 = bw * 2 + aw
    add_box(s, X0, 4.45, w2, 0.6, [runs(("来場を決める前：LINEでつながる入口がまだない", dict(size=11, bold=True, color=NAVY)))],
            fill=TINT, line=NAVY, dash="dash", align="ctr", anchor="ctr", name="範囲 これから")
    add_box(s, X0 + w2 + aw, 4.45, w2, 0.6,
            [runs(("来場を決めた後：今のLINEがご案内している範囲", dict(size=11, bold=True, color="FFFFFF")))],
            fill=NAVY, align="ctr", anchor="ctr", name="範囲 今")
    band(s, [("課題は配信の中身ではなく、", dict(size=13, bold=True, color=NAVY)),
             ("LINEでつながっている範囲にある", dict(size=13, bold=True, color=RED))], y=5.55, h=0.62)
    note(s, "※検討者の動きは、LINEヤフー社の前後検索データより", y=6.4)


# ---------------------------------------------------------------- 4 解決策（新規）
def p_ou(prs, src):
    s = new_page(prs, src, "解決策：資料請求の時点から、来場・契約までを追う",
                 "ここまで追い続けられることが、LINEの強み。今の来場のご案内は、そのまま活かす")
    rows = [[head_cell(""), head_cell("今のLINE"), head_cell("弊社のご提案")],
            [label_cell("追う相手"), cell("来場を決めたお客様", 11, color=GRAY),
             cell("資料請求の時点の、検討中のお客様から", 11.5, True, NAVY)],
            [label_cell("ゴール"), cell("来場予約", 11, color=GRAY), cell("来場の先の、契約", 11.5, True, NAVY)],
            [label_cell("友だちの入口"), cell("来場特典のキャンペーン", 11, color=GRAY),
             cell("＋資料請求の完了ページ、サイト上での情報のご案内", 11.5, True, NAVY)],
            [label_cell("配信"), cell("見学会・来場予約のご案内", 11, color=GRAY),
             cell("＋前後検索で見える、検討の順番に沿った配信", 11.5, True, NAVY)]]
    add_table(s, X0, 1.85, [1.7, 3.2, 4.83], [0.42] + [0.62] * 4, rows, name="比較")
    band(s, [("来場・契約の結果は、貴社の記録と合わせて確認し、", dict(size=13, bold=True, color=NAVY)),
             ("広告の改善に生かす", dict(size=13, bold=True, color=RED))], y=5.55, h=0.62)
    note(s, "※広告の改善は、別紙の集客改善シミュレーションとあわせてご提案", y=6.4)
    return s


# ---------------------------------------------------------------- 5 解決策（配信）
def p_haishin(s):
    set_lines(by_id(s, HDR[4]), ["検討の順番に合わせて配信し、「住みたいイメージ」を育てる"])
    drop(s, *range(7, 16))
    rows = [[head_cell("時期"), head_cell("検討者が調べていること"), head_cell("配信するコンテンツ"),
             head_cell("感じてもらうこと")],
            [label_cell("当日〜1日"), cell("大手ハウスメーカーの比較、土地探し", 10.5),
             cell("会社の特長・施工事例、街並みと土地の情報", 10.5), cell("この会社・この街がいい", 11, True, NAVY, align="ctr")],
            [label_cell("1〜7日後"), cell("住宅ローンのシミュレーション・審査、坪単価", 10.5),
             cell("費用・住宅ローンの考え方、補助金", 10.5), cell("これなら建てられる", 11, True, NAVY, align="ctr")],
            [label_cell("7〜15日後"), cell("断熱等級、間取り、吹き抜け、ルームツアー", 10.5),
             cell("青森の冬を乗り切る性能、間取り、ルームツアー動画、入居後の暮らし", 10.5),
             cell("この家に住みたい", 11, True, NAVY, align="ctr")],
            [label_cell("来場前"), cell("―", 10.5, align="ctr"), cell("見学会・来場予約のご案内、来場前日のお知らせ", 10.5),
             cell("見に行きたい", 11, True, NAVY, align="ctr")]]
    add_table(s, X0, 1.85, [1.35, 2.85, 3.45, 2.08], [0.4] + [0.62] * 4, rows, name="配信")
    flow(s, [("友だち追加\n資料請求・来場予約・サイト", None, NAVY), ("配信でファン化", NAVY, "FFFFFF"),
             ("タイミングが良い時に来場", NAVY, "FFFFFF"), ("商談・契約", RED, "FFFFFF")], 4.95, 0.72, size=11)
    note(s, "※時期と調べていることは、LINEヤフー社の前後検索データより／動画・写真は今の配信やInstagramの素材を活かす", y=6.4)


# ---------------------------------------------------------------- 6 今後の展開
def p_tenkai(s):
    for sid, (h1, h2) in [(8, ("① 入口", "資料請求・来場予約・サイトからLINEへ")),
                          (9, ("② 配信", "検討の順番に沿ったステップ配信")),
                          (10, ("③ 改善", "来場・契約の結果を見て、配信と広告を見直す"))]:
        ps = by_id(s, sid)._element.findall(f".//{A}p")
        ps[0].findall(f"{A}r")[0].find(f"{A}t").text = h1
        ps[1].findall(f"{A}r")[0].find(f"{A}t").text = h2
    rs = by_id(s, 14)._element.findall(f".//{A}r")
    assert rs[1].find(f"{A}t").text.startswith("今、LINE")
    rs[1].find(f"{A}t").text = "今、LINEでつながる範囲を広げる"


def main(base, out):
    prs = Presentation(base)
    sl = list(prs.slides)
    assert len(sl) == 7 and by_id(sl[4], 2).text_frame.text.startswith("現場体制")
    cover, genjo, kadai, haishin, taisei, tenkai, back = sl
    p_cover(cover)
    p_genjo(genjo)
    p_kadai(kadai)
    ou = p_ou(prs, kadai)
    p_haishin(haishin)
    p_tenkai(tenkai)
    for s in (cover, genjo, kadai, ou, haishin, tenkai):
        for sh in s.shapes:
            if sh.has_text_frame:
                meiryo(sh)
    keep = [cover, genjo, kadai, ou, haishin, tenkai, back]
    lst = prs.slides._sldIdLst
    by_part = {prs.part.related_part(e.rId): e for e in lst}
    for e in list(lst):
        lst.remove(e)
    keep_parts = {s.part for s in keep}
    for part, e in by_part.items():
        if part not in keep_parts:
            prs.part.drop_rel(e.rId)
    for s in keep:
        lst.append(by_part[s.part])
    prs.save(out)
    print("saved:", out, len(prs.slides), "枚")


if __name__ == "__main__":
    main(*sys.argv[1:3])
