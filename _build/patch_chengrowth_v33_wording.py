# -*- coding: utf-8 -*-
"""チェングロウス ver3.2 → ver3.3（2026-10-02）
・「友だちを貯める」→「LINEを追加してもらう期間」
・LINEアカウントはサイトが変わっても変わらない。ご契約後に、離脱防止ポップアップ・サンクスLINE誘導・Profile+の実装を進める形に。
  python3 _build/patch_chengrowth_v33_wording.py <ver3.2.pptx>
"""
import sys
from pathlib import Path
from pptx import Presentation

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "20260929_チェングロウス_LINE公式アカウント運用のご提案ver3.3.pptx"
prs = Presentation(sys.argv[1])


def shape(n, name):
    return next(sh for sh in prs.slides[n - 1].shapes if sh.name == name)


def set_par(n, name, k, text):
    p = shape(n, name).text_frame.paragraphs[k]
    runs = p.runs
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


set_par(23, "Rounded Rectangle 7", 1, "あいさつ・リッチメニュー・配信の設計、離脱防止ポップアップ・サンクスLINE誘導・Profile+の実装")
set_par(23, "Rounded Rectangle 9", 0, "LINEを追加してもらう期間")
set_par(23, "Rounded Rectangle 9", 1, "離脱防止ポップアップなどで、LINEを追加してもらう（この期間の数字はシミュレーションに含めない）")
set_par(23, "Rounded Rectangle 11", 1, "リニューアルしたサイトに「LINEで登録」を組み込み、計測と配信を本格化")
set_par(23, "TextBox 14", 0, "※LINE Profile+の審査期間は、申請内容により変わります／現行サイトへの実装可否、求人ボックス・Indeed経由の応募者への導線は、確認のうえ決定")
prs.save(str(OUT))
print("saved:", OUT.name, len(prs.slides), "枚")
