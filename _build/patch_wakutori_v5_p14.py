# -*- coding: utf-8 -*-
"""ワクトリ v4 → v5：P14「他社のLINE活用のまとめ」の中身を増やす（2026-10-30 殿村さん）。

    python -m _build.patch_wakutori_v5_p14 <v4.pptx> <v5.pptx>

  ・表を4列に：他社がやっていること／どこがやっているか／ワクトリの今／ご提案での打ち手（ページ）
  ・5行目「相談の入口を、最初から見せる」を足す（事実はP9の比較表のとおり）
  ・下に示唆の枠（薄い灰色の塗りだけ。左の線は付けない）：「アンケートで希望を聞けているため、その分析を用いた配信が重要」（殿村さんの文）
  ワクトリの今は、P9・P15と、殿村さんの画面（リッチメニュー：11月開始・スキー特集の入口）で確かめた事実だけ。ほかのページは触らない。
"""
import copy
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

COLS = [(0.42, 2.25), (2.71, 3.35), (6.10, 1.80), (7.94, 2.48)]
HEAD = ["他社がやっていること", "どこがやっているか", "ワクトリの今", "ご提案での打ち手"]
ROWS = [
    (["1. 追加の直後に、", "登録・相談へ誘導する"], ["リゾートバイトダイブ／", "リゾバ.com／アルファリゾート"],
     ["翌日に", "ENTRYの案内"], ["追加直後からの", "ステップ配信（P30）"]),
    (["2. 追加の後も、", "毎日配信する会社もある"], ["リゾートバイトダイブ（毎日1通）／", "リゾバ.com（1日2通）"],
     ["週1回の企画配信"], ["追加後15日の", "ステップ配信", "（P24・P30）"]),
    (["3. 経験・タイプを聞き、", "合う求人を出し分ける"], ["リゾバ.com（経験・タイプ別）／", "リゾートバイト.com（アンケート）"],
     ["アンケート3問で", "開始時期と", "エリアを伺う"], ["適正バイト診断と", "開始時期・エリア別の配信", "（P29・P31）"]),
    (["4. リッチメニューに、", "開始月別の求人を置く"], ["4社とも"],
     ["11月開始の求人と", "スキー特集の入口"], ["上段に開始月別の求人", "（P32）"]),
    (["5. 相談の入口を、", "最初から見せる"], ["リゾートバイトダイブ（電話15分）／", "リゾバ.com・アルファリゾート", "（チャット）"],
     ["専属エージェント", "（ENTRYの後）"], ["ステップ配信と", "リッチメニューに", "相談の入口（P24・P32）"]),
]
Y0, ROW_H, GAP = 2.44, 0.72, 0.04


def by_id(s, sid):
    return next(sh for sh in s.shapes if sh.shape_id == sid)


def put(el, lines):
    txb = el.find(qn("p:txBody"))
    ps = txb.findall(qn("a:p"))
    tmpl = copy.deepcopy(ps[0])
    for r in tmpl.findall(qn("a:r"))[1:] + tmpl.findall(qn("a:br")):
        tmpl.remove(r)
    for p in ps:
        txb.remove(p)
    for line in lines:
        p = copy.deepcopy(tmpl)
        p.find(qn("a:r")).find(qn("a:t")).text = line
        txb.append(p)


def clone(s, src, x, y, w, h, lines):
    el = copy.deepcopy(src._element)
    nid = max(int(e.get("id")) for e in s.shapes._spTree.iter(qn("p:cNvPr"))) + 1
    el.find(".//" + qn("p:cNvPr")).set("id", str(nid))
    s.shapes._spTree.append(el)
    sh = by_id(s, nid)
    sh.left, sh.top, sh.width, sh.height = Inches(x), Inches(y), Inches(w), Inches(h)
    put(el, lines)
    return sh


def main(src, out):
    prs = Presentation(src)
    s = prs.slides[13]
    assert "他社のLINE活用のまとめ" in by_id(s, 5).text_frame.text
    head = [by_id(s, 59), by_id(s, 60)]
    first = [by_id(s, 61), by_id(s, 62)]
    old = [59, 60, 61, 62, 63, 64, 65, 66, 67, 68]
    for k, (x, w) in enumerate(COLS):
        clone(s, head[0] if k == 0 else head[1], x, 2.0, w, 0.4, [HEAD[k]])
    for r, cells in enumerate(ROWS):
        y = Y0 + r * (ROW_H + GAP)
        for k, ((x, w), lines) in enumerate(zip(COLS, cells)):
            clone(s, first[0] if k == 0 else first[1], x, y, w, ROW_H, lines)
    for sid in old:
        el = by_id(s, sid)._element
        el.getparent().remove(el)
    # 示唆の枠（薄い灰色の塗りだけ）
    y = Y0 + len(ROWS) * (ROW_H + GAP) + 0.12
    box = s.shapes.add_shape(1, Inches(0.42), Inches(y), Inches(10.0), Inches(0.5))
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor.from_string("F2F2F2")
    box.line.fill.background()
    box.shadow.inherit = False
    st = box._element.find(qn("p:style"))
    if st is not None:
        box._element.remove(st)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.15)
    r = tf.paragraphs[0].add_run()
    r.text = "アンケートで希望を聞けているため、その分析を用いた配信が重要"
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string("1F2A5C")
    rpr = r._r.get_or_add_rPr()
    for t in ("a:latin", "a:ea", "a:cs"):
        rpr.append(rpr.makeelement(qn(t), {"typeface": "メイリオ"}))
    tf.paragraphs[0].alignment = 2
    prs.save(out)
    print("saved:", out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
