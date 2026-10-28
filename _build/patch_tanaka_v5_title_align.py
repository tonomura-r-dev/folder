# -*- coding: utf-8 -*-
"""タナカホーム v5（殿村さんのPC保存版）：各ページ左上の見出し（ページの名前）の位置をそろえる（2026-10-28 殿村さん「縦横関係そろえて」）。

    python -m _build.patch_tanaka_v5_title_align <v5（PC保存版）.pptx> <出力.pptx>

  基準＝11ページ「提案の流れ」の見出し（左 0.6・上 0.151インチ、上下中央）。
  7〜9ページ「前提調査」は位置が少しずつずれていて（左 0.507〜0.645）、頭に空白が入っていたので、基準の位置にそろえて空白を外す。
  28ページ（サンクスLINE誘導）の見出しのわずかなずれも基準に合わせる。文字・大きさ・幅は変えない。
"""
import sys

from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Emu


def head(s):
    return next(sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()
                and Emu(sh.top).inches < 0.6 and Emu(sh.left).inches < 3)


def main(src, out):
    prs = Presentation(src)
    slides = list(prs.slides)
    base = head(slides[10])
    assert base.text_frame.text == "提案の流れ"
    fixed = []
    for i, s in enumerate(slides[1:31], 2):
        try:
            h = head(s)
        except StopIteration:
            continue
        changed = h.left != base.left or h.top != base.top or h.text_frame.text != h.text_frame.text.lstrip()
        if not changed:
            continue
        h.left, h.top, h.height = base.left, base.top, base.height
        h.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        r = h.text_frame.paragraphs[0].runs[0]
        r.text = r.text.lstrip()
        fixed.append((i, h.text_frame.text))
    prs.save(out)
    print("saved:", out, fixed)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
