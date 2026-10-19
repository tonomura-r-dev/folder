# -*- coding: utf-8 -*-
"""すなおコンサルティング：成果報酬プランのページの差し込み枠に、成果報酬単価の表を入れる（2026-10-19 殿村さん「入れといて」）。

    python -m _build.patch_sunao_seika_tanka <今の資料.pptx> [出力.pptx]

  セミナー予約 1.2万円／サロン入会 2.5万円（build_sunao_v2.SEIKA）。ほかのページは触らない（出力を省くと上書き）。
"""
import sys

from pptx import Presentation
from pptx.util import Inches

from _build.build_sunao_v1 import all_meiryo
from _build.build_sunao_v2 import fill_seika


def main(src, out):
    prs = Presentation(src)
    hits = [s for s in prs.slides if any(sh.name == "差し込み 成果の地点・成果報酬単価" for sh in s.shapes)]
    assert len(hits) == 1, f"差し込み枠のページが {len(hits)} 枚"
    s = hits[0]
    for sh in s.shapes:
        if sh.name == "小見出し":
            sh.top = Inches(4.1)
    fill_seika(s)
    all_meiryo([s])
    prs.save(out)
    print("saved:", out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else sys.argv[1])
