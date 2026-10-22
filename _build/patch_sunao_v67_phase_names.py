# -*- coding: utf-8 -*-
"""すなお ver6.6 → ver6.7：「想定動線（フェーズごとのLINEの役割）」の表のフェーズ名を、上の8フェーズの帯に合わせる（2026-10-22 殿村さん）。

    python -m _build.patch_sunao_v67_phase_names <ver6.6.pptx> <ver6.7.pptx>

  表の1列目を帯と同じ呼び方にし、長くなった分だけ1列目を広げる（ほかの列は中身の長さに合わせて詰める。表の横幅は変えない）。
  ほかのページは変えない。
"""
import sys

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches

PHASES = ["① 接触（認知・流入）", "② サイト離脱（対策）", "③ 育成（興味喚起）", "④ リード獲得 (web/LINE)",
          "⑤ リード有効化", "⑥ 有効リードの再育成", "⑦ マネタイズ(成約/購入)", "⑧ LTV最大化・紹介"]
COLS = [2.35, 2.6, 2.55, 2.76]  # フェーズ／ユーザー状態・心理／離脱・停滞が起きやすい点／LINE OAの役割・施策（合計10.26＝元と同じ）


def main(src, out):
    prs = Presentation(src)
    s = list(prs.slides)[16]
    assert s.shapes[2].text_frame.text.startswith("想定動線（フェーズごとのLINEの役割）") or any(
        sh.has_text_frame and sh.text_frame.text.startswith("想定動線（フェーズごと") for sh in s.shapes)
    gf = next(sh for sh in s.shapes if sh.has_table)
    tbl = gf.table
    assert tbl.cell(0, 0).text == "フェーズ" and len(tbl.rows) == 9
    assert abs(sum(COLS) - sum(c.width for c in tbl.columns) / 914400) < 0.01
    for i, w in enumerate(COLS):
        tbl.columns[i].width = Inches(w)
    for r, name in enumerate(PHASES, start=1):
        p = tbl.cell(r, 0)._tc.find(qn("a:txBody")).find(qn("a:p"))
        runs = p.findall(qn("a:r"))
        for x in runs[1:]:
            p.remove(x)
        runs[0].find(qn("a:t")).text = name
    prs.save(out)
    print("saved:", out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
