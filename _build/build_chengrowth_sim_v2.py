import sys, json
from openpyxl import load_workbook
SRC='_templates/DYM_LINEOA_SIM_FMT_ver3.00.xlsx'
OUT=sys.argv[1]
CV=json.loads(sys.argv[2]) if len(sys.argv)>2 else None   # {"有":[...6], "無":[...6]}
wb=load_workbook(SRC)
base=wb['SIM1_']; comp=wb['SIM1_比較']
base.title='SIM1_通知メッセ有'
s2=wb.copy_worksheet(base); s2.title='SIM2_通知メッセ無'
c2=wb.copy_worksheet(comp); c2.title='SIM2_比較'
for ws,name in ((comp,'SIM1_通知メッセ有'),(c2,'SIM2_通知メッセ無')):
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value,str) and 'SIM1_!' in c.value:
                c.value=c.value.replace('SIM1_!',f"'{name}'!")
for ws,name in ((comp,'SIM1_通知メッセ有'),(c2,'SIM2_通知メッセ無')):
    ws['D7']=f"='{name}'!G27"; ws['E7']=f"='{name}'!J27"
order=['SIM1_通知メッセ有','SIM1_比較','SIM2_通知メッセ無','SIM2_比較','SIM考え方']
wb._sheets=[wb[n] for n in order]
cols='EFGHIJ'
mon=['4月','5月','6月','7月','8月','9月']
sd=[1.0175,1.002,0.9945,1.013,0.9845,1.0235]
ramp=[0.62,0.81,0.93,1.00,1.03,1.05]
apps=[40,41,43,45,44,47]
block=[0.406,0.391,0.378,0.362,0.359,0.344]
def fill(ws,notify,key):
    ws['A4']='株式会社チェングロウス御中_LINEOA施策提案【SIM】'+('（通知メッセージ有）' if notify else '（通知メッセージ無）')
    ws['C8']='求人応募（LINE経由）'; ws['D8']=0
    ws['C10']='-'; ws['D10']=0
    ws['H8']=6100; ws['H9']=5000; ws['G10']='サイト応募（広告経由）'; ws['H10']=45; ws['H11']=0
    ws['K8']=0; ws['P9']=False
    ws['I10']=None; ws['I11']=None
    for i,c in enumerate(cols):
        ws[f'{c}13']=f'{i+1}か月目（{mon[i]}）'
        ws[f'{c}14']=round(6100*sd[i]); ws[f'{c}15']=round(5000*sd[i]); ws[f'{c}16']=apps[i]
        ws[f'{c}17']=False; ws[f'{c}18']=True; ws[f'{c}19']=True; ws[f'{c}20']=notify; ws[f'{c}21']=False
        r=ramp[i]
        ws[f'{c}23']=f'=IF({c}18=FALSE,0,{c}14*0.4*0.125*0.125*{r})'
        ws[f'{c}24']=f'=IF({c}19=FALSE,0,{c}16*0.8*0.45*{r})'
        ws[f'{c}25']=f'=IF({c}20=FALSE,0,{c}16*(1-0.8*0.45*{r})*0.7*0.3)'
        ws[f'{c}26']=f'=SUM({c}22:{c}25)+{c}14*0.005*{r}'
        ws[f'{c}27']=block[i]
        ws[f'{c}30']=4
        ws[f'{c}69']=200000
        ws[f'{c}67']=50000
        ws[f'{c}70']=0
        ws[f'{c}71']=(f'=IF({c}20=FALSE,0,80000+ROUND({c}16*(1-0.8*0.45*{r}),0)*7)')
        ws[f'{c}72']=0
        if CV: ws[f'{c}36']=CV[key][i]
        ws[f'{c}38']=0
    ws['D69']=200000
    ws['D67']=100000
    ws['N67']='サンクスLINE誘導ツール：初期10万・月5万（FMTの目安）'
    ws['D70']=100000 if notify else 0
    ws['N23']='離脱防止：UU×表示40%×クリック12.5%×追加12.5%×立ち上げ係数'
    ws['N24']='完了画面：サイト応募×遷移80%×追加45%×立ち上げ係数'
    ws['N25']='通知メッセ：完了画面で追加しなかった応募者×情報一致70%×追加30%'
    ws['N26']='＋サイト内「LINEで登録」：UU×0.5%×立ち上げ係数'
    ws['N27']='ブロック率：40%前後から出し分け配信で徐々に低下（目標30%以内）'
    ws['N69']='LINE運用コンサル（初期20万・月20万）'
    ws['N70']='通知メッセージ初期費用（有のみ）'
    ws['N71']='通知メッセージ：ツール費月8万＋通数×7円（有のみ）'
fill(wb['SIM1_通知メッセ有'],True,'有'); fill(wb['SIM2_通知メッセ無'],False,'無')
from openpyxl.workbook.properties import CalcProperties
wb.calculation=CalcProperties(fullCalcOnLoad=True)
wb.save(OUT)
