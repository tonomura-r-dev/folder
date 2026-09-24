# -*- coding: utf-8 -*-
"""チェングロウスver1.1：「できること・できないこと｜応募の前後で追う」を1枚追加＋先方に数字を求める文言を削除（2026-09-24）。
- 汎用の「ミニアプリ」（旧S46）を作り替えて、S3（想定動線）の直後＝新S4へ移動
- それに伴いページ参照を+1（改善モデル P40〜43→P41〜44／都度発注プラン P31→P32）
- S2・S16の脚注から「実測値・実績値をいただき次第〜」を削除（殿村さん：先方に数字は求めない）
  python3 _build/patch_chengrowth_0924h.py <入力pptx> <出力pptx>
"""
import sys
from pptx import Presentation
from pptx.util import Cm, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn

SRC, OUT = sys.argv[1], sys.argv[2]
prs = Presentation(SRC)
NAVY, INK, GRAY = "1F285A", "333333", "7F7F7F"
LBLUE, BEIGE, LGREEN, LGRAY = "DDEBF7", "FFF2CC", "E2F0D9", "F2F2F2"


def meiryo(run):
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        e = rPr.find(qn(tag))
        if e is None:
            e = rPr.makeelement(qn(tag), {})
            rPr.append(e)
        e.set("typeface", "メイリオ")


def box(slide, x, y, w, h, fill, lines, color=INK, size=10.5, bold=False, anchor=MSO_ANCHOR.MIDDLE,
        align=PP_ALIGN.CENTER, line=None, lw=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE, ml=0.15):
    sp = slide.shapes.add_shape(shape, Cm(x), Cm(y), Cm(w), Cm(h))
    if fill:
        sp.fill.solid()
        sp.fill.fore_color.rgb = RGBColor.from_string(fill)
    else:
        sp.fill.background()
    if line:
        sp.line.color.rgb = RGBColor.from_string(line)
        sp.line.width = Pt(lw)
    else:
        sp.line.fill.background()
    sp.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sp.adjustments[0] = 0.08
    tf = sp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Cm(ml)
    tf.margin_top = tf.margin_bottom = Cm(0.06)
    for i, ln in enumerate(lines):
        t, sz, b, c = (ln, size, bold, color) if isinstance(ln, str) else ln
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = t
        r.font.size = Pt(sz)
        r.font.bold = b
        r.font.color.rgb = RGBColor.from_string(c)
        meiryo(r)
    return sp


def text(slide, x, y, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    return box(slide, x, y, w, h, None, lines, align=align, anchor=anchor, shape=MSO_SHAPE.RECTANGLE, ml=0.05)


def arrow(slide, x1, y1, x2, y2, color="8C8C8C"):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Cm(x1), Cm(y1), Cm(x2), Cm(y2))
    c.line.color.rgb = RGBColor.from_string(color)
    c.line.width = Pt(1.5)
    ln = c.line._get_or_add_ln()
    ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle"}))


# ---------- 旧S46（ミニアプリ）を作り替え ----------
s = prs.slides[45]
tree = s.shapes._spTree
tshape = None
for sh in list(s.shapes):
    if sh.has_text_frame and sh.top is not None and sh.top < Cm(1.2) and sh.text_frame.text.strip() and tshape is None:
        tshape = sh
        continue
    if sh.shape_type == 9 and abs(sh.top - Cm(3.86)) < Cm(0.1) and sh.width > Cm(20):
        continue
    tree.remove(sh._element)
rs = [r for p in tshape.text_frame.paragraphs for r in p.runs]
rs[0].text = "できること・できないこと｜応募の前後で追う"
for r in rs[1:]:
    r.text = ""
if not any(sh.shape_type == 9 for sh in s.shapes):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, 0, Cm(3.86), Cm(27.52), Cm(3.86))
    c.line.color.rgb = RGBColor.from_string("D9D9D9")

text(s, 1.2, 1.85, 25.1, 1.85, [("広告の単価や求人の中身は変えません。", 13, False, INK),
                                ("変えるのは「応募の手前」と「応募の後」だけ。そこをLINEで追いかけます。", 13, False, INK)],
     anchor=MSO_ANCHOR.MIDDLE)

# 左：変えないこと
box(s, 1.2, 4.3, 6.9, 0.9, "7F7F7F", [("変えないこと", 12.5, True, "FFFFFF")])
box(s, 1.2, 5.3, 6.9, 8.6, LGRAY, [
    ("・広告の単価", 11.5, True, INK), ("　Meta・Googleは今のまま", 10, False, GRAY), ("", 6, False, INK),
    ("・求人の中身・条件", 11.5, True, INK), ("　今の求人をそのまま活かす", 10, False, GRAY), ("", 6, False, INK),
    ("・サイト・応募フォーム", 11.5, True, INK), ("　大きな改修はしない", 10, False, GRAY), ("", 6, False, INK),
    ("・担当者の人数", 11.5, True, INK), ("　電話の件数を増やさない", 10, False, GRAY)],
    align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, ml=0.4)

# 右：応募の手前／応募の後
RX, RW = 8.5, 17.82
ROWS = [
    ("応募の手前：帰ろうとした人を残す", BEIGE,
     [("1", "求人ページを見て", "応募せずに帰る"), ("2", "離脱防止で", "年収相場をLINEで"),
      ("3", "友だち追加", "30秒 適職診断"), ("4", "14日間の配信で", "面談へご案内")]),
    ("応募の後：面談・就業まで追いかける", LGREEN,
     [("1", "応募完了画面から", "LINEへご案内"), ("2", "面談日程をLINEで", "決定＋前日に連絡"),
      ("3", "選考の連絡を", "LINEに一本化"), ("4", "内定後の不安に", "即答・辞退を防ぐ")]),
]
y = 4.3
for head, col, steps in ROWS:
    box(s, RX, y, RW, 0.9, col, [(head, 12.5, True, INK)], align=PP_ALIGN.LEFT, ml=0.35)
    n, gap = 4, 0.45
    sw = (RW - gap * (n - 1)) / n
    for k, (no, a, b) in enumerate(steps):
        sx = RX + k * (sw + gap)
        box(s, sx, y + 1.1, sw, 2.95, NAVY if k == n - 1 else "FFFFFF",
            [(no, 11, True, "FFFFFF" if k == n - 1 else "3467B2"),
             (a, 11, True, "FFFFFF" if k == n - 1 else NAVY), (b, 10.5, False, "FFFFFF" if k == n - 1 else INK)],
            line=None if k == n - 1 else "8EA9DB")
        if k < n - 1:
            arrow(s, sx + sw + 0.05, y + 2.57, sx + sw + gap - 0.05, y + 2.57)
    y += 4.8
box(s, 1.2, 14.35, 25.12, 1.0, NAVY, [("広告は今のまま。取りこぼしていた「応募の手前」と「応募の後」を、LINEで埋める。", 13, True, "FFFFFF")])

# ---------- S3の直後へ移動 ----------
lst = prs.slides._sldIdLst
ids = list(lst)
el = ids[45]
lst.remove(el)
ids[2].addnext(el)

# ---------- 文言：数字のお願いを削除／ページ参照を+1 ----------
REPL = [
    ("。UUはSimilarWeb推計のため、GA実測値をいただき次第置き換え", "。UUはSimilarWeb推計"),
    ("ほか仮置き値を含む。実績値をいただき次第、再試算", "ほか仮置き値を含む"),
    ("改善モデル①〜④（P40〜43）", "改善モデル①〜④（P41〜44）"),
    ("（P31 都度発注プラン）", "（P32 都度発注プラン）"),
]
cnt = 0
for sl in prs.slides:
    for sh in sl.shapes:
        if not sh.has_text_frame:
            continue
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                for a, b in REPL:
                    if a in r.text:
                        r.text = r.text.replace(a, b)
                        cnt += 1
prs.save(OUT)
print("saved", OUT, "replaced", cnt)
