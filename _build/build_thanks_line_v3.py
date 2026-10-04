# -*- coding: utf-8 -*-
"""サンクスLINEのご提案（4枚：機能／重要性／シーン1／シーン2）  2026-10-02
文面は `_drafts/サンクスLINE_資料化テキスト（2026-10-01）.md`（殿村さん確定版）のとおり。
土台は ver1.4 の5枚（先頭4枚の中身を作り直し、5枚目を削除）。
画像は使わない（殿村さん指示 2026-10-02）。すべて図形で図解
  python3 _build/build_thanks_line_v3.py <20260930_サンクスLINEのご提案ver1.4.pptx>
"""
import copy
import sys
from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.util import Cm, Pt

ROOT = Path(__file__).resolve().parent.parent
IMG = ROOT / "_images"
OUT = ROOT / "20261002_サンクスLINEのご提案.pptx"
NAVY, INK, GRAY, WHITE = RGBColor(0x1F, 0x28, 0x5A), RGBColor(0x33, 0x33, 0x33), RGBColor(0x7F, 0x7F, 0x7F), RGBColor(0xFF, 0xFF, 0xFF)
TITLE = RGBColor(0x00, 0x20, 0x60)
GREENTXT = RGBColor(0x0B, 0x7A, 0x3B)

prs = Presentation(sys.argv[1])
SW = prs.slide_width / 360000


def shp(si, name):
    return next(x for x in prs.slides[si].shapes if x.name == name)


# ---- 図形の見本（ver1.4から複製して使う）----
TPL = {
    "card": copy.deepcopy(shp(1, "Rounded Rectangle 4")._element),        # 薄い青・中央
    "bullet": copy.deepcopy(shp(1, "Rounded Rectangle 12")._element),     # 薄い青・紺枠・左揃え
    "green": copy.deepcopy(shp(2, "Rounded Rectangle 8")._element),       # LINEの緑
    "navy": copy.deepcopy(shp(2, "Rounded Rectangle 4")._element),        # 紺
    "lgreen": copy.deepcopy(shp(1, "Rounded Rectangle 10")._element),     # 薄い緑・緑枠
    "arrow_navy": copy.deepcopy(shp(1, "Right Arrow 5")._element),
    "arrow_green": copy.deepcopy(shp(1, "Right Arrow 7")._element),
    "note": copy.deepcopy(shp(3, "TextBox 19")._element),                 # 出典の注記
    "line": copy.deepcopy(shp(1, "Rounded Rectangle 12")._element),       # 白＋紺枠（中央）
    "ink": copy.deepcopy(shp(2, "Rounded Rectangle 4")._element),         # 黒
}


def set_runs(sh, lines, align=None):
    """lines = [(text, size, bold, color), ...] 1要素＝1段落。書式は元の最初のrunを引き継ぐ。"""
    tf = sh.text_frame
    base_p = tf.paragraphs[0]._p
    base_r = base_p.find(qn("a:r"))
    rpr0 = copy.deepcopy(base_r.find(qn("a:rPr"))) if base_r is not None else None
    for p in list(tf._txBody.findall(qn("a:p")))[1:]:
        tf._txBody.remove(p)
    for r in list(base_p.findall(qn("a:r"))) + list(base_p.findall(qn("a:br"))):
        base_p.remove(r)
    empty_p = copy.deepcopy(base_p)
    for i, (text, size, bold, color) in enumerate(lines):
        p = base_p if i == 0 else copy.deepcopy(empty_p)
        if i:
            tf._txBody.append(p)
        r = p.makeelement(qn("a:r"), {})
        if rpr0 is not None:
            r.append(copy.deepcopy(rpr0))
        t = r.makeelement(qn("a:t"), {})
        t.text = text
        r.append(t)
        end = p.find(qn("a:endParaRPr"))
        (end.addprevious(r) if end is not None else p.append(r))
    for para, (text, size, bold, color) in zip(tf.paragraphs, lines):
        if align is not None:
            para.alignment = align
        for run in para.runs:
            run.font.size, run.font.bold = Pt(size), bold
            run.font.color.rgb = color


def add(s, kind, x, y, w, h, lines=None, align=None):
    el = copy.deepcopy(TPL[kind])
    s.shapes._spTree.append(el)
    sh = s.shapes[-1]
    sh.left, sh.top, sh.width, sh.height = Cm(x), Cm(y), Cm(w), Cm(h)
    if kind == "line":                                   # 白＋紺の枠線
        sh.fill.solid()
        sh.fill.fore_color.rgb = WHITE
        sh.line.color.rgb = NAVY
        sh.line.width = Pt(1.5)
    if kind == "ink":
        sh.fill.solid()
        sh.fill.fore_color.rgb = RGBColor(0x33, 0x33, 0x33)
    if lines:
        set_runs(sh, lines, 2 if (kind == "line" and align is None) else align)
    return sh


def source(s, y, text):
    sh = add(s, "note", 3.0, y, 21.5, 0.6)
    set_runs(sh, [(text, 10, False, GRAY)])


def head(si, title, lead):
    s = prs.slides[si]
    for sh in list(s.shapes):
        if sh.name not in ("TextBox 1", "TextBox 2", "Connector 3"):
            sh._element.getparent().remove(sh._element)
    set_runs(shp(si, "TextBox 1"), [(title, 16, True, TITLE)])
    set_runs(shp(si, "TextBox 2"), [(lead, 16, False, INK)])
    t2 = shp(si, "TextBox 2")
    t2.top, t2.height = Cm(1.9), Cm(1.2)
    return s


def picture(s, name, x_w, y, crop_bottom=0.0):
    W = x_w
    H = W * 941 * (1 - crop_bottom) / 1672
    pic = s.shapes.add_picture(str(IMG / name), Cm((SW - W) / 2), Cm(y), Cm(W), Cm(H))
    if crop_bottom:
        pic.crop_bottom = crop_bottom
    return pic, H


# 色は紺・白・黒だけ（殿村さん指示 2026-10-02：緑は見づらい）
# 図解は「導線」だけ。それ以外は文字（殿村さん指示 2026-10-02）。枠は小さめ・周囲に余白。
X0, CWID = 3.0, 21.5          # 左端・内容の幅（左右に余白）


def bullets(s, y, items, h=7.0):
    """items = [(text, size, bold, color, space_after_pt), ...]"""
    sh = add(s, "note", X0, y, CWID, h)
    set_runs(sh, [(tx, sz, b, c) for tx, sz, b, c, _ in items])
    for para, it in zip(sh.text_frame.paragraphs, items):
        para.space_after = Pt(it[4])
    return sh


def flow(s, y, h, steps, bw=4.55, gap=1.1):
    x = X0
    for i, (kind, text) in enumerate(steps):
        col = WHITE if kind == "navy" else NAVY
        add(s, kind, x, y, bw, h, [(tx, 15, True, col) for tx in text.split("\n")])
        x += bw
        if i < len(steps) - 1:
            add(s, "arrow_navy", x + 0.05, y + h / 2 - 0.5, 1.0, 1.0)
            x += gap


# ================= 1枚目：機能 =================
s = head(0, "サンクスLINE｜機能", "申込み完了画面から、LINEへ移行")
Y, H = 5.4, 2.8
add(s, "line", X0, Y, 5.2, H, [("申込み完了画面", 16, True, NAVY)])
add(s, "arrow_navy", X0 + 5.25, Y + H / 2 - 0.5, 1.0, 1.0)
add(s, "navy", X0 + 6.4, Y, 4.4, H, [("LINEへ移行", 16, True, WHITE)])
add(s, "arrow_navy", X0 + 10.85, Y + H / 2 - 0.5, 1.0, 1.0)
RX = X0 + 12.0
add(s, "line", RX, Y, CWID - 12.0, 1.3, [("新規 → 友だち追加", 15, True, NAVY)])
add(s, "line", RX, Y + 1.5, CWID - 12.0, 1.3, [("既存の友だち → トーク画面", 15, True, NAVY)])
bullets(s, 9.9, [
    ("・入力内容を引き継ぎ、配信の出し分けに活用", 20, False, INK, 14),
    ("・今のアカウントにそのまま導入、APIツールとも併用可", 20, False, INK, 14),
    ("・費用：初期15万円／月額3万円〜", 20, True, NAVY, 0),
])

# ================= 2枚目：重要性（文字）=================
s = head(1, "サンクスLINE｜重要性", "申込みの直後に、LINEでつながる")
bullets(s, 6.0, [
    ("・申込み⇒来店までに一定数が離脱する", 20, True, NAVY, 22),
    ("・関心が一番高い申込み直後に、LINEでつながるのが最適", 20, True, NAVY, 22),
    ("・つながると、読まれる", 20, True, NAVY, 2),
    ("　（開封率 約70%・クリック率はメルマガの約5倍／バス会社）", 14, False, GRAY, 22),
    ("・つながると、動く", 20, True, NAVY, 2),
    ("　（予約後の来院率 40〜50%改善／美容クリニック）", 14, False, GRAY, 0),
], h=9.5)

# ================= 3枚目：シーン1 =================
s = head(2, "シーン1｜来店・来院の予約", "予約を、来店・来院につなげる")
flow(s, 5.4, 2.8, [("line", "予約完了"), ("navy", "LINEで\n予約確認"), ("navy", "前日に\nお知らせ"), ("line", "来店")])
bullets(s, 9.9, [
    ("・無断キャンセルを防ぎ、来店につなげる", 20, False, INK, 18),
    ("・実績：美容クリニック　来院率 40〜50%改善", 20, True, NAVY, 0),
])

# ================= 4枚目：シーン2 =================
s = head(3, "シーン2｜資料請求", "資料請求を、予約・商談につなげる")
flow(s, 5.4, 2.8, [("line", "資料請求"), ("navy", "LINEへ案内"), ("navy", "事例・動画を\n配信"), ("line", "見学・商談の\n予約")])
bullets(s, 9.9, [
    ("・資料請求のみで終わらせず、予約・商談につなげる", 20, False, INK, 18),
    ("・実績：皮膚科クリニック　予約数 約1.2倍", 20, True, NAVY, 14),
    ("・実績：就職支援　LINEで問い合わせた方の55%が面談を予約", 20, True, NAVY, 0),
])

# ---- 5枚目（進め方）を削除 ----
lst = prs.slides._sldIdLst
sid = lst[4]
prs.part.drop_rel(sid.get(qn("r:id")))
lst.remove(sid)

# ---- 図形IDの振り直し ----
for sl in prs.slides:
    for n, c in enumerate(sl.shapes._spTree.iter(qn("p:cNvPr")), start=2):
        c.set("id", str(n))

# ---- 字体をメイリオに統一（latin→ea→cs の順）----
AFTER = {qn(x) for x in ("a:sym", "a:hlinkClick", "a:hlinkMouseOver", "a:rtl", "a:extLst")}
for sl in prs.slides:
    for r in sl.shapes._spTree.iter(qn("a:r")):
        if r.find(qn("a:rPr")) is None:
            r.insert(0, r.makeelement(qn("a:rPr"), {"lang": "ja-JP"}))
    for tag in ("a:rPr", "a:endParaRPr", "a:defRPr"):
        for rpr in sl.shapes._spTree.iter(qn(tag)):
            fs = []
            for ft in ("a:latin", "a:ea", "a:cs"):
                e = rpr.find(qn(ft))
                if e is None:
                    e = rpr.makeelement(qn(ft), {})
                else:
                    rpr.remove(e)
                e.set("typeface", "メイリオ")
                fs.append(e)
            nxt = next((c for c in rpr if c.tag in AFTER), None)
            for e in fs:
                (nxt.addprevious(e) if nxt is not None else rpr.append(e))
prs.save(str(OUT))
print("saved:", OUT.name, len(prs.slides), "枚")
