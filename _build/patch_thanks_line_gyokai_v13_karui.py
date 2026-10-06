# -*- coding: utf-8 -*-
"""サンクスLINE誘導・業界別 ver1.2 → ver1.3（2026-10-06 殿村さん指示）。
新しい2枚目「軽いCVで接点をつくり、LINEで本CVまで引き上げる」を追加（全6枚）。既存5枚は変更しない。
  広告・LP → 簡易診断などの軽いCV（オレンジで強調）→｜→ サンクスLINE誘導 → LINE友だち追加 → LINEでナーチャリング → 本CV → 売上・LTV
上に AD領域／LINE領域 の帯（1枚目と同じ見せ方）、下に3ポイント＋結論1行。業界名・想定数値（60%・50%・30人）は出さない。
  python3 _build/patch_thanks_line_gyokai_v13_karui.py <業界別ver1.2.pptx> [出力]
"""
import copy
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

sys.path.insert(0, str(Path(__file__).resolve().parent))
from pptx_parts import *  # noqa

ROOT = Path(__file__).resolve().parent.parent
OUT = Path(sys.argv[2]) if len(sys.argv) > 2 else ROOT / "20261006_サンクスLINE誘導のご提案_業界別ver1.3.pptx"
ORANGE, ORANGE_BG, ORANGE_TX = "F59B21", "FEF0DE", "C46A00"  # 1枚目（全体座組）と同じ
PALE = "E4E8F6"

prs = Presentation(sys.argv[1])
ref = prs.slides[0]
assert ref.shapes[0].text_frame.text.startswith("広告で反響を獲得し"), "1枚目が全体座組ではない"
s = prs.slides.add_slide(ref.slide_layout)
for sh in list(s.shapes):
    sh._element.getparent().remove(sh._element)
for nm in ("TextBox 1", "TextBox 2", "Connector 3"):
    s.shapes._spTree.append(copy.deepcopy(next(x for x in ref.shapes if x.name == nm)._element))
by = {x.name: x for x in s.shapes}
set_paras(by["TextBox 1"], [("軽いCVで接点をつくり、LINEで本CVまで引き上げる", 0)])
set_paras(by["TextBox 2"], [("検討初期のユーザーとも早い段階で接点を持ち、LINEで継続的に育成", 0)])
lst = prs.slides._sldIdLst  # 2枚目へ移動
el = lst[-1]
lst.remove(el)
lst.insert(1, el)

# ---- 横型フロー（7ステップ）----
L, R = 2.2, 31.7
OV = 0.35
steps = [  # (行, 塗り, 文字色, 幅の比, 強調)
    (["広告・LP"], LGRAY, INK, 3.6, False),
    (["簡易診断などの", "軽いCV"], ORANGE, WHITE, 5.6, True),
    (["サンクス", "LINE誘導"], GREEN, WHITE, 4.2, False),
    (["LINE", "友だち追加"], GREEN_BG, GREEN_TX, 4.2, False),
    (["LINEで", "ナーチャリング"], GREEN_BG, GREEN_TX, 4.6, False),
    (["本CV"], PALE, NAVY, 3.6, False),
    (["売上・LTV"], NAVY, WHITE, 3.8, False),
]
k = (R - L + OV * (len(steps) - 1)) / sum(st[3] for st in steps)
Y, H = 6.3, 2.6
xs = []
x = L
for i, (ls, f, tc, wr, big) in enumerate(steps):
    w = wr * k
    y, h = (Y - 0.4, H + 0.8) if big else (Y, H)
    shape(s, MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON, x, y, w, h, None, fill=f, adj=0.22)
    pad = 0.3 if i == 0 else 0.65
    label(s, x + pad, y, w - pad - 0.55, h, [(t, 15 if big else 13, True, tc, 0) for t in ls], align=PP_ALIGN.CENTER,
          anchor=MSO_ANCHOR.MIDDLE)
    xs.append((x, w))
    x += w - OV

# 領域の帯と境目（1枚目と同じ見せ方）
ad_end = xs[1][0] + xs[1][1] - 0.15
line_start = xs[2][0] + 0.25
rect(s, L, 5.15, ad_end - L, 0.14, ORANGE)
label(s, L, 4.45, 6.0, 0.7, [("AD領域", 13, True, ORANGE_TX, 0)], anchor=MSO_ANCHOR.MIDDLE)
rect(s, line_start, 5.15, R - line_start, 0.14, GREEN)
label(s, line_start, 4.45, 6.0, 0.7, [("LINE領域", 13, True, GREEN_TX, 0)], anchor=MSO_ANCHOR.MIDDLE)
vline(s, (ad_end + line_start) / 2, 4.5, 12.0, LGRAY, 1.25, dash=True)


def tags(i, items, line, color):
    x0, w = xs[i]
    tw = min(w - 0.5, 4.4)
    tx = x0 + (w - tw) / 2 - (0 if i == 0 else 0.15)
    for j, t in enumerate(items):
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, tx, 9.75 + j * 0.85, tw, 0.68, [(t, 10.5, True, color, 0)], fill=WHITE, line=line,
              lw=0.75, adj=0.5, margins=(0.05, 0, 0.05, 0))


tags(1, ["簡易診断", "チェックコンテンツ", "資料DL"], ORANGE, ORANGE_TX)
tags(4, ["情報提供", "リマインド", "個別案内"], GREEN, GREEN_TX)
x5, w5 = xs[5]
label(s, x5 - 0.3, 9.75, w5 + 0.3, 2.4, [(t, 10.5, False, GRAY, 3) for t in ("問い合わせ", "面談・相談", "購入")], align=PP_ALIGN.CENTER)

# ---- 3ポイント（カードにせず、番号＋見出し＋1行を縦線で区切る）----
hline(s, L, R, 12.55, LGRAY, 0.75)
pts = [("CVハードルを下げる", "いきなり問い合わせ・購入を求めず、まず軽い接点をつくる"),
       ("早い段階でLINE接点を持つ", "検討初期のユーザーもLINE上で接点を確保"),
       ("本CVまで育成する", "継続的な情報提供で、問い合わせ・購入などへ引き上げる")]
cw = (R - L) / 3
for i, (h_, d) in enumerate(pts):
    cx = L + i * cw
    if i:
        vline(s, cx, 12.95, 15.0, LGRAY, 0.75)
    shape(s, MSO_SHAPE.OVAL, cx + 0.4, 13.0, 0.85, 0.85, [(str(i + 1), 12, True, WHITE, 0)], fill=NAVY, margins=(0, 0, 0, 0))
    label(s, cx + 1.45, 12.95, cw - 1.8, 0.95, [(h_, 13.5, True, NAVY, 0)], anchor=MSO_ANCHOR.MIDDLE)
    label(s, cx + 1.45, 13.9, cw - 1.8, 1.2, [(d, 10.5, False, INK, 0)])

# ---- 結論（1行だけ）----
label(s, L, 15.45, R - L, 1.3, [("今すぐ客だけでなく、検討初期のユーザーとも接点を持ち、本CVまで育成する。", 17, True, NAVY, 0)],
      align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

prs.save(str(OUT))
print("saved:", OUT.name, len(prs.slides), "枚")
