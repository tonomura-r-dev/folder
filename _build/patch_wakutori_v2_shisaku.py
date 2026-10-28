# -*- coding: utf-8 -*-
"""ワクトリ v1（殿村さんPC保存版・56枚）→ v2：施策の考え方の3ページ（23〜25）に中身を入れる（2026-10-28 殿村さん「派遣 登録 これでいいだろ、進めろ」）。

    python -m _build.patch_wakutori_v2_shisaku <PC保存版v1.pptx> <提案FMT_MAX版.pptx> <出力.pptx>

  v1では差し込み枠にしていた3ページを、提案FMT（MAX版）の15〜17ページの形に戻して中身を入れる。
  見出し・札・リード（v1で直したもの）は v1 のまま。FMTの決まった言葉・図形の形はそのまま（CLAUDE.md 2026-10-27）。
  23：前後検索「派遣 登録」（色分け画像 `python -m _build.make_wakutori_zengo_color`）。FMTの画像の上の小さい札（転職の例）は外す。凡例は3つ
  24：Day0〜15 のステップ配信（5枚）。ボタンはワクトリのリッチメニューの言葉だけ（おすすめ求人／お役立ち情報／解決したいQ&A／体験談／会員登録）
  25：Googleトレンド「リゾートバイト」「スキー場バイト」の月ごとの平均（`python -m _build.make_wakutori_trend_month`）と、時期ごとの企画配信
  FMTの黄色（画像の上の帯）は外す。赤はFMTのまま（CLAUDE.md 2026-10-22）。ほかのページは触らない。
"""
import copy
import io
import sys
from pathlib import Path

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches

from _build.patch_tanaka_v5 import copy_slide
from _build.pptx_parts import drop_slide

ROOT = Path(__file__).resolve().parent.parent
ZENGO = ROOT / "_images/wakutori_zengo_haken_touroku_color.png"
TREND = ROOT / "_images/wakutori_trend_month.png"


def by_id(s, sid):
    return next(sh for sh in s.shapes if sh.shape_id == sid)


def remove(s, sid):
    el = by_id(s, sid)._element
    el.getparent().remove(el)


def set_lines(sh, lines):
    """最初の段落・最初のrunの書式で、1行＝1段落に入れ直す"""
    txb = sh._element.find(qn("p:txBody"))
    ps = txb.findall(qn("a:p"))
    tmpl = copy.deepcopy(ps[0])
    for r in tmpl.findall(qn("a:r"))[1:] + tmpl.findall(qn("a:br")):
        tmpl.remove(r)
    for p in ps:
        txb.remove(p)
    for line in lines:
        p = copy.deepcopy(tmpl)
        p.find(qn("a:r")).find(qn("a:t")).text = line
        txb.append(p)


def set_image(s, sid, path):
    pic = by_id(s, sid)
    blip = pic._element.find(".//" + qn("a:blip"))
    old = blip.get(qn("r:embed"))
    _, rid = s.part.get_or_add_image_part(io.BytesIO(path.read_bytes()))
    blip.set(qn("r:embed"), rid)
    s.part.drop_rel(old)
    return pic


def take_header(new, old):
    """札（ステップ配信／企画配信）とリードは v1 のものを使う"""
    for sid in (8, 10):
        n, o = by_id(new, sid)._element, by_id(old, sid)._element
        n.addprevious(copy.deepcopy(o))
        n.getparent().remove(n)


def fill(sh, hexv, alpha):
    spPr = sh._element.find(qn("p:spPr"))
    old = spPr.find(qn("a:solidFill"))
    sf = spPr.makeelement(qn("a:solidFill"), {})
    c = sf.makeelement(qn("a:srgbClr"), {"val": hexv})
    c.append(c.makeelement(qn("a:alpha"), {"val": str(alpha)}))
    sf.append(c)
    old.addprevious(sf)
    spPr.remove(old)


# ── 23：前後検索「派遣 登録」 ──────────────────────────────────
LEGEND = [(270, "働き方・会社を比べる", "3467B2"), (271, "登録・応募の準備", "8E4EC6"), (272, "お金・保険を調べる", "00897B")]


def page_zengo(s):
    pic = set_image(s, 262, ZENGO)
    pic.name = "前後検索 派遣 登録（色分け）"
    for sid in [sh.shape_id for sh in s.shapes if 274 <= sh.shape_id <= 848]:   # 転職の画像に重ねていた札
        remove(s, sid)
    remove(s, 273)                                            # 凡例の4つ目（黄色）
    for sid, text, c in LEGEND:
        sh = by_id(s, sid)
        set_lines(sh, [text])
        fill(sh, c, 35000)
        sh.width = Inches(1.6)                                # 言葉が長いので横だけ広げる
    x = lambda px: pic.left + pic.width * px / 2000            # 画像の横の位置（2000px）
    band = by_id(s, 267)                                      # LINE追加想定の帯＝0日〜1日
    band.left, band.width = int(x(1018)), int(x(1257) - x(1018))
    by_id(s, 268).left = band.left - Inches(0.15)
    by_id(s, 269).left = band.left + Inches(0.21)
    pre, post = by_id(s, 263), by_id(s, 264)                  # 検索の前・後の帯を画像の軸に合わせる
    pre.left, pre.width = int(x(77)), int(band.left - Inches(0.12) - x(77))
    post.left, post.width = band.left, int(x(1984) - band.left)
    kw, cap = by_id(s, 852), by_id(s, 851)                    # キーワードの札は右上の空いたところへ
    set_lines(kw, ["派遣 登録"])
    kw.left, kw.top = Inches(6.15), Inches(2.45)
    cap.left, cap.top = kw.left + Inches(0.27), kw.top + Inches(0.33)


# ── 24：Day0〜15 のステップ配信 ─────────────────────────────
CARDS = [  # 見出しの各ID（Day・役割・メッセージの見出し・中身・検索データ・ボタン）, 文字
    ((185, 187, 189, 191, 197, 201),
     ("Day0", "追加直後：診断", ["ご登録ありがとう", "ございます！"],
      ["3問タップで", "開始時期・エリア・", "重視したいポイントを分岐", "・タグ付け", "（今のアンケートを発展）"],
      ["--", "（追加直後は検索データなし。アンケートの回答で配信を出し分け）"], "おすすめ求人")),
    ((207, 209, 211, 213, 219, 223),
     ("Day1〜2", "働き方の比較", ["リゾートバイトの", "働き方がわかる"],
      ["住み込みで働く流れと、", "食費無料・Wi-Fi完備など", "寮・生活の条件を紹介", "派遣・単発と比べて", "選ばれる理由を伝える"],
      ["検索の直前・当日に「派遣会社」「単発バイト」「リゾートバイト」などを検索。働き方を比べている"], "お役立ち情報")),
    ((229, 231, 235, 237, 247, 253),
     ("Day3〜4", "お金・保険の不安", ["働く前の", "お金の疑問Q&A"],
      ["保険料・年金はどうなる？", "扶養内で働ける金額は？", "など、働く前に気になる", "お金の疑問に", "先回りして答える"],
      ["検索の当日〜1日後に「扶養内で働ける金額 2025」「国民健康保険料」など、お金・保険の検索が続く"], "解決したいQ&A")),
    ((259, 261, 263, 265, 271, 275),
     ("Day6〜7", "会社の比較・選び方", ["お仕事探しの", "会社の選び方"],
      ["専属エージェントの", "サポートを、他社との", "違いとして伝える", "体験談で信頼を積み上げる"],
      ["当日〜3日後に「派遣会社 おすすめ企業」「バイトル」「スタッフサービス」「エン派遣」など、会社・サービスを比べる検索が続く"],
      "体験談")),
    ((281, 283, 285, 287, 293, 297),
     ("Day10〜15", "ENTRYの案内", ["かんたん1分", "ENTRYへ"],
      ["面接・履歴書作成ナシ", "専属エージェントが", "お仕事中も終了後も", "サポート", "開始時期・エリアに合う求人"],
      ["2〜7日後に「派遣 職場見学 落ちる確率」「履歴書 書き方」など、応募・面接の準備の検索が増える"], "会員登録")),
]
SUMMARY = (["検索当日の働き方の比較（Day1〜2）→ お金・保険の不安の解消（Day3〜4）→ 会社選び（Day6〜7）"],
           ["→ ", "ENTRYの案内", "（Day10〜15）と、", "段階的に信頼を積み上げながらENTRYへつなげる", "設計。"])


def page_step(s):
    for ids, (day, role, head, body, data, btn) in CARDS:
        for sid, val in zip(ids, (day, role, head, body, data, btn)):
            set_lines(by_id(s, sid), val if isinstance(val, list) else [val])
    ps = by_id(s, 915)._element.find(qn("p:txBody")).findall(qn("a:p"))
    for p, texts in zip(ps, SUMMARY):
        runs = p.findall(qn("a:r"))
        assert len(runs) == len(texts)
        for r, t in zip(runs, texts):
            r.find(qn("a:t")).text = t


# ── 25：Googleトレンドと時期ごとの企画配信 ───────────────────────
MONTHS = {3: 1, 7: 2, 13: 3, 17: 4, 21: 5, 24: 6, 27: 7, 31: 8, 49: 9, 51: 10, 41: 11, 55: 12}
NAVY_M, RED_M = {12, 1, 6, 7}, {9}
TEXTS25 = {
    12: ["リゾートバイト"],
    6: ["・10〜11月：「12月・1月開始のスキー場の求人」特集"],
    15: ["・「2026-2027 winter SKI Resort」の求人を、開始時期・エリアの答えで出し分け"],
    23: ["・冬の終わりに、春からの次のお仕事を案内"],
    39: ["夏（6〜7月）は「リゾートバイト」の検索が最も多い。"],
    42: ["海・離島・高原など、夏のエリアの求人特集。"],
    11: ["希望のエリア・開始時期の答えで出し分け。"],
    46: ["・9月：冬の求人の先行案内", "・4〜5月：夏の求人の先行案内",
         "・契約満了の前：次のお仕事の案内", "・「LINE友だち限定・先行公開」"],
}


def page_trend(s):
    set_image(s, 5, TREND).name = "Googleトレンド リゾートバイト・スキー場バイト（月ごと）"
    for sid, lines in TEXTS25.items():
        set_lines(by_id(s, sid), lines)
    remove(s, 14)                                             # 黄色の帯（グラフの中に薄い帯を入れた）
    remove(s, 16)
    red = by_id(s, 38)                                        # 赤の帯＝9月（冬の求人の先行案内）
    red.left, red.top, red.width, red.height = Inches(6.76), Inches(2.03), Inches(0.56), Inches(1.74)
    arw = by_id(s, 36)
    arw.left, arw.top, arw.width, arw.height = Inches(7.04), Inches(3.79), Inches(1.0), Inches(0.62)
    tmpl = {"red": by_id(s, 3)._element, "navy": by_id(s, 21)._element, "gray": by_id(s, 17)._element}
    tmpl = {k: copy.deepcopy(v) for k, v in tmpl.items()}
    for sid, m in MONTHS.items():
        el = by_id(s, sid)._element
        kind = "navy" if m in NAVY_M else "red" if m in RED_M else "gray"
        new = copy.deepcopy(tmpl[kind])
        new.find(".//" + qn("p:cNvPr")).set("id", el.find(".//" + qn("p:cNvPr")).get("id"))
        new.find(qn("p:spPr")).replace(new.find(qn("p:spPr")).find(qn("a:xfrm")),
                                       copy.deepcopy(el.find(qn("p:spPr")).find(qn("a:xfrm"))))
        for r in new.iter(qn("a:t")):
            r.text = ""
        new.find(".//" + qn("a:t")).text = f"{m}月"
        el.addprevious(new)
        el.getparent().remove(el)


def main(src, fmt, out):
    prs = Presentation(src)
    assert len(prs.slides) == 56
    f = Presentation(fmt)
    lst = prs.slides._sldIdLst
    olds = [prs.slides[p - 1] for p in (23, 24, 25)]
    for old, fpage, edit in zip(olds, (15, 16, 17), (page_zengo, page_step, page_trend)):
        assert by_id(old, 4).text_frame.text.startswith("施策の考え方")
        new = copy_slide(prs, f.slides[fpage - 1])            # 先に3枚とも足す（消してから足すと部品名が重なる）
        take_header(new, old)
        edit(new)
        el = list(lst)[-1]
        lst.remove(el)
        next(e for e in lst if int(e.get("id")) == old.slide_id).addprevious(el)
    for old in olds:
        drop_slide(prs, old)
    assert len(prs.slides) == 56
    prs.save(out)
    print("saved:", out)


if __name__ == "__main__":
    main(*sys.argv[1:4])
