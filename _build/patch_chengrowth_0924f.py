# -*- coding: utf-8 -*-
"""S20の4列目（根拠）を削除して3列に戻す＋S20/S21の行ID・列IDの重複を解消（PowerPointで列が空になる不具合）。
  python3 _build/patch_chengrowth_0924f.py <pptx>
"""
import sys
from pptx import Presentation
from pptx.util import Cm
from pptx.oxml.ns import qn

F = sys.argv[1]
prs = Presentation(F)
A16 = "{http://schemas.microsoft.com/office/drawing/2014/main}"


def fix_ids(tbl):
    for tag, attr, base in (("a:gridCol", "colId", 20000), ("a:tr", "rowId", 10000)):
        els = tbl.findall(".//" + qn(tag))
        for i, el in enumerate(els):
            idel = el.find(".//" + A16 + attr)
            if idel is not None:
                idel.set("val", str(base + i))


# S20：4列目を削除
gf = [sh for sh in prs.slides[19].shapes if getattr(sh, "has_table", False) and sh.has_table][0]
tbl = gf.table._tbl
grid = tbl.tblGrid
cols = grid.findall(qn("a:gridCol"))
if len(cols) == 4:
    grid.remove(cols[3])
    for tr in tbl.findall(qn("a:tr")):
        tr.remove(tr.findall(qn("a:tc"))[3])
for gc, w in zip(grid.findall(qn("a:gridCol")), (3.02, 5.03, 17.1)):
    gc.set("w", str(int(Cm(w))))
gf.width = Cm(25.15)
fix_ids(tbl)

# S21：行IDの重複を解消
gf21 = [sh for sh in prs.slides[20].shapes if getattr(sh, "has_table", False) and sh.has_table][0]
fix_ids(gf21.table._tbl)
prs.save(F)
print("ok")
