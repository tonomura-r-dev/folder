# -*- coding: utf-8 -*-
"""DYM共通資料（LINEOA_BUFFF_3.pptx）の会社紹介ページを、提案資料の先頭に差し込む。2026-10-13 殿村さん指示。
既定は共通資料の2・3・4枚目（受賞・実績／弊社ポジショニング／弊社の強み）。中身は一切変えずにそのまま写す。
同じスライドサイズ・同じレイアウト名（4_タイトルとコンテンツ）の資料どうしで使う。画像は差し込み先に新しく登録し直す。
  python3 _build/insert_company_intro.py <提案資料.pptx> <LINEOA_BUFFF_3.pptx> [出力.pptx] [--pages=2,3,4]
"""
import copy
import io
import sys
from pathlib import Path

from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

args = [a for a in sys.argv[1:] if not a.startswith("--")]
opts = dict(a[2:].split("=", 1) for a in sys.argv[1:] if a.startswith("--"))
DST, SRC = Path(args[0]), Path(args[1])
OUT = Path(args[2]) if len(args) > 2 else DST
PAGES = [int(x) for x in opts.get("pages", "2,3,4").split(",")]
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

src = Presentation(str(SRC))
dst = Presentation(str(DST))
assert (src.slide_width, src.slide_height) == (dst.slide_width, dst.slide_height), "スライドサイズが違う"
layouts = {l.name: l for l in dst.slide_layouts}

new_ids = []
for n in PAGES:
    ss = src.slides[n - 1]
    ds = dst.slides.add_slide(layouts[ss.slide_layout.name])
    for sh in list(ds.shapes):  # レイアウトから入る空のプレースホルダーは消す
        sh._element.getparent().remove(sh._element)
    # 画像などの関係を張り直す（旧rId → 新rId）
    rmap = {}
    for rId, rel in ss.part.rels.items():
        if rel.reltype == RT.SLIDE_LAYOUT:
            continue
        if rel.is_external:
            rmap[rId] = ds.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
        elif rel.reltype == RT.IMAGE:
            _, rmap[rId] = ds.part.get_or_add_image_part(io.BytesIO(rel.target_part.blob))
        else:
            raise SystemExit(f"未対応の関係: {rel.reltype}（{n}枚目）")
    tree = ds.shapes._spTree
    for el in ss.shapes._spTree:
        if el.tag.endswith("}nvGrpSpPr") or el.tag.endswith("}grpSpPr"):
            continue
        tree.append(copy.deepcopy(el))
    for el in tree.iter():
        for k, v in list(el.attrib.items()):
            if k.startswith("{%s}" % R_NS) and v in rmap:
                el.set(k, rmap[v])
    # 背景（スライドに直接設定されている場合）
    bg = ss._element.cSld.bg
    if bg is not None:
        ds._element.cSld.insert(0, copy.deepcopy(bg))
    new_ids.append(dst.slides._sldIdLst[-1])

# 先頭へ移動（元の順番のまま）
lst = dst.slides._sldIdLst
for i, el in enumerate(new_ids):
    lst.remove(el)
    lst.insert(i, el)

dst.save(str(OUT))
print("saved:", OUT.name, len(dst.slides), "枚（先頭に共通資料の", PAGES, "枚目を差し込み）")
