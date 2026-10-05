# -*- coding: utf-8 -*-
"""チェングロウス ver2.0 → ver2.1（2026-10-05 殿村さん指示）。
サンクスLINE誘導の説明スライドを1枚追加（15枚目＝「友だち獲得の2つの導線」の次）。
応募CVを増やす施策ではなく「応募者を友だち化し、面談・来店まで接点を残す」施策として見せる。SIMの係数は載せない。
  python3 _build/patch_chengrowth_bridge_v21_thanks.py <ver2.0.pptx>
"""
import copy
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

sys.path.insert(0, str(Path(__file__).resolve().parent))
from pptx_parts import *  # noqa

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "20261005_チェングロウス_サイトリニューアルとLINE同時導入のご提案ver2.1.pptx"
prs = Presentation(sys.argv[1])
ref = prs.slides[13]  # 友だち獲得の2つの導線
assert ref.shapes[0].text_frame.text.startswith("友だち獲得の2つの導線")
s = prs.slides.add_slide(ref.slide_layout)
for sh in list(s.shapes):
    sh._element.getparent().remove(sh._element)
for nm in ("TextBox 1", "TextBox 2", "Connector 3"):
    s.shapes._spTree.append(copy.deepcopy(next(x for x in ref.shapes if x.name == nm)._element))
by = {x.name: x for x in s.shapes}
set_paras(by["TextBox 1"], [("応募後もLINEでつながる｜サンクスLINE誘導", 0)])
set_paras(by["TextBox 2"], [("応募完了直後にLINEへ誘導し、面談・来店までの接点を残す", 0)])
PALE = "E4E8F6"

# メイン図：応募完了 → LINE友だち追加 → 応募受付・日程案内 → 前日リマインド → 面談・来店
W, STEP, Y0, H = 4.9, 4.55, 5.6, 2.8
steps = [(["応募完了"], LGRAY, INK, 14), (["LINE", "友だち追加"], GREEN, WHITE, 16), (["応募受付・", "日程案内"], PALE, NAVY, 13.5),
         (["前日", "リマインド"], PALE, NAVY, 13.5), (["面談・来店"], NAVY, WHITE, 14)]
for i, (ls, f, tc, sz) in enumerate(steps):
    x = 2.2 + i * STEP
    big = i == 1
    y, h = (Y0 - 0.45, H + 0.9) if big else (Y0, H)
    shape(s, MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON, x, y, W, h, None, fill=f, adj=0.25)
    label(s, x + (0.3 if i == 0 else 0.75), y, W - (1.0 if i == 0 else 1.45), h, [(t, sz, True, tc, 0) for t in ls],
          align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
label(s, 2.2, 4.35, 9.4, 0.7, [("サンクスページで「今後の連絡はLINEで」と案内", 10.5, False, GRAY, 0)], align=PP_ALIGN.CENTER,
      anchor=MSO_ANCHOR.MIDDLE)
# 下の区切り：友だちとして蓄積 ／ LINEで応募後フォロー
x1, x2 = 2.2 + STEP + 0.3, 2.2 + STEP + W - 0.3
rect(s, x1, 9.25, x2 - x1, 0.12, GREEN)
label(s, x1 - 0.4, 9.5, x2 - x1 + 0.8, 0.8, [("友だち化", 13, True, GREEN_TX, 0)], align=PP_ALIGN.CENTER)
x3, x4 = 2.2 + 2 * STEP + 0.3, 25.3
rect(s, x3, 9.25, x4 - x3, 0.12, NAVY)
label(s, x3, 9.5, x4 - x3, 0.8, [("LINEで応募後フォロー", 13, True, NAVY, 0)], align=PP_ALIGN.CENTER)

# 3つのポイント（カードにせず、番号＋短文を横に並べる）
pts = [("友だちとして蓄積", "応募で終わらせず、LINEに接点を残す"),
       ("応募後の連絡をLINEで", "受付・日時・場所・持ち物・日程変更を案内"),
       ("面談・来店までフォロー", "前日リマインドで来店まで伴走")]
for i, (h, d) in enumerate(pts):
    x = 2.2 + i * 7.75
    if i:
        vline(s, x - 0.2, 10.9, 13.2, LGRAY, 0.75)
    shape(s, MSO_SHAPE.OVAL, x + 0.1, 11.05, 1.0, 1.0, [(f"{i + 1}", 13, True, WHITE, 0)], fill=GREEN if i == 0 else NAVY,
          margins=(0, 0, 0, 0))
    label(s, x + 1.35, 10.9, 6.0, 2.5, [(h, 13.5, True, NAVY, 3), (d, 11, False, INK, 0)], anchor=MSO_ANCHOR.TOP)
conclusion(s, 14.0, "応募で接点を終わらせず、LINEで面談・来店までつなげる", 17)
label(s, 2.2, 16.6, 23.1, 0.6, [("※求人ボックス・Indeed経由については、着地先・応募完了先を確認したうえで導入可否を設計", 9, False, GRAY, 0)])

# 15枚目に移動
lst = prs.slides._sldIdLst
el = lst[-1]
lst.remove(el)
lst.insert(14, el)
prs.save(str(OUT))
print("saved:", OUT.name, len(prs.slides), "枚")
