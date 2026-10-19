# -*- coding: utf-8 -*-
"""塾ver3：前後検索をジムの形に（1キーワード1枚）＋設計図②の件数を整数に（2026-10-19 殿村さん）。

    python -m _build.juku_fmt.zengo4 <塾ver3（55枚）> <出力.pptx>

  殿村さん：「S8の表記いややな、ちゃんとジムの時の表記にしてほしい。前後検索と、色付けして、あとでこういう配信する、って」
           「前後検索一枚につき、スライド一枚使用してかまわない」
           「（設計図②に）CV件数少数は成立しないだろ、修正して」
  - 8〜10枚目（検索の流れ／中学受験・中高一貫／高校受験・英語）→ 8〜11枚目の4枚（中学受験 塾／中高一貫 塾／高校受験 塾／英語 塾）。
    1枚の中身＝色分けした前後検索の画像（左）＋4つの分類と、その検索ワード（右）＋「この流れで配信すること」（下・3つ）。
    ジムの「前後検索」（画像＋①〜④の分類）と「ステップ配信のブレスト」（画像＋何日目に何を送るか）を1枚にした形。
    配信は14日間のステップ配信（P31）の中身に合わせる。
  - 設計図②（38枚目）：1教室あたりだと件数が小数（0.3件など）になる → 10教室のモデル（新しい友だち 月2,500人）にして整数に。
    合計は資料請求5件・体験予約10件（100教室なら50件・100件は今までと同じ）。
  - 本文のページ参照：P8 → P8〜P11、P12以降は+1。
  ほかのページは1バイトも変えない（直したスライドのXMLと、足した4枚だけを差し込む）。
"""
import random
import re
import sys
import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn

import _build.juku_fmt.build as jb
from _build.arauchi_v4.build import slide_order
from _build.arauchi_v4.lib import add_box
from _build.arauchi_v8.pages import replace_in_paragraphs, shape_by_id
from _build.gym_curves.build import add_pic, keep_only, place, round_box, set_one
from _build.gym_fmt.build import fix_fonts
from _build.juku_fmt.build import CT_SLIDE, RT_SLIDE, rels_of
from _build.juku_fmt.meiko import BORDER, GREY, INK, NAVY, PALE, W, X0, title_of
from _build.juku_fmt.plan import set_lead

E = 914400
N_BASE = 55
FLOW, PAIR1, PAIR2, PLAN2, STEP = 8, 9, 10, 37, 30          # 55枚の番号
DROP = [FLOW, PAIR1, PAIR2]
ORDER = list(range(1, FLOW)) + [("N", k) for k in range(1, 5)] + list(range(PAIR2 + 1, N_BASE + 1))
NEW_POS = {key: i for i, key in enumerate(ORDER, 1)}
RT_IMAGE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"
COL = {"A": "3467B2", "B": "8E4EC6", "C": "C2185B", "D": "00897B"}       # 画像の色分けと同じ（黄・オレンジは使わない）
CAT = {"A": "① 志望校・偏差値", "B": "② 検定・勉強法", "C": "③ 塾選び・費用・時期", "D": "④ 塾名・料金・口コミ"}
IMG_W = 6.17
IMG_H = IMG_W * 1167 / 2000

# 1枚ずつの中身：(起点, 画像の元のページと並び, リード2行, 分類ごとの検索ワード, 配信3つ, 帯)
PAGES = [
    ("中学受験 塾", (PAIR1, 0),
     ["中学受験の保護者は、検索の前に志望校・偏差値・算数検定を調べ、当日は塾の名前に集中する。",
      "そのあとは、入室テスト・料金・口コミで比べている。"],
     {"A": "私立中学 偏差値ランキング／海城中学・女子学院／栄東中学校 入試／中学受験 偏差値40",
      "B": "算数検定／予習シリーズ／中学受験 社会 勉強方法／中学受験 親の関わり方",
      "C": "中学受験 塾選び／塾 ランキング／個別指導のみ／塾 費用／4年生から始める",
      "D": "日能研・早稲田アカデミー・サピックス／サピックス 入室テスト／日能研 費用／明光義塾 口コミ"},
     [("追加直後｜あいさつ＋30秒診断", "志望校・偏差値を調べている → 学年と志望を聞き、入試までの逆算を届ける"),
      ("5日後｜月謝＋講習費の年間の目安", "塾選び・費用・いつからを比べている → 費用の目安と始める時期を先に出す"),
      ("7日後｜事例＋体験のご案内①", "入室テスト・口コミで比べている → 同じタイプの子の事例と、体験授業の案内")],
     "塾の名前を調べる前（志望校・偏差値を調べる時期）に、LINEでつながっておく。"),
    ("中高一貫 塾", (PAIR1, 1),
     ["中高一貫の保護者は、都立中高一貫校・中等教育学校の名前と入試日程を、15日前から調べている。",
      "塾の名前（ena・栄光ゼミナール・日能研）は当日に集中し、そのあとは料金・合格実績で比べている。"],
     {"A": "都立中高一貫校／立川国際・三鷹・九段中等教育学校／都立高校 入試日程／都立推薦 日程",
      "B": "声の教育社（過去問）／四谷大塚 予習シリーズ",
      "C": "中学受験 塾選び／中学受験 塾 費用／中学受験 全落ち／中高一貫 塾",
      "D": "ena・栄光ゼミナール・日能研／栄光ゼミナール 料金／早稲田アカデミー 合格実績／森塾 口コミ"},
     [("追加直後｜あいさつ＋30秒診断", "学校名・入試日程を調べている → 志望校を聞き、入試までの逆算を届ける"),
      ("5日後｜月謝＋講習費の年間の目安", "塾の費用・全落ちを心配している → 年間の費用の目安と、準備の進め方"),
      ("7日後｜事例＋体験のご案内①", "料金・合格実績で比べている → 同じ志望の子の事例と、体験授業の案内")],
     "学校名と日程を調べ始める15日前から、志望校の情報でつながる。"),
    ("高校受験 塾", (PAIR2, 0),
     ["高校受験の保護者は、当日に「塾 いつから・費用・ランキング・勉強方法」をまとめて調べている。",
      "その前は志望校の偏差値、そのあとは塾の名前と料金で比べている。"],
     {"A": "都立高校／高校受験 偏差値／東京都 高校 偏差値／都立高校 入試",
      "B": "高校受験 勉強方法／数検3級／高校受験 塾なし／中学 英語",
      "C": "高校受験 塾 いつから／塾 費用／塾 ランキング／中学受験 高校受験 どっち／オンライン 塾",
      "D": "早稲田アカデミー・栄光ゼミナール・臨海セミナー／早稲アカ 個別 料金／臨海セミナー 口コミ"},
     [("3日後｜家でできる勉強のコツ", "勉強方法・塾なしで迷っている → 売り込まずに、家でできる勉強のコツ"),
      ("5日後｜月謝＋講習費の年間の目安", "いつから・費用を調べている → 年間の費用の目安と、始める時期"),
      ("14日後｜テストから逆算＋体験のご案内②", "塾の名前・料金で比べている → 次の定期テストから逆算した体験授業の案内")],
     "「いつから・費用」を調べる当日に、その答えをLINEで先に届ける。"),
    ("英語 塾", (PAIR2, 1),
     ["英語の保護者は、塾より先に英検（日程・級・合格発表）を調べている。",
      "当日は英語塾・個別指導塾の名前、そのあとは料金・口コミで比べている。"],
     {"A": "中高一貫校／大学受験／開成高校／高校受験",
      "B": "英検 日程／英検3級・2級・準1級／英検 合格発表／漢検／スタディサプリ／英単語",
      "C": "個別指導塾／英語塾 中高一貫校／高校受験 塾／小学生 英語 塾",
      "D": "武田塾・明光義塾・森塾／英検アカデミー 料金・口コミ／武田塾 料金"},
     [("追加直後｜あいさつ＋30秒診断", "英検の日程・級を調べている → 学年と目指す級を聞き、次の英検までの逆算を届ける"),
      ("3日後｜家でできる勉強のコツ", "英単語・勉強法を調べている → 家でできる英語の勉強のコツ"),
      ("7日後｜事例＋体験のご案内①", "料金・口コミで比べている → 同じ級を目指した子の事例と、体験授業の案内")],
     "英検の日程を調べるタイミングで、英検の情報からつながる。"),
]


# ---------------------------------------------------------------- 1) 組み立て（zipのまま）
def assemble(base, out):
    z = zipfile.ZipFile(base)
    order = slide_order(z)
    assert len(order) == N_BASE
    for n, t in [(FLOW, "前後検索｜検索の流れ"), (PAIR1, "前後検索｜中学受験・中高一貫の保護者"),
                 (PAIR2, "前後検索｜高校受験・英語の保護者"), (PLAN2, "施策の設計図②　効率改善（ステップ配信・企画配信）"),
                 (STEP, "ステップ配信｜14日間の配信")]:
        assert title_of(z, order[n - 1]) == t, (n, title_of(z, order[n - 1]))
    names = set(z.namelist())
    ct = z.read("[Content_Types].xml").decode("utf-8")
    prels = z.read("ppt/_rels/presentation.xml.rels").decode("utf-8")
    pres = z.read("ppt/presentation.xml").decode("utf-8")

    def rel_entry(name):
        m = re.search(r'<Relationship [^>]*Target="slides/%s"[^>]*/>' % re.escape(name.split("/")[-1]), prels)
        return m.group(0), re.search(r'Id="(rId\d+)"', m.group(0)).group(1)

    rid_of = {n: rel_entry(order[n - 1])[1] for n in range(1, N_BASE + 1)}
    next_rid = max(int(x) for x in re.findall(r'Id="rId(\d+)"', prels)) + 1
    next_slide = max(int(m.group(1)) for n in names for m in [re.match(r"ppt/slides/slide(\d+)\.xml$", n)] if m) + 1

    # 画像：元の2枚組のページの、絵の並び（rIdの順＝左・右）
    def images_of(no):
        xml = z.read(order[no - 1]).decode("utf-8")
        rel = {r.get("Id"): r.get("Target") for r in etree.fromstring(z.read(rels_of(order[no - 1])))}
        return [rel[r] for r in re.findall(r'<a:blip r:embed="(rId\d+)"', xml)]
    pics = {PAIR1: images_of(PAIR1), PAIR2: images_of(PAIR2)}
    assert all(len(v) == 2 for v in pics.values()), pics

    tpl_name = order[PAIR1 - 1]
    tpl = z.read(tpl_name).decode("utf-8")
    tpl = re.sub(r"<p:pic>.*?</p:pic>", "", tpl, flags=re.S)                 # 絵は外して、1枚ずつ足し直す
    layout = [r for r in etree.fromstring(z.read(rels_of(tpl_name))) if r.get("Type").endswith("/slideLayout")]
    assert len(layout) == 1
    used = set(re.findall(r'creationId[^>]*val="(\d+)"', "".join(z.read(o).decode("utf-8") for o in order)))
    rnd = random.Random(202610194)
    add = {}
    for k, (_, (src, idx), *_rest) in enumerate(PAGES, 1):
        dst = f"ppt/slides/slide{next_slide}.xml"
        next_slide += 1
        cid = str(rnd.randrange(10 ** 9, 4 * 10 ** 9))
        assert cid not in used
        xml, c = re.subn(r'(<p14:creationId [^>]*val=")\d+(")', rf"\g<1>{cid}\2", tpl, count=1)
        assert c == 1
        add[dst] = xml.encode("utf-8")
        rels = etree.Element("{http://schemas.openxmlformats.org/package/2006/relationships}Relationships")
        etree.SubElement(rels, "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship",
                         Id="rId1", Type=layout[0].get("Type"), Target=layout[0].get("Target"))
        etree.SubElement(rels, "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship",
                         Id="rId101", Type=RT_IMAGE, Target=pics[src][idx])
        add[rels_of(dst)] = etree.tostring(rels, xml_declaration=True, encoding="UTF-8", standalone=True)
        ct = ct.replace("</Types>", f'<Override PartName="/{dst}" ContentType="{CT_SLIDE}"/></Types>')
        rid = f"rId{next_rid}"
        next_rid += 1
        prels = prels.replace("</Relationships>", f'<Relationship Id="{rid}" Type="{RT_SLIDE}" Target="slides/{dst.split("/")[-1]}"/></Relationships>')
        rid_of[("N", k)] = rid

    gone = set()
    for no in DROP:                               # 外す3枚（ほかから参照されていないこと）
        old = order[no - 1]
        for r in etree.fromstring(z.read(rels_of(old))):
            assert r.get("Type").endswith(("/slideLayout", "/image")), r.get("Type")
        for n in names:
            if n.endswith(".rels") and n not in (rels_of(old), "ppt/_rels/presentation.xml.rels"):
                assert old.split("/")[-1] not in z.read(n).decode("utf-8"), n
        gone |= {old, rels_of(old)}
        prels = prels.replace(rel_entry(old)[0], "")
        ct, c = re.subn(r'<Override [^>]*PartName="/%s"[^>]*/>' % re.escape(old), "", ct)
        assert c == 1

    m = re.search(r"<p:sldIdLst>(.*?)</p:sldIdLst>", pres, re.S)
    entries = {re.search(r'r:id="(rId\d+)"', e).group(1): e for e in re.findall(r"<p:sldId [^>]*/>", m.group(1))}
    next_id = max(int(x) for x in re.findall(r'<p:sldId id="(\d+)"', m.group(1))) + 1
    lst = []
    for key in ORDER:
        if rid_of[key] in entries:
            lst.append(entries[rid_of[key]])
        else:
            lst.append(f'<p:sldId id="{next_id}" r:id="{rid_of[key]}"/>')
            next_id += 1
    pres = pres.replace(m.group(0), "<p:sldIdLst>" + "".join(lst) + "</p:sldIdLst>")

    app = z.read("docProps/app.xml").decode("utf-8")
    n_new, ppt_t = len(ORDER), "<vt:lpstr>PowerPoint プレゼンテーション</vt:lpstr>"
    assert f"<Slides>{N_BASE}</Slides>" in app and app.count(ppt_t) == N_BASE
    app = app.replace(f"<Slides>{N_BASE}</Slides>", f"<Slides>{n_new}</Slides>")
    app, c = re.subn(r"(<vt:lpstr>スライド タイトル</vt:lpstr></vt:variant><vt:variant><vt:i4>)\d+(</vt:i4>)", rf"\g<1>{n_new}\2", app)
    assert c == 1
    mv = re.search(r'<vt:vector size="(\d+)" baseType="lpstr">', app)
    app = app.replace(mv.group(0), f'<vt:vector size="{int(mv.group(1)) + n_new - N_BASE}" baseType="lpstr">')
    i = app.rfind(ppt_t) + len(ppt_t)
    app = app[:i] + ppt_t * (n_new - N_BASE) + app[i:]
    replace = {"ppt/presentation.xml": pres, "ppt/_rels/presentation.xml.rels": prels,
               "[Content_Types].xml": ct, "docProps/app.xml": app}

    # 外したページだけが使っていた画像（検索の流れの凡例など）は外す
    keep_targets = set()
    for n in names:
        if n.endswith(".rels") and n not in gone:
            keep_targets |= {t.split("/")[-1] for t in re.findall(r'Target="([^"]+)"', z.read(n).decode("utf-8"))}
    for n in add:
        if n.endswith(".rels"):
            keep_targets |= {t.split("/")[-1] for t in re.findall(r'Target="([^"]+)"', add[n].decode("utf-8"))}
    media_gone = {n for n in names if n.startswith("ppt/media/") and n.split("/")[-1] not in keep_targets}
    gone |= media_gone

    with zipfile.ZipFile(out, "w") as zo:
        for info in z.infolist():
            if info.filename in gone:
                continue
            data = replace.get(info.filename)
            zo.writestr(info, data.encode("utf-8") if data is not None else z.read(info.filename),
                        compress_type=info.compress_type)
        for name, data in add.items():
            zo.writestr(name, data, compress_type=zipfile.ZIP_DEFLATED)
    return sorted(media_gone)


# ---------------------------------------------------------------- 2) 中身
def page_zengo(s, key, lead, cats, deliver, band):
    keep_only(s, {2, 3, 4, 5, 6, 7, 8, 9, 10, 11, 12, 24, 25})
    set_one(s, 2, f"前後検索｜{key}")
    set_lead(s, lead)
    add_box(s, X0, 1.6, W, 0.28, [[(f"「{key}」の検索前後（LINEヤフー前後検索・実データ）", {})]], size=10.5, bold=True,
            color="FFFFFF", fill=NAVY, align="ctr", ins=(0.05, 0, 0.05, 0))
    add_pic(s, "rId101", f"前後検索_{key}", X0, 1.96, IMG_W, IMG_H)
    px, pw = X0 + IMG_W + 0.12, W - IMG_W - 0.12
    paras = []
    for k in "ABCD":
        paras.append([(CAT[k], {"size": 10.5, "bold": True, "color": COL[k]})])
        paras.append([(cats[k], {"size": 9, "color": INK})])
    add_box(s, px, 1.96, pw, IMG_H, paras, fill=PALE, line=BORDER, anchor="ctr", ins=(0.1, 0.05, 0.08, 0.05),
            line_spacing=105)
    y = 1.96 + IMG_H + 0.08
    lw = 1.05
    round_box(s, X0, y, lw, 0.72, [[("この流れで", {})], [("配信すること", {})]], adj=8000, size=10, bold=True,
              color="FFFFFF", fill=NAVY, align="ctr", ins=(0.04, 0, 0.04, 0), line_spacing=105)
    cw = (W - lw - 3 * 0.08) / 3
    for i, (head, body) in enumerate(deliver):
        x = X0 + lw + 0.08 + i * (cw + 0.08)
        add_box(s, x, y, cw, 0.72, [[(head, {"size": 9.5, "bold": True, "color": NAVY})], [(body, {"size": 9, "color": INK})]],
                fill="FFFFFF", line=BORDER, anchor="ctr", ins=(0.08, 0.03, 0.06, 0.03), line_spacing=105)
    band_sh = shape_by_id(s, 24)
    set_one(s, 24, band)
    place(band_sh, y=6.42, h=0.36)
    for r in band_sh._element.iter(qn("a:rPr"), qn("a:endParaRPr")):     # 帯の文字は12pt（元の型は大きすぎる）
        r.set("sz", "1200")
    set_one(s, 25, f"出典｜LINEヤフー前後検索（2026年10月取得）。横軸＝起点の検索の前後15日、縦軸＝検索した人数（UU）。"
                   f"色はDYMの分類（薄い文字は分類の対象外）。配信の中身はP{NEW_POS[STEP]}・P{NEW_POS[PLAN2]}")
    return s


def remap_refs(s):
    """本文中の「P◯」（55枚の番号）を新しい番号に。P8（検索の流れ）→ P8〜P11"""
    def one(m):
        a, b = int(m.group(1)), m.group(2)
        if b is not None:                                  # 範囲（P7〜P11 など）：終わりが8〜10なら P11 まで
            b = int(b)
            return f"P{a if a <= FLOW else NEW_POS[a]}〜P{11 if b in DROP else NEW_POS[b]}"
        if a == FLOW:
            return "P8〜P11"
        assert a not in (PAIR1, PAIR2), m.group(0)
        return f"P{NEW_POS[a]}"
    for t in s._element.iter(qn("a:t")):
        if t.text and re.search(r"P\d+", t.text):
            t.text = re.sub(r"P(\d+)(?:〜P(\d+))?", one, t.text)


def set_para(sh, idx, old, new):
    p = sh.text_frame.paragraphs[idx]
    assert p.text == old, (sh.shape_id, p.text)
    p.runs[0].text = new
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)


def integer_counts(s):
    """設計図②：10教室のモデル（新しい友だち 月2,500人）にして、件数を整数に"""
    by = {sh.shape_id: sh for sh in s.shapes}
    assert replace_in_paragraphs(s, "（新しい友だち 月250人）", "（新しい友だち 月2,500人＝10教室）") == 1
    for sid, old, new in ((27, "250人", "2,500人"), (36, "250人", "2,500人"), (54, "250人", "2,500人"),
                          (45, "約250人", "約2,500人"), (63, "約250人", "約2,500人"),
                          (33, "0.3件", "3件"), (51, "0.2件", "2件"), (60, "0.2件", "2件"), (69, "0.2件", "2件"),
                          (79, "約680人", "約6,800人"), (85, "0.5件", "5件"), (94, "0.1件", "1件")):
        idx = [p.text for p in by[sid].text_frame.paragraphs].index(old)
        set_para(by[sid], idx, old, new)
    assert replace_in_paragraphs(s, "（友だち 累計約1,000人・ブロックを除いて約680人）",
                                 "（友だち 累計約10,000人・ブロックを除いて約6,800人）") == 1
    assert replace_in_paragraphs(s, "約270人", "約2,700人") == 1
    assert replace_in_paragraphs(s, "1教室あたり月に 資料請求 約0.5件・体験予約 約1.0件",
                                 "10教室で月に 資料請求 約5件・体験予約 約10件") == 1
    assert replace_in_paragraphs(s, "※1教室あたりのモデル値（例）。",
                                 f"※10教室のモデル値（例）。新しい友だちは1教室あたり月250人（P{NEW_POS[36]}）×10教室。") == 1
    assert replace_in_paragraphs(s, "件数は小数点第1位まで", "件数は整数（端数は四捨五入）") == 1
    left = [t.text for t in s._element.iter(qn("a:t")) if t.text and re.search(r"\d\.\d+件|小数", t.text)]
    assert not left, left


def edit(prs):
    sl = prs.slides
    at = {key: sl[i - 1] for key, i in NEW_POS.items()}
    for n in range(1, N_BASE + 1):
        if n not in DROP:
            remap_refs(at[n])
    integer_counts(at[PLAN2])
    for k, (key, _, lead, cats, deliver, band) in enumerate(PAGES, 1):
        page_zengo(at[("N", k)], key, lead, cats, deliver, band)
    return sum(fix_fonts(at[k]) for k in [("N", k) for k in range(1, 5)] + [PLAN2])


def main(base, out):
    mid, tmp = out + ".mid.pptx", out + ".tmp.pptx"
    media_gone = assemble(base, mid)
    prs = Presentation(mid)
    fixed = edit(prs)
    prs.save(tmp)
    jb.ORDER = ORDER                      # repack は並びの長さだけ見る
    jb.repack(mid, tmp, out)
    Path(mid).unlink()
    Path(tmp).unlink()
    print("saved:", out, "／", len(ORDER), "枚／外した画像:", media_gone, "／メイリオにそろえた文字:", fixed)


if __name__ == "__main__":
    main(*sys.argv[1:3])
