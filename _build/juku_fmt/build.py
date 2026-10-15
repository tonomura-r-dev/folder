# -*- coding: utf-8 -*-
"""教育（塾）業界の資料ver3を、ジム・フィットネス業界の資料ver3の型（FMT）にそろえる（2026-10-15 殿村さん指示）。

    python -m _build.juku_fmt.build <塾ver3.pptx> <ジムver3.pptx（型の手本）> <出力.pptx>

  殿村さん：「ジム・フィットネスをFMTにして、塾の資料を修正してください」（大至急）。
  ジムの資料の型＝ アジェンダ構成（5章）→ 1 LINE公式アカウントを検討する背景（共通ページ：ハードル・重要な考え方・流入元別・施策展開図×3）＋現状分析
                 → [資料アジェンダ：配信の考え方] 全体設計想定・市場・ニーズ調査・前後検索・年間の配信
                 → [資料アジェンダ：施策の詳細（配信の効率化）] ステップ配信・構築（画面）・具体企画・施策の考え方・改善モデル①〜④・成果
                 → [資料アジェンダ：具体プラン提示] ご提供プラン
                 → [資料アジェンダ：ご提案にあたり（施策紹介・企画案）] 具体施策（18の施策・①〜④は非表示）→ 締め
  塾の中身（データ・文面・設計図）はそのまま。ジムの共通ページ22枚を取り込み、塾のページをジムの章の位置に並べ替え、見出しをジムの言い方に。
  塾のアジェンダ（7章・ページ番号つき）はジムのアジェンダ構成に置き換え。本文中のページ参照（P4〜P7 など）は新しい番号に直す。

  1) 組み立て：塾のzipに、ジムのスライド（XML＋画像）を取り込み、presentation.xml の並びを入れ替える（ノートは取り込まない）
  2) 文字の直し：python-pptx で直し、直したスライドのXMLだけを1)のzipに差し込む（python-pptxで保存し直さない）
"""
import copy
import re
import sys
import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn

from _build.arauchi_v4.build import c14n, normalize_decl, rel_set, slide_order
from _build.arauchi_v8.pages import replace_in_paragraphs
from _build.gym_fmt.build import fix_fonts, set_title, title_shape

# 新しい並び。("J", n)＝塾ver3のn枚目／("G", n)＝ジムver3のn枚目
ORDER = [("J", 1), ("G", 2),
         # 1 LINE公式アカウントを検討する背景
         ("G", 3), ("G", 4), ("G", 5), ("G", 6), ("G", 7), ("G", 8), ("J", 9), ("J", 10),
         # 2 配信の考え方 ── 市場・需要の実態
         ("G", 12), ("G", 13), ("J", 11), ("J", 3), ("J", 4), ("J", 5), ("J", 6), ("J", 7), ("J", 8), ("J", 24),
         # 3 施策の詳細（配信の効率化）── 設計と改善モデル
         ("G", 21), ("J", 12), ("J", 13), ("J", 20), ("J", 21), ("J", 22),
         ("J", 14), ("J", 15), ("J", 16), ("J", 17), ("J", 18), ("J", 19), ("J", 23), ("J", 25),
         ("G", 32), ("G", 33), ("G", 34), ("G", 35), ("G", 36), ("J", 26), ("J", 27), ("J", 28), ("J", 29), ("J", 30),
         # 4 具体プラン提示
         ("G", 38), ("J", 31), ("J", 32), ("J", 33),
         # 5 ご提案にあたり（施策紹介・企画案）
         ("G", 40), ("G", 41), ("G", 42), ("G", 43), ("G", 44), ("G", 45), ("J", 34), ("J", 35),
         ("J", 36)]
J_DROP = [2]                              # 塾のアジェンダ（ジムのアジェンダ構成に置き換え）
N_J = 36
assert sorted([n for k, n in ORDER if k == "J"] + J_DROP) == list(range(1, N_J + 1))
G_IMPORT = [n for k, n in ORDER if k == "G"]
NEW = {key: i for i, key in enumerate(ORDER, 1)}

# 見出し（塾の元の番号 → ジムの言い方）
TITLES_J = {
    3: "LINE公式アカウントを検討する背景｜市場データ",
    4: "LINE公式アカウントを検討する背景｜広告の実態",
    5: "前後検索｜誰が調べているか",
    6: "前後検索｜検索の流れ",
    7: "ニーズ調査（シーズナリティ）",
    9: "現状分析｜各社LINEの今",
    10: "現状分析｜LINEでできること",
    11: "全体設計想定",
    14: "構築①｜友だち追加の入口",
    15: "構築①｜つづき：友だちと一緒に",
    16: "構築②｜あいさつメッセージ",
    17: "構築③｜30秒ぴったり診断",
    18: "構築④｜リッチメニュー",
    19: "構築⑤｜配信の出し分け",
    20: "ステップ配信｜14日間の配信",
    21: "ステップ配信｜場面ごとのメッセージ（実文面）",
    22: "ステップ配信｜体験予約から入会まで",
    23: "具体企画の構成｜過去に問い合わせた保護者へ",
    25: "具体企画の構成｜第2の提案軸：友だちと一緒に",
    31: "ご提供プラン（運用構成）",
}
# 本文中のページ参照（塾の元の番号）がある所
REF_SLIDES = [8, 14, 17, 20, 21, 24, 25, 35]
# ジムの共通ページに残っていたジム固有の言葉 → 塾に
G_WORDS = {
    5: [("体づくりタイプ診断", "30秒ぴったり診断"), ("来店・入会への誘導", "体験予約・入会への誘導")],
    32: [("体づくりタイプ診断", "30秒ぴったり診断"), ("無断キャンセルが課題", "体験の無断欠席が課題"),
         ("⑤ 来店・体験", "⑤ 体験授業"), ("⑦ご入会", "⑦ご入会"),
         ("継続・PT", "継続・講習"), ("物販・友だち紹介", "友だち紹介"), ("退会防止・再入会", "退塾防止・再入塾")],
    34: [("実来店・直面談", "体験授業・面談")],
}

CT_SLIDE = "application/vnd.openxmlformats-officedocument.presentationml.slide+xml"
RT_SLIDE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide"
PR = "http://schemas.openxmlformats.org/package/2006/relationships"
MEDIA_CT = {"png": "image/png", "jpeg": "image/jpeg", "jpg": "image/jpeg", "gif": "image/gif", "svg": "image/svg+xml",
            "emf": "image/x-emf", "wmf": "image/x-wmf", "tif": "image/tiff", "tiff": "image/tiff"}


def rels_of(name):
    return name.replace("slides/", "slides/_rels/") + ".rels"


def layout_name(z, part):
    m = re.search(rb'<p:cSld name="([^"]*)"', z.read(part))
    return m.group(1).decode() if m else None


# ---------------------------------------------------------------- 1) 組み立て
def assemble(juku, gym, out):
    zj, zg = zipfile.ZipFile(juku), zipfile.ZipFile(gym)
    oj, og = slide_order(zj), slide_order(zg)
    assert len(oj) == N_J and len(og) == 46
    names = set(zj.namelist())
    ct = zj.read("[Content_Types].xml").decode("utf-8")
    prels = zj.read("ppt/_rels/presentation.xml.rels").decode("utf-8")
    pres = zj.read("ppt/presentation.xml").decode("utf-8")
    add, gone = {}, set()

    rid_nums = [int(x) for x in re.findall(r'Id="rId(\d+)"', prels)]
    next_rid = max(rid_nums) + 1
    slide_nums = [int(m.group(1)) for n in names for m in [re.match(r"ppt/slides/slide(\d+)\.xml$", n)] if m]
    next_slide = max(slide_nums) + 1
    rid_of = {}
    for key_n, name in enumerate(oj, 1):
        m = re.search(r'<Relationship [^>]*Target="slides/%s"[^>]*/>' % re.escape(name.split("/")[-1]), prels)
        rid_of[("J", key_n)] = re.search(r'Id="(rId\d+)"', m.group(0)).group(1)

    media_map = {}
    for gno in G_IMPORT:
        src = og[gno - 1]
        dst = f"ppt/slides/slide{next_slide}.xml"
        next_slide += 1
        rels = etree.fromstring(zg.read(rels_of(src)))
        for r in list(rels):
            t, target = r.get("Type"), r.get("Target")
            if t.endswith("/slideLayout"):
                lay = "ppt/slideLayouts/" + target.split("/")[-1]
                assert lay in names and layout_name(zj, lay) == layout_name(zg, lay), (lay, gno)
            elif t.endswith("/notesSlide"):
                rels.remove(r)
            elif t.endswith("/image"):
                assert r.get("TargetMode") != "External"
                media = "ppt/media/" + target.split("/")[-1]
                if media not in media_map:
                    new = "ppt/media/gym_" + target.split("/")[-1]
                    assert new not in names
                    media_map[media] = new
                    add[new] = zg.read(media)
                    ext = new.rsplit(".", 1)[-1].lower()
                    if f'Extension="{ext}"' not in ct and f'Extension="{ext.upper()}"' not in ct:
                        ct = ct.replace("<Default ", f'<Default Extension="{ext}" ContentType="{MEDIA_CT[ext]}"/><Default ', 1)
                r.set("Target", "../media/" + media_map[media].split("/")[-1])
            else:
                raise SystemExit(f"想定外のrels {t} (G{gno})")
        xml = zg.read(src)
        # ノートの rId がスライド本体から参照されていないこと
        for r_id in re.findall(rb'r:(?:id|embed|link)="(rId\d+)"', xml):
            assert rels.find(f"{{{PR}}}Relationship[@Id='{r_id.decode()}']") is not None, (gno, r_id)
        add[dst] = xml
        add[rels_of(dst)] = etree.tostring(rels, xml_declaration=True, encoding="UTF-8", standalone=True)
        ct = ct.replace("</Types>", f'<Override PartName="/{dst}" ContentType="{CT_SLIDE}"/></Types>')
        rid = f"rId{next_rid}"
        next_rid += 1
        prels = prels.replace("</Relationships>", f'<Relationship Id="{rid}" Type="{RT_SLIDE}" Target="slides/{dst.split("/")[-1]}"/></Relationships>')
        rid_of[("G", gno)] = rid

    # 塾のアジェンダを外す
    for jno in J_DROP:
        name = oj[jno - 1]
        for r in etree.fromstring(zj.read(rels_of(name))):
            assert r.get("Type").endswith("/slideLayout")
        gone |= {name, rels_of(name)}
        m = re.search(r'<Relationship [^>]*Target="slides/%s"[^>]*/>' % re.escape(name.split("/")[-1]), prels)
        prels = prels.replace(m.group(0), "")
        ct, n = re.subn(r'<Override [^>]*PartName="/%s"[^>]*/>' % re.escape(name), "", ct)
        assert n == 1

    # p:sldIdLst を新しい並びに
    m = re.search(r"<p:sldIdLst>(.*?)</p:sldIdLst>", pres, re.S)
    entries = {re.search(r'r:id="(rId\d+)"', e).group(1): e for e in re.findall(r"<p:sldId [^>]*/>", m.group(1))}
    next_id = max(int(x) for x in re.findall(r'<p:sldId id="(\d+)"', m.group(1))) + 1
    lst = []
    for key in ORDER:
        rid = rid_of[key]
        if rid in entries:
            lst.append(entries[rid])
        else:
            lst.append(f'<p:sldId id="{next_id}" r:id="{rid}"/>')
            next_id += 1
    pres = pres.replace(m.group(0), "<p:sldIdLst>" + "".join(lst) + "</p:sldIdLst>")

    # docProps/app.xml（スライド数・タイトル一覧）
    app = zj.read("docProps/app.xml").decode("utf-8")
    n_new = len(ORDER)
    app = re.sub(r"<Slides>\d+</Slides>", f"<Slides>{n_new}</Slides>", app)
    app = re.sub(r"(<vt:lpstr>スライド タイトル</vt:lpstr></vt:variant><vt:variant><vt:i4>)\d+(</vt:i4>)", rf"\g<1>{n_new}\2", app)
    have = app.count("<vt:lpstr>PowerPoint プレゼンテーション</vt:lpstr>")
    mv = re.search(r'<vt:vector size="(\d+)" baseType="lpstr">', app)
    app = app.replace(mv.group(0), f'<vt:vector size="{int(mv.group(1)) + n_new - have}" baseType="lpstr">')
    app = app.replace("<vt:lpstr>PowerPoint プレゼンテーション</vt:lpstr>",
                      "<vt:lpstr>PowerPoint プレゼンテーション</vt:lpstr>" * 1, 1)
    extra = "<vt:lpstr>PowerPoint プレゼンテーション</vt:lpstr>" * (n_new - have)
    i = app.rfind("<vt:lpstr>PowerPoint プレゼンテーション</vt:lpstr>")
    i = i + len("<vt:lpstr>PowerPoint プレゼンテーション</vt:lpstr>")
    app = app[:i] + extra + app[i:]

    with zipfile.ZipFile(out, "w") as zo:
        for info in zj.infolist():
            name = info.filename
            if name in gone:
                continue
            data = {"ppt/presentation.xml": pres.encode("utf-8"),
                    "ppt/_rels/presentation.xml.rels": prels.encode("utf-8"),
                    "[Content_Types].xml": ct.encode("utf-8"),
                    "docProps/app.xml": app.encode("utf-8")}.get(name)
            zo.writestr(info, data if data is not None else zj.read(name), compress_type=info.compress_type)
        for name, data in add.items():
            zo.writestr(name, data, compress_type=zipfile.ZIP_DEFLATED)


# ---------------------------------------------------------------- 2) 文字の直し
def remap_refs(s):
    """本文中の「P◯」（塾の元の番号）を新しい番号に"""
    for t in s._element.iter(qn("a:t")):
        if t.text and re.search(r"P\d+", t.text):
            t.text = re.sub(r"P(\d+)", lambda m: f"P{NEW[('J', int(m.group(1)))]}", t.text)


def edit(prs):
    sl = prs.slides
    at = {key: sl[i - 1] for key, i in NEW.items()}
    for jno, text in TITLES_J.items():
        set_title(title_shape(at[("J", jno)]), text)
    for jno in REF_SLIDES:
        remap_refs(at[("J", jno)])
    # 「P4〜P7」は新しい番号でも続き番号（P15〜P18）になる
    for gno, pairs in G_WORDS.items():
        s = at[("G", gno)]
        for old, new in pairs:
            if old != new:
                assert replace_in_paragraphs(s, old, new) >= 1, (gno, old)
    # ジムの背景（ハードル）のページ：ジムの実例の注記（FIT PLACE24・カーブス）を外す
    s = at[("G", 3)]
    hit = [sh for sh in s.shapes if sh.has_text_frame and "FIT PLACE24" in sh.text_frame.text]
    assert len(hit) == 1
    hit[0]._element.getparent().remove(hit[0]._element)
    return sum(fix_fonts(s) for s in sl)


def repack(base_path, built_path, out_path):
    zb, zn = zipfile.ZipFile(base_path), zipfile.ZipFile(built_path)
    bo, no = slide_order(zb), slide_order(zn)
    assert len(bo) == len(no) == len(ORDER)
    replaced = {}
    for b, n in zip(bo, no):
        if c14n(zb.read(b)) != c14n(zn.read(n)):
            replaced[b] = normalize_decl(zn.read(n))
        if rel_set(zb.read(rels_of(b))) != rel_set(zn.read(rels_of(n))):
            raise SystemExit(f"想定外のrelsの違い: {b}")
    with zipfile.ZipFile(out_path, "w") as zo:
        for info in zb.infolist():
            data = replaced.pop(info.filename) if info.filename in replaced else zb.read(info.filename)
            zo.writestr(info, data, compress_type=info.compress_type)
    assert not replaced


def main(juku, gym, out):
    mid, tmp = out + ".mid.pptx", out + ".tmp.pptx"
    assemble(juku, gym, mid)
    prs = Presentation(mid)
    fixed = edit(prs)
    prs.save(tmp)
    repack(mid, tmp, out)
    Path(mid).unlink()
    Path(tmp).unlink()
    print("saved:", out, "／", len(ORDER), "枚／メイリオにそろえた文字:", fixed)


if __name__ == "__main__":
    main(*sys.argv[1:4])
