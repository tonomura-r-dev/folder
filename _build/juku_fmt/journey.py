# -*- coding: utf-8 -*-
"""塾ver3に、殿村さんが取ってきた前後検索（2026-10-16）を入れる。

    python -m _build.juku_fmt.journey <塾ver3（50枚・plan.py の出力）> <出力.pptx>

  前後検索（LINEヤフー・画像で受領。_data/journey/教育塾/前後検索画像_2026-10-16/）
    起点＝中学受験 塾／中高一貫 塾／高校受験 塾／英語 塾（医学部 塾は大学受験なので使わない）
    殿村さん「体験授業、個別指導、高校受験 それぞれ単独で出てくる前後検索はなかったから、コレが最大限の譲歩」
  直すページ（ほかは1バイトも変えない）
    8枚目（前後検索｜検索の流れ）＝【データ待ち】3つを、実際の検索で埋める（4色＝志望校・偏差値／検定・勉強法／塾選び・費用・時期／塾名・料金・口コミ）
    24枚目（ステップ配信｜14日間の配信）＝「前後検索で確定させる」「取得中」を外す
    49枚目（出典）＝前後検索の行
  画像なので性別・年代は読めない → 7枚目（誰が調べているか）は前の前後検索（2022-05〜2023-05）のまま。
"""
import sys
import zipfile
from pathlib import Path

from pptx import Presentation

from _build.arauchi_v4.build import c14n, normalize_decl, slide_order
from _build.arauchi_v4.lib import TEXT, add_box
from _build.arauchi_v8.pages import replace_in_paragraphs, shape_by_id
from _build.gym_fmt.build import fix_fonts
from _build.juku_fmt.plan import corner

N = 50
S_FLOW, S_STEP, S_SRC = 8, 24, 49

# 色は8枚目の凡例の色のまま（文字色, 塗り）
CAT = {
    "A": ("志望校・偏差値", "2B5797", "D6E2F3"),
    "B": ("検定・勉強法", "7440A8", "E8DAF5"),
    "C": ("塾選び・費用・時期", "C0561A", "FBE0CF"),
    "D": ("塾名・料金・口コミ", "00796B", "CCE7E4"),
}
# 図から読んだ検索（時系列）。すべて図に出ている語のまま
BOX = {
    38: ("15〜3日前", [
        ("A", "中学受験 偏差値 2025"), ("A", "私立中学 偏差値ランキング"), ("A", "都立高校 入試日程"),
        ("A", "高校 偏差値 東京"), ("B", "英検 日程"), ("B", "漢検／算数検定"),
        ("B", "中学受験 社会 勉強方法"), ("C", "中学受験 高校受験 どっち")]),
    41: ("3日前〜当日", [
        ("C", "中学受験 塾選び"), ("C", "高校受験 塾 いつから"), ("C", "高校受験 塾 費用"),
        ("C", "中学受験 塾 ランキング"), ("C", "中学受験 個別指導のみ"), ("B", "塾なし 高校受験"),
        ("B", "高校受験 勉強方法"), ("D", "森塾 口コミ")]),
    44: ("当日", [
        ("D", "日能研／サピックス"), ("D", "早稲田アカデミー／ena"), ("D", "栄光ゼミナール／臨海セミナー"),
        (None, "その後（1〜15日後）"),
        ("D", "早稲田アカデミー 料金"), ("D", "サピックス 入室テスト"), ("D", "臨海セミナー 評判"),
        ("A", "都立高校 偏差値 2024")]),
}
LEGEND = [45, 46, 47, 48]
NOTE = 49


def flow(s):
    """8枚目：データ待ちの3つを、実際の検索で埋める"""
    assert replace_in_paragraphs(s, "最初に調べること → 迷っているときに調べること → 塾の名前、の順に並べる。",
                                 "最初は志望校・偏差値・検定、迷うと塾選び・費用・いつから。塾の名前はそのあと。") == 1
    for sid, (when, items) in BOX.items():
        old = shape_by_id(s, sid)
        assert "データ待ち" in old.text_frame.text
        x, y, w, h = (v / 914400 for v in (old.left, old.top, old.width, old.height))
        old._element.getparent().remove(old._element)
        paras = [[(when, {"size": 9, "bold": True, "color": "7F7F7F"})]]
        for cat, q in items:
            if cat is None:
                paras.append([(q, {"size": 9, "bold": True, "color": "7F7F7F"})])
            else:
                paras.append([("■ ", {"color": CAT[cat][1]}), (q, {})])
        corner(add_box(s, x, y, w, h, paras, size=10.5, color=TEXT, fill="FFFFFF", line="D9D9D9", anchor="t",
                       shape="round", ins=(0.12, 0.08, 0.08, 0.06), line_spacing=112, after=1), 5000)
    # 凡例：4色の名前を、検索に合わせて付け直す（幅を広げる）
    for k, (sid, cat) in enumerate(zip(LEGEND, "ABCD")):
        sh = shape_by_id(s, sid)
        r = list(sh.text_frame._txBody.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}t"))
        assert len(r) == 1
        r[0].text = CAT[cat][0]
        assert (sh.text_frame._txBody.find("{http://schemas.openxmlformats.org/drawingml/2006/main}bodyPr") is not None)
        bp = sh.text_frame._txBody.find("{http://schemas.openxmlformats.org/drawingml/2006/main}bodyPr")
        bp.set("lIns", "36000")
        bp.set("rIns", "36000")
        sh.left, sh.width = int((0.47 + k * 1.51) * 914400), int(1.45 * 914400)
    note = shape_by_id(s, NOTE)
    assert replace_in_paragraphs(s, "色分けは仮。データが届いたらクエリに合わせて決める",
                                 "※検索は首都圏が中心（塾の名前は首都圏の大手）") == 1
    note.left, note.width = int(6.53 * 914400), int(3.83 * 914400)
    assert replace_in_paragraphs(s, "塾の名前が出てくる前に、保護者はもう調べ始めている。",
                                 "塾の名前を調べる前に、保護者はもう志望校・偏差値・検定を調べている。") == 1
    assert replace_in_paragraphs(s, "その間、塾と保護者の接点はない。（データが届いたら確定）",
                                 "塾の名前を調べた後も、料金・入室テスト・口コミで比べている。その間、塾と保護者の接点はない。") == 1
    src = shape_by_id(s, 28)
    t = src.text_frame.text
    assert t.startswith("出典｜LINEヤフー 前後検索（直近1年）")
    assert replace_in_paragraphs(s, t, "出典｜LINEヤフー前後検索（2026年10月取得）起点KW：中学受験 塾／中高一貫 塾／高校受験 塾／英語 塾") == 1


def step(s):
    """24枚目：前後検索の「確定させる」「取得中」を外す"""
    assert replace_in_paragraphs(s, "（前後検索で確定させる）", "（前後検索・P8）") == 1
    assert replace_in_paragraphs(s, "※保護者が比べている間に何を調べているかは、前後検索（取得中）で確定させ、配信の中身を合わせる",
                                 "※前後検索（P8）：比べている間は、塾の費用・いつから・口コミを調べている。配信の中身をこれに合わせる") == 1


def sources(s):
    assert replace_in_paragraphs(s, "LINEヤフー前後検索 直近1年（取得中）KW：塾／個別指導／体験授業／高校受験",
                                 "LINEヤフー前後検索（2026年10月取得）起点KW：中学受験 塾／中高一貫 塾／高校受験 塾／英語 塾") == 1


def repack(base, built, out):
    zb, zn = zipfile.ZipFile(base), zipfile.ZipFile(built)
    bo, no = slide_order(zb), slide_order(zn)
    assert len(bo) == len(no) == N
    replaced = {b: normalize_decl(zn.read(n)) for b, n in zip(bo, no) if c14n(zb.read(b)) != c14n(zn.read(n))}
    changed = sorted(bo.index(b) + 1 for b in replaced)
    with zipfile.ZipFile(out, "w") as zo:
        for info in zb.infolist():
            data = replaced.pop(info.filename) if info.filename in replaced else zb.read(info.filename)
            zo.writestr(info, data, compress_type=info.compress_type)
    assert not replaced
    return changed


def main(base, out):
    prs = Presentation(base)
    assert len(prs.slides) == N
    flow(prs.slides[S_FLOW - 1])
    step(prs.slides[S_STEP - 1])
    sources(prs.slides[S_SRC - 1])
    fixed = sum(fix_fonts(prs.slides[i - 1]) for i in (S_FLOW, S_STEP, S_SRC))
    tmp = out + ".tmp.pptx"
    prs.save(tmp)
    changed = repack(base, tmp, out)
    Path(tmp).unlink()
    print("saved:", out, "／直したページ:", changed, "／メイリオにそろえた文字:", fixed)


if __name__ == "__main__":
    main(*sys.argv[1:3])
