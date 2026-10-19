# -*- coding: utf-8 -*-
"""塾ver3：先方に出す資料から、社内向けのメモを消す（2026-10-19 殿村さん「もちろん消してください。不要でしょう」）。

    python -m _build.juku_fmt.cleanup <塾ver3（54枚）> <出力.pptx>

  - 消す：上司確認・提出前の確認・e-Statでの照合・更新予定・旧版の話・社内のファイルの場所（_data/…）。
    出典ページの「提出前に確認すること」の枠は消し、左の「LINEヤフー公式データ」の枠を横いっぱいに。
  - 言い方を直す：「差し替える／書く／確定させる／確認する」（作る側への指示）→「〜します」（先方に見せる言い方）。
  ほかは1バイトも変えない（直したスライドのXMLだけ差し込む）。
"""
import sys
from pathlib import Path

from pptx import Presentation

import _build.juku_fmt.build as jb
from _build.arauchi_v8.pages import replace_in_paragraphs, shape_by_id

E = 914400
EDITS = {   # ページ: [(元の文字, 直した文字)]
    5: [("⚠ 市場規模は「定義」で2倍変わる。必ず併記すること", "市場規模は「定義」で約2倍変わる"),
        ("。※提案前にe-Statで統計表を開いて数値照合すること", "")],
    7: [("※最新の前後検索で更新予定", "")],
    11: [("）_data/trends/教育塾/ ※相対値", "）※相対値")],
    13: [("。★競合の実名掲載は上司確認", "")],
    25: [("文面の教室長名・対応時間・配信頻度は貴社の運用に合わせて差し替える",
          "※文面の教室長名・対応時間・配信頻度は、貴社の運用に合わせて調整します")],
    30: [("※事例の数字は貴社の実績に差し替え、学年・期間・通った回数を書く。",
          "※事例の数字は、貴社の実績（学年・期間・通った回数）に合わせて作成します。")],
    31: [("※体験に来ない割合・実施率の改善幅は、貴社の実績で確定させる（旧版の「約20%」「80%→92%」は出典未確認のため載せていない）",
          "※体験に来ない割合・改善の幅は、貴社の実績をもとに確認します")],
    32: [("要件定義で最新の条件を確認する", "要件定義の際に最新の条件を確認します")],
    33: [("週次／_data/trends/教育塾/）", "週次）"),
         ("。★貴社の月別の入会実績があれば、この表を合わせる", "")],
    53: [("出典のない数値は書きません。取得中のデータは、届き次第差し替えます。", "出典のない数値は書いていません。")],
}


def drop_check_box(s):
    """出典ページ：「提出前に確認すること」の枠を消し、左の枠を横いっぱいに"""
    box, left = shape_by_id(s, 32), shape_by_id(s, 31)
    assert box.text_frame.text.startswith("提出前に確認すること") and left.text_frame.text.startswith("LINEヤフー公式データ")
    box._element.getparent().remove(box._element)
    left.width = box.left + box.width - left.left


def main(base, out):
    prs = Presentation(base)
    assert len(prs.slides) == 54
    for n, pairs in EDITS.items():
        for old, new in pairs:
            assert replace_in_paragraphs(prs.slides[n - 1], old, new) == 1, (n, old)
    drop_check_box(prs.slides[52])
    tmp = out + ".tmp.pptx"
    prs.save(tmp)
    jb.ORDER = list(range(len(prs.slides)))       # repack は並びの長さだけ見る
    jb.repack(base, tmp, out)
    Path(tmp).unlink()
    print("saved:", out, "／直したページ:", sorted(EDITS) + [53])


if __name__ == "__main__":
    main(*sys.argv[1:3])
