# -*- coding: utf-8 -*-
"""塾ver3：オレンジ・黄色を使わない（殿村さん 2026-10-19「紺色か黒色か白色で」。ジムver3と同じ直し）。

    python -m _build.juku_fmt.recolor <塾ver3（51枚）> <出力.pptx>

  - 文字・線・塗りのオレンジ → 紺（1F285A）。強調の区別が要るところ（年間カレンダーの山の月・
    導入スケジュールのPhase 3・4・体制の「貴社｜教室」・企画のご参考の◎）は黒（333333）。
  - 薄い黄色・薄いオレンジの塗り → 紺を薄くした色（E4E9F2）。黄土色の枠線 → 紺。
  - 施策の考え方の①〜④の黄色い箱 → 紺＋白い文字（ジムの施策の考え方と同じ）。
  - S11のグラフ画像は、オレンジの系列を黒に近い灰色にして作り直したもの（_build/make_images_juku_v3.py）に差し替え。
  - S3（ハードル）・S44（具体施策18）は汎用ページ。ジムでも触っていないので、そのまま。
  - 下の要旨の帯の「左の縦線」は、塾には無い（確認済み）。
  ほかは1バイトも変えない（直したスライドのXMLと画像2つだけ、元のzipに差し込む）。
"""
import colorsys
import sys
import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn

import _build.juku_fmt.build as jb

ROOT = Path(__file__).resolve().parent.parent.parent
NAVY = "1F285A"
BLACK = "333333"
PALE = "E4E9F2"                 # 紺を薄くした色
PALEST = "F4F7FF"               # ごく薄い塗り（注意の箱）
SKIP = {3, 44}                  # 汎用ページ（ジムでも触っていない）
BLACK_SHAPES = {30: {6, 11, 21, 26},      # 年間の配信カレンダー：山の月（1月・2-3月・6月・7-8月）
                41: {13, 17},             # 導入スケジュール：Phase 3・4
                42: {14, 15},             # 体制：貴社｜教室
                49: {6, 10}}              # 企画のご参考：◎（○＝紺、△＝灰色と区別）
CHARTS = {"ppt/media/image20.png": "_images/juku_v3_trend_koushuu.png",   # S11 季節講習
          "ppt/media/image21.png": "_images/juku_v3_trend_juken.png"}     # S11 受験区分


def warm(v):
    r, g, b = (int(v[k:k + 2], 16) / 255 for k in (0, 2, 4))
    h, l, s = colorsys.rgb_to_hls(r, g, b)
    return s > 0.2 and 15 / 360 <= h <= 65 / 360 and 0.25 < l < 0.995


def text_white(sh):
    for r in sh._element.iter(qn("a:rPr"), qn("a:endParaRPr")):
        for f in r.findall(qn("a:solidFill")):
            r.remove(f)
        f = etree.Element(qn("a:solidFill"))
        etree.SubElement(f, qn("a:srgbClr"), val="FFFFFF")
        r.insert(0, f)


def recolor(prs):
    done = []
    for n, s in enumerate(prs.slides, 1):
        if n in SKIP:
            continue
        hit = False
        for sh in s.shapes:
            black = sh.shape_id in BLACK_SHAPES.get(n, ())
            white = False
            for c in sh._element.iter(qn("a:srgbClr")):
                v = c.get("val").upper()
                if v in ("BF4E0E", "ED7D31", "C55A11", "E0C36A"):
                    c.set("val", BLACK if black else NAVY)
                elif v == "FFC000":
                    c.set("val", NAVY)
                    white = True
                elif v in ("FFF2CC", "FCE4D6"):
                    c.set("val", PALE)
                elif v == "FFF7F0":
                    c.set("val", PALEST)
                else:
                    continue
                hit = True
            if white:
                text_white(sh)
        left = [(n, c.get("val")) for sh in s.shapes for c in sh._element.iter(qn("a:srgbClr")) if warm(c.get("val"))]
        assert not left, left
        if hit:
            done.append(n)
    for n, ids in BLACK_SHAPES.items():
        assert {sh.shape_id for sh in prs.slides[n - 1].shapes} >= ids, n
    return done


def main(base, out):
    prs = Presentation(base)
    assert len(prs.slides) == 51
    done = recolor(prs)
    tmp = out + ".tmp.pptx"
    prs.save(tmp)
    jb.ORDER = list(range(len(prs.slides)))       # repack は並びの長さだけ見る
    jb.repack(base, tmp, out)
    Path(tmp).unlink()
    z = zipfile.ZipFile(out)                     # S11のグラフ画像を差し替え
    rels = z.read("ppt/slides/_rels/slide9.xml.rels").decode()
    assert all(Path(k).name in rels for k in CHARTS)
    with zipfile.ZipFile(out + ".tmp", "w") as zo:
        for info in z.infolist():
            data = (ROOT / CHARTS[info.filename]).read_bytes() if info.filename in CHARTS else z.read(info.filename)
            zo.writestr(info, data, compress_type=info.compress_type)
    z.close()
    Path(out + ".tmp").replace(out)
    print("saved:", out, "／色を置き換えたページ:", done)


if __name__ == "__main__":
    main(*sys.argv[1:3])
