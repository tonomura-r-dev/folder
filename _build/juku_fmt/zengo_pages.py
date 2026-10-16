# -*- coding: utf-8 -*-
"""塾ver3：前後検索の画像（4色で色分け）を2枚足し、「ご予算の確認」を外す（2026-10-16 殿村さん）。

    python -m _build.juku_fmt.zengo_pages <塾ver3（50枚・journey.py の出力）> <出力.pptx>

  殿村さん：「黄・オレンジ色は使わないで欲しい」「せっかく画像おくったんやから、使って」「S37はいらない」
  - 8枚目（前後検索｜検索の流れ）のあとに2枚：中学受験・中高一貫／高校受験・英語（画像は _build/make_juku_zengo.py の出力）
  - 37枚目「ご予算の確認」を外す
  - 本文中のページ参照（P◯）を新しい番号に付け替える（出典の前後検索の行には、足した2枚も入れる）
  色：志望校・偏差値＝青／検定・勉強法＝紫／塾選び・費用・時期＝ローズ／塾名・料金・口コミ＝緑（黄・オレンジは使わない）
"""
import random
import re
import sys
import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.oxml.shapes.picture import CT_Picture

from _build.arauchi_v4.build import slide_order
from _build.arauchi_v4.lib import NAVY, TEXT, add_box
from _build.arauchi_v8.pages import replace_in_paragraphs
from _build.gym_fmt.build import fix_fonts, set_title
from _build.juku_fmt.build import CT_SLIDE, RT_SLIDE, rels_of, repack
from _build.juku_fmt.journey import CAT
from _build.juku_fmt.plan import clear_body, corner, set_header, set_lead

ROOT = Path(__file__).resolve().parent.parent.parent
N_BASE = 50
FLOW, TEMPLATE, BUDGET = 8, 13, 37
RT_IMAGE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"
PAGES = [   # 足す2枚（8枚目のあと）
    dict(title="前後検索｜中学受験・中高一貫の保護者",
         lead=["塾の名前を調べる前に、志望校・偏差値（青）を調べている。",
               "塾の名前（緑）は当日に集中し、そのあとは入室テスト・料金を調べている。"],
         charts=[("中学受験", "起点：中学受験 塾",
                  ["15〜3日前：私立中学の偏差値・学校名・算数検定", "当日〜：日能研・サピックスなど → 入室テスト・料金"]),
                 ("中高一貫", "起点：中高一貫 塾",
                  ["前後とも：都立中高一貫校・中等教育学校の名前", "当日：ena・栄光ゼミナールなど都立一貫に強い塾"])],
         bar="塾の名前を調べる前の、志望校・偏差値を調べている時期に、LINEでつながっておく。"),
    dict(title="前後検索｜高校受験・英語の保護者",
         lead=["高校受験は、当日に「塾 いつから・費用・ランキング」（ローズ）を調べている。",
               "英語は、先に英検（紫）を調べ、そのあと塾の名前（緑）にたどり着く。"],
         charts=[("高校受験", "起点：高校受験 塾",
                  ["15〜3日前：高校名・偏差値", "当日：塾 いつから・費用・ランキング → 塾の名前"]),
                 ("英語", "起点：英語 塾",
                  ["15〜3日前：英検 日程・英検3級／2級", "当日：英語塾・個別指導塾 → そのあと英検 合格発表"])],
         bar="「いつから・費用」や英検を調べている時期に、役に立つ情報をLINEで届ける。"),
]
N_NEW = len(PAGES)
ORDER = list(range(1, FLOW + 1)) + [("N", k) for k in range(1, N_NEW + 1)] + \
        [n for n in range(FLOW + 1, N_BASE + 1) if n != BUDGET]
NEW_POS = {key: i for i, key in enumerate(ORDER, 1)}
IMG_W, IMG_ASPECT = 4.83, 2000 / 1167


def title_of(z, name):
    m = re.search(r"<a:t>([^<]*)</a:t>", z.read(name).decode("utf-8"))
    return m.group(1) if m else ""


# ---------------------------------------------------------------- 1) 組み立て
def assemble(base, out):
    z = zipfile.ZipFile(base)
    order = slide_order(z)
    assert len(order) == N_BASE
    assert title_of(z, order[FLOW - 1]) == "前後検索｜検索の流れ"
    assert title_of(z, order[TEMPLATE - 1]) == "全体設計想定"
    assert title_of(z, order[BUDGET - 1]) == "ご予算の確認"
    names = set(z.namelist())
    ct = z.read("[Content_Types].xml").decode("utf-8")
    prels = z.read("ppt/_rels/presentation.xml.rels").decode("utf-8")
    pres = z.read("ppt/presentation.xml").decode("utf-8")

    def rel_of(name):
        m = re.search(r'<Relationship [^>]*Target="slides/%s"[^>]*/>' % re.escape(name.split("/")[-1]), prels)
        return m, re.search(r'Id="(rId\d+)"', m.group(0)).group(1)

    rid_of = {n: rel_of(order[n - 1])[1] for n in range(1, N_BASE + 1)}
    next_rid = max(int(x) for x in re.findall(r'Id="rId(\d+)"', prels)) + 1
    next_slide = max(int(m.group(1)) for n in names for m in [re.match(r"ppt/slides/slide(\d+)\.xml$", n)] if m) + 1
    tpl = z.read(order[TEMPLATE - 1]).decode("utf-8")
    tpl_rels = etree.fromstring(z.read(rels_of(order[TEMPLATE - 1])))
    for r in list(tpl_rels):
        if r.get("Type").endswith("/notesSlide"):
            tpl_rels.remove(r)
    assert [r.get("Type").split("/")[-1] for r in tpl_rels] == ["slideLayout"]
    used = set(re.findall(r'creationId[^>]*val="(\d+)"', "".join(z.read(o).decode("utf-8") for o in order)))
    rnd = random.Random(202610162)
    add = {}
    for k, page in enumerate(PAGES, 1):
        dst = f"ppt/slides/slide{next_slide}.xml"
        next_slide += 1
        cid = str(rnd.randrange(10 ** 9, 4 * 10 ** 9))
        assert cid not in used
        add[dst] = re.sub(r'(<p14:creationId [^>]*val=")\d+(")', rf"\g<1>{cid}\2", tpl, count=1).encode("utf-8")
        rels = etree.fromstring(etree.tostring(tpl_rels))
        for j, (key, _, _) in enumerate(page["charts"], 2):
            media = f"ppt/media/juku_zengo_{k}{j - 1}.png"          # パーツ名は英数字だけ（日本語の名前は開けない）
            assert media not in names
            add[media] = (ROOT / f"_images/juku_zengo_{key}.png").read_bytes()
            etree.SubElement(rels, "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship",
                             Id=f"rId{j}", Type=RT_IMAGE, Target=f"../media/{media.split('/')[-1]}")
        add[rels_of(dst)] = etree.tostring(rels, xml_declaration=True, encoding="UTF-8", standalone=True)
        ct = ct.replace("</Types>", f'<Override PartName="/{dst}" ContentType="{CT_SLIDE}"/></Types>')
        rid = f"rId{next_rid}"
        next_rid += 1
        prels = prels.replace("</Relationships>", f'<Relationship Id="{rid}" Type="{RT_SLIDE}" Target="slides/{dst.split("/")[-1]}"/></Relationships>')
        rid_of[("N", k)] = rid

    # ご予算の確認を外す
    old = order[BUDGET - 1]
    for r in etree.fromstring(z.read(rels_of(old))):
        assert r.get("Type").endswith("/slideLayout")
    for n in names:
        if n.endswith(".rels") and n not in (rels_of(old), "ppt/_rels/presentation.xml.rels"):
            assert old.split("/")[-1] not in z.read(n).decode("utf-8"), n
    gone = {old, rels_of(old)}
    m, _ = rel_of(old)
    prels = prels.replace(m.group(0), "")
    ct, k = re.subn(r'<Override [^>]*PartName="/%s"[^>]*/>' % re.escape(old), "", ct)
    assert k == 1

    m = re.search(r"<p:sldIdLst>(.*?)</p:sldIdLst>", pres, re.S)
    entries = {re.search(r'r:id="(rId\d+)"', e).group(1): e for e in re.findall(r"<p:sldId [^>]*/>", m.group(1))}
    next_id = max(int(x) for x in re.findall(r'<p:sldId id="(\d+)"', m.group(1))) + 1
    lst = []
    for key in ORDER:
        rid = rid_of[key]
        if rid in entries:
            lst.append(entries[rid])
        else:
            lst.append(f'<p:sldId id="{next_id}" r:id="{rid}"/>')
            next_id += 1
    pres = pres.replace(m.group(0), "<p:sldIdLst>" + "".join(lst) + "</p:sldIdLst>")

    app = z.read("docProps/app.xml").decode("utf-8")
    n_new, ppt_t = len(ORDER), "<vt:lpstr>PowerPoint プレゼンテーション</vt:lpstr>"
    assert f"<Slides>{N_BASE}</Slides>" in app and app.count(ppt_t) == N_BASE
    app = app.replace(f"<Slides>{N_BASE}</Slides>", f"<Slides>{n_new}</Slides>")
    app = re.sub(r"(<vt:lpstr>スライド タイトル</vt:lpstr></vt:variant><vt:variant><vt:i4>)\d+(</vt:i4>)", rf"\g<1>{n_new}\2", app)
    mv = re.search(r'<vt:vector size="(\d+)" baseType="lpstr">', app)
    app = app.replace(mv.group(0), f'<vt:vector size="{int(mv.group(1)) + n_new - N_BASE}" baseType="lpstr">')
    i = app.rfind(ppt_t) + len(ppt_t)
    app = app[:i] + ppt_t * (n_new - N_BASE) + app[i:]

    with zipfile.ZipFile(out, "w") as zo:
        for info in z.infolist():
            if info.filename in gone:
                continue
            data = {"ppt/presentation.xml": pres.encode("utf-8"),
                    "ppt/_rels/presentation.xml.rels": prels.encode("utf-8"),
                    "[Content_Types].xml": ct.encode("utf-8"),
                    "docProps/app.xml": app.encode("utf-8")}.get(info.filename)
            zo.writestr(info, data if data is not None else z.read(info.filename), compress_type=info.compress_type)
        for name, data in add.items():
            zo.writestr(name, data, compress_type=zipfile.ZIP_STORED if name.endswith(".png") else zipfile.ZIP_DEFLATED)


# ---------------------------------------------------------------- 2) 中身
def add_pic(s, rid, name, x, y, w, h):
    sp_tree = s.shapes._spTree
    nid = max(sh.shape_id for sh in s.shapes) + 1
    pic = CT_Picture.new_pic(nid, name, "", rid, int(x * 914400), int(y * 914400), int(w * 914400), int(h * 914400))
    sp_tree.append(pic)
    ln = etree.SubElement(pic.spPr, qn("a:ln"), w="9525")
    f = etree.SubElement(ln, qn("a:solidFill"))
    etree.SubElement(f, qn("a:srgbClr"), val="D9D9D9")


def zengo_page(s, page):
    clear_body(s)
    set_title([x for x in s.shapes if x.shape_id == 2][0], page["title"])
    set_lead(s, page["lead"])
    set_header(s)
    # 凡例（4色）
    for k, cat in enumerate("ABCD"):
        name, txt, fill = CAT[cat]
        corner(add_box(s, 0.47 + k * 1.51, 1.62, 1.45, 0.24, [[(name, {})]], size=9, bold=True, color=txt, fill=fill,
                       align="ctr", shape="round", ins=(0.04, 0, 0.04, 0)), 30000)
    add_box(s, 6.53, 1.62, 3.83, 0.24, [[("※背景の色はDYMの分類。薄い文字は分類の対象外", {})]], size=8.5, color="7F7F7F",
            ins=(0, 0, 0, 0))
    h = IMG_W / IMG_ASPECT
    for j, (key, cap, notes) in enumerate(page["charts"]):
        x = 0.47 + j * (IMG_W + 0.23)
        corner(add_box(s, x, 1.98, IMG_W, 0.3, [[(cap, {})]], size=11, bold=True, color="FFFFFF", fill=NAVY,
                       align="ctr", shape="round", ins=(0.04, 0, 0.04, 0)), 20000)
        add_pic(s, f"rId{j + 2}", f"前後検索_{key}", x, 2.34, IMG_W, h)
        add_box(s, x, 2.34 + h + 0.08, IMG_W, 0.62, [[("・" + t, {})] for t in notes], size=10.5, color=TEXT,
                anchor="t", ins=(0.04, 0.02, 0.04, 0.02), line_spacing=110)
    corner(add_box(s, 0.47, 6.0, 9.89, 0.5, [[(page["bar"], {})]], size=14, bold=True, color="FFFFFF", fill=NAVY,
                   align="ctr", shape="round"), 12000)
    add_box(s, 0.47, 6.83, 9.89, 0.3,
            [[("出典｜LINEヤフー前後検索（2026年10月取得）。横軸＝起点の検索の前後15日、縦軸＝検索した人数（UU）。色分けはDYMの分類", {})]],
            size=7.5, color="7F7F7F", anchor="t", ins=(0, 0, 0, 0))


def remap_refs(s):
    n = 0
    for t in s._element.iter(qn("a:t")):
        if t.text and re.search(r"P\d+", t.text):
            new = re.sub(r"P(\d+)", lambda m: f"P{NEW_POS[int(m.group(1))]}", t.text)
            n += new != t.text
            t.text = new
    return n


def edit(prs):
    sl = prs.slides
    at = {key: sl[i - 1] for key, i in NEW_POS.items()}
    for k, page in enumerate(PAGES, 1):
        zengo_page(at[("N", k)], page)
    # ページ参照（50枚の番号）を新しい番号に。ご予算の確認（37）を指す所は無いこと
    changed = []
    for key, s in at.items():
        if isinstance(key, int):
            for t in s._element.iter(qn("a:t")):
                assert not (t.text and re.search(rf"P{BUDGET}(?!\d)", t.text)), key
            if remap_refs(s):
                changed.append(NEW_POS[key])
    # 出典：前後検索の行に、足した2枚も入れる
    src = at[49]
    assert replace_in_paragraphs(src, f"P{NEW_POS[8]}・P{NEW_POS[10]}・P{NEW_POS[24]}",
                                 f"P{NEW_POS[8]}〜P{NEW_POS[('N', N_NEW)]}・P{NEW_POS[10]}・P{NEW_POS[24]}") == 1
    fixed = sum(fix_fonts(at[("N", k)]) for k in range(1, N_NEW + 1))
    return changed, fixed


def main(base, out):
    mid, tmp = out + ".mid.pptx", out + ".tmp.pptx"
    assemble(base, mid)
    prs = Presentation(mid)
    changed, fixed = edit(prs)
    prs.save(tmp)
    import _build.juku_fmt.build as jb
    jb.ORDER = ORDER                      # repack は並びの長さだけ見る
    repack(mid, tmp, out)
    Path(mid).unlink()
    Path(tmp).unlink()
    print("saved:", out, "／", len(ORDER), "枚／参照を直したページ:", changed, "／メイリオにそろえた文字:", fixed)


if __name__ == "__main__":
    main(*sys.argv[1:3])
