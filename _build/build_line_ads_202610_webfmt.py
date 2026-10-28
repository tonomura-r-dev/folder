"""
LINE広告トレンドレポート10月号（殿村さん版ver1＝正）を、WEB広告トレンドレポートと同じFMTに組み直す（2026-10-28 殿村さん指示）。

WEB広告のFMT（Criteo・X・LYの10月号）に合わせた点
- 表紙・目次なし。左上は「１０月アップデート情報」。分類タグは 検索／バナー／動画／ショッピング
- タイトルは【 】なしで見出しだけ。左にロゴ（LINE広告・VOOM＝LINE、LINEヤフー広告＝LY）
- サマリー2行 → 本文の上に1行の見出し → 本文 → 一番下に「Point」2行。フッターは出典URLだけ
- 右上の逆三角形は3段（認知／検討／獲得）。強調する段＝薄い青、ほか＝白（Criteo版と同じ）
対象業界は各ページのまま（LINE広告を運用中の広告主）。本文の枠・画像は ver1 から移す。

使い方: python3 _build/build_line_ads_202610_webfmt.py <WEBのFMT.pptx> <LINE広告ver1.pptx> <出力.pptx>
  WEBのFMT＝ _templates/DYM_WEB_TREND_FMT.pptx（LYの10月号の実ファイル）
"""
import io
import sys
import copy
import shutil
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.oxml.ns import qn
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_THEME_COLOR
import trend_report_engine as engine

WEB, LINE, OUT = sys.argv[1:4]
shutil.copyfile(WEB, OUT)
prs = Presentation(OUT)
src = Presentation(LINE)
src_pages = list(src.slides)[2:7]
TAGS = ("検索", "バナー", "動画", "ショッピング")
INDUSTRY = "LINE広告を運用中の広告主"


def text(sh):
    return sh.text_frame.text.strip() if sh.has_text_frame else ""


PAGES = [
    dict(headline="旧「LINE広告」の配信停止（10/26）", tags={"検索"}, logo="line", layout="img_cards",
         summary=["2026年10月26日（月）に、旧「LINE広告」の新規作成・編集・広告配信が停止された。",
                  "停止後の管理画面は閲覧のみ。2027年3月末頃に提供終了となる。"],
         sub="旧「LINE広告」は10/26に配信停止、2027年3月末頃に提供終了",
         point=["10/26以降、旧「LINE広告」では配信も編集もできない。",
                "比較に使うレポートは、提供終了（2027年3月末頃）までに保存しておく。"]),
    dict(headline="ディスプレイ広告 統合の概要", tags={"検索"}, logo="ly", layout="img_cards",
         summary=["2026年4月、LINE広告とYahoo!広告 ディスプレイ広告が統合され、「LINEヤフー広告 ディスプレイ広告」が始まった。",
                  "Yahoo!広告は自動で移行済み。LINE広告は自動では移らないため、移行が必要。"],
         sub="LINE広告とYahoo!広告 ディスプレイ広告が「LINEヤフー広告」に統合",
         point=["Yahoo!広告は自動で移行済み。LINE広告は自動では移らないため、",
                "移行ツールか手動でLINEヤフー広告へ移す必要がある。"]),
    dict(headline="移行の対象と方法", tags={"検索"}, logo="ly", layout="grid4",
         summary=["LINE広告を使っている広告主・代理店は、LINEヤフー広告への移行が必要。",
                  "移行の方法は、移行ツールを使う方法と手動の2つ。"],
         sub="LINE広告を使っている場合は、移行ツールか手動で移行する",
         point=["配信実績とオーディエンスの蓄積で最適化されるため、早めの移行が推奨されている。",
                "移行ツールを使うには、認証済みビジネスマネージャーとの連携が必須。"]),
    dict(headline="移行ツールで移せるもの・移せないもの", tags={"検索"}, logo="ly", layout="grid4",
         summary=["移行ツールは、LINE広告の配信データをできる限りLINEヤフー広告へコピーする機能。",
                  "タグ・コンバージョン設定や過去の実績など、移せないデータもある。"],
         sub="移せないもの（タグ・コンバージョン設定・過去の実績）は移行後に設定し直す",
         point=["移行後のキャンペーンは配信オフになる。",
                "タグ・コンバージョンを設定し直し、設定を確認してから配信を再開する。"]),
    dict(headline="LINE VOOM サービス終了に伴う広告配信の終了", tags={"動画"}, logo="line", layout="three_one",
         summary=["2026年9月30日（水）にLINE VOOMのサービスが終了し、VOOM面への広告配信も終了した。",
                  "終了前の配信実績は、終了後も配信面ごとに確認できる。"],
         sub="VOOM面への配信の設定によって、影響は3パターン",
         point=["VOOM面だけを指定していた広告は、配信されなくなる。",
                "配信先を見直し、不要になったVOOM関連の設定は削除する。"]),
]
FUNNEL = ["認知", "検討"]

# ---------- WEBのFMTの下準備 ----------
web = list(prs.slides)
tag_on = copy.deepcopy(next(x for x in web[1].shapes if text(x) == "ショッピング")._element.spPr)
tag_off = copy.deepcopy(next(x for x in web[1].shapes if text(x) == "動画")._element.spPr)
sld_ids = list(prs.slides._sldIdLst)
prs.part.drop_rel(sld_ids[1].rId)
prs.slides._sldIdLst.remove(sld_ids[1])

base = list(prs.slides)[0]
frame = next(x for x in base.shapes if x._element.tag == qn("p:graphicFrame") and not x.has_table)
FUNNEL_BOX = (frame.left, frame.top, frame.width, frame.height)
rel_ids = list(frame._element.find(".//{http://schemas.openxmlformats.org/drawingml/2006/diagram}relIds").attrib.values())
for rel in list(base.part.rels.values()):
    if rel.rId in rel_ids or rel.reltype.endswith("/diagramDrawing"):
        base.part.drop_rel(rel.rId)
frame._element.getparent().remove(frame._element)

for _ in range(len(PAGES) - 1):
    engine.duplicate_slide(prs, 0)
slides = list(prs.slides)

LINE_ICON = next(x for x in src_pages[0].shapes if x.shape_type == 13 and x.top < Inches(2.1)).image.blob


def funnel(slide, on):
    """WEB版と同じ3段の逆三角形（台形を180度回転＋文字は別のテキストボックス）。"""
    x0, y0, w, h = FUNNEL_BOX
    seg_h = h // 3
    for i, label in enumerate(["認知", "検討", "獲得"]):
        sw = int(w * (3 - i) / 3)
        sx = x0 + (w - sw) // 2
        sy = y0 + seg_h * i
        tz = slide.shapes.add_shape(MSO_SHAPE.TRAPEZOID, sx, sy, sw, seg_h)
        tz.rotation = 180
        tz.adjustments[0] = 0.60792
        tz.fill.solid()
        if label in on:
            tz.fill.fore_color.theme_color = MSO_THEME_COLOR.TEXT_2
            tz.fill.fore_color.brightness = 0.6
        else:
            tz.fill.fore_color.theme_color = MSO_THEME_COLOR.BACKGROUND_1
        tz.line.color.theme_color = MSO_THEME_COLOR.TEXT_2
        tz.line.width = Pt(2)
        tz.shadow.inherit = False
        tb = slide.shapes.add_textbox(sx, sy, sw, seg_h)
        tf = tb.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = label
        r.font.size = Pt(8)
        r.font.name = "Meiryo UI"
        r.font.color.theme_color = MSO_THEME_COLOR.TEXT_1


def set_tag(sh, on):
    el = sh._element
    new = copy.deepcopy(tag_on if on else tag_off)
    # WEBのFMTは「バナー」だけ枠線の太さが未指定（スタイル頼み）で線が見えないので、太さを明示する
    for ln in new.findall(qn("a:ln")):
        ln.set("w", "9525")
    old_x = el.spPr.find(qn("a:xfrm"))
    new_x = new.find(qn("a:xfrm"))
    new.replace(new_x, copy.deepcopy(old_x))
    el.replace(el.spPr, new)


def new_id(slide):
    return str(max(int(e.get("id")) for e in slide._element.iter(qn("p:cNvPr"))) + 1)


def copy_card(slide, card, x, y, w, h, head_pt=14, body_pt=11):
    el = copy.deepcopy(card._element)
    slide.shapes._spTree.append(el)
    el.find(".//" + qn("p:cNvPr")).set("id", new_id(slide))
    c = next(s for s in slide.shapes if s._element is el)
    c.left, c.top, c.width, c.height = Inches(x), Inches(y), Inches(w), Inches(h)
    for i, p in enumerate(c.text_frame.paragraphs):
        for r in p.runs:
            r.font.size = Pt(head_pt if i == 0 else body_pt)
        if i > 0:
            p.space_before = Pt(3)
    return c


L, W, TOP, BOTTOM, GAP = 0.32, 10.32, 3.16, 5.8, 0.15

for slide, pg, sp in zip(slides, PAGES, src_pages):
    for sh in list(slide.shapes):
        t = text(sh)
        if t == "-":
            engine.set_single_run_text(sh, INDUSTRY)
            sh.text_frame.paragraphs[0].runs[0].font.size = Pt(14)  # WEBの「-」は大きい字なので、文言が入る大きさに
        elif t in TAGS:
            set_tag(sh, t in pg["tags"])
        elif t.startswith("ビジネスIDをシステム"):
            p = sh.text_frame.paragraphs[0]
            p.runs[0].text = "\t" + pg["headline"]
            for r in list(p.runs[1:]):
                r._r.getparent().remove(r._r)
        elif t.startswith("未統合の"):
            engine.set_multiline_text(sh, pg["summary"])
        elif t.startswith("Yahoo! JAPANとLINE"):
            engine.set_single_run_text(sh, pg["sub"])
        elif t.startswith("ビジネスIDへのシステム統合"):
            engine.set_multiline_text(sh, pg["point"])
        elif t.startswith("https://"):
            url = next(text(x) for x in sp.shapes if text(x).startswith("出典：")).replace("出典：", "")
            engine.set_single_run_text(sh, url)
            sh.text_frame.paragraphs[0].runs[0].font.size = Pt(9)  # 長いURLが2行に折れないように
            sh.left, sh.width = Inches(3.3), Inches(6.6)
        elif Inches(3.05) <= sh.top < Inches(5.85):
            sh._element.getparent().remove(sh._element)
    if pg["logo"] == "line":
        ly = next(x for x in slide.shapes if x.shape_type == 13 and x.top < Inches(2.1))
        slide.shapes.add_picture(io.BytesIO(LINE_ICON), Inches(0.42), Inches(1.44), Inches(0.56), Inches(0.56))
        ly._element.getparent().remove(ly._element)
    funnel(slide, FUNNEL)

    cards = sorted([x for x in sp.shapes if x.shape_type == 1 and x.has_text_frame and x.top > Inches(3.2)
                    and text(x) != "■ 詳細"], key=lambda c: (round(c.top / 914400, 1), c.left))
    room = BOTTOM - TOP
    if pg["layout"] == "img_cards":
        pic = next(x for x in sp.shapes if x.shape_type == 13 and x.top > Inches(2.5))
        iw = 6.0
        ih = iw * pic.height / pic.width
        slide.shapes.add_picture(io.BytesIO(pic.image.blob), Inches(L), Inches(TOP + (room - ih) / 2), Inches(iw), Inches(ih))
        cx, cw, ch = L + iw + 0.2, W - iw - 0.2, (room - GAP) / 2
        for i, c in enumerate(cards[:2]):
            copy_card(slide, c, cx, TOP + i * (ch + GAP), cw, ch)
    elif pg["layout"] == "grid4":
        cw, ch = (W - 0.2) / 2, (room - GAP) / 2
        for i, c in enumerate(cards[:4]):
            copy_card(slide, c, L + (i % 2) * (cw + 0.2), TOP + (i // 2) * (ch + GAP), cw, ch,
                      body_pt=10.5 if pg["headline"].startswith("移行ツール") else 11)
    else:  # three_one
        tw, th = (W - 0.4) / 3, 1.12
        for i, c in enumerate(cards[:3]):
            copy_card(slide, c, L + i * (tw + 0.2), TOP, tw, th)
        copy_card(slide, cards[3], L, TOP + th + GAP, W, room - th - GAP)

prs.save(OUT)
print("saved", OUT, len(prs.slides))
