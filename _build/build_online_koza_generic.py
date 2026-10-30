# -*- coding: utf-8 -*-
"""オンライン講座・コンサル業界の業界汎用資料を、すなおコンサルティングの資料（ver7.3・63枚）から作る（2026-10-30 殿村さん「基本的にすなおコンサルが使える」）。

    python -m _build.build_online_koza_generic <すなおver7.3.pptx> <出力.pptx>

  すなおにしか無いもの（社名・講師名・サロン名・LPの文言・SIMの数字・他社の社名）を外し、業界向けの言葉にする。
  ・外すページ：P3 現状把握（すなおのLP・数値）／P21 概算お見積り（SIM）／P22 費用の比較（すなおの金額）／P25 成果報酬（すなおの単価）
  ・P23・P24 の料金プラン：DYMの目安（LINE運用コンサル 初期10万円〜・月額20万円〜／月3回、30万円〜／月5回）＋ポップアップ（初期1.5万円・月額3万円）
  ・P9 他社：社名を「A社・B社」に、画面の社名・URL・ロゴは伏せた画像（`_images/online_koza/tasha_A_aisatsu_masked.png`）に差し替え
  ・前後検索「投資セミナー」・診断・ステップ配信・参加後のフォローは、業界の言葉に置き換え
  補足資料（裏表紙の後ろ）は触らない。
"""
import copy
import re
import sys

from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches

from _build.pptx_parts import drop_slide, rich

A_T = qn("a:t")
GRAY = "595959"
IMG_A = "_images/online_koza/tasha_A_aisatsu_masked.png"

# 本編（P1〜P43）全体に当てる置き換え（ラン単位。当たらなくてもよい）
COMMON = [
    ("次回イベント・サロン入会", "次回のセミナー・講座の入会"),
    ("イベント参加・サロン入会", "セミナー参加・入会"),
    ("オンラインサロンへの入会", "講座・サロンへの入会"),
    ("サロン入会のお申込み方法", "入会のお申込み方法"),
    ("サロン入会を案内", "入会を案内"),
    ("サロン入会の段階", "講座・サロンの入会の段階"),
    ("サロン入会", "入会"),
    ("サロンの入会", "講座・サロンの入会"),
    ("サロンの 入会", "講座・サロンの 入会"),
    ("サロンの内容", "講座・サロンの内容"),
    ("サロンをご案内", "講座・サロンをご案内"),
    ("サロンを案内", "講座・サロンを案内"),
    ("サロンのご案内", "講座・サロンのご案内"),
    ("サロンの申込み", "講座・サロンの申込み"),
    ("サロンに入会", "講座・サロンに入会"),
    ("・サロン・サロン", "・サロン"),
    ("講座・講座・", "講座・"),
]

# ページごとの置き換え（ラン単位。必ず当たる）
PAGE = {
    1: [("株式会社すなおコンサルティング 御中", "オンライン講座・コンサル業界向け"),
        ("次のセミナー参加・入会へ", "次のセミナー参加・講座の入会へ")],
    2: [("現状把握", "業界の構造")],
    4: [("保険営業の方向けのセミナー・サロンが主軸（一般の方向けのマネーセミナーもあり）",
         "お仕事で成果を上げたい方・お金や暮らしを学びたい方"),
        ("【イベントは年4回】", "【イベント】　　　　"),
        ("イベントの間の時期にも", "セミナー・説明会の間の時期にも")],
    6: [("研修・学びの市場：営業職の学びのニーズが伸びている", "研修・学びの市場：学びのニーズが伸びている"),
        ("次のセミナー・講座・サロンの入会まで", "次のセミナー・講座の入会まで")],
    7: [("保険営業の方にとって、相談につながるチャンスが増えている。一般の方向けのマネーセミナーも、LINEで次回の案内を続けられる。",
         "お金・資産形成のセミナーや講座は、学びたい方の受け皿になれる。LINEで次回の案内を続けられる。")],
    8: [("ターゲット：保険営業の方が主軸。一般の方向けも考慮", "ターゲット：お仕事で学ぶ方と、ご自身のために学ぶ方"),
        ("保険営業の方（主軸）", "お仕事で成果を上げたい方"),
        ("営業力の向上（面談・提案の進め方）", "仕事の成果（集客・営業・時間の使い方）"),
        ("資産形成・お金の知識", "資産形成・お金の知識・副業"),
        ("生産性4倍セミナー → オンラインサロン「ハッピーライフ」", "無料セミナー → 講座・オンラインサロン"),
        ("保険営業の方向けのLPから友だち追加", "お仕事向けのLPから友だち追加"),
        ("保険営業の方向けサロンへの入会", "お仕事向けの講座への入会")],
    9: [("保険営業の方向けの他社も、LINEで特典や学びを届けている", "同じ業界の他社も、LINEで特典や学びを届けている"),
        ("野沢塾のあいさつ（LINE公式アカウント）", "A社のあいさつ（LINE公式アカウント）"),
        ("野沢塾", "A社"), ("R&C保険営業ラボ", "B社"),
        ("保険営業の方向けの講座", "お仕事向けの講座"),
        ("保険営業の方向けの無料の学習サービス", "お仕事向けの無料の学習サービス")],
    12: [("セミナー参加後に、サロンの", "セミナー参加後に、講座・サロンの")],
    17: [("営業力を上げたい・", "成果を上げたい・")],
    19: [("直後に、サロンの", "直後に、講座・サロンの")],
    20: [("期待成果・費用", "費用")],
    34: [("両想い営業・ライフプラン", "講座の考え方と"), ("ブロックを動画で紹介", "特徴を動画で紹介"),
         ("週2日で成果を出す", "セミナーで学べる"), ("面談メソッド", "要点3つ"),
         ("（成約／働く時間）", "（目的別）"),
         ("地方からも参加できる", "オンラインで参加できる"),
         ("「生産性4倍セミナーの特徴」", "「セミナーの特徴」")],
    35: [("営業のヒントを届ける", "学びのヒントを届ける")],
    36: [("営業のお悩みタイプ診断", "お悩みタイプ診断"),
         ("保険営業の方／一般の方", "お仕事で活かしたい方／ご自身のために学びたい方"),
         ("保険営業の方：成約を上げたい／働く時間を減らしたい", "お仕事の方：成果を上げたい／働く時間を減らしたい"),
         ("一般の方：資産形成を始めたい／保険を見直したい", "ご自身のため：資産形成を始めたい／副業・スキルを身につけたい"),
         ("成約タイプ", "成果タイプ"), ("面談で決めきれない方", "成果を上げたい方"),
         ("面談の進め方が変わった受講生の声", "成果が変わった受講生の声"),
         ("両想い営業の見どころ", "講座の見どころ"),
         ("講師紹介（週2日の営業で毎年TOT）", "講師紹介（講師の実績）"),
         ("講師メッセージ（働く時間を半分以下に）", "講師メッセージ（時間の使い方）"),
         ("一般の方向けのマネーセミナーのご案内", "一般の方向けのセミナー（マネー・副業など）のご案内")],
    37: [("両想い営業・ライフプランブロックの要点をおさらい", "当日の要点をおさらい"),
         ("「契約で頭がいっぱい」だった方が、個人保険でTOT", "同じ悩みから始めて、成果が出た方の声"),
         ("ライフプランブロックは、サロンの方だけが使える", "入会した方だけが受けられる内容"),
         ("ロープレ動画で面談を振り返り、TOTを達成", "続け方の工夫と、出た成果"),
         ("毎日遅くまで働いた頃から、働く時間を半分以下に", "講師が、この講座で伝えたいこと"),
         ("地方からも参加できる・録画はない・初心者でも分かる・参加は無料", "オンラインで参加できる・初心者でも分かる・参加は無料"),
         ("谷岡様の経歴：契約件数全国1位で表彰・週2日の営業で毎年TOT", "講師の経歴と実績")],
}

PLANS = ["プラン①\n初期設計のみ", "プラン②\n定期運用（月3回程度）", "プラン③\n定期運用（月5回程度）"]
FEE_ROWS = [  # （行の見出し, 3プラン）。すなおの「開催のない月／開催月」の2行ずつは1行にまとめる
    ("初期費用\n（ポップアップあり）", ["115,000円〜", "115,000円〜", "115,000円〜"]),
    ("月額コンサルティング費", ["なし", "200,000円〜", "300,000円〜"]),
    ("月額ツール費\n（ポップアップ）", ["30,000円〜", "30,000円〜", "30,000円〜"]),
    ("月額費用", ["30,000円〜", "230,000円〜", "330,000円〜"]),
]
FEE_NOTE = ["※初期費用＝運用設計・構築 100,000円〜＋ポップアップ 15,000円。LINE公式アカウントの利用料は別（配信数による）。契約期間は6か月〜",
            "※サンクスLINE誘導はオプション（初期150,000円・月額30,000円〜）",
            "※施策内容・配信本数・連携範囲等により、費用は個別にお見積もりいたします。"]
UNYO_KIKAKU = ["なし（貴社で運用）", "月3回程度", "月5回程度"]
# 言葉の途中で切れる折り返しを、文の切れ目で2行にする（旧ページ番号, 行, 列, 文）
CELL_FIX = [
    (6, 1, 0, "企業向け\n研修サービス市場"),
    (8, 0, 2, "ご自身のために学びたい方"),
    (8, 3, 1, "セミナーの見どころ・受講生の声・\n講座・サロンのご案内"),
    (12, 3, 3, "セミナー参加後に、\n講座・サロンの入会を案内"),
    (15, 3, 2, "・参加後に、振り返りと講座・サロンを案内\n・受講生の声で、入会の判断材料を届ける"),
]


def replace_runs(s, pairs, must):
    hit = {a: 0 for a, _ in pairs}
    for t in s._element.iter(A_T):
        if not t.text:
            continue
        for a, b in pairs:
            if a in t.text:
                t.text = t.text.replace(a, b)
                hit[a] += 1
    if must:
        miss = [a for a, n in hit.items() if not n]
        assert not miss, miss


def cell_text(cell, text):
    """セルの文字を差し替える（1段落目の書式で、\\nは段落を分ける）"""
    txb = cell._tc.find(qn("a:txBody"))
    ps = txb.findall(qn("a:p"))
    tmpl = copy.deepcopy(ps[0])
    for r in tmpl.findall(qn("a:r"))[1:] + tmpl.findall(qn("a:br")):
        tmpl.remove(r)
    for p in ps:
        txb.remove(p)
    for line in text.split("\n"):
        p = copy.deepcopy(tmpl)
        rs = p.findall(qn("a:r"))
        if rs:
            rs[0].find(A_T).text = line
        txb.append(p)


def drop_row(table_shape, k):
    tbl = table_shape.table._tbl
    tr = tbl.findall(qn("a:tr"))[k]
    h = int(tr.get("h"))
    tbl.remove(tr)
    table_shape.height = table_shape.height - h


def by_id(s, sid):
    return next(sh for sh in s.shapes if sh.shape_id == sid)


def head(s):
    return " ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame)[:120]


def main(src, out):
    prs = Presentation(src)
    S = {i + 1: s for i, s in enumerate(prs.slides)}
    assert len(S) == 63
    assert "現状把握" in head(S[3]) and "概算お見積り" in head(S[21]) and "費用の比較" in head(S[22]) and "成果報酬" in head(S[25])

    for n in range(1, 44):
        replace_runs(S[n], COMMON, must=False)
    for n, pairs in PAGE.items():
        replace_runs(S[n], pairs, must=True)

    # P4 サービス特性の「広告」の行はそのまま。P6 の4行目（保険営業の方の接触回数）を LINE の利用者に差し替え
    t6 = by_id(S[6], 50)
    cell_text(t6.table.cell(4, 0), "LINEの利用者")
    cell_text(t6.table.cell(4, 1), "LINEの国内月間利用者数は1億人（2025年12月時点）")
    cell_text(t6.table.cell(4, 2), "学びの案内も、LINEで届けられる相手が多い")
    note6 = by_id(S[6], 53)
    for t in note6._element.iter(A_T):
        if t.text and "接触回数" in t.text:
            t.text = re.sub(r"／接触回数：.*$", "／LINEの利用者：LINEヤフー発表（2026年1月29日。2025年12月時点の国内月間利用者数）", t.text)
    assert "LINEヤフー発表" in note6.text_frame.text, note6.text_frame.text

    # P9 他社の画面：社名・URL・ロゴを伏せた画像に
    pic = next(sh for sh in S[9].shapes if sh.shape_type == 13)
    rid = pic._element.find(".//" + qn("a:blip")).get(qn("r:embed"))
    S[9].part.related_part(rid)._blob = open(IMG_A, "rb").read()
    note9 = by_id(S[9], 55)
    for t in note9._element.iter(A_T):
        if t.text:
            t.text = t.text.replace("R&C株式会社のプレスリリースより", "プレスリリースより")
    assert "R&C" not in note9.text_frame.text

    # P36 スマホの診断のボタン
    for sid, txt in ((44, "自分のため"), (45, "仕事で")):
        sh = by_id(S[36], sid)
        sh.text_frame.paragraphs[0].runs[0].text = txt
        for r in sh.text_frame.paragraphs[0].runs[1:]:
            r.text = ""

    # P23 料金プラン（DYMの目安に。開催のない月／開催月の2行ずつを1行に）
    t23 = next(sh for sh in S[23].shapes if sh.has_table)
    drop_row(t23, 6)
    drop_row(t23, 3)
    tb = t23.table
    for c, name in enumerate(PLANS, 1):
        cell_text(tb.cell(0, c), name)
    for r, (label, vals) in enumerate(FEE_ROWS, 1):
        cell_text(tb.cell(r, 0), label)
        for c, v in enumerate(vals, 1):
            cell_text(tb.cell(r, c), v)
    y = Emu(t23.top + t23.height).inches + 0.15
    rich(S[23], 0.5 * 2.54, y * 2.54, 9.8 * 2.54, 0.7 * 2.54, [[(line, 9, False, GRAY)] for line in FEE_NOTE],
         anchor=MSO_ANCHOR.TOP)

    # P24 運用内容
    t24 = next(sh for sh in S[24].shapes if sh.has_table).table
    for c, name in enumerate(PLANS, 1):
        cell_text(t24.cell(0, c), name)
    assert t24.cell(3, 0).text.startswith("企画配信")
    for c, v in enumerate(UNYO_KIKAKU, 1):
        cell_text(t24.cell(3, c), v)

    for n, r, c, text in CELL_FIX:
        tb = next(sh for sh in S[n].shapes if sh.has_table).table
        cell_text(tb.cell(r, c), text)

    for n in (25, 22, 21, 3):
        drop_slide(prs, S[n])

    # すなおにしか無い言葉が本編に残っていないか
    NG = re.compile("すなお|谷岡|ハッピーライフ|野沢|R&C|TOT|両想い|ライフプランブロック|生産性4倍|保険営業|年4回")
    left = []
    for i, s in enumerate(prs.slides, 1):
        if i > 39:
            break
        for t in s._element.iter(A_T):
            if t.text and NG.search(t.text):
                left.append((i, t.text))
    assert not left, left
    prs.save(out)
    print("saved:", out, len(prs.slides))


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
