# ver2.2 → ver2.3：SIM3のCV地点を資料の定義に合わせる（CV①＝会員登録、CV②＝求人応募）
import sys
from openpyxl import load_workbook
from openpyxl.workbook.properties import CalcProperties
SRC,OUT=sys.argv[1],sys.argv[2]
wb=load_workbook(SRC)
ws=wb['SIM3_LINE登録2%']; cp=wb['SIM3_比較']
ramp=[0.62,0.81,0.93,1.00,1.03,1.05]
deliv=[1,2,3,3,4,6]   # 配信経由の応募（仮）
ws['C8']='会員登録（LINEで登録・Profile+）'
ws['C9']='求人応募（LINE経由）'
for i,c in enumerate('EFGHIJ'):
    ws[f'{c}36']=f'=ROUND({c}14*0.02*{ramp[i]},0)'
    ws[f'{c}37']=f'={deliv[i]}+ROUND({c}29*(1-SUM($E$24:{c}25)/{c}28)*0.3*2*0.015,0)'
    ws[f'{c}40']=f'=+{c}37'
ws['D40']='=+D37'
ws['N36']='CV①会員登録：UU×2.0%×立ち上げ係数（Profile+で入力ほぼゼロ）。CPA・CVRには含めない'
ws['N37']='CV②求人応募（LINE経由）＝配信経由（仮の数字）＋リッチメニュー経由（応募済みを除く有効な友だち×30%×2回×1.5%）'
ws['N40']='合計CVはCV②（応募）のみ。CPA・CVRは応募で計算'
cp['A17']='CV①_会員登録（参考・CPAに含めない）'
cp['A18']='CV②_求人応募（LINE経由）'
for c in 'CDE':
    cp[f'{c}16']=f'=+{c}18'
cp['A16']='CV数（+LINEOA施策）※応募のみ'
wb.calculation=CalcProperties(fullCalcOnLoad=True)
wb.save(OUT)
