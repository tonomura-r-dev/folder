# -*- coding: utf-8 -*-
"""チェングロウスver1.1：S6を「ステップ配信のブレスト」型に作り直し＋P16/P20をSIM ver1.1の数字に（2026-09-24）。
S6：左上＝前後検索（4色）＋横軸に「LINE追加の機会（当日〜翌日）」の緑帯
    右上＝凡例4色＋LINE追加率を上げるコンテンツ2つ（離脱防止で見せるもの）＋根拠
    下段＝LINEでのリード有効化率を上げるコンテンツ（日付カード。紺＝コンテンツ／緑＝面談オファー＝CV）
P16/P20：SIM ver1.1（CPF単価300円・コンサル費15万）＝半年89件・累計CPA 29,697円
  python3 _build/patch_chengrowth_0924e.py <入力pptx> <出力pptx>
"""
import copy
import sys
from pathlib import Path
from pptx import Presentation
from pptx.util import Cm, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
SRC, OUT = sys.argv[1], sys.argv[2]
prs = Presentation(SRC)
NAVY, GREEN, INK = "1F285A", "06C755", "333333"


def meiryo(run):
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        e = rPr.find(qn(tag))
        if e is None:
            e = rPr.makeelement(qn(tag), {})
            rPr.append(e)
        e.set("typeface", "メイリオ")


def box(slide, x, y, w, h, fill, lines, color="FFFFFF", size=10.5, bold=False, anchor=MSO_ANCHOR.MIDDLE,
        align=PP_ALIGN.CENTER, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    sp = slide.shapes.add_shape(shape, Cm(x), Cm(y), Cm(w), Cm(h))
    if fill:
        sp.fill.solid()
        sp.fill.fore_color.rgb = RGBColor.from_string(fill)
    else:
        sp.fill.background()
    if line:
        sp.line.color.rgb = RGBColor.from_string(line)
        sp.line.width = Pt(1)
    else:
        sp.line.fill.background()
    sp.shadow.inherit = False
    try:
        sp.adjustments[0] = 0.08
    except Exception:
        pass
    tf = sp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Cm(0.15)
    tf.margin_top = tf.margin_bottom = Cm(0.08)
    for i, ln in enumerate(lines):
        text, sz, b, c = (ln, size, bold, color) if isinstance(ln, str) else ln
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = text
        r.font.size = Pt(sz)
        r.font.bold = b
        r.font.color.rgb = RGBColor.from_string(c)
        meiryo(r)
    return sp


def arrow(slide, x1, y1, x2, y2, color="4472C4"):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Cm(x1), Cm(y1), Cm(x2), Cm(y2))
    c.line.color.rgb = RGBColor.from_string(color)
    c.line.width = Pt(1.5)
    ln = c.line._get_or_add_ln()
    tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle"})
    ln.append(tail)


# ================= S6 =================
s = prs.slides[5]
tree = s.shapes._spTree
keep = {2, 3, 4, 10}   # タイトル・リード・区切り線・出典
for sh in list(s.shapes):
    if sh.shape_id not in keep:
        tree.remove(sh._element)
lead = [sh for sh in s.shapes if sh.shape_id == 3][0]
tf = lead.text_frame
for p in tf.paragraphs[1:]:
    p._p.getparent().remove(p._p)
tf.paragraphs[0].runs[0].text = "「カーディーラー」の検索前後15日間。LINEの友だちにできる機会は、検索当日〜翌日だけ。"
for r in tf.paragraphs[0].runs[1:]:
    r.text = ""
lead.top, lead.height = Cm(1.75), Cm(1.9)

# 左上：前後検索
IL, IT, IW = 1.2, 4.05, 14.2
IH = IW * 1167 / 2000
s.shapes.add_picture(str(ROOT / "_images/chengrowth_zengo_4color.png"), Cm(IL), Cm(IT), width=Cm(IW))
px = lambda p: IL + p / 2000 * IW
py = lambda p: IT + p / 1167 * IH
gx0, gx1 = px(1033), px(1270)
g = box(s, (gx0 + gx1) / 2 - 1.3, py(1150), 2.6, 0.62, "C6EFCE", [("LINE追加", 9.5, True, "00873C")], shape=MSO_SHAPE.RECTANGLE)
band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(gx0), Cm(IT), Cm(gx1 - gx0), Cm(IH - 0.1))
band.fill.solid()
band.fill.fore_color.rgb = RGBColor.from_string("06C755")
band.fill._xPr.find(qn("a:solidFill"))[0].append(band.fill._xPr.makeelement(qn("a:alpha"), {"val": "14000"}))
band.line.fill.background()
band.shadow.inherit = False

# 右上：凡例＋LINE追加率を上げるコンテンツ
RX, RW = 15.9, 10.42
LEG = [("職業・働き方", "FBE0CF", "C0561A"), ("整備・工具・カー用品", "D6E2F3", "2B5797"),
       ("お金・手続き", "E8DAF5", "7440A8"), ("クルマ選び（車種・店舗）", "CCE7E4", "00796B")]
for i, (t, f, c) in enumerate(LEG):
    box(s, RX + (i % 2) * (RW / 2 + 0.05), 4.05 + (i // 2) * 0.72, RW / 2 - 0.05, 0.62, f, [(t, 9.5, True, c)])
box(s, RX, 5.7, RW, 0.62, "FFF2CC", [("LINE追加率を上げるコンテンツ（検索当日〜翌日）", 10.5, True, INK)])
CW = RW / 2 - 0.1
box(s, RX, 6.45, CW, 2.55, NAVY, [("離脱防止で…", 10, False, "FFFFFF"), ("「整備士の年収相場」", 11, True, "FFFFFF"),
                                   ("をLINEで見られます", 10.5, True, "FFFFFF")])
box(s, RX + CW + 0.2, 6.45, CW, 2.55, NAVY, [("この機会に", 10, False, "FFFFFF"), ("LINEで30秒", 11, True, "FFFFFF"),
                                             ("適職診断しませんか？", 10.5, True, "FFFFFF")])
box(s, RX, 9.1, CW, 1.95, "DDEBF7", [("当日の検索は「年収」", 10, True, NAVY), ("その場で答えを出す", 10, False, NAVY)])
box(s, RX + CW + 0.2, 9.1, CW, 1.95, "DDEBF7", [("電話なし・履歴書なし", 10, True, NAVY), ("応募前の人でも動きやすい", 10, False, NAVY)])
arrow(s, (gx0 + gx1) / 2, py(1120), RX - 0.1, 7.7)

# 下段：リード有効化コンテンツ
BY = IT + IH + 0.75
box(s, 1.2, BY, 25.12, 0.62, "FFF2CC", [("LINEでのリード有効化率を上げるコンテンツ（翌日〜14日・前後検索の関心の順番）", 10.5, True, INK)])
CARDS = [
    ("1日目", "職種図鑑", "整備・受付・営業の違い", NAVY),
    ("2日目", "働き方", "休み・残業・ホワイトな職場", NAVY),
    ("3日目", "面談のご案内①", "（CV）", GREEN),
    ("5日目", "他職種と比較", "営業・施工管理", NAVY),
    ("7日目", "資格別の年収相場", "キャリアマップ", NAVY),
    ("10日目", "現場のリアル", "冬の繁忙期", NAVY),
    ("14日目", "面談のご案内②", "（CV）", GREEN),
]
n = len(CARDS)
gap = 0.2
cw = (25.12 - gap * (n - 1)) / n
cy = BY + 0.75
ch = 17.2 - cy
for i, (d, t, sub, col) in enumerate(CARDS):
    box(s, 1.2 + i * (cw + gap), cy, cw, ch, col,
        [(d, 11, True, "FFFFFF"), (t, 10.5, True, "FFFFFF"), (sub, 9.5, False, "FFFFFF")])
arrow(s, px(1600), py(1120), px(1700), BY - 0.05)

foot = [sh for sh in s.shapes if sh.shape_id == 10][0]
foot.top = Cm(17.45)
tree.remove(foot._element)
tree.append(foot._element)

# 図形IDの重複を解消
seen, els = set(), [e for e in tree.iter() if e.tag == qn("p:cNvPr")]
mx = max(int(e.get("id")) for e in els)
for e in els:
    if int(e.get("id")) in seen:
        mx += 1
        e.set("id", str(mx))
    seen.add(int(e.get("id")))

# ================= P16・P20：SIM ver1.1 =================
REPL = {
    15: [("コンサル費 ¥70,000", "コンサル費 ¥150,000"),
         ("約4,650人", "約4,350人"), ("約2,500人", "約2,350人"), ("2,560件/月", "2,390件/月"), ("25件/月", "23件/月"),
         ("半年累計95件・累計CPA 22,768円", "半年累計89件・累計CPA 29,697円"),
         ("現状CPA 6万の約1/3・目標2〜2.5万の圏内", "現状の整備士CPA 6万円の約半分"),
         ("初期24.5万円・月額10.5万円〜の運用で、面談・応募を半年95件と見込みます。",
          "初期24.5万円・月額18.5万円〜の運用で、面談・応募を半年89件と見込みます。"),
         ("CPF単価280円ほか仮置き値を含む", "CPF単価300円（LINEヤフー公式事例）ほか仮置き値を含む")],
    19: [("面談・応募 95件（6ヶ月累計・SIM①）＝ 累計CPA 22,768円として算出",
          "面談・応募 89件（6ヶ月累計・SIM①）＝ 累計CPA 29,697円として算出")],
}
cnt = 0
for idx, pairs in REPL.items():
    for sh in prs.slides[idx].shapes:
        if not sh.has_text_frame:
            continue
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                for a, b in pairs:
                    if a in r.text:
                        r.text = r.text.replace(a, b)
                        cnt += 1
prs.save(OUT)
print("saved", OUT, "replaced", cnt)
