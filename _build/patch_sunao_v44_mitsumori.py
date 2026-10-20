# -*- coding: utf-8 -*-
"""すなお ver4.3 → ver4.4：費用対効果を、タナカホームの形（概算お見積り・料金プラン・運用内容）で出す（2026-10-20 殿村さん）。

    python -m _build.patch_sunao_v44_mitsumori <ver4.3.pptx> <タナカホーム v2.pptx> <SIM ver1.3.xlsx> [出力.pptx]

  殿村さん「費用対効果は詳細に出さなくていい。添付画像（タナカホームの概算お見積り・料金プラン・運用内容）の形式で」。
  ・想定の費用対効果 → 概算お見積り（タナカホーム v2 の「概算お見積り」のページを写して、すなおの中身に差し替える）
  ・想定の費用対効果（詳細）・弊社想定シミュレーション（2枚）は外す（月ごとの細かい数字は出さない）
  ・料金（ver3の表）→ 料金プラン（タナカの表の形）。中身は ver3 の3プランのまま
    （①初期設計のみ／②イベントに合わせた運用＝平常期 5〜10万・開催月 30万程度／③通年の運用・改善＝平常期 20万〜・開催月 30万〜）
  ・料金プランの後ろに、運用内容（タナカの表の形）を足す。成果報酬型プランの月額は料金プランの注記に書く
  数字は SIM ver1.3（概算お見積り＝プラン②。平常期のコンサルは10万円で計算）。
"""
import copy
import sys

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches

from _build.arauchi_v10.build import clone
from _build.build_sunao_v1 import all_meiryo
from _build.build_sunao_v4 import find, load_sim, n, put, set_cell
from _build.patch_tanaka_v5 import copy_slide

A16 = "{http://schemas.microsoft.com/office/drawing/2014/main}"
PLANS = ["プラン①\n初期設計のみ", "プラン②\nイベントに合わせた運用", "プラン③\n通年の運用・改善"]


def title_of(s):
    for sh in s.shapes:
        if sh.has_text_frame and sh.name.startswith("Google Shape") and sh.top < Inches(0.6):
            return sh.text_frame.text.strip()
    return ""


def fix_ids(tbl):
    for tag, attr, base in (("a:gridCol", "colId", 20000), ("a:tr", "rowId", 10000)):
        for i, el in enumerate(tbl.findall(".//" + qn(tag))):
            idel = el.find(".//" + A16 + attr)
            if idel is not None:
                idel.set("val", str(base + i))


def set_size(cell, size):
    for r in cell._tc.iter(qn("a:rPr"), qn("a:endParaRPr")):
        r.set("sz", str(int(size * 100)))


def numbers(sim):
    m = range(1, 7)
    ev = next(k for k in m if sim[12][k] == "イベント期")
    nv = next(k for k in m if sim[12][k] != "イベント期")
    tool = sum((sim[r][nv] or 0) for r in (66, 67, 68))          # ポップアップ＋サンクスLINE＋アカウント（月額）
    return dict(init=n(sim[73][0]), ev=n(sim[73][ev]), nv=n(sim[73][nv]), tool=n(tool),
                friends=n(sim[28][-1]), active=n(sim[29][-1]),
                seminar=sum(n(sim[36][k]) + n(sim[38][k]) for k in m), salon=sum(n(sim[37][k]) for k in m))


def p_mitsumori(s, d):
    """タナカホームの「概算お見積り」の形（初期／月次／6か月後）"""
    put(s, 339, f"概算お見積り：6か月で、追える人数{d['active']}人・セミナー予約{d['seminar']}件")
    put(s, 381, "プラン②｜イベントに合わせた運用で、集めた方を友だちにし、次のセミナー予約・入会へつなげる")
    put(s, 343, f"¥{d['init']:,}〜")
    put(s, 344, f"開催月 ¥{d['ev']:,}", f"平常期 ¥{d['nv']:,}")
    for (a, b), (k, v) in zip(((365, 366), (367, 368), (369, 370), (371, 372), (373, 374)),
                              ((["ポップアップ", "ツール"], ["導入"]), (["サンクス", "LINE誘導"], ["導入"]),
                               (["LINE公式", "アカウント開設"], ["対応"]), (["ステップ配信"], ["5通を構築"]),
                               (["配信の設計"], ["8フェーズで設計"]))):
        put(s, a, *k)
        put(s, b, *v, size=12 if len(v[0]) > 6 else None)
    for (a, b), (k, v) in zip(((349, 350), (351, 352), (353, 354), (355, 356), (357, 358)),
                              ((["企画配信"], ["開催月 4〜5回", "平常期 2回"]), (["セミナー前日の", "お知らせ"], ["予約ごとに配信"]),
                               (["ステップ配信"], ["5通の改善"]), (["友だちの入口"], ["運用"]),
                               (["個別チャット"], ["貴社で対応"]))):
        put(s, a, *k)
        put(s, b, *v, size=12 if len(v) > 1 else None)
    for sid in (359, 360):                                      # 6行目（月次定例会）は外す（プラン②にレポート・定例会は無い）
        e = find(s, sid)._element
        e.getparent().remove(e)
    put(s, 361, "友だち数")
    put(s, 375, f"{d['friends']:,}人")
    put(s, 362, "追える人数")
    put(s, 376, f"{d['active']:,}人")
    put(s, 364, "セミナー予約")
    put(s, 378, f"{d['seminar']}件/6か月")
    lab, val = clone(s, find(s, 364)), clone(s, find(s, 378))     # 4行目：サロン入会
    step = find(s, 364).top - find(s, 362).top
    for src, new in ((find(s, 364), lab), (find(s, 378), val)):
        new.top = src.top + step
    lab.text_frame.paragraphs[0].runs[0].text = "サロン入会"
    val.text_frame.paragraphs[0].runs[0].text = f"{d['salon']}件/6か月"


def p_ryokin(s, d):
    """タナカホームの「料金プラン」の形。中身は ver3 の3プラン"""
    put(s, 635, "料金プラン：3つのプランの費用")
    put(s, 636, "料金プランの比較")
    put(s, 637, "ご予算と運用の体制に合わせて、3つのプランをご用意")
    gf = next(sh for sh in s.shapes if sh.has_table)
    tbl = gf._element.find(".//" + qn("a:tbl"))
    trs = tbl.findall(qn("a:tr"))
    trs[2].addnext(copy.deepcopy(trs[2]))                      # 月額コンサル（平常期／セミナー開催月）の2行に
    trs[4].addnext(copy.deepcopy(trs[4]))                      # 月額費用（平常期／セミナー開催月）の2行に
    fix_ids(tbl)
    t = gf.table
    tool = f"{d['tool']:,}円〜"
    rows = [["項目"] + PLANS,
            ["初期費用"] + [f"{d['init']:,}円〜"] * 3,
            ["月額コンサルティング費\n（平常期）", "なし", "50,000〜100,000円", "200,000円〜"],
            ["月額コンサルティング費\n（セミナー開催月）", "なし", "300,000円程度", "300,000円〜"],
            ["月額ツール・LINE関連費"] + [tool] * 3,
            ["月額費用\n（平常期）", tool, f"{50000 + d['tool']:,}〜{100000 + d['tool']:,}円", f"{200000 + d['tool']:,}円〜"],
            ["月額費用\n（セミナー開催月）", tool, f"{300000 + d['tool']:,}円程度", f"{300000 + d['tool']:,}円〜"]]
    for r, vals in enumerate(rows):
        for c, v in enumerate(vals):
            set_cell(t.cell(r, c), v)
    for r in (5, 6):
        for c in (1, 2, 3):
            set_size(t.cell(r, c), 16)
    set_size(t.cell(5, 2), 14)                                 # 「115,000〜165,000円」を1行に収める
    heights = [0.7, 0.5, 0.6, 0.6, 0.5, 0.65, 0.65]
    for r, h in enumerate(heights):
        t.rows[r].height = Inches(h)
    gf.height = Inches(sum(heights))
    notes = ["※初期費用＝運用設計・構築 100,000円〜＋ポップアップツール 15,000円＋サンクスLINE誘導 150,000円＋LINE公式アカウント開設 30,000円",
             "※月額ツール・LINE関連費＝ポップアップツール 30,000円＋サンクスLINE誘導 30,000円〜＋LINE公式アカウント利用料 5,000円〜",
             f"※成果報酬型プラン：初期 {d['init']:,}円〜／月額 {d['tool']:,}円＋成果報酬（セミナー予約 12,000円／件）",
             "※契約期間は6か月〜。広告費は含まない。概算お見積りのページはプラン②の場合（平常期の月額コンサルティング費は100,000円で計算）",
             "※施策内容・配信本数・連携範囲等により、費用は個別にお見積もりいたします。"]
    shapes = [find(s, sid) for sid in (640, 641, 642, 643)]
    shapes.append(clone(s, shapes[-1]))
    y = gf.top + gf.height + Inches(0.12)
    for sh, text in zip(shapes, notes):
        sh.top = y
        sh.text_frame.paragraphs[0].runs[0].text = text
        for r in sh.text_frame.paragraphs[0].runs[1:]:
            r._r.getparent().remove(r._r)
        y += Inches(0.25)


def p_naiyou(s):
    """タナカホームの「運用内容」の形"""
    put(s, 2, "運用内容：プランごとの対応範囲")
    put(s, 4, "運用内容の比較")
    put(s, 5, "3つのプランの運用内容の比較")
    gf = next(sh for sh in s.shapes if sh.has_table)
    tbl = gf._element.find(".//" + qn("a:tbl"))
    tr = tbl.findall(qn("a:tr"))[2]                            # 「サイトのLINEボタン」の行は外す（すなおのLPには置かない）
    gf.height = gf.height - int(tr.get("h"))
    tbl.remove(tr)
    fix_ids(tbl)
    t = gf.table
    rows = [["項目"] + PLANS,
            ["ポップアップツール・\nサンクスLINE誘導", "導入・運用", "導入・運用", "導入・運用"],
            ["ステップ配信", "5通\n（構築のみ）", "5通\n（構築・改善）", "5通\n（構築・改善）"],
            ["企画配信", "なし（貴社で運用）", "開催月 月4〜5回\n平常期 月2回", "開催月 月5回\n平常期 月3回"],
            ["セミナー前日のお知らせ", "設定のみ", "設定・改善", "設定・改善"],
            ["結果の分析・改善\n（月次レポート・定例会）", "なし", "なし", "対応"],
            ["個別チャット", "貴社で対応", "貴社で対応", "貴社で対応"]]
    for r, vals in enumerate(rows):
        for c, v in enumerate(vals):
            set_cell(t.cell(r, c), v)
    last = t.rows[len(rows) - 1]                               # 最後の行（元はLINE広告の2行分）を他の行と同じ高さに
    gf.height = gf.height - last.height + t.rows[1].height
    last.height = t.rows[1].height


def move_to(prs, slide, index):
    lst = prs.slides._sldIdLst
    el = next(e for e in lst if prs.part.related_part(e.rId) is slide.part)
    lst.remove(el)
    lst.insert(index, el)


def drop(prs, slide):
    lst = prs.slides._sldIdLst
    el = next(e for e in lst if prs.part.related_part(e.rId) is slide.part)
    prs.part.drop_rel(el.rId)
    lst.remove(el)


def main(src, tanaka_path, sim_path, out):
    d = numbers(load_sim(sim_path))
    prs = Presentation(src)
    tanaka = Presentation(tanaka_path)
    tk = {title_of(s): s for s in tanaka.slides}
    slides = list(prs.slides)
    hiyou = next(s for s in slides if title_of(s) == "想定の費用対効果")
    detail = next(s for s in slides if any(sh.name.startswith("月ごとの見込み") for sh in s.shapes))
    sims = [s for s in slides if title_of(s) == "弊社想定シミュレーション"]
    ryokin = next(s for s in slides if any(sh.has_text_frame and sh.text_frame.text.strip() == "料金" for sh in s.shapes))
    assert len(sims) == 2

    mitsu = copy_slide(prs, next(v for k, v in tk.items() if k.startswith("概算お見積り")))
    plan = copy_slide(prs, tk["料金プラン：3つのプランの費用"])
    naiyou = copy_slide(prs, tk["運用内容：プランごとの対応範囲"])
    p_mitsumori(mitsu, d)
    p_ryokin(plan, d)
    p_naiyou(naiyou)
    all_meiryo([mitsu, plan, naiyou])

    order = list(prs.slides)
    move_to(prs, mitsu, order.index(hiyou))
    order = list(prs.slides)
    move_to(prs, plan, order.index(ryokin))
    move_to(prs, naiyou, order.index(ryokin) + 1)
    for s in [hiyou, detail, ryokin] + sims:
        drop(prs, s)
    prs.save(out)
    print("saved:", out, len(prs.slides), "枚", d)


if __name__ == "__main__":
    main(*sys.argv[1:4], sys.argv[4] if len(sys.argv) > 4 else sys.argv[1])
