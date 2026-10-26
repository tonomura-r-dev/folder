"""
10月号（正）の P.6 を OK/NG の表に置き換え、P.3〜P.7 の枠を文字量に合わせて詰めて、
「■ 詳細」の下の範囲で上下中央に置く（2026-10-26 殿村さん指示）。
使い方: python3 _build/patch_trend_202610_layout2.py <pptx>
"""
import sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
import trend_report_engine as engine

F = sys.argv[1]
prs = Presentation(F)
BOTTOM = 7.05


def text(sh):
    return sh.text_frame.text.strip() if sh.has_text_frame else ""


def text_w(s, pt):
    return sum((0.55 if ord(c) < 0x2000 else 1.0) for c in s) * pt / 72


def lines(s, pt, w):
    return max(1, -(-text_w(s, pt) // w))


def card_need(sh):
    tf = sh.text_frame
    usable = Emu(sh.width - tf.margin_left - tf.margin_right).inches - 0.1
    h = 0.0
    for i, p in enumerate(tf.paragraphs):
        if not p.runs:
            continue
        pt = p.runs[0].font.size.pt if p.runs[0].font.size else 13
        n = sum(lines(seg, pt, usable) for seg in p.text.split("\x0b"))
        lh = pt * (1.3 if i == 0 else 1.45) / 72
        sb = (p.space_before.pt if p.space_before is not None else 0) / 72
        h += n * lh + (sb if i > 0 else 0)
    return Inches(h + 0.32)


def body_of(slide):
    head = next(x for x in slide.shapes if text(x) == "■ 詳細")
    top = head.top + head.height
    return head, [x for x in slide.shapes if x.top >= top - Inches(0.02) and x.top < Inches(7.0)]


def is_card(x):
    return x.shape_type == 1 and getattr(x, "auto_shape_type", None) == MSO_SHAPE.ROUNDED_RECTANGLE


# ---------- P.6：表に置き換え ----------
s6 = list(prs.slides)[5]
assert "新商品のご案内" in next(text(x) for x in s6.shapes if text(x).startswith("【"))
_, body = body_of(s6)
for x in body:
    x._element.getparent().remove(x._element)
rows = [
    ("", "OK", "NG"),
    ("告知してよい場所",
     "広告主のLINE公式アカウント・アプリ・SNS、Web全般（LINEヤフー広告も可）、店頭POP・屋外広告",
     "LINEヤフーが持つアプリ（スタンプショップ、LINEミニアプリ、LINEヤフーの公式アカウントなど）"),
    ("コードを渡す条件",
     "Webでの行動（商品購入など）、店頭での行動（来店・購入）",
     "LINEアプリ内の行動（友だち追加・アンケート回答、ミニアプリ内の行動）、広告主アプリ内の行動（会員登録など）、アプリのダウンロードが必須のもの"),
]
colw = [1.9, 3.6, 4.14]
rowh = [0.42, 0.95, 1.15]
gf = s6.shapes.add_table(3, 3, Inches(0.6), Inches(3.5), Inches(sum(colw)), Inches(sum(rowh)))
tbl = gf.table
for j, w in enumerate(colw):
    tbl.columns[j].width = Inches(w)
for i, h in enumerate(rowh):
    tbl.rows[i].height = Inches(h)
for i, row in enumerate(rows):
    for j, val in enumerate(row):
        cell = tbl.cell(i, j)
        cell.fill.solid()
        if i == 0:
            cell.fill.fore_color.rgb = engine.NAVY
        elif j == 0:
            cell.fill.fore_color.rgb = engine.LIGHT_BLUE
        else:
            cell.fill.fore_color.rgb = engine.WHITE
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = cell.margin_right = Inches(0.12)
        p = cell.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER if (i == 0 or j == 0) else PP_ALIGN.LEFT
        r = p.add_run()
        r.text = val
        r.font.name = "メイリオ"
        if i == 0:
            r.font.size, r.font.bold, r.font.color.rgb = Pt(14), True, engine.WHITE
        elif j == 0:
            r.font.size, r.font.bold, r.font.color.rgb = Pt(14), True, engine.NAVY
        else:
            r.font.size, r.font.color.rgb = Pt(13), engine.BODY_GRAY
note = s6.shapes.add_textbox(Inches(0.6), Inches(3.5 + sum(rowh) + 0.08), Inches(9.64), Inches(0.35))
np_ = note.text_frame.paragraphs[0]
nr = np_.add_run()
nr.text = "※ スタンプをもらうまでの流れに、アプリへの遷移が少しでも入るとNG"
nr.font.size, nr.font.name, nr.font.color.rgb = Pt(12), "メイリオ", engine.BODY_GRAY

# ---------- P.3〜P.7：枠を詰めて上下中央へ ----------
for slide in list(prs.slides)[2:7]:
    head, body = body_of(slide)
    cards = [x for x in body if is_card(x)]
    # 同じ段のカードは高さをそろえて、上から詰め直す
    rows_ = {}
    for c in cards:
        rows_.setdefault(round(Emu(c.top).inches, 1), []).append(c)
    y = None
    for key in sorted(rows_):
        row = rows_[key]
        h = max(card_need(c) for c in row)
        top = row[0].top if y is None else y
        for c in row:
            c.top, c.height = top, h
        y = top + h + Inches(0.2)
    # まとまり全体を上下中央へ
    top = min(x.top for x in body)
    bottom = max(x.top + x.height for x in body)
    area_top = head.top + head.height + Inches(0.08)
    shift = int(area_top + (Inches(BOTTOM) - area_top - (bottom - top)) / 2 - top)
    for x in body:
        x.top = x.top + shift

engine.force_font(prs)
prs.save(F)
print("saved", F)
