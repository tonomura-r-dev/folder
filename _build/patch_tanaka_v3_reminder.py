# -*- coding: utf-8 -*-
"""タナカホーム v3：来場前のお知らせ（予約の確認・3日前・前日）は「貴社の事務の方が予約を確認してLINEで送る」形で書く（2026-10-21 殿村さん）。

    python -m _build.patch_tanaka_v3_reminder <v3.pptx> [出力.pptx]

  予約の仕組みとLINEの自動連携は前提にしない（CLAUDE.md 2026-10-20）。DYMは文面の型をつくる・直す。
  来場前のお知らせ：お知らせの枠を中身に合わせて縮め、その下に「送り方」の枠を足す。実績の行を少し下げる
  概算お見積り：初期に「来場前のお知らせ｜型をつくる」の行を足す。月次の「予約ごとに配信」→「型の改善」
  運用内容：来場前のお知らせ「設定のみ／設定・改善」→「型の作成のみ／型の作成・改善（送信は貴社）」
  スケジュール：12月の「来場前のお知らせ」→「来場前のお知らせの型」
  出力を省くと上書き。ほかのページは触らない。
"""
import copy
import sys

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches

from _build.patch_tanaka_v3_fix import find_table, replace_in_cell
from _build.patch_yokohama_shimizu_deck_plans import set_text


def by_id(s, sid):
    return next(sh for sh in s.shapes if sh.shape_id == sid)


def new_id(s):
    return str(max(x.shape_id for x in s.shapes) + 1)


def clone(s, sid, text, x, y, w, h, name):
    el = copy.deepcopy(by_id(s, sid)._element)
    s.shapes._spTree.append(el)
    sh = s.shapes[-1]
    sh._element.find(qn("p:nvSpPr")).find(qn("p:cNvPr")).set("id", new_id(s))
    sh.name = name
    if x is not None:
        sh.left, sh.width = Inches(x), Inches(w)
    sh.top, sh.height = Inches(y), Inches(h)
    set_text(sh.text_frame, text)
    return sh


def page_reminder(s):
    assert by_id(s, 2).text_frame.text.startswith("来場前のお知らせ")
    for sid in range(16, 21):                                 # お知らせの枠：1.55 → 1.3（中身は4行まで）
        by_id(s, sid).height = Inches(1.3)
    box = clone(s, 16, "送り方｜貴社の事務の方が予約を確認し、LINEでお送りする。文面の型は弊社が用意し、反応を見て直す",
                0.55, 3.98, 9.73, 0.42, "送り方")
    box.text_frame._txBody.find(qn("a:bodyPr")).set("anchor", "ctr")
    for r in box._element.iter(qn("a:rPr")):
        r.set("sz", "1200")
    for sid in (21, 22, 23):                                  # 実績の行は送り方の枠の下へ
        by_id(s, sid).top = Inches(4.6)


def page_estimate(s):
    assert by_id(s, 339).text_frame.text.startswith("概算お見積り")
    assert by_id(s, 352).text_frame.text == "予約ごとに配信"
    set_text(by_id(s, 352).text_frame, "型の改善")
    step = by_id(s, 373).top - by_id(s, 371).top              # 初期の実施内容：最後の行の下に1行足す
    top = by_id(s, 373).top + step
    for sid, text, name in ((373, "来場前の\nお知らせ", "初期 来場前のお知らせ"), (374, "型をつくる", "初期 型をつくる")):
        sh = clone(s, sid, text, None, 0, 0, 0, name)
        sh.top, sh.height = top, by_id(s, sid).height


def page_scope(s):
    shape = find_table(s, "項目")
    row = next(r for r in shape.table.rows if r.cells[0].text == "来場前のお知らせ")
    for c, text in zip(list(row.cells)[1:], ["型の作成のみ", "型の作成・改善\n（送信は貴社）", "型の作成・改善\n（送信は貴社）"]):
        set_text(c.text_frame, text)


def main(src, out):
    prs = Presentation(src)
    slides = list(prs.slides)
    assert len(slides) == 31
    page_reminder(slides[11])
    page_estimate(slides[12])
    page_scope(slides[14])
    replace_in_cell(find_table(slides[16], "11月").table.cell(1, 1), "来場前のお知らせ", "来場前のお知らせの型")
    prs.save(out)
    print("saved:", out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else sys.argv[1])
