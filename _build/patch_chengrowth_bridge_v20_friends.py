# -*- coding: utf-8 -*-
"""チェングロウス ver1.4（殿村さんPC保存版）→ ver2.0（2026-10-05 殿村さん指示・方針変更）。
・LINEの役割＝「来訪者を友だちとしてストックし、応募・面談・来店まで育てるCRM」。CV数は主役にしない
・友だち獲得の主要施策＝離脱防止／サンクスLINE誘導。LINE Profile+ は提案から外す（7枚目を削除）
・費用＝初期構築費（既存条件 20万円）＋4月以降の月額運用パッケージ 約30万円（離脱防止・サンクスLINE込み）
・11月〜3月でステップ配信まで構築完了。友だち獲得施策の先行稼働は「状況に応じて可能」
・Chapter表紙はあっさり（中央にChapter番号＋タイトルだけ）
  python3 _build/patch_chengrowth_bridge_v20_friends.py <ver1.4.pptx>
"""
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Cm

sys.path.insert(0, str(Path(__file__).resolve().parent))
from pptx_parts import *  # noqa

ROOT = Path(__file__).resolve().parent.parent
IMG = ROOT / "_images"
OUT = ROOT / "20261005_チェングロウス_サイトリニューアルとLINE同時導入のご提案ver2.0.pptx"
prs = Presentation(sys.argv[1])
S = list(prs.slides)
PALE = "E4E8F6"


def sl(n):
    return S[n - 1]


def shp(slide, name):
    return next(sh for sh in slide.shapes if sh.name == name)


def by_text(slide, start):
    return next(sh for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.startswith(start))


def lead(slide, *lines):
    set_paras(shp(slide, "TextBox 2"), [(t, 0) for t in lines])


def title(slide, t):
    set_paras(shp(slide, "TextBox 1"), [(t, 0)])


# ================= 1 表紙 =================
set_paras(shp(sl(1), "正方形/長方形 5"), [("―サイトに来た求職者を、LINEでストックする「同時導入」プラン―", 0)])

# ================= 2 アジェンダ（1行説明だけ差し替え）=================
s = sl(2)
for old, new in [("サイト改修後も残る課題と、LINEの役割", "その場の応募で終わらせず、LINEでストックする"),
                 ("タウンワークの実績と勝因", "LINEが求人を探す入口・継続接点になる"),
                 ("導線・画面・自動応答のイメージ", "離脱防止・サンクスLINEから応募・来店まで"),
                 ("友だち獲得の試算・費用・進め方", "友だち獲得・費用・11月からの進め方")]:
    sh = next(x for x in s.shapes if x.has_text_frame and old in x.text_frame.text)
    k = [p.text for p in sh.text_frame.paragraphs].index(old)
    items = [(p.text if i != k else new, i) for i, p in enumerate(sh.text_frame.paragraphs)]
    set_paras(sh, items)

# ================= Chapter表紙（あっさり）=================
CH = {3: (1, "なぜLINEを同時導入するのか"), 9: (2, "求人業界でのLINE活用"),
      13: (3, "自動車求人Naviでの活用イメージ"), 18: (4, "効果・費用・導入スケジュール")}
for no, (k, t) in CH.items():
    s = sl(no)
    assert any(sh.has_text_frame and sh.text_frame.text == t for sh in s.shapes), (no, t)
    for sh in list(s.shapes):
        sh._element.getparent().remove(sh._element)
    label(s, 2.2, 8.6, 23.1, 0.9, [(f"Chapter {k}", 15, True, GRAY, 0)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    label(s, 2.2, 9.5, 23.1, 2.0, [(t, 32, True, NAVY, 0)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

# ================= 6 解決策（LINEの役割）=================
s = sl(6)
title(s, "解決策｜LINEで「その後」まで追える状態をつくる")
lead(s, "応募に至らない方をLINEの友だちとして残し、応募・面談・来店まで追う")
set_paras(shp(s, "Rounded Rectangle 12"), [("サイトで応募", 0)])
set_paras(shp(s, "Rounded Rectangle 15"), [("求人を配信", 0), ("→ 応募・来店", 1)])
set_paras(shp(s, "Rounded Rectangle 17"), [("② 見込み求職者が貯まる", 0), ("今すぐ応募しない方もリストに残る", 1)])

# ================= 14 連動導線（作り直し）=================
s = sl(14)
keep_header_only(s)
title(s, "新サイトとLINEの連動導線（全体フロー）")
lead(s, "LINEの友だちを増やし、その後の応募・面談・来店まで接点を持ち続ける")
steps = [("広告・\n求人媒体", GRAY, WHITE), ("Web\nサイト", GRAY, WHITE), ("離脱防止\n・サンクス", NAVY, WHITE),
         ("友だち\n追加", GREEN, WHITE), ("ステップ\n配信", PALE, NAVY), ("応募", PALE, NAVY), ("面談・\n来店", PALE, NAVY)]
W, STEP, Y0 = 3.55, 3.26, 5.6
for i, (t, f, tc) in enumerate(steps):
    a, b = t.split("\n") if "\n" in t else (t, None)
    paras = [(a, 11.5, True, tc, 0)] + ([(b, 11.5, True, tc, 0)] if b else [])
    x = 2.2 + i * STEP
    shape(s, MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON, x, Y0, W, 2.6, None, fill=f, adj=0.28)
    label(s, x + (0.25 if i == 0 else 0.6), Y0, W - (0.85 if i == 0 else 1.2), 2.6, paras, align=PP_ALIGN.CENTER,
          anchor=MSO_ANCHOR.MIDDLE)
spans = [(0, 1, GRAY, "サイトに来訪", "広告・求人媒体から"),
         (2, 3, GREEN_TX, "友だちとしてストック", "離脱防止・サンクスLINE誘導から"),
         (4, 6, NAVY, "継続して接点を持つ", "ステップ配信・求人案内から応募・面談・来店へ")]
for a, b, c, h, d in spans:
    x1 = 2.2 + a * STEP + (0.15 if a else 0)
    x2 = 2.2 + b * STEP + W - 0.25
    rect(s, x1, 8.55, x2 - x1, 0.12, GREEN if c == GREEN_TX else c)
    label(s, x1, 8.85, x2 - x1, 2.0, [(h, 14, True, c, 4), (d, 11, False, INK, 0)])
conclusion(s, 12.6, "その場で応募しなかった方も、LINEでその後の応募・面談・来店まで追える", 15)
label(s, 2.2, 16.4, 23.1, 0.6, [("※求人ボックス・Indeed経由については、着地先・応募完了先を確認したうえでLINE導線を設計", 9, False, GRAY, 0)])

# ================= 15 友だち獲得の2つの導線（作り直し）=================
s = sl(15)
keep_header_only(s)
title(s, "友だち獲得の2つの導線｜離脱防止・サンクスLINE誘導")
lead(s, "サイトを離れる方・応募した方を、LINEの友だちとして残す")
s.shapes.add_picture(str(IMG / "chengrowth_v11_popup.png"), Cm(2.2), Cm(4.6), width=Cm(11.2))
s.shapes.add_picture(str(IMG / "chengrowth_v11_talk_thanks.png"), Cm(17.75), Cm(4.6), height=Cm(7.0))
label(s, 2.2, 11.9, 11.2, 2.4, [("① 離脱防止", 15, True, NAVY, 4),
                                ("サイトを離れる方に「条件に合う新着求人をLINEで受け取る」", 12, False, INK, 0)])
label(s, 14.1, 11.9, 11.2, 2.4, [("② サンクスLINE誘導", 15, True, NAVY, 4),
                                 ("応募完了後に「今後の連絡・求人情報をLINEで受け取る」", 12, False, INK, 0)])
label(s, 2.2, 16.6, 23.1, 0.6, [("※画面はイメージです（右：応募完了後にLINEを追加した後のトーク画面）", 9, False, GRAY, 0)])

# ================= 17 自動応答（会員連携の記載を外す）=================
set_paras(shp(sl(17), "Rounded Rectangle 7"), [("あいさつ（自動）", 0), ("アンケートを案内", 1)])

# ================= 19 シミュレーション（友だち数が主役）=================
s = sl(19)
title(s, "友だち獲得シミュレーション")
lead(s, "4月の公開から半年で、友だちは約542人まで積み上がる")
ch = shp(s, "Chart 5")
ch.left, ch.top, ch.width, ch.height = Cm(10.8), Cm(4.3), Cm(14.5), Cm(5.6)
rich(s, 2.2, 4.6, 8.2, 1.0, [[("半年後の友だち（累計）", 13, True, GRAY)]], align=PP_ALIGN.CENTER)
rich(s, 2.2, 5.6, 8.2, 3.2, [[("約", 24, True, NAVY), ("542", 66, True, NAVY), ("人", 24, True, NAVY)]], align=PP_ALIGN.CENTER)
conclusion(s, 13.0, "LINEで見込み求職者を蓄積し、応募・面談・来店につなげる資産をつくる", 15)
set_paras(shp(s, "TextBox 900"), [("前提：リニューアル前に貯まる友だちは含めない／サイト来訪 月約6,100人（広告のシミュレーション確定後に差し替え）", 0)])

# ================= 20 費用の考え方（LINEの費用の中身を新方針に）=================
s = sl(20)
set_paras(shp(s, "Chevron 8"), [("応募・来店", 0)])
set_paras(shp(s, "TextBox 14"), [("LINE独自構築費", 0), ("担い手：弊社（本提案）", 1), ("来訪者を友だちとして蓄積し、応募・面談・来店までつなげる", 2)])
set_paras(shp(s, "Rectangle 17"), [("LINE独自構築費のかかり方", 0), ("初期構築費：11月頃〜　｜　月額運用：4月〜", 1),
                                    ("離脱防止・サンクスLINE誘導は月額に含む", 2), ("LINE公式アカウント利用料：別途", 3)])
set_paras(shp(s, "TextBox 19"), [("LINEの費用を分けることで、友だちの増え方とその後の効果を単独で把握できる", 0)])

# ================= 21 概算お見積もり（作り直し）=================
s = sl(21)
keep_header_only(s)
title(s, "概算お見積もり")
lead(s, "11月頃からの初期構築費と、4月以降の月額運用パッケージ")
rich(s, 2.2, 4.5, 11.4, 1.0, [[("初期構築費", 15, True, GRAY), ("（11月頃〜3月）", 12, False, GRAY)]], align=PP_ALIGN.CENTER)
rich(s, 2.2, 5.5, 11.4, 2.6, [[("20", 54, True, NAVY), ("万円", 24, True, NAVY)]], align=PP_ALIGN.CENTER)
vline(s, 13.75, 4.7, 13.4, LGRAY, 1.0)
rich(s, 13.9, 4.5, 11.4, 1.0, [[("月額運用パッケージ", 15, True, GREEN_TX), ("（4月以降）", 12, False, GREEN_TX)]], align=PP_ALIGN.CENTER)
rich(s, 13.9, 5.5, 11.4, 2.6, [[("約", 24, True, NAVY), ("30", 54, True, NAVY), ("万円／月", 24, True, NAVY)]], align=PP_ALIGN.CENTER)
inc1 = ["アカウント開設・設計", "リッチメニュー・あいさつ・ステップ配信", "セグメント・計測・友だち獲得導線の設計", "各種設定・動作確認"]
inc2 = ["LINE公式アカウント運用・配信", "ステップ配信・リッチメニュー運用", "離脱防止・サンクスLINE誘導", "分析・改善・毎月のレポート"]
for x, head, items in [(2.9, "含むもの", inc1), (14.6, "含むもの（主要施策込み）", inc2)]:
    label(s, x, 8.5, 10.4, 4.8, [(head, 11, True, GRAY, 4)] + [("・" + t, 12, False, INK, 2) for t in items])
hline(s, 2.2, 25.3, 13.9, LGRAY, 0.75)
label(s, 2.2, 14.1, 23.1, 0.8, [("＋ LINE公式アカウント（開設3万円〜／利用料 月額0.5万円〜・配信通数で変動）", 12, False, INK, 0)],
      align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
label(s, 2.2, 16.4, 23.1, 0.6, [("体制：弊社＝設計・運用・制作・毎月のレポート／貴社＝求人情報・面談体制の共有、サイト制作会社様との連携", 9.5, False, GRAY, 0)])

# ================= 22 なぜ11月頃から =================
s = sl(22)
set_paras(shp(s, "TextBox 12"), [("公開前に", 0), ("構築・設定・検証", 1)])
set_paras(shp(s, "TextBox 13"), [("ステップ配信まで完成させておく", 0)])

# ================= 23 スケジュール（作り直し）=================
s = sl(23)
keep_header_only(s)
title(s, "スケジュール｜11月頃から初期構築、4月から本格運用")
lead(s, "11月〜3月でステップ配信まで構築し、4月は構築済みのLINEで本格運用")
TX, MW, EX = 2.2, 2.5, 25.3
GX = TX + 5 * MW
shape(s, MSO_SHAPE.PENTAGON, TX, 4.4, GX - TX + 0.3, 0.85, [("11月〜3月｜設計・構築", 12.5, True, WHITE, 0)], fill=NAVY, adj=0.5)
shape(s, MSO_SHAPE.CHEVRON, GX - 0.1, 4.4, EX - GX + 0.1, 0.85, [("4月〜｜運用・改善", 12.5, True, WHITE, 0)], fill=GREEN, adj=0.5,
      margins=(0.6, 0.05, 0.3, 0.05))
for i, m in enumerate(["11月", "12月", "1月", "2月", "3月"]):
    label(s, TX + i * MW, 5.35, MW, 0.6, [(m, 11.5, True, INK, 0)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
label(s, GX, 5.35, EX - GX, 0.6, [("4月〜", 11.5, True, GREEN_TX, 0)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
RH, GAP = 0.95, 0.12


def bar(x1, x2, yy, name, detail, fill, line, tc, dash=False):
    r = rect(s, x1, yy, x2 - x1, RH, fill, line=line, lw=1.0)
    if dash:
        r.line.dash_style = 4
    rich(s, x1 + 0.25, yy, x2 - x1 - 0.4, RH, [[(name + "　", 12, True, tc), (detail, 10.5, False, tc)]])


rows1 = [("設計", "アカウント・リッチメニュー・あいさつ・セグメント・計測", NAVY, None, WHITE),
         ("構築", "ステップ配信の制作・配信条件・自動応答の設定", NAVY, None, WHITE),
         ("検証", "動作テスト・公開前の確認", PALE, NAVY, NAVY)]
rows2 = [("獲得・配信", "ステップ／企画配信・求人案内", GREEN, None, WHITE),
         ("追客", "応募への追客・面談／来店のフォロー", GREEN_BG, GREEN, GREEN_TX),
         ("分析・改善", "KPI計測・月次分析・改善", GREEN_BG, GREEN, GREEN_TX)]
y = 6.1
bottom = y + 4 * (RH + GAP) + 1.1 + 3 * (RH + GAP)
for i in range(1, 5):
    vline(s, TX + i * MW, 5.45, bottom, "E7E7E7", 0.5)
for name, d, f, l, tc in rows1:
    bar(TX + 0.05, GX - 0.05, y, name, d, f, l, tc)
    y += RH + GAP
bar(TX + 0.05, GX - 0.05, y, "先行稼働", "状況に応じて、友だち獲得施策の一部を先に稼働", WHITE, GRAY, GRAY, dash=True)
y += RH + GAP
my = y
shape(s, MSO_SHAPE.DIAMOND, GX - 0.45, my + 0.1, 0.9, 0.9, fill=GREEN)
arrow_line(s, TX + 0.4, my + 0.55, GX - 0.6, my + 0.55, NAVY, 1.5, dash=True)
rich(s, GX + 0.6, my, EX - GX - 0.6, 1.1, [[("サイト公開　／　構築済みのLINEで本格運用", 13, True, GREEN_TX)]])
y = my + 1.1
for name, d, f, l, tc in rows2:
    bar(GX + 0.05, EX - 0.05, y, name, d, f, l, tc)
    y += RH + GAP
vline(s, GX, 5.35, bottom, GREEN, 2.5)
conclusion(s, bottom + 0.25, "初期構築で設計した内容を、そのまま運用へ", 16)
label(s, 2.2, bottom + 1.75, 23.1, 0.6, [("※先行稼働はサイト状況・実装状況に応じて判断（先行稼働分はシミュレーションに含めない）", 8.5, False, GRAY, 0)])

# ================= 7 LINE Profile+ のページを削除 =================
drop_slide(prs, sl(7))

prs.save(str(OUT))
print("saved:", OUT.name, len(prs.slides), "枚")
