# -*- coding: utf-8 -*-
"""株式会社セキュリティー京葉 LINE採用追客のご提案（簡易6枚・4:3）。2026-10-13。
面接に来なかった方・辞退した方をLINE（Lステップ）でリスト化し、ステップ配信で追う。月額30万円で提案（殿村さん）。
器はDYMの4:3マスターの既存PPTX（サンクスLINE Appendix ver1.2）。2枚目にリード行と区切り線を3枚目から写して6枚にする。
  python3 _build/build_keiyo_6p.py <Appendix ver1.2.pptx> [出力.pptx]
"""
import copy
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Cm

sys.path.insert(0, str(Path(__file__).resolve().parent))
from pptx_parts import (GRAY, INK, NAVY, WHITE, arrow_line, find, keep_header_only, rect,  # noqa: E402
                        rich, set_paras, shape)

ROOT = Path(__file__).resolve().parent.parent
SRC = Path(sys.argv[1])
OUT = Path(sys.argv[2]) if len(sys.argv) > 2 else ROOT / "20261013_セキュリティー京葉御中_LINE採用追客のご提案.pptx"

SLIDES = [
    ("課題｜一度外れた応募者が、そのまま離れている", "警備を含む保安職業の求人倍率は6.41倍。集めた応募者を、選考の途中で手放さない"),
    ("方針｜外れた方に通知メッセージで届け、追い続ける", "面接に来なかった方・辞退した方に、通知メッセージでLINEを届け、ステップ配信で追う"),
    ("ステップ配信｜働き方を少しずつ届ける", "給与・働き方・現場の様子を順番に届け、もう一度応募したくなる状態をつくる"),
    ("Lステップ｜外れた理由ごとに追い分ける", "LINE公式アカウントにLステップを入れ、リスト化・出し分け・面接予約までLINEで完結"),
    ("進め方｜アカウント開設から配信開始まで", "LINE公式アカウントの開設から、ステップ配信の開始までまとめて支援"),
    ("費用", "月額30万円で、Lステップの運用から配信の改善まで支援"),
]

X, W = 1.79, 23.97
Y0 = 4.5
S = 13


def text(s, y, runs, h=0.8, x=X, w=W):
    return rich(s, x, y, w, h, runs, anchor=MSO_ANCHOR.TOP)


def head(s, y, txt):
    rect(s, X, y + 0.14, 0.15, 0.5, NAVY)
    text(s, y, [[(txt, 13.5, True, NAVY)]], x=X + 0.35, w=W - 0.35)
    return y + 0.85


def rows(s, y, items, kw=4.6, h=0.72):
    for i, (k, v) in enumerate(items):
        text(s, y + i * h, [[(k, S, True, INK)]], h=h, x=X + 0.35, w=kw)
        text(s, y + i * h, [[(v, S, False, INK)]], h=h, x=X + 0.35 + kw, w=W - 0.35 - kw)
    return y + len(items) * h


def rows3(s, y, items, w1=3.2, w2=5.2, h=0.72):
    for i, (a, b, c) in enumerate(items):
        text(s, y + i * h, [[(a, S, True, INK)]], h=h, x=X + 0.35, w=w1)
        text(s, y + i * h, [[(b, S, False, INK)]], h=h, x=X + 0.35 + w1, w=w2)
        text(s, y + i * h, [[(c, 12, False, GRAY)]], h=h, x=X + 0.35 + w1 + w2, w=W - 0.35 - w1 - w2)
    return y + len(items) * h


def flow(s, y, items, h=1.35, gap=0.6, size=11.5):
    x0, ww = X + 0.35, W - 0.35
    bw = (ww - gap * (len(items) - 1)) / len(items)
    for i, t in enumerate(items):
        last = i == len(items) - 1
        x = x0 + i * (bw + gap)
        paras = [(ln, size, last, WHITE if last else INK, 0) for ln in t.split("\n")]
        shape(s, MSO_SHAPE.RECTANGLE, x, y, bw, h, paras, fill=NAVY if last else WHITE, line=NAVY, lw=1.0,
              margins=(0.08, 0.05, 0.08, 0.05))
        if not last:
            arrow_line(s, x + bw + 0.08, y + h / 2, x + bw + gap - 0.08, y + h / 2, NAVY, w=1.5)
    return y + h


prs = Presentation(str(SRC))
slides = list(prs.slides)
HDR = {"TextBox 1", "TextBox 2", "Connector 3"}
donor = next(s for s in slides if HDR <= {sh.name for sh in s.shapes})
# タイトルだけあるページに、リード行と区切り線を写して使えるようにする
for s in slides:
    names = {sh.name for sh in s.shapes}
    if "TextBox 1" in names and not HDR <= names:
        for nm in ("TextBox 2", "Connector 3"):
            if nm not in names:
                s.shapes._spTree.append(copy.deepcopy(find(donor, nm)._element))
keep = [i for i, s in enumerate(prs.slides) if HDR <= {sh.name for sh in s.shapes}][:6]
assert len(keep) == 6, keep
lst = prs.slides._sldIdLst
for i, el in reversed(list(enumerate(list(lst)))):
    if i not in keep:
        prs.part.drop_rel(el.rId)
        lst.remove(el)

for s, (title, lead) in zip(prs.slides, SLIDES):
    keep_header_only(s)
    t1, t2 = find(s, "TextBox 1"), find(s, "TextBox 2")
    t1.left, t1.top, t1.width, t1.height = Cm(1.52), Cm(0.38), Cm(18.1), Cm(0.94)
    t2.left, t2.top, t2.width, t2.height = Cm(1.79), Cm(2.27), Cm(25.31), Cm(0.96)
    set_paras(t1, [(title, 0)])
    set_paras(t2, [(lead, 0)])
    for t in (t1, t2):
        for r in t.text_frame._txBody.iter(qn("a:rPr")):
            r.set("sz", "1600")

s1, s2, s3, s4, s5, s6 = prs.slides

# 1. 課題
y = head(s1, Y0, "警備員の採用環境")
y = rows(s1, y, [("有効求人倍率", "保安職業 6.41倍（職業全体は 1.10倍）")])
text(s1, y + 0.05, [[("出典：労働政策研究・研修機構「職業別にみた有効求人倍率、平均求人賃金」（2025年度）", 10, False, GRAY)]],
     h=0.5, x=X + 0.35)
y = head(s1, y + 0.9, "選考の途中で起きていること")
y = flow(s1, y + 0.1, ["応募", "面接のご案内", "面接に来ない・\n辞退", "連絡が途切れ、\nそのまま離れる"])
y = head(s1, y + 0.6, "外れた方について")
rows(s1, y, [("共通点", "一度は応募している＝警備の仕事に関心がある方"),
             ("いまの対応", "電話・メールでの再連絡は手間がかかり、続けにくい")])

# 2. 方針
y = head(s2, Y0, "流れ")
y = flow(s2, y + 0.1, ["面接に来ない・\n辞退", "通知メッセージで\nLINEに届く", "友だち追加・\nリスト化", "ステップ配信で\n声をかけ続ける", "再応募・面接"])
y = head(s2, y + 0.6, "ポイント")
rows(s2, y, [("対象", "面接に来なかった方・辞退した方"),
             ("届け方", "通知メッセージ：友だちでなくても、応募時の電話番号宛てにLINEで届く"),
             ("手間", "配信は自動。担当者が1人ずつ連絡しなくてよい"),
             ("狙い", "中途・アルバイトで働ける方を、タイミングが合ったときに採用につなげる")])

# 3. ステップ配信（内容は採用ページの条件から）
y = head(s3, Y0, "配信の例")
y = rows3(s3, y, [
    ("登録直後", "お礼・働き方の紹介", "「ご応募ありがとうございました。改めて働き方をご紹介します」"),
    ("3日後", "給与・条件", "「1現場10,500円（2現場で19,500円）。直行直帰OK・夜勤なし」"),
    ("1週間後", "働く方の声", "「平均年齢64歳。84歳の方も活躍しています」"),
    ("2週間後", "短期・スポット勤務", "「年末年始だけ、などのスポット勤務もあります」"),
    ("1か月後", "現場・面接のご案内", "「柏・松戸・熊谷で現場があります。面接日はこちら」"),
    ("以降 月1回", "募集のお知らせ", "「今月の募集をお知らせします」"),
])
y = head(s3, y + 0.5, "そのほかに伝えられること")
rows(s3, y, [("待遇", "ガソリン代・駐車場代の全額支給／空調服・ネッククーラー支給"),
             ("安心材料", "資格取得の支援／月15日以上の勤務で社会保険に加入可能")])

# 4. Lステップ
y = head(s4, Y0, "できること")
rows(s4, y, [("リスト化", "「面接に来なかった」「辞退した」をタグで分けて管理"),
             ("出し分け", "タグごとに、別々のステップ配信を送る"),
             ("希望の聞き取り", "希望の勤務地・日数・働き方（中途・アルバイト・スポット）をLINEのフォームで聞く"),
             ("面接予約", "LINEの中で面接日を選んで予約できる"),
             ("反応の確認", "誰がどのリンクを押したかを確認し、次の連絡につなげる")])

# 5. 進め方
y = head(s5, Y0, "初期構築の流れ")
y = flow(s5, y + 0.1, ["LINE公式\nアカウント開設", "Lステップの\n導入・設定", "外れた方へ\n通知メッセージ", "ステップ配信\nの作成", "配信開始・\n改善"])
y = head(s5, y + 0.6, "友だちになってもらう方法")
rows(s5, y, [("これまでの方", "通知メッセージ（応募時の電話番号宛て）で届け、友だち追加につなげる"),
             ("これからの方", "応募の受付・面接の連絡でLINEをご案内し、外れても追える状態にする"),
             ("送る内容", "通知メッセージは日程の再調整・辞退の受付などのお知らせに使う"),
             ("募集の案内", "友だち追加の後、ステップ配信で送る")])

# 6. 費用
y = head(s6, Y0, "内訳")
y = rows(s6, y, [("初期構築", "10万円"),
                 ("月額", "30万円"),
                 ("初期構築の内容", "LINE公式アカウント開設・Lステップと通知メッセージの設定・配信の作成")])
y = head(s6, y + 0.5, "月額に含む支援")
y = rows(s6, y, [("配信", "通知メッセージ・ステップ配信・募集のお知らせの作成と運用"),
                 ("リスト管理", "外れた方のタグ付け・出し分けの管理"),
                 ("改善", "反応を見て、配信の内容・タイミングを見直す")])
text(s6, y + 0.3, [[("※Lステップ・LINE公式アカウント・通知メッセージの利用料は別途（実費）", 10.5, False, GRAY)]], h=0.6, x=X + 0.35)

# 字体をメイリオに統一（latin → ea → cs の順）
AFTER = {qn(x) for x in ("a:sym", "a:hlinkClick", "a:hlinkMouseOver", "a:rtl", "a:extLst")}
for s in prs.slides:
    for r in s.shapes._spTree.iter(qn("a:r")):
        if r.find(qn("a:rPr")) is None:
            r.insert(0, r.makeelement(qn("a:rPr"), {"lang": "ja-JP"}))
    for tag in ("a:rPr", "a:endParaRPr", "a:defRPr"):
        for rpr in s.shapes._spTree.iter(qn(tag)):
            for ft in ("a:latin", "a:ea", "a:cs"):
                for e in rpr.findall(qn(ft)):
                    rpr.remove(e)
            nxt = next((c for c in rpr if c.tag in AFTER), None)
            for ft in ("a:latin", "a:ea", "a:cs"):
                e = rpr.makeelement(qn(ft), {"typeface": "メイリオ"})
                if nxt is not None:
                    nxt.addprevious(e)
                else:
                    rpr.append(e)

prs.save(str(OUT))
print("saved:", OUT.name, len(prs.slides), "枚")
