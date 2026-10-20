# -*- coding: utf-8 -*-
"""すなお ver4.9 → ver5.0：事例紹介のページに、LINEヤフー公式の事例（保険・イベント）を入れる（2026-10-20 殿村さん「それで進んで」）。

    python -m _build.patch_sunao_v50_jirei <ver4.9.pptx> [出力.pptx]

  殿村さん「LINEヤフー公式の事例から」「保険・イベント系で」「保険相談窓口はいらない」。
  ・上段「生命保険会社」：第一生命（同じ内容をメールとLINEで配信 → 資料請求のCV率がLINEはメールの約8倍。2024年4月公開）／
    日本生命（営業職員経由で友だちになった方のブロック率 5.2%。自然流入は74.9%。2025年2月公開）
  ・下段「教室・スクール（見学会などのイベント集客）」：JPCスポーツ教室（見学会の案内をLINEで配信し、参加者の約7割が入会。2021年8月公開）
  ・社名は書かず業種名だけ。弊社の実績と読まれないよう「LINEヤフー for Business 掲載の事例」と注記する。
  ・数字の札はFMTの札（紺の塗り・白い文字）を写し直して使う（ver4では「--」の1段落にしていたため）。
  出典URL：https://www.lycbiz.com/jp/case-study/line-official-account/daiichiseimei/ ／ …/nissay/ ／ …/jpcsports/
"""
import copy
import sys

from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from _build.arauchi_v4.lib import add_box
from _build.arauchi_v10.build import runs
from _build.build_sunao_v4 import FMT, GRAY
from _build.patch_sunao_v4_colors import fix_slide

CARDS = {10: ("約8倍", ["資料請求の", "CV率", "（メール比）"]),
         12: ("5.2%", ["営業職員経由", "の友だちの", "ブロック率"]),
         16: ("約7割", ["見学会の参加者の", "入会率"])}


def by_id(s, sid):
    return next(sh for sh in s.shapes if sh.shape_id == sid)


def set_text(sh, text):
    p = sh.text_frame.paragraphs[0]
    p.runs[0].text = text
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)


def card(dst, src, number, labels):
    """FMTの札を写し直し、1段落目＝大きな数字、2段落目＝説明（改行は a:br）にする"""
    el = copy.deepcopy(src._element)
    old = by_id(dst, src.shape_id)._element
    old.getparent().replace(old, el)
    sh = by_id(dst, src.shape_id)
    p0, p1 = sh.text_frame.paragraphs[:2]
    p0.runs[0].text = number
    for r in p0.runs[1:]:
        r._r.getparent().remove(r._r)
    for r in p0.runs:
        r.font.size = Pt(18)
    rpr = copy.deepcopy(p1.runs[0]._r.find(qn("a:rPr")))
    for e in list(p1._p):
        if e.tag in (qn("a:r"), qn("a:br"), qn("a:fld")):
            p1._p.remove(e)
    end = p1._p.find(qn("a:endParaRPr"))
    for i, text in enumerate(labels):
        if i:
            br = etree.Element(qn("a:br"))
            br.append(copy.deepcopy(rpr))
            (end.addprevious(br) if end is not None else p1._p.append(br))
        r = etree.Element(qn("a:r"))
        r.append(copy.deepcopy(rpr))
        etree.SubElement(r, qn("a:t")).text = text
        (end.addprevious(r) if end is not None else p1._p.append(r))
    for r in p1._p.iter(qn("a:rPr")):
        r.set("sz", "1000")
    return sh


def main(src, out):
    prs = Presentation(src)
    s = next(x for x in prs.slides if any(sh.has_text_frame and sh.text_frame.text.strip() == "事例紹介" for sh in x.shapes)
             and any(sh.shape_id == 10 and sh.has_text_frame and sh.text_frame.text.strip() == "--" for sh in x.shapes))
    fmt = Presentation(FMT).slides[61]
    set_text(by_id(s, 3), "保険・イベント集客におけるLINE公式アカウントの活用事例")
    set_text(by_id(s, 8), "生命保険会社")
    set_text(by_id(s, 14), "教室・スクール（見学会などのイベント集客）")
    for sid, (number, labels) in CARDS.items():
        card(s, by_id(fmt, sid), number, labels)
    by_id(s, 16).left, by_id(s, 16).width = Inches(1.0), Inches(1.6)       # 上段の札と同じ位置・幅に
    for sid in (18, 20):                                                    # 下段の残りの札は使わない
        e = by_id(s, sid)._element
        e.getparent().remove(e)
    add_box(s, 0.4, 5.7, 9.9, 0.45,
            [runs(("※LINEヤフー for Business 掲載の事例（数値は各社調べ）。生命保険会社＝2024年4月・2025年2月公開、教室＝2021年8月公開。"
                   "5.2%は営業職員経由の友だち（自然に友だちになった方は74.9%）", dict(size=9, color=GRAY)))],
            anchor="t", ins=(0, 0, 0, 0), name="注記")
    print("colors:", fix_slide(s))
    prs.save(out)
    print("saved:", out, len(prs.slides), "枚")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else sys.argv[1])
