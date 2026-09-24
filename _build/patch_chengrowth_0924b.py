# -*- coding: utf-8 -*-
"""チェングロウスver1.1 殿村さんフィードバック反映（2026-09-24 夕）。PC保存版に1回だけ当てる。
- P6  前後検索：画像を大きく＋職業系クエリをオレンジ強調した版に差し替え、前/0-3日/後のゾーン帯を追加
- P7  検索トレンド：文字を大きくした画像に差し替え（3月=山/12月=谷を帯表示）、全幅に
- P11 ②③の本文が途中で切れていた不具合を修復（段落末尾タグの順序）
- P16 費用感：やる気スイッチver2_3 P20 の型（初期/月次/実施内容/6ヶ月後）を移植してSIM①の数字で埋める
- P20 14日間ステップ：同 P23 の型（Day/対象/配信内容の表）を移植
- P21 年間企画：同 P25 の型（時期/企画/内容の表）を移植
  python3 _build/patch_chengrowth_0924b.py <PC保存版pptx> <出力pptx>
"""
import copy
import sys
from pathlib import Path
from pptx import Presentation
from pptx.util import Cm, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
SRC, OUT = sys.argv[1], sys.argv[2]
YARUKI = ROOT / "202609_株式会社やる気スイッチグループ御中_LINE公式アカウント運用のご提案_ver2_3.pptx"
IMG_ZENGO = ROOT / "_images/chengrowth_zengo_colored.png"
IMG_TREND = ROOT / "_images/chengrowth_trend_seibishi_5y_L.png"

prs = Presentation(SRC)
yk = Presentation(str(YARUKI))
SHAPE_TAGS = ("sp", "cxnSp", "pic", "graphicFrame", "grpSp")


# ---------------- helpers ----------------
def by_id(slide, sid):
    for sh in slide.shapes:
        if sh.shape_id == sid:
            return sh
    raise KeyError(sid)


def meiryo(run):
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        e = rPr.find(qn(tag))
        if e is None:
            e = rPr.makeelement(qn(tag), {})
            rPr.append(e)
        e.set("typeface", "メイリオ")


def set_text(tf, lines, tmpl_para=0):
    """段落・runの書式を保ったまま全行差し替え。lines は str か (str, size_pt, bold, color)"""
    txBody = tf._txBody
    ps = txBody.findall(qn("a:p"))
    tp = ps[min(tmpl_para, len(ps) - 1)]
    tr = tp.find(qn("a:r"))
    p_tmpl = copy.deepcopy(tp)
    for r in p_tmpl.findall(qn("a:r")) + p_tmpl.findall(qn("a:br")) + p_tmpl.findall(qn("a:fld")):
        p_tmpl.remove(r)
    r_tmpl = copy.deepcopy(tr) if tr is not None else None
    for p in ps:
        txBody.remove(p)
    for ln in lines:
        text, size, bold, color = (ln, None, None, None) if isinstance(ln, str) else (list(ln) + [None] * 4)[:4]
        p = copy.deepcopy(p_tmpl)
        end = p.find(qn("a:endParaRPr"))
        if r_tmpl is not None:
            r = copy.deepcopy(r_tmpl)
        else:
            r = p.makeelement(qn("a:r"), {})
            r.append(r.makeelement(qn("a:rPr"), {"lang": "ja-JP"}))
            r.append(r.makeelement(qn("a:t"), {}))
        r.find(qn("a:t")).text = text
        rPr = r.find(qn("a:rPr"))
        if size is not None:
            rPr.set("sz", str(int(size * 100)))
        if bold is not None:
            rPr.set("b", "1" if bold else "0")
        if color is not None:
            for f in rPr.findall(qn("a:solidFill")):
                rPr.remove(f)
            sf = rPr.makeelement(qn("a:solidFill"), {})
            c = sf.makeelement(qn("a:srgbClr"), {"val": color})
            sf.append(c)
            rPr.insert(0, sf)
        if end is not None:
            end.addprevious(r)
        else:
            p.append(r)
        txBody.append(p)


def set_cell(cell, text):
    set_text(cell.text_frame, [text])


def transplant(dst, src, drop_pics=True):
    """dst の図形を全部消して src の図形を複製（画像は除く）"""
    tree = dst.shapes._spTree
    for el in list(tree):
        if el.tag.split("}")[-1] in SHAPE_TAGS:
            tree.remove(el)
    for el in src.shapes._spTree:
        tag = el.tag.split("}")[-1]
        if tag not in SHAPE_TAGS:
            continue
        if drop_pics and tag == "pic":
            continue
        tree.append(copy.deepcopy(el))


def textbox(slide, x, y, w, h, lines, size=11, color="333333", bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Cm(0.1)
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        text, sz, b, c = (ln, size, bold, color) if isinstance(ln, str) else ln
        r = p.add_run()
        r.text = text
        r.font.size = Pt(sz)
        r.font.bold = b
        r.font.color.rgb = RGBColor.from_string(c)
        meiryo(r)
    return tb


def bar(slide, x0, x1, y, h, fill, text):
    sp = slide.shapes.add_shape(MSO_SHAPE.PENTAGON, Cm(x0), Cm(y), Cm(x1 - x0), Cm(h))
    sp.fill.solid()
    sp.fill.fore_color.rgb = RGBColor.from_string(fill)
    sp.line.fill.background()
    sp.shadow.inherit = False
    tf = sp.text_frame
    tf.margin_left = tf.margin_right = Cm(0.1)
    tf.margin_top = tf.margin_bottom = Cm(0)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string("FFFFFF")
    meiryo(r)


# ================= P6 前後検索 =================
s = prs.slides[5]
old = by_id(s, 5)
old._element.getparent().remove(old._element)
L, T, W = 1.2, 4.15, 17.3
H = W * 1167 / 2000
s.shapes.add_picture(str(IMG_ZENGO), Cm(L), Cm(T), width=Cm(W))
px = lambda p: L + p / 2000 * W
yb = T + H + 0.12
bar(s, px(80), px(1033), yb, 0.72, "8C8C8C", "サイト訪問時：離脱防止バナーで受け止め")
bar(s, px(1033), px(1507), yb, 0.72, "06C755", "LINE友だち追加")
bar(s, px(1507), px(1975), yb, 0.72, "00873C", "LINE施策（配信）")
textbox(s, px(1330), T + 0.25, px(1990) - px(1330), 1.2,
        [("■ 職業・転職系", 11, True, "ED7D31"), ("■ 車種・店舗・ローン", 11, True, "5CBFAF")])
CARDS = [
    ("① 職業クエリは0日周辺に実在", ["「年収」「受付嬢」「営業マン」", "「ホワイト企業」等が起点前後に出現"]),
    ("② 大半は「買う人」の検索", ["車種名・店舗名・ローンが支配的。", "働きたい人はその中に埋もれている"]),
    ("③ 検索広告では狙い撃てない", ["サイトに来た瞬間に受け止める", "＝離脱防止とLINEの役割"]),
]
for (sid, top), (head, body) in zip(((6, 4.15), (7, 7.6), (8, 11.05)), CARDS):
    c = by_id(s, sid)
    c.left, c.top, c.width, c.height = Cm(18.8), Cm(top), Cm(7.52), Cm(3.3)
    tf = c.text_frame
    set_text(tf, [(head, 12, True, None)] + [(b, 10.5, False, None) for b in body], tmpl_para=1)
    tf.paragraphs[0].runs[0].font.color.rgb = RGBColor.from_string("1F285A")
band = by_id(s, 9)
band.top = Cm(yb + 0.95)
band.height = Cm(1.0)

# ================= P7 検索トレンド =================
s = prs.slides[6]
old = by_id(s, 5)
old._element.getparent().remove(old._element)
W7 = 25.12
H7 = W7 * 780 / 2400
s.shapes.add_picture(str(IMG_TREND), Cm(1.2), Cm(4.1), width=Cm(W7))
dy = (4.1 + H7 + 0.25) - 11.3
for sid in range(6, 18):
    sh = by_id(s, sid)
    sh.top = sh.top + Cm(dy)
b = by_id(s, 18)
b.top = b.top + Cm(dy)

# ================= P11 本文切れの修復 =================
# 前回の手改行外しで後半の文字が不正位置に入り、PC保存時に消えていた → 元の文章を入れ直す
P11 = {29: "登録後の自動ステップ配信と、セグメント別の企画配信で、開封〜CVまでの導線を最適化する。",
       44: "よくある質問への自動応答と、必要な場面での個別チャット対応でユーザー体験を向上させる。"}
for sid, t in P11.items():
    set_text(by_id(prs.slides[10], sid).text_frame, [t])

# ================= P16 費用感（やる気P20の型） =================
s = prs.slides[15]
foot = copy.deepcopy(by_id(s, 15)._element)       # SIM前提の脚注は残す
transplant(s, yk.slides[19])
s.shapes._spTree.append(foot)
set_text(by_id(s, 45).text_frame, ["想定の費用対効果"])
set_text(by_id(s, 40).text_frame, [
    "Meta広告の目的を「応募獲得」から「友だち獲得」へ切り替え（広告費は据え置き）、",
    "初期24.5万円・月額10.5万円〜の運用で、面談・応募を半年196件と見込みます。"])
set_text(by_id(s, 8).text_frame, ["¥245,000", ("構築20万＋離脱防止1.5万＋アカウント3万", 9, False, None)])
set_text(by_id(s, 10).text_frame, [("コンサル費 ¥70,000", 12, None, None), ("＋ 離脱防止 ¥30,000", 12, None, None), ("＋ アカウント費 ¥5,000〜", 12, None, None)])
LEFT = {12: "あいさつ＋診断", 14: "1回（4問）", 18: "リッチメニュー", 22: "3タブ×6枠",
        26: "ステップ", 30: "構築×10", 34: "離脱防止", 38: "ポップアップ設置"}
RIGHT = {25: "月次投稿", 29: "4回", 50: "月次定例会", 54: "1回",
         66: "ステップ改善", 70: "ABテスト随時"}
for sid, t in {**LEFT, **RIGHT}.items():
    set_text(by_id(s, sid).text_frame, [t])
set_text(by_id(s, 80).text_frame, ["友だち数"])
set_text(by_id(s, 11).text_frame, ["約4,650人"])
set_text(by_id(s, 82).text_frame, ["有効友だち"])
set_text(by_id(s, 16).text_frame, ["約2,500人"])
set_text(by_id(s, 84).text_frame, ["流入"])
set_text(by_id(s, 23).text_frame, ["2,560件/月"])
set_text(by_id(s, 27).text_frame, ["37件/月"])
note = by_id(s, 31)
note.left, note.top, note.width, note.height = Cm(18.74), Cm(10.85), Cm(8.06), Cm(3.0)
set_text(note.text_frame, [
    ("半年累計196件・累計CPA 11,036円", 11, True, None),
    ("現状CPA 6万・目標2〜2.5万を大きく下回る", 9.5, False, "333333"),
    ("参考：CPFなし＝半年42件・CPA 20,833円", 9, False, "7F7F7F")])
for sid in (32, 36):
    e = by_id(s, sid)._element
    e.getparent().remove(e)
textbox(s, 1.06, 15.75, 25.4, 0.95, [
    "※ 別途 CPF広告費 月20.8万円は、既存Meta予算（年250万円）からの転用＝追加出費ゼロ。",
    "※ アカウント費は配信数に応じて自動切替（5,000円→3ヶ月目以降15,000円）。"], size=10, color="595959")

# ================= P20 14日間ステップ（やる気P23の型） =================
s = prs.slides[19]
transplant(s, yk.slides[22])
set_text(by_id(s, 45).text_frame, ["実施施策【14日間ステップ配信】"])
set_text(by_id(s, 6).text_frame, [
    "友だち追加後14日間で、診断結果をもとに内容を出し分けながら面談へのご案内を行う",
    "（売り込みは2回だけ。未予約者には再アプローチ、面談予約者は別シナリオへ分岐）"])
STEP = [
    ("0日後", "全員", "あいさつ＋30秒診断（職種・資格・転職理由・温度感の4問）"),
    ("1日後", "診断未回答者", "診断リマインド（「電話はかかってきません」を再明記）"),
    ("3日後", "診断回答者", "資格別の年収相場（2級整備士でも100万円以上の幅）"),
    ("5日後", "同上", "資格別キャリアマップ（3級→2級→検査員で年収がどう変わるか）"),
    ("7日後", "未予約者", "面談のご案内①「履歴書、まだ要りません」夜の面談OK"),
    ("10日後", "未予約者", "転職体験レポート（同じ職種・同じ地域の実例）"),
    ("12日後", "未予約者", "地域別の求人ピックアップ（オートバックスG特集）"),
    ("14日後", "未予約者", "面談のご案内②＋以後は月数回の定常配信へ切替"),
]
tbl = by_id(s, 46).table
for i, row in enumerate(STEP, start=1):
    for j, v in enumerate(row):
        set_cell(tbl.cell(i, j), v)
set_text(by_id(s, 47).text_frame, [
    "面談・応募 196件（6ヶ月累計・SIM①）＝ 累計CPA 11,036円として算出",
    "面談予約で止めず、就業までLINEで追いきります",
    "（面談前日リマインド→当日案内→面談後フォロー→選考状況の連絡をLINEに一本化）"], tmpl_para=0)
set_text(by_id(s, 48).text_frame, ["出典：コンテンツ案v1（別紙・配信文面の全文）／前後検索KW「カーディーラー」／Googleトレンド「整備士」"])

# ================= P21 年間企画（やる気P25の型） =================
s = prs.slides[20]
transplant(s, yk.slides[24])
set_text(by_id(s, 45).text_frame, ["想定企画【年間】"])
set_text(by_id(s, 8).text_frame, ["年間を通じて、友だち増加・面談獲得・ブロック防止を狙う企画投稿を月別に配置する（月4本）"])
PLAN = [
    ("1-2月", "「冬のうちに相談だけ」訴求", "3月の山に向けて前倒し。在職中でも相談OKを打ち出す"),
    ("3月", "年度替わり求人特集", "4月入社求人を地域別に配信（3月＝検索の山）"),
    ("3-4月", "資格取得者向け特集", "2級・検査員の合格者へ資格別の年収相場を配信"),
    ("5月", "連休明けの見直し訴求", "GW明けの「このままでいいか」層を面談へ"),
    ("6月", "LINEリサーチでニーズ収集", "転職理由・希望条件を簡易アンケートで把握し企画に反映"),
    ("7-8月", "賞与後の転職検討特集", "夏のボーナス後に動く層へ年収相場・求人を配信"),
    ("9月", "資格取得・キャリア特集", "無資格・3級向けに整備士育成プロジェクトを案内"),
    ("10-11月", "年内入社ラストコール", "年内に動きたい層へ面談枠をご案内"),
    ("12月", "冬も相談だけOK訴求", "検索の谷でも検討は続く。配信は止めず関係を維持"),
]
tbl = by_id(s, 46).table
for i, row in enumerate(PLAN, start=1):
    for j, v in enumerate(row):
        set_cell(tbl.cell(i, j), v)
set_text(by_id(s, 47).text_frame, ["出典：Googleトレンド「整備士」（山3月・谷12月）／前後検索KW「カーディーラー」／コンテンツ案v1"])

# 図形IDの重複を解消（移植＋脚注の再追加で重なる可能性）
for sl in (prs.slides[15], prs.slides[19], prs.slides[20], prs.slides[5], prs.slides[6]):
    seen, mx = set(), 0
    els = [e for e in sl.shapes._spTree.iter() if e.tag == qn("p:cNvPr")]
    mx = max(int(e.get("id")) for e in els)
    for e in els:
        i = int(e.get("id"))
        if i in seen:
            mx += 1
            e.set("id", str(mx))
        seen.add(int(e.get("id")))

prs.save(OUT)
print("saved", OUT)
