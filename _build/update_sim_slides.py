# -*- coding: utf-8 -*-
"""v02: 広告SIMと数値を統一（参加率77%・来場単価13,136円）
対象は Slide 11（解決の方向性）と Slide 23（想定の費用対効果）のみ。
"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
TARGET = ROOT / "20260804_株式会社ビイサイドプランニング御中_公式LINEのご提案v02.pptx"

TNAVY = "002060"
NAVY = "1F285A"
RED = "C00000"
INK = "333333"
MUT = "808080"
WHITE = "FFFFFF"
PALE = "F4F7FF"
BORDER = "D9D9D9"

prs = Presentation(str(TARGET))
slides = list(prs.slides)
assert len(slides) == 28, len(slides)


def set_font(run, size, bold=None, color=INK, name="メイリオ"):
    f = run.font
    f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    f.name = name
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        e = rPr.find(qn(tag))
        if e is None:
            e = rPr.makeelement(qn(tag), {})
            rPr.append(e)
        e.set("typeface", name)
    f.color.rgb = RGBColor.from_string(color)


ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
ANCH = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}


def reset_tf(tf):
    for para in tf.paragraphs[1:]:
        para._p.getparent().remove(para._p)
    for r in list(tf.paragraphs[0].runs):
        r._r.getparent().remove(r._r)


def put_text(tf, paras, anchor="t", ml=0.08, mr=0.08, mt=0.04, mb=0.04, wrap=True):
    reset_tf(tf)
    tf.word_wrap = wrap
    tf.margin_left = Inches(ml); tf.margin_right = Inches(mr)
    tf.margin_top = Inches(mt); tf.margin_bottom = Inches(mb)
    tf.vertical_anchor = ANCH[anchor]
    first = True
    for p in paras:
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        para.alignment = ALIGN[p.get("align", "l")]
        if p.get("sa") is not None:
            para.space_after = Pt(p["sa"])
        if p.get("ls") is not None:
            para.line_spacing = p["ls"]
        for (t, sz, b, c) in p["runs"]:
            r = para.add_run(); r.text = t
            set_font(r, sz, b, c)
    return tf


def add_text(slide, x, y, w, h, paras, anchor="t", **kw):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    put_text(box.text_frame, paras, anchor=anchor, **kw)
    return box


def add_box(slide, x, y, w, h, fill=None, line=None, lw=1.0,
            shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08):
    sp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = RGBColor.from_string(line); sp.line.width = Pt(lw)
    sp.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sp.adjustments[0] = radius
        except Exception:
            pass
    return sp


def card(slide, x, y, w, h, head, body, hcol=NAVY, fill=PALE,
         hsz=15, bsz=12, line=None):
    sp = add_box(slide, x, y, w, h, fill=fill, line=line, lw=1.0)
    paras = [{"runs": [(head, hsz, True, hcol)], "sa": 4}]
    if body:
        paras.append({"runs": [(body, bsz, None, INK)], "ls": 1.15})
    put_text(sp.text_frame, paras, anchor="m", ml=0.16, mr=0.14, mt=0.08, mb=0.08)
    return sp


def clear_slide(slide):
    spTree = slide.shapes._spTree
    for el in list(spTree):
        if el.tag.split("}")[-1] in ("sp", "cxnSp", "pic", "graphicFrame", "grpSp"):
            spTree.remove(el)


def frame(slide, title, lead):
    clear_slide(slide)
    add_text(slide, 0.60, 0.13, 9.60, 0.45,
             [{"runs": [(title, 20, True, TNAVY)]}], anchor="m", ml=0.0, mr=0.0)
    add_text(slide, 0.60, 0.70, 9.60, 0.58,
             [{"runs": [(lead, 14, True, INK)], "ls": 1.1}], anchor="m", ml=0.0, mr=0.0)


# ================= Slide 11: 解決の方向性 =================
s = slides[10]
frame(s, "解決の方向性",
      "「予約後の接点」をLINEで設計し、参加率60%→77%、来場単価16,667円→13,136円を目指します。")
rows = [
    ("課題①", "直前辞退", "① Meta広告×LINE連携（オフラインCV最適化）",
     "来場・未来場データをMeta広告へ還元し、来場につながる層へ配信を最適化。"),
    ("課題②", "不安・迷い", "② 直前リマインド・安心コンテンツ配信",
     "前日・当日朝のリマインドと「服装自由・履歴書不要」等のQ&Aで参加辞退を抑止。"),
    ("課題③", "未応募層の離脱", "③ 属性セグメント配信による再アプローチ",
     "アンケート取得した希望職種・年代をもとに、次回フェア・個別求人へ再案内。"),
]
for i, (num, issue, act, desc) in enumerate(rows):
    y = 1.50 + i * 1.32
    sp = add_box(s, 0.60, y, 2.60, 1.15, fill=WHITE, line=RED, lw=1.0)
    put_text(sp.text_frame,
             [{"runs": [(num, 14, True, RED)], "align": "c", "sa": 2},
              {"runs": [(issue, 14, True, RED)], "align": "c"}],
             anchor="m", ml=0.05, mr=0.05, mt=0.0, mb=0.0)
    add_box(s, 3.32, y + 0.36, 0.52, 0.44, fill=NAVY, shape=MSO_SHAPE.RIGHT_ARROW)
    card(s, 3.95, y, 6.25, 1.15, act, desc, hcol=NAVY, fill=PALE, hsz=15, bsz=12)

tb = add_box(s, 0.60, 5.52, 9.60, 1.10, fill=WHITE, line=NAVY, lw=1.5)
put_text(tb.text_frame,
         [{"runs": [("目標効果（広告費は据え置き）", 12, True, NAVY)], "align": "c", "sa": 4},
          {"runs": [("予約→来場の参加率　", 12, None, INK), ("60%", 14, True, INK),
                    (" → ", 12, None, MUT), ("77%", 20, True, NAVY),
                    ("　／　来場単価　", 12, None, INK), ("16,667円", 14, True, INK),
                    (" → ", 12, None, MUT), ("13,136円", 20, True, NAVY)], "align": "c"}],
         anchor="m", ml=0.0, mr=0.0, mt=0.0, mb=0.0)
add_text(s, 0.60, 6.70, 9.60, 0.30,
         [{"runs": [("※参加率はMeta広告シミュレーション（当社作成）の想定値と統一。来場単価はLINE運用費を含む総額ベース。", 10, None, MUT)]}], ml=0.0)

# ================= Slide 23: 想定の費用対効果 =================
s = slides[22]
frame(s, "想定の費用対効果（サマリー）",
      "広告費80万円は据え置いたまま、来場単価を16,667円→13,136円（▲21%）へ改善します。")
kpis = [("初期費用", "¥215,000", "初期構築20万＋離脱防止1.5万"),
        ("月次固定費", "¥80,000", "コンサル5万＋離脱防止1.5万＋アカウント費1.5万"),
        ("契約期間", "3ヶ月〜", "秋開催4会場（9月末〜11月頭）を1クールとして検証")]
for i, (lab, val, note) in enumerate(kpis):
    x = 0.60 + i * 3.28
    sp = add_box(s, x, 1.40, 3.05, 1.50, fill=PALE)
    put_text(sp.text_frame,
             [{"runs": [(lab, 12, True, NAVY)], "align": "c", "sa": 4},
              {"runs": [(val, 20, True, NAVY)], "align": "c", "sa": 4},
              {"runs": [(note, 10, None, INK)], "align": "c", "ls": 1.15}],
             anchor="m", ml=0.12, mr=0.12, mt=0.06, mb=0.06)

rows = [
    ("月額広告費", "80万円", "80万円", "±0円（据え置き）"),
    ("予約単価（CPA）", "10,000円", "10,000円", "広告側は現状維持"),
    ("月間予約数", "80名", "87名", "+7名：LINE経由の申込増"),
    ("予約→来場の参加率", "60%", "77%", "+17pt：リマインド＋Meta連携"),
    ("月間来場者数", "48名", "67名", "+19名（+40%）"),
    ("来場単価（LINE費込）", "16,667円", "13,136円", "▲21%：ドタキャン層への広告費を排除"),
]
gf = s.shapes.add_table(7, 5, Inches(0.60), Inches(3.06), Inches(9.60), Inches(2.72))
tbl = gf.table
tbl.first_row = False
tbl.horz_banding = False
for i, w in enumerate([2.55, 1.45, 0.45, 1.65, 3.50]):
    tbl.columns[i].width = Inches(w)
tbl.rows[0].height = Inches(0.34)
for r in range(1, 7):
    tbl.rows[r].height = Inches(0.39)
for c, htext in enumerate(["指標", "現状の実績", "", "提案後（目標）", "改善インパクト"]):
    cell = tbl.cell(0, c)
    cell.fill.solid(); cell.fill.fore_color.rgb = RGBColor.from_string(NAVY)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = Inches(0.06); cell.margin_right = Inches(0.04)
    cell.margin_top = Inches(0.01); cell.margin_bottom = Inches(0.01)
    put_text(cell.text_frame, [{"runs": [(htext, 11, True, WHITE)], "align": "c"}],
             anchor="m", ml=0.02, mr=0.02, mt=0.0, mb=0.0)
for r, (met, cur, prop, drv) in enumerate(rows, start=1):
    band = WHITE if r % 2 == 1 else PALE
    hl = r in (5, 6)
    vals = [(met, 10.5, True, INK, "l"), (cur, 11, None, INK, "c"),
            ("→", 11, None, MUT, "c"),
            (prop, 14 if hl else 12, True, NAVY, "c"), (drv, 10, None, INK, "l")]
    for c, (t, sz, b, col, al) in enumerate(vals):
        cell = tbl.cell(r, c)
        cell.fill.solid(); cell.fill.fore_color.rgb = RGBColor.from_string(band)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = Inches(0.08); cell.margin_right = Inches(0.04)
        cell.margin_top = Inches(0.01); cell.margin_bottom = Inches(0.01)
        put_text(cell.text_frame, [{"runs": [(t, sz, b, col)], "align": al}],
                 anchor="m", ml=0.02, mr=0.02, mt=0.0, mb=0.0)

opt = add_box(s, 0.60, 5.92, 9.60, 0.72, fill=WHITE, line=BORDER, lw=1.0)
put_text(opt.text_frame,
         [{"runs": [("［参考］広告費を100万円へ増額した場合", 11, True, NAVY),
                    ("　来場75名・来場単価14,460円（▲13%）。", 11, None, INK),
                    ("増額分の増分単価は25,000円/名となり現状単価を上回るため、まずは予算据え置きでの検証を推奨します。", 11, None, INK)], "ls": 1.15}],
         anchor="m", ml=0.20, mr=0.16, mt=0.06, mb=0.06)
add_text(s, 0.60, 6.74, 9.60, 0.30,
         [{"runs": [("参照元：貴社ご提供の実績数値／参加率はMeta広告シミュレーション（当社作成）と統一。詳細は別添Excel（SIM）をご参照ください。", 10, None, MUT)]}], ml=0.0)

prs.save(str(TARGET))
print("saved:", TARGET)
