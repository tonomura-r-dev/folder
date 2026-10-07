# -*- coding: utf-8 -*-
"""サンクスLINE誘導・業界別 ver3.0 の1枚目を、殿村さん添付の参考画像どおりに組み直す（2026-10-07）。2〜5枚目は触らない。
  左パネル（薄青）：CV獲得まで｜施策全体のCPA改善：広告→LP→離脱防止×成果報酬施策→CV／「CVの取りこぼしを減らし、獲得チャネルを増やすことで施策全体のCPAを改善」
  右パネル（薄緑）：CV獲得後｜CPOの改善：CV→サンクスLINE誘導→LINEでナーチャリング→購入・契約などの最終成果／「CV後の歩留まりを改善し、CPOを改善」
  下：CPOの改善→許容CPAの引き上げ→広告配信の拡大（各サブ説明）＋戻り矢印「さらに成果が増えれば、再投資して継続的に拡大」
  最下：紺の帯「広告とあわせて実施することで、事業全体の効率化を図る。」（「事業全体の効率化」は黄色）
画像にリード行・区切り線が無いので、このページは外す。
  python3 _build/patch_thanks_line_gyokai_s1_ref.py <業界別ver3.0.pptx> [出力]
"""
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

sys.path.insert(0, str(Path(__file__).resolve().parent))
from pptx_parts import *  # noqa
from pptx_parts import _flat

OUT = Path(sys.argv[2]) if len(sys.argv) > 2 else Path(sys.argv[1])
PANEL_L, PANEL_R = "F2F5FB", "EEF7F1"
PALE, PALE2 = "DDE3F2", "E4E8F6"
GREEN_DK, YELLOW = "0E8A3F", "FFE400"
L, R = 2.2, 31.7
GAP = 0.7
PW = (R - L - GAP) / 2

prs = Presentation(sys.argv[1])
s = prs.slides[0]
assert s.shapes[0].text_frame.text.startswith("広告と一緒にサンクスLINE"), "1枚目が「広告と一緒に…」ではない"
for sh in list(s.shapes):
    if sh.name != "TextBox 1":
        sh._element.getparent().remove(sh._element)

PY, PH = 2.2, 9.9


def panel(x, fill, pill, pill_fill, head_runs, steps, caption):
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, PY, PW, PH, None, fill=fill, adj=0.04)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x + 0.8, PY + 0.55, 3.6, 0.9, [(pill, 11, True, WHITE, 0)], fill=pill_fill, adj=0.5,
          margins=(0.05, 0, 0.05, 0))
    rich(s, x + 4.6, PY + 0.2, PW - 5.0, 1.6, [head_runs], align=PP_ALIGN.LEFT)
    bx, bw, bh, gap = x + 1.3, PW - 2.6, 1.15, 0.5
    y = PY + 2.1
    for i, (t, f, c) in enumerate(steps):
        shape(s, MSO_SHAPE.RECTANGLE, bx, y, bw, bh, [(t, 13, True, c, 0)], fill=f, margins=(0.1, 0, 0.1, 0))
        if i < len(steps) - 1:
            tri = s.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, Cm(x + PW / 2 - 0.22), Cm(y + bh + 0.1), Cm(0.44), Cm(0.3))
            tri.rotation = 180
            tri.fill.solid()
            tri.fill.fore_color.rgb = rgb(c if f not in (WHITE, PALE, PALE2) else (NAVY if fill == PANEL_L else GREEN_DK))
            tri.line.fill.background()
            _flat(tri)
        y += bh + gap
    label(s, x + 0.6, y - 0.1, PW - 1.2, 1.4, [(t, 10.5, False, INK, 0) for t in caption], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


panel(L, PANEL_L, "CV獲得まで", NAVY,
      [("施策全体の ", 15, True, NAVY), ("CPA", 30, True, NAVY), (" 改善", 19, True, NAVY)],
      [("広告", PALE, INK), ("LP", PALE, INK), ("離脱防止 × 成果報酬施策", PALE, INK), ("CV", NAVY, WHITE)],
      ["CVの取りこぼしを減らし、獲得チャネルを増やすことで", "施策全体のCPAを改善"])
panel(L + PW + GAP, PANEL_R, "CV獲得後", GREEN_DK,
      [("CPO", 30, True, GREEN_DK), (" の改善", 19, True, GREEN_DK)],
      [("CV", WHITE, INK), ("サンクスLINE誘導", GREEN_BG, GREEN_TX), ("LINEでナーチャリング", GREEN_BG, GREEN_TX), ("購入・契約などの最終成果", GREEN_DK, WHITE)],
      ["CV後の歩留まりを改善し、CPOを改善"])

# ---- 下：3ステップ ----
CY, CH = PY + PH + 0.5, 1.9
ov = 0.3
cw = (R - L + 2 * ov) / 3
chev = [("CPOの改善", "1件の最終成果あたりの広告費を抑えられる", GREEN_BG, GREEN_TX),
        ("許容CPAの引き上げ", "1件の最終成果にかけられる広告費が増える", PALE2, NAVY),
        ("広告配信の拡大", "入札・予算・配信対象を広げやすくなる", PALE2, NAVY)]
for i, (t, sub, f, c) in enumerate(chev):
    x = L + i * (cw - ov)
    shape(s, MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON, x, CY, cw, CH, None, fill=f, adj=0.22)
    pad = 0.4 if i == 0 else 0.9
    label(s, x + pad, CY, cw - pad - 0.8, CH, [(t, 14, True, c, 2), (sub, 9.5, False, INK, 0)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

# 戻り矢印（右端→下→左端→上）
LY = CY + CH + 0.55
xr, xl = R - 0.5, L + 0.5
vline(s, xr, CY + CH + 0.1, LY, NAVY, 1.5)
hline(s, xl, xr, LY, NAVY, 1.5)
arrow_line(s, xl, LY, xl, CY + CH + 0.12, NAVY, 1.5)
lw = 11.0
rect(s, (L + R) / 2 - lw / 2, LY - 0.35, lw, 0.7, WHITE, [("さらに成果が増えれば、再投資して継続的に拡大", 10, True, INK, 0)])

# ---- 最下：紺の帯 ----
BY = LY + 0.65
rect(s, L, BY, R - L, 1.7, NAVY)
rich(s, L, BY, R - L, 1.7, [[("広告とあわせて実施することで、", 20, True, WHITE), ("事業全体の効率化", 20, True, YELLOW), ("を図る。", 20, True, WHITE)]],
     align=PP_ALIGN.CENTER)

prs.save(str(OUT))
print("saved:", OUT.name)
