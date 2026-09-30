# -*- coding: utf-8 -*-
"""チェングロウス本資料 ver2.5 に、SIM ver2.7（CV①の率 1.5%→1.0%、CV②は据え置き）の数字を反映する。

  python3 _build/patch_chengrowth_v25_visual.py   # 先に入れ替え版を作る
  python3 _build/patch_chengrowth_v25_sim27.py    # その上に数字を反映（ver2.5を上書き）

反映先：12枚目（②の表・注記）／15枚目（半年後の友だち数）／18枚目（シミュレーションの表・前提）
②の表は、SIMの9月の値から次の考え方で出す（ver2.6のときと同じ）：
  対象＝SIMの流入 ÷（通数×開封率×クリック率×反応率の補正）、クリック数＝流入×2.05（複数回のクリック）、
  リッチメニュー＝有効な友だち×30%×2回。動線00/01は、友だちの累計に占める離脱防止の割合で振り分ける。
  件数は SIMの9月の応募数（CV②）を クリック数の比で振り分け、整数にする（端数は大きい順に繰り上げ）。
"""
import math
from pathlib import Path

import openpyxl
from pptx import Presentation
from pptx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
DECK = ROOT / "20260929_株式会社チェングロウス御中_LINE公式アカウント運用のご提案ver2.5.pptx"
SIM = ROOT / "20260929_株式会社チェングロウス御中_LINEOA施策提案【SIM】ver2.7.xlsx"

v = openpyxl.load_workbook(SIM, data_only=True)["SIM3_LINE登録2%"]
M = range(5, 11)                                   # E〜J＝4〜9月
row = lambda r: [v.cell(r, c).value for c in M]
J = lambda r: v.cell(r, 10).value                  # 9月
P = lambda r: v.cell(r, 16).value

# ---- ②の表（9月）
share00 = sum(row(23)) / J(28)                     # 友だちの累計に占める離脱防止
kikaku_t = J(31) / (4 * P(22) * P(23) * P(36))
step_t = J(33) / (10 * P(26) * P(27) * P(36))
target = J(29)
clicks = {"step": J(33) * 2.05, "kikaku": J(31) * 2.05, "rich": target * 0.3 * 2}
total_click = round(J(35))
cv2 = J(37)
rate = cv2 / J(35)


def split(values, total):
    """合計を保ったまま整数にする（端数は大きい順に繰り上げ）"""
    fl = [math.floor(x) for x in values]
    for i in sorted(range(len(values)), key=lambda i: values[i] - fl[i], reverse=True)[:total - sum(fl)]:
        fl[i] += 1
    return fl


keys = [("00", "step"), ("00", "kikaku"), ("00", "rich"), ("01", "step"), ("01", "kikaku"), ("01", "rich")]
sh = {"00": share00, "01": 1 - share00}
raw_click = [clicks[k] * sh[d] for d, k in keys]
click_i = split(raw_click, total_click)
cnt_i = split([c / sum(click_i) * cv2 for c in click_i], cv2)
tgt = {"step": step_t, "kikaku": kikaku_t, "rich": target}
tlabel = {"step": ("離脱防止から\n追加した方", "LINEで登録した方\n"), "kikaku": ("離脱防止から\n追加した方", "LINEで登録した方\n"),
          "rich": ("有効な友だち\n", "有効な友だち\n")}
rate_s = f"{rate * 100:.1f}%"
print(f"離脱防止の割合 {share00:.1%} / 応募率 {rate_s} / クリック {click_i}={sum(click_i)} / 件数 {cnt_i}={sum(cnt_i)}")

prs = Presentation(str(DECK))
S = list(prs.slides)


def set_cell(cell, text):
    ts = list(cell._tc.iter(qn("a:t")))
    ts[0].text = text
    for t in ts[1:]:
        t.text = ""


def sub_all(slide, pairs):
    for t in slide.shapes._spTree.iter(qn("a:t")):
        for a, b in pairs:
            if t.text and a in t.text:
                t.text = t.text.replace(a, b)


# 12枚目：②の表
s = S[11]
tbl = next(x for x in s.shapes if x.has_table).table
for i, (d, k) in enumerate(keys, start=1):
    r = tbl.rows[i].cells
    who = tlabel[k][0 if d == "00" else 1]
    set_cell(r[3], f"{who}（約{round(tgt[k] * sh[d])}人）" if k != "rich" else f"{who}（約{round(tgt[k] * sh[d])}人）")
    head = r[5].text.split("→")[0]
    set_cell(r[5], f"{head}→ {click_i[i - 1]}回")
    set_cell(r[6], f"{rate_s}\n→ {cnt_i[i - 1]}件")
set_cell(tbl.rows[9].cells[5], f"{sum(click_i)}回")
set_cell(tbl.rows[9].cells[6], f"{sum(cnt_i)}件")
sub_all(s, [("応募率2.0%", f"応募率{rate_s}"),
            ("離脱防止29%・LINEで登録71%", f"離脱防止{share00 * 100:.0f}%・LINEで登録{100 - share00 * 100:.0f}%")])

# 15枚目：半年後の友だち数
friends9 = J(28)
sub_all(S[14], [("約700人", f"約{round(friends9, -1):.0f}人")])

# 18枚目：シミュレーション
s = S[17]
tbl = next(x for x in s.shapes if x.has_table).table
for j, val in enumerate(row(28), start=1):
    set_cell(tbl.rows[1].cells[j], f"{round(val)}人")
for j, val in enumerate(row(36), start=1):
    set_cell(tbl.rows[2].cells[j], f"{val}件")
cv1_total = sum(row(36))
set_cell(tbl.rows[2].cells[7], f"{cv1_total}件")
sub_all(s, [("501件", f"{cv1_total}件"), ("来訪者の1.5%", "来訪者の1.0%")])

prs.save(str(DECK))
print("saved:", DECK.name)
for i in (11, 14, 17):
    sl = Presentation(str(DECK)).slides[i]
    for x in sl.shapes:
        if x.has_table:
            for r in x.table.rows:
                print(i + 1, [c.text.replace("\n", "/") for c in r.cells])
        elif x.has_text_frame and any(k in x.text_frame.text for k in ("応募率", "離脱防止", "約", "件", "1.0%")):
            print(i + 1, x.text_frame.text.replace("\n", "/")[:120])
