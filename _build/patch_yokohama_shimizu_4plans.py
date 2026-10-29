# -*- coding: utf-8 -*-
"""横浜シミズ資料（19枚・殿村さんPC保存版が正）→ SIM ver2.2（殿村さんPC保存版b）の数字に合わせる（2026-10-29 殿村さん）。

    python -m _build.patch_yokohama_shimizu_4plans <資料.pptx> [出力.pptx]

  ・15ページ「料金プラン」の表に成果報酬プランを4列目として足し、4つのプランを1枚で比べられるようにする
    （成果報酬＝初期295,000円〜／月額運用費の代わりに 応募1件8,000円・会員登録1件2,500円／ツール・LINE関連費 65,000円〜）
  ・16ページ「料金プラン（成果報酬）」は消す（中身は15ページへ）。成果報酬の対象の注記は15ページへ移す
  ・14ページ「概算お見積り」（SIM①＝コンサルのプラン②）：6か月目 友だち1,102人・新規応募 月7件・再応募 月3件・会員登録 月10件
  表の列を写したら a16:colId を振り直す（CLAUDE.md 2026-09-24）。出力を省くと上書き。
"""
import copy
import sys

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches

from _build.pptx_parts import drop_slide

A16 = "{http://schemas.microsoft.com/office/drawing/2014/main}"
PLAN4 = [["プラン④", "成果報酬"], ["295,000円〜"], ["成果報酬", "応募 8,000円/件", "会員登録 2,500円/件"],
         ["65,000円〜"], ["65,000円〜", "＋成果報酬"]]
NOTES15 = [
    "※初期費用＝運用設計・構築 100,000円〜＋離脱防止 15,000円＋サンクスLINE誘導 150,000円＋LINE公式アカウント開設 30,000円",
    "※月額ツール・LINE関連費＝離脱防止 30,000円＋サンクスLINE誘導 30,000円〜＋LINE公式アカウント利用料 5,000円〜（ライトプラン想定）",
    "※成果報酬の対象＝LINEからの新規応募・再応募・会員登録（再応募は、勤務希望をLINEで受け付ける場合）",
    "※契約期間は6か月〜。概算お見積りのページはプラン②の場合",
    "※施策内容・配信本数・連携範囲等により、費用は個別にお見積もりいたします。",
]


def by_id(s, sid):
    return next(sh for sh in s.shapes if sh.shape_id == sid)


def set_lines(el, lines, size=None):
    """txBody を持つ要素（図形・表のセル）の文字を、最初の段落の書式で1行＝1段落に入れ直す"""
    txb = el.find(qn("p:txBody"))
    if txb is None:
        txb = el.find(qn("a:txBody"))
    ps = txb.findall(qn("a:p"))
    tmpl = copy.deepcopy(ps[0])
    for r in tmpl.findall(qn("a:r"))[1:] + tmpl.findall(qn("a:br")):
        tmpl.remove(r)
    for p in ps:
        txb.remove(p)
    for line in lines:
        p = copy.deepcopy(tmpl)
        r = p.find(qn("a:r"))
        r.find(qn("a:t")).text = line
        if size:
            r.find(qn("a:rPr")).set("sz", str(size))
        txb.append(p)


def page15(s):
    tshape = next(sh for sh in s.shapes if sh.has_table)
    tbl = tshape._element.find(".//" + qn("a:tbl"))
    grid = tbl.find(qn("a:tblGrid"))
    gcs = grid.findall(qn("a:gridCol"))
    gcs[-1].addnext(copy.deepcopy(gcs[-1]))
    for k, gc in enumerate(grid.findall(qn("a:gridCol"))):
        gc.set("w", str(int(Inches(1.8 if k == 0 else 2.05))))
        for e in gc.iter(A16 + "colId"):
            e.set("val", str(20000 + k))
    for tr, lines in zip(tbl.findall(qn("a:tr")), PLAN4):
        last = tr.findall(qn("a:tc"))[-1]
        new = copy.deepcopy(last)
        last.addnext(new)
        set_lines(new, lines, size=1300 if len(lines) == 3 else None)
    rows = tbl.findall(qn("a:tr"))                            # 言葉の途中で切れないよう、切れ目で改行
    set_lines(rows[0].findall(qn("a:tc"))[3], ["プラン③", "運用代行", "フルサポート"])
    set_lines(rows[3].findall(qn("a:tc"))[0], ["月額ツール・", "LINE関連費"])
    set_lines(by_id(s, 635)._element, ["料金プラン：4つのプランの費用"])
    set_lines(by_id(s, 637)._element, ["運用範囲と料金の形に合わせて、4つのプランをご用意"])
    notes = [by_id(s, i) for i in (640, 641, 642, 643)]
    extra = copy.deepcopy(notes[-1]._element)
    notes[-1]._element.addnext(extra)
    nid = max(int(e.get("id")) for e in s.shapes._spTree.iter(qn("p:cNvPr"))) + 1
    extra.find(".//" + qn("p:cNvPr")).set("id", str(nid))
    notes.append(next(sh for sh in s.shapes if sh.shape_id == nid))
    for k, (sh, text) in enumerate(zip(notes, NOTES15)):
        set_lines(sh._element, [text])
        sh.top = Inches(5.81 + 0.25 * k)


def page14(s):
    set_lines(by_id(s, 339)._element, ["概算お見積り：月215,000円で、6か月目に新規応募 月7件・再応募 月3件"])
    set_lines(by_id(s, 375)._element, ["1,102人"])
    set_lines(by_id(s, 376)._element, ["7件/月"])
    step = by_id(s, 364).top - by_id(s, 362).top                # 再応募の下に「会員登録」の行を足す
    nid = max(int(e.get("id")) for e in s.shapes._spTree.iter(qn("p:cNvPr"))) + 1
    for k, (sid, text) in enumerate(((364, "会員登録"), (378, "10件/月"))):
        el = copy.deepcopy(by_id(s, sid)._element)
        el.find(".//" + qn("p:cNvPr")).set("id", str(nid + k))
        by_id(s, sid)._element.addnext(el)
        sh = by_id(s, nid + k)
        sh.top = by_id(s, sid).top + step
        set_lines(sh._element, [text])


def main(src, out):
    prs = Presentation(src)
    assert len(prs.slides) == 19
    s14, s15, s16 = (prs.slides[k] for k in (13, 14, 15))
    assert by_id(s16, 635).text_frame.text.startswith("料金プラン（成果報酬）")
    page14(s14)
    page15(s15)
    drop_slide(prs, s16)
    assert len(prs.slides) == 18
    prs.save(out)
    print("saved:", out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else sys.argv[1])
