"""
10月号（殿村さんの表紙32pt版＝正）に、2026-10-21受領のPDF5本（P.3〜P.7）の内容を反映する。
ビルドし直さず、正のファイルの該当テキストだけを差し替える（レイアウト・書式はそのまま）。

使い方: python3 _build/patch_trend_202610_pdf.py <pptx>
"""
import sys
from pptx import Presentation
import trend_report_engine as engine

F = sys.argv[1]
prs = Presentation(F)
slides = list(prs.slides)


def find(slide, first_line):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.paragraphs[0].text.strip() == first_line:
            return sh
    raise SystemExit(f"見つからない: {first_line}")


def set_paras(sh, lines):
    paras = sh.text_frame.paragraphs
    assert len(paras) == len(lines), (sh.text_frame.text, lines)
    for p, line in zip(paras, lines):
        p.runs[0].text = line
        for r in list(p.runs[1:]):
            r._r.getparent().remove(r._r)


def desc(slide, lines):
    for sh in slide.shapes:
        if sh.has_text_frame and len(sh.text_frame.paragraphs) == 3 \
                and sh.text_frame.paragraphs[0].alignment == engine.PP_ALIGN.CENTER:
            engine.set_multiline_text(sh, lines)
            return
    raise SystemExit("説明バンドが見つからない")


def card(slide, old_head, head, body):
    set_paras(find(slide, old_head), [head, body])


# P.3 ネイティブアプリ遷移
s = slides[2]
desc(s, [
    "これまで一律禁止だった、LINE通知メッセージからネイティブアプリへの遷移が、",
    "LINEヤフーの審査で承認された場合に限り解禁された。",
    "2026年9月9日（水）からLBPMで審査受付を開始（審査費 10万円／1UX・税抜）。",
])
set_paras(find(s, "通知メッセージのリンクからネイティブアプリへは遷移できない"),
          ["一律でWebサイトへ遷移（ネイティブアプリへの遷移は禁止）"])
set_paras(find(s, "ベースプラン限定・条件付きで遷移可能に"),
          ["審査承認でネイティブアプリへ遷移可能に",
           "エッセンシャル／パフォーマンス／アンリミテッドプランが対象"])
card(s, "通知からアプリへ直接誘導できる", "シームレスな導線でUXが向上",
     "アプリがインストール済みならアプリへ直接遷移、未インストールなら通知内容に関連するWebページへ遷移する。")
card(s, "遷移には事前審査が必要", "主な遷移条件",
     "Universal links／App Linksであること、送信元と同一ドメインであること、アプリとWebで表示内容が同じであること。ストアへの自動遷移は不可。")

# P.4 UX追加申請
s = slides[3]
card(s, "発注画面・申請方法も変更", "既存案件は再作成が必要",
     "過去に作成した申請を編集しても受領されない。既存案件をコピーして再作成し、LBPMから改めて発注・申請する。")
card(s, "請求書の発行先が変わる", "アプリ遷移審査は有償オプションで選択",
     "2026年9月9日からの新しい発注画面では、申請フォームの入力後に金額計算画面が出る。ネイティブアプリ遷移は有償オプションから選ぶ。")

# P.5 分析 参照期間
s = slides[4]
set_paras(find(s, "リリース日：2027年2月頃に変更"),
          ["リリース日：2027年2月頃に変更", "日程は決まり次第、管理画面のお知らせで案内（変更の場合あり）"])
card(s, "対象はWeb版・管理アプリの両方", "影響は「分析」タブの実績値のみ",
     "Web版管理画面・管理アプリの「分析」タブが対象。それ以外のデータや、作成済みのメッセージには影響しない。")
card(s, "過去データの確認は変更前に", "過去データは変更前にダウンロード",
     "36カ月より前の実績値が必要な場合は事前に確認する。「分析」タブの一部データはダウンロードできる。")

# P.6 シリアルコードスタンプ
s = slides[5]
desc(s, [
    "2026年10月1日（木）リリースの新商品「LINEシリアルコードスタンプ」について、",
    "媒体資料（更新版）と進行ガイド（2026年10月版）が公開された。",
    "クリエイティブ基準の緩和と、独自のキャンペーン設計ルールがポイント。",
])
card(s, "媒体資料：体裁を統一", "クリエイティブ基準を3点緩和",
     "この商品に限り、飲酒表現（注記必須）・実写素材（CM・ポスターの転用可）・PR要素の個数（全体の1/4以内→制限なし）を緩和。")
card(s, "媒体資料：スペックの一部更新", "告知できる面・できない面",
     "広告主のLINE公式アカウント・アプリ・SNS・Web・店頭はOK。スタンプショップなどLINEヤフー保有アプリからの告知はNG。")
card(s, "進行ガイド：独自フロー", "シリアルコードの付与条件",
     "Web・オフライン（購入・来店など）での条件達成はOK。アプリ内での条件達成や、獲得までにアプリ遷移を含む導線はNG。")
card(s, "進行ガイド：設計・制作ルール", "コード発行は発注受領後",
     "発注受領後にLBPMで申請し、3営業日を目途に納品。追加発注は不可。印刷・配荷がある場合はその日数も見込む。")

# P.7 利用と請求
s = slides[6]
card(s, "対応", "利用中の機能は通常どおり",
     "顧客管理・飲食注文・サロン予約など、利用中の機能はメンテナンス中も使える。チャットProオプションは影響なし。")

engine.force_font(prs)
prs.save(F)
print("saved", F)
