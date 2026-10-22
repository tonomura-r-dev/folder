# -*- coding: utf-8 -*-
"""すなお ver6.5（殿村さんPC保存版・59枚）→ ver6.6：「配信の内容」の章を立てる（2026-10-22 殿村さん）。

    python -m _build.patch_sunao_v66_haishin_chapter <ver6.5.pptx> <ver6.6.pptx>

  ①「施策の詳細」のアジェンダのページを複製して「配信の内容」に変え、「LINEOA施策」の前に置く
  ②「配信の出し分け：予約した方／予約していない方」を、運用方針の章から「LINEOA施策」の直後へ移す
  並び：…サンクスLINE誘導 → 資料アジェンダ（配信の内容）→ LINEOA施策 → 配信の出し分け → ステップ配信 → 参加後のフォローアップ。
  ほかのページの中身は変えない（59枚 → 60枚）。
"""
import sys

from pptx import Presentation

from _build.patch_tanaka_v5 import copy_slide


def title(s):
    return next(sh.text_frame.text for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip())


def main(src, out):
    prs = Presentation(src)
    slides = list(prs.slides)
    assert len(slides) == 59
    agenda, split, lineoa = slides[25], slides[18], slides[30]
    assert [sh.text_frame.text for sh in agenda.shapes if sh.has_text_frame][1] == "施策の詳細"
    assert title(split).startswith("配信の出し分け")
    assert title(lineoa) == "LINEOA施策"

    new = copy_slide(prs, agenda)
    box = next(sh for sh in new.shapes if sh.has_text_frame and sh.text_frame.text == "施策の詳細")
    r = box.text_frame.paragraphs[0].runs
    r[0].text = "配信の内容"
    for x in r[1:]:
        x._r.getparent().remove(x._r)

    lst = prs.slides._sldIdLst
    ids = {s.slide_id: e for s, e in zip(prs.slides, list(lst))}
    e_new, e_split, e_lineoa = ids[new.slide_id], ids[split.slide_id], ids[lineoa.slide_id]
    for e in (e_new, e_split):
        lst.remove(e)
    e_lineoa.addprevious(e_new)
    e_lineoa.addnext(e_split)

    order = [title(s) if i != 1 else "" for i, s in enumerate(prs.slides)]
    k = order.index("LINEOA施策")
    assert order[k + 1].startswith("配信の出し分け") and order[k + 2].startswith("友だち追加直後")
    assert len(prs.slides) == 60
    prs.save(out)
    print("saved:", out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
