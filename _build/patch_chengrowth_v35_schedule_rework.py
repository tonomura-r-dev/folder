# -*- coding: utf-8 -*-
"""チェングロウス ver3.4 → ver3.5：23枚目、現行サイトに入れると作り直しになるため、
Profile+は申請だけ先／実装・サンクスLINE誘導はリニューアルのサイトに入れる整理に（2026-10-02）
  python3 _build/patch_chengrowth_v35_schedule_rework.py <ver3.4.pptx>
"""
import sys
from pathlib import Path
from pptx import Presentation

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "20260929_チェングロウス_LINE公式アカウント運用のご提案ver3.5.pptx"
prs = Presentation(sys.argv[1])


def shape(n, name):
    return next(sh for sh in prs.slides[n - 1].shapes if sh.name == name)


def set_par(n, name, k, text):
    p = shape(n, name).text_frame.paragraphs[k]
    runs = p.runs
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


set_par(23, "Rounded Rectangle 7", 1, "あいさつ・リッチメニュー・配信の設計、各ツールとLINEの連携、離脱防止の設置")
set_par(23, "Rounded Rectangle 11", 1, "リニューアルしたサイトに「LINEで登録」・サンクスLINE誘導を組み込み、計測と配信を本格化")
set_par(23, "TextBox 14", 0, "※LINE Profile+の審査期間は、申請内容により変わります／求人ボックス・Indeed経由の応募者への導線は、着地先を確認のうえ決定")
prs.save(str(OUT))
print("saved:", OUT.name, len(prs.slides), "枚")
