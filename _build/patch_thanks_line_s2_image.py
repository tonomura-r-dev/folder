# -*- coding: utf-8 -*-
"""サンクスLINEのご提案（3枚）：S2の流れの図を画像（_images/thanks_line_flow.png）に差し替える（2026-10-01）。
画像は殿村さんがChatGPTで生成したもの。強み3つは画像の下に見出しだけの帯で残す。
  python3 _build/patch_thanks_line_s2_image.py <サンクスLINEのご提案.pptx>   （上書き保存）
"""
import sys
from pathlib import Path
from pptx import Presentation
from pptx.util import Cm, Pt
from pptx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
IMG = ROOT / "_images" / "thanks_line_flow.png"
path = sys.argv[1]
prs = Presentation(path)
s = prs.slides[1]
by = lambda n: next(x for x in s.shapes if x.name == n)

# 流れの図（四角と矢印）を消す
for n in ["Rounded Rectangle 4", "Right Arrow 5", "Rounded Rectangle 6", "Right Arrow 7",
          "Rounded Rectangle 8", "Right Arrow 9", "Rounded Rectangle 10", "Rounded Rectangle 11"]:
    el = by(n)._element
    el.getparent().remove(el)

# 画像（16:9）を中央に
W = 21.0
H = W * 941 / 1672
pic = s.shapes.add_picture(str(IMG), Cm((prs.slide_width / 360000 - W) / 2), Cm(4.15), Cm(W), Cm(H))
pic.name = "サンクスLINEの流れ"

# 強み3つ：見出しだけの帯にする
CW, G, X0, Y, HH = (24.6 - 0.4) / 3, 0.2, 1.46, 4.15 + H + 0.25, 1.25
for i, n in enumerate(("Rounded Rectangle 12", "Rounded Rectangle 13", "Rounded Rectangle 14")):
    sh = by(n)
    sh.left, sh.top, sh.width, sh.height = Cm(X0 + i * (CW + G)), Cm(Y), Cm(CW), Cm(HH)
    txb = sh.text_frame._txBody
    for p in txb.findall(qn("a:p"))[1:]:
        txb.remove(p)
    for r in sh.text_frame.paragraphs[0].runs:
        r.font.size = Pt(15)
    sh.text_frame.paragraphs[0].line_spacing = 1.0
    sh.text_frame.vertical_anchor = 3   # 中央
    sh.text_frame.paragraphs[0].alignment = 2   # 左右も中央
prs.save(path)
print("saved:", path)
