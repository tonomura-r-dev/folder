# -*- coding: utf-8 -*-
"""学校法人業界（大学・専門学校の学生募集）LINEOA施策提案（50枚）

**原本は絶対に編集しない。** `_templates/ジム・フィットネス業界_LINEOA施策提案_ver3.pptx`
（46枚・構成の完成見本）をコピーし、業界依存ページを差し替え、
施策設計4点セット（①動線／②効率改善（配信）／③満足度改善（ユーザー）／④効率改善（管理側））
を4枚追加して50枚にする。

  python3 _build/build_gakkou.py
  python3 _build/qa_render.py 学校法人業界_LINEOA施策提案.pptx

【この資料の背骨】
高校生の進路検討は約18ヶ月。なのに学校の接点はオープンキャンパス当日の数時間だけ。
しかも勝負は年内（総合型9/1出願・11/1発表）に終わる。だから仕込みは春。
そして募集のチャンスは年1回。1人の入学は4年間で約400万円。

【実測データ（★このデッキの強み。捏造なし）】
- p16/p17 Googleトレンド（_data/trends/学校法人/・2026-09-15取得）
    OC＝5年連続7月ピーク／週次ピーク7/27週・6/29週で既に半分・動き出し5/25週
    総合型選抜＝5年で3.3倍（9月値30→100）・2023年9月に指定校推薦を逆転
    週次ピーク8/31週＝9/1出願解禁を含む週／専門学校＝年中フラット22〜40
- p19/p20 前後検索（_data/journey/学校法人/・読み取りメモ）

【未取得データ（★破線の差込枠。数値は捏造していない）】
- p9/p10/p11 競合の友だち数・リッチメニュー実機（page.line.me をPCで実査）
- p19 起点KW未確認／総合型選抜・専門学校の前後検索が未取得
- p18/p37/p39 社内実績（CV地点別 CPC・CVR・CPA）

【落とし穴メモ】
- put_text() は必ず reset_tf() を通す（既存テキストへの追記事故を防ぐ）
- 書式はジムver3に合わせる：タイトル16pt太字 #1F2A5C／リード15pt #646C82／
  本文・表12pt #1C2233／脚注9pt #646C82。**メイリオ固定**
- BUFFF由来の数値は3つだけラベルを直す（README_FMTの正しい使い方.md 参照）
"""
import re
import shutil
from pathlib import Path

from pptx import Presentation
from pptx.util import Cm, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "_templates" / "ジム・フィットネス業界_LINEOA施策提案_ver3.pptx"
OUT = ROOT / "学校法人業界_LINEOA施策提案.pptx"
IMG = ROOT / "_images"

# ---- 配色（ジムver3の実測値に合わせる）----
TNAVY = "1F2A5C"   # タイトル
LEADC = "646C82"   # リード・脚注
INK = "1C2233"     # 本文
NAVY = "1F285A"    # カード見出し
ORANGE = "ED7D31"
RED = "C00000"
MUT = "7F7F7F"
WHITE = "FFFFFF"
PALE = "F4F7FF"
PORANGE = "FCE4D6"
GREY = "F2F2F2"
BORDER = "D9D9D9"
GREEN = "06C755"   # LINE UI 専用
BEZEL = "2B2B2B"
SCREEN = "EFF2F7"
PRED = "FDF2F2"

# ---- レイアウト座標（cm・ジムver3準拠）----
SW, SH = 27.52, 19.05
TITLE_XY = (1.50, 0.30, 20.90, 1.00)
LEAD_XY = (1.10, 2.00, 25.4, 1.10)
CX0, CW = 1.10, 25.40
CY0 = 3.40
BAND_Y = 15.40
FOOT_Y = 17.05

shutil.copyfile(SRC, OUT)          # ★原本は触らない
prs = Presentation(str(OUT))
slides = list(prs.slides)
assert len(slides) == 46, len(slides)
assert (prs.slide_width, prs.slide_height) == (9906000, 6858000)


# ======================= helpers =======================
def set_font(run, size, bold=None, color=INK, name="メイリオ"):
    f = run.font
    f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    f.name = name
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        e = rPr.find(qn(tag))
        if e is None:
            e = rPr.makeelement(qn(tag), {})
            rPr.append(e)
        e.set("typeface", name)
    f.color.rgb = RGBColor.from_string(color)


ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
ANCH = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}


def reset_tf(tf):
    for para in tf.paragraphs[1:]:
        para._p.getparent().remove(para._p)
    for r in list(tf.paragraphs[0].runs):
        r._r.getparent().remove(r._r)


def put_text(tf, paras, anchor="t", ml=0.14, mr=0.14, mt=0.06, mb=0.06, wrap=True):
    reset_tf(tf)
    tf.word_wrap = wrap
    tf.margin_left, tf.margin_right = Cm(ml), Cm(mr)
    tf.margin_top, tf.margin_bottom = Cm(mt), Cm(mb)
    tf.vertical_anchor = ANCH[anchor]
    first = True
    for p in paras:
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        para.alignment = ALIGN[p.get("align", "l")]
        if p.get("sa") is not None:
            para.space_after = Pt(p["sa"])
        if p.get("ls") is not None:
            para.line_spacing = p["ls"]
        for t, sz, b, c in p["runs"]:
            r = para.add_run()
            r.text = t
            set_font(r, sz, b, c)
    return tf


def T(slide, x, y, w, h, paras, anchor="t", **kw):
    b = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    put_text(b.text_frame, paras, anchor=anchor, **kw)
    return b


def one(text, sz, b=None, c=INK, align="l", sa=None, ls=None):
    d = {"runs": [(text, sz, b, c)], "align": align}
    if sa is not None:
        d["sa"] = sa
    if ls is not None:
        d["ls"] = ls
    return d


def multi(runs, align="l", sa=None, ls=None):
    d = {"runs": runs, "align": align}
    if sa is not None:
        d["sa"] = sa
    if ls is not None:
        d["ls"] = ls
    return d


def box(slide, x, y, w, h, fill=None, line=None, lw=1.0,
        shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06, dash=None):
    sp = slide.shapes.add_shape(shape, Cm(x), Cm(y), Cm(w), Cm(h))
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = RGBColor.from_string(line)
        sp.line.width = Pt(lw)
        if dash:
            sp.line.dash_style = dash
    sp.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sp.adjustments[0] = radius
        except Exception:
            pass
    reset_tf(sp.text_frame)
    return sp


def card(slide, x, y, w, h, head, body, hcol=NAVY, fill=PALE,
         hsz=12, bsz=11, line=None, anchor="t", ls=1.20):
    sp = box(slide, x, y, w, h, fill=fill, line=line)
    paras = [one(head, hsz, True, hcol, sa=4)]
    for b in (body if isinstance(body, list) else [body]):
        if b:
            paras.append(one(b, bsz, None, INK, ls=ls, sa=2))
    put_text(sp.text_frame, paras, anchor=anchor, ml=0.24, mr=0.20, mt=0.16, mb=0.12)
    return sp


def clear_slide(slide):
    spTree = slide.shapes._spTree
    for el in list(spTree):
        if el.tag.split("}")[-1] in ("sp", "cxnSp", "pic", "graphicFrame", "grpSp"):
            spTree.remove(el)


def frame(slide, title, lead):
    """ジムver3準拠：タイトル16pt太字 #1F2A5C（y=0.30）／リード15pt #646C82（y=2.00・2行以内）"""
    clear_slide(slide)
    T(slide, *TITLE_XY, [one(title, 16, True, TNAVY)], anchor="m", ml=0, mr=0)
    T(slide, *LEAD_XY, [one(l, 15, None, LEADC, ls=1.22) for l in lead],
      anchor="m", ml=0, mr=0)


def foot(slide, text, y=None):
    """脚注。3行以上は FOOT_Y のままだとフッター帯（y≒17.8）に食い込むので上へ逃がす。"""
    lines = text.split("\n")
    if y is None:
        y = FOOT_Y - 0.52 * max(0, len(lines) - 2)
    T(slide, CX0, y, CW, 0.52 * len(lines) + 0.4,
      [one(l, 9, None, LEADC, ls=1.20) for l in lines], ml=0, mr=0)


def band(slide, y, text, col=INK, sz=14):
    """ジムver3の結論バンド：左に細い縦棒＋薄い帯＋本文"""
    box(slide, CX0, y, 0.20, 1.30, fill=ORANGE, shape=MSO_SHAPE.RECTANGLE)
    box(slide, CX0 + 0.10, y, CW - 0.20, 1.30, fill=GREY, shape=MSO_SHAPE.RECTANGLE)
    T(slide, CX0 + 0.50, y + 0.22, CW - 0.90, 0.90,
      [one(text, sz, True, col)], anchor="m", ml=0, mr=0)


def badge(slide, x, y, w, h, text, fill=PORANGE, col=ORANGE, sz=10, dash=None):
    sp = box(slide, x, y, w, h, fill=fill, line=col, lw=1.0, dash=dash)
    put_text(sp.text_frame, [one(text, sz, True, col, align="c")],
             anchor="m", ml=0.06, mr=0.06, mt=0, mb=0)
    return sp


def placeholder(slide, x, y, w, h, label, note):
    """★未取得データの差込枠（破線）。捏造しない。"""
    sp = box(slide, x, y, w, h, fill=WHITE, line=MUT, lw=1.25,
             dash=MSO_LINE_DASH_STYLE.DASH)
    paras = [one(label, 12, True, MUT, align="c", sa=5)]
    paras += [one(l, 10, None, MUT, align="c", ls=1.25) for l in note.split("\n")]
    put_text(sp.text_frame, paras, anchor="m", ml=0.3, mr=0.3, mt=0.1, mb=0.1)
    return sp


def table(slide, x, y, w, h, headers, rows, col_w=None,
          hsz=12, bsz=11, header_fill=NAVY, zebra=PALE, align=None, row_h=None,
          cell_col=None):
    n_r, n_c = len(rows) + 1, len(headers)
    gf = slide.shapes.add_table(n_r, n_c, Cm(x), Cm(y), Cm(w), Cm(h))
    tbl = gf.table
    if col_w:
        for i, cw in enumerate(col_w):
            tbl.columns[i].width = Cm(cw)
    if row_h:
        for r in tbl.rows:
            r.height = Cm(row_h)
    for j, htext in enumerate(headers):
        c = tbl.cell(0, j)
        c.fill.solid()
        c.fill.fore_color.rgb = RGBColor.from_string(header_fill)
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
        put_text(c.text_frame, [one(htext, hsz, True, WHITE,
                                    align=(align[j] if align else "c"))],
                 anchor="m", ml=0.14, mr=0.14, mt=0.03, mb=0.03)
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            c = tbl.cell(i, j)
            c.fill.solid()
            c.fill.fore_color.rgb = RGBColor.from_string(
                zebra if (zebra and i % 2 == 0) else WHITE)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            col = INK
            bold = None
            if cell_col and (i, j) in cell_col:
                col, bold = cell_col[(i, j)]
            put_text(c.text_frame,
                     [one(l, bsz, bold, col, align=(align[j] if align else "l"))
                      for l in str(val).split("\n")],
                     anchor="m", ml=0.14, mr=0.14, mt=0.03, mb=0.03)
    return gf


def pic(slide, name, x, y, w=None, h=None, note=""):
    p = IMG / name
    if not p.exists():
        placeholder(slide, x, y, w or 10, h or 5, "★ 画像未生成", name)
        return None
    kw = {}
    if w:
        kw["width"] = Cm(w)
    if h:
        kw["height"] = Cm(h)
    return slide.shapes.add_picture(str(p), Cm(x), Cm(y), **kw)


# ---- 既存ページの文言差し替え（デザインを壊さない）----
def all_shapes(container):
    for sh in container.shapes:
        yield sh
        if sh.shape_type == 6:  # GROUP
            for s2 in all_shapes(sh):
                yield s2


def by_id(slide, sid):
    for sh in all_shapes(slide):
        if sh.shape_id == sid:
            return sh
    raise KeyError(f"shape id={sid} not found")


def swap(slide, sid, text):
    """先頭runの書式を引き継いだまま本文を差し替える（改行は段落で表現）"""
    tf = by_id(slide, sid).text_frame
    p0 = tf.paragraphs[0]
    if not p0.runs:
        raise ValueError(f"shape id={sid} has no run to inherit")
    for para in list(tf.paragraphs[1:]):
        para._p.getparent().remove(para._p)
    for r in list(p0.runs[1:]):
        r._r.getparent().remove(r._r)
    lines = text.split("\n")
    p0.runs[0].text = lines[0]
    from copy import deepcopy
    from pptx.text.text import _Paragraph
    tmpl = deepcopy(p0._p)
    for line in lines[1:]:
        new_p = deepcopy(tmpl)
        p0._p.getparent().append(new_p)
        _Paragraph(new_p, tf).runs[0].text = line


def swap_cell(slide, sid, r, c, text):
    tbl = by_id(slide, sid).table
    cell = tbl.cell(r, c)
    tf = cell.text_frame
    p0 = tf.paragraphs[0]
    if not p0.runs:
        put_text(tf, [one(l, 11, None, INK) for l in text.split("\n")], anchor="m",
                 ml=0.14, mr=0.14, mt=0.03, mb=0.03)
        return
    for para in list(tf.paragraphs[1:]):
        para._p.getparent().remove(para._p)
    for run in list(p0.runs[1:]):
        run._r.getparent().remove(run._r)
    from copy import deepcopy
    from pptx.text.text import _Paragraph
    lines = text.split("\n")
    p0.runs[0].text = lines[0]
    tmpl = deepcopy(p0._p)
    for line in lines[1:]:
        new_p = deepcopy(tmpl)
        p0._p.getparent().append(new_p)
        _Paragraph(new_p, tf).runs[0].text = line


S = slides  # 0-based

# ============================================================
# p1 表紙
# ============================================================
swap(S[0], 7, "業界：学校法人（大学・専門学校）")

# ============================================================
# p2 アジェンダ（章名はそのまま。学校法人の副題だけ添える）
# ============================================================
swap(S[1], 7,
     "1. LINE公式アカウントを検討する背景\n"
     "2. 配信の考え方 ── 市場・需要の実態\n"
     "3. 施策の詳細（配信の効率化）── 設計と改善モデル\n"
     "4. 具体プラン提示\n"
     "5. ご提案にあたり（施策紹介・企画案）")

# ============================================================
# p3 検討背景ファネル
# ============================================================
swap(S[2], 4,
     "資料請求・オープンキャンパス予約に対して、”ハードルが高く感じる”高校生が増えている状況。\n"
     "離脱防止 >> 友だち追加 >> LINEで学部学科診断・OC予約 >> 出願 の導線設計でハードルを抑える。")
swap(S[2], 23, "「LINE学部診断・OC予約」にてLINEリードを獲得")
swap(S[2], 31, "志望領域・通学圏\nを確認")
swap(S[2], 47, "入学後の定着・中退防止\n在学生／保護者／卒業生へ")
swap(S[2], 59,
     "※ 高校生の進路検討は約18ヶ月（高2春〜高3秋）。募集のチャンスは年1回で、"
     "取りこぼした1人は翌年戻らない（1名の入学＝4年間で約400万円）")

# ============================================================
# p5 流入元 × LINE誘導 × 具体の施策方針
# ============================================================
swap(S[4], 86, "あいさつMSG＋\n特典(診断) 自動送付 / 30秒 学部学科診断")
swap(S[4], 98, "ステップ配信（7〜14通）\n→ オープンキャンパス予約への誘導")
swap(S[4], 116, "スクリーニングアンケート\n（4問）→ 学年・志望度で分類")
swap(S[4], 119, "サンクス自動MSG＋\nOC当日の流れ・持ち物案内")
swap(S[4], 120, "前日リマインド配信\n出願スケジュール案内")
swap(S[4], 122, "スクリーニング（4問）\n志望度・通学圏の把握")
swap(S[4], 123, "有効リード育成シナリオ\n→ OC参加・出願への誘導")
swap(S[4], 125, "一斉配信でリスト活性化\nアンケートで学年・志望把握")
swap(S[4], 126, "セグメント別ナーチャリング\n入試方式別・時期別の企画配信")
swap(S[4], 128, "LINE移行キャンペーン\n(資料請求者・OC参加者を友だち化)")
swap(S[4], 129, "入学後もLINEで接点維持\n中退防止・在学生・卒業生へ")

# ============================================================
# p8 企画投稿 × ステップ配信
# ============================================================
swap(S[7], 76,
     "◯LINE友だち追加によるメリット提示\n✓デジタルパンフ　✓OC優先予約\n\n"
     "◯LINE上でのコミュニケーション誘導\n✓LINEで学部診断　✓LINEでOC予約")

print("[1/4] 文言差し替え（p1,2,3,5,8）完了")

# ============================================================
# p9 現状分析｜大学・専門学校のLINE公式アカウント実態
#    ★友だち数は未実測。page.line.me をPCで実査するまで差込枠のまま
# ============================================================
s = S[8]
frame(s, "現状分析｜大学・専門学校のLINE公式アカウント実態",
      ["友だち数は page.line.me（LINEヤフー社公式ページ）で実測する。取得日を必ず併記する。"])
table(s, CX0, CY0, CW, 7.6,
      ["調査対象", "区分", "受験生向けか", "友だち数 実測", "リッチメニュー", "診断の有無"],
      [["大規模総合大（首都圏）", "大学", "―", "―", "―", "―"],
       ["大規模総合大（関西）", "大学", "―", "―", "―", "―"],
       ["中規模私立大（地方）", "大学", "―", "―", "―", "―"],
       ["女子大", "大学", "―", "―", "―", "―"],
       ["大手専門学校（ビジネス系）", "専門", "―", "―", "―", "―"],
       ["大手専門学校（IT・クリエイティブ）", "専門", "―", "―", "―", "―"],
       ["医療・看護系専門学校", "専門", "―", "―", "―", "―"],
       ["通信制高校グループ", "高校", "―", "―", "―", "―"]],
      col_w=[6.4, 2.2, 3.6, 4.6, 4.6, 4.0],
      align=["l", "c", "c", "r", "c", "c"], row_h=0.84)
card(s, CX0, 11.30, 8.20, 3.70, "① 大学は「在学生向け」が多い",
     ["受験生向けと在学生向けを分けているか、",
      "タグで出し分けているかを必ず確認する。",
      "混在していると配信精度が落ちる。"], fill=PALE)
card(s, CX0 + 8.60, 11.30, 8.20, 3.70, "② 見るのは友だち数ではない",
     ["学部数が多いほど「出し分け」の必要が高い。",
      "リッチメニューが学部別か機能別かで、",
      "その学校の設計思想が分かる。"], fill=PALE)
card(s, CX0 + 17.20, 11.30, 8.20, 3.70, "③ 空白は「診断」",
     ["塾業界の実測では、大手8社すべてが",
      "LINEを新規集客に使っていた。差がつくのは",
      "「診断で適合を見せているか」だけだった。"], fill=PORANGE, hcol=ORANGE)
band(s, BAND_Y, "友だちは広告で買える。でも運用設計がなければ、リストは資産にならない。")
foot(s, "※ 友だち数は未実測（2026-09-16時点）。page.line.me（公式）または実機で取得し、取得日を併記する。\n"
        "※ 第三者サイト（lineoa.jp）の数値は当たり付け専用。本資料には載せない。実名掲載の最終可否は学内確認事項。")

# ============================================================
# p10 現状分析｜LINE運用は4つの型に分類できる
# ============================================================
s = S[9]
frame(s, "現状分析｜LINE運用は4つの型に分類できる",
      ["「集めている学校」と「育てている学校」に、はっきり分かれる。"])
table(s, CX0, CY0, CW, 9.30,
      ["型", "運用の中身", "強み／弱み"],
      [["① 診断ドリブン型",
        "学部学科診断・適性診断で属性を取得し、学年・志望度別に配信を出し分ける",
        "本提案の型。実装している学校が少なく、空白地帯"],
       ["② 相談カウンター型",
        "有人チャットでの個別相談が中心。配信はほとんどしない",
        "相談に来た子には強い。来ない子は拾えない"],
       ["③ カタログ型",
        "学部一覧・OC日程・資料請求のリンクを並べただけ。押すとWebへ飛ぶ",
        "最も多い。LINEを使っているのではなくリンク集にしている"],
       ["④ 連絡網型",
        "在学生への休講連絡・事務連絡が中心。受験生向けには使っていない",
        "友だち数は多いが、募集には一切効いていない"]],
      col_w=[5.2, 12.0, 8.2], align=["l", "l", "l"], row_h=2.20)
band(s, BAND_Y, "目指すのは ①診断 × 出し分け × 入学後まで の三点セット。どれか1つだけでは必ず頭打ちになる。")
foot(s, "※ 型分類は塾業界（2026-09-15実測）で検証したフレーム。学校法人での各校の当てはめは実機調査（page.line.me・PC実査）で確定する。")

# ============================================================
# p11 現状分析｜各校のリッチメニュー・あいさつメッセージ
# ============================================================
s = S[10]
frame(s, "現状分析｜各校のリッチメニュー・あいさつメッセージ",
      ["実機の画面を並べると、型の違いが一目で伝わる。"])
labels = [("大規模総合大（首都圏）", "学部別か機能別か"),
          ("大規模総合大（関西）", "受験生／在学生の分離"),
          ("中規模私立大（地方）", "本モデル校に最も近い"),
          ("大手専門学校", "体験入学の予約導線"),
          ("医療・看護系専門", "保護者向け情報の出し方")]
for i, (nm, note) in enumerate(labels):
    x = CX0 + i * 5.10
    badge(s, x, CY0, 4.70, 1.10, nm, fill=PALE, col=NAVY, sz=10)
    placeholder(s, x, CY0 + 1.20, 4.70, 8.60, "▧ 実機トーク画面\n差込予定",
                "・リッチメニュー全体\n・あいさつメッセージ\n・2週間分の配信")
    badge(s, x, CY0 + 10.00, 4.70, 1.40, note, fill=GREY, col=INK, sz=10)
band(s, BAND_Y, "撮影は調査用の捨てアカウントで。取得日を画面ごとに記録する。")
foot(s, "※ 撮影対象：リッチメニュー全体／友だち追加直後のあいさつメッセージ／2週間分の配信トーク画面。\n"
        "※ クラウド環境からは page.line.me が参照できないため、PC側で取得する。")

# ============================================================
# p14 8フェーズを学校法人に置き換える
# ============================================================
swap(S[13], 3, "カスタマージャーニーマップにあわせてLINE投稿を配置。\n"
                "※一般的な8フェーズを、学校法人（大学・専門学校）に置き換えます。")
swap(S[13], 80, "④資料請求・友だち追加(web/LINE)")
swap(S[13], 70, "⑤ オープンキャンパス参加")
swap(S[13], 72, "⑦出願・入学手続")
swap(S[13], 76, "⑧入学後の定着・中退防止\n　在学生／保護者／卒業生 等")
swap(S[13], 83, "← --- 学校法人（大学・専門学校）の場合 --- →")

# ============================================================
# p15 市場の構造｜母数は減り、獲得コストだけが上がる
# ============================================================
s = S[14]
frame(s, "市場の構造｜18歳人口は30年で半減。私立大の6割が定員割れ",
      ["母数が減っても学校は減らない。＝1人を取り合う競争だけが激化する構造。",
       "だから入札で勝つ設計ではなく、取りこぼしを拾う設計に寄せる。"])
stats = [("【18歳人口】\n205万人 → 106万人\n（1992年 → 2024年）\n⇊\n30年で ▲48%", RED, PRED),
         ("【私立大の定員割れ】\n59.2%（354校／598校）\n（2024年度）\n⇊\n6割が定員を埋められない", RED, PRED),
         ("【大学進学率】\n約59%（過去最高圏）\n⇊\n進学率の伸びで\n母数減を吸収する余地はない", INK, GREY),
         ("【1名の入学の重さ】\n初年度納付金 約136万円\n4年間で 約400万円\n⇊\n取りこぼしの単価が桁違い", ORANGE, PORANGE)]
for i, (txt, col, fl) in enumerate(stats):
    x = CX0 + i * 6.42
    sp = box(s, x, CY0 + 0.30, 6.00, 6.40, fill=fl, line=col, lw=1.25)
    put_text(sp.text_frame, [one(l, 12.5, True, col, align="c", ls=1.30)
                             for l in txt.split("\n")],
             anchor="m", ml=0.18, mr=0.18, mt=0.12, mb=0.12)
card(s, CX0, 10.60, 12.40, 4.30, "母数の減り方は、これから加速する",
     ["2040年の18歳人口は80万人台まで減る見通し（中央教育審議会の推計）。",
      "いま在籍している小中高生の数がそのまま波及するため、",
      "競争密度は今がピークではない。"], fill=PALE)
card(s, CX0 + 13.00, 10.60, 12.40, 4.30, "だから「取りこぼしを拾う」が最も効く",
     ["広告の入札を強めても、母数が半分の市場では単価が上がるだけ。",
      "一度でも手を挙げた高校生（資料請求者・OC参加者）を",
      "出願まで落とさない仕組みのほうが、費用対効果が高い。"], fill=PORANGE, hcol=ORANGE)
band(s, BAND_Y, "入札で勝つ設計ではなく、取りこぼしを拾う設計へ。")
foot(s, "※ 出典：文部科学省「学校基本調査」／日本私立学校振興・共済事業団「私立大学・短期大学等入学志願動向」／"
        "文部科学省「私立大学等入学者に係る初年度学生納付金等調査」。【提出前に最新年度で原典照合すること】\n"
        "※ 2040年推計は中央教育審議会の答申の版によって82万人／88万人と異なる。引用する版を確定してから提出する。")

# ============================================================
# p16 ニーズ調査（シーズナリティ）｜勝負は年内に終わる
# ============================================================
s = S[15]
frame(s, "ニーズ調査（シーズナリティ）｜OCは7月末、総合型選抜は8月末",
      ["オープンキャンパスのピークと、総合型選抜の出願準備のピークは1ヶ月しか離れていない。",
       "＝ 夏のOCで接点を持てなかった学校は、8月末には候補から消えている。"])
pic(s, "gakkou_trend_oc_weekly.png", CX0, CY0, w=12.40)
pic(s, "gakkou_trend_nyushi_weekly.png", CX0 + 12.90, CY0, w=12.40)
cards16 = [("① 動き出しは5月下旬",
            ["「オープンキャンパス」の検索が20を超えるのが5/25週。",
             "6/29週にはもう年間ピークの半分（51）。",
             "6月に広報を始めたら、もう遅い。"]),
           ("② 総合型のピークは8/31週",
            ["9/1の出願解禁を含む週が年間最大（100）。",
             "指定校推薦は8月末と11月中旬の二山、",
             "学校推薦型は11月の単峰。3方式で時期が違う。"]),
           ("③ 仕込みの窓は5週間",
            ["5/25週〜6/29週の5週間が、友だち化の窓。",
             "ここまでに友だちになっていない高校生は、",
             "夏のOCの候補リストに入らない。"])]
for i, (h, b) in enumerate(cards16):
    card(s, CX0 + i * 8.60, 8.80, 8.20, 6.00, h, b,
         fill=PORANGE if i == 2 else PALE, hcol=ORANGE if i == 2 else NAVY)
foot(s, "※ 出典：Googleトレンド（日本・ウェブ検索・2026-09-15取得）。週次データは2025年（2024-12-29〜2025-12-28）の実測。\n"
        "※ 入試日程は文部科学省「大学入学者選抜実施要項」（総合型＝9/1出願・11/1発表／学校推薦型＝11/1出願・12/1発表／一般＝2/1〜）。")

# ============================================================
# p17 ニーズ調査（5年トレンド）｜総合型選抜が主戦場になった
# ============================================================
s = S[16]
frame(s, "ニーズ調査（5年トレンド）｜総合型選抜は5年で3.3倍",
      ["オープンキャンパスの形は5年間まったく変わらない＝予定を立てられる需要。",
       "一方で入試方式の主役は入れ替わっており、2023年9月に総合型が指定校推薦を抜いた。"])
pic(s, "gakkou_trend_oc_5y.png", CX0, CY0, w=12.40)
pic(s, "gakkou_trend_nyushi_5y.png", CX0 + 12.90, CY0, w=12.40)
card(s, CX0, 8.80, 8.20, 6.00, "① OCは5年連続で7月ピーク",
     ["2022→2026年の7月値は 100 / 94 / 80 / 85 / 90。",
      "谷は12〜1月（3〜5）でピークの約1/20。",
      "形が変わらない＝毎年同じ日程で仕込める。"], fill=PALE)
card(s, CX0 + 8.60, 8.80, 8.20, 6.00, "② 総合型選抜 5年で3.3倍",
     ["9月値の比較：総合型 30 → 100（3.3倍・5年連続増）。",
      "指定校推薦は 64 → 57（▲11%）で横ばい。",
      "2023年9月に逆転し、いまは総合型が1.75倍。"], fill=PORANGE, hcol=ORANGE)
card(s, CX0 + 17.20, 8.80, 8.20, 6.00, "③ 専門学校は年中フラット",
     ["「専門学校」は5年間ずっと22〜40のレンジ。",
      "大学（夏に集中）とは需要曲線がまったく違う。",
      "→ 大学と専門学校は別の配信設計が要る。"], fill=PALE)
foot(s, "※ 出典：Googleトレンド（日本・ウェブ検索・2021/9〜2026/9・2026-09-15取得）。2026年9月は月途中（9/15まで）の参考値。\n"
        "※ 「共通テスト」は1月に100・他月は10以下と桁が違うため、同一グラフに混ぜていない（混ぜると小さい方が潰れる）。")

# ============================================================
# p18 年間の配信設計マトリクス（時期 × 対象 × 訴求 × CV）
# ============================================================
s = S[17]
frame(s, "年間の配信設計マトリクス（時期 × 対象 × 訴求 × CV）",
      ["需要の波と入試日程が動かせない業界。時期ごとに「誰に・何を言い・どこへ着地させるか」を先に決める。"])
table(s, CX0, CY0, CW, 10.60,
      ["時期", "対象（学年・状態）", "訴求", "着地させるCV"],
      [["3〜4月", "新高2・新高3／進路意識が立ち上がる層",
        "「まだ決まっていなくていい」30秒 学部学科診断", "友だち追加・診断完了【CV①】"],
       ["5月下旬〜6月", "高3／OC検索が動き出す層　※検索実測：5/25週",
        "OC日程の先行案内・1人参加OK・当日の流れ", "オープンキャンパス予約【CV②】"],
       ["7〜8月", "OC予約者・参加者",
        "前日リマインド／参加後アンケート／学部の深掘り", "OC参加・再訪【CV②】"],
       ["9〜11月", "OC参加者・年内入試層　※総合型9/1・学校推薦型11/1",
        "志望理由書の書き方・面接10問・締切カウントダウン", "出願【CV③】"],
       ["12〜2月", "一般選抜層・併願検討層　※共通テスト1月中旬",
        "入試方式の違い・過去問・入試説明会", "出願【CV③】"],
       ["2〜3月", "合格者・入学手続前",
        "入学前オリエンテーション・先輩との交流・一人暮らし準備", "入学手続【CV③】・定着"]],
      col_w=[3.0, 7.4, 9.4, 5.6], align=["c", "l", "l", "l"], row_h=1.50)
foot(s, "※ 時期の根拠：Googleトレンド実測（2026-09-15取得）＋文部科学省「大学入学者選抜実施要項」。\n"
        "※ 専門学校はAOエントリー6/1・出願10/1解禁（全国専修学校各種学校総連合会の申し合わせ・要確認）のため、上記より前倒しになる。\n"
        "※ 各時期の広告単価・CVRは貴学の実績値で埋める（進学領域の公開統計は存在しない）。")

# ============================================================
# p19 前後検索｜検討は「偏差値」で動き、オープンキャンパスは後に来る
# ============================================================
s = S[18]
frame(s, "前後検索｜検討は「偏差値」で動く。オープンキャンパスは「後」に来る",
      ["LINEヤフー media Journey の実データ。0日前後の密集帯は、ほぼ全部「大学名＋偏差値」だった。",
       "学部でも、やりたいことでもない。"])
STG = [("検索前（−15〜−4日）", "上位〜中位の大学名", GREY, INK,
        ["東洋大学／立教大学／中央大学／明治大学",
         "明治学院大学／學習院大學／成蹊大学",
         "日東駒専／法政大学 偏差値／駒澤大学 偏差値",
         "成成明学 序列／オープンキャンパス 東京"]),
       ("0日（起点付近）", "中堅〜下位の大学名＋偏差値", PORANGE, ORANGE,
        ["東京国際大学 偏差値／亜細亜大学",
         "帝京平成大学／東京経済大学／駿河台大学",
         "大東亜帝国／共栄大学／聖学院大学",
         "定員割れの大学一覧（−2日）"]),
       ("検索後（+1〜+15日）", "オープンキャンパス・学部・キャンパス", PALE, NAVY,
        ["東洋大学 オープンキャンパス2025",
         "成蹊大学・日本大学 オープンキャンパス",
         "東洋大学 学部／専修大学 生田キャンパス",
         "城西大学駅伝部／国士舘大学 野球部（スポーツ推薦の導線）"])]
for i, (ttl, sub, fl, col, kws) in enumerate(STG):
    x = CX0 + i * 8.60
    sp = box(s, x, CY0, 8.20, 7.60, fill=fl, line=col, lw=1.25,
             dash=MSO_LINE_DASH_STYLE.DASH if i == 0 else None)
    paras = [one(ttl, 12, True, col, align="c", sa=3),
             one(sub, 11, True, INK, align="c", sa=7)]
    paras += [one("・" + k, 10.5, None, INK, ls=1.30, sa=2) for k in kws]
    put_text(sp.text_frame, paras, anchor="t", ml=0.24, mr=0.20, mt=0.20, mb=0.12)
# 友だち化マーカーは1列目の内側（下端）に置く。中央カラムに重ねない
box(s, CX0 + 0.35, CY0 + 6.30, 7.50, 1.00, fill=GREEN, radius=0.10)
T(s, CX0 + 0.35, CY0 + 6.30, 7.50, 1.00,
  [one("▶ 仕掛けるのはここ（友だち化）", 12, True, WHITE, align="c")],
  anchor="m", ml=0, mr=0)
card(s, CX0, 11.40, 12.40, 3.60, "① 志望校は「下方修正の過程」で決まる",
     ["−7日＝MARCH・成成明学 → 0日＝大東亜帝国・地方中小私大。",
      "時間が経つほど偏差値帯が下がっていく。上から順に降りてきている。"], fill=PALE)
card(s, CX0 + 13.00, 11.40, 12.40, 3.60, "② OCは入口ではなく「絞り込み後の確認」",
     ["大学名のOC検索が立つのは +1〜+7日。",
      "＝ OC予約を取る勝負は、偏差値で群に仕分けられる「前」に始まっている。"],
     fill=PORANGE, hcol=ORANGE)
foot(s, "※ 出典：LINEヤフー media Journey（2026-09-15取得）。【起点KWが未確認。確定後に本ページを更新する】\n"
        "※ 地域は埼玉西部〜北部の高校生が中心（川越・所沢・大宮・春日部・熊谷・坂戸・東松山）＝通学圏が検討の主軸であることの裏付け。\n"
        "※ 競合実名を含むため、社外提示は学内確認のうえ。")

# ============================================================
# p20 前後検索（補足）｜お金の確認は、決めた後に来る
# ============================================================
s = S[19]
frame(s, "前後検索（補足）｜学費・奨学金の検索は、決めた「後」に来る",
      ["教育3業界（塾・通信制高校・大学）の前後検索で、同じ構造が確認できた。",
       "＝ 入口の訴求を「安さ・学費」に置くと、届く前に終わる。"])
table(s, CX0, CY0, CW, 5.60,
      ["業界", "起点KW", "お金の検索が立つ位置", "示唆"],
      [["教育（塾）", "塾 費用", "+1.6日（決断の後）",
        "最初の問いは「いくら」ではなく「いつから」。関連度1位は「塾 いつから通わせる」4.98"],
       ["通信制高校", "通信制高校", "学校名の直後（+1日以降）",
        "前半15日は学校選びですらない（不登校・起立性調節障害＝医療と心理の悩み）"],
       ["大学", "（起点KW確認中）", "学校名を調べた後",
        "0日前後は「大学名＋偏差値」。学費は絞り込みの後に確認される"]],
      col_w=[3.6, 4.4, 5.4, 12.0], align=["c", "c", "c", "l"], row_h=1.70)
card(s, CX0, 9.80, 12.40, 5.00, "打ち手：入口は「いつ・自分に合うか」で開く",
     ["入口＝30秒の学部学科診断（名前も電話番号も入力しない）。",
      "学費・奨学金は、診断結果の「後半」に置く。",
      "先に費用を出すと、比較の土俵に乗る前に離脱する。"], fill=PORANGE, hcol=ORANGE)
card(s, CX0 + 13.00, 9.80, 12.40, 5.00, "ただし保護者には「先に」出す",
     ["この業界は決める人（高校生）と払う人（保護者）が違う。",
      "リッチメニューの保護者タブは、左上に",
      "「4年間の学費シミュレーション」を置く。順番が逆になる。"], fill=PALE)
band(s, BAND_Y, "本人には「合うか」を、保護者には「総額」を。同じアカウントで、出し分ける。")
foot(s, "※ 出典：LINEヤフー media Journey（塾＝2026-09-15取得・起点「塾 費用」／通信制高校・大学＝2026-09-15取得）。\n"
        "※ 大学は起点KWが未確認。確定後に本ページの数値を更新する。")

print("[2/4] 市場・データ章（p9,10,11,14,15,16,17,18,19,20）完了")


# ---- LINE UI モック用ヘルパー ----
def phone(slide, x, y, w, h, header, lines):
    """LINEトーク画面のモック。lines=[(kind,text)] kind: in/chip/btn/note"""
    box(slide, x, y, w, h, fill=BEZEL, radius=0.10)
    box(slide, x + 0.16, y + 0.46, w - 0.32, h - 0.62, fill=SCREEN, radius=0.02,
        shape=MSO_SHAPE.RECTANGLE)
    hd = box(slide, x + 0.16, y + 0.46, w - 0.32, 0.58, fill=GREEN,
             shape=MSO_SHAPE.RECTANGLE)
    put_text(hd.text_frame, [one(header, 8, True, WHITE, align="c")],
             anchor="m", ml=0.05, mr=0.05, mt=0, mb=0)
    cy = y + 1.16
    iw = w - 0.70
    for kind, text in lines:
        n = text.count("\n") + 1
        bh = 0.26 + 0.34 * n
        if kind == "in":
            sp = box(slide, x + 0.35, cy, iw, bh, fill=WHITE, line=BORDER, lw=0.75)
            put_text(sp.text_frame, [one(l, 8, None, INK, ls=1.16)
                                     for l in text.split("\n")],
                     anchor="m", ml=0.16, mr=0.12, mt=0.05, mb=0.05)
        elif kind == "chip":
            sp = box(slide, x + 0.35, cy, iw, bh, fill=WHITE, line=GREEN, lw=0.9)
            put_text(sp.text_frame, [one(l, 8, None, "0B7A3B", align="c", ls=1.14)
                                     for l in text.split("\n")],
                     anchor="m", ml=0.08, mr=0.08, mt=0.03, mb=0.03)
        elif kind == "btn":
            sp = box(slide, x + 0.35, cy, iw, bh, fill=GREEN)
            put_text(sp.text_frame, [one(l, 8.5, True, WHITE, align="c", ls=1.14)
                                     for l in text.split("\n")],
                     anchor="m", ml=0.08, mr=0.08, mt=0.03, mb=0.03)
        else:
            sp = slide.shapes.add_textbox(Cm(x + 0.35), Cm(cy), Cm(iw), Cm(bh))
            put_text(sp.text_frame, [one(l, 7.5, None, MUT, ls=1.14)
                                     for l in text.split("\n")],
                     anchor="m", ml=0.05, mr=0.05, mt=0, mb=0)
        cy += bh + 0.13
    return cy


def richmenu(slide, x, y, w, h, tab_names, active, btns, hi=0):
    """リッチメニュー1面（3列×2行＝6枠）＋上部タブ。hi=強調するボタンindex"""
    tab_h = 0.62
    box(slide, x, y, w, h, fill=BEZEL, radius=0.06)
    tw = (w - 0.14) / len(tab_names)
    for i, nm in enumerate(tab_names):
        on = (i == active)
        tb = box(slide, x + 0.07 + i * tw, y + 0.07, tw - 0.04, tab_h,
                 fill=GREEN if on else "3F3F3F", shape=MSO_SHAPE.RECTANGLE)
        put_text(tb.text_frame, [one(nm, 8.5, True, WHITE, align="c")],
                 anchor="m", ml=0, mr=0, mt=0, mb=0)
    gx, gy = x + 0.07, y + 0.07 + tab_h + 0.07
    gw, gh = w - 0.14, h - tab_h - 0.21
    cols, rows = 3, 2
    bw = (gw - 0.07 * (cols - 1)) / cols
    bh = (gh - 0.07 * (rows - 1)) / rows
    for i, label in enumerate(btns):
        r, c = divmod(i, cols)
        on = (i == hi)
        bb = box(slide, gx + c * (bw + 0.07), gy + r * (bh + 0.07), bw, bh,
                 fill=PORANGE if on else SCREEN, line=ORANGE if on else BORDER,
                 lw=1.0 if on else 0.75, radius=0.04)
        put_text(bb.text_frame,
                 [one(l, 8.5, True, ORANGE if on else NAVY, align="c", ls=1.10)
                  for l in label.split("\n")],
                 anchor="m", ml=0.05, mr=0.05, mt=0, mb=0)


def swap_text(slide, mapping):
    """スライド内の全runを走査して文言を置換（書式は維持）。★数値には使わない"""
    n = 0
    for sh in all_shapes(slide):
        frames = []
        if sh.has_text_frame:
            frames.append(sh.text_frame)
        if sh.has_table:
            for row in sh.table.rows:
                for c in row.cells:
                    frames.append(c.text_frame)
        for tf in frames:
            for para in tf.paragraphs:
                for r in para.runs:
                    t = r.text
                    for a, b in mapping.items():
                        if a in t:
                            t = t.replace(a, b)
                    if t != r.text:
                        r.text = t
                        n += 1
    return n


# 8フェーズの業界語（p22・p32 で共通して使う）
PHASE_MAP = {
    "⑤ 来店・体験": "⑤ オープンキャンパス参加",
    "⑦ご入会": "⑦出願・入学手続",
    "⑧継続・PT/物販・友だち紹介": "⑧入学後の定着・中退防止",
    "退会防止・再入会 等": "在学生／保護者／卒業生 等",
    "④体験予約(web/LINE)": "④資料請求・友だち追加(web/LINE)",
    "④体験予約(LINE)": "④資料請求・友だち追加(LINE)",
    "④体験予約(WEB)": "④資料請求(WEB)",
    "ジム・フィットネス業界": "学校法人（大学・専門学校）",
    "[LINEで体づくりタイプ診断] から": "[LINEで学部学科診断] から",
    "体験予約/見学予約 から": "資料請求／OC予約 から",
}

# ============================================================
# p22 8フェーズを学校法人に当てはめる
# ============================================================
swap_text(S[21], PHASE_MAP)
swap(S[21], 4,
     "一般的な8フェーズを、学校法人（大学・専門学校）に当てはめるとこうなる。\n"
     "④資料請求と⑤オープンキャンパス参加の間が最も落ちる。ここに「匿名で聞ける」入口を置く。")
swap(S[21], 73, "どの学校が自分に合うか分からない\nまだ漠然。偏差値で調べている")
swap(S[21], 75, "電話がかかってこないか不安\n1人で行って大丈夫か知りたい")
swap(S[13], 8, "一般的な8フェーズを、学校法人（大学・専門学校）に置き換えるとこうなります。\n"
                "④資料請求・友だち追加 → ⑤オープンキャンパス参加 → ⑦出願・入学手続 が、この業界のCV地点です。")

# ============================================================
# p23 前後検索から導いた、友だち追加の仕掛けどころ
# ============================================================
s = S[22]
frame(s, "前後検索から導いた仕掛けどころ｜名前を調べられる前に友だちになる",
      ["高校生が学校名を検索し始めた時点で、もう比較の土俵。そこで広告に出ても1校にしかなれない。",
       "仕掛けるのは、偏差値で群に仕分けられる手前。"])
TLX, TLW, TLY = CX0 + 0.6, CW - 1.2, CY0 + 1.90
box(s, TLX, TLY, TLW, 1.10, fill=GREY, shape=MSO_SHAPE.RECTANGLE)
box(s, TLX, TLY, TLW * 0.46, 1.10, fill=PALE, line=NAVY, lw=1.25,
    shape=MSO_SHAPE.RECTANGLE, dash=MSO_LINE_DASH_STYLE.DASH)
box(s, TLX + TLW * 0.62, TLY, TLW * 0.38, 1.10, fill=PORANGE, line=ORANGE, lw=1.5,
    shape=MSO_SHAPE.RECTANGLE)
T(s, TLX, TLY + 0.18, TLW * 0.46, 0.80,
  [one("学校名がまだ出てこない期間（＝ここが空白）", 11.5, True, NAVY, align="c")],
  anchor="m", ml=0, mr=0)
T(s, TLX + TLW * 0.62, TLY + 0.18, TLW * 0.38, 0.80,
  [one("比較の土俵（もう1校にしかなれない）", 11.5, True, ORANGE, align="c")],
  anchor="m", ml=0, mr=0)
for frac, lab in [(0.0, "−15日"), (0.46, "0日"), (0.62, "+2日"), (1.0, "+15日")]:
    T(s, TLX + TLW * frac - 1.1, TLY + 1.22, 2.2, 0.60,
      [one(lab, 10.5, True, MUT, align="c")], anchor="m", ml=0, mr=0)
box(s, TLX + TLW * 0.34, TLY - 1.40, 5.00, 1.15, fill=GREEN, radius=0.10)
T(s, TLX + TLW * 0.34, TLY - 1.40, 5.00, 1.15,
  [one("▶ ここで友だち化", 12, True, WHITE, align="c")], anchor="m", ml=0, mr=0)
sets = [("① 離脱防止ポップアップの文言", PALE, NAVY,
         ["まだ、学部で迷っていませんか？",
          "30秒の質問に答えるだけで、あなたに向いている",
          "学部と、その学部の1日が分かります。",
          "＃名前も電話番号も入力しません"]),
        ("② LINE広告（CPF）の文言", PALE, NAVY,
         ["やりたいこと、まだ決まってなくていい。",
          "文系・理系すら迷っている人向けの30秒診断。",
          "あなたの「向いてる学び方」から学部を提案します。",
          "→ タップ1回で友だち追加（入力ゼロ）"]),
        ("③ コストゼロで一番効く導線", PORANGE, ORANGE,
         ["オープンキャンパス当日の受付にQRコードを1枚。",
          "最も志望度が高い層が、その場で友だちになる。",
          "資料請求の完了画面・パンフ裏表紙・高校での説明会資料も同じ。",
          "→ 来週から実施できる"])]
for i, (h, fl, col, body) in enumerate(sets):
    card(s, CX0 + i * 8.60, 9.60, 8.20, 5.20, h, body, fill=fl, hcol=col)
foot(s, "※ 出典：LINEヤフー media Journey（2026-09-15取得）。起点KWの確定後に日数の目盛りを更新する。")

# ============================================================
# p24 ステップ配信｜追加から14日でオープンキャンパス予約を取り切る
# ============================================================
s = S[23]
frame(s, "ステップ配信｜追加から14日でオープンキャンパス予約を取り切る",
      ["友だち追加から15日でアクティブ率が急落する。起点KW検索の後15日で他校の指名検索が立つ。",
       "読まれる15日と、候補が決まる15日が重なるのが最初の14日。"])
table(s, CX0, CY0, CW, 10.20,
      ["Day", "対象", "内容", "通知プレビュー（15字）", "CV"],
      [["0", "診断回答直後の全員", "診断結果＋在学生の1日（動画3分）。売らない",
        "診断結果が出ました📋あな", "―"],
       ["2", "Q2「全然決まっていない」層", "やりたいことが決まっていない人の進路の決め方",
        "やりたいこと、まだ決ま", "―"],
       ["3", "全員", "在学生インタビュー（高校時代の成績・迷い・いま）",
        "高2の冬、模試D判定だ", "―"],
       ["6", "全員", "OCの日程と当日の流れ。1人参加OKを明記",
        "オープンキャンパスの日程", "OC予約【CV②】"],
       ["9", "全員（保護者にも転送を明記）", "4年間の学費と奨学金",
        "4年間でいくらかかるの", "―"],
       ["11", "全員", "就職・進路データ（学部別の実績）",
        "卒業生はどこに行ったの", "―"],
       ["14", "未CV者", "残席状況＋オンライン相談も選べる（出口を3つ）",
        "8/3の残席が少なくなっ", "OC予約【CV②】"]],
      col_w=[1.6, 5.4, 9.4, 5.4, 3.6],
      align=["c", "l", "l", "l", "c"], row_h=1.25,
      cell_col={(4, 4): (ORANGE, True), (7, 4): (ORANGE, True)})
band(s, BAND_Y, "CVオファーは2回置く（Day6・Day14）。1回で終わらせない。出口は3つ用意する。")
foot(s, "※ 14日の根拠：①友だち追加から15日前後でアクティブ率が急落する　②起点KW検索の後15日で他校の指名検索が立つ（前後検索）。\n"
        "※ Day14の出口3つ＝①OCに行く　②オンラインで在学生に聞く（20分）　③学部診断をもう一度やる。行けない子を落とさないため。")

# ============================================================
# p25 ステップ配信｜入学後14日で「辞めない形」を作る（第2の提案軸）
# ============================================================
s = S[24]
frame(s, "ステップ配信｜入学後14日で「辞めない形」を作る（第2の提案軸）",
      ["中退の予兆は行動に出る。履修登録の遅れ・ガイダンス欠席・配信への無反応。",
       "募集で1人増やすのと、中退を1人止めるのは、経営上まったく同じ価値。"])
table(s, CX0, CY0, CW, 6.40,
      ["Day", "対象", "内容", "狙い"],
      [["0（入学手続後）", "入学予定者", "入学前オリエンテーション案内／同じ学部の先輩の1日",
        "「ついていけるか」の不安を先に潰す"],
       ["7", "入学予定者", "履修登録の仕方（3ステップ・つまずきやすい点つき）",
        "履修登録の遅れ＝中退の第1予兆を防ぐ"],
       ["14", "入学予定者・保護者", "一人暮らし準備／奨学金の手続き期限／学生相談窓口の案内",
        "保護者の不安も同時に外す"],
       ["以降・常時", "在学生（無反応が続く層）", "「最近どう？」の自動声かけ→学生相談へトスアップ",
        "離脱予兆検知。幽霊化する前に止める"]],
      col_w=[3.6, 5.4, 9.6, 6.8], align=["c", "l", "l", "l"], row_h=1.50)
card(s, CX0, 10.40, 12.40, 4.40, "中退1人の損失＝残存年数分の学費",
     ["大学の年間中退率は約2%。4年間では入学者の約8%（12人に1人）。",
      "2年次で中退なら、残り2年分＝約200〜300万円が消える。",
      "＝ 中退を1人止めることは、入学者を1人増やすのと同じ。"], fill=PRED, hcol=RED)
card(s, CX0 + 13.00, 10.40, 12.40, 4.40, "同じアカウントで、入学後もそのまま使える",
     ["入学前教育／休講・災害時の一斉連絡／保護者向け通知／",
      "就職支援／卒業生・寄付の案内まで1本で回せる。",
      "※受験生向けと在学生向けは、タグで分離する設計にする。"], fill=PALE)
band(s, BAND_Y, "入学前の14日でOC予約を取り、入学後の14日で「辞めない形」を作る。この2本で1セット。")
foot(s, "※ 中退率は文部科学省の調査に基づく概数。【提出前に最新調査で原典照合すること】\n"
        "※ 中退率・在籍年数はLINE側では計測できない。貴学の学務システムの数値と突き合わせて効果を見る。")

# ============================================================
# p26 リッチメニュー設計案｜3タブ × 6枠
# ============================================================
s = S[25]
frame(s, "リッチメニュー設計案｜3タブ × 6枠（18ボタン）",
      ["「受験生」「OC・入試」「保護者」で、見せる画面を変える。",
       "左上に置くのは資料請求ではない。まだ決めていない人が押せる場所。"])
table(s, CX0, CY0, CW, 9.60,
      ["位置", "① 受験生・はじめての方", "② オープンキャンパス・入試", "③ 保護者の方へ"],
      [["左上（最重要）", "30秒 学部学科診断", "次回OCの日程・予約", "4年間の学費シミュレーション"],
       ["右上", "オープンキャンパス予約", "当日の流れ・持ち物", "奨学金・学費支援制度"],
       ["左中", "学部・学科を見る", "アクセス・キャンパスマップ", "就職実績・進路データ"],
       ["右中", "資料（デジタルパンフ）をもらう", "入試方式を調べる", "保護者向け説明会"],
       ["左下", "まだ決まっていない人へ", "出願スケジュール", "一人暮らしの費用の目安"],
       ["右下", "在学生に質問する", "過去問・入試説明会", "個別相談を申し込む"]],
      col_w=[4.0, 7.2, 7.2, 7.0], align=["c", "l", "l", "l"], row_h=1.45,
      cell_col={(1, 1): (ORANGE, True), (1, 2): (ORANGE, True), (1, 3): (ORANGE, True)})
card(s, CX0, 13.40, 8.20, 3.30, "決めていない子の逃げ道を2つ",
     ["「30秒 学部学科診断」と「まだ決まっていない人へ」。",
      "学部が決まっていない子は資料請求を押せない。"], fill=PALE)
card(s, CX0 + 8.60, 13.40, 8.20, 3.30, "タブ③は「払う人」専用",
     ["決める人（高校生）と払う人（保護者）が違う業界。",
      "保護者が一番知りたい総額を、左上に置く。"], fill=PORANGE, hcol=ORANGE)
card(s, CX0 + 17.20, 13.40, 8.20, 3.30, "年3回、差し替える前提",
     ["OC期（6〜8月）はタブ②を、出願期（9〜11月）は",
      "出願スケジュールをデフォルト表示に切り替える。"], fill=PALE)

# ============================================================
# p27 リッチメニュー実装イメージ｜3タブの実物
# ============================================================
s = S[26]
frame(s, "リッチメニュー実装イメージ｜3タブの実物",
      ["同じLINEでも、受験生・入試期・保護者で見える画面が切り替わる。"])
TABS = ["受験生", "OC・入試", "保護者の方へ"]
MENUS = [["30秒\n学部学科診断", "オープン\nキャンパス予約", "学部・学科を見る",
          "資料（デジタル\nパンフ）をもらう", "まだ決まって\nいない人へ", "在学生に質問する"],
         ["次回OCの\n日程・予約", "当日の流れ・\n持ち物", "アクセス・\nキャンパスマップ",
          "入試方式を調べる", "出願スケジュール", "過去問・\n入試説明会"],
         ["4年間の学費\nシミュレーション", "奨学金・\n学費支援制度", "就職実績・\n進路データ",
          "保護者向け説明会", "一人暮らしの\n費用の目安", "個別相談を\n申し込む"]]
CAPS = ["受験生のデフォルト表示", "OC期（6〜8月）に切替", "保護者タブ（常設）"]
for i in range(3):
    x = CX0 + i * 8.60
    badge(s, x, CY0, 8.20, 1.00, f"{'①②③'[i]} {TABS[i]}", fill=PALE, col=NAVY, sz=12)
    richmenu(s, x, CY0 + 1.15, 8.20, 6.00, TABS, i, MENUS[i], hi=0)
    badge(s, x, CY0 + 7.40, 8.20, 1.10, CAPS[i], fill=GREY, col=INK, sz=10.5)
band(s, 13.00, "大サイズ 2500×1686px／3列×2行（1マス 833×843px）。学年・志望度タグで自動切替。")
foot(s, "※ 画像はイメージです。実装時はキャンパス写真・スクールカラー・アイコンを貴学仕様に差し替えます。\n"
        "※ 左上（オレンジ）が最重要CTA。親指が最初に届く位置に「まだ決めていない人が押せるもの」を置く。", y=14.70)

# ============================================================
# p28 トーク画面設計案｜あいさつメッセージと30秒診断
# ============================================================
s = S[27]
frame(s, "トーク画面設計案｜あいさつメッセージと30秒 学部学科診断",
      ["1通目で資料請求は求めない。求めるのは診断への回答だけ。",
       "設問は、そのまま配信セグメントになるものだけにする。"])
gsp = box(s, CX0, CY0, 12.40, 8.60, fill=WHITE, line=BORDER, lw=1.0)
put_text(gsp.text_frame,
         [one("■ 友だち追加直後のあいさつメッセージ（自動）", 12, True, NAVY, sa=6)] +
         [one(l, 11, None, INK, ls=1.34, sa=1) for l in [
             "◯◯大学 受験生応援アカウントへ",
             "友だち追加ありがとうございます🌸",
             "",
             "入試広報課の[担当者名]です。",
             "進路のこと、気軽に相談してください。",
             "",
             "📌 このアカウントについて",
             "・こちらから確認のお電話をすることはありません",
             "・有人でのご返信：平日 9:00-17:00",
             "・保護者の方も、このままご覧いただけます",
             "",
             "まずはこちらから👇",
             "あなたに向いている学部が30秒で分かります。"]],
         anchor="t", ml=0.30, mr=0.24, mt=0.24, mb=0.16)
badge(s, CX0 + 0.40, CY0 + 7.50, 11.60, 0.90, "▶ 30秒 学部学科診断をはじめる",
      fill=GREEN, col=GREEN, sz=11.5)
T(s, CX0 + 0.40, CY0 + 7.50, 11.60, 0.90,
  [one("▶ 30秒 学部学科診断をはじめる", 11.5, True, WHITE, align="c")],
  anchor="m", ml=0, mr=0)
qsp = box(s, CX0 + 13.00, CY0, 12.40, 8.60, fill=PALE, line=BORDER, lw=1.0)
put_text(qsp.text_frame,
         [one("■ 30秒 学部学科診断（4問すべて選択式・自由入力ゼロ）", 12, True, NAVY, sa=6)] +
         [one(l, 11, None, INK, ls=1.34, sa=1) for l in [
             "Q1 いまの学年　→ タグ：高1／高2／高3／既卒",
             "　　配信の内容と頻度を学年で変える",
             "",
             "Q2 学部は決まっていますか？",
             "　　① だいたい決まっている　② 2〜3個に絞れた",
             "　　③ 全然決まっていない ←ここが一番多くて大丈夫です",
             "",
             "Q3 学校を選ぶとき、いちばん大事にしたいことは？",
             "　　学びたいこと／就職・資格／通いやすさ／学費・奨学金",
             "",
             "Q4 通学はどうなりそうですか？",
             "　　自宅から通う／一人暮らしもできる／まだ分からない"]],
         anchor="t", ml=0.30, mr=0.24, mt=0.24, mb=0.16)
card(s, CX0, 12.40, 12.40, 4.20, "この1行に一番時間をかけた",
     ["Q2の③に「←ここが一番多くて大丈夫です」と書いてある。",
      "これがあるだけで、迷っている子が正直に答える。",
      "正直に答えてくれれば、その子に必要なものを送れる。"], fill=PORANGE, hcol=ORANGE)
card(s, CX0 + 13.00, 12.40, 12.40, 4.20, "取れるセグメントが、そのまま配信設計になる",
     ["学年 × 決まり具合 × 重視する軸 × 通学圏 の4象限。",
      "「学費」を選んだ子には、保護者向けコンテンツも同時に送る。",
      "「一人暮らし可」なら遠方＝オンライン相談を案内する。"], fill=PALE)

# ============================================================
# p29 トーク画面イメージ｜友だち追加の直後に届く1通目
# ============================================================
s = S[28]
frame(s, "トーク画面イメージ｜友だち追加の直後に届く1通目",
      ["この1通で「誰か」「何をくれるか」「売り込まないか」の3つに答える。",
       "ここで離脱すると、以降の配信はすべて届かない。"])
phone(s, CX0 + 0.60, CY0, 7.40, 10.60, "◯◯大学 受験生応援",
      [("in", "友だち追加ありがとうございます🌸\n入試広報課の[担当者名]です。"),
       ("in", "こちらから確認のお電話を\nすることはありません。\n保護者の方もご覧いただけます。"),
       ("in", "まずは、あなたに向いている\n学部を30秒で診断しませんか？"),
       ("btn", "▶ 30秒 学部学科診断をはじめる"),
       ("note", "既読 9:02")])
phone(s, CX0 + 9.40, CY0, 7.40, 10.60, "◯◯大学 受験生応援",
      [("in", "Q2 学部は決まっていますか？"),
       ("chip", "だいたい決まっている"),
       ("chip", "2〜3個に絞れた"),
       ("chip", "全然決まっていない\n←ここが一番多くて大丈夫です"),
       ("note", "※ 4問すべて選択式。入力はゼロ")])
phone(s, CX0 + 18.20, CY0, 7.40, 10.60, "◯◯大学 受験生応援",
      [("in", "診断結果が出ました📋"),
       ("in", "「やりたいことが決まっていない」\n——実は、いま在学している学生の\n多くが高2のこの時期は同じでした。"),
       ("in", "あなたに向いていそうなのは\n“入ってから選べる”学び方です。"),
       ("btn", "▶ 在学生Aさんの1日を見る（3分）"),
       ("note", "費用の話は、この後に出す")])
for i, cap in enumerate(["① 追加直後：売らない。電話しないと明記する",
                         "② 診断：4問すべて選択式・入力はゼロ",
                         "③ 結果：共感 → 実例 → 費用 → CTA の順"]):
    badge(s, CX0 + 0.60 + i * 8.80, CY0 + 10.75, 7.40, 1.10, cap,
          fill=PALE, col=NAVY, sz=10.5)
band(s, BAND_Y, "1通目で予約は求めない。求めるのは診断への回答だけ。")
foot(s, "※ 診断結果は「共感 → 実例 → 学べること → 費用 → CTA」の順。費用を先に出すと逃げられる（前後検索で検証済み）。")

# ============================================================
# p30 企画配信カレンダー（年間）
# ============================================================
s = S[29]
frame(s, "企画配信カレンダー｜入試日程は動かせない。だから逆算で組む",
      ["検索の山（実測）と入試日程（文科省）を重ねて、配信する月を先に決めてしまう。"])
table(s, CX0, CY0, CW, 11.20,
      ["時期", "企画", "参照データ"],
      [["3〜4月", "新学年スタート企画：「まだ決まっていなくていい」学部学科診断キャンペーン",
        "入塾・進路検討の立ち上がり（高2春〜）"],
       ["5月下旬〜6月", "OC先行案内：日程公開＋1人参加OK＋当日の流れを先に送る　★仕込みの本番",
        "検索が動き出すのは5/25週。6/29週で既に年間ピークの半分（トレンド実測）"],
       ["7〜8月", "OC直前・当日・参加後の3点セット：前日リマインド／受付QR／参加後アンケート",
        "「オープンキャンパス」ピーク＝7/27週（5年連続で7月ピーク）"],
       ["9月", "総合型選抜の出願支援：志望理由書の書き方・面接10問・締切カウントダウン",
        "「総合型選抜」ピーク＝8/31週（9/1出願解禁を含む週）。5年で3.3倍"],
       ["10〜11月", "学校推薦型・指定校推薦の出願支援／保護者向け説明会の案内",
        "「学校推薦型選抜」は11月単峰。「指定校推薦」は11月中旬に第2の山"],
       ["12〜1月", "一般選抜層の受け皿：入試方式の違い・過去問・入試説明会",
        "「共通テスト」は1月中旬の2週間だけ動く（他月は10以下）"],
       ["2〜3月", "入学手続期の離脱防止：入学前オリエン・先輩との交流・一人暮らし準備",
        "一般選抜 2/1〜。併願による流出が起きる区間"]],
      col_w=[3.4, 13.0, 9.0], align=["c", "l", "l"], row_h=1.55)
foot(s, "※ 参照データ：Googleトレンド（2026-09-15取得・日本・ウェブ検索）／文部科学省「大学入学者選抜実施要項」。\n"
        "※ 専門学校はAOエントリー6/1・出願10/1解禁（全国専修学校各種学校総連合会の申し合わせ・要確認）のため、全体が前倒しになる。")

# ============================================================
# p31 通知メッセージ 利用シーン
# ============================================================
s = S[30]
frame(s, "通知メッセージ・飛び道具｜友だちでない高校生にも届かせる",
      ["電話番号があれば、友だちでなくてもLINEに届く。資料請求者・OC予約者の取りこぼしを拾う。"])
T(s, CX0, CY0, 12.40, 0.90, [one("■ 通知メッセージ 利用シーン", 12.5, True, NAVY)], ml=0, mr=0)
notif = [("OC予約確定のお知らせ", "持ち物・アクセス・当日の流れ → 参加（CV②）へ"),
         ("OC前日リマインド", "「来られなくなっても大丈夫」を先に書く → 無断キャンセルを連絡ありに変える"),
         ("出願受付・書類到着のお知らせ", "次の手続きの期限を同時に案内 → 出願完了（CV③）へ"),
         ("入学手続の期限リマインド", "併願による流出が起きる区間を守る")]
for i, (h, b) in enumerate(notif):
    card(s, CX0, CY0 + 1.10 + i * 2.55, 12.40, 2.30, h, [b], fill=PALE, hsz=11.5, bsz=10.5)
T(s, CX0 + 13.00, CY0, 12.40, 0.90, [one("■ 飛び道具（この業界ならではの企画軸）", 12.5, True, ORANGE)],
  ml=0, mr=0)
gadget = [("① 30秒 学部学科診断", "匿名・入力ゼロ。決まっていない子が押せる唯一の入口"),
          ("② 4年間の学費シミュレーター", "入力3項目・1分。保護者の最大の不安に、先に答える"),
          ("③ 志望理由書の下書き支援", "出願直前に立つのは「書き方」検索。情報ではなく手伝いを渡す"),
          ("④ 中退予兆検知（入学後）", "一定期間アクションのない在学生に自動で声かけ")]
for i, (h, b) in enumerate(gadget):
    card(s, CX0 + 13.00, CY0 + 1.10 + i * 2.55, 12.40, 2.30, h, [b],
         fill=PORANGE, hcol=ORANGE, hsz=11.5, bsz=10.5)
foot(s, "※ 通知メッセージは Sales Partner／Tech Partner 限定公開のAPIを使用する機能（初期10万円／月額8万円〜＋通数単価7円）。\n"
        "※ 送信対象は貴学が保有する電話番号リスト。利用にあたっては個人情報の取得同意の範囲を確認する。")

# ============================================================
# p32 施策の考え方（施策パターン①②）
# ============================================================
swap_text(S[31], PHASE_MAP)
swap_text(S[31], {
    "✓ Web予約は既に回っている": "✓ 資料請求は既に取れている",
    "✓ 無断キャンセルが課題": "✓ OCの無断キャンセルが課題",
    "✓ 入会後もLINEを活用したい": "✓ 入学後もLINEを活用したい",
})

# ============================================================
# p34・p36 改善モデル②④の業界語
# ============================================================
swap_text(S[33], {"短期：ナーチャリングにより、情報収集→実来店・直面談の機会設定":
                  "短期：ナーチャリングにより、情報収集→オープンキャンパス参加・個別相談の機会設定"})
swap_text(S[35], {"誕生日・入会記念日": "誕生日・入学記念日", "誕生日・入会日等": "誕生日・入学日等",
                  "年間LTVを平均20%向上": "在学中の接点が途切れにくくなる"})

# ============================================================
# p37 導入成果シミュレーション（Before / After）
# ============================================================
s = S[36]
frame(s, "導入成果シミュレーション（Before / After）",
      ["デジタル募集広報費 年1,000万円は据え置き。追加するのはLINE運用費 年360万円のみ。",
       "モデル校＝私立大学・入学定員600名・現在の入学者480名（充足率80%）。"])
table(s, CX0, CY0, CW, 10.60,
      ["指標", "Before（Web広告のみ）", "After（Web広告 × LINE OA）"],
      [["投下計（広告費＋LINE運用費）", "1,000万円", "1,360万円（＋LINE運用費360万円）"],
       ["LP流入", "100,000セッション", "100,000セッション（変えない）"],
       ["LP CVR（資料請求）", "2.5%", "2.5%（変えない）"],
       ["広告経由のCV①（資料請求）", "2,500件", "2,500件"],
       ["LINE友だち化（離脱防止）", "―", "2,925人（離脱97,500 ×★3%）"],
       ["LINE経由のCV①", "―", "585件（友だち ×★20%）"],
       ["CV①合計 ／ CPA", "2,500件 ／ 4,000円", "3,085件 ／ 4,408円"],
       ["→ OC予約率", "30%（750件）", "30%（926件）※率は変えない"],
       ["→ OC参加率", "70%（525名）", "80%（741名）※前日リマインド"],
       ["→ 出願率", "40%（210名）", "43%（319名）★"],
       ["→ 入学手続率", "57%（120名）", "57%（182名）"],
       ["入学者（Web経由）", "120名", "182名（＋62名）"],
       ["入学1名あたりの獲得コスト", "83,333円", "74,725円（▲10.3%）"]],
      col_w=[8.4, 8.0, 9.0], align=["l", "r", "r"], row_h=0.76,
      cell_col={(9, 1): (ORANGE, True), (9, 2): (ORANGE, True),
                (10, 1): (ORANGE, True), (10, 2): (ORANGE, True),
                (12, 1): (RED, True), (12, 2): (RED, True),
                (13, 1): (NAVY, True), (13, 2): (NAVY, True)})
card(s, CX0, 14.35, 8.20, 2.30, "① LP CVRは動かしていない",
     ["触るのは「LPから落ちた後」と「予約が入った後」の2箇所だけ。"], fill=PALE, bsz=10.5)
card(s, CX0 + 8.60, 14.35, 8.20, 2.30, "② 効くのは母数ではなく率",
     ["CV①は＋23%。でも入学者は＋52%。歩留まりが掛け算で効く。"], fill=PALE, bsz=10.5)
card(s, CX0 + 17.20, 14.35, 8.20, 2.30, "③ CPA単体は悪化して見える",
     ["見るのは入学1名あたり。83,333円→74,725円（▲10.3%）。"],
     fill=PORANGE, hcol=ORANGE, bsz=10.5)
foot(s, "★＝仮置き（①友だち化率3%　②友だち→CV①転換率20%　③出願率40→43%）。貴学の実績値で置き換える。\n"
        "※ 本試算はWeb経由の募集だけを切り出したもの。他経路と重複するため純増は試算より小さくなる。"
        "定員充足済みの学校では「歩留まり改善によるコスト削減」として読む。", y=16.90)

print("[3/4] 施策章（p22〜p37）完了")

# ============================================================
# p43 ご提供プラン ／ p45-49 具体施策 の業界語
# ============================================================
swap_text(S[38], {
    "体験希望はチャット受付 → 店舗スタッフへトスアップ":
        "OC・個別相談の希望はチャット受付 → 入試広報課へトスアップ",
})
for idx in (40, 41, 42, 43, 44):
    swap_text(S[idx], {
        "体験予約": "OC予約", "入会": "出願", "来店": "オープンキャンパス参加",
        "店舗": "キャンパス", "会員": "在学生", "ジム": "学校",
    })

# ============================================================
# ★追加4枚：施策設計4点セット（_templates/施策設計4点セット_FMT.pptx の型）
#   ①動線 ②効率改善（配信） ③満足度改善（ユーザー） ④効率改善（管理側）
# ============================================================
LAYOUT = next(l for l in prs.slide_layouts if l.name == "4_タイトルとコンテンツ")


def new_slide():
    sl = prs.slides.add_slide(LAYOUT)
    clear_slide(sl)
    return sl


# --- ① 動線 ---
s = new_slide()
frame(s, "施策設計①｜友だち追加の動線",
      ["新しく集めるのではなく、いま取れている資料請求者・OC参加者をLINEに載せ替える。"])
table(s, CX0, CY0, CW, 8.60,
      ["箇所", "訴求", "到達点", "費用（税抜）"],
      [["資料請求の完了画面\n（サンクスLINE誘導）", "学部学科診断で、自分に合う学部が30秒でわかる",
        "LINE追加", "コンサル発注なら無償"],
       ["オープンキャンパス受付のQR", "当日の流れ・持ち物を先に送る／受付はLINE画面を見せるだけ",
        "LINE追加", "無償"],
       ["LP離脱防止ポップアップ", "30秒の学部診断／名前も電話番号も入力しません",
        "LINE追加", "初期 1.5万円 ／ 月 3万円"],
       ["LINE広告（CPF）", "タップ1回で友だち追加・フォーム入力ゼロ",
        "LINE追加", "20万円〜（1ヶ月〜）"],
       ["通知メッセージ（保有リスト）", "資料請求者・OC予約者へ、友だちでなくても直接届く",
        "LINE追加", "初期 10万円 ／ 月 8万円〜＋通数 7円"],
       ["パンフ裏表紙・高校での説明会資料", "デジタルパンフをLINEで受け取れる",
        "LINE追加", "無償"]],
      col_w=[6.6, 10.4, 3.0, 5.4], align=["l", "l", "c", "l"], row_h=1.28,
      cell_col={(1, 3): (NAVY, True), (2, 3): (NAVY, True), (6, 3): (NAVY, True)})
band(s, 13.10, "最も志望度が高いのは「OCに来てくれた高校生」。受付のQR1枚が、コストゼロで一番効く。")
foot(s, "※ 費用は DYM 標準（税抜・6ヶ月〜）。オフライン動線・WEBサイト動線整備はコンサル発注時は無償で付帯。\n"
        "※ 通知メッセージは Sales Partner／Tech Partner 限定公開のAPIを使用する機能。料金プランにより通数単価・利用可能機能が異なる。",
     y=14.70)

# --- ② 効率改善（配信の効果算出） ---
s = new_slide()
frame(s, "施策設計②｜効率改善（配信の効果算出）",
      ["配信ごとに「対象 → 訴求 → 開封率×クリック率×CVR → 件数」を一本の式で示す。",
       "率は貴学の実績値を入れて確定する（進学領域の公開統計は存在しない）。"])
table(s, CX0, CY0, CW, 11.40,
      ["区分", "対象（想定数）", "訴求", "計算式", "件数"],
      [["ステップ配信", "Day0｜診断に回答した全員（○人）", "診断結果＋在学生の1日（動画3分）",
        "開封率×クリック率×CVR", "○件"],
       ["", "Day6｜Day0未反応者（○人）", "OC日程＋当日の流れ＋1人参加OK",
        "開封率×クリック率×CVR", "○件"],
       ["", "Day14｜未CV者（○人）", "残席状況＋オンライン相談も選べる",
        "開封率×クリック率×CVR", "○件"],
       ["企画配信", "高2の友だち（○人）", "学部の選び方／文系理系の決め方",
        "開封率×クリック率×CVR", "○件"],
       ["", "9月｜OC参加者（○人）", "志望理由書の書き方（3ステップ・記入例つき）",
        "開封率×クリック率×CVR", "○件"],
       ["", "11月｜出願検討層（○人）", "学校推薦型の締切カウントダウン",
        "開封率×クリック率×CVR", "○件"],
       ["", "2月｜合格者（○人）", "入学前オリエン・先輩との交流・一人暮らし準備",
        "開封率×クリック率×CVR", "○件"]],
      col_w=[3.4, 6.4, 8.2, 5.0, 2.4], align=["c", "l", "l", "c", "c"], row_h=1.55)
foot(s, "※ 配信の時期は Googleトレンド実測（2026-09-15取得）と文部科学省「大学入学者選抜実施要項」から逆算。\n"
        "※ 開封率・クリック率・CVR は貴学の実績で埋める。運用費は「施策設計①」の表を参照。")

# --- ③ 満足度改善（ユーザー） ---
s = new_slide()
frame(s, "施策設計③｜満足度改善（ユーザー）",
      ["情報を届けるだけでなく、困ったときに自分で解決できる導線も用意する。"])
table(s, CX0, CY0, CW, 6.40,
      ["区分", "内容"],
      [["QA自動化",
        "リッチメニュー「よくある質問」＋KW自動応答（学費／奨学金／アクセス／OC／出願締切）。24時間、待たせない"],
       ["個別チャット",
        "自動応答で解決しない相談は有人対応（Chat Plus／Call Plus・別途費用）。対応時間を必ず画面に明記する"],
       ["その他",
        "「診断をやり直す」「配信を止める」で自己解決。ブロックさせないための出口を、隠さず置く"]],
      col_w=[5.4, 20.0], align=["c", "l"], row_h=2.10)
card(s, CX0, 10.60, 12.40, 4.20, "この業界で最大の恐怖に、先回りする",
     ["高校生が資料請求をためらう最大の理由は「そのあと電話がかかってくること」。",
      "あいさつメッセージの1行目に",
      "「こちらから確認のお電話をすることはありません」と書く。コストはゼロ。"],
     fill=PORANGE, hcol=ORANGE)
card(s, CX0 + 13.00, 10.60, 12.40, 4.20, "保護者も同じアカウントで受け止める",
     ["決める人（高校生）と払う人（保護者）が違う業界。",
      "「保護者の方も、このままご覧いただけます」を明記し、",
      "リッチメニュー③に学費・奨学金・就職実績をまとめて置く。"], fill=PALE)
band(s, BAND_Y, "困ったときに自分で解決できるほど、ブロックされずに残る。")

# --- ④ 効率改善（管理側） ---
s = new_slide()
frame(s, "施策設計④｜効率改善（管理側）",
      ["人手を増やさずに、資料請求直後の初動と、募集ピーク期の問い合わせ対応を回す。"])
table(s, CX0, CY0, CW, 4.60,
      ["区分", "内容"],
      [["自動応答",
        "資料請求の直後にあいさつ＋診断4問を自動送信。夜間でも初動の速度が落ちない"],
       ["その他",
        "GAレポート連携で、友だち数・OC予約率・OC参加率・出願率を毎月数値で振り返る"]],
      col_w=[5.4, 20.0], align=["c", "l"], row_h=2.10)
card(s, CX0, 8.80, 12.40, 6.00, "電話が鳴り続けるのは7〜8月と1〜2月",
     ["「学費は」「アクセスは」「締切は」——入試広報課の電話が集中する時期は、",
      "そのまま募集のピークと重なる。",
      "",
      "この問い合わせをKW自動応答に逃がすと、",
      "その時間が募集活動そのものに戻る。",
      "",
      "※ 工数削減は時給換算で語らない。「ピーク期に募集の手が止まる」で語る。"],
     fill=PALE)
card(s, CX0 + 13.00, 8.80, 12.40, 6.00, "紙とメールのコストも落ちる",
     ["デジタルパンフをLINEで配信すれば、",
      "印刷費・郵送費・封入作業がゼロになる。",
      "",
      "しかも資料請求で取ったメールアドレスは、",
      "高校生が最も見ない経路。",
      "LINEなら当日中に約8割が開封される（LINEヤフー社公開値）。"],
     fill=PORANGE, hcol=ORANGE)
band(s, BAND_Y, "増員ではなく自動化で、募集ピーク期の初動速度を守る。")

# ---- 追加4枚を p37 の直後（38〜41枚目）へ移動 ----
xml_slides = prs.slides._sldIdLst
els = list(xml_slides)
for i, old in enumerate([46, 47, 48, 49]):
    xml_slides.remove(els[old])
    xml_slides.insert(37 + i, els[old])

# ============================================================
# 非表示スライドの状態を確認（ジムver3のまま維持する）
#   ジムver3では 20（前後検索・補足）と 42-45（具体施策4枚）が非表示だった。
#   学校法人版でもその状態を維持する（殿村さん指示・2026-09-16）。
#   ※出したくなったら del sl.element.attrib["show"] で表示に戻せる。
# ============================================================
hidden = [i for i, sl in enumerate(list(prs.slides), 1)
          if sl.element.get("show") == "0"]
print(f"[5/5] 非表示スライド（ジムver3の状態を維持）: {hidden}")

# ============================================================
# 最終チェック：ジム業界の語が残っていないか
# ============================================================
RESIDUE = re.compile(
    r"ジム|フィットネス|筋トレ|筋肉|ダイエット|痩せ|chocoZAP|カーブス|ライザップ|"
    r"エニタイム|ティップネス|FIT ?PLACE|フィットイージー|体づくり|脂肪|GLP|"
    r"リベルサス|マンジャロ|パーソナルトレー|入会金|退会")
hits = []
for i, sl in enumerate(list(prs.slides), 1):
    for sh in all_shapes(sl):
        frames = []
        if sh.has_text_frame:
            frames.append(sh.text_frame)
        if sh.has_table:
            for row in sh.table.rows:
                for c in row.cells:
                    frames.append(c.text_frame)
        for tf in frames:
            for m in RESIDUE.finditer(tf.text):
                hits.append((i, m.group(), tf.text.strip().replace("\n", " ")[:70]))

prs.save(str(OUT))
print(f"[4/4] 4点セット4枚を追加し、38〜41枚目へ移動。総枚数 = {len(prs.slides.__iter__.__self__._sldIdLst)}")
print(f"保存: {OUT.name}")
if hits:
    print(f"\n⚠️ ジム業界の語が {len(hits)} 件残っています：")
    seen = set()
    for i, w, ctx in hits:
        key = (i, w)
        if key in seen:
            continue
        seen.add(key)
        print(f"   p{i} 「{w}」 :: {ctx}")
else:
    print("\n✅ 残存語チェック：0件")

print("""
【商談前に必ず埋めるもの】
  p9/p10/p11  競合の友だち数・リッチメニュー実機（page.line.me をPCで実査）
  p15         18歳人口2040年推計（82万か88万か）・定員割れ比率・進学率・納付金の原典照合
  p19         前後検索の起点KW（未確認）／総合型選抜・専門学校の前後検索が未取得
  p18/p37/p39 社内実績（CV地点別 CPC・CVR・CPA）→ ★仮置き3つが消える
  p25         大学の中退率（文科省調査）の原典照合
""")
