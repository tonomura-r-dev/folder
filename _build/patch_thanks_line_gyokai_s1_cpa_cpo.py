# -*- coding: utf-8 -*-
"""サンクスLINE誘導・業界別 ver3.0（殿村さんPC保存版）の1枚目だけ組み直す（2026-10-07 殿村さん指示）。2〜5枚目は触らない。
伝えること（1点）：広告運用とサンクスLINEをセットで実施し、CV獲得前のCPA改善とCV獲得後のCPO改善を両方行い、事業単位の獲得効率改善・広告拡大につなげる。
  左＝CV獲得まで｜CPA改善：広告→LP→離脱防止×成果報酬施策→CV／補足「CVの取りこぼしを減らし、施策全体のCPA改善へ」
  右＝CV獲得後｜CPO改善：CV→サンクスLINE誘導→LINEで継続フォロー→購入・来店・面談などの最終成果／補足「CV後の歩留まりを改善し、CPO改善へ」
  下＝CPO改善→許容CPA（目安CPA）の引き上げ→広告配信の拡大（3ステップ）／結論1文。
白＋紺、緑はLINE側だけ。機能・費用・事例・業界別は入れない。
  python3 _build/patch_thanks_line_gyokai_s1_cpa_cpo.py <業界別ver3.0.pptx> [出力]
"""
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

sys.path.insert(0, str(Path(__file__).resolve().parent))
from pptx_parts import *  # noqa

ROOT = Path(__file__).resolve().parent.parent
OUT = Path(sys.argv[2]) if len(sys.argv) > 2 else Path(sys.argv[1])
PALE = "E4E8F6"
L, R = 2.2, 31.7
GAP = 1.2
COLW = (R - L - GAP) / 2

prs = Presentation(sys.argv[1])
s = prs.slides[0]
assert s.shapes[0].text_frame.text.startswith("広告と一緒にサンクスLINE"), "1枚目が「広告と一緒に…」ではない"
keep_header_only(s)
set_paras(find(s, "TextBox 2"), [("広告はCV獲得までのCPA、サンクスLINEはCV獲得後のCPOを改善", 0)])


def column(x, stage, kpi, accent, steps, note):
    # 見出し：CV獲得まで｜CPA改善（KPIを大きく）
    rich(s, x, 4.5, COLW, 1.3, [[(stage + "｜", 13, True, NAVY), (kpi, 24, True, NAVY)]], align=PP_ALIGN.CENTER)
    rect(s, x + COLW / 2 - 3.0, 5.85, 6.0, 0.12, accent)
    # 縦の流れ（4段）
    bw, bh, gap = 9.2, 1.1, 0.4
    bx = x + (COLW - bw) / 2
    y = 6.3
    for i, (t, f, c) in enumerate(steps):
        shape(s, MSO_SHAPE.RECTANGLE, bx, y, bw, bh, [(t, 12.5, True, c, 0)], fill=f, margins=(0.1, 0, 0.1, 0))
        if i < len(steps) - 1:
            arrow_line(s, x + COLW / 2, y + bh + 0.05, x + COLW / 2, y + bh + gap - 0.05, NAVY, 1.75)
        y += bh + gap
    # 補足（1行）
    label(s, x, y + 0.1, COLW, 0.8, [(note, 11, True, INK, 0)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    return y + 0.9


column(L, "CV獲得まで", "CPA改善", NAVY,
       [("広告", PALE, NAVY), ("LP", PALE, NAVY), ("離脱防止 × 成果報酬施策", NAVY, WHITE), ("CV", NAVY, WHITE)],
       "CVの取りこぼしを減らし、施策全体のCPA改善へ")
yb = column(L + COLW + GAP, "CV獲得後", "CPO改善", GREEN,
            [("CV", PALE, NAVY), ("サンクスLINE誘導", GREEN, WHITE), ("LINEで継続フォロー", GREEN_BG, GREEN_TX),
             ("購入・来店・面談などの最終成果", NAVY, WHITE)],
            "CV後の歩留まりを改善し、CPO改善へ")
vline(s, L + COLW + GAP / 2, 4.7, yb - 0.2, LGRAY, 1.0, dash=True)

# ---- CPO改善 → 許容CPAの引き上げ → 広告配信の拡大（小さな3ステップ）----
hline(s, L, R, yb + 0.1, LGRAY, 0.75)
CY, CH = yb + 0.45, 1.25
cw, ov = 8.6, 0.3
cx0 = (L + R) / 2 - (3 * cw - 2 * ov) / 2
chev = [("CPO改善", GREEN, WHITE), ("許容CPA（目安CPA）の引き上げ", PALE, NAVY), ("広告配信の拡大", NAVY, WHITE)]
caps = ["CV後の歩留まりが上がる", "1件の最終成果にかけられる広告費が増える", "入札・配信を広げやすくなる"]
for i, ((t, f, c), cap) in enumerate(zip(chev, caps)):
    x = cx0 + i * (cw - ov)
    shape(s, MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON, x, CY, cw, CH, None, fill=f, adj=0.25)
    pad = 0.3 if i == 0 else 0.6
    label(s, x + pad, CY, cw - pad - 0.5, CH, [(t, 11.5, True, c, 0)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    label(s, x + 0.1, CY + CH + 0.05, cw - 0.2, 0.6, [(cap, 9, False, GRAY, 0)], align=PP_ALIGN.CENTER)

# ---- 結論（1文）----
label(s, L, CY + CH + 1.0, R - L, 1.3, [("広告施策とサンクスLINEをセットで実施し、事業単位での獲得効率を改善", 17, True, NAVY, 0)],
      align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

prs.save(str(OUT))
print("saved:", OUT.name)
