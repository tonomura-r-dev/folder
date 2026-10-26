"""
10月号 P.8：下の※2行を「注意点」のカードに入れる（2026-10-26 殿村さん指示）。
変更点カードを複製して注意点カードにし、2枚を同じ幅で縦に並べて上下左右中央に置く。
使い方: python3 _build/patch_trend_202610_p8b.py <pptx>
"""
import sys
import copy
from pptx import Presentation
from pptx.util import Inches, Emu
from pptx.oxml.ns import qn

F = sys.argv[1]
prs = Presentation(F)
s8 = list(prs.slides)[7]


def text(sh):
    return sh.text_frame.text.strip() if sh.has_text_frame else ""


head = next(x for x in s8.shapes if text(x) == "■ 詳細")
card = next(x for x in s8.shapes if text(x).startswith("変更点"))
note = next(x for x in s8.shapes if text(x).startswith("※ ユニット名"))
note._element.getparent().remove(note._element)

el = copy.deepcopy(card._element)
card._element.addnext(el)
cn = el.find(".//" + qn("p:cNvPr"))
cn.set("id", str(max(int(e.get("id")) for e in s8._element.iter(qn("p:cNvPr"))) + 1))
cn.set("name", cn.get("name") + "_note")
card2 = next(x for x in s8.shapes if x._element is el)
BODY2 = ["注意点",
         "・ユニット名は送るときに付ける（1通に1つ・後から変更不可・月1,000種類まで）。",
         "・20人未満の数字は表示されない。送信数は取れないため、開封率の分母は自社で持っておく。"]
for p, t in zip(card2.text_frame.paragraphs, BODY2):
    p.runs[0].text = t


def text_w(s, pt):
    return sum((0.55 if ord(c) < 0x2000 else 1.0) for c in s) * pt / 72


tf = card.text_frame
lines = [p.text for c in (card, card2) for p in c.text_frame.paragraphs[1:]]
w = min(max(text_w(t, 14) for t in lines) + Emu(tf.margin_left + tf.margin_right).inches + 0.45, 9.64)
h = 0.32 + 20 * 1.3 / 72 + 2 * (14 * 1.45 + 6) / 72
gap = 0.2
area_top = Emu(head.top + head.height).inches + 0.08
top = area_top + (7.05 - area_top - (2 * h + gap)) / 2
left = (Emu(prs.slide_width).inches - w) / 2
for i, c in enumerate((card, card2)):
    c.left, c.top, c.width, c.height = Inches(left), Inches(top + i * (h + gap)), Inches(w), Inches(h)

prs.save(F)
print("saved", F, round(w, 2))
