# -*- coding: utf-8 -*-
"""あらうち資料v10：「競合分析」1枚を6ページ（課題）と7ページ（本質課題）の間に足す（2026-10-15 殿村さん指示）。

    python -m _build.arauchi_v10.add_kyougou <殿村さんがPowerPointで直したv10.pptx> <出力.pptx>

  タイトル（殿村さん決定）：競合分析：一律100円・追加条件なしは、100円賃貸だけの価値
  ・表＝競合3社（タダスム／39room／ぜろいちくん）と100円賃貸の、仲介手数料・手数料が決まるタイミング・LINE友だち数
  ・左「競合と並ぶ点」／右「100円賃貸だけの点」（先方の資料に「弱み」とは書かない）
  ・下「ポイントは問い合わせ率と契約率」
  出典は `_drafts/あらうち_商材・市場調査（2026-10-09）.md` の 2026-10-15 追記。

  作り方：python-pptx で新しいスライドを1枚作り、そのXMLだけを元のzipに足す
  （ほかの43枚・[Content_Types].xml などは PowerPoint が書いたまま。追加と並べ替えだけ手で反映）。
"""
import re
import sys
import zipfile
from pathlib import Path

from pptx import Presentation

from _build.arauchi_v4.lib import NAVY, RED, TEXT, add_box, add_table, move, set_paras
from _build.arauchi_v10.build import by_id, c14n, clone, meiryo, normalize_decl, rels_of, runs, slide_order

INSERT_AT = 7                     # 新しいページの番号
SRC_HEADER = 7                    # タイトル・リードの帯を写すページ（本質課題）
GRAY, TINT, LINE_C, TAKE = "7F7F7F", "F4F7FF", "D9D9D9", "EEF1F8"

TITLE = "競合分析：一律100円・追加条件なしは、100円賃貸だけの価値"
LEAD = "格安仲介はどこも「URLを送るだけ」。差をつけられるのは、価格の分かりやすさ"
ROWS = [("タダスム", "0円または最大50%", "送った物件ごとに判定", "22,830人（関東の窓口の1つ）"),
        ("39room", "0円／最大39,000円", "送った物件ごとに判定", "13,492人"),
        ("ぜろいちくん", "0円か1万円", "送った物件ごとに判定", "5,619人")]
SELF = ("100円賃貸", "一律100円（対応できる物件）", "送る前から分かる", "1,241人")
SAME = ["・「URLを送るだけ」の手軽さは、どこも同じ",
        "・「仲介手数料0円」をうたう競合もある",
        "・LINEの友だち数は、これから伸ばす段階"]
NOTE = "※各社の公式サイト・LINE公式アカウントの記載より（2026年10月15日確認）"


def build_slide(prs):
    src = prs.slides[SRC_HEADER - 1]
    s = prs.slides.add_slide(src.slide_layout)
    for sh in list(s.shapes):
        sh._element.getparent().remove(sh._element)
    t = clone(s, by_id(src, 649))
    set_paras(t, [TITLE])
    meiryo(t)
    move(t, w=8.6)
    for sid in (650, 651, 652):
        clone(s, by_id(src, sid))
    set_paras(clone(s, by_id(src, 653)), [LEAD])

    b = {k: (0.75, LINE_C) for k in "lrtb"}
    head = dict(fill=NAVY, color="FFFFFF", bold=True, size=11.5, borders=b)
    rows = [[dict(text=x, **head) for x in ("競合", "仲介手数料", "手数料が決まるタイミング", "LINE友だち数")]]
    for r in ROWS:
        rows.append([dict(text=r[0], size=11.5, bold=True, color=TEXT, fill=TINT, borders=b)]
                    + [dict(text=x, size=11.5, borders=b) for x in r[1:]])
    hl = dict(size=12, color=NAVY, bold=True, fill=TAKE, borders={k: (1.5, NAVY) for k in "lrtb"})
    rows.append([dict(text=x, **hl) for x in SELF])
    add_table(s, 0.55, 1.72, [1.55, 3.05, 2.35, 2.78], [0.38] * 5, rows, name="競合比較")

    w = (9.73 - 0.25) / 2
    x2 = 0.55 + w + 0.25
    add_box(s, 0.55, 3.78, w, 0.36, [runs(("競合と並ぶ点", dict(size=12.5, bold=True, color="FFFFFF")))],
            fill=GRAY, align="ctr", name="競合と並ぶ点 見出し")
    add_box(s, 0.55, 4.14, w, 1.55, [runs((x, dict(size=12, color=TEXT))) for x in SAME],
            line=LINE_C, anchor="ctr", line_spacing=130, ins=(0.15, 0.06, 0.12, 0.06), name="競合と並ぶ点")
    add_box(s, x2, 3.78, w, 0.36, [runs(("100円賃貸だけの点", dict(size=12.5, bold=True, color="FFFFFF")))],
            fill=NAVY, align="ctr", name="100円賃貸だけの点 見出し")
    add_box(s, x2, 4.14, w, 1.55, [
        runs(("・対応できる物件なら一律100円", dict(size=12, color=TEXT))),
        runs(("　（競合は物件ごとに「0円か数万円」）", dict(size=10.5, color=GRAY))),
        runs(("・オプション契約・付帯契約・違約・しばりなし", dict(size=12, color=TEXT))),
        runs(("→ 送る前から総額が読め、問い合わせのハードルが低い", dict(size=12, bold=True, color=NAVY)))],
        line=NAVY, anchor="ctr", line_spacing=130, ins=(0.15, 0.06, 0.12, 0.06), name="100円賃貸だけの点")
    add_box(s, 0.55, 5.86, 9.73, 0.66, [
        runs(("ポイントは問い合わせ率と契約率。", dict(size=13, bold=True, color=RED)),
             ("LINEの入口で2つの強みを最初に伝え、送った後もLINEで追う", dict(size=13, bold=True, color=NAVY)))],
        fill=TINT, align="ctr", anchor="ctr", name="ポイント")
    add_box(s, 0.55, 6.58, 9.73, 0.3, [runs((NOTE, dict(size=9, color=GRAY)))], anchor="ctr", ins=(0, 0, 0, 0), name="注記")
    return s


def repack(base_path, built_path, out_path):
    zb, zn = zipfile.ZipFile(base_path), zipfile.ZipFile(built_path)
    bo, no = slide_order(zb), slide_order(zn)
    n = len(bo)
    assert len(no) == n + 1, (len(bo), len(no))
    new_built = no[-1]                                   # python-pptx は末尾に足す
    for i, (b, x) in enumerate(zip(bo, no[:n]), 1):      # ほかのページは1バイトも変えない
        assert c14n(zb.read(b)) == c14n(zn.read(x)), f"S{i} が変わっている"
    names = set(zb.namelist())
    k = n + 1
    while f"ppt/slides/slide{k}.xml" in names:
        k += 1
    part = f"ppt/slides/slide{k}.xml"
    rels_xml = zn.read(rels_of(new_built)).decode("utf-8")
    assert rels_xml.count("<Relationship ") == 1 and "slideLayout" in rels_xml, rels_xml

    prels = zb.read("ppt/_rels/presentation.xml.rels").decode("utf-8")
    rid_nums = [int(x) for x in re.findall(r'Id="rId(\d+)"', prels)]
    rid = f"rId{max(rid_nums) + 1}"
    prels = prels.replace("</Relationships>",
                          f'<Relationship Id="{rid}" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
                          f'relationships/slide" Target="slides/slide{k}.xml"/></Relationships>')
    pres = zb.read("ppt/presentation.xml").decode("utf-8")
    sld = re.findall(r"<p:sldId [^>]*/>", pres)
    assert len(sld) == n
    new_id = max(int(re.search(r'id="(\d+)"', x).group(1)) for x in sld) + 1
    sld.insert(INSERT_AT - 1, f'<p:sldId id="{new_id}" r:id="{rid}"/>')
    lst = re.search(r"<p:sldIdLst>.*?</p:sldIdLst>", pres).group(0)
    pres = pres.replace(lst, "<p:sldIdLst>" + "".join(sld) + "</p:sldIdLst>")
    ct = zb.read("[Content_Types].xml").decode("utf-8")
    ct = ct.replace("</Types>", f'<Override PartName="/{part}" ContentType="application/vnd.openxmlformats-officedocument.'
                                f'presentationml.slide+xml"/></Types>')
    app = zb.read("docProps/app.xml").decode("utf-8")
    app = app.replace(f"<Slides>{n}</Slides>", f"<Slides>{n + 1}</Slides>", 1)
    head = f"<vt:lpstr>スライド タイトル</vt:lpstr></vt:variant><vt:variant><vt:i4>{n}</vt:i4>"
    if head in app:
        app = app.replace(head, head.replace(f"<vt:i4>{n}</vt:i4>", f"<vt:i4>{n + 1}</vt:i4>"))
        m = re.search(r'<vt:vector size="(\d+)" baseType="lpstr">', app)
        app = app.replace(m.group(0), f'<vt:vector size="{int(m.group(1)) + 1}" baseType="lpstr">')
        app = app.replace("</vt:vector></TitlesOfParts>", "<vt:lpstr>PowerPoint プレゼンテーション</vt:lpstr></vt:vector></TitlesOfParts>")

    with zipfile.ZipFile(out_path, "w") as zo:
        for info in zb.infolist():
            data = {"ppt/presentation.xml": pres.encode("utf-8"),
                    "ppt/_rels/presentation.xml.rels": prels.encode("utf-8"),
                    "[Content_Types].xml": ct.encode("utf-8"),
                    "docProps/app.xml": app.encode("utf-8")}.get(info.filename)
            zo.writestr(info, data if data is not None else zb.read(info.filename), compress_type=info.compress_type)
        zo.writestr(part, normalize_decl(zn.read(new_built)), compress_type=zipfile.ZIP_DEFLATED)
        zo.writestr(rels_of(part), rels_xml.encode("utf-8"), compress_type=zipfile.ZIP_DEFLATED)
    return part


def main(base, out):
    tmp = out + ".tmp.pptx"
    prs = Presentation(base)
    build_slide(prs)
    prs.save(tmp)
    part = repack(base, tmp, out)
    Path(tmp).unlink()
    print("saved:", out, "／ 追加したパーツ:", part, f"（{INSERT_AT}ページ目）")


if __name__ == "__main__":
    main(*sys.argv[1:3])
