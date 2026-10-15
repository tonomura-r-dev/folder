# -*- coding: utf-8 -*-
"""あらうち提案資料 v10（2026-10-15 殿村さん指示）。J.P.Returns v11 の構成（1スライド1要旨）に組み直す。

    python -m _build.arauchi_v10.build <v9.pptx（45枚・commit c9d4b30）> <出力.pptx>

  お手本：`_templates/構成お手本_20261009_J.P.Returns_アドアフィ×LINE×Meta広告集客のご提案v11.pdf`
  1. タイトルを「論点：結論」に（サマリー・本質課題はお手本どおり、結論をKEY TAKEAWAY枠に入れる）
  2. 並べ替え：課題を施策より前に、施策の詳細を費用より前に
  3. 足す：サマリー／本質課題／次のステップ／Appendixの区切り（v9の章の区切りページ12・25・30・37を作り直す）
  4. まとめる：v9のS10（目標）を S11（費用対効果）に入れて1枚に（S10は外す）。月ごとの件数は S11 に表で残す
  5. 「月17件」は新15ページだけ：v9のS8の右の列は「ゴール：フォーム送信（CV）」に
  6. 使わない施策・汎用ページは、中身を変えずに裏表紙の後ろ（Appendix）へ

  作り方は v9 と同じ：python-pptx で直したスライドのXMLだけを、元のv9のzipに差し込む
  （[Content_Types].xml・rels などは PowerPoint が書いたまま。S10の削除と並べ替えだけ手で反映）。
"""
import copy
import re
import sys
import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn

from _build.arauchi_v4.lib import NAVY, RED, TEXT, add_box, add_table, delete, move, set_paras

N_BASE = 45
DROP = [10]
NEW_ORDER = [1, 12, 2, 4, 5, 9, 25, 3, 7, 8, 6, 17, 15, 16, 11, 13, 14, 27, 28, 29, 30, 45,      # 本編（22枚）
             37, 18, 19, 20, 21, 22, 23, 24, 26, 31, 32, 33, 34, 35, 36, 38, 39, 40, 41, 42, 43, 44]  # Appendix
assert sorted(NEW_ORDER + DROP) == list(range(1, N_BASE + 1))
assert NEW_ORDER.index(11) + 1 == 15                       # 費用対効果＝新15ページ（S14の注記で参照）

MID = "3E5C99"
TINT = "F4F7FF"
TAKE = "EEF1F8"
LINE_C = "D9D9D9"
GRAY = "595959"
CV = [13, 14, 14, 15, 16, 17]                               # SIM（2026-10-14 16時保存）の問い合わせ・月ごと
assert sum(CV) == 89

TITLES = {   # v9の番号: (タイトルの図形id, 文字, 幅を広げるか)
    2: (35, "現状：標準機能のまま、追加後の導線を厚くできる", True),
    3: (91, "お客様の動線：LINE追加後の物件送信が最重要", True),
    4: (125, "LINEの現状：送る形は合っている。案内を足して厚くする", True),
    5: (138, "市場：家賃が上がる今、初期費用の負担感に応える", True),
    6: (170, "運用ツール：標準機能で運用する", True),
    7: (178, "全体像：サイト → LINE追加 → フォーム送信を標準機能でつなぐ", True),
    8: (234, "施策：新しい友だちと今の友だちの両方に届ける", True),
    9: (304, "課題：サイトの取りこぼし・送る前のためらい・追加後の導線不足", True),
    11: (339, "目標と費用対効果：月13.5万円で、6か月目に月17件（試算）", True),
    12: (633, "サマリー", False),
    13: (635, "料金プラン：3つのプランの費用", False),
    14: (2, "運用内容：プランごとの対応範囲", False),
    15: (385, "配信：送らなかった方を、ステップ配信と企画配信で追い続ける", True),
    16: (463, "問い合わせ対応：一次対応は自動、物件の確認に集中", True),
    17: (478, "新規友だち：離脱防止ポップアップとサイトのボタンで増やす", True),
    25: (649, "本質課題", False),
    27: (681, "効果計測：配信・リッチメニューのリンクごとに成果を数値化", True),
    28: (684, "体制：戦略から運用・制作まで一貫して対応", True),
    29: (718, "スケジュール：お申込みから運用開始までの5ステップ", True),
    30: (736, "次のステップ：運用開始に向けて伺いたいこと", True),
    37: (980, "Appendix", False),
}
TITLE_W = 8.40      # ロゴ（x≒9.3in）の手前まで


# ---------------------------------------------------------------- 小さな部品
def by_id(slide, sid):
    hit = [sh for sh in slide.shapes if sh.shape_id == sid]
    assert len(hit) == 1, (sid, len(hit))
    return hit[0]


def meiryo(shape):
    """図形の中の文字の字体をメイリオに（latin・ea）"""
    for tag in ("a:rPr", "a:endParaRPr"):
        for rpr in shape._element.iter(qn(tag)):
            for ft in ("a:latin", "a:ea"):
                e = rpr.find(qn(ft))
                if e is not None:
                    e.set("typeface", "メイリオ")


def drop_paras(shape, idxs):
    ps = shape.text_frame._txBody.findall(qn("a:p"))
    for i in sorted(idxs, reverse=True):
        ps[i].getparent().remove(ps[i])


def replace_text(shape, old, new):
    """段落の中で old を new に（run をまたいでいても、最初の run の書式で1つにまとめる）"""
    n = 0
    for p in shape.text_frame._txBody.findall(qn("a:p")):
        rs = p.findall(qn("a:r"))
        text = "".join(r.findtext(qn("a:t")) or "" for r in rs)
        if old not in text:
            continue
        rs[0].find(qn("a:t")).text = text.replace(old, new)
        for r in rs[1:]:
            p.remove(r)
        n += 1
    return n


def clone(dst, src_shape):
    el = copy.deepcopy(src_shape._element)
    for ext in list(el.iter(qn("a:extLst"))):
        if ext.getparent().tag == qn("p:cNvPr"):
            ext.getparent().remove(ext)
    nid = max(sh.shape_id for sh in dst.shapes) + 1
    el.find(".//" + qn("p:cNvPr")).set("id", str(nid))
    dst.shapes._spTree.append(el)
    return by_id(dst, nid)


def lead_band(dst, src11, text):
    """S11 のリード部分（帯・区切り線・文字）を写して、リードを1行入れる"""
    for sid in (337, 338, 340):
        clone(dst, by_id(src11, sid))
    lead = clone(dst, by_id(src11, 381))
    set_paras(lead, [text])
    return lead


def runs(*items):
    """[(文字, dict(size, bold, color)), ...] の1段落"""
    return [(t, o) for t, o in items]


def key_takeaway(s, lines, y=1.85, h=1.45):
    add_box(s, 0.55, y, 2.30, 0.36, [runs(("KEY TAKEAWAY", dict(size=12, bold=True, color="FFFFFF")))],
            fill=NAVY, align="ctr", name="KEY TAKEAWAY 見出し")
    add_box(s, 0.55, y + 0.36, 9.73, h, lines, fill=TAKE, line=NAVY, line_w=0.75, align="ctr",
            line_spacing=125, ins=(0.2, 0.08, 0.2, 0.08), name="KEY TAKEAWAY")


def three_points(s, label, items, y=4.05, body_h=1.85, foot=None):
    """お手本の「補足ポイント」：見出し＋3列（色の短い線・番号つき見出し・本文）"""
    add_box(s, 0.55, y, 4.0, 0.32, [runs((label, dict(size=13, bold=True, color=TEXT)))], ins=(0, 0, 0, 0),
            anchor="ctr", name="補足ポイント")
    gap, w = 0.265, (9.73 - 2 * 0.265) / 3
    for k, (head, body, color) in enumerate(items):
        x = 0.55 + k * (w + gap)
        add_box(s, x, y + 0.44, 0.85, 0.055, [[("", {})]], fill=color, name=f"線{k + 1}")
        add_box(s, x, y + 0.55, w, 0.36, [runs((head, dict(size=13, bold=True, color=color)))], ins=(0, 0, 0, 0),
                anchor="ctr", name=f"見出し{k + 1}")
        add_box(s, x, y + 0.95, w, body_h, [runs((body, dict(size=13, color=TEXT)))], ins=(0, 0.04, 0.05, 0.04),
                anchor="t", line_spacing=135, name=f"本文{k + 1}")
        if foot:
            add_box(s, x, y + 0.95 + body_h + 0.08, w, 0.50,
                    [runs(("→ " + foot[k], dict(size=11, bold=True, color=NAVY)))],
                    fill=TINT, line=LINE_C, ins=(0.1, 0.04, 0.08, 0.04), name=f"打ち手{k + 1}")


def clear_body(s, keep_title):
    for sh in list(s.shapes):
        if sh.shape_id != keep_title:
            delete(sh)


# ---------------------------------------------------------------- 新しいページ
def p_summary(s, src11):
    clear_body(s, 633)
    lead_band(s, src11, "標準機能のまま、LINEからの物件問い合わせを、フォームで計測しながら増やす")
    key_takeaway(s, [runs(("「LINEで受けるだけ」から", dict(size=20, bold=True, color=NAVY))),
                     runs(("「フォーム送信まで案内する」LINEへ", dict(size=20, bold=True, color=NAVY)))])
    three_points(s, "補足ポイント", [
        ("01 現状・課題", "全ページにLINEのボタンがあり、問い合わせの入口は整っている。追加後に物件の送信へ進む案内（リッチメニュー・ボタン）を厚くできる。", NAVY),
        ("02 提案", "離脱防止ポップアップで入口を増やし、あいさつ・リッチメニュー・ステップ配信で、フォームでの物件送信まで案内する。", MID),
        ("03 進め方", "約1か月で初期設計を構築し、運用を開始する。月次レポートと定例会で数字を確認し、配信を見直す。", RED),
    ])


def p_honshitsu(s, src11):
    clear_body(s, 649)
    lead_band(s, src11, "3つの課題は、ひとつの課題につながっている")
    key_takeaway(s, [runs(("受付はできるが、追える仕組みがない", dict(size=20, bold=True, color=NAVY))),
                     runs(("その受付も分かりづらく、何を送ればいいか・送った後どうなるか・対象の物件かが伝わっていない",
                           dict(size=13, color=NAVY)))])
    three_points(s, "3つの課題と打ち手", [
        ("01 集める", "サイトを離れる方に、もう一度LINEを案内する手段がない", NAVY),
        ("02 伝える", "何を送るか・送った後の流れ・対象の物件かが、送る前に分からない", MID),
        ("03 追う", "物件を送らなかった方へ、続けて案内する配信がない", RED),
    ], body_h=1.15, foot=["離脱防止ポップアップ", "あいさつ・リッチメニュー・自動応答", "ステップ配信・企画配信"])


def p_next(s, src11):
    clear_body(s, 736)
    lead_band(s, src11, "以下を伺ったうえで、初期設計（約1か月）に入る")
    items = [
        ("サイトの「LINEでお問い合わせ」ボタンのリンク先を、友だち追加用のURLに変えてよいか", "友だちの入口（サイトのボタン）"),
        ("対応できない物件の条件と、送られた物件を確認するご担当者・手順", "あいさつ・自動応答の案内"),
        ("LINEでの返信の目安（受付時間・返信までの時間）", "あいさつ・受付時間外の自動返信"),
        ("物件を送った方へのステップ配信を止めるか", "ステップ配信の設計"),
        ("文面・画像・リッチメニューを最終確認するご担当者", "公開までの進め方"),
    ]
    border = {k: (0.75, LINE_C) for k in "lrtb"}
    head = dict(fill=NAVY, color="FFFFFF", bold=True, size=12, borders=border)
    rows = [[dict(text="#", **head), dict(text="伺いたいこと", **head), dict(text="反映するところ", **head)]]
    for i, (q, use) in enumerate(items, 1):
        rows.append([dict(text=str(i), size=13, bold=True, color=NAVY, fill=TINT, borders=border),
                     dict(text=q, size=12, align="l", borders=border, ins=(0.12, 0.04, 0.1, 0.04)),
                     dict(text=use, size=12, align="l", borders=border, ins=(0.12, 0.04, 0.1, 0.04))])
    add_table(s, 0.55, 1.95, [0.55, 5.45, 3.73], [0.45] + [0.70] * len(items), rows, name="伺いたいこと")


def p_appendix(s):
    for sh in list(s.shapes):
        if sh.shape_id == 981:
            set_paras(sh, ["参考資料"])
            meiryo(sh)


# ---------------------------------------------------------------- 直すページ
def p02(s):
    drop_paras(by_id(s, 88), [1])                      # 2行目（タイトルと同じ内容）を外す


def p08(s):
    set_paras(by_id(s, 271), ["ゴール"])
    t = by_id(s, 272)
    set_paras(t, ["フォーム送信", "（CV）"])            # 「約17件」は新15ページだけに
    drop_paras(by_id(s, 293), [2])                     # 件数の注記


def p11(s):
    set_paras(by_id(s, 381), ["プラン①｜総合運用で、友だち追加の増加と導線改善により、LINE経由の物件問い合わせを増やす"])
    add_box(s, 7.45, 3.74, 3.10, 0.30, [runs(("月ごとの問い合わせ（試算）", dict(size=11, bold=True, color=NAVY)))],
            ins=(0.02, 0, 0, 0), name="月ごとの問い合わせ 見出し")
    gap = 0.06
    w = (3.10 - 2 * gap) / 3
    for i, n in enumerate(CV):
        x = 7.45 + (i % 3) * (w + gap)
        y = 4.06 + (i // 3) * 0.56
        add_box(s, x, y, w, 0.50, [runs((f"{i + 1}か月目", dict(size=9.5, color=GRAY))),
                                   runs((f"{n}件", dict(size=13, bold=True, color=NAVY)))],
                fill=TINT, line=LINE_C, align="ctr", ins=(0.02, 0.02, 0.02, 0.02), line_spacing=100,
                name=f"{i + 1}か月目")
    move(by_id(s, 382), y=5.30, h=1.10)                # 試算の前提を下へ


def p14(s):
    assert replace_text(by_id(s, 7), "11ページ", "15ページ") == 1


def p27(s):
    delete(by_id(s, 663))                              # タイトルの右の「効果計測」の札（タイトルと重なる言葉）


def apply(prs):
    sl = {i + 1: s for i, s in enumerate(prs.slides)}
    src11 = sl[11]
    for no, (sid, text, wide) in TITLES.items():
        t = by_id(sl[no], sid)
        set_paras(t, [text])
        meiryo(t)
        if wide:
            move(t, w=TITLE_W)
    p02(sl[2])
    p08(sl[8])
    p_summary(sl[12], src11)
    p14(sl[14])
    p_honshitsu(sl[25], src11)
    p27(sl[27])
    p_next(sl[30], src11)
    p_appendix(sl[37])
    p11(sl[11])                                        # リードを写したあとで S11 を直す


CHANGED = sorted({2, 8, 11, 14, 27} | set(TITLES) | {12, 25, 30, 37})


# ---------------------------------------------------------------- 差し込み
_NOBLANK = etree.XMLParser(remove_blank_text=True)
DECL = b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\r\n'
P_NS = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main"}
RID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"


def c14n(data):
    return etree.tostring(etree.fromstring(data, _NOBLANK), method="c14n")


def rel_set(data):
    return {(r.get("Id"), r.get("Type"), r.get("Target"), r.get("TargetMode")) for r in etree.fromstring(data)}


def slide_order(z):
    rels = {r.get("Id"): r.get("Target") for r in etree.fromstring(z.read("ppt/_rels/presentation.xml.rels"))}
    pres = etree.fromstring(z.read("ppt/presentation.xml"))
    return ["ppt/" + rels[s.get(RID)] for s in pres.findall(".//p:sldId", P_NS)]


def normalize_decl(data):
    return DECL + etree.tostring(etree.fromstring(data), encoding="UTF-8", xml_declaration=False)


def rels_of(name):
    d, f = name.rsplit("/", 1)
    return f"{d}/_rels/{f}.rels"


def repack(base_path, built_path, out_path):
    zb, zn = zipfile.ZipFile(base_path), zipfile.ZipFile(built_path)
    bo, no = slide_order(zb), slide_order(zn)
    assert len(bo) == len(no) == N_BASE, (len(bo), len(no))

    # 外すパーツ（S10 とそのノート）
    gone, n_notes = set(), 0
    for d in DROP:
        part = bo[d - 1]
        gone |= {part, rels_of(part)}
        for r in etree.fromstring(zb.read(rels_of(part))):
            t = r.get("Type").rsplit("/", 1)[-1]
            if t == "notesSlide":
                notes = "ppt/notesSlides/" + r.get("Target").rsplit("/", 1)[-1]
                gone |= {notes, rels_of(notes)}
                n_notes += 1
            else:
                assert t == "slideLayout", t
    for name in zb.namelist():
        if name.endswith(".rels") and name not in gone and name != "ppt/_rels/presentation.xml.rels":
            data = zb.read(name).decode("utf-8")
            for g in gone:
                assert "/" + g.rsplit("/", 1)[-1] + '"' not in data, (name, g)

    prels = zb.read("ppt/_rels/presentation.xml.rels").decode("utf-8")
    pres = zb.read("ppt/presentation.xml").decode("utf-8")
    ct = zb.read("[Content_Types].xml").decode("utf-8")
    for g in gone:
        if g.endswith(".rels"):
            continue
        ct, n = re.subn(r'<Override [^>]*PartName="/%s"[^>]*/>' % re.escape(g), "", ct)
        assert n == 1, g

    # 並べ替え（sldIdLst を新しい順に書き直す。S10 は外す）
    sld = re.findall(r'<p:sldId [^>]*/>', pres)
    assert len(sld) == N_BASE
    for d in DROP:
        rid = re.search(r'r:id="(rId\d+)"', sld[d - 1]).group(1)
        m = re.search(r'<Relationship [^>]*Id="%s"[^>]*/>' % rid, prels)
        assert m and bo[d - 1].endswith(re.search(r'Target="([^"]+)"', m.group(0)).group(1)), (d, rid)
        prels = prels.replace(m.group(0), "")
    lst = re.search(r"<p:sldIdLst>.*?</p:sldIdLst>", pres).group(0)
    pres = pres.replace(lst, "<p:sldIdLst>" + "".join(sld[k - 1] for k in NEW_ORDER) + "</p:sldIdLst>")

    app = zb.read("docProps/app.xml").decode("utf-8")
    nb, nn = N_BASE, N_BASE - len(DROP)
    app = app.replace(f"<Slides>{nb}</Slides>", f"<Slides>{nn}</Slides>", 1)
    m = re.search(r"<Notes>(\d+)</Notes>", app)
    app = app.replace(m.group(0), f"<Notes>{int(m.group(1)) - n_notes}</Notes>", 1)
    head = f"<vt:lpstr>スライド タイトル</vt:lpstr></vt:variant><vt:variant><vt:i4>{nb}</vt:i4>"
    assert head in app
    app = app.replace(head, head.replace(f"<vt:i4>{nb}</vt:i4>", f"<vt:i4>{nn}</vt:i4>"))
    m = re.search(r'<vt:vector size="(\d+)" baseType="lpstr">', app)
    app = app.replace(m.group(0), f'<vt:vector size="{int(m.group(1)) - len(DROP)}" baseType="lpstr">')
    for _ in DROP:
        assert "<vt:lpstr>PowerPoint プレゼンテーション</vt:lpstr>" in app
        app = app.replace("<vt:lpstr>PowerPoint プレゼンテーション</vt:lpstr>", "", 1)

    replaced, changed = {}, []
    for i, (b, n) in enumerate(zip(bo, no), 1):
        if i in DROP:
            continue
        if c14n(zb.read(b)) != c14n(zn.read(n)):
            replaced[b] = normalize_decl(zn.read(n))
            changed.append(i)
        if rel_set(zb.read(rels_of(b))) != rel_set(zn.read(rels_of(n))):
            raise SystemExit(f"想定外のrelsの違い: S{i}")
    assert changed == CHANGED, f"変わったスライド {changed} != {CHANGED}"

    with zipfile.ZipFile(out_path, "w") as zo:
        for info in zb.infolist():
            name = info.filename
            if name in gone:
                continue
            data = {"ppt/presentation.xml": pres.encode("utf-8"),
                    "ppt/_rels/presentation.xml.rels": prels.encode("utf-8"),
                    "[Content_Types].xml": ct.encode("utf-8"),
                    "docProps/app.xml": app.encode("utf-8")}.get(name)
            if data is None:
                data = replaced.pop(name) if name in replaced else zb.read(name)
            zo.writestr(info, data, compress_type=info.compress_type)
    assert not replaced
    return changed, sorted(gone)


def main(base, out):
    tmp = out + ".tmp.pptx"
    prs = Presentation(base)
    apply(prs)
    prs.save(tmp)
    changed, gone = repack(base, tmp, out)
    Path(tmp).unlink()
    print("saved:", out)
    print("直したスライド（v9の番号）:", changed)
    print("外したパーツ:", gone)


if __name__ == "__main__":
    main(*sys.argv[1:3])
