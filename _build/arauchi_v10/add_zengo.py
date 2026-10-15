# -*- coding: utf-8 -*-
"""あらうち資料v10：「前後検索「賃貸」」1枚を、配信のページ（14ページ）の前に足す（2026-10-15 殿村さん指示）。

    python -m _build.arauchi_v10.add_zengo <殿村さんの最新のv10.pptx> <出力.pptx>

  画像は `_build/make_arauchi_zengo_color.py` で作った4色版（`_images/arauchi_zengo_chintai_color.png`）。
  色は配信ステップの順番と同じ分け方（エリアで探す／ポータルサイトで物件を見る／不動産会社を比べる／初期費用・審査を調べる）。
  凡例は画像の上。右に時期ごとに見えること、下に配信への示唆。
  作り方は add_kyougou.py と同じ（新しいスライドのXMLと画像だけを元のzipに足す。ほかのページは1バイトも変えない）。
"""
import re
import sys
import zipfile
from pathlib import Path, PurePosixPath

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.util import Inches

from _build.arauchi_v4.lib import NAVY, RED, TEXT, add_box, move, set_paras
from _build.arauchi_v10.build import by_id, c14n, clone, meiryo, normalize_decl, rels_of, runs, slide_order

ROOT = Path(__file__).resolve().parents[2]
IMG = ROOT / "_images/arauchi_zengo_chintai_color.png"
IMGW, IMGH = 2400, 1400
INSERT_AT = 14                    # 新しいページの番号（配信のページの前）
SRC_HEADER = 7                    # タイトル・リードの帯を写すページ（競合分析）
GRAY, TINT = "7F7F7F", "F4F7FF"
COL = {"A": "3467B2", "B": "00897B", "C": "8E4EC6", "D": "D9661F"}

TITLE = "前後検索「賃貸」：物件を見た後は、会社を比べ、費用・審査を調べる"
LEAD = "「賃貸」を検索した方が、その前後15日間に検索した言葉を、時期ごとに色分け"
LEGEND = [("A", "エリアで探す"), ("B", "ポータルサイトで物件を見る"), ("C", "不動産会社を比べる"), ("D", "初期費用・審査を調べる")]
ROWS = [("A", "検索前", "検索語のほとんどが「エリア名＋賃貸」"),
        ("B", "当日", "大手ポータルサイトの検索が最も多い"),
        ("C", "当日〜3日後", "不動産会社名の検索が続く"),
        ("D", "3〜15日後", "「賃貸 審査 落ちる確率」「賃貸 初期費用 抑える」など")]
NOTE = "出典：LINEヤフー社提供の前後検索データ（検索起点：「賃貸」／2022年9月〜2023年9月）"
X0, Y0, W = 0.55, 1.98, 6.5       # 画像の位置（インチ）
H = W * IMGH / IMGW


def px(p):
    return X0 + p / IMGW * W


def py(p):
    return Y0 + p / IMGH * H


def alpha_fill(shape, color, alpha):
    sp_pr = shape._element.spPr
    for tag in ("a:solidFill", "a:noFill", "a:ln"):
        for e in sp_pr.findall(qn(tag)):
            sp_pr.remove(e)
    f = etree.SubElement(sp_pr, qn("a:solidFill"))
    c = etree.SubElement(f, qn("a:srgbClr"), val=color)
    etree.SubElement(c, qn("a:alpha"), val=str(alpha))
    ln = etree.SubElement(sp_pr, qn("a:ln"))
    etree.SubElement(ln, qn("a:noFill"))


def build_slide(prs):
    from pptx.enum.shapes import MSO_SHAPE
    src = prs.slides[SRC_HEADER - 1]
    s = prs.slides.add_slide(src.slide_layout)
    for sh in list(s.shapes):
        sh._element.getparent().remove(sh._element)
    t = clone(s, by_id(src, 2))
    set_paras(t, [TITLE])
    meiryo(t)
    move(t, w=9.0)
    for sid in (3, 4, 5):
        clone(s, by_id(src, sid))
    lead = clone(s, by_id(src, 6))
    set_paras(lead, [LEAD])
    meiryo(lead)

    # 凡例（画像の上）
    for (k, name), x, w in zip(LEGEND, (X0, 1.72, 3.92, 5.47), (1.12, 2.15, 1.5, 1.75)):
        add_box(s, x, 1.64, w, 0.3, [runs(("●", dict(size=10.5, color=COL[k])), (" " + name, dict(size=9, color=TEXT)))],
                anchor="ctr", ins=(0, 0, 0, 0), name=f"凡例 {name}")

    pic = s.shapes.add_picture(str(IMG), Inches(X0), Inches(Y0), width=Inches(W))
    pic.name = "前後検索 賃貸"
    pic.line.color.rgb = RGBColor.from_string("D9D9D9")
    pic.line.width = 9525

    # LINE追加の時期（当日〜3日後）：0日＝1260px・3日＝1819px
    gx0, gx1 = px(1260), px(1819)
    band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(gx0), Inches(py(20)), Inches(gx1 - gx0), Inches(py(1325) - py(20)))
    band.name = "LINE追加の時期"
    alpha_fill(band, "06C755", 14000)
    band.shadow.inherit = False
    add_box(s, (gx0 + gx1) / 2 - 0.55, Y0 + H + 0.04, 1.1, 0.26,
            [runs(("LINE追加", dict(size=9.5, bold=True, color="00873C")))],
            fill="C6EFCE", align="ctr", anchor="ctr", ins=(0, 0, 0, 0), name="LINE追加 札")

    # 右：時期ごとに見えること
    RX, RW, rh, gap = 7.3, 2.98, 0.92, 0.1
    for i, (k, when, body) in enumerate(ROWS):
        y = 1.68 + i * (rh + gap)
        add_box(s, RX, y, 0.07, rh, [runs(("", dict(size=8)))], fill=COL[k], ins=(0, 0, 0, 0), name=f"時期{i + 1} 色")
        add_box(s, RX + 0.17, y + 0.02, RW - 0.17, 0.32, [runs((when, dict(size=12, bold=True, color=COL[k])))],
                anchor="ctr", ins=(0, 0, 0, 0), name=f"時期{i + 1}")
        add_box(s, RX + 0.17, y + 0.34, RW - 0.17, rh - 0.36, [runs((body, dict(size=10.5, color=TEXT)))],
                anchor="t", line_spacing=115, ins=(0, 0, 0, 0), name=f"時期{i + 1} 本文")

    add_box(s, 0.55, 6.1, 9.73, 0.72, [
        runs(("会社を比べる時期に「仲介手数料100円」、", dict(size=12.5, bold=True, color=RED))),
        runs(("費用・審査を調べる時期に「初期費用の交渉・お申込みの準備」を届ける", dict(size=12.5, bold=True, color=NAVY)))],
        fill=TINT, align="ctr", anchor="ctr", line_spacing=115, name="示唆")
    add_box(s, 0.55, 6.86, 9.73, 0.26, [runs((NOTE, dict(size=9, color=GRAY)))], anchor="ctr", ins=(0, 0, 0, 0), name="出典")
    return s


def repack(base_path, built_path, out_path, at=INSERT_AT):
    zb, zn = zipfile.ZipFile(base_path), zipfile.ZipFile(built_path)
    bo, no = slide_order(zb), slide_order(zn)
    n = len(bo)
    assert len(no) == n + 1, (len(bo), len(no))
    new_built = no[-1]                                   # python-pptx は末尾に足す
    for i, (b, x) in enumerate(zip(bo, no[:n]), 1):      # ほかのページは1バイトも変えない
        assert c14n(zb.read(b)) == c14n(zn.read(x)), f"S{i} が変わっている"
    names = set(zb.namelist())
    k = n + 1
    while f"ppt/slides/slide{k}.xml" in names:
        k += 1
    part = f"ppt/slides/slide{k}.xml"
    rels_xml = zn.read(rels_of(new_built)).decode("utf-8")
    targets = re.findall(r'Target="([^"]+)"', rels_xml)
    media = []
    for tg in targets:
        if "slideLayout" in tg:
            continue
        assert "/media/" in tg, rels_xml
        full = str(PurePosixPath("ppt/slides") / tg).replace("ppt/slides/../", "ppt/")
        assert full not in names, full                   # 元のzipの画像と名前がぶつからないこと
        media.append(full)
    assert len(media) == 1, media

    prels = zb.read("ppt/_rels/presentation.xml.rels").decode("utf-8")
    rid_nums = [int(x) for x in re.findall(r'Id="rId(\d+)"', prels)]
    rid = f"rId{max(rid_nums) + 1}"
    prels = prels.replace("</Relationships>",
                          f'<Relationship Id="{rid}" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
                          f'relationships/slide" Target="slides/slide{k}.xml"/></Relationships>')
    pres = zb.read("ppt/presentation.xml").decode("utf-8")
    sld = re.findall(r"<p:sldId [^>]*/>", pres)
    assert len(sld) == n
    new_id = max(int(re.search(r'id="(\d+)"', x).group(1)) for x in sld) + 1
    sld.insert(at - 1, f'<p:sldId id="{new_id}" r:id="{rid}"/>')
    lst = re.search(r"<p:sldIdLst>.*?</p:sldIdLst>", pres).group(0)
    pres = pres.replace(lst, "<p:sldIdLst>" + "".join(sld) + "</p:sldIdLst>")
    ct = zb.read("[Content_Types].xml").decode("utf-8")
    ct = ct.replace("</Types>", f'<Override PartName="/{part}" ContentType="application/vnd.openxmlformats-officedocument.'
                                f'presentationml.slide+xml"/></Types>')
    if 'Extension="png"' not in ct:
        ct = ct.replace("<Default ", '<Default Extension="png" ContentType="image/png"/><Default ', 1)
    app = zb.read("docProps/app.xml").decode("utf-8")
    app = app.replace(f"<Slides>{n}</Slides>", f"<Slides>{n + 1}</Slides>", 1)
    head = f"<vt:lpstr>スライド タイトル</vt:lpstr></vt:variant><vt:variant><vt:i4>{n}</vt:i4>"
    if head in app:
        app = app.replace(head, head.replace(f"<vt:i4>{n}</vt:i4>", f"<vt:i4>{n + 1}</vt:i4>"))
        m = re.search(r'<vt:vector size="(\d+)" baseType="lpstr">', app)
        app = app.replace(m.group(0), f'<vt:vector size="{int(m.group(1)) + 1}" baseType="lpstr">')
        app = app.replace("</vt:vector></TitlesOfParts>", "<vt:lpstr>PowerPoint プレゼンテーション</vt:lpstr></vt:vector></TitlesOfParts>")

    with zipfile.ZipFile(out_path, "w") as zo:
        for info in zb.infolist():
            data = {"ppt/presentation.xml": pres.encode("utf-8"),
                    "ppt/_rels/presentation.xml.rels": prels.encode("utf-8"),
                    "[Content_Types].xml": ct.encode("utf-8"),
                    "docProps/app.xml": app.encode("utf-8")}.get(info.filename)
            zo.writestr(info, data if data is not None else zb.read(info.filename), compress_type=info.compress_type)
        zo.writestr(part, normalize_decl(zn.read(new_built)), compress_type=zipfile.ZIP_DEFLATED)
        zo.writestr(rels_of(part), rels_xml.encode("utf-8"), compress_type=zipfile.ZIP_DEFLATED)
        for m in media:
            zo.writestr(m, zn.read(m), compress_type=zipfile.ZIP_STORED)
    return part


def main(base, out):
    tmp = out + ".tmp.pptx"
    prs = Presentation(base)
    build_slide(prs)
    prs.save(tmp)
    part = repack(base, tmp, out)
    Path(tmp).unlink()
    print("saved:", out, "／ 追加したパーツ:", part, f"（{INSERT_AT}ページ目）")


if __name__ == "__main__":
    main(*sys.argv[1:3])
