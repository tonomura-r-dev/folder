# -*- coding: utf-8 -*-
"""すなお：事例紹介（保険・イベント集客）のページを作り直す（2026-10-20 殿村さん「S14は流石にひどいだろ」）。

    python -m _build.patch_sunao_v62_jirei <資料.pptx> <出力.pptx>

  左に小さな丸3つ・右半分が空、で何が言いたいか伝わらなかった。
  → 「業種／LINEでしたこと／結果（大きな数字）／貴社での活かし方」の表にする。数字・出典・業種名（社名は書かない）は元のまま。
  出典：LINEヤフー for Business 導入事例（第一生命 2024年4月／日本生命 2025年2月／JPCスポーツ教室 2021年8月）。ほかのページは触らない。
"""
import sys

from pptx import Presentation
from pptx.util import Inches

from _build.arauchi_v10.build import runs
from _build.arauchi_v4.lib import add_table
from _build.build_sunao_v1 import all_meiryo
from _build.patch_sunao_v53_kaizen import BAND, NAVY, TEXT, cell, clear, head, label, lead_band, note, set_one


def find(prs, text):
    hit = [s for s in prs.slides if any(sh.has_text_frame and text in sh.text_frame.text for sh in s.shapes)]
    assert len(hit) == 1, (text, len(hit))
    return hit[0]


def main(src, out):
    prs = Presentation(src)
    s = find(prs, "保険・イベント集客におけるLINE公式アカウントの活用事例")
    doshen = find(prs, "貴社のユーザー動線イメージ")
    t = clear(s)
    set_one(t, "事例紹介：保険・イベント集客でのLINEの成果")
    t.width = Inches(8.4)
    lead_band(s, doshen, "LINEヤフー for Business に掲載されている、保険・イベント集客の事例")
    rows = [[head("業種"), head("LINEでしたこと"), head("結果"), head("貴社での活かし方")]]
    data = [
        ("生命保険会社", "同じ内容をメールとLINEで配信", "約8倍", "資料請求のCV率（メール比）",
         "次回セミナーの案内を、LINEで届ける"),
        ("生命保険会社", "営業職員経由で、お客様にLINEの友だち追加を案内", "5.2%", "営業職員経由の友だちのブロック率（自然に友だちになった方は74.9%）",
         "接点のある方を友だちにし、配信を届け続ける"),
        ("教室・スクール", "見学会（イベント）の案内をLINEで配信", "約7割", "見学会の参加者の入会率",
         "セミナー参加後に、サロンの入会を案内する"),
    ]
    for i, (gyo, did, num, what, use) in enumerate(data):
        fill = BAND[i % 2]
        res = cell([runs((num, dict(size=22, bold=True, color=NAVY))), runs((what, dict(size=11, color=TEXT)))],
                   fill=fill, align="ctr")
        rows.append([label(gyo, fill), cell(did, fill=fill), res, cell(use, fill=fill, bold=True, color=NAVY)])
    add_table(s, 0.46, 1.95, [1.75, 2.6, 2.75, 2.8], [0.42, 1.2, 1.3, 1.2], rows, name="事例")
    note(s, 6.4, ["※LINEヤフー for Business 掲載の導入事例（数値は各社調べ）。生命保険会社＝2024年4月・2025年2月公開、教室・スクール＝2021年8月公開"])
    all_meiryo([s])
    prs.save(out)
    print("saved:", out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
