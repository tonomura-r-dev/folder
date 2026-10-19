# -*- coding: utf-8 -*-
"""株式会社くふう住まい（くふうイエタテフェア）LINE公式アカウント運用のご提案 v1。
広告の初訪資料（リスティング・SNS・DSP）に添える「LINEの部分」だけの簡潔な資料（2026-10-19 殿村さん）。
「注文住宅の汎用資料」と「サンクスLINEの資料」を組み合わせ、先方の課題（予約〜来場の来場率80%）に話を合わせる。

    python -m _build.build_kufu_v1 [出力.pptx]

  並び（本編）：
    表紙（注文住宅の汎用資料の表紙に「御中」とサブタイトル）
    現状と課題（新）＝与件の現状・目標の数字＋予約から当日まで LINE でつながっていない（LINE登録は当日の受付）
    広告と一緒にサンクスLINEを実施する理由／サンクスLINE｜機能・費用（サンクスLINE業界別 ver3.7）
    サンクスLINE｜導入・活用実績（サンクスLINE Appendix ver1.2）
    来場率と申込みCPA（新）＝来場率ごとの必要申込み数・CPA（与件の②③に答える計算。85%・90%は仮に）
    前後検索「注文住宅」（タナカホームv2の実データのページ。汎用資料の前後検索は推計なので使わない）
    全体設計想定／施策展開図（注文住宅の汎用資料）
    DYMのコンサル支援内容・費用（サンクスLINE業界別 ver3.7）
    裏表紙 → Appendix（タナカホームの Appendix＝正。CLAUDE.md 2026-10-19）
  汎用資料のステップ配信案・予算診断・シーズナリティは、根拠に推計の前後検索データを使っているので入れていない。
"""
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches

from _build.arauchi_v4.lib import TEXT, add_box, add_table, move, set_paras
from _build.arauchi_v10.build import by_id, clone, runs
from _build.build_sunao_v1 import all_meiryo
from _build.patch_tanaka_v4 import set_lines
from _build.patch_tanaka_v5 import copy_slide

ROOT = Path(__file__).resolve().parent.parent
CHUMON = ROOT / "注文住宅業界_LINEOA施策提案.pptx"
THANKS = ROOT / "20261007_サンクスLINE誘導のご提案_業界別ver3.7.pptx"
THANKS_APPX = ROOT / "20261009_サンクスLINE誘導のご提案_Appendix_ver1.2.pptx"
TANAKA = ROOT / "202610_タナカホーム株式会社御中_LINE公式アカウント運用のご提案_v2.pptx"
ARAUCHI = ROOT / "202610_株式会社あらうち御中_LINE公式アカウント運用のご提案_v10.pptx"
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "202610_株式会社くふう住まい御中_LINE公式アカウント運用のご提案.pptx"

NAVY = "1F2C5C"
LBLUE = "DCE6F2"
LGRAY = "F2F2F2"
GRAY = "7F7F7F"
RED = "C00000"
B = {k: (0.75, "BFBFBF") for k in "lrtb"}
X0, W = 0.7, 9.43


def cell(text, size=13, bold=False, color=TEXT, fill="FFFFFF", align="ctr"):
    return dict(text=text, size=size, bold=bold, color=color, fill=fill, align=align, borders=B,
                ins=(0.08, 0.03, 0.08, 0.03))


def head(text, size=13):
    return cell(text, size, True, "FFFFFF", NAVY)


def label(text, size=13):
    return cell(text, size, True, NAVY, LBLUE)


def note(s, text, y, h=0.2):
    add_box(s, X0, y, W, h, [runs((text, dict(size=9, color=GRAY)))], anchor="t", ins=(0, 0, 0, 0), name="注記")


def band(s, text, y, h=0.55, size=15):
    add_box(s, X0, y, W, h, [runs((text, dict(size=size, bold=True, color=NAVY)))], fill=LGRAY, align="ctr",
            anchor="ctr", name="示唆")


def new_page(prs, tmpl, title, lead):
    """サンクスLINEの資料と同じ見出し（タイトル＋リード＋区切り線）のページ。中身は空にする"""
    s = copy_slide(prs, tmpl)
    for sh in list(s.shapes):
        if sh.shape_id not in (3, 4, 5):
            sh._element.getparent().remove(sh._element)
    set_lines(by_id(s, 3), [title])
    set_lines(by_id(s, 4), [lead])
    return s


def lift(s, top=1.6, bottom=7.05, ratio=0.33):
    """中身を上寄りの中央へ下げる：空きの上1/3・下2/3（CLAUDE.md 2026-10-19 すなおで決めた位置）"""
    body = [sh for sh in s.shapes if sh.shape_id not in (3, 4, 5)]
    y0 = min(sh.top for sh in body)
    y1 = max(sh.top + sh.height for sh in body)
    dy = int(Inches(top) - y0 + (Inches(bottom) - Inches(top) - (y1 - y0)) * ratio)
    if dy > 0:
        for sh in body:
            sh.top = sh.top + dy


# ---------------------------------------------------------------- 表紙
def p_cover(s, ar_cover):
    sub = by_id(s, 6)
    set_paras(sub, ["予約から来場までをLINEでつなぎ、", "実来場を増やす"])
    move(sub, x=1.76, y=2.3, w=7.25, h=0.95)
    for p in sub.text_frame.paragraphs:                        # 元の表紙と同じ中央揃え
        p.alignment = PP_ALIGN.CENTER
    for sid in (3, 7):                                         # 「- 2026最新版 -」「業界：注文住宅」
        e = by_id(s, sid)._element
        e.getparent().remove(e)
    set_paras(clone(s, by_id(ar_cover, 31)), ["株式会社くふう住まい 御中"])


# ---------------------------------------------------------------- 現状と課題（新）
def p_kadai(prs, tmpl):
    s = new_page(prs, tmpl, "現状と課題｜予約から来場まで、LINEでつながっていない",
                 "来場者獲得単価を、今の37,500円から24,000円以下へ。申込みCPAと来場率の両方で改善する")
    rows = [[head(""), head("現状"), head("目標")],
            [label("申込みCPA"), cell("30,000円"), cell("19,000円以下")],
            [label("申込み件数"), cell("約40件"), cell("63件以上")],
            [label("来場率"), cell("80%"), cell("80%以上")],
            [label("実来場者数"), cell("約32人"), cell("50人以上")],
            [label("来場者獲得単価"), cell("約37,500円", bold=True, color=NAVY), cell("24,000円以下", bold=True, color=NAVY)]]
    add_table(s, X0, 1.85, [1.9, 1.45, 1.45], [0.42] + [0.46] * 5, rows, name="現状と目標")
    note(s, "※1開催あたり（広告予算120万円）", y=4.62)
    fx, fw = 6.0, 4.13
    steps = [("WEBで来場予約", NAVY, "FFFFFF"), ("予約確認メール", LBLUE, NAVY),
             ("当日の受付でLINE登録", LBLUE, NAVY)]
    y = 1.85
    for k, (t, fill, color) in enumerate(steps):
        add_box(s, fx, y, fw, 0.5, [runs((t, dict(size=14, bold=True, color=color)))], fill=fill, align="ctr",
                anchor="ctr", name=f"流れ{k + 1}")
        y += 0.5
        if k == 1:
            add_box(s, fx, y + 0.05, fw, 0.85, [runs(("予約から当日まで、LINEでつながっていない", dict(size=13, bold=True, color=RED))),
                                               runs(("キャンセル 20%", dict(size=16, bold=True, color=RED)))],
                    line=RED, align="ctr", anchor="ctr", line_spacing=120, ins=(0.08, 0.03, 0.08, 0.03), name="切れる")
            y += 0.95
        elif k < 2:
            add_box(s, fx, y, fw, 0.3, [runs(("▼", dict(size=12, color=NAVY)))], align="ctr", anchor="ctr",
                    ins=(0, 0, 0, 0), name="矢印")
            y += 0.3
    note(s, "※くふうイエタテフェアのサイトより（2026年10月19日確認）", y=4.62)
    move(list(s.shapes)[-1], x=fx, w=fw)
    band(s, "予約の直後にLINEでつながれば、当日までの案内で来場率を上げられる", y=5.1)
    lift(s)
    return s


# ---------------------------------------------------------------- 来場率と申込みCPA（新）
def p_raijo(prs, tmpl):
    s = new_page(prs, tmpl, "来場率と申込みCPA｜来場率が上がると、必要なCPAが緩む",
                 "予算120万円で実来場50人を目指す場合の、来場率ごとの必要な申込み数とCPA")
    rows = [[head("来場率"), head("来場50人に\n必要な申込み"), head("必要な申込みCPA\n（予算120万円）"),
             head("CPA30,000円のままの\n来場者獲得単価")]]
    for rate, apps, cpa, cur in (("80%（今）", "63件", "19,048円", "37,500円"),
                                 ("85%", "59件", "20,339円", "35,294円"),
                                 ("90%", "56件", "21,429円", "33,333円")):
        rows.append([label(rate, 14), cell(apps, 14), cell(cpa, 15, True, NAVY), cell(cur, 14)])
    add_table(s, X0, 1.85, [1.9, 2.2, 2.55, 2.78], [0.7, 0.55, 0.55, 0.55], rows, name="来場率とCPA")
    for k, t in enumerate(["※来場率85%・90%は仮に置いた数字。申込み1件＝1人として計算（ご提示いただいた条件のまま）",
                           "※必要な申込みは小数点以下を切り上げ、CPAは予算120万円÷必要な申込みで計算"]):
        note(s, t, y=4.35 + k * 0.2)
    lift(s)
    return s


def main(out):
    prs = Presentation(CHUMON)
    sl = list(prs.slides)
    cover, zentai, tenkai, back = sl[0], sl[2], sl[19], sl[35]
    assert by_id(cover, 2).text_frame.text == "LINE公式アカウント運用のご提案"
    assert "全体設計想定" in "".join(sh.text_frame.text for sh in zentai.shapes if sh.has_text_frame)
    assert "施策展開図" in "".join(sh.text_frame.text for sh in tenkai.shapes if sh.has_text_frame)
    th = list(Presentation(THANKS).slides)
    assert by_id(th[0], 3).text_frame.text == "広告と一緒にサンクスLINEを実施する理由"
    assert by_id(th[3], 3).text_frame.text == "実例｜予約後の来院率"
    assert by_id(th[5], 3).text_frame.text == "DYMのコンサル支援内容・費用"
    ta = list(Presentation(THANKS_APPX).slides)
    assert "導入・活用実績" in "".join(sh.text_frame.text for sh in ta[2].shapes if sh.has_text_frame)
    tk = list(Presentation(TANAKA).slides)
    assert by_id(tk[7], 2).text_frame.text.startswith("前後検索「注文住宅」")
    ar_cover = Presentation(ARAUCHI).slides[0]

    main_pages = [cover]
    main_pages.append(p_kadai(prs, th[3]))
    main_pages += [copy_slide(prs, th[0]), copy_slide(prs, th[1]), copy_slide(prs, ta[2])]
    main_pages.append(p_raijo(prs, th[3]))
    main_pages += [copy_slide(prs, tk[7]), zentai, tenkai, copy_slide(prs, th[5]), back]
    appx = [copy_slide(prs, tk[i - 1]) for i in range(19, 32)]   # タナカホームの Appendix（扉〜都度発注プラン）
    p_cover(cover, ar_cover)

    keep = main_pages + appx
    lst = prs.slides._sldIdLst
    by_part = {prs.part.related_part(e.rId): e for e in lst}
    for e in list(lst):
        lst.remove(e)
    keep_parts = {s.part for s in keep}
    for part, e in by_part.items():
        if part not in keep_parts:
            prs.part.drop_rel(e.rId)
    for s in keep:
        lst.append(by_part[s.part])
    all_meiryo(keep)
    prs.save(out)
    print("saved:", Path(out).name, len(prs.slides), "枚")


if __name__ == "__main__":
    main(OUT)
