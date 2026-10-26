"""
LINE広告トレンドレポート10月号（殿村さん版＝正）の S5（LINE VOOM）を、
VOOM終了の資料（20260930_YDA_LINE_VOOM_Service_Termination）に沿って組み直す。
影響3パターン（上段3枚）＋やるべきこと（下段1枚）。見出しと目次のP.7も直す。
S3・S4など他のページには触らない。
使い方: python3 _build/patch_line_ads_202610_voom.py <pptx>
"""
import sys
import copy
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.oxml.ns import qn
from pptx.enum.shapes import MSO_SHAPE
import trend_report_engine as engine

F = sys.argv[1]
prs = Presentation(F)
sl = list(prs.slides)
s = next(x for x in sl if any(sh.has_text_frame and sh.text_frame.text.strip().startswith("【LINE VOOM】")
                              for sh in x.shapes))
TITLE = "【LINE VOOM】サービス終了に伴う広告配信の終了"


def text(sh):
    return sh.text_frame.text.strip() if sh.has_text_frame else ""


head = next(x for x in s.shapes if text(x) == "■ 詳細")
body = [x for x in s.shapes if x.top >= head.top + head.height - Inches(0.02) and x.top < Inches(7.0)]
cards = [x for x in body if x.shape_type == 1 and x.auto_shape_type == MSO_SHAPE.ROUNDED_RECTANGLE]
assert len(cards) == 2, len(cards)
for x in body:
    if x not in cards:
        x._element.getparent().remove(x._element)

mx = max(int(e.get("id")) for e in s._element.iter(qn("p:cNvPr")))
for k in range(2):
    el = copy.deepcopy(cards[k]._element)
    cards[-1]._element.addnext(el)
    cn = el.find(".//" + qn("p:cNvPr"))
    cn.set("id", str(mx + 1 + k))
    cn.set("name", cn.get("name") + f"_v{k}")
    cards.append(next(x for x in s.shapes if x._element is el))


def setc(c, h, b, hp, bp):
    p0, p1 = c.text_frame.paragraphs
    for p, t in ((p0, h), (p1, b)):
        p.runs[0].text = t
        for r in list(p.runs[1:]):
            r._r.getparent().remove(r._r)
        for br in list(p._p.findall(qn("a:br"))):
            br.getparent().remove(br)
    for r in p0.runs:
        r.font.size = Pt(hp)
    for r in p1.runs:
        r.font.size = Pt(bp)


def add_break(c, second):
    p1 = c.text_frame.paragraphs[1]
    r = p1.runs[0]
    br = p1._p.makeelement(qn("a:br"), {})
    rpr = r._r.find(qn("a:rPr"))
    if rpr is not None:
        br.append(copy.deepcopy(rpr))
    r2 = copy.deepcopy(r._r)
    r2.find(qn("a:t")).text = second
    r._r.addnext(br)
    br.addnext(r2)


L, W, GAP = 0.6, 9.64, 0.2
cw = (W - 2 * GAP) / 3
y1, h1 = 3.4, 1.45
three = [
    ("① VOOM面だけ指定", "終了後は配信されない。配信設定がONのままでも、配信できる面がなくなるため。"),
    ("② VOOM面＋他の面", "VOOM面以外への配信は、そのまま続く。"),
    ("③ VOOM面を除外", "VOOM終了による影響はない。"),
]
for i, (c, (h, b)) in enumerate(zip(cards[:3], three)):
    c.left, c.top, c.width, c.height = Inches(L + i * (cw + GAP)), Inches(y1), Inches(cw), Inches(h1)
    setc(c, h, b, 17, 12)
c = cards[3]
c.left, c.top, c.width, c.height = Inches(L), Inches(y1 + h1 + 0.25), Inches(W), Inches(1.55)
setc(c, "やるべきこと",
     "LINEヤフー広告：配信先リストの中のVOOM関連の設定は自動では消えないため、必要に応じて削除する。", 19, 13)
add_break(c, "LINE広告：配信先の「LINE VOOM」は9月30日に選択肢ごと管理画面から消えるため、VOOMだけを選んでいた広告は配信先を見直す。")

engine.set_single_run_text(next(x for x in s.shapes if text(x).startswith("【")), TITLE)
d = next(x for x in s.shapes if x.has_text_frame and len(x.text_frame.paragraphs) == 3
         and x.text_frame.paragraphs[0].alignment == engine.PP_ALIGN.CENTER)
engine.set_multiline_text(d, [
    "LINEのリニューアルに伴い、2026年9月30日（水）にLINE VOOMのサービスが終了した。",
    "LINEヤフー広告 ディスプレイ広告（運用型）・LINE広告ともに、VOOM面への広告配信も終了。",
    "終了前の配信実績は、終了後も配信面ごとに確認できる。",
])

toc = sl[1]
ts = next(x for x in toc.shapes if text(x).startswith("【"))
lines = [p.text for p in ts.text_frame.paragraphs]
lines = [TITLE if ln.startswith("【LINE VOOM】") else ln for ln in lines]
engine.set_multiline_text(ts, lines)

engine.force_font(prs)
prs.save(F)
print("saved", F)
