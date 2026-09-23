# -*- coding: utf-8 -*-
"""コンサル費②を15万に下げた数値を資料に反映（Slide 24）"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
TARGET = ROOT / "20260804_株式会社ビイサイドプランニング御中_公式LINEのご提案v08.pptx"

NAVY = "1F285A"; INK = "333333"; MUT = "808080"; WHITE = "FFFFFF"

prs = Presentation(str(TARGET))
slides = list(prs.slides)


def set_font(run, size, bold=None, color=INK, name="メイリオ"):
    f = run.font; f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    f.name = name
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        e = rPr.find(qn(tag))
        if e is None:
            e = rPr.makeelement(qn(tag), {}); rPr.append(e)
        e.set("typeface", name)
    f.color.rgb = RGBColor.from_string(color)


ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
ANCH = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}


def reset_tf(tf):
    for para in tf.paragraphs[1:]:
        para._p.getparent().remove(para._p)
    for r in list(tf.paragraphs[0].runs):
        r._r.getparent().remove(r._r)


def put_text(tf, paras, anchor="m", ml=0.04, mr=0.04, mt=0.0, mb=0.0):
    reset_tf(tf); tf.word_wrap = True
    tf.margin_left = Inches(ml); tf.margin_right = Inches(mr)
    tf.margin_top = Inches(mt); tf.margin_bottom = Inches(mb)
    tf.vertical_anchor = ANCH[anchor]
    first = True
    for p in paras:
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        para.alignment = ALIGN[p.get("align", "l")]
        if p.get("sa") is not None:
            para.space_after = Pt(p["sa"])
        if p.get("ls") is not None:
            para.line_spacing = p["ls"]
        for (t, sz, b, c) in p["runs"]:
            r = para.add_run(); r.text = t
            set_font(r, sz, b, c)


# Slide 24（想定の費用対効果）を特定
target = None
for s in slides:
    txt = "\n".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame)
    if "想定の費用対効果" in txt and "6か月後" in txt:
        target = s
        break
assert target is not None, "Slide 24が見つからない"

done = []
for sh in target.shapes:
    if not sh.has_text_frame:
        continue
    t = sh.text_frame.text.strip()
    if t.startswith("25.98万円"):
        put_text(sh.text_frame, [{"runs": [("20.98万円", 28, True, NAVY)], "align": "c"}])
        done.append("月次 25.98→20.98万円")
    elif t.startswith("259,800円"):
        put_text(sh.text_frame, [{"runs": [("209,800円／月", 9.5, None, INK)], "align": "c"}])
        done.append("コスト 259,800→209,800円")
    elif t.startswith("12,286"):
        put_text(sh.text_frame, [{"runs": [("10,232円", 9.5, None, INK)], "align": "c"}])
        done.append("CPA 12,286→10,232円")
    elif "費用内訳" in t:
        put_text(sh.text_frame,
                 [{"runs": [("＜費用内訳＞", 11, True, NAVY)], "sa": 3},
                  {"runs": [("初期：構築20万＋離脱防止1.5万＋ツール2万", 9, None, INK)], "ls": 1.15, "sa": 2},
                  {"runs": [("月次：コンサル15万＋ツール2.98万＋アカウント1.5万＋離脱防止1.5万", 9, None, INK)], "ls": 1.15}],
                 anchor="t", ml=0.0, mr=0.0)
        done.append("費用内訳")
    elif "LINEで増える来場者の獲得単価" in t:
        put_text(sh.text_frame,
                 [{"runs": [("LINEで増える来場者の獲得単価　", 12, True, WHITE),
                            ("約11,200円", 19, True, WHITE),
                            ("　（現状の来場単価 16,000〜18,000円 → 30〜38%減）", 11.5, None, WHITE)],
                   "align": "c"}])
        done.append("獲得単価 13,500→11,200円")

for d in done:
    print("updated:", d)
prs.save(str(TARGET))
print("saved:", TARGET)
