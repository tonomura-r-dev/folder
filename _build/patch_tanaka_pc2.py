# -*- coding: utf-8 -*-
"""タナカホーム：配信・試算・料金プランを、あらうちv10のページ（殿村さんが見やすいと指定した形）で作り直す（2026-10-16）。
patch_tanaka_pc1.main から呼ぶ（単体では使わない）。

  あらうちv10の 15（配信）・17（目標と費用対効果）・18（料金プラン）・19（運用内容）を写し、中身をタナカホームに差し替える。
  - 配信：前後検索「注文住宅」の順番（当日＝会社を比べる／1〜7日＝費用・住宅ローン／7〜15日＝断熱・間取り・ルームツアー）で
          ステップ配信 全5通（SIM ver1.2の5通）＋企画配信 月4回程度
  - 試算：プラン①（総合運用）の初期・月次と、SIM ver1.2 の6か月後の数字・月ごとの資料請求と来場予約
  - 料金プラン：①総合運用 26万円〜／②総合運用＋友だち獲得（LINE広告）46万円〜／③初期導入 6万円〜（初期は3つとも26.5万円〜）
  - 運用内容：3つのプランの対応範囲
  「松竹梅」の文字は使わない（殿村さん）。良い数字は赤い文字にしない（紺に変える）。
"""
import copy

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

from _build.arauchi_v4.lib import NAVY, TEXT, add_box, add_table
from _build.arauchi_v10.build import TINT, by_id, meiryo, runs
from _build.build_tanaka_v1 import GRAY, head_cell, cell, label_cell, note
from _build.patch_tanaka_v4 import set_lines
from _build.patch_tanaka_v5 import ROOT, SIM, copy_slide

ARAUCHI = ROOT / "202610_株式会社あらうち御中_LINE公式アカウント運用のご提案_v10.pptx"


def lines(s, sid, texts):
    set_lines(by_id(s, sid), texts if isinstance(texts, list) else texts.split("\n"))


def set_cell(c, text):
    """表のセルの文字を差し替える（最初の段落・最初の文字列の書式を使う）"""
    tx = c._tc.find(qn("a:txBody"))
    ps = tx.findall(qn("a:p"))
    tmpl_p = ps[0]
    tmpl_r = copy.deepcopy(tx.find(".//" + qn("a:r")))
    for p in ps:
        tx.remove(p)
    for ln in text.split("\n"):
        p = copy.deepcopy(tmpl_p)
        for e in p.findall(qn("a:r")) + p.findall(qn("a:br")):
            p.remove(e)
        r = copy.deepcopy(tmpl_r)
        r.find(qn("a:t")).text = ln
        end = p.find(qn("a:endParaRPr"))
        if end is not None:
            end.addprevious(r)
        else:
            p.append(r)
        tx.append(p)


def color_runs(shape, rgb):
    for r in shape.text_frame._txBody.iter(qn("a:r")):
        rpr = r.find(qn("a:rPr"))
        for f in rpr.findall(qn("a:solidFill")):
            rpr.remove(f)
        sf = rpr.makeelement(qn("a:solidFill"), {})
        sf.append(sf.makeelement(qn("a:srgbClr"), {"val": rgb}))
        rpr.insert(0, sf)


# ---------------------------------------------------------------- 配信（あらうち15）
def p_haishin(prs, src):
    s = copy_slide(prs, src)
    lines(s, 385, "解決策：LINEの配信でナーチャリングを")
    lines(s, 386, "前後検索で見える検討の順番に合わせた、14日間のステップ配信（全5通）と月4回程度の企画配信で、来場予約を後押し")
    lines(s, 387, "[ステップ配信・全5通｜新しい友だちへ追加後に自動配信]")
    steps = [((388, 393, 395, 397), ("当日〜1日後", "施工事例・会社の特長\n土地探しのご相談", "比べる会社の\n候補に入る", "資料請求へ")),
             ((398, 403, 405, 407), ("2〜7日後", "住宅ローン・坪単価の考え方\n平屋・2階建ての価格の目安", "費用の不安を解消", "来場予約へ")),
             ((408, 413, 415, 417), ("8〜14日後", "青森の冬を快適にする断熱・間取り\nルームツアー動画", "住みたいイメージを\n育てる", "来場予約へ"))]
    kikaku = [((419, 421, 423, 425, 427), ("友だち全員", "絞り込みなし", "見学会・\nモデルハウス公開", "見に行く理由を\nつくる", "来場予約へ")),
              ((428, 430, 432, 434, 436), ("来場予約前の方", "来場予約済みを除く\n（チャットタグ）", "予算の簡易診断\n平屋か2階建てか", "希望を聞き分ける", "来場予約へ")),
              ((437, 439, 441, 443, 445), ("費用に関心がある方", "住宅ローンの配信を\nクリックした方", "資金計画の\n個別相談のご案内", "費用の不安を解消", "来場予約へ")),
              ((446, 448, 450, 452, 454), ("来場済み・\n契約前の方", "来場済みのタグ", "お客様の声\n構造見学会", "契約前の後押し", "商談・契約へ"))]
    for ids, texts in steps + kikaku:
        for sid, t in zip(ids, texts):
            lines(s, sid, t)
    return s


# ---------------------------------------------------------------- 試算（あらうち17）
def p_sim(prs, src):
    s = copy_slide(prs, src)
    doc, visit, total = SIM["doc"], SIM["visit"], SIM["total"]
    lines(s, 339, f"試算：月26万円で、6か月目に資料請求 月{doc[5]}件・来場予約 月{visit[5]}件")
    lines(s, 381, "プラン①｜総合運用で、友だちの入口を増やし、検討の順番に沿った配信で資料請求・来場予約を増やす")
    lines(s, 343, "¥265,000〜")
    lines(s, 344, "¥260,000〜")
    init = [("離脱防止", "導入"), ("サンクス\nLINE誘導", "導入"), ("サイトの\nLINEボタン", "設置"),
            ("ステップ配信", "5通を構築"), ("配信の設計", "前後検索をもとに設計")]
    for (a, b), (k, v) in zip(((365, 366), (367, 368), (369, 370), (371, 372), (373, 374)), init):
        lines(s, a, k)
        lines(s, b, v)
    monthly = [("企画配信", "月4回程度"), ("来場前のお知らせ", "予約ごとに配信"), ("ステップ配信", "5通の改善"),
               ("月次レポート", "1回"), ("月次定例会", "1回"), ("友だちの入口", "運用")]
    for (a, b), (k, v) in zip(((349, 350), (351, 352), (353, 354), (355, 356), (357, 358), (359, 360)), monthly):
        lines(s, a, k)
        lines(s, b, v)
    for (k_id, v_id), (k, v) in zip(((361, 375), (362, 376), (364, 378)),
                                    (("友だち数", f"{total[5]:,}人"), ("資料請求", f"{doc[5]}件/月"), ("来場予約", f"{visit[5]}件/月"))):
        lines(s, k_id, k)
        lines(s, v_id, v)
        color_runs(by_id(s, v_id), NAVY)                    # 良い数字は赤い文字にしない
    lines(s, 383, "月ごとの資料請求・来場予約（件）")
    for sid in range(384, 390):
        e = by_id(s, sid)._element
        e.getparent().remove(e)
    rows = [[head_cell("か月目", 8.5)] + [head_cell(str(m), 9.5) for m in range(1, 7)],
            [label_cell("資料請求", 9)] + [cell(str(v), 11, True, NAVY, align="ctr") for v in doc],
            [label_cell("来場予約", 9)] + [cell(str(v), 11, True, NAVY, align="ctr") for v in visit]]
    for r in rows:
        for c in r:
            c["ins"] = (0.02, 0.02, 0.02, 0.02)
    add_table(s, 7.45, 4.06, [0.82] + [0.38] * 6, [0.36, 0.4, 0.4], rows, name="月ごとの資料請求・来場予約")
    add_box(s, 7.45, 5.35, 3.1, 0.75,
            [runs(("※試算の前提：サイトのUU 月13,228／今の友だち417人。LINE公式アカウント利用料は別", dict(size=8, color=GRAY)))],
            anchor="t", line_spacing=115, ins=(0, 0, 0, 0), name="注記")
    return s


# ---------------------------------------------------------------- 料金プラン（あらうち18）
HEAD = ["項目", "プラン①\n総合運用", "プラン②\n総合運用＋友だち獲得", "プラン③\n初期導入"]


def p_plan(prs, src):
    s = copy_slide(prs, src)
    lines(s, 635, "料金プラン：3つのプランの費用")
    t = [sh for sh in s.shapes if sh.has_table][0].table
    data = [HEAD,
            ["初期費用", "265,000円〜", "265,000円〜", "265,000円〜"],
            ["月額コンサルティング費", "200,000円〜", "200,000円〜", "なし"],
            ["月額ツール・LINE関連費", "60,000円〜", "60,000円〜\n＋LINE広告\n200,000円〜", "60,000円〜"],
            ["月額費用", "260,000円〜", "460,000円〜", "60,000円〜"]]
    for r, vals in zip(t.rows, data):
        for c, v in zip(r.cells, vals):
            set_cell(c, v)
    e = by_id(s, 640)._element
    e.getparent().remove(e)
    for k, txt in enumerate(["※初期費用＝運用設計・構築 100,000円〜＋離脱防止 15,000円＋サンクスLINE誘導 150,000円",
                             "※月額ツール・LINE関連費＝離脱防止 30,000円＋サンクスLINE誘導 30,000円〜。LINE公式アカウント利用料は別（貴社のご契約のまま）",
                             "※契約期間は6か月〜。試算のページはプラン①の場合",
                             "※施策内容・配信本数・連携範囲等により、費用は個別にお見積もりいたします。"]):
        add_box(s, 0.42, 5.75 + k * 0.27, 10.0, 0.27, [runs((txt, dict(size=9, color=GRAY)))], anchor="t",
                ins=(0, 0, 0, 0), name="注記")
    return s


# ---------------------------------------------------------------- 運用内容（あらうち19）
def p_naiyo(prs, src):
    s = copy_slide(prs, src)
    t = [sh for sh in s.shapes if sh.has_table][0].table
    data = [HEAD,
            ["離脱防止・\nサンクスLINE誘導", "導入・運用", "導入・運用", "導入・運用"],
            ["サイトのLINEボタン", "設置", "設置", "設置"],
            ["ステップ配信", "14日間・5通\n（構築・改善）", "14日間・5通\n（構築・改善）", "14日間・5通\n（構築のみ）"],
            ["企画配信", "月4回程度", "月4回程度", "なし（貴社で運用）"],
            ["来場前のお知らせ", "設定・改善", "設定・改善", "設定のみ"],
            ["結果の分析・改善\n（月次レポート・定例会）", "対応", "対応", "なし"],
            ["LINE広告\n（友だち追加）", "なし", "運用", "なし"]]
    for r, vals in zip(t.rows, data):
        for c, v in zip(r.cells, vals):
            set_cell(c, v)
    return s


def rebuild(prs, old_haishin, old_sim, old_hiyo):
    """配信・試算・費用のページを差し替える。新しいページは古いページの位置に入れる"""
    ar = Presentation(str(ARAUCHI))
    asl = list(ar.slides)
    assert by_id(asl[14], 385).text_frame.text.startswith("配信")
    assert by_id(asl[16], 339).text_frame.text.startswith("目標と費用対効果")
    assert by_id(asl[17], 635).text_frame.text.startswith("料金プラン")
    assert by_id(asl[18], 2).text_frame.text.startswith("運用内容")
    new = {"haishin": [p_haishin(prs, asl[14])], "sim": [p_sim(prs, asl[16])],
           "hiyo": [p_plan(prs, asl[17]), p_naiyo(prs, asl[18])]}
    for group in new.values():
        for s in group:
            for sh in s.shapes:
                if sh.has_text_frame:
                    meiryo(sh)
    lst = prs.slides._sldIdLst
    by_part = {prs.part.related_part(e.rId): e for e in lst}
    for old, key in ((old_haishin, "haishin"), (old_sim, "sim"), (old_hiyo, "hiyo")):
        old_el = by_part[old.part]
        pos = list(lst).index(old_el)
        for k, s in enumerate(new[key]):
            el = by_part[s.part]
            lst.remove(el)
            lst.insert(pos + 1 + k, el)
        lst.remove(old_el)
        prs.part.drop_rel(old_el.rId)
