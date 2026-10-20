# -*- coding: utf-8 -*-
"""すなお ver5.4 → ver6.0：営業提案資料としての組み直し（2026-10-20 殿村さんの修正指示）。

    python -m _build.patch_sunao_v60 <ver5.4.pptx> <ver6.0.pptx>

  中心メッセージ：広告で獲得した見込み客をLINE公式アカウントに蓄積し、継続的な情報配信で、次回セミナー参加・サロン入会を増やす。
  章立て：①現状把握・課題 ②市場・ユーザー・競合 ③なぜLINE公式アカウントなのか ④運用方針・全体導線 ⑤期待成果・費用
          ⑥施策の詳細 ⑦運用体制・スケジュール ／ 裏表紙 ／ ⑧補足資料・その他サービス（裏表紙の後ろ）
  ・現状を断定する言い方（手段が無い・接点が切れる 等）を「強化する」言い方に直す
  ・新規7枚：ターゲット／全体導線／費用の比較（サンクスLINE誘導あり・なし）／施策の詳細の扉／LINE施策の全体像／ステップ配信（5通）／確認事項
  ・外す7枚：トラック分岐の重複2枚／「--」のDay0〜15 2枚／休眠リード2枚／中身の無い「事例紹介」1枚
  ・前後検索の3枚はデータ取得まで非表示（中身は残す）
  費用・数値・出典は変えない（サンクスLINE誘導なしの金額は、殿村さん指定の単純計算）。変更一覧は `_drafts/すなお_ver6.0_変更一覧.md`。
"""
import sys

from pptx import Presentation
from pptx.util import Inches, Pt

from _build.arauchi_v4.lib import add_box, add_table
from _build.arauchi_v10.build import runs
from _build.build_sunao_v1 import all_meiryo
from _build.build_sunao_v4 import set_cell
from _build.patch_sunao_v41 import set_text
from _build.patch_sunao_v53_kaizen import (BAND, GRAY, NAVY, TAKE, TEXT, cell, clear, head, label, lead_band, lines, note,
                                           set_one, shiji, title_of)
from _build.patch_tanaka_v5 import copy_slide

LINE_G = "06C755"
LIGHT = "F5F7FA"


def by_id(s, sid):
    return next(sh for sh in s.shapes if sh.shape_id == sid)


def title(s):
    t = title_of(s)
    return t.text_frame.text.strip() if t is not None else ""


def run_replace(sh, old, new):
    """図形の中の、old を含む文字列（a:r）だけを書き換える（太字の見出しなど、同じ段落のほかの書式は残す）"""
    for p in sh.text_frame.paragraphs:
        for r in p.runs:
            if old in r.text:
                r.text = r.text.replace(old, new)
                return
    raise AssertionError((sh.shape_id, old))


def drop(s, *ids):
    for sid in ids:
        e = by_id(s, sid)._element
        e.getparent().remove(e)


def new_page(prs, base, src3, ttl, lead=None):
    """業界STUDYのページを写して中身を空にし、タイトル（と想定動線のページのリード帯）を入れる"""
    s = copy_slide(prs, base)
    t = clear(s)
    set_one(t, ttl)
    if lead:
        lead_band(s, src3, lead)
    return s


def divider(prs, base, text):
    s = copy_slide(prs, base)
    set_divider(s, text)
    return s


def set_divider(s, text):
    sh = by_id(s, 2)
    set_one(sh, text)
    sh.left, sh.width = Inches(0.92), Inches(9.0)


def table(s, x, y, widths, heights, rows, name):
    return add_table(s, x, y, widths, heights, rows, name=name)


# ---------------------------------------------------------------- ① 現状把握・課題
def p_doshen(s):
    """想定動線（殿村さんの版）：ターゲットに一般の方向けを足し、「接点が切れやすい」を言い換える"""
    box = by_id(s, 10)
    run_replace(box, "：保険営業の方の営業力向上を目的としたセミナー・サロン",
                "：保険営業の方向けのセミナー・サロンが主軸（一般の方向けのマネーセミナーもあり）")
    run_replace(box, "：イベントとイベントの間の時期は、接点が切れやすい",
                "：イベントの間の時期にも、継続して情報を届けることが重要")


def p_genjo(s):
    assert set_text(by_id(s, 43), "申し込まずに離れた方を、次のイベントまで案内する手段が無い。",
                    "次のイベントまで、接点を持ち続ける仕組みを強化できる。")
    assert set_text(by_id(s, 53), "参加して入会していない方へ、", "LINE公式アカウントを、")
    assert set_text(by_id(s, 53), "入会まで案内を続ける手段が無い状況。", "見込み客との継続的な接点として活用できる。")
    set_one(by_id(s, 76), "イベント開催")
    set_one(by_id(s, 116), "年4回")
    drop(s, 92, 177, 169, 191)                                       # 中身の無い「※」の2行


# ---------------------------------------------------------------- ② 市場・ユーザー・競合
def p_study(s):
    assert set_text(by_id(s, 52), "イベントのたびに集め直すのではなく、一度つながった方をLINEで次のイベントまで追い続ける。",
                    "広告で集めた方と、LINEで次のイベントまでつながり続ける。")


def p_target(s):
    rows = [[head("項目"), head("保険営業の方（主軸）"), head("一般の方")]]
    data = [
        ("主な関心", "営業力の向上（面談・提案の進め方）", "資産形成・お金の知識"),
        ("主な申込先", "生産性4倍セミナー → オンラインサロン「ハッピーライフ」", "資産形成・マネーセミナー等"),
        ("LINEで届ける内容", "セミナーの見どころ・受講生の声・サロンのご案内", "マネーセミナー等のご案内"),
        ("出し分けの基本", "保険営業の方向けのLPから友だち追加", "一般の方向けの申込みページ等から友だち追加"),
    ]
    for i, (k, a, b) in enumerate(data):
        fill = BAND[i % 2]
        rows.append([label(k, fill), cell(a, fill=fill), cell(b, fill=fill)])
    table(s, 0.46, 1.8, [2.1, 4.1, 3.7], [0.4] + [0.62] * 4, rows, "ターゲット")
    shiji(s, 4.95, 1.2, lines("同じLINE公式アカウントで運用する場合は、友だち追加の流入経路（どのページから追加したか）で出し分ける。",
                               "一般の方向けセミナーの参加者には、保険営業の方向けサロンへの入会を前提にした案内はしない。",
                               size=13, color=TEXT))
    note(s, 6.45, ["※流入経路で分けられない方には、どちらにも当てはまる内容（開催のお知らせなど）を基本に配信"])


# ---------------------------------------------------------------- ③ なぜLINE公式アカウントなのか
def p_why(s):
    set_one(by_id(s, 4), "広告で新しい接点をつくり、LINEで継続して情報を届ける")
    drop(s, 5, 6)
    rows = [[cell("", fill="FFFFFF", borders={k: (0.75, "FFFFFF") for k in "lrtb"}), head("広告"), head("LINE公式アカウント")]]
    data = [
        ("役割", "新しい見込み客との接点を獲得する", "一度つながった見込み客に、継続して情報を届ける"),
        ("主に届く相手", "まだつながっていない方", "友だち追加した方"),
        ("届けられる時期", "広告を配信している期間", "イベント期も、イベントのない期間も"),
        ("積み上がるもの", "セミナー予約などの反響", "友だち（次回以降も案内できる相手）"),
    ]
    for i, (k, a, b) in enumerate(data):
        fill = BAND[i % 2]
        rows.append([label(k, fill), cell(a, fill=fill, align="ctr"), cell(b, fill=fill, align="ctr", bold=True, color=NAVY)])
    table(s, 0.55, 1.65, [2.0, 3.85, 3.88], [0.4] + [0.5] * 4, rows, "広告とLINEの役割")
    add_box(s, 0.55, 4.25, 3.4, 0.34, [runs(("LINE公式アカウントのメリット", dict(size=13, bold=True, color="FFFFFF")))],
            fill=NAVY, align="ctr", anchor="ctr", name="メリット 見出し")
    add_box(s, 0.55, 4.59, 9.73, 1.9,
            lines("・友だち追加された方に、次回以降のイベントを案内できる",
                  "・イベントのない期間も、継続して接点を持てる",
                  "・興味や参加状況に応じた情報配信を設計できる",
                  "・新規集客だけに頼らない導線を構築できる",
                  "・中長期的に、広告への依存度を下げられる可能性がある", size=14, color=TEXT),
            fill=TAKE, line=NAVY, anchor="ctr", ins=(0.3, 0.06, 0.2, 0.06), line_spacing=130, name="メリット")


# ---------------------------------------------------------------- ④ 運用方針・全体導線
def p_yoken(s):
    drop(s, *[sh.shape_id for sh in s.shapes if sh.has_table])
    fill = BAND[1]
    rows = [
        [label("プロモーション目的", "FFFFFF"),
         cell("広告で獲得した見込み客をLINE公式アカウントに蓄積し、継続的な情報配信で、次回セミナーへの参加とオンラインサロンへの入会を増やす",
              size=14, bold=True, color=NAVY)],
        [label("中長期の目的", fill), cell("入会後の継続利用・紹介にもLINEを活用し、LTVの向上をめざす", size=14, fill=fill)],
        [label("配信目標", "FFFFFF"),
         cell(lines("友だち獲得を増やし、ブロック率を抑えながら、配信可能な人数を積み上げる",
                    "（6か月の想定は、概算お見積りに記載）", size=14))],
        [label("手法（ご提案）", fill),
         cell(lines("① 離脱防止ポップアップと予約完了後の友だち追加誘導で、友だちを集める",
                    "② 友だち追加直後・予約した方・参加後の段階ごとに、案内を変える",
                    "③ 平常期は月2回、セミナー開催月は月4〜5回の企画配信で、つながりを保つ", size=14), fill=fill)],
        [label("＋αの施策", "FFFFFF"),
         cell(lines("過去のセミナー予約者・参加者へのご案内（通知メッセージ）",
                    "サロンの申込み完了ページからのLINE誘導（サンクスLINE誘導）", size=14)
              + lines("※通知メッセージは、予約時の電話番号がある場合", size=11, color=GRAY))],
    ]
    for r in rows:
        r[0]["size"] = 14
        r[0]["align"] = "ctr"
    heights = [0.85, 0.55, 0.8, 1.15, 1.05]
    top = 0.7 + (7.0 - 0.7 - sum(heights)) / 3
    table(s, 0.46, top, [2.2, 7.7], heights, rows, "要件定義")


def p_hoshin(s):
    set_one(title_of(s), "運用方針")
    set_one(by_id(s, 49), "3つのポイントで、見込み客との接点を継続する仕組みを強化する")
    set_one(by_id(s, 51), "イベント開催時だけでなく、平常期にも継続して情報を届ける")
    drop(s, 52)
    rows = [[head("段階"), head("強化するポイント"), head("打ち手"), head("伸ばす数字")]]
    data = [
        ("01 集める", "LPを見た方・予約した方と、LINEでつながる",
         ("・離脱防止ポップアップで友だちに", "・予約完了後に、友だち追加をご案内"), "友だち数"),
        ("02 届け続ける", "平常期・開催月それぞれの時期に合わせて届ける",
         ("・平常期 月2回／開催月 月4〜5回の企画配信", "・予約した方へ、前日・当日のリマインド"), "予約・参加数"),
        ("03 つなげる", "参加後、関心が高いうちにサロンを案内する",
         ("・参加後に、振り返りとサロンの内容をご案内", "・受講生の声で、入会の判断材料を届ける"), "入会数"),
    ]
    for i, (k, now, acts, kpi) in enumerate(data):
        fill = BAND[i % 2]
        rows.append([label(k, fill), cell(now, fill=fill), cell(lines(*acts, size=13), fill=fill),
                     cell(kpi, fill=fill, bold=True, color=NAVY, align="ctr")])
    table(s, 0.46, 3.2, [1.7, 2.7, 4.2, 1.3], [0.4] + [0.9] * 3, rows, "3つのポイント")


def p_flow(s):
    steps = ["広告", "セミナーLP", "LINE友だち追加", "継続的な情報配信", "セミナー予約・参加", "オンラインサロン入会"]
    desc = ["Meta広告などで、新しい見込み客と接点をつくる",
            "セミナーの内容を知り、予約する",
            "離脱防止ポップアップ／予約完了後の友だち追加のご案内",
            ["あいさつ・見どころ動画・参加特典", "企画配信：平常期 月2回、開催月 月4〜5回"],
            "前日・当日のリマインド／開催情報・参加方法のご案内",
            "参加後の振り返り・サロンの内容紹介・入会のご案内"]
    w, gap, x0, y = 1.45, 0.24, 0.46, 2.25
    add_box(s, x0, 1.8, 2 * w + gap, 0.34, [runs(("広告（新規集客）", dict(size=13, bold=True, color="FFFFFF")))],
            fill=NAVY, align="ctr", anchor="ctr", name="広告の領域")
    add_box(s, x0 + 2 * (w + gap), 1.8, 4 * w + 3 * gap, 0.34,
            [runs(("LINE公式アカウント（継続的な顧客育成）", dict(size=13, bold=True, color="FFFFFF")))],
            fill=LINE_G, align="ctr", anchor="ctr", name="LINEの領域")
    for k, (st, d) in enumerate(zip(steps, desc)):
        x = x0 + k * (w + gap)
        ad = k < 2
        add_box(s, x, y, w, 0.75, [runs((st, dict(size=13, bold=True, color="FFFFFF" if ad else NAVY)))],
                fill=NAVY if ad else "FFFFFF", line=None if ad else LINE_G, line_w=1.5, align="ctr", anchor="ctr",
                ins=(0.05, 0.03, 0.05, 0.03), name=f"導線 {st}")
        add_box(s, x, y + 0.85, w, 1.55, lines(*(d if isinstance(d, list) else [d]), size=12, color=TEXT), fill=LIGHT, anchor="t",
                ins=(0.08, 0.08, 0.08, 0.06), line_spacing=115, name=f"導線の説明 {st}")
        if k < 5:
            add_box(s, x + w, y, gap, 0.75, [runs(("▶", dict(size=12, color=GRAY)))], align="ctr", anchor="ctr",
                    ins=(0, 0, 0, 0), name="矢印")
    xl = x0 + 2 * (w + gap)
    add_box(s, xl, 4.85, 4 * w + 3 * gap, 0.5,
            [runs(("中長期：入会後の継続利用・紹介にも、LINEでの情報配信を活用", dict(size=13, bold=True, color=NAVY)))],
            line=GRAY, dash="dash", align="ctr", anchor="ctr", name="中長期")
    note(s, 6.45, ["※一般の方向けセミナーから友だち追加した方は、流入経路で分け、一般の方向けの案内を中心に配信"])


def p_phase_table(s):
    set_one(title_of(s), "想定動線（フェーズごとのLINEの役割）")
    t = next(sh for sh in s.shapes if sh.has_table).table
    set_cell(t.cell(0, 2), "離脱・停滞が起きやすい点")
    rows = {t.cell(r, 0).text.strip()[:1]: r for r in range(1, len(t.rows))}
    set_cell(t.cell(rows["②"], 2), "次の開催まで関心が続きにくい")
    set_cell(t.cell(rows["④"], 3), "予約後の友だち追加・前日の案内")
    set_cell(t.cell(rows["⑥"], 2), "次のイベントまで期間が空く")
    set_cell(t.cell(rows["⑧"], 2), "継続利用・紹介まで届きにくい")
    set_cell(t.cell(rows["⑧"], 3), "イベント案内・紹介（中長期）")


def p_track(s):
    set_one(title_of(s), "配信の出し分け：予約した方／予約していない方")
    assert set_text(by_id(s, 11), "次のイベントの案内を、広告より先に届けられる", "次のイベントの案内を、開催が決まり次第届けられる")
    box = by_id(s, 13)
    assert set_text(box, "友だちの入口（ポップアップツール／サンクスLINE誘導）で自動で分け、予約した方は当日の参加へ、",
                    "友だち追加の流入経路（離脱防止ポップアップ／予約完了後のご案内）で分け、予約した方は当日の参加へ、")
    assert set_text(box, "予約していない方は平常期の配信で次のイベントへつなげる。",
                    "予約していない方は平常期の配信で次のイベントへつなげる。予約・参加・入会の状況での出し分けは、連携の可否を確認のうえ設計。")
    box.height = Inches(1.15)


# ---------------------------------------------------------------- ⑤ 期待成果・費用
def p_mitsumori(s):
    set_one(by_id(s, 339), "概算お見積り：6か月目の配信可能な人数90人・セミナー予約11件（想定）")
    set_one(by_id(s, 346), "6か月後（想定）")
    set_one(by_id(s, 362), "配信可能な人数")
    ps = by_id(s, 351).text_frame.paragraphs
    ps[0].runs[0].text, ps[1].runs[0].text = "前日・当日の", "リマインド"
    add_box(s, 7.38, 4.3, 3.17, 1.6,
            lines("※SIMによる想定値（実績ではない）", "※配信可能な人数＝ブロックされずに届く友だちの数",
                  "※セミナー予約・サロン入会は、LINEの配信・リンクからの件数として想定（広告からの直接の予約は含まない）",
                  size=10, color=GRAY),
            anchor="t", ins=(0.05, 0, 0, 0), line_spacing=110, name="想定の注記")


def p_hikaku(s):
    for sh in list(s.shapes):
        if sh.shape_id not in (635, 636, 637):
            sh._element.getparent().remove(sh._element)
    set_one(by_id(s, 635), "費用の比較：サンクスLINE誘導あり（基本案）・なし（参考）")
    set_one(by_id(s, 636), "費用の比較")
    set_one(by_id(s, 637), "基本案はプラン②。ツールを調整する場合の参考も記載")
    big = dict(size=15, bold=True, color=NAVY, align="ctr")
    rows = [[head("項目"), head("基本案：サンクスLINE誘導あり"), head("参考：サンクスLINE誘導なし")]]
    data = [
        ("初期費用", cell("295,000円〜", **big), cell("145,000円〜", **big)),
        ("月額費用（平常期）", cell("165,000円", **big), cell("135,000円〜", **big)),
        ("月額費用（セミナー開催月）", cell("365,000円程度", **big), cell("335,000円程度", **big)),
        ("予約完了後の友だち追加",
         cell("サンクスLINE誘導ツールで、LINEへ自動遷移（フォームの入力内容の引き継ぎも可能）", size=12),
         cell("予約完了ページに、通常の友だち追加ボタン等を設置（自動遷移・入力内容の引き継ぎはなし）", size=12)),
        ("成果の想定",
         cell("概算お見積りのとおり（SIMによる想定値）", size=12),
         cell("予約した方の友だち追加率が下がる可能性があるため、想定は別途見直し", size=12)),
    ]
    for i, (k, a, b) in enumerate(data):
        fill = BAND[i % 2]
        a["fill"] = b["fill"] = fill
        lb = label(k, fill)
        lb["align"] = "ctr"
        rows.append([lb, a, b])
    table(s, 0.42, 1.48, [2.7, 3.65, 3.65], [0.42, 0.5, 0.5, 0.5, 0.8, 0.7], rows, "費用の比較")
    note(s, 5.2, ["※月額費用＝コンサルティング費（平常期 100,000円／セミナー開催月 300,000円程度）＋ツール・LINE関連費 65,000円〜（参考は 35,000円〜）",
                  "※参考の金額は、基本案からサンクスLINE誘導（初期 150,000円・月額 30,000円）を除いた単純計算",
                  "※通常の友だち追加ボタン等の設置にかかる費用・実装条件は、予約フォームの仕様を確認のうえお見積り",
                  "※広告費は含まない。施策内容・配信本数・連携範囲等により、費用は個別にお見積もりいたします。"])


def p_unyo(s):
    t = next(sh for sh in s.shapes if sh.has_table).table
    r = next(i for i in range(len(t.rows)) if t.cell(i, 0).text.startswith("セミナー前日のお知らせ"))
    set_cell(t.cell(r, 0), "セミナー前日・当日のリマインド")


def p_seika(s):
    set_one(by_id(s, 5), "別案：成果件数に応じた費用で、リスクを抑えて運用を始められる")
    note(s, 6.05, ["※成果地点の定義・重複計上の有無・計測方法は、ご契約前に確認のうえ決定"])


# ---------------------------------------------------------------- ⑥ 施策の詳細
def p_shisaku(s):
    rows = [[head("段階"), head("主な施策"), head("ねらい")]]
    data = [
        ("① 友だち獲得", "離脱防止ポップアップ／セミナー予約完了後の友だち追加誘導", "LPを離れる方・予約した方とつながる"),
        ("② 友だち追加直後", "あいさつメッセージ・見どころ動画・参加特典（必要に応じて、簡易的な興味分野の把握）",
         "セミナーの中身を伝え、予約につなげる"),
        ("③ 予約した方", "前日・当日のリマインド／開催情報・参加方法のご案内", "当日の参加を後押しする"),
        ("④ 参加後", "セミナー内容の振り返り／オンラインサロンの内容紹介・入会のご案内", "関心が高いうちに、入会を案内する"),
        ("⑤ 平常期", "月2回の企画配信／次回イベントのご案内", "イベントのない期間もつながりを保つ"),
        ("⑥ セミナー開催月", "月4〜5回の企画配信", "開催に合わせて、予約・参加を後押しする"),
    ]
    for i, (k, a, b) in enumerate(data):
        fill = BAND[i % 2]
        rows.append([label(k, fill), cell(a, fill=fill, size=12), cell(b, fill=fill, size=12)])
    table(s, 0.46, 1.75, [2.0, 4.6, 3.3], [0.4] + [0.62] * 6, rows, "LINE施策の全体像")
    note(s, 6.0, ["※③のリマインドは、企画配信の回数とは別に、予約ごとに配信",
                  "※③④は、予約・参加した方を特定できる範囲で実施（予約完了後の友だち追加など）。予約システム等との連携は、可否を確認のうえ設計"])


def p_step(s):
    days = ["追加直後", "1〜2日後", "4〜5日後", "8〜9日後", "12〜15日後"]
    heads = ["あいさつ・参加特典", "見どころ動画", "受講生の声", "よくある質問", "次回セミナーのご案内"]
    bodies = ["参加特典をご案内。必要に応じて、興味分野を1問で確認（保険営業の方向け／一般の方向け）",
              "両想い営業・ライフプランブロックなど、セミナーの見どころを動画で紹介",
              "面談の進め方が変わった受講生の声を紹介",
              "地方からも参加できる・録画はない・初心者でも分かる・参加は無料",
              "開催日程と予約ページをご案内（募集期間外は、案内開始時にお届け）"]
    w, gap, x0, y = 1.86, 0.15, 0.46, 1.95
    for k, (d, h, b) in enumerate(zip(days, heads, bodies)):
        x = x0 + k * (w + gap)
        add_box(s, x, y, w, 0.42, [runs((d, dict(size=13, bold=True, color="FFFFFF")))], fill=NAVY, align="ctr",
                anchor="ctr", name=f"ステップ {d}")
        add_box(s, x, y + 0.5, w, 0.5, [runs((h, dict(size=13, bold=True, color=NAVY)))], fill=TAKE, align="ctr",
                anchor="ctr", ins=(0.04, 0.02, 0.04, 0.02), name=f"配信 {d}")
        add_box(s, x, y + 1.08, w, 1.75, [runs((b, dict(size=12, color=TEXT)))], fill=LIGHT, anchor="t",
                ins=(0.08, 0.08, 0.08, 0.06), line_spacing=115, name=f"内容 {d}")
    note(s, 5.05, ["※配信内容は、貴社LPの「生産性4倍セミナーの特徴」「受講生の声」「よくある質問」をもとにした例",
                   "※一般の方向けを選んだ方には、一般の方向けの案内を中心に配信",
                   "※前後検索のデータを取得後、内容・順番を見直す"])


def p_follow(s):
    set_one(by_id(s, 3), "参加後のフォローアップ：関心が高いうちに、入会の判断材料を届ける")
    set_one(by_id(s, 5), "参加直後と参加後30〜60日に、サロンの内容と入会方法をお届けする")
    set_one(by_id(s, 4), "配信内容（例）")


# ---------------------------------------------------------------- ⑦ 運用体制・スケジュール
def p_kakunin(s):
    rows = [[head("項目"), head("確認させていただきたいこと"), head("確認後の対応")]]
    data = [
        ("予約・顧客情報との連携", "予約システム・フォーム・顧客管理の情報と、LINEの友だちを紐づけられるか",
         "可能な場合：予約・参加・入会の状況に応じた配信を設計／難しい場合：流入経路などで分けた配信・全体配信を基本にする"),
        ("予約完了後の友だち追加", "サンクスLINE誘導ツールを使うか、通常の友だち追加ボタン等にするか",
         "ボタン等の設置費用・実装条件は、予約フォームの仕様を確認のうえお見積り"),
        ("成果の数え方", "セミナー予約・サロン入会を、LINEからの件数とするか、広告からの予約も含めるか",
         "概算お見積りの想定（LINEの配信・リンクからの件数）を、定義に合わせて見直す"),
        ("成果報酬型プラン（別案）", "成果地点の定義・重複計上の有無・計測方法",
         "確認のうえ、成果報酬単価（12,000円／件）での条件を決定"),
        ("一般の方向けセミナー", "同じLINE公式アカウントで運用するか、流入経路を分けられるか", "出し分けの方法を決定"),
    ]
    for i, (k, a, b) in enumerate(data):
        fill = BAND[i % 2]
        lb = label(k, fill)
        lb["size"] = 12
        rows.append([lb, cell(a, fill=fill, size=12), cell(b, fill=fill, size=12)])
    table(s, 0.46, 1.75, [2.55, 3.7, 3.65], [0.4, 0.95, 0.7, 0.8, 0.6, 0.6], rows, "確認事項")
    note(s, 6.0, ["※データ連携が必要な高度な施策は、別途要件定義・お見積りのうえ実施"])


# ---------------------------------------------------------------- 組み立て
def main(src, out):
    prs = Presentation(src)
    S = list(prs.slides)                                             # ver5.4 の並び（1始まりで S[n-1]）

    def o(n, expect):
        s = S[n - 1]
        assert expect in title(s) or expect in " ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame), (n, expect, title(s))
        return s

    cover, d_genjo, doshen, phase_t, genjo = o(1, "御中"), o(2, "現状ステータスの整理"), o(3, "想定動線"), o(4, "想定動線"), o(5, "現状把握")
    zengo = [o(6, "前後検索"), o(7, "前後検索"), o(8, "前後検索")]
    study, axis, d_market = o(9, "業界STUDY"), o(10, "訴求軸"), o(11, "施策の詳細（配信の効率化）")
    yoken, hoshin, mitsumori = o(12, "要件定義"), o(13, "本質課題"), o(14, "概算お見積り")
    kangae, kangae15, zentai, track = o(15, "重要な考え方"), o(16, "重要な考え方"), o(17, "全体設計"), o(18, "想定の費用対効果")
    dels = [o(19, "想定の費用対効果"), o(20, "Day0〜15"), o(21, "Day0〜15"), o(22, "想定の費用対効果"),
            o(23, "休眠リード"), o(24, "休眠リード"), o(60, "事例紹介")]
    follow, d_hoshin, teikyo = o(25, "長期ニーズトラック"), o(26, "施策の詳細（新規友だち獲得）"), o(27, "提供施策について")
    popup1, popup2, thanks, tsuchi1, tsuchi2 = o(28, "ポップアップ"), o(29, "ポップアップ"), o(30, "サンクスLINE誘導"), o(31, "通知メッセージ"), o(32, "通知メッセージ")
    d_hiyo, seika, service, ryokin, unyo, tsudo = o(33, "費用感"), o(34, "成果報酬型プラン"), o(35, "サービス内容"), o(36, "料金プラン"), o(37, "運用内容"), o(38, "都度発注")
    d_taisei, report, keisoku, taisei, schedule = o(39, "サポート体制"), o(40, "報告FMT"), o(41, "効果計測"), o(42, "体制図"), o(43, "運用スケジュール")
    d_why, why = o(44, "LINE公式アカウント運用のメリット"), o(45, "なぜLINEか")
    haikei = [o(n, "検討する背景") for n in range(46, 52)]
    d_hosoku = o(52, "中長期施策")
    sonota = [o(n, "提供施策") if n in (53, 58, 59) else o(n, "代表的な活用機能") for n in range(53, 60)]
    jirei, back = o(61, "保険・イベント集客"), o(62, "株式会社DYM")

    # 新しいページ（業界STUDYの器・想定動線のリード帯を使う）
    target = new_page(prs, study, doshen, "ターゲット：保険営業の方が主軸。一般の方向けも考慮", "対象によって、関心のある内容と申込先が異なる")
    flow = new_page(prs, study, doshen, "全体導線：広告で集め、LINEで育て、セミナー参加・サロン入会へ",
                    "広告による新規集客と、LINEによる継続的な顧客育成を組み合わせる")
    shisaku = new_page(prs, study, doshen, "LINE施策の全体像：6つの段階で、友だち獲得から入会まで案内",
                       "友だち追加の時期と状況に合わせて、届ける内容を変える")
    step = new_page(prs, study, doshen, "友だち追加直後のステップ配信（5通・案）", "追加直後の関心が高いうちに、セミナーの中身と参加の判断材料を届ける")
    kakunin = new_page(prs, study, doshen, "導入前に確認させていただきたい事項", "連携の可否と成果の数え方によって、配信の設計と費用が変わるため、事前に確認する")
    hikaku = copy_slide(prs, ryokin)
    d_shisaku = divider(prs, d_genjo, "施策の詳細")

    # 中身の直し
    p_doshen(doshen); p_genjo(genjo); p_study(study); p_target(target); p_why(why)
    p_yoken(yoken); p_hoshin(hoshin); p_flow(flow); p_phase_table(phase_t); p_track(track)
    p_mitsumori(mitsumori); p_hikaku(hikaku); p_unyo(unyo); p_seika(seika)
    p_shisaku(shisaku); p_step(step); p_follow(follow); p_kakunin(kakunin)
    for s, t in ((d_genjo, "現状把握・課題"), (d_market, "市場・ユーザー・競合"), (d_why, "なぜLINE公式アカウントなのか"),
                 (d_hoshin, "運用方針・全体導線"), (d_hiyo, "期待成果・費用"), (d_taisei, "運用体制・スケジュール"),
                 (d_hosoku, "補足資料・その他サービス")):
        set_divider(s, t)
    p_service_fix(service)
    for s in zengo:                                                  # 前後検索はデータ取得まで非表示
        s._element.set("show", "0")

    order = ([cover, d_genjo, doshen, genjo]
             + [d_market, study, target, axis] + zengo
             + [d_why, why, jirei]
             + [d_hoshin, yoken, hoshin, flow, phase_t, zentai, track]
             + [d_hiyo, mitsumori, hikaku, ryokin, unyo, seika]
             + [d_shisaku, shisaku, popup1, popup2, thanks, kangae15, step, follow]
             + [d_taisei, report, keisoku, taisei, kakunin, schedule]
             + [back]
             + [d_hosoku, kangae] + haikei + [teikyo, tsuchi1, tsuchi2, service, tsudo] + sonota)
    assert len(order) == len(set(id(x) for x in order))
    lst = prs.slides._sldIdLst
    el_of = {prs.part.related_part(e.rId): e for e in lst}
    keep = [el_of[s.part] for s in order]
    for e in list(lst):
        if e not in keep:
            assert prs.part.related_part(e.rId) in {d.part for d in dels}, "外す予定でないページ"
            prs.part.drop_rel(e.rId)
        lst.remove(e)
    for e in keep:
        lst.append(e)
    all_meiryo([target, flow, shisaku, step, kakunin, hikaku, d_shisaku, why, yoken, hoshin, mitsumori, seika, track])
    prs.save(out)
    print("saved:", out, len(prs.slides), "枚")


def p_service_fix(s):
    """サービス内容（DYMプラン・補足資料へ）：03と05が同じ文だったのと、提案内容とずれていたところを直す"""
    for old, new in (("LINEステップ配信シナリオ（4-8本）", "LINEステップ配信シナリオ（5通）"),
                     ("友だち追加直後〜成約誘導まで", "友だち追加直後〜セミナー予約まで"),
                     ("季節・時期ごとの企画配信", "セミナー開催月・平常期の企画配信"),
                     ("繁忙期に連動・各ユーザー層をターゲットとしたキャンペーン案", "開催月は月4〜5回、平常期は月2回の配信案")):
        sh = next(x for x in s.shapes if x.has_text_frame and x.text_frame.text.strip() == old)
        set_one(sh, new)
    hits = [x for x in s.shapes if x.has_text_frame and x.text_frame.text.strip() == "アンケート設計（直接関与 / 診断型）"]
    subs = [x for x in s.shapes if x.has_text_frame and x.text_frame.text.strip() == "リード有効化＋長期追客セグメント戦略"]
    assert len(hits) == 2 and len(subs) == 2
    hits.sort(key=lambda x: x.top)
    subs.sort(key=lambda x: x.top)
    set_one(hits[0], "興味分野の把握（簡易アンケート）")
    set_one(subs[0], "流入経路・興味分野に応じた配信の出し分け")
    set_one(hits[1], "効果計測・改善")
    set_one(subs[1], "配信ごとの開封・クリック・予約を確認し、次の配信に反映")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
