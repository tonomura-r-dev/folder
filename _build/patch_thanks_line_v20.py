# -*- coding: utf-8 -*-
"""サンクスLINEのご提案 ver1.4 → ver2.0（2026-10-01）
元の5枚のうち S1・S2 ＋ 実績（新規）＋ 費用（元S5）の4枚に組み直す。
実績ページ＝弊社の運用実績（美容クリニック）＋ LINEヤフー公式の導入事例2件。
  python3 _build/patch_thanks_line_v20.py <ver1.4.pptx>
数字の出典（2026-10-01 に元ページで確認）：
  https://www.lycbiz.com/jp/case-study/line-official-account/acne-clinic/
    「ステップ配信」でクリニックの予約数も約120％まで増加（ステップ配信の実施前と比較）
  https://www.lycbiz.com/jp/case-study/line-official-account/sugitama/
    １年で「LINEで予約」経由の予約数は約6倍増加（2021年12月と2022年12月を比較）
"""
import copy
import sys
from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.util import Cm, Pt

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "20261001_サンクスLINEのご提案ver2.0.pptx"
prs = Presentation(sys.argv[1])
NAVY, GREEN, INK, GRAY = RGBColor(0x1F, 0x28, 0x5A), RGBColor(0x06, 0xC7, 0x55), RGBColor(0x33, 0x33, 0x33), RGBColor(0x7F, 0x7F, 0x7F)

# ---- 実績ページ：元S3を土台に作り直す ----
s = prs.slides[2]
by = lambda n: next(x for x in s.shapes if x.name == n)
tpl_navy = copy.deepcopy(by("Rounded Rectangle 4")._element)   # 紺の帯
tpl_card = copy.deepcopy(by("Rounded Rectangle 6")._element)   # 薄い青のカード
tpl_green = copy.deepcopy(by("Rounded Rectangle 8")._element)  # 緑の帯
note_src = copy.deepcopy(prs.slides[3].shapes[-1]._element)     # 元S4の※注記（テキストボックス）
keep = {"TextBox 1", "TextBox 2", "Connector 3"}
for sh in list(s.shapes):
    if sh.name not in keep:
        sh._element.getparent().remove(sh._element)


def set_runs(tf_shape, lines):
    """lines = [(text, size, bold, color), ...] 1要素＝1段落。書式は元の最初のrunを引き継ぐ。"""
    tf = tf_shape.text_frame
    base_p = tf.paragraphs[0]._p
    base_r = base_p.find(qn("a:r"))
    rpr0 = copy.deepcopy(base_r.find(qn("a:rPr"))) if base_r is not None else None
    for p in list(tf._txBody.findall(qn("a:p")))[1:]:
        tf._txBody.remove(p)
    for r in list(base_p.findall(qn("a:r"))) + list(base_p.findall(qn("a:br"))):
        base_p.remove(r)
    empty_p = copy.deepcopy(base_p)
    for i, (text, size, bold, color) in enumerate(lines):
        p = base_p if i == 0 else copy.deepcopy(empty_p)
        if i:
            tf._txBody.append(p)
        r = p.makeelement(qn("a:r"), {})
        if rpr0 is not None:
            r.append(copy.deepcopy(rpr0))
        t = r.makeelement(qn("a:t"), {})
        t.text = text
        r.append(t)
        end = p.find(qn("a:endParaRPr"))
        (end.addprevious(r) if end is not None else p.append(r))
    for para, (text, size, bold, color) in zip(tf.paragraphs, lines):
        for run in para.runs:
            run.font.size, run.font.bold = Pt(size), bold
            run.font.color.rgb = color


def add(tpl, x, y, w, h, lines):
    el = copy.deepcopy(tpl)
    s.shapes._spTree.append(el)
    sh = s.shapes[-1]
    sh.left, sh.top, sh.width, sh.height = Cm(x), Cm(y), Cm(w), Cm(h)
    set_runs(sh, lines)
    return sh


# タイトル・リード
set_runs(by("TextBox 1"), [("実績｜LINEで追うと、予約・来店が伸びる", 16, True, RGBColor(0x00, 0x20, 0x60))])
set_runs(by("TextBox 2"), [("弊社の運用実績と、LINEヤフー公式の導入事例", 16, False, INK)])

CASES = [
    ("弊社実績", "美容クリニック",
     ["予約の完了画面から", "LINEへ誘導（サンクスLINE）"],
     "40〜50%改善", "予約後の来院率"),
    ("LINEヤフー公式事例", "皮膚科クリニック",
     ["友だち追加後に、", "ステップ配信で予約を案内"],
     "予約数 約120%", "ステップ配信の実施前と比べて"),
    ("LINEヤフー公式事例", "寿司居酒屋（全国73店舗）",
     ["「LINEで予約」と", "来店を促すメッセージ配信"],
     "予約数 約6倍", "「LINEで予約」経由・1年で"),
]
X0, W, G = 1.46, 7.93, 0.4
for i, (tag, who, what, big, sub) in enumerate(CASES):
    x = X0 + i * (W + G)
    add(tpl_card if i else tpl_green, x, 4.75, W, 0.85,
        [(tag, 13, True, NAVY if i else RGBColor(0xFF, 0xFF, 0xFF))])
    add(tpl_navy, x, 5.75, W, 1.15, [(who, 16, True, RGBColor(0xFF, 0xFF, 0xFF))])
    add(tpl_card, x, 7.05, W, 2.6, [(t, 14, False, INK) for t in what])
    add(tpl_green, x, 9.8, W, 2.9, [(sub, 12, False, RGBColor(0xFF, 0xFF, 0xFF)), (big, 22, True, RGBColor(0xFF, 0xFF, 0xFF))])

add(tpl_navy, X0, 13.3, 3 * W + 2 * G, 1.4,
    [("申込みの直後にLINEでつながり、配信で予約・来店へつなげる", 17, True, RGBColor(0xFF, 0xFF, 0xFF))])
s.shapes._spTree.append(note_src)
note = s.shapes[-1]
note.left, note.top, note.width, note.height = Cm(X0), Cm(15.0), Cm(3 * W + 2 * G), Cm(0.8)
set_runs(note, [("出典：LINEヤフー for Business 導入事例（アクネクリニック／鮨 酒 肴 杉玉）", 10, False, GRAY)])
# 複製した図形のIDが重ならないよう振り直す
for n, c in enumerate(s.shapes._spTree.iter(qn("p:cNvPr")), start=2):
    c.set("id", str(n))

# ---- 元S4（広告とLINEの役割）を削除 ----
sldIdLst = prs.slides._sldIdLst
sid = sldIdLst[3]
prs.part.drop_rel(sid.get(qn("r:id")))
sldIdLst.remove(sid)

# ---- 字体をメイリオに統一（patch_yaruki_v32_iikiri.py と同じ手順）----
AFTER = {qn(x) for x in ("a:sym", "a:hlinkClick", "a:hlinkMouseOver", "a:rtl", "a:extLst")}
for sl in prs.slides:
    for r in sl.shapes._spTree.iter(qn("a:r")):
        if r.find(qn("a:rPr")) is None:
            r.insert(0, r.makeelement(qn("a:rPr"), {"lang": "ja-JP"}))
    for tag in ("a:rPr", "a:endParaRPr", "a:defRPr"):
        for rpr in sl.shapes._spTree.iter(qn(tag)):
            fs = []
            for ft in ("a:latin", "a:ea", "a:cs"):
                e = rpr.find(qn(ft))
                if e is None:
                    e = rpr.makeelement(qn(ft), {})
                else:
                    rpr.remove(e)
                e.set("typeface", "メイリオ")
                fs.append(e)
            nxt = next((c for c in rpr if c.tag in AFTER), None)
            for e in fs:
                (nxt.addprevious(e) if nxt is not None else rpr.append(e))
prs.save(str(OUT))
print("saved:", OUT.name, len(prs.slides), "枚")
