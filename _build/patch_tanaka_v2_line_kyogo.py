# -*- coding: utf-8 -*-
"""タナカホーム：競合の2ページを「会社の強みの比較」から「他社がLINEをどう運用しているか」に作り直す（2026-10-19 殿村さん指示）。

    python -m _build.patch_tanaka_v2_line_kyogo <殿村さんのPC保存版v1.pptx> <出力.pptx>

  直すのは「競合の強み：…」（5枚目）と「競合の取り組みに対して：…」（6枚目）だけ。ほかのページは触らない。
  ・5枚目：大進建設・一条工務店のLINE公式アカウントの実際の画面（2026年10月19日に友だち追加して殿村さんが撮影）で、
           友だち追加の入口と特典／あいさつ／リッチメニューを見せる。ステップ配信は後日の画面を入れる差し込み枠
  ・6枚目：2社の運用を受けて「タナカホームはLINEでこう運用する」を表で示す
  スクショは加工しない。一条工務店のあいさつに出る個人の名前だけ灰色で伏せた（`_images/tanaka_line_ichijo_*.png`）。
  元の画面は `_data/tanaka/競合LINE_20261019/`。
"""
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from _build.arauchi_v4.lib import NAVY, TEXT, add_box, add_table, set_paras
from _build.arauchi_v10.build import TINT, by_id, runs
from _build.build_sunao_v1 import all_meiryo
from _build.build_tanaka_v1 import GRAY, head_cell, cell, label_cell, note

IMG = "_images/"
LINE_C = "BFBFBF"


def drop(s, *ids):
    for sid in ids:
        e = by_id(s, sid)._element
        e.getparent().remove(e)


def slot_label(s, x, w, text):
    add_box(s, x, 2.05, w, 0.46, [runs((t, dict(size=12, bold=True, color=NAVY))) for t in text.split("\n")],
            align="ctr", anchor="b", ins=(0, 0, 0, 0.02), line_spacing=100, name=f"見出し {text}")


def caption(s, x, w, text):
    add_box(s, x, 5.76, w, 0.78, [runs((text, dict(size=12, color=TEXT)))], align="l", anchor="t",
            ins=(0.02, 0.02, 0.02, 0), line_spacing=115, name="画面の説明")


def shot(s, path, x, w, top=2.56, h=3.14):
    """スマホの画面（277×600）を、枠の幅に合わせて縦横比のまま置く。細い灰色の線で囲む"""
    pic = s.shapes.add_picture(path, Inches(x), Inches(top), Inches(w), Inches(w * 600 / 277))
    pic.line.color.rgb = RGBColor(0xBF, 0xBF, 0xBF)
    pic.line.width = Pt(0.75)
    return pic


def step_slot(s, x, w, top=2.56, h=3.14):
    add_box(s, x, top, w, h, [runs(("ステップ配信の画面", dict(size=12, bold=True, color=GRAY))),
                              runs(("（差し込み）", dict(size=12, color=GRAY)))],
            line=GRAY, dash="dash", align="ctr", anchor="ctr", line_spacing=130, name="差し込み ステップ配信")


# ---------------------------------------------------------------- 5枚目：競合のLINE運用（実際の画面）
def p_screens(s):
    set_paras(by_id(s, 2), ["競合のLINE運用：大進建設は特典で来場へ、一条工務店は情報でつなぐ"])
    set_paras(by_id(s, 6), ["友だち追加の入口と特典・あいさつ・リッチメニューを、各社の実際の画面で見る"])
    drop(s, 7, 9, 10)                                   # 会社の強みの表・画像の説明・注記
    half, sw = 4.85, 1.45
    gap = (half - 3 * sw) / 2
    for k, (x0, name) in enumerate(((0.45, "大進建設"), (5.53, "一条工務店"))):
        add_box(s, x0, 1.62, half, 0.36, [runs((name, dict(size=13, bold=True, color="FFFFFF")))], fill=NAVY,
                align="ctr", anchor="ctr", name=f"会社 {name}")
    xs = [[0.45 + i * (sw + gap) for i in range(3)], [5.53 + i * (sw + gap) for i in range(3)]]
    # 大進建設：入口と特典（サイトのポップアップ・元のページの画像をそのまま使う）／あいさつ／ステップ配信
    x = xs[0]
    slot_label(s, x[0], sw, "入口と特典\n（サイト）")
    pop = by_id(s, 8)
    pop.left, pop.top, pop.width, pop.height = Inches(x[0]), Inches(2.56 + (3.14 - sw) / 2), Inches(sw), Inches(sw)
    caption(s, x[0], sw, "LINE登録で最大4万円分")
    slot_label(s, x[1], sw, "あいさつ")
    shot(s, IMG + "tanaka_line_daishin_aisatsu.png", x[1], sw)
    caption(s, x[1], sw, "来場特典・相談・来場予約をLINEで受付")
    slot_label(s, x[2], sw, "追加後の配信")
    step_slot(s, x[2], sw)
    # 一条工務店：あいさつ（画面の下はメニュー「TOP」）／リッチメニュー「コンテンツ」／ステップ配信
    x = xs[1]
    slot_label(s, x[0], sw, "あいさつ")
    shot(s, IMG + "tanaka_line_ichijo_top.png", x[0], sw)
    caption(s, x[0], sw, "アンケートへ案内。TOPに来場キャンペーン")
    slot_label(s, x[1], sw, "リッチメニュー")
    shot(s, IMG + "tanaka_line_ichijo_contents.png", x[1], sw)
    caption(s, x[1], sw, "コンテンツ：性能・暮らし・実例など")
    slot_label(s, x[2], sw, "追加後の配信")
    step_slot(s, x[2], sw)
    add_box(s, 5.375, 1.62, 0.01, 4.95, [[("", {})]], fill="D9D9D9", name="区切り")
    note(s, "※各社のLINE公式アカウントより（2026年10月19日確認）。サイトのポップアップは大進建設の公式サイトより（2026年10月16日確認）",
         y=6.62, h=0.28)


# ---------------------------------------------------------------- 6枚目：だからタナカホームはLINEでこう運用する
def p_tanaka(s):
    set_paras(by_id(s, 6), ["今の来場特典に加え、検討の順番に沿った情報で、資料請求から来場・契約まで追う"])
    drop(s, 7, 9)                                       # 前の表・セキスイハイムの注記
    T = dict(size=12, bold=True, color=NAVY, fill=TINT)
    rows = [[head_cell("", 12.5), head_cell("大進建設", 12.5), head_cell("一条工務店", 12.5),
             head_cell("タナカホームの運用（ご提案）", 12.5)],
            [label_cell("入口と特典", 12.5), cell("サイトのポップアップ「LINE登録で最大4万円分」", 12),
             cell("サイトの公式SNSの一覧から友だち追加", 12),
             cell("今の来場特典に加え、サイトを閉じるとき・資料請求の完了直後にも入口をつくる", **T)],
            [label_cell("あいさつ", 12.5), cell("来場特典・相談・来場予約をLINEで受け付け、資料請求とイベントのリンクを案内", 12),
             cell("1分のアンケートに答えると、オーナー様の実例や家づくりの知識が届く", 12),
             cell("来場予約・資料請求の案内に加え、平屋か2階建てか・予算の簡易診断へ案内する", **T)],
            [label_cell("リッチメニュー", 12.5), cell("友だち追加直後の画面では表示なし（相談・来場予約はメッセージで受付）", 12),
             cell("「TOP」（展示場・カタログ請求・来場キャンペーン）と「コンテンツ」（性能・暮らし・実例）の2つ", 12),
             cell("今の「予約／ホームページ／公式SNS」に、施工事例・ルームツアー・資金相談を足す", **T)],
            [label_cell("追加後の配信", 12.5), cell("", 12), cell("", 12),
             cell("前後検索の順番に沿ったステップ配信（全5通）で、来場予約まで案内する", **T)]]
    cw = [1.45, 2.55, 2.6, 3.13]
    rh = [0.42, 0.9, 0.9, 0.9, 0.78]
    add_table(s, 0.55, 1.72, cw, rh, rows, name="LINEの運用")
    y4 = 1.72 + sum(rh[:4])
    for x, w in ((0.55 + cw[0], cw[1]), (0.55 + cw[0] + cw[1], cw[2])):
        add_box(s, x + 0.06, y4 + 0.06, w - 0.12, rh[4] - 0.12,
                [runs(("ステップ配信の内容（差し込み）", dict(size=12, color=GRAY)))], line=GRAY, dash="dash",
                align="ctr", anchor="ctr", name="差し込み ステップ配信")
    band = by_id(s, 8)
    band.top = Inches(5.85)
    note(s, "※各社のLINE公式アカウントより（2026年10月19日確認）。入口・ポップアップは各社の公式サイトより（2026年10月16日確認）",
         y=6.6, h=0.28)


def main(base, out):
    prs = Presentation(base)
    sl = list(prs.slides)
    s5, s6 = sl[4], sl[5]
    assert by_id(s5, 2).text_frame.text.startswith("競合の強み")
    assert by_id(s6, 2).text_frame.text.startswith("競合の取り組みに対して")
    assert by_id(s5, 8).name == "大進建設のポップアップ"
    p_screens(s5)
    p_tanaka(s6)
    all_meiryo([s5, s6])
    prs.save(out)
    print("saved:", out, len(prs.slides), "枚")


if __name__ == "__main__":
    main(*sys.argv[1:3])
