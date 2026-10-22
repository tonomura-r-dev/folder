# -*- coding: utf-8 -*-
"""人材業界（派遣・バイト・単発）：P12 他社分析を、転職系から派遣・バイトの会社に差し替える（2026-10-22 殿村さん）。

    python -m _build.patch_haken_v3_p12 <入力.pptx> <出力.pptx>

  入力は `patch_haken_v2_chapters` を当てた版。調査は殿村さんがGPTで実施（2026-10-22）。
  友だち数は page.line.me のプロフィール表示値を 2026-10-22 にこちらで確認し直した値
  （プロフィールに会社の公式サイトのリンクが出ていることも確認。バイトル・タイミーは殿村さんがLINEアプリで確認した値。マイナビバイトはIDを特定できず）。
  求職者向けのアカウントを特定できなかったリクルートスタッフィング、2017年の発表しかないパーソルテンプスタッフは載せない。
  ほかのページは変えない。
"""
import sys

from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn

from _build.pptx_parts import rect, rich

NAVY, INK, SUB, ROW = "1E2761", "333333", "64748B", "F3F4F8"
IN = 2.54

TITLE = "［第2章］他社分析｜派遣・バイト大手のLINE運用"
LEAD = ["大手各社は、LINEを「求人を探す入口」か「希望を伺って求人を届ける場」として使っている。",
        "応募の手続きは、サイト・アプリ・マイページで行う会社が多い。"]
HEAD = ["会社（区分）", "友だち数", "LINEでできること（各社の案内）", "LINEの使い方"]
ROWS = [  # 会社, 区分, 友だち数, できること, 使い方
    ("スタッフサービス", "派遣", "5,248,970", "希望に合う求人情報が届く。ジョブタイプ診断、/在宅・週4日以内など条件別の求人", "探す入口＋診断"),
    ("アデコ", "派遣", "1,010,975", "追加後のアンケート（約1分）で希望を伺い、/条件に合う求人を配信。求人への応募も可能", "希望を伺って届ける"),
    ("ランスタッド", "派遣", "76,424", "希望条件に合う求人の案内。入社書類の提出、/就業中のチャット相談・勤怠報告まで", "届ける＋就業中の連絡"),
    ("マンパワーグループ", "派遣", "11,030", "求人情報・キャンペーン・セミナーの配信。/応募はJOBNETのマイページから", "マイページへつなぐ"),
    ("バイトル", "バイト", "15,848,943", "おすすめ求人や、あなたに合った情報を届ける。/トークのメニューから条件・キーワードで探す", "探す入口＋届ける"),
    ("マイナビバイト", "バイト", "―", "学生・主婦（夫）・フリーター向けの情報、/履歴書・面接対策、保存した求人の通知", "情報・通知の配信"),
    ("タイミー", "単発", "199,061", "あなたに合ったおすすめの求人を届ける。/仕事探し・申込みはアプリで行う", "届ける＋アプリへ"),
]
POINTS = [[("希望を伺って求人を届けるのはアデコ・ランスタッド、探す入口はスタッフサービス・バイトル。", 13, False, INK)],
          [("タイミー・マンパワーグループは、LINEで求人を届け、応募はアプリ・マイページで行う。", 13, True, NAVY)]]
BAND = "追加の直後に希望を伺い、合う求人を届ける形は、派遣大手がすでに取り入れている"
NOTE = ("出典｜各社の公式サイトのLINE紹介ページ・LINE公式アカウントのプロフィール（2026年10月22日確認）。友だち数は表示値。"
        "マンパワーグループは「JOBNET」のアカウント。マイナビバイトは友だち数を確認できず")
COLS = [(0.55, 2.25), (2.85, 1.25), (4.2, 4.05), (8.35, 1.95)]   # (左端, 幅) インチ


def set_paras(sh, texts):
    txb = sh.text_frame._txBody
    ps = txb.findall(qn("a:p"))
    import copy
    while len(ps) < len(texts):
        ps[-1].addnext(copy.deepcopy(ps[-1]))
        ps = txb.findall(qn("a:p"))
    for p in ps[len(texts):]:
        txb.remove(p)
    for p, t in zip(txb.findall(qn("a:p")), texts):
        runs = p.findall(qn("a:r"))
        for r in runs[1:]:
            p.remove(r)
        runs[0].find(qn("a:t")).text = t


def main(src, out):
    prs = Presentation(src)
    s = list(prs.slides)[11]
    shapes = {sh.name: sh for sh in s.shapes}
    assert shapes["TextBox 1"].text_frame.text.startswith("［第2章］他社分析")
    keep = {"TextBox 1", "TextBox 2", "Rectangle 40", "TextBox 41", "TextBox 42"}
    for name, sh in shapes.items():
        if name not in keep:
            sh._element.getparent().remove(sh._element)
    set_paras(shapes["TextBox 1"], [TITLE])
    set_paras(shapes["TextBox 2"], LEAD)
    set_paras(shapes["TextBox 41"], [BAND])
    set_paras(shapes["TextBox 42"], [NOTE])

    x0, w = 0.47, 9.89
    hy, hh, rh = 1.5, 0.34, 0.5
    rect(s, x0 * IN, hy * IN, w * IN, hh * IN, NAVY)
    for (x, cw), t in zip(COLS, HEAD):
        rich(s, x * IN, hy * IN, cw * IN, hh * IN, [[(t, 11.5, True, "FFFFFF")]],
             align=PP_ALIGN.LEFT if t != "友だち数" else PP_ALIGN.CENTER)
    for k, (name, kind, n, can, use) in enumerate(ROWS):
        y = hy + hh + rh * k
        if k % 2:
            rect(s, x0 * IN, y * IN, w * IN, rh * IN, ROW)
        (x1, w1), (x2, w2), (x3, w3), (x4, w4) = COLS
        rich(s, x1 * IN, y * IN, w1 * IN, rh * IN, [[(name, 12, True, INK), ("　" + kind, 10, False, SUB)]])
        rich(s, x2 * IN, y * IN, w2 * IN, rh * IN, [[(n, 12, False, INK)]], align=PP_ALIGN.CENTER)
        rich(s, x3 * IN, y * IN, w3 * IN, rh * IN, [[(t, 11, False, INK)] for t in can.split("/")])   # 意味の切れ目で改行
        rich(s, x4 * IN, y * IN, w4 * IN, rh * IN, [[(use, 11.5, True, NAVY)]])
    bottom = hy + hh + rh * len(ROWS)
    rich(s, x0 * IN, (bottom + 0.08) * IN, w * IN, 0.6 * IN, POINTS, anchor=MSO_ANCHOR.MIDDLE)
    for nm in ("Rectangle 40", "TextBox 41"):          # 結論の帯：本文の直下へ
        sh = shapes[nm]
        sh.top = int((bottom + 0.78 + (0.04 if nm == "TextBox 41" else 0)) * 914400)
    assert shapes["Rectangle 40"].top / 914400 + 0.41 < 6.8
    prs.save(out)
    print("saved:", out, round(bottom, 2))


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
