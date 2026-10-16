# -*- coding: utf-8 -*-
"""タナカホーム「LINE公式アカウント運用のご提案」：殿村さんのPC保存版（2026-10-16・28枚）に、まとめて手を入れる。

    python -m _build.patch_tanaka_pc1 <PC保存版.pptx> <出力.pptx>

  ① 「資料請求の時点から」→「資料請求以前から」（提案の流れの表・解決策のタイトル）
  ② 競合の強み・競合の取り組みに対して（新・2枚）：LINE公式アカウントの現状の次に入れる
  ③ 前後検索「注文住宅」：言葉を4つに色分けした画像（_build/make_tanaka_zengo_color.py）＋時期ごとの分析
  ④ 解決策：追う範囲の違いを帯で見せ、「これで可能になること」を3つ
  ⑤ 配信・試算・費用：あらうちv10の形（配信／目標と費用対効果／料金プラン／運用内容）に作り直す → _build/patch_tanaka_pc2.py
    （p_hiyo の表1枚の版は使わなくなった。「松竹梅」の文字は出さない）
  ⑥ SEOの言い方を「SEO」にそろえる（提案の流れ・現状と打ち手）
  殿村さんが消した注記・直した文言は戻さない。ほかのページは触らない。
"""
import copy
import sys

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Inches

from _build.arauchi_v4.lib import NAVY, RED, TEXT, add_box, add_table
from _build.arauchi_v10.build import TINT, LINE_C, MID, by_id, meiryo, runs
from _build.build_tanaka_v1 import GRAY, head_cell, cell, label_cell, note, band
from _build.build_tanaka_v2 import heading, flow, X0, W
from _build.patch_tanaka_v4 import new_page, set_lines, HDR
from _build.patch_tanaka_v5 import IMG, SOFT, pic, caption
from _build.arauchi_v10.add_zengo import alpha_fill
from _build.patch_tanaka_pc2 import rebuild

COMP = IMG.parent / "_data/tanaka"


def keep_header(s):
    for sh in list(s.shapes):
        if sh.shape_id not in HDR:
            sh._element.getparent().remove(sh._element)


def replace_text(s, old, new):
    n = 0
    for t in s._element.iter(qn("a:t")):
        if t.text and old in t.text:
            t.text = t.text.replace(old, new)
            n += 1
    return n


# ---------------------------------------------------------------- ② 競合のLINE（新）
def p_kyogo(prs, src):
    """競合の強み（2026-10-16 殿村さんが集めた競合情報のうち、公式サイト・LINEのプロフィールで確認できたものだけ）"""
    s = new_page(prs, src, "競合の強み：一条工務店は性能と実績、大進建設は全館空調と特典",
                 "青森で比べられる2社は、性能に加えて、実績・特典でも選ぶ理由をつくっている")
    rows = [[head_cell(""), head_cell("一条工務店"), head_cell("大進建設"), head_cell("タナカホーム")],
            [label_cell("主な訴求", 10.5), cell("「家は、性能。」", 10.5, True, NAVY),
             cell("全館空調\n7つの機能と保証", 10.5, True, NAVY),
             cell("地元を知り尽くした地場の会社\nW断熱・トリプルガラス樹脂窓（標準）", 10, True, NAVY)],
            [label_cell("実績・拠点", 10.5), cell("青森県戸建 着工棟数 No.1（2024年度）\n県内の建築実績 4,167棟", 9.5),
             cell("青森を中心に、宮城にも拠点", 9.5), cell("八戸で1973年に創業（50年以上）", 9.5)],
            [label_cell("LINEの取り組み", 10.5), cell("サイトの公式SNSの一覧から、公式LINEへ", 9.5),
             cell("LINE登録で最大4万円分\n（はじめての資料請求と来場予約で1万円分、来場後の事前審査で最大3万円分）", 9.5, True, NAVY),
             cell("来場特典に上乗せで友だち登録\n北国の暮らしのお役立ち情報を発信", 9.5)]]
    add_table(s, X0, 1.85, [1.35, 2.0, 2.0, 2.05], [0.4, 0.85, 0.75, 1.15], rows, name="競合の強み")
    pic(s, COMP / "大進建設_LINE登録キャンペーン_20261016.png", 8.1, 1.85, 2.18, 2.18, name="大進建設のポップアップ")
    caption(s, 8.0, 4.08, 2.38, "大進建設のサイトのポップアップ")
    note(s, "※各社の公式サイト・LINE公式アカウントより（2026年10月16日確認）。一条工務店の着工棟数は住宅産業研究所調べ（2024年度）",
         y=5.4, h=0.45)
    return s


def p_taiou(prs, src):
    """競合の取り組みに対して、タナカホームがLINEで行うこと"""
    s = new_page(prs, src, "競合の取り組みに対して：LINEで、タナカホームを選ぶ理由を届ける",
                 "性能・診断・特典で比べられる中、地元の会社ならではの情報で、選ばれる理由をつくる")
    rows = [[head_cell("競合の取り組み"), head_cell(""), head_cell("タナカホームのLINEで行うこと")],
            [cell(["性能を前面に打ち出す", "一条工務店「家は、性能。」"], 10.5, color=TEXT),
             cell("▶", 12, True, NAVY, align="ctr"),
             cell("青森の冬の暮らし（暖房費・雪・断熱）を、地元の施工事例とあわせて届ける", 11, True, NAVY)],
            [cell(["LINEで診断・間取り作成を受け付ける", "セキスイハイム「カタログ無料診断」「無料間取り作成」"], 10.5),
             cell("▶", 12, True, NAVY, align="ctr"),
             cell("予算の簡易診断、平屋か2階建てかの診断で希望を聞き分け、届ける情報を分ける", 11, True, NAVY)],
            [cell(["特典で来場の先まで追う", "大進建設「LINE登録で最大4万円分」"], 10.5),
             cell("▶", 12, True, NAVY, align="ctr"),
             cell("特典の金額ではなく、検討の順番に沿った情報で、契約まで追う", 11, True, NAVY)]]
    add_table(s, X0, 1.85, [4.2, 0.45, 5.08], [0.4, 0.95, 0.95, 0.95], rows, name="競合の取り組みに対して")
    band(s, [("ポイントは、比べられている間に、", dict(size=13, bold=True, color=NAVY)),
             ("タナカホームを選ぶ理由を届けられること", dict(size=13, bold=True, color=RED))], y=5.3, h=0.62)
    note(s, "※セキスイハイムはLINE公式アカウントのプロフィールより（2026年10月16日確認）", y=6.05)
    return s


# ---------------------------------------------------------------- ③ 前後検索「注文住宅」
def p_zengo(s):
    keep_header(s)
    iw = 6.9
    ih = iw * 1167 / 2000
    pic(s, IMG / "tanaka_zengo_chumon_color.png", X0, 1.75, iw, ih, name="前後検索 注文住宅（色分け）")
    # LINE追加の時期（検索後0〜1日）：0日＝1027px・1日＝1260px（画像は2000×1167px）
    px = lambda v: X0 + v / 2000 * iw
    py = lambda v: 1.75 + v / 1167 * ih
    gx0, gx1 = px(1027), px(1260)
    zone = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(gx0), Inches(py(15)), Inches(gx1 - gx0), Inches(py(1105) - py(15)))
    zone.name = "LINE追加の時期"
    alpha_fill(zone, "06C755", 16000)
    zone.shadow.inherit = False
    add_box(s, (gx0 + gx1) / 2 - 0.5, 1.75 + ih + 0.04, 1.0, 0.26,
            [runs(("LINE追加", dict(size=9.5, bold=True, color="00873C")))],
            fill="C6EFCE", align="ctr", anchor="ctr", ins=(0, 0, 0, 0), name="LINE追加 札")
    x, w = X0 + iw + 0.2, W - iw - 0.2
    legend = [("3467B2", "住宅会社を比べる"), ("D9661F", "費用・住宅ローンを調べる"),
              ("8E4EC6", "間取り・性能を調べる"), ("00897B", "土地を調べる")]
    for k, (c, t) in enumerate(legend):
        y = 1.75 + k * 0.27
        add_box(s, x, y + 0.06, 0.15, 0.15, [runs(("", dict(size=8)))], fill=c, ins=(0, 0, 0, 0), name="凡例の色")
        add_box(s, x + 0.2, y, w - 0.2, 0.27, [runs((t, dict(size=9.5, color=TEXT)))], anchor="ctr",
                ins=(0, 0, 0, 0), name="凡例")
    items = [("当日", "住宅会社名が中心", "一条工務店・積水ハウス・住友林業など。住宅展示場・土地探しも"),
             ("1〜7日後", "住宅ローンと坪単価が増える", "シミュレーション・審査・控除、坪単価。比べる会社はさらに広がる"),
             ("7〜15日後", "間取り・性能が中心", "断熱等級・吹き抜け・パントリー・ルームツアー。住宅ローンは金利の比較へ")]
    y = 2.9
    for h1, h2, t in items:
        add_box(s, x, y, w, 0.28, [runs((h1, dict(size=10, bold=True, color="FFFFFF")))], fill=NAVY, align="ctr",
                anchor="ctr", name=f"分析 {h1}")
        add_box(s, x, y + 0.28, w, 0.7, [runs((h2, dict(size=10, bold=True, color=NAVY))), runs((t, dict(size=9)))],
                line=LINE_C, anchor="t", line_spacing=112, ins=(0.08, 0.05, 0.08, 0.04), name=f"分析 {h1} 本文")
        y += 1.06
    band(s, [("会社を比べる → 費用 → 間取り・性能 の順に、", dict(size=12.5, bold=True, color=NAVY)),
             ("約2週間かけて検討が進む", dict(size=12.5, bold=True, color=RED))], y=6.12, h=0.5)
    note(s, "※LINEヤフー社の前後検索データより。「注文住宅」を検索した方が、前後15日間に検索した言葉（縦軸＝検索した人数）。色分けは弊社",
         y=6.66, h=0.3, size=8.5)


# ---------------------------------------------------------------- ④ 解決策：追う範囲
def p_ou(s):
    keep_header(s)
    steps = ["サイトに来る", "資料請求", "来場予約", "来場", "商談", "契約"]
    n, aw = len(steps), 0.2
    bw = (W - aw * (n - 1)) / n
    flow(s, [(t, NAVY, "FFFFFF") for t in steps], 1.85, 0.5, arrow_w=aw, size=11)
    col = lambda k: X0 + k * (bw + aw)
    add_box(s, X0, 2.55, 1.6, 0.48, [runs(("今のLINE", dict(size=10.5, bold=True, color=GRAY)))], align="ctr",
            anchor="ctr", name="帯の名前 今")
    add_box(s, col(2), 2.55, col(3) + bw - col(2), 0.48,
            [runs(("来場予約・来場のご案内", dict(size=11, bold=True, color=GRAY)))], fill=SOFT, line=LINE_C,
            align="ctr", anchor="ctr", name="帯 今のLINE")
    add_box(s, X0, 3.13, W, 0.55, [runs(("弊社のご提案：資料請求の前から、契約まで追う", dict(size=12.5, bold=True, color="FFFFFF")))],
            fill=NAVY, align="ctr", anchor="ctr", name="帯 弊社のご提案")
    heading(s, X0, 3.95, W, "これで可能になること")
    items = [("01", "来場を決める前の方を取りこぼさない", "サイトに来た方・資料請求をした方が友だちになり、比べている間も情報を届けられる"),
             ("02", "来場予約のキャンセルを減らす", "予約から来場までをLINEでつなぎ、前日にお知らせする"),
             ("03", "契約につながる集客が分かる", "契約までの結果を貴社の記録と合わせて確認し、広告の改善に生かす")]
    gw = 0.2
    cw = (W - gw * 2) / 3
    for k, (no, h1, t) in enumerate(items):
        x = X0 + k * (cw + gw)
        add_box(s, x, 4.45, cw, 1.45, [runs((no, dict(size=16, bold=True, color=RED))),
                                       runs((h1, dict(size=11, bold=True, color=NAVY))),
                                       runs((t, dict(size=10, color=TEXT)))],
                line=LINE_C, anchor="t", line_spacing=120, ins=(0.12, 0.08, 0.12, 0.06), name=f"可能になること{no}")
    note(s, "※広告の改善は、別紙の集客改善シミュレーションとあわせてご提案", y=6.4)


# ---------------------------------------------------------------- ⑤ 費用：松竹梅
def p_hiyo(s):
    keep_header(s)
    set_lines(by_id(s, HDR[0]), ["費用：運用の範囲に合わせて、3つのプランから選べる"])
    set_lines(by_id(s, HDR[4]), ["おすすめは、構築から運用・改善まで一貫して行う「構築＋運用プラン」"])

    def plan_head(name, sub, fill):
        return dict(text=[name, sub] if sub else name, fill=fill, color="FFFFFF", bold=True, size=11, borders=head_cell("")["borders"])

    mark = lambda v, on=False: cell(v, 11, True, NAVY if on else TEXT, TINT if on else "FFFFFF", "ctr")
    rows = [[head_cell(""), plan_head("構築プラン", "", MID), plan_head("構築＋運用プラン", "（おすすめ）", NAVY),
             plan_head("構築＋運用＋広告プラン", "", MID)],
            [label_cell("初期費用", 11), mark("26.5万円〜"), mark("26.5万円〜", True), mark("26.5万円〜")],
            [label_cell("月額費用", 11), cell("6万円〜", 14, True, NAVY, align="ctr"),
             cell("26万円〜", 16, True, NAVY, TINT, "ctr"), cell("46万円〜", 14, True, NAVY, align="ctr")],
            [label_cell("月額の内訳", 10.5), cell("離脱防止 3万円\nサンクスLINE誘導 3万円〜", 9.5, align="ctr"),
             cell("構築プランの内容\n＋LINE運用コンサル 20万円〜", 9.5, fill=TINT, align="ctr"),
             cell("構築＋運用プランの内容\n＋LINE広告（友だち追加）20万円〜", 9.5, align="ctr")],
            [label_cell("友だちの入口", 10.5), mark("○"), mark("○", True), mark("○")],
            [label_cell("ステップ配信の構築", 10.5), mark("○"), mark("○", True), mark("○")],
            [label_cell("企画配信・\n来場前のお知らせ", 10), cell("貴社で運用", 10, color=GRAY, align="ctr"), mark("○", True), mark("○")],
            [label_cell("結果の分析・改善", 10.5), cell("―", 11, color=GRAY, align="ctr"), mark("○", True), mark("○")],
            [label_cell("LINE広告で\n友だちを増やす", 10), cell("―", 11, color=GRAY, align="ctr"),
             cell("―", 11, color=GRAY, fill=TINT, align="ctr"), mark("○")]]
    add_table(s, X0, 1.75, [2.05, 2.56, 2.56, 2.56], [0.55, 0.4, 0.48, 0.55, 0.36, 0.36, 0.48, 0.36, 0.48], rows, name="3つのプラン")
    note(s, "※初期費用＝運用設計・構築 10万円〜＋離脱防止 1.5万円＋サンクスLINE誘導 15万円。LINE公式アカウント利用料は別"
            "（貴社のご契約のまま）。契約期間は6か月〜。試算のページは構築＋運用プランの場合", y=6.25, h=0.42, size=8.5)
    note(s, "※施策内容・配信本数・連携範囲等により、費用は個別にお見積もりいたします。", y=6.7, size=8.5)


def main(base, out):
    prs = Presentation(base)
    sl = list(prs.slides)
    assert len(sl) == 28
    nagare, line, kadai, zengo, ou, hiyo = sl[1], sl[3], sl[4], sl[5], sl[6], sl[11]
    assert by_id(zengo, 2).text_frame.text.startswith("前後検索「注文住宅」")
    assert by_id(hiyo, 2).text_frame.text.startswith("費用")
    assert replace_text(nagare, "資料請求の時点から", "資料請求以前から") == 1
    assert replace_text(ou, "資料請求の時点から", "資料請求以前から") >= 1
    # ⑥ SEOは「地域SEO」「地域特化のSEO」ではなく、ただの「SEO」にそろえる（2026-10-16 殿村さん）
    assert replace_text(nagare, "広告と地域SEOで", "広告とSEOで") == 1
    assert replace_text(sl[2], "地域特化のSEO", "SEO") == 2
    kyogo = p_kyogo(prs, kadai)
    taiou = p_taiou(prs, kadai)
    p_zengo(zengo)
    p_ou(ou)
    for s in (kyogo, taiou, zengo, ou):
        for sh in s.shapes:
            if sh.has_text_frame:
                meiryo(sh)
    lst = prs.slides._sldIdLst
    els = list(lst)[-2:]                                 # 競合の強み・競合の取り組みに対して（いま末尾）を「LINE公式アカウントの現状」の次へ
    for k, el in enumerate(els):
        lst.remove(el)
        lst.insert(4 + k, el)
    rebuild(prs, sl[8], sl[10], hiyo)                    # 配信・試算・費用 → あらうちv10の形（_build/patch_tanaka_pc2.py）
    prs.save(out)
    print("saved:", out, len(prs.slides), "枚")


if __name__ == "__main__":
    main(*sys.argv[1:3])
