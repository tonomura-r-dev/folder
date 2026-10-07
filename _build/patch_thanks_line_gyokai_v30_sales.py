# -*- coding: utf-8 -*-
"""サンクスLINE誘導・業界別 ver2.0（殿村さんPC保存版）→ ver3.0（2026-10-07 殿村さん指示）。
「広告提案時にサンクスLINEをセット提案する営業資料」として5枚に組み直す。4枚目（業界別）だけそのまま残す。
  1 広告と一緒にサンクスLINEを実施する理由（CV獲得まで＝CPA改善／CV獲得後＝CPO改善）
  2 サンクスLINE｜機能・費用（フォーム→サンクス画面→自動遷移→未追加／追加済み＋入力内容の引き継ぎ＋費用）
  3 実例｜予約後の来院率（美容クリニック 課題→導入→結果。数字は1つだけ。LINEヤフー事例は数字なしの参考）
  4 ターゲット業界別（変更なし）
  5 DYMのコンサル支援内容・費用（①導入→②配信設計→③効果計測→④改善）
画面画像は DYM共通資料（LINEOA_BUFFF_3.pptx 31枚目）の実画面をロゴ・社名だけ消して流用（_images/thanks_tool_*.png）。
費用は承認済みの「初期15万円・月額3万円〜＋設定地点追加」だけ。運用コンサル費用は別途お見積もり。
  python3 _build/patch_thanks_line_gyokai_v30_sales.py <業界別ver2.0.pptx> [出力]
"""
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Cm
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

sys.path.insert(0, str(Path(__file__).resolve().parent))
from pptx_parts import *  # noqa

ROOT = Path(__file__).resolve().parent.parent
IMG = ROOT / "_images"
OUT = Path(sys.argv[2]) if len(sys.argv) > 2 else ROOT / "20261007_サンクスLINE誘導のご提案_業界別ver3.0.pptx"
ORANGE, ORANGE_BG, ORANGE_TX = "F59B21", "FEF0DE", "C46A00"
PALE, OFF = "E4E8F6", "F2F2F2"
L, R = 2.2, 31.7
CW = R - L
OV = 0.35

prs = Presentation(sys.argv[1])
S = list(prs.slides)
assert len(S) == 5 and S[3].shapes[0].text_frame.text.startswith("ターゲット業界別"), "ver2.0（5枚・4枚目が業界別）ではない"


def head(s, title, lead):
    set_paras(find(s, "TextBox 1"), [(title, 0)])
    set_paras(find(s, "TextBox 2"), [(lead, 0)])


def flow(s, x0, x1, y, h, steps, size=13):
    """steps = [(行, 塗り, 文字色, 幅の比)]。x0〜x1 に矢羽根を OV ずつ重ねて並べる。返り値 [(x, w)]"""
    k = (x1 - x0 + OV * (len(steps) - 1)) / sum(st[3] for st in steps)
    out, x = [], x0
    for i, (ls, f, tc, wr) in enumerate(steps):
        w = wr * k
        shape(s, MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON, x, y, w, h, None, fill=f, adj=0.22)
        pad = 0.3 if i == 0 else 0.65
        label(s, x + pad, y, w - pad - 0.55, h, [(t, size, True, tc, 0) for t in ls], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        out.append((x, w))
        x += w - OV
    return out


def chip(s, x, y, w, h, text, fill, color, size=11):
    return shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, [(text, size, True, color, 0)], fill=fill, adj=0.5, margins=(0.1, 0, 0.1, 0))


def pic(s, name, x, y, h):
    im_w, im_h = {"form": (351, 479), "thanks": (252, 448), "add": (296, 526), "talk": (296, 526)}[name]
    w = h * im_w / im_h
    p = s.shapes.add_picture(str(IMG / f"thanks_tool_{name}.png"), Cm(x), Cm(y), Cm(w), Cm(h))
    p.line.color.rgb = rgb(LGRAY)
    p.line.width = Pt(0.75)
    return w


# ================= 1 広告と一緒に実施する理由 =================
s = S[0]
keep_header_only(s)
head(s, "広告と一緒にサンクスLINEを実施する理由", "広告はCVまで、サンクスLINEはCVの後を担う")
LW = 4.6          # 左の段階ラベル
FX0 = L + LW + 0.3
FX1 = 24.4        # 矢羽根の右端
RX, RW = 25.6, R - 25.6  # 結果の箱


def stage(y, h, name, zone, zone_fill, zone_tx):
    label(s, L, y, LW, h * 0.55, [(name, 16, True, NAVY, 0)], anchor=MSO_ANCHOR.BOTTOM)
    chip(s, L + 0.1, y + h * 0.58, 2.6, 0.7, zone, zone_fill, zone_tx, 10.5)


# CV獲得まで（AD領域）→ CPA改善
Y1, H1 = 5.0, 2.9
stage(Y1, H1, "CV獲得まで", "AD領域", ORANGE_BG, ORANGE_TX)
flow(s, FX0, FX1, Y1, H1, [(["広告"], ORANGE, WHITE, 1), (["LP"], ORANGE_BG, ORANGE_TX, 1), (["離脱防止施策"], ORANGE_BG, ORANGE_TX, 1.15),
                           (["CV"], ORANGE, WHITE, 1)], size=14)
shape(s, MSO_SHAPE.RECTANGLE, RX, Y1, RW, H1, [("CPA改善", 20, True, ORANGE_TX, 2), ("CV1件あたりの広告費", 9.5, False, GRAY, 0)],
      fill=WHITE, line=ORANGE, lw=2.0)
label(s, FX0, Y1 + H1 + 0.1, FX1 - FX0, 0.7, [("離脱しそうなユーザーを引き止め、CVの取りこぼしを減らす", 10, False, GRAY, 0)],
      align=PP_ALIGN.CENTER)

hline(s, L, R, 9.55, LGRAY, 0.75)

# CV獲得後（LINE領域）→ CPO改善
Y2, H2 = 10.3, 2.9
stage(Y2, H2, "CV獲得後", "LINE領域", GREEN_BG, GREEN_TX)
xs = flow(s, FX0, FX1, Y2, H2, [(["CV"], PALE, NAVY, 0.8), (["サンクスLINE誘導"], GREEN, WHITE, 1.25), (["LINEで継続フォロー"], GREEN_BG, GREEN_TX, 1.25),
                                (["来店・面談・購入"], NAVY, WHITE, 1.4)], size=13)
shape(s, MSO_SHAPE.RECTANGLE, RX, Y2, RW, H2, [("CPO改善", 20, True, WHITE, 2), ("来店・面談・購入1件あたりの費用", 9.5, False, "D9DCEA", 0)],
      fill=NAVY)
label(s, xs[2][0], Y2 + H2 + 0.1, xs[2][1] + xs[3][1] - OV, 0.7, [("リマインド・情報提供で、CVの後の離脱を防ぐ", 10, False, GRAY, 0)],
      align=PP_ALIGN.CENTER)

# 結論（1行だけ）
label(s, L, 15.4, CW, 1.4, [("広告のCPA改善だけではなく、CV後まで含めて事業全体の獲得効率を改善する", 17, True, NAVY, 0)],
      align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

# ================= 2 機能・費用 =================
s = S[1]
keep_header_only(s)
head(s, "サンクスLINE｜機能・費用", "フォーム送信後のLINE誘導を自動化するツール")
PH = 5.6   # 画面画像の高さ
PY = 6.4   # 画面画像の上端
CY = 5.5   # 見出しチップの上端
# 左：WEBフォーム → サンクス画面
w_form = PH * 351 / 479
w_th = PH * 252 / 448
x_form = L
x_th = x_form + w_form + 1.4
chip(s, x_form, CY, w_form, 0.7, "WEBフォーム", OFF, INK)
pic(s, "form", x_form, PY, PH)
arrow_line(s, x_form + w_form + 0.25, PY + PH / 2, x_th - 0.25, PY + PH / 2, NAVY, 2.0)
chip(s, x_th, CY, w_th, 0.7, "サンクス画面", OFF, INK)
pic(s, "thanks", x_th, PY, PH)
label(s, x_th + 0.15, PY + 0.33, w_th - 0.3, 0.45, [("WEBサイト名", 7, True, GRAY, 0)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
label(s, x_form, PY + PH + 0.1, w_form, 0.7, [("入力・送信", 10.5, True, INK, 0)], align=PP_ALIGN.CENTER)
label(s, x_th, PY + PH + 0.1, w_th, 0.7, [("CV完了", 10.5, True, INK, 0)], align=PP_ALIGN.CENTER)

# 中央：自動遷移＋DYM開発ツール
GX0 = x_th + w_th + 0.6
GX1 = 20.0
gw = GX1 - GX0 - 0.6
arrow_line(s, x_th + w_th + 0.15, PY + PH / 2, GX0 + 0.3 - 0.1, PY + PH / 2, NAVY, 2.0)
shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, GX0 + 0.3, PY + PH / 2 - 1.35, gw, 2.7,
      [("CV完了後 約2〜3秒で", 12, True, WHITE, 2), ("LINEへ自動遷移", 14, True, WHITE, 0)], fill=GREEN, adj=0.15, margins=(0.15, 0, 0.15, 0))
chip(s, GX0 + 0.3 + gw / 2 - 2.0, PY + PH / 2 + 1.6, 4.0, 0.75, "DYM開発ツール", NAVY, WHITE, 10.5)
arrow_line(s, GX0 + 0.3 + gw + 0.15, PY + PH / 2, GX1 + 0.3, PY + PH / 2, NAVY, 2.0)

# 右：動作イメージ（未追加／追加済み）
DX0 = GX1 + 0.5
rect(s, DX0, 4.5, R - DX0, 0.8, NAVY, [("動作イメージ", 12, True, WHITE, 0)])
colw = (R - DX0) / 2
cases = [("LINE未追加", "add", "友だち登録ページへ遷移"), ("LINE追加済み", "talk", "トーク画面へ遷移")]
w_ln = PH * 296 / 526
for i, (nm, img, cap) in enumerate(cases):
    cx = DX0 + i * colw
    if i:
        vline(s, cx, CY, PY + PH + 0.8, LGRAY, 0.75)
    chip(s, cx + 0.3, CY, colw - 0.6, 0.7, nm, GREEN_BG, GREEN_TX)
    px = cx + (colw - w_ln) / 2
    pic(s, img, px, PY, PH)
    if img == "add":
        label(s, px + 0.2, PY + 1.35, w_ln - 0.4, 1.2, [("ロゴ", 8, True, GRAY, 0)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        label(s, px + 0.1, PY + 2.65, w_ln - 0.2, 0.45, [("アカウント名", 7.5, True, INK, 0)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    else:
        tb = label(s, px + 0.6, PY + 0.22, 1.6, 0.4, [("アカウント名", 5.5, True, INK, 0)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        tb.text_frame.word_wrap = False
    label(s, cx, PY + PH + 0.1, colw, 0.7, [(cap, 10.5, True, INK, 0)], align=PP_ALIGN.CENTER)

# 下：入力内容の引き継ぎ（分岐に紐づけない全体の機能）
BY = PY + PH + 1.2
shape(s, MSO_SHAPE.RECTANGLE, L, BY, CW, 1.0, [("フォーム入力内容をLINEトーク入力画面へ引き継ぎ可能", 12.5, True, NAVY, 0)], fill=PALE)

# 費用
FY = BY + 1.5
label(s, L, FY, 2.2, 1.6, [("費用", 13, True, NAVY, 0)], anchor=MSO_ANCHOR.MIDDLE)
rect(s, L + 2.2, FY, 6.4, 1.6, NAVY)
rich(s, L + 2.2, FY, 6.4, 1.6, [[("初期費用 ", 11, True, WHITE), ("15万円", 18, True, WHITE)]], align=PP_ALIGN.CENTER)
rect(s, L + 8.9, FY, 7.6, 1.6, NAVY)
rich(s, L + 8.9, FY, 7.6, 1.6, [[("月額 ", 11, True, WHITE), ("3万円〜", 18, True, WHITE), ("　+設定地点追加", 10, True, WHITE)]], align=PP_ALIGN.CENTER)
label(s, L + 16.9, FY, R - (L + 16.9), 1.6, [("※既存APIツールとの併用についても、問題なく稼働が可能。", 10, False, GRAY, 0)], anchor=MSO_ANCHOR.MIDDLE)

# ================= 3 実例 =================
s = S[2]
keep_header_only(s)
head(s, "実例｜予約後の来院率", "予約直後にLINEでつながり、前日のお知らせで来院につなげた")
label(s, L, 4.5, 10, 0.9, [("美容クリニック", 17, True, NAVY, 0)], anchor=MSO_ANCHOR.MIDDLE)
cols = [("課題", 7.6), ("導入・活用方法", 11.6), ("結果", 8.3)]
gap = (CW - sum(w for _, w in cols)) / 2
TY, TH = 5.7, 8.6
x = L
colx = []
for i, (nm, w) in enumerate(cols):
    rect(s, x, TY, w, 0.85, NAVY if i == 2 else PALE, [(nm, 12, True, WHITE if i == 2 else NAVY, 0)])
    shape(s, MSO_SHAPE.RECTANGLE, x, TY + 0.85, w, TH - 0.85, None, fill=WHITE, line=LGRAY, lw=1.0)
    colx.append((x, w))
    if i < 2:
        arrow_line(s, x + w + 0.2, TY + TH / 2, x + w + gap - 0.2, TY + TH / 2, NAVY, 2.0)
    x += w + gap
# 課題
cx, cw = colx[0]
label(s, cx + 0.5, TY + 1.4, cw - 1.0, TH - 2.0, [("予約から来院日まで、接点がない", 11.5, True, INK, 10),
                                                   ("予約後に来院しない方が一定数いる", 11.5, True, INK, 0)], anchor=MSO_ANCHOR.MIDDLE)
# 導入・活用方法（縦のステップ）
cx, cw = colx[1]
steps = [("予約完了画面", PALE, NAVY), ("サンクスLINE誘導", GREEN, WHITE), ("LINE友だち追加", GREEN_BG, GREEN_TX),
         ("予約確認・前日のお知らせをLINEで配信", NAVY, WHITE)]
sh_ = (TH - 0.85 - 0.7) / len(steps)
for j, (t, f, c) in enumerate(steps):
    y = TY + 0.85 + 0.35 + j * sh_
    shape(s, MSO_SHAPE.RECTANGLE, cx + 0.8, y + 0.15, cw - 1.6, sh_ - 0.6, [(t, 11.5, True, c, 0)], fill=f)
    if j < len(steps) - 1:
        arrow_line(s, cx + cw / 2, y + sh_ - 0.42, cx + cw / 2, y + sh_ + 0.12, NAVY, 1.5)
# 結果
cx, cw = colx[2]
label(s, cx, TY + 1.3, cw, 0.9, [("予約後の来院率", 13, True, INK, 0)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
label(s, cx, TY + 2.4, cw, 3.4, [("40〜50%", 40, True, NAVY, 0)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
label(s, cx, TY + 5.8, cw, 1.2, [("改善", 22, True, NAVY, 0)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
# 参考（LINEヤフー公式事例・数字は載せない）
hline(s, L, R, 15.0, LGRAY, 0.75)
label(s, L, 15.2, CW, 0.6, [("参考｜LINEでつながった後の活用例（LINEヤフー公式の導入事例）", 10, True, GRAY, 0)])
label(s, L, 15.85, CW / 2 - 0.3, 1.0, [("観光バス：高速バスの予約後、乗車前日の案内をLINEで配信（琴平バス）", 10, False, INK, 0)], anchor=MSO_ANCHOR.MIDDLE)
label(s, L + CW / 2, 15.85, CW / 2, 1.0, [("皮膚科クリニック：友だち追加後のステップ配信で、カウンセリング予約へ（アクネクリニック）", 10, False, INK, 0)],
      anchor=MSO_ANCHOR.MIDDLE)
label(s, L, 17.0, CW, 0.5, [("※参考の2例はサンクスLINE誘導の成果ではなく、LINEでつながった後の活用例", 9, False, GRAY, 0)])

# ================= 5 DYMのコンサル支援内容・費用 =================
s = S[4]
keep_header_only(s)
head(s, "DYMのコンサル支援内容・費用", "ツール導入だけではなく、CPO改善に向けたLINE運用まで支援")
FY0, FH = 5.0, 2.4
xs = flow(s, L, R, FY0, FH, [(["① サンクスLINE導入"], GREEN, WHITE, 1), (["② LINE配信設計"], NAVY, WHITE, 1), (["③ 効果計測"], NAVY, WHITE, 1),
                               (["④ 改善"], NAVY, WHITE, 1)], size=13.5)
details = [["導線設計", "設定支援"],
           ["あいさつメッセージ", "ステップ配信", "リマインド", "セグメント配信"],
           ["友だち追加数", "LINE経由CV", "来店・面談・購入"],
           ["配信内容", "配信タイミング", "導線", "を継続改善"]]
DY = FY0 + FH + 0.5
for (x, w), items in zip(xs, details):
    for j, t in enumerate(items):
        if t.startswith("を"):
            label(s, x + 0.4, DY + j * 0.95, w - 1.0, 0.8, [(t, 10.5, False, GRAY, 0)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            continue
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x + 0.4, DY + j * 0.95, w - 1.0, 0.8, [(t, 10.5, True, INK, 0)], fill=WHITE, line=LGRAY, lw=0.75,
              adj=0.5, margins=(0.05, 0, 0.05, 0))
# ④→② を繰り返す（④の下から線を下ろし、②の下へ戻す）
BOT = DY + 4 * 0.95          # 詳細チップの下端
RY = BOT + 0.7
c2 = xs[1][0] + xs[1][1] / 2 - 0.1
c4 = xs[3][0] + xs[3][1] / 2 - 0.1
vline(s, c4, BOT, RY, NAVY, 1.5)
hline(s, c2, c4, RY, NAVY, 1.5)
arrow_line(s, c2, RY, c2, BOT + 0.05, NAVY, 1.5)
label(s, c2 + 0.3, RY + 0.05, 10.0, 0.6, [("計測結果をもとに、配信・導線を見直す", 10, False, GRAY, 0)], anchor=MSO_ANCHOR.MIDDLE)
# 費用の考え方（数字は2枚目にだけ置く）
PY5 = RY + 1.3
hline(s, L, R, PY5, LGRAY, 0.75)
label(s, L, PY5 + 0.3, 2.2, 1.8, [("費用", 13, True, NAVY, 0)], anchor=MSO_ANCHOR.MIDDLE)
rows = [("① サンクスLINE誘導ツール", "ツール費用（初期・月額）※2枚目に記載", GREEN, WHITE),
        ("②〜④ LINE運用コンサル", "配信本数・範囲に応じて別途お見積もり", NAVY, WHITE)]
rw = (CW - 2.2 - 0.6) / 2
for i, (a, b, f, c) in enumerate(rows):
    x = L + 2.2 + i * (rw + 0.6)
    rect(s, x, PY5 + 0.3, 5.8, 1.8, f, [(a, 11, True, c, 0)])
    label(s, x + 5.9, PY5 + 0.3, rw - 5.9, 1.8, [(b, 10.5, True, INK, 0)], anchor=MSO_ANCHOR.MIDDLE)

prs.save(str(OUT))
print("saved:", OUT.name, len(prs.slides), "枚")
