# -*- coding: utf-8 -*-
"""人材業界（派遣・バイト・単発）：入口ごとの14日ステップのページを足し、P24・P25の Day2・Day6 を差し替える（2026-10-22 殿村さん）。

    python -m _build.patch_haken_v4_entry_steps <PC保存版ver2.pptx> <出力.pptx>

  入力は殿村さんのPC保存版ver2（`_data/haken/PC保存版_20261022b_人材業界（派遣・バイト・単発）_ver2.pptx`・34枚）。
  ① P24（14日ステップの全体像）：Day2「相場を示す」（前後検索のデータに無いキーワード）→「比べ方を伝える」、
     Day6「相談先を伝える」→「しくみを伝える」（当日「派遣 3年ルール」）。リードの2行目もこの順に合わせる
  ② P25（実文面 Day0・Day6）：Day6の説明と、画面イメージの Day6 の吹き出し・ボタンを「派遣のしくみ」に
     （画像は `_build/make_haken_day6_image.py` で作った `_images/haken_day6_shikumi.png`）
  ③ P24の直後に「入口ごとの14日ステップ」を足す（34枚 → 35枚）。入口①〜⑤ × Day0・2・6・10・14 の5回で、応募（CV）までつなげる。
     ①は殿村さん指定の順（面接・見学の準備を先に）。図形で組む（表にしない）
"""
import sys

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn

from _build.patch_tanaka_v5 import copy_slide
from _build.pptx_parts import arrow_line, rect, rich, shape

NAVY, INK, SUB, LGRAY, CARD = "1E2761", "333333", "64748B", "D9D9D9", "F4F7FF"
IN = 2.54
IMAGE = "_images/haken_day6_shikumi.png"

P24 = [("相場を示す", "比べ方を伝える"), ("希望エリア・職種の時給の相場", "派遣会社の比べ方"),
       ("+2日「派遣 時給 相場」（仮説）", "同じ日に複数の会社を比べている"),
       ("相談先を伝える", "しくみを伝える"), ("合わなかったときの相談窓口", "派遣のしくみ（3年ルール・契約期間）"),
       ("「派遣 契約期間内 辞める」", "当日「派遣 3年ルール」"),
       ("前後検索で見えた「比較 → 相場 → 続けられるか」の順に、疑問に答える。",
        "前後検索で見えた「比較 → しくみ → 職場見学」の順に、疑問に答える。")]
P25 = [("「合わなかったら相談できる」と先に伝えることで、信頼につなげる。", "働き方のしくみを先に伝えることで、信頼につなげる。"),
       ("「合わなかったときの相談先」だけを伝える。", "3年ルール・契約期間のしくみだけを伝える。"),
       ("前後検索でも「契約期間内 辞める」が調べられている。", "前後検索でも「派遣 3年ルール」が調べられている。")]

TITLE = "［第4章］配信設計｜入口ごとの14日ステップ"
LEAD = ["どの入口から友だちになっても、Day0・2・6・10・14の5回で応募（CV）までつなげる。",
        "入口ごとに、届ける順番と中身を変える。"]
NOTE = "※ Day0〜14は友だち追加からの日数。①の応募した方には、面接・見学の準備を先にお届けする。中身は前後検索（P10・P11）で調べられている順に合わせる"
DAYS = ["Day0", "Day2", "Day6", "Day10", "Day14"]
SHIKUMI = "派遣のしくみ/（3年ルール・/契約期間）"
MENSETSU = "面接・見学の/流れと持ち物"
ROWS = [  # (入口, ゴール, [Day0, Day2, Day6, Day10, Day14])  「/」で改行
    ("① 応募完了画面", "面接・見学に来る",
     ["求人診断→/面接・見学/までの流れ", "面接・見学の/持ち物・服装", "合う求人/3件", SHIKUMI, "新着のまとめ/＋「条件を/変える」"]),
    ("② 離脱防止ポップアップ", "応募",
     ["求人診断→/合う求人3件", "派遣会社の/比べ方", SHIKUMI, MENSETSU, "新着のまとめ/＋応募のご案内"]),
    ("③ LINE広告（CPF）", "応募",
     ["できること→/求人診断→/合う求人3件", "登録から就業/までのサポート", SHIKUMI, MENSETSU, "人気の求人の/まとめ/＋応募のご案内"]),
    ("④ 求人誌・店頭のQR", "応募",
     ["求人診断→/掲載エリアの/求人3件", "掲載エリアの/新着", "扶養内・/掛け持ちで/働くときの目安", MENSETSU, "新着のまとめ/＋応募のご案内"]),
    ("⑤ 通知メッセージ", "再応募",
     ["求人診断/（次に探す/時期も）", "時期が近い方に/合う求人3件", "更新しないとき/の次の探し方", "前回と近い/職種の新着", "新着のまとめ/＋「条件を/変える」"]),
]


def ptext(p):
    return "".join(t.text or "" for t in p.iter(qn("a:t")))


def set_runs(p, text):
    runs = p.findall(qn("a:r"))
    for r in runs[1:]:
        p.remove(r)
    runs[0].find(qn("a:t")).text = text


def replace(s, pairs):
    ps = [p for sh in s.shapes if sh.has_text_frame for p in sh.text_frame._txBody.iter(qn("a:p"))]
    for old, new in pairs:
        hits = [p for p in ps if ptext(p) == old]
        assert len(hits) == 1, old
        set_runs(hits[0], new)


def by_name(s, name):
    return next(sh for sh in s.shapes if sh.name == name)


def set_paras(sh, texts):
    ps = sh.text_frame._txBody.findall(qn("a:p"))
    assert len(ps) == len(texts), (sh.name, len(ps))
    for p, t in zip(ps, texts):
        set_runs(p, t)


def build(s):
    keep = {"TextBox 1", "TextBox 2", "TextBox 41"}
    for sh in list(s.shapes):
        if sh.name not in keep:
            sh._element.getparent().remove(sh._element)
    set_paras(by_name(s, "TextBox 1"), [TITLE])
    set_paras(by_name(s, "TextBox 2"), LEAD)
    set_paras(by_name(s, "TextBox 41"), [NOTE])

    x0, w = 0.47, 9.89
    lw, gap, ag = 1.95, 0.1, 0.14                          # 入口の欄／入口と1列目の間／列の間（三角）
    cw = (w - lw - gap - ag * 4) / 5
    xs = [x0 + lw + gap + (cw + ag) * k for k in range(5)]
    hy, y0, rh, rg = 1.42, 1.86, 0.88, 0.08
    rich(s, x0 * IN, hy * IN, lw * IN, 0.3 * IN, [[("入口（ゴール）", 11, True, NAVY)]], align=PP_ALIGN.CENTER)
    for x, d in zip(xs, DAYS):
        rich(s, x * IN, hy * IN, cw * IN, 0.3 * IN, [[(d, 12, True, NAVY)]], align=PP_ALIGN.CENTER)
    arrow_line(s, xs[0] * IN, (hy + 0.33) * IN, (xs[-1] + cw) * IN, (hy + 0.33) * IN, NAVY, w=1.25)
    for r, (entry, goal, cells) in enumerate(ROWS):
        y = y0 + (rh + rg) * r
        rect(s, x0 * IN, y * IN, lw * IN, rh * IN, NAVY,
             [(entry, 11, True, "FFFFFF", 2), ("ゴール：" + goal, 9.5, False, "D6DBEA", 0)])
        for k, (x, text) in enumerate(zip(xs, cells)):
            first, last = k == 0, k == 4
            box = rect(s, x * IN, y * IN, cw * IN, rh * IN, CARD if first else "FFFFFF",
                       line=NAVY if (first or last) else LGRAY, lw=1.25 if last else 0.75)
            box._element.getparent().remove(box._element)        # 枠だけ先に置き、文字は別の箱で（行間をそろえるため）
            s.shapes._spTree.append(box._element)
            rich(s, (x + 0.03) * IN, y * IN, (cw - 0.06) * IN, rh * IN,
                 [[(t, 11.5, last, NAVY if last else INK)] for t in text.split("/")],
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            if k < 4:
                tri = shape(s, MSO_SHAPE.ISOSCELES_TRIANGLE, (x + cw + ag / 2 - 0.04) * IN, (y + rh / 2 - 0.07) * IN,
                            0.08 * IN, 0.14 * IN, fill=NAVY)
                tri.rotation = 90
    bottom = y0 + (rh + rg) * len(ROWS) - rg
    assert bottom <= 6.75, bottom
    return bottom


def main(src, out):
    prs = Presentation(src)
    slides = list(prs.slides)
    assert len(slides) == 34
    s24, s25 = slides[23], slides[24]
    assert by_name(s24, "TextBox 1").text_frame.text.endswith("14日ステップの全体像")
    replace(s24, P24)
    replace(s25, P25)
    pic = next(sh for sh in s25.shapes if sh.shape_type == 13)
    x, y, cx, cy = pic.left, pic.top, pic.width, pic.height
    new_pic = s25.shapes.add_picture(IMAGE, x, y, cx, cy)
    pic._element.addprevious(new_pic._element)                 # 元の画像と同じ重なり順に
    pic._element.getparent().remove(pic._element)

    new = copy_slide(prs, s24)
    lst = prs.slides._sldIdLst
    ids = list(lst)
    e = ids[-1]
    lst.remove(e)
    ids[23].addnext(e)
    bottom = build(new)
    assert len(prs.slides) == 35
    prs.save(out)
    print("saved:", out, round(bottom, 2))


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
