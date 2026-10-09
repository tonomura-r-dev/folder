# -*- coding: utf-8 -*-
"""展示会 LINE運用方針（テキストベース5枚・4:3）。文言は _drafts/展示会LINE_スライド要旨（2026-10-09）.md のとおり。
器はDYMの4:3マスターの既存PPTX（サンクスLINE Appendix 等）。ヘッダー（タイトル・リード・区切り線）だけ残して中身を入れる。
  python3 _build/build_tenjikai_line_5p.py <4:3のDYM資料.pptx> [出力.pptx]
"""
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Cm, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptx_parts import (GRAY, INK, NAVY, WHITE, arrow_line, find, keep_header_only, rect, rich,  # noqa: E402
                        set_paras, shape)

ROOT = Path(__file__).resolve().parent.parent
SRC = Path(sys.argv[1])
OUT = Path(sys.argv[2]) if len(sys.argv) > 2 else ROOT / "20261009_展示会LINE運用方針.pptx"

SLIDES = [
    ("課題｜同じ方に広告が出すぎている",
     ["少ない方に広告が重なり、効率が落ちている"]),
    ("方針｜エリア予算を見直し、浮いた予算をLINEへ",
     ["オーディエンスに対して予算が多いエリアを削り、新規予算10〜30万円をつくる"]),
    ("LINE広告（通常）｜Metaで届いていない方に届ける",
     ["媒体を変え、Metaとは違うユーザー層に展示会を案内する"]),
    ("サンクスLINE｜申し込んだ方を来場まで追う",
     ["申込完了画面からLINEへ案内し、前日・当日のお知らせで来場につなげる"]),
    ("友だち登録後の配信と仮の数値",
     ["今の友だちとサンクスLINEで増えた友だちへ、展示会ごとに配信で案内する"]),
]

X, W = 1.79, 23.97          # 本文の左端・幅（リード行と同じ左端）
Y0 = 4.5                    # 本文の開始位置（区切り線 3.86 の下）
B = 16                      # 本文の文字サイズ


def text(s, y, runs_by_para, h=1.0, x=X, w=W):
    return rich(s, x, y, w, h, runs_by_para, anchor=MSO_ANCHOR.TOP)


def table_rows(s, y, rows, kw=5.4, h=0.95, size=B):
    """左＝項目（紺・太字）、右＝内容。行ごとに同じ高さで並べる"""
    for i, (k, v) in enumerate(rows):
        text(s, y + i * h, [[(k, size, True, NAVY)]], h=h, w=kw)
        text(s, y + i * h, [[(v, size, False, INK)]], h=h, x=X + kw, w=W - kw)
    return y + len(rows) * h


prs = Presentation(str(SRC))
# ヘッダー一式（TextBox 1/2・Connector 3）がそろっているページを5枚だけ残す
lst = prs.slides._sldIdLst
keep = [i for i, s in enumerate(prs.slides)
        if {"TextBox 1", "TextBox 2", "Connector 3"} <= {sh.name for sh in s.shapes}][:5]
assert len(keep) == 5, keep
for i, el in reversed(list(enumerate(list(lst)))):
    if i not in keep:
        prs.part.drop_rel(el.rId)
        lst.remove(el)

for s, (title, lead) in zip(prs.slides, SLIDES):
    keep_header_only(s)
    t1, t2 = find(s, "TextBox 1"), find(s, "TextBox 2")
    t1.left, t1.top, t1.width, t1.height = Cm(1.52), Cm(0.38), Cm(18.1), Cm(0.94)
    t2.left, t2.top, t2.width, t2.height = Cm(1.79), Cm(2.27), Cm(25.31), Cm(0.96 if len(lead) == 1 else 1.5)
    set_paras(t1, [(title, 0)])
    set_paras(t2, [(l, 0) for l in lead])
    for r in t1.text_frame._txBody.iter(qn("a:rPr")):
        r.set("sz", "1600")
    for r in t2.text_frame._txBody.iter(qn("a:rPr")):
        r.set("sz", "1600")

s1, s2, s3, s4, s5 = prs.slides
S = 13  # 本文の文字サイズ


def head(s, y, txt):
    """見出し（紺の縦棒＋紺の太字）"""
    rect(s, X, y + 0.14, 0.15, 0.5, NAVY)
    text(s, y, [[(txt, 13.5, True, NAVY)]], h=0.8, x=X + 0.35, w=W - 0.35)
    return y + 0.85


def rows(s, y, items, kw=4.6, h=0.72, size=S):
    for i, it in enumerate(items):
        k, v = it[0], it[1]
        text(s, y + i * h, [[(k, size, True, INK)]], h=h, x=X + 0.35, w=kw)
        text(s, y + i * h, [[(v, size, False, INK)]], h=h, x=X + 0.35 + kw, w=W - 0.35 - kw)
    return y + len(items) * h


def rows3(s, y, items, w1=3.2, w2=5.6, h=0.72, size=S):
    for i, (a, b, c) in enumerate(items):
        text(s, y + i * h, [[(a, size, True, INK)]], h=h, x=X + 0.35, w=w1)
        text(s, y + i * h, [[(b, size, False, INK)]], h=h, x=X + 0.35 + w1, w=w2)
        text(s, y + i * h, [[(c, 12, False, GRAY)]], h=h, x=X + 0.35 + w1 + w2, w=W - 0.35 - w1 - w2)
    return y + len(items) * h


def flow(s, y, items, h=1.35, gap=0.75, size=11.5):
    """横一列の流れ図。最後の箱だけ紺の塗り"""
    x0, ww = X + 0.35, W - 0.35
    bw = (ww - gap * (len(items) - 1)) / len(items)
    for i, t in enumerate(items):
        last = i == len(items) - 1
        x = x0 + i * (bw + gap)
        paras = [(ln, size, last, WHITE if last else INK, 0) for ln in t.split("\n")]
        shape(s, MSO_SHAPE.RECTANGLE, x, y, bw, h, paras, fill=NAVY if last else WHITE, line=NAVY, lw=1.0,
              margins=(0.1, 0.05, 0.1, 0.05))
        if not last:
            arrow_line(s, x + bw + 0.1, y + h / 2, x + bw + gap - 0.1, y + h / 2, NAVY, w=1.5)
    return y + h


# 1. 課題
y = head(s1, Y0, "現状の数値")
y = rows(s1, y, [("フリークエンシー", "5.24（同じ方が月に5回以上、同じ広告を見ている）"),
                 ("CPA", "15,000円以上（予約1件あたりの広告費）")])
y = head(s1, y + 0.5, "なぜ効率が落ちるのか")
y = flow(s1, y + 0.1, ["オーディエンスが少ない", "予算を使い切るため、\n同じ方に繰り返し表示", "予約は増えず、\nCPAだけ上がる"])
y = head(s1, y + 0.6, "オーディエンスの少ないエリア（例）")
rows(s1, y, [("三鷹", "10,700〜12,600"), ("高松", "上限25,000"), ("日向", "4,400")])

# 2. 方針
y = head(s2, Y0, "見直しの基準")
y = rows(s2, y, [("削る順番", "オーディエンス数に対して予算が多いエリアから削る"),
                 ("例", "銀座店 30万円 → 20万円")])
y = head(s2, y + 0.5, "予算の流れ")
y = flow(s2, y + 0.1, ["Meta広告\n（各エリアの予算）", "見直しで浮いた予算\n（全体で10〜30万円）", "LINE広告（通常）\n20万円"])
y = head(s2, y + 0.6, "費用について")
rows(s2, y, [("広告費の総額", "変わらない（今のMeta広告の予算の中で振り替える）"),
             ("LINE広告", "20万円は全店舗の合計")])

# 3. LINE広告（通常）
y = head(s3, Y0, "LINEの特徴")
y = rows(s3, y, [("利用者数", "国内の月間利用者数 1億人以上（2025年12月時点・LINEヤフー公表）"),
                 ("届く方", "Instagram・Facebookをあまり使わない方にも届く"),
                 ("絞り込み", "地域・年齢・性別で配信先を絞れる")])
y = head(s3, y + 0.5, "仮の数値（月）")
flow(s3, y + 0.1, ["予算\n20万円", "クリック 2,000\n（CPC 100円）", "予約 20件\n（予約率 1.0%）", "予約1件あたり\n10,000円"])

# 4. サンクスLINE
y = head(s4, Y0, "流れ")
y = flow(s4, y + 0.1, ["申込フォーム", "申込完了画面", "LINE\n友だち追加", "前日・当日の\nお知らせ", "来場"],
         h=1.5, gap=0.6)
y = head(s4, y + 0.6, "仮の数値（申込100件あたり）")
y = rows(s4, y, [("LINE友だち", "30人（友だち登録率 30%）"),
                 ("来場率", "LINEでつながった方 50% → 70%"),
                 ("来場", "50人 → 56人")])
y = head(s4, y + 0.5, "実績・費用")
rows(s4, y, [("実績", "弊社運用の美容クリニックで、予約後の来院率 40〜50%改善"),
             ("費用", "初期15万円・月額3万円〜")])

# 5. 配信と仮の数値
y = head(s5, Y0, "配信のタイミングと内容")
y = rows3(s5, y, [("登録直後", "概要・予約", "「〇月〇日から〇〇で展示会を開催します。ご予約はこちら」"),
                  ("2週間前", "見どころ", "「今回の見どころをご紹介します」"),
                  ("3日前", "リマインド・アクセス", "「ご来場まであと3日です。会場へのアクセスはこちら」"),
                  ("当日朝", "開催のお知らせ", "「本日開催です。お気をつけてお越しください」"),
                  ("終了後", "お礼・次回案内", "「ご来場ありがとうございました。次回もご案内します」")])
y = head(s5, y + 0.5, "仮の数値（月）")
y = rows(s5, y, [("予約", "20件（LINE広告）"),
                 ("費用", "23万円（LINE広告 20万円＋サンクスLINE 3万円）"),
                 ("予約1件あたり", "11,500円（現状：15,000円以上）")])

# 字体をメイリオに統一（latin → ea → cs の順）
AFTER = {qn(x) for x in ("a:sym", "a:hlinkClick", "a:hlinkMouseOver", "a:rtl", "a:extLst")}
for s in prs.slides:
    for r in s.shapes._spTree.iter(qn("a:r")):
        if r.find(qn("a:rPr")) is None:
            r.insert(0, r.makeelement(qn("a:rPr"), {"lang": "ja-JP"}))
    for tag in ("a:rPr", "a:endParaRPr", "a:defRPr"):
        for rpr in s.shapes._spTree.iter(qn(tag)):
            for ft in ("a:latin", "a:ea", "a:cs"):
                for e in rpr.findall(qn(ft)):
                    rpr.remove(e)
            nxt = next((c for c in rpr if c.tag in AFTER), None)
            for ft in ("a:latin", "a:ea", "a:cs"):
                e = rpr.makeelement(qn(ft), {"typeface": "メイリオ"})
                if nxt is not None:
                    nxt.addprevious(e)
                else:
                    rpr.append(e)

prs.save(str(OUT))
print("saved:", OUT.name, len(prs.slides), "枚")
