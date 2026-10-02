# -*- coding: utf-8 -*-
"""チェングロウス ver3.0 → ver3.1（2026-10-02）
確定した進め方：11月から構築（初期費用のみ）／4月のリニューアルと同時に運用を本格化（月額の固定費）。
11月〜3月は友だちを貯める時期（数字はシミュレーションに含めない）。サンクスLINEは自社サイトの応募のみ。
求人ボックス・Indeed経由の応募者への導線を図で追加（15枚目の次）。費用と体制（25枚目）は触らない。
  python3 _build/patch_chengrowth_v31_schedule.py <ver3.0.pptx>
"""
import copy
import sys
from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.util import Cm, Pt

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "20260929_チェングロウス_LINE公式アカウント運用のご提案ver3.1.pptx"
prs = Presentation(sys.argv[1])
NAVY, INK, WHITE, GRAY = RGBColor(0x1F, 0x28, 0x5A), RGBColor(0x33, 0x33, 0x33), RGBColor(0xFF, 0xFF, 0xFF), RGBColor(0x7F, 0x7F, 0x7F)
LINEG, LIGHT, LGREEN = RGBColor(0x06, 0xC7, 0x55), RGBColor(0xE6, 0xE9, 0xF0), RGBColor(0xE8, 0xF8, 0xEE)


def walk(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type == 6:
            yield from walk(sh.shapes)


def shape(n, name):
    return next(sh for sh in walk(prs.slides[n - 1].shapes) if sh.name == name)


def set_par(n, name, k, text):
    p = shape(n, name).text_frame.paragraphs[k]
    runs = p.runs
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


# ---------- 文言の直し（番号はver3.0のもの）----------
set_par(1, "正方形/長方形 5", 0, "―LINEの初期構築と、サイトリニューアルでの連携のご提案―")
set_par(11, "Rounded Rectangle 13", 0, "LINE側は先に構築し、サイト側はリニューアルに合わせてつなぐ")
set_par(12, "TextBox 2", 0, "サイトリニューアルに合わせて、会員登録をLINEで完結できる形へ")
set_par(12, "Rounded Rectangle 7", 0, "リニューアル後")
set_par(15, "TextBox 13", 0,
        "※月約17人＝サイトからの応募（月約47件）×完了画面の表示80%×友だち追加45%の想定／求人ボックス・Indeed経由の応募は完了画面を触れないため対象外／画面はイメージです")
# 16 注記・24 シミュレーション：4月開始（4〜9月）に戻す
for p in shape(16, "TextBox 8").text_frame.paragraphs:
    for r in p.runs:
        r.text = r.text.replace("6か月目の値", "9月の値")
set_par(24, "TextBox 1", 0, "成果シミュレーション（4〜9月）")
set_par(24, "TextBox 2", 1, "友だちが増えるほど応募も増え、7月に応募単価が2.5万円を下回る見込み")
set_par(24, "Rounded Rectangle 7", 0, "9月の応募単価")
tbl = shape(24, "Table 4").table
for j in range(1, 7):
    runs = tbl.cell(0, j).text_frame.paragraphs[0].runs
    assert runs[0].text == f"{j}か月目", runs[0].text
    runs[0].text = f"{j + 3}月"
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
for p in shape(24, "TextBox 8").text_frame.paragraphs:
    for r in p.runs:
        r.text = r.text.replace("前提：会員登録のLINE化の開始月を1か月目とする／", "前提：11〜3月に貯まる友だちは含めない／")
# 20 先に進める理由
set_par(20, "Rounded Rectangle 4", 3, "半年で約540人の友だちに")
set_par(20, "TextBox 9", 0, "11月 ────── 友だちが増え続ける ────── 翌3月（転職のピーク）")
set_par(20, "Rounded Rectangle 14", 0, "11月から構築を始め、4月のリニューアルと同時に運用を本格化することをご提案します。")
# 21 要件
set_par(21, "TextBox 1", 0, "リニューアルに入れる要件3点")
set_par(21, "Rounded Rectangle 10", 2, "サイト制作会社様と仕様をすり合わせて進めます。")
# 22 スケジュール
set_par(22, "TextBox 1", 0, "スケジュール｜11月から構築、4月に運用を本格化")
set_par(22, "Rounded Rectangle 4", 0, "11月〜")
set_par(22, "Rounded Rectangle 5", 0, "お申込み・申請")
set_par(22, "Rounded Rectangle 5", 1, "アカウント設定、LINE Profile+の申請、サイト・配信ツールとLINEをつなぐ仕様のすり合わせ")
set_par(22, "Rounded Rectangle 6", 0, "11〜3月")
set_par(22, "Rounded Rectangle 7", 0, "初期構築（この期間は初期費用のみ）")
set_par(22, "Rounded Rectangle 7", 1, "あいさつ・リッチメニュー・配信の設計、各ツールとLINEの連携、離脱防止ポップアップの設置")
set_par(22, "Rounded Rectangle 8", 0, "11〜3月")
set_par(22, "Rounded Rectangle 9", 0, "友だちを貯める")
set_par(22, "Rounded Rectangle 9", 1, "離脱防止ポップアップなどで友だちを増やす（この期間の数字はシミュレーションに含めない）")
set_par(22, "Rounded Rectangle 10", 0, "4月〜")
set_par(22, "Rounded Rectangle 11", 0, "リニューアルと同時に運用を本格化（月額の固定費）")
set_par(22, "Rounded Rectangle 11", 1, "「LINEで登録」を開始。計測と配信を本格化")
set_par(22, "TextBox 14", 0, "※LINE Profile+の審査期間は、申請内容により変わります／求人ボックス・Indeed経由の応募者への導線は、着地先を確認のうえ決定")

# ---------- 新スライド：求人ボックス・Indeed経由の導線（15枚目の次）----------
src = prs.slides[6]          # 7枚目（広告とLINEの流れ）の部品を借りる
parts = {sh.name: sh for sh in src.shapes}
T = {k: copy.deepcopy(parts[k]._element) for k in
     ("TextBox 1", "TextBox 2", "Connector 3", "TextBox 4", "Rounded Rectangle 5", "Right Arrow 6", "Rounded Rectangle 24", "TextBox 25")}
new = prs.slides.add_slide(prs.slides[14].slide_layout)
for sh in list(new.shapes):
    sh._element.getparent().remove(sh._element)


def set_runs(sh, lines, align=None):
    """lines=[(text,size,bold,color)]。書式は最初のrunを引き継ぐ"""
    tf = sh.text_frame
    base_p = tf.paragraphs[0]._p
    base_r = base_p.find(qn("a:r"))
    rpr0 = copy.deepcopy(base_r.find(qn("a:rPr"))) if base_r is not None else None
    for p in list(tf._txBody.findall(qn("a:p")))[1:]:
        tf._txBody.remove(p)
    for r in list(base_p.findall(qn("a:r"))) + list(base_p.findall(qn("a:br"))):
        base_p.remove(r)
    empty_p = copy.deepcopy(base_p)
    for i, (text, size, bold, color) in enumerate(lines):
        p = base_p if i == 0 else copy.deepcopy(empty_p)
        if i:
            tf._txBody.append(p)
        r = p.makeelement(qn("a:r"), {})
        if rpr0 is not None:
            r.append(copy.deepcopy(rpr0))
        t = r.makeelement(qn("a:t"), {})
        t.text = text
        r.append(t)
        end = p.find(qn("a:endParaRPr"))
        (end.addprevious(r) if end is not None else p.append(r))
    for para, (text, size, bold, color) in zip(tf.paragraphs, lines):
        if align is not None:
            para.alignment = align
        for run in para.runs:
            run.font.size, run.font.bold = Pt(size), bold
            run.font.color.rgb = color


def add(key, x, y, w, h, lines=None, fill=None, align=None):
    el = copy.deepcopy(T[key])
    new.shapes._spTree.append(el)
    sh = new.shapes[-1]
    sh.left, sh.top, sh.width, sh.height = Cm(x), Cm(y), Cm(w), Cm(h)
    if fill is not None:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if lines:
        set_runs(sh, lines, align)
    return sh


for k in ("TextBox 1", "TextBox 2", "Connector 3"):
    new.shapes._spTree.append(copy.deepcopy(T[k]))
t1, t2 = new.shapes[0], new.shapes[1]
set_runs(t1, [("求人ボックス・Indeed経由の応募者への導線", 16, True, RGBColor(0x00, 0x20, 0x60))])
set_runs(t2, [("サンクスLINEは自社サイトの応募が対象。外部経由は、求人ページ上の導線でつなぐ", 16, False, INK)])
t2.top, t2.height = Cm(2.4), Cm(0.8)


def box(x, y, w, h, text, fill, color):
    return add("Rounded Rectangle 5", x, y, w, h, [(tx, 13, True, color) for tx in text.split("\n")], fill, 2)


def arrow(x, y, h, color):
    a = add("Right Arrow 6", x, y + h / 2 - 0.3, 0.7, 0.6)
    a.fill.solid()
    a.fill.fore_color.rgb = color


def label(x, y, text):
    add("TextBox 4", x, y, 16.0, 0.8, [(text, 14, True, NAVY)])


# レーンA：自社サイト
label(2.2, 4.4, "■ 自社サイトから応募する方")
YA, HA = 5.3, 1.5
names = [("求人ページ", LIGHT, INK, 3.4), ("応募フォーム", LIGHT, INK, 3.4), ("完了画面", LIGHT, INK, 3.4),
         ("サンクスLINE\n（オプション）", RGBColor(0x0B, 0x7A, 0x3B), WHITE, 5.2), ("LINE追加", LINEG, WHITE, 3.6)]
x = 2.2
for i, (tx, f, c, w) in enumerate(names):
    box(x, YA, w, HA, tx, f, c)
    x += w
    if i < 4:
        arrow(x + 0.15, YA, HA, GRAY)
        x += 1.0

# レーンB：求人ボックス・Indeed
label(2.2, 7.7, "■ 求人ボックス・Indeedから来る方")
Y1, Y2, HB = 8.7, 11.0, 1.7
box(2.2, Y1, 4.2, Y2 + HB - Y1, "求人ボックス\nIndeed", LIGHT, INK)
arrow(6.5, Y1, HB, GRAY)
arrow(6.5, Y2, HB, GRAY)
box(7.4, Y1, 5.2, HB, "貴社の求人ページに\n着地する場合", LIGHT, INK)
arrow(12.75, Y1, HB, GRAY)
box(13.7, Y1, 5.6, HB, "離脱防止ポップアップ\n「LINEで登録」", NAVY, WHITE)
arrow(19.45, Y1, HB, GRAY)
box(20.4, Y1, 4.9, HB, "LINE追加", LINEG, WHITE)
box(7.4, Y2, 5.2, HB, "応募が先方の画面で\n完結する場合", LIGHT, INK)
arrow(12.75, Y2, HB, GRAY)
box(13.7, Y2, 5.6, HB, "貴社の画面を通らない", LIGHT, INK)
arrow(19.45, Y2, HB, GRAY)
box(20.4, Y2, 4.9, HB, "導線を置けない\n（対象外）", GRAY, WHITE)
add("Rounded Rectangle 24", 2.2, 13.8, 23.1, 1.2,
    [("貴社の求人ページに来た方は、離脱防止と「LINEで登録」で友だち追加へつなげる", 15, True, WHITE)], NAVY, 2)
add("TextBox 25", 2.2, 16.8, 23.1, 0.8,
    [("※求人ボックス・Indeedからの着地先と、応募の完結先は、掲載方法により異なるため、確認のうえ決定", 10, False, GRAY)])

# 15枚目の次へ移動
lst = prs.slides._sldIdLst
item = lst[-1]
lst.remove(item)
lst.insert(15, item)

for sl in prs.slides:
    for n, c in enumerate(sl.shapes._spTree.iter(qn("p:cNvPr")), start=2):
        c.set("id", str(n))

# 字体をメイリオに統一（latin→ea→cs）
AFTER = {qn(x) for x in ("a:sym", "a:hlinkClick", "a:hlinkMouseOver", "a:rtl", "a:extLst")}
for sl in prs.slides:
    for r in sl.shapes._spTree.iter(qn("a:r")):
        if r.find(qn("a:rPr")) is None:
            r.insert(0, r.makeelement(qn("a:rPr"), {"lang": "ja-JP"}))
    for tag in ("a:rPr", "a:endParaRPr", "a:defRPr"):
        for rpr in sl.shapes._spTree.iter(qn(tag)):
            fs = []
            for ft in ("a:latin", "a:ea", "a:cs"):
                e = rpr.find(qn(ft))
                if e is None:
                    e = rpr.makeelement(qn(ft), {})
                else:
                    rpr.remove(e)
                e.set("typeface", "メイリオ")
                fs.append(e)
            nxt = next((c for c in rpr if c.tag in AFTER), None)
            for e in fs:
                (nxt.addprevious(e) if nxt is not None else rpr.append(e))
prs.save(str(OUT))
print("saved:", OUT.name, len(prs.slides), "枚")
