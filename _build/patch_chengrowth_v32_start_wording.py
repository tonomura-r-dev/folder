# -*- coding: utf-8 -*-
"""チェングロウス ver3.1 → ver3.2：「11月」は確定ではないので、開始時期の書き方を「ご契約後」に直す（2026-10-02）
  python3 _build/patch_chengrowth_v32_start_wording.py <ver3.1.pptx>
"""
import sys
from pathlib import Path
from pptx import Presentation

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "20260929_チェングロウス_LINE公式アカウント運用のご提案ver3.2.pptx"
prs = Presentation(sys.argv[1])


def shape(n, name):
    return next(sh for sh in prs.slides[n - 1].shapes if sh.name == name)


def set_par(n, name, k, text):
    p = shape(n, name).text_frame.paragraphs[k]
    runs = p.runs
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def sub(n, name, old, new):
    for p in shape(n, name).text_frame.paragraphs:
        for r in p.runs:
            if old in r.text:
                r.text = r.text.replace(old, new)
                return
    raise SystemExit(f"not found: {n} {name} {old}")


set_par(21, "TextBox 9", 0, "開始 ────── 友だちが増え続ける ────── 翌3月（転職のピーク）")
set_par(21, "Rounded Rectangle 14", 0, "ご契約後すぐに構築を始め、4月のリニューアルと同時に運用を本格化することをご提案します。")
set_par(23, "TextBox 1", 0, "スケジュール｜ご契約後に構築、4月に運用を本格化")
set_par(23, "Rounded Rectangle 4", 0, "ご契約後")
set_par(23, "Rounded Rectangle 6", 0, "〜3月")
set_par(23, "Rounded Rectangle 8", 0, "〜3月")
sub(25, "TextBox 8", "11〜3月に貯まる友だち", "リニューアル前に貯まる友だち")
prs.save(str(OUT))
print("saved:", OUT.name, len(prs.slides), "枚")
