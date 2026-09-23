# -*- coding: utf-8 -*-
"""学校法人業界_LINEOA施策提案（50枚・B版）に実データを反映して ver3 を作る。

ベース：2026-09-15/16作成の50枚版（別セッション作・黄マーカー付き）
反映（ラリー2）：
- P9  現状分析表：友だち数未実測 → 大手8校の実測（page.line.me公式・2026-09-20）
- P15 市場の構造：定員割れ59.2%→令和8年度速報46.2%／納付金136万→150.8万・4年400万→約470万
- P20 前後検索補足：専門学校の新発見（奨学金が起点7日前にも立つ）を追記
- P3  注記の「4年間で約400万円」→「約470万円」

  python3 _build/patch_gakko_ver3.py
"""
import shutil
from pathlib import Path
from pptx import Presentation

ROOT = Path(__file__).resolve().parent.parent
SRC = "/tmp/claude-0/-home-user-folder/dadd2f6b-abde-5965-beb8-93809cb827e0/scratchpad/cand_B.pptx"
OUT = str(ROOT / "学校法人業界_LINEOA施策提案.pptx")

shutil.copyfile(SRC, OUT)
prs = Presentation(OUT)
slides = list(prs.slides)
assert len(slides) == 50, len(slides)


def walk_replace(shape, mapping):
    if shape.shape_type == 6:
        for c in shape.shapes:
            walk_replace(c, mapping)
        return
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            for cell in row.cells:
                for para in cell.text_frame.paragraphs:
                    for r in para.runs:
                        for k, v in mapping.items():
                            if k in r.text:
                                r.text = r.text.replace(k, v)
        return
    if shape.has_text_frame:
        for para in shape.text_frame.paragraphs:
            for r in para.runs:
                for k, v in mapping.items():
                    if k in r.text:
                        r.text = r.text.replace(k, v)


def replace_on(slide, mapping):
    for sh in slide.shapes:
        walk_replace(sh, mapping)


def set_cell(cell, text):
    """セルの先頭runの書式を保ったままテキスト差し替え"""
    tf = cell.text_frame
    runs = [r for p in tf.paragraphs for r in p.runs]
    if runs:
        runs[0].text = text
        for r in runs[1:]:
            r.text = ""
        # 2段落目以降は空に
        for p in tf.paragraphs[1:]:
            for r in p.runs:
                r.text = ""
    else:
        tf.paragraphs[0].add_run().text = text


# ============ P3 注記（4年間で約400万円 → 約470万円） ============
replace_on(slides[2], {"1名の入学＝4年間で約400万円": "1名の入学＝4年間で約470万円"})

# ============ P9 現状分析表：実測8校を充当 ============
s = slides[8]
ROWS = [
    ["立命館大学 入学センター", "大学", "○", "139,406", "資料請求・OC申込", "―"],
    ["近畿大学", "大学", "○", "137,996", "イベント・クーポン", "―"],
    ["龍谷大学 入試部", "大学", "○", "95,199", "OC予約・入試対策講座", "―"],
    ["日本工学院", "専門", "○", "42,075", "OC・体験入学", "未確認"],
    ["NSGカレッジリーグ", "専門", "○", "34,659", "診断・資料請求", "○"],
    ["大原学園", "専門", "○", "24,428", "個別相談・AO連絡", "―"],
    ["HAL東京", "専門", "○", "14,661", "OC情報・チャット", "―"],
    ["大阪モード学園", "専門", "○", "3,206", "OC情報", "―"],
]
for sh in s.shapes:
    if getattr(sh, "has_table", False) and sh.has_table:
        tbl = sh.table
        for i, row_vals in enumerate(ROWS, start=1):
            for j, v in enumerate(row_vals):
                if j < len(tbl.columns):
                    set_cell(tbl.cell(i, j), v)
replace_on(s, {
    "友だち数は page.line.me（LINEヤフー社公式ページ）で実測する。取得日を必ず併記する。":
        "友だち数は page.line.me（LINEヤフー社公式ページ）で実測済み（2026-09-20取得）。",
    "※ 友だち数は未実測（2026-09-16時点）。page.line.me（公式）または実機で取得し、取得日を併記する。":
        "※ 友だち数は 2026-09-20 に page.line.me（公式）で実測。機能はWeb公開情報から（リッチメニュー内部は外部から未確認のため断定しない）。",
})

# ============ P15 市場の構造：最新値へ ============
s = slides[14]
replace_on(s, {
    "市場の構造｜18歳人口は30年で半減。私立大の6割が定員割れ":
        "市場の構造｜18歳人口は30年で半減。私立大の約半数が定員割れ",
    "59.2%（354校／598校）": "46.2%（275校／595校）",
    "（2024年度）": "（令和8年度・速報値）",
    "6割が定員を埋められない": "約半数が定員を埋められない",
    "初年度納付金 約136万円": "初年度納付金 約150.8万円",
    "4年間で 約400万円": "4年間で 約470万円",
    "2040年の18歳人口は80万人台まで減る見通し（中央教育審議会の推計）。":
        "2040年の18歳人口は約80万人まで減る見通し（国推計）。2026年の約111万人は一時的な踊り場で、以降は本格減少。",
})
# 注記に速報の補足を追記
for sh in s.shapes:
    if sh.has_text_frame and "原典照合" in sh.text_frame.text:
        for para in sh.text_frame.paragraphs:
            for r in para.runs:
                if "原典照合すること】" in r.text:
                    r.text = r.text.replace(
                        "原典照合すること】",
                        "原典照合すること】※定員割れ46.2%は令和8年度速報（前年53.2%・定員削減等の影響含む）",
                    )

# ============ P20 前後検索補足：専門学校の新発見 ============
s = slides[19]
replace_on(s, {
    "前後検索（補足）｜学費・奨学金の検索は、決めた「後」に来る":
        "前後検索（補足）｜学費の検索は「後」に来る。ただし専門学校は「前」にも来る",
    "教育3業界（塾・通信制高校・大学）の前後検索で、同じ構造が確認できた。":
        "教育4業界（塾・通信制高校・大学・専門学校）の前後検索で確認。ただし専門学校は例外がある。",
    "＝ 入口の訴求を「安さ・学費」に置くと、届く前に終わる。":
        "＝ 専門学校は「奨学金」が起点の7日前にも最大級で立つ。お金の不安が検討を前後から挟む。",
    "（起点KW確認中）": "大学名・偏差値系",
})
replace_on(s, {
    "※ 大学は起点KWが未確認。確定後に本ページの数値を更新する。":
        "※ 専門学校・柔道整復師ほか8KWを2026-09-23に追加取得。専門学校は奨学金・教育ローンが起点前にも立つ（詳細は分析メモ）。",
})

# ============ P20を表示に戻す（起点KW確認待ちで非表示だったが、データ受領済み） ============
sl20 = slides[19]._element
if sl20.get("show") == "0":
    del sl20.attrib["show"]

prs.save(OUT)
print("saved:", OUT, f"（{len(Presentation(OUT).slides)}枚）")
