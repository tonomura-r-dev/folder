# -*- coding: utf-8 -*-
"""人材業界（派遣・バイト・単発）：入口①〜⑤の14日ステップを1入口1枚で5枚足す（2026-10-22 殿村さん）。

    python -m _build.patch_haken_v5_entry_pages <殿村さんのPC保存版ver3.pptx> <出力.pptx>

  入力は殿村さんのPC保存版ver3（`_data/haken/PC保存版_20261022_人材業界（派遣・バイト・単発）_ver3.pptx`・35枚）。
  P25「入口ごとの14日ステップ」を複製して、その直後に5枚（P26〜P30）を入れる。ほかのページは変えない。
  1枚の中身：上の段＝入口／登録する方／ゴール（CV）、中央＝Day0→Day2→Day6→Day10→Day14→CV（図形と矢印）、下＝注記1行。
  5枚とも部品の位置と大きさは同じ（同じ関数で置く）。色は紺・白・黒・灰色だけ。
"""
import sys

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Pt

from _build.patch_tanaka_v5 import copy_slide
from _build.pptx_parts import rect, rich, shape, style_run

NAVY, INK, GRAY, LGRAY, WHITE = "1E2761", "333333", "7F7F7F", "D9D9D9", "FFFFFF"
IN = 2.54
NOTE = "※ Day0〜14は友だち追加からの日数。前後検索の根拠はP10・P11"
DAYS = ["Day0", "Day2", "Day6", "Day10", "Day14"]

PAGES = [  # (タイトル, リード, 入口, 登録する方, ゴール, [(配信内容, ねらい, ボタン)×5], CV)
    ("［第4章］配信設計｜入口①　応募完了画面からの14日ステップ",
     "応募した方には、面接・職場見学の準備を先にお届けし、来ていただくまでつなげる。",
     "応募完了画面（サンクスLINE誘導）", "応募した方", "面接・職場見学に来る",
     [("求人診断 → 面接・見学までの流れ", "応募直後の不安を減らし、日程を確かめていただく", "日程を確認する"),
      ("面接・見学の持ち物・服装", "当日に迷わず来ていただく", "持ち物を確認する"),
      ("ほかにも合う求人3件", "比べている方に、選べる求人を残す", "求人を見る"),
      ("派遣のしくみ（3年ルール・契約期間）", "働く前の疑問に答え、辞退を防ぐ", "しくみを見る"),
      ("新着のまとめ＋「条件を変える」", "条件が合わない方も、次の応募につなげる", "条件を変える")],
     "面接・職場見学に来る"),
    ("［第4章］配信設計｜入口②　離脱防止ポップアップからの14日ステップ",
     "LPを離れた方には、比べるときの疑問に順に答え、14日目に応募をご案内する。",
     "離脱防止ポップアップ", "応募せずにLPを離れた方", "応募",
     [("求人診断 → 合う求人3件", "関心が高いうちに、合う求人を届ける", "求人を見る"),
      ("派遣会社の比べ方", "同じ日に複数の会社を比べている方に、選ぶ基準を渡す", "比べ方を見る"),
      ("派遣のしくみ（求人は送らない回）", "働き方の疑問に答え、信頼につなげる", "しくみを見る"),
      ("面接・見学の流れと持ち物", "応募の後の流れを見せ、応募しやすくする", "流れを見る"),
      ("新着のまとめ", "関心が高いうちに応募を後押しする", "応募する")],
     "応募"),
    ("［第4章］配信設計｜入口③　LINE広告からの14日ステップ",
     "まだ貴社を知らない方には、できることとサポートを先に伝え、14日目に応募をご案内する。",
     "LINE広告（CPF）", "まだ貴社を知らない方", "応募",
     [("できることの紹介 → 求人診断 → 合う求人3件", "友だちでいる理由を最初に伝える", "求人を見る"),
      ("登録から就業までのサポート", "貴社を選ぶ理由を伝える", "サポートを見る"),
      ("派遣のしくみ（3年ルール・契約期間）", "働き方の疑問に答える", "しくみを見る"),
      ("面接・見学の流れと持ち物", "応募の後の流れを見せる", "流れを見る"),
      ("人気の求人のまとめ", "応募を後押しする", "応募する")],
     "応募"),
    ("［第4章］配信設計｜入口④　求人誌・店頭のQRコードからの14日ステップ",
     "求人誌・店頭で知った方には、掲載エリアの求人を中心に届け、14日目に応募をご案内する。",
     "求人誌・店頭のQRコード", "求人誌・店頭で貴社を知った方", "応募",
     [("求人診断（エリアは掲載地域から）→ 掲載エリアの求人3件", "見ていた地域の求人から届ける", "求人を見る"),
      ("掲載エリアの新着", "近くの求人を続けて届ける", "新着を見る"),
      ("扶養内・掛け持ちで働くときの目安", "働く条件の疑問に答える", "目安を見る"),
      ("面接・見学の流れと持ち物", "応募の後の流れを見せる", "流れを見る"),
      ("新着のまとめ", "応募を後押しする", "応募する")],
     "応募"),
    ("［第4章］配信設計｜入口⑤　通知メッセージからの14日ステップ",
     "過去に応募した方には、次に探す時期を伺い、時期に合わせて再応募をご案内する。",
     "通知メッセージ（既存リスト）", "過去に応募した方", "再応募",
     [("求人診断（次に探す時期も伺う）", "ご案内の時期を、探し始める時期に合わせる", "診断に答える"),
      ("時期が近い方に合う求人3件", "探し始めた方に、他社より先に届ける", "求人を見る"),
      ("更新しないときの次の探し方", "契約満了前の疑問に答える", "探し方を見る"),
      ("前回と近い職種の新着", "経験を生かせる求人を届ける", "新着を見る"),
      ("新着のまとめ＋「条件を変える」", "条件を見直して、再応募につなげる", "応募する")],
     "再応募"),
]

# 位置（インチ）。5枚とも同じ値を使う
X0, W = 0.47, 9.89
BAND_Y, BAND_H, BAND_GAP = 1.85, 0.62, 0.12
DAY_Y, DAY_H = 2.67, 0.32
CARD_Y, CARD_H = 3.03, 2.28
COLS, AG = 6, 0.2
CW = (W - AG * (COLS - 1)) / COLS
TEXT_H, AIM_Y, AIM_H = 1.0, 1.1, 0.62        # カードの上端からの位置
BTN_W, BTN_H, BTN_B = CW - 0.14, 0.36, 0.12         # ボタンの幅・高さ・下端からの余白


def by_name(s, name):
    return next(sh for sh in s.shapes if sh.name == name)


def set_one(sh, text, size=None):
    """1段落にして、最初の文字の書式で文言だけ差し替える"""
    txb = sh.text_frame._txBody
    ps = txb.findall(qn("a:p"))
    for p in ps[1:]:
        txb.remove(p)
    runs = ps[0].findall(qn("a:r"))
    for r in runs[1:]:
        ps[0].remove(r)
    runs[0].find(qn("a:t")).text = text
    if size:
        runs[0].get_or_add_rPr().set("sz", str(int(size * 100)))


def c(v):
    return v * IN


def build(s, page):
    title, lead, entry, who, goal, steps, cv = page
    keep = {"TextBox 1", "TextBox 2", "TextBox 41"}
    for sh in list(s.shapes):
        if sh.name not in keep:
            sh._element.getparent().remove(sh._element)
    set_one(by_name(s, "TextBox 1"), title)
    set_one(by_name(s, "TextBox 2"), lead)
    set_one(by_name(s, "TextBox 41"), NOTE, size=10)

    # 上の段：入口／登録する方／ゴール（CV）
    bw = (W - BAND_GAP * 2) / 3
    for k, (lab, val) in enumerate([("入口", entry), ("登録する方", who), ("ゴール（CV）", goal)]):
        x = X0 + (bw + BAND_GAP) * k
        rect(s, c(x), c(BAND_Y), c(bw), c(BAND_H), WHITE, line=NAVY, lw=1.0)
        rich(s, c(x + 0.05), c(BAND_Y), c(bw - 0.1), c(BAND_H),
             [[(lab, 10, True, NAVY)], [(val, 12, True, INK)]], align=PP_ALIGN.CENTER)

    # 中央：Day0 → … → Day14 → CV
    xs = [X0 + (CW + AG) * k for k in range(COLS)]
    for k, x in enumerate(xs):
        last = k == COLS - 1
        rich(s, c(x), c(DAY_Y), c(CW), c(DAY_H), [[("CV" if last else DAYS[k], 13, True, NAVY)]],
             align=PP_ALIGN.CENTER)
        if last:
            rect(s, c(x), c(CARD_Y), c(CW), c(CARD_H), NAVY)
            rich(s, c(x + 0.05), c(CARD_Y), c(CW - 0.1), c(CARD_H), [[(cv, 14, True, WHITE)]],
                 align=PP_ALIGN.CENTER)
        else:
            text, aim, btn = steps[k]
            rect(s, c(x), c(CARD_Y), c(CW), c(CARD_H), WHITE, line=LGRAY, lw=0.75)
            rich(s, c(x + 0.04), c(CARD_Y + 0.08), c(CW - 0.08), c(TEXT_H), [[(text, 12, True, INK)]],
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.TOP)
            rich(s, c(x + 0.04), c(CARD_Y + AIM_Y), c(CW - 0.08), c(AIM_H), [[(aim, 10.5, False, INK)]],
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.TOP)
            shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, c(x + (CW - BTN_W) / 2), c(CARD_Y + CARD_H - BTN_B - BTN_H),
                  c(BTN_W), c(BTN_H), [(btn, 11, True, NAVY, 0)], fill=WHITE, line=NAVY, lw=1.0,
                  margins=(0.05, 0.02, 0.05, 0.02))
        if k < COLS - 1:
            tri = shape(s, MSO_SHAPE.ISOSCELES_TRIANGLE, c(x + CW + AG / 2 - 0.06), c(CARD_Y + CARD_H / 2 - 0.09),
                        c(0.12), c(0.18), fill=NAVY)
            tri.rotation = 90


def main(src, out):
    prs = Presentation(src)
    slides = list(prs.slides)
    assert len(slides) == 35
    base = slides[24]
    assert by_name(base, "TextBox 1").text_frame.text == "［第4章］配信設計｜入口ごとの14日ステップ"
    lst = prs.slides._sldIdLst
    anchor = list(lst)[24]
    for page in PAGES:
        new = copy_slide(prs, base)
        e = list(lst)[-1]
        lst.remove(e)
        anchor.addnext(e)
        anchor = e
        build(new, page)
    assert len(prs.slides) == 40
    prs.save(out)
    print("saved:", out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
