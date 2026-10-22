# -*- coding: utf-8 -*-
"""すなお ver6.8 → ver7.0：「想定動線（フェーズごとのLINEの役割）」を①〜④／⑤〜⑧の2ページに分ける（2026-10-22 殿村さん「テキストが多くて見づらい」）。

    python -m _build.patch_sunao_v70_s17_2p <ver6.8.pptx> <ver7.0.pptx>

  ver6.8の表（9行×4列）のページを複製して2枚にし、それぞれ表を外して4フェーズを左から右へ並べる。
  1フェーズ＝フェーズ名の札＋「心理」「離脱課題」「LINEの役割」（LINEの役割だけ薄い青の枠で強調）。文言は表の中身そのまま。
  上の8フェーズの帯は残し、そのページで扱わないフェーズの札は文字を薄くする。
  ver6.9（P17を1枚のブロックにした版）は使わない。ほかのページは変えない（60枚 → 61枚）。
"""
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

from _build.patch_tanaka_v5 import copy_slide
from _build.pptx_parts import CARD, INK, LGRAY, NAVY, hline, rect, rich, shape

SUB = "64748B"
TAG = "EEF0F4"
DIM = "BFBFBF"
IN = 2.54
TITLE = "想定動線（フェーズごとのLINEの役割）"
PAGES = [
    (range(0, 4), TITLE + "①〜④", "LPを離れる方を友だちにし、セミナーの予約・参加までつなげる"),
    (range(4, 8), TITLE + "⑤〜⑧", "参加後の関心が高いうちにサロン入会を案内し、次のイベント・紹介までつなげる"),
]
MARKS = "①②③④⑤⑥⑦⑧"
# 1行に収まらない文は、意味の切れ目で改行する（「収/集中」「経/つ」のように1〜2文字だけ次の行に落ちないように）
BREAKS = ["営業力を上げたい・/情報収集中", "日程が合わず、/申し込まない", "離脱防止ポップアップ/（推奨）",
          "次の開催まで/関心が続きにくい", "ポップアップから/LINE友だち化",
          "どんなセミナーか/分からない", "中身が見えず、/予約に進まない", "見どころ・特典を/ステップ配信",
          "予約後の友だち追加・/前日の案内",
          "参加して/関心が高まっている", "入会を迷ったまま/時間が経つ",
          "次のイベントまで/期間が空く", "入会後の中身・/成果が見えない",
          "継続利用・紹介まで/届きにくい", "イベント案内・紹介/（中長期）"]
SPLIT = {x.replace("/", ""): x.split("/") for x in BREAKS}


def by_text(s, start):
    return next(sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.startswith(start))


def set_text(sh, text):
    p = sh.text_frame.paragraphs[0]
    p.runs[0].text = text
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)


def build(s, idx, title, lead, risk_row=True):
    """risk_row=False で「離脱課題」を外し、「LINEの役割」を「心理」のすぐ下に上げる（ver7.2〜）"""
    gf = next(sh for sh in s.shapes if sh.has_table)
    t = gf.table
    assert t.cell(0, 0).text == "フェーズ" and len(t.rows) == 9
    rows = [[t.cell(r, c).text for c in range(4)] for r in range(1, 9)]
    x0, w = gf.left / 914400, gf.width / 914400
    gf._element.getparent().remove(gf._element)
    set_text(by_text(s, TITLE), title)
    set_text(by_text(s, "8フェーズごとの"), lead)

    for k, mark in enumerate(MARKS):                       # 帯：このページで扱わないフェーズは文字を薄く
        if k in idx:
            continue
        for p in by_text(s, mark).text_frame.paragraphs:
            for r in p.runs:
                r.font.color.rgb = RGBColor.from_string(DIM)

    gap = 0.3
    cw = (w - gap * 3) / 4
    y0, hh = 3.5, 0.62                                     # 札の上端／札の高さ
    for n, k in enumerate(idx):
        name, mind, risk, role = rows[k]
        x = x0 + (cw + gap) * n
        cut = min([i for i in (name.find("（"), name.find("(")) if i > 0] or [len(name)])
        head, sub = name[:cut].strip(), name[cut:].strip()
        rect(s, x * IN, y0 * IN, cw * IN, hh * IN, TAG,
             [(head, 13, True, NAVY, 0)] + ([(sub, 10.5, True, NAVY, 0)] if sub else []))
        if n < 3:                                          # フェーズの間に小さな三角
            shape(s, MSO_SHAPE.ISOSCELES_TRIANGLE, (x + cw + gap / 2 - 0.08) * IN, (y0 + hh / 2 - 0.1) * IN,
                  0.16 * IN, 0.2 * IN, fill=NAVY).rotation = 90
        y = y0 + hh + 0.12
        for j, (lab, val) in enumerate((("心理", mind), ("離脱課題", risk))[:2 if risk_row else 1]):
            rich(s, x * IN, y * IN, cw * IN, 0.24 * IN, [[(lab, 10, False, SUB)]], anchor=MSO_ANCHOR.BOTTOM)
            rich(s, x * IN, (y + 0.24) * IN, cw * IN, 0.5 * IN, [[(v, 13, False, INK)] for v in SPLIT.get(val, [val])],
                 anchor=MSO_ANCHOR.TOP)
            y += 0.82
            hline(s, (x + 0.05) * IN, (x + cw - 0.05) * IN, (y - 0.06) * IN, LGRAY, 0.5)
        rect(s, x * IN, (y + 0.02) * IN, cw * IN, 0.86 * IN, CARD, line=NAVY, lw=0.75,
             paras=[("LINEの役割", 10, True, NAVY, 2)] + [(v, 13, True, NAVY, 0) for v in SPLIT.get(role, [role])],
             align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, margins=(0.2, 0.15, 0.2, 0.1))
    return y + 0.88


def main(src, out):
    prs = Presentation(src)
    slides = list(prs.slides)
    assert len(slides) == 60
    first = slides[16]
    assert by_text(first, TITLE)
    second = copy_slide(prs, first)
    lst = prs.slides._sldIdLst
    ids = list(lst)
    new = ids[-1]
    lst.remove(new)
    ids[16].addnext(new)
    for s, (idx, title, lead) in zip((first, second), PAGES):
        bottom = build(s, idx, title, lead)
        assert bottom <= 7.0, bottom
    assert len(prs.slides) == 61
    used = "|".join(sh.text_frame.text.replace("\n", "") for s in (first, second) for sh in s.shapes if sh.has_text_frame)
    assert all(k in used for k in SPLIT), [k for k in SPLIT if k not in used]   # 改行の表と表の中身がずれていないか
    prs.save(out)
    print("saved:", out, round(bottom, 2))


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
