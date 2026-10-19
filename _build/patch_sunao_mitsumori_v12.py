# -*- coding: utf-8 -*-
"""すなおコンサルティング：「概算お見積り」のページの数字だけを、SIM ver1.2（LP UU 仮に3,000）に差し替える。

    python -m _build.patch_sunao_mitsumori_v12 <今の資料.pptx> <SIM ver1.2.xlsx> [出力.pptx]

  変えるのは、タイトルの追える人数・6か月後の数字（追える人数／セミナー予約／サロン入会）・初期と月次の金額だけ。
  ほかのページ・ほかの文字は触らない（出力を省くと上書き）。
"""
import sys

from pptx import Presentation
from pptx.oxml.ns import qn

from _build.arauchi_v10.build import by_id
from _build.build_sunao_v2 import NAVY, load_sim
from _build.patch_tanaka_v4 import set_lines


def main(src, sim_path, out):
    sim = load_sim(sim_path)
    n = lambda v: int(v + 0.5)                                        # noqa: E731
    active6 = n(sim["active"][-1])
    ev = [k for k, p in enumerate(sim["phase"]) if p == "イベント期"]
    nv = [k for k, p in enumerate(sim["phase"]) if p != "イベント期"]
    total = sim["total"]
    prs = Presentation(src)
    s = next(s for s in prs.slides if any(sh.shape_id == 2 and sh.has_text_frame
                                          and sh.text_frame.text.startswith("概算お見積り") for sh in s.shapes))
    vals = {2: [f"概算お見積り：6か月目に追える人数{active6:,}人"],
            343: [f"¥{n(total[0]):,}〜"],
            344: [f"イベント期 ¥{n(total[1 + ev[0]]):,}〜", f"平常期 ¥{n(total[1 + nv[0]]):,}〜"],
            375: [f"{active6:,}人"],
            376: [f"{sum(n(x) for x in sim['cv1']):,}件/6か月"],
            378: [f"{sum(n(x) for x in sim['cv2']):,}件/6か月"]}
    for sid, lines in vals.items():
        sh = by_id(s, sid)
        before = sh.text_frame.text
        set_lines(sh, lines)
        if sid in (375, 376, 378):
            for clr in sh._element.iter(qn("a:srgbClr")):              # 良い数字は赤い文字にしない
                clr.set("val", NAVY)
        print(sid, before.replace("\n", " / "), "→", " / ".join(lines))
    prs.save(out)
    print("saved:", out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2], sys.argv[3] if len(sys.argv) > 3 else sys.argv[1])
