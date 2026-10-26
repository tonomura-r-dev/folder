"""
10月号 P.9 の見栄えを直す（2026-10-26 殿村さん「見栄えが悪い」）。
左＝数字3つを縦に、右＝対象条件の表（P.6と同じ表の作り）＋注意点カード、の2列にする。
使い方: python3 _build/patch_trend_202610_p9b.py <pptx>
"""
import sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
import trend_report_engine as engine

F = sys.argv[1]
prs = Presentation(F)
s9 = list(prs.slides)[8]


def text(sh):
    return sh.text_frame.text.strip() if sh.has_text_frame else ""


head = next(x for x in s9.shapes if text(x) == "■ 詳細")
tiles = sorted([x for x in s9.shapes if text(x).startswith(("6円", "9円", "2P"))], key=lambda x: x.left)
cond = next(x for x in s9.shapes if text(x).startswith("対象条件"))
caution = next(x for x in s9.shapes if text(x).startswith("注意点"))
cond._element.getparent().remove(cond._element)

L, W = 0.6, 9.64
LW, GAP = 3.0, 0.25
RX, RW = L + LW + GAP, W - LW - GAP
area_top = Emu(head.top + head.height).inches + 0.08

ROWS = [
    ("対象期間", "2026/12/1（火）〜2027/3/31（水）\n※2027/3/31までに掲載が終わる案件"),
    ("対象メニュー", "Friends Swipe AD／Friends Video AD\n※リーチ最適化配信を含む。デリタゲ配信は対象外"),
    ("対象の配信", "ALL配信のみ"),
    ("対象クライアント", "全クライアント"),
    ("最低発注金額", "200万円〜"),
]
HEAD_H = 0.32
row_h = [0.5 if "\n" in v else 0.3 for _, v in ROWS]
TBL_H = HEAD_H + sum(row_h)
CARD_H = 1.38
TOTAL = TBL_H + 0.15 + CARD_H
top = area_top + (7.05 - area_top - TOTAL) / 2

colw = [1.55, RW - 1.55]
gf = s9.shapes.add_table(len(ROWS) + 1, 2, Inches(RX), Inches(top), Inches(RW), Inches(TBL_H))
tbl = gf.table
for j, w in enumerate(colw):
    tbl.columns[j].width = Inches(w)
tbl.rows[0].height = Inches(HEAD_H)
for i, h in enumerate(row_h):
    tbl.rows[i + 1].height = Inches(h)
hc = tbl.cell(0, 0)
hc.merge(tbl.cell(0, 1))


def put(cell, val, size, bold, color, fill, align):
    cell.fill.solid()
    cell.fill.fore_color.rgb = fill
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = cell.margin_right = Inches(0.12)
    cell.margin_top = cell.margin_bottom = Inches(0.03)
    tf = cell.text_frame
    for k, line in enumerate(val.split("\n")):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line
        r.font.size, r.font.bold, r.font.color.rgb = Pt(size if k == 0 else size - 1), bold, color


put(hc, "対象条件", 14, True, engine.WHITE, engine.NAVY, PP_ALIGN.CENTER)
for i, (k, v) in enumerate(ROWS, start=1):
    put(tbl.cell(i, 0), k, 12, True, engine.NAVY, engine.LIGHT_BLUE, PP_ALIGN.CENTER)
    put(tbl.cell(i, 1), v, 12, False, engine.BODY_GRAY, engine.WHITE, PP_ALIGN.LEFT)

caution.left, caution.top, caution.width, caution.height = Inches(RX), Inches(top + TBL_H + 0.15), Inches(RW), Inches(CARD_H)

th = (TOTAL - 2 * 0.15) / 3
for i, t in enumerate(tiles):
    t.left, t.top, t.width, t.height = Inches(L), Inches(top + i * (th + 0.15)), Inches(LW), Inches(th)
    for r in t.text_frame.paragraphs[0].runs:
        r.font.size = Pt(28)
    for r in t.text_frame.paragraphs[1].runs:
        r.font.size = Pt(13)
    for r in t.text_frame.paragraphs[2].runs:
        r.font.size = Pt(10)

engine.force_font(prs)
prs.save(F)
print("saved", F, round(top, 2), round(TOTAL, 2))
