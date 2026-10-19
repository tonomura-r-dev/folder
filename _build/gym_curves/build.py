# -*- coding: utf-8 -*-
"""ジム・フィットネス業界ver3：カーブスを「運用の成功事例」として入れる（2026-10-19 殿村さん）。

    python -m _build.gym_curves.build <ジムver3（46枚・殿村さん保存版 2026-09-09）> <出力.pptx>

  指示：業界資料は、運用がうまくいっている企業の事例をメインで扱う。フィットネスはカーブスの配信をもとに修正。
        タウンワーク（チェングロウスの資料）の時のように、運用の参考事例・成功事例として扱う。
  - 10枚目＝運用の成功事例｜カーブス（友だち592万人・2年で2.6倍・会員数の約6.7倍＋実際の配信の流れ）
  - 11枚目＝カーブスの実際の画面（殿村さんのスクショ 4枚。_data/参考資料/カーブスLINE_20261019/）
  - 21枚目（区切り：施策の詳細）のあとに1枚＝カーブスの型を分解すると、3つの仕組み
      ① あいさつで無料体験と期限を出す（バナー画像。動画は使っていない）
      ② 配信で、動画を使うこともある（全員か、対象を絞っているかは外からは分からない）
      ③ 特典を毎月替えて、締切前に告知する
  - 24・28・29枚目＝あいさつに無料体験と期限のバナーを入れる形に（診断は迷っている人向けに残す）
  - 9枚目＝「配信実態は8/31の実機調査で別途確認」→ カーブスの配信はP10・P11
  「4枚目（当てはめると）はいらない」「動画は、あいさつの段階では未配信」「動画を流すこともある、くらいに」（殿村さん）
"""
import random
import re
import sys
import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.oxml.shapes.picture import CT_Picture

from _build.arauchi_v4.build import slide_order
from _build.arauchi_v4.lib import add_box, make_para
from _build.arauchi_v8.pages import replace_in_paragraphs, shape_by_id
from _build.gym_fmt.build import fix_fonts, set_title
from _build.juku_fmt.build import CT_SLIDE, RT_SLIDE, rels_of
from _build.juku_fmt.plan import corner, set_lead

ROOT = Path(__file__).resolve().parent.parent.parent
SHOTS = ROOT / "_data/参考資料/カーブスLINE_20261019"
N_BASE = 46
CASE, SCREENS, DIVIDER = 10, 11, 21
STEP, TALK, TALK_IMG, REAL = 24, 28, 29, 9
ORDER = list(range(1, DIVIDER + 1)) + [("N", 1)] + list(range(DIVIDER + 1, N_BASE + 1))
NEW_POS = {key: i for i, key in enumerate(ORDER, 1)}
RT_IMAGE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"

NAVY, INK, GREY = "1F2A5C", "1C2233", "646C82"
PALE, BORDER, ACCENT = "EFF1F7", "C6CADD", "BF4E0E"
X0, W = 0.42, 9.99

# 実際の画面（11枚目）：見出し・画像・説明2行・ぼかす所（スクショに写っている殿村さんの名前）
SCREEN_SHOTS = [
    ("1　あいさつ（7/28）", "01_あいさつ_0728.jpg", ["友だち追加の直後に", "無料体験と期限のバナー"], (39, 97, 150, 112)),
    ("2　配信（7/28）", "02_会員の声の動画_0728.jpg", ["会員の声の動画", "＋期限と予約ボタン"], None),
    ("3　配信（9/26）", "06_締切前の告知_0926.jpg", ["締切前の告知", "動画＋カルーセル"], (39, 102, 136, 117)),
    ("4　リッチメニュー（10月）", "08_リッチメニュー_10月.jpg", ["その月の特典を表示", "＋無料スタンプの配布"], None),
]
SHOT_ASPECT = 277 / 600


def shot_bytes(path, blur):
    """スクショを読み、名前の写っている所をぼかす（元の画像ファイルは変えない）"""
    if blur is None:
        return path.read_bytes()
    from io import BytesIO
    from PIL import Image, ImageFilter
    im = Image.open(path).convert("RGB")
    im.paste(im.crop(blur).filter(ImageFilter.GaussianBlur(4)), blur)
    buf = BytesIO()
    im.save(buf, "JPEG", quality=92)
    return buf.getvalue()


def title_of(z, name):
    m = re.search(r"<a:t>([^<]*)</a:t>", z.read(name).decode("utf-8"))
    return m.group(1) if m else ""


# ---------------------------------------------------------------- 1) 組み立て（zipのまま）
def assemble(base, out):
    z = zipfile.ZipFile(base)
    order = slide_order(z)
    assert len(order) == N_BASE
    for n, t in [(CASE, "現状分析｜LINE運用は4つの型に分類できる"), (SCREENS, "現状分析｜各社のリッチメニュー・あいさつメッセージ"),
                 (DIVIDER, "資料アジェンダ"), (STEP, "ステップ配信｜追加から14日で体験予約を取り切る"),
                 (TALK, "トーク画面設計案｜あいさつメッセージと30秒診断"), (TALK_IMG, "トーク画面イメージ｜友だち追加の直後に届く1通目")]:
        assert title_of(z, order[n - 1]) == t, (n, title_of(z, order[n - 1]))
    assert "施策の詳細" in z.read(order[DIVIDER - 1]).decode("utf-8")
    names = set(z.namelist())
    ct = z.read("[Content_Types].xml").decode("utf-8")
    prels = z.read("ppt/_rels/presentation.xml.rels").decode("utf-8")
    pres = z.read("ppt/presentation.xml").decode("utf-8")
    assert '<Default Extension="jpeg"' in ct

    def rel_of(name):
        m = re.search(r'<Relationship [^>]*Target="slides/%s"[^>]*/>' % re.escape(name.split("/")[-1]), prels)
        return re.search(r'Id="(rId\d+)"', m.group(0)).group(1)

    rid_of = {n: rel_of(order[n - 1]) for n in range(1, N_BASE + 1)}
    next_rid = max(int(x) for x in re.findall(r'Id="rId(\d+)"', prels)) + 1
    next_slide = max(int(m.group(1)) for n in names for m in [re.match(r"ppt/slides/slide(\d+)\.xml$", n)] if m) + 1
    add, replace = {}, {}

    # 足す1枚（型の分解）＝10枚目の型（見出し・説明・帯・注記）を写す
    tpl_name = order[CASE - 1]
    tpl = z.read(tpl_name).decode("utf-8")
    tpl_rels = z.read(rels_of(tpl_name))
    assert [r.get("Type").split("/")[-1] for r in etree.fromstring(tpl_rels)] == ["slideLayout"]
    used = set(re.findall(r'creationId[^>]*val="(\d+)"', "".join(z.read(o).decode("utf-8") for o in order)))
    cid = str(random.Random(202610191).randrange(10 ** 9, 4 * 10 ** 9))
    assert cid not in used
    dst = f"ppt/slides/slide{next_slide}.xml"
    new_xml, k = re.subn(r'(<p14:creationId [^>]*val=")\d+(")', rf"\g<1>{cid}\2", tpl, count=1)
    assert k == 1
    add[dst] = new_xml.encode("utf-8")
    add[rels_of(dst)] = tpl_rels
    ct = ct.replace("</Types>", f'<Override PartName="/{dst}" ContentType="{CT_SLIDE}"/></Types>')
    rid = f"rId{next_rid}"
    prels = prels.replace("</Relationships>", f'<Relationship Id="{rid}" Type="{RT_SLIDE}" Target="slides/{dst.split("/")[-1]}"/></Relationships>')
    rid_of[("N", 1)] = rid

    # 11枚目：スクショ4枚の画像を足す（パーツ名は英数字だけ）
    s11 = order[SCREENS - 1]
    rels = etree.fromstring(z.read(rels_of(s11)))
    for j, (_, fn, _, blur) in enumerate(SCREEN_SHOTS, 1):
        media = f"ppt/media/curves_{j}.jpeg"
        assert media not in names
        add[media] = shot_bytes(SHOTS / fn, blur)
        etree.SubElement(rels, "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship",
                         Id=f"rId{100 + j}", Type=RT_IMAGE, Target=f"../media/curves_{j}.jpeg")
    replace[rels_of(s11)] = etree.tostring(rels, xml_declaration=True, encoding="UTF-8", standalone=True)

    m = re.search(r"<p:sldIdLst>(.*?)</p:sldIdLst>", pres, re.S)
    entries = {re.search(r'r:id="(rId\d+)"', e).group(1): e for e in re.findall(r"<p:sldId [^>]*/>", m.group(1))}
    next_id = max(int(x) for x in re.findall(r'<p:sldId id="(\d+)"', m.group(1))) + 1
    lst = [entries[rid_of[key]] if rid_of[key] in entries else f'<p:sldId id="{next_id}" r:id="{rid_of[key]}"/>' for key in ORDER]
    pres = pres.replace(m.group(0), "<p:sldIdLst>" + "".join(lst) + "</p:sldIdLst>")

    app = z.read("docProps/app.xml").decode("utf-8")
    n_new, ppt_t = len(ORDER), "<vt:lpstr>PowerPoint プレゼンテーション</vt:lpstr>"
    assert f"<Slides>{N_BASE}</Slides>" in app and app.count(ppt_t) == N_BASE
    app = app.replace(f"<Slides>{N_BASE}</Slides>", f"<Slides>{n_new}</Slides>")
    app, k = re.subn(r"(<vt:lpstr>スライド タイトル</vt:lpstr></vt:variant><vt:variant><vt:i4>)\d+(</vt:i4>)", rf"\g<1>{n_new}\2", app)
    assert k == 1
    mv = re.search(r'<vt:vector size="(\d+)" baseType="lpstr">', app)
    app = app.replace(mv.group(0), f'<vt:vector size="{int(mv.group(1)) + n_new - N_BASE}" baseType="lpstr">')
    i = app.rfind(ppt_t) + len(ppt_t)
    app = app[:i] + ppt_t * (n_new - N_BASE) + app[i:]
    replace.update({"ppt/presentation.xml": pres.encode("utf-8"), "ppt/_rels/presentation.xml.rels": prels.encode("utf-8"),
                    "[Content_Types].xml": ct.encode("utf-8"), "docProps/app.xml": app.encode("utf-8")})

    with zipfile.ZipFile(out, "w") as zo:
        for info in z.infolist():
            data = replace.get(info.filename)
            zo.writestr(info, data if data is not None else z.read(info.filename), compress_type=info.compress_type)
        for name, data in add.items():
            zo.writestr(name, data, compress_type=zipfile.ZIP_STORED if name.endswith(".jpeg") else zipfile.ZIP_DEFLATED)


# ---------------------------------------------------------------- 2) 中身
def keep_only(s, ids):
    for sh in list(s.shapes):
        if sh.shape_id not in ids:
            sh._element.getparent().remove(sh._element)


def set_one(s, sid, text):
    set_title(shape_by_id(s, sid), text)


def place(sh, x=None, y=None, w=None, h=None):
    for attr, v in (("left", x), ("top", y), ("width", w), ("height", h)):
        if v is not None:
            setattr(sh, attr, int(v * 914400))


def round_box(s, x, y, w, h, paras, adj=8000, **kw):
    return corner(add_box(s, x, y, w, h, paras, shape="round", **kw), adj)


def chip(s, x, y, w, h, text, size=11, fill=NAVY):
    return round_box(s, x, y, w, h, [[(text, {})]], adj=20000, size=size, bold=True, color="FFFFFF", fill=fill,
                     align="ctr", ins=(0.04, 0, 0.04, 0))


def add_pic(s, rid, name, x, y, w, h):
    nid = max(sh.shape_id for sh in s.shapes) + 1
    pic = CT_Picture.new_pic(nid, name, "", rid, int(x * 914400), int(y * 914400), int(w * 914400), int(h * 914400))
    s.shapes._spTree.append(pic)
    ln = etree.SubElement(pic.spPr, qn("a:ln"), w="9525")
    etree.SubElement(etree.SubElement(ln, qn("a:solidFill")), qn("a:srgbClr"), val=BORDER)


def page_case(s):
    """10枚目：運用の成功事例｜カーブス"""
    keep_only(s, {2, 3, 5, 6, 7, 8})
    set_one(s, 2, "運用の成功事例｜カーブス")
    set_lead(s, ["友だち592万人。ジム・フィットネスでも最大級のLINE公式アカウントです。",
                 "毎月「期限つきの無料体験」を、あいさつ・配信・リッチメニューで案内し続けています。"])
    kw, gap = 3.15, (W - 3 * 3.15) / 2
    for i, (lab, big, sub) in enumerate([
            ("友だち数", "592万人", "2026年8月の実測"),
            ("2年での伸び", "×2.6", "2024年8月 約228万人 → 2026年8月"),
            ("会員数に対する友だち数", "約6.7倍", "会員 89.0万人（2026年5月末）に対して")]):
        round_box(s, X0 + i * (kw + gap), 1.5, kw, 1.5,
                  [[(lab, {"size": 11, "bold": True, "color": NAVY})],
                   [(big, {"size": 30, "bold": True, "color": ACCENT})],
                   [(sub, {"size": 9.5, "color": GREY})]], fill=PALE, line=BORDER, align="ctr", line_spacing=100)
    add_box(s, X0, 3.2, W, 0.3, [[("実際の配信の流れ（2026年7〜10月・弊社で確認）", {})]], size=12, bold=True, color=NAVY,
            ins=(0, 0, 0, 0))
    flow = [("7/28｜あいさつ", ["5回無料体験の案内", "期限（7/31まで）と", "先着30名をバナーで"]),
            ("7/28｜配信", ["会員の声の動画", "「半年で-6キロ」", "＋期限と予約ボタン"]),
            ("7月下旬｜配信", ["締切前の告知", "動画", "＋カルーセル"]),
            ("8/27｜配信", ["会員の声の動画", "＋9月の特典", "1か月無料", "（9/30まで）"]),
            ("9/26｜配信", ["締切前の告知", "PEANUTSとのコラボ", "＋カルーセル"]),
            ("10月｜リッチメニュー", ["10回無料体験", "（10/31まで）", "＋無料スタンプ配布"])]
    cw = (W - 5 * 0.1) / 6
    for i, (head, lines) in enumerate(flow):
        x = X0 + i * (cw + 0.1)
        add_box(s, x, 3.87, cw, 1.5, [[(t, {})] for t in lines], size=10, color=INK, fill="FFFFFF", line=BORDER,
                align="ctr", anchor="ctr", ins=(0.04, 0.04, 0.04, 0.04), line_spacing=110)
        add_box(s, x, 3.55, cw, 0.32, [[(head, {})]], size=9.5, bold=True, color="FFFFFF", fill=NAVY, align="ctr",
                ins=(0.02, 0, 0.02, 0))
    add_box(s, X0, 5.5, W, 0.28,
            [[("※ 特典は月ごとに替わる（7月＝5回無料／9月＝1か月無料／10月＝10回無料）。配信では毎回、期限と「見学だけでもOK」をセットにしている。", {})]],
            size=9.5, color=GREY, ins=(0, 0, 0, 0))
    set_one(s, 7, "大規模でも、やっていることは「期限つきの無料体験」を毎月案内し続けるシンプルな型。")
    set_one(s, 8, "※ 友だち数：page.line.me（LINEヤフー社公式ページ）2026/8/17実測・過去値は同ページの保存版（2024/8）。"
                  "会員数：カーブスホールディングス 決算補足資料（2026年8月期 第3四半期・2026年5月末）。配信：弊社確認（2026年7〜10月）。")
    place(shape_by_id(s, 8), y=6.74)


def page_screens(s):
    """11枚目：カーブスの実際の画面（スクショ4枚）"""
    keep_only(s, {2, 3, 22})
    set_one(s, 2, "カーブスの実際の画面")
    set_lead(s, ["あいさつで期限つきの無料体験を案内し、その後の配信で動画や締切前の告知を重ねています。",
                 "リッチメニューにも、その月の特典を表示しています。"])
    colw = 2.3
    gap = (W - 4 * colw) / 3
    ih = 4.1
    iw = ih * SHOT_ASPECT
    for j, (head, _, cap, _) in enumerate(SCREEN_SHOTS):
        x = X0 + j * (colw + gap)
        chip(s, x, 1.55, colw, 0.32, head, size=11)
        add_pic(s, f"rId{101 + j}", f"カーブス_{head[3:]}", x + (colw - iw) / 2, 1.96, iw, ih)
        add_box(s, x, 6.12, colw, 0.5, [[(t, {})] for t in cap], size=10, color=INK, align="ctr", anchor="t",
                ins=(0, 0, 0, 0), line_spacing=105)
    set_one(s, 22, "画面：カーブス LINE公式アカウント（2026年7〜10月 弊社確認）")
    place(shape_by_id(s, 22), y=6.8)


def page_pattern(s):
    """足す1枚：カーブスの型を分解すると、3つの仕組み"""
    keep_only(s, {2, 3, 5, 6, 7, 8})
    set_one(s, 2, "カーブスの型を分解すると、3つの仕組み")
    p_step, p_talk, p_img = NEW_POS[STEP], NEW_POS[TALK], NEW_POS[TALK_IMG]
    set_lead(s, ["やっていることは3つ。どれも同じ型を、毎月くり返しています。",
                 f"この型を、このあとのステップ配信（P{p_step}）とあいさつ（P{p_talk}〜P{p_img}）に取り入れます。"])
    kw, gap = 3.15, (W - 3 * 3.15) / 2
    cols = [
        ("① あいさつで、無料体験と期限を出す",
         [["友だち追加の直後に、期限つきの無料体験をバナー画像で案内する。"],
          ["「お急ぎください」と添えて、今すぐ始めたい人を逃さない。"],
          ["あいさつは、文章とバナー画像だけ（動画は使っていない）。"]],
         ["7/28 あいさつ", "5回無料体験・7/31まで・各店先着30名"]),
        ("② 配信で、動画を使うこともある",
         [["会員の声（「半年で-6キロ」など）の動画で、「自分にもできそう」と感じてもらう。"],
          [("※全員に送っているか、対象を絞って送っているかは、弊社の確認範囲では判別できない。", {"size": 10, "color": GREY})]],
         ["7/28・8/27：会員の声の動画", "7月下旬・9/26：締切前の告知にも動画"]),
        ("③ 特典を毎月替えて、締切前に告知",
         [["その月の特典に「期限・見学だけでもOK」をセットにして、締切の数日前に告知する。"],
          ["文面は、毎月同じ型を使っている。"],
          ["リッチメニューにも、その月の特典を表示。"]],
         ["7月＝5回無料／9月＝1か月無料", "10月＝10回無料（10/31まで）"]),
    ]
    for i, (head, body, ex) in enumerate(cols):
        x = X0 + i * (kw + gap)
        chip(s, x, 1.52, kw, 0.42, head, size=12)
        paras = [[(t, {})] if isinstance(t, str) else [t] for p in body for t in p]
        round_box(s, x, 2.04, kw, 2.25, paras, adj=5000, size=11.5, color=INK, fill=PALE, line=BORDER, anchor="t",
                  ins=(0.1, 0.1, 0.08, 0.06), after=4, line_spacing=105)
        round_box(s, x, 4.42, kw, 1.2,
                  [[("実際の例", {"size": 10, "bold": True, "color": ACCENT})]] + [[(t, {})] for t in ex],
                  adj=5000, size=10.5, color=INK, fill="FFFFFF", line=BORDER, anchor="t", ins=(0.1, 0.08, 0.08, 0.04),
                  line_spacing=105)
    set_one(s, 7, "仕掛けを増やすより、同じ型を毎月くり返す。だから配信を止めずに続けられる。")
    set_one(s, 8, f"※ 弊社確認（カーブス LINE公式アカウント・2026年7〜10月の配信）。実際の画面はP{NEW_POS[SCREENS]}。")


def set_cell(cell, text):
    p = cell.text_frame._txBody.findall(qn("a:p"))
    assert len(p) == 1
    runs = p[0].findall(qn("a:r"))
    runs[0].find(qn("a:t")).text = text
    for r in runs[1:]:
        p[0].remove(r)


def page_step(s):
    """24枚目：Day0のあいさつに、無料体験と期限のバナーを入れる"""
    t = [x for x in s.shapes if x.shape_id == 4][0].table
    row = t.rows[1].cells
    assert [c.text for c in row] == ["Day0", "属性取得", "あいさつ＋30秒 体づくりタイプ診断", "診断回答", "【30秒診断】あなたに合う通い方"]
    set_cell(row[1], "体験の案内＋属性取得")
    set_cell(row[2], "あいさつ＋無料体験（期限つき）のバナー＋30秒 体づくりタイプ診断")
    set_cell(row[3], "体験予約／診断回答")
    set_cell(row[4], "【〇月〇日まで】無料体験のご案内")
    d6 = t.rows[4].cells
    assert d6[0].text == "Day6" and d6[1].text == "1回目のオファー"
    set_cell(d6[1], "2回目のオファー")      # Day0（あいさつ）が1回目になった
    d8 = t.rows[5].cells
    assert d8[0].text == "Day8" and d8[2].text == "同年代・同目的の会員の声／続いている人の共通点"
    set_cell(d8[2], "同年代・同目的の会員の声（動画など）／続いている人の共通点")
    for sid in (5, 6, 7):                 # 下の3つの枠：1行ふえるぶん高さを足す（注記の手前まで）
        place(shape_by_id(s, sid), h=1.05)
    assert replace_in_paragraphs(s, "① CV地点は1つ、オファーは2回", "① CV地点は1つ、オファーは3回") == 1
    assert replace_in_paragraphs(s, "Day6とDay14の2回訴求し、1回で終わらせない。", "Day0・Day6・Day14の3回訴求し、1回で終わらせない。") == 1


LEAD_TALK = "1通目で無料体験と期限を見せる。迷っている人には、30秒診断を用意する。"


def page_talk(s):
    """28枚目：あいさつメッセージ＝文章＋無料体験と期限のバナー（診断は迷っている人向け）"""
    assert shape_by_id(s, 3).text_frame.text == "1通目で予約は求めない。求めるのは診断への回答だけ。"
    set_one(s, 3, LEAD_TALK)
    bub = shape_by_id(s, 5)
    assert bub.text_frame.text.startswith("【〇〇フィットネス △△店】")
    tb = bub.text_frame._txBody
    for p in tb.findall(qn("a:p")):
        tb.remove(p)
    for line in ["【〇〇フィットネス △△店】", "友だち追加ありがとうございます💪", "",
                 "〇月〇日まで、無料体験を受付中です（見学だけでもOK）。", "",
                 "「続けられるか不安」という方は、まずは", "▶ 30秒でわかる「体づくりタイプ診断」から。", "",
                 "※ こちらから営業のお電話をすることはありません。", "ご質問はこのままトークで送ってください(10:00–20:00)"]:
        tb.append(make_para([(line, {})], 11, False, INK, "l"))
    place(bub, h=2.78)
    add_box(s, 0.42, 4.5, 4.8, 0.22, [[("▼ バナー画像（期限・特典は毎月替える）", {})]], size=9, color=GREY, ins=(0.02, 0, 0, 0))
    round_box(s, 0.42, 4.73, 4.8, 0.78,
              [[("〇月〇日まで｜無料体験＋入会特典", {"size": 13, "bold": True})],
               [("各店先着〇名・見学だけでもOK", {"size": 10})]],
              adj=10000, color="FFFFFF", fill=NAVY, ins=(0.14, 0.04, 1.5, 0.04), line_spacing=105)
    round_box(s, 3.78, 4.92, 1.28, 0.4, [[("予約する ▶", {})]], adj=30000, size=11, bold=True, color=NAVY, fill="FFFFFF",
              align="ctr", ins=(0.02, 0, 0.02, 0))
    for sid in (6, 7, 8):                 # 下の3つの枠は、バナーの下へ
        place(shape_by_id(s, sid), y=5.64)
    tip = shape_by_id(s, 7)               # 下の3つの枠の真ん中（右の見出しにも「4問すべて選択式」があるので枠を指定）
    assert tip.text_frame.text == "4問すべて選択式\n自由入力はゼロ", tip.text_frame.text
    for p, new in zip(tip.text_frame._txBody.findall(qn("a:p")), ["無料体験と期限は", "バナーで見せる"]):
        runs = p.findall(qn("a:r"))
        runs[0].find(qn("a:t")).text = new
        for r in runs[1:]:
            p.remove(r)


def page_talk_img(s):
    """29枚目：見出しと②を直す（スマホの画像は差し替えていない）"""
    assert shape_by_id(s, 3).text_frame.text == "1通目で予約は求めない。求めるのは診断への回答だけ。"
    set_one(s, 3, LEAD_TALK)
    assert replace_in_paragraphs(s, "予約ではなく診断を求める", "無料体験と期限は、バナーで見せる") == 1
    assert replace_in_paragraphs(s, "検討初期の人に予約を求めると落ちます", "今すぐの人は、そのまま体験予約へ") == 1
    assert replace_in_paragraphs(s, "求めるのは4問の回答だけ。", "迷っている人には、4問の診断を用意します。") == 1
    # スマホの画像は直す前（診断だけの画面）のままなので、注記で補う（画像はClaudeが作らない）
    assert replace_in_paragraphs(s, "※ 画面はイメージです。", "※ 画面はイメージです（診断の部分。無料体験のバナーは前ページ）。") == 1


def page_real(s):
    """9枚目：「配信実態は別途確認」を、カーブスのページへの案内に"""
    assert replace_in_paragraphs(s, "※配信実態は8/31の実機調査で別途確認",
                                 f"※カーブスの配信はP{NEW_POS[CASE]}・P{NEW_POS[SCREENS]}") == 1


def edit(prs):
    sl = prs.slides
    at = {key: sl[i - 1] for key, i in NEW_POS.items()}
    page_case(at[CASE])
    page_screens(at[SCREENS])
    page_pattern(at[("N", 1)])
    page_step(at[STEP])
    page_talk(at[TALK])
    page_talk_img(at[TALK_IMG])
    page_real(at[REAL])
    return sum(fix_fonts(s) for s in sl)   # 字体は全部メイリオ（元の資料のメイリオ指定なしの文字も）


# ---------------------------------------------------------------- 3) 差し込み＋使わなくなった画像を外す
def prune(path):
    z = zipfile.ZipFile(path)
    names = z.namelist()
    data = {n: z.read(n) for n in names}
    infos = {i.filename: i for i in z.infolist()}
    order = slide_order(z)
    s11 = order[NEW_POS[SCREENS] - 1]
    xml = data[s11].decode("utf-8")
    rel = etree.fromstring(data[rels_of(s11)])
    dropped = []
    for r in list(rel):
        if r.get("Type") == RT_IMAGE and f'"{r.get("Id")}"' not in xml:
            dropped.append(r.get("Target"))
            rel.remove(r)
    data[rels_of(s11)] = etree.tostring(rel, xml_declaration=True, encoding="UTF-8", standalone=True)
    targets = set()
    for n in names:
        if n.endswith(".rels"):
            base = n.replace("_rels/", "").rsplit("/", 1)[0] if "/" in n.replace("_rels/", "") else ""
            for r in etree.fromstring(data[n]):
                if r.get("TargetMode") == "External":
                    continue
                t = r.get("Target")
                full = t.lstrip("/") if t.startswith("/") else str(Path(base, t)).replace("\\", "/")
                parts = []
                for seg in full.split("/"):
                    if seg == "..":
                        parts.pop()
                    elif seg and seg != ".":
                        parts.append(seg)
                targets.add("/".join(parts))
    gone = [n for n in names if n.startswith("ppt/media/") and n not in targets]
    with zipfile.ZipFile(path + ".tmp", "w") as zo:
        for n in names:
            if n in gone:
                continue
            zo.writestr(infos[n], data[n], compress_type=infos[n].compress_type)
    Path(path + ".tmp").replace(path)
    return dropped, gone


def main(base, out):
    mid, tmp = out + ".mid.pptx", out + ".tmp.pptx"
    assemble(base, mid)
    prs = Presentation(mid)
    fixed = edit(prs)
    prs.save(tmp)
    import _build.juku_fmt.build as jb
    jb.ORDER = ORDER                      # repack は並びの長さだけ見る
    jb.repack(mid, tmp, out)
    Path(mid).unlink()
    Path(tmp).unlink()
    dropped, gone = prune(out)
    print("saved:", out, "／", len(ORDER), "枚／外した画像:", gone, "／メイリオにそろえた文字:", fixed)


if __name__ == "__main__":
    main(*sys.argv[1:3])
