# -*- coding: utf-8 -*-
"""ジムver3：殿村さんが文字を直した版（2026-10-19）をベースに、下の要旨の帯の「左の縦線」だけを外す。

    python -m _build.gym_curves.drop_bars <殿村さん保存版（2026-10-19）> <出力.pptx>

  殿村さんは S9・S10 で縦線だけを消し、背景の四角と文字の位置はそのまま残していた（「縦線消した、了解」）。
  → 残りのページも同じく、縦線（幅0.1in未満の細い四角）だけを外す。背景の四角・文字は動かさない。
  ほかは1バイトも変えない（直したスライドのXMLだけ、元のzipに差し込む）。
"""
import sys
from pathlib import Path

from pptx import Presentation

import _build.juku_fmt.build as jb

E = 914400


def drop_bars(prs):
    done = []
    for n, s in enumerate(prs.slides, 1):
        shp = list(s.shapes)
        for bar in shp:
            if bar.width is None or bar.height is None:
                continue
            if not (bar.width < 0.1 * E and 0.3 * E < bar.height < 0.9 * E) or (bar.has_text_frame and bar.text_frame.text.strip()):
                continue
            box = [m for m in shp if m is not bar and m.width and abs(m.top - bar.top) < 0.03 * E
                   and abs(m.left - (bar.left + bar.width)) < 0.05 * E]
            if len(box) != 1:
                continue
            bar._element.getparent().remove(bar._element)
            done.append(n)
    return done


def main(base, out):
    prs = Presentation(base)
    done = drop_bars(prs)
    tmp = out + ".tmp.pptx"
    prs.save(tmp)
    jb.ORDER = list(range(len(prs.slides)))       # repack は並びの長さだけ見る
    jb.repack(base, tmp, out)
    Path(tmp).unlink()
    print("saved:", out, "／縦線を外したページ:", done)


if __name__ == "__main__":
    main(*sys.argv[1:3])
