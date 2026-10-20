# -*- coding: utf-8 -*-
"""すなお ver5.0 → ver5.1：空いていた枠のうち、LPから埋められる3か所を埋める（2026-10-20 殿村さん「埋めて欲しい」）。

    python -m _build.patch_sunao_v51_ume <ver5.0.pptx> [出力.pptx]

  1. 長期ニーズトラック「D. 参加後フォローアップ」：参加直後／参加後30〜60日 のねらいと配信内容（各3つ）。
     中身はLP（https://utage-system.com/p/ZPdd7xiWJJIl）の「受講生の声」「講師メッセージ」「生産性4倍セミナーの特徴」から。
     見出しの下から下端までの空きの上1/3に寄せる（箱は作らず、項目の間に細い灰色の線）
  2. 想定動線 ⑧ファン化・紹介の「離脱・停滞の課題」："--" → 「入会後、LINEの接点が切れる」
  3. 現状把握の「他SNS」："--" → 「Facebookグループ」
     （LPに「Youtubeなどのメディアにも一切出ていない」とあるが、申込みフォームの「何を見てお申込みいただきましたか？」に
       「facebook無料グループ」があるため、「なし」ではなくこちらにした）
  ほかのページは触らない（出力を省くと上書き）。
"""
import sys

from pptx import Presentation
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from _build.arauchi_v4.lib import add_box
from _build.arauchi_v10.build import runs
from _build.build_sunao_v4 import GRAY, NAVY, TEXT
from _build.patch_sunao_v41 import set_text
from _build.patch_sunao_pc44_to_v49 import slide

FOLLOW = [
    ("熱が高いうちに、サロンをご案内", [
        ("セミナーの振り返り", "両想い営業・ライフプランブロックの要点をおさらい"),
        ("受講生の声", "「契約で頭がいっぱい」だった方が、個人保険でTOT"),
        ("サロンのご案内", "ライフプランブロックは、サロンの方だけが使える"),
    ]),
    ("迷っている方へ、成果と続け方を届ける", [
        ("受講生の声", "ロープレ動画で面談を振り返り、TOTを達成"),
        ("講師メッセージ", "毎日遅くまで働いた頃から、働く時間を半分以下に"),
        ("次のセミナーのご案内", "次回の日程と、サロン入会のお申込み方法"),
    ]),
]
COL_X = (0.5, 5.51)
COL_W = 4.89
ITEM_H, ITEM_GAP = 0.62, 0.12


def by_id(s, sid):
    return next(sh for sh in s.shapes if sh.shape_id == sid)


def line(s, x, y, w):
    ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y), Inches(x + w), Inches(y))
    ln.line.color.rgb = RGBColor.from_string("D9D9D9")
    ln.line.width = Pt(0.75)


def p_follow(s):
    head, chevs = by_id(s, 4), (by_id(s, 8), by_id(s, 11))
    assert head.text_frame.text.strip() == "D. 参加後フォローアップ"
    items_h = ITEM_H * 3 + ITEM_GAP * 2
    group = 0.35 + 0.05 + 0.75 + 0.25 + items_h + 0.2 + 0.25      # 見出し〜注記
    top = 1.1 + (7.05 - 1.1 - group) / 3                           # リードの下からの空きの上1/3
    head.top = Inches(top)
    y_chev = top + 0.4
    for ch, (aim, items), x in zip(chevs, FOLLOW, COL_X):
        ch.top = Inches(y_chev)
        p1 = ch.text_frame.paragraphs[1]
        assert p1.runs[0].text == "--"
        p1.runs[0].text = aim
        p1.runs[0].font.size = Pt(12)
        y = y_chev + 0.75 + 0.25
        for i, (title, desc) in enumerate(items):
            if i:
                line(s, x + 0.1, y - ITEM_GAP / 2, COL_W - 0.2)
            add_box(s, x + 0.1, y, COL_W - 0.2, ITEM_H,
                    [runs((title, dict(size=13, bold=True, color=NAVY))),
                     runs((desc, dict(size=12, color=TEXT)))],
                    anchor="ctr", ins=(0.05, 0.02, 0.05, 0.02), name="配信内容")
            y += ITEM_H + ITEM_GAP
    add_box(s, 0.5, y_chev + 0.75 + 0.25 + items_h + 0.2, 9.9, 0.25,
            [runs(("※配信内容は、貴社LPの「受講生の声」「講師メッセージ」をもとにした例", dict(size=9, color=GRAY)))],
            anchor="t", ins=(0, 0, 0, 0), name="注記")


def main(src, out):
    prs = Presentation(src)

    # 1. 長期ニーズトラック：参加後フォローアップ
    p_follow(slide(prs, "長期ニーズトラック：シーズナリティ連動の中長期設計"))

    # 2. 想定動線 ⑧の「離脱・停滞の課題」
    s = [x for x in prs.slides if any(sh.has_table and sh.table.cell(0, 0).text == "フェーズ" for sh in x.shapes)]
    assert len(s) == 1, len(s)
    t = next(sh for sh in s[0].shapes if sh.has_table).table
    r = next(i for i in range(len(t.rows)) if t.cell(i, 0).text.startswith("⑧"))
    assert t.cell(r, 2).text == "--", t.cell(r, 2).text
    assert set_text(t.cell(r, 2), "--", "入会後、LINEの接点が切れる")

    # 3. 現状把握の「他SNS」
    s = slide(prs, "現状把握")
    label = next(sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip() == "他SNS")
    val = next(sh for sh in s.shapes if sh.has_text_frame and abs(sh.top - label.top) < Inches(0.05)
               and sh.left > label.left and sh.left < label.left + Inches(2))
    assert val.text_frame.text.strip() == "--", val.text_frame.text
    val.text_frame.paragraphs[0].runs[0].text = "Facebookグループ"

    prs.save(out)
    print("saved:", out, len(prs.slides), "枚")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else sys.argv[1])
