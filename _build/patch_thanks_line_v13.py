# -*- coding: utf-8 -*-
"""サンクスLINE簡易資料 ver1.2 → ver1.3（2026-10-01 営業戦略に合わせて直す）
戦略：LINE公式アカウントを持っている会社に、広告のCV（資料請求・来店予約）の「後」をLINEで育てる提案をADの提案に乗せる。
  S1 帯の文末を言い切りに
  S3 資料請求→「予約・来店へ」、予約→「来店につなげる」（購入→リピートは残す：殿村さん）
  S4 小数をやめて円の整数に／LINE費用を月5万円で計算し直す
  S5 費用を 初期15万円・月額5万円 に（殿村さん）／リードを「お持ちのLINE公式アカウントに〜」に／文末を言い切りに
  python3 _build/patch_thanks_line_v13.py <ver1.2.pptx>
"""
import sys
from pathlib import Path
from pptx import Presentation
from pptx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "20260930_サンクスLINEのご提案（簡易版）ver1.3.pptx"
AD, LINE_FEE, N_AD, N_LINE = 1_000_000, 50_000, 20, 24
REP = {
    1: [("同じ広告費のまま成果が増えます", "同じ広告費のまま成果が増える")],
    3: [("来店・商談へ", "予約・来店へ"), ("キャンセルを減らす", "来店につなげる")],
    4: [("1件あたり 5.0万円", f"1件あたり {AD // N_AD:,}円"),
        ("1件あたり 約4.3万円", f"1件あたり {round((AD + LINE_FEE) / N_LINE):,}円"),
        ("1件あたり＝（広告費＋LINE費用 月3万円）÷成約件数", "1件あたり＝（広告費＋LINE費用 月5万円）÷成約件数。初期費用は含まない")],
    5: [("進め方｜広告のご提案と、あわせて始められます", "進め方｜広告のご提案と、あわせて開始可能"),
        ("広告の予算に、月3万円〜を足すだけ。", "お持ちのLINE公式アカウントに、完了画面からの導線を足すだけ"),
        ("初期 10万円", "初期 15万円"), ("月額 3万円〜", "月額 5万円"),
        ("LINEへの導線設計をご提案します", "LINEへの導線設計をご提案")],
}
prs = Presentation(sys.argv[1])
done = []
for no, pairs in REP.items():
    s = prs.slides[no - 1]
    for p in s.shapes._spTree.iter(qn("a:p")):
        ts = [t for t in p.iter(qn("a:t"))]
        for old, new in pairs:
            hit = next((t for t in ts if t.text and old in t.text), None)
            if hit is not None:                       # 1つのrunの中なら書式ごと残す
                hit.text = hit.text.replace(old, new); done.append((no, old)); continue
            full = "".join(t.text or "" for t in ts)
            if old in full:                           # runをまたぐときは段落ごと
                ts[0].text = full.replace(old, new)
                for t in ts[1:]:
                    t.text = ""
                done.append((no, old))
miss = [(n, o) for n, ps in REP.items() for o, _ in ps if (n, o) not in done]
assert not miss, miss
prs.save(str(OUT))
print("saved:", OUT.name)
