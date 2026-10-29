# -*- coding: utf-8 -*-
"""横浜シミズ資料（18枚）：16ページ「運用内容」のプラン②の列にプラン④（成果報酬）も入れる（2026-10-29 殿村さん「②のところに④も付け加えて」）。
運用内容はプラン②と④で同じ。見出しの文字だけ変える（列は増やさない）。出力を省くと上書き。

    python -m _build.patch_yokohama_shimizu_p16_plan4 <資料.pptx> [出力.pptx]
"""
import sys

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches

from _build.patch_yokohama_shimizu_4plans import by_id, set_lines


def main(src, out):
    prs = Presentation(src)
    s = prs.slides[15]
    assert by_id(s, 2).text_frame.text.startswith("運用内容")
    tbl = by_id(s, 6)._element.find(".//" + qn("a:tbl"))
    head = tbl.findall(qn("a:tr"))[0].findall(qn("a:tc"))
    assert "プラン②" in "".join(t.text or "" for t in head[2].iter(qn("a:t")))
    set_lines(head[2], ["プラン②・プラン④", "運用代行・効率改善", "／成果報酬"])     # 1行だと「成果／報酬」で切れる
    tr0 = tbl.findall(qn("a:tr"))[0]                          # 見出しが3行になったので、行を少し高く
    add = Inches(0.86) - int(tr0.get("h"))
    tr0.set("h", str(int(Inches(0.86))))
    by_id(s, 6).height = by_id(s, 6).height + add
    set_lines(by_id(s, 5)._element, ["4つのプランの運用内容の比較（プラン②と④は同じ）"])
    prs.save(out)
    print("saved:", out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else sys.argv[1])
