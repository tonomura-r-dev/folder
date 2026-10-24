# -*- coding: utf-8 -*-
"""提案FMT（MAX版）S12「施策全体図」に、黄金ルート（4ステップ）の図を入れる（2026-10-24 殿村さん）。

    python -m _build.patch_fmt_max_s12_golden <殿村さん作業中のFMT.pptx> <出力.pptx>

  1. 集客・獲得：広告・SEOの離脱者（約99%）を、ポップアップバナー・通知メッセージでLINEへ
  2. 即時分類　：友だち追加直後（Day0）のアンケート・診断で、属性・検討度を分ける
  3. 分岐育成　：短期集中トラック（Day0〜15・15日の壁を超える）／中長期トラック（シーズナリティ連動の企画配信）
  4. 成果回収　：不安・手続きの壁を解消して CV（予約・面談・購入）へ → CPA・CPO改善
  色は紺・白・黒（緑はLINEの部分だけ）。上の一文（S7のコピーのまま）も差し替える。ほかのページは変えない。
"""
import sys

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn

from _build.pptx_parts import arrow_line, hline, rect, rich, shape, vline

NAVY, INK, SUB, LINE, LIGHT, GREEN, BLACK, WHITE = "1F2A5C", "333333", "777777", "D0D5E0", "EEF1F8", "06C755", "1A1A1A", "FFFFFF"
IN = 2.54
LEAD = "集客・獲得 → 即時分類 → 分岐育成 → 成果回収 の4ステップで描く、LINE施策の全体図"
STEPS = ["STEP1　集客・獲得", "STEP2　即時分類", "STEP3　分岐育成", "STEP4　成果回収"]
COLS = [(0.40, 2.25), (2.95, 2.25), (5.60, 2.35), (8.35, 2.08)]      # 列の左端・幅（インチ）
KPI = ["見る数字：友だち追加数", "見る数字：回答率", "見る数字：開封率・クリック率", "見る数字：CV数・CPA・CPO"]
BAND = "取りこぼしていた約99%を、LINEで追い続けてCVに変える"


def i(v):
    return v * IN


def box(s, x, y, w, h, paras, fill=WHITE, line=NAVY):
    return shape(s, MSO_SHAPE.RECTANGLE, i(x), i(y), i(w), i(h), paras, fill=fill, line=line, lw=1.25,
                 margins=(0.15, 0.05, 0.15, 0.05))


def arrow(s, x1, y1, x2, y2):
    arrow_line(s, i(x1), i(y1), i(x2), i(y2), NAVY, w=1.75)


def line(s, x1, y1, x2, y2):
    (hline(s, i(x1), i(x2), i(y1), NAVY, 1.75) if y1 == y2 else vline(s, i(x1), i(y1), i(y2), NAVY, 1.75))


def build(s):
    lead = next(sh for sh in s.shapes if sh.name == "テキスト ボックス 6")
    p = lead.text_frame.paragraphs[0]
    p.runs[0].text = LEAD
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)

    # ステップの帯（左から右へ。最後だけ黒）
    for k, ((x, w), t) in enumerate(zip(COLS, STEPS)):
        kind = MSO_SHAPE.PENTAGON if k == 0 else MSO_SHAPE.CHEVRON
        x0 = x if k == 0 else x - 0.22
        w0 = w + (0.22 if k == 0 else 0.44)
        sh = shape(s, kind, i(x0), i(1.98), i(w0), i(0.46), [(t, 13, True, WHITE, 0)],
                   fill=BLACK if k == 3 else NAVY, adj=0.3, margins=(0.3, 0.02, 0.2, 0.02))

    (c1, w1), (c2, w2), (c3, w3), (c4, w4) = COLS
    top, mid, low = 2.70, 3.65, 4.52                                   # 上の箱／中心の箱／下の補足
    # 1. 集客・獲得
    box(s, c1, top, w1, 0.62, [("広告・SEOの来訪者のうち", 11, False, INK, 0), ("離脱する 約99%", 14, True, NAVY, 0)], fill=LIGHT, line=None)
    arrow(s, c1 + w1 / 2, top + 0.64, c1 + w1 / 2, mid - 0.02)
    box(s, c1, mid, w1, 0.80, [("ポップアップバナー", 13, True, NAVY, 0), ("通知メッセージ", 13, True, NAVY, 0)])
    rich(s, i(c1), i(low), i(w1), i(0.55), [[("サイトを離れる前に表示", 10.5, False, SUB)], [("保有リストにも届ける", 10.5, False, SUB)]],
         align=PP_ALIGN.CENTER)
    arrow(s, c1 + w1, mid + 0.40, c2, mid + 0.40)
    # 2. 即時分類
    box(s, c2, top, w2, 0.62, [("最も反応が高い", 11, False, INK, 0), ("友だち追加の直後", 14, True, NAVY, 0)], fill=LIGHT, line=None)
    arrow(s, c2 + w2 / 2, top + 0.64, c2 + w2 / 2, mid - 0.02)
    rect(s, i(c2), i(mid), i(w2), i(0.30), GREEN, [("LINE友だち追加", 11, True, WHITE, 0)], align=PP_ALIGN.CENTER)
    rect(s, i(c2), i(mid + 0.30), i(w2), i(0.50), NAVY, [("Day0 アンケート・診断", 13, True, WHITE, 0)], align=PP_ALIGN.CENTER)
    rich(s, i(c2), i(low), i(w2), i(0.55), [[("属性・検討度を分け、", 10.5, False, SUB)], [("次の配信を出し分ける", 10.5, False, SUB)]],
         align=PP_ALIGN.CENTER)
    # 3. 分岐育成（2つのトラック。短期は15日を過ぎたら中長期に合流）
    t_top, t_low, th = 2.70, 4.10, 1.15
    box(s, c3, t_top, w3, th, [("短期集中トラック", 13, True, NAVY, 2), ("Day0〜15｜検討度が高い方", 10.5, False, INK, 0),
                              ("15日の壁を超えるステップ配信", 10.5, True, INK, 0)])
    box(s, c3, t_low, w3, th, [("中長期トラック", 13, True, NAVY, 2), ("15日以降｜検討度が中・低の方", 10.5, False, INK, 0),
                              ("シーズナリティ連動の企画配信", 10.5, True, INK, 0)])
    ex = (c2 + w2 + c3) / 2 - 0.04                                       # 分かれ道の縦線
    line(s, c2 + w2, mid + 0.40, ex, mid + 0.40)
    line(s, ex, t_top + th / 2, ex, t_low + th / 2)
    arrow(s, ex, t_top + th / 2, c3, t_top + th / 2)
    arrow(s, ex, t_low + th / 2, c3, t_low + th / 2)
    arrow(s, c3 + w3 / 2, t_top + th, c3 + w3 / 2, t_low)
    rich(s, i(c3 + w3 / 2 + 0.06), i(t_top + th + 0.01), i(1.0), i(0.23), [[("15日後に合流", 9, False, SUB)]])
    # 4. 成果回収
    mx = (c3 + w3 + c4) / 2 - 0.05                                       # 合流の縦線
    line(s, c3 + w3, t_top + th / 2, mx, t_top + th / 2)
    line(s, c3 + w3, t_low + th / 2, mx, t_low + th / 2)
    line(s, mx, t_top + th / 2, mx, t_low + th / 2)
    arrow(s, mx, mid + 0.40, c4, mid + 0.40)
    box(s, c4, top, w4, 0.62, [("不安・手続きの壁を", 12, True, NAVY, 0), ("ひとつずつ解消", 12, True, NAVY, 0)], fill=LIGHT, line=None)
    arrow(s, c4 + w4 / 2, top + 0.64, c4 + w4 / 2, mid - 0.02)
    box(s, c4, mid, w4, 0.80, [("CV", 15, True, WHITE, 0), ("予約・面談・購入", 11, False, WHITE, 0)], fill=NAVY, line=None)
    arrow(s, c4 + w4 / 2, mid + 0.82, c4 + w4 / 2, low + 0.20)
    box(s, c4, low + 0.22, w4, 0.55, [("CPA・CPO 改善", 14, True, WHITE, 0)], fill=BLACK, line=None)

    # 見る数字・結論の帯
    hline(s, i(0.40), i(10.43), i(5.55), LINE, 0.75)
    for k, ((x, w), t) in enumerate(zip(COLS, KPI)):
        if k == 3:                                                     # 文字が長いので、列の間の空きまで広げる
            x, w = x - 0.25, w + 0.25
        rich(s, i(x), i(5.60), i(w), i(0.34), [[(t, 10.5, False, SUB)]], align=PP_ALIGN.CENTER)
    rect(s, i(0.40), i(6.15), i(10.03), i(0.48), NAVY, [(BAND, 14, True, WHITE, 0)], align=PP_ALIGN.CENTER)


def main(src, out):
    prs = Presentation(src)
    slides = list(prs.slides)
    assert len(slides) == 43
    s = slides[11]
    assert any(sh.has_text_frame and sh.text_frame.text == "施策全体図" for sh in s.shapes)
    build(s)
    prs.save(out)
    print("saved:", out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
