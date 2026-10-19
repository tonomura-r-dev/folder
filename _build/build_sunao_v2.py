# -*- coding: utf-8 -*-
"""株式会社すなおコンサルティング LINE公式アカウント運用のご提案（骨組み）を、DYM共通資料 LINEOA_BUFFF_3 のFMTで作り直す
（2026-10-16 殿村さん「これをFMTとして、修正してもらえたら」）。

    python -m _build.build_sunao_v2 <LINEOA_BUFFF_3.pptx（FMT）> <あらうちv10.pptx（表紙の「御中」だけ借りる）> <出力.pptx>

  中身は v1（`_build/build_sunao_v1.py`・殿村さん確定の構成）と同じ。見た目をFMTに合わせた：
   ・ヘッダー＝タイトルだけ（リード帯なし）。その下に紺のラベル＋太字のメッセージ（FMTの「施策展開図」と同じ形）
   ・ページは殿村さんの要旨の項目だけ（章の扉・配信の回数のページは入れない。2026-10-16）
   ・運用スケジュール＝FMTのページ（「約1か月」は書かない。見出しはあらうちv10と同じ「初期設計が整い次第、順次進行」）
   ・成果報酬プラン＝FMTの「成果報酬型プラン」の式（成果件数×成果報酬単価＋固定費用）を使い、成果の地点・単価は差し込み枠
   ・表紙・裏表紙＝FMTのもの（表紙に「御中」とサブタイトル2行を入れ、「- 2026最新版 -」は外す）
  SIM（概算お見積り）と成果報酬の条件は、まだ決まっていないので破線の差し込み枠のまま。
"""
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from _build.arauchi_v4.lib import TEXT, add_box, add_table, move, set_paras
from _build.arauchi_v10.build import by_id, clone, runs
from _build.build_sunao_v1 import all_meiryo
from _build.patch_tanaka_v5 import copy_slide

NAVY = "002060"      # FMTのラベル・見出しの紺
LBLUE = "DCE6F2"     # FMTの薄い青の箱
LGRAY = "F2F2F2"     # FMTの薄いグレーの箱
GRAY = "7F7F7F"
RED = "C00000"
LINE_C = "BFBFBF"
W = 9.73
X0 = 0.55
B = {k: (0.75, LINE_C) for k in "lrtb"}
IMG = "_images/"

# FMTのページ番号（1始まり）
F_COVER, F_DIVIDER, F_HEAD, F_SEIKA, F_SCHEDULE, F_BACK = 1, 5, 17, 37, 46, 62
F_POPUP, F_THANKS = 27, 31            # Appendix：ツール費用（ポップアップツール・サンクスLINE誘導）


# ---------------------------------------------------------------- 部品
def head_cell(text, size=14):
    return dict(text=text, fill=NAVY, color="FFFFFF", bold=True, size=size, borders=B)


def cell(text, size=14, bold=False, color=TEXT, fill="FFFFFF", align="l"):
    return dict(text=text, size=size, bold=bold, color=color, fill=fill, align=align, borders=B,
                ins=(0.1, 0.04, 0.08, 0.04))


def label_cell(text, size=14):
    return cell(text, size=size, bold=True, color=NAVY, fill=LBLUE, align="ctr")


def note(s, text, y, h=0.28):
    add_box(s, X0, y, W, h, [runs((text, dict(size=9, color=GRAY)))], anchor="t", ins=(0, 0, 0, 0), name="注記")


def band(s, text, y, h=0.75, size=16):
    add_box(s, X0, y, W, h, [runs((text, dict(size=size, bold=True, color=NAVY)))], fill=LGRAY, align="ctr",
            anchor="ctr", name="示唆")


def heading(s, x, y, w, text, h=0.45):
    add_box(s, x, y, w, h, [runs((text, dict(size=14, bold=True, color="FFFFFF")))], fill=NAVY, align="ctr",
            anchor="ctr", name=f"見出し {text}")


def flow(s, items, y, h, size=14, x0=X0, w_total=W, aw=0.34):
    """横一列の流れ図：items=[(文字, 強調するか)]。強調＝紺の塗り、通常＝薄い青"""
    n = len(items)
    bw = (w_total - aw * (n - 1)) / n
    for k, (text, strong) in enumerate(items):
        x = x0 + k * (bw + aw)
        add_box(s, x, y, bw, h, [runs((t, dict(size=size, bold=True, color="FFFFFF" if strong else NAVY)))
                                 for t in text.split("\n")],
                fill=NAVY if strong else LBLUE, align="ctr", anchor="ctr", line_spacing=115,
                ins=(0.06, 0.04, 0.06, 0.04), name=f"流れ{k + 1}")
        if k < n - 1:
            add_box(s, x + bw, y, aw, h, [runs(("▶", dict(size=14, color=NAVY)))], align="ctr", anchor="ctr",
                    ins=(0, 0, 0, 0), name="矢印")
    return bw


def slot(s, x, y, w, h, text):
    """まだ中身が決まっていない所の破線の差し込み枠"""
    add_box(s, x, y, w, h, [runs((text, dict(size=14, bold=True, color=GRAY)))], line=GRAY, dash="dash",
            align="ctr", anchor="ctr", name=f"差し込み {text}")


def new_page(prs, head, title, label, message):
    """FMTの本文ページ：タイトル＋紺のラベル＋太字のメッセージ（「施策展開図」の形）"""
    s = prs.slides.add_slide(head.slide_layout)
    for sh in list(s.shapes):
        sh._element.getparent().remove(sh._element)
    t = clone(s, by_id(head, 4))
    set_paras(t, [title])
    lb = clone(s, by_id(head, 9))
    set_paras(lb, [label])
    lw = max(1.97, 0.45 + 0.2 * len(label))
    move(lb, x=0, y=0.85, w=lw, h=0.55)
    ms = clone(s, by_id(head, 3))
    set_paras(ms, message.split("\n"))
    two = "\n" in message or len(message) * 0.235 > 10.55 - lw
    move(ms, x=lw + 0.2, y=0.81 if two else 0.94, w=10.6 - lw - 0.2, h=0.64 if two else 0.37)
    return s


def divider(prs, src, text):
    s = copy_slide(prs, src)
    set_paras(by_id(s, 2), [text])
    return s


# ---------------------------------------------------------------- 表紙
def p_cover(s, ar_cover):
    set_paras(by_id(s, 6), ["広告で集めた方をLINEでナーチャリングをし、", "セミナー参加・入会へ"])
    move(by_id(s, 6), x=1.46, y=2.2, w=7.91, h=1.0)
    by_id(s, 3)._element.getparent().remove(by_id(s, 3)._element)          # 「- 2026最新版 -」は外す
    gochu = clone(s, by_id(ar_cover, 31))
    set_paras(gochu, ["株式会社すなおコンサルティング 御中"])


# ---------------------------------------------------------------- 要旨
def p_summary(prs, head):
    s = new_page(prs, head, "サマリー", "提案の要旨", "広告で集めた方をLINEにため、\n次のイベント参加数・入会数を伸ばし、LTVを最大化する")
    add_box(s, X0, 1.65, 4.0, 0.35, [runs(("構築する仕組み", dict(size=13, bold=True, color=NAVY)))], ins=(0, 0, 0, 0),
            name="小見出し")
    flow(s, [("広告", False), ("イベント\n＋LINEにたまる", False), ("ナーチャリング", False), ("次のイベント\n・入会", True)],
         2.05, 1.3, size=16)
    cw, gap = (W - 0.3 * 2) / 3, 0.3
    for k, (h, body) in enumerate([("現状と課題", "申し込まずに離れた方・入会していない方を、次のイベントまで追う場所が無い"),
                                   ("施策", "ポップアップツールとサンクスLINE誘導で友だちにし、興味あり・予約・参加の段階ごとに配信する"),
                                   ("目標と費用", "追える人数を増やし・減らさず、入会へ。\n費用は、3つのプランから選択可能")]):
        x = X0 + k * (cw + gap)
        heading(s, x, 3.8, cw, h)
        add_box(s, x, 4.25, cw, 2.2, [runs((t, dict(size=15, color=TEXT))) for t in body.split("\n")], fill=LGRAY,
                anchor="ctr", line_spacing=140, ins=(0.15, 0.1, 0.12, 0.1), name=f"補足{k + 1}")


# ---------------------------------------------------------------- 0 市場環境（2026-10-19 殿村さん「業界の課題を示すのに役に立つ数値感」）
def p_kankyo(prs, head):
    s = new_page(prs, head, "保険営業の方の環境", "市場環境", "保険の加入は今も営業職員が中心。一方で、お客様と会う回数は減っている")
    lw = 4.75
    heading(s, X0, 1.75, lw, "生命保険の営業職員数（各年度末・名）")
    vals = [(2020, 248601), (2021, 242947), (2022, 247081), (2023, 240903), (2024, 241507)]
    base, hmax, sw, bw = 4.95, 2.05, lw / 5, 0.62                     # 棒は0から（横ばいがそのまま見える）
    for k, (yr, v) in enumerate(vals):
        x = X0 + k * sw
        h = hmax * v / 250000
        add_box(s, x + (sw - bw) / 2, base - h, bw, h, [[("", {})]], fill=NAVY if k == 4 else "8E9BC4",
                name=f"棒 {yr}")
        add_box(s, x, base - h - 0.32, sw, 0.3, [runs((f"{v:,}", dict(size=12, bold=True, color=NAVY)))],
                align="ctr", anchor="b", ins=(0, 0, 0, 0), name=f"値 {yr}")
        add_box(s, x, base + 0.05, sw, 0.3, [runs((f"{yr}年度", dict(size=12, color=TEXT)))], align="ctr",
                anchor="t", ins=(0, 0, 0, 0), name=f"年度 {yr}")
    add_box(s, X0, base, lw, 0.01, [[("", {})]], fill=GRAY, name="基線")
    add_box(s, X0, 5.4, lw, 0.4, [runs(("約24万人で横ばい", dict(size=15, bold=True, color=NAVY)))], align="ctr",
            anchor="ctr", ins=(0, 0, 0, 0), name="棒の示唆")
    rx, rw = X0 + lw + 0.25, W - lw - 0.25
    for y, h, num, lines in ((1.75, 1.75, "56.7%", ["生命保険の加入経路のうち", "「生命保険会社の営業職員」", "（2024年調査）"]),
                             (3.75, 2.05, "55.5%", ["コロナ前と比べ、お客様との", "接触回数が「減少した」営業職員", "（2023年調査）"])):
        heading(s, rx, y, rw, "保険の加入経路" if num == "56.7%" else "お客様との接触回数")
        add_box(s, rx, y + 0.5, 1.75, h - 0.5, [runs((num, dict(size=32, bold=True, color=NAVY)))], align="ctr",
                anchor="ctr", ins=(0, 0, 0, 0), name=f"数字 {num}")
        add_box(s, rx + 1.75, y + 0.5, rw - 1.75, h - 0.5,
                [runs((lines[0], dict(size=13, color=TEXT))), runs((lines[1], dict(size=13, bold=True, color=NAVY))),
                 runs((lines[2], dict(size=12, color=GRAY)))], anchor="ctr", line_spacing=125,
                ins=(0.05, 0, 0, 0), name=f"説明 {num}")
    band(s, "会える回数が減る中で、成果を出す営業のやり方が求められている", y=5.95, h=0.5, size=15)
    for k, t in enumerate(["※営業職員数：生命保険協会「2025年版 生命保険の動向」（加盟会社の年度末登録営業職員数）",
                           "※加入経路：生命保険文化センター「2024（令和6）年度 生命保険に関する全国実態調査」（直近の加入が2019年以降の世帯）",
                           "※接触回数：矢野経済研究所「生命保険の営業職員アンケート調査」（2023年4〜6月、生命保険の営業職員400名）"]):
        note(s, t, y=6.52 + k * 0.15, h=0.15)


# ---------------------------------------------------------------- 1 現状と課題
def p_kadai(prs, head):
    s = new_page(prs, head, "現状と課題", "現状", "イベントの後、離れた方・未入会の方とのつながりが切れる")
    bw = flow(s, [("広告", False), ("LP", False), ("予約", False), ("イベント参加", False), ("サロン入会", True)],
              1.85, 1.0, size=15)
    for k, (h1, body) in ((1, ("申し込まずに離れた方", "次のイベントまで案内する手段が無い")),
                          (3, ("参加して入会していない方", "入会まで案内を続ける手段が無い"))):
        cx = X0 + k * (bw + 0.34) + bw / 2
        add_box(s, cx - 0.2, 2.9, 0.4, 0.32, [runs(("▼", dict(size=13, color=RED)))], align="ctr", anchor="ctr",
                ins=(0, 0, 0, 0), name="矢印下")
        add_box(s, cx - 1.8, 3.25, 3.6, 1.5, [runs((h1, dict(size=15, bold=True, color=RED))),
                                              runs((body, dict(size=14, color=TEXT)))],
                line=RED, align="ctr", anchor="ctr", line_spacing=130, ins=(0.1, 0.06, 0.1, 0.06), name=f"切れる{k}")
    add_box(s, X0, 5.3, W, 0.85, [runs(("集めた方を、次のイベントの集客に", dict(size=16, bold=True, color=NAVY)),
                                       ("活かしきれていない", dict(size=16, bold=True, color=RED)))],
            fill=LGRAY, align="ctr", anchor="ctr", name="示唆")


# ---------------------------------------------------------------- 1 本質課題
def p_honshitsu(prs, head):
    s = new_page(prs, head, "本質課題", "本質課題", "ためて追い続ける場所が無い。この場所を、LINE公式アカウントでつくる")
    cw, gap = (W - 0.3 * 2) / 3, 0.3
    items = [("01 集める", "離れる方・予約した方を、友だちにする手段が無い", "ポップアップツール\nサンクスLINE誘導"),
             ("02 追う", "次のイベントまで、情報を届ける手段が無い", "平常期の配信"),
             ("03 つなげる", "参加した方に、入会を案内する手段が無い", "参加後の入会の案内")]
    for k, (h, kadai, uchite) in enumerate(items):
        x = X0 + k * (cw + gap)
        heading(s, x, 1.75, cw, h)
        add_box(s, x, 2.2, cw, 1.95, [runs((kadai, dict(size=15, color=TEXT)))], fill=LGRAY, align="ctr", anchor="ctr",
                line_spacing=140, ins=(0.15, 0.1, 0.15, 0.1), name=f"課題{k + 1}")
        add_box(s, x + cw / 2 - 0.25, 4.2, 0.5, 0.45, [runs(("▼", dict(size=16, color=NAVY)))], align="ctr",
                anchor="ctr", ins=(0, 0, 0, 0), name="矢印下")
        add_box(s, x, 4.7, cw, 1.9, [runs((t, dict(size=16, bold=True, color=NAVY))) for t in uchite.split("\n")],
                fill=LBLUE, align="ctr", anchor="ctr", line_spacing=130, name=f"打ち手{k + 1}")


# ---------------------------------------------------------------- 2 なぜLINEか
def p_line(prs, head):
    s = new_page(prs, head, "なぜLINEか", "LINE公式アカウント", "広告は毎回ゼロから、LINEは積み上がる")
    rows = [[head_cell(""), head_cell("広告"), head_cell("LINE公式アカウント")],
            [label_cell("届く相手"), cell("毎回、新しく集める"), cell("一度友だちになった方に、何度でも届く", bold=True, color=NAVY)],
            [label_cell("届き方"), cell("広告を見たときだけ"), cell("通知で届き、見てもらいやすい", bold=True, color=NAVY)],
            [label_cell("イベントの後"), cell("つながりが切れる"), cell("友だちとして、毎月積み上がる", bold=True, color=NAVY)]]
    add_table(s, X0, 1.75, [2.0, 3.4, W - 5.4], [0.5, 1.1, 1.1, 1.1], rows, name="広告とLINE")
    band(s, "この仕組みが回れば、今後の広告費を抑えられる可能性もある", y=5.8, h=0.8)


# ---------------------------------------------------------------- 2 気をつける点と対策
def p_chuui(prs, head):
    s = new_page(prs, head, "LINEで気をつける点", "気をつける点と対策", "気をつける点は、運用の設計で抑える")
    rows = [[head_cell("気をつける点"), head_cell("内容"), head_cell("対策")],
            [label_cell("ブロック"), cell("配信が多い・合わない内容だと、ブロックされる"),
             cell("段階に合わせて、配信の中身と回数を変える", bold=True, color=NAVY)],
            [label_cell("配信費"), cell("配信の数に応じて、LINEの費用がかかる"),
             cell("イベントの無い時期は、配信の回数を絞る", bold=True, color=NAVY)],
            [label_cell("時間"), cell("友だちがたまるまで、時間がかかる"),
             cell("次のイベントから入口を置き、イベントごとに積み上げる", bold=True, color=NAVY)]]
    add_table(s, X0, 1.75, [2.0, 3.6, W - 5.6], [0.5, 1.3, 1.3, 1.3], rows, name="気をつける点")


# ---------------------------------------------------------------- 2 他社のLINE運用（2026-10-19 殿村さん「企業の強みより、他社がやっているLINEの運用方法」）
def p_kyogo(prs, head):
    """野沢塾＝殿村さんが友だち追加したあいさつの画面（名前とパスワードだけ灰色で伏せた。元は `_data/sunao/競合LINE_20261019/`）。
    R&C保険営業ラボは招待制で友だち追加できないので、プレスリリースの記載だけ"""
    s = new_page(prs, head, "他社のLINE運用", "他社の運用", "保険営業の方向けの他社も、LINEで特典を渡し、友だちに学びを届けている")
    sw = 1.62
    add_box(s, X0, 1.72, sw, 0.34, [runs(("野沢塾のあいさつ", dict(size=12, bold=True, color=NAVY)))], align="ctr",
            anchor="b", ins=(0, 0, 0, 0.02), name="見出し 画面")
    pic = s.shapes.add_picture(IMG + "sunao_line_nozawa_aisatsu.png", Inches(X0), Inches(2.1), Inches(sw),
                               Inches(sw * 600 / 277))
    pic.line.color.rgb = RGBColor(0xBF, 0xBF, 0xBF)
    pic.line.width = Pt(0.75)
    tx = X0 + sw + 0.3
    cw = (X0 + W - tx - 1.25) / 2
    rows = [[head_cell("", 13), head_cell("野沢塾", 13), head_cell("R&C保険営業ラボ", 13)],
            [label_cell("サービス", 13), cell("保険営業の方向けの講座\nオンラインサロン（月2回）", 12),
             cell("保険営業の方向けの無料の学習サービス\n（2026年9月開始）", 12)],
            [label_cell("入口", 13), cell("サイトの「公式LINE登録」ボタン", 12), cell("社員の紹介URLから\n友だち追加（招待制）", 12)],
            [label_cell("特典", 13), cell("「営業力向上トークスクリプト3本」\nの動画", 12, True, NAVY),
             cell("LINEそのものが、無料の学習サービス", 12, True, NAVY)],
            [label_cell("追加後の\n配信", 13), cell("あいさつで、特典のURLとパスワードを届ける。「最新情報を定期的に配信」と案内", 12),
             cell("登録時に5問のアンケート。答えに合わせて、資料・動画・セミナーのアーカイブ・トークスクリプトを段階的に配信", 12)]]
    add_table(s, tx, 1.72, [1.25, cw, cw], [0.45, 0.8, 0.7, 0.7, 1.2], rows, name="他社のLINE運用")
    band(s, "すなおコンサルティング様は、イベントの予約・参加に合わせた案内で、入会へつなげる", y=5.85, h=0.55, size=15)
    for k, t in enumerate(["※野沢塾：LINE公式アカウント（2026年10月19日確認）・公式サイト（2026年10月18日確認）より。画面の名前とパスワードは伏せている",
                           "※R&C保険営業ラボ：R&C株式会社のプレスリリースより。招待制のため、画面は未掲載"]):
        note(s, t, y=6.5 + k * 0.16, h=0.16)


# ---------------------------------------------------------------- 3 施策：友だちの入口
def p_iriguchi(prs, head):
    s = new_page(prs, head, "友だちの入口", "新規友だち追加", "ポップアップツールとサンクスLINE誘導で、離れる方と予約した方の両方を友だちにする")
    cw, gap = (W - 0.3) / 2, 0.3
    cols = [("ポップアップツール（離脱防止）",
             [("対象", "LPで申し込まずに離れる方"), ("タイミング", "LPを離れようとしたとき"),
              ("案内", "次のセミナーの案内を、LINEで受け取れる")]),
            ("サンクスLINE誘導",
             [("対象", "セミナーを予約した方"), ("タイミング", "予約完了の画面"),
              ("案内", "当日の案内・リマインドを、LINEで届ける")])]
    for k, (h, items) in enumerate(cols):
        x = X0 + k * (cw + gap)
        heading(s, x, 1.75, cw, h)
        add_table(s, x, 2.25, [1.4, cw - 1.4], [1.05] * 3, [[label_cell(a), cell(b)] for a, b in items],
                  name=f"入口{k + 1}")
    band(s, "どちらも、次のイベントまで追える友だちになる", y=5.75, h=0.8)


# ---------------------------------------------------------------- 3 施策：段階ごとの配信
def p_dankai(prs, head):
    s = new_page(prs, head, "配信の設計", "段階ごとの配信", "興味あり・予約・参加の段階ごとに、案内を変える")
    flow(s, [("興味あり", False), ("予約", False), ("参加", False), ("サロン入会", True)], 1.75, 0.65, size=15)
    rows = [[head_cell("段階"), head_cell("対象"), head_cell("届ける内容"), head_cell("ねらい")],
            [label_cell("興味あり"), cell("予約していない方\n（ポップアップツールから）", 13),
             cell("セミナーの見どころ動画・参加特典"), cell("予約へ", bold=True, color=NAVY, align="ctr")],
            [label_cell("予約"), cell("予約した方\n（サンクスLINE誘導から）", 13),
             cell("日程ごとに、前日・当日のお知らせ"), cell("当日の参加へ", bold=True, color=NAVY, align="ctr")],
            [label_cell("参加"), cell("イベントに参加した方", 13),
             cell("イベント直後に、\nサロンの申込みフォームを直接案内\n（申込みはいつでも可能）"),
             cell("入会へ", bold=True, color=NAVY, align="ctr")]]
    add_table(s, X0, 2.65, [1.5, 2.75, 3.98, 1.5], [0.5, 1.05, 1.05, 1.05], rows, name="段階ごとの配信")
    note(s, "※配信の内容は、イベントの内容に合わせて設計", y=6.45)


# ---------------------------------------------------------------- 3 施策：配信の回数
def p_kaisu(prs, head):
    s = new_page(prs, head, "配信の回数", "配信スケジュール", "平常期は月2回、イベント期は月4〜5回")
    flow(s, [("平常期\n月2回", False), ("イベント期\n月4〜5回", True), ("平常期\n月2回", False), ("次のイベント期\n月4〜5回", True)],
         1.75, 0.85, size=14)
    half = (W - 1.6) / 2
    rows = [[head_cell(""), head_cell("平常期（イベントの無い月）"), head_cell("イベント期（イベントの前後）")],
            [label_cell("回数"), cell("月2回", 16, True, NAVY, align="ctr"), cell("月4〜5回", 16, True, NAVY, align="ctr")],
            [label_cell("内容"), cell("定期配信1回＋企画・動画1回"), cell("告知・リマインド・イベント直後のフォロー")],
            [label_cell("ねらい"), cell("つながりを保ち、追える人数を減らさない"), cell("次のイベントの参加・入会へつなげる")]]
    add_table(s, X0, 2.85, [1.6, half, half], [0.5, 0.8, 1.0, 1.0], rows, name="配信の回数")
    note(s, "※イベントを年4回とした場合", y=6.45)


# ---------------------------------------------------------------- 4 目標
def p_mokuhyo(prs, head):
    s = new_page(prs, head, "目標", "指標", "追える人数を増やし・減らさず、次のイベント参加・入会へ")
    flow(s, [("追える人数", False), ("次のイベントの予約・参加", False), ("サロン入会", True)], 1.75, 0.65, size=15)
    rows = [[head_cell("指標"), head_cell("見る時期"), head_cell("内容")],
            [label_cell("追える人数"), cell("平常期", align="ctr"), cell("ブロックされずに届く友だちの数。増やし、減らさない")],
            [label_cell("予約・参加"), cell("イベント期", align="ctr"), cell("LINEからの予約と、当日の参加")],
            [label_cell("サロン入会"), cell("イベントの後", align="ctr"), cell("参加した方のうち、入会した方")]]
    add_table(s, X0, 2.65, [2.2, 1.9, W - 4.1], [0.5, 1.1, 1.1, 1.1], rows, name="指標")


# ---------------------------------------------------------------- 4 概算お見積り
def p_mitsumori(prs, head):
    s = new_page(prs, head, "概算お見積り", "見込み", "平常期・イベント期の友だち数と、予約・入会の見込み")
    slot(s, X0, 1.75, W, 4.9, "概算お見積り（SIM）の表")


# ---------------------------------------------------------------- 4 料金
def p_ryokin(prs, head):
    """3つのプラン（左から安い順）。ツール費は込みにしない。金額はAppendixに入れるので、このページには書かない
    （2026-10-16 殿村さん）。「松竹梅」の文字は使わない"""
    s = new_page(prs, head, "料金", "DYMプラン", "ご予算と運用の体制に合わせて、3つのプランから選択可能")
    big = dict(size=17, bold=True, color=NAVY, align="ctr")
    rows = [[head_cell(""), head_cell("プラン①\n初期設計のみ"), head_cell("プラン②\nイベントに合わせた運用"),
             head_cell("プラン③\n通年の運用・改善")],
            [label_cell("内容", 13), cell("入口と配信の仕組みを構築\n配信は貴社で運用", 12.5),
             cell("弊社で配信\n平常期 月2回\nイベント期 月4〜5回", 12.5),
             cell("弊社で配信＋レポート・定例会\n平常期 月3回\nイベント期 月5回", 12.5)],
            [label_cell("初期費用", 13), cell("10万円〜", **big), cell("10万円〜", **big), cell("10万円〜", **big)],
            [label_cell("平常期\n（月額）", 13), cell("0円", **big), cell("5〜10万円", **big), cell("20万円〜", **big)],
            [label_cell("イベント期\n（月額）", 13), cell("0円", **big), cell("30万円程度", **big), cell("30万円〜", **big)]]
    cw = (W - 1.6) / 3
    add_table(s, X0, 1.75, [1.6, cw, cw, cw], [0.7, 1.15, 0.85, 0.85, 0.85], rows, name="料金")
    for k, t in enumerate(["※広告費・ツール費・LINE公式アカウントの利用料は含まない。いずれも6か月〜の契約",
                           "※施策内容・配信本数・連携範囲等により、費用は個別にお見積り"]):
        note(s, t, y=6.3 + k * 0.26, h=0.26)


# ---------------------------------------------------------------- 4 運用スケジュール（FMTのページ）
def p_schedule(prs, src):
    s = copy_slide(prs, src)
    # 初期構築の期間（約1か月）は書かない（2026-10-15・10-16 殿村さん。あらうちv10と同じ書き方）
    set_paras(by_id(s, 19), ["サイト・LINEへの実装は、初期設計が整い次第、順次進行"])
    for sid in (21, 22):                                      # 「約1か月」の札と点線の矢印
        e = by_id(s, sid)._element
        e.getparent().remove(e)
    set_paras(by_id(s, 20), ["※貴社の作業・確認にかかる時間により、前後する可能性あり"])      # 文末は言い切り
    return s


# ---------------------------------------------------------------- 4 成果報酬プラン（FMTの「成果報酬型プラン」）
def p_seika(prs, src):
    s = copy_slide(prs, src)
    # 右の配信図（離脱防止・CPF広告）と、下の追加作業の費用表は外す（この提案に無い施策・料金が入っているため）
    for sid in (8, 10, 11, 12, 13, 14, 16, 20, 41, 48, 50, 51, 54, 57, 61, 62, 63, 65, 66, 67, 68, 69):
        e = by_id(s, sid)._element
        e.getparent().remove(e)
    move(by_id(s, 24), y=1.95)
    move(by_id(s, 23), y=2.3)
    add_box(s, 0.6, 4.0, 9.69, 0.35, [runs(("すなおコンサルティング様の場合", dict(size=13, bold=True, color=NAVY)))],
            ins=(0, 0, 0, 0), name="小見出し")
    slot(s, 0.6, 4.4, 9.69, 2.4, "成果の地点・成果報酬単価")
    return s


# ---------------------------------------------------------------- Appendix：ツール費用（FMTのページ）
def p_tool_popup(prs, src):
    return copy_slide(prs, src)                                # 初期1.5万円・月額3万円（FMTのまま）


def p_tool_thanks(prs, src):
    """サンクスLINE誘導は 初期15万円・月額3万円〜（2026-10-01 殿村さん）。社内向けの価格改定の注記は外し、実績の数字は紺に"""
    s = copy_slide(prs, src)
    set_paras(by_id(s, 26), ["初期：15万円"])
    e = by_id(s, 10)._element                                 # 「※暫定措置として…2026年7月～の価格改定の予定」
    e.getparent().remove(e)
    for clr in by_id(s, 50)._element.iter(qn("a:srgbClr")):  # 良い数字は赤い文字にしない
        if clr.get("val") == "FF0000":
            clr.set("val", NAVY)
    return s


# 殿村さんの要旨の項目だけで作る（2026-10-16「メインとなる部分だけで資料化して」）。章の扉・配信の回数のページは入れない
# 2026-10-19：市場環境（保険営業の方の環境）と他社のLINE運用を追加。「申込み→当日参加」の離脱は確認不足のため書かない
ORDER = ["summary", "kankyo", "kadai", "honshitsu", "line", "chuui", "kyogo", "iriguchi", "dankai",
         "mokuhyo", "mitsumori", "ryokin", "schedule", "seika",
         ("div", "Appendix"), "tool_popup", "tool_thanks"]          # ツール費用はAppendixに（2026-10-17 殿村さん）
PAGES = dict(summary=p_summary, kankyo=p_kankyo, kadai=p_kadai, honshitsu=p_honshitsu, line=p_line, chuui=p_chuui,
             kyogo=p_kyogo, iriguchi=p_iriguchi, dankai=p_dankai, kaisu=p_kaisu, mokuhyo=p_mokuhyo,
             mitsumori=p_mitsumori, ryokin=p_ryokin)


def main(fmt, arauchi, out):
    prs = Presentation(fmt)
    sl = list(prs.slides)
    cover, div_src, head, seika_src, sch_src, back = (sl[i - 1] for i in (F_COVER, F_DIVIDER, F_HEAD, F_SEIKA,
                                                                          F_SCHEDULE, F_BACK))
    assert by_id(cover, 2).text_frame.text == "LINE公式アカウント運用のご提案"
    assert by_id(div_src, 4).text_frame.text == "資料アジェンダ"
    assert by_id(head, 4).text_frame.text == "施策展開図"
    assert by_id(seika_src, 9).text_frame.text == "成果報酬型プラン"
    assert by_id(sch_src, 18).text_frame.text == "運用スケジュール"
    assert "貴社の明るい未来" in by_id(back, 18).text_frame.text
    popup_src, thanks_src = sl[F_POPUP - 1], sl[F_THANKS - 1]
    assert by_id(popup_src, 22).text_frame.text == "新規友だちの追加（ポップアップ / 離脱防止）"
    assert by_id(popup_src, 15).text_frame.text == "初期：1.5万円"
    assert by_id(thanks_src, 22).text_frame.text == "新規友だちの追加（サンクスLINE誘導）"
    ar_cover = Presentation(arauchi).slides[0]
    new = []
    for item in ORDER:
        if isinstance(item, tuple):
            new.append(divider(prs, div_src, item[1]))
        elif item == "schedule":
            new.append(p_schedule(prs, sch_src))
        elif item == "seika":
            new.append(p_seika(prs, seika_src))
        elif item == "tool_popup":
            new.append(p_tool_popup(prs, popup_src))
        elif item == "tool_thanks":
            new.append(p_tool_thanks(prs, thanks_src))
        else:
            PAGES[item](prs, head)
            new.append(prs.slides[-1])
    p_cover(cover, ar_cover)
    keep = [cover] + new + [back]
    lst = prs.slides._sldIdLst
    by_part = {prs.part.related_part(e.rId): e for e in lst}
    for e in list(lst):
        lst.remove(e)
    keep_parts = {s.part for s in keep}
    for part, e in by_part.items():
        if part not in keep_parts:
            prs.part.drop_rel(e.rId)
    for s in keep:
        lst.append(by_part[s.part])
    all_meiryo(keep)
    prs.save(out)
    print("saved:", out, len(prs.slides), "枚")


if __name__ == "__main__":
    main(*sys.argv[1:4])
