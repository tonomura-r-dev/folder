# -*- coding: utf-8 -*-
"""ワクトリ v2 → v3：施策詳細の8枚（27〜34「具体施策のイメージ」）に中身を入れる（2026-10-28 殿村さん「そのテキストベースで」）。

    python -m _build.patch_wakutori_v3_gutai <v2.pptx> <v3.pptx>

  FMTの枠（左の大きい枠・右の4つの枠）はそのまま。左＝施策名の下に流れ（4〜5段）、右＝目的／対象／中身／見る数字。
  33（サンクスLINE誘導）は両プランともオプションなので、施策名に「（オプション）」と費用を書く。
  画像は使わない（図形と文字だけ）。色は紺・白・灰色。ほかのページは触らない。
"""
import sys

from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

from _build.patch_wakutori_v2_shisaku import by_id, set_lines
from _build.pptx_parts import rect, rich

NAVY, INK, LIGHT = "1F2A5C", "333333", "EEF1F8"
IN = 2.54
RIGHT = ["目的", "対象", "中身", "見る数字"]

PAGES = [  # 施策名, 流れ, (目的, 対象, 中身, 見る数字)
    ("友だちの入口", ["サイト・広告から来る", "求人ページ・トップのLINEボタン", "友だち追加", "あいさつ・診断"],
     ("サイトに来た方を、LINEでつながる友だちにする",
      "サイトに来た方（月 約8.6万UU）",
      "ボタンの訴求は「LINEで新着求人・開始月別の求人が届く」",
      "サイトのボタンからの友だち追加数")),
    ("ポップアップ", ["求人を見る", "応募せずに離れようとする", "ポップアップ「LINEで新着求人を受け取る」", "友だち追加"],
     ("応募せずに離れる方を、LINEで追える友だちにする",
      "求人ページを見て、応募せずに離れる方",
      "離れるタイミングで表示。開始時期・エリアに合う求人がLINEで届くと訴求",
      "表示回数・クリック率・友だち追加数（月 約680〜710人）")),
    ("適正バイト診断", ["友だち追加", "3問タップ（開始時期・エリア・重視したいポイント）", "結果に合う求人を表示", "答えでタグ付け"],
     ("追加の直後に、合う求人と配信の出し分けを決める",
      "新しい友だち全員",
      "今のアンケート（開始時期・エリア）に「重視したいポイント（立地・稼げる・出会い多数・住環境）」を足す",
      "回答率・結果からの求人クリック")),
    ("ステップ配信", ["Day0　診断", "Day1〜2　働き方", "Day3〜4　お金の疑問", "Day6〜7　会社の選び方", "Day10〜15　ENTRY"],
     ("追加から15日で、ENTRYまで信頼を積み上げる",
      "その月の新しい友だち",
      "前後検索「派遣 登録」の順番に合わせた5通（24ページ）",
      "各通のクリック率・ENTRY数")),
    ("開始時期・エリア別の配信", ["診断の答え（開始時期・エリア）", "タグで分ける", "合う求人だけを配信", "応募"],
     ("全員への配信をやめ、合う求人だけを届けてブロックを減らす",
      "開始時期・エリアに答えた友だち",
      "例：12月開始×山・高原 → スキー場の求人／1〜2か月先×海・離島 → 海・離島の求人",
      "ブロック率・クリック率（全員への配信と比べる）")),
    ("リッチメニュー", ["メニューを開く", "開始月別の求人", "求人を見る", "会員登録"],
     ("配信を待たずに、自分で求人とENTRYにたどり着ける",
      "すべての友だち",
      "今の2タブを活かし、上段に開始月別の求人と会員登録。体験談・Q&Aは2つ目のタブ",
      "項目ごとのタップ数")),
    ("サンクスLINE誘導（オプション）", ["マイページ登録（応募）", "完了画面", "LINE友だち追加", "受付と次の流れ・専属エージェントの案内"],
     ("登録した方とLINEでつながり、カウンセリング・お仕事開始まで追う",
      "マイページ登録を完了した方",
      "完了画面からLINEへ。登録の受付・次の流れ・専属エージェントを案内（初期15万円・月額3万円〜）",
      "登録した方のうち友だちになった数")),
    ("契約満了前の次のお仕事案内", ["お仕事開始", "満了が近づく", "次のお仕事の案内", "次の応募"],
     ("1回で終わらせず、次のお仕事につなげる",
      "お仕事を始めた方（診断の開始時期から配信の時期を合わせる）",
      "満了の前に、次の季節のお仕事を案内（冬→春、夏→秋）",
      "再応募数")),
]
RIGHT_Y = [2.03, 3.22, 4.43, 5.66]


def flow(s, steps):
    h, gap, area_top, area_bottom = 0.52, 0.2, 2.95, 6.62     # 行の高さは文字に合わせ、余りは上1/3・下2/3（CLAUDE.md 2026-10-19）
    used = h * len(steps) + gap * (len(steps) - 1)
    top = area_top + (area_bottom - area_top - used) / 3
    for k, step in enumerate(steps):
        y = top + k * (h + gap)
        rect(s, 0.5 * IN, y * IN, 0.45 * IN, h * IN, NAVY, [(str(k + 1), 13, True, "FFFFFF", 0)], align=PP_ALIGN.CENTER)
        rect(s, 0.95 * IN, y * IN, 5.6 * IN, h * IN, LIGHT, [(step, 12.5, True, NAVY, 0)], align=PP_ALIGN.LEFT,
             margins=(0.3, 0.05, 0.2, 0.05))
        if k < len(steps) - 1:
            rich(s, 0.5 * IN, (y + h - 0.02) * IN, 0.45 * IN, (gap + 0.04) * IN, [[("▼", 8, False, NAVY)]],
                 align=PP_ALIGN.CENTER)


def main(src, out):
    prs = Presentation(src)
    for n, (name, steps, right) in zip(range(27, 35), PAGES):
        s = prs.slides[n - 1]
        head = next(sh for sh in s.shapes if sh.name.startswith("TextBox") and sh.has_text_frame and sh.top > 1.8 * 914400)
        old_name = head.text_frame.paragraphs[0].text
        assert old_name[:4] == name[:4], (n, old_name)
        ps = head.text_frame.paragraphs
        ps[0].runs[0].text = name
        flow(s, steps)
        for (y, label), text in zip(zip(RIGHT_Y, RIGHT), right):
            rich(s, 6.97 * IN, (y + 0.06) * IN, 3.46 * IN, 1.01 * IN,
                 [[(label, 12, True, NAVY)], [(text, 10.5, False, INK)]], anchor=MSO_ANCHOR.TOP)
    prs.save(out)
    print("saved:", out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
