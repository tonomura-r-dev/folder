"""
LINE広告トレンドレポート10月号（殿村さん版＝正）を更新する（2026-10-26 殿村さん指示）。
- S3：3段の流れを公式のスケジュール画像に差し替え、カード2枚は右側に縦並び
- S4：1枚（カード4枚）を3枚に分ける（統合の概要／移行の対象と方法／移行ツールで移せるもの・移せないもの）
- 目次を P.3／P.4-6／P.7 に直す
使い方: python3 _build/patch_line_ads_202610_images.py <pptx>
"""
import sys
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches
from PIL import Image
import trend_report_engine as engine

ROOT = Path(__file__).resolve().parent.parent
IMG = ROOT / "_images" / "line_ads_202610"
F = sys.argv[1]
prs = Presentation(F)


def text(sh):
    return sh.text_frame.text.strip() if sh.has_text_frame else ""


def card_set(sh, head, body):
    p0, p1 = sh.text_frame.paragraphs
    for p, t in ((p0, head), (p1, body)):
        p.runs[0].text = t
        for r in list(p.runs[1:]):
            r._r.getparent().remove(r._r)


def place(sh, x, y, w, h):
    sh.left, sh.top, sh.width, sh.height = Inches(x), Inches(y), Inches(w), Inches(h)


def add_image_left(slide, path, top=3.4, bottom=6.95, left=0.6, width=5.7):
    with Image.open(path) as im:
        ratio = im.height / im.width
    h = width * ratio
    y = top + ((bottom - top) - h) / 2
    slide.shapes.add_picture(str(path), Inches(left), Inches(y), Inches(width), Inches(h))


def banner(slide, title):
    sh = next(s for s in slide.shapes if text(s).startswith("【"))
    engine.set_single_run_text(sh, title)


def desc(slide, lines):
    sh = next(s for s in slide.shapes if s.has_text_frame and len(s.text_frame.paragraphs) == 3
              and s.text_frame.paragraphs[0].alignment == engine.PP_ALIGN.CENTER)
    engine.set_multiline_text(sh, lines)


RIGHT_X, RIGHT_W = 6.5, 3.74

# ---------- S3：流れ → スケジュール画像 ----------
s3 = list(prs.slides)[2]
flow = {"2026年6月30日", "2026年10月26日", "2027年3月末頃", "▶"}
chips = [s for s in s3.shapes if text(s) in flow]
chip_lefts = {c.left for c in chips if text(c) != "▶"}
chip_top = min(c.top for c in chips)
boxes = [s for s in s3.shapes if s.has_text_frame and s.left in chip_lefts and s not in chips
         and chip_top < s.top < chip_top + Inches(0.6)]
for s in chips + boxes:
    s._element.getparent().remove(s._element)
cards = sorted([s for s in s3.shapes if text(s).startswith(("APIと一括登録も停止", "請求はあとから更新あり"))],
               key=lambda c: c.left)
place(cards[0], RIGHT_X, 3.4, RIGHT_W, 1.68)
place(cards[1], RIGHT_X, 5.27, RIGHT_W, 1.68)
add_image_left(s3, IMG / "schedule.png")

# ---------- S4：1枚 → 3枚 ----------
s4_index = 3
base_id = prs.slides._sldIdLst[s4_index]
for k in range(2):
    engine.duplicate_slide(prs, s4_index)
    lst = prs.slides._sldIdLst
    new_id = lst[-1]
    lst.remove(new_id)
    lst.insert(s4_index + 1 + k, new_id)
s41, s42, s43 = list(prs.slides)[3:6]


def grid_cards(slide):
    cs = [s for s in slide.shapes if s.has_text_frame and len(s.text_frame.paragraphs) == 2
          and s.top > Inches(3.2) and s.height >= Inches(1.0)]
    return sorted(cs, key=lambda c: (c.top, c.left))


# S4-1 統合の概要
banner(s41, "【LINEヤフー広告】ディスプレイ広告 統合の概要")
desc(s41, [
    "2026年4月、LINE広告とYahoo!広告 ディスプレイ広告が統合され、",
    "「LINEヤフー広告 ディスプレイ広告」の提供が始まった。",
    "Yahoo!広告は自動で移行済み。LINE広告は自動では移らないため、移行が必要。",
])
g = grid_cards(s41)
for c in g[2:]:
    c._element.getparent().remove(c._element)
card_set(g[0], "Yahoo!広告は自動で移行済み",
         "Yahoo!広告 ディスプレイ広告は、2026年4月の統合で自動的にLINEヤフー広告へ移行。追加の対応は不要。")
card_set(g[1], "LINE広告は自動では移らない",
         "画面や仕様はそのまま引き継がれない。アカウントや配信設定は、移行ツールか手動で移す。")
place(g[0], RIGHT_X, 3.4, RIGHT_W, 1.68)
place(g[1], RIGHT_X, 5.27, RIGHT_W, 1.68)
add_image_left(s41, IMG / "integration.png")

# S4-2 移行の対象と方法
banner(s42, "【LINEヤフー広告】移行の対象と方法")
desc(s42, [
    "LINE広告を使っている広告主・代理店は、LINEヤフー広告への移行が必要。",
    "移行の方法は、移行ツールを使う方法と手動の2つ。",
    "配信実績とオーディエンスの蓄積で最適化されるため、早めの移行が推奨されている。",
])
g = grid_cards(s42)
card_set(g[0], "移行が必要な方", "LINE広告を使っている広告主と代理店。代理店が管理しているアカウントも対象。")
card_set(g[1], "移行が不要な方", "Yahoo!広告だけを使っている場合は、原則として対応不要。")
card_set(g[2], "方法①：移行ツール", "できる限り自動で移せる。認証済みのビジネスマネージャーとの連携が必須。")
card_set(g[3], "方法②：手動", "Yahoo!広告のアカウントがあれば、そこへ移すか新しく作る。なければLINEヤフー広告で新しく作る。")

# S4-3 移行ツールで移せるもの・移せないもの
banner(s43, "【LINEヤフー広告】移行ツールで移せるもの・移せないもの")
desc(s43, [
    "移行ツールは、LINE広告の配信データをできる限りLINEヤフー広告へコピーする機能。",
    "タグ・コンバージョン設定や過去の実績など、移せないデータもある。",
    "移行後はキャンペーンが配信オフになるため、設定を確認してから配信を再開する。",
])
g = grid_cards(s43)
card_set(g[0], "移せるもの",
         "アカウント・グループ、クレジットカード情報と権限（移すか選べる）、キャンペーン・広告グループ・広告の設定、画像・動画、オーディエンス。")
card_set(g[1], "移せないもの",
         "「商品フィードから販売」のキャンペーン、LINE Tag・コンバージョン設定、過去の実績・レポート（自動入札の学習データ含む）、外部ツールとのAPI連携。")
card_set(g[2], "事前準備",
         "LINE広告と接続しているLINE公式アカウントを、同じ認証済みビジネスマネージャーに接続。操作する方はビジネスIDを統合しておく。")
card_set(g[3], "対象アカウントと移行後",
         "日本向け配信・運用型・審査承認済みのアカウントが対象。移行後のキャンペーンは配信オフなので、設定を確認してからオンにする。")

# ---------- 目次 ----------
toc = list(prs.slides)[1]
page_sh = next(s for s in toc.shapes if text(s).startswith("P."))
title_sh = next(s for s in toc.shapes if text(s).startswith("【"))
engine.set_multiline_text(page_sh, ["P.3", "P.4-6", "P.7"])
engine.set_multiline_text(title_sh, [
    "【LINE広告】旧「LINE広告」の配信停止（10/26）",
    "【LINEヤフー広告】統合の概要・移行の対象と方法・移行ツール",
    "【LINE VOOM】サービス終了に伴う広告枠の終了",
])

engine.force_font(prs)
prs.save(F)
print("saved", F, "slides", len(prs.slides))
