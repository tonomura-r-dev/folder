# -*- coding: utf-8 -*-
"""株式会社あらうち（100円賃貸・仲介手数料100円）LINEOA施策SIM。新FMT ver3.00 で SNS予算30万円／100万円の2パターン（2026-10-09）。
与件：目標CPA 1,000円（問い合わせ）／現状広告なし／ターゲット＝都内在住の女性オフィスワーカー／LINE友だち1,240人（2026-10-09 page.line.me）。
設計（2026-09-25に提案した案B）：
  問い合わせ①＝友だち追加 → LINEで物件URLを送付（100円賃貸の問い合わせはURLを送るだけで成立）＝新規友だち×URL送付率
  問い合わせ②＝配信（企画4本＋ステップ10本）のクリック×CVR 0.4〜0.8%
  友だち獲得＝SNS予算を全額 LINE友だち追加広告（CPF）に充当（300円/人・仮置き）。離脱防止はUUが小さく割高なのでOFF。
  費用＝LINE運用コンサル 初期20万・月10万＋LINE公式アカウント費（FMTの自動式）。成果報酬は置かない（コンサル運営）。
  python3 _build/build_arauchi_sim.py [出力.xlsx]
"""
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.workbook.properties import CalcProperties

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "_templates/DYM_LINEOA_SIM_FMT_ver3.00.xlsx"
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "株式会社あらうち御中_LINEOA施策提案【SIM】ver1.0.xlsx"

CPF_UNIT = 300                      # 友だち追加広告の単価（円/人）仮置き：LINEヤフー公式事例（UZUZ）の約300円
URL_RATE = 0.25                     # 新規友だちのうち、LINEで物件URLを送ってくる率（問い合わせ①）仮置き
CVR2 = [0.004, 0.005, 0.006, 0.007, 0.008, 0.008]   # 問い合わせ②＝配信クリック×CVR（0.3〜1.0%の範囲）
BLOCK0, BLOCK_K = 0.412, 1.015      # ブロック率：用済み型＋CPF流入で高め。初月41.2%→月1.5%ずつ上昇
UU, FRIENDS0, BLOCK_NOW = 3000, 1240, 0.35   # サイトUU（仮置き）／現状の友だち数（実測）／現状ブロック率（仮置き）
CONSULT_INIT, CONSULT_M = 200000, 100000     # LINE運用コンサル：初期／月額
BUDGETS = (("SIM1_SNS予算30万円", 300000, "30万円"), ("SIM2_SNS予算100万円", 1000000, "100万円"))

wb = load_workbook(SRC)
base, comp = wb["SIM1_"], wb["SIM1_比較"]
base.title = BUDGETS[0][0]
s2 = wb.copy_worksheet(base); s2.title = BUDGETS[1][0]
c2 = wb.copy_worksheet(comp); c2.title = "SIM2_比較"
for ws, name in ((comp, BUDGETS[0][0]), (c2, BUDGETS[1][0])):
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and "SIM1_!" in c.value:
                c.value = c.value.replace("SIM1_!", f"'{name}'!")
    ws["D7"] = f"='{name}'!G27"
    ws["E7"] = f"='{name}'!J27"
wb._sheets = [wb[n] for n in (BUDGETS[0][0], "SIM1_比較", BUDGETS[1][0], "SIM2_比較", "SIM考え方")]
COLS = "EFGHIJ"


def fill(ws, budget, label):
    ws["A4"] = f"株式会社あらうち御中_LINEOA施策提案【SIM】（SNS予算 {label}／月）"
    ws["C8"], ws["D8"] = "問い合わせ①（LINEで物件URL送付）", 0
    ws["C9"], ws["D9"] = "問い合わせ②（配信経由）", 0
    ws["C10"], ws["D10"], ws["C11"], ws["D11"] = "-", 0, "-", 0
    ws["H8"], ws["H9"], ws["H10"], ws["H11"], ws["K8"], ws["P9"] = UU, 0, 0, 0, 0, False
    ws["P30"] = CPF_UNIT
    ws["D27"], ws["D28"] = BLOCK_NOW, FRIENDS0
    ws["D29"] = "=+D28*(1-D27)"   # FMTは0固定だったので式に（比較シートの「実装前」に反映される）
    for k in ("D66", "D67", "D70", "D71", "D72"):
        ws[k] = 0
    ws["D69"] = CONSULT_INIT
    ws["B71"] = "CPF広告費（SNS予算）"
    for i, c in enumerate(COLS):
        prev = "D" if i == 0 else COLS[i - 1]
        ws[f"{c}13"] = f"{i + 1}か月目"
        ws[f"{c}17"], ws[f"{c}18"], ws[f"{c}19"], ws[f"{c}20"], ws[f"{c}21"] = True, False, False, False, False
        ws[f"{c}22"] = f"=IF({c}17=FALSE,0,{c}71/$P$30)"
        ws[f"{c}27"] = BLOCK0 if i == 0 else f"=ROUND({prev}27*{BLOCK_K},3)"
        ws[f"{c}30"], ws[f"{c}32"] = 4, 10
        ws[f"{c}36"] = f"=ROUND({c}26*{URL_RATE},0)"
        ws[f"{c}37"] = f"=ROUND({c}35*{CVR2[i]},0)"
        ws[f"{c}38"], ws[f"{c}39"] = 0, 0
        ws[f"{c}46"] = f"=(SUM({c}65:{c}67)+{c}71)/({c}28-{prev}28)"
        ws[f"{c}56"] = f"=(SUM({c}65:{c}67)+{c}71)/({c}28-{prev}28)"
        ws[f"{c}66"], ws[f"{c}67"], ws[f"{c}69"], ws[f"{c}70"], ws[f"{c}71"], ws[f"{c}72"] = 0, 0, CONSULT_M, 0, budget, 0
    # 根拠メモ（N列＝PDF範囲外）
    ws["N8"] = "UUはSimilarWeb未取得のため仮置き3,000/月（LINEの自然増 約15人/月÷サイト誘導0.5%から逆算）。ターゲット＝都内在住の女性オフィスワーカー"
    ws["N22"] = f"CPF：SNS予算÷{CPF_UNIT}円/人（LINE友だち追加広告・仮置き。LINEヤフー公式事例UZUZの約300円）"
    ws["N23"] = "離脱防止はOFF：UU3,000では月24人・約1,250円/人でCPFより割高"
    ws["N27"] = "ブロック率：用済み型（引っ越しが済めば離脱）＋CPF流入で高め。41.2%→月1.5%ずつ上昇"
    ws["N36"] = f"問い合わせ①＝新規友だち×{int(URL_RATE * 100)}%（LINEで物件URLを送る＝問い合わせ。仮置き）"
    ws["N37"] = "問い合わせ②＝配信クリック×0.4〜0.8%（運用が育つ想定で月ごとに上げる）"
    ws["N41"] = "CVRが高く出るのは①がクリック経由でないため（分母はクリックのみ）。②だけなら0.4〜0.8%"
    ws["N69"] = "LINE運用コンサル：初期20万（あいさつ・リッチメニュー・URL受付の自動応答・CPF設定）／月10万（CPF運用＋企画4本＋ステップ10本）"
    ws["N71"] = "SNS予算は全額CPF（LINE友だち追加広告）に充当。広告運用手数料は含まない（AD側のSIMに合わせる）"


for name, budget, label in BUDGETS:
    fill(wb[name], budget, label)
wb.calculation = CalcProperties(fullCalcOnLoad=True)
wb.save(str(OUT))
print("saved:", OUT.name, wb.sheetnames)
