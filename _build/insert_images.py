# -*- coding: utf-8 -*-
"""Slide 11/13/14/15/16 を、生成したLINEモック画像入りに再構成"""
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
TARGET = ROOT / "20260804_株式会社ビイサイドプランニング御中_公式LINEのご提案v03.pptx"
IMG = ROOT / "_img"

TNAVY = "002060"; NAVY = "1F285A"; RED = "C00000"; INK = "333333"
MUT = "808080"; WHITE = "FFFFFF"; PALE = "F4F7FF"; BORDER = "D9D9D9"

prs = Presentation(str(TARGET))
slides = list(prs.slides)


def set_font(run, size, bold=None, color=INK, name="メイリオ"):
    f = run.font; f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    f.name = name
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        e = rPr.find(qn(tag))
        if e is None:
            e = rPr.makeelement(qn(tag), {}); rPr.append(e)
        e.set("typeface", name)
    f.color.rgb = RGBColor.from_string(color)


ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
ANCH = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}


def reset_tf(tf):
    for para in tf.paragraphs[1:]:
        para._p.getparent().remove(para._p)
    for r in list(tf.paragraphs[0].runs):
        r._r.getparent().remove(r._r)


def put_text(tf, paras, anchor="t", ml=0.06, mr=0.06, mt=0.03, mb=0.03, wrap=True):
    reset_tf(tf); tf.word_wrap = wrap
    tf.margin_left = Inches(ml); tf.margin_right = Inches(mr)
    tf.margin_top = Inches(mt); tf.margin_bottom = Inches(mb)
    tf.vertical_anchor = ANCH[anchor]
    first = True
    for p in paras:
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        para.alignment = ALIGN[p.get("align", "l")]
        if p.get("sa") is not None:
            para.space_after = Pt(p["sa"])
        if p.get("ls") is not None:
            para.line_spacing = p["ls"]
        for (t, sz, b, c) in p["runs"]:
            r = para.add_run(); r.text = t
            set_font(r, sz, b, c)
    return tf


def add_text(slide, x, y, w, h, paras, anchor="t", **kw):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    put_text(box.text_frame, paras, anchor=anchor, **kw)
    return box


def add_box(slide, x, y, w, h, fill=None, line=None, lw=1.0,
            shape=MSO_SHAPE.RECTANGLE, radius=None):
    sp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = RGBColor.from_string(line); sp.line.width = Pt(lw)
    sp.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE and radius is not None:
        try:
            sp.adjustments[0] = radius
        except Exception:
            pass
    return sp


def clear_slide(slide):
    spTree = slide.shapes._spTree
    for el in list(spTree):
        if el.tag.split("}")[-1] in ("sp", "cxnSp", "pic", "graphicFrame", "grpSp"):
            spTree.remove(el)


def frame(slide, title, lead):
    clear_slide(slide)
    add_text(slide, 0.60, 0.13, 9.60, 0.45,
             [{"runs": [(title, 20, True, TNAVY)]}], anchor="m", ml=0.0, mr=0.0)
    add_text(slide, 0.60, 0.68, 9.60, 0.55,
             [{"runs": [(lead, 14, True, INK)], "ls": 1.05}], anchor="m", ml=0.0, mr=0.0)


def phone(slide, name, x, y, h):
    """画像を高さ指定で配置（幅は元比率から算出）。左上xは中心指定ではなく実座標"""
    p = IMG / name
    iw, ih = Image.open(p).size
    w = h * iw / ih
    slide.shapes.add_picture(str(p), Inches(x), Inches(y), Inches(w), Inches(h))
    return w


def caption(slide, x, y, w, text):
    add_text(slide, x, y, w, 0.24,
             [{"runs": [(text, 8.5, None, MUT)], "align": "c"}], anchor="m", ml=0.0, mr=0.0)


def bar(slide, y, text, sz=13.5, h=0.48):
    b = add_box(slide, 0.60, y, 9.60, h, fill=NAVY)
    put_text(b.text_frame, [{"runs": [(text, sz, True, WHITE)], "align": "c"}],
             anchor="m", ml=0.0, mr=0.0, mt=0.0, mb=0.0)


# ============ Slide 11：メッセージ設計 Before / After ============
s = slides[10]
frame(s, "施策① メッセージ設計 Before / After",
      "「予約して終わり」から「予約後に接点を持ち続ける」へ。同じ予約数でも、来場率が変わります。")
cols = [
    (0.60, "Before（現状）", RED, "s11_before.png",
     ["・予約完了後はメールでの一方的な案内のみ",
      "・当日までの数週間、接点が実質ゼロ",
      "・服装や持ち物への不安が解消されないまま当日を迎える"]),
    (5.65, "After（改善後）", NAVY, "s11_after.png",
     ["・友だち追加直後にアンケート（希望職種・年代）",
      "・「服装自由・履歴書不要」等の安心Q&Aを自動送付",
      "・前日・当日朝にリマインドで来場を後押し",
      "・下部のリッチメニューから会場MAPへ常時アクセス"]),
]
for x, label, col, img, lines in cols:
    hb = add_box(s, x, 1.36, 4.55, 0.46, fill=col)
    put_text(hb.text_frame, [{"runs": [(label, 12.5, True, WHITE)], "align": "c"}],
             anchor="m", ml=0.0, mr=0.0, mt=0.0, mb=0.0)
    add_text(s, x, 1.98, 2.55, 3.60,
             [{"runs": [(t, 11, None, INK)], "ls": 1.15, "sa": 7} for t in lines],
             anchor="t", ml=0.02)
    w = phone(s, img, x + 2.70, 1.98, 3.60)
    caption(s, x + 2.70, 5.62, w, "▲ 実際のトーク画面")
bar(s, 5.98, "予約後の空白期間を埋めることが、来場率を左右する", sz=14, h=0.50)

# ============ Slide 13：属性取得アンケート ============
s = slides[12]
frame(s, "具体施策① 友だち追加直後の属性取得アンケート",
      "「30秒でわかる希望条件診断」で回答ハードルを下げ、取得した属性を来場後の追客まで活用します。")
pw = phone(s, "s13_survey.png", 0.60, 1.40, 4.18)
caption(s, 0.60, 5.62, pw, "▲ ボタン選択式で30秒")
add_text(s, 2.95, 1.36, 3.30, 0.30, [{"runs": [("アンケート設問（案）", 12.5, True, NAVY)]}], ml=0.0)
qs = [("Q1", "希望職種は？", "事務／製造・軽作業／販売・サービス／営業"),
      ("Q2", "就職・転職の希望時期は？", "すぐにでも／3ヶ月以内／半年以内／情報収集中"),
      ("Q3", "重視する条件は？", "給与／勤務地／休日／未経験歓迎")]
for i, (no, q, opt) in enumerate(qs):
    y = 1.74 + i * 1.32
    sp = add_box(s, 2.95, y, 3.30, 1.14, fill=PALE)
    put_text(sp.text_frame,
             [{"runs": [(no + "　", 11, True, NAVY), (q, 12, True, INK)], "sa": 4},
              {"runs": [(opt, 10, None, INK)], "ls": 1.15}],
             anchor="m", ml=0.16, mr=0.12, mt=0.06, mb=0.06)
add_text(s, 6.55, 1.36, 3.65, 0.30, [{"runs": [("取得した属性の活用", 12.5, True, NAVY)]}], ml=0.0)
uses = [("① タグ付けによる自動仕分け", "回答内容をタグとして自動付与し、希望職種・時期別にリスト化。"),
        ("② セグメント配信", "「事務希望×すぐにでも」など、条件に合う出展企業だけを案内。"),
        ("③ イベント後の追客", "未応募の求職者へ、次回フェア・個別求人をタグ別に再案内。")]
for i, (h, b) in enumerate(uses):
    y = 1.74 + i * 1.32
    sp = add_box(s, 6.55, y, 3.65, 1.14, fill=WHITE, line=NAVY, lw=1.0)
    put_text(sp.text_frame,
             [{"runs": [(h, 12, True, NAVY)], "sa": 4},
              {"runs": [(b, 10.5, None, INK)], "ls": 1.15}],
             anchor="m", ml=0.16, mr=0.12, mt=0.06, mb=0.06)
bar(s, 6.00, "アンケートは「聞く」ためではなく、配信を当てるための設計", sz=14, h=0.50)

# ============ Slide 14：リッチメニュー ============
s = slides[13]
frame(s, "具体施策② リッチメニュー設計（フェーズ別の出し分け）",
      "検討フェーズに合わせて出し分け、「いま押してほしいボタン」を常に左上に置きます。")
pw = phone(s, "s14_richmenu.png", 0.75, 1.38, 4.32)
caption(s, 0.75, 5.74, pw, "▲ トーク画面下部に常時表示")
pats = [
    ("① 通常時（情報収集フェーズ）", "左上＝出展企業一覧",
     "まずは比較検討。どんな企業が来るのかを見せ、興味を醸成する。"),
    ("② 週末・土日の夜（決断フェーズ）", "左上＝事前予約",
     "転職意欲のピーク。予約を左上に移し、申込への最短距離をつくる。"),
    ("③ フェア直前・3日前〜当日（来場フェーズ）", "左上＝会場MAP・アクセス",
     "当日の迷いを消す。「行けるか不安」を潰し、来場率を守る。"),
]
for i, (name, top, body) in enumerate(pats):
    y = 1.38 + i * 1.46
    sp = add_box(s, 3.55, y, 6.65, 1.30, fill=PALE)
    put_text(sp.text_frame,
             [{"runs": [(name, 12.5, True, NAVY), ("　／　", 11, None, MUT),
                        (top, 11.5, True, INK)], "sa": 4},
              {"runs": [(body, 11, None, INK)], "ls": 1.15}],
             anchor="m", ml=0.18, mr=0.14, mt=0.06, mb=0.06)
add_text(s, 3.55, 5.78, 6.65, 0.30,
         [{"runs": [("切替は①⇄②が毎週自動（金曜夜→日曜深夜）、③は各会場の3日前に切替。画像は3枚制作すれば以降の追加費用は不要です。", 9.5, None, MUT)], "ls": 1.15}], ml=0.0)
bar(s, 6.22, "メニューは固定せず、求職者の検討フェーズに合わせて入れ替える", sz=13.5, h=0.48)

# ============ Slide 15：曜日別 ============
s = slides[14]
frame(s, "具体施策③ 曜日別コンテンツ配信方針（転職熱量のピーク活用）",
      "転職意欲のピークは日曜夜。週末2段構えで「種まき」→「決断の後押し」を設計します。")
days = [
    (0.60, "土曜 夜　｜　種まき", "s15_sat.png", "注目企業TOP5",
     "腰を据えて情報収集する時間帯。出展企業をカルーセルで比較提示し、興味を醸成する。"),
    (5.65, "日曜 夜　｜　決断の後押し", "s15_sun.png", "背中を押す1通",
     "「明日からまた仕事…」の心理が最も高まるタイミング。予約への最短導線だけを提示する。"),
]
for x, label, img, head, body in days:
    hb = add_box(s, x, 1.36, 4.55, 0.46, fill=NAVY)
    put_text(hb.text_frame, [{"runs": [(label, 12.5, True, WHITE)], "align": "c"}],
             anchor="m", ml=0.0, mr=0.0, mt=0.0, mb=0.0)
    w = phone(s, img, x, 1.96, 3.30)
    add_text(s, x + w + 0.18, 1.96, 4.55 - w - 0.18, 3.30,
             [{"runs": [(head, 14, True, NAVY)], "sa": 8},
              {"runs": [(body, 10.5, None, INK)], "ls": 1.2}],
             anchor="m", ml=0.02)
nos = [("月曜朝は配信しない", "意欲は高いが出勤直後で行動に移せない。開封されても流れて終わるため対象外。"),
       ("平日日中は最小限", "勤務中で反応が取れない。お役立ちコラム等の軽い接点にとどめる。")]
for i, (h, b) in enumerate(nos):
    x = 0.60 + i * 4.95
    sp = add_box(s, x, 5.42, 4.55, 0.92, fill=WHITE, line=BORDER, lw=1.0)
    put_text(sp.text_frame,
             [{"runs": [(h, 11, True, NAVY)], "sa": 3},
              {"runs": [(b, 9.5, None, INK)], "ls": 1.1}],
             anchor="m", ml=0.14, mr=0.10, mt=0.05, mb=0.05)
bar(s, 6.48, "「いつ送るか」で反応は変わる。週末2段設計で来場予約を積み上げる", sz=13, h=0.46)

# ============ Slide 16：時間帯別 ============
s = slides[15]
frame(s, "具体施策④ 時間帯別コンテンツ配信方針（朝・昼・夜）",
      "1日の中でも求職者の状態は変わります。時間帯ごとに「読める内容」を出し分けます。")
IW = 6.55
p = IMG / "s16_timeband.png"
iw, ih = Image.open(p).size
IH = IW * ih / iw
s.shapes.add_picture(str(p), Inches((10.83 - IW) / 2), Inches(1.36), Inches(IW), Inches(IH))
ytxt = 1.36 + IH + 0.10
labels = [("朝　7時台", "新着参加企業・会場アクセス"),
          ("昼　12時台", "お役立ちコラム（面接対策）"),
          ("夜　19〜21時", "事前予約・特典受け取り")]
for i, (t, b) in enumerate(labels):
    x = 0.60 + i * 3.28
    sp = add_box(s, x, ytxt, 3.05, 0.78, fill=PALE)
    put_text(sp.text_frame,
             [{"runs": [(t, 11.5, True, NAVY)], "align": "c", "sa": 3},
              {"runs": [(b, 10, None, INK)], "align": "c", "ls": 1.1}],
             anchor="m", ml=0.10, mr=0.10, mt=0.04, mb=0.04)
add_text(s, 0.60, ytxt + 0.86, 9.60, 0.44,
         [{"runs": [("通常週は週2〜3通、フェア直前週は最大週5通まで。開封率・ブロック率を見ながら通数を調整し、「送りすぎ」による離脱を防ぎます。", 10.5, None, INK)], "ls": 1.15}], ml=0.0)
bar(s, 6.52, "夜19〜21時が予約獲得の主戦場。朝・昼は関係構築に使う", sz=13, h=0.46)

prs.save(str(TARGET))
print("saved:", TARGET)
