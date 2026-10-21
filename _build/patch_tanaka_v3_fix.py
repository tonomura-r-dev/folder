# -*- coding: utf-8 -*-
"""タナカホーム v3（殿村さんのPC保存版）：2点だけ直す（2026-10-21 殿村さん「1，2は直して」）。

    python -m _build.patch_tanaka_v3_fix <PC保存版v3.pptx> <出力.pptx>

  1. 提案の流れ：Appendixの行を今のAppendixの中身に合わせる
     「LINE公式アカウントのメリット／中長期の施策／スポットでのご依頼」→「LINE導入の背景／CRM・API連携／その他の活用機能／提供施策／運用プラン」
  2. サイトのLINEボタンは先方のサイトなので、DYMの作業として書かない（CLAUDE.md 2026-10-21。横浜シミズと同じ）
     概算お見積り：初期の実施内容「サイトのLINEボタン｜設置」の行を外し、下の2行を1行ずつ上げる
     運用内容：表の「サイトのLINEボタン｜設置」の行を外す
     スケジュール：12月の「構築：離脱防止・サイトのLINEボタン・サンクスLINE誘導、…」から「サイトのLINEボタン・」を外す
  友だちの入口のページの「サイトのLINEボタン」（入口の1つ）はそのまま。ほかのページは触らない。
"""
import sys

from pptx import Presentation
from pptx.oxml.ns import qn

APPENDIX_OLD = "LINE公式アカウントのメリット／中長期の施策／スポットでのご依頼"
APPENDIX_NEW = "LINE導入の背景／CRM・API連携／その他の活用機能／提供施策／運用プラン"


def by_id(s, sid):
    return next(sh for sh in s.shapes if sh.shape_id == sid)


def replace_in_cell(cell, old, new):
    """runの中の文字を置き換える（書式はそのまま）。1か所だけ"""
    hit = 0
    for r in cell._tc.iter(qn("a:t")):
        if r.text and old in r.text:
            r.text = r.text.replace(old, new)
            hit += 1
    assert hit == 1, (old, hit)


def find_table(s, first_cell):
    return next(sh for sh in s.shapes if sh.has_table and sh.table.cell(0, 0).text.startswith(first_cell))


def main(src, out):
    prs = Presentation(src)
    slides = list(prs.slides)
    assert len(slides) == 31

    # 1. 提案の流れ
    s2 = slides[1]
    replace_in_cell(find_table(s2, "Appendix").table.cell(0, 1), APPENDIX_OLD, APPENDIX_NEW)

    # 2-1. 概算お見積り
    s13 = slides[12]
    assert by_id(s13, 339).text_frame.text.startswith("概算お見積り")
    assert by_id(s13, 369).text_frame.text == "サイトの\nLINEボタン"
    tops = {sid: by_id(s13, sid).top for sid in (369, 371)}
    for sid in (369, 370):
        el = by_id(s13, sid)._element
        el.getparent().remove(el)
    for sids, top in (((371, 372), tops[369]), ((373, 374), tops[371])):
        for sid in sids:
            by_id(s13, sid).top = top

    # 2-2. 運用内容
    s15 = slides[14]
    shape = find_table(s15, "項目")
    trs = shape.table._tbl.findall(qn("a:tr"))
    assert "".join(t.text or "" for t in trs[2].iter(qn("a:t"))).startswith("サイトのLINEボタン")
    shape.height -= int(trs[2].get("h"))
    shape.table._tbl.remove(trs[2])

    # 2-3. スケジュール
    s17 = slides[16]
    replace_in_cell(find_table(s17, "11月").table.cell(1, 1), "サイトのLINEボタン・", "")

    prs.save(out)
    print("saved:", out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
