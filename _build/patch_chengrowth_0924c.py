# -*- coding: utf-8 -*-
"""チェングロウスver1.1（PC保存版 2026-09-24 夜）に当てる。
- P6 前後検索：緑（買う人系）が薄すぎて見えない → 薄め処理を弱めた画像に差し替え、凡例の色を点の色に合わせる
- P16・P20：SIM v3（CVR 0.3〜1.0%前後）の数字に更新（95件・累計CPA 22,768円・6ヶ月目25件/月）
  python3 _build/patch_chengrowth_0924c.py <入力pptx> <出力pptx>
"""
import sys
from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor

ROOT = Path(__file__).resolve().parent.parent
SRC, OUT = sys.argv[1], sys.argv[2]
prs = Presentation(SRC)

# P6 画像差し替え（同じ位置・同じ幅）
s = prs.slides[5]
pic = [sh for sh in s.shapes if sh.shape_type == 13][0]
l, t, w = pic.left, pic.top, pic.width
new = s.shapes.add_picture(str(ROOT / "_images/chengrowth_zengo_colored.png"), l, t, width=w)
pic._element.addprevious(new._element)   # 重なり順を元の位置に
pic._element.getparent().remove(pic._element)
for sh in s.shapes:
    if sh.has_text_frame and "車種・店舗・ローン" in sh.text_frame.text:
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                if "車種" in r.text:
                    r.font.color.rgb = RGBColor.from_string("1FA38F")

# P16・P20 数字の更新（runごと）
REPL = {
    15: [("面談・応募を半年196件と見込みます。", "面談・応募を半年95件と見込みます。"),
         ("37件/月", "25件/月"),
         ("半年累計196件・累計CPA 11,036円", "半年累計95件・累計CPA 22,768円"),
         ("現状CPA 6万・目標2〜2.5万を大きく下回る", "現状CPA 6万の約1/3・目標2〜2.5万の圏内"),
         ("参考：CPFなし＝半年42件・CPA 20,833円", "参考：CPFなし＝半年4件")],
    19: [("面談・応募 196件（6ヶ月累計・SIM①）＝ 累計CPA 11,036円として算出",
          "面談・応募 95件（6ヶ月累計・SIM①）＝ 累計CPA 22,768円として算出")],
}
n = 0
for idx, pairs in REPL.items():
    for sh in prs.slides[idx].shapes:
        if not sh.has_text_frame:
            continue
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                for a, b in pairs:
                    if a in r.text:
                        r.text = r.text.replace(a, b); n += 1
prs.save(OUT)
print("replaced", n)
