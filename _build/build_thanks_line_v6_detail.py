# -*- coding: utf-8 -*-
"""サンクスLINEのご提案（事例入り・4枚）  2026-10-02／2026-10-05 4枚に整理
1 機能（DYM共通資料 31枚目＝動作画面入り）／2 重要性／3 シーン1・2（申込み後のフォロー＝実績3つ）／4 シーン3（購入・来店の後＝リピート）
実績は「LINEヤフー公式の導入事例」と「弊社実績」だけ。各数字は2026-10-02に元ページで再確認済み。
  python3 _build/build_thanks_line_v6_detail.py <20260930_サンクスLINEのご提案ver1.4.pptx> <LINEOA_BUFFF_3.pptx>
"""
import copy
import io
import sys
from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Cm, Pt

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "20261002_サンクスLINEのご提案_事例入り.pptx"
NAVY, INK, GRAY, WHITE = RGBColor(0x1F, 0x28, 0x5A), RGBColor(0x33, 0x33, 0x33), RGBColor(0x7F, 0x7F, 0x7F), RGBColor(0xFF, 0xFF, 0xFF)
TITLE = RGBColor(0x00, 0x20, 0x60)

prs = Presentation(sys.argv[1])
buff = Presentation(sys.argv[2]).slides[0 if "事例入り" in sys.argv[2] else 30]


def shp(si, name):
    return next(x for x in prs.slides[si].shapes if x.name == name)


TPL = {
    "navy": copy.deepcopy(shp(2, "Rounded Rectangle 4")._element),
    "line": copy.deepcopy(shp(1, "Rounded Rectangle 12")._element),
    "arrow": copy.deepcopy(shp(1, "Right Arrow 5")._element),
    "note": copy.deepcopy(shp(3, "TextBox 19")._element),
    "ink": copy.deepcopy(shp(2, "Rounded Rectangle 4")._element),
}


def set_runs(sh, lines, align=None):
    tf = sh.text_frame
    base_p = tf.paragraphs[0]._p
    base_r = base_p.find(qn("a:r"))
    rpr0 = copy.deepcopy(base_r.find(qn("a:rPr"))) if base_r is not None else None
    for p in list(tf._txBody.findall(qn("a:p")))[1:]:
        tf._txBody.remove(p)
    for r in list(base_p.findall(qn("a:r"))) + list(base_p.findall(qn("a:br"))):
        base_p.remove(r)
    empty_p = copy.deepcopy(base_p)
    for i, (text, size, bold, color) in enumerate(lines):
        p = base_p if i == 0 else copy.deepcopy(empty_p)
        if i:
            tf._txBody.append(p)
        r = p.makeelement(qn("a:r"), {})
        if rpr0 is not None:
            r.append(copy.deepcopy(rpr0))
        t = r.makeelement(qn("a:t"), {})
        t.text = text
        r.append(t)
        end = p.find(qn("a:endParaRPr"))
        (end.addprevious(r) if end is not None else p.append(r))
    for para, (text, size, bold, color) in zip(tf.paragraphs, lines):
        if align is not None:
            para.alignment = align
        for run in para.runs:
            run.font.size, run.font.bold = Pt(size), bold
            run.font.color.rgb = color


def add(s, kind, x, y, w, h, lines=None, align=None, fill=None, border=None):
    el = copy.deepcopy(TPL[kind])
    s.shapes._spTree.append(el)
    sh = s.shapes[-1]
    sh.left, sh.top, sh.width, sh.height = Cm(x), Cm(y), Cm(w), Cm(h)
    if kind == "line":
        sh.fill.solid()
        sh.fill.fore_color.rgb = WHITE
        sh.line.color.rgb = NAVY
        sh.line.width = Pt(1.5)
    if kind == "ink":
        sh.fill.solid()
        sh.fill.fore_color.rgb = INK
    if lines:
        set_runs(sh, lines, 2 if (kind in ("line", "navy", "ink") and align is None) else align)
    return sh


def head(si, title, lead):
    s = prs.slides[si]
    for sh in list(s.shapes):
        if sh.name not in ("TextBox 1", "TextBox 2", "Connector 3"):
            sh._element.getparent().remove(sh._element)
    set_runs(shp(si, "TextBox 1"), [(title, 16, True, TITLE)])
    leads = lead if isinstance(lead, (list, tuple)) else [lead]
    set_runs(shp(si, "TextBox 2"), [(x, 18, False, INK) for x in leads])
    t2 = shp(si, "TextBox 2")
    t2.top, t2.height = Cm(1.9), Cm(1.2)
    t1 = shp(si, "TextBox 1")
    t1.top, t1.height = Cm(0.38), Cm(0.94)  # 全ページでタイトルの位置をそろえる
    return s


def source(s, text):
    sh = add(s, "note", 3.0, 17.1, 21.5, 0.7)
    set_runs(sh, [(text, 10, False, GRAY)])


def flow(s, y, steps, bw=4.6, gap=1.0, h=1.7):
    x = 3.0
    for i, (kind, text) in enumerate(steps):
        col = WHITE if kind == "navy" else NAVY
        add(s, kind, x, y, bw, h, [(tx, 13, True, col) for tx in text.split("\n")])
        x += bw
        if i < len(steps) - 1:
            add(s, "arrow", x + 0.05, y + h / 2 - 0.4, 0.9, 0.8)
            x += gap


def band(s, y, text):
    add(s, "ink", 3.0, y, 21.5, 1.2, [(text, 15, True, WHITE)])


def case_cards(s, y, h, cards, num=None):
    """cards = [(区分, 業界名, 取り組み, 指標, 数字, 補足)]  実績の出所・社名は書かない（殿村さん指示 2026-10-04）。
    上（業界名・取り組み）と下（指標・数字・補足）を別の枠にして、数字の高さをカード間でそろえる。枠線・塗りなし。"""
    n = len(cards)
    gap = 0.4
    w = (21.5 - gap * (n - 1)) / n
    top_h = 2.9
    for i, (kind, name, what, label, num_text, note) in enumerate(cards):
        x = 3.0 + i * (w + gap)
        parts = [
            (y, top_h, [(name, 15, True, INK), (what, 13, False, INK)], (6, 0)),
            (y + top_h + 0.1, h - top_h - 0.1, [(label, 13, False, INK), (num_text, num or (20 if n == 3 else 32), True, NAVY)]
             + [(tx, 11.5 if k == 0 else 10, False, GRAY) for k, tx in enumerate(note if isinstance(note, (list, tuple)) else [note])],
             (4, 2) + (2,) * 5),
        ]
        for yy, hh, lines, sas in parts:
            c = add(s, "line", x, yy, w, hh, lines, 1)
            c.line.fill.background()
            c.fill.background()
            c.shadow.inherit = False
            c._element.find(qn("p:style")).find(qn("a:effectRef")).set("idx", "0")
            tf = c.text_frame
            tf.vertical_anchor = MSO_ANCHOR.TOP
            tf.margin_top, tf.margin_left, tf.margin_right, tf.margin_bottom = Cm(0.2), Cm(0.5), Cm(0.5), Cm(0.1)
            for p, sa in zip(tf.paragraphs, sas):
                p.space_after = Pt(sa)


# ================= 重要性（ver1.4 の1枚目を作り替え）=================
s = head(0, "サンクスLINE｜重要性", ["申込み直後の“熱量が高い瞬間”に、LINEでつながる", "申込み直後からLINEで接点を持ち、来店・商談までの離脱を防止。"])
t = add(s, "note", 3.0, 4.2, 21.5, 2.9)
set_runs(t, [("申込み⇒来店までに、一定数が離脱する。", 17, True, NAVY), ("広告費を増やしても、この離脱層は減らない。", 17, True, NAVY),
             ("申込み完了の画面は、関心が一番高い瞬間。ここでLINEにつなげる。", 17, True, NAVY)])
for p in t.text_frame.paragraphs:
    p.space_after = Pt(4)
case_cards(s, 7.3, 8.0, [
    ("公式", "読まれる", "バス会社：申込み直後のLINE通知メッセージは、よく読まれ、クリックもされる", "メッセージの開封率", "約70%", "クリック率はメルマガの約5倍"),
    ("自社", "動く", "美容クリニック：予約完了画面からLINEへ誘導し、予約確認・前日のお知らせを配信", "予約後の来院率", "40〜50%改善", ""),
])
band(s, 15.7, "関心が一番高い申込み直後に、LINEでつながるのが最適")

# ================= シーン1・2を1枚に統合（ver1.4 の2枚目を使う。3枚目は後で削除）=================
s = head(1, "実例｜来店・来院の予約／資料請求", ["申込み後のLINEフォローで、来店・面談までつなげる", "予約確認・リマインド・ステップ配信で、申込み後の離脱を防止"])
flow(s, 4.8, [("line", "申込み完了"), ("navy", "LINEで\nつながる"), ("navy", "確認・リマインド・\n情報配信"), ("line", "来店・面談・\n商談")])
case_cards(s, 7.3, 8.0, [
    ("公式", "美容室", "予約後のフォロー", "次回予約客数", "20%増", ["約200名 → 約240名", "リピート率も88% → 91%（導入前後の半年比較。LINEミニアプリを併用）"]),
    ("公式", "皮膚科クリニック", "友だち追加後のステップ配信", "LINE経由の予約数", "約1.2倍", "ステップ配信実施前との比較"),
    ("公式", "就職支援", "LINEチャットで面談へ誘導", "問い合わせた方のうち", "55%", "が面談予約。他チャネルと比べても高い水準"),
], num=28)
band(s, 15.7, "申込み後の離脱を防ぎ、来店・面談・商談につなげる")

# ================= シーン3（ver1.4 の4枚目）=================
s = head(3, "実例｜購入・来店の後", ["来店・購入後のLINEフォローで、リピートを伸ばす", "アンケート・クーポン・次回来店タイミングの配信で、再来店を後押し"])
flow(s, 4.8, [("line", "購入・来店"), ("navy", "翌日に\nアンケート"), ("navy", "適切なタイミングで\nクーポン・案内"), ("line", "再来店・\nリピート")])
case_cards(s, 7.3, 8.0, [
    ("公式", "居酒屋", "来店翌日にアンケートを配信。再来店の少し前にクーポンを配信（LINEミニアプリを併用）", "リピーターの売上割合", "7.6% → 12.9%", "2025年1月と3月の比較"),
    ("公式", "焼肉店", "来店翌日の11時に、アンケートとクーポンを自動配信（LINEミニアプリを併用）", "リピーター率", "19.4% → 40%超", "2024年1月と2025年1月の比較"),
], num=28)
band(s, 15.7, "購入・来店の後も、LINEで次の来店につなげる")

source(s, "※サンクスLINE単体の効果ではなく、LINEでつながった後の施策の例")

lst = prs.slides._sldIdLst

# ---- 先頭に「機能」（共通資料31枚目）を入れる ----
new = prs.slides.add_slide(prs.slides[1].slide_layout)
for sh in list(new.shapes):
    sh._element.getparent().remove(sh._element)
for el in buff.shapes._spTree:
    if el.tag in (qn("p:nvGrpSpPr"), qn("p:grpSpPr")):
        continue
    el = copy.deepcopy(el)
    for blip in el.iter(qn("a:blip")):
        rid = blip.get(qn("r:embed"))
        if rid:
            img = buff.part.related_part(rid)
            _, new_rid = new.part.get_or_add_image_part(io.BytesIO(img.blob))
            blip.set(qn("r:embed"), new_rid)
    new.shapes._spTree.append(el)
# ---- ver1.4の3枚目（旧シーン2）・5枚目を削除（機能ページの追加後に行う：部品名の重複を避けるため）----
for k in (4, 2):
    sid = lst[k]
    prs.part.drop_rel(sid.get(qn("r:id")))
    lst.remove(sid)
for sl in list(prs.slides)[:-1]:
    for n, c in enumerate(sl.shapes._spTree.iter(qn("p:cNvPr")), start=2):
        c.set("id", str(n))
item = lst[-1]
lst.remove(item)
lst.insert(0, item)
s1 = prs.slides[0]


def by(name):
    return [sh for sh in s1.shapes if sh.name == name]


def par(sh, k, text):
    p = sh.text_frame.paragraphs[k]
    p.runs[0].text = text
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)


# 共通31枚目の代わりに前回出力の1枚目を渡した場合は、編集済みなので手を加えない
if "事例入り" not in sys.argv[2]:
    par(by("Google Shape;752;p28")[0], 0, "サンクスLINE｜機能")
    for sh in by("Google Shape;51;p19"):
        if "初期" in sh.text_frame.text:
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    r.text = r.text.replace("10", "15")
    for sh in by("テキスト ボックス 9"):
        sh._element.getparent().remove(sh._element)
    for sh in by("テキスト ボックス 13"):
        sh.top = Cm(3.9)
    for sh in s1.shapes:
        if sh.has_text_frame and "40-50%" in sh.text_frame.text:
            for p in sh.text_frame.paragraphs:
                if "40-50%" in p.text:
                    runs = p.runs
                    runs[0].text = "美容クリニックでは、予約後来院率が40〜50%改善"
                    for r in runs[1:]:
                        r._r.getparent().remove(r._r)
                    runs[0].font.color.rgb = NAVY
                    runs[0].font.bold = True
    for sh in by("テキスト ボックス 1042"):
        sh.left, sh.top = Cm(1.0), Cm(15.0)

# ---- 字体メイリオ・最小8pt ----
AFTER = {qn(x) for x in ("a:sym", "a:hlinkClick", "a:hlinkMouseOver", "a:rtl", "a:extLst")}
for sl in prs.slides:
    for r in sl.shapes._spTree.iter(qn("a:r")):
        if r.find(qn("a:rPr")) is None:
            r.insert(0, r.makeelement(qn("a:rPr"), {"lang": "ja-JP"}))
    for tag in ("a:rPr", "a:endParaRPr", "a:defRPr"):
        for rpr in sl.shapes._spTree.iter(qn(tag)):
            if rpr.get("sz") and int(rpr.get("sz")) < 800:
                rpr.set("sz", "800")
            fs = []
            for ft in ("a:latin", "a:ea", "a:cs"):
                e = rpr.find(qn(ft))
                if e is None:
                    e = rpr.makeelement(qn(ft), {})
                else:
                    rpr.remove(e)
                e.set("typeface", "メイリオ")
                fs.append(e)
            nxt = next((c for c in rpr if c.tag in AFTER), None)
            for e in fs:
                (nxt.addprevious(e) if nxt is not None else rpr.append(e))
prs.save(str(OUT))
print("saved:", OUT.name, len(prs.slides), "枚")
