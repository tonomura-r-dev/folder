"""
10月号 P.9（Friends AD 冬季キャンペーン）を見やすく組み直す（2026-10-26 殿村さん指示）。
- 数字の枠は 6円／9円／2P の3つ（200万円〜は対象条件へ）
- 表と「発注時の注意」をやめて、「対象条件」「注意点」のカードを左右に並べる
- 逆三角形＝認知
使い方: python3 _build/patch_trend_202610_p9.py <pptx>
"""
import sys
import copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.oxml.ns import qn
import trend_report_engine as engine

F = sys.argv[1]
prs = Presentation(F)
s9 = list(prs.slides)[8]


def text(sh):
    return sh.text_frame.text.strip() if sh.has_text_frame else ""


assert any(text(x).startswith("【LINEポイント Friends AD】") for x in s9.shapes)
head = next(x for x in s9.shapes if text(x) == "■ 詳細")
tiles = sorted([x for x in s9.shapes if text(x).startswith(("6円", "9円", "2P", "200万円"))], key=lambda x: x.left)
tbl = next(x for x in s9.shapes if x.has_table)
note = next(x for x in s9.shapes if text(x).startswith("発注時の注意"))
for x in (tiles[3], tbl):
    x._element.getparent().remove(x._element)
tiles = tiles[:3]

SW = Emu(prs.slide_width).inches
L, W, GAP = 0.6, 9.64, 0.2
area_top = Emu(head.top + head.height).inches + 0.08

# ---- 数字の枠（3つ） ----
TILE_H = 0.9
tw = (W - 2 * GAP) / 3

# ---- カード（左右） ----
mx = max(int(e.get("id")) for e in s9._element.iter(qn("p:cNvPr")))
el = copy.deepcopy(note._element)
note._element.addnext(el)
cn = el.find(".//" + qn("p:cNvPr"))
cn.set("id", str(mx + 1))
cn.set("name", cn.get("name") + "_cond")
cond = next(x for x in s9.shapes if x._element is el)
caution = note

COND = ["対象条件", [
    "・対象期間：2026/12/1（火）〜2027/3/31（水）\x0b　※2027/3/31までに掲載が終わる案件",
    "・対象メニュー：Friends Swipe AD／Video AD\x0b　※リーチ最適化配信を含む。デリタゲ配信は対象外",
    "・対象の配信：ALL配信のみ",
    "・対象クライアント：全クライアント",
    "・最低発注金額：200万円〜",
]]
CAUTION = ["注意点", [
    "・発注時、備考欄に次の文言を記載\x0b　「FY26 Winter Campaign適用」",
    "・配信中・発注受領済みの案件には適用できない",
    "・他の割引・キャンペーンとは併用できない",
    "・記載のない条件は、最新の媒体資料に準拠",
]]


def fill(card, heading, lines):
    tf = card.text_frame
    p0, p1 = tf.paragraphs[:2]
    p0.runs[0].text = heading
    for r in p0.runs[1:]:
        r._r.getparent().remove(r._r)
    tmpl = copy.deepcopy(p1._p)
    for p in tf.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    prev = p0._p
    for t in lines:
        np_ = copy.deepcopy(tmpl)
        prev.addnext(np_)
        prev = np_
    for p, t in zip(tf.paragraphs[1:], lines):
        for r in p.runs[1:]:
            r._r.getparent().remove(r._r)
        for br in p._p.findall(qn("a:br")):
            br.getparent().remove(br)
        parts = t.split("\x0b")
        p.runs[0].text = parts[0]
        p.runs[0].font.size = Pt(12)
        if len(parts) > 1:
            r0 = p.runs[0]._r
            br = p._p.makeelement(qn("a:br"), {})
            rpr = r0.find(qn("a:rPr"))
            if rpr is not None:
                br.append(copy.deepcopy(rpr))
            r2 = copy.deepcopy(r0)
            r2.find(qn("a:t")).text = parts[1]
            r0.addnext(br)
            br.addnext(r2)
        p.space_before = Pt(0)
    for r in p0.runs:
        r.font.size = Pt(18)
    tf._txBody.find(qn('a:bodyPr')).set('anchor', 't')


fill(cond, *COND)
fill(caution, *CAUTION)

CARD_H = 0.3 + 18 * 1.3 / 72 + 7 * (12 * 1.5 / 72)
total = TILE_H + 0.15 + CARD_H
top = area_top + (7.05 - area_top - total) / 2
for i, t in enumerate(tiles):
    for r in t.text_frame.paragraphs[0].runs:
        r.font.size = Pt(24)
    t.left, t.top, t.width, t.height = Inches(L + i * (tw + GAP)), Inches(top), Inches(tw), Inches(TILE_H)
cw = (W - GAP) / 2
for i, c in enumerate((cond, caution)):
    c.left, c.top, c.width, c.height = Inches(L + i * (cw + GAP)), Inches(top + TILE_H + 0.15), Inches(cw), Inches(CARD_H)

engine.set_funnel_highlight(s9, ["認知"])
engine.force_font(prs)
prs.save(F)
print("saved", F, round(CARD_H, 2), round(top, 2))
