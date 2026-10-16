# -*- coding: utf-8 -*-
"""タナカホーム v5：v4（7枚）に、殿村さんの「提案の流れ」（2026-10-16）に対応するページを足し、Appendixを付ける。

    python -m _build.patch_tanaka_v5 <v4.pptx> <あらうちv10.pptx（Appendixの元）> <出力.pptx>

  本編：表紙／提案の流れ／
        1 現状と打ち手（v4）・LINE公式アカウントの現状（新）／
        2 課題（v4を手直し）・前後検索「注文住宅」（新）／
        3 解決策 追う（v4）・友だちの入口（新）・配信（v4）・来場前のお知らせ（新）・試算（新・SIM ver1.2）・費用（新）／
        4 今後の展開（v4）・スケジュール（新・Googleトレンド）／裏表紙
  Appendix（あらうちv10から写す）：LINE公式アカウントのメリット（34〜37）／中長期の施策（39〜45）／スポットでのご依頼（31）
    31（都度発注）はサンクスLINE誘導ツールを 初期15万円・月額3万円〜、離脱防止ツールの月額を3万円に直す。
    38（「業態平均で最大3%」＝出典なし）・33（定期購入・クロスセルの話）は入れない。
"""
import copy
import csv
import io
import sys
from collections import defaultdict
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.parts.image import ImagePart
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from _build.arauchi_v4.lib import NAVY, RED, TEXT, add_box, add_table
from _build.arauchi_v10.build import TINT, LINE_C, by_id, meiryo, runs
from _build.build_tanaka_v1 import GRAY, head_cell, cell, label_cell, note, band
from _build.build_tanaka_v2 import heading, flow, X0, W
from _build.patch_tanaka_v4 import new_page, set_lines, set_title_lead, HDR

ROOT = Path(__file__).resolve().parent.parent
IMG = ROOT / "_images"
LIGHT = "C9CEDF"
SOFT = "F2F2F2"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

# SIM ver1.2（殿村さんのPC版・佐村さんに確認中）を LibreOffice で再計算した値（1〜6か月目・整数に丸め）
SIM = dict(new=[177, 177, 179, 180, 181, 182], total=[594, 772, 950, 1130, 1311, 1493],
           active=[428, 551, 673, 793, 912, 1030], click=[477, 557, 637, 613, 676, 738],
           doc=[4, 6, 6, 8, 9, 11], visit=[2, 3, 4, 6, 7, 8])
APPX = [24, 34, 35, 36, 37, 39, 40, 41, 42, 43, 44, 45, 31]   # あらうちv10のページ番号（この順に並べる）


def pic(s, path, x, y, w, h=None, name=None):
    p = s.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h) if h else None)
    if name:
        p.name = name
    return p


def caption(s, x, y, w, text):
    add_box(s, x, y, w, 0.28, [runs((text, dict(size=9, color=GRAY)))], align="ctr", anchor="t", ins=(0, 0.02, 0, 0),
            name="画像の説明")


# ---------------------------------------------------------------- 2 提案の流れ
def p_nagare(prs, src):
    s = new_page(prs, src, "提案の流れ", "")
    set_lines(list(s.shapes)[4], ["資料請求の時点からLINEでつながり、来場・契約まで案内する。", "青森店で確立し、他店舗へ広げることも見据える"])
    items = [("1. 現状と打ち手", "来場予約は目標50〜60件に対し35〜40件。広告と地域SEOで質の高い見込み客を集める。LINEは来場予約のご案内が中心"),
             ("2. 課題", "来場を決める前の検討中の方と、LINEでつながれていない。検討期間の約2週間の接点が、来場を左右する"),
             ("3. 解決策", "資料請求の時点から来場・契約までを追う。前後検索に沿って配信し、来場前日のお知らせでキャンセルも減らす"),
             ("4. 今後の展開", "青森店で型をつくり、全店に広げる")]
    rows = [[label_cell(k, 12), cell(v, 11.5)] for k, v in items]
    add_table(s, X0, 1.85, [2.1, W - 2.1], [0.78] * 4, rows, name="提案の流れ")
    rows = [[cell("Appendix", 11, True, GRAY, SOFT, "ctr"),
             cell("LINE公式アカウントのメリット／中長期の施策／スポットでのご依頼", 11, color=GRAY)]]
    add_table(s, X0, 5.3, [2.1, W - 2.1], [0.5], rows, name="Appendix")
    return s


# ---------------------------------------------------------------- 3 現状と打ち手：LINEの行は次のページへ
def p_genjo(s):
    gf = [sh for sh in s.shapes if sh.has_table][0]
    tbl = gf._element.find(".//" + qn("a:tbl"))
    trs = tbl.findall(qn("a:tr"))
    assert "LINE" in "".join(t.text or "" for t in trs[-1].iter(qn("a:t")))
    tbl.remove(trs[-1])
    for tr in trs[:-1]:
        tr.set("h", str(Inches(0.624)))
    gf.height = Inches(0.624 * 5)


# ---------------------------------------------------------------- 4 LINE公式アカウントの現状（新）
def p_line(prs, src):
    s = new_page(prs, src, "LINE公式アカウントの現状：来場を決めた方へのご案内は整っている",
                 "実写真の配信と来場特典のご案内で、来場予約までの流れができている")
    rows = [[label_cell("友だち数"), cell("417人", 11.5, True, NAVY)],
            [label_cell("友だちの入口"), cell(["サイトのポップアップ「LINEお友だち登録キャンペーン」",
                                             "来場特典 最大7,000円分（予約来場 5,000円分＋友だち登録 2,000円分）"], 10.5)],
            [label_cell("配信"), cell(["見学会・来場予約のご案内が中心",
                                       "実写真の画像とカルーセルで、見学会の中身が伝わる作り"], 10.5)],
            [label_cell("リッチメニュー"), cell("予約／ホームページ／公式SNS", 10.5)],
            [label_cell("1:1チャット"), cell("担当者が返信する設定。個別のご相談を受けられる", 10.5)]]
    add_table(s, X0, 1.85, [1.5, 4.55], [0.5, 0.78, 0.78, 0.5, 0.62], rows, name="LINEの現状")
    pic(s, IMG / "tanaka_line_talk.png", 6.95, 1.85, 1.5, 3.25, name="LINEの配信")
    caption(s, 6.75, 5.15, 1.9, "LINEの配信")
    pic(s, IMG / "tanaka_line_popup.png", 8.65, 1.85, 1.63, 1.69, name="サイトのポップアップ")
    caption(s, 8.55, 3.6, 1.83, "サイトのポップアップ")
    note(s, "※2026年10月時点で、貴社のLINE公式アカウントとサイトを弊社で確認", y=6.4)
    return s


# ---------------------------------------------------------------- 5 課題：2週間の話は次のページ（前後検索）へ
def p_kadai(s):
    set_lines(by_id(s, HDR[4]), ["サイトに来た方の多くは、来場予約も資料請求もせずに離れている"])
    for sid, text in [(14, "来場を決めるまでの間、複数の会社を比べ、費用や性能を調べ続ける"),
                      (16, "LINEで見学会・来場予約をご案内できている")]:
        ps = by_id(s, sid)._element.findall(".//" + qn("a:p"))
        rs = ps[1].findall(qn("a:r"))
        for r in rs[1:]:
            ps[1].remove(r)
        rs[0].find(qn("a:t")).text = text


# ---------------------------------------------------------------- 6 前後検索「注文住宅」（新）
def p_zengo(prs, src):
    s = new_page(prs, src, "前後検索「注文住宅」：当日に大手を比べ、その後は住宅ローン・性能を調べる",
                 "検討期間の約2週間の接点が、来場を左右する")
    iw, ih = 7.2, 7.2 * 1167 / 2000
    ix, iy = (10.83 - iw) / 2, 1.68
    pic(s, IMG / "tanaka_zengo_chumon.png", ix, iy, iw, ih, name="前後検索 注文住宅")
    px = lambda v: ix + v / 2000 * iw                     # 画像の横のピクセル → インチ（-15日=97／0日=1027／1日=1260／7日=1723／15日=1955）
    bars = [(97, 1027, "検索前", "中古一戸建て・平屋の間取り・家を建てる費用", False),
            (1027, 1260, "当日", "大手ハウスメーカー・相場・展示場", True),
            (1260, 1723, "1〜7日後", "住宅ローン・ほかの住宅会社・間取り", True),
            (1723, 1955, "7〜15日後", "金利・断熱等級・ルームツアー", True)]
    by = iy + ih + 0.04
    for a, b, h1, t, on in bars:
        x, w = px(a) + 0.02, px(b) - px(a) - 0.04
        col = "FFFFFF" if on else GRAY
        add_box(s, x, by, w, 0.62, [runs((h1, dict(size=8.5, bold=True, color=col))), runs((t, dict(size=8, color=col)))],
                fill=NAVY if on else SOFT, align="ctr", anchor="ctr", line_spacing=105, ins=(0.03, 0.02, 0.03, 0.02),
                name=f"時期 {h1}")
    note(s, "※LINEヤフー社の前後検索データより。「注文住宅」を検索した方が、前後15日間に検索した言葉（縦軸＝検索した人数）",
         y=by + 0.68)
    return s


# ---------------------------------------------------------------- 8 友だちの入口（新）
def p_iriguchi(prs, src):
    s = new_page(prs, src, "友だちの入口：来場特典に加え、家づくりの情報で友だちになっていただく",
                 "サイトを離れる前、資料請求の直後など、検討中の方が友だちになる理由を用意する")
    g = dict(color=GRAY)
    rows = [[head_cell(""), head_cell("入口"), head_cell("タイミング"), head_cell("友だちになる理由")],
            [cell("今", 11, True, GRAY, SOFT, "ctr"), cell("来場特典のキャンペーン", 10.5, **g),
             cell("来場を決めたとき", 10.5, **g), cell("来場特典の上乗せ", 10.5, **g)],
            [cell("追加", 11, True, NAVY, TINT, "ctr"), cell("離脱防止", 11, True, NAVY),
             cell("サイトを閉じようとしたとき", 10.5), cell("家づくりの情報（例：平屋の間取り、費用の目安）をLINEでお届け", 10.5)],
            [cell("追加", 11, True, NAVY, TINT, "ctr"), cell("サイトのLINEボタン", 11, True, NAVY),
             cell("施工事例・見学会のページを見ているとき", 10.5), cell("見学会や新しい施工事例のお知らせ", 10.5)],
            [cell("追加", 11, True, NAVY, TINT, "ctr"), cell("サンクスLINE誘導", 11, True, NAVY),
             cell("資料請求・来場予約が完了した直後", 10.5), cell("資料の続き（ルームツアー・費用・土地）、来場前のお知らせ", 10.5)]]
    add_table(s, X0, 1.85, [0.8, 2.2, 2.55, 4.18], [0.4, 0.5, 0.55, 0.55, 0.55], rows, name="友だちの入口")
    # 実績（「来場前のお知らせ」の実績と同じ位置・大きさ。数字はサンクスLINE誘導のご提案 Appendix「導入・活用実績」より）
    y = 4.55
    add_box(s, X0, y, 1.9, 1.0, [runs(("弊社の導入実績", dict(size=11, bold=True, color="FFFFFF")))], fill=NAVY,
            align="ctr", anchor="ctr", name="実績 見出し")
    add_box(s, X0 + 1.9, y, 4.6, 1.0, [runs(("人材業界", dict(size=11, bold=True, color=NAVY))),
                                       runs(("サンクスLINE誘導ツールを導入", dict(size=10.5))),
                                       runs(("友だち増加数（3か月）1,532人 → 2,103人", dict(size=10.5)))],
            fill=TINT, anchor="ctr", line_spacing=120, ins=(0.15, 0.06, 0.1, 0.06), name="実績 施策")
    add_box(s, X0 + 6.5, y, W - 6.5, 1.0, [runs(("友だち増加数", dict(size=10.5, bold=True, color=NAVY))),
                                          runs(("導入前比137.3%", dict(size=20, bold=True, color=NAVY))),
                                          runs(("前年比127.5%", dict(size=12, bold=True, color=NAVY)))],
            fill=TINT, align="ctr", anchor="ctr", line_spacing=100, name="実績 数字")
    note(s, "※導入前の3か月・前年の同じ期間と比べた友だち増加数（他の施策や時期の影響も含む）", y=5.65)
    note(s, "※来場予約の完了ページでは、ウェブ予約特典の事前アンケートに回答いただいた後に、LINEをご案内する（特典の条件を損なわないため）",
         y=5.95, h=0.45)
    return s


# ---------------------------------------------------------------- 10 来場前のお知らせ（新）
def p_remind(prs, src):
    s = new_page(prs, src, "来場前のお知らせ：予約から来場までをLINEでつなぎ、キャンセルを減らす",
                 "日程の変更もLINEのチャットで受け付け、キャンセルではなく変更につなげる")
    steps = [("予約の完了", "事前アンケートの回答後に、LINEの友だち追加をご案内"),
             ("予約の確認", "日時・展示場・担当者をLINEでお知らせ"),
             ("3日前", "当日の見どころ・流れ・所要時間をご案内"),
             ("前日", "日時・場所・地図・駐車場をお知らせ。日程の変更はチャットで受付"),
             ("来場の後", "お礼と、次のご相談（商談）のご案内")]
    flow(s, [(h, NAVY, "FFFFFF") for h, _ in steps], 1.85, 0.6, arrow_w=0.25, size=11.5)
    n, aw = len(steps), 0.25
    bw = (W - aw * (n - 1)) / n
    for k, (_, t) in enumerate(steps):
        add_box(s, X0 + k * (bw + aw), 2.53, bw, 1.55, [runs((t, dict(size=11, color=TEXT)))], line=LINE_C,
                anchor="t", line_spacing=125, ins=(0.1, 0.1, 0.1, 0.06), name=f"お知らせ{k + 1}")
    # 実績（業種 → 施策 → 数字）
    y = 4.55
    add_box(s, X0, y, 1.9, 1.0, [runs(("弊社運用の実績", dict(size=11, bold=True, color="FFFFFF")))], fill=NAVY,
            align="ctr", anchor="ctr", name="実績 見出し")
    add_box(s, X0 + 1.9, y, 4.6, 1.0, [runs(("美容クリニック", dict(size=11, bold=True, color=NAVY))),
                                       runs(("予約の完了後にLINEへご案内し、予約の確認・前日のお知らせを配信", dict(size=10.5)))],
            fill=TINT, anchor="ctr", line_spacing=120, ins=(0.15, 0.06, 0.1, 0.06), name="実績 施策")
    add_box(s, X0 + 6.5, y, W - 6.5, 1.0, [runs(("予約後の来院率", dict(size=10.5, bold=True, color=NAVY))),
                                          runs(("40〜50%改善", dict(size=22, bold=True, color=NAVY)))],
            fill=TINT, align="ctr", anchor="ctr", line_spacing=105, name="実績 数字")
    note(s, "※業種・予約の条件が異なるため、数字は参考値", y=5.65)
    return s


# ---------------------------------------------------------------- 11 試算（新）
def p_sim(prs, src):
    s = new_page(prs, src, f"試算：6か月目に、LINE経由の資料請求 月{SIM['doc'][5]}件・来場予約 月{SIM['visit'][5]}件",
                 "サイトから友だちを増やし、検討の順番に沿った配信で、資料請求・来場予約につなげる")
    f = lambda v: f"{v:,}"
    head = [head_cell("")] + [head_cell(f"{i}か月目", 11) for i in range(1, 7)]
    rows = [head]
    for label, key, strong in [("新しく友だちになる方（人）", "new", False), ("友だち数（累計・人）", "total", False),
                               ("配信が届く友だち（人）", "active", False), ("配信からのクリック", "click", False),
                               ("資料請求（LINE経由・件）", "doc", True), ("来場予約（LINE経由・件）", "visit", True)]:
        rows.append([label_cell(label, 10.5)] + [cell(f(v), 12 if strong else 11, strong, NAVY if strong else TEXT,
                                                      TINT if strong else "FFFFFF", "ctr") for v in SIM[key]])
    add_table(s, X0, 1.85, [2.65] + [(W - 2.65) / 6] * 6, [0.4] + [0.5] * 6, rows, name="試算")
    note(s, "※前提：サイトのUU 月13,228／今の友だち417人／友だちの入口＝離脱防止・サイトのLINEボタン・サンクスLINE誘導"
            "／配信＝企画配信 月4本＋ステップ配信5通／ブロック率は28%から月ごとに上昇", y=5.4, h=0.45)
    note(s, "※配信が届く友だち＝ブロックしていない友だち。資料請求・来場予約は、LINEの配信からのクリックを通ったものだけを数える",
         y=5.88, h=0.45)
    return s


# ---------------------------------------------------------------- 12 費用（新）
def p_hiyo(prs, src):
    s = new_page(prs, src, "費用：初期26.5万円〜・月額26万円〜（6か月〜）",
                 "LINE運用コンサルに、友だちの入口をつくる2つのツールを組み合わせる")
    rows = [[head_cell("項目"), head_cell("初期"), head_cell("月額"), head_cell("内容")],
            [label_cell("LINE運用コンサル", 11), cell("10万円〜", 12, True, NAVY, align="ctr"),
             cell("20万円〜", 12, True, NAVY, align="ctr"),
             cell("入口と配信の設計、配信の制作（企画配信・ステップ配信・来場前のお知らせ）、結果の分析と改善", 10.5)],
            [label_cell("離脱防止", 11), cell("1.5万円", 12, True, NAVY, align="ctr"), cell("3万円", 12, True, NAVY, align="ctr"),
             cell("サイトを閉じようとした方へのLINEのご案内", 10.5)],
            [label_cell("サンクスLINE誘導", 11), cell("15万円", 12, True, NAVY, align="ctr"),
             cell("3万円〜", 12, True, NAVY, align="ctr"), cell("資料請求・来場予約の完了ページからLINEへのご案内", 10.5)],
            [cell("合計", 11.5, True, "FFFFFF", NAVY, "ctr"), cell("26.5万円〜", 13, True, "FFFFFF", NAVY, "ctr"),
             cell("26万円〜", 13, True, "FFFFFF", NAVY, "ctr"), cell("", 10.5, fill=NAVY)],
            [label_cell("LINE公式アカウント\n利用料", 10), cell("―", 11, align="ctr", color=GRAY),
             cell("LINEヤフー社の料金", 9.5, align="ctr", color=GRAY), cell("貴社のご契約のまま（配信数に応じた料金）", 10.5, color=GRAY)]]
    add_table(s, X0, 1.85, [2.0, 1.45, 1.45, 4.83], [0.4, 0.72, 0.55, 0.55, 0.55, 0.62], rows, name="費用")
    note(s, "※契約期間は6か月〜", y=5.45)
    note(s, "※施策内容・配信本数・連携範囲等により、費用は個別にお見積もりいたします。", y=5.75)
    return s


# ---------------------------------------------------------------- 14 スケジュール（新）
def trend_avg():
    d = defaultdict(list)
    for r in list(csv.reader(open(ROOT / "_data/trends/tanaka_aomori_hiraya_10year_monthly.csv")))[1:]:
        if r[0] >= "2021-10":
            d[int(r[0][5:7])].append(float(r[1]))
    return [round(sum(d[m]) / len(d[m])) for m in range(1, 13)]


def p_schedule(prs, src):
    s = new_page(prs, src, "スケジュール：検索が最も増える年明けまでに、LINEの仕組みを整える",
                 "11月から設計・構築を進め、1月から運用を始める")
    heading(s, X0, 1.85, 4.45, "「平屋」の検索（青森県・月別）")
    vals = trend_avg()
    cd = CategoryChartData()
    cd.categories = [f"{m}月" for m in range(1, 13)]
    cd.add_series("平屋", vals)
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(X0), Inches(2.3), Inches(4.45), Inches(3.0), cd)
    gf.name = "検索の月別"
    ch = gf.chart
    ch.has_legend = False
    ch.has_title = False
    ch.font.size = Pt(9)
    ch.font.name = "メイリオ"
    ch.font.color.rgb = RGBColor.from_string(TEXT)
    pl = ch.plots[0]
    pl.gap_width = 60
    pl.has_data_labels = True
    dl = pl.data_labels
    dl.font.size = Pt(9)
    dl.font.name = "メイリオ"
    dl.number_format = "0"
    dl.number_format_is_linked = False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    ser = pl.series[0]
    for i in range(12):
        pt = ser.points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = RGBColor.from_string(NAVY if i in (0, 1) else LIGHT)
    va = ch.value_axis
    va.visible = False
    va.has_major_gridlines = False
    va.maximum_scale = 100
    va.minimum_scale = 0
    ca = ch.category_axis
    ca.tick_labels.font.size = Pt(9)
    ca.format.line.color.rgb = RGBColor.from_string(LINE_C)
    heading(s, 5.25, 1.85, W - 4.7, "進め方")
    rows = [[label_cell("11月", 11), cell("設計：友だちの入口、配信の内容と順番、使う写真・動画の確認", 10.5)],
            [label_cell("12月", 11), cell("構築：離脱防止・サイトのLINEボタン・サンクスLINE誘導、ステップ配信、来場前のお知らせ", 10.5)],
            [label_cell("1月〜", 11), cell("運用開始：検索が増える時期に、友だちを集めて配信する", 10.5, True, NAVY)],
            [label_cell("3月〜", 11), cell("改善：来場・契約の結果を見て、配信と広告を見直す。他店舗への展開を準備", 10.5)]]
    add_table(s, 5.25, 2.3, [0.95, W - 4.7 - 0.95], [0.75] * 4, rows, name="進め方")
    note(s, "※Googleトレンドより。2016年10月〜2026年9月で最も多い月を100とした値を、直近5年（2021年10月〜2026年9月）の月ごとに平均。"
            "全国の「注文住宅」も1〜2月が最も多い", y=5.5, h=0.45)
    note(s, "※スケジュールは、ご契約の時期により前後する", y=5.98)
    return s


# ---------------------------------------------------------------- Appendix（あらうちv10から写す）
def copy_slide(dst, ss):
    layouts = {l.name: l for l in dst.slide_layouts}
    ds = dst.slides.add_slide(layouts[ss.slide_layout.name])
    for sh in list(ds.shapes):
        sh._element.getparent().remove(sh._element)
    rmap = {}
    for rId, rel in ss.part.rels.items():
        if rel.reltype in (RT.SLIDE_LAYOUT, RT.NOTES_SLIDE):
            continue
        if rel.is_external:
            rmap[rId] = ds.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
        elif rel.reltype == RT.IMAGE:
            try:
                _, rmap[rId] = ds.part.get_or_add_image_part(io.BytesIO(rel.target_part.blob))
            except Exception:                              # EMF・SVG など（PILで読めない画像）はそのまま新しい名前で登録
                tp = rel.target_part
                ext = str(tp.partname).rsplit(".", 1)[-1]
                name = dst.part.package.next_image_partname(ext)
                rmap[rId] = ds.part.relate_to(ImagePart(name, tp.content_type, dst.part.package, tp.blob), RT.IMAGE)
        else:
            raise SystemExit(f"未対応の関係: {rel.reltype}")
    tree = ds.shapes._spTree
    for el in ss.shapes._spTree:
        if el.tag.endswith("}nvGrpSpPr") or el.tag.endswith("}grpSpPr"):
            continue
        tree.append(copy.deepcopy(el))
    for el in tree.iter():
        for k, v in list(el.attrib.items()):
            if k.startswith("{%s}" % R_NS) and v in rmap:
                el.set(k, rmap[v])
    bg = ss._element.cSld.bg
    if bg is not None:
        ds._element.cSld.insert(0, copy.deepcopy(bg))
    return ds


def put(c, text, tmpl, ppr):
    """表のセルの文字を差し替える（書式は tmpl＝文字列 a:r の写しから）"""
    txb = c._tc.find(qn("a:txBody"))
    ps = txb.findall(qn("a:p"))
    for p in ps[1:]:
        txb.remove(p)
    p = ps[0]
    for r in p.findall(qn("a:r")) + p.findall(qn("a:br")):
        p.remove(r)
    old = p.find(qn("a:pPr"))
    if old is not None:
        p.remove(old)
    p.insert(0, copy.deepcopy(ppr))
    r = copy.deepcopy(tmpl)
    r.find(qn("a:t")).text = text
    end = p.find(qn("a:endParaRPr"))
    if end is not None:
        end.addprevious(r)
    else:
        p.append(r)


def fix_spot(s):
    tbl = [sh for sh in s.shapes if sh.has_table][0].table
    rows = {r.cells[4].text: r for r in tbl.rows}
    src_tc = rows["サンクスLINE誘導ツール"].cells[5]._tc
    tmpl = copy.deepcopy(src_tc.find(".//" + qn("a:r")))
    ppr = copy.deepcopy(src_tc.find(".//" + qn("a:pPr")))
    for c, t in [(rows["サンクスLINE誘導ツール"].cells[5], "¥150,000"), (rows["サンクスLINE誘導ツール"].cells[6], "¥30,000~"),
                 (rows["離脱防止ツール"].cells[6], "¥30,000")]:
        put(c, t, tmpl, ppr)
    hit = [sh for sh in s.shapes if sh.has_text_frame and "受けさせていただきます" in sh.text_frame.text]
    assert len(hit) == 1
    rs = list(hit[0]._element.iter(qn("a:r")))
    for r in rs[1:]:
        r.getparent().remove(r)
    rs[0].find(qn("a:t")).text = "貴社リソースの状況に応じて、都度発注でのご依頼も各単価で対応可能"


def main(base, appx_src, out):
    prs = Presentation(base)
    sl = list(prs.slides)
    assert len(sl) == 7 and by_id(sl[3], 2).text_frame.text.startswith("解決策：資料請求の時点から")
    cover, genjo, kadai, ou, haishin, tenkai, back = sl
    src = kadai
    nagare = p_nagare(prs, src)
    p_genjo(genjo)
    line = p_line(prs, src)
    zengo = p_zengo(prs, src)
    iriguchi = p_iriguchi(prs, src)
    remind = p_remind(prs, src)
    sim = p_sim(prs, src)
    hiyo = p_hiyo(prs, src)
    sched = p_schedule(prs, src)
    p_kadai(kadai)
    for s in (nagare, line, zengo, iriguchi, remind, sim, hiyo, sched, kadai):
        for sh in s.shapes:
            if sh.has_text_frame:
                meiryo(sh)
    ar = Presentation(appx_src)
    asl = list(ar.slides)
    assert asl[23].shapes[0].text_frame.text == "Appendix" if asl[23].shapes[0].has_text_frame else True
    # 表紙：DYMの標準の形（あらうちv10の表紙＝サブタイトル2行＋タイトル「LINE公式アカウント運用のご提案」）にそろえる（2026-10-16 殿村さん）
    cover = copy_slide(prs, asl[0])
    set_lines(by_id(cover, 31), ["タナカホーム株式会社 御中"])
    set_lines(by_id(cover, 26), ["LINE公式アカウント運用のご提案"])
    set_lines(by_id(cover, 27), ["集めた見込み客を逃さずナーチャリングをし、", "来場・契約へ変えるLINE仕組みの構築"])
    for sid in (31, 26, 27):
        meiryo(by_id(cover, sid))
    appx = []
    for n in APPX:
        ds = copy_slide(prs, asl[n - 1])
        if n == 31:
            fix_spot(ds)
        appx.append(ds)
    keep = [cover, nagare, genjo, line, kadai, zengo, ou, iriguchi, haishin, remind, sim, hiyo, tenkai, sched, back] + appx
    lst = prs.slides._sldIdLst
    by_part = {prs.part.related_part(e.rId): e for e in lst}
    for e in list(lst):
        lst.remove(e)
    for s in keep:
        lst.append(by_part[s.part])
    prs.save(out)
    print("saved:", out, len(prs.slides), "枚")


if __name__ == "__main__":
    main(*sys.argv[1:4])
