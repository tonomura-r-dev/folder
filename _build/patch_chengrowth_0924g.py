# -*- coding: utf-8 -*-
"""チェングロウスver1.1：人材（転職）版PDF（41p）の考え方を反映（2026-09-24 殿村さん「全部進めて」）。
- S4  本提案の考え方：下段に「LINE導入の3つのメリット」を追加（上の図は詰めて残す）
- S13 汎用の想定動線 → カスタマージャーニー8フェーズ（整備士の転職版）
- S14 業界平均値 → 施策パターン3つ（A 会員登録型／B 日程調整型★本命／C LINE内完結型）
- S16 費用対効果：損益分岐の一言を追加
- S22 その他 → 通知メッセージの使いどころ＋特殊動線
- S38/S40/S42/S43（汎用の検討背景）→ 改善モデル①〜④（整備士版）。S39・S41は残し、並べ替え
  python3 _build/patch_chengrowth_0924g.py <入力pptx> <出力pptx>
"""
import sys
from pptx import Presentation
from pptx.util import Cm, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn

SRC, OUT = sys.argv[1], sys.argv[2]
prs = Presentation(SRC)
NAVY, INK, GRAY = "1F285A", "333333", "7F7F7F"
LBLUE, BEIGE, LGREEN, LGRAY, GREEN = "DDEBF7", "FFF2CC", "E2F0D9", "F2F2F2", "06C755"
ORANGE_L, BLUE_L = "F8CBAD", "BDD7EE"


def meiryo(run):
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        e = rPr.find(qn(tag))
        if e is None:
            e = rPr.makeelement(qn(tag), {})
            rPr.append(e)
        e.set("typeface", "メイリオ")


def box(slide, x, y, w, h, fill, lines, color=INK, size=10.5, bold=False, anchor=MSO_ANCHOR.MIDDLE,
        align=PP_ALIGN.CENTER, line=None, lw=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE, ml=0.15):
    sp = slide.shapes.add_shape(shape, Cm(x), Cm(y), Cm(w), Cm(h))
    if fill:
        sp.fill.solid()
        sp.fill.fore_color.rgb = RGBColor.from_string(fill)
    else:
        sp.fill.background()
    if line:
        sp.line.color.rgb = RGBColor.from_string(line)
        sp.line.width = Pt(lw)
    else:
        sp.line.fill.background()
    sp.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sp.adjustments[0] = 0.08
    tf = sp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Cm(ml)
    tf.margin_top = tf.margin_bottom = Cm(0.06)
    for i, ln in enumerate(lines):
        text, sz, b, c = (ln, size, bold, color) if isinstance(ln, str) else ln
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = text
        r.font.size = Pt(sz)
        r.font.bold = b
        r.font.color.rgb = RGBColor.from_string(c)
        meiryo(r)
    return sp


def text(slide, x, y, w, h, lines, size=11, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    return box(slide, x, y, w, h, None, lines, color=color, size=size, bold=bold, align=align, anchor=anchor,
               shape=MSO_SHAPE.RECTANGLE, ml=0.05)


def arrow(slide, x1, y1, x2, y2, color="8C8C8C", w=1.25):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Cm(x1), Cm(y1), Cm(x2), Cm(y2))
    c.line.color.rgb = RGBColor.from_string(color)
    c.line.width = Pt(w)
    ln = c.line._get_or_add_ln()
    ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle"}))


def clear(slide, title):
    """タイトルと区切り線だけ残して空にし、タイトルを差し替える"""
    tree = slide.shapes._spTree
    tshape = None
    for sh in list(slide.shapes):
        is_title = sh.has_text_frame and sh.top is not None and sh.top < Cm(1.2) and sh.text_frame.text.strip()
        is_div = sh.shape_type == 9 and abs(sh.top - Cm(3.86)) < Cm(0.1) and sh.width > Cm(20)
        if is_title and tshape is None:
            tshape = sh
            continue
        if is_div:
            continue
        tree.remove(sh._element)
    if tshape is None:
        text(slide, 1.52, 0.38, 24.4, 0.94, [(title, 16, True, "002060")], anchor=MSO_ANCHOR.MIDDLE)
        return
    tf = tshape.text_frame
    for p in tf.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    runs = tf.paragraphs[0].runs
    runs[0].text = title
    for r in runs[1:]:
        r.text = ""
    if not any(sh.shape_type == 9 for sh in slide.shapes):
        c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, 0, Cm(3.86), Cm(27.52), Cm(3.86))
        c.line.color.rgb = RGBColor.from_string("D9D9D9")


def lead(slide, lines):
    text(slide, 1.2, 1.85, 25.1, 1.85, [(l, 13, False, INK) for l in lines], anchor=MSO_ANCHOR.MIDDLE)


def band(slide, y, msg, h=0.95):
    box(slide, 1.2, y, 25.12, h, NAVY, [(msg, 13, True, "FFFFFF")])


def by_id(slide, sid):
    return [sh for sh in slide.shapes if sh.shape_id == sid][0]


PHASES = ["①接触", "②離脱", "③興味喚起", "④友だち追加", "⑤面談", "⑥求人紹介", "⑦内定承諾", "⑧定着・紹介"]

# ================= S4 本提案の考え方：3つのメリット =================
s = prs.slides[3]
moves = {5: 7.25}
for sid, top in moves.items():
    by_id(s, sid).top = Cm(top)
for sid in range(6, 11):
    by_id(s, sid).top = by_id(s, sid).top - Cm(0.6)
for sid in range(11, 21):
    by_id(s, sid).top = by_id(s, sid).top - Cm(1.6)
text(s, 1.2, 13.05, 25.1, 0.7, [("LINE導入で得られる3つのメリット", 12.5, True, NAVY)])
MERITS = [
    ("① 広告費のロスを減らす", "サイトに来て応募せず帰る整備士を、離脱防止からLINEで拾う"),
    ("② 面談・内定の辞退を防ぐ", "面談前日のリマインド、内定後の不安にLINEですぐ答える"),
    ("③ 連絡をLINEに一本化", "電話は平日10〜18時のみ。在職中の整備士とは夜のLINEでつながる"),
]
cw = (25.12 - 0.4) / 3
for i, (h, b) in enumerate(MERITS):
    box(s, 1.2 + i * (cw + 0.2), 13.85, cw, 3.3, LBLUE, [(h, 12, True, NAVY), (b, 10.5, False, INK)])

# ================= S13 カスタマージャーニー =================
s = prs.slides[12]
clear(s, "全体設計｜整備士の転職カスタマージャーニー")
lead(s, ["整備士の転職の流れにあわせて、LINEの施策を配置します。",
         "主戦場は④友だち追加ではなく⑤面談。そこから先も同じLINEで就業・定着まで運びます。"])
n, gap = 8, 0.15
cw = (25.12 - gap * (n - 1)) / n
x0 = lambda i: 1.2 + i * (cw + gap)
box(s, x0(0), 4.2, cw * 4 + gap * 3, 0.65, BEIGE, [("WEB → LINE：どう友だちにするか", 10.5, True, INK)])
box(s, x0(4), 4.2, cw * 4 + gap * 3, 0.65, LGREEN, [("LINE → 就業：どう面談・就業につなげるか", 10.5, True, INK)])
STATES = ["年収・職種が気になり検索する", "応募はまだ。そのまま帰る", "年収・働き方を知りたい", "30秒の診断なら答える",
          "電話は無理。夜なら話せる", "自分に合う求人か見極めたい", "本当にこの会社でいいか不安", "入社後のギャップが不安"]
ACTS = ["CPF広告（Meta予算を転用）", "離脱防止ポップアップ", "年収相場・職種図鑑の配信", "あいさつ＋30秒適職診断",
        "LINEで面談予約＋前日リマインド", "地域別の求人ピックアップ", "不安に即答する自動応答・体験談", "定着フォロー・友だち紹介"]
for i, ph in enumerate(PHASES):
    fill = NAVY if i == 4 else (ORANGE_L if i < 4 else BLUE_L)
    col = "FFFFFF" if i == 4 else INK
    box(s, x0(i), 5.05, cw, 1.1, fill, [(ph, 10, True, col)], shape=MSO_SHAPE.CHEVRON if i else MSO_SHAPE.PENTAGON, ml=0.1)
text(s, 1.2, 6.35, 6, 0.55, [("ユーザーの状態", 10, True, GRAY)])
text(s, 1.2, 9.55, 6, 0.55, [("LINEの施策", 10, True, GRAY)])
for i in range(n):
    box(s, x0(i), 6.9, cw, 2.5, LGRAY, [(STATES[i], 10, False, INK)])
    hl = i == 4
    box(s, x0(i), 10.1, cw, 3.0, NAVY if hl else LBLUE, [(ACTS[i], 10.5, True, "FFFFFF" if hl else NAVY)])
band(s, 13.75, "①〜④で友だちにし、⑤面談を成果（CV）に。⑥〜⑧も同じLINEで就業・定着まで運ぶ。")
text(s, 1.2, 15.0, 25.1, 0.6, [("※ 各フェーズの具体策は、改善モデル①〜④（P40〜43）", 9.5, False, GRAY)])

# ================= S14 施策パターン3つ =================
s = prs.slides[13]
clear(s, "施策の考え方｜どこでLINEにつなぐか（3パターン）")
lead(s, ["LINEにつなぐ地点で、3つの型に分かれます。",
         "チェングロウスは「電話なしで面談まで進める」B 日程調整型が本命です。"])
PAT = [
    ("パターンA", "会員登録型", ["求人応募・登録", "面談", "求人紹介", "内定・就業"],
     "友だち化を早く済ませ、休眠層に再アプローチできる", "求人応募の完了画面からLINEへ（サンクスLINE誘導）", False),
    ("パターンB ★本命", "日程調整型", ["30秒診断（LINE）", "LINEで面談予約", "面談・求人紹介", "内定・就業"],
     "電話なしで面談まで進める。在職中の整備士に合う", "本提案の主軸（3日目・14日目に面談をご案内）", True),
    ("パターンC", "LINE内完結型", ["診断（LINE）", "LINE内で求人紹介", "興味があれば面談", "内定・就業"],
     "人手を増やさずに、多くの友だちをさばける", "中長期（地域別の求人配信が育ってから）", False),
]
y = 4.3
for name, typ, flow, val, cg, star in PAT:
    h = 3.55
    box(s, 1.2, y, 25.12, h, "FFFFFF", [], line=GREEN if star else "D9D9D9", lw=2.5 if star else 1.0)
    box(s, 1.45, y + 0.3, 4.0, h - 0.6, NAVY if star else "5A6378",
        [(name, 11, True, "FFFFFF"), (typ, 13, True, "FFFFFF")])
    fw = 2.85
    for k, st in enumerate(flow):
        fx = 5.75 + k * (fw + 0.35)
        box(s, fx, y + 0.35, fw, 1.25, "FFE699" if (star and k in (0, 1)) else LGRAY, [(st, 10, True, INK)])
        if k < 3:
            arrow(s, fx + fw + 0.02, y + 0.97, fx + fw + 0.33, y + 0.97)
    text(s, 5.75, y + 1.8, 12.4, 1.4, [(val, 10.5, False, INK)], anchor=MSO_ANCHOR.MIDDLE)
    box(s, 18.75, y + 0.3, 7.3, h - 0.6, LGREEN if star else LBLUE,
        [("チェングロウスでは", 9.5, True, GRAY), (cg, 10.5, True, NAVY)])
    y += h + 0.25
band(s, 15.95, "まずBで「電話なしで面談まで」をつくり、Aで応募者を、Cで友だち資産を広げる。")

# ================= S16 損益分岐 =================
s = prs.slides[15]
box(s, 18.74, 14.0, 8.06, 1.55, "FFFFFF", [("損益分岐：半年で約3人の就業で回収", 11.5, True, NAVY),
                                            ("総額264万円 ÷ 1就業あたり売上120万円（仮定）", 9, False, GRAY)],
    line=GREEN, lw=2.0)

# ================= S22 通知メッセージ =================
s = prs.slides[21]
clear(s, "通知メッセージと特殊動線の使いどころ")
lead(s, ["通知メッセージは、友だち追加していない人にも電話番号あてに届きます。",
         "過去リストの掘り起こしではなく、「面談に来てもらう」ために使います。"])
text(s, 1.2, 4.25, 20, 0.6, [("■ 通知メッセージの利用シーン", 12, True, NAVY)])
NOTI = [("面談予約の確定通知", "来場率を上げる → 面談実施（CV）"), ("面談前日のリマインド", "当日キャンセルを防ぐ"),
        ("内定通知・入社日の確定", "満足度を上げる → 友だち紹介"), ("求人フェア・相談会の告知", "保有リストへ再接触")]
cw = (25.12 - 0.6) / 4
for i, (h, b) in enumerate(NOTI):
    box(s, 1.2 + i * (cw + 0.2), 4.95, cw, 2.7, "FFFFFF", [(h, 11.5, True, NAVY), (b, 10, False, INK)], line="8EA9DB")
text(s, 1.2, 8.1, 20, 0.6, [("■ その他の特殊動線（LINEで答えるだけ）", 12, True, NAVY)])
SPEC = [("3問で簡易キャリア診断", "答えるだけ。心理ハードルを最小に"), ("「あなたに合う職場」診断", "数問で求人を提示 → 面談予約へ"),
        ("資格別の年収シミュレーション", "参加型で友だち獲得"), ("整備士向けキャリア相談会", "特化型の入口に")]
for i, (h, b) in enumerate(SPEC):
    box(s, 1.2 + i * (cw + 0.2), 8.8, cw, 2.7, NAVY, [(h, 11.5, True, "FFFFFF"), (b, 10, False, "FFFFFF")])
text(s, 1.2, 11.8, 25.1, 1.0, [
    ("※ 通知メッセージは、予約確認・リマインドなど本人の行動に紐づく連絡に使う機能です（広告目的の一斉送信は不可）。", 9.5, False, GRAY),
    ("※ 導入費用は初期20万円〜（P31 都度発注プラン）。", 9.5, False, GRAY)])
band(s, 13.3, "面談の「予約」で終わらせず、「実施」まで届ける。")

# ================= 改善モデル①〜④ =================
MODELS = [
    (37, "改善モデル①｜友だち・リードを増やす", {0, 1, 2, 3},
     ["広告・サイトの集客地点に「LINE追加」を加え、", "応募まで行かない整備士もリストとして残します。"],
     [("CPF完結型", "Meta広告の目的を友だち獲得に変える。タップするだけで友だちになり、そのままLINE内で30秒診断まで進む。サイトに移動させないので離脱が起きない。",
       ["CPF広告をタップ", "友だち追加", "診断・配信が起動"], "既存のMeta予算（年250万円）の転用＝追加出費ゼロ"),
      ("離脱防止ポップアップ", "求人ページを見て帰ろうとした瞬間に「整備士の年収相場をLINEで見られます」を表示。「応募する」より手前の、軽い一歩を用意する。",
       ["離脱を検知", "ポップアップ表示", "LINE追加"], "関心が高い「検索当日」の訪問を、その場で拾う"),
      ("サンクスLINE誘導", "求人応募・転職支援サービス登録の完了画面で「選考の連絡はLINEで届きます」と案内。応募した人をLINEに乗せ、連絡を一本化する。",
       ["応募完了", "完了画面で案内", "LINE追加"], "電話がつながらない在職中の人にも連絡が届く")],
     "応募しないで帰る訪問者を、LINEの友だちとして残す。"),
    (39, "改善モデル②｜面談に引き上げる", {4, 5},
     ["友だちになった人を「まず話を聞くだけ」の面談へ。", "売り込みは2回だけにします。"],
     [("14日ステップ×面談オファー", "年収相場→職種図鑑→働き方の順に届け、関心が消える前の3日目と14日目に面談をご案内。「履歴書不要・夜OK・Webで30分」とハードルを下げる。",
       ["1〜2日目 情報", "3日目 面談①", "14日目 面談②"], "前後検索：職業の関心は当日〜2日後に集中"),
      ("3問アンケート×スコアリング", "「いまの温度感」を3問で聞き、「3ヶ月以内に動きたい」「今すぐ相談したい」と答えた人だけを担当者が即日フォロー。全員に電話しない。",
       ["3問を配信", "温度感を判定", "担当者へ連携"], "担当者の手間を増やさず、熱い人から面談へ"),
      ("季節企画で休眠を起こす", "検索の山（3月）の2〜3週前に「年度替わりの転職、今から準備する人が増えています」と配信。売り込みではなく「タイミングの案内」として届ける。",
       ["山の2〜3週前", "「今だから」配信", "休眠層が再反応"], "整備士の検索は山3月・谷12月（Googleトレンド）")],
     "面談の「予約」をゴールにせず、「実施」まで運ぶ。"),
    (41, "改善モデル③｜内定辞退を防ぐ", {6},
     ["内定から承諾までの不安にLINEですぐ答え、", "辞退・音信不通を防ぎます。"],
     [("不安に即答する自動応答", "「入社日は選べる？」「給与は交渉できる？」「休みは本当に取れる？」など内定後によくある質問に、LINEの自動応答で即答。答えきれないものは担当者へ。",
       ["質問を受信", "自動で即答", "担当者の面談へ"], "夜間・休日でも止まらない"),
      ("日程調整の自動化", "「相談したい」と送った瞬間に候補日を表示し、タップで確定。前日にリマインドを自動送付。電話やメールの往復で気持ちが冷めるのを防ぐ。",
       ["「相談したい」", "候補日から選ぶ", "前日リマインド"], "面談・内定者面談の当日キャンセルを減らす"),
      ("入社後がわかる体験談", "「入社3ヶ月の先輩整備士の1日」「工場の設備・工具」など、入社後が想像できる体験談を配信。入社後どうなるかの不安を、実例で解消する。",
       ["体験談を配信", "不安を解消", "内定承諾"], "前後検索：検索の直後に「ホワイト企業」など働き方の検索")],
     "内定の後もLINEでつながっているから、辞退を防げる。"),
    (42, "改善モデル④｜定着・紹介につなげる", {7},
     ["入社後3ヶ月の定着を支え、", "友だち紹介と、次の転職時の再利用につなげます。"],
     [("定着フォロー", "入社日からの日数をタグで管理し、1週間後・1ヶ月後・3ヶ月後に「困っていることはありませんか？」を自動配信。早期離職の兆しがあれば担当者が介入。",
       ["入社日をタグ化", "1週・1ヶ月・3ヶ月", "早期離職を防ぐ"], "定着が、紹介・再利用の土台になる"),
      ("友だち紹介キャンペーン", "定着を確認できたタイミング（入社3ヶ月が目安）で「整備士仲間をご紹介ください」を配信。LINEでシェアでき、紹介経由の人はミスマッチが少ない。",
       ["定着を確認", "紹介のご案内", "紹介経由の登録"], "広告費ゼロで新しい友だちが増える"),
      ("長期離脱層への再接触", "半年以上反応がない人には、売り込みをせず整備士の年収動向・資格情報を2〜3回届けて関係を戻す。そのうえで再登録・紹介へ。",
       ["情報だけを届ける", "関係を戻す", "再登録・紹介"], "ブロックされない接し方で、次の転職時に思い出してもらう")],
     "一度つながった整備士と、次の転職まで関係を続ける。"),
]
for idx, title, active, ld, cards, msg in MODELS:
    s = prs.slides[idx]
    clear(s, title)
    lead(s, ld)
    n, gap = 8, 0.12
    pw = (25.12 - gap * (n - 1)) / n
    for i, ph in enumerate(PHASES):
        on = i in active
        box(s, 1.2 + i * (pw + gap), 4.15, pw, 0.8, NAVY if on else "E7E6E6", [(ph, 9, True, "FFFFFF" if on else "A6A6A6")],
            shape=MSO_SHAPE.CHEVRON if i else MSO_SHAPE.PENTAGON, ml=0.1)
    cw = (25.12 - 0.5) / 3
    for k, (h, desc, flow, note) in enumerate(cards):
        cx = 1.2 + k * (cw + 0.25)
        box(s, cx, 5.3, cw, 9.45, "F7F9FC", [], line="D9D9D9")
        box(s, cx, 5.3, cw, 0.95, NAVY, [(f"{k + 1}  {h}", 12, True, "FFFFFF")], align=PP_ALIGN.LEFT, ml=0.3)
        text(s, cx + 0.25, 6.45, cw - 0.5, 3.9, [(desc, 11.5, False, INK)])
        fw = (cw - 0.5 - 0.5) / 3
        for j, st in enumerate(flow):
            fx = cx + 0.25 + j * (fw + 0.25)
            box(s, fx, 10.55, fw, 1.5, "3467B2", [(f"{j + 1}", 9.5, True, "FFFFFF"), (st, 10, True, "FFFFFF")])
        box(s, cx + 0.25, 12.3, cw - 0.5, 2.2, LGREEN, [("チェングロウスでは", 9.5, True, GRAY), (note, 11, True, NAVY)])
    band(s, 15.2, msg)

# 中扉の文言
s = prs.slides[36]
for sh in s.shapes:
    if sh.has_text_frame and "運用のメリット" in sh.text_frame.text:
        rs = [r for p in sh.text_frame.paragraphs for r in p.runs]
        rs[0].text = "LINE公式アカウント運用のメリットと改善モデル"
        for r in rs[1:]:
            r.text = ""

# 並べ替え：37(中扉), 39, 41, 38(改善①), 40(改善②), 42(改善③), 43(改善④)
lst = prs.slides._sldIdLst
ids = list(lst)
order = [36, 38, 40, 37, 39, 41, 42]          # 0始まりの元の位置
block = [ids[i] for i in order]
for e in block:
    lst.remove(e)
anchor = ids[35]
for e in block:
    anchor.addnext(e)
    anchor = e

prs.save(OUT)
print("saved", OUT, len(prs.slides))
