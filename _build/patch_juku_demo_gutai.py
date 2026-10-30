# -*- coding: utf-8 -*-
"""塾 新FMT版：具体施策のイメージ（P29〜P36）を「左＝デモの画面／右＝ほかの案3つ」にする（2026-10-30 殿村さん「ワクトリの施策の部分を例に」）。

    python -m _build.patch_juku_demo_gutai <入力.pptx> <出力.pptx>

  形はワクトリv6（_build/patch_wakutori_v6_demo.py）と同じ。
  ・左：施策名＋1行の見出し、塾のデモ（`_data/juku/塾_LINE施策デモ_20261030.html`）の ①ユーザーが見る画面＋②裏側の処理
    （`_images/juku/demo_p29〜36.png`）、下に小さく画面の注記
  ・右：4つの枠（目的／対象／中身／見る数字）を、3つの枠（入口／LINEの仕組み／配信内容）にする
  上の8フェーズの矢印・札・タイトルは触らない。
"""
import copy
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Emu, Inches, Pt

from _build.pptx_parts import style_run

NAVY, INK, GRAY = "1F2A5C", "333333", "7F7F7F"
IMG = "_images/juku/demo_p{}.png"
LABELS = ["入口", "LINEの仕組み", "配信内容"]
PAGES = {  # 札, 施策名, 見出しの1行, 右の3つ（各2行）
    29: ("LINEの友だちへ", "離脱防止ポップアップ", "サイトを離れる方を、資料と診断でLINEへ",
         [["チラシ・教室掲示のQRや、", "友だちからの紹介にも置く"],
          ["入口ごとにURLを分けて記録し、", "あいさつの文面を出し分ける"],
          ["追加の特典に「学年別の", "テスト対策プリント」を届ける"]]),
    30: ("露出の強化", "季節講習の企画配信", "資料請求・体験の方に、冬期講習を案内",
         [["体験に来た保護者へ、翌日に", "「お友だちも一緒に」を届ける"],
          ["診断の結果を、", "お友だちに送れるようにする"],
          ["講習の案内に、お友だちと", "一緒に参加できる体験枠を付ける"]]),
    31: ("CVの最大化", "30秒ぴったり診断", "4問の回答で、お子さまに合う学び方を返す",
         [["リッチメニューや配信から、", "いつでも診断をやり直せる"],
          ["回答（学年・悩み・性格・時期）", "をタグにして、次の配信を変える"],
          ["結果に合う学び方と、", "同じタイプのお子さまの例"]]),
    32: ("CPOの改善（CV質）", "体験の前日のご確認", "持ち物・アクセスを届け、来塾につなげる",
         [["予約はLINEの中でも、", "サイトのフォームでも受け付ける"],
          ["予約した方は、リッチメニューを", "「体験予約済みの方」に切り替える"],
          ["担当講師の紹介・持ち物・", "教室までのアクセス"]]),
    33: ("休眠顧客の有効化", "季節の企画で再案内", "入会しなかった方に、冬期講習で再検討の場を",
         [["資料請求だけの方にも、", "通知メッセージで再案内する"],
          ["体験に来た・来ていないで", "タグを分けて配信する"],
          ["夏期・冬期講習や新学年の前に、", "学年に合うコースを届ける"]]),
    34: ("マネタイズ最大化", "回答に合わせた出し分け", "同じタイプの例を届け、体験予約につなげる",
         [["資料請求をして、まだ体験を", "予約していない方に届ける"],
          ["「情報収集だけ」の方には、", "案内を控えて情報だけを届ける"],
          ["よくある質問", "（入会を断ってもいい？総額は？）"]]),
    35: ("リピート施策", "入会後のご案内", "入会の登録で、通塾中の方向けの配信に切り替え",
         [["入会の登録で、リッチメニューを", "通塾中の方向けに切り替える"],
          ["学年・コースでタグを付け、", "講習の案内を出し分ける"],
          ["定期テスト前の学習のヒントと、", "季節講習の案内"]]),
    36: ("コミュニティ運営", "友だちへの紹介", "トークを転送するだけで、お友だちに紹介できる",
         [["通塾中の保護者に、紹介の", "トークを転送してもらう"],
          ["紹介から追加した方にタグを付け、", "紹介元が分かるようにする"],
          ["教室からのお知らせや、", "保護者の声を届ける"]]),
}
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr"


def inch(v):
    return Emu(v).inches


def add_text(s, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for k, (text, size, bold, color, after) in enumerate(paras):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        if after:
            p.space_after = Pt(after)
        r = p.add_run()
        r.text = text
        style_run(r, size, bold, color)
    return tb


def main(src, out):
    prs = Presentation(src)
    for n, (tag, name, sub, ideas) in PAGES.items():
        s = prs.slides[n - 1]
        assert any(sh.has_text_frame and sh.text_frame.text == tag for sh in s.shapes), (n, tag)
        big = next(sh for sh in s.shapes if abs(inch(sh.left) - 0.29) < 0.05 and abs(inch(sh.top) - 2.03) < 0.05
                   and inch(sh.width) > 6)
        right_boxes = sorted([sh for sh in s.shapes if inch(sh.left) > 6.8 and inch(sh.top) > 1.9
                              and not (sh.has_text_frame and sh.text_frame.text.strip())], key=lambda x: inch(x.top))
        right_texts = [sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip() and inch(sh.left) > 6.8
                       and inch(sh.top) > 1.9]
        flow = [sh for sh in s.shapes if sh is not big and inch(sh.left) < 6.8 and inch(sh.top) > 2.5]
        assert len(right_boxes) == 4 and len(right_texts) == 4 and len(flow) == 11, (n, len(right_boxes), len(right_texts), len(flow))
        tmpl = right_boxes[0]
        tmpl.left = Inches(6.85)
        for sh in flow + right_texts + right_boxes[1:]:
            sh._element.getparent().remove(sh._element)

        # 左：施策名＋見出し、デモの画面（①＋②）、注記
        add_text(s, 0.45, 2.15, 6.1, 0.75, [(name, 16, True, NAVY, 0), (sub, 12, False, INK, 0)])
        h = 3.62
        w = h * 1756 / 1478
        pic = s.shapes.add_picture(IMG.format(n), Inches(0.29 + (6.44 - w) / 2), Inches(2.9), Inches(w), Inches(h))
        pic.line.color.rgb = RGBColor.from_string("D9D9D9")
        pic.line.width = Pt(0.75)
        add_text(s, 0.45, 6.53, 6.1, 0.24,
                 [("画面：塾向けに作成したデモ（①ユーザーが見る画面／②裏側の処理）。教室名・文面はデモ用のサンプル", 8.5, False, GRAY, 0)])

        # 右：3つの枠（入口／LINEの仕組み／配信内容）
        top, bottom, gap = 2.03, 6.80, 0.1
        bh = (bottom - top - gap * 2) / 3
        tmpl.top, tmpl.height = Inches(top), Inches(bh)
        boxes = [tmpl]
        for k in (1, 2):
            s.shapes._spTree.append(copy.deepcopy(tmpl._element))
            sh = s.shapes[-1]
            sh.top = Inches(top + k * (bh + gap))
            boxes.append(sh)
        for box, label, lines in zip(boxes, LABELS, ideas):
            add_text(s, inch(box.left) + 0.12, inch(box.top) + 0.1, inch(box.width) - 0.24, bh - 0.2,
                     [(label, 13.5, True, NAVY, 4)] + [(x, 12, False, INK, 0) for x in lines], anchor=MSO_ANCHOR.MIDDLE)

        used = set()
        nid = max(int(e.get("id")) for e in s.shapes._spTree.iter(P)) + 1
        for e in s.shapes._spTree.iter(P):
            if e.get("id") in used:
                e.set("id", str(nid))
                nid += 1
            used.add(e.get("id"))
    prs.save(out)
    print("saved:", out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
