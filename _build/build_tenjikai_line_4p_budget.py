# -*- coding: utf-8 -*-
"""展示会 LINE運用方針（4枚・4:3）。2026-10-13 佐村さん「振り替えは最低限。前提予算を取る提案に」を受けて作り直し。
課題 → 方針（LINE広告の予算枠を全店舗共通で1つ：広告費 月30万円＋運用手数料20%。2026-10-13 10万→30万に変更）→ 運用（毎月の改善。試す期間は置かない）→ LINEの友だち（入口2つ：CPF・サンクスLINE）。枠30万＝通常20万＋CPF10万。
振り替え（銀座店 Meta30万→20万＋LINE10万）は資料に入れない（一度載せたが殿村さん判断で外した 2026-10-13）。断られたときの切り返し用。
  python3 _build/build_tenjikai_line_4p_budget.py <Appendix ver1.2.pptx> [出力.pptx]
"""
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Cm

sys.path.insert(0, str(Path(__file__).resolve().parent))
from pptx_parts import (CARD, GRAY, INK, NAVY, WHITE, arrow_line, find, keep_header_only, rect,  # noqa: E402
                        rich, set_paras, shape)

ROOT = Path(__file__).resolve().parent.parent
SRC = Path(sys.argv[1])
OUT = Path(sys.argv[2]) if len(sys.argv) > 2 else ROOT / "20261013_展示会LINE運用方針_予算枠のご提案.pptx"

SLIDES = [
    ("課題｜同じ方に広告が出すぎている", "少ない方に広告が重なり、予約1件に15,000円以上かかっている"),
    ("方針｜LINE広告の予算枠を、全店舗共通で1つ設ける", "店舗ごとに積まず、1つの枠でMetaが届いていない方に展示会を案内する"),
    ("運用｜毎月の数字を見て、予約の取れるエリアに寄せる", "エリア・訴求ごとの予約数を見て、予約1件あたりの広告費を下げていく"),
    ("LINEの友だち｜見込みの方と申し込んだ方をLINEで追う", "2つの入口から友だちを増やし、展示会の案内から来場まで追う"),
]

X, W = 1.79, 23.97
Y0 = 4.5
S = 13


def text(s, y, runs, h=0.8, x=X, w=W):
    return rich(s, x, y, w, h, runs, anchor=MSO_ANCHOR.TOP)


def head(s, y, txt):
    rect(s, X, y + 0.14, 0.15, 0.5, NAVY)
    text(s, y, [[(txt, 13.5, True, NAVY)]], x=X + 0.35, w=W - 0.35)
    return y + 0.85


def rows(s, y, items, kw=4.6, h=0.72):
    for i, (k, v) in enumerate(items):
        text(s, y + i * h, [[(k, S, True, INK)]], h=h, x=X + 0.35, w=kw)
        text(s, y + i * h, [[(v, S, False, INK)]], h=h, x=X + 0.35 + kw, w=W - 0.35 - kw)
    return y + len(items) * h


def flow(s, y, items, h=1.35, gap=0.75, size=11.5):
    x0, ww = X + 0.35, W - 0.35
    bw = (ww - gap * (len(items) - 1)) / len(items)
    for i, t in enumerate(items):
        last = i == len(items) - 1
        x = x0 + i * (bw + gap)
        paras = [(ln, size, last, WHITE if last else INK, 0) for ln in t.split("\n")]
        shape(s, MSO_SHAPE.RECTANGLE, x, y, bw, h, paras, fill=NAVY if last else WHITE, line=NAVY, lw=1.0,
              margins=(0.1, 0.05, 0.1, 0.05))
        if not last:
            arrow_line(s, x + bw + 0.1, y + h / 2, x + bw + gap - 0.1, y + h / 2, NAVY, w=1.5)
    return y + h


prs = Presentation(str(SRC))
lst = prs.slides._sldIdLst
keep = [i for i, s in enumerate(prs.slides)
        if {"TextBox 1", "TextBox 2", "Connector 3"} <= {sh.name for sh in s.shapes}][:4]
assert len(keep) == 4, keep
for i, el in reversed(list(enumerate(list(lst)))):
    if i not in keep:
        prs.part.drop_rel(el.rId)
        lst.remove(el)

for s, (title, lead) in zip(prs.slides, SLIDES):
    keep_header_only(s)
    t1, t2 = find(s, "TextBox 1"), find(s, "TextBox 2")
    t1.left, t1.top, t1.width, t1.height = Cm(1.52), Cm(0.38), Cm(18.1), Cm(0.94)
    t2.left, t2.top, t2.width, t2.height = Cm(1.79), Cm(2.27), Cm(25.31), Cm(0.96)
    set_paras(t1, [(title, 0)])
    set_paras(t2, [(lead, 0)])
    for t in (t1, t2):
        for r in t.text_frame._txBody.iter(qn("a:rPr")):
            r.set("sz", "1600")

s1, s2, s3, s4 = prs.slides

# 1. 課題（3枚版と同じ）
y = head(s1, Y0, "現状の数値")
y = rows(s1, y, [("フリークエンシー", "5.24（同じ方が月に5回以上、同じ広告を見ている）"),
                 ("CPA", "15,000円以上（予約1件あたりの広告費）")])
y = head(s1, y + 0.5, "なぜ効率が落ちるのか")
y = flow(s1, y + 0.1, ["オーディエンスが少ない", "予算を使い切るため、\n同じ方に繰り返し表示", "予約は増えず、\nCPAだけ上がる"])
y = head(s1, y + 0.6, "オーディエンスの少ないエリア（例）")
rows(s1, y, [("三鷹", "10,700〜12,600"), ("高松", "上限25,000"), ("日向", "4,400")])

# 2. 方針：予算枠（全店舗共通）
y = head(s2, Y0, "費用（月）")
# 2026-10-13 殿村さん：枠30万円を「予約を取る20万」と「見込みの方を友だちとして貯める10万（CPF）」に分ける
op = 0.9
items = [("LINE広告（通常）", "20万円", "予約を取る"), ("LINE広告（CPF）", "10万円", "見込みの方を友だちに"),
         ("運用手数料", "6万円", "広告費の20%"), ("合計", "36万円", "")]
n_it = len(items)
bw = (W - 0.35 - op * (n_it - 1)) / n_it
x0 = X + 0.35
for i, (lab, val, sub) in enumerate(items):
    x = x0 + i * (bw + op)
    last = i == n_it - 1
    paras = [(lab, 12, True, WHITE if last else NAVY, 2), (val, 22, True, WHITE if last else NAVY, 2)]
    if sub:
        paras.append((sub, 11, False, WHITE if last else GRAY, 0))
    shape(s2, MSO_SHAPE.RECTANGLE, x, y + 0.1, bw, 2.3, paras, fill=NAVY if last else CARD, line=NAVY, lw=1.0,
          margins=(0.1, 0.05, 0.1, 0.05))
    if i < n_it - 1:
        text(s2, y + 0.75, [[("＝" if i == n_it - 2 else "＋", 22, True, NAVY)]], h=1.0, x=x + bw, w=op)
y = head(s2, y + 3.0, "店舗ごとに積まない理由")
y = rows(s2, y, [("金額", "1店舗分のMeta予算と同じ30万円で、全店舗のエリアに配信"),
                 ("配信", "LINE広告は1つのアカウントで、展示会を開く全エリアに配信できる"),
                 ("配分", "予約の取れるエリアに、広告費を寄せて調整する"),
                 ("LINE広告", "Metaとは別の媒体。国内の月間利用者数 1億人以上（LINEヤフー公表）")])
text(s2, y + 0.35, [[("仮に通常20万円 → クリック2,000（CPC 100円）→ 予約20件（予約率1.0%）→ 予約1件あたりの広告費 10,000円", 11, False, GRAY)],
                    [("仮にCPF10万円 → 友だち500人（1人200円）→ ブロックを除いて約250人が残り、展示会のたびに案内できる", 11, False, GRAY)]],
     h=1.0)
# 振り替え（今の予算の中で始める場合）は一度載せたが、殿村さん判断で外した（2026-10-13）。切り返し用として手元に置く

# 3. 運用：毎月の改善（試す期間・止める基準は置かない。殿村さん 2026-10-13）
y = head(s3, Y0, "毎月の流れ")
y = flow(s3, y + 0.1, ["LINE広告を配信", "エリア・訴求ごとの\n予約数を確認", "予約の取れる\nエリアに寄せる", "訴求・クリエイティブ\nを見直す"], h=1.6, gap=0.6)
y = head(s3, y + 0.6, "毎月お出しするもの")
rows(s3, y, [("レポート", "エリア別・訴求別の予約数と、予約1件あたりの広告費（Metaと並べて比較）"),
             ("次の打ち手", "予算を寄せるエリア・訴求の見直し案")])

# 4. LINEの友だち：入口2つ（CPF・サンクスLINE）→ 配信 → 予約・来場（2026-10-13 殿村さん。メモの⑤）
y = head(s4, Y0, "入口① CPF｜まだ申し込んでいない方")
y = flow(s4, y + 0.1, ["LINE広告（CPF）", "LINE\n友だち追加", "展示会ごとに\nエリア別の開催案内", "予約"], h=1.5, gap=0.6)
y = head(s4, y + 0.5, "入口② サンクスLINE｜申し込んだ方")
y = flow(s4, y + 0.1, ["申込完了画面", "LINE\n友だち追加", "前日・当日の\nお知らせ", "来場"], h=1.5, gap=0.6)
y = head(s4, y + 0.5, "実績・費用")
rows(s4, y, [("実績", "弊社運用の美容クリニックで、予約後の来院率 40〜50%改善"),
             ("費用", "CPFは2枚目の枠に含む／サンクスLINEは初期15万円・月額3万円〜")])

# 本文を上下の真ん中に寄せる（殿村さん 2026-10-13「情報量を真ん中に寄せたい」）
# 本文の範囲＝区切り線（3.86cm）〜フッターの上（18.2cm）。ヘッダー3つ（タイトル・リード・区切り線）は動かさない
HDR_NAMES = {"TextBox 1", "TextBox 2", "Connector 3"}
AREA_TOP, AREA_BOTTOM = 3.86, 18.2
for sl in prs.slides:
    body = [sh for sh in sl.shapes if sh.name not in HDR_NAMES]
    if not body:
        continue
    top = min(sh.top for sh in body)
    bottom = max(sh.top + sh.height for sh in body)
    shift = int((Cm(AREA_TOP) + Cm(AREA_BOTTOM)) / 2 - (top + bottom) / 2)
    if shift > 0:
        for sh in body:
            sh.top = sh.top + shift

# 字体をメイリオに統一（latin → ea → cs の順）
AFTER = {qn(x) for x in ("a:sym", "a:hlinkClick", "a:hlinkMouseOver", "a:rtl", "a:extLst")}
for s in prs.slides:
    for r in s.shapes._spTree.iter(qn("a:r")):
        if r.find(qn("a:rPr")) is None:
            r.insert(0, r.makeelement(qn("a:rPr"), {"lang": "ja-JP"}))
    for tag in ("a:rPr", "a:endParaRPr", "a:defRPr"):
        for rpr in s.shapes._spTree.iter(qn(tag)):
            for ft in ("a:latin", "a:ea", "a:cs"):
                for e in rpr.findall(qn(ft)):
                    rpr.remove(e)
            nxt = next((c for c in rpr if c.tag in AFTER), None)
            for ft in ("a:latin", "a:ea", "a:cs"):
                e = rpr.makeelement(qn(ft), {"typeface": "メイリオ"})
                if nxt is not None:
                    nxt.addprevious(e)
                else:
                    rpr.append(e)

prs.save(str(OUT))
print("saved:", OUT.name, len(prs.slides), "枚")
