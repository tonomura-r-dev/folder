# -*- coding: utf-8 -*-
"""チェングロウス ver4.0 → ver4.1：19枚目（ステップ配信の中身）をカードから表に変更（2026-10-02 殿村さん指示）
  python3 _build/patch_chengrowth_v41_steptable.py <ver4.0.pptx>
"""
import sys
from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Cm, Pt

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "20260929_チェングロウス_LINE公式アカウント運用のご提案ver4.1.pptx"
NAVY, INK, WHITE = RGBColor(0x1F, 0x28, 0x5A), RGBColor(0x33, 0x33, 0x33), RGBColor(0xFF, 0xFF, 0xFF)
BAND, HOT = RGBColor(0xF4, 0xF7, 0xFF), RGBColor(0xE8, 0xF8, 0xEE)
GREEN = RGBColor(0x0B, 0x7A, 0x3B)
prs = Presentation(sys.argv[1])
s = prs.slides[18]


def set_par(sh, text):
    p = sh.text_frame.paragraphs[0]
    p.runs[0].text = text
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)


for sh in list(s.shapes):
    if sh.name == "Rounded Rectangle 4":
        sh._element.getparent().remove(sh._element)
for sh in s.shapes:
    if sh.name == "TextBox 1":
        set_par(sh, "②効率改善｜ステップ配信の内容（10通の例）")
    if sh.name == "TextBox 2":
        set_par(sh, "友だち追加から14日間で、全10通を自動で配信")

ROWS = [("0日", "あいさつ（希望の職種を聞く）"), ("1日", "職種の違い"), ("2日", "働き方"), ("3日", "求人のご案内①"), ("4日", "資格取得支援"),
        ("5日", "他の職種との比較"), ("7日", "資格別の年収"), ("10日", "現場の様子"), ("12日", "地域別の新着求人"), ("14日", "求人のご案内②")]
RH = 1.05
gf = s.shapes.add_table(len(ROWS) + 1, 2, Cm(2.2), Cm(4.9), Cm(18.0), Cm(RH * (len(ROWS) + 1)))
tbl = gf.table
tbl.columns[0].width = Cm(3.4)
tbl.columns[1].width = Cm(14.6)


def cell(c, text, size, bold, color, fill, align):
    c.fill.solid()
    c.fill.fore_color.rgb = fill
    c.vertical_anchor = MSO_ANCHOR.MIDDLE
    c.margin_left = c.margin_right = Cm(0.4)
    tf = c.text_frame
    tf.text = text
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.runs[0]
    r.font.size, r.font.bold = Pt(size), bold
    r.font.color.rgb = color
    rpr = r._r.get_or_add_rPr()
    for ft in ("a:latin", "a:ea", "a:cs"):
        e = rpr.makeelement(qn(ft), {"typeface": "メイリオ"})
        rpr.append(e)


cell(tbl.cell(0, 0), "日", 14, True, WHITE, NAVY, PP_ALIGN.CENTER)
cell(tbl.cell(0, 1), "配信の内容", 14, True, WHITE, NAVY, PP_ALIGN.CENTER)
for i, (d, t) in enumerate(ROWS, start=1):
    hot = "ご案内" in t
    fill = HOT if hot else (BAND if i % 2 else WHITE)
    col = GREEN if hot else INK
    cell(tbl.cell(i, 0), d, 14, True, col, fill, PP_ALIGN.CENTER)
    cell(tbl.cell(i, 1), t + ("　← 応募を案内" if hot else ""), 14, hot, col, fill, PP_ALIGN.LEFT)
for r in tbl.rows:
    r.height = Cm(RH)
tblPr = tbl._tbl.find(qn("a:tblPr"))
tblPr.set("bandRow", "0")
tblPr.set("firstRow", "0")
sid = tblPr.find(qn("a:tableStyleId"))
if sid is not None:
    sid.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"   # スタイルなし（罫線・色は自前）
for sl in prs.slides:
    for n, c in enumerate(sl.shapes._spTree.iter(qn("p:cNvPr")), start=2):
        c.set("id", str(n))
prs.save(str(OUT))
print("saved:", OUT.name, len(prs.slides), "枚")
