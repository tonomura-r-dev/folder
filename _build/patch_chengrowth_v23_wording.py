# -*- coding: utf-8 -*-
"""ver2.2（殿村さんPC保存版）→ ver2.3：残っていた言い回しの修正
  python3 _build/patch_chengrowth_v23_wording.py
"""
from pptx import Presentation
SRC = "20260929_株式会社チェングロウス御中_LINE公式アカウント運用のご提案ver2.2.pptx"
OUT = "20260929_株式会社チェングロウス御中_LINE公式アカウント運用のご提案ver2.3.pptx"
REP = [
 (12, "3月の山に向けて前倒し", "3月のピークに向けて前倒し"),
 (19, "求人に対して人が大きく足りない「超売り手市場」です。", "求人に対して人材が大きく不足している「超売り手市場」です。"),
 (20, "資格を持つ人が集まる", "資格を持つ方が集まる"),
 (20, "一年中、転職を考える人がいる。", "一年を通して、転職を検討する方がいる。"),
 (20, "LINEで登録者として残しておくことが大切です。", "LINEで登録者としてつながっておくことが大切です。"),
]


def runs(shape):
    if shape.shape_type == 6:
        for c in shape.shapes:
            yield from runs(c)
    elif getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            for cell in row.cells:
                for p in cell.text_frame.paragraphs:
                    yield from p.runs
    elif shape.has_text_frame:
        for p in shape.text_frame.paragraphs:
            yield from p.runs


prs = Presentation(SRC)
S = list(prs.slides)
for no, a, b in REP:
    hit = 0
    for sh in S[no - 1].shapes:
        for r in runs(sh):
            if a in r.text:
                r.text = r.text.replace(a, b); hit += 1
    print(("OK " if hit else "未置換 ") + f"{no}: {a}")
prs.save(OUT)
