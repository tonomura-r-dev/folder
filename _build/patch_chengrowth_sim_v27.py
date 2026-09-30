# -*- coding: utf-8 -*-
"""チェングロウスSIM ver2.6（殿村さん修正版）→ ver2.7：CV①（LINEで登録）の率を 1.5% → 1.0% に下げる。

- CV②（応募）は手入力のまま触らない（殿村さん指示）
- 友だち数の増加（26行）も同じ率を使っているので一緒に直す（CV①＝LINEで登録＝友だち追加）
- チェックボックス（xl/featurePropertyBag）を消さないため、openpyxlで保存せずXMLを直接書き換える
- 数式のキャッシュ値は、LibreOfficeで再計算したコピーの値を入れる（Excelでは開いた時に再計算される）

  python3 _build/patch_chengrowth_sim_v27.py <ver2.6のxlsx>
"""
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
SRC = Path(sys.argv[1])
OUT = ROOT / "20260929_株式会社チェングロウス御中_LINEOA施策提案【SIM】ver2.7.xlsx"
SHEET = "xl/worksheets/sheet1.xml"   # SIM3_LINE登録2%
SN = "SIM3_LINE登録2%"
OLD, NEW = "*0.015*", "*0.01*"
TEXT = [("LINEで登録1.5%", "LINEで登録1.0%"), ("UU×1.5%×立ち上げ係数", "UU×1.0%×立ち上げ係数")]

with zipfile.ZipFile(SRC) as z:
    files = {n: z.read(n) for n in z.namelist()}
    infos = {i.filename: i for i in z.infolist()}

sheet = files[SHEET].decode("utf-8")
assert sheet.count(OLD) == 12, sheet.count(OLD)          # 26行・36行 × E〜J
sheet = sheet.replace(OLD, NEW)

ss = files["xl/sharedStrings.xml"].decode("utf-8")
for a, b in TEXT:
    assert a in ss, a
    ss = ss.replace(a, b)
files["xl/sharedStrings.xml"] = ss.encode("utf-8")

wb = files["xl/workbook.xml"].decode("utf-8")
if "fullCalcOnLoad" not in wb:
    wb = re.sub(r"<calcPr ", '<calcPr fullCalcOnLoad="1" ', wb, count=1)
files["xl/workbook.xml"] = wb.encode("utf-8")

# LibreOfficeで再計算した値を、数式セルのキャッシュ値に入れる（全シートの数式セル）
files[SHEET] = sheet.encode("utf-8")
tmp = Path(tempfile.mkdtemp())
draft = tmp / "draft.xlsx"
# 再計算させるため、下書きでは数式セルのキャッシュ値を消す（残っているとLibreOfficeが古い値を使う）
strip = re.compile(r'(<f[^>]*>[^<]*</f>|<f[^>]*/>)<v>[^<]*</v>')
with zipfile.ZipFile(draft, "w", zipfile.ZIP_DEFLATED) as z:
    for n, b in files.items():
        if n.startswith("xl/worksheets/sheet"):
            b = strip.sub(r"\1", b.decode("utf-8")).encode("utf-8")
        z.writestr(infos[n], b)
subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", str(tmp / "out"), str(draft)],
               capture_output=True, timeout=180)
calc = openpyxl.load_workbook(tmp / "out" / "draft.xlsx", data_only=True)
names = [s.title for s in openpyxl.load_workbook(draft, read_only=True).worksheets]

for idx, name in enumerate(names, 1):
    part = f"xl/worksheets/sheet{idx}.xml"
    xml = files[part].decode("utf-8")
    ws = calc[name]

    def fill(m):
        ref, attrs, f = m.group(1), m.group(2), m.group(3)
        v = ws[ref].value
        if isinstance(v, bool) or v is None or isinstance(v, str):
            return m.group(0)              # 文字列・真偽の数式は元のキャッシュのまま
        attrs = re.sub(r'\s+t="[^"]*"', "", attrs)
        return f'<c r="{ref}"{attrs}>{f}<v>{repr(float(v)) if isinstance(v, float) else v}</v></c>'

    xml = re.sub(r'<c r="([A-Z]+\d+)"([^>]*)>(<f[^>]*>[^<]*</f>|<f[^>]*/>)<v>[^<]*</v></c>', fill, xml)
    files[part] = xml.encode("utf-8")

with zipfile.ZipFile(OUT, "w", zipfile.ZIP_DEFLATED) as z:
    for n, b in files.items():
        z.writestr(infos[n], b)
shutil.rmtree(tmp)

v = openpyxl.load_workbook(OUT, data_only=True)[SN]
row = lambda r: [round(v.cell(r, c).value, 4 if r == 41 else 1) if isinstance(v.cell(r, c).value, float) else v.cell(r, c).value for c in range(5, 12)]
for r, lab in [(26, "友だち増加"), (28, "友だち数"), (29, "ターゲット"), (35, "想定Click"), (36, "CV①"), (37, "CV②"), (41, "CVR"), (59, "CPA")]:
    print(lab, row(r))
print("checkbox:", any("featurePropertyBag" in n for n in zipfile.ZipFile(OUT).namelist()))
print("saved:", OUT.name)
