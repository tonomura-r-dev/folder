# -*- coding: utf-8 -*-
"""清水グループ様 LINE公式アカウント運用のご提案 v1（簡易版の資料。2026-10-19 殿村さん「簡易的でいい」）。

    python -m _build.build_shimizu_v1 <人材業界資料.pptx> <あらうちv10.pptx> <SIM.xlsx>

  ベース＝人材業界の業界資料（派遣・バイト・単発）。業界資料のページを清水グループ向けに直して使い、課題・全体像・概算は新しく作る。
  料金はAIの案の金額で確定（比べるために業界資料の金額の版も作ったが、殿村さん「Bで進めていい」2026-10-19）。
  並び：表紙／アジェンダ／現状と課題／解決策（職種分岐型）／友だちの入口／リッチメニュー／面接に来る方を増やす／
        概算お見積り／運用プラン（松竹梅）／運用サポート体制／裏表紙（簡易なのでAppendixは付けない）
  使わない業界資料のページ：派遣の契約満了・前後検索・あいさつ（「営業のお電話はいたしません」は派遣向け）・通知メッセージ（歩留まりのページと重なる）ほか
"""
import copy
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches

from _build.arauchi_v4.lib import add_box, move, set_paras
from _build.arauchi_v10.build import by_id, clone, runs
from _build.build_sunao_v1 import all_meiryo
from _build.build_sunao_v2 import load_sim
from _build.patch_tanaka_v5 import copy_slide

ROOT = Path(__file__).resolve().parent.parent
NAVY = "002060"
LBLUE = "DCE6F2"
TEXT = "333333"
A = qn("a:p")

PLANS = {
    "": dict(                                                  # 2026-10-19 殿村さん「Bで進めていい」＝AIの案の金額
        rows=[("松：D 運用代行フルサポート", "50万円", "35万円", "配信の作成・画像制作・月次の分析と改善まで代行"),
              ("竹：A コンサル基本", "30万円", "20万円", "月次レポート・改善のご提案・運用のレクチャー"),
              ("梅：C 初期設計・構築のみ", "60万円", "0円", "初期構築・職種別の自動応答・リッチメニュー")],
        consult_init=300000),
}


# ---------------------------------------------------------------- 部品
def put(s, sid, lines):
    """段落ごとの書式を残して文字を差し替える（i番目の行は元のi番目の段落を型にする。足りなければ最後の段落）"""
    if isinstance(lines, str):
        lines = [lines]
    body = by_id(s, sid)._element.find(qn("p:txBody"))
    ps = body.findall(A)
    for p in ps:
        body.remove(p)
    for i, ln in enumerate(lines):
        p = copy.deepcopy(ps[min(i, len(ps) - 1)])
        rs = p.findall(qn("a:r"))
        for r in rs[1:]:
            p.remove(r)
        for br in p.findall(qn("a:br")):
            p.remove(br)
        rs[0].find(qn("a:t")).text = ln
        body.append(p)


def drop(s, *ids):
    for sid in ids:
        e = by_id(s, sid)._element
        e.getparent().remove(e)


def fill_row(s, sid, y, h):
    """1行用の文字の箱を、行の高さいっぱいにして上下中央に（2行になっても行からはみ出さない）"""
    sh = by_id(s, sid)
    move(sh, y=y, h=h)
    sh._element.find(".//" + qn("a:bodyPr")).set("anchor", "ctr")


def shift(s, ids, dy):
    for sid in ids:
        sh = by_id(s, sid)
        sh.top = sh.top + Inches(dy)


def body_page(prs, src, title, lead):
    """業界資料と同じ見出し（タイトル＋リード）だけ残した空のページ"""
    s = copy_slide(prs, src)
    for sh in list(s.shapes):
        if sh.shape_id not in (2, 3):
            sh._element.getparent().remove(sh._element)
    put(s, 2, title)
    put(s, 3, lead)
    return s


def box(s, x, y, w, h, text, strong=False, size=14):
    add_box(s, x, y, w, h, [runs((t, dict(size=size, bold=True, color="FFFFFF" if strong else NAVY)))
                            for t in text.split("\n")],
            fill=NAVY if strong else LBLUE, align="ctr", anchor="ctr", line_spacing=110,
            ins=(0.06, 0.04, 0.06, 0.04), name=text.split("\n")[0])


def arrow(s, x, y, w, h, mark="▶"):
    add_box(s, x, y, w, h, [runs((mark, dict(size=14, color=NAVY)))], align="ctr", anchor="ctr",
            ins=(0, 0, 0, 0), name="矢印")


# ---------------------------------------------------------------- 表紙・アジェンダ
def p_cover(s, ar_cover):
    sub = by_id(s, 6)
    set_paras(sub, ["グループ全体の求人に対応する、", "職種分岐型のLINE採用の仕組み"])
    move(sub, x=1.76, y=2.3, w=7.25, h=0.95)
    for p in sub.text_frame.paragraphs:
        p.alignment = PP_ALIGN.CENTER
    drop(s, 3, 7)                                              # 「- 2026最新版 -」「業界：人材（派遣・バイト・単発）」
    set_paras(clone(s, by_id(ar_cover, 31)), ["清水グループ 御中"])


def p_agenda(s):
    put(s, 5, ["サイトに来ている方・応募した方を、LINEでつなぎ止める。",
               "飲食・イベント・警備・正社員まで、グループ全体の求人に職種を分けて案内する提案。"])
    for (a, b), (h, t) in zip(((9, 10), (14, 15), (19, 20), (24, 25)),
                              (("第1章｜現状と課題", "応募の前後で、求職者とのつながりが切れている"),
                               ("第2章｜解決策", "希望条件で職種を分け、グループ全体の求人につなぐ"),
                               ("第3章｜施策", "友だちの入口・リッチメニュー・面接に来る方を増やす"),
                               ("第4章｜費用", "概算お見積り・運用プラン・運用サポート体制"))):
        put(s, a, h)
        put(s, b, t)
    drop(s, *range(26, 36), 41)                                # 第5章・第6章の行と、下の★の注記
    shift(s, range(6, 26), 0.6)                                # 4行を上寄りの中央へ


# ---------------------------------------------------------------- 第1章 現状と課題（業界資料の「歩留まり」の行の形）
def p_kadai(prs, src):
    s = copy_slide(prs, src)
    put(s, 2, "［第1章］現状と課題｜応募の前後で、求職者とのつながりが切れている")
    put(s, 3, "求人サイトには月約9,000人が来ている。それでも、応募の前・応募の後・不採用の後でつながりが切れている。")
    rows = [(6, 7, 4, "① 応募の前", "他社と並べて検討し、「あとで応募しよう」と離れる"),
            (11, 12, 9, "② 応募フォーム", "入力の手間で、途中で離れる"),
            (16, 17, 14, "③ 応募の後", "日程の連絡に時間がかかり、面接に来ない方が出る"),
            (21, 22, 19, "④ 不採用・辞退の後", "その職種で合わなかった方との接点が切れる")]
    for lab, txt, rect, h, t in rows:
        put(s, lab, h)
        put(s, txt, t)
        r = by_id(s, rect)
        move(by_id(s, txt), w=10.28 - 3.62)
        fill_row(s, txt, r.top / 914400, r.height / 914400)
    drop(s, 8, 13, 18, 23, 26)                                 # 右の列（対策は次のページから）と下の注記
    put(s, 25, "広告を増やさずに、すでにサイトに来ている方をLINEでつなぎ止める")
    add_box(s, 0.47, 6.83, 9.89, 0.3, [runs(("※ 求人サイトの月間UU（8,947）は、貴社からいただいた数値", dict(size=9, color="7F7F7F")))],
            anchor="t", ins=(0, 0, 0, 0), name="注記")
    return s


# ---------------------------------------------------------------- 第2章 解決策（職種分岐型）
def p_zentai(prs, src):
    s = body_page(prs, src, "［第2章］解決策｜希望条件で職種を分け、グループ全体の求人につなぐ",
                  "飲食店だけでなく、イベント・警備・施設・正社員まで、ひとつのLINEで職種ごとに案内を分ける。")
    x0, w, aw = 0.47, 9.89, 0.3
    bw = (w - aw * 3) / 4
    y = 1.75
    for k, (t, strong) in enumerate((("求人サイト\n月約9,000人", False), ("LINE\n友だち追加", False),
                                     ("希望条件を選ぶ\n働き方・職種・エリア", False), ("職種ごとに\n案内を分ける", True))):
        x = x0 + k * (bw + aw)
        box(s, x, y, bw, 0.8, t, strong)
        if k < 3:
            arrow(s, x + bw, y, aw, 0.8)
    arrow(s, x0 + 3 * (bw + aw), y + 0.8, bw, 0.35, "▼")
    cw = (w - aw * 2) / 3
    y2 = 2.95
    for k, t in enumerate(("飲食店スタッフ", "イベント設営・警備・\n施設スタッフ", "正社員・総合職")):
        box(s, x0 + k * (cw + aw), y2, cw, 0.75, t)
    add_box(s, x0, y2 + 0.8, w, 0.32, [runs(("職種ごとに、新着の求人・面接のお知らせ・よくある質問を届ける", dict(size=12, color=TEXT)))],
            align="ctr", anchor="ctr", ins=(0, 0, 0, 0), name="補足")
    add_box(s, x0, 4.3, w, 0.45, [runs(("不採用・辞退でも、グループ内の別の職種を案内できる", dict(size=16, bold=True, color="FFFFFF")))],
            fill=NAVY, align="ctr", anchor="ctr", name="示唆")
    shift(s, [sh.shape_id for sh in s.shapes if sh.shape_id not in (2, 3)], 0.55)   # 下の空きを減らし、上寄りの中央へ
    return s


# ---------------------------------------------------------------- 第3章 施策（業界資料のページを直す）
def p_iriguchi(s):
    put(s, 2, "［第3章］施策｜友だちの入口")
    put(s, 3, "広告を増やさずに、いまサイトに来ている方・応募した方をLINEにつなぐ。")
    for sid, t in ((6, "① 応募完了画面"), (7, "応募した直後に友だち追加を案内"), (8, "最も確実。すでに応募した方"),
                   (11, "② 離脱防止ポップアップ"), (12, "求人サイトを離れようとした方にLINEを案内"), (13, "応募フォームより軽い入口"),
                   (16, "③ サイトのLINEボタン"), (17, "求人ページにLINEの入口を置く"), (18, "まだ応募を決めていない方"),
                   (21, "④ 店舗・現場のQRコード"), (22, "店舗・現場のポスターやチラシから直接"), (23, "すでにある接点を使う")):
        put(s, sid, t)
    drop(s, 24, 25, 26, 27, 28, 31)                            # ⑤ 既存リストへの通知メッセージと注記
    put(s, 30, "職種が決まっていない方も、LINEに残しておける")
    shift(s, (29, 30), 4.95 - 5.63)


def p_menu(s):
    put(s, 2, "［第3章］施策｜リッチメニュー（3タブ × 6枠）")
    put(s, 14, "職種を変えて、別の求人へ")
    put(s, 17, "面接・持ち物に先回り")
    put(s, 23, "★ 働いている間は配信を止める")
    put(s, 25, "「通知を止める」「今は働いている」を置くから、ブロックされずにつながり続ける")
    put(s, 26, "※ 画面はイメージ。タブ構成とボタン名は運用しながら調整する")


def p_budomari(s):
    put(s, 2, "［第3章］施策｜面接に来る方を増やす")
    put(s, 3, "応募数も広告費も増やさずに、面接に来る方を増やす。")
    for sid, t in ((6, "① 前日にお知らせ"), (7, "日時・場所・道順・持ち物"), (8, "無断欠席が減る"),
                   (11, "② 不安に先回り"), (12, "面接の流れ・職場の様子を送る"), (13, "面接前の不安を消す"),
                   (16, "③ 日程をその場で確定"), (17, "LINEで候補日を選んでもらう"), (18, "電話の往復で冷めるのを防ぐ"),
                   (21, "④ 合わなければ別の職種へ"), (22, "不採用・辞退の翌日に、グループ内の別の求人"), (23, "気持ちが高いうちに次の応募へ")):
        put(s, sid, t)
    put(s, 25, "目指すのは、応募した方がグループのどこかで働き始めること")
    put(s, 26, "※ 前日のお知らせは、個別チャットか通知メッセージ（別途費用）で送る。友だち追加の翌日以降に届くステップ配信では、面接日に合わせられないため")


# ---------------------------------------------------------------- 第4章 概算お見積り（あらうちv10・タナカホームと同じ形）
def p_mitsumori(prs, src, ar_sim, sim, initial_total):
    n = lambda v: int(v + 0.5)                                        # noqa: E731
    active6 = n(sim["active"][-1])
    s = body_page(prs, src, f"［第4章］概算お見積り：6か月目に追える人数{active6:,}人",
                  "竹（コンサル基本）で運用した場合の費用と、6か月後の見込み。")
    skip = {337, 338, 339, 340, 381} | set(range(383, 390))          # あらうちの見出し・リード／月ごとの小さい表
    tree = s.shapes._spTree
    for el in ar_sim.shapes._spTree:
        if el.tag.endswith("}nvGrpSpPr") or el.tag.endswith("}grpSpPr"):
            continue
        cnv = el.find(".//" + qn("p:cNvPr"))
        if cnv is not None and int(cnv.get("id")) in skip:
            continue
        tree.append(copy.deepcopy(el))
    total = sim["total"]
    put(s, 343, f"¥{initial_total:,}〜")
    put(s, 344, f"¥{n(total[1]):,}〜")
    for (a, b), (k, v) in zip(((365, 366), (367, 368), (369, 370), (371, 372), (373, 374)),
                              (("離脱防止\nポップアップ", "導入"), ("サンクス\nLINE誘導", "導入"), ("ステップ配信", "3通を構築"),
                               ("職種ごとの\n配信", "希望条件で設計"), ("LINE公式\nアカウント", "開設"))):
        put(s, a, k.split("\n"))
        put(s, b, v)
    for (a, b), (k, v) in zip(((349, 350), (351, 352), (353, 354), (355, 356), (357, 358), (359, 360)),
                              (("企画配信", "月3回"), ("ステップ配信", "3通の改善"), ("職種ごとの\n配信", "新着の案内"),
                               ("友だちの入口", "運用・改善"), ("定例会", "月1回"), ("効果の確認", "応募率・面接率"))):
        put(s, a, k.split("\n"))
        put(s, b, v)
    friends6 = n(sum(sim["add"]))                                      # 友だちの数＝毎月の増加の合計
    apply6 = sum(n(x) for x in sim["cv1"])
    for (a, b), (k, v) in zip(((361, 375), (362, 376), (364, 378)),
                              (("追える人数", f"{active6:,}人"), ("友だちの数", f"{friends6:,}人"),
                               ("応募", f"{apply6:,}件/6か月"))):
        put(s, a, k)
        put(s, b, v)
        for clr in by_id(s, b)._element.iter(qn("a:srgbClr")):       # 良い数字は赤い文字にしない
            clr.set("val", NAVY)
    add_box(s, 0.47, 6.83, 9.89, 0.35,
            [runs(("※ 求人サイトの月間UU（8,947）をもとに、友だちの入口ごとの一般的な割合で計算。サイトからの応募は仮に月134件。"
                   "追える人数＝ブロックされずに届く友だちの数。費用は税抜", dict(size=9, color="7F7F7F")))],
            anchor="t", ins=(0, 0, 0, 0), name="注記")
    return s


def p_ryokin(s, rows):
    put(s, 2, "［第4章］運用プラン（6か月〜／税抜）")
    put(s, 3, ["体制と予算に合わせて、3つのプランから選択可能。", "初期費用は初月のみ。2か月目以降は月額のみ。"])
    for ids, vals in zip(((9, 10, 11, 12), (14, 15, 16, 17), (18, 19, 20, 21)), rows):
        for sid, v in zip(ids, vals):
            put(s, sid, v)
    drop(s, 22, 23, 24, 25, 26)                                # 4行目（運用代行・効率改善）
    put(s, 32, ["★通知メッセージ　初期10万＋月8万〜＋通数7円", "ポップアップ（離脱防止）初期1.5万＋月3万",
                "サンクスLINE誘導 初期15万＋月3万〜", "アカウント開設 1万〜／ID連携 20万〜＋月5万〜"])
    put(s, 34, f"まずは{rows[2][0].split('：')[0]}（{rows[2][0].split(' ', 1)[1]}・月額0円）から始めることも可能")
    put(s, 35, "※ すべて税抜・6か月〜の契約。開発やツールの導入は、別途費用がかかる")


def p_support(s):
    put(s, 6, "［第4章］運用サポート体制")
    put(s, 7, ["① 月1回の定例会", "GAレポート連携で、応募率・面接実施率を毎月確認"])
    put(s, 9, ["② 企画配信の制作", "年間カレンダーに沿った配信を、貴社に代わって作成"])
    put(s, 14, "構築だけでなく、運用しながら改善する体制")


def main(base, arauchi, sim_path):
    sim = load_sim(sim_path)
    # load_sim は1枚目のシート（業界資料の金額）を読む。月額・友だち・応募は両案同じで、違うのは初期費用だけ
    for label, plan in PLANS.items():
        prs = Presentation(base)
        sl = list(prs.slides)
        assert by_id(sl[0], 2).text_frame.text == "LINE公式アカウント運用のご提案"
        assert by_id(sl[8], 2).text_frame.text.endswith("友だち追加の動線")
        assert by_id(sl[17], 2).text_frame.text.startswith("［第4章］歩留まり")
        assert by_id(sl[21], 2).text_frame.text.startswith("［第6章］費用プラン")
        cover, agenda, iriguchi, menu, budomari, support, ryokin, back = (sl[i - 1] for i in (1, 2, 9, 12, 18, 20, 22, 23))
        ar = Presentation(arauchi)
        ar_cover, ar_sim = ar.slides[0], ar.slides[16]
        assert by_id(ar_sim, 339).text_frame.text.startswith("目標と費用対効果")

        kadai = p_kadai(prs, budomari)
        zentai = p_zentai(prs, budomari)
        init_total = plan["consult_init"] + 15000 + 150000 + 30000
        mitsumori = p_mitsumori(prs, budomari, ar_sim, sim, init_total)
        p_cover(cover, ar_cover)
        p_agenda(agenda)
        p_iriguchi(iriguchi)
        p_menu(menu)
        p_budomari(budomari)
        p_ryokin(ryokin, plan["rows"])
        p_support(support)

        keep = [cover, agenda, kadai, zentai, iriguchi, menu, budomari, mitsumori, ryokin, support, back]
        lst = prs.slides._sldIdLst
        by_part = {prs.part.related_part(e.rId): e for e in lst}
        for e in list(lst):
            lst.remove(e)
        keep_parts = {s.part for s in keep}
        for part, e in by_part.items():
            if part not in keep_parts:
                prs.part.drop_rel(e.rId)
        for s in keep:
            e = by_part[s.part]
            e.attrib.pop("show", None)
            s._element.attrib.pop("show", None)
            lst.append(e)
        all_meiryo(keep)
        out = ROOT / "202610_清水グループ御中_LINE公式アカウント運用のご提案.pptx"
        prs.save(out)
        print("saved:", out.name, len(prs.slides), "枚")


if __name__ == "__main__":
    main(*sys.argv[1:4])
