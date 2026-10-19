# -*- coding: utf-8 -*-
"""くふう住まい：殿村さんのPC保存版（本編11枚）に「住宅・不動産領域におけるLINEの活用事例」を足す（2026-10-19 殿村さん）。

    python -m _build.patch_kufu_add_jirei <PC保存版.pptx> <事例の元.pptx> <出力.pptx>

  元＝殿村さんから受け取った注文住宅の資料のS30（LINEヤフー公開の4事例）。「サンクスLINE｜導入・活用実績」の次に入れる。
  元ページ（lycbiz.com）と突き合わせて直したところ（殿村さんOK）：
   ・オープンハウスの数字はLINE広告（クリエイティブ改善）の成果 → 見出しを「オープンハウスグループ（LINE広告）」、説明もLINE広告の内容に。
     タイトルも「LINE公式アカウントの活用事例」→「LINEの活用事例」（LINE広告の事例が入るため）
   ・LIFULL HOME'Sの「作業時間 月8時間→2時間」は元ページにそのままの記述が無いので外し、残り2つの枠を横に広げる
   ・リードの「数値は同社の公開値をそのまま転記…丸めたり、脚色したりはしません」と下の帯「公開値のみ…二次情報は使用していません」を消す
   ・数字のオレンジ（ED7D31）は紺に（黄色・オレンジは使わない。CLAUDE.md 2026-10-19）
   ・出典の確認日を 2026年10月19日に
  ほかの11枚は触らない。
"""
import sys

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches

from _build.arauchi_v10.build import by_id
from _build.build_sunao_v1 import all_meiryo
from _build.patch_tanaka_v4 import set_lines
from _build.patch_tanaka_v5 import copy_slide

NAVY = "1F285A"


def drop(s, *ids):
    for sid in ids:
        e = by_id(s, sid)._element
        e.getparent().remove(e)


def main(base, src_path, out):
    prs = Presentation(base)
    sl = list(prs.slides)
    assert len(sl) == 11
    assert any(sh.has_text_frame and "導入・活用実績" in sh.text_frame.text for sh in sl[4].shapes)
    src = Presentation(src_path).slides[29]
    assert by_id(src, 2).text_frame.text == "住宅・不動産領域におけるLINE公式アカウントの活用事例"

    s = copy_slide(prs, src)
    set_lines(by_id(s, 2), ["住宅・不動産領域におけるLINEの活用事例"])
    set_lines(by_id(s, 3), ["LINEヤフー社が公開している、住宅・不動産領域の活用事例"])
    set_lines(by_id(s, 21), ["オープンハウスグループ（LINE広告）"])
    set_lines(by_id(s, 23), ["LINE広告の運用をインハウス化し、エリアごとの顧客データと営業担当の声をクリエイティブに反映。"])
    drop(s, 44, 45, 46, 61, 62)                          # LIFULLの「作業時間」の枠／下の帯
    # LIFULLの残り2つの枠を、カードの幅いっぱいに並べ直す
    x0, w = 0.78, 2.12
    for k, (box, lab, val) in enumerate(((38, 39, 40), (41, 42, 43))):
        x = x0 + k * (w + 0.1)
        by_id(s, box).left, by_id(s, box).width = Inches(x), Inches(w)
        for sid in (lab, val):
            by_id(s, sid).left, by_id(s, sid).width = Inches(x + 0.04), Inches(w - 0.08)
    for clr in s._element.iter(qn("a:srgbClr")):         # 数字のオレンジ → 紺
        if clr.get("val") == "ED7D31":
            clr.set("val", NAVY)
    set_lines(by_id(s, 63), ["※出典：LINEヤフー for Business 導入事例（lycbiz.com）／2026年10月19日確認。"
                             "ミサワホーム北越は「その他業界の活用事例」ページに掲載。"])
    all_meiryo([s])

    lst = prs.slides._sldIdLst                           # 「サンクスLINE｜導入・活用実績」の次へ
    new = lst[-1]
    lst.remove(new)
    lst.insert(5, new)
    prs.save(out)
    print("saved:", out, len(prs.slides), "枚")


if __name__ == "__main__":
    main(*sys.argv[1:4])
