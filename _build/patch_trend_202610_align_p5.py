"""
10月号 P.3〜P.13 の本文のまとまりを、P.5 に合わせる（2026-10-26 殿村さん指示）。
- 横：まとまり全体を P.5 の枠と同じ 左端・幅 にする（中の枠・表は比率どおり伸縮。表は列幅も）
- 縦：一番上の枠の上端を P.5 の一番上の枠と同じ位置にする（間隔・高さはそのまま）
P.14 は触らない。
使い方: python3 _build/patch_trend_202610_align_p5.py <pptx>
"""
import sys
from pptx import Presentation
from pptx.util import Inches, Emu

F = sys.argv[1]
prs = Presentation(F)
slides = list(prs.slides)


def text(sh):
    return sh.text_frame.text.strip() if sh.has_text_frame else ""


def body(slide):
    head = next(x for x in slide.shapes if text(x) == "■ 詳細")
    return [x for x in slide.shapes if x.top >= head.top + head.height - Inches(0.02) and x.top < Inches(7.0)]


ref = body(slides[4])
REF_L = min(x.left for x in ref)
REF_R = max(x.left + x.width for x in ref)
REF_T = min(x.top for x in ref)
print("P.5", round(Emu(REF_L).inches, 2), round(Emu(REF_R - REF_L).inches, 2), round(Emu(REF_T).inches, 2))

for n in range(3, 14):
    if n == 5:
        continue
    shapes = body(slides[n - 1])
    l = min(x.left for x in shapes)
    r = max(x.left + x.width for x in shapes)
    t = min(x.top for x in shapes)
    k = (REF_R - REF_L) / (r - l)
    for x in shapes:
        new_left = REF_L + int((x.left - l) * k)
        new_w = int(x.width * k)
        if x.has_table:
            cols = x.table.columns
            total = sum(c.width for c in cols)
            acc = 0
            for i, c in enumerate(cols):
                w = new_w - acc if i == len(cols) - 1 else int(c.width * new_w / total)
                c.width = w
                acc += w
        x.left, x.width = new_left, new_w
        x.top = x.top + (REF_T - t)
    print(f"P.{n}", round(k, 3), round(Emu(REF_T - t).inches, 2))

prs.save(F)
print("saved", F)
