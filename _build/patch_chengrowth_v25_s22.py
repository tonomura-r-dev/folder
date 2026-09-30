# -*- coding: utf-8 -*-
"""チェングロウス本資料 ver2.5（殿村さんPC保存版）の22枚目「前後検索（近い業種：工場求人ナビ）」を、
カーディーラー（23枚目）と同じく色分けした画像に差し替え、凡例（3色）を付ける。

  python3 _build/make_chengrowth_zengo_kojo3.py      # 色分け画像を作る
  python3 _build/patch_chengrowth_v25_s22.py         # ver2.5 を上書き

凡例のチップは23枚目のチップ（職業・働き方／整備・工具／お金・手続き）を複製して、文字だけ差し替える。
"""
import copy
from pathlib import Path

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Emu

ROOT = Path(__file__).resolve().parent.parent
DECK = ROOT / "20260929_株式会社チェングロウス御中_LINE公式アカウント運用のご提案ver2.5.pptx"
IMG = ROOT / "_images/chengrowth_zengo_kojo_3color.png"
CM = 360000

prs = Presentation(str(DECK))
s22, s23 = prs.slides[21], prs.slides[22]
by = lambda s, n: next(x for x in s.shapes if x.name == n)

# 23枚目のチップ：13＝オレンジ、14＝青、15＝紫
CHIPS = [("Rounded Rectangle 14", "条件"),
         ("Rounded Rectangle 13", "評判・不安"),
         ("Rounded Rectangle 15", "応募の準備")]

pic = by(s22, "Picture 4")
left, top, width, height = pic.left, pic.top, pic.width, pic.height   # 1.2, 4.4, 15.8, 9.22 cm
chip_h, gap = int(0.62 * CM), int(0.2 * CM)
cw = (width - gap * 2) // 3

# 凡例チップ（画像の上）
next_id = max(int(x.shape_id) for x in s22.shapes) + 1
for i, (name, text) in enumerate(CHIPS):
    el = copy.deepcopy(by(s23, name)._element)
    el.find(".//" + qn("p:cNvPr")).set("id", str(next_id + i))
    el.find(".//" + qn("p:cNvPr")).set("name", f"Legend {i + 1}")
    off = el.find(".//" + qn("a:off")); ext = el.find(".//" + qn("a:ext"))
    off.set("x", str(left + i * (cw + gap))); off.set("y", str(top))
    ext.set("cx", str(cw)); ext.set("cy", str(chip_h))
    ts = list(el.iter(qn("a:t")))
    ts[0].text = text
    for t in ts[1:]:
        t.text = ""
    s22.shapes._spTree.append(el)

# 画像：色分け版に差し替え、チップの下に収まるよう縮小（縦横比は維持・横は中央寄せ）
new_top = top + chip_h + int(0.15 * CM)
new_h = top + height - new_top
new_w = int(new_h * width / height)
new_left = left + (width - new_w) // 2
_, rid = s22.part.get_or_add_image_part(str(IMG))
pic._element.find(".//" + qn("a:blip")).set(qn("r:embed"), rid)
pic.left, pic.top, pic.width, pic.height = Emu(new_left), Emu(new_top), Emu(new_w), Emu(new_h)

# 出典に色分けの注記を足す
src = by(s22, "TextBox 9")
ts = list(src._element.iter(qn("a:t")))
if "色分け" not in "".join(t.text for t in ts):
    ts[-1].text += "。色分けはDYM分類（サイト名・求人検索は色なし）"

prs.save(str(DECK))
print("saved:", DECK.name)
