# -*- coding: utf-8 -*-
"""「Lステップとは」1枚（機能＋費用感）をつくる（2026-10-14 殿村さん指示）。

    python -m _build.lstep.build [出力.pptx]

1. DYM汎用FMT（PowerPointで保存したもの）の2枚目に、中身を組む（python-pptx。メモリ上だけ）
2. FMTのzipから、2枚目以外のスライド（とノート・孤立したパーツ）を外し、2枚目のXMLだけ差し替える
   （python-pptxでzipごと保存し直さない＝PowerPointの「修復」対策。arauchi_v8 と同じ考え方）

出典：Lステップ公式サイト（https://linestep.jp/lp/01/）の料金表・機能表と、サービス資料（linestep.pdf）の機能解説。
サービス資料の料金ページは画像で文字が読めないため、料金・配信数・使えるプランは公式サイトの表を使った。
"""
import posixpath
import re
import sys
import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn

from _build.arauchi_v3.common import delete, set_paras
from _build.arauchi_v4.build import normalize_decl, slide_order
from _build.arauchi_v4.lib import NAVY, TEXT, add_box

ROOT = Path(__file__).resolve().parents[2]
BASE = ROOT / "_templates" / "DYM_LINEOA_FMT.pptx"
OUT = ROOT / "20261014_Lステップとは_機能と費用感.pptx"
KEEP = 1                       # FMTの2枚目（0始まり）
SRC_URL = "https://linestep.jp/lp/01/"

# ---------------------------------------------------------------- 色（あらうち資料v8のS18と同じ）
C_NAVY, C_BLUE, C_SLATE = "1F2A5C", "4A6FA5", "5A6378"
C_BODY, C_LINE, C_MUTE = "F5F7FA", "D1D8E0", "5A6378"
WHITE = "FFFFFF"

X0, W = 0.47, 9.76            # 内容の左端と全体の幅

# ---------------------------------------------------------------- 中身（文言）
# 3版目（2026-10-14 殿村さん指示）：機能＝1版目の中身（各4項目＋ひとこと）に戻す／費用感＝プランごとのひとことを消し、カードを縦横とも小さく
LEAD = [
    "LINE公式アカウントの機能を拡張する、LINE専用のツールです。",
    "月額0円・初期費用なしから始められます。",
]
FEATURES = [  # (色, 見出し, ひとこと, [(機能名, 説明), ...])
    (C_NAVY, "① 配信を自動化", "お客様の行動に合わせて自動で届ける",
     [("ステップ配信", "順番に自動で配信"), ("セグメント配信", "狙った方に配信"),
      ("リマインダ配信", "予約前に自動通知"), ("自動応答", "質問に自動で返信")]),
    (C_BLUE, "② お客様を管理", "友だちごとの情報と対応状況を一元管理",
     [("友だち情報", "属性や行動を記録"), ("個別トーク", "1対1でチャット対応"),
      ("回答フォーム", "アンケート・申込受付"), ("予約管理", "LINEで予約を受付")]),
    (C_SLATE, "③ 効果を分析", "どこから来て何に反応したかを把握",
     [("流入経路分析", "どこから追加したか"), ("URLクリック測定", "誰が押したか"),
      ("ファネル分析", "離脱ポイントを把握"), ("行動スコアリング", "関心度を点数化")]),
]
PLANS = [  # (名前, 色, 月額, 配信数)
    ("フリープラン", "7A8294", "0", "月200通まで"),
    ("スタートプラン", C_BLUE, "5,000", "月5,000通まで"),
    ("スタンダードプラン", "2F4A82", "21,780", "月30,000通まで"),
    ("プロプラン", C_NAVY, "32,780", "月50,000通まで"),
]
CALLOUT = "※別途、LINE公式\nアカウントの\n月額費用が\nかかります。"      # 折り返し位置を決めてある（「ます。」だけ残らないように）
NOTES = [
    [("※初期費用0円／表示価格は税込／解約は3か月後（3回課金後）から可能です。", {})],
    [("出典：", {}), ("Lステップ公式サイト", dict(link=SRC_URL)), ("・サービス資料", {})],
]


# ---------------------------------------------------------------- スライドを組む
def _spacing(shape, befores, default=0):
    """段落ごとの「段落前の間隔」（pt）。befores＝{段落の番号: pt}"""
    for k, pt in enumerate(shape._element.findall(".//" + qn("a:p"))):
        sb = etree.Element(qn("a:spcBef"))
        etree.SubElement(sb, qn("a:spcPts"), val=str(int(befores.get(k, default) * 100)))
        pt.find(qn("a:pPr")).insert(0, sb)


def compose(slide):
    keep = {4: "lead", 5: "line", 45: "title"}
    for sh in list(slide.shapes):
        if sh.shape_id not in keep:
            delete(sh)
    shp = {keep[sh.shape_id]: sh for sh in slide.shapes}

    set_paras(shp["title"], ["Lステップとは"])
    for rpr in shp["title"]._element.iter(qn("a:rPr")):
        rpr.set("lang", "ja-JP")
        rpr.attrib.pop("altLang", None)
    set_paras(shp["lead"], [[(LEAD[0], False)], [(LEAD[1], True)]])

    def rr(shape, adj):
        shape.adjustments[0] = adj
        return shape

    def chip(y, label, sub):
        rr(add_box(slide, X0, y, 1.10, 0.32, [label], size=14, bold=True, color=WHITE, fill=C_NAVY, line=C_NAVY,
                   align="ctr", shape="round", ins=(0.05, 0.02, 0.05, 0.02), name=f"見出し：{label}"), 0.25)
        add_box(slide, X0 + 1.20, y, W - 1.20, 0.32, [sub], size=13, bold=True, color=NAVY, align="l",
                ins=(0.04, 0.02, 0.04, 0.02), name=f"見出しの説明：{label}")

    # ---- 機能 ----
    chip(1.64, "機能", "Lステップの3つの特徴｜配信の自動化・顧客管理・効果分析")
    cw, gap = 3.05, (W - 3 * 3.05) / 2
    for i, (col, title, sub, items) in enumerate(FEATURES):
        x = X0 + i * (cw + gap)
        rr(add_box(slide, x, 2.02, cw, 0.62, [[(title, dict(size=14, bold=True))], [(sub, dict(size=10.5, bold=False))]],
                   size=14, bold=True, color=WHITE, fill=col, line=col, align="l", shape="round",
                   ins=(0.14, 0.03, 0.10, 0.03), name=f"機能の見出し{i + 1}"), 0.12)
        paras = [[("・" + n, dict(bold=True, color=NAVY)), ("｜" + g, {})] for n, g in items]
        body = add_box(slide, x, 2.68, cw, 1.40, paras, size=11, color=TEXT, fill=C_BODY, line=C_LINE, align="l",
                       anchor="t", shape="round", ins=(0.12, 0.09, 0.08, 0.04), name=f"機能の中身{i + 1}")
        rr(body, 0.08)
        _spacing(body, {0: 0}, default=4)

    # ---- 費用感 ----
    chip(4.40, "費用感", "月額料金（税込）｜上位プランほど機能が増えます")
    pw, pg = 1.85, 0.12
    for i, (name, col, price, cnt) in enumerate(PLANS):
        x = X0 + i * (pw + pg)
        rr(add_box(slide, x, 4.78, pw, 0.32, [name], size=12.5, bold=True, color=WHITE, fill=col, line=col, align="ctr",
                   shape="round", ins=(0.04, 0.02, 0.04, 0.02), name=f"プラン名{i + 1}"), 0.2)
        paras = [
            [(price, dict(size=20, bold=True, color=NAVY)), ("円", dict(size=12, bold=True, color=NAVY)), ("／月", dict(size=10.5, color=C_MUTE))],
            [(cnt, dict(size=11.5, bold=True))],
        ]
        box = add_box(slide, x, 5.14, pw, 0.86, paras, size=11.5, color=TEXT, fill=C_BODY, line=C_LINE, align="ctr",
                      anchor="t", shape="round", ins=(0.06, 0.07, 0.06, 0.04), name=f"プランの中身{i + 1}")
        rr(box, 0.08)
        _spacing(box, {0: 0, 1: 2})
    cx = X0 + 4 * pw + 3 * pg + 0.22                                   # 右の空いたところに「別途」の注意
    rr(add_box(slide, cx, 4.78, X0 + W - cx, 1.22, [CALLOUT], size=11, bold=True, color=NAVY, fill=C_BODY, line=C_LINE,
               align="ctr", anchor="ctr", shape="round", ins=(0.08, 0.06, 0.08, 0.06), name="別途費用の注意"), 0.08)

    # ---- 注記・出典 ----
    add_box(slide, X0, 6.18, W, 0.60, NOTES, size=11, color="444444", align="l", anchor="t",
            ins=(0.04, 0.02, 0.04, 0.02), name="注記・出典")


# ---------------------------------------------------------------- パッケージ（FMTのzipから必要なパーツだけを残す）
def _rels_of(files, part):
    d, f = posixpath.split(part)
    return posixpath.join(d, "_rels", f + ".rels")


def _reachable(files):
    """_rels/.rels から関係をたどって、つながっているパーツ名を返す"""
    seen, todo = set(), ["_rels/.rels"]
    while todo:
        rels_name = todo.pop()
        if rels_name in seen or rels_name not in files:
            continue
        seen.add(rels_name)
        src_dir = posixpath.dirname(posixpath.dirname(rels_name))              # …/_rels/x.rels → …
        for r in etree.fromstring(files[rels_name]):
            if r.get("TargetMode") == "External":
                continue
            tgt = r.get("Target")
            part = tgt.lstrip("/") if tgt.startswith("/") else posixpath.normpath(posixpath.join(src_dir, tgt))
            if part in files and part not in seen:
                seen.add(part)
                todo.append(_rels_of(files, part))
    return seen


def package(base_path, built_path, out_path):
    zb, zn = zipfile.ZipFile(base_path), zipfile.ZipFile(built_path)
    files = {n: zb.read(n) for n in zb.namelist()}
    order = list(files)

    b_name, n_name = slide_order(zb)[KEEP], slide_order(zn)[KEEP]
    b_rels, n_rels = _rels_of(files, b_name), _rels_of(files, n_name)

    # ---- 中身を組んだスライドのXMLと、そのrels（ノートへの関係は外す）
    files[b_name] = normalize_decl(zn.read(n_name))
    base_layout = [r.get("Target") for r in etree.fromstring(zb.read(b_rels)) if r.get("Type").endswith("/slideLayout")]
    rels_root = etree.fromstring(zn.read(n_rels))
    for r in list(rels_root):
        if r.get("Type").endswith("/notesSlide"):
            rels_root.remove(r)
    assert [r.get("Target") for r in rels_root if r.get("Type").endswith("/slideLayout")] == base_layout
    files[b_rels] = etree.tostring(rels_root, encoding="UTF-8", xml_declaration=True, standalone=True)

    # ---- presentation.xml / rels：残すのは1枚だけ
    pres = files["ppt/presentation.xml"].decode("utf-8")
    entries = re.findall(r"<p:sldId [^>]*/>", pres)
    keep_entry = entries[KEEP]
    keep_rid = re.search(r'r:id="(rId\d+)"', keep_entry).group(1)
    pres = re.sub(r"<p:sldIdLst>.*?</p:sldIdLst>", lambda m: "<p:sldIdLst>" + keep_entry + "</p:sldIdLst>", pres, flags=re.S)
    files["ppt/presentation.xml"] = pres.encode("utf-8")
    prels = etree.fromstring(files["ppt/_rels/presentation.xml.rels"])
    for r in list(prels):
        if r.get("Type").endswith("/slide") and r.get("Id") != keep_rid:
            prels.remove(r)
    files["ppt/_rels/presentation.xml.rels"] = etree.tostring(prels, encoding="UTF-8", xml_declaration=True, standalone=True)

    # ---- つながっていないパーツ（外したスライド・ノート・それだけが使う画像）を落とす
    live = _reachable(files) | {"[Content_Types].xml"}
    gone = [n for n in order if n not in live]
    for n in gone:
        del files[n]

    # ---- [Content_Types].xml：外したパーツの Override を消す
    ct = files["[Content_Types].xml"].decode("utf-8")
    for n in gone:
        ct = re.sub(r'<Override PartName="/%s"[^>]*/>' % re.escape(n), "", ct)
    files["[Content_Types].xml"] = ct.encode("utf-8")

    # ---- docProps/app.xml：スライド数・タイトル一覧
    app = files["docProps/app.xml"].decode("utf-8")
    app = re.sub(r"<Slides>\d+</Slides>", "<Slides>1</Slides>", app)
    app = re.sub(r"<Notes>\d+</Notes>", "<Notes>0</Notes>", app)
    app = re.sub(r"<Words>\d+</Words>|<Paragraphs>\d+</Paragraphs>", "", app)
    m = re.search(r"<vt:lpstr>スライド タイトル</vt:lpstr></vt:variant><vt:variant><vt:i4>(\d+)</vt:i4>", app)
    n_old = int(m.group(1))
    app = app.replace(m.group(0), m.group(0).replace(f"<vt:i4>{n_old}</vt:i4>", "<vt:i4>1</vt:i4>"))
    t = re.search(r'<TitlesOfParts><vt:vector size="(\d+)" baseType="lpstr">(.*?)</vt:vector></TitlesOfParts>', app, flags=re.S)
    parts = re.findall(r"<vt:lpstr>.*?</vt:lpstr>", t.group(2))
    parts = parts[:len(parts) - n_old] + ["<vt:lpstr>PowerPoint プレゼンテーション</vt:lpstr>"]
    app = app.replace(t.group(0), '<TitlesOfParts><vt:vector size="%d" baseType="lpstr">%s</vt:vector></TitlesOfParts>' % (len(parts), "".join(parts)))
    files["docProps/app.xml"] = app.encode("utf-8")

    with zipfile.ZipFile(out_path, "w", zipfile.ZIP_DEFLATED) as zo:
        for n in order:
            if n in files:
                zo.writestr(n, files[n])
    return gone


def main(out=None):
    out = str(out or OUT)
    tmp = out + ".tmp.pptx"
    prs = Presentation(str(BASE))
    compose(prs.slides[KEEP])
    prs.save(tmp)
    gone = package(str(BASE), tmp, out)
    Path(tmp).unlink()
    print("saved:", out)
    print("外したパーツ:", len(gone))


if __name__ == "__main__":
    main(*sys.argv[1:2])
