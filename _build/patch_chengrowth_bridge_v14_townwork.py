# -*- coding: utf-8 -*-
"""チェングロウス ver1.3 → ver1.4（2026-10-05 殿村さん指示）。
10枚目（タウンワーク事例）を殿村さんの文面どおりに作り直す：応募数＋79%を主役に、ポイント3つ＋結論1行。
  python3 _build/patch_chengrowth_bridge_v14_townwork.py <ver1.3.pptx>
"""
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

sys.path.insert(0, str(Path(__file__).resolve().parent))
from pptx_parts import *  # noqa

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "20261005_チェングロウス_サイトリニューアルとLINE同時導入のご提案ver1.4.pptx"
prs = Presentation(sys.argv[1])
s = prs.slides[9]
for sh in list(s.shapes):
    if sh.name not in ("TextBox 1", "TextBox 2", "Connector 3", "TextBox 14"):
        sh._element.getparent().remove(sh._element)
by = {sh.name: sh for sh in s.shapes}
set_paras(by["TextBox 1"], [("【事例】求人業界におけるLINE活用", 0)])
set_paras(by["TextBox 2"], [("タウンワークでは、LINEから求人を探せる導線を強化", 0)])
set_paras(by["TextBox 14"], [("出典：LINEヤフー for Business 導入事例（タウンワーク・2019年公開）", 0)])

# 左：応募数 ＋79%（主役）
label(s, 1.8, 5.9, 10.6, 1.0, [("応募数", 20, True, GRAY, 0)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
label(s, 1.8, 6.8, 10.6, 3.2, [("＋79%", 68, True, NAVY, 0)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
vline(s, 12.6, 5.4, 12.4, LGRAY, 1.0)
# 右：応募増加につながったポイント（縦の3ステップ）
label(s, 13.6, 5.0, 11.7, 0.9, [("応募増加につながったポイント", 15, True, NAVY, 0)], anchor=MSO_ANCHOR.MIDDLE)
pts = ["リッチメニューで条件を選ぶ", "条件に合う求人を表示", "そのまま求人ページへ移動"]
cy0, step = 7.0, 2.1
vline(s, 14.3, cy0 + 0.3, cy0 + 2 * step + 0.3, NAVY, 2.0)
for i, t in enumerate(pts):
    cy = cy0 + i * step
    shape(s, MSO_SHAPE.OVAL, 13.7, cy - 0.3, 1.2, 1.2, [(f"{i + 1}", 16, True, WHITE, 0)], fill=NAVY, margins=(0, 0, 0, 0))
    label(s, 15.4, cy - 0.35, 9.9, 1.3, [(t, 17, True, INK, 0)], anchor=MSO_ANCHOR.MIDDLE)
# 結論
hline(s, 2.2, 25.3, 14.0, NAVY, 1.5)
label(s, 2.2, 14.15, 23.1, 1.3, [("LINEを「求人を探す入口」にしたことが、応募増加へ", 16, True, NAVY, 0)],
      align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
prs.save(str(OUT))
print("saved:", OUT.name, len(prs.slides), "枚")
