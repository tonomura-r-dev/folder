# -*- coding: utf-8 -*-
"""サンクスLINE誘導｜4枚のざっくり提案資料（16:9）。2026-10-05 殿村さん指示。
1枚目＝何か／2枚目＝なぜ必要か／3枚目＝どう売上化するか／4枚目＝どうLTVにつながるか。
数字は想定例だけ（100CV × 誘導可能率60% × 友だち追加率50% ＝ 30友だち）。改善率は書かない。「倍」は使わない。
DYMテンプレ（チェングロウス資料のレイアウト）を16:9に広げて使う。
  python3 _build/build_thanks_line_4p_cv.py <DYMテンプレを使った既存pptx>
"""
import copy
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Cm

sys.path.insert(0, str(Path(__file__).resolve().parent))
from pptx_parts import *  # noqa

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "20261005_サンクスLINE誘導のご提案.pptx"
PALE = "E4E8F6"

prs = Presentation(sys.argv[1])
OLD_W = prs.slide_width
W = Cm(33.867)
DX = W - OLD_W
SW = W / 360000  # cm
L, R = 2.2, SW - 2.2
CW = R - L
ref = next(s for s in prs.slides if any(sh.name == "Connector 3" for sh in s.shapes))
layout = ref.slide_layout

# ---- 16:9 に広げる（レイアウトの上部ライン・マスターのロゴ／フッター）----
prs.slide_width = W
for sh in layout.shapes:
    if sh.name == "直線コネクタ 4":
        sh.width = W
for sh in layout.slide_master.shapes:
    if sh.name in ("図 1", "図 4", "スライド番号プレースホルダ 1"):
        sh.left = sh.left + DX
    elif sh.name == "正方形/長方形 5":
        sh.width = W


def new_slide(title, lead):
    s = prs.slides.add_slide(layout)
    for sh in list(s.shapes):
        sh._element.getparent().remove(sh._element)
    for nm in ("TextBox 1", "TextBox 2", "Connector 3"):
        s.shapes._spTree.append(copy.deepcopy(next(x for x in ref.shapes if x.name == nm)._element))
    by = {x.name: x for x in s.shapes}
    by["TextBox 1"].width = by["TextBox 1"].width + DX
    by["TextBox 2"].width = by["TextBox 2"].width + DX
    by["Connector 3"].width = W
    set_paras(by["TextBox 1"], [(title, 0)])
    set_paras(by["TextBox 2"], [(lead, 0)])
    return s


def flow(s, x0, x1, y, h, steps, size=13, ov=0.35, big=None):
    """steps = [(行のリスト, 塗り, 文字色)]。big=強調する段の番号（少し大きく）"""
    n = len(steps)
    w = (x1 - x0 + (n - 1) * ov) / n
    for i, (ls, f, tc) in enumerate(steps):
        x = x0 + i * (w - ov)
        yy, hh = (y - 0.4, h + 0.8) if i == big else (y, h)
        shape(s, MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON, x, yy, w, hh, None, fill=f, adj=0.25)
        pad_l = 0.3 if i == 0 else 0.7
        label(s, x + pad_l, yy, w - pad_l - 0.6, hh, [(t, size + (1.5 if i == big else 0), True, tc, 0) for t in ls],
              align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    return w, ov


def note(s, text="※数値は想定例"):
    label(s, L, 17.2, CW, 0.6, [(text, 9, False, GRAY, 0)])


def chips(s, x, y, items, per_row, cw, ch=0.95, gap=0.3, fill=CARD, tc=NAVY, size=11.5):
    for i, t in enumerate(items):
        r, c = divmod(i, per_row)
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x + c * (cw + gap), y + r * (ch + 0.25), cw, ch, [(t, size, True, tc, 0)],
              fill=fill, adj=0.5, margins=(0.15, 0, 0.15, 0))


# ================= 1 サンクスLINE誘導とは =================
s = new_slide("CV後の接点をLINEに残す｜サンクスLINE誘導とは", "サイト・LPでCVした方をLINEの友だちにし、その後の接点をつくる")
flow(s, L, R, 4.9, 2.6, [(["サイト・LP", "来訪"], LGRAY, INK), (["CV"], LGRAY, INK), (["完了画面"], LGRAY, INK),
                         (["サンクス", "LINE誘導"], NAVY, WHITE), (["LINE", "友だち追加"], GREEN, WHITE), (["継続", "フォロー"], PALE, NAVY)],
     size=13, big=4)
w1 = (CW + 5 * 0.35) / 6
label(s, L + (w1 - 0.35) * 1 - 1.5, 7.6, w1 + 3.0, 0.9, [("予約・申込・問い合わせ・会員登録など", 9.5, False, GRAY, 0)], align=PP_ALIGN.CENTER)
# 数値：100CV → 60人 → 30人
cols = [("100", "CV", INK, ""), ("60", "人", NAVY, "サンクスLINE誘導の対象"), ("30", "人", GREEN_TX, "LINE友だち追加")]
bx = [L + 0.5, L + 10.6, L + 20.7]
for (num, unit, c, cap), x in zip(cols, bx):
    rich(s, x, 9.3, 8.3, 2.4, [[(num, 48, True, c), (unit, 20, True, c)]], align=PP_ALIGN.CENTER)
    if cap:
        label(s, x, 11.6, 8.3, 0.7, [(cap, 11.5, True, c, 0)], align=PP_ALIGN.CENTER)
for x, t in [(L + 8.5, "× 誘導可能率 60%"), (L + 18.6, "× 友だち追加率 50%")]:
    arrow_line(s, x + 0.2, 10.5, x + 2.4, 10.5, NAVY, 2.0)
    label(s, x - 1.2, 11.0, 5.0, 0.7, [(t, 10.5, True, NAVY, 0)], align=PP_ALIGN.CENTER)
hline(s, L, R, 13.4, LGRAY, 0.75)
rich(s, L, 13.7, CW, 1.2, [[("100CV", 15, True, INK), ("　×　誘導可能率60%　×　友だち追加率50%　＝　", 14, False, INK),
                            ("30友だち", 18, True, GREEN_TX)]], align=PP_ALIGN.CENTER)
note(s, "※数値は想定例（誘導可能率・友だち追加率は想定値）")

# ================= 2 CVはゴールではなく、売上化のスタート地点 =================
s = new_slide("CVはゴールではなく、売上化のスタート地点", "CV後をLINEでつなぐことで、来店・購入・成約まで届ける")
LW = 4.6
label(s, L, 4.6, LW, 2.2, [("サンクスLINE", 12, True, GRAY, 0), ("なし", 16, True, GRAY, 0)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
flow(s, L + LW + 0.3, R, 4.7, 2.0, [(["CV"], LGRAY, INK), (["その後の", "接点なし"], LGRAY, INK), (["離脱"], LGRAY, INK),
                                    (["来店・購入に", "つながらない"], "F2F2F2", GRAY)], size=12.5)
label(s, L, 7.6, LW, 2.2, [("サンクスLINE", 12, True, NAVY, 0), ("あり", 16, True, GREEN_TX, 0)], align=PP_ALIGN.CENTER,
      anchor=MSO_ANCHOR.MIDDLE)
flow(s, L + LW + 0.3, R, 7.7, 2.0, [(["CV"], PALE, NAVY), (["LINE", "友だち化"], GREEN, WHITE), (["リマインド・", "フォロー"], PALE, NAVY),
                                    (["来店・購入", "・成約"], NAVY, WHITE), (["再来店・", "再購入"], NAVY, WHITE)], size=12.5)
label(s, L + LW + 0.3, 10.0, CW - LW - 0.3, 0.7, [("例：100CVのうち30人をLINE友だち化できれば、その30人と継続して接点を持てる（想定）", 11, False, INK, 0)])
hline(s, L, R, 11.2, LGRAY, 0.75)
pts = [("実売上化", "CV後の離脱を防ぎ、来店・購入へ"), ("顧客ストック化", "獲得したCVを、LINE上の接点として蓄積"),
       ("LTV向上", "継続フォローで、再来店・再購入へ")]
pw = CW / 3
for i, (h, d) in enumerate(pts):
    x = L + i * pw
    if i:
        vline(s, x, 11.7, 15.6, LGRAY, 0.75)
    rect(s, x + 0.6, 11.8, 2.0, 0.14, GREEN if i == 1 else NAVY)
    label(s, x + 0.5, 12.15, pw - 1.0, 3.4, [(h, 17, True, NAVY, 6), (d, 12, False, INK, 0)])
note(s)

# ================= 3 具体施策① CV後の離脱を防ぎ、来店・成約につなげる =================
s = new_slide("具体施策①｜CV後の離脱を防ぎ、来店・成約につなげる", "CV直後から、来店・面談・購入までの行動をLINEで後押し")
flow(s, L, R, 4.9, 2.6, [(["CV"], LGRAY, INK), (["LINE追加"], GREEN, WHITE), (["受付完了", "メッセージ"], PALE, NAVY),
                         (["前日", "リマインド"], PALE, NAVY), (["日程変更・", "不安解消"], PALE, NAVY), (["来店・面談", "・購入"], NAVY, WHITE)],
     size=13, big=1)
# 左：対象になる人数
rich(s, L, 9.0, 11.5, 1.0, [[("100CVのうち", 13, True, GRAY)]], align=PP_ALIGN.CENTER)
rich(s, L, 9.9, 11.5, 2.6, [[("30", 54, True, GREEN_TX), ("人", 22, True, GREEN_TX)]], align=PP_ALIGN.CENTER)
label(s, L, 12.5, 11.5, 1.6, [("がLINE友だち化（想定）", 13, True, NAVY, 3), ("この30人に、来店・面談・購入までフォローできる", 11.5, False, INK, 0)],
      align=PP_ALIGN.CENTER)
vline(s, L + 12.3, 9.0, 14.6, LGRAY, 0.75)
label(s, L + 13.2, 9.0, 15.0, 0.8, [("施策の例", 13, True, NAVY, 0)], anchor=MSO_ANCHOR.MIDDLE)
chips(s, L + 13.2, 10.1, ["受付完了メッセージ", "前日・当日リマインド", "日程変更の導線", "よくある質問", "不安解消コンテンツ", "来店前のご案内"],
      per_row=3, cw=5.1, ch=1.1)
note(s, "※数値は想定例（改善率を示すものではない）")

# ================= 4 具体施策② LINEで育成し、LTV向上につなげる =================
s = new_slide("具体施策②｜LINEで育成し、LTV向上につなげる", "初回で終わらせず、継続フォローで再来店・再購入へ")
stairs = [("初回来店・", "購入・成約"), ("LINE", "継続接点"), ("フォロー", "配信"), ("再来店・", "再購入"), ("アップセル／", "クロスセル"), ("LTV", "向上")]
n = len(stairs)
gw, gap = (CW - 0.3 * (n - 1)) / n, 0.3
base = 10.6
for i, (a, b) in enumerate(stairs):
    h = 2.0 + i * 0.55
    x = L + i * (gw + gap)
    f, tc = (GREEN, WHITE) if i == n - 1 else ((NAVY, WHITE) if i >= 3 else (PALE, NAVY))
    rect(s, x, base - h, gw, h, f)
    label(s, x, base - h, gw, h, [(a, 12.5, True, tc, 0), (b, 12.5, True, tc, 0)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
arrow_line(s, L + 0.5, base + 0.45, R - 0.3, base + 0.45, NAVY, 1.5)
label(s, L, 11.15, 10.0, 0.8, [("施策の例", 13, True, NAVY, 0)], anchor=MSO_ANCHOR.MIDDLE)
chips(s, L, 12.05, ["来店後フォロー", "再来店促進", "再購入促進", "関連サービス案内", "クーポン配信", "セグメント配信"], per_row=6,
      cw=(CW - 0.3 * 5) / 6, ch=0.95, size=11)
eq = [("広告でCVを増やす", LGRAY, INK), ("サンクスLINEで\n売上化と継続接点を高める", PALE, NAVY), ("実売上・LTV向上", NAVY, WHITE)]
ew = (CW - 2 * 1.4) / 3
for i, (t, f, tc) in enumerate(eq):
    x = L + i * (ew + 1.4)
    lines = t.split("\n")
    rect(s, x, 13.7, ew, 2.0, f, paras=[(ln, 13.5, True, tc, 0) for ln in lines])
    if i < 2:
        label(s, x + ew, 13.7, 1.4, 2.0, [("＋" if i == 0 else "＝", 24, True, NAVY, 0)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
note(s, "※100CV → 30友だち化（想定）の30人は、初回売上だけでなく継続フォローの対象になる")

# ---- 元のスライドを外す ----
new_ids = {x.slide_id for x in list(prs.slides)[-4:]}
for sl in list(prs.slides):
    if sl.slide_id not in new_ids:
        drop_slide(prs, sl)
prs.save(str(OUT))
print("saved:", OUT.name, len(prs.slides), "枚", round(prs.slide_width / 360000, 2), "x", round(prs.slide_height / 360000, 2))
