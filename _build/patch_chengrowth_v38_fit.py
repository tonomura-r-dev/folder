# -*- coding: utf-8 -*-
"""チェングロウス ver3.7 → ver3.8（2026-10-02 殿村さん指示）
・S11,12,15,18,19,21,22：枠の大きさに文字が合っていない → 文字をさらに大きくする
・S17：情報を詰め込みすぎ → 3枚に分ける（動線00／動線01・02／ステップ配信の中身）。費用の列は注記へ。
  python3 _build/patch_chengrowth_v38_fit.py <ver3.7.pptx>
"""
import copy
import sys
from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.util import Cm, Pt

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "20260929_チェングロウス_LINE公式アカウント運用のご提案ver3.8.pptx"
EMU = 360000
NAVY, INK, WHITE = RGBColor(0x1F, 0x28, 0x5A), RGBColor(0x33, 0x33, 0x33), RGBColor(0xFF, 0xFF, 0xFF)
LINEG = RGBColor(0x06, 0xC7, 0x55)
prs = Presentation(sys.argv[1])

CARD = copy.deepcopy(next(sh for sh in prs.slides[20].shapes if sh.name == "Rounded Rectangle 4")._element)

# ---------- A：文字を大きくする（ver3.7の番号）----------
G = {11: 1.3, 12: 1.3, 15: 1.25, 18: 1.2, 19: 1.2, 21: 1.3, 22: 1.3}


def scale(el, f):
    for tag in ("a:rPr", "a:endParaRPr", "a:defRPr"):
        for e in el.iter(qn(tag)):
            sz = e.get("sz")
            if sz:
                e.set("sz", str(int(round(int(sz) * f / 50.0)) * 50))


for n, f in G.items():
    for sh in prs.slides[n - 1].shapes:
        if sh.name in ("TextBox 1", "Connector 3", "TextBox 2"):
            continue
        if sh.top >= 16.5 * EMU and sh.height < 1.2 * EMU:      # 注記は据え置き
            continue
        small = sh.height < 1.4 * EMU and sh.width < 12 * EMU    # 幅の狭い見出し枠は折り返し防止
        narrow = sh.width < 4.5 * EMU                            # 動線ラベルなど幅の狭い枠
        banner = sh.width > 20 * EMU and sh.height < 1.4 * EMU   # 下の帯は1行に収める
        scale(sh._element, 1.0 if small else ((1.1 if n == 15 else 0.95) if narrow else (1.0 if banner else f)))


def adj(n, name, top=None, height=None, dy=None):
    sh = next(x for x in prs.slides[n - 1].shapes if x.name == name)
    if top is not None:
        sh.top = Cm(top)
    if dy is not None:
        sh.top = sh.top + Cm(dy)
    if height is not None:
        sh.height = Cm(height)


# 11枚目：カードと下の枠を大きくし、帯を下げる
for nm in ("Rounded Rectangle 5", "Rounded Rectangle 7", "Rounded Rectangle 9"):
    adj(11, nm, height=5.6)
adj(11, "TextBox 10", top=12.0)
adj(11, "Rounded Rectangle 11", top=12.8, height=3.2)
adj(11, "Rounded Rectangle 12", top=12.8, height=3.2)
adj(11, "Rounded Rectangle 13", top=16.3)
# 21枚目：カードを高くし、下のタイムラインを下げる
for nm in ("Rounded Rectangle 4", "Rounded Rectangle 5", "Rounded Rectangle 6"):
    adj(21, nm, height=6.4)
for nm in ("Rounded Rectangle 7", "Rounded Rectangle 8", "TextBox 9", "Rounded Rectangle 10", "Rounded Rectangle 11", "Rounded Rectangle 12", "TextBox 13"):
    adj(21, nm, dy=0.5)
adj(21, "Rounded Rectangle 14", top=16.3)
# 22枚目：カードを高くし、下の枠と帯を下げる
for nm in ("Rounded Rectangle 5", "Rounded Rectangle 7", "Rounded Rectangle 9"):
    adj(22, nm, height=6.4)
adj(22, "Rounded Rectangle 10", top=12.9, height=2.8)
adj(22, "Rounded Rectangle 11", top=16.0)


def merge_paras(n, names, start=1):
    """カード本文の手動改行をなくし、「。」で終わるまでを1段落にまとめる（折り返しの崩れ防止）"""
    for nm in names:
        sh = next(x for x in prs.slides[n - 1].shapes if x.name == nm)
        paras = sh.text_frame.paragraphs
        texts = [p.text for p in paras]
        buf, out = "", []
        for tx in texts[start:]:
            buf += tx
            if tx.endswith("。"):
                out.append(buf)
                buf = ""
        if buf:
            out.append(buf)
        body = paras[start:]
        for p in body[1:]:
            p._p.getparent().remove(p._p)
        first = body[0]
        runs = first.runs
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
        runs[0].text = out[0]
        prev = first._p
        for tx in out[1:]:
            newp = copy.deepcopy(first._p)
            prev.addnext(newp)
            for r in newp.findall(qn("a:r"))[1:]:
                newp.remove(r)
            newp.find(qn("a:r")).find(qn("a:t")).text = tx
            prev = newp


merge_paras(11, ("Rounded Rectangle 5", "Rounded Rectangle 7", "Rounded Rectangle 9"))
def set_body(n, name, lines):
    sh = next(x for x in prs.slides[n - 1].shapes if x.name == name)
    paras = sh.text_frame.paragraphs
    for p in paras[2:]:
        p._p.getparent().remove(p._p)
    paras = sh.text_frame.paragraphs
    for i, tx in enumerate(lines):
        if i + 1 >= len(paras):
            newp = copy.deepcopy(paras[-1]._p)
            paras[-1]._p.addnext(newp)
            paras = sh.text_frame.paragraphs
        pp = paras[i + 1]
        for r in pp.runs[1:]:
            r._r.getparent().remove(r._r)
        pp.runs[0].text = tx


set_body(21, "Rounded Rectangle 4", ["友だち・会員・クリックの履歴は始めた日から蓄積。", "半年で約540人の友だちに（弊社シミュレーション）"])
set_body(21, "Rounded Rectangle 5", ["会員登録のLINE化をリニューアルの要件に入れれば、", "サイトの二度手間を避けられる。"])
set_body(21, "Rounded Rectangle 6", ["整備士の検索は一年中あり、ピークは3月。", "秋までに友だちを増やしておけば、ピークの時期にご案内できる。"])
for nm, tx in (("Rounded Rectangle 5", "② 作り直しを避ける"), ("Rounded Rectangle 4", "① データを初日から蓄積"), ("Rounded Rectangle 6", "③ 3月のピークに間に合う")):
    sh = next(x for x in prs.slides[20].shapes if x.name == nm)
    r0 = sh.text_frame.paragraphs[0].runs
    r0[0].text = tx
    for r in r0[1:]:
        r._r.getparent().remove(r._r)
for nm in ("Rounded Rectangle 4", "Rounded Rectangle 5", "Rounded Rectangle 6"):   # 21枚目のカード見出しは14ptで1行に
    sh = next(x for x in prs.slides[20].shapes if x.name == nm)
    for r in sh.text_frame.paragraphs[0].runs:
        r.font.size = Pt(14)
# 11枚目：サイト側の箱を2行に
sh = next(x for x in prs.slides[10].shapes if x.name == "Rounded Rectangle 11")
sh.text_frame.paragraphs[1].runs[0].text = "条件別の検索結果ページ"
for r in sh.text_frame.paragraphs[1].runs[1:]:
    r._r.getparent().remove(r._r)
merge_paras(22, ("Rounded Rectangle 5", "Rounded Rectangle 7", "Rounded Rectangle 9"))


def set_par(sh, k, text):
    p = sh.text_frame.paragraphs[k]
    runs = p.runs
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def set_runs(sh, lines, align=None):
    tf = sh.text_frame
    base_p = tf.paragraphs[0]._p
    base_r = base_p.find(qn("a:r"))
    rpr0 = copy.deepcopy(base_r.find(qn("a:rPr"))) if base_r is not None else None
    for p in list(tf._txBody.findall(qn("a:p")))[1:]:
        tf._txBody.remove(p)
    for r in list(base_p.findall(qn("a:r"))) + list(base_p.findall(qn("a:br"))):
        base_p.remove(r)
    empty_p = copy.deepcopy(base_p)
    for i, (text, size, bold, color) in enumerate(lines):
        p = base_p if i == 0 else copy.deepcopy(empty_p)
        if i:
            tf._txBody.append(p)
        r = p.makeelement(qn("a:r"), {})
        if rpr0 is not None:
            r.append(copy.deepcopy(rpr0))
        t = r.makeelement(qn("a:t"), {})
        t.text = text
        r.append(t)
        end = p.find(qn("a:endParaRPr"))
        (end.addprevious(r) if end is not None else p.append(r))
    for para, (text, size, bold, color) in zip(tf.paragraphs, lines):
        if align is not None:
            para.alignment = align
        for run in para.runs:
            run.font.size, run.font.bold = Pt(size), bold
            run.font.color.rgb = color


# ---------- B：17枚目を3枚に分ける ----------
orig = prs.slides[16]
oby = {sh.name: sh for sh in orig.shapes}
layout = orig.slide_layout
T_TITLE, T_LEAD, T_LINE, T_NOTE = (copy.deepcopy(oby[k]._element) for k in ("TextBox 1", "TextBox 2", "Connector 3", "TextBox 8"))
TBL = copy.deepcopy(oby["Table 4"]._element)
COLW = [1.2, 3.0, 3.5, 4.1, 5.3, 3.3, 2.7]       # 費用の列をなくし、幅を広げる（計23.1cm）


def table_slide(title, keep, subtotal=None):
    new = prs.slides.add_slide(layout)
    for sh in list(new.shapes):
        sh._element.getparent().remove(sh._element)
    for el in (T_TITLE, T_LEAD, T_LINE):
        new.shapes._spTree.append(copy.deepcopy(el))
    set_par(new.shapes[0], 0, title)
    frame = copy.deepcopy(TBL)
    new.shapes._spTree.append(frame)
    shp = new.shapes[-1]
    tbl = shp.table
    t = tbl._tbl
    trs = t.findall(qn("a:tr"))
    total_tr = copy.deepcopy(trs[9])
    # 費用の列を消す
    t.find(qn("a:tblGrid")).remove(t.find(qn("a:tblGrid")).findall(qn("a:gridCol"))[7])
    for tr in trs:
        tr.remove(tr.findall(qn("a:tc"))[7])
    total_tr.remove(total_tr.findall(qn("a:tc"))[7])
    for i, tr in enumerate(trs):
        if i not in keep:
            t.remove(tr)
    if subtotal:
        t.append(total_tr)
        cells = total_tr.findall(qn("a:tc"))
        for idx, text in subtotal.items():
            tc = cells[idx]
            runs = list(tc.iter(qn("a:t")))
            runs[0].text = text
            for r in runs[1:]:
                r.text = ""
    for gc, w in zip(t.find(qn("a:tblGrid")).findall(qn("a:gridCol")), COLW):
        gc.set("w", str(int(w * EMU)))
    for k, tr in enumerate(t.findall(qn("a:tr"))):
        tr.set("h", str(int((1.3 if k == 0 else 1.7) * EMU)))
    for tag in ("a:rPr", "a:endParaRPr", "a:defRPr"):
        for e in t.iter(qn(tag)):
            e.set("sz", "1100")
    shp.left, shp.top, shp.width = Cm(2.2), Cm(4.9), Cm(23.1)
    new.shapes._spTree.append(copy.deepcopy(T_NOTE))
    return new


A = table_slide("②効率改善｜動線00「離脱防止」から", {0, 1, 2, 3}, {1: "小計", 5: "221回", 6: "5件"})
B = table_slide("②効率改善｜動線01「LINEで登録」・動線02「サンクスLINE」", {0, 4, 5, 6, 7, 8, 9})
lst = prs.slides._sldIdLst
ids = list(lst)[-2:]
for it in ids:
    lst.remove(it)
lst.insert(16, ids[0])
lst.insert(17, ids[1])

# 元の17枚目（いまは19番目）をステップ配信の中身のページにする
C = prs.slides[18]
cby = {sh.name: sh for sh in C.shapes}
for k in ("Table 4", "TextBox 7", "TextBox 8"):
    cby[k]._element.getparent().remove(cby[k]._element)
set_par(cby["TextBox 1"], 0, "②効率改善｜ステップ配信の中身（全10通の例）")
set_par(cby["TextBox 2"], 0, "0日に あいさつ（職種を聞く）、1日後からは ステップ配信")
steps = [("0日", "あいさつ\n職種を聞く"), ("1日", "職種の違い"), ("2日", "働き方"), ("3日", "求人の\nご案内①"), ("4日", "資格取得支援"),
         ("5日", "他の職種\nとの比較"), ("7日", "資格別の年収"), ("10日", "現場の様子"), ("12日", "地域別の\n新着求人"), ("14日", "求人の\nご案内②")]
BW, BH, GAP = 3.5, 4.4, 0.3
for i, (day, text) in enumerate(steps):
    r, c = divmod(i, 5)
    el = copy.deepcopy(CARD)
    C.shapes._spTree.append(el)
    sh = C.shapes[-1]
    sh.left, sh.top, sh.width, sh.height = Cm(2.2 + c * (BW + GAP)), Cm(5.2 + r * (BH + 0.7)), Cm(BW), Cm(BH)
    hot = "ご案内" in text
    if hot:
        sh.fill.solid()
        sh.fill.fore_color.rgb = LINEG
    lines = [(day, 20, True, WHITE if hot else NAVY)] + [(tx, 13, False, WHITE if hot else INK) for tx in text.split("\n")]
    set_runs(sh, lines, 2)
cby["Picture 5"].left, cby["Picture 5"].top = Cm(21.0), Cm(5.2)
cby["Picture 6"].left, cby["Picture 6"].top = Cm(21.0), Cm(8.3)
note = copy.deepcopy(T_NOTE)
C.shapes._spTree.append(note)
set_par(C.shapes[-1], 0, "※画面はイメージです")

for sl in prs.slides:
    for n, c in enumerate(sl.shapes._spTree.iter(qn("p:cNvPr")), start=2):
        c.set("id", str(n))
AFTER = {qn(x) for x in ("a:sym", "a:hlinkClick", "a:hlinkMouseOver", "a:rtl", "a:extLst")}
for sl in prs.slides:
    for r in sl.shapes._spTree.iter(qn("a:r")):
        if r.find(qn("a:rPr")) is None:
            r.insert(0, r.makeelement(qn("a:rPr"), {"lang": "ja-JP"}))
    for tag in ("a:rPr", "a:endParaRPr", "a:defRPr"):
        for rpr in sl.shapes._spTree.iter(qn(tag)):
            fs = []
            for ft in ("a:latin", "a:ea", "a:cs"):
                e = rpr.find(qn(ft))
                if e is None:
                    e = rpr.makeelement(qn(ft), {})
                else:
                    rpr.remove(e)
                e.set("typeface", "メイリオ")
                fs.append(e)
            nxt = next((c for c in rpr if c.tag in AFTER), None)
            for e in fs:
                (nxt.addprevious(e) if nxt is not None else rpr.append(e))
prs.save(str(OUT))
print("saved:", OUT.name, len(prs.slides), "枚")
