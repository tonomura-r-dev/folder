# -*- coding: utf-8 -*-
"""チェングロウス提案書ver1.1のP6（前後検索★差込枠）に実データを組み込む。

起点KW=カーディーラー（2026-09-23受領）。他ページは触らない（PC保存版が正）。
  python3 _build/patch_chengrowth_s6.py
"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Cm, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
OUT = str(ROOT / "20260922_株式会社チェングロウス御中_LINE公式アカウント運用のご提案ver1.1.pptx")
IMG = "/tmp/claude-0/-home-user-folder/dadd2f6b-abde-5965-beb8-93809cb827e0/scratchpad/cardealer.png"

TNAVY = "002060"; NAVY = "1F285A"; ORANGE = "ED7D31"; INK = "333333"
MUT = "7F7F7F"; WHITE = "FFFFFF"; PALE = "F4F7FF"; BORDER = "D9D9D9"
SW = 27.52
TITLE_XY = (1.52, 0.38, 24.4, 0.94)
LEAD_XY = (1.20, 1.80, 25.1, 1.90)
DIV_Y = 3.86
CX0, CW = 1.20, 25.12
CY0 = 4.30
FOOT_Y = 17.35

prs = Presentation(OUT)
s = prs.slides[5]
assert "前後検索" in s.shapes[0].text_frame.text or any(
    sh.has_text_frame and "前後検索" in sh.text_frame.text for sh in s.shapes)


def set_font(run, size, bold=None, color=INK, name="メイリオ"):
    f = run.font
    f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    f.name = name
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        e = rPr.find(qn(tag))
        if e is None:
            e = rPr.makeelement(qn(tag), {})
            rPr.append(e)
        e.set("typeface", name)
    f.color.rgb = RGBColor.from_string(color)


ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
ANCH = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}


def reset_tf(tf):
    for para in tf.paragraphs[1:]:
        para._p.getparent().remove(para._p)
    for r in list(tf.paragraphs[0].runs):
        r._r.getparent().remove(r._r)


def put_text(tf, paras, anchor="t", ml=0.14, mr=0.14, mt=0.06, mb=0.06):
    reset_tf(tf)
    tf.word_wrap = True
    tf.margin_left = Cm(ml); tf.margin_right = Cm(mr)
    tf.margin_top = Cm(mt); tf.margin_bottom = Cm(mb)
    tf.vertical_anchor = ANCH[anchor]
    first = True
    for p in paras:
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        para.alignment = ALIGN[p.get("align", "l")]
        if p.get("sa") is not None:
            para.space_after = Pt(p["sa"])
        if p.get("ls") is not None:
            para.line_spacing = p["ls"]
        for t, sz, b, c in p["runs"]:
            r = para.add_run(); r.text = t
            set_font(r, sz, b, c)


def T(slide, x, y, w, h, paras, anchor="t", **kw):
    b = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    put_text(b.text_frame, paras, anchor=anchor, **kw)


def one(text, sz, b=None, c=INK, align="l", sa=None, ls=None):
    d = {"runs": [(text, sz, b, c)], "align": align}
    if sa is not None: d["sa"] = sa
    if ls is not None: d["ls"] = ls
    return d


def box(slide, x, y, w, h, fill=None, line=None, lw=1.0):
    sp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Cm(x), Cm(y), Cm(w), Cm(h))
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = RGBColor.from_string(line); sp.line.width = Pt(lw)
    sp.shadow.inherit = False
    try: sp.adjustments[0] = 0.06
    except Exception: pass
    reset_tf(sp.text_frame)
    return sp


# --- クリアして再構築 ---
spTree = s.shapes._spTree
for el in list(spTree):
    if el.tag.split("}")[-1] in ("sp", "cxnSp", "pic", "graphicFrame", "grpSp"):
        spTree.remove(el)

T(s, *TITLE_XY, [one("市場分析（前後検索クエリ）", 16, True, TNAVY)], anchor="m", ml=0, mr=0)
T(s, *LEAD_XY,
  [one("LINEヤフーの前後検索データ（起点KW：カーディーラー）。", 12, None, INK, ls=1.28),
   one("職業への関心は、クルマを「買う人」の検索の中に埋もれている。", 12, None, INK, ls=1.28)],
  anchor="m", ml=0, mr=0)
ln = s.shapes.add_connector(1, Cm(0), Cm(DIV_Y), Cm(SW), Cm(DIV_Y))
ln.line.color.rgb = RGBColor.from_string(BORDER)
ln.line.width = Pt(1.0)

# 左：散布図（2000x1167 → w13.6, h7.94）
iw = 13.6
s.shapes.add_picture(IMG, Cm(CX0), Cm(CY0 + 0.3), width=Cm(iw))

# 右：所見3カード
rx = CX0 + iw + 0.5
rw = CW - iw - 0.5
cards = [
    ("① 職業クエリは0日周辺に実在する",
     "「カーディーラー 年収」「受付嬢」\n「ディーラー 営業マン」が起点前後に出現"),
    ("② ただし大半は「買う人」の検索",
     "車種名・店舗名・マイカーローンが支配的。\n働きたい人は、その中に埋もれている"),
    ("③ だから検索広告では狙い撃てない",
     "職業関心層だけを抽出する配信は困難。\nサイトに来た瞬間に受け止める＝\n離脱防止とLINEの役割"),
]
cy = CY0 + 0.3
for h, b in cards:
    bb = box(s, rx, cy, rw, 2.45, fill=PALE)
    put_text(bb.text_frame,
             [one(h, 11, True, NAVY, sa=3)] + [one(l, 9.5, None, INK, ls=1.3) for l in b.split("\n")],
             anchor="t", ml=0.25, mr=0.2, mt=0.15, mb=0.1)
    cy += 2.65

bb = box(s, CX0, CY0 + 8.6, CW, 1.1, fill=NAVY)
put_text(bb.text_frame,
         [one("検索で狙えない層は、来た瞬間に受け止める。それが本提案の入口設計。", 12.5, True, WHITE, align="c")],
         anchor="m", ml=0.3, mr=0.3)
T(s, CX0, FOOT_Y, CW, 0.9,
  [one("出典：LINEヤフー 前後検索データ（起点KW：カーディーラー／2026-09-23取得）。横軸=検索起点からの日数、縦軸=UU（相対）", 7.5, None, MUT, ls=1.15)],
  ml=0, mr=0)

prs.save(OUT)
print("saved:", OUT)
