"""
LINE広告トレンドレポート10月号（WEB形式・殿村さんのver2＝正）の本文カードの改行を直す（2026-10-28 殿村さん指摘）。
「。」「）」だけが次の行に落ちないよう、区切りのいい位置（、。の後・（の前など）で改行し、行の長さをそろえる。
使い方: python3 _build/patch_line_ads_202610_breaks.py <pptx>
"""
import sys
from pptx import Presentation
from pptx.util import Inches
import trend_report_engine as engine

F = sys.argv[1]
prs = Presentation(F)
for slide in prs.slides:
    for sh in slide.shapes:
        if sh.shape_type == 1 and sh.has_text_frame and Inches(3.05) <= sh.top < Inches(5.85):
            engine.apply_smart_breaks(sh)
prs.save(F)
print("saved", F)
