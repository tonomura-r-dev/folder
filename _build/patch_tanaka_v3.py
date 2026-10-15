"""タナカホーム v2 → v3：表紙タイトル（殿村さん確定）＋解決策に「資料請求→来場」を追加。
python -m _build.patch_tanaka_v3 <v2.pptx> <出力>"""
import sys, copy
from pptx import Presentation
from pptx.util import Inches
from lxml import etree

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
src, out = sys.argv[1], sys.argv[2]
prs = Presentation(src)
s1, s3, s4 = prs.slides[0], prs.slides[2], prs.slides[3]
by = lambda s, i: [x for x in s.shapes if x.shape_id == i][0]

def set_lines(sh, lines):
    ps = sh._element.findall(f".//{A}p")
    tmpl = ps[0]
    for p in ps[1:]:
        p.getparent().remove(p)
    body = tmpl.getparent()
    body.remove(tmpl)
    for ln in lines:
        p = copy.deepcopy(tmpl)
        rs = p.findall(f"{A}r")
        for r in rs[1:]:
            p.remove(r)
        rs[0].find(f"{A}t").text = ln
        body.append(p)

# 1 表紙：タイトル2行、同じ意味のサブタイトルは外す
t = by(s1, 26)
set_lines(t, ["集めた見込み客を逃さずナーチャリングをし、", "来場・契約へ変えるLINE仕組みの構築"])
bp = t._element.find(f".//{A}bodyPr")
bp.set("wrap", "square"); bp.set("anchor", "b")
for c in list(bp): bp.remove(c)
etree.SubElement(bp, f"{A}noAutofit")
t.left, t.top, t.width, t.height = Inches(0.6), Inches(2.5), Inches(9.63), Inches(1.3)
sub = by(s1, 27)
sub._element.getparent().remove(sub._element)

# 4 解決策：入口を資料請求に、下に帯
set_lines(by(s4, 8), ["資料請求・来場予約", "完了後にLINE追加"])
band = copy.deepcopy(by(s3, 18)._element)
band.find(".//{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr").set("id", "20")
rs = band.findall(f".//{A}r")
rs[0].find(f"{A}t").text = "資料請求の時点から、"
rs[1].find(f"{A}t").text = "来場・契約まで追える"
r3 = copy.deepcopy(rs[0]); r3.find(f"{A}t").text = "のがLINEの強み"
rs[1].addnext(r3)
note = by(s4, 15)._element
note.addprevious(band)

prs.save(out)
print("saved", out)
