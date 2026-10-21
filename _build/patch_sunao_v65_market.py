# -*- coding: utf-8 -*-
"""すなお ver6.4 → ver6.5：「ニーズ調査（業界STUDY）」を、市場分析の2枚に差し替える（2026-10-21 殿村さん確定の文面）。

    python -m _build.patch_sunao_v65_market <ver6.4.pptx> <ver6.5.pptx>

  6枚目（ニーズ調査 業界STUDY）を複製して2枚にし、表（項目／ファクト／意味すること）を4行に、示唆・注記を差し替える。
  形は元のページのまま（タイトル・リード・表・示唆・注記）。表の文字は13pt → 12pt（4行に収めるため）。
  数字の出典と確認の記録は `_drafts/すなおコンサルティング_商材・市場調査（2026-10-17）.md`。ほかのページは変えない（59枚 → 60枚）。
"""
import copy
import math
import sys
import unicodedata

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from _build.patch_tanaka_v5 import copy_slide

PAGES = [
    {
        "title": "研修・学びの市場：営業職の学び直しのニーズが伸びている",
        "lead": "企業の研修への投資も、個人がお金を払って学ぶ市場も伸びている",
        "rows": [
            ("企業向け研修サービス市場", "2025年度 6,075億円（前年度比3.7%増）。2026年度は6,300億円の予測",
             "企業の人材育成への投資は増えている"),
            ("eラーニング市場", "2025年度 3,924億円（前年度比2.7%増）。法人向け 1,294億円（4.7%増）・個人向け 2,630億円（1.7%増）。個人向けは資格関連の伸びで3年ぶりのプラス",
             "個人がお金を払って学ぶ市場も、再び伸び始めている"),
            ("オンラインサロン", "最大手のDMMオンラインサロンの流通総額は2024年度 約60億円。シェアは61.9%（2021年度は51.4%）",
             "有料のオンラインコミュニティで学ぶ形が定着している"),
            ("保険営業の方の接触回数", "コロナ前と比べ、接触回数が「減少した」営業職員が55.5%（2023年調査）",
             "限られた面談で成果を出すやり方を学ぶ必要がある"),
        ],
        "insight": "学びたい方は増えている。そのため、セミナーで出会った方とLINEでつながり、次のセミナー・サロンの入会まで案内を続けることが重要。",
        "note": "※企業向け研修サービス市場：矢野経済研究所（2026年8月24日）／eラーニング市場：矢野経済研究所（2026年4月22日。2025年度は見込み）／"
                "オンラインサロン：DMM.com発表（2025年11月14日。調査：デロイト トーマツ ミック経済研究所）／"
                "接触回数：矢野経済研究所「生命保険の営業職員アンケート調査」（2023年4〜6月、生命保険の営業職員400名）",
    },
    {
        "title": "お金への関心は高まる一方で、学ぶ機会は少ない",
        "lead": "資産形成を始める方は増えているのに、金融について学んでいない方が約9割",
        "rows": [
            ("NISA口座数", "2023年末 2,125万 → 2024年末 2,559万 → 2025年末 2,821万口座（政府目標：2027年末までに3,400万口座）",
             "資産形成を始める方は増え続けている"),
            ("NISAの累計買付額", "2023年末 35兆円 → 2024年末 53兆円 → 2025年末 71兆円。2027年末の政府目標56兆円をすでに超えた",
             "実際にお金が動いている"),
            ("金融経済教育", "「受けたと認識している」方は8.7%。一方で67.5%が「学校で行うべき」と回答",
             "学ぶ機会は少なく、学びたいニーズは大きい"),
            ("知識・計画", "金融リテラシーの正答率は53.8%（前回から1.9ポイント低下）。老後の資金計画をたてている方は40.1%、商品性を理解せずに投資信託を買った方は29.3%",
             "相談に乗れる方が求められている"),
        ],
        "insight": "お金への関心が高まる一方で、学ぶ機会は少ない。保険営業の方にとって、相談につながるチャンスが増えている。一般の方向けのマネーセミナーも、LINEで次回の案内を続けられる。",
        "note": "※NISA口座数・買付額：金融庁「NISAの利用状況」（2025年12月末時点）／"
                "金融経済教育・正答率・資金計画・投資信託：金融経済教育推進機構（J-FLEC）「金融リテラシー調査（2025年）」（2026年3月27日公表、30,000人）",
    },
]
COLS = [1.9, 4.7, 3.3]          # 項目／ファクト／意味すること（合計9.9インチ＝元と同じ）
CELL_PT = 12
TABLE_TOP = 1.8
NOTE_TOP = 6.45


def width_em(s):
    return sum(0.55 if unicodedata.east_asian_width(c) in ("Na", "H") else 1.0 for c in s)


def lines(text, width_in, pt, pad_in=0.22):
    cap = (width_in - pad_in) * 72 / pt * 0.95
    return max(1, math.ceil(width_em(text) / cap))


def by_id(s, sid):
    return next(sh for sh in s.shapes if sh.shape_id == sid)


def set_runs(p, text, size=None):
    runs = p.findall(qn("a:r"))
    for r in runs[1:]:
        p.remove(r)
    runs[0].find(qn("a:t")).text = text
    if size:
        runs[0].find(qn("a:rPr")).set("sz", str(size * 100))
        end = p.find(qn("a:endParaRPr"))
        if end is not None:
            end.set("sz", str(size * 100))


def set_shape(sh, texts):
    """段落ごとに文字を入れる（足りない段落は最後の段落を複製、余る段落は消す）"""
    ps = sh.text_frame._txBody.findall(qn("a:p"))
    while len(ps) < len(texts):
        ps[-1].addnext(copy.deepcopy(ps[-1]))
        ps = sh.text_frame._txBody.findall(qn("a:p"))
    for p in ps[len(texts):]:
        p.getparent().remove(p)
    for p, t in zip(sh.text_frame._txBody.findall(qn("a:p")), texts):
        set_runs(p, t)


def fill(s, page):
    set_shape(by_id(s, 45), [page["title"]])
    set_shape(by_id(s, 49), [page["lead"]])
    gf = next(sh for sh in s.shapes if sh.has_table)
    tbl = gf.table
    trs = tbl._tbl.findall(qn("a:tr"))
    while len(trs) < 1 + len(page["rows"]):
        trs[-1].addnext(copy.deepcopy(trs[-1]))
        trs = tbl._tbl.findall(qn("a:tr"))
    for i, w in enumerate(COLS):
        tbl.columns[i].width = Inches(w)
    total = tbl.rows[0].height
    for r, row in enumerate(page["rows"], start=1):
        n = max(lines(t, w, CELL_PT) for t, w in zip(row, COLS))
        h = Pt(n * CELL_PT * 1.3 + 10)
        tbl.rows[r].height = h
        total += h
        for c, text in enumerate(row):
            p = tbl.cell(r, c)._tc.find(qn("a:txBody")).find(qn("a:p"))
            set_runs(p, text, CELL_PT)
    gf.top, gf.height = Inches(TABLE_TOP), total

    line, box = by_id(s, 51), by_id(s, 52)
    set_shape(box, ["示唆", page["insight"]])
    top = gf.top + total + Inches(0.18)
    n = lines(page["insight"], 9.83, 12, pad_in=0.3)
    h = Pt(13 * 1.3 + n * 12 * 1.3 + 14)
    box.top, box.height = top, h
    box.left, box.width = gf.left, sum(Inches(w) for w in COLS)   # 線を外した分、表と左右をそろえる
    line._element.getparent().remove(line._element)       # 示唆の左の紺の線は外す（殿村さん）
    note = by_id(s, 53)
    set_shape(note, [page["note"]])
    assert top + h <= Inches(NOTE_TOP), (page["title"], (top + h) / 914400)


def main(src, out):
    prs = Presentation(src)
    slides = list(prs.slides)
    assert len(slides) == 59
    s6 = slides[5]
    assert by_id(s6, 45).text_frame.text == "ニーズ調査（業界STUDY）"
    s7 = copy_slide(prs, s6)
    lst = prs.slides._sldIdLst                               # 複製したページを6枚目の直後へ
    ids = list(lst)
    new = ids[-1]
    lst.remove(new)
    ids[5].addnext(new)
    fill(s6, PAGES[0])
    fill(s7, PAGES[1])
    assert len(prs.slides) == 60
    prs.save(out)
    print("saved:", out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
