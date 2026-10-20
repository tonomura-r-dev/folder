# -*- coding: utf-8 -*-
"""すなお ver4 → ver4.1（2026-10-20 殿村さん）。ver4（色直し済み）に当てる。

    python -m _build.patch_sunao_v41 <ver4.pptx> [出力.pptx]

  ・想定動線（8フェーズの表）：殿村さん「見づらい」
    2つに分かれていた表（見出しが2回出る）を1つに、文字を8〜9pt→12pt、本文は左寄せ・上下中央、
    ⑧の行の濃い灰色を外して行の色を交互に、文を短く（1行に収める）。
  ・想定の費用対効果（詳細）のトラック分岐（3枚）：殿村さん「枠が無駄が多い」
    左右の箱の高さを中身に合わせて詰める。上の帯が空なので、見出しの位置はほぼそのまま。
    薄くする重ね（片側を薄くする四角）も同じだけ動かす。
  ・成果報酬はセミナー予約だけ（佐村さん「成果報酬の方でカウントに入れないで」）。サロン入会は計測できないので
    成果報酬の地点から外す（資料内のCVとしては残す）。成果報酬型プランの「今回の成果単価の費用感」からサロン入会の行を外し、
    弊社想定シミュレーションの「※予約・入会で設定」を「※セミナー予約で設定」に。
  ほかのページは触らない（出力を省くと上書き）。
"""
import copy
import sys

from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

A16 = "{http://schemas.microsoft.com/office/drawing/2014/main}"
NAVY = "1F2A5C"
TEXT = "333333"
BAND = ("FFFFFF", "F5F7FA")

PHASES = ["① 接触", "② サイト離脱", "③ 育成（追加直後）", "④ リード獲得",
          "⑤ リード有効化", "⑥ 再育成", "⑦ マネタイズ", "⑧ ファン化・紹介"]
ROWS = [
    ("営業力を上げたい・情報収集中", "日程が合わず、申し込まない", "離脱防止ポップアップを表示"),
    ("今回は申し込まない", "次の案内を届ける手段が無い", "ポップアップからLINE友だち化"),
    ("どんなセミナーか分からない", "中身が見えず、予約に進まない", "見どころ・特典をステップ配信"),
    ("セミナーを予約した", "日が空き、参加を忘れる", "サンクスLINE誘導・前日の案内"),
    ("参加して関心が高まっている", "入会を迷ったまま時間が経つ", "参加直後にサロンを案内"),
    ("未入会／予約したが不参加", "次のイベントまで接点が無い", "月2回の配信で次のイベントへ"),
    ("サロンの入会を検討", "入会後の中身・成果が見えない", "入会した方の声を配信"),
    ("サロンに入会", "--", "次のイベントの案内・紹介"),
]
HEAD = ["フェーズ", "ユーザー状態・心理", "離脱・停滞の課題", "LINE OAの役割・施策"]


def slide_by_title(prs, title, start=0):
    out = []
    for i, s in enumerate(prs.slides):
        for sh in s.shapes:
            if sh.has_text_frame and sh.name.startswith("Google Shape;156") and sh.text_frame.text.strip() == title:
                out.append(i)
    return [prs.slides[i] for i in out if i >= start]


def find(s, sid):
    return next(sh for sh in s.shapes if sh.shape_id == sid)


def fix_ids(tbl):
    for tag, attr, base in (("a:gridCol", "colId", 20000), ("a:tr", "rowId", 10000)):
        for i, el in enumerate(tbl.findall(".//" + qn(tag))):
            idel = el.find(".//" + A16 + attr)
            if idel is not None:
                idel.set("val", str(base + i))


def cell_xml(text, fill, bold, color, size, algn):
    tc = etree.Element(qn("a:tc"))
    tx = etree.SubElement(tc, qn("a:txBody"))
    etree.SubElement(tx, qn("a:bodyPr"))
    etree.SubElement(tx, qn("a:lstStyle"))
    p = etree.SubElement(tx, qn("a:p"))
    etree.SubElement(p, qn("a:pPr"), algn=algn)
    r = etree.SubElement(p, qn("a:r"))
    rpr = etree.SubElement(r, qn("a:rPr"), lang="ja-JP", sz=str(size * 100), b="1" if bold else "0")
    sf = etree.SubElement(rpr, qn("a:solidFill"))
    etree.SubElement(sf, qn("a:srgbClr"), val=color)
    for t in ("a:latin", "a:ea", "a:cs"):
        etree.SubElement(rpr, qn(t), typeface="メイリオ")
    etree.SubElement(r, qn("a:t")).text = text
    pr = etree.SubElement(tc, qn("a:tcPr"), marL=str(Inches(0.1)), marR=str(Inches(0.08)),
                          marT=str(Inches(0.05)), marB=str(Inches(0.05)), anchor="ctr")
    f = etree.SubElement(pr, qn("a:solidFill"))
    etree.SubElement(f, qn("a:srgbClr"), val=fill)
    return tc


def p_doshen_table(s):
    """2つの表を1つにまとめて作り直す（見出しは1回だけ）"""
    t1, t2 = find(s, 18), find(s, 20)
    tbl = t1._element.find(".//" + qn("a:tbl"))
    trs = tbl.findall(qn("a:tr"))
    for tr in trs[1:]:
        tbl.remove(tr)
    tmpl = trs[0]
    widths = [1.95, 2.77, 2.77, 2.77]
    for gc, w in zip(tbl.find(qn("a:tblGrid")).findall(qn("a:gridCol")), widths):
        gc.set("w", str(Inches(w)))
    row_h = Inches(0.36)

    def make_row(cells, h):
        tr = copy.deepcopy(tmpl)
        tr.set("h", str(h))
        for tc in tr.findall(qn("a:tc")):
            tr.remove(tc)
        ext = tr.find(qn("a:extLst"))
        for tc in cells:
            (ext.addprevious(tc) if ext is not None else tr.append(tc))
        return tr

    head = make_row([cell_xml(t, NAVY, True, "FFFFFF", 12, "ctr") for t in HEAD], Inches(0.36))
    tbl.replace(tmpl, head)
    for i, (ph, vals) in enumerate(zip(PHASES, ROWS)):
        fill = BAND[i % 2]
        cells = [cell_xml(ph, fill, True, NAVY, 12, "l")] + [cell_xml(v, fill, False, TEXT, 12, "l") for v in vals]
        tbl.append(make_row(cells, row_h))
    fix_ids(tbl)
    t1.left, t1.top = Inches(0.42), Inches(3.5)
    t1.width = Inches(sum(widths))
    t1.height = Inches(0.36) + row_h * 8
    t2._element.getparent().remove(t2._element)


def p_track(s):
    """左右の箱を中身の高さに詰める。上の帯（0.7〜1.5）は空なので、帯の上からの空きの上1/3に見出しを置く"""
    BH, CH, GAP = 2.65, 0.89, 0.2
    group = 0.42 + 0.42 + BH + GAP + CH                       # 見出し〜結論の高さ
    top = 0.7 + (7.05 - 0.7 - group) / 3                      # 空きの上1/3・下2/3
    old = [1.22, 1.64, 2.06, 5.95, 6.11, 7.00]
    new = [top, top + 0.42, top + 0.84, top + 0.84 + BH, top + 0.84 + BH + GAP, top + 0.84 + BH + GAP + CH]

    def f(y):
        if y <= old[0]:
            return y - old[0] + new[0]
        for (a, b), (c, d) in zip(zip(old, old[1:]), zip(new, new[1:])):
            if y <= b:
                return c + (y - a) * (d - c) / (b - a)
        return y - old[-1] + new[-1]

    for sh in s.shapes:
        if sh.shape_id in (3, 4, 45):                         # 上の帯・タイトルはそのまま
            continue
        y0 = Emu(sh.top).inches
        y1 = y0 + Emu(sh.height).inches
        sh.top = Inches(f(y0))
        sh.height = Inches(f(y1) - f(y0))


def set_text(sh, old, new):
    """図形・セルの中の文字を置き換える（最初の文字の書式のまま）"""
    el = sh._tc if hasattr(sh, "_tc") else sh._element
    for p in el.iter(qn("a:p")):
        rs = p.findall(qn("a:r"))
        text = "".join(r.findtext(qn("a:t")) or "" for r in rs)
        if old in text:
            rs[0].find(qn("a:t")).text = text.replace(old, new)
            for r in rs[1:]:
                p.remove(r)
            return True
    return False


def p_seika(prs):
    """成果報酬はセミナー予約だけ"""
    plan = next(s for s in prs.slides if any(sh.has_text_frame and sh.text_frame.text.strip() == "成果報酬型プラン"
                                              for sh in s.shapes))
    gf = next(sh for sh in plan.shapes if sh.has_table and sh.table.cell(0, 0).text == "成果の地点")
    tbl = gf._element.find(".//" + qn("a:tbl"))
    for tr in tbl.findall(qn("a:tr")):
        if "".join(tr.itertext()).startswith("サロン入会"):
            gf.height = gf.height - int(tr.get("h"))
            tbl.remove(tr)
    fix_ids(tbl)
    sim = next(s for s in prs.slides if any(sh.has_table and "※予約" in sh.table.cell(0, 0).text
                                             for sh in s.shapes if sh.has_table))
    cell = next(sh for sh in sim.shapes if sh.has_table and "※予約" in sh.table.cell(0, 0).text).table.cell(0, 0)
    import re
    t = cell.text
    m = re.search(r"※予約 ([\d,]+)円・入会 [\d,]+円で設定", t)
    assert set_text(cell, m.group(0), f"※セミナー予約 {m.group(1)}円／件で設定"), t


def main(src, out):
    prs = Presentation(src)
    doshen = slide_by_title(prs, "想定動線")
    p_doshen_table(doshen[1])
    p_seika(prs)
    tracks = [s for s in slide_by_title(prs, "想定の費用対効果（詳細）")
              if any(sh.has_text_frame and sh.text_frame.text.strip() == "トラック分岐の設計思想" for sh in s.shapes)]
    assert len(tracks) == 3, len(tracks)
    for s in tracks:
        p_track(s)
    prs.save(out)
    print("saved:", out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else sys.argv[1])
