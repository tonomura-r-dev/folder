# -*- coding: utf-8 -*-
"""ジム・フィットネス 新FMT版：具体施策のイメージ（P28〜P35）を「左＝デモの画面／右＝ほかの案3つ」にする（2026-10-30 殿村さん。塾と同じ形）。

    python -m _build.patch_fitness_demo_gutai <入力.pptx> <出力.pptx>

  形はワクトリv6（_build/patch_wakutori_v6_demo.py）と同じ。
  ・左：施策名＋1行の見出し、ジムのデモ（`_data/fitness/ジム_LINE施策デモ_20261030.html`）の ①ユーザーが見る画面＋②裏側の処理
    （`_images/fitness/demo_p28〜35.png`）、下に小さく画面の注記
  ・右：4つの枠（目的／対象／中身／見る数字）を、3つの枠（入口／LINEの仕組み／配信内容）にする
  上の8フェーズの矢印・札・タイトルは触らない。
"""
import copy
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Emu, Inches, Pt

from _build.pptx_parts import style_run

NAVY, INK, GRAY = "1F2A5C", "333333", "7F7F7F"
IMG = "_images/fitness/demo_p{}.png"
LABELS = ["入口", "LINEの仕組み", "配信内容"]
PAGES = {  # 札, 施策名, 見出しの1行, 右の3つ（各2行）
    28: ("LINEの友だちへ", "離脱防止ポップアップ", "サイトを離れる方を、30秒診断でLINEへ",
         [["料金ページのLINEボタンや、", "店頭のQRにも置く"],
          ["入口ごとにURLを分けて記録し、", "あいさつの文面を出し分ける"],
          ["追加の特典に、無料体験の", "ご案内と期限を届ける"]]),
    29: ("露出の強化", "季節の企画配信", "需要の波に合わせ、冬に始める方を後押し",
         [["閑散期に追加した方にも、", "1月の前にまとめて届ける"],
          ["会員とそれ以外で、", "届ける企画を分ける"],
          ["期限つきの無料体験と、", "会員の声の動画"]]),
    30: ("CVの最大化", "30秒 体づくりタイプ診断", "4問の回答で、合う始め方をお伝え",
         [["リッチメニューや配信から、", "いつでも診断をやり直せる"],
          ["回答（目的・経験・頻度）を", "タグにして、次の配信を変える"],
          ["性格タイプ・ダイエットタイプ・", "続けられるかの診断を足す"]]),
    31: ("CPOの改善（CV質）", "体験の前日のご確認", "持ち物・当日の流れを届け、来店につなげる",
         [["予約はLINEの中でも、", "サイトのフォームでも受け付ける"],
          ["日程の変更・キャンセルを、", "トークで受け付ける"],
          ["「見学だけOK・入会を", "迫りません」を添える"]]),
    32: ("休眠顧客の有効化", "体験後の再案内", "入会を迷う方に、比べて決めてよいと伝える",
         [["体験だけの方や休会した方にも、", "通知メッセージで再案内する"],
          ["来店した・していないで", "タグを分けて配信する"],
          ["1月・新生活の前に、", "期限つきの特典を届ける"]]),
    33: ("マネタイズ最大化", "タイプ別の出し分け", "診断のタイプに合う、会員の変化を届ける",
         [["体験を予約した方にも、", "同じタイプの例を届ける"],
          ["診断のタイプと予約の有無で、", "配信を出し分ける"],
          ["「週2回なら1回あたり」の", "料金と、よくある質問"]]),
    34: ("リピート施策", "入会後14日の配信", "入会直後のつまずきを、先回りして防ぐ",
         [["入会の登録で、リッチメニューを", "会員向けに切り替える"],
          ["来店が2週間途切れた方を、", "タグで見つけてフォローする"],
          ["フォームチェック・", "体組成測定の予約の案内"]]),
    35: ("コミュニティ運営", "お友だちへの紹介", "紹介のメッセージを、転送するだけで紹介できる",
         [["リッチメニューの", "「お友だち紹介」から送れる"],
          ["紹介から追加した方にタグを付け、", "紹介元が分かるようにする"],
          ["会員の声や、", "店舗のイベントのお知らせ"]]),
}
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr"


def inch(v):
    return Emu(v).inches


def add_text(s, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for k, (text, size, bold, color, after) in enumerate(paras):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        if after:
            p.space_after = Pt(after)
        r = p.add_run()
        r.text = text
        style_run(r, size, bold, color)
    return tb


def main(src, out):
    prs = Presentation(src)
    for n, (tag, name, sub, ideas) in PAGES.items():
        s = prs.slides[n - 1]
        assert any(sh.has_text_frame and sh.text_frame.text == tag for sh in s.shapes), (n, tag)
        big = next(sh for sh in s.shapes if abs(inch(sh.left) - 0.29) < 0.05 and abs(inch(sh.top) - 2.03) < 0.05
                   and inch(sh.width) > 6)
        right_boxes = sorted([sh for sh in s.shapes if inch(sh.left) > 6.8 and inch(sh.top) > 1.9
                              and not (sh.has_text_frame and sh.text_frame.text.strip())], key=lambda x: inch(x.top))
        right_texts = [sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip() and inch(sh.left) > 6.8
                       and inch(sh.top) > 1.9]
        flow = [sh for sh in s.shapes if sh is not big and inch(sh.left) < 6.8 and inch(sh.top) > 2.5]
        assert len(right_boxes) == 4 and len(right_texts) == 4 and len(flow) == 11, (n, len(right_boxes), len(right_texts), len(flow))
        tmpl = right_boxes[0]
        tmpl.left = Inches(6.85)
        for sh in flow + right_texts + right_boxes[1:]:
            sh._element.getparent().remove(sh._element)

        # 左：施策名＋見出し、デモの画面（①＋②）、注記
        add_text(s, 0.45, 2.15, 6.1, 0.75, [(name, 16, True, NAVY, 0), (sub, 12, False, INK, 0)])
        h = 3.62
        w = h * 1756 / 1478
        pic = s.shapes.add_picture(IMG.format(n), Inches(0.29 + (6.44 - w) / 2), Inches(2.9), Inches(w), Inches(h))
        pic.line.color.rgb = RGBColor.from_string("D9D9D9")
        pic.line.width = Pt(0.75)
        add_text(s, 0.45, 6.53, 6.1, 0.24,
                 [("画面：ジム向けに作成したデモ（①ユーザーが見る画面／②裏側の処理）。店舗名・文面はデモ用のサンプル", 8.5, False, GRAY, 0)])

        # 右：3つの枠（入口／LINEの仕組み／配信内容）
        top, bottom, gap = 2.03, 6.80, 0.1
        bh = (bottom - top - gap * 2) / 3
        tmpl.top, tmpl.height = Inches(top), Inches(bh)
        boxes = [tmpl]
        for k in (1, 2):
            s.shapes._spTree.append(copy.deepcopy(tmpl._element))
            sh = s.shapes[-1]
            sh.top = Inches(top + k * (bh + gap))
            boxes.append(sh)
        for box, label, lines in zip(boxes, LABELS, ideas):
            add_text(s, inch(box.left) + 0.12, inch(box.top) + 0.1, inch(box.width) - 0.24, bh - 0.2,
                     [(label, 13.5, True, NAVY, 4)] + [(x, 12, False, INK, 0) for x in lines], anchor=MSO_ANCHOR.MIDDLE)

        used = set()
        nid = max(int(e.get("id")) for e in s.shapes._spTree.iter(P)) + 1
        for e in s.shapes._spTree.iter(P):
            if e.get("id") in used:
                e.set("id", str(nid))
                nid += 1
            used.add(e.get("id"))
    prs.save(out)
    print("saved:", out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
