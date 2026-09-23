# -*- coding: utf-8 -*-
"""学校法人ver2.1のS11（現状分析｜各校のリッチメニュー・あいさつ）に実機スクショを組み込む。

殿村さん実機撮影（2026-09-23・8校）。画像は _data/school_screenshots/ に全11枚保存済み。
スライドには代表5校を掲載（立命館・近畿・龍谷・日本工学院・HAL東京）。
大原（個別トーク型）・NSG（あいさつ無し）・大阪モード学園は注記で言及。
※ このスクリプトは S11 だけを触る。他ページはPC保存版が正。
  python3 _build/patch_gakko_s11.py
"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Cm, Pt
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
OUT = str(ROOT / "学校法人業界_LINEOA施策提案ver2.1.pptx")
IMGDIR = ROOT / "_data/school_screenshots"

NAVY = "1F285A"; DARK = "1C2233"; MUT = "646C82"; BORDER = "D9D9D9"

prs = Presentation(OUT)
s = prs.slides[10]
assert "各校のリッチメニュー" in s.shapes[0].text_frame.text

shapes = list(s.shapes)

COLS = [  # (見出し, 画像, 下ラベル1行目, 2行目)
    ("立命館大学 入学センター", "立命館大学入学センター_あいさつ_リッチメニュー.png",
     "王道の3導線型", "情報サイト・入試イベント・資料請求"),
    ("近畿大学", "近畿大学_自動応答_リッチメニュー.png",
     "個別問合せは受けない", "自動応答でメール・電話へ誘導"),
    ("龍谷大学 入試部", "龍谷大学入試部_あいさつ2_リッチメニュー.png",
     "友だち限定の特典型", "入試対策講座をLINE登録者だけに配布"),
    ("日本工学院", "日本工学院_リッチメニュー.png",
     "OC日程の直載せ型", "次回日程＋チャット質問窓口"),
    ("HAL東京", "HAL東京_リッチメニュー.png",
     "チャット相談型", "入学担当が質問に回答"),
]
HEAD_IDX = [2, 5, 8, 11, 14]
BOX_IDX = [3, 6, 9, 12, 15]
LABEL_IDX = [4, 7, 10, 13, 16]


def set_run(r, text, size=None, bold=None, color=None):
    r.text = text
    if size is not None:
        r.font.size = Pt(size)
    if bold is not None:
        r.font.bold = bold
    if color is not None:
        r.font.color.rgb = RGBColor.from_string(color)
    rPr = r._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        e = rPr.find(qn(tag))
        if e is None:
            e = rPr.makeelement(qn(tag), {})
            rPr.append(e)
        e.set("typeface", "メイリオ")


def set_text(shape, lines):
    """lines = [(text, size, bold, color), ...] 先頭段落の書式を流用しつつ全置換"""
    tf = shape.text_frame
    for para in tf.paragraphs[1:]:
        para._p.getparent().remove(para._p)
    p0 = tf.paragraphs[0]
    for r in list(p0.runs)[1:]:
        r._r.getparent().remove(r._r)
    if not p0.runs:
        p0.add_run()
    set_run(p0.runs[0], *lines[0])
    for ln in lines[1:]:
        p = tf.add_paragraph()
        p.alignment = p0.alignment
        set_run(p.add_run(), *ln)


# 1) 見出し5枠
for idx, (head, *_rest) in zip(HEAD_IDX, COLS):
    set_text(shapes[idx], [(head, 10, True, NAVY)])

# 2) 差込枠を消してスクショを配置
IMG_H = 8.4
IMG_W = IMG_H * 1170 / 2532  # 実機スクショ比率
for idx, (_h, img, *_l) in zip(BOX_IDX, COLS):
    sp = shapes[idx]
    col_x = sp.left / 360000
    col_w = sp.width / 360000
    top = sp.top / 360000
    sp._element.getparent().remove(sp._element)
    pic = s.shapes.add_picture(str(IMGDIR / img),
                               Cm(col_x + (col_w - IMG_W) / 2), Cm(top + 0.1),
                               height=Cm(IMG_H))
    pic.line.color.rgb = RGBColor.from_string(BORDER)
    pic.line.width = Pt(0.75)

# 3) 下ラベル（所見）
for idx, (_h, _i, l1, l2) in zip(LABEL_IDX, COLS):
    set_text(shapes[idx], [(l1, 9.5, True, DARK), (l2, 8, False, MUT)])

# 4) 帯のメッセージ
set_text(shapes[19], [("8校とも導線は「OC・資料請求」まで。保護者に向けた発信はゼロ＝ここが空白地帯。",
                       14, True, DARK)])

# 5) 注記
set_text(shapes[20], [
    ("※ 2026-09-23 実機取得（友だち追加直後のあいさつメッセージ＋リッチメニュー）。友だち数は page.line.me 実測（2026-09-20）。",
     9, False, MUT),
    ("※ 他3校：大原学園＝先生と個別トーク型／NSGカレッジリーグ＝登録直後の配信なし／大阪モード学園＝OC・資料請求・Instagram導線。",
     9, False, MUT),
])

prs.save(OUT)
print("saved:", OUT)
