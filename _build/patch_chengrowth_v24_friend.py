# -*- coding: utf-8 -*-
"""ver2.3 → ver2.4：「登録と同時に友だち追加も完了」を正確な表現に直す
（LINEでログインの友だち追加は、利用者がチェックを外すと追加されないため）
  python3 _build/patch_chengrowth_v24_friend.py
"""
from pptx import Presentation
SRC = "20260929_株式会社チェングロウス御中_LINE公式アカウント運用のご提案ver2.3.pptx"
OUT = "20260929_株式会社チェングロウス御中_LINE公式アカウント運用のご提案ver2.4.pptx"
prs = Presentation(SRC)
S = list(prs.slides)


def runs(sh):
    if sh.shape_type == 6:
        for c in sh.shapes:
            yield from runs(c)
    elif getattr(sh, "has_table", False) and sh.has_table:
        for r in sh.table.rows:
            for c in r.cells:
                for p in c.text_frame.paragraphs:
                    yield from p.runs
    elif sh.has_text_frame:
        for p in sh.text_frame.paragraphs:
            yield from p.runs


done = []
for no, a, b in [(9, "・登録と同時にLINEの友だちになる", "・登録と同じ画面で、LINEの友だち追加もあわせてご案内"),
                 (15, "登録と同時に友だち追加も完了。", "同じ画面で友だち追加もあわせてご案内。")]:
    for sh in S[no - 1].shapes:
        for r in runs(sh):
            if a in r.text:
                r.text = r.text.replace(a, b); done.append(no)
# 10枚目の表のセルは run が「登録と友だち追加が」「改行」「同時に完了」に分かれている
for sh in S[9].shapes:
    if getattr(sh, "has_table", False) and sh.has_table:
        for row in sh.table.rows:
            for c in row.cells:
                if "登録と友だち追加が" in c.text and "同時に完了" in c.text:
                    p = c.text_frame.paragraphs[0]
                    rs = p.runs
                    rs[0].text = "登録と同じ画面で、"
                    rs[-1].text = "友だち追加もあわせてご案内"
                    done.append(10)
prs.save(OUT)
print("修正したスライド:", done)
