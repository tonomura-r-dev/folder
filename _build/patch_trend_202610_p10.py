"""
10月号 P.10（レストランオプション 機材の価格改定）を直す（2026-10-26 殿村さん指示）。
- カード4枚 → 「価格（改定前 → 改定後）」「適用時期」の2枚（P.3の変更点カードを複製）
- 分類タグ＝その他、逆三角形は色なし
使い方: python3 _build/patch_trend_202610_p10.py <pptx>
"""
import sys
import copy
from pptx import Presentation
from pptx.util import Inches, Emu
from pptx.oxml.ns import qn
import trend_report_engine as engine

F = sys.argv[1]
prs = Presentation(F)
slides = list(prs.slides)
s3, s10 = slides[2], slides[9]


def text(sh):
    return sh.text_frame.text.strip() if sh.has_text_frame else ""


assert any(text(x).startswith("【レストランオプション】") for x in s10.shapes)
head = next(x for x in s10.shapes if text(x) == "■ 詳細")
for x in [x for x in s10.shapes if x.top >= head.top + head.height - Inches(0.02) and x.top < Inches(7.0)]:
    x._element.getparent().remove(x._element)

CARDS = [
    ("価格（改定前 → 改定後）", ["・プリンター：37,000円 → 41,700円", "・キャッシュドロアー：16,000円 → 17,700円"]),
    ("適用時期", ["・2027年2月上旬の予定", "・改定日以降に購入した機材から、改定後の価格になる"]),
]
src = next(x for x in s3.shapes if x.top > Inches(3.0) and text(x).startswith("変更点"))
cards = []
for k, (h, body) in enumerate(CARDS):
    el = copy.deepcopy(src._element)
    s10.shapes._spTree.append(el)
    cn = el.find(".//" + qn("p:cNvPr"))
    cn.set("id", str(max(int(e.get("id")) for e in s10._element.iter(qn("p:cNvPr"))) + 1))
    cn.set("name", cn.get("name") + f"_p10_{k}")
    c = next(x for x in s10.shapes if x._element is el)
    p0, p1, p2 = c.text_frame.paragraphs
    p0.runs[0].text = h
    for p, t in ((p1, body[0]), (p2, body[1])):
        p.runs[0].text = t
        for r in list(p.runs[1:]):
            r._r.getparent().remove(r._r)
        for br in p._p.findall(qn("a:br")):
            br.getparent().remove(br)
    cards.append(c)


def text_w(s, pt):
    return sum((0.55 if ord(ch) < 0x2000 else 1.0) for ch in s) * pt / 72


tf = cards[0].text_frame
w = max(max(text_w(t, 14) for _, b in CARDS for t in b), text_w(CARDS[0][0], 20))
w = min(w + Emu(tf.margin_left + tf.margin_right).inches + 0.45, 9.64)
h = 0.32 + 20 * 1.3 / 72 + 2 * (14 * 1.45 + 6) / 72
gap = 0.2
area_top = Emu(head.top + head.height).inches + 0.08
top = area_top + (7.05 - area_top - (2 * h + gap)) / 2
left = (Emu(prs.slide_width).inches - w) / 2
for i, c in enumerate(cards):
    c.left, c.top, c.width, c.height = Inches(left), Inches(top + i * (h + gap)), Inches(w), Inches(h)

engine.set_tag_highlight(s10, "その他")
engine.set_funnel_highlight(s10, [])
engine.force_font(prs)
prs.save(F)
print("saved", F, round(w, 2))
