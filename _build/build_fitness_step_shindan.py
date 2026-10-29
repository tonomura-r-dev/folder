# -*- coding: utf-8 -*-
"""ジム・フィットネス業界：ステップ配信（追加から14日）に診断を足して厚くした版を、1枚の試作で作る（2026-10-29 殿村さん）。

    python -m _build.build_fitness_step_shindan <ジム・フィットネス業界_LINEOA施策提案ver3.pptx> <出力.pptx> [--full]
    （--full を付けると、元の資料の該当3ページを直した全体版を出す。付けないと3枚の試作）

  ver3の25ページ「ステップ配信｜追加から14日で体験予約を取り切る」だけを残し、表に診断の配信を3行足す
  （Day3 性格タイプ診断／Day5 ダイエットタイプ診断／Day10 続けられるか診断）。Day8 の会員の声は診断のタイプ別に変える。
  「MBTI」の名前は使わない（登録商標）。足した行は薄い紺で塗る。表の行を写したら a16:rowId を振り直す（CLAUDE.md 2026-09-24）。
  実際の資料に入れるかは、この試作を見て殿村さんが決める。
  2026-10-29 追加：リッチメニューの27・28ページも入れる（タブ②＝30秒 体づくり診断／性格タイプ診断／ダイエットタイプ診断／
  続けられるか診断／1分 料金診断／よくある質問・相談。初心者メニューを外し、よくある質問とチャット相談を1マスに）。
  画像は殿村さんが生成：_images/fitness/richmenu_sekkei_v2.png（27）／richmenu_tab2_v2.png（28のタブ②）。
"""
import io
from pathlib import Path
import copy
import sys

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from _build.pptx_parts import drop_slide

NEW_FILL = "DCE3F2"
ROW_FILL = ["FFFFFF", "EFF1F7"]
ROWS = [  # (元の行（写す書式）, 追加か, Day, 目的, 配信コンテンツ, CTA, 通知プレビュー冒頭)。\n は改行
    (1, False, "Day0", "体験の案内\n＋属性取得", "あいさつ＋無料体験（期限つき）のバナー＋\n30秒 体づくりタイプ診断", "体験予約／\n診断回答", "【〇月〇日まで】無料体験のご案内"),
    (2, False, "Day2", "初心者の不安", "タイプ別「最初の1ヶ月はこれだけ」\n＋マシンの使い方動画", "動画視聴", "【初めての方へ】まずはマシン3種類"),
    (2, True, "Day3", "続け方を見つける", "性格タイプ診断（3問・4タイプ）→\n計画派・気分派・仲間派・記録派ごとの通い方", "性格タイプ診断", "【3問】あなたに合うジムの通い方は？"),
    (3, False, "Day4", "費用の納得", "1分料金シミュレーター（週2回なら1回いくら？）", "料金診断", "【1回いくら？】週2回なら…"),
    (2, True, "Day5", "目的別の深掘り", "ダイエットタイプ診断（4問：食事・間食・睡眠・運動）\n→ タイプ別「ジムでやること＋家での1つ」", "ダイエット\nタイプ診断", "【4問】あなたのダイエットタイプは？"),
    (4, False, "Day6", "2回目のオファー", "無料体験のご案内＋\n「見学だけでもOK・入会を迫りません」", "体験予約 CV①", "【見学だけOK】今週の空き枠"),
    (5, False, "Day8", "社会的証明", "同じタイプの会員の声（診断のタイプ別）\n／続いている人の共通点", "事例閲覧", "【コツコツ型の方へ】3ヶ月後の変化"),
    (2, True, "Day10", "不安の見える化", "続けられるか診断（3問）→ 結果に合わせた\n体験の回り方と、最初の2週間の予定", "続けられるか診断", "【3問】あなたは続けられる？"),
    (6, False, "Day11", "疑問の除去", "不安Q&A（勧誘・解約・シャワー・混雑・持ち物）\n＋チャット相談", "よくある\n質問・相談", "【解約もOK】よくあるご質問"),
    (7, False, "Day14", "最終オファー", "入会金無料・事務手数料0円＋体験枠の残数", "体験予約（再）", "【残3枠】今月の体験ご予約"),
]
LEAD = ["友だち追加から15日でアクティブ率が急落する。その14日間に、答えるだけの診断を4回挟み、",
        "反応を切らさずに体験予約へつなげる。"]
CARDS = [
    ("① 診断は4回、どれも1タップ", ["Day0・3・5・10に答えるだけの配信"]),
    ("② 答えはそのまま出し分けに", ["診断のタイプを配信のタグにする"]),
    ("③ 診断の結果から体験予約へ", ["結果の最後に「無料体験を予約」"]),
]
NOTE = ["※ 薄い紺の行が、今回足した診断の配信。Day5は目的に「ダイエット」を選んだ方だけ（他の目的の方は送らない）。",
        "※ 性格タイプ診断・ダイエットタイプ診断は、リッチメニュー「② 診断・シミュレーター」からも押せる（次の2ページ）。"]
NOTE_FULL = ["※ Day5は、目的に「ダイエット」を選んだ方だけに送る。",
             "※ 性格タイプ診断・ダイエットタイプ診断は、リッチメニュー「② 診断・シミュレーター」からも押せる。"]
ROOT = Path(__file__).resolve().parent.parent
IMG27 = ROOT / "_images/fitness/richmenu_sekkei_v2.png"
IMG28 = ROOT / "_images/fitness/richmenu_tab2_v2.png"


def swap_image(slide, sid, path):
    """画像を差し替え、元の枠の中に縦横比を保って中央に置く"""
    from PIL import Image
    pic = next(sh for sh in slide.shapes if sh.shape_id == sid)
    blip = pic._element.find(".//" + qn("a:blip"))
    old = blip.get(qn("r:embed"))
    _, rid = slide.part.get_or_add_image_part(io.BytesIO(path.read_bytes()))
    blip.set(qn("r:embed"), rid)
    slide.part.drop_rel(old)
    w, h = Image.open(path).size
    bx, by_, bw, bh = pic.left, pic.top, pic.width, pic.height
    scale = min(bw / w, bh / h)
    nw, nh = int(w * scale), int(h * scale)
    pic.left, pic.top, pic.width, pic.height = bx + (bw - nw) // 2, by_ + (bh - nh) // 2, nw, nh
    sr = pic._element.find(".//" + qn("a:srcRect"))
    if sr is not None:
        sr.getparent().remove(sr)


def set_cell_text(tc, text, sz="950"):
    p = tc.find(qn("a:txBody")).find(qn("a:p"))
    rs = p.findall(qn("a:r"))
    for r in rs[1:]:
        p.remove(r)
    r0 = rs[0]
    r0.find(qn("a:rPr")).set("sz", sz)
    for k, part in enumerate(text.split("\n")):
        if k == 0:
            r0.find(qn("a:t")).text = part
            continue
        br = p.makeelement(qn("a:br"), {})
        br.append(copy.deepcopy(r0.find(qn("a:rPr"))))
        r = copy.deepcopy(r0)
        r.find(qn("a:t")).text = part
        p.append(br)
        p.append(r)


def set_fill(tc, hexv):
    pr = tc.find(qn("a:tcPr"))
    sf = pr.find(qn("a:solidFill"))
    sf.find(qn("a:srgbClr")).set("val", hexv)


def set_paras(sh, lines):
    txb = sh._element.find(qn("p:txBody"))
    ps = txb.findall(qn("a:p"))
    tmpl = copy.deepcopy(ps[0])
    for r in tmpl.findall(qn("a:r"))[1:] + tmpl.findall(qn("a:br")):
        tmpl.remove(r)
    for p in ps:
        txb.remove(p)
    for line in lines:
        p = copy.deepcopy(tmpl)
        p.find(qn("a:r")).find(qn("a:t")).text = line
        txb.append(p)


def set_card(sh, head, body):
    txb = sh._element.find(qn("p:txBody"))
    ps = txb.findall(qn("a:p"))
    h, b = copy.deepcopy(ps[0]), copy.deepcopy(ps[1])
    for p in (h, b):
        for r in p.findall(qn("a:r"))[1:] + p.findall(qn("a:br")):
            p.remove(r)
    for p in ps:
        txb.remove(p)
    h.find(qn("a:r")).find(qn("a:t")).text = head
    txb.append(h)
    for line in body:
        p = copy.deepcopy(b)
        p.find(qn("a:r")).find(qn("a:t")).text = line
        txb.append(p)


def main(src, out, full=False):
    prs = Presentation(src)
    def find(head):
        return next(s for s in prs.slides if any(sh.has_text_frame and sh.text_frame.text.startswith(head) for sh in sh_list(s)))
    keep = [find("ステップ配信｜追加から14日"), find("リッチメニュー設計案"), find("リッチメニュー実装イメージ")]
    ids = {k.slide_id for k in keep}
    if not full:                                              # 試作は3枚だけ。full＝元の資料に入れる（ほかのページは触らない）
        for s in [s for s in prs.slides if s.slide_id not in ids]:
            drop_slide(prs, s)
    for sh in [sh for sh in keep[1].shapes if 4 <= sh.shape_id <= 27]:   # 画像の裏に隠れていた古い枠（新しい画像は背が低く、はみ出して見える）
        sh._element.getparent().remove(sh._element)
    swap_image(keep[1], 33, IMG27)
    swap_image(keep[2], 33, IMG28)
    for k in (keep[1], keep[2]):                               # 画像サイズ（2500×1686px など）の表示は外す（殿村さん）
        set_paras(next(sh for sh in k.shapes if sh.shape_id == 30), ["入会タグで①→③に自動切替。"])
    box, txt = (next(sh for sh in keep[1].shapes if sh.shape_id == i) for i in (29, 30))   # 27は1行になったので枠を28と同じ高さに
    box.height = Inches(0.52)
    txt.top, txt.height = box.top + Inches(0.11), Inches(0.3)
    s = keep[0]
    by = {sh.shape_id: sh for sh in s.shapes}

    tbl = by[4]._element.find(".//" + qn("a:tbl"))
    trs = tbl.findall(qn("a:tr"))
    src_rows = [copy.deepcopy(tr) for tr in trs]
    for tr in trs[1:]:
        tbl.remove(tr)
    row_h = Inches(0.42)
    trs[0].set("h", str(int(Inches(0.4))))
    for k, (base, new, *texts) in enumerate(ROWS):
        tr = copy.deepcopy(src_rows[base])
        tr.set("h", str(int(row_h)))
        tcs = tr.findall(qn("a:tc"))
        for tc, t in zip(tcs, texts):
            set_cell_text(tc, t)
        if full:                                              # 元の資料：塗り分けはせず、白と灰色の交互にそろえる
            for tc in tcs:
                v = tc.find(qn("a:tcPr")).find(qn("a:solidFill")).find(qn("a:srgbClr"))
                if v.get("val") != "1F2A5C":
                    v.set("val", ROW_FILL[k % 2])
        elif new:
            for tc in tcs:
                set_fill(tc, NEW_FILL)
        tbl.append(tr)
    for n, e in enumerate(tbl.iter("{http://schemas.microsoft.com/office/drawing/2014/main}rowId")):
        e.set("val", str(10000 + n))
    by[4].height = int(Inches(0.4) + row_h * len(ROWS))

    set_paras(by[3], LEAD)
    top = by[4].top + by[4].height + Inches(0.1)
    for sid, (head, body) in zip((5, 6, 7), CARDS):
        set_card(by[sid], head, body)
        by[sid].top, by[sid].height = top, Inches(0.52)
    set_paras(by[8], NOTE_FULL if full else NOTE)
    by[8].top, by[8].height = top + Inches(0.6), Inches(0.36)
    prs.save(out)
    print("saved:", out)


def sh_list(s):
    return list(s.shapes)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2], "--full" in sys.argv)
