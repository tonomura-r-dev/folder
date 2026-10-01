# -*- coding: utf-8 -*-
"""サンクスLINEの費用と実績（1枚）（2026-10-01）
ver2.1の3枚目だけを1枚のPPTX/PDFにする。月額は「3万円〜」（殿村さん指示）。
（以下、ver2.1の説明）
3枚で完結：S1・S2 ＋「実績と費用」1枚（元S3を土台に作り直し。元S4・S5は削除）。
実績＝弊社の運用実績（美容クリニック）＋ LINEヤフー公式の導入事例5件。
  python3 _build/build_thanks_line_hiyo_jisseki.py <ver1.4.pptx>
数字の出典（2026-10-01 に元ページで確認）：
  https://www.lycbiz.com/jp/case-study/line-official-account/acne-clinic/
    「ステップ配信」でクリニックの予約数も約120％まで増加（ステップ配信の実施前と比較）
  https://www.lycbiz.com/jp/case-study/line-official-account/sugitama/
    １年で「LINEで予約」経由の予約数は約6倍増加（2021年12月と2022年12月を比較）※ver2.1では未掲載
  https://www.lycbiz.com/jp/case-study/line-official-account/hokepon/
    LINE通知メッセージで予約日時のリマインド → 予約からの面談実施率が5ポイント増加
  https://www.lycbiz.com/jp/case-study/line-official-account/kotobus/
    LINE通知メッセージを受信したユーザーの友だち追加率は約70%
  https://www.lycbiz.com/jp/case-study/line-official-account/uzuz/
    LINE経由の問い合わせのうち、面談予約率が55%
  https://www.lycbiz.com/jp/case-study/line-official-account/saintmarccafe/
    メッセージを開封した人の約40%が開封当日に来店
"""
import copy
import sys
from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.util import Cm, Pt

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "20261001_サンクスLINEの費用と実績.pptx"
prs = Presentation(sys.argv[1])
NAVY, GREEN, INK, GRAY = RGBColor(0x1F, 0x28, 0x5A), RGBColor(0x06, 0xC7, 0x55), RGBColor(0x33, 0x33, 0x33), RGBColor(0x7F, 0x7F, 0x7F)

# ---- 実績ページ：元S3を土台に作り直す ----
s = prs.slides[2]
by = lambda n: next(x for x in s.shapes if x.name == n)
tpl_navy = copy.deepcopy(by("Rounded Rectangle 4")._element)   # 紺の帯
tpl_card = copy.deepcopy(by("Rounded Rectangle 6")._element)   # 薄い青のカード
tpl_green = copy.deepcopy(by("Rounded Rectangle 8")._element)  # 緑の帯
note_src = copy.deepcopy(prs.slides[3].shapes[-1]._element)     # 元S4の※注記（テキストボックス）
fee_note = copy.deepcopy(prs.slides[4].shapes[-2]._element)     # 元S5の※注記（APIツール併用可…）
keep = {"TextBox 1", "TextBox 2", "Connector 3"}
for sh in list(s.shapes):
    if sh.name not in keep:
        sh._element.getparent().remove(sh._element)


def set_runs(tf_shape, lines):
    """lines = [(text, size, bold, color), ...] 1要素＝1段落。書式は元の最初のrunを引き継ぐ。"""
    tf = tf_shape.text_frame
    base_p = tf.paragraphs[0]._p
    base_r = base_p.find(qn("a:r"))
    rpr0 = copy.deepcopy(base_r.find(qn("a:rPr"))) if base_r is not None else None
    for p in list(tf._txBody.findall(qn("a:p")))[1:]:
        tf._txBody.remove(p)
    for r in list(base_p.findall(qn("a:r"))) + list(base_p.findall(qn("a:br"))):
        base_p.remove(r)
    empty_p = copy.deepcopy(base_p)
    for i, (text, size, bold, color) in enumerate(lines):
        p = base_p if i == 0 else copy.deepcopy(empty_p)
        if i:
            tf._txBody.append(p)
        r = p.makeelement(qn("a:r"), {})
        if rpr0 is not None:
            r.append(copy.deepcopy(rpr0))
        t = r.makeelement(qn("a:t"), {})
        t.text = text
        r.append(t)
        end = p.find(qn("a:endParaRPr"))
        (end.addprevious(r) if end is not None else p.append(r))
    for para, (text, size, bold, color) in zip(tf.paragraphs, lines):
        for run in para.runs:
            run.font.size, run.font.bold = Pt(size), bold
            run.font.color.rgb = color


def add(tpl, x, y, w, h, lines):
    el = copy.deepcopy(tpl)
    s.shapes._spTree.append(el)
    sh = s.shapes[-1]
    sh.left, sh.top, sh.width, sh.height = Cm(x), Cm(y), Cm(w), Cm(h)
    set_runs(sh, lines)
    return sh


# タイトル・リード
set_runs(by("TextBox 1"), [("実績と費用｜LINEで追うと、来店・面談が伸びる", 16, True, RGBColor(0x00, 0x20, 0x60))])
set_runs(by("TextBox 2"), [("弊社の運用実績と、LINEヤフー公式の導入事例", 16, False, INK)])

WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CASES = [   # (見出し, 何をしたか, 数字, 弊社か)
    ("弊社実績｜美容クリニック", "完了画面からLINEへ（サンクスLINE）", "来院率 40〜50%改善", True),
    ("保険見直し本舗", "予約日時をLINEでリマインド", "面談実施率 ＋5ポイント", False),
    ("琴平バス（観光バス）", "予約の通知をLINEで受け取った方のうち", "友だち追加 約70%", False),
    ("皮膚科クリニック", "友だち追加後のステップ配信", "予約数 約120%", False),
    ("UZUZ（就職支援）", "LINEからの問い合わせ", "面談予約率 55%", False),
    ("サンマルクカフェ", "メッセージを開封した方のうち", "当日に来店 約40%", False),
]
X0, W, G = 1.46, 7.93, 0.4
for i, (who, what, big, own) in enumerate(CASES):
    x, y = X0 + (i % 3) * (W + G), 4.55 + (i // 3) * 3.85
    add(tpl_green if own else tpl_navy, x, y, W, 0.9, [(who, 14, True, WHITE)])
    add(tpl_card, x, y + 1.0, W, 2.6, [(what, 11, False, INK), (big, 18, True, NAVY)])

# 費用（元S5の内容）
Y = 12.45
add(tpl_navy, X0, Y, 3.0, 1.3, [("費用", 16, True, WHITE)])
add(tpl_green, X0 + 3.2, Y, 5.4, 1.3, [("初期 15万円", 18, True, WHITE)])
add(tpl_green, X0 + 8.8, Y, 5.4, 1.3, [("月額 3万円〜", 18, True, WHITE)])
s.shapes._spTree.append(fee_note)
fn = s.shapes[-1]
fn.left, fn.top, fn.width, fn.height = Cm(X0 + 14.4), Cm(Y), Cm(3 * W + 2 * G - 14.4), Cm(1.3)
set_runs(fn, [("※APIツール（Lステップ等）と併用可", 11, False, GRAY), ("※完了画面を増やす場合は追加費用", 11, False, GRAY)])
fn.text_frame.vertical_anchor = 3  # 中央

add(tpl_navy, X0, 14.1, 3 * W + 2 * G, 1.3,
    [("広告のご提案とあわせて、LINEへの導線設計からご提案可能", 17, True, WHITE)])
s.shapes._spTree.append(note_src)
note = s.shapes[-1]
note.left, note.top, note.width, note.height = Cm(X0), Cm(15.6), Cm(3 * W + 2 * G), Cm(0.8)
set_runs(note, [("出典：LINEヤフー for Business 導入事例（保険見直し本舗／琴平バス／アクネクリニック／UZUZ／サンマルクカフェ）", 10, False, GRAY)])
# 複製した図形のIDが重ならないよう振り直す
for n, c in enumerate(s.shapes._spTree.iter(qn("p:cNvPr")), start=2):
    c.set("id", str(n))

# ---- 元S4（広告とLINEの役割）・元S5（進め方・費用）を削除 ----
sldIdLst = prs.slides._sldIdLst
for sid in list(sldIdLst)[3:] + list(sldIdLst)[:2]:   # 元S4・S5に加えてS1・S2も外し、1枚にする
    prs.part.drop_rel(sid.get(qn("r:id")))
    sldIdLst.remove(sid)

# ---- 字体をメイリオに統一（patch_yaruki_v32_iikiri.py と同じ手順）----
AFTER = {qn(x) for x in ("a:sym", "a:hlinkClick", "a:hlinkMouseOver", "a:rtl", "a:extLst")}
for sl in prs.slides:
    for r in sl.shapes._spTree.iter(qn("a:r")):
        if r.find(qn("a:rPr")) is None:
            r.insert(0, r.makeelement(qn("a:rPr"), {"lang": "ja-JP"}))
    for tag in ("a:rPr", "a:endParaRPr", "a:defRPr"):
        for rpr in sl.shapes._spTree.iter(qn(tag)):
            fs = []
            for ft in ("a:latin", "a:ea", "a:cs"):
                e = rpr.find(qn(ft))
                if e is None:
                    e = rpr.makeelement(qn(ft), {})
                else:
                    rpr.remove(e)
                e.set("typeface", "メイリオ")
                fs.append(e)
            nxt = next((c for c in rpr if c.tag in AFTER), None)
            for e in fs:
                (nxt.addprevious(e) if nxt is not None else rpr.append(e))
prs.save(str(OUT))
print("saved:", OUT.name, len(prs.slides), "枚")
