# ver2.1 → ver2.2：応募済み（サンクスLINEから友だち）をCVの計算から外す
import sys
from openpyxl import load_workbook
from openpyxl.workbook.properties import CalcProperties
SRC,OUT=sys.argv[1],sys.argv[2]
wb=load_workbook(SRC)
for n in ('SIM1_通知メッセ有','SIM2_通知メッセ無','SIM3_LINE登録2%'):
    ws=wb[n]
    for c in 'EFGHIJ':
        # 有効な友だちのうち、応募済み（24行・25行の累計）を除いた人だけがリッチメニューから応募する
        ws[f'{c}37']=f'=ROUND({c}29*(1-SUM($E$24:{c}25)/{c}28)*0.3*2*0.015,0)'
    ws['N37']='リッチメニュー経由の応募：応募済みを除いた有効な友だち×月に使う率30%×2回検索×応募率1.5%'
    ws['N24']='完了画面：サイト応募×遷移80%×追加45%×立ち上げ係数（応募済みの人。リマインドのみでCVには数えない）'
wb['SIM3_LINE登録2%']['N36']='配信経由の応募（応募済みを除く・仮）'
for i,c in enumerate('EFGHIJ'):
    wb['SIM3_LINE登録2%'][f'{c}36']=[1,2,3,3,4,6][i]
wb.calculation=CalcProperties(fullCalcOnLoad=True)
wb.save(OUT)
