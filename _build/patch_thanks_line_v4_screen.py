# -*- coding: utf-8 -*-
"""サンクスLINEのご提案（4枚）1枚目：「申込み完了 → LINE起動」の動作画面を入れる（2026-10-02 佐村さん指摘）
画面は前の資料（20261001版）に入れていた画像 `_images/thanks_line_flow.png` の最初の2コマ。
  python3 _build/patch_thanks_line_v4_screen.py <20261002_サンクスLINEのご提案.pptx>   （同じファイル名で上書き）
"""
import copy
import sys
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.util import Cm, Pt

path = sys.argv[1]
prs = Presentation(path)
s = prs.slides[0]
keep = ("TextBox 1", "TextBox 2", "Connector 3", "TextBox 19")
for sh in list(s.shapes):
    if sh.name not in keep:
        sh._element.getparent().remove(sh._element)
W = 12.8
H = W * 660 / 797
s.shapes.add_picture("_images/thanks_line_screen_start.png", Cm(2.4), Cm(5.2), Cm(W), Cm(H))
body = next(sh for sh in s.shapes if sh.name == "TextBox 19")
body.left, body.top, body.width, body.height = Cm(16.1), Cm(5.6), Cm(9.2), Cm(9.5)
tf = body.text_frame
ps = tf.paragraphs
plain = ps[0]
plain._p.addprevious(copy.deepcopy(plain._p))
ps = tf.paragraphs


def settext(p, t):
    p.runs[0].text = t
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)


settext(ps[0], "・新規は友だち追加、既存の友だちはトーク画面へ")
settext(ps[1], "・入力内容を引き継ぎ、配信の出し分けに活用")
settext(ps[2], "・今のアカウントにそのまま導入、APIツールとも併用可")
for p in tf.paragraphs:
    for r in p.runs:
        r.font.size = Pt(16)
    p.space_after = Pt(14)
# 注記
note = copy.deepcopy(body._element)
s.shapes._spTree.append(note)
nt = s.shapes[-1]
nt.left, nt.top, nt.width, nt.height = Cm(3.0), Cm(16.8), Cm(21.5), Cm(0.6)
tfn = nt.text_frame
for p in list(tfn._txBody.findall(qn("a:p")))[1:]:
    tfn._txBody.remove(p)
settext(tfn.paragraphs[0], "※画面はイメージです")
for r in tfn.paragraphs[0].runs:
    r.font.size, r.font.bold = Pt(10), False
    r.font.color.rgb = RGBColor(0x7F, 0x7F, 0x7F)
for n, c in enumerate(s.shapes._spTree.iter(qn("p:cNvPr")), start=2):
    c.set("id", str(n))
prs.save(path)
print("saved", path)
