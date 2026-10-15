# -*- coding: utf-8 -*-
"""ジム・フィットネス業界の資料（ver3）を、EC業界の資料の型にそろえる（2026-10-15 殿村さん指示）。

    python -m _build.gym_fmt.build <ジムver3.pptx> <EC業界.pptx（型の手本）> <出力.pptx>

  殿村さん：「ECの業界資料を基に、FMTを整えてほしい」→（確認）「塾ver3をECの型にそろえる」→「塾を添付したジム・フィットネスに」。
  → 中身はそのまま。EC資料の型＝ ①見出しを「章｜〜」の形に ②アジェンダを章の箱（1〜5を2ページ目、6〜9を中盤）に
     ③章の順番（市場と実態 → なぜLINEか → 全体設計 → 構築 → ニーズ調査 → [アジェンダ] → 配信設計 → 具体施策 → 成果とプラン → 締め）にページを並べ替え
     ④区切りだけのページ（資料アジェンダ：施策の詳細／具体プラン提示／ご提案にあたり）は外す（中盤のアジェンダに集約）
  ⑤見出しの箱の位置・高さ・色をEC（x0.6・y0.15・高さ0.37in・紺002060）にそろえる（ずれていた12枚）

  入力：_data/gym/ジム・フィットネス業界_LINEOA施策提案ver3_殿村さん保存版（2026-09-09）.pptx
        _data/ec/EC業界_LINEOA施策提案_殿村さん保存版（2026-09-14）.pptx（43枚・型の手本。S2・S23のアジェンダを写す）
  PowerPointで保存したzipに、直したスライドのXMLだけ差し込み、presentation.xml の並び順を入れ替え、外すページはパーツごと外す。
"""
import copy
import re
import sys
import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches

from _build.arauchi_v4.build import c14n, normalize_decl, rel_set, slide_order
from _build.arauchi_v8.pages import replace_in_paragraphs, shape_by_id

N_BASE = 46
DROP = [21, 38, 40]                      # 区切りだけのページ（施策の詳細／具体プラン提示／ご提案にあたり）
# 新しい並び（元の番号）。12＝「配信の考え方」の区切りページを中盤のアジェンダに作り替えて使う
ORDER = [1, 2,
         15, 9, 10, 11,                  # 1 市場と実態
         3, 4,                           # 2 なぜLINEか
         13, 14, 22, 5,                  # 3 全体設計
         32, 6, 8, 7, 26, 27, 28, 29,    # 4 構築
         16, 17, 18, 19, 20,             # 5 ニーズ調査（20は非表示のまま）
         12,                             # 資料アジェンダ（6〜9）
         23, 24, 25, 30, 31, 33, 34, 35, 36,   # 6 配信設計
         41, 42, 43, 44, 45,             # 7 具体施策（42〜45は非表示のまま）
         37, 39,                         # 8 成果とプラン
         46]                             # 9 締め
assert sorted(ORDER + DROP) == list(range(1, N_BASE + 1))
NEW = {old: i for i, old in enumerate(ORDER, 1)}          # 元の番号 → 新しい番号


def new_no(old):
    return NEW[old]


# 見出し（元の番号 → 新しい見出し）。「現状分析｜」「具体施策｜18の施策…」はもともと「章｜〜」の形なので触らない
TITLES = {
    15: "市場｜獲得単価の推移",
    3: "なぜLINEか｜アプリ登録のハードルが上がっている",
    4: "なぜLINEか｜LINE施策における重要な考え方",
    13: "全体設計｜カスタマージャーニー8フェーズ",
    14: "全体設計｜ジム・フィットネスの8フェーズ",
    22: "全体設計｜ユーザーニーズ・インサイトの変化と施策方針",
    5: "全体設計｜流入元別のCV率と、LINEで打てる施策",
    32: "構築｜施策の考え方（誰を、どのCVへ）",
    6: "構築｜施策展開図（初期・月次）",
    8: "構築｜友だち追加と育成の2軸",
    7: "構築｜友だち追加後15日でアクティブ率が落ちる",
    26: "構築｜リッチメニュー設計案（3タブ × 6枠）",
    27: "構築｜リッチメニュー実装イメージ（3タブの実物）",
    28: "構築｜トーク画面設計案（あいさつメッセージと30秒診断）",
    29: "構築｜トーク画面イメージ（友だち追加の直後に届く1通目）",
    16: "ニーズ調査｜シーズナリティ",
    17: "ニーズ調査｜5年トレンド（参考データ）",
    18: "ニーズ調査｜年間の配信設計マトリクス（時期 × 対象 × 訴求 × CV）",
    19: "ニーズ調査｜前後検索",
    20: "ニーズ調査｜前後検索（補足）ダイエット需要は“医療”へ流出している",
    23: "配信設計｜ステップ配信のブレスト",
    24: "配信設計｜ステップ配信① 追加から14日で体験予約を取り切る",
    25: "配信設計｜ステップ配信② 入会後14日で「やめない形」を作る",
    30: "配信設計｜企画配信（需要の波に合わせた企画）",
    31: "配信設計｜通知メッセージと参加型の企画",
    33: "配信設計｜改善モデル①",              # 右に「導入推奨／中長期で導入検討」の札があるので副題は付けない
    34: "配信設計｜改善モデル②",
    35: "配信設計｜改善モデル③",
    36: "配信設計｜改善モデル④",
    42: "具体施策①｜低ハードルCVでリードを大量獲得",
    43: "具体施策②｜保有リストの再活性化",
    44: "具体施策③｜リピート・クロスセルの計画実行",
    45: "具体施策④｜チャット対応の簡略化・自動化",
    37: "成果｜導入成果シミュレーション（Before / After）",
    39: "費用プラン｜ご提供プラン（運用構成）",
}
TITLE_NAVY = "002060"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def title_shape(s):
    """左上（x<0.8in・y<0.3in）の見出しの箱"""
    c = [sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()
         and sh.top is not None and sh.top < Inches(0.3) and sh.left < Inches(0.8)]
    assert len(c) == 1, [(sh.shape_id, sh.text_frame.text) for sh in c]
    return c[0]


def set_title(sh, text):
    """最初のrunの書式のまま、見出しの文字を差し替える（ほかのrun・段落は消す）"""
    tb = sh.text_frame._txBody
    ps = tb.findall(qn("a:p"))
    for p in ps[1:]:
        tb.remove(p)
    p = ps[0]
    runs = p.findall(qn("a:r"))
    runs[0].find(qn("a:t")).text = text
    for el in list(p):
        if el.tag in (qn("a:r"), qn("a:br"), qn("a:fld")) and el is not runs[0]:
            p.remove(el)


def normalize_title_box(sh):
    """ECの見出しの箱にそろえる（x0.6・y0.15・高さ0.37in・上下中央・紺）"""
    sh.left, sh.top, sh.height = Inches(0.6), Inches(0.15), Inches(0.37)
    sh.text_frame._txBody.find(qn("a:bodyPr")).set("anchor", "ctr")
    for r in sh._element.iter(qn("a:rPr")):
        for c in r.iter(qn("a:srgbClr")):
            if c.get("val") == "1F2A5C":
                c.set("val", TITLE_NAVY)


def set_run_texts(sh, para_idx, texts):
    """段落の run を前から順に差し替える（数が合わないときは止める）"""
    p = sh.text_frame._txBody.findall(qn("a:p"))[para_idx]
    runs = p.findall(qn("a:r"))
    assert len(runs) == len(texts), (sh.shape_id, para_idx, [r.find(qn("a:t")).text for r in runs])
    for r, t in zip(runs, texts):
        r.find(qn("a:t")).text = t


def copy_slide_body(dst, src, skip_ids=()):
    """src の図形を丸ごと dst に写す（dst の図形は全部消す）。どちらも画像・リンクなし（レイアウトのrelsだけ）"""
    tree = dst.shapes._spTree
    for el in list(tree):
        if el.tag not in (qn("p:nvGrpSpPr"), qn("p:grpSpPr")):
            tree.remove(el)
    for el in src.shapes._spTree:
        if el.tag in (qn("p:nvGrpSpPr"), qn("p:grpSpPr")):
            continue
        nv = el.find(".//" + qn("p:cNvPr"))
        if nv is not None and int(nv.get("id")) in skip_ids:
            continue
        assert not [a for e in el.iter() for a in e.attrib if a.startswith("{%s}" % R_NS)], "rels付きの図形"
        tree.append(copy.deepcopy(el))


# ---------------------------------------------------------------- アジェンダ（ECのS2・S23を写して中身を差し替える）
AGENDA1 = [  # (番号の図形, 章名, 中身)
    (8, 9, "市場と実態", 11, "検索需要と獲得単価 → 各社LINEの実態・4つの型"),
    (13, 14, "なぜLINEか", 16, "登録のハードル・LINE施策の考え方"),
    (18, 19, "全体設計", 21, "8フェーズ・ニーズの変化・流入元別CV率"),
    (23, 24, "構築の型", 26, "考え方・展開図・2軸・追加後15日・画面設計"),
    (3, 6, "ニーズ調査", 10, "シーズナリティ・年間マトリクス・前後検索"),
]


def agenda1(s, ec):
    copy_slide_body(s, ec.slides[1])
    lead = shape_by_id(s, 5)
    set_run_texts(lead, 0, ["検索需要は横ばい、店舗は増え続けている", "。"])
    set_run_texts(lead, 1, ["「", "集める", "」", "⇒", "「体験予約から継続まで、LINEでつなぐ」"])
    for _num, name_id, name, desc_id, desc in AGENDA1:
        set_run_texts(shape_by_id(s, name_id), 0, [name])
        set_run_texts(shape_by_id(s, desc_id), 0, [desc])
    set_run_texts(shape_by_id(s, 27), 0, [f"※後半（6 配信設計／7 具体施策／8 成果とプラン／9 締め）はS{new_no(12)}に記載"])


def pages(a, b=None):
    return f"S{new_no(a)}" if b is None else f"S{new_no(a)}-{new_no(b)}"


def agenda2(s, ec):
    copy_slide_body(s, ec.slides[22], skip_ids=(6,))       # ECのS23は同じ見出しの文字が2つ重なっている（#5・#6）→1つに
    lead = shape_by_id(s, 5)
    set_run_texts(lead, 0, ["ここからは「何を、いつ、どう配信するか」の具体。"])
    set_run_texts(lead, 1, ["CVは 体験予約（CV①）→ 来店（CV②）→ 継続（CV③）の3段。"])
    rows = [
        (9, "配信設計", 10, pages(23, 36), 11, "ステップ配信・企画・通知メッセージ・改善モデル①〜④"),
        (14, "具体施策", 15, pages(41), 16, "18の施策をジャーニー上に配置"),
        (19, "成果とプラン", 20, pages(37, 39), 21, "導入成果シミュレーション・ご提供プラン"),
        (24, "締め", 25, pages(46), 26, "貴社の明るい未来"),
    ]
    for name_id, name, pg_id, pg, desc_id, desc in rows:
        set_run_texts(shape_by_id(s, name_id), 0, [name])
        set_run_texts(shape_by_id(s, pg_id), 0, [pg])
        set_run_texts(shape_by_id(s, desc_id), 0, [desc])
    set_run_texts(shape_by_id(s, 27), 0, ["※CV①＝体験予約／CV②＝来店／CV③＝継続"])


def apply(prs, ec):
    sl = prs.slides
    agenda1(sl[1], ec)
    agenda2(sl[11], ec)
    for old, text in TITLES.items():
        sh = title_shape(sl[old - 1])
        set_title(sh, text)
        if abs(sh.top - Inches(0.15)) > Inches(0.01) or abs(sh.left - Inches(0.6)) > Inches(0.01):
            normalize_title_box(sh)
    # 非表示のS20：「S13〜14の『ダイエット検索の5年半減』」→ 5年トレンドのページ（元S17）の新しい番号
    assert replace_in_paragraphs(sl[19], "S13〜14の", f"S{new_no(17)}の") == 1
    # 配信設計｜ステップ配信のブレスト：見出しが長くなり、右の「パーソナルジム 女性の検索前後データ」がロゴに重なるので、
    # 出典の「※ LINEヤフー media Journey（2025-2026）」を2行にして幅を詰め、札を左へ寄せる（ページ下は図で埋まっているので移さない）
    s23 = sl[22]
    src, chip, tail = shape_by_id(s23, 5), shape_by_id(s23, 479), shape_by_id(s23, 481)
    gap = tail.left - chip.left
    p0 = src.text_frame._txBody.findall(qn("a:p"))[0]
    assert "".join(t.text for t in p0.iter(qn("a:t"))).replace(" ", "") == "※LINEヤフーmediaJourney（2025-2026）"
    p1 = copy.deepcopy(p0)
    p0.addnext(p1)
    for para, text in ((p0, "※ LINEヤフー media Journey"), (p1, "（2025-2026）")):
        runs = para.findall(qn("a:r"))
        runs[0].find(qn("a:t")).text = text
        for r in runs[1:]:
            para.remove(r)
    bp = src.text_frame._txBody.find(qn("a:bodyPr"))
    bp.set("wrap", "square")
    bp.set("lIns", "0")
    bp.set("rIns", "0")
    src.left, src.top, src.width, src.height = Inches(4.34), Inches(0.13), Inches(1.80), Inches(0.36)
    chip.left = Inches(6.18)
    tail.left = chip.left + gap
    fixed = sum(fix_fonts(sl[i - 1]) for i in ORDER)
    print("メイリオにそろえた文字:", fixed)


# ---------------------------------------------------------------- 字体：全部メイリオ（CLAUDE.md 2026-10-13。元のver3に指定なし・+mn-ea が98か所あった）
FONT = "メイリオ"
_AFTER_LATIN = ("a:ea", "a:cs", "a:sym", "a:hlinkClick", "a:hlinkMouseOver", "a:rtl", "a:extLst")


def _set_face(rpr, tag, before_tags):
    el = rpr.find(qn(tag))
    if el is None:
        el = etree.Element(qn(tag))
        nxt = next((c for c in rpr if c.tag in [qn(t) for t in before_tags]), None)
        (nxt.addprevious(el) if nxt is not None else rpr.append(el))
    for k in list(el.attrib):
        del el.attrib[k]
    el.set("typeface", FONT)


def fix_fonts(slide):
    n = 0
    for r in slide._element.iter(qn("a:r")):
        if not "".join((t.text or "") for t in r.iter(qn("a:t"))).strip():
            continue
        rpr = r.find(qn("a:rPr"))
        if rpr is None:
            rpr = etree.Element(qn("a:rPr"), lang="ja-JP", altLang="en-US")
            r.insert(0, rpr)
        lat, ea = rpr.find(qn("a:latin")), rpr.find(qn("a:ea"))
        if lat is not None and lat.get("typeface") == FONT and ea is not None and ea.get("typeface") == FONT:
            continue
        _set_face(rpr, "a:latin", _AFTER_LATIN)
        _set_face(rpr, "a:ea", _AFTER_LATIN[1:])
        n += 1
    return n


# ---------------------------------------------------------------- 差し込み・並べ替え・削除
def repack(base_path, built_path, out_path):
    zb, zn = zipfile.ZipFile(base_path), zipfile.ZipFile(built_path)
    base_order, built_order = slide_order(zb), slide_order(zn)
    assert len(base_order) == len(built_order) == N_BASE
    drops = [base_order[d - 1] for d in DROP]
    gone = set()
    for drop in drops:
        drop_rels = drop.replace("slides/", "slides/_rels/") + ".rels"
        gone |= {drop, drop_rels}
        for r in etree.fromstring(zb.read(drop_rels)):
            assert r.get("Type").endswith("/slideLayout"), r.get("Type")
    for name in zb.namelist():
        if name.endswith(".rels") and name not in gone and name != "ppt/_rels/presentation.xml.rels":
            data = zb.read(name).decode("utf-8")
            for drop in drops:
                assert 'slides/%s"' % drop.split("/")[-1] not in data, name

    prels = zb.read("ppt/_rels/presentation.xml.rels").decode("utf-8")
    pres = zb.read("ppt/presentation.xml").decode("utf-8")
    ct = zb.read("[Content_Types].xml").decode("utf-8")
    rid_of = {}
    for name in base_order:
        m = re.search(r'<Relationship [^>]*Target="slides/%s"[^>]*/>' % re.escape(name.split("/")[-1]), prels)
        rid_of[name] = re.search(r'Id="(rId\d+)"', m.group(0)).group(1)
    for drop in drops:
        m = re.search(r'<Relationship [^>]*Target="slides/%s"[^>]*/>' % re.escape(drop.split("/")[-1]), prels)
        prels = prels.replace(m.group(0), "")
        ct, n = re.subn(r'<Override [^>]*PartName="/%s"[^>]*/>' % re.escape(drop), "", ct)
        assert n == 1, drop
    # p:sldIdLst を新しい並びに
    m = re.search(r"<p:sldIdLst>(.*?)</p:sldIdLst>", pres, re.S)
    ids = {re.search(r'r:id="(rId\d+)"', e).group(1): e for e in re.findall(r"<p:sldId [^>]*/>", m.group(1))}
    assert len(ids) == N_BASE
    new_list = "".join(ids[rid_of[base_order[old - 1]]] for old in ORDER)
    pres = pres.replace(m.group(0), f"<p:sldIdLst>{new_list}</p:sldIdLst>")

    app = zb.read("docProps/app.xml").decode("utf-8")
    nb, nn = N_BASE, len(ORDER)
    assert f"<Slides>{nb}</Slides>" in app
    app = app.replace(f"<Slides>{nb}</Slides>", f"<Slides>{nn}</Slides>")
    head = f"<vt:lpstr>スライド タイトル</vt:lpstr></vt:variant><vt:variant><vt:i4>{nb}</vt:i4>"
    assert head in app
    app = app.replace(head, head.replace(f"<vt:i4>{nb}</vt:i4>", f"<vt:i4>{nn}</vt:i4>"))
    mv = re.search(r'<vt:vector size="(\d+)" baseType="lpstr">', app)
    app = app.replace(mv.group(0), f'<vt:vector size="{int(mv.group(1)) - len(DROP)}" baseType="lpstr">')
    for _ in DROP:
        app = app.replace("<vt:lpstr>PowerPoint プレゼンテーション</vt:lpstr>", "", 1)

    replaced, changed = {}, []
    for i, (b_name, n_name) in enumerate(zip(base_order, built_order), 1):
        if i in DROP:
            continue
        if c14n(zb.read(b_name)) != c14n(zn.read(n_name)):
            replaced[b_name] = normalize_decl(zn.read(n_name))
            changed.append(i)
        b_rels = b_name.replace("slides/", "slides/_rels/") + ".rels"
        n_rels = n_name.replace("slides/", "slides/_rels/") + ".rels"
        if rel_set(zb.read(b_rels)) != rel_set(zn.read(n_rels)):
            raise SystemExit(f"想定外のrelsの違い: {b_rels}")

    with zipfile.ZipFile(out_path, "w") as zo:
        for info in zb.infolist():
            name = info.filename
            if name in gone:
                continue
            data = {"ppt/presentation.xml": pres.encode("utf-8"),
                    "ppt/_rels/presentation.xml.rels": prels.encode("utf-8"),
                    "[Content_Types].xml": ct.encode("utf-8"),
                    "docProps/app.xml": app.encode("utf-8")}.get(name)
            if data is None:
                data = replaced.pop(name) if name in replaced else zb.read(name)
            zo.writestr(info, data, compress_type=info.compress_type)
    assert not replaced
    return changed, sorted(gone)


def main(base, ec_path, out):
    tmp = out + ".tmp.pptx"
    prs, ec = Presentation(base), Presentation(ec_path)
    assert len(prs.slides) == N_BASE and len(ec.slides) == 43
    apply(prs, ec)
    prs.save(tmp)
    changed, gone = repack(base, tmp, out)
    Path(tmp).unlink()
    print("saved:", out)
    print("差し替えたスライド（元の番号）:", changed)
    print("外したパーツ:", gone)


if __name__ == "__main__":
    main(*sys.argv[1:4])
