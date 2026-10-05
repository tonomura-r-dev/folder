# -*- coding: utf-8 -*-
"""チェングロウス ver1.2 → ver1.3（2026-10-05 殿村さん指示）。
・Chapter表紙を4枚追加し、Chapterごとに並べ替える（Ch1=3,9,4,10,11／Ch2=6,7,8／Ch3=12〜15／Ch4=16,17,18,5,19）
・2Pアジェンダを4項目・ページ番号なしに
・全ページの文字量を削る（リードは1行・結論帯の重複を外す・箇条書きは3つ程度・説明は1行）
・5（なぜ11月頃から）／18（費用）／19（スケジュール）は作り直し
数字・費用・出典・事例は ver1.2 のまま（費用のオプション表記だけ殿村さん指定の形に）。
  python3 _build/patch_chengrowth_bridge_v13_chapters.py <ver1.2.pptx>
"""
import copy
import io
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Cm, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
from pptx_parts import *  # noqa

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "20261005_チェングロウス_サイトリニューアルとLINE同時導入のご提案ver1.3.pptx"
prs = Presentation(sys.argv[1])
S = list(prs.slides)  # ver1.2 の19枚（添字＝番号-1）
PALE = "E4E8F6"


def sl(n):
    return S[n - 1]


def shp(slide, name):
    return next(sh for sh in slide.shapes if sh.name == name)


def remove(slide, *names):
    for nm in names:
        e = shp(slide, nm)._element
        e.getparent().remove(e)


def set_paras(sh, items):
    """items = [(text, 元の段落番号)]。元の段落の書式をコピーして文言だけ差し替える。"""
    txb = sh.text_frame._txBody
    olds = list(txb.findall(qn("a:p")))
    news = []
    for text, k in items:
        p = copy.deepcopy(olds[k])
        rs = p.findall(qn("a:r"))
        for r in rs[1:]:
            p.remove(r)
        rs[0].find(qn("a:t")).text = text
        news.append(p)
    for p in olds:
        txb.remove(p)
    for p in news:
        txb.append(p)


def lead(slide, *lines):
    set_paras(shp(slide, "TextBox 2"), [(t, 0) for t in lines])


def rich(s, x, y, w, h, runs_by_para, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE):
    """runs_by_para = [[(text, size, bold, color), ...], ...]（1段落に書式違いの文字を並べる）"""
    tb = s.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Cm(0.1)
    tf.margin_top = tf.margin_bottom = Cm(0.03)
    for i, runs in enumerate(runs_by_para):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = 1.05
        p.space_after = Pt(2)
        for text, size, bold, color in runs:
            r = p.add_run()
            r.text = text
            style_run(r, size, bold, color)
    return tb


def arrow_line(s, x1, y1, x2, y2, color, w=2.0, dash=False):
    c = s.shapes.add_connector(1, Cm(x1), Cm(y1), Cm(x2), Cm(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(w)
    if dash:
        c.line.dash_style = 4
    ln = c.line._get_or_add_ln()
    ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
    st = c._element.find(qn("p:style"))
    if st is not None:
        st.find(qn("a:effectRef")).set("idx", "0")
    return c


def conclusion(s, y, text, size=15):
    hline(s, 2.2, 25.3, y, NAVY, 1.5)
    label(s, 2.2, y + 0.15, 23.1, 1.3, [(text, size, True, NAVY, 0)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


# ================= 本文ページの文字を削る =================
# 3 サイトリニューアルの次の課題
s = sl(3)
lead(s, "サイトを新しくしても、応募しない「検討層・離脱層」は残る")
set_paras(shp(s, "Rounded Rectangle 6"), [("・見やすさ・求人の探しやすさ", 0), ("・応募フォームの入力しやすさ", 1), ("→ 応募する方が増える", 3)])
set_paras(shp(s, "Rounded Rectangle 9"), [("・比較検討中で、まだ応募しない方", 0), ("・フォームの途中で離れた方", 1), ("→ その後の接点が残らない", 2)])
set_paras(shp(s, "TextBox 10"), [("整備士は、複数の経路を比べて仕事を探している（仕事を探した経路・複数回答）", 0)])
remove(s, "Rounded Rectangle 14")

# 4 解決策
s = sl(4)
lead(s, "応募に至らない方を、LINEの友だちとしてストックする")
for nm, t, b in [("Rounded Rectangle 16", "① 後から案内できる", "新着求人・条件に合う求人をいつでも届けられる"),
                 ("Rounded Rectangle 17", "② 会員の予備軍が貯まる", "今すぐ応募しない方もリストに残る"),
                 ("Rounded Rectangle 18", "③ 興味が分かる", "配信のクリックで、興味のある職種が分かる")]:
    set_paras(shp(s, nm), [(t, 0), (b, 1)])
remove(s, "Rounded Rectangle 19")

# 6 タウンワーク事例（リードは2行のまま。出典の友だち数がここに掛かるため）
s = sl(6)
set_paras(shp(s, "Rounded Rectangle 8"), [("1　リッチメニューに人気の検索条件", 0), ("単発・日払い・在宅など", 1)])
set_paras(shp(s, "Rounded Rectangle 10"), [("2　トークで条件を選ぶ", 0), ("条件に合う求人がカードで並ぶ", 2)])
set_paras(shp(s, "Rounded Rectangle 12"), [("3　そのまま求人ページへ", 0), ("検索結果から、サイトで応募", 1)])
remove(s, "Rounded Rectangle 13")

# 7 タウンワークの勝因
s = sl(7)
lead(s, "勝因は、ワンタップで求人を探せる常設の入口（リッチメニュー）")
set_paras(shp(s, "TextBox 6"), [("求人を受け取るための会員連携を案内", 0)])
set_paras(shp(s, "TextBox 9"), [("時間帯などの条件がワンタップ", 0)])
set_paras(shp(s, "TextBox 12"), [("条件を入れた検索結果ページ", 0)])
set_paras(shp(s, "TextBox 15"), [("リクルートIDでログイン", 0)])

# 8 自動車求人サイトへの応用
s = sl(8)
lead(s, "資格・職種・こだわり条件は、LINEのワンタップ検索と相性が良い")
set_paras(shp(s, "Rounded Rectangle 13"), [("掲載中の求人は、全国2,134件", 0), ("（うち正社員・整備士946件、正社員・営業180件）", 1)])
remove(s, "Rounded Rectangle 14")

# 9 現状のボトルネック
s = sl(9)
lead(s, "現行サイトは「その場の応募」のみで、求職者データが残らない")
for nm, t, b in [("Rounded Rectangle 13", "① 検討層の受け皿がない", "比較検討中の方と接点を持てない"),
                 ("Rounded Rectangle 14", "② 登録が使われにくい", "氏名・経歴・パスワードの入力が壁に"),
                 ("Rounded Rectangle 15", "③ 名簿が残らない", "長期的に育てるリストがない")]:
    set_paras(shp(s, nm), [(t, 0), (b, 1)])
remove(s, "Rounded Rectangle 16")

# 10 エントリーハードルの低減
s = sl(10)
lead(s, "「LINEでログイン」とLINE Profile+で、登録は「確認して押すだけ」")
set_paras(shp(s, "Rounded Rectangle 7"), [("利用の条件", 0), ("日本の法人のみ", 1), ("項目ごとに申請・LINEヤフー社の審査あり", 2),
                                          ("ユーザーの同意が必要", 3)])
remove(s, "Rounded Rectangle 8")

# 11 ストック顧客へのアプローチ
s = sl(11)
lead(s, "蓄積した友だちに、条件に合う求人をLINEで直接届ける")
for nm, t, b in [("Rounded Rectangle 5", "① 友だち（ストック）", "応募前の方とも接点が残る"),
                 ("Rounded Rectangle 7", "② クリックで自動分類", "整備士・営業・未経験OKなどに分かれる"),
                 ("Rounded Rectangle 9", "③ 興味に合う求人を配信", "新着求人・条件に合う求人を届ける")]:
    set_paras(shp(s, nm), [(t, 0), (b, 1)])
remove(s, "Rounded Rectangle 14")

# 12 連動導線
s = sl(12)
lead(s, "新サイトのあらゆる接点から、自然にLINEへ誘導")
remove(s, "Rounded Rectangle 28")

# 13 ユーザー体験①
s = sl(13)
lead(s, "ポップアップ・ボタンで、サイトからその場でLINEへ")
set_paras(shp(s, "TextBox 7"), [("離脱防止ポップアップ（動線00）", 0), ("離脱しそうな方に「新着求人をLINEでお届け」", 1)])
set_paras(shp(s, "TextBox 8"), [("「LINEで登録」ボタン（動線01）", 0), ("会員登録と友だち追加が同時に完了", 1)])

# 14 ユーザー体験②
s = sl(14)
lead(s, "自動車求人に特化したリッチメニューを、トーク画面に常設")
set_paras(shp(s, "TextBox 6"), [("リッチメニュー（トーク画面の下に常設）", 0)])
set_paras(shp(s, "Rounded Rectangle 8"), [("LINE内で完結", 0), ("① 職種をタップ", 1), ("② 条件に合う求人がカードで届く", 2),
                                          ("③ 求人ページからサイトで応募", 3)])

# 15 自動応答・アンケート
s = sl(15)
lead(s, "友だち追加直後のアンケートで、資格・希望時期を把握")
set_paras(shp(s, "Rounded Rectangle 15"), [("回答に合わせた配信の例", 0), ("整備士2級・すぐ → 整備士の新着求人を優先", 1),
                                           ("未経験・情報収集中 → 資格取得支援の求人", 2), ("友だち追加から14日間、全10通のステップ配信", 3)])
remove(s, "Rounded Rectangle 17")

# 16 シミュレーション
lead(sl(16), "新サイトの来訪（月約6,100人）に応じて、友だちは半年で542人まで増加")

# ================= 2 アジェンダ =================
s = sl(2)
keep_header_only(s)
lead(s, "Webサイト改善の「先」にある、求職者を逃さないLINE連携の全体像")
AG = [("01", "なぜLINEを同時導入するのか", "サイト改修後も残る課題と、LINEの役割"),
      ("02", "求人業界でのLINE活用", "タウンワークの実績と勝因"),
      ("03", "自動車求人Naviでの活用イメージ", "導線・画面・自動応答のイメージ"),
      ("04", "効果・費用・導入スケジュール", "友だち獲得の試算・費用・進め方")]
for i, (no, t, d) in enumerate(AG):
    y = 4.9 + i * 2.85
    hline(s, 3.0, 24.5, y, LGRAY, 0.75)
    label(s, 3.0, y + 0.3, 3.0, 2.2, [(no, 30, True, NAVY, 0)], anchor=MSO_ANCHOR.MIDDLE)
    label(s, 6.3, y + 0.3, 18.2, 2.2, [(t, 19, True, NAVY, 3), (d, 12.5, False, GRAY, 0)], anchor=MSO_ANCHOR.MIDDLE)
hline(s, 3.0, 24.5, 4.9 + 4 * 2.85, LGRAY, 0.75)

# ================= 5 なぜ11月頃から（作り直し）=================
s = sl(5)
keep_header_only(s)
set_paras(shp(s, "TextBox 1"), [("なぜ4月を待たず、11月頃から初期構築を進めるのか", 0)])
lead(s, "4月から成果を出すために、11月頃から準備を始める")
arrow_line(s, 3.0, 8.6, 24.8, 8.6, NAVY, 2.5)
pts = [(6.6, "11月頃〜", NAVY, "①", "サイト制作段階から\nLINE連携を設計", "公開後の大きな改修を避けられる"),
       (13.75, "〜3月", NAVY, "②", "公開前に\n申請・実装・検証", "公開日からすぐ施策を動かせる"),
       (20.9, "4月〜", GREEN, "③", "4月から\nそのまま本格運用へ", "運用設計をやり直す必要がない")]
for cx, when, col, no, head, body in pts:
    label(s, cx - 3.4, 6.3, 6.8, 0.8, [(when, 13, True, GREEN_TX if col == GREEN else GRAY, 0)], align=PP_ALIGN.CENTER,
          anchor=MSO_ANCHOR.MIDDLE)
    shape(s, MSO_SHAPE.OVAL, cx - 0.8, 7.8, 1.6, 1.6, [(no, 18, True, WHITE, 0)], fill=col, margins=(0, 0, 0, 0))
    hl = head.split("\n")
    label(s, cx - 3.4, 9.9, 6.8, 1.9, [(hl[0], 15, True, NAVY, 0), (hl[1], 15, True, NAVY, 0)], align=PP_ALIGN.CENTER,
          anchor=MSO_ANCHOR.TOP)
    label(s, cx - 3.5, 11.9, 7.0, 0.9, [(body, 11.5, False, INK, 0)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.TOP)
conclusion(s, 14.3, "初期構築から運用・改善まで一貫して支援することで、サイト公開後もスムーズに本格運用へ移行", 14)

# ================= 18 概算お見積もり（作り直し）=================
s = sl(18)
keep_header_only(s)
set_paras(shp(s, "TextBox 1"), [("概算お見積もり", 0)])
lead(s, "LINE運用コンサルの固定費を基本に、必要な施策だけ追加")
label(s, 2.2, 4.4, 23.1, 0.8, [("基本固定費（LINE運用コンサル）", 14, True, NAVY, 0)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
rich(s, 2.2, 5.3, 11.4, 3.6, [[("初期", 16, True, GRAY)], [("20", 54, True, NAVY), ("万円", 24, True, NAVY)],
                               [("11月頃〜の設計・初期構築", 12, False, GRAY)]], align=PP_ALIGN.CENTER)
vline(s, 13.75, 5.6, 8.7, LGRAY, 1.0)
rich(s, 13.9, 5.3, 11.4, 3.6, [[("4月以降", 16, True, GREEN_TX)], [("月額 ", 24, True, NAVY), ("20", 54, True, NAVY), ("万円", 24, True, NAVY)],
                                [("配信・分析・改善・毎月のレポート", 12, False, GRAY)]], align=PP_ALIGN.CENTER)
label(s, 2.2, 9.1, 23.1, 0.8, [("＋ LINE公式アカウント利用料（開設3万円〜／月額0.5万円〜・配信通数で変動）　　＋ 必要に応じて追加施策", 12, False, INK, 0)],
      align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
label(s, 2.2, 10.5, 12.0, 0.7, [("オプション（必要に応じて）", 13, True, NAVY, 0)], anchor=MSO_ANCHOR.MIDDLE)
label(s, 17.3, 10.5, 3.8, 0.7, [("初期", 10.5, True, GRAY, 0)], align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
label(s, 21.5, 10.5, 3.8, 0.7, [("月額", 10.5, True, GRAY, 0)], align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
opts = [("離脱防止ポップアップ", "ページを離れる方をLINEへ誘導", "1.5万円", "3万円"),
        ("サンクスLINE誘導", "応募完了後にLINEへ誘導", "15万円", "3万円〜"),
        ("LINE Profile+／LINE連携開発", "登録の自動入力・会員連携", "別途お見積もり", None)]
y = 11.25
for name, desc, a, b in opts:
    hline(s, 2.2, 25.3, y, LGRAY, 0.75)
    label(s, 2.2, y + 0.1, 8.6, 1.0, [(name, 13.5, True, INK, 0)], anchor=MSO_ANCHOR.MIDDLE)
    label(s, 10.8, y + 0.1, 6.6, 1.0, [(desc, 11, False, GRAY, 0)], anchor=MSO_ANCHOR.MIDDLE)
    if b is None:
        label(s, 17.3, y + 0.1, 8.0, 1.0, [(a, 13.5, True, NAVY, 0)], align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    else:
        label(s, 17.3, y + 0.1, 3.8, 1.0, [(a, 13.5, True, NAVY, 0)], align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
        label(s, 21.5, y + 0.1, 3.8, 1.0, [(b, 13.5, True, NAVY, 0)], align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    y += 1.2
hline(s, 2.2, 25.3, y, LGRAY, 0.75)
label(s, 2.2, y + 0.6, 23.1, 0.8, [("体制：弊社＝設計・運用・制作・毎月のレポート／貴社＝求人情報・面談体制の共有、サイト制作会社様との連携", 9.5, False, GRAY, 0)])

# ================= 19 スケジュール（作り直し）=================
s = sl(19)
keep_header_only(s)
set_paras(shp(s, "TextBox 1"), [("スケジュール｜11月頃から初期構築、4月から本格運用", 0)])
lead(s, "11月〜3月で設計・構築し、4月のサイト公開と同時に運用へ")
TX, MW, EX = 2.2, 2.5, 25.3
GX = TX + 5 * MW
shape(s, MSO_SHAPE.PENTAGON, TX, 4.4, GX - TX + 0.3, 0.85, [("11月〜3月｜設計・構築", 12.5, True, WHITE, 0)], fill=NAVY, adj=0.5)
shape(s, MSO_SHAPE.CHEVRON, GX - 0.1, 4.4, EX - GX + 0.1, 0.85, [("4月〜｜運用・改善", 12.5, True, WHITE, 0)], fill=GREEN, adj=0.5,
      margins=(0.6, 0.05, 0.3, 0.05))
for i, m in enumerate(["11月", "12月", "1月", "2月", "3月"]):
    label(s, TX + i * MW, 5.35, MW, 0.6, [(m, 11.5, True, INK, 0)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
label(s, GX, 5.35, EX - GX, 0.6, [("4月〜", 11.5, True, GREEN_TX, 0)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
RH, GAP = 1.05, 0.15
rows1 = [("設計", "アカウント・リッチメニュー・配信シナリオ・計測", NAVY, None, WHITE),
         ("連携仕様", "サイト制作会社と、LINE連携の仕様を調整", NAVY, None, WHITE),
         ("実装・検証", "開発・実装・公開前テスト", PALE, NAVY, NAVY)]
rows2 = [("配信", "ステップ配信・企画配信・応募への追客", GREEN, None, WHITE),
         ("分析", "KPI計測・月次レポート", GREEN_BG, GREEN, GREEN_TX),
         ("改善", "導線・配信・セグメントを改善", GREEN_BG, GREEN, GREEN_TX)]
y = 6.15
bottom = y + 3 * (RH + GAP) + 1.2 + 3 * (RH + GAP)
for i in range(1, 5):
    vline(s, TX + i * MW, 5.45, bottom, "E7E7E7", 0.5)


def bar(x1, x2, yy, name, detail, fill, line, tc):
    rect(s, x1, yy, x2 - x1, RH, fill, line=line, lw=1.0)
    rich(s, x1 + 0.25, yy, x2 - x1 - 0.4, RH, [[(name + "　", 12.5, True, tc), (detail, 11, False, tc)]])


for name, d, f, l, tc in rows1:
    bar(TX + 0.05, GX - 0.05, y, name, d, f, l, tc)
    y += RH + GAP
my = y
shape(s, MSO_SHAPE.DIAMOND, GX - 0.45, my + 0.15, 0.9, 0.9, fill=GREEN)
arrow_line(s, TX + 0.4, my + 0.6, GX - 0.6, my + 0.6, NAVY, 1.5, dash=True)
rich(s, GX + 0.6, my, EX - GX - 0.6, 1.2, [[("サイト公開　／　LINE本格稼働", 14, True, GREEN_TX)]])
y = my + 1.2
for name, d, f, l, tc in rows2:
    bar(GX + 0.05, EX - 0.05, y, name, d, f, l, tc)
    y += RH + GAP
vline(s, GX, 5.35, bottom, GREEN, 2.5)
conclusion(s, bottom + 0.3, "初期構築で設計した内容を、そのまま運用へ", 16)
label(s, 2.2, bottom + 1.85, 23.1, 1.0, [
    ("※11月頃〜3月は設計・構築が中心。必要に応じて一部施策の先行稼働も検討可能（先行稼働分はシミュレーションに含めない）／LINE Profile+の審査期間は申請内容により変動", 8.5, False, GRAY, 0)])

# ================= Chapter表紙 =================
cover = sl(1)
cov_pics = [sh for sh in cover.shapes if sh.name in ("図 8", "図 9")]  # 表紙の上部（紺帯を隠す白）とロゴ
CH = [("01", "なぜLINEを同時導入するのか", "サイトリニューアルだけでは残る「検討層・離脱層」をLINEでストック"),
      ("02", "求人業界でのLINE活用", "大手求人サービスの事例から、自動車求人Naviで使える型を整理"),
      ("03", "自動車求人Naviでの活用イメージ", "サイトとLINEをつなぎ、会員登録から応募までの導線を設計"),
      ("04", "効果・費用・導入スケジュール", "11月頃から初期構築を進め、4月から本格運用へ")]
chap = []
for k, (no, title, sub) in enumerate(CH):
    c = prs.slides.add_slide(cover.slide_layout)
    for ph in list(c.placeholders):
        ph._element.getparent().remove(ph._element)
    for p in cov_pics:
        c.shapes.add_picture(io.BytesIO(p.image.blob), p.left, p.top, p.width, p.height)
    label(c, 15.5, 10.6, 10.5, 6.6, [(no, 170, True, PALE, 0)], align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.BOTTOM)
    rect(c, 2.6, 6.2, 0.28, 5.4, NAVY)
    label(c, 3.4, 6.1, 14.0, 0.9, [(f"Chapter {int(no)}", 18, True, GREEN_TX, 0)], anchor=MSO_ANCHOR.MIDDLE)
    label(c, 3.4, 7.1, 20.0, 1.9, [(title, 32, True, NAVY, 0)], anchor=MSO_ANCHOR.MIDDLE)
    hline(c, 3.5, 15.5, 9.4, NAVY, 2.0)
    label(c, 3.4, 9.75, 21.0, 1.0, [(sub, 14, False, INK, 0)], anchor=MSO_ANCHOR.MIDDLE)
    for j in range(4):
        rect(c, 3.5 + j * 2.9, 13.4, 2.6, 0.2, NAVY if j == k else LGRAY)
        label(c, 3.5 + j * 2.9, 13.7, 2.6, 0.6, [(f"0{j + 1}", 9.5, j == k, NAVY if j == k else GRAY, 0)], anchor=MSO_ANCHOR.TOP)
    chap.append(c)

# ================= 並べ替え =================
order = [sl(1), sl(2), chap[0], sl(3), sl(9), sl(4), sl(10), sl(11), chap[1], sl(6), sl(7), sl(8),
         chap[2], sl(12), sl(13), sl(14), sl(15), chap[3], sl(16), sl(17), sl(18), sl(5), sl(19)]
lst = prs.slides._sldIdLst
want = [o.slide_id for o in order]
elems = {int(e.get("id")): e for e in lst}
assert sorted(want) == sorted(elems)
for e in list(lst):
    lst.remove(e)
for sid in want:
    lst.append(elems[sid])

# 字体メイリオ・最小8pt
for slide in prs.slides:
    for tag in ("rPr", "endParaRPr"):
        for rpr in slide.shapes._spTree.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}" + tag):
            if rpr.get("sz") and int(rpr.get("sz")) < 800:
                rpr.set("sz", "800")
prs.save(str(OUT))
print("saved:", OUT.name, len(prs.slides), "枚")
