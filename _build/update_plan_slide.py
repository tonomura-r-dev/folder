# -*- coding: utf-8 -*-
"""Slide 23を田苑酒造型の費用対効果スライドに差し替え／Slide 11から広告費ベースの記載を除去"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
TARGET = ROOT / "20260804_株式会社ビイサイドプランニング御中_公式LINEのご提案v02.pptx"

TNAVY = "002060"
NAVY = "1F285A"
RED = "C00000"
INK = "333333"
MUT = "808080"
WHITE = "FFFFFF"
PALE = "F4F7FF"
GREY = "F2F2F2"
BORDER = "D9D9D9"

prs = Presentation(str(TARGET))
slides = list(prs.slides)


def set_font(run, size, bold=None, color=INK, name="メイリオ"):
    f = run.font
    f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    f.name = name
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        e = rPr.find(qn(tag))
        if e is None:
            e = rPr.makeelement(qn(tag), {})
            rPr.append(e)
        e.set("typeface", name)
    f.color.rgb = RGBColor.from_string(color)


ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
ANCH = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}


def reset_tf(tf):
    for para in tf.paragraphs[1:]:
        para._p.getparent().remove(para._p)
    for r in list(tf.paragraphs[0].runs):
        r._r.getparent().remove(r._r)


def put_text(tf, paras, anchor="t", ml=0.06, mr=0.06, mt=0.02, mb=0.02, wrap=True):
    reset_tf(tf)
    tf.word_wrap = wrap
    tf.margin_left = Inches(ml); tf.margin_right = Inches(mr)
    tf.margin_top = Inches(mt); tf.margin_bottom = Inches(mb)
    tf.vertical_anchor = ANCH[anchor]
    first = True
    for p in paras:
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        para.alignment = ALIGN[p.get("align", "l")]
        if p.get("sa") is not None:
            para.space_after = Pt(p["sa"])
        if p.get("ls") is not None:
            para.line_spacing = p["ls"]
        for (t, sz, b, c) in p["runs"]:
            r = para.add_run(); r.text = t
            set_font(r, sz, b, c)
    return tf


def add_text(slide, x, y, w, h, paras, anchor="t", **kw):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    put_text(box.text_frame, paras, anchor=anchor, **kw)
    return box


def add_box(slide, x, y, w, h, fill=None, line=None, lw=1.0, shape=MSO_SHAPE.RECTANGLE):
    sp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = RGBColor.from_string(line); sp.line.width = Pt(lw)
    sp.shadow.inherit = False
    return sp


def clear_slide(slide):
    spTree = slide.shapes._spTree
    for el in list(spTree):
        if el.tag.split("}")[-1] in ("sp", "cxnSp", "pic", "graphicFrame", "grpSp"):
            spTree.remove(el)


def kv(slide, x, y, label, value, lw=1.33, vw=1.61, h=0.39, vsz=10):
    """田苑型の「ラベル｜内容」タグ"""
    b1 = add_box(slide, x, y, lw, h, fill=NAVY)
    put_text(b1.text_frame, [{"runs": [(label, 11, True, WHITE)], "align": "c"}],
             anchor="m", ml=0.02, mr=0.02, mt=0.0, mb=0.0)
    b2 = add_box(slide, x + lw + 0.04, y, vw, h, fill=PALE)
    put_text(b2.text_frame, [{"runs": [(value, vsz, None, INK)], "align": "c", "ls": 1.0}],
             anchor="m", ml=0.04, mr=0.04, mt=0.0, mb=0.0)
    return b1, b2


# ================= Slide 23: 想定の費用対効果（田苑型） =================
s = slides[22]
clear_slide(s)
add_text(s, 0.60, 0.15, 7.12, 0.40,
         [{"runs": [("想定の費用対効果", 16, True, TNAVY)]}], anchor="m", ml=0.0, mr=0.0)
add_text(s, 0.42, 0.78, 10.0, 0.40,
         [{"runs": [("コンサル運営プラン（企画投稿5本／ステップ配信／リッチメニュー運用）", 13, True, INK)]}],
         anchor="m", ml=0.0, mr=0.0)

# --- 3ブロックの見出しと金額 ---
heads = [(0.29, "初期"), (3.65, "月次"), (7.38, "6か月後")]
for x, lab in heads:
    hb = add_box(s, x, 1.78, 3.17, 0.45, fill=NAVY)
    put_text(hb.text_frame, [{"runs": [(lab, 12, True, WHITE)], "align": "c"}],
             anchor="m", ml=0.0, mr=0.0, mt=0.0, mb=0.0)
for x, amt in [(0.29, "23.5万円"), (3.65, "25.98万円")]:
    ab = add_box(s, x, 2.30, 3.17, 0.95, fill=WHITE, line=NAVY, lw=1.5)
    put_text(ab.text_frame, [{"runs": [(amt, 28, True, NAVY)], "align": "c"}],
             anchor="m", ml=0.0, mr=0.0, mt=0.0, mb=0.0)

# --- 初期：実施内容 ---
add_text(s, 0.49, 3.32, 2.77, 0.30,
         [{"runs": [("＜実施内容＞", 12, True, NAVY)]}], anchor="m", ml=0.0, mr=0.0)
init_items = [
    ("プロフ設定", "動線設計・整備"),
    ("タグ設定", "アンケート／セグメント設計"),
    ("ツール導入", "Lstep想定"),
    ("外部ツール", "離脱防止ポップアップ"),
    ("ステップ構築", "Day0〜14の6ステップ"),
]
for i, (lab, val) in enumerate(init_items):
    kv(s, 0.36, 3.66 + i * 0.47, lab, val)

# --- 月次：実施内容 ---
add_text(s, 3.83, 3.32, 2.77, 0.30,
         [{"runs": [("＜実施内容＞", 12, True, NAVY)]}], anchor="m", ml=0.0, mr=0.0)
kv(s, 3.73, 3.66, "企画投稿", "5回／月")
b1 = add_box(s, 3.73, 4.13, 1.33, 1.49, fill=NAVY)
put_text(b1.text_frame, [{"runs": [("修正・改善", 11, True, WHITE)], "align": "c"}],
         anchor="m", ml=0.02, mr=0.02, mt=0.0, mb=0.0)
b2 = add_box(s, 5.10, 4.13, 1.61, 1.49, fill=PALE)
put_text(b2.text_frame,
         [{"runs": [(t, 10, None, INK)], "ls": 1.05, "sa": 2} for t in
          ["・配信セグメント", "・ステップ配信", "・リッチメニュー", "・離脱防止バナー"]],
         anchor="m", ml=0.10, mr=0.06, mt=0.04, mb=0.04)
kv(s, 3.73, 5.69, "レポート", "1回／月")
kv(s, 3.73, 6.16, "定例会", "1回／月")

# --- 6か月後：着地見込み ---
res_items = [
    ("友だち数", "2,601人\n（アクティブ約1,849人）"),
    ("メッセ数", "15,000〜18,000通／月"),
    ("コスト", "259,800円／月"),
    ("CV", "135件（半年累計）"),
    ("CPA", "13,287円"),
]
for i, (lab, val) in enumerate(res_items):
    y = 2.30 + i * 0.47
    b1 = add_box(s, 7.53, y, 1.33, 0.39, fill=NAVY)
    put_text(b1.text_frame, [{"runs": [(lab, 11, True, WHITE)], "align": "c"}],
             anchor="m", ml=0.02, mr=0.02, mt=0.0, mb=0.0)
    b2 = add_box(s, 8.90, y, 1.53, 0.39, fill=PALE)
    lines = val.split("\n")
    put_text(b2.text_frame,
             [{"runs": [(t, 9.5 if len(lines) == 1 else 8.5, None, INK)], "align": "c", "ls": 1.0}
              for t in lines],
             anchor="m", ml=0.03, mr=0.03, mt=0.0, mb=0.0)

# --- 内訳の注記 ---
add_text(s, 7.38, 4.72, 3.17, 0.90,
         [{"runs": [("＜費用内訳＞", 11, True, NAVY)], "sa": 3},
          {"runs": [("初期：構築20万＋離脱防止1.5万＋ツール2万", 9, None, INK)], "ls": 1.15, "sa": 2},
          {"runs": [("月次：コンサル20万＋ツール2.98万＋アカウント1.5万＋離脱防止1.5万", 9, None, INK)], "ls": 1.15}],
         anchor="t", ml=0.0, mr=0.0)

# --- 成果報酬プランの併記 ---
opt = add_box(s, 0.36, 6.68, 6.35, 0.52, fill=WHITE, line=BORDER, lw=1.0)
put_text(opt.text_frame,
         [{"runs": [("［成果報酬プランの場合］", 10.5, True, NAVY),
                    ("　初期23.5万円＋月次5.98万円＋成果報酬（事前申込7,500円／来場9,000円）＝半年171.0万円・CPA 12,665円", 10, None, INK)], "ls": 1.1}],
         anchor="m", ml=0.14, mr=0.10, mt=0.03, mb=0.03)
add_text(s, 7.38, 6.68, 3.17, 0.52,
         [{"runs": [("※CV＝フェア事前申込＋来場引き上げの合計", 9, None, MUT)], "ls": 1.1}],
         anchor="m", ml=0.0, mr=0.0)

# ================= Slide 11: 広告費ベースの記載を除去（位置で判定） =================
s = slides[10]
for sh in s.shapes:
    if not sh.has_text_frame:
        continue
    t = sh.text_frame.text
    top_in = sh.top / 914400
    if "参加率" not in t and "来場" not in t and "※" not in t:
        continue
    if top_in < 1.2 and "設計" in t:
        # リード文
        put_text(sh.text_frame,
                 [{"runs": [("「予約後の接点」をLINEで設計し、予約から来場への参加率を60%→77%（＋17pt）へ引き上げます。", 14, True, INK)], "ls": 1.1}],
                 anchor="m", ml=0.0, mr=0.0)
    elif top_in > 5.0 and "目標効果" in t:
        # 目標効果ボックス
        put_text(sh.text_frame,
                 [{"runs": [("目標効果", 12, True, NAVY)], "align": "c", "sa": 4},
                  {"runs": [("予約→来場の参加率　", 13, None, INK), ("60%", 15, True, INK),
                            (" → ", 13, None, MUT), ("77%", 24, True, NAVY),
                            ("　（＋17pt）", 12, None, INK)], "align": "c"}],
                 anchor="m", ml=0.0, mr=0.0, mt=0.0, mb=0.0)
    elif top_in > 6.4 and "※" in t:
        # 脚注
        put_text(sh.text_frame,
                 [{"runs": [("※参加率はMeta広告シミュレーション（当社作成）の想定値と統一。提案後の数値は本施策の実行を前提とした目標値です。", 10, None, MUT)]}],
                 ml=0.0, mr=0.0)

prs.save(str(TARGET))
print("saved:", TARGET)
