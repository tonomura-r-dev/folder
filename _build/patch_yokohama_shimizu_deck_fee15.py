# -*- coding: utf-8 -*-
"""横浜シミズ：運用の月額を15万円にしたSIM（殿村さんのPC保存版）に合わせて、資料の費用と「月額20万円以内」の文言を直す（2026-10-20 殿村さん「取っ払って」）。

    python -m _build.patch_yokohama_shimizu_deck_fee15 <資料.pptx> [出力.pptx]

  S14：タイトル「月額20万円以内」→「月額215,000円〜」／リードの「想定予算の月額20万円以内に収めています。」を外す／
       運用費 100,000円→150,000円、合計 165,000円〜→215,000円〜、合計の行の「（予算：月額20万円以内）」を外す
  S15：まとめの「月額20万円以内（ツール費込み）で始め、」→「月額215,000円〜（ツール費込み）で始め、」
  S2の「想定予算 月額20万円（ツール費込み）」は先方からいただいた情報なので残す。出力を省くと上書き。
"""
import sys

from pptx import Presentation
from pptx.oxml.ns import qn


def replace_in(frame_or_cell, old, new):
    """runの中の文字を置き換える（書式はそのまま）。見つからなければエラー"""
    hit = 0
    for p in frame_or_cell.text_frame.paragraphs:
        for r in p.runs:
            if old in r.text:
                r.text = r.text.replace(old, new)
                hit += 1
    assert hit == 1, (old, hit)


def by_id(s, sid):
    return next(sh for sh in s.shapes if sh.shape_id == sid)


def main(src, out):
    prs = Presentation(src)
    slides = list(prs.slides)
    s14, s15 = slides[13], slides[14]
    replace_in(by_id(s14, 2), "月額20万円以内（ツール費込み）", "月額215,000円〜（ツール費込み）")
    replace_in(by_id(s14, 3), "165,000円〜（税抜）。想定予算の月額20万円以内に収めています。", "215,000円〜（税抜）。")
    rows = by_id(s14, 21).table.rows
    assert rows[1].cells[0].text.startswith("DYMの運用費")
    replace_in(rows[1].cells[1], "100,000円", "150,000円")
    replace_in(rows[5].cells[0], "合計（予算：月額20万円以内）", "合計")
    replace_in(rows[5].cells[1], "165,000円〜", "215,000円〜")
    replace_in(by_id(s15, 28), "月額20万円以内（ツール費込み）で始め、", "月額215,000円〜（ツール費込み）で始め、")
    prs.save(out)
    print("saved:", out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else sys.argv[1])
