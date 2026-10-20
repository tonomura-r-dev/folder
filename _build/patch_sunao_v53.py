# -*- coding: utf-8 -*-
"""すなお ver5.2 → ver5.3（2026-10-20 殿村さん「go」）。

    python -m _build.patch_sunao_v53 <ver5.2.pptx> <ver5.3.pptx>

  1. 業界STUDY・訴求軸・要件定義・本質課題の4枚を作り直す（patch_sunao_v53_kaizen）
  2. ニーズ調査（シーズナリティ）を外す（殿村さん「情報量なさすぎ。入れないほうがええか」。MDRTへの差し替えも「配信は定期的なので、2〜4月に厚くとはならない」で不採用）
  ほかのページは触らない。想定動線（ユーザー動線の図）は殿村さんの版のファイル待ちのため、ここでは触らない。
"""
import sys

from pptx import Presentation

from _build import patch_sunao_v53_kaizen as kaizen


def drop_slide(prs, title):
    hit = [s for s in prs.slides if kaizen.title_of(s) is not None and kaizen.title_of(s).text_frame.text.strip() == title]
    assert len(hit) == 1, (title, len(hit))
    lst = prs.slides._sldIdLst
    for e in list(lst):
        if prs.part.related_part(e.rId) is hit[0].part:
            prs.part.drop_rel(e.rId)
            lst.remove(e)


def main(src, out):
    kaizen.main(src, out)
    prs = Presentation(out)
    drop_slide(prs, "ニーズ調査（シーズナリティ）")
    prs.save(out)
    print("saved:", out, len(prs.slides), "枚")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
