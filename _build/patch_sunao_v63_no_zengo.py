# -*- coding: utf-8 -*-
"""すなお ver6.2 → ver6.3：前後検索を使わないことにしたので、その名残を外す（2026-10-21 殿村さん）。

    python -m _build.patch_sunao_v63_no_zengo <ver6.2.pptx> <ver6.3.pptx>

  ①前後検索の3枚（9〜11枚目・非表示）を消す（62枚 → 59枚）
  ②「友だち追加直後のステップ配信（5通・案）」の注記から「※前後検索のデータを取得後、内容・順番を見直す」の行を外す
  ほかのページ・文言は変えない。
"""
import sys

from pptx import Presentation

from _build.pptx_parts import drop_slide

NOTE_LINE = "※前後検索のデータを取得後、内容・順番を見直す"


def main(src, out):
    prs = Presentation(src)
    slides = list(prs.slides)
    assert len(slides) == 62
    zengo = slides[8:11]
    for s in zengo:
        assert s._element.get("show") == "0"
        assert any(sh.has_text_frame and "前後検索" in sh.text_frame.text for sh in s.shapes)
    step = slides[33]
    note = next(sh for sh in step.shapes if sh.name == "注記")
    ps = [p for p in note.text_frame.paragraphs if p.text == NOTE_LINE]
    assert len(ps) == 1
    ps[0]._p.getparent().remove(ps[0]._p)
    for s in zengo:
        drop_slide(prs, s)
    assert len(prs.slides) == 59
    prs.save(out)
    print("saved:", out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
