# -*- coding: utf-8 -*-
"""タナカホーム v4（殿村さんのPC保存版・41枚）→ v5：前の方に足した汎用資料のページに見出しの飾りを足し、章の区切りページを作る（2026-10-28 殿村さん）。

    python -m _build.patch_tanaka_v5_header_chapters <v4（PC保存版）.pptx> <v5.pptx>

  ・2〜10ページ（汎用資料から足したページ。レイアウト「2_タイトルとコンテンツ」）には、左上の紺の四角と、その下の線が無い。
    11ページ以降のレイアウト「4_タイトルとコンテンツ」にある2つ（紺の四角・線）を、2〜5・7〜10ページの一番後ろに写す。6ページ（章の区切り）は除く
  ・6ページ「アジェンダ 01 施策背景」と同じ形で、章の区切りページを作る。章は11ページ「提案の流れ」の4つに合わせる
      02 現状と打ち手（現状と打ち手の前）／03 課題（課題の前）／04 解決策（解決策の前）／05 今後の展開（今後の展開の前）
  ほかのページは変えない。
"""
import copy
import sys

from pptx import Presentation
from pptx.oxml.ns import qn

from _build.patch_tanaka_v5 import copy_slide

CHAPTERS = [("02", "現状と打ち手", "現状と打ち手："), ("03", "課題", "課題："),
            ("04", "解決策", "解決策：資料請求以前から"), ("05", "今後の展開", "今後の展開：")]


def title_of(s):
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            return sh.text_frame.text
    return ""


def add_header_deco(slide, deco):
    """レイアウトの紺の四角・線を、スライドの一番後ろ（ほかの図形の下）に入れる"""
    tree = slide.shapes._spTree
    ids = [int(e.get("id")) for e in tree.iter(qn("p:cNvPr"))]
    nid = max(ids + [1]) + 1
    for k, el in enumerate(deco):
        e = copy.deepcopy(el)
        e.find(".//" + qn("p:cNvPr")).set("id", str(nid + k))
        tree.insert(2 + k, e)                            # nvGrpSpPr・grpSpPr の直後＝一番後ろ


def main(src, out):
    prs = Presentation(src)
    slides = list(prs.slides)
    assert len(slides) == 41
    lay4 = next(l for l in prs.slide_layouts if l.name == "4_タイトルとコンテンツ")
    deco = [sh._element for sh in lay4.shapes if sh.name in ("正方形/長方形 6", "直線コネクタ 4")]
    assert len(deco) == 2
    for i in (1, 2, 3, 4, 6, 7, 8, 9):                   # 2〜5・7〜10ページ（6ページは章の区切りなので除く）
        assert slides[i].slide_layout.name == "2_タイトルとコンテンツ", i + 1
        add_header_deco(slides[i], deco)

    chap = slides[5]
    assert [sh.text_frame.text for sh in chap.shapes if sh.has_text_frame][:3] == ["アジェンダ", "01", "施策背景"]
    lst = prs.slides._sldIdLst
    for num, name, head in CHAPTERS:
        target = next(s for s in prs.slides if title_of(s).startswith(head))
        new = copy_slide(prs, chap)
        texts = [sh for sh in new.shapes if sh.has_text_frame]
        texts[1].text_frame.paragraphs[0].runs[0].text = num
        texts[2].text_frame.paragraphs[0].runs[0].text = name
        el = list(lst)[-1]
        lst.remove(el)
        tid = next(e for e in lst if int(e.get("id")) == target.slide_id)
        tid.addprevious(el)
    assert len(prs.slides) == 45
    prs.save(out)
    print("saved:", out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
