# -*- coding: utf-8 -*-
"""Slide 13（ツナグバの直後）を、同業＋来場型の運用事例スライドに作り替える"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
TARGET = ROOT / "20260804_株式会社ビイサイドプランニング御中_公式LINEのご提案v05.pptx"

TNAVY = "002060"; NAVY = "1F285A"; INK = "333333"
MUT = "808080"; WHITE = "FFFFFF"; PALE = "F4F7FF"; BORDER = "D9D9D9"

prs = Presentation(str(TARGET))
slides = list(prs.slides)


def set_font(run, size, bold=None, color=INK, name="メイリオ"):
    f = run.font; f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    f.name = name
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        e = rPr.find(qn(tag))
        if e is None:
            e = rPr.makeelement(qn(tag), {}); rPr.append(e)
        e.set("typeface", name)
    f.color.rgb = RGBColor.from_string(color)


ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
ANCH = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}


def reset_tf(tf):
    for para in tf.paragraphs[1:]:
        para._p.getparent().remove(para._p)
    for r in list(tf.paragraphs[0].runs):
        r._r.getparent().remove(r._r)


def put_text(tf, paras, anchor="t", ml=0.06, mr=0.06, mt=0.03, mb=0.03, wrap=True):
    reset_tf(tf); tf.word_wrap = wrap
    tf.margin_left = Inches(ml); tf.margin_right = Inches(mr)
    tf.margin_top = Inches(mt); tf.margin_bottom = Inches(mb)
    tf.vertical_anchor = ANCH[anchor]
    first = True
    for p in paras:
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        para.alignment = ALIGN[p.get("align", "l")]
        if p.get("sa") is not None:
            para.space_after = Pt(p["sa"])
        if p.get("ls") is not None:
            para.line_spacing = p["ls"]
        for (t, sz, b, c) in p["runs"]:
            r = para.add_run(); r.text = t
            set_font(r, sz, b, c)


def add_text(slide, x, y, w, h, paras, anchor="t", **kw):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    put_text(box.text_frame, paras, anchor=anchor, **kw)
    return box


def add_box(slide, x, y, w, h, fill=None, line=None, lw=1.0, shape=MSO_SHAPE.RECTANGLE):
    sp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = RGBColor.from_string(line); sp.line.width = Pt(lw)
    sp.shadow.inherit = False
    return sp


def clear_slide(slide):
    spTree = slide.shapes._spTree
    for el in list(spTree):
        if el.tag.split("}")[-1] in ("sp", "cxnSp", "pic", "graphicFrame", "grpSp"):
            spTree.remove(el)


# ================= Slide 13：運用事例2社 =================
s = slides[12]
clear_slide(s)
add_text(s, 0.60, 0.13, 9.60, 0.45,
         [{"runs": [("参考：他社の運用事例（同業界／予約→来場型）", 20, True, TNAVY)]}],
         anchor="m", ml=0.0, mr=0.0)
add_text(s, 0.60, 0.68, 9.60, 0.55,
         [{"runs": [("「求職者にLINEが届くこと」と「予約者を来場につなげること」を、それぞれ実運用で確認しています。", 14, True, INK)], "ls": 1.05}],
         anchor="m", ml=0.0, mr=0.0)

cases = [
    (0.60, "同業界での実績", "日研トータルソーシング「e仕事」",
     "工場・製造業の人材派遣／求人サイト",
     [("友だち数", "約4,000人 → 40,000人超"),
      ("ブロック率", "業界比で「とても低い」"),
      ("CV地点", "応募・面接予約・Web登録")],
     "トーク上で属性に答えるだけで応募・面接予約が完了。"
     "「来場時の履歴書が不要」になり、応募の摩擦をLINE内で外している。",
     "出典：mobilus.co.jp（導入事例）"),
    (5.65, "「予約→来場」での実績", "ホットヨガスタジオLAVA",
     "ホットヨガスタジオ（全国チェーン）",
     [("CV地点", "体験レッスンの来場"),
      ("主な施策", "予約完了通知＋前日リマインド"),
      ("連携", "ID連携で予約・変更がトーク内完結")],
     "体験予約者へ「予約完了」と「前日のお知らせ」をLINE通知メッセージで配信。"
     "予約から当日来場までをLINEが並走する設計。",
     "出典：lava-intl.co.jp（LINE公式アカウント案内）"),
]
for x, role, name, biz, rows, body, src in cases:
    hb = add_box(s, x, 1.36, 4.55, 0.44, fill=NAVY)
    put_text(hb.text_frame, [{"runs": [(role, 12, True, WHITE)], "align": "c"}],
             anchor="m", ml=0.0, mr=0.0, mt=0.0, mb=0.0)
    add_text(s, x, 1.88, 4.55, 0.60,
             [{"runs": [(name, 14, True, NAVY)], "sa": 2},
              {"runs": [(biz, 10, None, MUT)]}], anchor="t", ml=0.04)
    for i, (k, v) in enumerate(rows):
        y = 2.56 + i * 0.52
        kb = add_box(s, x, y, 1.42, 0.44, fill=PALE)
        put_text(kb.text_frame, [{"runs": [(k, 10, True, NAVY)], "align": "c"}],
                 anchor="m", ml=0.03, mr=0.03, mt=0.0, mb=0.0)
        add_text(s, x + 1.52, y, 3.03, 0.44,
                 [{"runs": [(v, 11, True, INK)]}], anchor="m", ml=0.02)
    bb = add_box(s, x, 4.20, 4.55, 1.10, fill=WHITE, line=BORDER, lw=1.0)
    put_text(bb.text_frame, [{"runs": [(body, 10.5, None, INK)], "ls": 1.2}],
             anchor="m", ml=0.16, mr=0.12, mt=0.06, mb=0.06)
    add_text(s, x, 5.34, 4.55, 0.26,
             [{"runs": [(src, 8.5, None, MUT)]}], anchor="m", ml=0.02)

# 位置づけの明示（誇張しないための注記）
nb = add_box(s, 0.60, 5.72, 9.60, 0.72, fill=PALE)
put_text(nb.text_frame,
         [{"runs": [("本提案での位置づけ　", 11, True, NAVY),
                    ("上記2社は「同じ構造の運用が実在すること」を示す事例です。来場率の改善幅そのものは公表されていないため、"
                     "本提案の目標値77%は、当社の配信実績（予約→来場の引き上げ率 6〜7割→8割程度）を根拠としています。",
                     10, None, INK)], "ls": 1.15}],
         anchor="m", ml=0.18, mr=0.14, mt=0.05, mb=0.05)

bar = add_box(s, 0.60, 6.58, 9.60, 0.46, fill=NAVY)
put_text(bar.text_frame,
         [{"runs": [("求職者への到達は同業界で、予約から来場への接続は他業種で、それぞれ実証されている", 13, True, WHITE)], "align": "c"}],
         anchor="m", ml=0.0, mr=0.0, mt=0.0, mb=0.0)

prs.save(str(TARGET))
print("saved:", TARGET)
