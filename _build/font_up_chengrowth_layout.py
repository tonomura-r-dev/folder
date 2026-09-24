# -*- coding: utf-8 -*-
"""font_up_chengrowth.py の後に実行。文字拡大で溢れた枠を広げる（下の余白に逃がす）。
  python3 _build/font_up_chengrowth_layout.py <pptx>
"""
import sys
from pptx import Presentation
from pptx.util import Cm, Pt
from pptx.oxml.ns import qn

F = sys.argv[1]
prs = Presentation(F)

def by_id(slide, sid):
    def walk(shapes):
        for sh in shapes:
            if sh.shape_id == sid: return sh
            if sh.shape_type == 6:
                r = walk(sh.shapes)
                if r is not None: return r
    r = walk(slide.shapes)
    assert r is not None, sid
    return r

def geo(si, sid, top=None, h=None, left=None, w=None):
    sh = by_id(prs.slides[si - 1], sid)
    if top is not None: sh.top = Cm(top)
    if h is not None: sh.height = Cm(h)
    if left is not None: sh.left = Cm(left)
    if w is not None: sh.width = Cm(w)

# s3 LINE導入後の6枠を高く、矢印と帯を下げる
for sid in (14, 16, 18, 20, 22, 24): geo(3, sid, h=2.1)
for sid in (15, 17, 19, 21, 23): geo(3, sid, top=7.85)
geo(3, 25, top=10.0)

# s6 所見カード3枚を高く、帯を下げる
for sid, t in ((6, 4.6), (7, 7.7), (8, 10.8)): geo(6, sid, top=t, h=2.9)
geo(6, 9, top=14.0)

# s8 表が伸びた分、帯を下げる
geo(8, 6, top=11.3)

# s11 カード本文：手入力の改行を外して自然に折り返す（「の」だけが次行に落ちるのを防ぐ）
for sid in (17, 29, 44, 69):
    tf = by_id(prs.slides[10], sid).text_frame
    for br in tf._txBody.findall(".//" + qn("a:br")):
        br.getparent().remove(br)
    ps = tf.paragraphs
    if len(ps) > 1:
        first = ps[0]._p
        for p in ps[1:]:
            for r in p._p.findall(qn("a:r")):
                first.append(r)
            p._p.getparent().remove(p._p)

# s12 結論ボックスを高く
geo(12, 15, h=2.9)

# s15 「※中長期施策として展開予定」は元のサイズへ（②の行が折り返さないように）
for p in by_id(prs.slides[14], 34).text_frame.paragraphs:
    for r in p.runs:
        if "中長期" in r.text and r.font.size is not None:
            r.font.size = Pt(r.font.size.pt - 1)

# s16 費用・効果カードを高く
for sid in (5, 8): geo(16, sid, h=6.6)
geo(16, 7, h=5.4)
geo(16, 14, h=3.0)

# s17 事例カードを高く、帯を下げる
for sid in (5, 10): geo(17, sid, h=6.0)
for sid in (9, 14): geo(17, sid, h=2.5)
geo(17, 15, top=10.9)

# s20 Day本文を高く、帯を下げる
for sid in (6, 9, 12, 15): geo(20, sid, h=4.2)
geo(20, 16, top=9.9)

prs.save(F)
print("layout ok")

# --- 2回目の調整 ---
prs = Presentation(F)
for sid in (5, 8): geo(16, sid, h=7.4)
geo(16, 14, h=3.8)
# s20 Day本文は+1ptに留める（6cm幅の枠で1〜2文字だけ折り返すため）
for sid in (6, 9, 12, 15):
    for p in by_id(prs.slides[19], sid).text_frame.paragraphs:
        for r in p.runs:
            if r.font.size is not None and r.text.strip() and r.font.size.pt > 3:
                r.font.size = Pt(r.font.size.pt - 1)
prs.save(F)
print("layout2 ok")
