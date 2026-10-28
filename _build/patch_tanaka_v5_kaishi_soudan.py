# -*- coding: utf-8 -*-
"""タナカホーム v5：31ページ（スケジュール）の運用開始「1月から」は弊社の提案で、時期は相談できることを伝える（2026-10-28 殿村さん）。

    python -m _build.patch_tanaka_v5_kaishi_soudan <v5.pptx> <出力.pptx>

  ・リード「11月から設計・構築を進め、1月から運用を始める」→「11月から設計・構築を進め、1月からの運用開始をご提案（開始時期はご相談可能）」
  ・注記「※スケジュールは、ご契約の時期により前後する」→「※運用開始の時期は、貴社のご都合に合わせてご相談可能。スケジュールは、ご契約の時期により前後する」
  ほかは変えない。
"""
import sys

from pptx import Presentation

REP = [("11月から設計・構築を進め、1月から運用を始める", "11月から設計・構築を進め、1月からの運用開始をご提案（開始時期はご相談可能）"),
       ("※スケジュールは、ご契約の時期により前後する", "※運用開始の時期は、貴社のご都合に合わせてご相談可能。スケジュールは、ご契約の時期により前後する")]


def main(src, out):
    prs = Presentation(src)
    s = prs.slides[30]
    assert s.shapes[0].text_frame.text.startswith("スケジュール：")
    for old, new in REP:
        sh = next(x for x in s.shapes if x.has_text_frame and x.text_frame.text == old)
        p = sh.text_frame.paragraphs[0]
        p.runs[0].text = new
        for r in p.runs[1:]:
            r._r.getparent().remove(r._r)
    prs.save(out)
    print("saved:", out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
