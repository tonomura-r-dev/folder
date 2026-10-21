"""
10月号（正）のカード（見出し＋本文の角丸の箱）の高さを、中の文字量に合わせて詰める。
（2026-10-21 殿村さん指摘：枠の余白が多すぎる）
同じ行に並ぶカードは高さをそろえ、2段のときは下の段を上に詰める。
使い方: python3 _build/patch_trend_202610_fit.py <pptx>
"""
import sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.shapes import MSO_SHAPE

F = sys.argv[1]
prs = Presentation(F)
GAP = Inches(0.2)
PAD = Inches(0.14)


def text_w(s, pt):
    return sum((0.55 if ord(c) < 0x2000 else 1.0) for c in s) * pt / 72  # inch


def lines(s, pt, width_in):
    return max(1, -(-text_w(s, pt) // width_in))


def need_h(sh):
    p0, p1 = sh.text_frame.paragraphs
    tf = sh.text_frame
    usable = Emu(sh.width - tf.margin_left - tf.margin_right).inches - 0.05
    hp = p0.runs[0].font.size.pt
    bp = p1.runs[0].font.size.pt
    h = lines(p0.text, hp, usable) * hp * 1.3 / 72 + 6 / 72 + lines(p1.text, bp, usable) * bp * 1.15 * 1.25 / 72
    return Inches(h) + 2 * PAD


for slide in list(prs.slides)[2:]:
    cards = [sh for sh in slide.shapes
             if sh.shape_type == 1 and sh.has_text_frame
             and sh.auto_shape_type == MSO_SHAPE.ROUNDED_RECTANGLE
             and len(sh.text_frame.paragraphs) == 2 and sh.height >= Inches(1.0)]
    if not cards:
        continue
    rows = {}
    for c in cards:
        rows.setdefault(c.top, []).append(c)
    y = None
    for top in sorted(rows):
        row = rows[top]
        h = max(need_h(c) for c in row)
        new_top = top if y is None else y
        for c in row:
            c.top = new_top
            c.height = h
        y = new_top + h + GAP

prs.save(F)
print("saved", F)
