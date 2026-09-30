# -*- coding: utf-8 -*-
"""サンクスLINE簡易資料 ver1.1（殿村さんPC修正版）→ ver1.2
費用がS2とS5の2か所にあるので、S2の費用の枠を消し、S5に費用を残す（殿村さん指示）。
S2の枠にあった「APIツールと併用できる」は、S5の※注記に移す。

  python3 _build/patch_thanks_line_v12.py <ver1.1.pptx>
"""
import sys
from pathlib import Path

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Cm

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "20260930_サンクスLINEのご提案（簡易版）ver1.2.pptx"
prs = Presentation(sys.argv[1])
S = list(prs.slides)
DIV_Y, FOOT_Y = 3.86, 17.35


def by(s, n):
    return next(x for x in s.shapes if x.name == n)


def center(s):
    """本文（区切り線〜注記の間）を上下中央へ"""
    body = [sh for sh in s.shapes if Cm(DIV_Y + 0.1) <= sh.top < Cm(FOOT_Y - 0.1)]
    top = min(sh.top for sh in body)
    bottom = max(sh.top + sh.height for sh in body)
    dy = int((Cm(DIV_Y + 0.4) + Cm(FOOT_Y - 0.3)) / 2 - (top + bottom) / 2)
    for sh in body:
        sh.top = sh.top + dy


# S2：費用の枠を消す
s2 = S[1]
for n in ("Rounded Rectangle 15", "Rounded Rectangle 16", "TextBox 17", "Connector 18", "TextBox 19"):
    el = by(s2, n)._element
    el.getparent().remove(el)
center(s2)

# S5：※注記にAPI併用を足す
note = by(S[4], "TextBox 10")
ts = list(note._element.iter(qn("a:t")))
full = "".join(t.text or "" for t in ts)
ts[0].text = "※APIツール（Lステップ等）と併用可／完了画面を増やす場合は追加費用"
for t in ts[1:]:
    t.text = ""
for pp in note._element.iter(qn("a:p")):         # 中央揃え
    ppr = pp.find(qn("a:pPr"))
    if ppr is None:
        ppr = pp.makeelement(qn("a:pPr"), {}); pp.insert(0, ppr)
    ppr.set("algn", "ctr")

# S5：費用のチップを大きく（幅・高さ・文字）
s5 = S[4]
c1, c2, nt, band = (by(s5, n) for n in ("Rounded Rectangle 8", "Rounded Rectangle 9", "TextBox 10", "Rounded Rectangle 11"))
y = c1.top
CX0, CW, BW, GAP = 1.2, 25.12, 6.5, 0.3
x1 = CX0 + (CW - (BW * 2 + GAP)) / 2            # 2つの箱をまとめて左右中央に
for sh, x in ((c1, x1), (c2, x1 + BW + GAP)):
    sh.left, sh.width, sh.height = Cm(x), Cm(BW), Cm(1.3)
    for r in sh._element.iter(qn("a:rPr")):
        r.set("sz", "1600")
nt.left, nt.top, nt.width, nt.height = Cm(CX0), y + Cm(1.4), Cm(CW), Cm(0.7)   # 注記は箱の下に中央揃え
band.top = y + Cm(1.3) + Cm(1.1)
center(s5)

# S3・S4・S5：図形を左右中央に寄せる（横幅を0.88倍に縮め、左右の余白を広げる。矢印は大きさを変えず位置だけ寄せる）
MID, F = Cm(27.52) / 2, 0.88
for sl in (S[2], S[3], S[4]):
    for sh in sl.shapes:
        if not (Cm(DIV_Y + 0.1) <= sh.top < Cm(FOOT_Y - 0.1)):
            continue
        cx = sh.left + sh.width / 2
        new_cx = MID + (cx - MID) * F
        if "Arrow" not in sh.name:
            sh.width = int(sh.width * F)
        sh.left = int(new_cx - sh.width / 2)

# S4：結果の箱が狭く折り返すので、申込みの箱から1cm分を回す（右側の図形を1cm左へ）
s4, D1 = S[3], Cm(1.0)
g = lambda n: by(s4, n)
for n in ("Rounded Rectangle 4", "Rounded Rectangle 7", "Rounded Rectangle 13"):
    g(n).width = g(n).width - D1
for n in ("Right Arrow 8", "Rounded Rectangle 9", "Right Arrow 10", "Right Arrow 14", "Rounded Rectangle 15", "Right Arrow 16"):
    g(n).left = g(n).left - D1
for n in ("Rounded Rectangle 5", "Rounded Rectangle 11", "Rounded Rectangle 17"):
    g(n).left = g(n).left - D1
    g(n).width = g(n).width + D1

prs.save(str(OUT))
print("saved:", OUT.name)
