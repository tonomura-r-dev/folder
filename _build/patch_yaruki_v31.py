# -*- coding: utf-8 -*-
"""やる気スイッチ ver3.1（殿村さんPC保存版）を直して、正しいファイル名で保存する（2026-09-30）。

  S8   前後検索「子ども 習い事」を、チェングロウスのカーディーラー版と同じ作りに
       （左：色分け画像＋LINE追加の緑帯／右：凡例＋LINE追加率を上げるコンテンツ／下：登録後に体験予約へつなげるコンテンツ）
  S9   前後検索「習い事」を、チェングロウスの工場求人ナビ版と同じ作りに
       （左：凡例＋色分け画像／右：検索前・当日・後の3枚／下：配信の順番の結論）
  S10  文字を大きくした版に（ver3.1 は大きくする前の版だったため）
  ver2_5 で直したもの：S5「アフィリエイト」／S21 リード・注記／S24・S31 線の色／S38 古い注記の削除／
       CPO改善のご提案を【オプション】にして都度発注プランの後ろへ・カードの順番
  S39  【オプション】QA自動化（チャットボット）は削除（殿村さん：いらない）

  python3 _build/patch_yaruki_v31.py <ver3.1.pptx>
"""
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
_src = (ROOT / "_build/build_chengrowth_v11.py").read_text(encoding="utf-8")
_helpers = _src[_src.index("# ================= helpers"):_src.index("# ============================================================\n# 流用するページの素材")]
_head = _src[_src.index("import shutil"):_src.index("shutil.copyfile(SRC, OUT)")]
exec(_head)
exec(_helpers)
from copy import deepcopy
from pptx.enum.shapes import MSO_CONNECTOR

SRC = Path(sys.argv[1])
OUT = ROOT / "202609_株式会社やる気スイッチグループ御中_LINE公式アカウント運用のご提案_ver3.1.pptx"
prs = Presentation(str(SRC))
S = list(prs.slides)
_v24 = (ROOT / "_build/patch_yaruki_v24.py").read_text(encoding="utf-8")
exec(_v24[_v24.index("def by(s, name, nth=0):"):_v24.index("def insert_after(anchor, new):")])   # by / set_lines など
DGREEN, GRAY = "0B7A3B", "7F7F7F"
# 分類の色（チェングロウスの凡例と同じ配色）：塗り・文字
LEG = [("習い事の比較・検討", "D6E2F3", "2B5797"), ("子育て・学び", "FBE0CF", "C0561A"),
       ("家族のお出かけ・楽しみ", "E8DAF5", "7440A8")]
IMGW, IMGH = 1790, 1044


def rbox(slide, x, y, w, h, fill, lines, anchor="m", align="c", line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    """lines = [(文字, pt, 太字, 色), ...]"""
    sp = box(slide, x, y, w, h, fill=fill, line=line, shape=shape, radius=0.08)
    put_text(sp.text_frame, [one(t, sz, b, c, align=align, ls=1.15) for t, sz, b, c in lines], anchor=anchor,
             ml=0.15, mr=0.15, mt=0.08, mb=0.08)
    return sp


def link(slide, x1, y1, x2, y2, color="4472C4"):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Cm(x1), Cm(y1), Cm(x2), Cm(y2))
    c.line.color.rgb = RGBColor.from_string(color)
    c.line.width = Pt(1.5)
    ln = c.line._get_or_add_ln()
    ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle"}))


def keep_only(slide, names):
    drop(*[sh for sh in slide.shapes if sh.name not in names])


def lead1(slide, text):
    ld = by(slide, "Text 7")
    ld.top, ld.height = Cm(1.8), Cm(1.85)
    set_lines(ld, [text], sz=14)


def source(slide, kw):
    T(slide, 1.2, 17.45, 20.0, 0.5, [one(f"出典：LINEヤフー社提供の前後検索データ（検索起点：「{kw}」）", 8, None, MUT)],
      anchor="m", ml=0, mr=0)


# ============================================================
# S8 前後検索「子ども 習い事」（カーディーラー版と同じ作り）
# ============================================================
s = S[7]
keep_only(s, ("Google Shape;156;p7", "Text 7"))
lead1(s, "「子ども 習い事」の検索前後15日間。友だち追加の機会は、比較が集中する検索当日〜翌日")
IL, IT, IW = 1.2, 4.05, 14.2
IH = IW * IMGH / IMGW
s.shapes.add_picture(str(IMG / "yaruki_zengo_kodomo_color.png"), Cm(IL), Cm(IT), width=Cm(IW))
px = lambda p: IL + p / IMGW * IW
py = lambda p: IT + p / IMGH * IH
gx0, gx1 = px(917), px(1131)                      # 0日（検索起点）〜1日
band_ = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(gx0), Cm(IT), Cm(gx1 - gx0), Cm(IH - 0.1))
band_.fill.solid(); band_.fill.fore_color.rgb = RGBColor.from_string("06C755")
band_.fill._xPr.find(qn("a:solidFill"))[0].append(band_.fill._xPr.makeelement(qn("a:alpha"), {"val": "14000"}))
band_.line.fill.background(); band_.shadow.inherit = False
rbox(s, (gx0 + gx1) / 2 - 1.3, IT + IH + 0.05, 2.6, 0.62, "C6EFCE", [("LINE追加", 9.5, True, "00873C")],
     shape=MSO_SHAPE.RECTANGLE)
RX, RW = 15.9, 10.42
for i, (t, f, c) in enumerate(LEG):
    rbox(s, RX + (i % 2) * (RW / 2 + 0.05), 4.05 + (i // 2) * 0.72, RW / 2 - 0.05, 0.62, f, [(t, 9.5, True, c)])
T(s, RX + RW / 2 + 0.15, 4.77, RW / 2 - 0.15, 0.62, [one("色なし：その他", 9, None, MUT)], anchor="m", ml=0, mr=0)
rbox(s, RX, 5.7, RW, 0.62, "FFF2CC", [("LINE追加率を上げるコンテンツ（検索当日〜翌日）", 10.5, True, INK)])
CW2 = RW / 2 - 0.1
rbox(s, RX, 6.45, CW2, 2.55, NAVY, [("離脱防止で…", 10, False, WHITE), ("「習い事ランキング」", 11, True, WHITE),
                                   ("をLINEで見られます", 10.5, True, WHITE)])
rbox(s, RX + CW2 + 0.2, 6.45, CW2, 2.55, NAVY, [("この機会に", 10, False, WHITE), ("LINEで30秒", 11, True, WHITE),
                                              ("適性診断しませんか？", 10.5, True, WHITE)])
rbox(s, RX, 9.1, CW2, 1.95, "DDEBF7", [("当日の検索は「ランキング」", 10, True, NAVY), ("その場で比べる材料を示す", 10, False, NAVY)])
rbox(s, RX + CW2 + 0.2, 9.1, CW2, 1.95, "DDEBF7", [("電話なし・すぐ結果がわかる", 10, True, NAVY),
                                                  ("保護者の方も気軽に登録", 10, False, NAVY)])
link(s, (gx0 + gx1) / 2, py(700), RX - 0.1, 7.7)
BY = IT + IH + 0.85
rbox(s, 1.2, BY, 25.12, 0.62, "FFF2CC",
     [("登録後に体験予約へつなげるコンテンツ（翌日〜14日・「習い事」の検索後に多い「評判」「月謝」の確認に合わせる）", 10.5, True, INK)])
CARDS = [("1日目", "診断リマインド", "まだの方へ", NAVY),
         ("2日目", "診断結果", "タイプ別おすすめ", NAVY),
         ("4日目", "在籍生の声", "成長エピソード", NAVY),
         ("6日目", "体験のご案内①", "（CV）", GREEN),
         ("9日目", "伸びる力", "非認知能力・知育", NAVY),
         ("12日目", "よくある疑問", "月謝・送迎・振替", NAVY),
         ("14日目", "体験のご案内②", "（CV）", GREEN)]
n, gap = len(CARDS), 0.2
cw = (25.12 - gap * (n - 1)) / n
cy = BY + 0.75
for i, (d, t, sub, col) in enumerate(CARDS):
    rbox(s, 1.2 + i * (cw + gap), cy, cw, 17.2 - cy, col, [(d, 11, True, WHITE), (t, 10.5, True, WHITE), (sub, 9.5, False, WHITE)])
link(s, px(1450), py(1000), px(1560), BY - 0.05)
source(s, "子ども 習い事")

# ============================================================
# S9 前後検索「習い事」（工場求人ナビ版と同じ作り）
# ============================================================
s = S[8]
keep_only(s, ("Google Shape;156;p7", "Text 7"))
lead1(s, "「習い事」で探す保護者の動き方には、一定の型があります。")
L0, T0, W0 = CX0, 4.4, 15.8
chip_h, gap = 0.62, 0.2
cw = (W0 - gap * 2) / 3
for i, (t, f, c) in enumerate(LEG):
    rbox(s, L0 + i * (cw + gap), T0, cw, chip_h, f, [(t, 10, True, c)])
top = T0 + chip_h + 0.15
w = W0
h = w * IMGH / IMGW
s.shapes.add_picture(str(IMG / "yaruki_zengo_naraigoto_color.png"), Cm(L0 + (W0 - w) / 2), Cm(top), width=Cm(w))
gx0, gx1 = L0 + 921 / IMGW * w, L0 + 1134 / IMGW * w          # LINE追加：0日（検索起点）〜1日
band_ = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(gx0), Cm(top), Cm(gx1 - gx0), Cm(h - 0.1))
band_.fill.solid(); band_.fill.fore_color.rgb = RGBColor.from_string("06C755")
band_.fill._xPr.find(qn("a:solidFill"))[0].append(band_.fill._xPr.makeelement(qn("a:alpha"), {"val": "14000"}))
band_.line.fill.background(); band_.shadow.inherit = False
rbox(s, (gx0 + gx1) / 2 - 1.3, top + h + 0.05, 2.6, 0.62, "C6EFCE", [("LINE追加", 9.5, True, "00873C")],
     shape=MSO_SHAPE.RECTANGLE)
T(s, L0 + W0 - 5.5, top + h + 0.1, 5.5, 0.5, [one("色なし：大人の習い事・趣味、その他", 8.5, None, MUT, align="r")], anchor="m", ml=0, mr=0)
for i, (hd, body, col, fill) in enumerate([
        ("検索前（〜前日）：子どもの悩み・習い事の効果", "療育／集団行動が苦手な子供／子供 イライラする／9歳の壁／そろばん 効果", NAVY, PALE),
        ("検索当日：習い事の種類・ランキング", "習い事 ランキング／子供の習い事／体操教室／そろばん／スポーツ 習い事", ORANGE, PORANGE),
        ("検索後（翌日〜15日）：教室名・評判・月謝の確認", "忍者ナイン 評判／くもん 月謝／小学生 習い事 いくつ／幼児教室／小学校一年生", NAVY, PALE)]):
    card(s, 17.3, 4.4 + i * 3.4, 9.02, 3.2, hd, body, hcol=col, fill=fill, hsz=12, bsz=10.5)
band(s, 15.35, "登録した後の配信は「子どもの悩み → 習い事の比べ方 → 評判・月謝」の順番が最適", sz=14, h=1.3)
source(s, "習い事")

# ============================================================
# S10 前後検索のまとめ（項目別）：文字を大きくした版
# ============================================================
_s10 = _v24[_v24.index("# S10 前後検索のまとめ（項目別）"):_v24.index("# ============================================================\n# 並べ替え")]
_s10 = _s10.replace('''drop(*[sh for sh in s.shapes if sh.name not in ("Google Shape;156;p7", "Text 7")])''',
                    '''keep_only(s, ("Google Shape;156;p7", "Text 7"))''')
CAT = [("習い事の比較・検討", "3467B2", "E7EEF8"), ("子育て・学び", "D9661F", "FBEBDD"), ("家族のお出かけ・楽しみ", "8E4EC6", "F1E9F8")]
exec(_s10.split("\n", 1)[1])

# ============================================================
# ver2_5 で直したもの
# ============================================================
lab = by(S[4], "TextBox 5")                                   # S5：略さない
lab.left, lab.width = Cm(0.7), Cm(3.55)
set_lines(lab, ["広告・", "アフィリエイト"], sz=11)

s = S[20]                                                     # S21：リードを表に合わせる・注記
set_lines(by(s, "テキスト ボックス 39"), ["運用が育つにつれて体験予約・入会が増え、6か月目のCPA・CPOは1か月目より大きく下がります。",
                                         "（4か月目は固定費が上がるため、一時的に上がります）"])
ft = by_text(s, "※CPA＝")
if "サンクスLINE誘導" not in ft.text_frame.text:
    ts = list(ft._element.iter(qn("a:t")))
    ts[-1].text += "サンクスLINE誘導（オプション）の費用は含みません。"

for n_ in ("17", "27", "28", "29", "53", "67"):                # S24：ぼかした側の線を薄く
    by(S[23], "コネクタ: カギ線 " + n_).line.color.rgb = RGBColor.from_string("D9D9D9")
s = S[30]                                                     # S31：同上＋残す線は手前に
for n_ in ("52", "61", "62", "63", "64"):
    by(s, "コネクタ: カギ線 " + n_).line.color.rgb = RGBColor.from_string("D9D9D9")
el = by(s, "コネクタ: カギ線 67")._element
el.getparent().remove(el)
s.shapes._spTree.append(el)

drop(by_text(S[37], "暫定措置として導入済アカウント"))          # S38：古い注記

s = S[41]                                                     # CPO改善のご提案 → オプション
set_lines(by(s, "Google Shape;156;p7"), ["【オプション】CPO改善のご提案｜サンクスLINE誘導"])
put_text(by(s, "Rounded Rectangle 56").text_frame,
         [one("体験前後のフォロー", 12.5, True, NAVY, sa=5), one("体験前の案内で来場を、体験後のフォローで入会を後押しします", 10.5, None, INK, ls=1.25, sa=2)],
         anchor="m", ml=0.35, mr=0.3, mt=0.25, mb=0.15)
nt = by(s, "TextBox 60")
nt.height = Cm(1.1)
put_text(nt.text_frame, [one("※APIツール（Lステップ等）と併用可／完了画面を増やす場合は追加費用", 10, None, MUT, align="c"),
                         one("※「想定の費用対効果」には含まれない、追加のご提案です", 10, None, MUT, align="c")], anchor="m", ml=0, mr=0)

# ============================================================
# S39 削除・CPO改善のご提案を都度発注プランの後ろへ
# ============================================================
lst = prs.slides._sldIdLst
ids = list(lst)
cpo, tsudo, qa = ids[41], ids[42], ids[38]
assert "QA自動化" in "".join(sh.text_frame.text for sh in S[38].shapes if sh.has_text_frame)
assert "都度発注" in "".join(sh.text_frame.text for sh in S[42].shapes if sh.has_text_frame)
prs.part.drop_rel(qa.rId)
lst.remove(qa)
lst.remove(cpo)
lst.insert(list(lst).index(tsudo) + 1, cpo)

# 図形IDの重複を解消
for sl in prs.slides:
    tree = sl.shapes._spTree
    seen, els = set(), [e for e in tree.iter() if e.tag == qn("p:cNvPr")]
    mx = max(int(e.get("id")) for e in els)
    for e in els:
        if int(e.get("id")) in seen:
            mx += 1
            e.set("id", str(mx))
        seen.add(int(e.get("id")))

prs.save(str(OUT))
print("saved:", OUT.name, len(Presentation(str(OUT)).slides), "枚")
