# -*- coding: utf-8 -*-
"""数字の辻褄合わせ3点
① Slide 22: 来場単価の答え（約13,700円）を追記
② Slide 19: 期待効果を77%に統一
③ Slide 10: 脚注の根拠を他社実績に差し替え
"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
TARGET = ROOT / "20260804_株式会社ビイサイドプランニング御中_公式LINEのご提案v04.pptx"

NAVY = "1F285A"; INK = "333333"; MUT = "808080"; WHITE = "FFFFFF"; PALE = "F4F7FF"

prs = Presentation(str(TARGET))
slides = list(prs.slides)


def set_font(run, size, bold=None, color=INK, name="メイリオ"):
    f = run.font; f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    f.name = name
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        e = rPr.find(qn(tag))
        if e is None:
            e = rPr.makeelement(qn(tag), {}); rPr.append(e)
        e.set("typeface", name)
    f.color.rgb = RGBColor.from_string(color)


ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
ANCH = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}


def reset_tf(tf):
    for para in tf.paragraphs[1:]:
        para._p.getparent().remove(para._p)
    for r in list(tf.paragraphs[0].runs):
        r._r.getparent().remove(r._r)


def put_text(tf, paras, anchor="t", ml=0.06, mr=0.06, mt=0.03, mb=0.03, wrap=True):
    reset_tf(tf); tf.word_wrap = wrap
    tf.margin_left = Inches(ml); tf.margin_right = Inches(mr)
    tf.margin_top = Inches(mt); tf.margin_bottom = Inches(mb)
    tf.vertical_anchor = ANCH[anchor]
    first = True
    for p in paras:
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        para.alignment = ALIGN[p.get("align", "l")]
        if p.get("sa") is not None:
            para.space_after = Pt(p["sa"])
        if p.get("ls") is not None:
            para.line_spacing = p["ls"]
        for (t, sz, b, c) in p["runs"]:
            r = para.add_run(); r.text = t
            set_font(r, sz, b, c)


def find(slide, needle):
    for sh in slide.shapes:
        if sh.has_text_frame and needle in sh.text_frame.text:
            return sh
    return None


# ---------- ① Slide 22：来場単価の答えを追記 ----------
s = slides[21]
sh = find(s, "CV＝フェア事前申込")
assert sh is not None, "S22 注記が見つからない"
# 既存の注記は右下に残しつつ、成果報酬の帯の上に「答え」を1本置く
opt = find(s, "成果報酬プランの場合")
assert opt is not None, "S22 成果報酬帯が見つからない"
# 「レポート」「定例会」の2行を1行に統合し、下部に帯のスペースを作る
rep_lab = rep_val = tei_lab = tei_val = None
for sh in list(s.shapes):
    if not sh.has_text_frame:
        continue
    t = sh.text_frame.text.strip()
    if t == "レポート":
        rep_lab = sh
    elif t == "定例会":
        tei_lab = sh
    elif t.startswith("1回／月") and sh.left / 914400 > 4.5:
        (rep_val if rep_val is None else None)
        if rep_val is None:
            rep_val = sh
        else:
            tei_val = sh
# 上の行を「レポート・定例会 / 各1回／月」にまとめ、下の行は削除
keep_lab = rep_lab or tei_lab
keep_val = rep_val or tei_val
if keep_lab is not None:
    put_text(keep_lab.text_frame, [{"runs": [("レポート・定例会", 9.5, True, WHITE)], "align": "c"}],
             anchor="m", ml=0.02, mr=0.02, mt=0.0, mb=0.0)
if keep_val is not None:
    put_text(keep_val.text_frame, [{"runs": [("各1回／月", 10, None, INK)], "align": "c"}],
             anchor="m", ml=0.02, mr=0.02, mt=0.0, mb=0.0)
for sh in (tei_lab if keep_lab is not tei_lab else None,
           tei_val if keep_val is not tei_val else None):
    if sh is not None:
        sh._element.getparent().remove(sh._element)

# 下部3要素の位置を明示的に決め直す
opt.left, opt.top = Inches(0.42), Inches(6.52)
opt.width, opt.height = Inches(6.55), Inches(0.44)
sh_note = find(s, "CV＝フェア事前申込")
if sh_note is not None:
    sh_note.left, sh_note.top = Inches(7.20), Inches(6.50)
    sh_note.width, sh_note.height = Inches(3.30), Inches(0.48)

ans = s.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                         Inches(0.36), Inches(5.96), Inches(9.87), Inches(0.48))
ans.fill.solid(); ans.fill.fore_color.rgb = RGBColor.from_string(NAVY)
ans.line.fill.background(); ans.shadow.inherit = False
put_text(ans.text_frame,
         [{"runs": [("LINEで増える来場者の獲得単価　", 12, True, WHITE),
                    ("約13,700円", 19, True, WHITE),
                    ("　（現状の来場単価 16,000〜18,000円 → 15〜24%減）", 11.5, None, WHITE)],
           "align": "c"}],
         anchor="m", ml=0.0, mr=0.0, mt=0.0, mb=0.0)
# 注記に算出根拠を足す
put_text(sh.text_frame,
         [{"runs": [("※CV＝フェア事前申込＋来場引き上げの合計", 9, None, MUT)], "sa": 2},
          {"runs": [("※獲得単価＝月次259,800円÷来場増19名", 9, None, MUT)]}],
         anchor="m", ml=0.0, mr=0.0)

# ---------- ② Slide 19：期待効果とリード文を77%に統一 ----------
s = slides[18]
lead = find(s, "8割まで改善")
if lead is not None:
    put_text(lead.text_frame,
             [{"runs": [("「未来場者データの自動除外」により、予約からの来場引き上げ率を77%まで改善します。", 14, True, INK)], "ls": 1.05}],
             anchor="m", ml=0.0, mr=0.0)
sh = find(s, "期待効果")
assert sh is not None, "S19 期待効果が見つからない"
put_text(sh.text_frame,
         [{"runs": [("期待効果", 12, True, NAVY)], "sa": 4},
          {"runs": [("予約→来場の引き上げ率　", 11.5, None, INK), ("60〜70%", 13, True, INK),
                    (" → ", 11.5, None, MUT), ("77%", 19, True, NAVY)], "sa": 3},
          {"runs": [("他クライアント事例では8割程度まで改善。本提案は保守的に77%で試算しています。", 9.5, None, INK)], "ls": 1.1, "sa": 2},
          {"runs": [("参照元：社内配信実績", 9, None, MUT)]}],
         anchor="m", ml=0.18, mr=0.14, mt=0.05, mb=0.05)

# ---------- ③ Slide 10：脚注の根拠を差し替え ----------
s = slides[9]
sh = find(s, "Meta広告シミュレーション")
assert sh is not None, "S10 脚注が見つからない"
put_text(sh.text_frame,
         [{"runs": [("※参加率は他クライアントでの実績（予約→来場の引き上げ率 6〜7割→8割程度）を参考に、保守的に77%で設定しています。提案後の数値は本施策の実行を前提とした目標値です。", 9.5, None, MUT)], "ls": 1.05}],
         ml=0.0, mr=0.0)

prs.save(str(TARGET))
print("saved:", TARGET)
