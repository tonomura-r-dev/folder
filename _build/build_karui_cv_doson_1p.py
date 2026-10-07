# -*- coding: utf-8 -*-
"""1枚もの「ハードルの低いCVからLINEにつなぎ、最終成果へ」（2026-10-07 殿村さん指示・新規・単体で完結）。
土台＝業界別資料（16:9・DYMヘッダー/フッター）の1枚を残して中身を作り直す。
  広告→簡易ページ→ハードルの低いCV｜サンクスLINE誘導→LINE友だち追加→LINEでナーチャリング→購入・契約などの最終成果
  左＝AD・獲得領域（紺。2026-10-07 殿村さん指示でオレンジ→紺）／右＝LINE領域（緑）。簡易ページの下に低ハードルCVの例、ナーチャリングの下に施策例。
  下部1文「簡易ページの作成から、サンクスLINE誘導・ナーチャリングまで一連で設計」。
想定値・CPA/CPO・機能詳細・料金・実績・効果数値・前後検索は入れない。
  python3 _build/build_karui_cv_doson_1p.py <業界別ver3.6.pptx>
"""
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

sys.path.insert(0, str(Path(__file__).resolve().parent))
from pptx_parts import *  # noqa

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "20261007_簡易ページからサンクスLINEへの導線ver1.1.pptx"
ORANGE, ORANGE_BG, ORANGE_TX = "F59B21", "FEF0DE", "C46A00"
PALE = "E4E8F6"
L, R = 2.2, 31.7
OV = 0.35

prs = Presentation(sys.argv[1])
# 業界別ページ（ヘッダー一式あり）だけ残す
keep = next(i for i, s in enumerate(prs.slides) if s.shapes[0].text_frame.text.startswith("ターゲット業界別"))
lst = prs.slides._sldIdLst
for i, el in reversed(list(enumerate(list(lst)))):
    if i != keep:
        prs.part.drop_rel(el.rId)
        lst.remove(el)
s = prs.slides[0]
keep_header_only(s)
set_paras(find(s, "TextBox 1"), [("ハードルの低いCVからLINEにつなぎ、最終成果へ", 0)])
set_paras(find(s, "TextBox 2"), [("低ハードルのCVで接点をつくり、LINEで最終成果まで引き上げる", 0)])

# ---- 横の導線（7ステップ）----
steps = [(["広告"], PALE, NAVY, 3.0),
         (["簡易ページ"], NAVY, WHITE, 3.8),
         (["ハードルの", "低いCV"], NAVY, WHITE, 3.8),
         (["サンクス", "LINE誘導"], GREEN, WHITE, 3.8),
         (["LINE", "友だち追加"], GREEN_BG, GREEN_TX, 3.8),
         (["LINEで", "ナーチャリング"], GREEN_BG, GREEN_TX, 4.9),
         (["購入・契約などの", "最終成果"], NAVY, WHITE, 5.5)]
k = (R - L + OV * (len(steps) - 1)) / sum(st[3] for st in steps)
Y, H = 7.2, 3.3
xs = []
x = L
for i, (ls, f, tc, wr) in enumerate(steps):
    w = wr * k
    shape(s, MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON, x, Y, w, H, None, fill=f, adj=0.22)
    pad = 0.3 if i == 0 else 0.7
    label(s, x + pad, Y, w - pad - 0.6, H, [(t, 13, True, tc, 0) for t in ls], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    xs.append((x, w))
    x += w - OV

# ---- 領域の帯と境目 ----
ad_end = xs[2][0] + xs[2][1] - 0.15
line_start = xs[3][0] + 0.25
rect(s, L, 6.0, ad_end - L, 0.14, NAVY)
label(s, L, 5.25, 8.0, 0.75, [("AD・獲得領域", 13, True, NAVY, 0)], anchor=MSO_ANCHOR.MIDDLE)
rect(s, line_start, 6.0, R - line_start, 0.14, GREEN)
label(s, line_start, 5.25, 8.0, 0.75, [("LINE領域", 13, True, GREEN_TX, 0)], anchor=MSO_ANCHOR.MIDDLE)
vline(s, (ad_end + line_start) / 2, 5.3, 14.6, LGRAY, 1.25, dash=True)


def tags(i, items, line, color):
    x0, w = xs[i]
    tw = 4.8
    tx = x0 + (w - tw) / 2 - (0 if i == 0 else 0.15)
    for j, t in enumerate(items):
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, tx, 11.1 + j * 0.85, tw, 0.68, [(t, 10, True, color, 0)], fill=WHITE, line=line,
              lw=0.75, adj=0.5, margins=(0.05, 0, 0.05, 0))


tags(1, ["簡易診断", "チェックコンテンツ", "資料DL"], NAVY, NAVY)
tags(5, ["情報提供", "リマインド", "セグメント配信", "個別フォロー"], GREEN, GREEN_TX)

# ---- 下部メッセージ（1文）----
label(s, L, 15.6, R - L, 1.4, [("簡易ページの作成から、サンクスLINE誘導・ナーチャリングまで一連で設計", 17, True, NAVY, 0)],
      align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

prs.save(str(OUT))
print("saved:", OUT.name, len(prs.slides), "枚")
