# -*- coding: utf-8 -*-
"""チェングロウス ver2.9 → ver3.0：内容ページの余白を増やす（2026-10-02 殿村さん指示）
枠を小さくし、スライド全体の白い余白を増やす。タイトル・見出し線・表紙・扉・弊社実績・裏表紙・費用と体制（25枚目）は触らない。
本文の図形を、ページ中央を基準に S 倍に縮小（位置・大きさ・文字の大きさ・図形内の余白。全体を縮小して見せるのと同じ）。
  python3 _build/patch_chengrowth_v30_margin.py <ver2.9.pptx> [S]
"""
import sys
from pathlib import Path
from pptx import Presentation
from pptx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "20260929_チェングロウス_LINE公式アカウント運用のご提案ver3.0.pptx"
S0 = float(sys.argv[2]) if len(sys.argv) > 2 else 0.92
S = S0
S_BY_SLIDE = {16: 0.96}          # 表が大きいページは縮小を弱める（表の行が詰まって重なるため）
TARGET = [4, 5, 6, 7, 8, 9, 11, 12, 13, 15, 16, 17, 18, 20, 21, 22, 24, 27, 28, 29, 30]
EMU = 360000
prs = Presentation(sys.argv[1])
CX = prs.slide_width / 2
CY = (3.9 + 18.2) / 2 * EMU          # 見出し線〜フッターの中央

for n in TARGET:
    S = S_BY_SLIDE.get(n, S0)
    s = prs.slides[n - 1]
    for sh in s.shapes:
        if sh.name in ("TextBox 1", "Connector 3"):
            continue
        if sh.name == "TextBox 2":                 # 見出し線の上のリード文：左端だけ本文に揃える
            new_left = CX + (sh.left - CX) * S
            sh.width = int(sh.width * S)
            sh.left = int(new_left)
            continue
        L, T, W, H = sh.left, sh.top, sh.width, sh.height
        sh.left = int(CX + (L - CX) * S)
        sh.top = int(CY + (T - CY) * S)
        sh.width = int(W * S)
        sh.height = int(H * S)
        if getattr(sh, "has_table", False) and sh.has_table:
            tbl = sh.table
            for col in tbl.columns:
                col.width = int(col.width * S)
            for row in tbl.rows:
                row.height = int(row.height * S)

def scale_text(el):
    for tag in ("a:rPr", "a:endParaRPr", "a:defRPr"):
        for e in el.iter(qn(tag)):
            sz = e.get("sz")
            if sz:
                e.set("sz", str(max(600, int(round(int(sz) * S / 50.0)) * 50)))
    for b in el.iter(qn("a:bodyPr")):
        for a in ("lIns", "tIns", "rIns", "bIns"):
            if b.get(a):
                b.set(a, str(int(int(b.get(a)) * S)))
    for c in el.iter(qn("a:tcPr")):
        for a in ("marL", "marR", "marT", "marB"):
            if c.get(a):
                c.set(a, str(int(int(c.get(a)) * S)))


for n in TARGET:
    S = S_BY_SLIDE.get(n, S0)
    for sh in prs.slides[n - 1].shapes:
        if sh.name in ("TextBox 1", "Connector 3", "TextBox 2"):
            continue
        scale_text(sh._element)
prs.save(str(OUT))
print("saved:", OUT.name, "S =", S)
