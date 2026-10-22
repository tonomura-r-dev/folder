# -*- coding: utf-8 -*-
"""人材業界（派遣・バイト・単発）：章の組み替えと「友だち追加の動線」の1対1の設計（2026-10-22 殿村さん）。

    python -m _build.patch_haken_v2_chapters <入力.pptx> <出力.pptx>

  入力は殿村さんのPC保存版（2026-10-21c ＝ ルートの `人材業界（派遣・バイト・単発）_LINE公式アカウント運用のご提案.pptx`）。
  ① 第3章＝タウンワークのアカウント運用事例（P13・P14を第2章から第3章へ。P15はそのまま）
  ② 第4章＝新規で獲得した方への配信設計（今の第3章の構築 P16〜P21 を第4章に入れる。ページの並びは変えない）
  ③ アジェンダ（P2）の第2〜4章の説明を合わせる
  ④ P18「友だち追加の動線」を、友だち追加の場面 → 登録する方 → その方に配信する内容 の1対1の表に組み直す。
     別途費用がかかる施策には＊を付け、下に各機能の追加費用を書く（金額は費用プランのページと同じ）
  P12（他社分析）は別のスクリプトで直す。ほかのページは変えない。
"""
import copy
import sys

from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn

from _build.pptx_parts import rect, rich

NAVY, INK, CARD = "1E2761", "333333", "F4F7FF"
IN = 2.54

RETITLE = {13: ("［第2章］", "［第3章］"), 14: ("［第2章］", "［第3章］")} | {n: ("［第3章］", "［第4章］") for n in range(16, 22)}
AGENDA = [("契約満了後の再応募を軸にした全体設計・データ・他社事例", "契約満了後の再応募を軸にした全体設計・データ・他社分析"),
          ("第3章｜新規獲得", "第3章｜参考事例"),
          ("タウンワークの型・友だち追加の動線・あいさつ・リッチメニュー", "タウンワークのLINE運用（公式事例・実際の画面・3つの仕組み）"),
          ("14日間のステップ配信・配信文面・満了前のご案内・通知メッセージ",
           "友だち追加の動線・あいさつ・リッチメニュー\n14日間のステップ配信・満了前のご案内・通知メッセージ")]

TITLE18 = "［第4章］構築｜友だち追加の動線と、その後の配信"
LEAD18 = "どこで友だちになったかで、その方に配信する内容を決める。"
ROWS = [  # (場面, 登録する方, 配信する内容)
    (["① 応募完了画面＊", "（サンクスLINE誘導）"], "応募した方",
     "あいさつ → 希望条件5問 → 14日間のステップ配信\n面接・職場見学の前日のご案内"),
    (["② 離脱防止ポップアップ＊"], "応募せずに\nLPを離れた方", "条件に合う新着求人と、\nリッチメニューの「探す入口」をご案内"),
    (["③ LINE広告（CPF）＊"], "まだ貴社を\n知らない方", "あいさつで「できること」をご案内\n→ 希望条件5問 → 新着求人"),
    (["④ 求人誌・店頭のQRコード"], "求人誌・店頭で\n貴社を知った方", "掲載エリアの新着求人をご案内"),
    (["⑤ 通知メッセージ＊"], "過去に応募した方\n（既存リスト）", "希望条件の再確認 → 次のお仕事のご案内"),
]
COST = [[("＊ 別途費用がかかるもの", 12, True, NAVY)],
        [("① サンクスLINE誘導：初期15万円・月額3万円〜　　② 離脱防止ポップアップ：初期1.5万円・月額3万円", 11.5, False, INK)],
        [("③ LINE広告（CPF）：広告費が別途かかる　　⑤ 通知メッセージ：初期10万円・月額8万円〜＋通数7円", 11.5, False, INK)]]


def first_text(s):
    return next(sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip())


def set_runs(p, text):
    runs = p.findall(qn("a:r"))
    for r in runs[1:]:
        p.remove(r)
    runs[0].find(qn("a:t")).text = text


def paras(s):
    for sh in s.shapes:
        if sh.has_text_frame:
            yield from sh.text_frame._txBody.iter(qn("a:p"))


def ptext(p):
    return "".join(t.text or "" for t in p.iter(qn("a:t")))


def rebuild18(s):
    keep = {"TextBox 1", "TextBox 2", "TextBox 30"} | {f"Rectangle {n}" for n in (3, 4, 8, 9, 13, 14, 18, 19, 23, 24)}
    shapes = {sh.name: sh for sh in s.shapes}
    for name, sh in shapes.items():
        if name not in keep:
            sh._element.getparent().remove(sh._element)
    set_runs(shapes["TextBox 1"].text_frame._txBody.find(qn("a:p")), TITLE18)
    set_runs(shapes["TextBox 2"].text_frame._txBody.find(qn("a:p")), LEAD18)

    x0, w, lw = 0.47, 9.89, 2.68                       # 行の左端・横幅・左の紺の欄の幅（元のまま）
    wx, ww = x0 + lw + 0.1, 2.15                       # 登録する方
    cx = wx + ww + 0.1                                 # 配信する内容
    cw = x0 + w - cx - 0.1
    hy = 1.33
    for x, wd, t in ((x0, lw, "友だち追加の場面"), (wx, ww, "登録する方"), (cx, cw, "その方に配信する内容")):
        rich(s, x * IN, hy * IN, wd * IN, 0.3 * IN, [[(t, 11, True, NAVY)]],
             align=PP_ALIGN.CENTER if x == x0 else PP_ALIGN.LEFT, anchor=MSO_ANCHOR.BOTTOM)
    for k, (scene, who, body) in enumerate(ROWS):
        y = 1.69 + 0.77 * k                            # 元の行の位置
        rich(s, x0 * IN, y * IN, lw * IN, 0.66 * IN,
             [[(scene[0], 13, True, "FFFFFF")]] + [[(t, 10.5, True, "FFFFFF")] for t in scene[1:]],
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        rich(s, wx * IN, y * IN, ww * IN, 0.66 * IN, [[(t, 12, True, NAVY)] for t in who.split("\n")],
             anchor=MSO_ANCHOR.MIDDLE)
        rich(s, cx * IN, y * IN, cw * IN, 0.66 * IN, [[(t, 12, False, INK)] for t in body.split("\n")],
             anchor=MSO_ANCHOR.MIDDLE)
    rect(s, x0 * IN, 5.6 * IN, w * IN, 1.0 * IN, CARD)
    rich(s, (x0 + 0.15) * IN, 5.62 * IN, (w - 0.3) * IN, 0.96 * IN, COST, anchor=MSO_ANCHOR.MIDDLE)


def main(src, out):
    prs = Presentation(src)
    slides = list(prs.slides)
    assert len(slides) == 33
    for no, (old, new) in RETITLE.items():
        p = first_text(slides[no - 1]).text_frame.paragraphs[0]._p
        t = ptext(p)
        assert t.startswith(old), (no, t)
        set_runs(p, new + t[len(old):])
    ag = slides[1]
    for old, new in AGENDA:
        hits = [p for p in paras(ag) if ptext(p) == old]
        assert len(hits) == 1, old
        lines = new.split("\n")
        set_runs(hits[0], lines[0])
        for t in reversed(lines[1:]):                  # 2行に分けるところは、同じ書式の段落を足す
            q = copy.deepcopy(hits[0])
            set_runs(q, t)
            hits[0].addnext(q)
    box = next(sh for sh in ag.shapes if sh.has_text_frame and sh.text_frame.text.startswith("友だち追加の動線・あいさつ"))
    box.top, box.height = int(3.7 * 914400), int(0.64 * 914400)   # 2行になるので、第2・3章の説明と同じ高さに
    box.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    assert first_text(slides[17]).text_frame.text.endswith("友だち追加の動線")
    rebuild18(slides[17])
    prs.save(out)
    print("saved:", out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
