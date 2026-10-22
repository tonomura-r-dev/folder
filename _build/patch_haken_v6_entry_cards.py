# -*- coding: utf-8 -*-
"""人材業界（派遣・バイト・単発）：入口①〜⑤の14日ステップ5枚を、殿村さんの見本（施策の考え方（短期：ステップ配信））の形で作り直す（2026-10-22）。

    python -m _build.patch_haken_v6_entry_cards <殿村さんのPC保存版ver3.pptx> <出力.pptx>

  入力は殿村さんのPC保存版ver3（35枚）。P25の直後に5枚（P26〜P30）を入れる（v5 と同じ入れ方。v5 は使わない）。
  1枚の中身：上の段＝入口／登録する方／ゴール（CV）
             中央＝Day0・2・6・10・14 の5枚のカード（見出し帯［Day＋役割］→ 配信内容 → ねらい → ［検索データ］→ ボタン）
             下＝注記1行
  色は紺・白・黒（＋灰色）だけ（殿村さん）。Day14 の見出し帯だけ黒にして、最後の案内だと分かるようにする。
  検索データは P10・P11 の前後検索だけを使う（殿村さん）。根拠のないカードは「—」と、何をもとに出し分けるかだけ書く。
"""
import sys

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

from _build.patch_haken_v5_entry_pages import by_name, c, set_one
from _build.patch_tanaka_v5 import copy_slide
from _build.pptx_parts import rect, rich, shape

NAVY, INK, BLACK, GRAY, LGRAY, PANEL, WHITE = "1E2761", "333333", "333333", "7F7F7F", "D9D9D9", "F2F2F2", "FFFFFF"
NOTE = "※ Day0〜14は友だち追加からの日数。検索データは前後検索（P10・P11）より"
DAYS = ["Day0", "Day2", "Day6", "Day10", "Day14"]
NO_DATA = "—（{}）"

PAGES = [  # (タイトル, リード, 入口, 登録する方, ゴール, [(役割, 配信内容, ねらい, 検索データ, ボタン)×5])
    ("［第4章］配信設計｜入口①　応募完了画面からの14日ステップ",
     "応募した方には、面接・職場見学の準備を先にお届けし、来ていただくまでつなげる。",
     "応募完了画面（サンクスLINE誘導）", "応募した方", "面接・職場見学に来る",
     [("追加直後：診断", "求人診断 → 面接・見学までの流れ", "応募直後の不安を減らし、日程を確かめていただく",
       "0〜+2日に派遣会社名が並ぶ（エン派遣・アデコ・パソナ）", "日程を確認する"),
      ("来ていただく準備", "面接・見学の持ち物・服装", "当日に迷わず来ていただく",
       "+1〜+15日「派遣 職場見学」", "持ち物を確認する"),
      ("ほかの求人も", "ほかにも合う求人3件", "比べている方に、選べる求人を残す",
       "同じ日に、複数の会社を比べている", "求人を見る"),
      ("しくみの説明", "派遣のしくみ（3年ルール・契約期間）", "働く前の疑問に答え、辞退を防ぐ",
       "当日「派遣 3年ルール」", "しくみを見る"),
      ("次の応募へ", "新着のまとめ＋「条件を変える」", "条件が合わない方も、次の応募につなげる",
       "5日後「派遣 更新しない」＝次の仕事を考え始めている", "条件を変える")]),
    ("［第4章］配信設計｜入口②　離脱防止ポップアップからの14日ステップ",
     "LPを離れた方には、比べるときの疑問に順に答え、14日目に応募をご案内する。",
     "離脱防止ポップアップ", "応募せずにLPを離れた方", "応募",
     [("追加直後：診断", "求人診断 → 合う求人3件", "関心が高いうちに、合う求人を届ける",
       "0〜+2日に派遣会社名が並ぶ（エン派遣・アデコ・パソナ）", "求人を見る"),
      ("会社の比べ方", "派遣会社の比べ方", "同じ日に複数の会社を比べている方に、選ぶ基準を渡す",
       "同じ日に、複数の会社を比べている", "比べ方を見る"),
      ("しくみの説明", "派遣のしくみ（求人は送らない回）", "働き方の疑問に答え、信頼につなげる",
       "当日「派遣 3年ルール」", "しくみを見る"),
      ("応募の準備", "面接・見学の流れと持ち物", "応募の後の流れを見せ、応募しやすくする",
       "+1〜+15日「派遣 職場見学」", "流れを見る"),
      ("応募のご案内", "新着のまとめ", "関心が高いうちに応募を後押しする",
       NO_DATA.format("Day0〜10の反応をもとにご案内"), "応募する")]),
    ("［第4章］配信設計｜入口③　LINE広告からの14日ステップ",
     "まだ貴社を知らない方には、できることとサポートを先に伝え、14日目に応募をご案内する。",
     "LINE広告（CPF）", "まだ貴社を知らない方", "応募",
     [("できることの紹介", "できることの紹介 → 求人診断 → 合う求人3件", "友だちでいる理由を最初に伝える",
       "−7〜−1日「バイト 求人」「ハローワーク 求人」で探し始めている", "求人を見る"),
      ("選ぶ理由", "登録から就業までのサポート", "貴社を選ぶ理由を伝える",
       "0〜+2日に派遣会社名が並ぶ（エン派遣・アデコ・パソナ）", "サポートを見る"),
      ("しくみの説明", "派遣のしくみ（3年ルール・契約期間）", "働き方の疑問に答える",
       "当日「派遣 3年ルール」", "しくみを見る"),
      ("応募の準備", "面接・見学の流れと持ち物", "応募の後の流れを見せる",
       "+1〜+15日「派遣 職場見学」", "流れを見る"),
      ("応募のご案内", "人気の求人のまとめ", "応募を後押しする",
       NO_DATA.format("Day0〜10の反応をもとにご案内"), "応募する")]),
    ("［第4章］配信設計｜入口④　求人誌・店頭のQRコードからの14日ステップ",
     "求人誌・店頭で知った方には、掲載エリアの求人を中心に届け、14日目に応募をご案内する。",
     "求人誌・店頭のQRコード", "求人誌・店頭で貴社を知った方", "応募",
     [("地域の求人", "求人診断（エリアは掲載地域から）→ 掲載エリアの求人3件", "見ていた地域の求人から届ける",
       NO_DATA.format("求人誌・店頭で見た地域から始める"), "求人を見る"),
      ("近くの新着", "掲載エリアの新着", "近くの求人を続けて届ける",
       "0〜+2日「短期バイト」「単発バイト アプリ」", "新着を見る"),
      ("働く条件", "扶養内・掛け持ちで働くときの目安", "働く条件の疑問に答える",
       "+1〜+15日「ダブルワーク 扶養内 いくらまで」", "目安を見る"),
      ("応募の準備", "面接・見学の流れと持ち物", "応募の後の流れを見せる",
       "+1〜+15日「派遣 職場見学」", "流れを見る"),
      ("応募のご案内", "新着のまとめ", "応募を後押しする",
       NO_DATA.format("Day0〜10の反応をもとにご案内"), "応募する")]),
    ("［第4章］配信設計｜入口⑤　通知メッセージからの14日ステップ",
     "過去に応募した方には、次に探す時期を伺い、時期に合わせて再応募をご案内する。",
     "通知メッセージ（既存リスト）", "過去に応募した方", "再応募",
     [("時期を伺う", "求人診断（次に探す時期も伺う）", "ご案内の時期を、探し始める時期に合わせる",
       "当日「派遣 3年ルール」＝先の働き方も気にしている", "診断に答える"),
      ("先にお届け", "時期が近い方に合う求人3件", "探し始めた方に、他社より先に届ける",
       "同じ日に、複数の会社を比べている", "求人を見る"),
      ("満了前の疑問", "更新しないときの次の探し方", "契約満了前の疑問に答える",
       "5日後「派遣 更新しない」", "探し方を見る"),
      ("経験を生かす", "前回と近い職種の新着", "経験を生かせる求人を届ける",
       NO_DATA.format("前回の応募職種をもとに出し分け"), "新着を見る"),
      ("再応募のご案内", "新着のまとめ＋「条件を変える」", "条件を見直して、再応募につなげる",
       NO_DATA.format("Day0〜10の反応をもとにご案内"), "応募する")]),
]

# 位置（インチ）。5枚とも同じ値を使う
X0, W = 0.47, 9.89
BAND_Y, BAND_H, BAND_GAP = 1.5, 0.6, 0.12
COLS, GAP = 5, 0.14
CW = (W - GAP * (COLS - 1)) / COLS
CARD_Y, CARD_H = 2.32, 3.92
HEAD_H = 0.66                                   # 見出し帯
MSG_Y, MSG_H = 0.76, 0.78                       # 配信内容（カード上端から）
AIM_Y, AIM_H = 1.58, 0.62                       # ねらい
DATA_Y, DATA_H = 2.28, 0.94                     # 検索データの灰色の枠
BTN_H, BTN_B = 0.4, 0.12                        # ボタンの高さ・下端からの余白
PAD = 0.1


def build(s, page):
    title, lead, entry, who, goal, steps = page
    keep = {"TextBox 1", "TextBox 2", "TextBox 41"}
    for sh in list(s.shapes):
        if sh.name not in keep:
            sh._element.getparent().remove(sh._element)
    set_one(by_name(s, "TextBox 1"), title)
    set_one(by_name(s, "TextBox 2"), lead)
    set_one(by_name(s, "TextBox 41"), NOTE, size=10)

    bw = (W - BAND_GAP * 2) / 3
    for k, (lab, val) in enumerate([("入口", entry), ("登録する方", who), ("ゴール（CV）", goal)]):
        x = X0 + (bw + BAND_GAP) * k
        rect(s, c(x), c(BAND_Y), c(bw), c(BAND_H), WHITE, line=NAVY, lw=1.0)
        rich(s, c(x + 0.05), c(BAND_Y), c(bw - 0.1), c(BAND_H),
             [[(lab, 10, True, NAVY)], [(val, 12, True, INK)]], align=PP_ALIGN.CENTER)

    for k, (role, msg, aim, data, btn) in enumerate(steps):
        x = X0 + (CW + GAP) * k
        head = BLACK if k == COLS - 1 else NAVY
        rect(s, c(x), c(CARD_Y), c(CW), c(CARD_H), WHITE, line=LGRAY, lw=0.75)
        rect(s, c(x), c(CARD_Y), c(CW), c(HEAD_H), head)
        rich(s, c(x), c(CARD_Y), c(CW), c(HEAD_H),
             [[(DAYS[k], 16, True, WHITE)], [(role, 10.5, False, WHITE)]], align=PP_ALIGN.CENTER)
        rich(s, c(x + PAD), c(CARD_Y + MSG_Y), c(CW - PAD * 2), c(MSG_H), [[(msg, 12, True, NAVY)]],
             anchor=MSO_ANCHOR.TOP)
        rich(s, c(x + PAD), c(CARD_Y + AIM_Y), c(CW - PAD * 2), c(AIM_H), [[(aim, 10.5, False, INK)]],
             anchor=MSO_ANCHOR.TOP)
        rect(s, c(x + PAD), c(CARD_Y + DATA_Y), c(CW - PAD * 2), c(DATA_H), PANEL, line=LGRAY, lw=0.5)
        rich(s, c(x + PAD + 0.04), c(CARD_Y + DATA_Y + 0.04), c(CW - PAD * 2 - 0.08), c(DATA_H - 0.08),
             [[("［検索データ］", 9, True, GRAY)], [(data, 10, False, INK)]], anchor=MSO_ANCHOR.TOP)
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, c(x + PAD), c(CARD_Y + CARD_H - BTN_B - BTN_H),
              c(CW - PAD * 2), c(BTN_H), [(btn, 11.5, True, NAVY, 0)], fill=WHITE, line=NAVY, lw=1.25,
              margins=(0.05, 0.02, 0.05, 0.02))


def main(src, out):
    prs = Presentation(src)
    slides = list(prs.slides)
    assert len(slides) == 35
    base = slides[24]
    assert by_name(base, "TextBox 1").text_frame.text == "［第4章］配信設計｜入口ごとの14日ステップ"
    lst = prs.slides._sldIdLst
    anchor = list(lst)[24]
    for page in PAGES:
        new = copy_slide(prs, base)
        e = list(lst)[-1]
        lst.remove(e)
        anchor.addnext(e)
        anchor = e
        build(new, page)
    assert len(prs.slides) == 40
    prs.save(out)
    print("saved:", out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
