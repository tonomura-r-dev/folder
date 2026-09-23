# -*- coding: utf-8 -*-
"""Slide 18（ステップ配信）を、LINE風トークの実文面つきに作り直す"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
TARGET = ROOT / "20260804_株式会社ビイサイドプランニング御中_公式LINEのご提案v03.pptx"

TNAVY = "002060"
NAVY = "1F285A"
INK = "333333"
MUT = "808080"
WHITE = "FFFFFF"
PALE = "F4F7FF"
BORDER = "D9D9D9"
CHATBG = "E9EDF2"      # LINEのトーク背景
LGREEN = "8DE055"      # 自分側の吹き出し
LINEGRN = "06C755"     # LINEグリーン（ボタン文字）

prs = Presentation(str(TARGET))
slides = list(prs.slides)


def set_font(run, size, bold=None, color=INK, name="メイリオ"):
    f = run.font
    f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    f.name = name
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        e = rPr.find(qn(tag))
        if e is None:
            e = rPr.makeelement(qn(tag), {})
            rPr.append(e)
        e.set("typeface", name)
    f.color.rgb = RGBColor.from_string(color)


ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
ANCH = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}


def reset_tf(tf):
    for para in tf.paragraphs[1:]:
        para._p.getparent().remove(para._p)
    for r in list(tf.paragraphs[0].runs):
        r._r.getparent().remove(r._r)


def put_text(tf, paras, anchor="t", ml=0.05, mr=0.05, mt=0.02, mb=0.02, wrap=True):
    reset_tf(tf)
    tf.word_wrap = wrap
    tf.margin_left = Inches(ml); tf.margin_right = Inches(mr)
    tf.margin_top = Inches(mt); tf.margin_bottom = Inches(mb)
    tf.vertical_anchor = ANCH[anchor]
    first = True
    for p in paras:
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        para.alignment = ALIGN[p.get("align", "l")]
        if p.get("sa") is not None:
            para.space_after = Pt(p["sa"])
        if p.get("ls") is not None:
            para.line_spacing = p["ls"]
        for (t, sz, b, c) in p["runs"]:
            r = para.add_run(); r.text = t
            set_font(r, sz, b, c)
    return tf


def add_text(slide, x, y, w, h, paras, anchor="t", **kw):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    put_text(box.text_frame, paras, anchor=anchor, **kw)
    return box


def add_box(slide, x, y, w, h, fill=None, line=None, lw=1.0,
            shape=MSO_SHAPE.RECTANGLE, radius=None):
    sp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = RGBColor.from_string(line); sp.line.width = Pt(lw)
    sp.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE and radius is not None:
        try:
            sp.adjustments[0] = radius
        except Exception:
            pass
    return sp


def clear_slide(slide):
    spTree = slide.shapes._spTree
    for el in list(spTree):
        if el.tag.split("}")[-1] in ("sp", "cxnSp", "pic", "graphicFrame", "grpSp"):
            spTree.remove(el)


def bubble(slide, x, y, w, lines, sz=8.5, h=None):
    """相手側の白い吹き出し"""
    hh = h if h else 0.20 + 0.155 * len(lines)
    sp = add_box(slide, x, y, w, hh, fill=WHITE,
                 shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.16)
    put_text(sp.text_frame,
             [{"runs": [(t, sz, None, INK)], "ls": 0.95} for t in lines],
             anchor="m", ml=0.09, mr=0.07, mt=0.02, mb=0.02)
    return y + hh


def btncard(slide, x, y, w, labels, sz=8.5):
    """ボタンテンプレート（白カード＋緑文字ボタン＋極細区切り）"""
    rowh = 0.235
    hh = rowh * len(labels)
    add_box(slide, x, y, w, hh, fill=WHITE, line=BORDER, lw=0.75,
            shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.10)
    for i, lb in enumerate(labels):
        yy = y + i * rowh
        if i:
            ln = add_box(slide, x + 0.06, yy, w - 0.12, 0.008, fill=BORDER)
            ln.line.fill.background()
        t = add_text(slide, x, yy, w, rowh,
                     [{"runs": [(lb, sz, True, LINEGRN)], "align": "c"}],
                     anchor="m", ml=0.02, mr=0.02, mt=0.0, mb=0.0)
    return y + hh


def richmenu(slide, x, y, w, cells, h=0.85):
    """トーク画面下部のリッチメニュー（3列×2行）。cells=[(ラベル, 主役か), ...]×6"""
    gap = 0.03
    cw = (w - gap * 4) / 3
    ch = (h - gap * 3) / 2
    add_box(slide, x, y, w, h, fill=WHITE)
    for i, (lb, star) in enumerate(cells):
        r, c = divmod(i, 3)
        cx = x + gap + c * (cw + gap)
        cy = y + gap + r * (ch + gap)
        cell = add_box(slide, cx, cy, cw, ch, fill=NAVY if star else "DDE5F5")
        put_text(cell.text_frame,
                 [{"runs": [(lb, 7, True, WHITE if star else NAVY)], "align": "c"}],
                 anchor="m", ml=0.01, mr=0.01, mt=0.0, mb=0.0)


def mybubble(slide, x_right, y, w, text, sz=8.5):
    """自分側の黄緑の吹き出し（右寄せ）＋既読"""
    hh = 0.26
    sp = add_box(slide, x_right - w, y, w, hh, fill=LGREEN,
                 shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.18)
    put_text(sp.text_frame, [{"runs": [(text, sz, None, INK)], "align": "c"}],
             anchor="m", ml=0.04, mr=0.04, mt=0.0, mb=0.0)
    add_text(slide, x_right - w - 0.42, y, 0.38, hh,
             [{"runs": [("既読", 6.5, None, MUT)], "align": "r"}],
             anchor="m", ml=0.0, mr=0.0)
    return y + hh


# ================= Slide 18 =================
s = slides[17]
clear_slide(s)
add_text(s, 0.60, 0.13, 9.60, 0.45,
         [{"runs": [("具体施策⑥ ステップ配信設計（Day 0 〜 Day 14）", 20, True, TNAVY)]}],
         anchor="m", ml=0.0, mr=0.0)
add_text(s, 0.60, 0.68, 9.60, 0.55,
         [{"runs": [("アクティブ率が最も高い追加直後に価値を渡し、全8通で「登録→予約」まで自動で運びます。", 14, True, INK)], "ls": 1.05}],
         anchor="m", ml=0.0, mr=0.0)

# --- 8通のタイムライン ---
steps = [("Day 0", "特典＋診断"), ("Day 1", "診断結果"), ("Day 3", "不安解消"),
         ("Day 5", "参加者の声"), ("Day 7", "予約オファー"), ("Day 10", "特典で後押し"),
         ("Day 14", "ラストコール")]
cw, gap = 1.15, 0.26
for i, (d, label) in enumerate(steps):
    x = 0.60 + i * (cw + gap)
    hl = d in ("Day 0", "Day 7", "Day 14")
    hb = add_box(s, x, 1.32, cw, 0.34, fill=NAVY if hl else "8A93A8")
    put_text(hb.text_frame, [{"runs": [(d, 10, True, WHITE)], "align": "c"}],
             anchor="m", ml=0.0, mr=0.0, mt=0.0, mb=0.0)
    add_text(s, x, 1.68, cw, 0.26,
             [{"runs": [(label, 8.5, None, INK)], "align": "c"}],
             anchor="m", ml=0.0, mr=0.0)
    if i < len(steps) - 1:
        add_box(s, x + cw + 0.06, 1.42, 0.14, 0.16, fill=MUT, shape=MSO_SHAPE.RIGHT_ARROW)

# --- 主要3通の実文面（LINE風） ---
cols = [
    (0.60, "Day 0", "登録直後（2通）"),
    (4.02, "Day 7", "日曜の夜19〜21時"),
    (7.44, "Day 14", "ラストコール"),
]
PW = 3.18
for cx, day, when in cols:
    hb = add_box(s, cx, 2.08, PW, 0.42, fill=NAVY)
    put_text(hb.text_frame,
             [{"runs": [(day + "　", 11, True, WHITE), (when, 8.5, None, WHITE)], "align": "c"}],
             anchor="m", ml=0.03, mr=0.03, mt=0.0, mb=0.0)
    add_box(s, cx, 2.50, PW, 3.40, fill=CHATBG)

# --- 各パネル下部のリッチメニュー（フェーズごとに出し分け） ---
menus = [
    (0.60, [("出展企業", True), ("会場MAP", False), ("開催日程", False),
            ("事前予約", False), ("相談", False), ("FAQ", False)], "通常時"),
    (4.02, [("事前予約", True), ("特典受取", False), ("出展企業", False),
            ("会場MAP", False), ("相談", False), ("FAQ", False)], "週末"),
    (7.44, [("会場MAP", True), ("当日の流れ", False), ("持ち物", False),
            ("出展企業", False), ("事前予約", False), ("FAQ", False)], "フェア直前"),
]
for mx, cells, mlabel in menus:
    richmenu(s, mx + 0.06, 4.99, PW - 0.12, cells)
    add_text(s, mx, 5.90, PW, 0.24,
             [{"runs": [("▲ リッチメニュー：" + mlabel + "パターン", 8, None, MUT)], "align": "c"}],
             anchor="m", ml=0.0, mr=0.0)

# Day 0
y = 2.62
y = bubble(s, 0.72, y, 2.42, ["友だち追加ありがとうございます！", "登録特典をお送りします。"])
y = btncard(s, 0.72, y + 0.06, 2.42, ["受け取る"])
y = bubble(s, 0.72, y + 0.10, 2.42, ["あなたに合う企業をご案内するため、", "3つだけ教えてください（30秒）"])
y = btncard(s, 0.72, y + 0.06, 2.42, ["事務", "製造・軽作業", "販売・サービス"])

# Day 7
y = 2.70
y = bubble(s, 4.14, y, 2.42,
           ["次回フェアの予約を受付中です。", "",
            "▼9月26日（土）東近江会場", "　13:00〜16:00", "　入場無料・履歴書不要"])
y = bubble(s, 4.14, y + 0.08, 2.42,
           ["事前予約の方には『事前エントリー", "シート』をお渡ししています。"])
y = btncard(s, 4.14, y + 0.08, 2.42, ["予約する"], sz=9.5)

# Day 14
y = 2.72
y = bubble(s, 7.56, y, 2.42,
           ["フェアまであと3日です。", "",
            "1日で複数社と直接話せる機会は", "貴重です。まだ間に合います。"])
y = btncard(s, 7.56, y + 0.10, 2.42, ["予約する"], sz=9.5)
y = mybubble(s, 9.98, y + 0.16, 1.30, "今回は見送る", sz=8)

# --- 分岐の設計 ---
bb = add_box(s, 0.60, 6.20, 9.60, 0.62, fill=PALE)
put_text(bb.text_frame,
         [{"runs": [("分岐の設計　", 10.5, True, NAVY),
                    ("予約した時点でステップを停止し、開催日起点のリマインドへ合流　／　Day 14まで未予約は長期育成へ切替　／　「今回は見送る」で離脱を検知しタグ振分け", 10, None, INK)], "ls": 1.15}],
         anchor="m", ml=0.18, mr=0.14, mt=0.04, mb=0.04)

add_text(s, 0.60, 6.86, 9.60, 0.26,
         [{"runs": [("※開催前のリマインド（3日前・前日・当日朝）は、登録日ではなく開催日起点の配信として別途実施します。", 9, None, MUT)]}], ml=0.0)

prs.save(str(TARGET))
print("saved:", TARGET)
