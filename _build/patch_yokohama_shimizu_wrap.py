# -*- coding: utf-8 -*-
"""横浜シミズ資料（33枚）：1語だけ次の行に落ちる・言葉の途中で切れる折り返しを、文の切れ目の改行（a:br）で2行に整える（2026-10-29 殿村さん）。

    python -m _build.patch_yokohama_shimizu_wrap <資料.pptx> [出力.pptx]

  対象は今回変えたページ（14・17・24・32）。文字・大きさ・色は変えない（「静岡」の行は「湘南」を外して1行に、32の追加施策の枠は言い方を短くして2行ずつに収める）。
  出力を省くと上書き。
"""
import copy
import sys

from pptx import Presentation
from pptx.oxml.ns import qn

from _build.patch_yokohama_shimizu_chapters import head


def br(txb, after):
    """txBody の中で after を含むランを after の直後で切り、間に改行（a:br）を入れる"""
    for r in txb.iter(qn("a:r")):
        t = r.find(qn("a:t"))
        if t.text and after in t.text:
            k = t.text.index(after) + len(after)
            rest = t.text[k:]
            assert rest, after
            t.text = t.text[:k]
            r2 = copy.deepcopy(r)
            r2.find(qn("a:t")).text = rest
            b = r.makeelement(qn("a:br"), {})
            rpr = r.find(qn("a:rPr"))
            if rpr is not None:
                b.append(copy.deepcopy(rpr))
            r.addnext(b)
            b.addnext(r2)
            return
    raise SystemExit(f"見つからない: {after}")


def sp(s, sid):
    return next(sh for sh in s.shapes if sh.shape_id == sid)._element


def cell(s, r, c):
    return next(sh for sh in s.shapes if sh.has_table).table.cell(r, c)._tc


def main(src, out):
    prs = Presentation(src)
    sl = list(prs.slides)
    s14, s17, s24, s32 = (sl[n - 1] for n in (14, 17, 24, 32))
    assert head(s14).text_frame.text.startswith("業界課題②")
    assert head(s17).text_frame.text.startswith("競合分析②")
    assert head(s24).text_frame.text.startswith("施策③")
    assert head(s32).text_frame.text.startswith("今後の展開")

    for sid, after in ((23, "つながらず、"), (27, "服装の"), (31, "手段が残らず、"),
                       (25, "応募いただき、"), (29, "一斉に、")):
        br(sp(s14, sid), after)

    for (r, c), after in (((1, 1), "手順を"), ((1, 2), "ふりがな）"), ((1, 3), "LINEでの"), ((1, 4), "届くことを"),
                          ((2, 1), "スタッフ登録・"), ((2, 3), "フォーム"), ((3, 2), "未確認"), ((3, 3), "起床連絡・"),
                          ((4, 1), "「条件に合う"), ((5, 1), "お仕事一覧から"), ((5, 3), "エントリーを"),
                          ((6, 1), "Q&A・"), ((6, 3), "あり"), ((6, 4), "勤務の案内・")):
        br(cell(s17, r, c), after)
    br(sp(s17, 23), "活用して、")
    br(sp(s17, 23), "導入事例）。")
    br(sp(s17, 20), "弊社確認）／")

    t = next(sh for sh in s24.shapes if sh.has_table).table
    r0 = t.cell(1, 1).text_frame.paragraphs[0].runs[0]
    assert r0.text.startswith("神奈川（横浜・川崎・湘南など）")
    r0.text = "神奈川（横浜・川崎など）／静岡"
    br(cell(s24, 2, 2), "質問・")

    br(sp(s32, 25), "再稼働の計測")
    box = next(sh for sh in s32.shapes if sh.shape_id == 29)       # 枠が狭いので、言い方を短くして2行ずつに収める
    ps = box.text_frame.paragraphs
    assert ps[1].runs[0].text.startswith("社員採用") and ps[2].runs[0].text.startswith("社員とアルバイト")
    ps[1].runs[0].text = "社員採用（中途・新卒）にも、同じアカウントを使える可能性があります。"
    ps[2].runs[0].text = "社員とアルバイトで案内を分ける想定です（今回のご提案には含めていません）。"
    br(box._element, "にも、")
    br(box._element, "想定です")
    prs.save(out)
    print("saved:", out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else sys.argv[1])
