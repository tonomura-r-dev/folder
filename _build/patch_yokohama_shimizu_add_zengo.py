# -*- coding: utf-8 -*-
"""横浜シミズ：前後検索「単発バイト」のページを、施策③と配信イメージの間に1枚足す（2026-10-21 殿村さん「前後検索も入れたい」）。

    python -m _build.patch_yokohama_shimizu_add_zengo <横浜シミズ資料.pptx> <タナカホーム資料.pptx> [出力.pptx]

  タナカホームの前後検索「注文住宅」のページをそのままコピーし、画像と言葉だけ入れ替える（形は同じ。CLAUDE.md 2026-10-19）。
  画像は _images/yokohama_shimizu_zengo_tanpatsu_color.png（`python3 _build/make_yokohama_shimizu_zengo_color.py` で作る。
  元データ＝殿村さんから受領した LINEヤフーの前後検索「単発バイト」）。
  凡例は3つ（求人サービスを比べる／応募・面接の準備／給与・働き方を調べる）。橙・赤は使わない（下の一言の赤い文字も紺に）。
  注記に「色分けは弊社」は書かない。ほかのページは触らない。出力を省くと上書き。
"""
import io
import sys
from pathlib import Path

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches

from _build.patch_tanaka_v5 import copy_slide
from _build.patch_yokohama_shimizu_deck_plans import set_text

ROOT = Path(__file__).resolve().parent.parent
IMG = ROOT / "_images/yokohama_shimizu_zengo_tanpatsu_color.png"
TANAKA_PAGE = 8
COLORS = ["3467B2", "8E4EC6", "00897B"]
TEXTS = {
    2: "前後検索「単発バイト」｜当日に何社も比べ、その後は給与・働き方を調べる",
    6: "応募の前後に調べていることを、LINEで先に届ける",
    11: "求人サービスを比べる",
    13: "応募・面接の準備",
    15: "給与・働き方を調べる",
    18: "検索前",
    19: "求人サイトで探し始める\nバイト 求人・タウンワークなど",
    20: "当日",
    21: "単発のサービスを比べる\nタイミー・バイトル・短期バイトなど。履歴書の書き方も",
    22: "1〜15日後",
    23: "給与・働き方を調べる\n日払い・確定申告・ダブルワークの扶養。面接の服装も",
    24: "求人サービスを比べる → 応募・面接の準備 → 給与・働き方 の順に調べている",
    25: "※LINEヤフー社の前後検索データより。「単発バイト」を検索した方が、前後15日間に検索した言葉（縦軸＝検索した人数）。",
}


def by_id(s, sid):
    return next(sh for sh in s.shapes if sh.shape_id == sid)


def set_paras(sh, lines):
    """段落ごとに、その段落の最初のrunの書式で文字を入れる（見出し＋本文の2段落）"""
    ps = sh._element.find(qn("p:txBody")).findall(qn("a:p"))
    assert len(ps) == len(lines)
    for p, text in zip(ps, lines):
        runs = p.findall(qn("a:r"))
        for r in runs[1:]:
            p.remove(r)
        runs[0].find(qn("a:t")).text = text


def main(src, tanaka, out):
    prs = Presentation(src)
    assert len(prs.slides) == 17
    assert by_id(prs.slides[10], 2).text_frame.text.startswith("施策③")
    tk = Presentation(tanaka)
    ts = tk.slides[TANAKA_PAGE - 1]
    assert by_id(ts, 2).text_frame.text.startswith("前後検索「注文住宅」")
    s = copy_slide(prs, ts)

    pic = by_id(s, 7)                                         # 画像の差し替え
    blip = pic._element.find(".//" + qn("a:blip"))
    old = blip.get(qn("r:embed"))
    _, rid = s.part.get_or_add_image_part(io.BytesIO(IMG.read_bytes()))
    blip.set(qn("r:embed"), rid)
    s.part.drop_rel(old)
    pic.name = "前後検索 単発バイト（色分け）"

    for sid in (16, 17):                                      # 凡例は3つ（4つ目を外す）
        el = by_id(s, sid)._element
        el.getparent().remove(el)
    for sid, color in zip((10, 12, 14), COLORS):
        by_id(s, sid)._element.find(qn("p:spPr")).find(qn("a:solidFill")).find(qn("a:srgbClr")).set("val", color)
    step = by_id(s, 12).top - by_id(s, 10).top                # 凡例が1つ減った分、時期の枠を上げる
    for sid in range(18, 24):
        by_id(s, sid).top -= step

    band, tag = by_id(s, 8), by_id(s, 9)                      # 0日〜1日（画像の 1037px〜1270px）に合わせる
    band.left = int(pic.left + pic.width * 1037 / 2000)
    band.width = int(pic.width * (1270 - 1037) / 2000)
    tag.left = int(band.left + band.width / 2 - tag.width / 2)

    for sid, t in TEXTS.items():
        sh = by_id(s, sid)
        if sid in (19, 21, 23):
            set_paras(sh, t.split("\n"))
        elif sid == 24:                                       # 赤い文字のrunを外し、紺の1本にする
            set_text(sh.text_frame, t)
        else:
            set_text(sh.text_frame, t)
    for r in by_id(s, 24)._element.iter(qn("a:srgbClr")):
        assert r.get("val") != "C00000"

    lst = prs.slides._sldIdLst                                # 施策③のあと（12枚目）に入れる
    ids = list(lst)
    new = ids[-1]
    lst.remove(new)
    ids[10].addnext(new)
    assert len(prs.slides) == 18
    prs.save(out)
    print("saved:", out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2], sys.argv[3] if len(sys.argv) > 3 else sys.argv[1])
