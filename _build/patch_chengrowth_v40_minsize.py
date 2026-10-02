# -*- coding: utf-8 -*-
"""チェングロウス ver3.8 → 8pt未満をなくし、図形内の文字を10〜12ptにそろえる（2026-10-02 殿村さん指示）
  ・注記・出典（下の細い文字）は8pt以上
  ・図形（カード・枠）・表の文字は10pt以上、本文の枠は12ptを基本に（10〜11.5pt → 12pt）
  ・見出し・リード・大きな数字は変えない。2枚目（弊社実績）は変えない。
  このあと patch_chengrowth_v39_fit_boxes.py で枠の高さを合わせる。
  python3 _build/patch_chengrowth_v40_minsize.py <ver3.8.pptx> <出力.pptx>
"""
import sys
from pptx import Presentation
from pptx.oxml.ns import qn

E = 360000
prs = Presentation(sys.argv[1])


def walk(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type == 6:
            yield from walk(sh.shapes)


changed = 0
for i, sl in enumerate(prs.slides, 1):
    if i in (1, 2, len(prs.slides)):
        continue
    for sh in walk(sl.shapes):
        if sh.name in ("TextBox 1", "TextBox 2", "Connector 3"):
            continue
        is_table = getattr(sh, "has_table", False) and sh.has_table
        is_auto = sh.shape_type == 1
        is_note = sh.shape_type == 17 and (sh.top >= 16.0 * E or sh.height < 1.0 * E)
        for tag in ("a:rPr", "a:endParaRPr", "a:defRPr"):
            for r in sh._element.iter(qn(tag)):
                sz = r.get("sz")
                if not sz:
                    continue
                v = int(sz)
                new = v
                if is_note:
                    new = max(v, 800)
                elif is_table:
                    new = max(v, 1000)
                elif is_auto:
                    if v < 1000:
                        new = 1000
                    elif v < 1200 and not (sh.height < 1.3 * E and sh.width < 12 * E):   # 小さな見出し枠は折り返し防止で10〜11ptのまま
                        new = 1200
                else:
                    if v < 1000:
                        new = 1000
                if new != v:
                    r.set("sz", str(new))
                    changed += 1
prs.save(sys.argv[2])
print("変更した箇所:", changed)
