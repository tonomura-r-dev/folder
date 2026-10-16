# -*- coding: utf-8 -*-
"""株式会社すなおコンサルティング LINE公式アカウント運用のご提案（骨組み・2026-10-16 殿村さん指示「骨組だけの資料作成をして」）。

    python -m _build.build_sunao_v1 <あらうちv10.pptx（器）> <出力.pptx>

  器：あらうち v10（4:3・DYMのヘッダー／表紙／裏表紙）。表紙と裏表紙だけ残し、本文は全部作り直す。
  構成は殿村さん確定版（`_drafts/すなおコンサルティング_与件整理（2026-10-15）.md` の「提案の概要・確定版」）。
  SIM（概算お見積り）と成果報酬の条件は、まだ決まっていないので破線の差し込み枠にしてある。
"""
import sys

from pptx import Presentation

from _build.arauchi_v4.lib import NAVY, RED, TEXT, add_box, add_table, set_paras
from _build.arauchi_v10.build import MID, TINT, LINE_C, by_id, key_takeaway, meiryo, runs, three_points
from _build.build_tanaka_v1 import GRAY, COVER, BACK, HDR_SRC, HDR, new_page, head_cell, cell, label_cell, note, band
from _build.build_tanaka_v2 import heading, flow

W = 9.73
X0 = 0.55


def slot(s, x, y, w, h, text):
    """まだ中身が決まっていない所の破線の差し込み枠"""
    add_box(s, x, y, w, h, [runs((text, dict(size=14, bold=True, color=GRAY)))], line=GRAY, dash="dash",
            align="ctr", anchor="ctr", name=f"差し込み {text}")


# ---------------------------------------------------------------- 表紙
def p_cover(s):
    set_paras(by_id(s, 31), ["株式会社すなおコンサルティング 御中"])
    set_paras(by_id(s, 26), ["LINE公式アカウント運用のご提案"])
    set_paras(by_id(s, 27), ["広告で集めた方をLINEでナーチャリングをし、", "セミナー参加・入会へ"])
    for sid in (31, 26, 27):
        meiryo(by_id(s, sid))


# ---------------------------------------------------------------- 要旨
def p_summary(prs, src):
    s = new_page(prs, src, "サマリー", "広告で集めた方をLINEにため、次のイベント参加・入会（LTV）を最大化する")
    key_takeaway(s, [runs(("広告 → イベント＋LINEにたまる → ナーチャリング → 次のイベント・入会", dict(size=18, bold=True, color=NAVY))),
                     runs(("この流れを構築する", dict(size=14, color=NAVY)))])
    three_points(s, "補足ポイント", [
        ("01 現状と課題", "申し込まずに離れた方・入会していない方を、次のイベントまで追う場所が無い。", NAVY),
        ("02 施策", "ポップアップツールとサンクスLINE誘導で友だちにし、興味あり・予約・参加の段階ごとに配信する。", MID),
        ("03 目標と費用", "追える人数を増やし・減らさず、入会へ。平常期 月5〜10万円／イベント期 月30万円程度。", RED),
    ])


# ---------------------------------------------------------------- 1 現状と課題
def p_kadai(prs, src):
    s = new_page(prs, src, "現状と課題：イベントの後、離れた方・未入会の方とのつながりが切れる",
                 "イベントで集めた方を、次のイベントまで追う仕組みが無い")
    flow(s, [("広告", NAVY, "FFFFFF"), ("LP", NAVY, "FFFFFF"), ("予約", NAVY, "FFFFFF"),
             ("イベント参加", NAVY, "FFFFFF"), ("サロン入会", NAVY, "FFFFFF")], 1.95, 0.85, size=13)
    bw = (W - 0.32 * 4) / 5
    for k, (head, body) in ((1, ("申し込まずに離れた方", "次のイベントまで、案内する手段が無い")),
                            (3, ("参加して入会していない方", "入会まで、案内を続ける手段が無い"))):
        cx = X0 + k * (bw + 0.32) + bw / 2
        add_box(s, cx - 0.2, 2.85, 0.4, 0.32, [runs(("▼", dict(size=13, color=RED)))], align="ctr", anchor="ctr",
                ins=(0, 0, 0, 0), name="矢印下")
        add_box(s, cx - 1.65, 3.2, 3.3, 1.2, [runs((head, dict(size=13, bold=True, color=RED))),
                                             runs((body, dict(size=12, color=TEXT)))],
                line=RED, align="ctr", anchor="ctr", line_spacing=125, ins=(0.1, 0.06, 0.1, 0.06), name=f"切れる{k}")
    band(s, [("次のイベントでは、", dict(size=14, bold=True, color=NAVY)),
             ("また広告で集め直している", dict(size=14, bold=True, color=RED))], y=4.95, h=0.75)


# ---------------------------------------------------------------- 1 本質課題
def p_honshitsu(prs, src):
    s = new_page(prs, src, "本質課題", "イベントのたびに集めた方が、手元に残っていない")
    key_takeaway(s, [runs(("ためて追い続ける場所が無い", dict(size=20, bold=True, color=NAVY))),
                     runs(("この場所を、LINE公式アカウントでつくる", dict(size=14, color=NAVY)))])
    three_points(s, "3つの課題と打ち手", [
        ("01 集める", "離れる方・予約した方を、友だちにする手段が無い", NAVY),
        ("02 追う", "次のイベントまで、情報を届ける手段が無い", MID),
        ("03 つなげる", "参加した方に、入会を案内する手段が無い", RED),
    ], body_h=1.0, foot=["ポップアップツール・サンクスLINE誘導", "平常期の配信", "参加後の入会の案内"])


# ---------------------------------------------------------------- 2 なぜLINEか
def p_line(prs, src):
    s = new_page(prs, src, "なぜLINEか：広告は毎回ゼロから、LINEは積み上がる",
                 "一度つながった方に、次のイベントまで何度でも届けられる")
    rows = [[head_cell("", 13), head_cell("広告", 13), head_cell("LINE公式アカウント", 13)],
            [label_cell("届く相手", 13), cell("毎回、新しく集める", 13), cell("一度友だちになった方に、何度でも届く", 13, True, NAVY)],
            [label_cell("届き方", 13), cell("広告を見たときだけ", 13), cell("通知で届き、見てもらいやすい", 13, True, NAVY)],
            [label_cell("イベントの後", 13), cell("つながりが切れる", 13), cell("友だちとして、毎月積み上がる", 13, True, NAVY)]]
    add_table(s, X0, 1.85, [2.0, 3.4, W - 5.4], [0.48, 0.85, 0.85, 0.85], rows, name="広告とLINE")
    band(s, [("この仕組みが回れば、", dict(size=14, bold=True, color=NAVY)),
             ("今後の広告費を抑えられる可能性もある", dict(size=14, bold=True, color=RED))], y=5.3, h=0.75)


# ---------------------------------------------------------------- 3 施策：友だちの入口
def p_iriguchi(prs, src):
    s = new_page(prs, src, "友だちの入口：ポップアップツールとサンクスLINE誘導の2つ",
                 "申し込まずに離れる方と、予約した方の両方を、LINEの友だちにする")
    cw, gap = (W - 0.3) / 2, 0.3
    cols = [("ポップアップツール（離脱防止）",
             [("対象", "LPで申し込まずに離れる方"), ("タイミング", "LPを離れようとしたとき"),
              ("案内", "次のセミナーの案内を、LINEで受け取れる")]),
            ("サンクスLINE誘導",
             [("対象", "セミナーを予約した方"), ("タイミング", "予約完了の画面"),
              ("案内", "当日の案内・リマインドを、LINEで届ける")])]
    for k, (head, items) in enumerate(cols):
        x = X0 + k * (cw + gap)
        heading(s, x, 1.85, cw, head)
        rows = [[label_cell(a, 13), cell(b, 13)] for a, b in items]
        add_table(s, x, 2.3, [1.35, cw - 1.35], [0.8] * 3, rows, name=f"入口{k + 1}")
    band(s, [("どちらも、次のイベントまで", dict(size=14, bold=True, color=NAVY)),
             ("追える友だちになる", dict(size=14, bold=True, color=RED))], y=5.05, h=0.75)


# ---------------------------------------------------------------- 3 施策：段階ごとの配信
def p_dankai(prs, src):
    s = new_page(prs, src, "配信：興味あり・予約・参加の段階ごとに、案内を変える",
                 "段階に合わせた配信で、次のイベント参加と入会へつなげる")
    flow(s, [("興味あり", NAVY, "FFFFFF"), ("予約", NAVY, "FFFFFF"), ("参加", NAVY, "FFFFFF"),
             ("サロン入会", RED, "FFFFFF")], 1.85, 0.6, size=13)
    rows = [[head_cell("段階", 13), head_cell("対象", 13), head_cell("届ける内容", 13), head_cell("ねらい", 13)],
            [label_cell("興味あり", 13), cell("予約していない方\n（ポップアップツールから）", 12),
             cell("セミナーの見どころ動画・参加特典", 13), cell("予約へ", 13, True, NAVY, align="ctr")],
            [label_cell("予約", 13), cell("予約した方\n（サンクスLINE誘導から）", 12),
             cell("前日のリマインド・当日の案内", 13), cell("当日の参加へ", 13, True, NAVY, align="ctr")],
            [label_cell("参加", 13), cell("イベントに参加した方", 12),
             cell("イベント直後に、サロン入会の案内", 13), cell("入会へ", 13, True, NAVY, align="ctr")]]
    add_table(s, X0, 2.75, [1.5, 2.75, 3.68, 1.8], [0.48, 0.85, 0.85, 0.85], rows, name="段階ごとの配信")
    note(s, "※配信の内容は、イベントの内容に合わせて設計", y=6.0)


# ---------------------------------------------------------------- 3 施策：配信の回数
def p_kaisu(prs, src):
    s = new_page(prs, src, "配信の回数：平常期は月2回、イベント期は月4〜5回",
                 "イベントの無い時期もつながりを保ち、イベント期に参加・入会へつなげる")
    flow(s, [("平常期\n月2回", None, NAVY), ("イベント期\n月4〜5回", NAVY, "FFFFFF"), ("平常期\n月2回", None, NAVY),
             ("次のイベント期\n月4〜5回", NAVY, "FFFFFF")], 1.85, 0.8, size=13)
    rows = [[head_cell("", 13), head_cell("平常期（イベントの無い月）", 13), head_cell("イベント期（イベントの前後）", 13)],
            [label_cell("回数", 13), cell("月2回", 14, True, NAVY, align="ctr"), cell("月4〜5回", 14, True, NAVY, align="ctr")],
            [label_cell("内容", 13), cell("定期配信1回＋企画・動画1回", 13), cell("告知・リマインド・イベント直後のフォロー", 13)],
            [label_cell("ねらい", 13), cell("つながりを保ち、追える人数を減らさない", 13), cell("次のイベントの参加・入会へつなげる", 13)]]
    add_table(s, X0, 2.95, [1.6, (W - 1.6) / 2, (W - 1.6) / 2], [0.48, 0.65, 0.8, 0.8], rows, name="配信の回数")
    note(s, "※イベントを年4回とした場合", y=6.0)


# ---------------------------------------------------------------- 4 目標
def p_mokuhyo(prs, src):
    s = new_page(prs, src, "目標：追える人数を増やし・減らさず、次のイベント参加・入会へ",
                 "時期ごとに、見る指標を分ける")
    flow(s, [("追える人数", NAVY, "FFFFFF"), ("次のイベントの予約・参加", NAVY, "FFFFFF"), ("サロン入会", RED, "FFFFFF")],
         1.85, 0.6, size=13)
    rows = [[head_cell("指標", 13), head_cell("見る時期", 13), head_cell("内容", 13)],
            [label_cell("追える人数", 13), cell("平常期", 13, align="ctr"), cell("ブロックされずに届く友だちの数。増やし、減らさない", 13)],
            [label_cell("予約・参加", 13), cell("イベント期", 13, align="ctr"), cell("LINEからの予約と、当日の参加", 13)],
            [label_cell("サロン入会", 13), cell("イベントの後", 13, align="ctr"), cell("参加した方のうち、入会した方", 13)]]
    add_table(s, X0, 2.75, [2.2, 1.9, W - 4.1], [0.48, 0.8, 0.8, 0.8], rows, name="指標")


# ---------------------------------------------------------------- 4 概算お見積り
def p_mitsumori(prs, src):
    s = new_page(prs, src, "概算お見積り：友だち数・予約・入会の見込み",
                 "平常期とイベント期の、友だち数と予約・入会の見込み")
    slot(s, X0, 1.85, W, 4.5, "概算お見積り（SIM）の表")


# ---------------------------------------------------------------- 4 料金
def p_ryokin(prs, src):
    s = new_page(prs, src, "料金：平常期 月5〜10万円、イベント期 月30万円程度",
                 "イベントの有無に合わせて、月ごとの費用を変える")
    rows = [[head_cell("区分", 13), head_cell("月額", 13), head_cell("内容", 13)],
            [label_cell("平常期\n（イベントの無い月）", 12.5), cell("5〜10万円", 16, True, NAVY, align="ctr"),
             cell("月2回の配信（ツール費を含む）", 13)],
            [label_cell("イベント期\n（イベントの前後）", 12.5), cell("30万円程度", 16, True, NAVY, align="ctr"),
             cell("月4〜5回の配信（告知・リマインド・イベント直後のフォロー）", 13)],
            [label_cell("年間の目安", 12.5), cell("約160〜200万円", 16, True, NAVY, align="ctr"),
             cell("イベント期4か月＋平常期8か月", 13)]]
    add_table(s, X0, 1.85, [2.3, 2.6, W - 4.9], [0.48, 0.85, 0.85, 0.85], rows, name="料金")
    for k, t in enumerate(["※広告費は含まない",
                           "※イベントを年4回、イベント期を各1か月として計算",
                           "※施策内容・配信本数・連携範囲等により、費用は個別にお見積り"]):
        note(s, t, y=5.35 + k * 0.28, h=0.28)


# ---------------------------------------------------------------- 4 スケジュール
def p_schedule(prs, src):
    s = new_page(prs, src, "スケジュール：初期構築に約1か月、構築後に運用を開始",
                 "お申込みから約1か月で、友だちの入口と配信の準備を整える")
    flow(s, [("お申込み", None, NAVY), ("初期構築（約1か月）", NAVY, "FFFFFF"), ("運用開始", NAVY, "FFFFFF")],
         1.85, 0.75, size=13)
    bw = (W - 0.32 * 2) / 3
    for k, items in enumerate([["ご契約", "運用方針のすり合わせ"],
                               ["アカウントの設定", "ポップアップツールの設置", "サンクスLINE誘導の設置", "段階ごとの配信の設計"],
                               ["平常期・イベント期の配信", "配信の結果を確認し、改善"]]):
        add_box(s, X0 + k * (bw + 0.32), 2.75, bw, 1.9, [runs(("・" + t, dict(size=13, color=TEXT))) for t in items],
                line=LINE_C, anchor="t", line_spacing=130, ins=(0.12, 0.1, 0.1, 0.08), name=f"内容{k + 1}")


# ---------------------------------------------------------------- 4 成果報酬プラン
def p_seika(prs, src):
    s = new_page(prs, src, "成果報酬プラン：リスクを抑えて始められるプラン",
                 "固定費のプランに加えて、成果に応じて費用が決まるプランも選択可能")
    cw, gap = (W - 0.3) / 2, 0.3
    heading(s, X0, 1.85, cw, "固定費のプラン")
    add_box(s, X0, 2.3, cw, 2.6, [runs(("平常期　月5〜10万円", dict(size=16, bold=True, color=NAVY))),
                                  runs(("イベント期　月30万円程度", dict(size=16, bold=True, color=NAVY)))],
            line=NAVY, align="ctr", anchor="ctr", line_spacing=160, name="固定費")
    heading(s, X0 + cw + gap, 1.85, cw, "成果報酬プラン")
    slot(s, X0 + cw + gap, 2.3, cw, 2.6, "成果の地点・単価")


ORDER = [p_summary, p_kadai, p_honshitsu, p_line, p_iriguchi, p_dankai, p_kaisu, p_mokuhyo, p_mitsumori,
         p_ryokin, p_schedule, p_seika]


def main(base, out):
    prs = Presentation(base)
    slides = list(prs.slides)
    src, cover, back = slides[HDR_SRC - 1], slides[COVER - 1], slides[BACK - 1]
    assert by_id(src, HDR[0]).text_frame.text == "サマリー"
    assert by_id(cover, 26).text_frame.text == "LINE公式アカウント運用のご提案"
    assert any(sh.has_text_frame and sh.text_frame.text.startswith("株式会社DYM") for sh in back.shapes)
    new = []
    for f in ORDER:
        f(prs, src)
        new.append(prs.slides[-1])
    p_cover(cover)
    keep = [cover] + new + [back]
    lst = prs.slides._sldIdLst
    by_part = {prs.part.related_part(e.rId): e for e in lst}
    for e in list(lst):
        lst.remove(e)
    keep_parts = {s.part for s in keep}
    for part, e in by_part.items():
        if part not in keep_parts:
            prs.part.drop_rel(e.rId)
    for s in keep:
        lst.append(by_part[s.part])
    prs.save(out)
    print("saved:", out, len(prs.slides), "枚")


if __name__ == "__main__":
    main(*sys.argv[1:3])
