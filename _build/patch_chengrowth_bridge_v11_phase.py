# -*- coding: utf-8 -*-
"""チェングロウス19枚 ver1.0 → ver1.1（2026-10-05 殿村さん指示）。
「11月頃から初期構築 → 4月のサイト公開後に本格運用・改善」という時間軸に直す。
直すのは 5（なぜ11月頃から）・12（サンクスLINE表記・Indeed注記）・17（費用の考え方）・18（見積もり）・19（スケジュール）。
事例・SIM・市場データ・デザインは触らない。
  python3 _build/patch_chengrowth_bridge_v11_phase.py <ver1.0.pptx>
"""
import copy
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Cm, Pt

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "20261005_チェングロウス_サイトリニューアルとLINE同時導入のご提案ver1.1.pptx"

NAVY, INK, GRAY = "1F285A", "333333", "7F7F7F"
CARD, GREEN, GREEN_BG, GREEN_TX = "F4F7FF", "06C755", "E8F8EE", "0B7A3B"
WHITE = "FFFFFF"

prs = Presentation(sys.argv[1])
S = list(prs.slides)


def rgb(h):
    return RGBColor.from_string(h)


def style_run(r, size, bold=False, color=INK):
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = rgb(color)
    rpr = r._r.get_or_add_rPr()
    for t in ("a:latin", "a:ea", "a:cs"):
        for e in rpr.findall(qn(t)):
            rpr.remove(e)
    for t in ("a:latin", "a:ea", "a:cs"):
        rpr.append(rpr.makeelement(qn(t), {"typeface": "メイリオ"}))


def fill_text(tf, paras, align=PP_ALIGN.LEFT):
    """paras = [(text, size, bold, color, space_after_pt, bullet)]"""
    first = True
    for p_ in paras:
        text, size = p_[0], p_[1]
        bold = p_[2] if len(p_) > 2 else False
        color = p_[3] if len(p_) > 3 else INK
        sa = p_[4] if len(p_) > 4 else 4
        bullet = p_[5] if len(p_) > 5 else False
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.space_after = Pt(sa)
        p.line_spacing = 1.08
        if bullet:
            ppr = p._p.get_or_add_pPr()
            ppr.set("marL", str(Cm(0.42)))
            ppr.set("indent", str(-Cm(0.42)))
            ppr.append(ppr.makeelement(qn("a:buFont"), {"typeface": "Arial"}))
            ppr.append(ppr.makeelement(qn("a:buChar"), {"char": "•"}))
        r = p.add_run()
        r.text = text
        style_run(r, size, bold, color)


def box(s, x, y, w, h, paras, fill=CARD, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, adj=0.06, line=None,
        margins=(0.4, 0.25, 0.4, 0.15)):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Cm(x), Cm(y), Cm(w), Cm(h))
    sh.adjustments[0] = adj
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(fill)
    if line:
        sh.line.color.rgb = rgb(line)
        sh.line.width = Pt(1.25)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    st = sh._element.find(qn("p:style"))
    if st is not None:
        st.find(qn("a:effectRef")).set("idx", "0")
    tf = sh.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = [Cm(m) for m in margins]
    fill_text(tf, paras, align)
    return sh


def label(s, x, y, w, h, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Cm(0.1)
    tf.margin_top = tf.margin_bottom = Cm(0.05)
    fill_text(tf, paras, align)
    return tb


def arrow(s, x, y, w, h, color=NAVY):
    a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Cm(x), Cm(y), Cm(w), Cm(h))
    a.fill.solid()
    a.fill.fore_color.rgb = rgb(color)
    a.line.fill.background()
    a.shadow.inherit = False
    st = a._element.find(qn("p:style"))
    if st is not None:
        st.find(qn("a:effectRef")).set("idx", "0")
    return a


def band(s, y, text, h=1.2):
    return box(s, 2.2, y, 23.11, h, [(text, 14, True, WHITE, 0)], fill=NAVY, align=PP_ALIGN.CENTER,
               anchor=MSO_ANCHOR.MIDDLE, adj=0.1, margins=(0.4, 0, 0.4, 0))


def find(slide, name):
    return next(sh for sh in slide.shapes if sh.name == name)


def set_para(sh, k, text):
    while len(sh.text_frame.paragraphs) <= k:
        last = sh.text_frame.paragraphs[-1]._p
        last.addnext(copy.deepcopy(last))
    p = sh.text_frame.paragraphs[k]
    runs = p.runs
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def drop_para(sh, k):
    p = sh.text_frame.paragraphs[k]
    p._p.getparent().remove(p._p)


def retitle(slide, title, lead1, lead2=None):
    set_para(find(slide, "TextBox 1"), 0, title)
    t2 = find(slide, "TextBox 2")
    set_para(t2, 0, lead1)
    if lead2 is None:
        if len(t2.text_frame.paragraphs) > 1:
            drop_para(t2, 1)
    else:
        set_para(t2, 1, lead2)


def keep_header_only(slide):
    for sh in list(slide.shapes):
        if sh.name not in ("TextBox 1", "TextBox 2", "Connector 3"):
            sh._element.getparent().remove(sh._element)


def replace_par(slide, old, new):
    n = 0
    for sh in slide.shapes:
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                if p.text == old:
                    p.runs[0].text = new
                    for r in p.runs[1:]:
                        r._r.getparent().remove(r._r)
                    n += 1
    assert n, old


def group_card(s, x, y, w, h, title, items, fill=CARD, tcolor=NAVY, size=11):
    paras = [(title, 12.5, True, tcolor, 3)] + [(t, size, False, INK, 1.5, True) for t in items]
    return box(s, x, y, w, h, paras, fill=fill, margins=(0.4, 0.2, 0.3, 0.1))


# ================= 5 なぜ4月を待たず、11月頃から初期構築を進めるのか =================
s5 = S[4]
keep_header_only(s5)
retitle(s5, "なぜ4月を待たず、11月頃から初期構築を進めるのか",
        "サイト制作の段階から、LINE連携を組み込む",
        "4月の公開時点で、すぐ本格運用を始められる状態をつくる")
cols5 = [
    ("① 制作の段階からLINE連携を組み込める",
     "公開後にLINEを足すのではなく、制作段階から連携を要件として整理",
     ["LINE Profile+／「LINEで登録」", "会員情報とLINEの連携", "サンクスLINE誘導", "条件別検索ページとの連携", "計測設計"],
     "公開後に大きな改修をするリスクを減らせる"),
    ("② 公開前に申請・設計・実装・検証を終えられる",
     "11月頃から進める内容",
     ["アカウント・リッチメニュー・配信シナリオの設計", "Profile+の申請／会員連携の仕様整理", "サイト制作会社との仕様調整", "実装・動作確認"],
     "公開後にゼロから構築せず、公開時点ですぐ動かせる"),
    ("③ 初期構築から本格運用へスムーズに移行できる",
     "初期構築のうちに、運用に必要なことまで設計",
     ["どの導線から友だちを獲得するか", "どの情報を取得するか", "どの条件でユーザーを分けるか", "どのような配信を行うか", "何をKPIとして計測するか"],
     "公開後に運用設計をやり直す必要がない"),
]
for i, (h, sub, items, res) in enumerate(cols5):
    x = 2.2 + i * 7.9
    paras = [(h, 13, True, NAVY, 4), (sub, 10.5, False, GRAY, 5)] + [(t, 11, False, INK, 1.5, True) for t in items] + \
            [(res, 11, True, GREEN_TX, 0)]
    box(s5, x, 4.6, 7.3, 8.0, paras, margins=(0.4, 0.3, 0.3, 0.2))
    # 結果行は下に寄せるため、上の枠の上に重ねず同じ枠内に置く
box(s5, 2.2, 12.95, 10.9, 1.9, [("11月頃〜3月｜初期構築・実装準備", 13, True, WHITE, 3),
                               ("運用を見据え、導線・会員連携・配信・計測まで設計", 11, False, WHITE, 0)],
    fill=NAVY, anchor=MSO_ANCHOR.MIDDLE, margins=(0.4, 0.1, 0.3, 0.1))
arrow(s5, 13.2, 13.45, 0.8, 0.9)
box(s5, 14.2, 12.95, 11.1, 1.9, [("4月〜｜サイト公開・本格運用・改善", 13, True, WHITE, 3),
                                ("設計した内容を引き継ぎ、配信・分析・改善を継続", 11, False, WHITE, 0)],
    fill=GREEN, anchor=MSO_ANCHOR.MIDDLE, margins=(0.4, 0.1, 0.3, 0.1))
box(s5, 2.2, 15.2, 23.11, 1.3, [("初期構築から運用・改善まで一貫して支援できるため、サイト公開後もスムーズに運用へ移行できます", 13, True, WHITE, 0)], fill=NAVY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, adj=0.1, margins=(0.3, 0, 0.3, 0))

# ================= 12 サンクスLINEの表記・求人媒体の注記 =================
s12 = S[11]
for sh in s12.shapes:
    if sh.has_text_frame and sh.text_frame.text.startswith("サンクスLINE"):
        set_para(sh, 0, "応募完了後の")
        set_para(sh, 1, "LINE誘導")
        break
else:
    raise SystemExit("サンクスLINEの枠が見つからない")
n12 = next(sh for sh in s12.shapes if sh.has_text_frame and sh.text_frame.text.startswith("※キープ画面"))
set_para(n12, 1, "※求人ボックス・Indeed経由については、着地先・応募完了先を確認したうえでLINE導線を設計")

# ================= 17 費用の考え方（初期＝11月頃から／月額＝4月以降）=================
s17 = S[16]
for sh in s17.shapes:
    if sh.has_text_frame:
        for p in sh.text_frame.paragraphs:
            if p.text.startswith("初期費用＋月額の固定費"):
                p.runs[0].text = ("初期費用（11月頃〜）＋月額費用（4月〜）／"
                                  "LINE公式アカウントの利用料は配信通数で変動（別途）／LINE連携の開発は都度お見積もり")
                for r in p.runs[1:]:
                    r._r.getparent().remove(r._r)

# ================= 18 概算お見積もり =================
s18 = S[17]
retitle(s18, "概算お見積もり",
        "初期構築から運用・改善まで、一貫して支援するための費用",
        "初期費用は11月頃からの構築、月額費用は4月以降の運用に対応")
tbl = next(sh for sh in s18.shapes if sh.has_table).table
for ri, row in enumerate(tbl.rows):
    t0 = row.cells[0].text
    if t0.startswith("サンクスLINE"):
        c0 = row.cells[0]
        c0.text_frame.paragraphs[0].runs[0].text = "サンクスLINE誘導"
        for r in c0.text_frame.paragraphs[0].runs[1:]:
            r._r.getparent().remove(r._r)
        c1 = row.cells[1]
        for r in row.cells[2].text_frame.paragraphs[0].runs:
            r._r.getparent().remove(r._r)
        c1.merge(row.cells[2])
        c1.text_frame.paragraphs[0].runs[0].text = "実装内容に応じて別途お見積もり"
        for r in c1.text_frame.paragraphs[0].runs[1:]:
            r._r.getparent().remove(r._r)
        break
else:
    raise SystemExit("サンクスLINEの行が見つからない")
box(s18, 1.2, 13.95, 12.4, 2.2, [("初期費用｜11月頃〜3月（LINE運用コンサル 20万円）", 12, True, NAVY, 3),
                                ("設計・初期構築・実装準備・サイト連携の仕様整理・配信シナリオ設計", 11, False, INK, 0)],
    margins=(0.4, 0.25, 0.3, 0.1))
box(s18, 13.9, 13.95, 12.4, 2.2, [("月額費用｜4月以降（LINE運用コンサル 20万円）", 12, True, NAVY, 3),
                                 ("配信・分析・改善・レポーティング・継続的な運用支援", 11, False, INK, 0)],
    margins=(0.4, 0.25, 0.3, 0.1))
bnd = next(sh for sh in s18.shapes if sh.has_text_frame and sh.text_frame.text.startswith("会員登録から応募まで"))
bnd.top, bnd.height = Cm(16.35), Cm(1.0)
set_para(bnd, 0, "初期構築から運用・改善まで、会員登録から応募までをLINEでつなげる仕組みを一貫して支援")

# ================= 19 スケジュール =================
s19 = S[18]
keep_header_only(s19)
retitle(s19, "スケジュール｜11月頃から初期構築、4月の公開と同時に本格運用へ",
        "11月頃から初期構築に入り、4月のサイト公開後すぐに本格運用へ移行")
LX, RX, CW = 2.2, 14.3, 11.0
box(s19, LX, 4.4, CW, 1.7, [("Phase 1｜11月頃〜3月", 14, True, WHITE, 2), ("初期構築・実装準備　運用を見据えて設計", 11.5, False, WHITE, 0)],
    fill=NAVY, anchor=MSO_ANCHOR.MIDDLE, margins=(0.4, 0.1, 0.3, 0.1))
arrow(s19, 13.3, 4.8, 0.9, 0.9)
box(s19, RX, 4.4, CW, 1.7, [("Phase 2｜4月〜", 14, True, WHITE, 2), ("サイト公開・本格運用・改善　実データで継続", 11.5, False, WHITE, 0)],
    fill=GREEN, anchor=MSO_ANCHOR.MIDDLE, margins=(0.4, 0.1, 0.3, 0.1))
rows = [  # (高さ, Phase1, Phase2)
    (2.6, ("設計", ["LINE公式アカウントの設計", "リッチメニュー・あいさつメッセージの設計", "ステップ配信シナリオ・計測の設計"]),
     ("公開・獲得", ["リニューアルサイト公開", "「LINEで登録」の本稼働", "友だち・会員の獲得"])),
    (3.6, ("サイト連携の設計", ["LINE Profile+の申請", "「LINEで登録」・会員連携の仕様整理", "サンクスLINE誘導の仕様整理",
                          "条件別検索ページとの連携設計", "サイト制作会社との仕様調整"]),
     ("配信・運用", ["ステップ配信・企画配信", "リッチメニュー運用・条件別配信", "応募への追客", "会員登録・応募のKPI計測"])),
    (2.05, ("実装・検証", ["開発・実装・公開前の動作検証"]),
     ("分析・改善", ["月次分析・レポーティング", "導線・配信内容・セグメントの改善"])),
]
y = 6.3
for h, (t1, i1), (t2, i2) in rows:
    group_card(s19, LX, y, CW, h, t1, i1, size=10.5)
    group_card(s19, RX, y, CW, h, t2, i2, fill=GREEN_BG, tcolor=GREEN_TX, size=10.5)
    y += h + 0.12
band(s19, y + 0.1, "初期構築で設計した導線・配信・計測を引き継ぎ、4月から本格運用・改善へスムーズに移行", h=1.0)
label(s19, 2.2, y + 1.2, 23.1, 1.6, [
    ("※11月頃〜3月は、定期配信・月次分析・改善運用は原則行わず、設計・構築に集中。必要に応じて一部施策の先行稼働も検討可能（先行稼働分は、シミュレーションに含めない）", 9, False, GRAY, 1),
    ("※LINE Profile+の審査期間は、申請内容により変わります／実装の時期は、サイト制作の進行に応じて決定", 9, False, GRAY, 1),
    ("※求人ボックス・Indeed経由については、着地先・応募完了先を確認したうえでLINE導線を設計", 9, False, GRAY, 0)])

# ---------------- 字体メイリオ・最小8pt ----------------
AFTER = {qn(x) for x in ("a:sym", "a:hlinkClick", "a:hlinkMouseOver", "a:rtl", "a:extLst")}
for sl in prs.slides:
    for tag in ("a:rPr", "a:endParaRPr", "a:defRPr"):
        for rpr in sl.shapes._spTree.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}" + tag[2:]):
            if rpr.get("sz") and int(rpr.get("sz")) < 800:
                rpr.set("sz", "800")
            fs = []
            for ft in ("a:latin", "a:ea", "a:cs"):
                e = rpr.find(qn(ft))
                if e is None:
                    e = rpr.makeelement(qn(ft), {})
                else:
                    rpr.remove(e)
                e.set("typeface", "メイリオ")
                fs.append(e)
            nxt = next((c for c in rpr if c.tag in AFTER), None)
            for e in fs:
                (nxt.addprevious(e) if nxt is not None else rpr.append(e))
prs.save(str(OUT))
print("saved:", OUT.name, len(prs.slides), "枚")
