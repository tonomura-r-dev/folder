# -*- coding: utf-8 -*-
"""チェングロウス ver2.8 → ver2.9（2026-10-02）
「4月のリニューアルと同時に開始」前提をやめ、LINEの初期構築を先に進め、サイト連携は実装を進めてつなぐ形に直す。
  先に進める：アカウント設定／Profile+申請／友だち追加の導線／リッチメニュー／あいさつ／配信設計／サンクスLINE誘導（完了画面を触れる経路のみ）
  実装を進める：会員登録のLINE化（現行サイトに実装できればリニューアルを待たない。リニューアル時は新サイトに組み込む）
費用と体制（25枚目）は触らない（固定費の金額と内訳が決まるまで）。
  python3 _build/patch_chengrowth_v29_initial_build.py <ver2.8.pptx>
"""
import sys
from pathlib import Path
from pptx import Presentation

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "20260929_チェングロウス_LINE公式アカウント運用のご提案ver2.9.pptx"
prs = Presentation(sys.argv[1])


def walk(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type == 6:
            yield from walk(sh.shapes)


def shape(n, name):
    return next(sh for sh in walk(prs.slides[n - 1].shapes) if sh.name == name)


def set_par(n, name, k, text):
    """段落kの文字を差し替える（最初のrunの書式を残し、他のrunは消す）"""
    p = shape(n, name).text_frame.paragraphs[k]
    runs = p.runs
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def sub_runs(n, name, old, new):
    """run内の部分置換（書式を保つ）。run内に無ければ失敗"""
    for p in shape(n, name).text_frame.paragraphs:
        for r in p.runs:
            if old in r.text:
                r.text = r.text.replace(old, new)
                return
    raise SystemExit(f"not found: {n} {name} {old}")


# 1 表紙
set_par(1, "正方形/長方形 5", 0, "―LINEの初期構築と、サイト連携のご提案―")
# 11 タウンワークの型
set_par(11, "Rounded Rectangle 13", 0, "LINE側は先に構築し、サイト側は実装を進めてつなぐ")
# 12 貴社に当てはめると
set_par(12, "TextBox 2", 0, "会員登録をLINEで完結できる形へ")
set_par(12, "Rounded Rectangle 7", 0, "実装後")
# 15 友だち追加の動線
set_par(15, "Rounded Rectangle 11", 0, "応募の完了画面（貴社が触れる経路）")
set_par(15, "TextBox 13", 0,
        "※月約17人＝サイトからの応募（月約47件）×完了画面の表示80%×友だち追加45%の想定／求人ボックス・Indeed経由は、完了画面の扱いを確認のうえ決定／画面はイメージです")
# 16 クリック数の注記
sub_runs(16, "TextBox 8", "9月の値", "6か月目の値")
# 20 先に進める理由
set_par(20, "TextBox 1", 0, "LINEの構築を先に進めることをおすすめする理由")
set_par(20, "TextBox 2", 1, "先に始めれば、翌3月の転職のピークを、増えた友だちとともに抑制")
set_par(20, "Rounded Rectangle 5", 1, "会員登録のLINE化を")
set_par(20, "Rounded Rectangle 5", 2, "リニューアルの要件に入れておけば、")
set_par(20, "Rounded Rectangle 5", 3, "サイトの改修が二度手間にならない。")
set_par(20, "Rounded Rectangle 7", 0, "先に開始")
set_par(20, "TextBox 9", 0, "開始 ────── 友だちが増え続ける ────── 翌3月（転職のピーク）")
set_par(20, "Rounded Rectangle 14", 0, "LINEの構築を先に進め、サイトとの連携は実装を進めてつなぐことをご提案します。")
# 21 要件3点
set_par(21, "TextBox 1", 0, "サイトに入れる要件3点")
set_par(21, "Rounded Rectangle 10", 2, "サイト制作会社様と仕様をすり合わせて進めます。現行サイトに実装できる場合は、リニューアルを待たず先行します。")
# 22 スケジュール
set_par(22, "TextBox 1", 0, "スケジュール｜初期構築を先に進め、サイトと連携")
set_par(22, "Rounded Rectangle 6", 0, "決定後")
set_par(22, "Rounded Rectangle 7", 1, "アカウント設定、LINE Profile+の申請、サイト制作会社様と仕様のすり合わせ")
set_par(22, "Rounded Rectangle 8", 0, "申請と並行")
set_par(22, "Rounded Rectangle 9", 1, "あいさつ・リッチメニュー・配信の設計、友だち追加の導線、サンクスLINE誘導（完了画面を触れる経路のみ）")
set_par(22, "Rounded Rectangle 10", 0, "実装")
set_par(22, "Rounded Rectangle 11", 0, "サイトへの実装と運用開始")
set_par(22, "Rounded Rectangle 11", 1, "会員登録のLINE化を実装。現行サイトで可能な範囲から先行し、リニューアル時は新サイトに組み込む")
set_par(22, "TextBox 14", 0, "※LINE Profile+の審査期間は、申請内容により変わります／現行サイトへの実装可否、求人ボックス・Indeed経由の応募の扱いは、確認のうえ決定")
# 24 シミュレーション
set_par(24, "TextBox 1", 0, "成果シミュレーション（運用開始から6か月）")
set_par(24, "TextBox 2", 1, "友だちが増えるほど応募も増え、4か月目に応募単価が2.5万円を下回る見込み")
set_par(24, "Rounded Rectangle 7", 0, "6か月目の応募単価")
tbl = shape(24, "Table 4").table
for j, m in enumerate(["4月", "5月", "6月", "7月", "8月", "9月"], start=1):
    c = tbl.cell(0, j)
    assert c.text.strip() == m, c.text
    set_par_c = c.text_frame.paragraphs[0].runs
    set_par_c[0].text = f"{j}か月目"
    for r in set_par_c[1:]:
        r._r.getparent().remove(r._r)
sub_runs(24, "TextBox 8", "前提：", "前提：会員登録のLINE化の開始月を1か月目とする／")
prs.save(str(OUT))
print("saved:", OUT.name, len(prs.slides), "枚")
