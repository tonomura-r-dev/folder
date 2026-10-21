"""
10月号（正）に、2026-10-21受領の追加資料（P.8 ユニット統計／P.10 レストラン価格／P.11 CRM事例）を反映する。
使い方: python3 _build/patch_trend_202610_pdf2.py <pptx>
"""
import sys
from pptx import Presentation
import trend_report_engine as engine

F = sys.argv[1]
prs = Presentation(F)
slides = list(prs.slides)


def find(slide, first_line):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.paragraphs[0].text.strip() == first_line:
            return sh
    raise SystemExit(f"見つからない: {first_line}")


def set_paras(sh, lines):
    paras = sh.text_frame.paragraphs
    assert len(paras) == len(lines), (sh.text_frame.text, lines)
    for p, line in zip(paras, lines):
        p.runs[0].text = line
        for r in list(p.runs[1:]):
            r._r.getparent().remove(r._r)


def desc(slide, lines):
    for sh in slide.shapes:
        if sh.has_text_frame and len(sh.text_frame.paragraphs) == 3 \
                and sh.text_frame.paragraphs[0].alignment == engine.PP_ALIGN.CENTER:
            engine.set_multiline_text(sh, lines)
            return
    raise SystemExit("説明バンドが見つからない")


def card(slide, old_head, head, body):
    set_paras(find(slide, old_head), [head, body])



# P.8 ユニット単位の統計
s = slides[7]
desc(s, [
    "2026年9月15日（火）から、ユニットごとの統計情報を取得するAPIが",
    "LINE通知メッセージでも利用できるようになった。",
    "アンケートで要望の多かった機能。開封数・URLタップ数などをユニットごとに確認できる。",
])
set_paras(find(s, "LINE通知メッセージではユニット単位の統計を取得できない"),
          ["プライバシー保護のため、開封・タップの統計は取得できない"])
set_paras(find(s, "ユニット単位の統計情報をAPIで取得可能に"),
          ["ユニット単位の統計情報をAPIで取得可能に",
           "開封・URLタップ・動画再生の人数をユニットごとに確認"])
card(s, "要望の多かった機能", "送信時にユニット名を付ける",
     "送信時にユニット名（1通に1つ）を指定して集計する。後から付けたり変えたりはできない。月1,000種類まで。")
card(s, "通知ごとの反応を比較できる", "少ない数は表示されない",
     "値が20未満（実人数20人未満）の項目は表示されない。送信数は含まれないため、開封率の分母は別で持っておく。")

# P.10 レストランオプション
s = slides[9]
card(s, "価格", "価格",
     "プリンター：37,000円→41,700円／キャッシュドロアー：16,000円→17,700円。")
card(s, "改定時期", "改定時期",
     "2027年2月上旬予定。改定日以降に購入した機材から新価格が適用される。")

# P.11 CRMオプション事例
s = slides[10]
desc(s, [
    "CRMオプションの導入事例（タオル美術館様）が、媒体資料Ver2.0に追加された。",
    "誕生月を聞くフォームで回答数が約2.0倍、",
    "好きなキャラクターに合わせた出し分け配信でクリック率が約2.8倍に。",
])
card(s, "① アンケート回答率UP", "① フォーム回答数 約2.0倍",
     "リサーチ機能から、誕生月を聞くフォーム（誕生月クーポン付き）に変更。店頭チラシ・リッチメニュー・あいさつメッセージで回答を促した。")
card(s, "② 配信メッセージのCTR改善", "② 配信のクリック率 約2.8倍",
     "好きなキャラクターをフォームで聞いてタグを付け、コラボ商品の配信を出し分け。従来の一斉配信と比べてクリック率が約2.8倍。")

engine.force_font(prs)
prs.save(F)
print("saved", F)
