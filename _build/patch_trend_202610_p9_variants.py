"""
10月号 P.9 の見せ方を3案つくる（2026-10-26 殿村さん「複数案出して」）。
入力は数字3つ＋対象条件カード＋注意点カードの状態（patch_trend_202610_p9.py 適用後）。
使い方: python3 _build/patch_trend_202610_p9_variants.py <入力pptx> <出力pptx> <A|B|C>
  A = 左に数字を縦に3つ／右に対象条件の表＋注意点カード
  B = 上に数字3つ／下の左に対象条件の表、右に注意点カード
  C = 上に数字3つ／下に表1つ（対象条件と注意点をまとめる）
"""
import sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
import trend_report_engine as engine

SRC, OUT, V = sys.argv[1], sys.argv[2], sys.argv[3]
prs = Presentation(SRC)
s9 = list(prs.slides)[8]
BOTTOM = 7.08


def text(sh):
    return sh.text_frame.text.strip() if sh.has_text_frame else ""


head = next(x for x in s9.shapes if text(x) == "■ 詳細")
tiles = sorted([x for x in s9.shapes if text(x).startswith(("6円", "9円", "2P"))], key=lambda x: x.left)
cond = next(x for x in s9.shapes if text(x).startswith("対象条件"))
caution = next(x for x in s9.shapes if text(x).startswith("注意点"))
cond._element.getparent().remove(cond._element)
area_top = Emu(head.top + head.height).inches + 0.08
L, W = 0.6, 9.64

COND = [
    ("対象期間", "2026/12/1（火）〜2027/3/31（水）", "※2027/3/31までに掲載が終わる案件"),
    ("対象メニュー", "Friends Swipe AD／Friends Video AD", "※リーチ最適化配信を含む。デリタゲ配信は対象外"),
    ("対象の配信", "ALL配信のみ", None),
    ("対象クライアント", "全クライアント", None),
    ("最低発注金額", "200万円〜", None),
]
NOTES = ["発注時、備考欄に「FY26 Winter Campaign適用」と記載",
         "配信中・発注受領済みの案件には適用できない",
         "他の割引・キャンペーンとは併用できない",
         "記載のない条件は、最新の媒体資料に準拠"]


def put(cell, lines, size, bold, color, fill, align):
    cell.fill.solid()
    cell.fill.fore_color.rgb = fill
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = cell.margin_right = Inches(0.12)
    cell.margin_top = cell.margin_bottom = Inches(0.02)
    tf = cell.text_frame
    for k, line in enumerate(lines):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line
        sub = line.startswith("※")
        r.font.size = Pt(size - 1 if sub else size)
        r.font.bold, r.font.color.rgb = bold, color


def row_h(ls, size):
    return 0.12 + 0.19 * len(ls) * size / 12


def table(x, y, w, labw, rows, size=12, one_line=False):
    """rows: [(label, [lines])]"""
    hs = [row_h(ls, size) for _, ls in rows]
    gf = s9.shapes.add_table(len(rows), 2, Inches(x), Inches(y), Inches(w), Inches(sum(hs)))
    t = gf.table
    t.columns[0].width, t.columns[1].width = Inches(labw), Inches(w - labw)
    for i, ((lab, ls), h) in enumerate(zip(rows, hs)):
        t.rows[i].height = Inches(h)
        put(t.cell(i, 0), [lab], size, True, engine.NAVY, engine.LIGHT_BLUE, PP_ALIGN.CENTER)
        put(t.cell(i, 1), ls, size, False, engine.BODY_GRAY, engine.WHITE, PP_ALIGN.LEFT)
    return sum(hs)


def set_caution(x, y, w, h, size):
    tf = caution.text_frame
    for p, t in zip(tf.paragraphs[1:], NOTES):
        for r in p.runs[1:]:
            r._r.getparent().remove(r._r)
        for br in p._p.findall(engine.qn("a:br")) if hasattr(engine, "qn") else []:
            br.getparent().remove(br)
        p.runs[0].text = "・" + t
        p.runs[0].font.size = Pt(size)
    caution.left, caution.top, caution.width, caution.height = Inches(x), Inches(y), Inches(w), Inches(h)


def style_tiles(n, l, s):
    for t in tiles:
        for r in t.text_frame.paragraphs[0].runs:
            r.font.size = Pt(n)
        for r in t.text_frame.paragraphs[1].runs:
            r.font.size = Pt(l)
        for r in t.text_frame.paragraphs[2].runs:
            r.font.size = Pt(s)


def tiles_row(y, h):
    tw = (W - 0.4) / 3
    for i, t in enumerate(tiles):
        t.left, t.top, t.width, t.height = Inches(L + i * (tw + 0.2)), Inches(y), Inches(tw), Inches(h)


if V == "A":
    LW, GAP = 2.8, 0.25
    RX, RW = L + LW + GAP, W - LW - GAP
    rows = [(k, [a] + ([b] if b else [])) for k, a, b in COND]
    est = sum(row_h(ls, 11) for _, ls in rows)
    card_h = 0.28 + 0.33 + 4 * 0.2
    total = est + 0.12 + card_h
    top = area_top + (BOTTOM - area_top - total) / 2
    table(RX, top, RW, 1.75, rows, 11)
    set_caution(RX, top + est + 0.12, RW, card_h, 11)
    th = (total - 0.3) / 3
    style_tiles(28, 13, 10)
    for i, t in enumerate(tiles):
        t.left, t.top, t.width, t.height = Inches(L), Inches(top + i * (th + 0.15)), Inches(LW), Inches(th)
elif V == "B":
    TH = 0.85
    rows = [(k, [a] + ([b] if b else [])) for k, a, b in COND]
    est = sum(row_h(ls, 11) for _, ls in rows)
    total = TH + 0.15 + est
    top = area_top + (BOTTOM - area_top - total) / 2
    style_tiles(24, 12, 9)
    tiles_row(top, TH)
    LWT = 5.3
    table(L, top + TH + 0.15, LWT, 1.65, rows, 11)
    set_caution(L + LWT + 0.2, top + TH + 0.15, W - LWT - 0.2, est, 11)
else:  # C
    TH = 0.8
    rows = [(k, [a + ("　" + b if b else "")]) for k, a, b in COND] + [("注意点", ["・" + t for t in NOTES])]
    est = sum(row_h(ls, 12) for _, ls in rows)
    total = TH + 0.15 + est
    top = area_top + (BOTTOM - area_top - total) / 2
    style_tiles(24, 12, 9)
    tiles_row(top, TH)
    table(L, top + TH + 0.15, W, 1.75, rows, 12)
    caution._element.getparent().remove(caution._element)

engine.force_font(prs)
prs.save(OUT)
print("saved", OUT)
