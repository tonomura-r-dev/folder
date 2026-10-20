# -*- coding: utf-8 -*-
"""すなお：想定動線（ユーザー動線の図）の2点を直す（2026-10-20 殿村さん）。

    python -m _build.patch_sunao_s3_doshen <資料.pptx> <出力.pptx>

  1. 上の8フェーズ：③ 育成（興味喚起）が ②・④ の箱に隠れて見えない。
     8つの箱の高さが段の間隔より大きく重なっていたので、段の間隔＝箱の高さ（0.26インチ）にそろえて重ならないようにする。
     ④→⑤ の点線の矢印は、④の上から⑤の下へまっすぐ上がる線に引き直す（③の右端を通らない位置）。
  2. 下の「・またサービス特性として」「・貴社サービスとして」：2つあって違いが分からない（殿村さん）。
     「貴社の強み」と「LINEで補えるところ」の左右2列にし、下に結論を1行。中身（数字・事実）は元の文のまま。
  あわせて「<貴社のユーザー動線イメージ>」が2行に折れていたので幅を広げる。ほかのページは触らない。
"""
import sys

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from _build.arauchi_v4.lib import add_box
from _build.arauchi_v10.build import runs
from _build.build_sunao_v4 import NAVY, TEXT

ROW_Y, ROW_H = 2.20, 0.26
PHASE_ROW = {"①": 0, "②": 1, "③": 2, "④": 3, "⑤": 0, "⑥": 1, "⑦": 2, "⑧": 3}
STRONG = ["広告で、1回のイベントで81件のセミナー予約を集めている",
          "セミナーの参加後に、サロンの入会を案内できる",
          "対象は保険営業の方（営業力向上のセミナー・サロン）"]
LINE_HELP = ["イベントは年4回。間の時期は、集めた方との接点が切れやすい",
             "セミナーは無料・オンライン。金額を考えるのはサロン入会のときで、参加後に迷いやすい"]
CONCLUSION = "集めた方と、次のイベント・入会まで接点を持ち続け、強みを入会数につなげる。"


def find_slide(prs):
    hit = [s for s in prs.slides
           if any(sh.has_text_frame and "貴社のユーザー動線イメージ" in sh.text_frame.text for sh in s.shapes)]
    assert len(hit) == 1, len(hit)
    return hit[0]


def p_phases(s):
    boxes = {}
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip()[:1] in PHASE_ROW and Inches(2.0) < sh.top < Inches(3.4):
            boxes[sh.text_frame.text.strip()[:1]] = sh
    assert len(boxes) == 8, sorted(boxes)
    for k, sh in boxes.items():
        sh.top = Inches(ROW_Y + ROW_H * PHASE_ROW[k])
        sh.height = Inches(ROW_H)
        bp = sh.text_frame._txBody.find(qn("a:bodyPr"))
        for e in list(bp):
            if e.tag in (qn("a:spAutoFit"), qn("a:normAutofit"), qn("a:noAutofit")):
                bp.remove(e)
        bp.set("tIns", "0")
        bp.set("bIns", "0")
        bp.set("anchor", "ctr")
        for p in sh.text_frame.paragraphs[1:]:                     # ⑧の空の2段落目
            if not p.text.strip():
                p._p.getparent().remove(p._p)
    # ③を前面に（念のため）：②・④より後ろに並べ替え
    tree = s.shapes._spTree
    tree.remove(boxes["③"]._element)
    tree.append(boxes["③"]._element)

    # ④→⑤ の点線：④の上（③の右端より右）から⑤の下へまっすぐ
    old = next(sh for sh in s.shapes if sh.shape_type == 9 and sh.name.startswith("コネクタ: カギ線"))
    x = Emu(boxes["⑤"].left).inches + 1.0                           # ⑤の中（③の右端4.45より右・⑥の左端より左）
    y_top = ROW_Y + ROW_H                                            # ⑤の下
    y_bot = ROW_Y + ROW_H * 3                                        # ④の上
    ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y_bot), Inches(x), Inches(y_top))
    src_ln = old._element.find(qn("p:spPr")).find(qn("a:ln"))
    sp = ln._element.find(qn("p:spPr"))
    for e in sp.findall(qn("a:ln")):
        sp.remove(e)
    sp.append(etree.fromstring(etree.tostring(src_ln)))
    old._element.getparent().remove(old._element)


def p_label(s):
    lab = next(sh for sh in s.shapes if sh.has_text_frame and "貴社のユーザー動線イメージ" in sh.text_frame.text)
    lab.width = Inches(3.2)
    lab.text_frame.paragraphs[0].alignment = 1                       # 左寄せ


def p_service(s):
    box = next(sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.startswith("・またサービス特性として"))
    tf = box.text_frame
    for p in tf.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    for r in tf.paragraphs[0].runs:
        r._r.getparent().remove(r._r)
    x0, y0 = Emu(box.left).inches, Emu(box.top).inches
    w0, h0 = Emu(box.width).inches, Emu(box.height).inches
    col_w = (w0 - 0.5) / 2
    xs = (x0 + 0.2, x0 + 0.3 + col_w)
    for x, head, items in ((xs[0], "貴社の強み", STRONG), (xs[1], "LINEで補えるところ", LINE_HELP)):
        add_box(s, x, y0 + 0.12, col_w, 0.3, [runs((head, dict(size=13, bold=True, color=NAVY)))],
                anchor="ctr", ins=(0, 0, 0, 0), name=head)
        add_box(s, x, y0 + 0.45, col_w, 1.05, [runs(("・" + t, dict(size=12, color=TEXT))) for t in items],
                anchor="t", ins=(0, 0, 0, 0), line_spacing=110, name=head + " 内容")
    div = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x0 + 0.15 + col_w + 0.1), Inches(y0 + 0.15),
                                 Inches(x0 + 0.15 + col_w + 0.1), Inches(y0 + 1.45))
    div.line.color.rgb = RGBColor.from_string("BFBFBF")
    div.line.width = Pt(0.75)
    sep = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x0 + 0.2), Inches(y0 + 1.55),
                                 Inches(x0 + w0 - 0.2), Inches(y0 + 1.55))
    sep.line.color.rgb = RGBColor.from_string("BFBFBF")
    sep.line.width = Pt(0.75)
    add_box(s, x0 + 0.2, y0 + 1.6, w0 - 0.4, h0 - 1.65, [runs((CONCLUSION, dict(size=13, bold=True, color=NAVY)))],
            anchor="ctr", align="ctr", ins=(0, 0, 0, 0), name="結論")


def main(src, out):
    prs = Presentation(src)
    s = find_slide(prs)
    p_phases(s)
    p_label(s)
    p_service(s)
    prs.save(out)
    print("saved:", out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
