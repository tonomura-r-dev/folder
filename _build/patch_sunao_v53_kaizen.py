# -*- coding: utf-8 -*-
"""すなお：伝わらないページ4枚の作り直し（2026-10-20 殿村さん「何が言いたいのか伝えきれてない」）。

    python -m _build.patch_sunao_v53_kaizen <資料.pptx> <出力.pptx>

  1スライド1メッセージ。型はJ.P.Returnsのお手本（リード1行 → 表 or KEY TAKEAWAY → 示唆／注記）。中身の数字・事実は元のまま。
  ・ニーズ調査（業界STUDY）：営業職員数の棒グラフを外す（言っているのは「横ばい」だけ）。
    3つの数字を「ファクト／意味すること」の表にし、下に貴社への示唆（LPの「週2日で成果を出す面談メソッド」とつなぐ）。
  ・ニーズ調査（業界STUDY）＞訴求軸：左だけ画像・真ん中が空のバランスを直す。
    左＝他社2社の表＋その下に「貴社の訴求軸」、右＝野沢塾の画面。
  ・要件定義：詰まりすぎ。目的／配信目標／手法（ご提案）／＋αの施策の4行の表にし、手法の補足（細かい打ち手）は本質課題のページへ。
  ・本質課題（元の「要件定義」2枚目）：KEY TAKEAWAY＋「段階／いま起きていること／打ち手／伸ばす数字」の表。
    「平常期の配信で次のイベントまでつながる」は 02 追う の打ち手に移した（元は 03 つなげる にあり、課題とずれていた）。
  タイトル以外の図形は作り直す。ほかのページは触らない。
"""
import copy
import io
import sys

from pptx import Presentation
from pptx.util import Inches

from _build.arauchi_v4.lib import add_box, add_table
from _build.arauchi_v10.build import runs

NAVY = "1F2A5C"
TEXT = "333333"
GRAY = "7F7F7F"
TAKE = "EEF1F8"
BAND = ("FFFFFF", "F5F7FA")
LINE = "D9D9D9"


def title_of(s):
    for sh in s.shapes:
        if sh.has_text_frame and sh.top < Inches(0.6) and sh.text_frame.text.strip():
            return sh
    return None


def find(prs, pred):
    hit = [s for s in prs.slides if pred(s)]
    assert len(hit) == 1, len(hit)
    return hit[0]


def text_of(s):
    return " ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame)


def clear(s):
    t = title_of(s)
    for sh in list(s.shapes):
        if sh.shape_id != t.shape_id:
            sh._element.getparent().remove(sh._element)
    return t


def set_one(sh, text):
    tf = sh.text_frame
    for p in tf.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    rs = tf.paragraphs[0].runs
    rs[0].text = text
    for r in rs[1:]:
        r._r.getparent().remove(r._r)


def lead_band(dst, src, text):
    """想定動線のページのリード（白い帯・区切り線・16ptの文字）を写して、リードを1行入れる"""
    for sid in (3, 4, 2, 9):
        el = copy.deepcopy(next(sh for sh in src.shapes if sh.shape_id == sid)._element)
        nid = max(sh.shape_id for sh in dst.shapes) + 1
        el.find(".//{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr").set("id", str(nid))
        dst.shapes._spTree.append(el)
    set_one(next(sh for sh in dst.shapes if sh.shape_id == max(x.shape_id for x in dst.shapes)), text)


def cell(text, **kw):
    d = dict(text=text, size=13, color=TEXT, align="l", fill="FFFFFF", ins=(0.12, 0.04, 0.1, 0.04),
             borders={k: (0.75, LINE) for k in "lrtb"})
    d.update(kw)
    return d


def head(text, align="ctr"):
    return cell(text, bold=True, color="FFFFFF", fill=NAVY, align=align, borders={k: (0.75, NAVY) for k in "lrtb"})


def label(text, fill):
    return cell(text, bold=True, color=NAVY, fill=fill)


def lines(*items, **kw):
    return [runs((t, dict(kw))) for t in items]


def shiji(s, y, h, body):
    """ページ下の「示唆」の枠（左に紺の線）"""
    add_box(s, 0.46, y, 0.07, h, [[("", {})]], fill=NAVY, name="示唆 線")
    add_box(s, 0.53, y, 9.83, h, [runs(("示唆", dict(size=13, bold=True, color=NAVY)))] + body,
            fill="F5F7FA", anchor="ctr", ins=(0.2, 0.06, 0.2, 0.06), line_spacing=120, name="示唆")


def note(s, y, texts):
    add_box(s, 0.46, y, 9.9, 0.2 * len(texts) + 0.05, lines(*texts, size=9, color=GRAY),
            anchor="t", ins=(0, 0, 0, 0), name="注記")


# ---------------------------------------------------------------- ニーズ調査（業界STUDY）
def p_study(s, src):
    clear(s)
    lead_band(s, src, "保険営業の方は毎年約24万人。会える回数が減るなか、成果を出すやり方へのニーズは続く")
    rows = [[head("項目"), head("ファクト"), head("意味すること")]]
    data = [
        ("営業職員数", "約24万人で、5年間ほぼ横ばい（2020年度 248,601人 → 2024年度 241,507人）",
         "毎年ほぼ同じ規模の方が、営業を続けている"),
        ("保険の加入経路", "加入経路のうち「生命保険会社の営業職員」が56.7%（2024年調査）",
         "保険の加入は、今も営業職員が中心。営業力を上げたいニーズは続く"),
        ("お客様との接触回数", "コロナ前と比べ、接触回数が「減少した」営業職員が55.5%（2023年調査）",
         "会える回数が減り、限られた面談で成果を出すやり方が求められている"),
    ]
    for i, (k, f, m) in enumerate(data):
        fill = BAND[i % 2]
        rows.append([label(k, fill), cell(f, fill=fill), cell(m, fill=fill)])
    add_table(s, 0.46, 1.8, [2.1, 3.9, 3.9], [0.4] + [0.85] * 3, rows, name="業界STUDY")
    shiji(s, 4.85, 1.0, lines("貴社の「週2日で成果を出す面談メソッド」は、このニーズに合う。",
                              "イベントのたびに集め直すのではなく、一度つながった方をLINEで次のイベントまで追い続ける。",
                              size=13, color=TEXT))
    note(s, 6.45, ["※営業職員数：生命保険協会「2025年版 生命保険の動向」／加入経路：生命保険文化センター「2024（令和6）年度 生命保険に関する全国実態調査」",
                   "※接触回数：矢野経済研究所「生命保険の営業職員アンケート調査」（2023年4〜6月、生命保険の営業職員400名）"])


# ---------------------------------------------------------------- ニーズ調査（業界STUDY）＞訴求軸
def p_axis(s, src):
    pic = next(sh for sh in s.shapes if sh.shape_type == 13)
    blob, ratio = pic.image.blob, pic.width / pic.height
    clear(s)
    lead_band(s, src, "保険営業の方向けの他社も、LINEで特典や学びを届けている")
    rows = [[head("他社"), head("サービス"), head("LINEでしていること")]]
    data = [
        ("野沢塾", "保険営業の方向けの講座（月2回）", "サイトのボタンから友だち追加。あいさつで特典の動画を届ける"),
        ("R&C保険営業ラボ", "保険営業の方向けの無料の学習サービス（2026年9月開始）",
         "登録時の5問のアンケートに合わせて、資料・動画を段階的に配信"),
    ]
    for i, (k, sv, ln) in enumerate(data):
        fill = BAND[i % 2]
        rows.append([label(k, fill), cell(sv, fill=fill), cell(ln, fill=fill)])
    add_table(s, 0.46, 1.8, [1.75, 2.35, 3.0], [0.4, 1.0, 1.0], rows, name="他社のLINE")
    add_box(s, 0.46, 4.55, 2.0, 0.34, [runs(("貴社の訴求軸", dict(size=13, bold=True, color="FFFFFF")))],
            fill=NAVY, align="ctr", anchor="ctr", name="貴社の訴求軸 見出し")
    add_box(s, 0.46, 4.89, 7.1, 0.9,
            lines("イベントの予約・参加に合わせて、案内を変える", "参加後は、サロンの入会へつなげる",
                  size=14, bold=True, color=NAVY),
            fill=TAKE, line=NAVY, align="ctr", anchor="ctr", line_spacing=125, name="貴社の訴求軸")
    ph = 3.75
    pw = ph * ratio
    px = 7.75 + (10.37 - 7.75 - pw) / 2
    s.shapes.add_picture(io.BytesIO(blob), Inches(px), Inches(1.8), Inches(pw), Inches(ph))
    add_box(s, 7.65, 1.8 + ph + 0.05, 2.8, 0.25, [runs(("野沢塾のあいさつ（LINE公式アカウント）", dict(size=10, color=GRAY)))],
            align="ctr", anchor="t", ins=(0, 0, 0, 0), name="画像の説明")
    note(s, 6.45, ["※野沢塾：LINE公式アカウント（2026年10月19日確認）・公式サイト（2026年10月18日確認）より。画面の名前とパスワードは伏せている",
                   "※R&C保険営業ラボ：R&C株式会社のプレスリリースより。招待制のため、画面は未掲載"])


# ---------------------------------------------------------------- 要件定義
def p_yoken(s, active6):
    clear(s)
    fill = BAND[1]
    rows = [
        [label("プロモーション目的", "FFFFFF"),
         cell("広告で集めた方をLINEにため、次のイベント参加数・入会数を伸ばし、LTVを最大化する",
              size=15, bold=True, color=NAVY)],
        [label("配信目標", fill),
         cell(lines(f"6か月目に追える人数{active6}人（ブロックされずに届く友だち）",
                    "イベント期はセミナー予約・参加、イベントの後はサロン入会へつなげる", size=14), fill=fill)],
        [label("手法（ご提案）", "FFFFFF"),
         cell(lines("① ポップアップツールとサンクスLINE誘導で、友だちを集める",
                    "② 興味あり・予約・参加の段階ごとに、案内を変える",
                    "③ 平常期の配信で、次のイベントまでつながりを保つ", size=14))],
        [label("＋αの施策", fill),
         cell(lines("過去のセミナー予約者・参加者へのご案内（通知メッセージ）",
                    "サロンの申込み完了ページからのLINE誘導（サンクスLINE誘導）", size=14)
              + lines("※通知メッセージは、予約時の電話番号がある場合", size=11, color=GRAY), fill=fill)],
    ]
    for r in rows:
        r[0]["size"] = 14
        r[0]["align"] = "ctr"
    heights = [0.8, 0.95, 1.3, 1.1]
    top = 0.7 + (7.0 - 0.7 - sum(heights)) / 3                    # タイトルの下からの空きの上1/3
    add_table(s, 0.46, top, [2.2, 7.7], heights, rows, name="要件定義")


# ---------------------------------------------------------------- 本質課題
def p_honshitsu(s, src):
    t = clear(s)
    set_one(t, "本質課題")
    lead_band(s, src, "3つの課題は、ひとつの課題につながっている")
    add_box(s, 0.46, 1.75, 2.3, 0.34, [runs(("KEY TAKEAWAY", dict(size=12, bold=True, color="FFFFFF")))],
            fill=NAVY, align="ctr", anchor="ctr", name="KEY TAKEAWAY 見出し")
    add_box(s, 0.46, 2.09, 9.9, 0.8, [runs(("イベントが終わると、集めた方との接点が切れる", dict(size=20, bold=True, color=NAVY)))],
            fill=TAKE, line=NAVY, align="ctr", anchor="ctr", name="KEY TAKEAWAY")
    rows = [[head("段階"), head("いま起きていること"), head("打ち手"), head("伸ばす数字")]]
    data = [
        ("01 集める", "LPを離れる方・予約した方を、友だちにする手段が無い",
         ("・ポップアップツールで、離れる方を友だちに", "・サンクスLINE誘導で、予約した方を友だちに"), "友だち数"),
        ("02 追う", "次のイベントまで、情報を届ける手段が無い",
         ("・平常期の配信（月2回）で、つながりを保つ", "・見どころ動画・前日のお知らせで、予約・参加へ"), "予約・参加数"),
        ("03 つなげる", "参加した方に、入会を案内する手段が無い",
         ("・イベント直後に、サロンの申込みを案内", "・入会した方の声を、迷っている方へ届ける"), "入会数"),
    ]
    for i, (k, now, acts, kpi) in enumerate(data):
        fill = BAND[i % 2]
        rows.append([label(k, fill), cell(now, fill=fill), cell(lines(*acts, size=13), fill=fill),
                     cell(kpi, fill=fill, bold=True, color=NAVY, align="ctr")])
    add_table(s, 0.46, 3.2, [1.35, 2.85, 4.4, 1.3], [0.4] + [0.9] * 3, rows, name="3つの課題と打ち手")


def main(src, out, active6=90):
    prs = Presentation(src)
    doshen = find(prs, lambda s: "貴社のユーザー動線イメージ" in text_of(s))
    p_study(find(prs, lambda s: title_of(s) is not None and title_of(s).text_frame.text.strip() == "ニーズ調査（業界STUDY）"), doshen)
    p_axis(find(prs, lambda s: title_of(s) is not None and title_of(s).text_frame.text.strip().startswith("ニーズ調査（業界STUDY）　＞")), doshen)
    p_yoken(find(prs, lambda s: title_of(s) is not None and title_of(s).text_frame.text.strip() == "要件定義"
                 and "プロモーション目的" in text_of(s)), active6)
    p_honshitsu(find(prs, lambda s: title_of(s) is not None and title_of(s).text_frame.text.strip() == "要件定義"
                     and "本質課題" in text_of(s)), doshen)
    prs.save(out)
    print("saved:", out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
