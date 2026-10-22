# -*- coding: utf-8 -*-
"""すなお ver6.7 → ver6.8（2026-10-22 殿村さん）。

    python -m _build.patch_sunao_v68_kaisai <ver6.7.pptx> <ver6.8.pptx>

  ①「平常期」を「開催のない月」に全部変える（本編・Appendixとも。言い回しは文ごとに合わせる）
  ②「なぜLINEか」の表をやめ、広告とLINE公式アカウントを左右に並べて比べる形にする
    （右のLINEの方を広く・紺で強調。間に「友だち追加」の矢印。下の「LINE公式アカウントのメリット」はそのまま）
    表の「イベント期も、イベントのない期間も」も「開催月も、開催のない月も」にそろえる。
  ほかの文言・ページは変えない。
"""
import sys

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn

from _build.pptx_parts import CARD, INK, LGRAY, NAVY, WHITE, hline, rect, rich, shape

SUB = "64748B"
IN = 2.54  # インチ → cm

EDITS = {  # ページ番号: [(元の段落, 新しい段落)]
    14: [("6か月（平常期とセミナー開催月で、配信内容と工数が変わる）", "6か月（セミナー開催月と開催のない月で、配信内容と工数が変わる）"),
         ("③ 平常期 月1〜2回、開催月 月4〜5回の企画配信で、つながりを保つ", "③ 開催のない月は月1〜2回、開催月は月4〜5回の企画配信で、つながりを保つ")],
    15: [("平常期・開催月それぞれの", "開催月・開催のない月"),
         ("時期に合わせて届ける", "それぞれに合わせて届ける"),
         ("・企画配信：平常期 月1〜2回／開催月 月4〜5回", "・企画配信：開催のない月は月1〜2回／開催月は月4〜5回")],
    16: [("企画配信：平常期 月1〜2回／開催月 月4〜5回", "企画配信：開催のない月は月1〜2回／開催月は月4〜5回")],
    18: [("平常期の配信", "開催のない月の配信")],
    20: [("平常期 ¥135,000〜", "開催のない月 ¥135,000〜"),
         ("平常期 1〜2回", "開催のない月 1〜2回"),
         ("（追加した場合の合計：初期 ¥295,000〜／開催月 ¥365,000程度／平常期 ¥165,000〜）",
          "（追加した場合の合計：初期 ¥295,000〜／開催月 ¥365,000程度／開催のない月 ¥165,000〜）")],
    21: [("月額費用（平常期）", "月額費用（開催のない月）"),
         ("※月額費用＝コンサルティング費（平常期 100,000円／セミナー開催月 300,000円程度）＋ツール・LINE関連費（推奨構成 35,000円〜＝ポップアップ 30,000円＋LINE公式アカウント利用料 5,000円〜）",
          "※月額費用＝コンサルティング費（開催のない月 100,000円／セミナー開催月 300,000円程度）＋ツール・LINE関連費（推奨構成 35,000円〜＝ポップアップ 30,000円＋LINE公式アカウント利用料 5,000円〜）"),
         ("※ポップアップ・サンクスLINE誘導はどちらも選択制。ポップアップも導入しない場合の参考：初期 130,000円〜／平常期 105,000円〜／開催月 305,000円程度（別途の実装・制作費等は含まない）",
          "※ポップアップ・サンクスLINE誘導はどちらも選択制。ポップアップも導入しない場合の参考：初期 130,000円〜／開催のない月 105,000円〜／開催月 305,000円程度（別途の実装・制作費等は含まない）")],
    22: [("（平常期）", "（開催のない月）")],
    23: [("平常期 月1〜2回", "開催のない月は月1〜2回"),
         ("平常期 月3回", "開催のない月は月3回"),
         ("※平常期の配信は、次のイベントまで関係を保つためのもの（配信の回数を増やすこと自体は目的にしない）",
          "※開催のない月の配信は、次のイベントまで関係を保つためのもの（配信の回数を増やすこと自体は目的にしない）")],
    26: [("⑤ 平常期", "⑤ 開催のない月")],
    32: [("平常期の層（予約していない方）", "開催のない月の層（予約していない方）"),
         ("平常期の配信で、次のイベントまでつながりを保てる", "開催のない月の配信で、次回までつながりを保てる"),
         ("予約していない方は平常期の配信で次のイベントへつなげる。流入経路だけでは予約・参加・入会の状況をすべて把握できないため、状況に応じた出し分けは、連携の仕様を確認のうえ設計。",
          "予約していない方は開催のない月の配信で次のイベントへつなげる。流入経路だけでは予約・参加・入会の状況をすべて把握できないため、状況に応じた出し分けは、連携の仕様を確認のうえ設計。")],
    52: [("セミナー開催月・平常期の企画配信", "セミナー開催月・開催のない月の企画配信"),
         ("開催月は月4〜5回、平常期は月1〜2回の配信案", "開催月は月4〜5回、開催のない月は月1〜2回の配信案")],
}

# なぜLINEか：左右の比較（ラベル, 広告, LINE公式アカウント）
ROWS = [("役割", "新しい見込み客との接点を獲得する", "一度つながった見込み客に、継続して情報を届ける"),
        ("主に届く相手", "まだつながっていない方", "友だち追加した方"),
        ("届けられる時期", "広告を配信している期間", "開催月も、開催のない月も"),
        ("積み上がるもの", "セミナー予約などの反響", "友だち（次回以降も案内できる相手）")]


def paras(s):
    def walk(shs):
        for sh in shs:
            if sh.shape_type == 6:
                yield from walk(sh.shapes)
            if sh.has_text_frame:
                yield from sh.text_frame._txBody.iter(qn("a:p"))
            if getattr(sh, "has_table", False) and sh.has_table:
                yield from sh._element.iter(qn("a:p"))
    return list(walk(s.shapes))


def ptext(p):
    return "".join(t.text or "" for t in p.iter(qn("a:t")))


def set_para(p, text):
    runs = p.findall(qn("a:r"))
    for r in runs[1:]:
        p.remove(r)
    runs[0].find(qn("a:t")).text = text


def compare(s):
    """表を外し、広告（左）とLINE公式アカウント（右）を並べる。位置はインチで考えてcmに直す"""
    tbl = next(sh for sh in s.shapes if sh.has_table)
    x0, y0 = 0.55, 1.6
    lw, gap, rw = 3.85, 0.9, 4.98                    # 左（広告）／矢印／右（LINE）。合計9.73＝元の表の幅
    hh, item = 0.4, 0.5                              # 見出しの高さ／1項目の高さ
    body = item * len(ROWS) + 0.1
    tbl._element.getparent().remove(tbl._element)

    for x, w, head, strong in ((x0, lw, "広告", False), (x0 + lw + gap, rw, "LINE公式アカウント", True)):
        rect(s, x * IN, (y0 + hh) * IN, w * IN, body * IN, CARD if strong else WHITE,
             line=NAVY if strong else LGRAY, lw=1.25 if strong else 0.75)
        rect(s, x * IN, y0 * IN, w * IN, hh * IN, NAVY if strong else "EEF0F4",
             [(head, 13, True, WHITE if strong else NAVY, 0)])
        for i, row in enumerate(ROWS):
            ty = y0 + hh + 0.05 + i * item
            if i:
                hline(s, (x + 0.15) * IN, (x + w - 0.15) * IN, ty * IN, LGRAY, 0.5)
            rich(s, (x + 0.1) * IN, (ty + 0.03) * IN, (w - 0.2) * IN, (item - 0.06) * IN,
                 [[(row[0], 10, False, SUB)], [(row[2] if strong else row[1], 13, strong, NAVY if strong else INK)]],
                 anchor=MSO_ANCHOR.MIDDLE)

    ax = x0 + lw + 0.12
    ay = y0 + hh + body / 2
    rich(s, (x0 + lw) * IN, (ay - 0.55) * IN, gap * IN, 0.3 * IN, [[("友だち追加", 10, True, NAVY)]],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.BOTTOM)
    shape(s, MSO_SHAPE.RIGHT_ARROW, ax * IN, (ay - 0.2) * IN, (gap - 0.24) * IN, 0.4 * IN, fill=NAVY)
    return y0 + hh + body


def main(src, out):
    prs = Presentation(src)
    slides = list(prs.slides)
    assert len(slides) == 60
    for no, pairs in EDITS.items():
        ps = paras(slides[no - 1])
        for old, new in pairs:
            hits = [p for p in ps if ptext(p) == old]
            assert hits, (no, old)
            for p in hits:
                set_para(p, new)

    why = slides[10]
    assert any(sh.has_text_frame and sh.text_frame.text == "なぜLINEか" for sh in why.shapes)
    bottom = compare(why)
    head = next(sh for sh in why.shapes if sh.name == "メリット 見出し")
    assert bottom + 0.15 <= head.top / 914400, bottom

    left = [i + 1 for i, s in enumerate(prs.slides) if any("平常" in ptext(p) for p in paras(s))]
    assert not left, left
    prs.save(out)
    print("saved:", out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
