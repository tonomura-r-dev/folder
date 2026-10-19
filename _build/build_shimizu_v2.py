# -*- coding: utf-8 -*-
"""株式会社横浜シミズ様 LINE公式アカウント運用のご提案 v2（草案）。中途・新卒・アルバイトの3本の柱で組み直した版。

    python -m _build.build_shimizu_v2 <人材業界資料.pptx> <あらうちv10.pptx> <SIM.xlsx>

  2026-10-19 殿村さん「その三本で組んでください」。要件は `_drafts/清水グループ_与件整理（2026-10-19）.md` の末尾。
  ・社員（中途・新卒）とアルバイトの両方。職種＝イベントスタッフ（アルバイト）／運営部／制作部／建築士・施工管理技士
  ・アルバイト＝働いた後もつながり、次のイベントに呼ぶ（営業が先方と話している「管理と、連絡が取れない層への再アプローチ」）
  ・中途＝今すぐ転職しない方も残し、募集を届け続ける／新卒＝説明会の予約とリマインド
  ・CV＝応募。面接の日程・リマインドは「LINEでできる」程度
  ・いまのLINE公式アカウントは「LINEで応募」だけ（入力は手入力）
  並び：表紙／アジェンダ／現状／課題／解決策／施策①アルバイト／施策②中途／施策③新卒／友だちの入口／リッチメニュー／
        概算お見積り／運用プラン／運用サポート体制／裏表紙
  v1の「面接に来る方を増やす」は施策①③に入れたので外した。概算お見積り・運用プラン・運用サポート体制はv1と同じ作り方。
"""
import copy
import sys
from pathlib import Path

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches

from _build.arauchi_v4.lib import add_box, move, set_paras
from _build.arauchi_v10.build import by_id, clone, runs
from _build.build_shimizu_v1 import (LBLUE, NAVY, PLANS, TEXT, arrow, body_page, box, drop, fill_row, p_mitsumori,
                                     p_ryokin, p_support, put)
from _build.build_sunao_v1 import all_meiryo
from _build.build_sunao_v2 import load_sim
from _build.patch_tanaka_v5 import copy_slide
from pptx.enum.text import PP_ALIGN

ROOT = Path(__file__).resolve().parent.parent
GRAY = "7F7F7F"
LINE = "BFC7D5"
X0, W = 0.47, 9.89                                            # 本文の左端と幅


# ---------------------------------------------------------------- 部品
def center(s, top=1.1, bottom=6.8, ratio=0.33):
    """タイトル・リード・注記より下の中身を、まとめて上寄りの中央へ（空きの上1/3・下2/3。すなおと同じ）"""
    body = [sh for sh in s.shapes if sh.shape_id not in (2, 3) and sh.name != "注記"]
    y0 = min(sh.top for sh in body)
    y1 = max(sh.top + sh.height for sh in body)
    dy = int(Inches(top) - y0 + (Inches(bottom) - Inches(top) - (y1 - y0)) * ratio)
    for sh in body:
        sh.top = sh.top + dy


def note(s, text):
    add_box(s, X0, 6.83, W, 0.3, [runs((text, dict(size=9, color=GRAY)))], anchor="t", ins=(0, 0, 0, 0), name="注記")


def band(s, y, text):
    add_box(s, X0, y, W, 0.45, [runs((text, dict(size=16, bold=True, color="FFFFFF")))],
            fill=NAVY, align="ctr", anchor="ctr", name="示唆")


def flow(s, y, items, h=0.8, aw=0.3):
    """横一本の流れ図。items＝[(文字, 強調), ...]"""
    bw = (W - aw * (len(items) - 1)) / len(items)
    for k, (t, strong) in enumerate(items):
        x = X0 + k * (bw + aw)
        box(s, x, y, bw, h, t, strong)
        if k < len(items) - 1:
            arrow(s, x + bw, y, aw, h)
    return bw


def points(s, y, items, label_w=2.6, h=0.55, gap=0.1):
    """表の形の「ポイント」：左に紺の見出し、右に1〜2行の説明"""
    for k, (lab, txt) in enumerate(items):
        yy = y + k * (h + gap)
        add_box(s, X0, yy, label_w, h, [runs((lab, dict(size=13, bold=True, color="FFFFFF")))],
                fill=NAVY, align="ctr", anchor="ctr", ins=(0.06, 0.03, 0.06, 0.03), name=lab)
        add_box(s, X0 + label_w, yy, W - label_w, h, [runs((txt, dict(size=13, color=TEXT)))],
                line=LINE, anchor="ctr", ins=(0.15, 0.03, 0.1, 0.03), name=lab + "の説明")


def rows_page(prs, src, title, lead, data, header=None, band_text=None, note_text=None):
    """業界資料の「歩留まり」の行の形（左に紺の見出し・右に文字）を、行の数と列の数に合わせて使う。
    data＝[(見出し, 2列目, 3列目 or None), ...]。3列目が無いときは2列目を右端まで広げる"""
    s = copy_slide(prs, src)
    put(s, 2, title)
    put(s, 3, lead)
    two = all(d[2] is None for d in data)
    for i in range(4):
        rect, lrect, lab, c2, c3 = (4 + 5 * i + k for k in range(5))
        if i >= len(data):
            drop(s, rect, lrect, lab, c2, c3)
            continue
        h, t2, t3 = data[i]
        r = by_id(s, rect)
        y, hh = r.top / 914400, r.height / 914400
        put(s, lab, h)
        fill_row(s, lab, y, hh)
        put(s, c2, t2)
        if two:
            move(by_id(s, c2), w=10.28 - 3.62)
            drop(s, c3)
        else:
            put(s, c3, t3)
            fill_row(s, c3, y, hh)
        fill_row(s, c2, y, hh)
    last = by_id(s, 4 + 5 * (len(data) - 1))
    yb = (last.top + last.height) / 914400 + 0.25
    if band_text:
        put(s, 25, band_text)
        move(by_id(s, 24), y=yb)
        move(by_id(s, 25), y=yb + 0.04)
    else:
        drop(s, 24, 25)
    drop(s, 26)
    if header:
        for x, w, t in zip((0.47, 3.62, 7.01), (2.99, 3.31, 3.27), header):
            add_box(s, x, 1.69 - 0.32, w, 0.28, [runs((t, dict(size=11, bold=True, color=GRAY)))],
                    align="ctr" if x < 1 else "l", anchor="b", ins=(0, 0, 0, 0), name="列の見出し")
    if note_text:
        note(s, note_text)
    center(s)
    return s


# ---------------------------------------------------------------- 表紙・アジェンダ
def p_cover(s, ar_cover):
    sub = by_id(s, 6)
    set_paras(sub, ["中途・新卒・アルバイトの採用を、", "ひとつのLINEでつなぐ仕組み"])
    move(sub, x=1.76, y=2.3, w=7.25, h=0.95)
    for p in sub.text_frame.paragraphs:
        p.alignment = PP_ALIGN.CENTER
    drop(s, 3, 7)                                              # 「- 2026最新版 -」「業界：人材（派遣・バイト・単発）」
    set_paras(clone(s, by_id(ar_cover, 31)), ["株式会社横浜シミズ 御中"])


def p_agenda(s):
    put(s, 5, ["中途・新卒・アルバイトの求職者と、LINEでつながり続ける。",
               "いまある「LINEで応募」の仕組みを、応募の前と働いた後まで広げる提案。"])
    for (a, b), (h, t) in zip(((9, 10), (14, 15), (19, 20), (24, 25)),
                              (("第1章｜現状と課題", "採用の入口は3つ。つながりが切れるところが、それぞれ違う"),
                               ("第2章｜解決策", "ひとつのLINEで、中途・新卒・アルバイトを分けて追う"),
                               ("第3章｜施策", "アルバイト・中途・新卒の施策と、友だちの入口・リッチメニュー"),
                               ("第4章｜費用", "概算お見積り・運用プラン・運用サポート体制"))):
        put(s, a, h)
        put(s, b, t)
    drop(s, *range(26, 36), 41)
    for sid in range(6, 26):
        sh = by_id(s, sid)
        sh.top = sh.top + Inches(0.6)


# ---------------------------------------------------------------- 第1章
def p_genjo(prs, src):
    return rows_page(
        prs, src, "［第1章］現状｜採用の入口は、中途・新卒・アルバイトの3つ",
        "職種も、募集している場所も、3つで違う。",
        [("中途（キャリア採用）", ["運営部・制作部・", "建築士・施工管理技士（正社員）"], "自社の採用ページ"),
         ("新卒（27卒）", ["運営部・制作部（正社員）", "会社説明会"], "自社の採用ページ"),
         ("アルバイト", ["会場スタッフ（神奈川・静岡）", "スタジアム・サッカー会場など"],
          ["グループのアルバイト求人サイト", "（LINEで応募できる）"])],
        header=("対象", "募集している職種", "募集の場所"),
        band_text="LINEで応募を受け付ける仕組みは、すでにある",
        note_text="※ 2026年10月時点の、貴社の採用ページとアルバイト求人サイトの掲載内容より")


def p_kadai(prs, src):
    return rows_page(
        prs, src, "［第1章］課題｜対象ごとに、つながりが切れるところが違う",
        "応募の前・説明会の前・働いた後のどこかで、求職者が離れていく。",
        [("① アルバイト", "イベントごとの仕事が多く、働き終わると連絡が途切れやすい", None),
         ("② 中途（資格職）", ["建築士・施工管理技士は、人数が少ない職種",
                              "今すぐ転職しない方とのつながりが残らない"], None),
         ("③ 新卒", ["説明会の予約から当日までに、間があく",
                    "その間に参加を忘れたり、他社に気持ちが移ったりする"], None)],
        band_text="一度つながった方と、つながり続ける場所が要る")


# ---------------------------------------------------------------- 第2章
def p_kaiketsu(prs, src):
    s = body_page(prs, src, "［第2章］解決策｜ひとつのLINEで、中途・新卒・アルバイトを分けて追う",
                  "友だち追加の後に「社員／アルバイト」と職種を選んでもらい、届ける中身を分ける。")
    y = 1.75
    bw = flow(s, y, (("採用ページ\nアルバイト求人サイト", False), ("LINE\n友だち追加", False),
                     ("社員／アルバイト\n職種を選ぶ", False), ("対象ごとに\n届ける中身を分ける", True)))
    arrow(s, X0 + 3 * (bw + 0.3), y + 0.8, bw, 0.35, "▼")
    aw = 0.3
    cw = (W - aw * 2) / 3
    for k, t in enumerate(("アルバイト\n働いた後の再アプローチ", "中途\n資格職とのつながりを残す", "新卒\n説明会の予約とリマインド")):
        box(s, X0 + k * (cw + aw), y + 1.2, cw, 0.8, t)
    band(s, y + 2.3, "いまのLINE公式アカウント（LINEで応募）を、そのまま活かせる")
    center(s)
    return s


# ---------------------------------------------------------------- 第3章 施策①〜③（流れ図＋ポイント）
def pillar(prs, src, title, lead, steps, sub, pts, note_text=None):
    s = body_page(prs, src, title, lead)
    y = 1.75
    flow(s, y, steps)
    yy = y + 0.85
    if sub:
        add_box(s, X0, yy, W, 0.32, [runs((sub, dict(size=13, bold=True, color=NAVY)))],
                align="r", anchor="ctr", ins=(0, 0, 0, 0), name="補足")
        yy += 0.32
    points(s, yy + 0.3, pts)
    if note_text:
        note(s, note_text)
    center(s)
    return s


def p_arubaito(prs, src):
    return pillar(
        prs, src, "［第3章］施策①アルバイト｜働いた後もつながり、次のイベントに呼ぶ",
        "応募から働いた後まで、同じLINEでやりとりを続ける。",
        (("LINEで\n応募", False), ("面接の日程\n前日のお知らせ", False), ("会場で\n働く", False),
         ("次のイベントの\n募集を届ける", False), ("もう一度\n応募", True)),
        "イベントごとに、くり返し呼べる",
        (("勤務地・会場で分ける", "神奈川・静岡、スタジアム・サッカー会場など、働きたい場所の募集だけを届ける"),
         ("働く日の連絡", "集合場所・持ち物・当日の連絡を、LINEで送る"),
         ("しばらく働いていない方", "新しい会場・時期の募集を届け、もう一度働くきっかけをつくる")),
        "※ 面接の前日のお知らせは、個別チャットか通知メッセージ（別途費用）で送る")


def p_chuto(prs, src):
    return pillar(
        prs, src, "［第3章］施策②中途｜今すぐ転職しない方も残し、募集を届け続ける",
        "建築士・施工管理技士のように人数の少ない職種ほど、一度つながった方が貴重になる。",
        (("採用ページを\n見る", False), ("LINE友だち追加\n職種を選ぶ", False), ("職種ごとに\n仕事・現場を紹介", False),
         ("新しい募集を\nお知らせ", False), ("応募", True)),
        None,
        (("職種で分ける", "運営部・制作部・建築士・施工管理技士で、届ける中身を分ける"),
         ("届ける中身の例", "ステージ・会場づくりの事例、現場の1日、働く方の声"),
         ("募集を出した日に", "その職種を選んだ方へ、すぐに届ける")),
        "※ 届ける中身は例。貴社と相談して決める")


def p_shinsotsu(prs, src):
    return pillar(
        prs, src, "［第3章］施策③新卒｜説明会の予約から当日まで、LINEでつなぐ",
        "予約した方が当日に来るように、予約から説明会までの間を埋める。",
        (("説明会の\n案内", False), ("LINEから\n予約", False), ("前日の\nお知らせ", False),
         ("説明会に\n参加", False), ("選考の案内\n応募", True)),
        None,
        (("予約の後", "会社・仕事の紹介を届け、当日までに関心を高める"),
         ("前日", "予約した日ごとに分けて、日時・場所・参加方法を届ける"),
         ("来られなかった方", "次の説明会の日程を案内する")))


# ---------------------------------------------------------------- 第3章 友だちの入口・リッチメニュー
def p_iriguchi(s, lead_src):
    put(s, 2, "［第3章］施策｜友だちの入口")
    lead = by_id(s, 3)                                         # リードの位置・大きさを、ほかのページ（歩留まりの形）とそろえる
    lead._element.getparent().replace(lead._element, copy.deepcopy(by_id(lead_src, 3)._element))
    put(s, 3, "広告を増やさずに、すでにある接点からLINEにつなぐ。")
    # ⑤の行は業界資料ではオレンジの強調（★）。①〜④と同じ紺の行に置き換える（4行目を写して下にずらす）
    dy = by_id(s, 24).top - by_id(s, 19).top
    drop(s, 24, 25, 26, 27, 28)
    for sid in range(19, 24):
        el = copy.deepcopy(by_id(s, sid)._element)
        el.find(".//" + qn("p:cNvPr")).set("id", str(sid + 5))
        s.shapes._spTree.append(el)
        sh = by_id(s, sid + 5)
        sh.top = sh.top + dy
    data = (("① 応募完了画面", "応募した直後に友だち追加を案内", "すでに応募した方"),
            ("② 離脱防止ポップアップ", "採用ページ・求人サイトを離れる前に案内", "応募を迷っている方"),
            ("③ サイトのLINEボタン", "採用ページ・求人サイトにLINEの入口を置く", "まだ応募を決めていない方"),
            ("④ 会場のQRコード", "働いた日に、会場で友だち追加を案内", "アルバイトで働いた方"),
            ("⑤ これまで働いた方", "登録のある連絡先へ、友だち追加を案内", "もう一度つながるきっかけ"))
    for i, (h, t2, t3) in enumerate(data):
        rect, lab, c2, c3 = 4 + 5 * i, 6 + 5 * i, 7 + 5 * i, 8 + 5 * i
        r = by_id(s, rect)
        y, hh = r.top / 914400, r.height / 914400
        for sid, t in ((lab, h), (c2, t2), (c3, t3)):
            put(s, sid, t)
            fill_row(s, sid, y, hh)
    put(s, 30, "今すぐ応募しない方も、働き終わった方も、LINEに残しておける")
    drop(s, 31)
    center(s)


def p_menu(prs, src):
    s = body_page(prs, src, "［第3章］施策｜リッチメニュー（社員・アルバイトでタブを分ける）",
                  "タブを切り替えるだけで、社員とアルバイトで見る中身が変わる。")
    mw, gap, th, ch = 4.6, W - 4.6 * 2, 0.45, 0.9
    menus = (("社員の方", ("募集中の職種", "新卒・説明会", "仕事・現場の紹介", "よくある質問", "応募する", "通知の設定")),
             ("アルバイトの方", ("募集中のお仕事", "勤務地で探す", "働く日の連絡", "よくある質問", "LINEで応募", "通知の設定")))
    y = 1.75
    for m, (active, buttons) in enumerate(menus):
        x = X0 + m * (mw + gap)
        for k, tab in enumerate(("社員の方", "アルバイトの方")):
            on = tab == active
            add_box(s, x + k * mw / 2, y, mw / 2, th, [runs((tab, dict(size=13, bold=True, color="FFFFFF" if on else GRAY)))],
                    fill=NAVY if on else "E7E9EE", align="ctr", anchor="ctr", ins=(0, 0, 0, 0), name="タブ")
        for k, b in enumerate(buttons):
            add_box(s, x + (k % 3) * mw / 3, y + th + (k // 3) * ch, mw / 3, ch,
                    [runs((b, dict(size=13, bold=True, color=NAVY)))], fill="FFFFFF", line=LINE,
                    align="ctr", anchor="ctr", ins=(0.04, 0.03, 0.04, 0.03), name=b)
    band(s, y + th + ch * 2 + 0.35, "「通知の設定」を置くから、ブロックされずにつながり続ける")
    note(s, "※ ボタンの名前と並びはイメージ。運用しながら調整する")
    center(s)
    return s


def main(base, arauchi, sim_path):
    sim = load_sim(sim_path)
    plan = PLANS[""]
    prs = Presentation(base)
    sl = list(prs.slides)
    assert by_id(sl[0], 2).text_frame.text == "LINE公式アカウント運用のご提案"
    assert by_id(sl[8], 2).text_frame.text.endswith("友だち追加の動線")
    assert by_id(sl[17], 2).text_frame.text.startswith("［第4章］歩留まり")
    assert by_id(sl[21], 2).text_frame.text.startswith("［第6章］費用プラン")
    cover, agenda, iriguchi, budomari, support, ryokin, back = (sl[i - 1] for i in (1, 2, 9, 18, 20, 22, 23))
    ar = Presentation(arauchi)
    ar_cover, ar_sim = ar.slides[0], ar.slides[16]
    assert by_id(ar_sim, 339).text_frame.text.startswith("目標と費用対効果")

    genjo = p_genjo(prs, budomari)
    kadai = p_kadai(prs, budomari)
    kaiketsu = p_kaiketsu(prs, budomari)
    arubaito = p_arubaito(prs, budomari)
    chuto = p_chuto(prs, budomari)
    shinsotsu = p_shinsotsu(prs, budomari)
    menu = p_menu(prs, budomari)
    init_total = plan["consult_init"] + 15000 + 150000 + 30000
    mitsumori = p_mitsumori(prs, budomari, ar_sim, sim, init_total)
    p_cover(cover, ar_cover)
    p_agenda(agenda)
    p_iriguchi(iriguchi, budomari)
    p_ryokin(ryokin, plan["rows"])
    p_support(support)

    keep = [cover, agenda, genjo, kadai, kaiketsu, arubaito, chuto, shinsotsu, iriguchi, menu,
            mitsumori, ryokin, support, back]
    lst = prs.slides._sldIdLst
    by_part = {prs.part.related_part(e.rId): e for e in lst}
    for e in list(lst):
        lst.remove(e)
    keep_parts = {s.part for s in keep}
    for part, e in by_part.items():
        if part not in keep_parts:
            prs.part.drop_rel(e.rId)
    for s in keep:
        e = by_part[s.part]
        e.attrib.pop("show", None)
        s._element.attrib.pop("show", None)
        lst.append(e)
    all_meiryo(keep)
    out = ROOT / "202610_株式会社横浜シミズ御中_LINE公式アカウント運用のご提案.pptx"
    prs.save(out)
    print("saved:", out.name, len(prs.slides), "枚")


if __name__ == "__main__":
    main(*sys.argv[1:4])
