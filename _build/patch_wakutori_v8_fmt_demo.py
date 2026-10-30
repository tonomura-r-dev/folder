# -*- coding: utf-8 -*-
"""ワクトリ v6（殿村さんPC保存版）→ v8：具体施策のイメージ＝提案FMT（最強資料）の8枚＋左にデモ画面＋右に3つの枠（2026-10-30 殿村さん）。

    python -m _build.patch_wakutori_v8_fmt_demo <PC保存版v6.pptx> <FMT具体施策枠のみ.pptx> <出力.pptx>

  ・ページ・札・上の矢印・枠は、FMTの8枚（LINEの友だちへ〜コミュニティ運営）のまま
  ・左のメインの枠：施策名＋1行、デモ（`_data/wakutori/ワクトリ_LINE施策デモ_20261030v3.html`）の①ユーザーが見る画面＋②裏側の処理
    （`_images/wakutori/demo8_p24〜31.png`）、下に画面の注記
  ・右：FMTの4つの枠を3つにして、入口／LINEの仕組み／配信内容（ほかにもこういう案）。見出しに前置きは付けない
  ・施策の考え方（ステップ配信2枚・企画配信1枚＝旧P23〜P25）は、FMTの並びどおり具体施策の後ろへ。P14の（P◯）は振り直す
  v7（FMTの中身の形）は使わない。
"""
import copy
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Emu, Inches, Pt
from PIL import Image

from _build.build_juku_fmt1029 import fmt_fix
from _build.patch_fmt_max_gutai import IN, NAVY, INK
from _build.patch_tanaka_v5 import copy_slide
from _build.patch_wakutori_v7_fmt_gutai import P14_FIX
from _build.pptx_parts import rich

GRAY = "7F7F7F"
IMG = "_images/wakutori/demo8_p{}.png"
LABELS = ["入口", "LINEの仕組み", "配信内容"]
PAGES = [  # 札, (施策名, 1行), [入口, LINEの仕組み, 配信内容]（「|」で改行）
    ("LINEの友だちへ", ("ポップアップ", "サイトを離れる方を、LINEの友だちに"),
     ["SNSのプロフィール・投稿や|体験談のページにも入口を置く",
      "入口ごとにURLを分けて記録し、|あいさつの文面を出し分ける",
      "追加の特典に「寮生活の|持ち物リスト」などを届ける"]),
    ("露出の強化", ("企画配信", "季節に合わせて、届く求人を切り替える"),
     ["特集から、リッチメニューの|開始月別の求人へつなげる",
      "季節ごとに、リッチメニューの|特集の枠を入れ替える",
      "「LINE友だち限定」の先行公開や|締切が近い求人の特集"]),
    ("CVの最大化", ("ステップ配信", "追加後15日で、応募（ENTRY）へ"),
     ["診断に答えていない方には、|翌日に答えるお願いを届ける",
      "押した内容（お金・寮など）で|タグを付け、次の配信を変える",
      "寮の写真や1日の流れなど、|不安ごとの配信を足す"]),
    ("CPOの改善（CV質）", ("サンクスLINE誘導（オプション）", "応募の完了画面から、そのままLINEへ"),
     ["お問い合わせなど、ほかの|フォームの完了画面にも置く",
      "ENTRYで伺った内容と同じ質問は|省いて、診断を短くする",
      "面談・お仕事開始までの流れ、|持ち物・よくある質問を届ける"]),
    ("休眠顧客の有効化", ("休眠の掘り起こし", "反応のない方に、開始時期を聞き直す"),
     ["リッチメニューにも、|「希望を変える」ボタンを置く",
      "反応のない期間（30日・60日）で|タグを分けて、送る回数を変える",
      "シーズンの前に、開始月別の|新着求人だけを短く届ける"]),
    ("マネタイズ最大化", ("開始時期・エリア別の配信", "合う方にだけ、合う求人を届ける"),
     ["求人を押したが応募していない方に、|翌日にその求人を届け直す",
      "重視したいポイント（立地・|稼げる・出会い・住環境）も聞く",
      "求人と一緒に、|専属エージェントへの相談の入口も届ける"]),
    ("リピート施策", ("次のシーズンのお仕事", "お仕事を終えた方に、次の求人を先に届ける"),
     ["お仕事が終わる前に、|次の希望を聞くアンケートを送る",
      "勤務した時期・エリアで|タグを付け、次の季節に出し分ける",
      "前回と同じエリア・違うエリアの|求人を並べて届ける"]),
    ("コミュニティ運営", ("体験談・紹介", "くり返し働く方とのつながりを強くする"),
     ["リッチメニューに、体験談・|紹介の入口を置く",
      "紹介した方にタグを付け、|紹介からの応募を数える",
      "経験者の体験談を、|新しい方へのステップ配信に入れる"]),
]


def inch(v):
    return Emu(v).inches


def fill(s, n, page):
    tag, (name, line), ideas = page
    assert any(sh.has_text_frame and sh.text_frame.text == tag for sh in s.shapes), tag
    # 右：FMTの4つの枠 → 3つ
    right = sorted([sh for sh in s.shapes if inch(sh.left) > 6.8 and inch(sh.top) > 1.9], key=lambda x: inch(x.top))
    assert len(right) == 4, len(right)
    right[3]._element.getparent().remove(right[3]._element)
    top, bottom, gap = 2.03, 6.80, 0.1
    bh = (bottom - top - gap * 2) / 3
    for k, box in enumerate(right[:3]):
        box.left, box.top, box.height = Inches(6.85), Inches(top + k * (bh + gap)), Inches(bh)
        rich(s, 6.97 * IN, (top + k * (bh + gap) + 0.06) * IN, 3.46 * IN, 1.2 * IN,
             [[(LABELS[k], 13, True, NAVY)]] + [[(t, 12, False, INK)] for t in ideas[k].split("|")], anchor=MSO_ANCHOR.TOP)
    # 左：施策名＋1行、デモの画面、注記
    rich(s, 0.45 * IN, 2.15 * IN, 6.1 * IN, 0.7 * IN, [[(name, 16, True, NAVY)], [(line, 12, False, INK)]], anchor=MSO_ANCHOR.TOP)
    path = IMG.format(n)
    w0, h0 = Image.open(path).size
    h = 3.6
    w = h * w0 / h0
    pic = s.shapes.add_picture(path, Inches(0.29 + (6.44 - w) / 2), Inches(2.88), Inches(w), Inches(h))
    pic.line.color.rgb = RGBColor.from_string("D9D9D9")
    pic.line.width = Pt(0.75)
    rich(s, 0.45 * IN, 6.52 * IN, 6.1 * IN, 0.24 * IN,
         [[("画面：ワクトリ様向けに作成したデモ（①ユーザーが見る画面／②裏側の処理）。求人はデモ用のサンプル", 8.5, False, GRAY)]],
         anchor=MSO_ANCHOR.TOP)


def main(src, fmt_empty, out):
    prs = Presentation(src)
    emp = Presentation(fmt_empty)
    assert len(prs.slides) == 55
    W = {i + 1: s for i, s in enumerate(prs.slides)}
    for n in range(27, 34):
        assert any(sh.has_text_frame and sh.text_frame.text == "具体施策のイメージ" for sh in W[n].shapes), n
    E = []
    for k, page in enumerate(PAGES):
        s = copy_slide(prs, emp.slides[17 + k])
        fmt_fix(s)
        fill(s, 24 + k, page)
        E.append(s)

    order = [W[n] for n in range(1, 23)] + [W[26]] + E + [W[23], W[24], W[25]] + [W[n] for n in range(34, 56)]
    assert len(order) == 56
    done = set()
    for t in W[14]._element.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}t"):
        for a, b in P14_FIX:
            if t.text == a:
                t.text = b
                done.add(a)
    assert done == {a for a, _ in P14_FIX}, done

    lst = prs.slides._sldIdLst
    ids = {id(s._element): sid for s, sid in zip(prs.slides, list(lst))}
    keep = [ids[id(s._element)] for s in order]
    keep_set = set(id(x) for x in keep)
    for sid in list(lst):
        lst.remove(sid)
        if id(sid) not in keep_set:
            prs.part.drop_rel(sid.rId)
    for sid in keep:
        lst.append(sid)
    prs.save(out)
    print("saved:", out, len(prs.slides))


if __name__ == "__main__":
    main(*sys.argv[1:4])
