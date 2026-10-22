# -*- coding: utf-8 -*-
"""すなお ver6.8 → ver6.9：「想定動線（フェーズごとのLINEの役割）」の表をやめる（2026-10-22 殿村さん「S17も」）。

    python -m _build.patch_sunao_v69_s17 <ver6.8.pptx> <ver6.9.pptx>

  上の8フェーズの帯（カスタマージャーニーマップ）はそのまま。下の表（9行×4列）を外し、
  左に①〜④、右に⑤〜⑧を縦に並べたブロックにする（帯の左右の並びと同じ）。
  1ブロック＝フェーズ名の札＋3行（心理／離脱課題／LINE）。文言は表の中身そのまま（ラベルはリードの言葉に合わせた）。
  ほかのページは変えない。
"""
import sys

from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

from _build.pptx_parts import INK, LGRAY, NAVY, hline, rect, rich

SUB = "64748B"
TAG = "EEF0F4"
IN = 2.54
LABELS = ["心理", "離脱課題", "LINE"]


def main(src, out):
    prs = Presentation(src)
    s = list(prs.slides)[16]
    assert any(sh.has_text_frame and sh.text_frame.text.startswith("想定動線（フェーズごと") for sh in s.shapes)
    gf = next(sh for sh in s.shapes if sh.has_table)
    t = gf.table
    assert t.cell(0, 0).text == "フェーズ" and len(t.rows) == 9
    rows = [[t.cell(r, c).text for c in range(4)] for r in range(1, 9)]
    x0, y0, w = gf.left / 914400, gf.top / 914400 - 0.05, gf.width / 914400
    gf._element.getparent().remove(gf._element)

    gap_x, h, gap_y = 0.2, 0.82, 0.08
    cw = (w - gap_x) / 2
    tw, lw = 1.5, 0.62                                   # フェーズ名の札／ラベルの幅
    for k, (name, mind, risk, role) in enumerate(rows):
        x = x0 + (cw + gap_x) * (k // 4)
        y = y0 + (h + gap_y) * (k % 4)
        cut = min([i for i in (name.find("（"), name.find(" (")) if i > 0] or [len(name)])
        head, sub = name[:cut].strip(), name[cut:].strip()   # かっこの前で改行（「接触（認知・/流入）」のように途中で切れないように）
        rect(s, x * IN, y * IN, tw * IN, h * IN, TAG,
             [(head, 12, True, NAVY, 0)] + ([(sub, 10.5, True, NAVY, 0)] if sub else []))
        for j, (lab, val) in enumerate(zip(LABELS, (mind, risk, role))):
            ly = y + 0.05 + j * 0.24
            line = j == 2                                # LINEの行だけ紺の太字
            rich(s, (x + tw + 0.08) * IN, ly * IN, lw * IN, 0.24 * IN, [[(lab, 9, line, NAVY if line else SUB)]],
                 anchor=MSO_ANCHOR.MIDDLE)
            rich(s, (x + tw + 0.08 + lw) * IN, ly * IN, (cw - tw - 0.08 - lw) * IN, 0.24 * IN,
                 [[(val, 12, line, NAVY if line else INK)]], anchor=MSO_ANCHOR.MIDDLE)
        if k % 4 != 3:
            hline(s, x * IN, (x + cw) * IN, (y + h + gap_y / 2) * IN, LGRAY, 0.5)
    bottom = y0 + 4 * h + 3 * gap_y
    assert bottom <= 7.05, bottom
    prs.save(out)
    print("saved:", out, round(bottom, 2))


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
