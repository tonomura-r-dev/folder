# -*- coding: utf-8 -*-
"""チェングロウス資料の図形部品（17P・19Pを図解で組み直すとき用）。メイリオ固定。"""
import copy

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Cm, Pt

NAVY, INK, GRAY, LGRAY = "1F285A", "333333", "7F7F7F", "D9D9D9"
CARD, GREEN, GREEN_BG, GREEN_TX = "F4F7FF", "06C755", "E8F8EE", "0B7A3B"
WHITE = "FFFFFF"


def rgb(h):
    return RGBColor.from_string(h)


def style_run(r, size, bold=False, color=INK):
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = rgb(color)
    rpr = r._r.get_or_add_rPr()
    for t in ("a:latin", "a:ea", "a:cs"):
        for e in rpr.findall(qn(t)):
            rpr.remove(e)
    for t in ("a:latin", "a:ea", "a:cs"):
        rpr.append(rpr.makeelement(qn(t), {"typeface": "メイリオ"}))


def fill_text(tf, paras, align=PP_ALIGN.LEFT):
    """paras = [(text, size, bold, color, space_after_pt)]"""
    first = True
    for p_ in paras:
        text, size = p_[0], p_[1]
        bold = p_[2] if len(p_) > 2 else False
        color = p_[3] if len(p_) > 3 else INK
        sa = p_[4] if len(p_) > 4 else 3
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.space_after = Pt(sa)
        p.line_spacing = 1.1
        r = p.add_run()
        r.text = text
        style_run(r, size, bold, color)


def _flat(sh):
    sh.shadow.inherit = False
    st = sh._element.find(qn("p:style"))
    if st is not None:
        st.find(qn("a:effectRef")).set("idx", "0")


def shape(s, kind, x, y, w, h, paras=None, fill=None, line=None, lw=1.25, align=PP_ALIGN.CENTER,
          anchor=MSO_ANCHOR.MIDDLE, adj=None, margins=(0.2, 0.05, 0.2, 0.05)):
    sh = s.shapes.add_shape(kind, Cm(x), Cm(y), Cm(w), Cm(h))
    if adj is not None:
        sh.adjustments[0] = adj
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(fill)
    if line:
        sh.line.color.rgb = rgb(line)
        sh.line.width = Pt(lw)
    else:
        sh.line.fill.background()
    _flat(sh)
    tf = sh.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = [Cm(m) for m in margins]
    if paras:
        fill_text(tf, paras, align)
    return sh


def rect(s, x, y, w, h, fill, paras=None, **kw):
    return shape(s, MSO_SHAPE.RECTANGLE, x, y, w, h, paras, fill=fill, **kw)


def label(s, x, y, w, h, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Cm(0.1)
    tf.margin_top = tf.margin_bottom = Cm(0.05)
    fill_text(tf, paras, align)
    return tb


def _noeffect(c):
    st = c._element.find(qn("p:style"))
    if st is not None:
        st.find(qn("a:effectRef")).set("idx", "0")


def vline(s, x, y1, y2, color, w=1.0, dash=False):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Cm(x), Cm(y1), Cm(x), Cm(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(w)
    if dash:
        c.line.dash_style = 4  # MSO_LINE.DASH
    _noeffect(c)
    return c


def hline(s, x1, x2, y, color, w=0.75):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Cm(x1), Cm(y), Cm(x2), Cm(y))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(w)
    _noeffect(c)
    return c


def find(slide, name):
    return next(sh for sh in slide.shapes if sh.name == name)


def set_para(sh, k, text):
    while len(sh.text_frame.paragraphs) <= k:
        last = sh.text_frame.paragraphs[-1]._p
        last.addnext(copy.deepcopy(last))
    p = sh.text_frame.paragraphs[k]
    runs = p.runs
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def retitle(slide, title, lead1, lead2=None):
    set_para(find(slide, "TextBox 1"), 0, title)
    t2 = find(slide, "TextBox 2")
    set_para(t2, 0, lead1)
    if lead2 is None:
        while len(t2.text_frame.paragraphs) > 1:
            p = t2.text_frame.paragraphs[1]
            p._p.getparent().remove(p._p)
    else:
        set_para(t2, 1, lead2)


def keep_header_only(slide):
    for sh in list(slide.shapes):
        if sh.name not in ("TextBox 1", "TextBox 2", "Connector 3"):
            sh._element.getparent().remove(sh._element)


def set_paras(sh, items):
    """items = [(text, 元の段落番号)]。元の段落の書式をコピーして文言だけ差し替える。"""
    txb = sh.text_frame._txBody
    olds = list(txb.findall(qn("a:p")))
    news = []
    for text, k in items:
        p = copy.deepcopy(olds[k])
        rs = p.findall(qn("a:r"))
        for r in rs[1:]:
            p.remove(r)
        rs[0].find(qn("a:t")).text = text
        news.append(p)
    for p in olds:
        txb.remove(p)
    for p in news:
        txb.append(p)
