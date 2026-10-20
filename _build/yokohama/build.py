# -*- coding: utf-8 -*-
"""株式会社横浜シミズ向け LINE公式アカウント運用のご提案（2026-10-20 殿村さん）。

    python -m _build.yokohama.build <人材（派遣・バイト・単発）業界_LINEOA施策提案.pptx> \
        <_data/yokohama_shimizu/20261008_サンクスLINE誘導のご提案_Appendix_ver1.1.pptx> <出力.pptx>

  依頼：人材（派遣・バイト・単発）の業界資料＋サンクスLINE誘導の資料＋競合3社のLINEの画面＋公式サイトから、
        横浜シミズ向けに約15ページの提案資料へ組み直す。月額20万円（ツール費込み）・初期費用は別途。6か月のSIMを載せる。
  作り方：人材の資料のzipから、表紙・型のページ（LINEOA実績）・裏表紙だけを残し、型のページを複製して中身を組む。
         サンクスLINE誘導の資料の1枚目（仕組み）と3枚目（実績）は、画像ごと取り込んで横浜シミズ向けに直す。
         python-pptxで中身を組み、元のzipに、直したスライドのXMLだけを差し込む（PowerPointの「修復」を避ける）。
"""
import copy
import re
import sys
import zipfile
from io import BytesIO
from pathlib import Path

from lxml import etree
from PIL import Image, ImageFilter
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

import _build.juku_fmt.build as jb
from _build.arauchi_v4.build import slide_order
from _build.arauchi_v4.lib import add_box, add_table
from _build.arauchi_v8.pages import replace_in_paragraphs, shape_by_id
from _build.gym_curves.build import add_pic, keep_only, place, round_box
from _build.gym_fmt.build import fix_fonts, set_title
from _build.juku_fmt.build import CT_SLIDE, RT_SLIDE
from _build.juku_fmt.plan import corner, set_lead

ROOT = Path(__file__).resolve().parent.parent.parent
DATA = ROOT / "_data/yokohama_shimizu"
RT_IMAGE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"
RT_LAYOUT = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideLayout"
PKG = "{http://schemas.openxmlformats.org/package/2006/relationships}"

# ---------------------------------------------------------------- 並び（B＝人材の資料、T＝型のページの複製、A＝サンクスLINEの資料）
B_COVER, B_TPL, B_CLOSE = 1, 31, 35
ORDER = [("B", B_COVER), ("T", 2), ("T", 3), ("T", 4), ("T", 5), ("T", 6), ("T", 7), ("A", 1), ("A", 3),
         ("T", 10), ("T", 11), ("T", 12), ("T", 13), ("T", 14), ("T", 15), ("B", B_CLOSE)]
POS = {key: i for i, key in enumerate(ORDER, 1)}
A_DROP = {1: {41, 43, 50, 52, 54}, 3: {7, 8, 9, 10, 11, 12, 13, 14, 15}}      # 1枚目＝動画のQRコード、3枚目＝クリニックの事例

# 競合のLINEの画面（P5）。名前の写っている所はぼかす（元の画像ファイルは変えない）
SHOTS = [("競合LINE_27.jpg", None), ("競合LINE_28.jpg", None),            # ムーヴ
         ("競合LINE_33.jpg", (100, 262, 420, 314)), ("競合LINE_29.jpg", None),   # ワントゥワン
         ("競合LINE_30.jpg", None), ("競合LINE_32.jpg", (28, 68, 152, 94))]      # ステッドワン
SHOT_PAGE = ("T", 5)

NAVY, NAVY2, INK, GREY, RED = "323366", "1E2761", "333333", "595959", "FF0000"
PALE, PALE2, BORDER, GREEN = "EEF2F9", "F5F6FA", "D9D9D9", "06C755"
X0, W = 0.47, 9.89
BAND_Y = 6.12

# ---------------------------------------------------------------- SIM（6か月・すべて仮定。整数で出す）
UU = 8947                 # サイトUU（集計期間は未確認 → 月間と仮置き）
APPLY_RATE = 0.010        # 応募率（UU→応募）
ADD_RATE = 0.80 * 0.40    # サンクスLINE誘導：遷移率80%×追加率40%（DYMの基準の下限）
BLOCK = 0.30              # ブロック率
REAPPLY = 0.10            # 再応募率（前月末の配信可能なスタッフ数に対する、月あたり）
WORK = 0.70               # 勤務決定率（再応募→再稼働）


def r(x):
    return int(x + 0.5)


def simulate():
    apply_n = r(UU * APPLY_RATE)
    add_n = r(apply_n * ADD_RATE)
    rows = {"apply": [], "add": [], "friends": [], "active": [], "reapply": [], "work": []}
    friends, prev_active = 0, 0
    for m in range(6):
        friends += add_n
        active = r(friends * (1 - BLOCK))
        reapply = r(prev_active * REAPPLY)
        rows["apply"].append(apply_n)
        rows["add"].append(add_n)
        rows["friends"].append(friends)
        rows["active"].append(active)
        rows["reapply"].append(reapply)
        rows["work"].append(r(reapply * WORK))
        prev_active = active
    return rows


SIM = simulate()


# ---------------------------------------------------------------- 1) 組み立て（zipのまま）
def rels_path(name):
    d, f = name.rsplit("/", 1)
    return f"{d}/_rels/{f}.rels"


def resolve(base, target):
    parts = base.split("/")[:-1] + target.split("/")
    out = []
    for p in parts:
        if p == "..":
            out.pop()
        elif p not in (".", ""):
            out.append(p)
    return "/".join(out)


def title_text(xml):
    return "".join(re.findall(r"<a:t>([^<]*)</a:t>", xml)[:2])


def shot_bytes(name, blur):
    im = Image.open(DATA / name).convert("RGB")
    if blur:
        im.paste(im.crop(blur).filter(ImageFilter.GaussianBlur(6 if im.width > 400 else 5)), blur)
    im = im.resize((554, 1200), Image.LANCZOS)
    buf = BytesIO()
    im.save(buf, "JPEG", quality=90)
    return buf.getvalue()


def drop_shapes(xml, ids):
    root = etree.fromstring(xml.encode("utf-8"))
    tree = root.find(qn("p:cSld")).find(qn("p:spTree"))
    n = 0
    for el in list(tree):
        nv = el.find(".//" + qn("p:cNvPr"))
        if nv is not None and int(nv.get("id")) in ids:
            tree.remove(el)
            n += 1
    assert n == len(ids), (n, ids)
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True).decode("utf-8")


def assemble(base, appx, out):
    zb, za = zipfile.ZipFile(base), zipfile.ZipFile(appx)
    ob, oa = slide_order(zb), slide_order(za)
    assert title_text(zb.read(ob[B_TPL - 1]).decode()).startswith("［第5章］LINEOA実績")
    assert "貴社の明るい未来" in zb.read(ob[B_CLOSE - 1]).decode()
    assert title_text(za.read(oa[0]).decode()).startswith("サンクスLINE｜機能・費用")
    assert title_text(za.read(oa[2]).decode()).startswith("サンクスLINE｜導入・活用実績")
    layout2 = "ppt/slideLayouts/slideLayout2.xml"
    assert '<p:cSld name="4_タイトルとコンテンツ">' in zb.read(layout2).decode("utf-8")

    data = {n: zb.read(n) for n in zb.namelist()}
    infos = {i.filename: i for i in zb.infolist()}
    prels = data["ppt/_rels/presentation.xml.rels"].decode("utf-8")
    pres = data["ppt/presentation.xml"].decode("utf-8")
    ct = data["[Content_Types].xml"].decode("utf-8")

    keep_b = {ob[B_COVER - 1], ob[B_CLOSE - 1]}
    rid_b = {}
    for name in ob:                                   # 使わないスライドを外す（presentation の rels・並び・種類の宣言から）
        m = re.search(r'<Relationship [^>]*Target="slides/%s"[^>]*/>' % re.escape(name.split("/")[-1]), prels)
        rid = re.search(r'Id="(rId\d+)"', m.group(0)).group(1)
        if name in keep_b:
            rid_b[name] = rid
            continue
        prels = prels.replace(m.group(0), "")
        pres = re.sub(r'<p:sldId [^>]*r:id="%s"/>' % rid, "", pres)
        for part in (name, rels_path(name)):
            data.pop(part, None)
        ct, c = re.subn(r'<Override [^>]*PartName="/%s"[^>]*/>' % re.escape(name), "", ct)
        assert c == 1, name

    next_rid = max(int(x) for x in re.findall(r'Id="rId(\d+)"', prels)) + 1
    new_parts, new_rids = {}, {}

    def add_slide(key, xml, rels_root):
        nonlocal next_rid, ct, prels
        name = f"ppt/slides/slide{200 + POS[key]}.xml"
        data[name] = xml.encode("utf-8")
        data[rels_path(name)] = etree.tostring(rels_root, xml_declaration=True, encoding="UTF-8", standalone=True)
        ct = ct.replace("</Types>", f'<Override PartName="/{name}" ContentType="{CT_SLIDE}"/></Types>')
        rid = f"rId{next_rid}"
        next_rid += 1
        prels = prels.replace("</Relationships>", f'<Relationship Id="{rid}" Type="{RT_SLIDE}" Target="slides/{name.split("/")[-1]}"/></Relationships>')
        new_rids[key] = rid
        new_parts[name] = key

    tpl_xml = zb.read(ob[B_TPL - 1]).decode("utf-8")
    tpl_rels = etree.fromstring(zb.read(rels_path(ob[B_TPL - 1])))
    for rr in list(tpl_rels):
        if not rr.get("Type").endswith("/slideLayout"):
            tpl_rels.remove(rr)
    for key in ORDER:
        if key[0] == "T":
            rels = copy.deepcopy(tpl_rels)
            if key == SHOT_PAGE:
                for j, (fn, blur) in enumerate(SHOTS, 1):
                    media = f"ppt/media/ys_shot_{j}.jpeg"
                    data[media] = shot_bytes(fn, blur)
                    etree.SubElement(rels, PKG + "Relationship", Id=f"rId{100 + j}", Type=RT_IMAGE,
                                     Target=f"../media/ys_shot_{j}.jpeg")
            add_slide(key, tpl_xml, rels)
        elif key[0] == "A":
            src = oa[key[1] - 1]
            xml = drop_shapes(za.read(src).decode("utf-8"), A_DROP[key[1]])
            rels = etree.fromstring(za.read(rels_path(src)))
            for rr in list(rels):
                typ, tgt = rr.get("Type"), rr.get("Target")
                if typ == RT_LAYOUT:
                    rr.set("Target", "../slideLayouts/slideLayout2.xml")
                elif f'"{rr.get("Id")}"' not in xml:                      # 外した図形だけが使っていた画像・リンク
                    rels.remove(rr)
                elif typ == RT_IMAGE:
                    ext = tgt.rsplit(".", 1)[-1].lower().replace("jpg", "jpeg")
                    stem = tgt.rsplit("/", 1)[-1].rsplit(".", 1)[0]
                    media = f"ppt/media/ys_app{key[1]}_{stem}.{ext}"
                    data[media] = za.read(resolve(src, tgt))
                    assert f'Extension="{ext}"' in ct, ext
                    rr.set("Target", f"../media/{media.split('/')[-1]}")
                else:
                    assert rr.get("TargetMode") == "External", typ
            add_slide(key, xml, rels)

    # 並び
    m = re.search(r"<p:sldIdLst>(.*?)</p:sldIdLst>", pres, re.S)
    entries = {re.search(r'r:id="(rId\d+)"', e).group(1): e for e in re.findall(r"<p:sldId [^>]*/>", m.group(1))}
    next_id = max(int(x) for x in re.findall(r'<p:sldId id="(\d+)"', m.group(1))) + 1
    lst = []
    for key in ORDER:
        if key[0] == "B":
            lst.append(entries[rid_b[ob[key[1] - 1]]])
        else:
            lst.append(f'<p:sldId id="{next_id}" r:id="{new_rids[key]}"/>')
            next_id += 1
    pres = pres.replace(m.group(0), "<p:sldIdLst>" + "".join(lst) + "</p:sldIdLst>")

    # どこからも参照されないパーツ（外したスライドのノート・画像など）を外す
    data["ppt/_rels/presentation.xml.rels"] = prels.encode("utf-8")
    seen, todo = set(), ["_rels/.rels"]
    reach = {"[Content_Types].xml", "_rels/.rels"}
    while todo:
        rp = todo.pop()
        if rp in seen or rp not in data:
            continue
        seen.add(rp)
        src = rp.replace("_rels/", "").removesuffix(".rels")
        for rr in etree.fromstring(data[rp]):
            if rr.get("TargetMode") == "External":
                continue
            t = rr.get("Target")
            t = t.lstrip("/") if t.startswith("/") else (resolve(src, t) if src else t)
            reach.add(t)
            if rels_path(t) in data:
                reach.add(rels_path(t))
                todo.append(rels_path(t))
    gone = sorted(n for n in data if n not in reach)
    for n in gone:
        data.pop(n)
        ct = re.sub(r'<Override [^>]*PartName="/%s"[^>]*/>' % re.escape(n), "", ct)

    app = data["docProps/app.xml"].decode("utf-8")
    nb = len(ob)
    nn = len(ORDER)
    app = app.replace(f"<Slides>{nb}</Slides>", f"<Slides>{nn}</Slides>")
    app = re.sub(r"<HiddenSlides>\d+</HiddenSlides>", "<HiddenSlides>0</HiddenSlides>", app)
    notes = sum(1 for n in data if re.match(r"ppt/notesSlides/notesSlide\d+\.xml$", n))
    app = re.sub(r"<Notes>\d+</Notes>", f"<Notes>{notes}</Notes>", app)
    head = f"<vt:lpstr>スライド タイトル</vt:lpstr></vt:variant><vt:variant><vt:i4>{nb}</vt:i4>"
    assert head in app
    app = app.replace(head, head.replace(f"<vt:i4>{nb}</vt:i4>", f"<vt:i4>{nn}</vt:i4>"))
    mv = re.search(r'<vt:vector size="(\d+)" baseType="lpstr">(.*?)</vt:vector>', app, re.S)
    items = re.findall(r"<vt:lpstr>.*?</vt:lpstr>", mv.group(2))
    items = items[:len(items) - nb] + ["<vt:lpstr>PowerPoint プレゼンテーション</vt:lpstr>"] * nn
    app = app.replace(mv.group(0), f'<vt:vector size="{len(items)}" baseType="lpstr">' + "".join(items) + "</vt:vector>")
    data["docProps/app.xml"] = app.encode("utf-8")
    data["ppt/presentation.xml"] = pres.encode("utf-8")
    data["[Content_Types].xml"] = ct.encode("utf-8")

    with zipfile.ZipFile(out, "w") as zo:
        for n, d in data.items():
            if n in infos:
                zo.writestr(infos[n], d, compress_type=infos[n].compress_type)
            else:
                zo.writestr(n, d, compress_type=zipfile.ZIP_STORED if n.endswith((".jpeg", ".png")) else zipfile.ZIP_DEFLATED)
    return gone


# ---------------------------------------------------------------- 2) 部品
def runs_size(sh, pt):
    for p in sh.text_frame.paragraphs:
        for rn in p.runs:
            rn.font.size = Pt(pt)


def lead(s, l1, l2=None):
    set_lead(s, [l1] + ([l2] if l2 else []))
    sh = shape_by_id(s, 3)
    place(sh, h=0.64 if l2 else 0.32)
    if l2:
        p2 = sh.text_frame.paragraphs[1]
        for rn in p2.runs:
            rn.font.color.rgb = RGBColor.from_string(RED)


def base_page(s, title, l1, l2=None, band=None, foot=None):
    keep_only(s, {2, 3, 18, 20})
    set_title(shape_by_id(s, 2), title)
    lead(s, l1, l2)
    if band is None:
        el = shape_by_id(s, 18)._element
        el.getparent().remove(el)
    else:
        place(shape_by_id(s, 18), x=X0, y=BAND_Y, w=W, h=0.41)
        add_box(s, X0, BAND_Y, W, 0.41, [[(band, {})]], size=13, bold=True, color="FFFFFF", align="ctr",
                ins=(0.05, 0, 0.05, 0))
    ft = shape_by_id(s, 20)
    if foot:
        set_title(ft, foot)
        runs_size(ft, 9)
        place(ft, y=6.66, h=0.5)
    else:
        ft._element.getparent().remove(ft._element)


def chip(s, x, y, w, h, text, size=11, fill=NAVY2, color="FFFFFF"):
    return round_box(s, x, y, w, h, [[(text, {})]], adj=20000, size=size, bold=True, color=color, fill=fill,
                     align="ctr", ins=(0.04, 0, 0.04, 0))


def card(s, x, y, w, h, paras, fill="FFFFFF", line=BORDER, size=11, anchor="t", ins=(0.12, 0.08, 0.1, 0.06), **kw):
    return round_box(s, x, y, w, h, paras, adj=5000, size=size, color=INK, fill=fill, line=line, anchor=anchor,
                     ins=ins, **kw)


def arrow(s, x1, y1, x2, y2, color=NAVY2, w=2.0):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    st = c._element.find(qn("p:style"))
    if st is not None:
        c._element.remove(st)
    c.line.color.rgb = RGBColor.from_string(color)
    c.line.width = Pt(w)
    etree.SubElement(c._element.spPr.find(qn("a:ln")), qn("a:tailEnd"), type="triangle", w="med", len="med")
    return c


def head(text, opt=None):
    return [(text, dict({"bold": True, "color": NAVY}, **(opt or {})))]


def cell(text, **kw):
    d = {"text": text, "size": kw.pop("size", 10.5)}
    d.update(kw)
    return d


# ---------------------------------------------------------------- 3) 各ページ
def p_cover(s):
    sh = shape_by_id(s, 2)
    set_title(sh, "LINE公式アカウント運用のご提案")
    set_title(shape_by_id(s, 6), "応募から勤務連絡、次回イベントへの再応募までをLINEでつなぐ")
    for sid in (3, 7):
        el = shape_by_id(s, sid)._element
        el.getparent().remove(el)
    add_box(s, 0.32, 0.29, 5.43, 0.37, [[("株式会社横浜シミズ 御中", {})]], size=16, bold=True, color=GREY,
            ins=(0.05, 0, 0.05, 0))
    add_box(s, 1.46, 4.3, 7.91, 0.4, [[("2026年10月　株式会社DYM", {})]], size=12, color=GREY, align="ctr",
            ins=(0, 0, 0, 0))


def p_current(s):
    base_page(s, "現状把握｜横浜シミズの事業と、想定している課題",
              "神奈川・静岡全域のイベント会場で、運営・警備・設営を担う、登録制アルバイトの会社です。",
              "イベントごとに人を集める事業だからこそ、勤務後の接点が次の人員確保につながると考えます。",
              band="一度働いた方に、次のイベントでも戻ってきていただく仕組みをつくる",
              foot="出典｜株式会社横浜シミズ 公式サイト（事業内容・会社情報・アルバイト募集）・採用サイト（shimizuocto-saiyou.net）"
                   "　2026年10月20日 弊社確認")
    card(s, X0, 1.5, 5.6, 4.45, [[("", {})]], fill=PALE2)
    chip(s, X0 + 0.15, 1.62, 3.3, 0.32, "事業の特徴（公式サイト・採用サイトより）", size=10.5)
    items = ["神奈川県・静岡県全域で、舞台設営・美術施工・\n運営管理・警備をワンストップで提供",
             "運営管理は、少人数から1日1,000人以上の規模まで動員",
             "アルバイトは登録制。勤務できる日を自己申告し、\n合う業務に就く",
             "一度採用されると、ほかの業務にも勤務できる（面接は1回）",
             "勤務地は日産スタジアム・横浜アリーナ・横浜スタジアム・\nKアリーナ横浜など多数",
             "応募は採用サイトから。面接予約の完了はメールで届く"]
    box = add_box(s, X0 + 0.12, 2.1, 5.36, 3.75, [[("・" + t, {})] for t in items], size=12, color=INK,
                  anchor="t", ins=(0.06, 0.04, 0.06, 0.04), after=8, line_spacing=105)
    from _build.juku_fmt.plan import hang
    hang(box, 0.17)
    x2, w2 = 6.25, 4.11
    rows = [[cell("いただいている情報", bold=True, color="FFFFFF", fill=NAVY2, gridSpan=2, align="l"),
             cell("", hMerge=True, fill=NAVY2)]]
    for k, v in [("サイトUU", "8,947（集計期間：未確認）"), ("月間の応募数", "未確認"), ("広告の実施状況", "未確認"),
                 ("LINE公式アカウントの\n現在の使い方", "未確認"), ("想定予算", "月額20万円（ツール費込み）\n初期費用は別途")]:
        rows.append([cell(k.split("\n"), bold=True, color=NAVY, fill=PALE, align="l"),
                     cell(v.split("\n"), color=RED if v == "未確認" else INK, align="l")])
    add_table(s, x2, 1.5, [1.75, w2 - 1.75], [0.34, 0.4, 0.4, 0.4, 0.55, 0.55], rows)
    card(s, x2, 4.25, w2, 1.7,
         [head("想定している課題（仮説）"),
          [("イベントの勤務が終わると、スタッフの方との連絡が途切れやすい。", {})],
          [("次のイベントでも働いていただける接点を残すことが、継続的な人員確保の鍵と考えます。", {})]],
         size=11, after=3, line=NAVY2)


def p_issue1(s):
    base_page(s, "業界課題①｜イベント・単発の仕事は、毎回の人員確保が前提",
              "単発・短期の仕事は、イベントのたびに必要な人数を集め直す必要があります。",
              "新規の応募だけに頼ると、募集・採用の負担がイベントごとにくり返されます。",
              band="比べて選ばれる仕事だからこそ、一度つながった方との接点が次の人員確保を左右する",
              foot="出典｜LINEヤフー 媒体データ（検索の前後行動）。起点キーワード「単発バイト」・直近1年間"
                   "（2025年3月17日〜2026年3月22日）。分類は弊社")
    cards = [("① イベントのたびに、必要な人数を集め直す", "開催日・会場・規模によって、必要な人数と業務が変わります"),
             ("② 新規の応募だけに頼ると、負担がくり返される", "募集・面接・登録の手間が、イベントごとにかかります"),
             ("③ 勤務が終わると、接点が薄れる", "次の募集を届ける手段が残らず、また新規で集めることになります")]
    for i, (h, b) in enumerate(cards):
        card(s, X0, 1.5 + i * 1.5, 4.55, 1.35, [head(h, {"size": 12}), [(b, {})]], size=11.5, after=4,
             anchor="ctr", line=NAVY2)
    x2, w2 = 5.22, 5.14
    card(s, x2, 1.5, w2, 4.35, [[("", {})]], fill=PALE2)
    add_box(s, x2 + 0.15, 1.58, w2 - 0.3, 0.55,
            [head("「単発バイト」を検索した方は、前後に何を調べているか", {"size": 12})],
            size=12, anchor="ctr", ins=(0, 0, 0, 0))
    rows = [("-7〜-1日", "スキマ時間で、すぐ稼ぎたい", "副業 在宅 スキマ時間／在宅ワーク 求人／バイト 求人"),
            ("0〜+2日", "同じ日に、複数のサービスを並べて比べる", "タイミー／バイトル／短期バイト／単発バイト アプリ"),
            ("+1〜+15日", "続けられるか・扶養や税金は大丈夫か", "副業 確定申告 20万円以下／ダブルワーク 扶養内 いくらまで")]
    for i, (d, h, q) in enumerate(rows):
        y = 2.2 + i * 1.08
        chip(s, x2 + 0.15, y + 0.05, 1.05, 0.3, d, size=10)
        add_box(s, x2 + 1.3, y, w2 - 1.45, 0.95,
                [[(h, {"bold": True, "color": NAVY, "size": 11.5})], [(q, {"color": GREY, "size": 10})]],
                size=11, anchor="t", ins=(0.02, 0.02, 0.02, 0.02), after=2)
    add_box(s, x2 + 0.15, 5.4, w2 - 0.3, 0.4,
            [[("★ 検索行動のデータです。応募者ご本人の行動と同一ではありません。", {})]], size=9.5, color=GREY,
            ins=(0, 0, 0, 0))


def p_issue2(s):
    base_page(s, "業界課題②｜応募後の連絡・勤務前の案内・次の募集に手間がかかる",
              "応募から勤務までの連絡、勤務前の案内、次の募集。どれも人数が増えるほど手間が増えます。",
              "連絡をLINEに集めることで、届きやすく、次の募集にもつなげやすくなります。",
              band="連絡の手間を減らしながら、次の募集が届く状態をつくる",
              foot="出典｜LINE国内MAU・メッセージ開封率：LINEヤフー公式値。※応募から勤務までの実施率・再応募率は"
                   "公開データがなく、貴社での計測が必要です")
    cols = [("応募のあと", "メールや電話がつながらず、面接・登録の案内が届きにくい",
             "応募の直後に友だちになっていただき、面接・登録の案内をLINEで届ける"),
            ("勤務の前", "集合場所・時間・持ち物・服装の案内を、人数分くり返し送る",
             "同じイベントの方へ一斉に、会場・業務ごとに絞って案内する"),
            ("勤務のあと", "次の募集を届ける手段が残らず、新規の募集に頼りがち",
             "希望に合う次回イベントの募集を届け、再応募につなげる")]
    cw, gap = 3.15, (W - 3 * 3.15) / 2
    for i, (h, a, b) in enumerate(cols):
        x = X0 + i * (cw + gap)
        chip(s, x, 1.5, cw, 0.34, h, size=12)
        card(s, x, 1.95, cw, 1.15, [[("起きやすいこと（一般的な傾向）", {"bold": True, "color": GREY, "size": 10})],
                                    [(a, {})]], fill=PALE2, size=11.5, after=3)
        add_box(s, x + cw / 2 - 0.25, 3.13, 0.5, 0.3, [[("▼", {})]], size=14, color=NAVY2, align="ctr",
                ins=(0, 0, 0, 0))
        card(s, x, 3.45, cw, 1.15, [[("LINEでできること", {"bold": True, "color": NAVY, "size": 10})], [(b, {})]],
             line=NAVY2, size=11.5, after=3)
    sw = (W - 0.22) / 2
    for i, (lab, big, sub) in enumerate([("メッセージ開封率（当日中）", "約8割", "受信直後 約2割／3〜6時間で約5割"),
                                         ("LINE国内MAU", "1億人突破", "2025年12月末時点")]):
        x = X0 + i * (sw + 0.22)
        card(s, x, 4.8, sw, 1.12, [[("", {})]], fill=PALE)
        add_box(s, x + 0.15, 4.85, sw * 0.55, 1.02,
                [[(lab, {"bold": True, "color": NAVY, "size": 11.5})], [(sub, {"color": GREY, "size": 10})]],
                size=11, anchor="ctr", ins=(0, 0, 0, 0), after=2)
        add_box(s, x + sw * 0.55, 4.85, sw * 0.45 - 0.15, 1.02, [[(big, {})]], size=24, bold=True, color=NAVY,
                align="r", anchor="ctr", ins=(0, 0, 0, 0))


def p_comp_screens(s):
    base_page(s, "競合分析①｜イベント・単発の人材会社のLINE運用（実際の画面）",
              "3社とも、LINEを応募の受付だけでなく、登録・選考・勤務のやり取りに使っています。",
              "画面から確認できたことだけを記載しています（2026年10月 弊社確認）。",
              foot="画面：各社のLINE公式アカウント（2026年10月 弊社確認）。個人名はぼかしています")
    groups = [("ムーヴ｜登録の入口をLINEに", ["・登録の手順を4つに分けて案内", "・登録／ID連携のボタンを表示",
                                          "・メニュー：仕事情報・Q&A・問い合わせ"]),
              ("ワントゥワン｜選考の案内をLINEで", ["・追加の直後に、氏名をLINEで送ってもらう", "・録画面接のURLを案内",
                                              "・合否はメール、質問はLINEで受付"]),
              ("ステッドワン｜勤務のやり取りをLINEで", ["・このLINEで仕事をやり取りすると宣言", "・新規案件・勤務の可否・集合詳細を送る",
                                                 "・スタッフからは起床連絡・勤怠報告"])]
    gw, gap = 3.15, (W - 3 * 3.15) / 2
    iw = 1.45
    ih = iw * 600 / 277
    for i, (h, bullets) in enumerate(groups):
        x = X0 + i * (gw + gap)
        chip(s, x, 1.5, gw, 0.34, h, size=11)
        x_img = x + (gw - 2 * iw - 0.12) / 2
        for j in range(2):
            k = i * 2 + j + 1
            add_pic(s, f"rId{100 + k}", f"競合LINE_{k}", x_img + j * (iw + 0.12), 1.95, iw, ih)
        add_box(s, x, 1.95 + ih + 0.08, gw, 5.98 - (1.95 + ih + 0.08), [[(b, {})] for b in bullets], size=10.5,
                color=INK, anchor="t", ins=(0.04, 0.02, 0.02, 0.02), after=1)


def p_comp_table(s):
    base_page(s, "競合分析②｜LINEの使い方の比較と、横浜シミズで目指す形",
              "応募受付だけでなく、スタッフとの継続的な接点にLINEを活用する余地があります。",
              "3社の画面で、勤務の連絡と次の案件までLINEで扱っていたのは、ステッドワンでした。",
              band="応募の受付から、勤務の連絡・次の募集までをLINEでつなぐ",
              foot="画面：各社のLINE公式アカウント（2026年10月 弊社確認）／タウンワーク：LINEヤフー for Business 導入事例"
                   "「バイト応募件数約80％増！リッチメニューを生かしたタウンワークのLINE活用事例」（2019年1月公開）")
    U = "画面では未確認"
    data = [("友だち追加後の案内", "登録までの手順を4つで案内", "氏名（漢字・ふりがな）の送信を依頼", "このLINEでのやり取りを説明",
             "勤務に必要な情報が届くことを最初に伝える"),
            ("応募・登録の導線", "LINEからスタッフ登録・ID連携", "録画面接のURLを案内", "採用手続きフォーム（外部）へ",
             "応募完了からそのままLINEへ（サンクスLINE誘導）"),
            ("勤務連絡への活用", U, U + "（合否はメール）", "集合詳細・起床連絡・勤怠報告", "集合・持ち物・前日のリマインド"),
            ("募集情報の配信", "登録後に「条件に合う仕事も案内」と記載", U, "新規案件をLINEで送る", "エリア・会場・希望で絞って届ける"),
            ("再応募への導線", "「お仕事一覧から申込み」と記載", U, "案件へのエントリーをLINEで送る", "次回イベントの募集から再応募へ"),
            ("リッチメニュー", "仕事情報・Q&A・問い合わせ＋登録", U, "「案件一覧」あり（ほかは未確認）", "仕事を探す・勤務の案内・よくある質問など")]
    hdr = ["評価項目", "ムーヴ", "ワントゥワン", "ステッドワン", "横浜シミズで目指す形"]
    rows = [[cell(t, bold=True, color="FFFFFF", fill=NAVY2 if k < 4 else NAVY, size=10.5) for k, t in enumerate(hdr)]]
    for row in data:
        cells = [cell(row[0], bold=True, color=NAVY, fill=PALE, size=10.5)]
        for k, t in enumerate(row[1:], 1):
            cells.append(cell(t, size=10, color=GREY if t.startswith(U) else INK,
                              fill="E4E9F2" if k == 4 else "FFFFFF", bold=(k == 4)))
        rows.append(cells)
    add_table(s, X0, 1.5, [1.45, 2.0, 1.85, 2.05, 2.54], [0.36] + [0.56] * 6, rows)
    card(s, X0, 5.32, W, 0.66,
         [[("参考｜タウンワーク（求人メディア）は、LINEを求人検索の入口に活用し、リッチメニュー経由の応募数が＋79%"
            "（2019年公開の導入事例）。直接の競合ではなく、LINEを応募の入口にした参考事例です。", {})]],
         fill=PALE2, size=10, anchor="ctr", ins=(0.12, 0.04, 0.12, 0.04))


def p_overview(s):
    base_page(s, "提案の全体像｜応募から再応募までを、LINEで循環させる",
              "応募の時点で友だちになっていただき、勤務の連絡と次の募集をLINEでお届けします。",
              "イベントが終わっても接点が切れず、次のイベントの人員確保につながる仕組みです。",
              band="単発の勤務を、イベントごとに途切れない関係に変える")

    def node(x, y, w, h, tag, name, desc, fill="FFFFFF", line=NAVY2):
        paras = []
        if tag:
            paras.append([(tag, {"bold": True, "color": GREEN if tag.startswith("施策①") else NAVY, "size": 10})])
        paras.append([(name, {"bold": True, "color": NAVY, "size": 14})])
        if desc:
            paras.append([(desc, {"color": INK, "size": 10.5})])
        return card(s, x, y, w, h, paras, fill=fill, line=line, line_w=1.25, align="ctr", anchor="ctr",
                    ins=(0.08, 0.04, 0.08, 0.04), after=1)

    node(X0, 1.65, 2.05, 1.05, None, "新規応募", "採用サイトから応募", fill=PALE2)
    arrow(s, X0 + 1.025, 2.72, X0 + 1.025, 3.33)
    node(X0, 3.35, 2.05, 1.25, "施策① 友だち獲得", "LINE友だち追加", "応募完了からLINEへ", fill="E9F8EF", line=GREEN)
    lx1, lx2, nw, nh = 3.25, 7.7, 2.66, 1.15
    ty, by = 1.6, 4.0
    arrow(s, X0 + 2.07, 3.97, lx1 - 0.02, 2.35)
    node(lx1, ty, nw, nh, "施策② 勤務連絡・管理", "勤務連絡", "集合・持ち物・前日のリマインド")
    node(lx2, ty, nw, nh, None, "イベント勤務", "現場で勤務")
    node(lx2, by, nw, nh, "施策③ 再応募・再稼働", "次回募集", "エリア・会場・希望で絞って案内")
    node(lx1, by, nw, nh, "施策③ 再応募・再稼働", "再応募・継続", "次のイベントにも勤務")
    arrow(s, lx1 + nw + 0.05, ty + nh / 2, lx2 - 0.05, ty + nh / 2)
    arrow(s, lx2 + nw / 2, ty + nh + 0.05, lx2 + nw / 2, by - 0.05)
    arrow(s, lx2 - 0.05, by + nh / 2, lx1 + nw + 0.05, by + nh / 2)
    arrow(s, lx1 + nw / 2, by - 0.05, lx1 + nw / 2, ty + nh + 0.05)
    add_box(s, lx1 + nw / 2 + 0.2, ty + nh + 0.3, lx2 - lx1 - 0.4, by - ty - nh - 0.6,
            [[("イベントのたびに", {})], [("くり返す", {})]], size=12, bold=True, color=GREY, align="ctr",
            ins=(0, 0, 0, 0))
    add_box(s, lx1, 5.3, lx2 + nw - lx1, 0.62,
            [[("応募の時点で友だちになり、イベントが終わっても接点が残る。登録スタッフの皆さまに、次の募集が届きます。", {})]],
            size=11, color=INK, align="ctr", anchor="ctr", ins=(0.05, 0, 0.05, 0))


def p_thanks(s):
    """サンクスLINE誘導の資料の1枚目（仕組み）を取り込み、横浜シミズの導線を足す"""
    set_title(shape_by_id(s, 22), "施策①｜友だち獲得：応募完了から、そのままLINEへ")
    assert replace_in_paragraphs(s, "CV完了後約2~3秒で", "応募後 約2〜3秒で") == 1
    card(s, X0, 5.89, 5.3, 1.25,
         [head("横浜シミズでの導線", {"size": 11}),
          [("採用サイトで応募 → 面接予約の完了 → 自動でLINEへ → 友だち追加 → 勤務の案内・次回の募集", {})],
          [("応募を済ませた、働く意思のある方を、そのまま友だちにできます。", {"color": NAVY})],
          [("※採用サイト（応募フォーム）の完了画面に設置できるかは確認事項です。", {"color": GREY, "size": 9})]],
         fill="FFFFFF", line=NAVY2, size=10, after=1, ins=(0.1, 0.05, 0.08, 0.04))


def p_record(s):
    """サンクスLINE誘導の資料の3枚目（実績）を取り込み、クリニックの事例を横浜シミズの活用イメージに替える"""
    set_title(shape_by_id(s, 3), "施策①｜サンクスLINE誘導の導入実績（人材業界）")
    set_title(shape_by_id(s, 4), "人材業界での導入実績（DYM）と、横浜シミズでの活用イメージ")
    add_box(s, 0.84, 1.78, 4.0, 0.3, [head("横浜シミズでの活用イメージ", {"size": 12})], size=12,
            ins=(0, 0, 0, 0))
    steps = [("応募完了", NAVY2), ("LINE\n友だち追加", GREEN), ("面接・登録の\n案内", NAVY2), ("勤務前の\n案内", NAVY2),
             ("次回イベントの\n募集", NAVY2)]
    bw, gap, x = 1.52, 0.33, 0.95
    for i, (t, f) in enumerate(steps):
        add_box(s, x + i * (bw + gap), 2.15, bw, 0.56, [[(ln, {})] for ln in t.split("\n")], size=10.5, bold=True,
                color="FFFFFF", fill=f, align="ctr", anchor="ctr", ins=(0.03, 0, 0.03, 0))
        if i:
            arrow(s, x + i * (bw + gap) - gap + 0.04, 2.43, x + i * (bw + gap) - 0.04, 2.43, w=1.5)
    add_box(s, 0.95, 2.78, 9.0, 0.36,
            [[("応募の直後から、面接・勤務の案内と次の募集を同じLINEでお届けします。", {})]], size=10.5,
            color=INK, ins=(0, 0, 0, 0))
    ft = shape_by_id(s, 2)
    set_title(ft, "※出典：人材業界｜株式会社DYM 社内集計（サンクスLINE誘導ツール導入前後3か月、2025年・2026年7〜9月の友だち増加数）")
    p = ft.text_frame.paragraphs[0]
    p2 = copy.deepcopy(p._p)
    p._p.addnext(p2)
    ft.text_frame.paragraphs[1].runs[0].text = "※他社での実績であり、横浜シミズ様での効果を示すものではありません。"
    runs_size(ft, 9)
    place(ft, x=0.47, y=6.62, w=9.6, h=0.5)


def p_work(s):
    base_page(s, "施策②｜勤務連絡・管理：勤務に必要な情報をLINEで届ける",
              "集合場所・時間・持ち物・服装・注意事項を、勤務の前にLINEでお届けします。",
              "「LINEを追加すると、勤務に必要な情報が届く」ことが、友だちでいる理由になります。",
              band="届く理由があるから友だちが残り、残るから次の募集が届く")
    cw, gap = 3.15, (W - 3 * 3.15) / 2
    cols = [("勤務の前にお届けする情報", [("勤務が決まったら", "会場・集合場所・集合時間・持ち物・服装"),
                                     ("勤務の前日", "リマインド（集合時間・持ち物・天候による変更など）"),
                                     ("勤務の当日", "注意事項・当日の連絡先")]),
            ("届け方（運用に合わせて選びます）", [("一斉案内", "同じイベントに勤務する方へ、まとめて案内"),
                                          ("タグで対象者別", "会場・業務ごとに、対象の方だけに案内"),
                                          ("個別連絡（チャット）", "勤務の変更・欠席・質問など、1対1のやり取り")]),
            ("期待できること", [("スタッフの皆さま", "必要な情報をLINEで見返せる。通知で気づける。質問しやすい"),
                              ("横浜シミズ様（管理側）", "一斉に送れる。案内の型をつくり、毎回の作成の手間を減らす"),
                              ("横浜シミズ様（管理側）", "電話・メールでの個別の確認を減らせる（期待効果）")])]
    for i, (h, items) in enumerate(cols):
        x = X0 + i * (cw + gap)
        chip(s, x, 1.5, cw, 0.36, h, size=11.5)
        for j, (t, b) in enumerate(items):
            card(s, x, 1.98 + j * 1.12, cw, 1.0, [[(t, {"bold": True, "color": NAVY, "size": 11.5})], [(b, {})]],
                 fill=PALE2 if i < 2 else "FFFFFF", line=BORDER if i < 2 else NAVY2, size=11, after=2)
    add_box(s, X0, 5.38, W, 0.6,
            [[("※勤務の予定とLINEの自動連携は前提にしていません。一斉案内・タグ・個別連絡の組み合わせで始め、"
               "自動連携が必要な場合は別途確認します。", {})]], size=10, color=GREY, ins=(0.02, 0, 0.02, 0))


def p_reapply(s):
    base_page(s, "施策③｜再応募・再稼働：希望に合う募集だけを届ける",
              "エリア・会場・業務・曜日などをタグで管理し、対象を絞ってイベントの募集をお届けします。",
              "合わない募集が届かないので、ブロックを抑え、配信の手間も減らせます。",
              band="取れる情報から始めて、配信のたびに絞り込みの精度を上げる")
    rows = [[cell(t, bold=True, color="FFFFFF", fill=NAVY2, size=10.5) for t in ("タグ", "例", "取り方")]]
    for t, e, g in [("エリア", "神奈川（横浜・川崎・湘南など）／静岡", "追加後の質問（選択式）"),
                    ("会場", "日産スタジアム・横浜アリーナ・横浜スタジアムなど", "追加後の質問・勤務の案内時"),
                    ("業務", "案内・誘導／販売／警備／設営・撤去", "追加後の質問"),
                    ("曜日・時間帯", "土日祝／平日／夜", "追加後の質問・メニューで変更"),
                    ("勤務の有無", "勤務あり／まだ勤務なし", "勤務の案内を送った方に付ける")]:
        rows.append([cell(t, bold=True, color=NAVY, fill=PALE, size=10.5), cell(e, size=10.5, align="l"),
                     cell(g, size=10, align="l")])
    add_table(s, X0, 1.5, [1.25, 2.45, 1.7], [0.36] + [0.6] * 5, rows)
    add_box(s, X0, 4.95, 5.4, 1.0,
            [[("※取れる情報から、段階的にタグを付けます。最初から高度なシステム連携は前提にしません。", {})],
             [("※勤務の履歴・希望条件をいまどこで管理しているかは確認事項です。", {})]],
            size=10, color=GREY, anchor="t", ins=(0.02, 0.02, 0.02, 0), after=2)
    x2, w2 = 6.12, 4.24
    chip(s, x2, 1.5, w2, 0.36, "配信の出し分けの例", size=11.5)
    for i, (who, what) in enumerate([("横浜エリアを希望", "横浜周辺のイベントの募集"), ("特定の会場に関心", "その会場のイベントの募集"),
                                     ("土日に勤務したい", "休日のイベントの募集"), ("一度勤務した方", "次回の募集を先にご案内")]):
        y = 2.0 + i * 0.86
        add_box(s, x2, y, 1.75, 0.66, [[(who, {})]], size=11, bold=True, color=NAVY, fill=PALE, align="ctr",
                anchor="ctr", ins=(0.04, 0, 0.04, 0))
        arrow(s, x2 + 1.8, y + 0.33, x2 + 2.12, y + 0.33, w=1.75)
        card(s, x2 + 2.17, y, w2 - 2.17, 0.66, [[(what, {})]], size=11, anchor="ctr", line=NAVY2, align="ctr",
             ins=(0.04, 0, 0.04, 0))
    add_box(s, x2, 5.45, w2, 0.5, [[("※配信設計の例です。", {})]], size=10, color=GREY, anchor="t",
            ins=(0.02, 0.02, 0.02, 0))


def p_image(s):
    base_page(s, "配信イメージ｜勤務前の案内・次回の募集・リッチメニュー",
              "横浜シミズ向けの配信の文面と、リッチメニューの構成案です。",
              "文面・項目は例です。運用を始める前にすり合わせて確定します。",
              band="勤務に必要な情報と次の募集が、同じLINEに届く")
    cw, gap = 3.15, (W - 3 * 3.15) / 2
    talks = [("勤務前の案内（前日）", "対象：そのイベントに勤務する方（タグ）",
              ["【明日の勤務のご案内】", "〇月〇日（〇）〇〇スタジアム", "集合：〇〇:〇〇　〇〇ゲート前", "持ち物：〇〇・〇〇",
               "服装：〇〇", "遅れる場合・体調がすぐれないときは、このLINEでご連絡ください。"], None),
             ("次回イベントの募集", "対象：横浜エリア・土日希望・勤務ありの方",
              ["【スタッフ募集｜横浜アリーナ】", "〇月〇日（〇）〇〇時〜〇〇時", "案内・誘導スタッフ　時給〇〇円〜",
               "前回のご勤務、ありがとうございました。次のイベントもお待ちしています。"], "募集の詳細を見る")]
    for i, (h, target, lines, btn) in enumerate(talks):
        x = X0 + i * (cw + gap)
        chip(s, x, 1.5, cw, 0.36, h, size=11.5)
        add_box(s, x, 1.95, cw, 3.55, [[("", {})]], fill="D8E2F1")
        bubble = card(s, x + 0.18, 2.1, cw - 0.36, 2.55 if btn else 3.2,
                      [[(t, {"bold": k == 0, "color": NAVY if k == 0 else INK})] for k, t in enumerate(lines)],
                      size=10.5, after=2, line=None, ins=(0.12, 0.08, 0.1, 0.06))
        if btn:
            add_box(s, x + 0.35, 4.75, cw - 0.7, 0.45, [[(btn, {})]], size=11, bold=True, color="FFFFFF", fill=GREEN,
                    align="ctr", anchor="ctr", ins=(0, 0, 0, 0))
        add_box(s, x, 5.55, cw, 0.4, [[(target, {})]], size=10, color=GREY, ins=(0.02, 0, 0.02, 0))
    x = X0 + 2 * (cw + gap)
    chip(s, x, 1.5, cw, 0.36, "リッチメニュー（構成案）", size=11.5)
    add_box(s, x, 1.95, cw, 3.55, [[("", {})]], fill="F2F2F2")
    menu = [("仕事を探す", "募集中のイベント"), ("勤務の案内", "集合・持ち物"), ("希望条件を変える", "エリア・曜日"),
            ("よくある質問", "給与・服装など"), ("お問い合わせ", "チャットで相談"), ("横浜シミズとは", "会社・現場の紹介")]
    mw, mh = (cw - 0.3) / 2, 1.0
    for k, (t, d) in enumerate(menu):
        cx = x + 0.1 + (k % 2) * (mw + 0.1)
        cy = 2.1 + (k // 2) * (mh + 0.1)
        add_box(s, cx, cy, mw, mh, [[(t, {"bold": True, "size": 11.5})], [(d, {"size": 9.5})]], size=11,
                color="FFFFFF" if k == 0 else NAVY, fill=NAVY2 if k == 0 else "FFFFFF", line=None if k == 0 else BORDER,
                align="ctr", anchor="ctr", ins=(0.04, 0.02, 0.04, 0.02), after=1)
    add_box(s, x, 5.55, cw, 0.4, [[("一番押してほしい「仕事を探す」を左上に", {})]], size=10, color=GREY,
            ins=(0.02, 0, 0.02, 0))


def p_sim(s):
    base_page(s, "目標・6か月SIM｜友だち獲得から再応募・再稼働まで（試算）",
              "サイトUU 8,947を月間と仮置きし、仮定の率から6か月の見込みを試算しました。",
              "応募数と各率は仮定です。実数をいただいたら置き換えて試算し直します。")
    names = [("新規応募数", "UU × 応募率 1.0%", "apply", True), ("LINE友だち追加数", "応募 × 追加率 32%", "add", True),
             ("累計の友だち数", "前月 ＋ 友だち追加", "friends", False),
             ("配信できるスタッフ数", "累計 ×（1 − ブロック率 30%）", "active", False),
             ("再応募件数", "前月の配信できる数 × 10%", "reapply", True), ("再稼働人数", "再応募 × 勤務決定率 70%", "work", True)]
    hdr = ["項目", "計算（仮定）"] + [f"{m}か月目" for m in range(1, 7)] + ["6か月計"]
    rows = [[cell(t, bold=True, color="FFFFFF", fill=NAVY2, size=10.5) for t in hdr]]
    for name, calc, key, total in names:
        vals = SIM[key]
        hl = key in ("reapply", "work")
        rows.append([cell(name, bold=True, color=NAVY, fill=PALE, size=10.5, align="l"),
                     cell(calc, size=10, color=GREY, align="l")]
                    + [cell(f"{v:,}", size=11, bold=hl and m == 5, color=NAVY if hl else INK,
                            fill="E4E9F2" if hl else "FFFFFF") for m, v in enumerate(vals)]
                    + [cell(f"{sum(vals):,}" if total else "—", size=11, bold=True, color=NAVY if hl else INK,
                            fill="E4E9F2" if hl else PALE2)])
    add_table(s, X0, 1.5, [1.75, 2.62] + [0.74] * 6 + [1.08], [0.36] + [0.42] * 6, rows)
    cw = (W - 0.3) / 3
    boxes = [("いただいている数値", ["サイトUU 8,947", "※集計期間は未確認のため、月間と仮置き", "※応募数・広告の状況は未確認"]),
             ("仮定（DYMが設定）", ["応募率 1.0%（UU→応募）", "友だち追加率 32%（遷移80%×追加40%）", "ブロック率 30%／再応募率 月10%",
                                  "勤務決定率 70%"]),
             ("試算に含めていないもの", ["今いる登録スタッフの友だち追加（人数未確認）", "求人媒体経由の応募（同じ完了画面なら増える）",
                                      "勤務連絡の効率化（時間の削減は数値化しない）"])]
    for i, (h, items) in enumerate(boxes):
        x = X0 + i * (cw + 0.15)
        card(s, x, 4.55, cw, 1.52, [head(h, {"size": 11})] + [[(t if t.startswith("※") else "・" + t, {})] for t in items],
             size=9.5, after=1,
             fill=PALE2 if i != 1 else "FFFFFF", line=NAVY2 if i == 1 else BORDER, ins=(0.1, 0.05, 0.06, 0.03))
    add_box(s, X0, 6.15, W, 0.4,
            [[("※試算であり、成果を保証するものではありません。広告経由と自然流入は区別していません。"
               "人数・件数は小数点以下を四捨五入しています。", {})]], size=9.5, color=GREY, ins=(0.02, 0, 0.02, 0))


def p_cost(s):
    base_page(s, "運用プラン・費用｜月額20万円以内（ツール費込み）で運用します",
              "月額はツール費を含めて13.5万円〜（税抜）。想定予算の月額20万円以内に収めています。",
              "初期費用は月額とは別にお見積りします（6か月〜のご契約）。",
              foot="※税抜。施策内容・配信本数・連携範囲により、費用は個別にお見積りします。LINE公式アカウントの利用料は配信数により変わります。")
    lw = 5.15
    rows = [[cell("月額費用", bold=True, color="FFFFFF", fill=NAVY2, align="l"), cell("金額", bold=True, color="FFFFFF", fill=NAVY2)]]
    for k, v, f in [("DYMの運用費（運用代行・効率改善プラン／月6本の配信）", "10万円", "FFFFFF"),
                    ("サンクスLINE誘導（ツール費）", "3万円〜", "FFFFFF"),
                    ("LINE公式アカウント利用料（ライトプラン想定）", "5,000円〜", "FFFFFF")]:
        rows.append([cell(k, size=10.5, align="l", fill=f), cell(v, size=11.5, bold=True, color=NAVY, fill=f)])
    rows.append([cell("合計（予算：月額20万円以内）", bold=True, color=NAVY, fill="E4E9F2", align="l"),
                 cell("13.5万円〜", size=14, bold=True, color=NAVY, fill="E4E9F2")])
    add_table(s, X0, 1.5, [lw - 1.35, 1.35], [0.36, 0.5, 0.42, 0.5, 0.5], rows)
    rows = [[cell("初期費用（別途）", bold=True, color="FFFFFF", fill=GREY, align="l"), cell("金額", bold=True, color="FFFFFF", fill=GREY)]]
    for k, v in [("LINEの初期構築（あいさつ・リッチメニュー・タグ設計・勤務前の案内の型）", "10万円〜"),
                 ("サンクスLINE誘導の導入", "15万円")]:
        rows.append([cell(k, size=10.5, align="l"), cell(v, size=11.5, bold=True, color=NAVY)])
    rows.append([cell("合計", bold=True, color=NAVY, fill=PALE, align="l"), cell("25万円〜", size=12, bold=True, color=NAVY, fill=PALE)])
    add_table(s, X0, 4.0, [lw - 1.35, 1.35], [0.34, 0.55, 0.4, 0.4], rows)
    x2, w2 = X0 + lw + 0.25, W - lw - 0.25
    chip(s, x2, 1.5, w2, 0.36, "役割分担", size=11.5)
    card(s, x2, 1.95, w2, 1.22, [head("DYM", {"size": 11.5})] + [[("・" + t, {})] for t in
                                ["配信の設計・作成（月6本）／タグ・セグメント配信の設定",
                                 "サンクスLINE誘導の導入・設定／月次レポート・改善提案"]], size=10.5, after=2, line=NAVY2)
    card(s, x2, 3.25, w2, 1.32, [head("横浜シミズ様", {"size": 11.5})] + [[("・" + t, {})] for t in
                                ["勤務・募集の情報のご提供（会場・日時・集合など）", "個別のチャット対応（変更・欠席の連絡など）",
                                 "応募フォームへの設置のご協力"]], size=10.5, after=2)
    card(s, x2, 4.7, w2, 1.38,
         [head("分析・改善（毎月）", {"size": 11.5}),
          [("見る数字：友だち数・ブロック率・配信の開封とクリック・再応募件数", {})],
          [("→ 配信の内容・対象・タイミングを見直し、配信本数・ツールも成果を見て調整します", {})]],
         fill=PALE2, size=10.5, after=2)


def p_next(s):
    base_page(s, "今後の展開｜横浜シミズで運用モデルをつくり、検証する",
              "まず横浜シミズで、友だち獲得・勤務連絡・再応募の運用モデルをつくり、成果と運用の負荷を確かめます。",
              "成果をふまえ、シミズグループの他社でも使えるモデルとして整理することを目指します。",
              band="応募から勤務・次の募集までをLINEでつなぎ、継続的な人員確保の仕組みに")
    from _build.juku_fmt.plan import phase_shape
    steps = [("STEP1｜構築（1〜2か月目）", ["サンクスLINE誘導の設置", "あいさつ・リッチメニューの設計", "タグの設計", "勤務前の案内の型づくり"]),
             ("STEP2｜運用・検証（3〜6か月目）", ["次回募集のセグメント配信", "再応募・再稼働の計測", "運用の負荷の確認と改善"]),
             ("STEP3｜展開の検討（7か月目〜）", ["成果と手順を整理", "グループ他社で使えるかを検討", "（展開は成果を見て判断）"])]
    sw = (W - 0.1) / 3
    for i, (h, items) in enumerate(steps):
        x = X0 + i * (sw + 0.05)
        runs_size(phase_shape(s, x, 1.5, sw, 0.5, h, NAVY2 if i < 2 else GREY, "FFFFFF", i == 0), 10)
        card(s, x + 0.05, 2.1, sw - 0.2, 1.45, [[("・" + t, {})] for t in items], size=11, after=3,
             fill=PALE2 if i < 2 else "FFFFFF")
    card(s, X0, 3.75, 6.05, 2.2,
         [head("ご提案のまとめ", {"size": 12}),
          [("① 友だち獲得：応募完了から、そのままLINEへ（サンクスLINE誘導）", {})],
          [("② 勤務連絡・管理：集合・持ち物・前日のリマインドをLINEで", {})],
          [("③ 再応募・再稼働：タグで絞って、希望に合う次回の募集を届ける", {})],
          [("月額20万円以内（ツール費込み）で始め、成果を見て調整します。", {"color": NAVY, "bold": True})]],
         line=NAVY2, size=11, after=4)
    card(s, X0 + 6.2, 3.75, W - 6.2, 2.2,
         [head("追加施策（ご参考）", {"size": 11.5}),
          [("社員採用（中途・新卒）にも、同じLINE公式アカウントを使える可能性があります。", {})],
          [("社員とアルバイトで案内を分ける設計を想定しています（今回のご提案には含めていません）。", {"color": GREY})]],
         fill=PALE2, size=10.5, after=3)


PAGES = {("B", B_COVER): p_cover, ("T", 2): p_current, ("T", 3): p_issue1, ("T", 4): p_issue2,
         ("T", 5): p_comp_screens, ("T", 6): p_comp_table, ("T", 7): p_overview, ("A", 1): p_thanks, ("A", 3): p_record,
         ("T", 10): p_work, ("T", 11): p_reapply, ("T", 12): p_image, ("T", 13): p_sim, ("T", 14): p_cost,
         ("T", 15): p_next}


def edit(prs):
    sl = prs.slides
    fixed = 0
    for key, fn in PAGES.items():
        s = sl[POS[key] - 1]
        fn(s)
        fixed += fix_fonts(s)
    return fixed


def main(base, appx, out):
    mid, tmp = out + ".mid.pptx", out + ".tmp.pptx"
    gone = assemble(base, appx, mid)
    prs = Presentation(mid)
    fixed = edit(prs)
    prs.save(tmp)
    jb.ORDER = ORDER                      # repack は並びの長さだけ見る
    jb.repack(mid, tmp, out)
    Path(mid).unlink()
    Path(tmp).unlink()
    print("saved:", out, "／", len(ORDER), "枚／外した不要パーツ:", len(gone), "／メイリオにそろえた文字:", fixed)
    for k, v in SIM.items():
        print(f"  SIM {k}: {v}  計 {sum(v)}")


if __name__ == "__main__":
    main(*sys.argv[1:4])
