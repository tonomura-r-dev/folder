#!/usr/bin/env python3
"""
チェングロウス アジェンダv8（20枚版）の簡易スケルトン資料を作る。
v7（56枚）を1枚あたり複数要素にまとめて圧縮した版。箇条書き中心。
スライド5（前後検索）には、近い業種（工場求人ナビ）の前後検索チャートを参考画像として貼る。
出力: 20260928_株式会社チェングロウス御中_LINE公式アカウント運用のご提案_構成案v8（20枚）.pptx
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

NAVY = RGBColor(0x1F, 0x28, 0x5A)
TITLE_NAVY = RGBColor(0x00, 0x20, 0x60)
GRAY = RGBColor(0x33, 0x33, 0x33)
LIGHT_BG = RGBColor(0xF4, 0xF7, 0xFF)
BORDER = RGBColor(0xD9, 0xD9, 0xD9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
AMBER_BG = RGBColor(0xFF, 0xF3, 0xCD)
AMBER_LINE = RGBColor(0xE0, 0xB0, 0x00)
AMBER_TEXT = RGBColor(0x99, 0x66, 0x00)
FONT = "メイリオ"
SW, SH = Inches(10.83), Inches(7.5)

REF_IMG = "_data/zengo/工場求人ナビ_前後検索_20260928.png"

# (chapter, title, bullets:list[str], note)
SLIDES = [
    ("0", "表紙", ["自動車求人Navi LINE公式アカウント運用のご提案"], ""),
    ("0", "アジェンダ", ["業界の課題 → 貴社の現状 → 全体設計 → 施策 → 効果・費用 → 弊社の運用の順に進める"], ""),

    ("1", "業界の課題①　人が足りない", [
        "整備士は求人倍率5.46倍。事業者の約6割が人手不足を感じている",
        "整備学校の入学者は約4割減、平均年齢は47.2歳",
        "営業は12.55倍でさらに不足、受付は0.96倍で人が集まる（主役は整備士）",
        "採用側の悩み：「応募が少ない」52.4%、「条件に合う人が来ない」42.9%",
    ], ""),
    ("1", "業界の課題②　求職者は比べながら動いている", [
        "エージェント・総合・特化の3経路を併用（特化サイトの84%は有資格者）",
        "決め手の1位は給与",
        "検索は一年中ある（山は3月、谷は12月。差は約1.4倍）",
    ], ""),
    ("2", "前後検索でわかること", [
        "「カーディーラー」の前後検索：受付嬢・年収などへの関心が検索当日〜翌日に出る",
        "参考：近い業種（工場求人・期間工）の前後検索の型",
        "　検索前＝求人サイトや条件で広く比較／検索当日＝ブランド名・具体的な求人で意思決定に近づく",
        "　検索後＝勤務条件・評判・志望動機を確認",
        "「自動車整備士」の前後検索は未取得。取得後に差し替える",
    ], "参考画像：工場求人ナビの前後検索（近い業種。自動車整備士のデータではない）"),
    ("2", "まとめ：だから「来た人を逃さない受け皿」が要る", [
        "総合大手はLINEを大規模に活用。整備士特化ではまだ少ない",
        "先に始めた方が有利",
    ], ""),

    ("3", "貴社の現状", [
        "強み：オートバックスグループ／転職支援／資格取得支援",
        "今の入口は「応募」か「転職支援の申込」の2つだけ",
        "整備士の応募単価は6万円（目標2〜2.5万円）。登録はほぼ発生していない",
        "変えられないこと：市場の倍率・広告単価／変えられること：来た人を残す・登録を軽くする・応募まで追う",
    ], ""),

    ("4", "最適解は、実績のある型をお手本にすること", [
        "お手本：タウンワーク（リッチメニューで条件検索 → 応募数79%増・LINEヤフー公式）",
        "貴社向けの工夫①：LINEだけで登録が完了する（タウンワークはリクルートIDが必要）",
        "貴社向けの工夫②：条件は整備士向け（職種・資格・エリア）",
        "貴社向けの工夫③：通知メッセージでLINE外の応募者もLINEに合流させる",
    ], ""),
    ("4", "全体設計｜業界課題 → だからこの施策", [
        "業界課題の一つひとつに、だから何をするかを対応させる",
        "8フェーズのうち、取りこぼしが起きている「サイト離脱」と「登録から応募まで」にLINEを置く",
    ], ""),
    ("4", "CVの置き方", [
        "CV①＝「LINEで登録」で完了（LINEアカウントが会員証になる）",
        "CV②＝求人応募。職種・資格・エリアは登録後にLINEで聞く",
        "規模の近い参考：保育box（LINEの応募率がメルマガの5倍。1.25%と0.25%）",
    ], ""),
    ("4", "サイトリニューアルに入れること", [
        "①「LINEで登録・ログイン」",
        "②会員名簿にLINEの番号を保存し、条件で選んで送れるようにする",
        "③条件別の検索結果ページ",
        "開発はDYMで対応可能（都度見積もり）",
        "全体像：広告で呼び込み → サイトで受け止め → LINEで残して登録・応募・面談まで運ぶ",
    ], ""),

    ("5", "友だち追加の動線", [
        "動線00 離脱防止ポップアップ：帰ろうとした人に新着求人をLINEで案内し、応募しなかった人も残す",
        "動線01「LINEで登録」ボタン：登録と友だち追加が同時にできる",
        "動線02 完了画面：応募後の連絡をLINEに寄せ、不採用・辞退時の再案内の余地も残す",
    ], ""),
    ("5", "Profile+と通知メッセージ", [
        "Profile+で入力を減らす：氏名・電話番号等が自動で入り、確認して押すだけ（利用には申請と契約が必要）",
        "LINE経由の人：Profile+で登録してLINEで連絡",
        "LINEを通らない応募者（Indeed・求人ボックス等）：通知メッセージで「応募受付完了」を届け、その場で友だち追加してもらう",
    ], ""),
    ("5", "あいさつ・リッチメニュー・分岐", [
        "あいさつは全員共通：1通目「LINEで登録」、2通目「すぐ転職したい？」",
        "リッチメニュー：職種・地域・条件をワンタップで検索。未経験・資格取得支援の入口も置く",
        "「まだ」の人はタグで絞り、興味のある内容だけ届けて配信数を抑える",
    ], ""),
    ("5", "配信で応募まで運ぶ", [
        "登録から応募までのステップ配信（前後検索の型を参考に順番を設計）",
        "会員名簿の条件で選んで、合う新着求人だけを届ける",
        "年間の企画配信は3月の山に向けて前倒し。応募しなかった人はクリック履歴で絞って再度声をかける",
    ], ""),
    ("5", "通知メッセージ・自動応答", [
        "「応募受付完了」「会員登録完了」の公式ひな形を使用（販促は不可）",
        "面談日程への利用は申込時に確認",
        "資格・受験料・勤務地などのよくある質問は自動応答で返す",
    ], ""),

    ("6", "効果測定・費用対効果・スケジュール", [
        "CV①（登録）とCV②（応募）を分けて測定",
        "成果シミュレーション：広告SIMの数字に合わせて見込みを示す（差し込み待ち）",
        "費用対効果：投資額を何ヶ月で回収できるか（数字はSIM確定後）",
        "スケジュール：11〜12月に予算決定 → 申請 → 4月にサイトと同時開設 → 翌3月の山までに友だちを貯める",
    ], "差し込み待ち：DYM広告チームの広告SIM（Meta30万・リス70万）確認中"),

    ("7", "LINEヤフー公式の実績と次のご確認事項", [
        "求人・人材の事例：保育box（応募率5倍）／トライト（面接設定率127%）／ランスタッド（返答率20%→80%）",
        "LINEで会員登録が増えた事例：登録と同時の連携率 約9割／新規会員登録 約3.5倍",
        "次のご確認：取得する項目・通知メッセージの運用・面談体制が決まれば構築へ",
    ], ""),

    ("8", "【汎用】弊社のLINE公式アカウント運用について", [
        "なぜLINEか：企業の活用が増加／若年層は電話・メールを避ける傾向",
        "施策展開図：新規の友だち追加と、既存友だちの育成・応募への着地の2動線",
        "改善モデル：①友だち・リードを増やす ②面談に引き上げる",
        "運用プラン：サービス内容（予算型コンサル）／都度発注（開発・Profile+・通知メッセージ等）",
        "弊社体制図",
    ], ""),
    ("0", "裏表紙", ["株式会社DYM"], ""),
]


def set_font(run, size, bold=False, color=GRAY):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    for tag in ("latin", "ea", "cs"):
        el = rPr.makeelement(f"{{http://schemas.openxmlformats.org/drawingml/2006/main}}{tag}", {"typeface": FONT})
        rPr.append(el)


def add_textbox(slide, l, t, w, h, text, size, bold=False, color=GRAY, align=PP_ALIGN.LEFT, anchor=None):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    if anchor:
        tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_font(run, size, bold, color)
    return box


def add_bullets(slide, l, t, w, h, bullets, size=15, color=GRAY):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, b in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(8)
        indent = b.startswith("　")
        run = p.add_run()
        run.text = ("　" if not indent else "") + ("・" + b if not indent else b)
        set_font(run, size - 2 if indent else size, False, color)
    return box


def add_rect(slide, l, t, w, h, fill, line=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line:
        shp.line.color.rgb = line
        shp.line.width = Pt(0.75)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def build():
    prs = Presentation()
    prs.slide_width = SW
    prs.slide_height = SH
    blank = prs.slide_layouts[6]
    total = len(SLIDES)

    for i, (chap, title, bullets, note) in enumerate(SLIDES, 1):
        slide = prs.slides.add_slide(blank)
        add_rect(slide, 0, 0, SW, SH, WHITE)

        if title in ("表紙", "裏表紙"):
            add_rect(slide, 0, 0, SW, SH, NAVY)
            add_textbox(slide, Inches(0.8), Inches(3.0), Inches(9.2), Inches(1.0),
                        bullets[0], 30, True, WHITE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
            if title == "表紙":
                add_textbox(slide, Inches(0.8), Inches(4.0), Inches(9.2), Inches(0.5),
                            "構成案（簡易スケルトン・20枚版）", 14, False, RGBColor(0xCC, 0xCC, 0xCC), PP_ALIGN.CENTER)
            continue

        add_rect(slide, Inches(0.5), Inches(0.3), Inches(1.4), Inches(0.4), NAVY)
        add_textbox(slide, Inches(0.5), Inches(0.3), Inches(1.4), Inches(0.4),
                    f"{chap}章" if chap != "0" else "導入", 12, True, WHITE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        add_textbox(slide, Inches(9.6), Inches(0.3), Inches(0.9), Inches(0.4),
                    f"{i}/{total}", 11, False, GRAY, PP_ALIGN.RIGHT)

        add_textbox(slide, Inches(0.5), Inches(0.82), Inches(9.8), Inches(0.65),
                    title, 19, True, TITLE_NAVY)
        add_rect(slide, Inches(0.5), Inches(1.5), Inches(9.8), Pt(1.5), NAVY)

        has_ref_img = (title == "前後検索でわかること")
        body_h = Inches(3.0) if has_ref_img else Inches(4.9)
        add_rect(slide, Inches(0.5), Inches(1.65), Inches(9.8), body_h, LIGHT_BG, BORDER)
        add_bullets(slide, Inches(0.8), Inches(1.85), Inches(9.2), body_h - Inches(0.3), bullets, size=14)

        if has_ref_img:
            img_top = Inches(1.65) + body_h + Inches(0.15)
            try:
                slide.shapes.add_picture(REF_IMG, Inches(0.5), img_top, width=Inches(9.8))
            except Exception as e:
                add_textbox(slide, Inches(0.5), img_top, Inches(9.8), Inches(0.4),
                            f"[参考画像 読み込み失敗: {e}]", 11, False, RGBColor(0xCC, 0, 0))

        if note:
            note_top = Inches(6.75)
            add_rect(slide, Inches(0.5), note_top, Inches(9.8), Inches(0.55), AMBER_BG, AMBER_LINE)
            add_textbox(slide, Inches(0.7), note_top, Inches(9.4), Inches(0.55),
                        f"⚠ {note}", 11, True, AMBER_TEXT, PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)

    out = "20260928_株式会社チェングロウス御中_LINE公式アカウント運用のご提案_構成案v8（20枚）.pptx"
    prs.save(out)
    print(f"saved: {out} ({total} slides)")


if __name__ == "__main__":
    build()
