# -*- coding: utf-8 -*-
"""横浜シミズ資料（32枚）：新規応募の起点を「サイトに来た方をLINEにつなぎ、LINEから応募」に直す（2026-10-29 殿村さん）。

    python -m _build.patch_yokohama_shimizu_site_line <資料.pptx> [出力.pptx]

  SIMは「サイトのLINEボタン・離脱防止で友だち → 配信から応募フォームで応募＝新規応募／通知メッセージで既存スタッフ → 再応募」。
  資料は「採用サイトで応募 → 完了画面からLINE（サンクスLINE）」が起点になっていて、SIMと合っていなかった。
  ・14 業界課題②：「応募のあと」のLINEでできること＝LINEから応募いただき、面接・登録の案内もLINEで
  ・19 提案の全体像：採用サイト → 施策① LINE友だち追加（LINEボタン・離脱防止）→ 新規応募（LINEから応募フォームへ）→ 勤務連絡…
  ・20（新）施策①｜友だち獲得：サイトに来た方を、LINEにつなぐ（求人を見ている／応募せずに離れる／サイトで応募した の3つの入口）
  ・21（旧20）サンクスLINE誘導のページは、タイトルを「サイトで応募した方も、そのままLINEへ」に
  マイナビの話は資料に入れない（横浜シミズのサイトだけの話）。出力を省くと上書き。
"""
import copy
import sys

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches

from _build.patch_tanaka_v5 import copy_slide
from _build.patch_yokohama_shimizu_chapters import head
from _build.pptx_parts import (CARD, GRAY, GREEN, GREEN_BG, INK, LGRAY, NAVY, WHITE, arrow_line, label, set_para,
                               shape)

IN = 2.54


def by_id(s, sid):
    return next(sh for sh in s.shapes if sh.shape_id == sid)


def title_of(s):
    return head(s).text_frame.text


def clone(s, sid):
    el = copy.deepcopy(by_id(s, sid)._element)
    nid = max(int(e.get("id")) for e in s.shapes._spTree.iter(qn("p:cNvPr"))) + 1
    el.find(".//" + qn("p:cNvPr")).set("id", str(nid))
    by_id(s, sid)._element.addnext(el)
    return by_id(s, nid)


def place(sh, x, y, w, h):
    sh.left, sh.top, sh.width, sh.height = Inches(x), Inches(y), Inches(w), Inches(h)


def p14(s):
    box = by_id(s, 25)
    assert box.text_frame.text.startswith("LINEでできること")
    set_para(box, 1, "LINEから応募いただき、面接・登録の案内もLINEで届ける")


def p19(s):
    set_para(by_id(s, 3), 0, "サイトに来た方に友だちになっていただき、応募から次の募集までLINEでつなぎます。")
    a, b = by_id(s, 22), by_id(s, 24)
    set_para(a, 0, "採用サイト")
    set_para(a, 1, "求人を見る")
    set_para(b, 2, "LINEボタン・離脱防止")
    place(a, 0.47, 1.6, 2.05, 0.85)
    place(by_id(s, 23), 1.5, 2.47, 0, 0.26)
    place(b, 0.47, 2.75, 2.05, 1.05)
    c = clone(s, 22)                                          # 新規応募（LINEから応募フォームへ）
    set_para(c, 0, "新規応募")
    set_para(c, 1, "LINEから応募フォームへ")
    place(c, 0.47, 4.1, 2.05, 0.9)
    place(clone(s, 23), 1.5, 3.82, 0, 0.26)
    d = by_id(s, 25)                                          # 斜めの矢印：新規応募 → 勤務連絡
    d.left, d.top, d.width, d.height = Inches(2.54), Inches(2.35), Inches(0.69), Inches(2.2)
    set_para(by_id(s, 35), 0, "応募の前から友だちになるので、イベントが終わっても次の募集が届きます。")


def iriguchi(prs, src, before):
    s = copy_slide(prs, src)
    for sh in list(s.shapes):
        if sh.name not in ("TextBox 1", "TextBox 2", "Rectangle 17", "Rectangle 20"):
            sh._element.getparent().remove(sh._element)
    set_para(by_id(s, 2), 0, "施策①｜友だち獲得：サイトに来た方を、LINEにつなぐ")
    set_para(by_id(s, 3), 0, "求人を見る・離れる・応募する、どの場面でもLINEにつながる入口を用意します。")
    set_para(by_id(s, 3), 1, "友だちになった方には、LINEから応募フォームへご案内します。")

    y0, h, gap = 1.95, 0.85, 0.2
    rows = [("求人を見ている", "サイトのLINEボタン", "新着の募集がLINEで届く"),
            ("応募せずに離れる", "離脱防止（ポップアップ）", "希望の募集をLINEで受け取る"),
            ("サイトで応募した", "サンクスLINE誘導", "面接・登録の案内がLINEで届く")]
    total = h * 3 + gap * 2
    rr = dict(adj=0.08, margins=(0.15, 0.05, 0.15, 0.05))
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.47 * IN, y0 * IN, 1.5 * IN, total * IN,
          [("採用サイトに", 13, True, NAVY, 0), ("来た方", 13, True, NAVY, 0)], fill=CARD, line=NAVY, lw=1.0, **rr)
    for k, (scene, name, why) in enumerate(rows):
        y = y0 + k * (h + gap)
        cy = (y + h / 2) * IN
        arrow_line(s, 2.0 * IN, cy, 2.28 * IN, cy, NAVY, 1.5)
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 2.3 * IN, y * IN, 2.0 * IN, h * IN,
              [(scene, 12.5, True, INK, 0)], fill=WHITE, line=LGRAY, lw=1.0, **rr)
        arrow_line(s, 4.33 * IN, cy, 4.6 * IN, cy, NAVY, 1.5)
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 4.62 * IN, y * IN, 2.6 * IN, h * IN,
              [(name, 13, True, WHITE, 2), (why, 11, False, WHITE, 0)], fill=NAVY, **rr)
        arrow_line(s, 7.25 * IN, cy, 7.6 * IN, cy, NAVY, 1.5)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 7.62 * IN, y0 * IN, 1.25 * IN, total * IN,
          [("LINE", 13, True, NAVY, 0), ("友だち追加", 13, True, NAVY, 0)], fill=GREEN_BG, line=GREEN, lw=1.5, **rr)
    mid = (y0 + total / 2) * IN
    arrow_line(s, 8.9 * IN, mid, 9.1 * IN, mid, NAVY, 1.5)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 9.12 * IN, y0 * IN, 1.24 * IN, total * IN,
          [("LINEで", 12, True, NAVY, 0), ("応募・", 12, True, NAVY, 0), ("勤務の連絡・", 12, True, NAVY, 0),
           ("次の募集", 12, True, NAVY, 0)], fill=WHITE, line=NAVY, lw=1.0, adj=0.08, margins=(0.05, 0.05, 0.05, 0.05))
    bottom = y0 + total
    label(s, 0.47 * IN, (bottom + 0.08) * IN, 9.89 * IN, 0.3 * IN,
          [("※サンクスLINE誘導は、サイトの応募フォームから応募した方を、完了画面からLINEへご案内する仕組みです（次ページ）。",
            9, False, GRAY, 0)])
    for name in ("Rectangle 17", "Rectangle 20"):
        sh = next(x for x in s.shapes if x.name == name)
        sh.top = Inches(bottom + 0.5)
    bar = next(x for x in s.shapes if x.name == "Rectangle 20")
    set_para(bar, 0, "サイトに来た方を、取りこぼさず友だちにする")

    lst = prs.slides._sldIdLst                                # 最後に足されたページを、サンクスLINE誘導のページの前へ
    el = list(lst)[-1]
    lst.remove(el)
    next(e for e in lst if int(e.get("id")) == before.slide_id).addprevious(el)
    return s


def main(src, out):
    prs = Presentation(src)
    sl = list(prs.slides)
    assert len(sl) == 32
    assert title_of(sl[13]).startswith("業界課題②")
    assert title_of(sl[18]).startswith("提案の全体像")
    assert title_of(sl[19]).startswith("施策①｜友だち獲得：応募完了から")
    p14(sl[13])
    p19(sl[18])
    iriguchi(prs, sl[18], sl[19])
    set_para(head(sl[19]), 0, "施策①｜友だち獲得：サイトで応募した方も、そのままLINEへ")
    prs.save(out)
    print("saved:", out, len(prs.slides))


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else sys.argv[1])
