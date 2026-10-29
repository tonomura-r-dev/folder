# -*- coding: utf-8 -*-
"""横浜シミズ資料（殿村さんPC保存版 2026-10-29・27枚）：章の区切りページを足し、左上のタイトルの位置をそろえる（殿村さん）。

    python -m _build.patch_yokohama_shimizu_chapters <PC保存版.pptx> [出力.pptx]

  ・章の区切り：6ページ「01 施策背景」と同じ形のページを写して、02〜06を足す
      02 現状と課題（現状把握の前）／03 競合分析（競合分析①の前）／04 ご提案内容（提案の全体像の前）
      05 料金プラン（概算お見積りの前）／06 今後の展開（今後の展開の前）
  ・左上のタイトル：汎用資料のページ（弊社紹介など）と提案のページで、枠の位置・余白が違い、文字の出だしがずれていた。
    2ページ「弊社紹介」のタイトルと同じ位置・余白・色にそろえる。文字・大きさは変えない。表紙・章の区切り・裏表紙は触らない。
  出力を省くと上書き。
"""
import copy
import sys

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Emu

from _build.patch_tanaka_v5 import copy_slide

CHAPTERS = [("02", "現状と課題", "現状把握｜"), ("03", "競合分析", "競合分析①｜"), ("04", "ご提案内容", "提案の全体像｜"),
            ("05", "料金プラン", "概算お見積り"), ("06", "今後の展開", "今後の展開｜")]


def head(s):
    c = [sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()
         and Emu(sh.top).inches < 0.45 and Emu(sh.left).inches < 1.0]
    return c[0] if c else None


def is_chapter(s):
    t = [sh.text_frame.text for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
    return len(t) >= 2 and t[0].isdigit() and len(t[0]) == 2


def main(src, out):
    prs = Presentation(src)
    slides = list(prs.slides)
    base = head(slides[1])
    assert base.text_frame.text.strip() == "弊社紹介"
    bp0 = base.text_frame._txBody.find(qn("a:bodyPr"))
    color0 = next(base.text_frame._txBody.iter(qn("a:srgbClr")), None)
    fixed = 0
    for s in slides[1:-1]:
        if is_chapter(s):
            continue
        h = head(s)
        if h is None:
            continue
        h.left, h.top, h.height = base.left, base.top, base.height
        bp = h.text_frame._txBody.find(qn("a:bodyPr"))
        for k in ("lIns", "rIns", "tIns", "bIns"):
            if bp0.get(k) is None:
                bp.attrib.pop(k, None)
            else:
                bp.set(k, bp0.get(k))
        bp.set("anchor", "ctr")
        if color0 is not None:
            for c in h.text_frame._txBody.iter(qn("a:srgbClr")):
                c.set("val", color0.get("val"))
        fixed += 1

    chap = next(s for s in slides if is_chapter(s))
    lst = prs.slides._sldIdLst
    for num, name, title in CHAPTERS:
        target = next(s for s in prs.slides if head(s) is not None and head(s).text_frame.text.startswith(title))
        new = copy_slide(prs, chap)
        texts = [sh for sh in new.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        texts[0].text_frame.paragraphs[0].runs[0].text = num
        texts[1].text_frame.paragraphs[0].runs[0].text = name
        el = list(lst)[-1]
        lst.remove(el)
        next(e for e in lst if int(e.get("id")) == target.slide_id).addprevious(el)
    prs.save(out)
    print("saved:", out, "titles", fixed, "slides", len(prs.slides))


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else sys.argv[1])
