# -*- coding: utf-8 -*-
"""タナカホーム v3 → v4：サンクスLINE誘導をプランから外してオプションにする（2026-10-27 殿村さん・佐村さん）。

    python -m _build.patch_tanaka_v4_thanks_option <v3（殿村さんのPC保存版）.pptx> <提案FMT_MAX版_20261024d.pptx> <v4.pptx>

  佐村さん「サンクスLINE誘導で資料請求→来場につなげましょうはOptionで入れといて、最悪そこだけの導入だけでもOK」。
  数字は SIM ver1.2 の B_サンクス無（殿村さんのPC保存版：初期の運用設計・構築 15万円）。
  ・P13 概算お見積り：初期 ¥165,000〜／月次 ¥230,000〜／6か月後の友だち 1,453人（資料請求 11件・来場予約 8件はそのまま）。
    初期の実施内容の「サンクスLINE誘導」は「オプション」に
  ・P14 料金プラン：初期 ①115,000円〜／②③165,000円〜、月額ツール 30,000円〜（離脱防止）、月額 ①30,000円〜／②230,000円〜／③430,000円〜。
    注記にオプション（追加する場合 初期＋150,000円・月額＋30,000円〜、サンクスLINE誘導のみの導入も可能）を1行足す
  ・P15 運用内容：「離脱防止・サンクスLINE誘導」→「離脱防止」＋注記「サンクスLINE誘導はオプション」
  ・P17 スケジュール：構築の行から「サンクスLINE誘導」を外し、オプションと書く
  ・運用内容（P15）の後ろ＝P16に、サンクスLINE誘導の機能・費用のページ（提案FMTのS35＝動画のQR付き）を1枚足す（本編に置く）
  ほかのページは変えない。
"""
import copy
import sys

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches

from _build.patch_tanaka_v5 import copy_slide


def sid(s, n):
    return next(sh for sh in s.shapes if sh.shape_id == n)


def set_lines(txbody_owner, lines):
    """最初の段落・最初の文字の書式のまま、行ごとに段落を作り直す（図形・表のセルどちらも）"""
    txb = txbody_owner._txBody if hasattr(txbody_owner, "_txBody") else txbody_owner.text_frame._txBody
    ps = txb.findall(qn("a:p"))
    base = ps[0]
    for p in ps:
        txb.remove(p)
    for line in lines:
        p = copy.deepcopy(base)
        rs = p.findall(qn("a:r"))
        for r in rs[1:]:
            p.remove(r)
        for br in p.findall(qn("a:br")):
            p.remove(br)
        rs[0].find(qn("a:t")).text = line
        txb.append(p)


def cell(tb, r, c, text):
    set_lines(tb.cell(r, c).text_frame, text.split("\n"))


def p13(s):
    set_lines(sid(s, 339).text_frame, ["概算お見積り：月23万円で、6か月目に資料請求 月11件・来場予約 月8件"])
    set_lines(sid(s, 343).text_frame, ["¥165,000〜"])
    set_lines(sid(s, 344).text_frame, ["¥230,000〜"])
    set_lines(sid(s, 375).text_frame, ["1,453人"])
    set_lines(sid(s, 368).text_frame, ["オプション"])


def p14(s):
    tb = next(sh for sh in s.shapes if sh.has_table).table
    assert tb.cell(1, 0).text == "初期費用" and tb.cell(4, 0).text == "月額費用"
    for c, v in ((1, "115,000円〜"), (2, "165,000円〜"), (3, "165,000円〜")):
        cell(tb, 1, c, v)
    cell(tb, 3, 1, "30,000円〜")
    cell(tb, 3, 2, "30,000円〜")
    cell(tb, 3, 3, "30,000円〜\n＋LINE広告\n200,000円〜")
    for c, v in ((1, "30,000円〜"), (2, "230,000円〜"), (3, "430,000円〜")):
        cell(tb, 4, c, v)
    notes = [sid(s, n) for n in (640, 641, 642, 643)]
    set_lines(notes[0].text_frame, ["※初期費用＝運用設計・構築（プラン① 100,000円〜／プラン②・③ 150,000円〜）＋離脱防止 15,000円"])
    set_lines(notes[1].text_frame, ["※月額ツール・LINE関連費＝離脱防止 30,000円。LINE公式アカウント利用料は別（貴社のご契約のまま）"])
    opt = copy.deepcopy(notes[1]._element)                   # オプションの1行（太字・紺）を、2行目の後ろに足す
    notes[1]._element.addnext(opt)
    o = next(sh for sh in s.shapes if sh._element is opt)
    set_lines(o.text_frame, ["※サンクスLINE誘導（オプション）：追加する場合は初期＋150,000円・月額＋30,000円〜。サンクスLINE誘導のみの導入も可能"])
    for r in o.text_frame.paragraphs[0].runs:
        r.font.bold = True
    for k, sh in enumerate([notes[0], notes[1], o, notes[2], notes[3]]):
        sh.top = Inches(5.66 + 0.25 * k)


def p15(s, note_src):
    tb = next(sh for sh in s.shapes if sh.has_table).table
    assert tb.cell(1, 0).text.startswith("離脱防止")
    cell(tb, 1, 0, "離脱防止")
    e = copy.deepcopy(note_src._element)
    s.shapes._spTree.append(e)
    n = next(sh for sh in s.shapes if sh._element is e)
    set_lines(n.text_frame, ["※サンクスLINE誘導はオプション（どのプランにも追加可能。サンクスLINE誘導のみの導入も可能）"])
    tbl = next(sh for sh in s.shapes if sh.has_table)
    n.top = tbl.top + tbl.height + Inches(0.12)


def p17(s):
    hit = 0
    for t in s._element.iter(qn("a:t")):
        if t.text and "構築：離脱防止・サンクスLINE誘導、" in t.text:
            t.text = t.text.replace("構築：離脱防止・サンクスLINE誘導、", "構築：離脱防止、").rstrip() + "（サンクスLINE誘導はオプション）"
            hit += 1
    assert hit == 1, hit


def main(src, fmt_path, out):
    prs = Presentation(src)
    slides = list(prs.slides)
    assert len(slides) == 31
    assert sid(slides[12], 339).text_frame.text.startswith("概算お見積り")
    p13(slides[12])
    p14(slides[13])
    p15(slides[14], sid(slides[13], 642))
    p17(slides[16])
    fmt = Presentation(fmt_path)
    s35 = fmt.slides[34]
    assert any(sh.has_text_frame and sh.text_frame.text == "新規友だちの追加（サンクスLINE誘導）" for sh in s35.shapes)
    copy_slide(prs, s35)                                    # 本編の運用内容（P15）の後ろに入れる（殿村さん「Appendixではなく15と16の間に」）
    lst = prs.slides._sldIdLst
    new = list(lst)[-1]
    lst.remove(new)
    list(lst)[14].addnext(new)
    left = [t.text for sl in prs.slides for t in sl._element.iter(qn("a:t")) if t.text and "サンクスLINE誘導 150,000円" in t.text]
    assert not left, left
    assert len(prs.slides) == 32
    prs.save(out)
    print("saved:", out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2], sys.argv[3])
