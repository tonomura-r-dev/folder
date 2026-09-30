# -*- coding: utf-8 -*-
"""チェングロウス本資料 ver2.4 → ver2.5：11〜13枚目を施策ビジュアル（①〜④）の4枚に入れ替える。

- ベース：殿村さんがPCで保存した ver2.4（本資料）と ビジュアル ver1.1
- ビジュアル側は先方向けに2点だけ直してから入れる
  ・「SIM ver2.6」→「弊社シミュレーション」
  ・ステップ配信は最短1日後から（公式マニュアル）なので、0日目はあいさつメッセージと書き分ける
- 入れ替え後：22枚＋裏表紙 → 23枚＋裏表紙（11〜14枚目が①〜④）

  python3 _build/patch_chengrowth_v25_visual.py
"""
import copy
import io
from pathlib import Path

from pptx import Presentation
from pptx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
MAIN = ROOT / "20260929_株式会社チェングロウス御中_LINE公式アカウント運用のご提案ver2.4.pptx"
VIS = ROOT / "20260929_チェングロウス_LINE施策のビジュアル（①〜④）ver1.1.pptx"
OUT = ROOT / "20260929_株式会社チェングロウス御中_LINE公式アカウント運用のご提案ver2.5.pptx"

RT_IMAGE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"

# ビジュアル側の文言直し（run単位の置き換え）
SUB = [
    ("SIM ver2.6", "弊社シミュレーション"),
    ("0〜14日後\n（全10通）", "0日 あいさつ／\n1〜14日後（全10通）"),
    ("ステップの中身（全10通の例）", "配信の中身（全10通の例）"),
    ("あいさつ・職種を聞く", "あいさつメッセージ（職種を聞く）"),
]
PARA = {}
STEP_HEAD = "ステップの中身（全10通の例）"   # この段落だけ、最後に「1日後からはステップ配信」を足す


def set_para(p, text):
    runs = p.findall(qn("a:r"))
    runs[0].find(qn("a:t")).text = text
    for r in runs[1:]:
        p.remove(r)


def fix_text(slide):
    hit = []
    for p in slide.shapes._spTree.iter(qn("a:p")):
        full = "".join(t.text or "" for t in p.iter(qn("a:t")))
        if full.startswith(STEP_HEAD):
            last = list(p.iter(qn("a:t")))[-1]
            last.text += "（1日後からはステップ配信）"
        if full in PARA:
            set_para(p, PARA[full]); hit.append(full[:12])
            continue
        for t in p.iter(qn("a:t")):
            for a, b in SUB:
                if t.text and a in t.text:
                    t.text = t.text.replace(a, b); hit.append(a)
    return hit


main = Presentation(str(MAIN))
vis = Presentation(str(VIS))
layout = {l.name: l for l in main.slide_layouts}

new_ids = []
for vs in vis.slides:
    print("直し:", fix_text(vs))
    ns = main.slides.add_slide(layout[vs.slide_layout.name])
    tree = ns.shapes._spTree
    for el in list(tree):
        if el.tag not in (qn("p:nvGrpSpPr"), qn("p:grpSpPr")):
            tree.remove(el)
    for el in vs.shapes._spTree:
        if el.tag in (qn("p:nvGrpSpPr"), qn("p:grpSpPr")):
            continue
        el = copy.deepcopy(el)
        for node in el.iter():
            for attr in (qn("r:embed"), qn("r:link"), qn("r:id")):
                rid = node.get(attr)
                if not rid:
                    continue
                rel = vs.part.rels[rid]
                assert rel.reltype == RT_IMAGE, rel.reltype
                _, new_rid = ns.part.get_or_add_image_part(io.BytesIO(rel.target_part.blob))
                node.set(attr, new_rid)
        tree.append(el)
    new_ids.append(main.slides._sldIdLst[-1])

# 11〜13枚目を外して、その位置に①〜④を入れる
lst = main.slides._sldIdLst
old = list(lst)
for el in old[10:13]:
    main.part.drop_rel(el.rId)
    lst.remove(el)
for el in new_ids:
    lst.remove(el)
for i, el in enumerate(new_ids):
    lst.insert(10 + i, el)

main.save(str(OUT))
p = Presentation(str(OUT))
print("saved:", OUT.name, len(p.slides), "枚")
for i, s in enumerate(p.slides, 1):
    t = next((sh.text_frame.text for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()), "")
    print(i, t.split("\n")[0][:40])
