# -*- coding: utf-8 -*-
"""塾の新FMT版：「06 次回ご提案に向けて」4枚（P53〜P56）の上の一文（要旨）を書き直し、文字を大きくする（2026-10-22 殿村さん）。

    python -m _build.patch_juku_maxfmt_06_leads <殿村さんPC保存版.pptx> <出力.pptx>

  ③と④が同じ文（現状の規模・活用度・費用から…）だったのを、ページごとの中身に合わせて言い切りの1文に。
  下の帯（①④）と同じことを言わないようにした。13pt → 16pt。ほかのページ・図形は変えない。
"""
import sys

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

LEADS = {  # ページ: (今の文, 新しい文)
    53: ("新規／既存／ファン化／効率化のうち、何を、どの程度重視するか",
         "新規・既存・ファン化・効率化のうち、貴塾がいま最も伸ばしたいものを決める"),
    54: ("流入量・CV・CV地点から、LINEで拾えるユーザーの規模を見積もる",
         "サイトの流入・CVの数から、LINEの友だちにできる保護者の規模を見積もる"),
    55: ("現状の規模・活用度・費用から、改善余地を見極める",
         "すでにいる友だちと配信の実績を、提案の出発点にする"),
    56: ("現状の規模・活用度・費用から、改善余地を見極める",
         "年間の募集スケジュールや配信の確認体制など、運用の進め方に関わる点を伺う"),
}
SIZE = Pt(16)
TOP = Inches(1.05)          # 2026-10-22 殿村さん「もう少し上に寄せて」（元は1.40インチ）


def main(src, out):
    prs = Presentation(src)
    slides = list(prs.slides)
    assert len(slides) == 71
    for no, (old, new) in LEADS.items():
        hits = [sh for sh in slides[no - 1].shapes if sh.has_text_frame and sh.text_frame.text == old]
        assert len(hits) == 1, (no, old)
        p = hits[0].text_frame.paragraphs[0]
        runs = p.runs
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
        runs[0].text = new
        runs[0].font.size = SIZE
        end = p._p.find(qn("a:endParaRPr"))
        if end is not None:
            end.set("sz", "1600")
        hits[0].top = TOP
    prs.save(out)
    print("saved:", out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
