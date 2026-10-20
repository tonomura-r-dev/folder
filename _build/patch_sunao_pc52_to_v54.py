# -*- coding: utf-8 -*-
"""すなお：殿村さんのPC保存版 ver5.2（想定動線を殿村さんが直した版）→ ver5.4（2026-10-20 殿村さん「見て、直して」）。

    python -m _build.patch_sunao_pc52_to_v54 <PC保存のver5.2.pptx> <ver5.4.pptx>

  PC保存版とリポジトリのver5.2の違いは、想定動線（ユーザー動線の図）の1枚だけ（ターゲット／商材・サービス／イベントは年4回／広告）。
  1. ver5.3の直し（業界STUDY・訴求軸・要件定義・本質課題の作り直し、シーズナリティを外す）をそのまま当てる（patch_sunao_v53）
  2. 想定動線の締め：「という強みを持っており、」→「という特性から、」
     （見出しが「サービス特性として」で、中に強みではない「イベントの間は接点が切れやすい」もあるため）
  想定動線のほかの部分は殿村さんの版のまま。
"""
import sys

from pptx import Presentation

from _build import patch_sunao_v53


def main(src, out):
    patch_sunao_v53.main(src, out)
    prs = Presentation(out)
    hit = [r for s in prs.slides for sh in s.shapes if sh.has_text_frame
           for p in sh.text_frame.paragraphs for r in p.runs if r.text == "という強みを持っており、"]
    assert len(hit) == 1, len(hit)
    hit[0].text = "という特性から、"
    prs.save(out)
    print("saved:", out, len(prs.slides), "枚")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
