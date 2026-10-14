# -*- coding: utf-8 -*-
"""あらうち提案資料 v8 の自動チェック。

    python -m _build.arauchi_v8.checks <v7.pptx（元）> <v8.pptx>

数字・表記・並び・パッケージ・グラフの寸法・文字の長さ（1行に収まるはずの箱）を見る。
（見た目のはみ出しは、LibreOfficeの画像と、行間を1.29倍にした確認用コピーの画像で別に見る）
"""
import re
import sys
import zipfile

from lxml import etree
from pptx import Presentation
from pptx.util import Emu

from _build.arauchi_v4.build import c14n, slide_order
from _build.arauchi_v8.build import DROP_INDEX
from _build.arauchi_v8.pages import C1_BAR, C2_BAR, CV1, CV2, TOTAL

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
NSA = {"a": A}
FAILS, OKS = [], []


def check(cond, msg):
    (OKS if cond else FAILS).append(msg)
    if not cond:
        print("NG:", msg)


def slide_text(slide):
    out = []
    for p in slide._element.iter("{%s}p" % A):
        t = "".join(x.text or "" for x in p.iter("{%s}t" % A))
        if t:
            out.append(t)
    return "\n".join(out)


def est_width_in(text, pt):
    """文字の幅の目安（インチ）。全角＝1em、半角の英数字＝0.58em"""
    em = 0.0
    for ch in text:
        if ord(ch) < 0x80:
            em += 0.3 if ch == " " else (0.85 if ch == "%" else 0.58)
        else:
            em += 1.0
    return em * pt / 72


def main(base, out):
    pb, po = Presentation(base), Presentation(out)
    sb, so = list(pb.slides), list(po.slides)
    check(len(sb) == 48 and len(so) == 47, f"スライド数 48→47（{len(sb)}→{len(so)}）")

    # ---- どのスライドが変わったか（元の番号で）
    zb, zo = zipfile.ZipFile(base), zipfile.ZipFile(out)
    ob, oo = slide_order(zb), slide_order(zo)
    changed = []
    for i, (n_out) in enumerate(oo):
        j = i if i < DROP_INDEX else i + 1
        if c14n(zb.read(ob[j])) != c14n(zo.read(n_out)):
            changed.append(j + 1)
    check(changed == [8, 9, 11, 12, 13, 15, 16, 18, 19, 21, 27], f"変えたスライド（元の番号）＝8,9,11,12,13,15,16,18,19,21,27 / 実際 {changed}")
    check(ob[DROP_INDEX] not in zo.namelist(), "プラン④のスライド（元S17）のパーツが残っていない")

    texts = [slide_text(s) for s in so]
    base_titles = [slide_text(s).split("\n")[0] for s in sb]

    # ---- 消すべき言葉
    banned = ["成果報酬", "プラン④", "有効CV", "8,000", "8000", "184,000", "専用フォーム", "別途実費", "実費のみ", "松竹梅", "おすすめ",
              "スコープ", "17ページ", "4つのプラン", "4プラン", "コンサル型", "成果連動", "月額費用完全無料", "完全無料",
              "出典は6ページ", "都度発注プラン", "ご依頼依頼", "運用コンサル費用に含みます", "お申込みの後押し"]
    for w in banned:
        hits = [i + 1 for i, t in enumerate(texts) if w in t]
        check(not hits, f"「{w}」がどのページにも無い {hits if hits else ''}")
    for w in ["23件", "90件", "月7件", "10件/月", "7件→", "13件＋10件", "累計90"]:
        hits = [i + 1 for i, t in enumerate(texts) if w in t]
        check(not hits, f"旧SIMの数字「{w}」が残っていない {hits if hits else ''}")

    # ---- 新しい数字
    s12, s13 = texts[11], texts[12]
    check("1か月目 月10件 → 6か月目 月33件（試算）" in s12, "S12 の配信目標の文")
    check("友だち追加の増加と導線改善で、6か月目にLINE経由の問い合わせ月33件を見込みます（試算）。" in s13, "S13 の見出し（なぜ増えるか＋33件）")
    check(est_width_in("友だち追加の増加と導線改善で、6か月目にLINE経由の問い合わせ月33件を見込みます（試算）。", 14) <= 9.92 - 0.2 - 0.3, "S13 見出し2行目が1行に収まる")
    check("月33件" in s13 and "20件/月" in s13 and "33件/月" in s13 and "13件/月" in s13 and "（6か月累計126件）" in s13, "S13 の問い合わせ（13/20/33件・累計126件）")
    check("約1,377人" in s13 and "¥215,000" in s13 and "¥135,000" in s13, "S13 の友だち数・初期・月次の費用はそのまま")
    check("紹介による追加成果は含みません" in s13 and "月4回程度" in s13, "S13 の試算の前提（紹介を含まない）・企画配信 月4回程度")
    check("21.5万円" in s13 and "13.5万円" in s13, "S13 の見出しの費用")

    # ---- S12 の棒グラフ：値・寸法
    def walk(shapes):
        for sh in shapes:
            yield sh
            if sh.shape_type == 6:
                yield from walk(sh.shapes)
    shapes12 = {sh.name: sh for sh in walk(so[11].shapes)}
    check("グラフ（問い合わせ件数・月間の試算）" in shapes12 and shapes12["グラフ（問い合わせ件数・月間の試算）"].shape_type == 6, "S12 のグラフが1つのグループになっている")
    scale = Emu(914400 * 0.55).emu / 33          # 33件＝0.55in
    for i, (c1, c2, t) in enumerate(zip(CV1, CV2, TOTAL), 1):
        lab = shapes12[f"グラフ_合計_{i}か月目"].text_frame.text
        check(lab in (str(t), f"{t}件"), f"S12 グラフ {i}か月目の合計ラベル {lab} == {t}")
        check(shapes12[f"グラフ_内訳_{i}か月目"].text_frame.text == f"{c1}＋{c2}", f"S12 グラフ {i}か月目の内訳 {c1}＋{c2}")
        for key, v in (("①", c1), ("②", c2)):
            h = shapes12[f"グラフ_問い合わせ{key}_{i}か月目"].height
            check(abs(h - v * scale) <= 2, f"S12 グラフ {i}か月目 問い合わせ{key} の高さ＝{v}件分（誤差 {abs(h - v * scale):.1f}EMU）")
        b1, b2 = shapes12[f"グラフ_問い合わせ①_{i}か月目"], shapes12[f"グラフ_問い合わせ②_{i}か月目"]
        check(abs((b1.top + b1.height) - (shapes12["グラフ_基準線"].top)) <= 2, f"S12 グラフ {i}か月目 ①の底が基準線に載っている")
        check(abs((b2.top + b2.height) - b1.top) <= 2, f"S12 グラフ {i}か月目 ②が①の真上に載っている")
    check("126件（試算）" in shapes12["グラフ_累計"].text_frame.text, "S12 グラフの累計126件")
    check(sum(TOTAL) == 126 and sum(CV1) == 55 and sum(CV2) == 71, "合計 126 ＝ ①55 ＋ ②71")

    # ---- S15・S16 の表
    def table_of(slide):
        for sh in slide.shapes:
            if getattr(sh, "has_table", False) and sh.has_table:
                return [[c.text.replace("\n", "").replace("\x0b", "") for c in r.cells] for r in sh.table.rows]
        raise KeyError

    t15, t16 = table_of(so[14]), table_of(so[15])
    check(all(len(r) == 4 for r in t15) and all(len(r) == 4 for r in t16), "S15・S16 の表は4列（項目＋3プラン）")
    check([t15[0][i] for i in (1, 2, 3)] == ["プラン①総合運用", "プラン②獲得重視", "プラン③初期導入"], "S15 の見出し")
    check([t16[0][i] for i in (1, 2, 3)] == ["プラン①総合運用", "プラン②獲得重視", "プラン③初期導入"], "S16 の見出し")
    row15 = {r[0]: r[1:] for r in t15}
    row16 = {r[0]: r[1:] for r in t16}
    check(row15["初期費用"] == ["215,000円"] * 3, "S15 初期費用 215,000円×3")
    check(row15["月額コンサルティング費"] == ["100,000円", "50,000円", "なし"], "S15 月額コンサルティング費")
    check(row15["月額ツール・LINE関連費"][:2] == ["35,000円", "35,000円"], "S15 月額ツール・LINE関連費 35,000円×2")
    check(row15["月額費用"][:2] == ["135,000円", "85,000円"], "S15 月額費用 135,000／85,000")
    check(100000 + 30000 + 5000 == 135000 and 50000 + 35000 == 85000, "月額＝内訳の合計（100,000+30,000+5,000／50,000+35,000）")
    check(row16["月額費用"][:2] == row15["月額費用"][:2], "S15 と S16 の月額費用が一致")
    check(row16["ステップ配信"] == ["14日間・10通", "なし", "なし"], "S16 ステップ配信")
    check(row16["企画配信（定期配信）"] == ["月4回程度", "月1〜2回程度", "なし"], "S16 配信頻度")
    check(row16["月額費用"][2] == "継続利用するツールのツール利用料＋DYM手数料", "S16 プラン③の月額費用（条件つき）")
    t15txt, t16txt = texts[14], texts[15]
    check("ツール利用料" in t15txt and "DYM手数料" in t15txt and "月額のコンサルティング費用はありません" in t15txt, "S15 プラン③の説明（ツール利用料＋DYM手数料・月額コンサルなし）")
    check("無料" not in t15txt + t16txt and not re.search(r"(?<![\d,])0円", t15txt + t16txt), "S15・S16 にプラン③を無料と読ませる表現が無い")
    check("13ページの試算（6か月目に月33件）" in t16txt, "S16 の注記の13ページ・33件")

    # ---- レビューで直したところ
    t8, t9, t18, t19 = texts[7], texts[8], texts[16], texts[17]
    check("<参考：CV以降の流れ（個別対応はあらうち様）／当社の計測対象外>" in t8 and "当社の定期対応項目(一部抽出)" in t8 and "（プラン①｜総合運用）" in t8,
          "S8 下段の見出し（個別対応はあらうち様）・右列（当社の定期対応項目）・上段（プラン①）")
    check("LINE上の接点" in t9 and "既存友だち（今の友だち 1,241人）" in t9 and "＋追加後の新しい友だちにも配信" not in t9, "S9 列の見出し・既存友だちのセル")
    check("リッチメニュー\n→ 物件の送信（CV）" in t9, "S9 既存友だちの導線の終点がCV")
    check("配信目標　問い合わせ＝LINE上での物件情報の送信（CV）" in s12 and "手法　（プラン①｜総合運用の場合）" in s12, "S12 見出しにCVの定義・手法はプラン①の場合")
    check("（プラン①｜総合運用）" in t18 and "月4回程度の企画配信" in t18, "S17 リード文にプラン①")
    check("役割分担　※配信運用・改善支援は月額コンサル費に含みます" in t19, "S18 役割分担の注記（プラン③に当てはまらない書き方を直した）")
    t26 = texts[25]
    check("都度発注メニュー" in t26 and "ご依頼も。" in t26, "S26 都度発注メニュー・誤字")
    stars = [sh for sh in so[19].shapes if sh.name.startswith("星")]
    check(len(stars) == 1, f"S20 ★は離脱防止バナーの1つだけ（{len(stars)}）")
    check("継続利用するツールのツール利用料＋DYM手数料" in table_of(so[15])[1][3].replace("\x0b", ""), "S16 プラン③の月額費用に条件（継続利用するツール）")
    w = est_width_in("役割分担", 16) + est_width_in("　※配信運用・改善支援は月額コンサル費に含みます", 12)
    note_w = est_width_in("※物件情報の送信（CV）以降は、当社の計測対象外です", 12)
    note_left = 5.50 + 4.72 - 0.04 - note_w                                    # 右寄せにした注記の文字の左端
    check(0.47 + 0.1 + w + 0.3 <= note_left, f"S18 役割分担の見出しの右端 {0.57 + w:.2f}in が、右寄せの注記の左端 {note_left:.2f}in に届かない")
    check(est_width_in("<参考：CV以降の流れ（個別対応はあらうち様）／当社の計測対象外>", 14) <= 6.8 - 0.2 - 0.3, "S8 下段の見出しが1行に収まる")
    check(est_width_in("※プラン①の試算。①＝友だち追加直後の物件の送信、②＝配信経由。新規・既存の内訳は②をクリック数の比で振り分け", 11) <= 9.62 - 0.22 - 0.3, "S9 注記が1行に収まる")

    # ---- ページの参照
    for i, t in enumerate(texts, 1):
        for m in re.finditer(r"(\d+)ページ", t):
            n = int(m.group(1))
            title = texts[n - 1].split("\n")[0]
            ok = (n == 13 and "想定の費用対効果" in title) or (n == 6 and "市場分析" in title)
            check(ok, f"S{i} の「{m.group(0)}」は参照先 S{n}（{title}）で合っている")

    # ---- 文字の長さ（1行に収まるはずの箱）
    s9 = {sh.shape_id: sh for sh in so[8].shapes}
    lead = s9[293].text_frame.paragraphs
    for k, (p, pt) in enumerate(zip(lead, (13, 13, 11))):
        w = est_width_in(p.text, pt)
        check(w <= 9.62 - 0.22 - 0.2, f"S9 見出し文{k + 1}行目が1行に収まる（{w:.2f}in／9.2in）")
    check(est_width_in("1か月目 月10件 → 6か月目 月33件（試算）", 18) <= 6.58 - 0.2 - 0.3, "S12 配信目標の文が1行に収まる")
    check(est_width_in("（6か月累計126件）", 14) <= 2.40 - 0.04 - 0.1, "S13 累計の文が1行に収まる")

    # ---- 字体・サイズ
    from _build.fontcheck import scan
    bad = scan(out)
    check(not bad or sum(sum(c.values()) for c in bad.values()) == 0, "メイリオ以外・指定なしの文字が0か所")
    small = []
    for i in (8, 9, 11, 12, 13, 15, 16, 18, 19, 21):        # 27（都度発注）は元の表の6〜9pt注記がそのまま残っているので除く
        sl = so[i - 1 if i < 17 else i - 2]
        for rpr in sl._element.iter("{%s}rPr" % A):
            sz = rpr.get("sz")
            if sz and int(sz) < 1000:
                small.append((i, sz))
    check(not small, f"変えたページに10pt未満の文字が無い {small[:3]}")

    # ---- パッケージ
    names = set(zo.namelist())
    ct = zo.read("[Content_Types].xml").decode("utf-8")
    parts = set(re.findall(r'PartName="/([^"]+)"', ct))
    check(parts <= names, f"[Content_Types] の Override が全部実在する {sorted(parts - names)[:3]}")
    slides = [n for n in names if re.match(r"ppt/slides/slide\d+\.xml$", n)]
    check(len(slides) == 47, f"スライドのパーツが47（{len(slides)}）")
    app = zo.read("docProps/app.xml").decode("utf-8")
    check("<Slides>47</Slides>" in app and '<vt:vector size="55" baseType="lpstr">' in app, "docProps/app.xml のスライド数")
    prels = zo.read("ppt/_rels/presentation.xml.rels").decode("utf-8")
    check(ob[DROP_INDEX].split("/")[-1] not in prels, "presentation.xml.rels に削除したスライドが無い")
    ids_ok = True
    for s in so:
        ids = [e.get("id") for e in s._element.iter("{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr")]
        ids_ok &= len(ids) == len(set(ids))
    check(ids_ok, "各スライドの図形IDが重複していない")

    print(f"\nOK {len(OKS)} / NG {len(FAILS)}")
    return not FAILS


if __name__ == "__main__":
    sys.exit(0 if main(*sys.argv[1:3]) else 1)
