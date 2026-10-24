# -*- coding: utf-8 -*-
"""提案FMT（MAX版）S42「提案に向けて明確化したいポイント④ その他」の本文を入れる（2026-10-24 殿村さん）。

    python -m _build.patch_fmt_max_s42 <入力.pptx> <出力.pptx> table|text

  table：S40・S41 と同じ表（項目／確認内容／数値・内容）。S41 の表の図形を写して5行に。
  text ：表にせず、「項目：確認内容」を5行の文字で並べる。
  管理画面の共有は下の帯にあるので、表には入れない。本文は14pt（S39〜S42 の本文を14ptにそろえる予定のため）。
"""
import copy
import sys

from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches

from _build.pptx_parts import rich

NAVY, INK = "1F2A5C", "333333"
IN = 2.54
ROWS = [
    ("配信の確認体制", "確認する方・承認の流れ・確認にかかる日数"),
    ("窓口・素材", "社内の窓口、写真・事例などの素材"),
    ("年間の予定", "繁忙期・キャンペーン・イベントの時期"),
    ("予約・顧客管理", "使っているシステムと、LINE連携の可否"),
    ("表現の決まり", "業法・社内ルールで使えない表現"),
]
SIZE = "1400"


def set_text(sh, text, size=SIZE):
    p = sh.text_frame._txBody.find(qn("a:p"))
    runs = p.findall(qn("a:r"))
    for r in runs[1:]:
        p.remove(r)
    runs[0].find(qn("a:t")).text = text
    runs[0].find(qn("a:rPr")).set("sz", size)
    end = p.find(qn("a:endParaRPr"))
    if end is not None:
        end.set("sz", size)


def table(s, src):
    by = {sh.name: sh for sh in src.shapes}
    band = next(sh for sh in s.shapes if sh.name == "Rectangle 24")
    step = Inches(0.665)
    new = []
    for k in (3, 4, 5):                                          # 見出しの3つ
        e = copy.deepcopy(by[f"Rectangle {k}"]._element)
        new.append(e)
    for r, (item, what) in enumerate(ROWS):
        for c in range(3):
            e = copy.deepcopy(by[f"Rectangle {6 + c}"]._element)  # 1行目の3つを写して、行ごとに下へずらす
            off = e.find(".//" + qn("a:off"))
            off.set("y", str(int(off.get("y")) + step * r))
            new.append(e)
    for e in new:
        band._element.addprevious(e)
    shapes = [sh for sh in s.shapes if sh._element in new]
    for sh in shapes[:3]:
        set_text(sh, sh.text_frame.text)
    for r, (item, what) in enumerate(ROWS):
        a, b, c = shapes[3 + r * 3: 6 + r * 3]
        set_text(a, item)
        set_text(b, what)
        set_text(c, "（ご記入）")


def text(s):
    paras = [[(f"{k}. {item}：", 14, True, NAVY), (what, 14, False, INK)] for k, (item, what) in enumerate(ROWS, 1)]
    tb = rich(s, 0.6 * IN, 2.15 * IN, 9.6 * IN, 3.2 * IN, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP)
    for p in tb.text_frame.paragraphs:
        p.space_after = Emu(14 * 12700)


def main(src, out, kind):
    prs = Presentation(src)
    slides = list(prs.slides)
    assert len(slides) == 43
    s = slides[41]
    assert any(sh.has_text_frame and sh.text_frame.text.endswith("その他") for sh in s.shapes)
    table(s, slides[40]) if kind == "table" else text(s)
    prs.save(out)
    print("saved:", out, kind)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2], sys.argv[3])
