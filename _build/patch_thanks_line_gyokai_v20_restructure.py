# -*- coding: utf-8 -*-
"""サンクスLINE誘導・業界別 ver1.3 → ver2.0（2026-10-06 殿村さん指示）。
1〜4枚目（全体座組／軽いCV／機能説明／重要性）が重複していたので3枚に再編。5・6枚目（業界別／事例）は触らない。全5枚。
  1 全体座組   ＝ 広告→サイト・LP→CV｜サンクスLINE誘導→LINE友だち追加→LINEで継続追客→本CV・売上→LTV最大化
  2 機能・導線 ＝ 広告・LP→軽いCV→サンクスLINE誘導→友だち追加→継続追客→本CV ＋ 想定値（100→60→30）
  3 重要性     ＝ なし（広告→CV→離脱）／あり（…→サンクスLINE→継続追客→本CV・再購入・継続利用→LTV）
  python3 _build/patch_thanks_line_gyokai_v20_restructure.py <殿村さん保存のver1.3.pptx> [出力]
"""
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

sys.path.insert(0, str(Path(__file__).resolve().parent))
from pptx_parts import *  # noqa

ROOT = Path(__file__).resolve().parent.parent
OUT = Path(sys.argv[2]) if len(sys.argv) > 2 else ROOT / "20261006_サンクスLINE誘導のご提案_業界別ver2.0.pptx"
ORANGE, ORANGE_BG, ORANGE_TX = "F59B21", "FEF0DE", "C46A00"
PALE, OFF = "E4E8F6", "F2F2F2"
L, R = 2.2, 31.7
OV = 0.35

prs = Presentation(sys.argv[1])
S = list(prs.slides)
assert S[0].shapes[0].text_frame.text.startswith("広告で反響を獲得し")
assert S[3].shapes[0].text_frame.text.startswith("CVはゴールではなく")
assert S[4].shapes[0].text_frame.text.startswith("ターゲット業界別")


def head(s, title, lead):
    keep_header_only(s)
    by = {x.name: x for x in s.shapes}
    set_paras(by["TextBox 1"], [(title, 0)])
    set_paras(by["TextBox 2"], [(lead, 0)])


def flow(s, x0, x1, y, h, steps, size=13):
    """steps = [(行, 塗り, 文字色, 幅の比, 強調)]。x0〜x1 に比率で割り付け。戻り値 = [(x, w)]"""
    k = (x1 - x0 + OV * (len(steps) - 1)) / sum(st[3] for st in steps)
    out, x = [], x0
    for i, (ls, f, tc, wr, big) in enumerate(steps):
        w = wr * k
        yy, hh = (y - 0.4, h + 0.8) if big else (y, h)
        shape(s, MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON, x, yy, w, hh, None, fill=f, adj=0.22)
        pad = 0.3 if i == 0 else 0.65
        label(s, x + pad, yy, w - pad - 0.55, hh, [(t, size + (2 if big else 0), True, tc, 0) for t in ls],
              align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        out.append((x, w))
        x += w - OV
    return out


def zones(s, ad_end, line_start, y_top=4.45, y_bottom=12.0):
    rect(s, L, y_top + 0.7, ad_end - L, 0.14, ORANGE)
    label(s, L, y_top, 6.0, 0.7, [("AD領域", 13, True, ORANGE_TX, 0)], anchor=MSO_ANCHOR.MIDDLE)
    rect(s, line_start, y_top + 0.7, R - line_start, 0.14, GREEN)
    label(s, line_start, y_top, 6.0, 0.7, [("LINE領域", 13, True, GREEN_TX, 0)], anchor=MSO_ANCHOR.MIDDLE)
    vline(s, (ad_end + line_start) / 2, y_top + 0.05, y_bottom, LGRAY, 1.25, dash=True)


def under(s, xw, items, y, color=GRAY, size=10.5, shift=0.0):
    x, w = xw
    label(s, x - 0.3 + shift, y, w + 0.3, 0.75 * len(items) + 0.3, [(t, size, False, color, 3) for t in items], align=PP_ALIGN.CENTER)


def tags(s, xw, items, y, line, color):
    x, w = xw
    tw = min(w - 0.5, 4.4)
    tx = x + (w - tw) / 2 - 0.15
    for j, t in enumerate(items):
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, tx, y + j * 0.85, tw, 0.68, [(t, 10.5, True, color, 0)], fill=WHITE, line=line,
              lw=0.75, adj=0.5, margins=(0.05, 0, 0.05, 0))


def conclusion_line(s, text, y=15.3):
    label(s, L, y, R - L, 1.3, [(text, 17, True, NAVY, 0)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


# ================= 1 全体座組 =================
s = S[0]
head(s, "広告で反響を獲得し、LINEで売上・LTVまで最大化", "集客からCV後のフォローまで、一気通貫で顧客接点を設計")
steps = [(["広告"], ORANGE, WHITE, 3.4, False), (["サイト・LP"], ORANGE_BG, ORANGE_TX, 3.6, False), (["CV"], ORANGE_BG, ORANGE_TX, 3.2, False),
         (["サンクス", "LINE誘導"], GREEN, WHITE, 4.0, False), (["LINE", "友だち追加"], GREEN_BG, GREEN_TX, 3.8, False),
         (["LINEで", "継続追客"], GREEN_BG, GREEN_TX, 3.9, False), (["本CV・売上"], PALE, NAVY, 3.8, False), (["LTV最大化"], NAVY, WHITE, 3.8, False)]
xs = flow(s, L, R, 6.3, 2.7, steps, size=12.5)
zones(s, xs[2][0] + xs[2][1] - 0.15, xs[3][0] + 0.25, y_bottom=13.6)
under(s, xs[0], ["検索広告", "SNS広告", "その他WEB広告"], 9.5)
under(s, xs[2], ["申込", "問い合わせ", "購入", "資料請求 など"], 9.5, shift=-0.2)
under(s, xs[5], ["情報提供", "リマインド", "セグメント配信", "個別フォロー"], 9.5, shift=-0.2)
under(s, xs[6], ["面談／商談", "購入", "再購入", "継続利用"], 9.5, shift=-0.2)
conclusion_line(s, "CVをゴールにせず、広告で獲得したユーザーをLINEで育成し、その後の売上までつなげる")

# ================= 2 機能・おすすめ導線 =================
s = S[1]
head(s, "軽いCVからLINEにつなぎ、本CVまで引き上げる", "サンクスLINE誘導＝CV完了画面から、LINE友だち追加へ案内する仕組み")
steps = [(["広告・LP"], LGRAY, INK, 3.4, False), (["簡易診断などの", "軽いCV"], ORANGE, WHITE, 5.4, True),
         (["サンクス", "LINE誘導"], GREEN, WHITE, 4.2, False), (["LINE", "友だち追加"], GREEN_BG, GREEN_TX, 4.2, False),
         (["LINEで", "継続追客"], GREEN_BG, GREEN_TX, 4.6, False), (["本CV"], NAVY, WHITE, 3.8, False)]
xs = flow(s, L, R, 6.3, 2.6, steps)
zones(s, xs[1][0] + xs[1][1] - 0.15, xs[2][0] + 0.25, y_bottom=12.9)
tags(s, xs[1], ["簡易診断", "チェックコンテンツ", "資料DL"], 9.75, ORANGE, ORANGE_TX)
tags(s, xs[4], ["情報提供", "リマインド", "セグメント配信", "個別フォロー"], 9.75, GREEN, GREEN_TX)
under(s, xs[5], ["問い合わせ", "面談／相談", "購入"], 9.75)
# 想定値（小さく）
hline(s, L, R, 13.6, LGRAY, 0.75)
label(s, L, 13.9, 4.0, 2.2, [("想定値", 11, True, GRAY, 0)], anchor=MSO_ANCHOR.MIDDLE)
nums = [("CV", "100件", INK), ("誘導可能", "60件", NAVY), ("LINE友だち追加", "30件", GREEN_TX)]
nx, nw = 7.0, 6.2
for i, (a, b, c) in enumerate(nums):
    x = nx + i * (nw + 2.0)
    label(s, x, 13.9, nw, 2.2, [(a, 11, False, GRAY, 2), (b, 22, True, c, 0)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    if i < 2:
        arrow_line(s, x + nw + 0.3, 15.1, x + nw + 1.7, 15.1, NAVY, 1.75)
        label(s, x + nw - 0.6, 15.3, 3.2, 0.6, [("× 60%" if i == 0 else "× 50%", 9.5, False, GRAY, 0)], align=PP_ALIGN.CENTER)
label(s, L, 16.35, R - L, 0.6, [("※誘導可能率・友だち追加率は想定値", 9, False, GRAY, 0)])

# ================= 3 重要性 =================
s = S[2]
head(s, "広告で獲得した反響を、売上につなげ切る", "CVしたユーザーの全員が、そのまま売上になるわけではない")
LW = 4.6
x0 = L + LW + 0.3
label(s, L, 5.3, LW, 2.6, [("サンクスLINE", 12, True, GRAY, 0), ("なし", 18, True, GRAY, 0)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
ari = [(["広告"], PALE, NAVY, 3.2, False), (["CV"], PALE, NAVY, 3.0, False), (["サンクス", "LINE誘導"], GREEN, WHITE, 4.0, False),
       (["LINEで", "継続追客"], GREEN_BG, GREEN_TX, 4.0, False), (["本CV／再購入", "／継続利用"], PALE, NAVY, 5.0, False), (["LTV最大化"], NAVY, WHITE, 4.0, False)]
k = (R - x0 + OV * 5) / sum(a[3] for a in ari)
nashi = [(["広告"], LGRAY, INK, 3.2, False), (["CV"], LGRAY, INK, 3.0, False), (["離脱／", "初回で終了"], OFF, GRAY, 4.0, False)]
x1 = x0 + sum(a[3] for a in nashi) * k - OV * 2
xs = flow(s, x0, x1, 5.3, 2.6, nashi, size=12.5)
lx = x1 + 0.6
label(s, lx, 5.1, R - lx, 0.7, [("CV後に取りこぼしている層", 11, True, GRAY, 0)], anchor=MSO_ANCHOR.MIDDLE)
drops = ["問い合わせ後に離脱", "比較検討中に離脱", "初回購入だけで終了", "来店／面談まで至らない"]
for j, t in enumerate(drops):
    cx = lx + (j % 2) * ((R - lx) / 2)
    cy = 5.95 + (j // 2) * 0.95
    shape(s, MSO_SHAPE.OVAL, cx, cy + 0.25, 0.3, 0.3, None, fill=GRAY)
    label(s, cx + 0.45, cy, (R - lx) / 2 - 0.5, 0.8, [(t, 11.5, False, INK, 0)], anchor=MSO_ANCHOR.MIDDLE)
hline(s, L, R, 8.9, LGRAY, 0.75)
label(s, L, 10.0, LW, 2.6, [("サンクスLINE", 12, True, GRAY, 0), ("あり", 18, True, GREEN_TX, 0)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
xs2 = flow(s, x0, R, 10.0, 2.6, ari, size=12.5)
# LINEが受け止める範囲（取りこぼし層 → LTVまで）
bx1, bx2 = xs2[2][0] + 0.3, R - 0.3
rect(s, bx1, 12.95, bx2 - bx1, 0.12, GREEN)
label(s, bx1, 13.15, bx2 - bx1, 0.7, [("LINE領域", 11, True, GREEN_TX, 0)], align=PP_ALIGN.CENTER)
conclusion_line(s, "広告で獲得した反響をLINEに残し、CV後の取りこぼしを減らして売上・LTVにつなげる")

# ================= 旧4枚目（重要性の旧版）を削除 =================
drop_slide(prs, S[3])

prs.save(str(OUT))
print("saved:", OUT.name, len(prs.slides), "枚")
