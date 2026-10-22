# -*- coding: utf-8 -*-
"""教育（塾）業界：ver3.1 の中身を、新しいFMT（MAX版・43枚・6章）の並びに入れ直した新しい資料を作る（2026-10-22 殿村さん）。

    python -m _build.build_juku_maxfmt <MAX版FMT.pptx> <塾ver3.1.pptx> <出力.pptx>

  土台はMAX版FMT。FMTのままのページ（弊社紹介・実績・ポジショニング、LINE公式アカウントの基礎、提供施策、確認したいこと、裏表紙）は残し、
  中身が空・例文のページ（市況の変化、ユーザーファネル、施策全体図、施策一覧、ステップ配信・企画配信の例、具体施策のイメージ8枚、料金プラン）は
  塾ver3.1のページに入れ替える。入れ場所が無い塾のページは、裏表紙の後ろの Appendix に置く。
  直すところ：表紙の小見出し／施策パターンのリード（FMTの「あああ」）／サンクスLINE誘導の初期費用（10万→15万円）／
  確認したいことの表の黄色の欄（白に）。赤はFMTのまま（殿村さん）。塾のページの中の「P7」などのページ番号は、新しい並びに合わせて直す。
  施策の詳細の塾のページには、FMTの「具体施策のイメージ」と同じ分類の札（露出の強化 など）を付ける。
"""
import re
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.util import Inches

from _build.patch_tanaka_v5 import copy_slide
from _build.pptx_parts import drop_slide, rect

NAVY = "1F2A5C"
IN = 2.54

# 新しい並び。("F", n)＝FMTのn枚目、("J", n)＝塾ver3.1のn枚目、("J", n, "札")＝札を付ける、("APPX",)＝Appendixの区切り
ORDER = [
    ("F", 1), ("F", 2), ("F", 3), ("F", 4),
    ("F", 5), ("J", 5), ("J", 7), ("F", 7), ("F", 8),                                   # 01 施策背景
    ("F", 9), ("J", 19), ("J", 43), ("F", 11), ("J", 24), ("J", 18), ("J", 22), ("J", 23),
    ("J", 8), ("J", 9), ("J", 10), ("J", 11), ("J", 31), ("J", 32), ("J", 12), ("J", 35),  # 02 LINE施策全体図
    ("F", 18), ("J", 15), ("J", 16), ("J", 17), ("J", 21),
    ("J", 25, "露出の強化"), ("J", 26, "露出の強化"), ("J", 36, "露出の強化"),
    ("J", 27, "CVの最大化"), ("J", 28, "CVの最大化"), ("J", 29, "CVの最大化"),
    ("J", 30, "マネタイズ最大化"), ("J", 33, "CPOの改善（CV質）"), ("J", 34, "休眠顧客の有効化"),  # 03 施策詳細
    ("F", 27), ("J", 44),                                                                # 04 料金形態
    ("F", 29), ("F", 30), ("F", 31), ("F", 32), ("F", 33), ("F", 34), ("F", 35), ("F", 36), ("F", 37), ("J", 47),  # 05 サポート体制
    ("F", 38), ("F", 39), ("F", 40), ("F", 41), ("F", 42),                              # 06 次回ご提案に向けて
    ("F", 43),
    ("APPX",), ("J", 3), ("J", 4), ("J", 13), ("J", 14), ("J", 37), ("J", 38), ("J", 39), ("J", 40), ("J", 41),
    ("J", 46), ("J", 49), ("J", 54), ("J", 55),
]
COVER_SUB = "保護者の「迷っている期間」に寄り添い、資料請求から入会までをつなぐ"
COVER_TAG = "業界：教育（学習塾）"
F11_LEAD = "LINEの施策は、目的ごとに5つのパターンに分けられる。"
J13_FIX = ("P7〜P12のデータから、取り組むべきことを整理する。", "本編の検索データから、取り組むべきことを整理する。")


def ptext(p):
    return "".join(t.text or "" for t in p.iter(qn("a:t")))


def all_paras(s):
    for sh in s.shapes:
        if sh.has_text_frame or getattr(sh, "has_table", False) and sh.has_table or sh.shape_type == 6:
            yield from sh._element.iter(qn("a:p"))


def set_runs(p, text):
    runs = p.findall(qn("a:r"))
    for r in runs[1:]:
        p.remove(r)
    runs[0].find(qn("a:t")).text = text


def replace_one(s, old, new):
    hits = [p for p in all_paras(s) if ptext(p) == old]
    assert len(hits) == 1, (old, len(hits))
    set_runs(hits[0], new)


def remap_pages(s, newno):
    """塾のページの中の「P7」「P8〜P11」を新しいページ番号に（a:t ごとに置き換える）"""
    for t in s._element.iter(qn("a:t")):
        if t.text and re.search(r"P\d", t.text):
            t.text = re.sub(r"P(\d{1,2})", lambda m: "P%d" % newno[int(m.group(1))], t.text)


def divider_title(s):
    """章の区切りの章名の箱（上の「アジェンダ」より下にある、いちばん幅の広い文字の箱）"""
    return max((sh for sh in s.shapes if sh.has_text_frame and sh.top > Inches(1)), key=lambda sh: sh.width)


def tag(s, text):
    sh = rect(s, 6.45 * IN, 0.2 * IN, 1.95 * IN, 0.32 * IN, NAVY, [(text, 11, True, "FFFFFF", 0)])
    return sh


def main(fmt_path, juku_path, out):
    prs = Presentation(fmt_path)
    J = Presentation(juku_path)
    fs, js = list(prs.slides), list(J.slides)
    assert len(fs) == 43 and len(js) == 56

    # FMTのページの直し
    replace_one(fs[0], "最上位パートナーの知見を活用し、獲得に強い動線を構築する", COVER_SUB)
    replace_one(fs[0], "MAX版", COVER_TAG)
    replace_one(fs[10], "あああああああああああああああ", F11_LEAD)
    replace_one(fs[34], "初期：10万円", "初期：15万円")
    note = next(sh for sh in fs[34].shapes if sh.has_text_frame and "価格改定の予定" in sh.text_frame.text)
    note._element.getparent().remove(note._element)    # 社内向けの注記（2026年7月〜の価格改定）は外す（10/02 サンクスLINE資料と同じ）
    for n in (39, 40, 41, 42):
        for e in fs[n - 1]._element.iter(qn("a:srgbClr")):
            if e.get("val").upper() == "FFF8E1":
                e.set("val", "FFFFFF")
    for s in fs:                                        # 紺の札の上の黄色い文字（+通数単価・+設定地点追加・※CPF広告※）は白に
        for e in s._element.iter(qn("a:srgbClr")):
            if e.get("val").upper() == "FFFF00":
                e.set("val", "FFFFFF")
    sub = next(sh for sh in fs[0].shapes if sh.has_text_frame and sh.text_frame.text == COVER_SUB)
    sub.left, sub.width = Inches(1.46), Inches(7.91)    # 表紙の小見出しを下線の幅まで広げて1行に
    for n in (5, 9, 18, 27, 29, 38):                    # 章の区切りの章名が2行に折れないように
        divider_title(fs[n - 1]).width = Inches(8.0)

    # 新しい並びのスライドを用意する
    lst = prs.slides._sldIdLst
    fid = {k + 1: e for k, e in enumerate(list(lst))}
    seq, new_of_j, tags = [], {}, {}
    for item in ORDER:
        if item[0] == "F":
            seq.append(fid[item[1]])
        elif item[0] == "APPX":
            d = copy_slide(prs, fs[4])                 # アジェンダの区切り（01 施策背景）を写して Appendix に（番号と下線は外す）
            for sh in list(d.shapes):
                if (sh.has_text_frame and sh.text_frame.text == "01") or sh.shape_type == 9 or "コネクタ" in sh.name:
                    sh._element.getparent().remove(sh._element)
            replace_one(d, "施策背景", "Appendix（参考資料）")
            t = divider_title(d)
            t.left, t.width = Inches(0.66), Inches(9.0)
            seq.append(list(lst)[-1])
        else:
            d = copy_slide(prs, js[item[1] - 1])
            e = list(lst)[-1]
            seq.append(e)
            new_of_j[item[1]] = d
            if len(item) == 3:
                tags[item[1]] = item[2]
    used_f = {item[1] for item in ORDER if item[0] == "F"}
    for k, s in enumerate(fs, start=1):                 # 使わないFMTのページを消す
        if k not in used_f:
            drop_slide(prs, s)
    for e in seq:                                       # 並べ替え
        lst.remove(e)
        lst.append(e)
    order_ids = [s.slide_id for s in prs.slides]
    newno = {j: order_ids.index(d.slide_id) + 1 for j, d in new_of_j.items()}

    for j, d in new_of_j.items():
        if j == 13:
            replace_one(d, *J13_FIX)
        remap_pages(d, newno)
        if j in tags:
            tag(d, tags[j])
    prs.save(out)
    print("saved:", out, len(prs.slides), "枚")
    print("塾→新:", dict(sorted(newno.items())))


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2], sys.argv[3])
