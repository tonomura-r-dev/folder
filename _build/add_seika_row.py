# -*- coding: utf-8 -*-
"""Slide 24に成果報酬プランの1行を併記"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
TARGET = ROOT / "20260804_株式会社ビイサイドプランニング御中_公式LINEのご提案v08.pptx"

NAVY = "1F285A"; INK = "333333"; MUT = "808080"; WHITE = "FFFFFF"; BORDER = "D9D9D9"

prs = Presentation(str(TARGET))
slides = list(prs.slides)


def set_font(run, size, bold=None, color=INK, name="メイリオ"):
    f = run.font; f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    f.name = name
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        e = rPr.find(qn(tag))
        if e is None:
            e = rPr.makeelement(qn(tag), {}); rPr.append(e)
        e.set("typeface", name)
    f.color.rgb = RGBColor.from_string(color)


ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
ANCH = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}


def reset_tf(tf):
    for para in tf.paragraphs[1:]:
        para._p.getparent().remove(para._p)
    for r in list(tf.paragraphs[0].runs):
        r._r.getparent().remove(r._r)


def put_text(tf, paras, anchor="m", ml=0.05, mr=0.05, mt=0.02, mb=0.02):
    reset_tf(tf); tf.word_wrap = True
    tf.margin_left = Inches(ml); tf.margin_right = Inches(mr)
    tf.margin_top = Inches(mt); tf.margin_bottom = Inches(mb)
    tf.vertical_anchor = ANCH[anchor]
    first = True
    for p in paras:
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        para.alignment = ALIGN[p.get("align", "l")]
        if p.get("sa") is not None:
            para.space_after = Pt(p["sa"])
        if p.get("ls") is not None:
            para.line_spacing = p["ls"]
        for (t, sz, b, c) in p["runs"]:
            r = para.add_run(); r.text = t
            set_font(r, sz, b, c)


s = slides[23]

# 帯（獲得単価）と注記の位置を上に詰めて、成果報酬の行を作る
band = note = None
for sh in s.shapes:
    if not sh.has_text_frame:
        continue
    t = sh.text_frame.text
    if "LINEで増える来場者の獲得単価" in t:
        band = sh
    elif "CV＝フェア事前申込" in t:
        note = sh

assert band is not None and note is not None

band.top = Inches(5.70)
band.height = Inches(0.46)
note.left, note.top = Inches(7.20), Inches(6.32)
note.width, note.height = Inches(3.30), Inches(0.70)

# 成果報酬プランの1行
box = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.36), Inches(6.26), Inches(6.60), Inches(0.80))
box.fill.background()
box.line.color.rgb = RGBColor.from_string(BORDER)
box.line.width = Pt(1.0)
box.shadow.inherit = False
put_text(box.text_frame,
         [{"runs": [("［成果報酬プラン］", 11, True, NAVY),
                    ("　初期23.5万円＋月次12.98万円＋成果報酬（フェア事前申込 8,000円／件）", 10.5, None, INK)], "sa": 3},
          {"runs": [("半年総額 146.98万円・CPA 10,067円　／　来場引き上げ分は成果報酬の対象外", 10, None, INK)], "ls": 1.1}],
         anchor="m", ml=0.18, mr=0.14)

prs.save(str(TARGET))
print("saved:", TARGET)
