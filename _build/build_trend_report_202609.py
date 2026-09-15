"""
LINE公式アカウント 媒体最新情報 2026年9月 トレンドレポート ビルドスクリプト。

トンマナは実際のDYM営業資料FMT（_templates/DYM_LINEOA_BUFFF_FMT.pptx）から抽出。
このFMTをコピーし、必要な9枚だけ残して clear_slide() → 作り直す経路
（build_special_plan.py と同じ。スライドの新規追加はしない＝CLAUDE.md参照）。

ヘッダーのネイビー角アイコン・区切り線・DYMロゴ・フッターの著作権表記／ページ番号は
スライドレイアウト側に定義されており、clear_slide() で個々のシェイプを消しても
自動的に継承表示される（要検証済み：2026-09、test_inherit.jpg）。
なので自前では描画しない。

採否は上長確認済み（2026-09）：
  ・［提供終了］リサーチ機能終了 → 先月既出のため除外
  ・既読API（旧）新規受付終了 → 除外。新しい既読APIの一般提供開始をメインに書き直し
  ・早期発注特典①の締切 → 期限切れのため除外
  ・LINE Sales Promotion Manager → 優先度低のため掲載順を最後に

使い方: python3 _build/build_trend_report_202609.py
"""

import shutil
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
SRC = str(ROOT / "_templates" / "DYM_LINEOA_BUFFF_FMT.pptx")
OUT = str(ROOT / "LINEOA媒体最新情報_2026年9月トレンドレポート.pptx")

FONT = "メイリオ"

# 実FMTから抽出した配色（2026-09、LINEOA_BUFFF_3.pptxを解析）
TITLE_NAVY = RGBColor(0x00, 0x20, 0x60)   # ヘッダー見出し・本文タイトル色
CHIP_NAVY = RGBColor(0x10, 0x25, 0x3F)    # セクション見出しチップ・フラグチップ
BODY_GRAY = RGBColor(0x34, 0x34, 0x34)    # 本文
META_GRAY = RGBColor(0x49, 0x48, 0x48)    # メタ情報（対象業界・フェーズ等）
RED = RGBColor(0xFF, 0x00, 0x00)          # 強調・注意（FMT実測）
LIGHT_BOX = RGBColor(0xF2, 0xF2, 0xF2)    # メインメッセージ等のボックス塗り
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BORDER_GRAY = RGBColor(0xD9, 0xD9, 0xD9)
SOURCE_GRAY = RGBColor(0x80, 0x80, 0x80)

SLIDE_W = Inches(10.83)
SLIDE_H = Inches(7.5)

N_SLIDES = 9  # 表紙1 + 目次1 + トピック7

# ヘッダーの角アイコン（レイアウト側で幅0.49in・高さ0.6in）と被らない開始位置
CONTENT_LEFT = Inches(0.62)
CONTENT_RIGHT = Inches(10.33)
CONTENT_W = CONTENT_RIGHT - CONTENT_LEFT


def set_font_all(text_frame, name=FONT):
    for para in text_frame.paragraphs:
        for run in para.runs:
            run.font.name = name
            rPr = run._r.get_or_add_rPr()
            for tag in ("a:latin", "a:ea", "a:cs"):
                el = rPr.find(qn(tag))
                if el is None:
                    el = rPr.makeelement(qn(tag), {})
                    rPr.append(el)
                el.set("typeface", name)


def clear_slide(slide):
    """スライド自身が持つシェイプだけを消す。レイアウト/マスター側の
    ヘッダーアイコン・区切り線・DYMロゴ・フッターは継承表示のため残る。"""
    spTree = slide.shapes._spTree
    for el in list(spTree):
        if el.tag.split("}")[-1] in ("sp", "cxnSp", "pic", "graphicFrame", "grpSp"):
            spTree.remove(el)


def load_trimmed_fmt(n_slides):
    """FMTをコピーし、先頭n_slides枚だけ残す（追加はしない）。"""
    shutil.copyfile(SRC, OUT)
    prs = Presentation(OUT)
    sldIdLst = prs.slides._sldIdLst
    ids = list(sldIdLst)
    keep_ids = ids[:n_slides]
    for sldId in ids:
        if sldId in keep_ids:
            continue
        prs.part.drop_rel(sldId.rId)
        sldIdLst.remove(sldId)
    slides = list(prs.slides)
    assert len(slides) == n_slides, len(slides)
    for slide in slides:
        clear_slide(slide)
    return prs, slides


def add_rect(slide, x, y, w, h, fill=None, line=None, line_w=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shp.shadow.inherit = False
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = line_w or Pt(0.75)
    return shp


def add_text(slide, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
             line_spacing=None, word_wrap=True):
    """runs: list of paragraphs; each paragraph is a list of (text, size, bold, color) tuples."""
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = word_wrap
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for i, para_runs in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if line_spacing:
            p.line_spacing = line_spacing
        for text, size, bold, color in para_runs:
            r = p.add_run()
            r.text = text
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = color
    set_font_all(tf)
    return box


def add_bullets(slide, x, y, w, h, items, size=10, color=BODY_GRAY, bullet="・",
                 line_spacing=1.2):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = line_spacing
        p.space_after = Pt(3)
        r = p.add_run()
        r.text = f"{bullet}{item}"
        r.font.size = Pt(size)
        r.font.color.rgb = color
    set_font_all(tf)
    return box


def chip(slide, x, y, w, h, text, fill=CHIP_NAVY, color=WHITE, size=9.5, bold=True):
    c = add_rect(slide, x, y, w, h, fill=fill)
    add_text(slide, x, y, w, h, [[(text, size, bold, color)]],
              align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    return c


def page_title(slide, text, size=17):
    """ヘッダー見出し（レイアウト側の角アイコンの右、区切り線の上）。"""
    add_text(slide, CONTENT_LEFT, Inches(0.1), Inches(8.1), Inches(0.42),
              [[(text, size, True, TITLE_NAVY)]], anchor=MSO_ANCHOR.MIDDLE)


def meta_row(slide, industry=None, phase=None, flag_agency=None, flag_noshare=None):
    """区切り線のすぐ下：号数チップ＋対象業界／フェーズ／フラグ（空欄可）。"""
    y = Inches(0.7)
    chip(slide, CONTENT_LEFT, y, Inches(0.85), Inches(0.28), "9月号", size=9)
    ind = industry or "－"
    ph = phase or "－"
    ag = flag_agency or "－"
    ns = flag_noshare or "－"
    text = f"対象業界：{ind}　　フェーズ：{ph}　　代理店限定：{ag}　　展開禁止：{ns}"
    add_text(slide, CONTENT_LEFT + Inches(1.0), y, Inches(8.8), Inches(0.28),
              [[(text, 9.5, False, META_GRAY)]], anchor=MSO_ANCHOR.MIDDLE)


def source_line(slide, url):
    if not url:
        return
    add_text(slide, CONTENT_LEFT, Inches(6.86), Inches(9.5), Inches(0.24),
              [[(f"出典：{url}", 7.5, False, SOURCE_GRAY)]], anchor=MSO_ANCHOR.MIDDLE)


def cover_slide(slide):
    add_text(slide, Inches(0.7), Inches(2.5), Inches(9), Inches(0.4),
              [[("LINE公式アカウント", 14, False, META_GRAY)]])
    add_text(slide, Inches(0.7), Inches(2.85), Inches(9.3), Inches(1.0),
              [[("媒体最新情報　9月 トレンドレポート", 30, True, TITLE_NAVY)]],
              line_spacing=1.15)
    add_rect(slide, Inches(0.7), Inches(3.95), Inches(9.0), Pt(1.5), fill=TITLE_NAVY)
    chip(slide, Inches(0.7), Inches(4.15), Inches(2.4), Inches(0.4), "2026年9月号",
         fill=LIGHT_BOX, color=META_GRAY, size=11, bold=False)
    add_text(slide, Inches(0.7), Inches(6.5), Inches(6), Inches(0.4),
              [[("株式会社DYM（DYM × LINEOA）", 12, True, TITLE_NAVY)]])


def agenda_slide(slide, topics):
    page_title(slide, "目次")
    y = Inches(1.0)
    row_h = Inches(0.62)
    for i, t in enumerate(topics, start=1):
        chip(slide, CONTENT_LEFT, y, Inches(0.44), Inches(0.42), str(i), size=12)
        add_text(slide, CONTENT_LEFT + Inches(0.62), y, Inches(9.0), Inches(0.42),
                  [[(t, 13, False, BODY_GRAY)]], anchor=MSO_ANCHOR.MIDDLE)
        y += row_h


def topic_slide(slide, data):
    page_title(slide, data["title"])
    meta_row(slide, industry=data.get("industry"), phase=data.get("phase"),
             flag_agency=data.get("flag_agency"), flag_noshare=data.get("flag_noshare"))

    # メインメッセージ
    msg_y = Inches(1.1)
    add_rect(slide, CONTENT_LEFT, msg_y, CONTENT_W, Inches(0.62), fill=LIGHT_BOX)
    add_text(slide, CONTENT_LEFT + Inches(0.2), msg_y, CONTENT_W - Inches(0.4), Inches(0.62),
              [[(data["message"], 12, True, TITLE_NAVY)]],
              anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.1)

    # 4区分ボックス（2x2）
    box_top = Inches(1.9)
    box_h = Inches(2.4)
    gap = Inches(0.15)
    hgap = Inches(0.2)
    box_w = (CONTENT_W - hgap) / 2
    positions = [
        (CONTENT_LEFT, box_top),
        (CONTENT_LEFT + box_w + hgap, box_top),
        (CONTENT_LEFT, box_top + box_h + gap),
        (CONTENT_LEFT + box_w + hgap, box_top + box_h + gap),
    ]
    sections = [
        ("概要・変更点", data.get("overview", [])),
        ("料金体系・出稿条件", data.get("pricing", [])),
        ("スケジュール・導入フロー", data.get("schedule", [])),
        ("注意点・影響", data.get("caution", [])),
    ]
    chip_h = Inches(0.34)
    for (x, y), (label, items) in zip(positions, sections):
        add_rect(slide, x, y, box_w, box_h, fill=WHITE, line=BORDER_GRAY, line_w=Pt(0.75))
        chip(slide, x, y, box_w, chip_h, label, size=10.5)
        items = items or ["－"]
        add_bullets(slide, x + Inches(0.18), y + chip_h + Inches(0.1),
                    box_w - Inches(0.36), box_h - chip_h - Inches(0.2), items)

    source_line(slide, data.get("source"))


TOPICS_LIST = [
    "1. LINEミニアプリ お気に入り追加促進機能",
    "2. 既読API 一般提供開始のお知らせ",
    "3. LINE公式アカウント「分析」データ参照期間 過去36カ月に変更",
    "4. プロモーションスタンプ 効果改善アップデート",
    "5. 友だち追加オプション 複数アカウント対応",
    "6. LINEオープンキャンペーン 26年10-12月 期間限定キャンペーン",
    "7. LINE Sales Promotion Manager 利用停止期間",
]

TOPICS = [
    dict(
        title="LINEミニアプリ お気に入り追加促進機能",
        message="特定条件を満たしたユーザーに、ミニアプリのお気に入り追加を促すポップアップが自動表示される新機能。",
        overview=["トリガー①：マイミニアプリの「履歴」からアクセス",
                   "トリガー②：7日以内に2日以上アクセス",
                   "条件を満たすとお気に入り追加を促すポップアップが表示"],
        pricing=["－"],
        schedule=["－（リリース日は要確認）"],
        caution=["リピート利用ユーザーの再訪導線が強化される。ミニアプリ運用中クライアントへの案内材料になる"],
        source="https://workers-hub.box.com/s/9uo4m0fiqyxmddz5br3xe9691w0q4ht7",
    ),
    dict(
        title="既読API「一般提供開始」のお知らせ",
        message="事前申請不要の新しい既読API（Messaging API）が一般提供開始。トークの既読管理をAPI側から自動化できる。",
        overview=["新しい既読API（Messaging API「mark as read」）が一般提供開始",
                   "事前の利用申請が不要になり、通常のMessaging API開発フローの中でそのまま利用可能に"],
        pricing=["該当なし（Messaging APIの利用範囲内）"],
        schedule=["一般提供開始済み",
                   "仕様：developers.line.biz/ja/docs/messaging-api/mark-as-read/"],
        caution=["CS・チャットツールと連携した運用で、既読管理の自動化に活用できる",
                   "旧既読API（LBPM経由の申請制）を使っていた案件は新APIへの移行を検討"],
        source="https://workers-hub.box.com/s/nxgvv2ejax65j6cadz7gdsdalxd8zioa",
    ),
    dict(
        title="LINE公式アカウント「分析」データ参照期間 過去36カ月に変更",
        message="「分析」タブの参照期間が過去36カ月までに変更予定。※社外告知は延期中、日程未確定につき最新情報を要確認。",
        overview=["Web版管理画面／管理アプリの「分析」タブで確認できる実績値の参照期間を過去36カ月までに変更"],
        pricing=["該当なし"],
        schedule=["リリース日：2026年11月頃（予定）",
                   "社外告知：当初9/2予定だったが内容変更の可能性ありいったん延期。改めて案内予定"],
        caution=["36カ月より前の実績データが必要なクライアントには変更前の確認・保存を案内",
                   "告知日程が流動的なため最新情報を要確認"],
        source="https://workers-hub.box.com/s/ms8iztyci2dn5z2qna88rklxwxp7o5ja",
    ),
    dict(
        title="プロモーションスタンプ 効果改善アップデート",
        message="LINEキーボードのサジェスト機能で無料スタンプが表示される条件が緩和され、全ユーザーに表示可能に。DL数は平均+5.63%増加。",
        overview=["サジェスト面での無料スタンプ表示ユーザー条件を緩和し全ユーザーに表示されるよう変更",
                   "配信中の案件を含む全案件に自動適用（広告主・案件側の対応は不要）"],
        pricing=["追加費用なし（自動適用）"],
        schedule=["リリース日：2026年8月28日（金）"],
        caution=["A/Bテスト結果：1案件あたり平均DL数+5.63%増",
                   "直近18ヶ月DL履歴のないユーザーでも約2倍のDL、普段DLしない層への到達も改善",
                   "運用中クライアントへの実績報告・追加提案の材料に使える"],
        source="https://workers-hub.box.com/s/rgfhp6hxyusj7domi1irjx68jld23k9r",
    ),
    dict(
        title="友だち追加オプション 複数アカウント対応",
        message="1つのLINEミニアプリに連携するLINE公式アカウントを複数設定できるようになる。",
        overview=["LINEミニアプリと連携するLINE公式アカウントを複数設定可能に"],
        pricing=["－"],
        schedule=["リリース予定日：9月2日"],
        caution=["複数ブランド・複数拠点アカウントを運用するクライアントの友だち獲得導線が広がる"],
        source="https://workers-hub.box.com/s/ko2jjwgxk5fs18ejrs97fvfqv3bm3089",
    ),
    dict(
        title="LINEオープンキャンペーン 26年10-12月 期間限定キャンペーン",
        message="2026年10月20日〜12月25日に終了する案件限定で、特別価格・特別期間・特別通数の3特典が適用される。",
        overview=["対象：2026年10月20日開始〜12月25日までに終了する案件限定",
                   "景品はLINEポイント"],
        pricing=["特典①特別価格：基本費用（キャンペーン価格）1,500万円",
                   "特典②特別掲載期間：キャンペーン掲載期間15日間",
                   "特典③特別通数：告知メッセージ（1通）の通数150万通増量（最大1,050万通）"],
        schedule=["キャンペーン期間：2026年10月20日〜12月25日終了の案件が対象"],
        caution=["期間限定・条件限定のため、対象クライアントへは早めの提案・申込誘導が必要"],
        source="https://workers-hub.box.com/s/odilqpe44tp398zlfko148yggcgj68ny",
    ),
    dict(
        title="LINE Sales Promotion Manager 利用停止期間",
        message="データベース管理システムのスケール最適化のため、LINEで応募の管理画面が一時利用停止となる。",
        overview=["データベース管理システムのスケール最適化実施のため対象期間中は管理画面が利用停止"],
        pricing=["該当なし"],
        schedule=["社外への案内開始日・リリース日：2026年9月2日",
                   "－（停止期間の具体的な日時は要確認）"],
        caution=["停止期間中は応募管理・当選者確認等の作業ができない",
                   "稼働中キャンペーンがあるクライアントには事前周知が必要"],
        source="https://workers-hub.box.com/s/mo3y54qld9vj4kung0c4040z41mfhoip",
    ),
]


def main():
    prs, slides = load_trimmed_fmt(N_SLIDES)

    cover_slide(slides[0])
    agenda_slide(slides[1], TOPICS_LIST)

    for slide, data in zip(slides[2:], TOPICS):
        topic_slide(slide, data)

    prs.save(OUT)
    print(f"saved: {OUT}")


if __name__ == "__main__":
    main()
