# -*- coding: utf-8 -*-
"""チェングロウス ver2.6 → ver2.7：表紙の次に「弊社実績」（LINEヤフー Partner Award）を追加（2026-10-02）
受賞スライドの原本は 20260825_テン・エンタープライズ御中…pptx の2枚目（DYM共通の弊社実績ページ）。
  python3 _build/patch_chengrowth_v27_jisseki.py <ver2.6.pptx> <弊社実績の原本.pptx>
"""
import copy, sys
from pathlib import Path
from pptx import Presentation
from pptx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "20260929_チェングロウス_LINE公式アカウント運用のご提案ver2.7.pptx"
prs = Presentation(sys.argv[1])
src = Presentation(sys.argv[2]).slides[1]

layout = prs.slides[1].slide_layout
new = prs.slides.add_slide(layout)
for sh in list(new.shapes):
    sh._element.getparent().remove(sh._element)
for el in src.shapes._spTree:
    if el.tag in (qn("p:nvGrpSpPr"), qn("p:grpSpPr")):
        continue
    el = copy.deepcopy(el)
    # 画像の参照を新スライド側に張り直す
    for blip in el.iter(qn("a:blip")):
        rid = blip.get(qn("r:embed"))
        if rid:
            img = src.part.related_part(rid)
            _, new_rid = new.part.get_or_add_image_part(__import__("io").BytesIO(img.blob))
            blip.set(qn("r:embed"), new_rid)
    new.shapes._spTree.append(el)

# 2枚目に移動
lst = prs.slides._sldIdLst
item = lst[-1]
lst.remove(item)
lst.insert(1, item)

for sl in prs.slides:
    for n, c in enumerate(sl.shapes._spTree.iter(qn("p:cNvPr")), start=2):
        c.set("id", str(n))

AFTER = {qn(x) for x in ("a:sym", "a:hlinkClick", "a:hlinkMouseOver", "a:rtl", "a:extLst")}
sl = prs.slides[1]
for r in sl.shapes._spTree.iter(qn("a:r")):
    if r.find(qn("a:rPr")) is None:
        r.insert(0, r.makeelement(qn("a:rPr"), {"lang": "ja-JP"}))
for tag in ("a:rPr", "a:endParaRPr", "a:defRPr"):
    for rpr in sl.shapes._spTree.iter(qn(tag)):
        fs = []
        for ft in ("a:latin", "a:ea", "a:cs"):
            e = rpr.find(qn(ft))
            if e is None:
                e = rpr.makeelement(qn(ft), {})
            else:
                rpr.remove(e)
            e.set("typeface", "メイリオ")
            fs.append(e)
        nxt = next((c for c in rpr if c.tag in AFTER), None)
        for e in fs:
            (nxt.addprevious(e) if nxt is not None else rpr.append(e))
prs.save(str(OUT))
print("saved:", OUT.name, len(prs.slides), "枚")
