# -*- coding: utf-8 -*-
"""教育（塾）業界 新FMT版（70枚・殿村さんPC保存版）を、10/29の提案FMT（MAX版・53枚）の形にそろえる（2026-10-29 殿村さん）。

    python -m _build.build_juku_fmt1029 <塾のPC保存版.pptx> <FMT最新.pptx> <FMT具体施策枠のみ.pptx> <出力.pptx>

  殿村さんの決定：A 料金形態＝FMTの4枚／B サポート体制＝FMTの5枚／C 具体施策のイメージ8枚を塾の構築に合わせて入れる。
  ・01：FMTの弊社紹介2枚目（P3）と実例紹介（P10）を足す
  ・03：アジェンダ03の直後に具体施策のイメージ8枚（中身は塾の構築①〜⑤・体験予約から入会まで・過去に問い合わせた保護者・年間の配信カレンダーから）
  ・04：ご提供プラン（旧P41）→ FMTの料金4枚
  ・05：体制（旧P51）と新規友だちの追加など（旧P43〜P50）→ FMTのサポート体制5枚
  ・06：提案に向けて明確化したいポイント①〜④ → FMTの表の形（記入欄・中央）。リードは塾用のまま
  ・新規友だちの追加〜提供施策（FMTのP45〜P52）は06の後ろ・裏表紙の前（FMTと同じ並び）
  FMTから入れたページは、黄色（FFF8E1の欄・FFFF00の文字）を白に、「貴社」を「貴塾」に、転職ユーザーを保護者に。
  料金表のサンクスLINE誘導は初期15万円・月額3万円〜、離脱防止は初期1.5万円・月額3万円（ポップアップのページと同じ）。
  塾のページの中の「P◯」は新しい並びに振り直す。
"""
import re
import sys

from lxml import etree
from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

from _build.patch_fmt_max_gutai import RIGHT, RIGHT_Y, flow, IN, NAVY, INK
from _build.patch_tanaka_v5 import copy_slide
from _build.pptx_parts import rich

A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"

GUTAI = [  # 札, [入口, LINE, 配信, ゴール], (目的, 対象, 中身, 見る数字)  右の文は「|」で改行
    ("LINEの友だちへ",
     [("サイト・チラシ・教室・友だちの紹介", "離脱防止ポップアップ／チラシ・教室のQR／|友だちからの紹介"),
      ("友だち追加", "資料と診断を受け取る"), "あいさつメッセージ・30秒ぴったり診断", "友だちとしてつながり続ける"],
     ("塾を決めていない保護者とも、|LINEでつながる",
      "塾を探している保護者|（一番の悩みは「合う塾が分からない」）",
      "離脱防止ポップアップ／|チラシ・教室掲示のQR／|友だちからの紹介",
      "友だち追加数／入口ごとの追加数")),
    ("露出の強化",
     ["友だち全体", ("企画配信", "年間の配信カレンダー（講習・定期テスト・新学年）"),
      "友だちと一緒の体験枠をLINEで案内", "友だちの保護者の友だち追加・体験予約"],
     ("講習・イベントを広く知らせ、|保護者どうしの口コミをLINEにつなげる",
      "友だち全体|（体験・入会の前の保護者も含む）",
      "年間の配信カレンダーに沿った企画配信／|友だちと一緒の体験枠／|診断結果を友だちに送れる",
      "配信到達数・開封率／|紹介からの友だち追加数")),
    ("CVの最大化",
     ["友だち追加から14日以内の保護者", "30秒ぴったり診断（4問）",
      ("14日間の配信", "体験のご案内は7日後・14日後の2通"), "資料請求（CV①）・体験予約（CV②）"],
     ("比べている14日間に、|資料請求・体験予約へ進んでいただく",
      "友だち追加から14日以内の保護者",
      "費用・時期・口コミの疑問に答える配信／|診断の回答で合う学び方をお伝え／|右下に無料体験・見学の予約",
      "資料請求数／体験予約数")),
    ("CPOの改善（CV質）",
     ["体験を予約した保護者", "予約の確認・カレンダー登録",
      ("体験の前日までの配信", "講師の紹介・持ち物とアクセス・日程の変更方法"), "体験への来塾・入会"],
     ("体験を予約した後、当日来られる方を|増やし、入会につなげる",
      "体験を予約した保護者",
      "予約の直後から翌日まで5回の自動配信／|日程の変更はトークで受け付け／|翌日に診断結果と入会のご相談",
      "体験の来塾率／キャンセル率／|体験から入会の割合")),
    ("休眠顧客の有効化",
     ["資料請求・体験だけで止まった保護者", "通知メッセージ（電話番号でLINEに届く）",
      ("講習の時期の再案内", "夏期講習・冬期講習・新学年の前"), "体験予約・講習の申し込み"],
     ("止まっている保護者に、|講習の時期にもう一度ご案内する",
      "資料請求だけで止まった保護者／|体験に来たが入会しなかった保護者",
      "夏期講習・冬期講習・新学年の前に再案内／|いつ・何で接点があったかを書いた文面",
      "再案内からの体験予約数／|講習の申し込み数")),
    ("マネタイズ最大化",
     ["診断に答えた保護者", ("診断タグ＋行動タグ", "学年・悩み・性格・時期／押したボタン・予約の有無"),
      "保護者ごとに「うちの子の話」になる配信", "体験予約・入会"],
     ("保護者ごとに合う内容を届け、|入会まで進めていただく",
      "診断に答えた保護者／|資料請求したが体験予約していない保護者",
      "回答と行動で配信を出し分け／|資料請求だけの方に体験の案内を追加／|「情報収集だけ」の方は週1回の情報のみ",
      "体験予約数／入会数／ブロック率")),
    ("リピート施策",
     ["入会した保護者", "学年・コースでタグ付け",
      ("講習・定期テストの時期の配信", "年間の配信カレンダー"), "講習の受講・通塾の継続"],
     ("入会の後も、講習・テストの時期に|必要なご案内を届ける",
      "入会した保護者",
      "夏期講習・冬期講習・春期講習の案内／|定期テストの前後の学習のご案内／|学年・コースに合わせた出し分け",
      "講習の申し込み数／継続率")),
    ("コミュニティ運営",
     ["通塾中の保護者", "通塾中の保護者のタグ付け",
      ("教室からのお知らせ・紹介のご案内", "友だちと一緒に（LINEの転送で紹介）"), "通塾の継続・紹介"],
     ("通塾中の保護者とのつながりを強くし、|紹介につなげる",
      "通塾中の保護者",
      "教室・講習のお知らせ／|LINEの転送で友だちに紹介／|紹介の特典の案内（中身は貴塾で決める）",
      "紹介からの友だち追加数／|紹介からの体験予約数")),
]

TEXT_FIX = [("貴社", "貴塾"), ("貴塾保有リスト をターゲット", "貴塾保有リストをターゲット"), ("転職ユーザーのカスタマージャーニー", "保護者のカスタマージャーニー")]


def fill_gutai(s, page):
    tag, steps, right = page
    assert any(sh.has_text_frame and sh.text_frame.text == tag for sh in s.shapes), tag
    flow(s, steps)
    for (y, label), text in zip(zip(RIGHT_Y, RIGHT), right):
        rich(s, 6.97 * IN, (y + 0.06) * IN, 3.46 * IN, 1.01 * IN,
             [[(label, 13, True, NAVY)]] + [[(line, 12, False, INK)] for line in text.split("|")], anchor=MSO_ANCHOR.TOP)


def fmt_fix(s):
    """FMTから入れたページ：黄色を白に、言葉を塾向けに、料金表のツール費をそろえる"""
    root = s._element
    for el in root.iter("{%s}srgbClr" % A_NS):
        v = el.get("val", "").upper()
        if v == "FFF8E1":
            el.set("val", "FFFFFF")
        elif v == "FFFF00" and el.getparent().getparent().tag.endswith("}rPr"):
            el.set("val", "FFFFFF")
    for t in root.iter("{%s}t" % A_NS):
        if t.text:
            for a, b in TEXT_FIX:
                t.text = t.text.replace(a, b)
    for sh in s.shapes:
        if sh.has_table:
            for row in sh.table.rows:
                cs = row.cells
                for i, c in enumerate(cs):
                    if c.text == "サンクスLINE誘導ツール" and i + 2 < len(cs):
                        set_cell(cs[i + 1], "¥150,000"), set_cell(cs[i + 2], "¥30,000~")
                    if c.text == "離脱防止ツール" and i + 2 < len(cs):
                        set_cell(cs[i + 1], "¥15,000"), set_cell(cs[i + 2], "¥30,000")


def set_cell(c, text):
    ps = c.text_frame.paragraphs
    rs = ps[0].runs
    rs[0].text = text
    for r in rs[1:]:
        r.text = ""


def set_lead(s, text):
    sh = next(x for x in s.shapes if x.name == "Rectangle 2")
    rs = sh.text_frame.paragraphs[0].runs
    rs[0].text = text
    for r in rs[1:]:
        r.text = ""


def renumber(s, mapping):
    """塾のページの中の「P◯」「P◯〜P◯」を新しい番号に"""
    for t in s._element.iter("{%s}t" % A_NS):
        if t.text and re.search(r"P\d", t.text):
            def rep(m):
                n = int(m.group(1))
                assert n in mapping, ("消えたページを指している", n, t.text)
                return "P%d" % mapping[n]
            t.text = re.sub(r"P(\d+)", rep, t.text)


def main(juku_path, fmt_path, fmt_empty_path, out):
    prs = Presentation(juku_path)
    fmt, emp = Presentation(fmt_path), Presentation(fmt_empty_path)
    assert len(prs.slides) == 70 and len(fmt.slides) == 53 and len(emp.slides) == 53
    J = {i + 1: s for i, s in enumerate(prs.slides)}
    leads = {n: next(x for x in J[n].shapes if x.name == "Rectangle 2").text_frame.text for n in (53, 54, 55, 56)}

    def add(src):
        s = copy_slide(prs, src)
        fmt_fix(s)
        return s

    F = {n: add(fmt.slides[n - 1]) for n in (3, 10, 30, 31, 32, 33, 35, 36, 37, 38, 39, 41, 42, 43, 44,
                                             45, 46, 47, 48, 49, 50, 51, 52)}
    E = {n: add(emp.slides[n - 1]) for n in range(18, 26)}
    for n, page in zip(range(18, 26), GUTAI):
        fill_gutai(E[n], page)
    for fn, jn in zip((41, 42, 43, 44), (53, 54, 55, 56)):
        set_lead(F[fn], leads[jn])

    order = ([J[1], J[2], F[3]] + [J[n] for n in range(3, 10)] + [F[10]] + [J[n] for n in range(10, 27)]
             + [E[n] for n in range(18, 26)] + [J[n] for n in range(27, 41)] + [F[n] for n in (30, 31, 32, 33)]
             + [J[42]] + [F[n] for n in (35, 36, 37, 38, 39)] + [J[52]] + [F[n] for n in (41, 42, 43, 44)]
             + [F[n] for n in range(45, 53)] + [J[57]] + [J[n] for n in range(58, 71)])
    assert len(order) == 87
    # 旧番号 → 新番号（塾のページだけ）
    mapping = {}
    for i, s in enumerate(order, 1):
        for n, js in J.items():
            if js is s:
                mapping[n] = i
    mapping.update({53: order.index(F[41]) + 1, 54: order.index(F[42]) + 1, 55: order.index(F[43]) + 1,
                    56: order.index(F[44]) + 1, 41: order.index(F[30]) + 1, 51: order.index(F[38]) + 1})
    for n in (43, 44, 45, 46, 47, 48, 49, 50):          # 新規友だちの追加など → FMTの同じページ
        mapping[n] = order.index(F[n + 2]) + 1
    for n, s in J.items():
        if any(s._element is o._element for o in order):
            renumber(s, mapping)

    lst = prs.slides._sldIdLst
    ids = {id(s._element): sid for s, sid in zip(prs.slides, list(lst))}
    keep = [ids[id(s._element)] for s in order]
    keep_set = set(id(x) for x in keep)
    for sid in list(lst):
        lst.remove(sid)
        if id(sid) not in keep_set:
            prs.part.drop_rel(sid.rId)
    for sid in keep:
        lst.append(sid)
    assert len(prs.slides) == 87
    prs.save(out)
    print("saved:", out)


if __name__ == "__main__":
    main(*sys.argv[1:5])
