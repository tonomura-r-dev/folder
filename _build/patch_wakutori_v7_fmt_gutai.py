# -*- coding: utf-8 -*-
"""ワクトリ v6（殿村さんPC保存版）→ v7：具体施策のイメージを、10/29の提案FMT（最強資料）の8枚の形にそろえる（2026-10-30 殿村さん）。

    python -m _build.patch_wakutori_v7_fmt_gutai <PC保存版v6.pptx> <FMT具体施策枠のみ.pptx> <出力.pptx>

  ・旧P27〜P33（ワクトリ独自の札・左にデモ画面・右に3つの枠）を外し、FMTの8枚
    （LINEの友だちへ／露出の強化／CVの最大化／CPOの改善（CV質）／休眠顧客の有効化／マネタイズ最大化／リピート施策／コミュニティ運営）
    を入れる。左＝入口→LINE→配信→ゴール、右＝目的／対象／中身／見る数字（形は `_build/patch_fmt_max_gutai.py` と同じ）
  ・施策の考え方（ステップ配信2枚・企画配信1枚＝旧P23〜P25）は、FMTの並びどおり具体施策の後ろへ
  ・P14の表の「（P◯）」を新しい番号に振り直す
  FMTのページの黄色の文字は白に（塾と同じ）。ほかのページは触らない。
"""
import sys

from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR

from _build.build_juku_fmt1029 import fmt_fix
from _build.patch_fmt_max_gutai import RIGHT, RIGHT_Y, flow, IN, NAVY, INK
from _build.patch_tanaka_v5 import copy_slide
from _build.pptx_parts import rich

GUTAI = [  # 札, [入口, LINE, 配信, ゴール], (目的, 対象, 中身, 見る数字)  「|」で改行
    ("LINEの友だちへ",
     [("サイト来訪・広告", "サイトのLINEボタン／ポップアップ／|応募完了画面（サンクスLINE誘導）"),
      "友だち追加", "あいさつメッセージ・適正バイト診断", "友だちとしてつながり続ける"],
     ("サイトに来た方・応募した方を、|LINEの友だちとして残す",
      "サイトに来た方（月 約8.6万人）／|応募せずに離れる方／応募した方",
      "求人ページ・トップのLINEボタン／|離れるときのポップアップ／|応募完了画面からのLINE誘導",
      "友だち追加数／入口ごとの追加数")),
    ("露出の強化",
     ["友だち全体", ("リッチメニュー", "開始月別の求人・季節の特集・相談の入口"),
      ("企画配信", "季節の特集（冬はスキー場、夏は海）"), "求人ページ・特集ページへの来訪"],
     ("季節の求人を広く知らせ、|求人を見る機会を増やす",
      "友だち全体（反応の少ない方も含む）",
      "季節に合わせた特集の企画配信／|リッチメニューに開始月別の求人／|特集の入口は季節ごとに入れ替え",
      "配信到達数／開封率／|求人ページへのクリック数")),
    ("CVの最大化",
     ["友だち追加から15日以内の方", ("Day0 適正バイト診断", "エントリー状況・開始時期・エリア"),
      ("ステップ配信（Day0〜15）", "働き方→お金の疑問→会社の選び方→ENTRY"), "応募（ENTRY）"],
     ("友だち追加の直後の反応が高い時期に、|応募（ENTRY）につなげる",
      "友だち追加から15日以内の方",
      "前後検索の順に疑問に答える5通／|診断の答えで合う求人を出し分け／|応募したら配信を止める",
      "応募数／各通のクリック率")),
    ("CPOの改善（CV質）",
     [("応募（ENTRY）の完了画面", "サンクスLINE誘導"), "友だち追加・応募済みのタグ付け",
      ("受付と次の流れのご案内", "担当のおすすめ求人・専属エージェントのご案内"), "お仕事開始"],
     ("応募した方を、お仕事開始まで|LINEで追う",
      "応募（ENTRY）した方",
      "応募の受付と次の流れのご案内／|担当のおすすめ求人／|寮・持ち物など不安に答える配信",
      "応募した方の友だち追加数／|応募からお仕事開始の割合")),
    ("休眠顧客の有効化",
     [("一定期間、反応のない友だち", "開封・クリックなし"), "タグで反応のない方を絞り込む",
      ("次のシーズンの企画配信", "冬はスキー場、夏は海・開始月別の新着"), "再び応募（ENTRY）へ"],
     ("反応が止まった友だちに、|次のシーズンの求人をもう一度届ける",
      "追加から15日を過ぎ、反応のない方／|応募まで進まなかった方",
      "反応の有無でタグを分けて配信／|シーズンの前に開始月別の求人を案内／|アンケートで開始時期を聞き直す",
      "再開封率／再クリック率／|休眠からの応募数")),
    ("マネタイズ最大化",
     ["診断に答えた方・求人を見ている方", ("診断タグ＋行動タグ", "開始時期・エリア／押した求人・関心（お金・寮）"),
      ("合う求人だけの配信", "開始時期・エリア別の特集"), "応募（ENTRY）"],
     ("検討度の高い方を取りこぼさず、|合う求人で応募まで進める",
      "開始時期・エリアに答えた方／|求人を押したが応募していない方",
      "開始時期×エリアで合う求人だけ配信／|押した関心（お金・寮）に答える配信／|専属エージェントへの相談の入口",
      "応募数／クリック率／ブロック率")),
    ("リピート施策",
     ["お仕事を終えた方", "勤務の実績・時期・エリアでタグ付け",
      ("次のシーズンのお仕事の案内", "冬→夏、夏→冬の開始月別の求人"), "次のお仕事への応募（ENTRY）"],
     ("お仕事を終えた方に、|次のシーズンのお仕事を届ける",
      "勤務の実績を確認できた方",
      "お仕事の終わる時期を起点にした配信／|前回と同じエリア・違うエリアの求人／|開始月別の新着の先行案内",
      "再応募数／リピート率")),
    ("コミュニティ運営",
     [("お仕事をした方", "くり返し働いている方"), ("経験者のタグ付け", "リッチメニューに体験談・紹介の入口"),
      ("体験談・友だちへの紹介のご案内", "LINEの転送で友だちに紹介"), "継続利用・紹介"],
     ("お仕事をした方とのつながりを強くし、|次のお仕事と紹介につなげる",
      "お仕事をした方／|くり返し働いている方",
      "経験者の体験談を届ける・集める／|LINEの転送で友だちに紹介／|紹介の特典の案内（中身は貴社で決める）",
      "紹介からの友だち追加数／|紹介からの応募数")),
]

P14_FIX = [("ステップ配信（P30）", "ステップ配信（P26）"), ("（P24・P30）", "（P26・P33）"),
           ("（P29・P31）", "（P26・P29）"), ("（P32）", "（P25）"), ("相談の入口（P24・P32）", "相談の入口（P33・P25）")]


def fill(s, page):
    tag, steps, right = page
    assert any(sh.has_text_frame and sh.text_frame.text == tag for sh in s.shapes), tag
    flow(s, steps)
    for (y, label), text in zip(zip(RIGHT_Y, RIGHT), right):
        rich(s, 6.97 * IN, (y + 0.06) * IN, 3.46 * IN, 1.01 * IN,
             [[(label, 13, True, NAVY)]] + [[(line, 12, False, INK)] for line in text.split("|")], anchor=MSO_ANCHOR.TOP)


def main(src, fmt_empty, out):
    prs = Presentation(src)
    emp = Presentation(fmt_empty)
    assert len(prs.slides) == 55
    W = {i + 1: s for i, s in enumerate(prs.slides)}
    assert W[26].shapes[0].has_text_frame and "アジェンダ" in "".join(sh.text_frame.text for sh in W[26].shapes if sh.has_text_frame)
    for n in range(27, 34):
        assert any(sh.has_text_frame and sh.text_frame.text == "具体施策のイメージ" for sh in W[n].shapes), n
    E = []
    for n, page in zip(range(18, 26), GUTAI):
        s = copy_slide(prs, emp.slides[n - 1])
        fmt_fix(s)
        fill(s, page)
        E.append(s)

    order = [W[n] for n in range(1, 23)] + [W[26]] + E + [W[23], W[24], W[25]] + [W[n] for n in range(34, 56)]
    assert len(order) == 56
    s14 = W[14]
    done = set()
    for t in s14._element.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}t"):
        for a, b in P14_FIX:
            if t.text == a:
                t.text = b
                done.add(a)
    assert done == {a for a, _ in P14_FIX}, done

    lst = prs.slides._sldIdLst
    ids = {id(s._element): sid for s, sid in zip(prs.slides, list(lst))}
    keep = [ids[id(s._element)] for s in order]
    keep_set = set(id(x) for x in keep)
    for sid in list(lst):
        lst.remove(sid)
        if id(sid) not in keep_set:
            prs.part.drop_rel(sid.rId)
    for sid in keep:
        lst.append(sid)
    prs.save(out)
    print("saved:", out, len(prs.slides))


if __name__ == "__main__":
    main(*sys.argv[1:4])
