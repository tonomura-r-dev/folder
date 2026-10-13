# -*- coding: utf-8 -*-
"""株式会社セキュリティー京葉 LINE採用追客のご提案（6枚・4:3）修正版。2026-10-13 殿村さんの修正指示。
- 通知メッセージは使わない。入口は過去の応募者＝SMS・メール等で友だち追加をご案内
- サンクスLINE誘導は押さない（殿村さん）。2枚目の注記に「連携できる場合はご案内可能」と1行だけ
- 友だち追加後はLステップのタグ（面接未参加・辞退・再応募希望・面接予約済み・希望勤務地・勤務条件）とセグメント配信で追客
- KPI＝CV①応募数（再応募を含む）・CV②面接予約数（目標数値・CVRは置かない）
- 費用＝初期構築費100,000円・月額費用300,000円（LINE公式アカウント・Lステップの利用料を含む）
- 画面全体を使う：大きい数字・流れ図・タイムライン・タグ図
器はDYMの4:3マスターの既存PPTX（サンクスLINE Appendix ver1.2）。
  python3 _build/build_keiyo_6p_v2.py <Appendix ver1.2.pptx> [出力.pptx]
"""
import copy
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Cm

sys.path.insert(0, str(Path(__file__).resolve().parent))
from pptx_parts import (CARD, GRAY, GREEN, INK, LGRAY, NAVY, WHITE, arrow_line, find,  # noqa: E402
                        hline, keep_header_only, rect, rich, set_paras, shape, vline)

ROOT = Path(__file__).resolve().parent.parent
SRC = Path(sys.argv[1])
OUT = Path(sys.argv[2]) if len(sys.argv) > 2 else ROOT / "20261013_セキュリティー京葉御中_LINE採用追客のご提案.pptx"

SLIDES = [
    ("課題｜一度外れた応募者が、そのまま離れている", "警備を含む保安職業の求人倍率は6.41倍。\n一度応募した方との接点を維持・再構築する導線を設計"),
    ("方針｜過去の応募者をLINEでつなぎ、追い続ける", "面接に来なかった方・辞退した方をLINEでつなぎ、応募・面接予約まで追う"),
    ("ステップ配信｜14日間で応募・面接予約を目指す", "友だち追加から14日間、働き方を順番に届け、応募・面接予約につなげる"),
    ("Lステップ｜応募者の状況ごとに配信を変える", "タグで応募者の状況を管理し、状況・希望に合わせた配信で面接予約までつなげる"),
    ("進め方｜アカウント開設から運用・改善まで", "LINE公式アカウントの開設から、運用・改善までまとめて支援"),
    ("費用", "初期構築費10万円・月額費用30万円で、開設から運用・改善まで支援"),
]

X0, X1 = 1.79, 25.76
XI = X0 + 0.35            # 本文の左端（見出しの縦棒の右）
WI = X1 - XI              # 本文の幅
Y0 = 4.4
LBLUE = "8E9BC4"          # 比較用の薄い紺


def text(s, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    return rich(s, x, y, w, h, runs, align=align, anchor=anchor)


def head(s, y, txt):
    rect(s, X0, y + 0.14, 0.15, 0.52, NAVY)
    text(s, XI, y, WI, 0.8, [[(txt, 15, True, NAVY)]])
    return y + 0.9


def box(s, x, y, w, h, paras, fill=WHITE, line=NAVY, lw=1.0, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
        rounded=True, margins=(0.2, 0.1, 0.2, 0.1)):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    return shape(s, kind, x, y, w, h, paras, fill=fill, line=line, lw=lw, align=align, anchor=anchor,
                 adj=0.08 if rounded else None, margins=margins)


def P(t, size, bold=False, color=INK, sa=2):
    return (t, size, bold, color, sa)


def arrow(s, x1, y1, x2, y2, color=NAVY, w=2.0):
    arrow_line(s, x1, y1, x2, y2, color, w=w)


prs = Presentation(str(SRC))
slides = list(prs.slides)
HDR = {"TextBox 1", "TextBox 2", "Connector 3"}
donor = next(s for s in slides if HDR <= {sh.name for sh in s.shapes})
for s in slides:
    names = {sh.name for sh in s.shapes}
    if "TextBox 1" in names and not HDR <= names:
        for nm in ("TextBox 2", "Connector 3"):
            if nm not in names:
                s.shapes._spTree.append(copy.deepcopy(find(donor, nm)._element))
keep = [i for i, s in enumerate(prs.slides) if HDR <= {sh.name for sh in s.shapes}][:6]
assert len(keep) == 6, keep
lst = prs.slides._sldIdLst
for i, el in reversed(list(enumerate(list(lst)))):
    if i not in keep:
        prs.part.drop_rel(el.rId)
        lst.remove(el)

for s, (title, lead) in zip(prs.slides, SLIDES):
    keep_header_only(s)
    t1, t2 = find(s, "TextBox 1"), find(s, "TextBox 2")
    t1.left, t1.top, t1.width, t1.height = Cm(1.52), Cm(0.38), Cm(18.1), Cm(0.94)
    t2.left, t2.top, t2.width, t2.height = Cm(1.79), Cm(2.27), Cm(25.31), Cm(0.96)
    set_paras(t1, [(title, 0)])
    set_paras(t2, [(ln, 0) for ln in lead.split("\n")])
    for t in (t1, t2):
        for r in t.text_frame._txBody.iter(qn("a:rPr")):
            r.set("sz", "1600")

s1, s2, s3, s4, s5, s6 = prs.slides

# ───────── 1. 課題 ─────────
y = head(s1, Y0, "警備員の採用環境")
text(s1, XI, y + 0.1, 9.0, 0.7, [[("保安職業の有効求人倍率", 14, False, INK)]])
text(s1, XI, y + 0.65, 9.0, 2.4, [[("6.41", 54, True, NAVY), ("倍", 26, True, NAVY)]])
# 横棒の比較
bx, bmax = 15.2, 9.6
for i, (lab, v, col) in enumerate([("保安職業", 6.41, NAVY), ("職業全体", 1.10, LBLUE)]):
    yy = y + 0.45 + i * 1.25
    text(s1, 11.6, yy, 3.5, 0.8, [[(lab, 14, True, INK)]], anchor=MSO_ANCHOR.MIDDLE)
    L = bmax * v / 6.41
    rect(s1, bx, yy + 0.08, L, 0.65, col)
    text(s1, bx + L + 0.15, yy, 3.0, 0.8, [[(f"{v:.2f}倍", 14, True, col if col == NAVY else INK)]],
         anchor=MSO_ANCHOR.MIDDLE)
text(s1, XI, y + 3.35, WI, 0.5,
     [[("出典：労働政策研究・研修機構「職業別にみた有効求人倍率、平均求人賃金」（2025年度）", 10.5, False, GRAY)]])
y = head(s1, y + 4.15, "選考の途中で起きていること")
items = ["応募", "面接のご案内", "面接に来ない・\n辞退", "連絡が途切れ、\nそのまま離れる"]
gap = 0.75
bw = (WI - gap * 3) / 4
for i, t in enumerate(items):
    last = i == 3
    x = XI + i * (bw + gap)
    box(s1, x, y + 0.1, bw, 2.0, [P(l, 15, last, WHITE if last else INK, 0) for l in t.split("\n")],
        fill=NAVY if last else WHITE)
    if not last:
        arrow(s1, x + bw + 0.1, y + 1.1, x + bw + gap - 0.1, y + 1.1)
y += 2.6
cw = (WI - 0.6) / 2
for i, (lab, body) in enumerate([("再アプローチ対象の特徴", "一度は警備の仕事に関心を示し、\n応募した方"),
                                 ("再アプローチの課題",
                                  "一人ずつの連絡は手間がかかり、\n次の募集を届け続ける仕組みがない")]):
    x = XI + i * (cw + 0.6)
    box(s1, x, y, cw, 2.9, None, fill=CARD, line=LGRAY)
    rect(s1, x, y, 0.18, 2.9, NAVY)
    text(s1, x + 0.5, y + 0.25, cw - 0.8, 0.7, [[(lab, 13, True, NAVY)]])
    text(s1, x + 0.5, y + 0.95, cw - 0.8, 1.8, [[(l, 15, True, INK)] for l in body.split("\n")])

# ───────── 2. 方針（過去の応募者 → 友だち追加 → Lステップ → CV） ─────────
head(s2, Y0, "追客の流れ")
cy = 9.9
ew, eh = 8.0, 6.4
ey = cy - eh / 2
box(s2, X0, ey, ew, eh, None, fill=WHITE, line=NAVY, lw=1.25)
rect(s2, X0, ey, ew, 1.0, NAVY, [P("過去の応募者", 15, True, WHITE, 0)])
text(s2, X0 + 0.45, ey + 1.35, ew - 0.9, 1.6, [[(l, 16, True, NAVY)] for l in ("面接に来なかった方・", "辞退した方")])
text(s2, X0 + 0.45, ey + 3.25, ew - 0.9, 2.8,
     [[(l, 14.5, False, INK)] for l in ("SMS・メール等で", "LINEの友だち追加を", "ご案内")])
fx, fw, fh = X0 + ew + 0.75, 3.5, 2.6
arrow(s2, X0 + ew + 0.08, cy, fx - 0.08, cy)
box(s2, fx, cy - fh / 2, fw, fh, [P("LINE", 15, True, WHITE, 0), P("友だち追加", 15, True, WHITE, 0)],
    fill=GREEN, line=GREEN)
lx, lw_ = fx + fw + 0.7, 5.8
arrow(s2, fx + fw + 0.08, cy, lx - 0.08, cy)
box(s2, lx, cy - 2.5, lw_, 5.0,
    [P("Lステップで追客", 15.5, True, NAVY, 10), P("・タグで状況を管理", 13.5, False, INK, 4),
     P("・状況に合わせて配信", 13.5, False, INK, 4), P("・LINEで面接予約", 13.5, False, INK, 0)],
    fill=WHITE, line=NAVY, lw=1.5, align=PP_ALIGN.LEFT, margins=(0.4, 0.2, 0.3, 0.2))
cx = lx + lw_ + 0.75
cvw = X1 - cx
# CV① → CV② の順番（応募の次に面接予約）。Lステップから CV① へ入り、CV① から CV② へ下りる
cv1y, cv2y, cvh = cy - 2.5, cy + 0.4, 2.1
arrow(s2, lx + lw_ + 0.08, cv1y + cvh / 2, cx - 0.08, cv1y + cvh / 2)
for yy, lab, main, sub in [(cv1y, "CV①", "応募数", "（再応募を含む）"), (cv2y, "CV②", "面接予約数", "")]:
    paras = [P(lab, 12, True, WHITE, 0), P(main, 18, True, WHITE, 0)] + ([P(sub, 11, False, WHITE, 0)] if sub else [])
    box(s2, cx, yy, cvw, cvh, paras, fill=NAVY, line=NAVY, margins=(0.1, 0.05, 0.1, 0.05))
arrow(s2, cx + cvw / 2, cv1y + cvh + 0.08, cx + cvw / 2, cv2y - 0.08)
text(s2, X0, 15.45, X1 - X0, 1.2,
     [[("※過去の応募者へのご案内は、連絡先の利用目的・ご連絡の許諾を確認のうえ実施", 11.5, False, GRAY)],
      [("※今後の応募者も、応募完了後にLINEへご案内可能（採用管理サービスと連携できる場合）", 11.5, False, GRAY)]])

# ───────── 3. ステップ配信（14日間のタイムライン → 14日以降は別施策の企画配信） ─────────
y = head(s3, Y0, "14日間のステップ配信（例）")
steps = [("登録直後", "お礼・\n働き方の紹介", "興味喚起", "「ご応募ありがとうございました。改めて働き方をご紹介します」"),
         ("3日後", "給与・\n勤務条件の紹介", "条件面の理解", "「1現場10,500円（2現場で19,500円）。直行直帰OK・夜勤なし」"),
         ("7日後", "働く方の声・\n安心材料", "不安の解消", "「平均年齢64歳。84歳の方も活躍中。資格取得の支援もあります」"),
         ("10日後", "スポット勤務・\n希望条件の確認", "応募ハードルの\n低減", "「スポット勤務もあります。ご希望の勤務地・日数を教えてください」"),
         ("14日後", "再応募・\n面接予約の案内", "CV①・CV②への\n誘導", "「ご希望に合う現場・日程をご案内します。応募・面接予約はこちら」")]
g = 0.35
cw = (WI - g * 4) / 5
ly = y + 0.35
ch = 5.75
hline(s3, XI + cw / 2, XI + 4 * (cw + g) + cw / 2, ly, NAVY, 2.5)
for i, (tm, ttl, aim, msg) in enumerate(steps):
    x = XI + i * (cw + g)
    last = i == 4
    tc, bc = (WHITE, WHITE) if last else (NAVY, INK)
    shape(s3, MSO_SHAPE.OVAL, x + cw / 2 - 0.3, ly - 0.3, 0.6, 0.6, None, fill=GREEN if last else NAVY, line=None)
    box(s3, x, ly + 0.55, cw, 0.85, [P(tm, 14, True, WHITE, 0)], fill=GREEN if last else NAVY,
        line=GREEN if last else NAVY, margins=(0.05, 0.02, 0.05, 0.02))
    box(s3, x, ly + 1.6, cw, ch, None, fill=NAVY if last else WHITE, line=NAVY if last else LGRAY,
        lw=1.5 if last else 1.0)
    text(s3, x + 0.2, ly + 1.75, cw - 0.4, 1.3, [[(l, 13.5, True, tc)] for l in ttl.split("\n")])
    text(s3, x + 0.2, ly + 3.1, cw - 0.4, 1.45, [[("狙い", 10.5, True, WHITE if last else GRAY)]]
         + [[(l, 12, True, WHITE if last else "0B7A3B")] for l in aim.split("\n")])
    text(s3, x + 0.2, ly + 4.9, cw - 0.4, 2.4, [[(msg, 12, False, bc)]])
# 14日以降＝別施策（ステップ配信とは分けて見せる）
by = ly + 1.6 + ch + 0.35
box(s3, XI, by, WI, 1.25, None, fill=WHITE, line=GRAY, lw=1.0)
box(s3, XI + 0.25, by + 0.2, 5.4, 0.85, [P("14日以降（企画配信）", 13, True, WHITE, 0)], fill=GRAY, line=GRAY,
    margins=(0.05, 0.02, 0.05, 0.02))
text(s3, XI + 5.95, by, WI - 6.2, 1.25,
     [[("新しい現場・募集が出るたびに、希望勤務地・勤務条件のタグごとに配信", 13.5, False, INK)]],
     anchor=MSO_ANCHOR.MIDDLE)
y = head(s3, by + 1.45, "そのほかに伝えられること")
cw2 = (WI - 0.6) / 2
for i, (lab, body) in enumerate([("待遇", "ガソリン代・駐車場代の全額支給\n空調服・ネッククーラー支給"),
                                 ("安心材料", "資格取得の支援\n月15日以上の勤務で社会保険に加入")]):
    x = XI + i * (cw2 + 0.6)
    box(s3, x, y + 0.05, cw2, 1.9, None, fill=CARD, line=LGRAY)
    rect(s3, x, y + 0.05, 0.18, 1.9, NAVY)
    text(s3, x + 0.5, y + 0.15, 2.4, 1.7, [[(lab, 14, True, NAVY)]], anchor=MSO_ANCHOR.MIDDLE)
    text(s3, x + 2.8, y + 0.15, cw2 - 2.95, 1.7, [[(l, 13, False, INK)] for l in body.split("\n")],
         anchor=MSO_ANCHOR.MIDDLE)

# ───────── 4. Lステップ（タグ → 配信） ─────────
y = head(s4, Y0, "タグごとに配信を変える")
tags = [("面接未参加", "面接日程の再案内と、LINEでの面接予約"),
        ("辞退", "時期をあけて、新しい募集・現場のお知らせ"),
        ("再応募希望", "LINEの中で面接日を選んで予約"),
        ("面接予約済み", "面接前日のリマインド・会場までのアクセス"),
        ("希望勤務地・勤務条件", "希望に合う現場・働き方（中途・アルバイト・スポット）のご案内")]
tw, rh, rg = 6.4, 1.25, 0.28
text(s4, XI, y, tw, 0.6, [[("タグ（応募者の状況）", 12, True, GRAY)]], align=PP_ALIGN.CENTER)
text(s4, XI + tw + 0.9, y, 10.0, 0.6, [[("配信の内容", 12, True, GRAY)]])
y += 0.6
for i, (tg, dl) in enumerate(tags):
    yy = y + i * (rh + rg)
    box(s4, XI, yy, tw, rh, [P(tg, 14, True, "0B7A3B", 0)], fill="E8F8EE", line=GREEN, lw=1.5)
    arrow(s4, XI + tw + 0.12, yy + rh / 2, XI + tw + 0.78, yy + rh / 2)
    box(s4, XI + tw + 0.9, yy, WI - tw - 0.9, rh, [P(dl, 14, False, INK, 0)], fill=WHITE, line=LGRAY,
        align=PP_ALIGN.LEFT, margins=(0.4, 0.05, 0.3, 0.05))
y = head(s4, y + 5 * (rh + rg) + 0.25, "Lステップでできること")
feats = [("タグ管理", "応募者の状況を記録"), ("セグメント配信", "状況・希望ごとに送り分け"),
         ("面接予約", "LINEの中で日程を選択"), ("反応の確認", "押されたリンクを確認")]
fw = (WI - 0.3 * 3) / 4
for i, (a, b) in enumerate(feats):
    box(s4, XI + i * (fw + 0.3), y + 0.05, fw, 1.75, [P(a, 14, True, NAVY, 4), P(b, 12, False, INK, 0)],
        fill=CARD, line=LGRAY)
text(s4, XI, y + 2.0, WI, 0.6,
     [[("※採用管理サービスとの連携方法・タグ付けの具体的な方法は、導入時にご相談のうえ決定", 11.5, False, GRAY)]])

# ───────── 5. 進め方 ─────────
y = head(s5, Y0, "導入の流れ")
stp = [("LINE公式\nアカウント開設", "LINE公式アカウントの開設・初期設定"),
       ("Lステップ\n構築", "タグ・セグメント・面接予約の設計と設定"),
       ("友だち獲得の\n導線整備", "過去の応募者へ\nSMS・メール等で\nLINEの友だち\n追加をご案内"),
       ("配信設定", "ステップ配信・募集のお知らせの作成"),
       ("運用・改善", "応募数・面接予約数を見て、配信の内容・タイミングを見直す")]
g5 = 0.6
sw = (WI - g5 * 4) / 5
for i, (h_, d) in enumerate(stp):
    x = XI + i * (sw + g5)
    last = i == 4
    shape(s5, MSO_SHAPE.OVAL, x + sw / 2 - 0.42, y + 0.05, 0.84, 0.84, [P(str(i + 1), 14, True, WHITE, 0)],
          fill=GREEN if last else NAVY, line=None, margins=(0, 0, 0, 0))
    box(s5, x, y + 1.1, sw, 1.9, [P(l, 13.5, True, WHITE, 0) for l in h_.split("\n")],
        fill=NAVY, line=NAVY, margins=(0.08, 0.05, 0.08, 0.05))
    if not last:
        arrow(s5, x + sw + 0.08, y + 2.05, x + sw + g5 - 0.08, y + 2.05)
    box(s5, x, y + 3.2, sw, 4.2, None, fill=WHITE, line=LGRAY)
    text(s5, x + 0.12, y + 3.4, sw - 0.24, 3.8, [[(l, 12, False, INK)] for l in d.split("\n")])
by = y + 7.75
x4r = XI + 3 * (sw + g5) + sw
hline(s5, XI, x4r, by, NAVY, 1.5)
vline(s5, XI, by - 0.25, by, NAVY, 1.5)
vline(s5, x4r, by - 0.25, by, NAVY, 1.5)
text(s5, XI, by + 0.1, x4r - XI, 0.7, [[("初期構築", 15, True, NAVY)]], align=PP_ALIGN.CENTER)
x5 = XI + 4 * (sw + g5)
hline(s5, x5, x5 + sw, by, GREEN, 1.5)
vline(s5, x5, by - 0.25, by, GREEN, 1.5)
vline(s5, x5 + sw, by - 0.25, by, GREEN, 1.5)
text(s5, x5, by + 0.1, sw, 0.7, [[("月額の運用", 15, True, "0B7A3B")]], align=PP_ALIGN.CENTER)
y = head(s5, by + 1.15, "改善で見る指標")
kg = 1.4                  # CV① → CV② の矢印を置く間隔
kw = (WI - kg) / 2
arrow(s5, XI + kw + 0.25, y + 0.8, XI + kw + kg - 0.25, y + 0.8, w=2.5)
for i, (lab, main) in enumerate([("CV①", "応募数（再応募を含む）"), ("CV②", "面接予約数")]):
    x = XI + i * (kw + kg)
    box(s5, x, y + 0.05, kw, 1.5, None, fill=CARD, line=LGRAY)
    box(s5, x + 0.3, y + 0.3, 1.8, 1.0, [P(lab, 14, True, WHITE, 0)], fill=NAVY, line=NAVY, margins=(0, 0, 0, 0))
    text(s5, x + 2.4, y + 0.05, kw - 2.6, 1.5, [[(main, 16, True, NAVY)]], anchor=MSO_ANCHOR.MIDDLE)

# ───────── 6. 費用 ─────────
y = head(s6, Y0, "内訳")
fw6 = (WI - 0.8) / 2
blocks = [("初期構築費", "100,000円", NAVY,
           ["LINE公式アカウントの開設", "Lステップの構築（タグ・配信の設計）", "友だち獲得の導線整備", "ステップ配信の作成"]),
          ("月額費用", "300,000円", NAVY,
           ["ステップ配信・募集のお知らせの運用", "タグ付け・リストの管理", "応募数・面接予約数を見た改善",
            "LINE公式アカウント・Lステップの利用料"])]
for i, (lab, amt, col, inc) in enumerate(blocks):
    x = XI + i * (fw6 + 0.8)
    box(s6, x, y + 0.05, fw6, 5.0, None, fill=CARD, line=NAVY, lw=1.5)
    text(s6, x, y + 0.45, fw6, 0.9, [[(lab, 18, True, NAVY)]], align=PP_ALIGN.CENTER)
    text(s6, x, y + 1.5, fw6, 3.0, [[(amt, 46, True, NAVY)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s6, x + 0.3, y + 5.5, fw6, 0.7, [[("含まれるもの", 14, True, GRAY)]])
    paras = []
    for it in inc:
        bold = "利用料" in it
        paras.append([("・" + it, 14, bold, NAVY if bold else INK)])
    text(s6, x + 0.3, y + 6.25, fw6 - 0.3, 4.5, paras)

# 字体をメイリオに統一（latin → ea → cs の順）
AFTER = {qn(x) for x in ("a:sym", "a:hlinkClick", "a:hlinkMouseOver", "a:rtl", "a:extLst")}
for s in prs.slides:
    for r in s.shapes._spTree.iter(qn("a:r")):
        if r.find(qn("a:rPr")) is None:
            r.insert(0, r.makeelement(qn("a:rPr"), {"lang": "ja-JP"}))
    for tag in ("a:rPr", "a:endParaRPr", "a:defRPr"):
        for rpr in s.shapes._spTree.iter(qn(tag)):
            for ft in ("a:latin", "a:ea", "a:cs"):
                for e in rpr.findall(qn(ft)):
                    rpr.remove(e)
            nxt = next((c for c in rpr if c.tag in AFTER), None)
            for ft in ("a:latin", "a:ea", "a:cs"):
                e = rpr.makeelement(qn(ft), {"typeface": "メイリオ"})
                if nxt is not None:
                    nxt.addprevious(e)
                else:
                    rpr.append(e)

prs.save(str(OUT))
print("saved:", OUT.name, len(prs.slides), "枚")
