# -*- coding: utf-8 -*-
"""横浜シミズ資料（33枚）：友だちの入口を「応募前の友だち追加＋サンクスLINEは補助」に統一し、勤務済みの判定を直す（2026-10-29 殿村さん）。

    python -m _build.patch_yokohama_shimizu_v33_fix <資料.pptx> [出力.pptx]

  ・17 競合分析②：横浜シミズ欄の「応募・登録の導線」＝3つの入口（LINEボタン・離脱防止／応募後はサンクスLINE誘導）
  ・22 サンクスLINEの実績：横浜シミズでの使い方を「応募前に友だちにならなかった方の受け皿（補助の入口）」に
  ・32 今後の展開：STEP1の設置・ご提案のまとめ①を、応募前の友だち追加を含む形に
    （サイトのLINEボタンは貴社のサイトなので、DYMの構築の項目には書かない）
  ・24 施策③：「勤務の案内を送った方」に勤務タグを付けない。勤務実績を確認できた方にだけ「勤務済み」
  ・26 配信イメージ：「前回のご勤務ありがとうございました」の文面は、勤務実績を確認できた方にだけ送る
  出力を省くと上書き。
"""
import sys

from pptx import Presentation
from pptx.util import Inches

from _build.pptx_parts import set_para
from _build.patch_yokohama_shimizu_chapters import head


def by_id(s, sid):
    return next(sh for sh in s.shapes if sh.shape_id == sid)


def cell_lines(cell, lines):
    sh = type("T", (), {"text_frame": cell.text_frame})
    for k, line in enumerate(lines):
        set_para(sh, k, line)
    ps = cell.text_frame.paragraphs
    for p in ps[len(lines):]:
        p._p.getparent().remove(p._p)


def table(s):
    return next(sh for sh in s.shapes if sh.has_table).table


def main(src, out):
    prs = Presentation(src)
    sl = list(prs.slides)
    assert len(sl) == 33
    s17, s22, s24, s26, s32 = (sl[n - 1] for n in (17, 22, 24, 26, 32))
    assert head(s17).text_frame.text.startswith("競合分析②")
    assert head(s22).text_frame.text.startswith("施策①｜サンクスLINE誘導の導入実績")
    assert head(s24).text_frame.text.startswith("施策③")
    assert head(s26).text_frame.text.startswith("配信イメージ")
    assert head(s32).text_frame.text.startswith("今後の展開")

    t = table(s17)
    assert t.cell(2, 4).text.startswith("応募完了から")
    cell_lines(t.cell(2, 4), ["サイトのLINEボタン・離脱防止で、", "応募の前からLINEへ", "（応募後はサンクスLINE誘導）"])

    set_para(by_id(s22, 4), 0, "人材業界での導入実績（DYM）と、横浜シミズでの使い方（補助の入口）")
    lab = by_id(s22, 63)
    set_para(lab, 0, "横浜シミズでの使い方｜応募前に友だちにならなかった方の受け皿")
    lab.width = Inches(7.0)
    box = by_id(s22, 64)
    set_para(box, 0, "LINE未追加で")
    set_para(box, 1, "応募完了")
    set_para(by_id(s22, 65), 0, "LINE友だち追加")
    set_para(by_id(s22, 65), 1, "（サンクスLINE）")
    set_para(by_id(s22, 73), 0, "応募の直後から、ほかの友だちと同じLINEで、面接・勤務の案内と次の募集をお届けします。")

    t = table(s24)
    assert t.cell(5, 0).text == "勤務の有無"
    cell_lines(t.cell(5, 0), ["勤務実績"])
    cell_lines(t.cell(5, 1), ["勤務済み／まだ勤務なし"])
    cell_lines(t.cell(5, 2), ["勤務実績を確認", "できた方だけ"])
    set_para(by_id(s24, 34), 0, "勤務実績のある方")
    note = by_id(s24, 23)
    set_para(note, 2, "※勤務の案内を送っただけでは「勤務済み」にしません。")
    set_para(note, 3, "※勤務の履歴・希望条件をいまどこで管理しているかは確認事項です。")

    set_para(note, 1, "※取れる情報から、段階的にタグを付けます（システム連携は前提にしません）。")

    set_para(by_id(s26, 30), 0, "対象：横浜・土日希望・勤務実績のある方")
    msg = by_id(s26, 28)                                      # 言葉の途中で切れないよう、文の切れ目で段落を分ける
    assert msg.text_frame.paragraphs[3].text.startswith("前回のご勤務")
    for k, line in enumerate(["前回のご勤務、", "ありがとうございました。", "次のイベントもお待ちしています。"]):
        set_para(msg, 3 + k, line)
    pre = by_id(s26, 24)
    assert pre.text_frame.paragraphs[5].text.startswith("遅れる場合")
    for k, line in enumerate(["遅れる場合・", "体調がすぐれないときは、", "このLINEでご連絡ください。"]):
        set_para(pre, 5 + k, line)

    set_para(by_id(s32, 23), 0, "・離脱防止・サンクスLINE誘導の設置")
    set_para(by_id(s32, 28), 1, "① 友だち獲得：応募の前からLINEへ（応募後はサンクスLINE誘導）")
    prs.save(out)
    print("saved:", out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else sys.argv[1])
