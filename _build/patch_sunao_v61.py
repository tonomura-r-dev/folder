# -*- coding: utf-8 -*-
"""すなお ver6.0（殿村さんのPC保存版）→ ver6.1：契約期間・配信頻度・ツールの選択制・費用・SIMの整合（2026-10-20 殿村さんの修正指示）。

    python -m _build.patch_sunao_v61 <ver6.0.pptx> <ver6.1.pptx>

  ・契約期間は6か月（「通年」「6か月〜」をやめる）。平常期の企画配信は月1〜2回（セミナー開催月は月4〜5回のまま）
  ・ツールは両方とも選択制。離脱防止ポップアップ＝導入推奨、サンクスLINE誘導＝追加オプション
  ・費用の主役は推奨構成（ポップアップあり）：初期145,000円〜／平常期135,000円〜／開催月335,000円程度
    サンクスLINE誘導を追加：初期＋150,000円・月額＋30,000円〜（合計 295,000円〜／165,000円〜／365,000円程度）
    コンサル費（平常期100,000円・開催月300,000円程度）は変えない。プラン②の平常期は「50,000〜100,000円」→「100,000円」
  ・SIMの90人・11件・4件は変えない。サンクスLINE誘導ありの想定だと明記し、推奨構成では友だち数が少なくなる想定と注記
  ・指定のないページは触らない。スライドの枚数・並びも変えない（62枚）
  「初期構築は約1か月」は入れていない（殿村さんの以前の指示「初期構築の期間を資料に書かない」のため。入れるなら指示待ち）。
"""
import copy
import sys

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from _build.arauchi_v4.lib import add_box, add_table
from _build.arauchi_v10.build import runs
from _build.build_sunao_v1 import all_meiryo
from _build.build_sunao_v4 import set_cell
from _build.patch_sunao_v41 import set_text
from _build.patch_sunao_v53_kaizen import BAND, GRAY, NAVY, TAKE, TEXT, cell, head, label, lines, note, set_one, title_of

LINE_W = "FFFFFF"


def by_id(s, sid):
    return next(sh for sh in s.shapes if sh.shape_id == sid)


def by_name(s, name):
    hit = [sh for sh in s.shapes if sh.name == name]
    assert len(hit) == 1, (name, len(hit))
    return hit[0]


def by_text(s, text):
    hit = [sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip() == text]
    assert len(hit) == 1, (text, len(hit))
    return hit[0]


def drop(sh):
    sh._element.getparent().remove(sh._element)


def set_lines(sh, texts):
    """1段落目の書式（最初の文字列）のまま、段落を texts の行に置き換える"""
    tx = sh.text_frame._txBody
    ps = tx.findall(qn("a:p"))
    tmpl = copy.deepcopy(ps[0])
    for p in ps:
        tx.remove(p)
    for t in texts:
        p = copy.deepcopy(tmpl)
        rs = p.findall(qn("a:r"))
        for e in rs[1:] + p.findall(qn("a:br")):
            p.remove(e)
        rs[0].find(qn("a:t")).text = t
        tx.append(p)


def table_of(s):
    return next(sh for sh in s.shapes if sh.has_table).table


def pill(s, x, y, w, text):
    add_box(s, x, y, w, 0.53, [runs((text, dict(size=13, bold=True, color=NAVY)))], fill=LINE_W, line=NAVY, line_w=1.5,
            align="ctr", anchor="ctr", ins=(0.05, 0.02, 0.05, 0.02), name="ツールの位置づけ")


# ---------------------------------------------------------------- 各ページ
def s01_cover(s):
    set_lines(by_id(s, 6), ["広告で集めた見込み客をLINEで育成し、", "次のセミナー参加・サロン入会へ"])


def s03_doshen(s):
    set_lines(by_id(s, 9), ["広告で新しい見込み客を集め、LINEで継続的に育成し、", "次回イベント・サロン入会へつなげる動線を、8フェーズで整理する。"])


def s04_genjo(s):
    assert set_text(by_id(s, 43), "次のイベントまで、接点を持ち続ける仕組みを強化できる。", "接点を次回イベントまで継続できる仕組みを強化する。")


def s13_why(s):
    set_one(by_id(s, 4), "広告で新規見込み客を集め、LINEで関係を積み上げる")
    old = by_name(s, "メリット")
    x, y, w, h = (old.left, old.top, old.width, old.height)
    drop(old)
    strong = dict(size=14, bold=True, color=NAVY)
    plain = dict(size=14, color=TEXT)
    box = add_box(s, 0, 0, 1, 1,
                  [runs(("・友だち追加した方に、継続して情報を届けられる（次回以降のイベントも案内できる）", strong)),
                   runs(("・イベントのない時期も、見込み客との接点を持てる", strong)),
                   runs(("・通知を通じて、見込み客との接点をつくれる", plain)),
                   runs(("・友だちとして、見込み客を蓄積できる", plain)),
                   runs(("・中長期的に、広告への依存度を下げられる可能性がある", plain))],
                  fill=TAKE, line=NAVY, anchor="ctr", ins=(0.3, 0.06, 0.2, 0.06), line_spacing=130, name="メリット")
    box.left, box.top, box.width, box.height = x, y, w, h


def s16_yoken(s):
    for sh in [x for x in s.shapes if x.has_table]:
        drop(sh)
    fill = BAND[1]
    rows = [
        [label("プロモーション目的", "FFFFFF"),
         cell("広告で獲得した見込み客をLINE公式アカウントに蓄積し、継続的な情報配信で、次回セミナーへの参加とオンラインサロンへの入会を増やす",
              size=14, bold=True, color=NAVY)],
        [label("中長期の目的", fill), cell("入会後の継続利用・紹介にもLINEを活用し、LTVの最大化をめざす", size=14, fill=fill)],
        [label("契約期間", "FFFFFF"), cell("6か月（平常期とセミナー開催月で、配信内容と工数が変わる）", size=14)],
        [label("配信目標", fill),
         cell(lines("友だち獲得を増やし、ブロック率を抑えながら、配信可能な人数を積み上げる",
                    "（6か月の想定は、概算お見積りに記載）", size=14), fill=fill)],
        [label("手法（ご提案）", "FFFFFF"),
         cell(lines("① 離脱防止ポップアップ（導入推奨）で、LPを離れる方を友だちに",
                    "　 予約完了後は、サンクスLINE誘導（追加オプション）または通常のご案内",
                    "② 友だち追加直後・予約した方・参加後の段階ごとに、案内を変える",
                    "③ 平常期 月1〜2回、開催月 月4〜5回の企画配信で、つながりを保つ", size=14))],
        [label("＋αの施策", fill),
         cell(lines("過去のセミナー予約者・参加者へのご案内（通知メッセージ）",
                    "サロンの申込み完了ページからのLINE誘導（サンクスLINE誘導）", size=14)
              + lines("※通知メッセージは、予約時の電話番号がある場合", size=11, color=GRAY), fill=fill)],
    ]
    for r in rows:
        r[0]["size"] = 14
        r[0]["align"] = "ctr"
    heights = [0.8, 0.45, 0.45, 0.75, 1.4, 1.0]
    top = 0.7 + (7.0 - 0.7 - sum(heights)) / 3
    add_table(s, 0.46, top, [2.2, 7.7], heights, rows, name="要件定義")


def s17_hoshin(s):
    set_one(by_id(s, 49), "一度集めた見込み客を蓄積し、継続的にアプローチできる仕組みを強化する")
    t = table_of(s)
    set_cell(t.cell(1, 2), "・離脱防止ポップアップ（推奨）で友だちに\n・予約者はサンクスLINE誘導（追加オプション）")
    set_cell(t.cell(2, 2), "・企画配信：平常期 月1〜2回／開催月 月4〜5回\n・予約した方へ、前日・当日のリマインド")
    for r in (1, 2, 3):
        for rpr in t.cell(r, 2)._tc.iter(qn("a:rPr")):
            rpr.set("sz", "1200")


def s18_flow(s):
    set_lines(by_name(s, "導線の説明 LINE友だち追加"), ["離脱防止ポップアップ（推奨）", "予約完了後のご案内（サンクスLINE誘導は追加オプション）"])
    set_lines(by_name(s, "導線の説明 継続的な情報配信"), ["あいさつ・見どころ動画・参加特典", "企画配信：平常期 月1〜2回／開催月 月4〜5回"])
    for sh in s.shapes:                                              # 説明の箱を少し高くし、中長期の帯を下げる（6つとも同じ高さ）
        if sh.name.startswith("導線の説明"):
            sh.height = Inches(1.85)
    by_name(s, "中長期").top = Inches(5.2)


def s19_phase(s):
    t = table_of(s)
    rows = {t.cell(r, 0).text.strip()[:1]: r for r in range(1, len(t.rows))}
    set_cell(t.cell(rows["①"], 3), "離脱防止ポップアップ（推奨）")
    set_cell(t.cell(rows["⑥"], 3), "月1〜2回の配信で次回へ")


def s20_zentai(s):
    set_lines(by_id(s, 43), ["予約完了後の", "友だち追加"])
    set_one(by_id(s, 62), "・前日・当日のリマインド")
    by_id(s, 62).width = Inches(1.9)
    add_box(s, 6.9, 4.55, 3.24, 0.6,
            [runs(("※予約完了後の友だち追加は、サンクスLINE誘導（追加オプション）または通常の友だち追加のご案内で実施",
                   dict(size=10, color=GRAY)))], anchor="t", ins=(0.05, 0, 0, 0), name="予約後の友だち追加の注記")


def s21_track(s):
    box = by_id(s, 13)
    assert set_text(box, "予約していない方は平常期の配信で次のイベントへつなげる。予約・参加・入会の状況での出し分けは、連携の可否を確認のうえ設計。",
                    "予約していない方は平常期の配信で次のイベントへつなげる。流入経路だけでは予約・参加・入会の状況をすべて把握できないため、状況に応じた出し分けは、連携の仕様を確認のうえ設計。")
    box.height = Inches(1.35)


def s23_mitsumori(s):
    set_one(by_id(s, 339), "概算お見積り：推奨構成（ポップアップあり）での6か月の運用")
    set_one(by_id(s, 343), "¥145,000〜")
    set_lines(by_id(s, 344), ["開催月 ¥335,000程度", "平常期 ¥135,000〜"])
    set_lines(by_id(s, 350), ["開催月 4〜5回", "平常期 1〜2回"])
    set_one(by_id(s, 366), "導入（推奨）")
    set_one(by_id(s, 368), "追加オプション")
    set_one(by_id(s, 346), "6か月後（サンクスLINE誘導あり）")
    by_id(s, 346).text_frame.paragraphs[0].runs[0].font.size = 12 * 12700
    set_lines(by_name(s, "想定の注記"), [
        "※SIMによる想定値（実績ではない）。ポップアップ＋サンクスLINE誘導の場合",
        "※推奨構成（ポップアップのみ）では、予約完了後の友だち追加が減るため、友だち数・配信可能な人数は少なくなる想定",
        "※セミナー予約・サロン入会は、LINEの配信・リンクからの件数。友だち数ほどの差は見込まないが、同数とは限らない",
        "※配信可能な人数＝ブロックされずに届く友だちの数"])
    note_box = by_name(s, "想定の注記")
    note_box.height = Inches(2.5)
    add_box(s, 0.29, 6.62, 6.53, 0.45,
            [runs(("追加オプション｜サンクスLINE誘導：初期＋¥150,000・月額＋¥30,000〜", dict(size=11, bold=True, color=NAVY))),
             runs(("（追加した場合の合計：初期 ¥295,000〜／開催月 ¥365,000程度／平常期 ¥165,000〜）", dict(size=10, color=TEXT)))],
            fill="F5F7FA", line=NAVY, line_w=0.75, anchor="ctr", ins=(0.1, 0.02, 0.1, 0.02), name="追加オプション")


def s24_hikaku(s):
    for sh in list(s.shapes):
        if sh.shape_id not in (635, 636, 637):
            drop(sh)
    set_one(by_id(s, 635), "費用の比較：推奨構成と、サンクスLINE誘導を追加する場合")
    set_one(by_id(s, 637), "ポップアップの導入を推奨。サンクスLINE誘導は追加の選択肢")
    big = dict(size=15, bold=True, color=NAVY, align="ctr")
    rows = [[head("項目"), head("推奨構成：ポップアップあり"), head("追加構成：＋サンクスLINE誘導")]]
    data = [
        ("初期費用", cell("145,000円〜", **big), cell("295,000円〜（＋150,000円）", **big)),
        ("月額費用（平常期）", cell("135,000円〜", **big), cell("165,000円〜（＋30,000円〜）", **big)),
        ("月額費用（セミナー開催月）", cell("335,000円程度", **big), cell("365,000円程度（＋30,000円〜）", **big)),
        ("友だちの入口",
         cell("離脱防止ポップアップで、LPを離れる方を友だちに。予約完了ページには通常の友だち追加ボタン等を設置（費用・実装条件は仕様確認のうえ）", size=12),
         cell("上記に加え、予約完了後にLINEへ自動遷移（フォームの入力内容の引き継ぎも可能）", size=12)),
        ("追加するメリット", cell("－", size=12, align="ctr"),
         cell("予約した方を友だちにしやすくなり、特にセミナー開催時期の友だち追加数の上積みが期待できる", size=12)),
    ]
    for i, (k, a, b) in enumerate(data):
        fill = BAND[i % 2]
        a["fill"] = fill
        b["fill"] = fill
        lb = label(k, fill)
        lb["align"] = "ctr"
        rows.append([lb, a, b])
    add_table(s, 0.42, 1.48, [2.75, 3.5, 3.75], [0.42, 0.5, 0.5, 0.5, 0.95, 0.7], rows, name="費用の比較")
    note(s, 5.2, ["※月額費用＝コンサルティング費（平常期 100,000円／セミナー開催月 300,000円程度）＋ツール・LINE関連費（推奨構成 35,000円〜＝ポップアップ 30,000円＋LINE公式アカウント利用料 5,000円〜）",
                  "※初期費用（推奨構成）＝運用設計・構築 100,000円〜＋ポップアップ 15,000円＋LINE公式アカウント開設 30,000円",
                  "※ポップアップ・サンクスLINE誘導はどちらも選択制。ポップアップも導入しない場合の参考：初期 130,000円〜／平常期 105,000円〜／開催月 305,000円程度（別途の実装・制作費等は含まない）",
                  "※サンクスLINE誘導の最低利用期間・停止条件（セミナー開催月だけの利用可否）は、確認のうえご案内",
                  "※契約期間は6か月。広告費は含まない。施策内容・配信本数・連携範囲等により、費用は個別にお見積もりいたします。"])


def s25_ryokin(s):
    t = table_of(s)
    set_cell(t.cell(0, 3), "プラン③\n継続的な運用・改善")
    set_cell(t.cell(1, 0), "初期費用\n（ポップアップあり）")
    for c in (1, 2, 3):
        set_cell(t.cell(1, c), "145,000円〜")
    set_cell(t.cell(2, 2), "100,000円")
    set_cell(t.cell(4, 0), "月額ツール・LINE関連費\n（ポップアップあり）")
    for c in (1, 2, 3):
        set_cell(t.cell(4, c), "35,000円〜")
    for c, v in ((1, "35,000円〜"), (2, "135,000円〜"), (3, "235,000円〜")):
        set_cell(t.cell(5, c), v)
    for c, v in ((1, "35,000円〜"), (2, "335,000円程度"), (3, "335,000円〜")):
        set_cell(t.cell(6, c), v)
    texts = ["※初期費用＝運用設計・構築 100,000円〜＋ポップアップ 15,000円＋LINE公式アカウント開設 30,000円",
             "※月額ツール・LINE関連費＝ポップアップ 30,000円＋LINE公式アカウント利用料 5,000円〜",
             "※ツールは選択制。サンクスLINE誘導を追加する場合は、初期＋150,000円・月額＋30,000円〜（ポップアップを導入しない場合は、初期−15,000円・月額−30,000円）",
             "※成果報酬型プラン：初期 145,000円〜／月額 35,000円〜＋成果報酬（セミナー予約 12,000円／件）。ツールの追加は上記と同じ",
             "※契約期間は6か月。広告費は含まない。施策内容・配信本数・連携範囲等により、費用は個別にお見積もりいたします。"]
    for sid, txt in zip((640, 641, 642, 643, 644), texts):
        set_one(by_id(s, sid), txt)


def s26_unyo(s):
    t = table_of(s)
    set_cell(t.cell(0, 3), "プラン③\n継続的な運用・改善")
    set_cell(t.cell(1, 0), "ポップアップ（推奨）・\nサンクスLINE誘導（追加）")
    for c in (1, 2, 3):
        set_cell(t.cell(1, c), "選択制・導入内容に\n応じて対応")
    set_cell(t.cell(3, 2), "開催月 月4〜5回\n平常期 月1〜2回")
    note(s, 5.6, ["※月額のコンサルティング費には、企画配信の制作に加え、LINEアカウントの運用・導線の管理・ステップ配信の改善・配信結果の分析・次回施策への改善提案を含む",
                  "※平常期の配信は、次のイベントまで関係を保つためのもの（配信の回数を増やすこと自体は目的にしない）",
                  "※前日・当日のリマインドは、企画配信とは別に、予約ごとに配信。契約期間は6か月"])


def s29_shisaku(s):
    t = table_of(s)
    set_cell(t.cell(1, 1), "離脱防止ポップアップ（導入推奨）／予約完了後の友だち追加：サンクスLINE誘導（追加オプション）または通常のご案内")
    set_cell(t.cell(5, 1), "月1〜2回の企画配信／次回イベントのご案内")
    t.rows[1].height = Inches(0.8)
    by_name(s, "注記").top = Inches(6.15)


def s30_popup(s):
    pill(s, 4.1, 1.44, 1.5, "任意・導入推奨")


def s32_thanks(s):
    pill(s, 7.95, 0.92, 1.6, "追加オプション")
    set_one(by_id(s, 9), "予約後のLINE誘導を自動化し、開催時期の友だち数を上積み")
    for r in by_id(s, 9).text_frame.paragraphs[0].runs:
        r.font.size = Pt(13)
    add_box(s, 6.0, 6.86, 4.4, 0.22,
            [runs(("※最低利用期間・停止条件は確認のうえご案内", dict(size=10, color=GRAY)))],
            anchor="t", ins=(0, 0, 0, 0), name="追加オプションの注記")


def s35_follow(s):
    n = by_id(s, 22)
    set_lines(n, ["※参加した方を判別できる場合に、その情報を活用して配信（自動で判別できるかは、予約システム等との連携の可否による）",
                  "※配信内容は、貴社LPの「受講生の声」「講師メッセージ」「よくある質問」「講師紹介」をもとにした例"])
    n.top, n.height = Inches(6.47), Inches(0.45)


def s40_kakunin(s):
    drop(next(sh for sh in s.shapes if sh.has_table))
    drop(by_name(s, "注記"))
    lead = next(sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.startswith("連携の可否と成果の数え方"))
    set_one(lead, "契約前に確認する事項と、運用開始後に調整できる事項を分けて整理する")
    pre = [("予約・参加情報との紐づけ", "予約システム・フォーム・顧客管理の情報と、LINEの友だちを紐づけられるか（難しい場合は流入経路での出し分け・全体配信が基本）"),
           ("サンクスLINE誘導の導入", "導入するかどうか。導入する場合は、最低利用期間・停止条件（セミナー開催月だけの利用可否）"),
           ("通常の友だち追加ボタン", "予約完了ページへの設置可否・費用・実装条件"),
           ("SIMと成果の計測方法", "セミナー予約・サロン入会を、LINEからの件数とするか、広告からの予約も含めるか"),
           ("成果報酬型プラン（別案）", "成果地点の定義・重複計上の有無・計測方法（成果報酬単価 12,000円／件）")]
    post = [("一般の方向けセミナー", "同じLINE公式アカウントでの配信の出し分け（流入経路の分け方・配信内容）"),
            ("平常期の配信", "月1〜2回の範囲での回数・内容"),
            ("興味分野の確認", "友だち追加直後の簡易アンケートを入れるかどうか（必要最小限の1問程度）")]
    rows = [[head("区分"), head("項目"), head("確認させていただきたいこと")]]
    for grp, items, fill in (("契約前に確認", pre, "FFFFFF"), ("運用開始後に調整", post, "F5F7FA")):
        for i, (k, v) in enumerate(items):
            g = cell(grp, bold=True, color=NAVY, fill=fill, align="ctr", size=12)
            if i == 0:
                g["rowSpan"] = len(items)
            else:
                g = dict(g, text="", vMerge=True)
            lb = label(k, fill)
            lb["size"] = 12
            rows.append([g, lb, cell(v, fill=fill, size=12)])
    add_table(s, 0.46, 1.75, [1.5, 2.5, 5.9], [0.4, 0.62, 0.5, 0.45, 0.5, 0.5, 0.5, 0.45, 0.45], rows, name="確認事項")
    note(s, 6.35, ["※データ連携が必要な高度な施策は、別途要件定義・お見積りのうえ実施"])


def s41_schedule(s):
    side = copy.deepcopy(by_id(s, 3)._element)
    nid = max(sh.shape_id for sh in s.shapes) + 1
    side.find(".//" + qn("p:cNvPr")).set("id", str(nid))
    s.shapes._spTree.append(side)
    sh = by_id(s, nid)
    sh.top = Inches(3.0)
    set_lines(sh, ["ツールの選択", "実装条件の確認", "（サンクスLINE誘導の", "契約・停止条件）"])
    set_one(by_id(s, 41), "ツール導入・連携の設定 / タグ確認")
    n = by_id(s, 20)
    set_one(n, "※契約期間は6か月。貴社作業・確認にかかる時間により前後する可能性がございます。")
    n.left, n.width = Inches(3.9), Inches(6.67)


def s54_service(s):
    set_one(by_text(s, "開催月は月4〜5回、平常期は月2回の配信案"), "開催月は月4〜5回、平常期は月1〜2回の配信案")


EDITS = {1: (s01_cover, "御中"), 3: (s03_doshen, "想定動線"), 4: (s04_genjo, "現状把握"), 13: (s13_why, "なぜLINEか"),
         16: (s16_yoken, "要件定義"), 17: (s17_hoshin, "運用方針"), 18: (s18_flow, "全体導線"), 19: (s19_phase, "想定動線（フェーズ"),
         20: (s20_zentai, "全体設計"), 21: (s21_track, "配信の出し分け"), 23: (s23_mitsumori, "概算お見積り"),
         24: (s24_hikaku, "費用の比較"), 25: (s25_ryokin, "料金プラン"), 26: (s26_unyo, "運用内容"),
         29: (s29_shisaku, "LINE施策の全体像"), 30: (s30_popup, "ポップアップ"), 31: (s30_popup, "ポップアップ"),
         32: (s32_thanks, "サンクスLINE誘導"), 35: (s35_follow, "参加後のフォローアップ"), 40: (s40_kakunin, "確認させていただきたい"),
         41: (s41_schedule, "運用スケジュール"), 54: (s54_service, "弊社からご提供できる運用プラン")}


def main(src, out):
    prs = Presentation(src)
    S = list(prs.slides)
    assert len(S) == 62
    for n, (fn, expect) in EDITS.items():
        s = S[n - 1]
        t = title_of(s)
        txt = t.text_frame.text if t is not None else ""
        assert expect in txt or expect in " ".join(x.text_frame.text for x in s.shapes if x.has_text_frame), (n, expect, txt)
        fn(s)
    all_meiryo([S[n - 1] for n in EDITS])
    prs.save(out)
    print("saved:", out, len(prs.slides), "枚")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
