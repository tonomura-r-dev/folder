# -*- coding: utf-8 -*-
"""チェングロウス ver1.4 → ver1.5（2026-10-05 殿村さん指示）。
Chapter表紙4枚（3・9・13・18枚目）を、あっさりした区切りページに統一する。
テンプレートのヘッダー（紺の角・上部ライン・DYMロゴ）とフッターはそのまま見せ、
中央に小さく「Chapter N」＋大きく章タイトルだけを置く（番号の飾り・サブコピー・進み具合のバーは外す）。
  python3 _build/patch_chengrowth_bridge_v15_chapter_simple.py <ver1.4.pptx>
"""
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

sys.path.insert(0, str(Path(__file__).resolve().parent))
from pptx_parts import *  # noqa

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "20261005_チェングロウス_サイトリニューアルとLINE同時導入のご提案ver1.5.pptx"
prs = Presentation(sys.argv[1])
CH = {3: (1, "なぜLINEを同時導入するのか"), 9: (2, "求人業界でのLINE活用"),
      13: (3, "自動車求人Naviでの活用イメージ"), 18: (4, "効果・費用・導入スケジュール")}
for no, (k, title) in CH.items():
    s = prs.slides[no - 1]
    assert any(sh.has_text_frame and sh.text_frame.text == title for sh in s.shapes), (no, title)
    for sh in list(s.shapes):
        sh._element.getparent().remove(sh._element)
    label(s, 2.2, 8.6, 23.1, 0.9, [(f"Chapter {k}", 15, True, GRAY, 0)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    label(s, 2.2, 9.5, 23.1, 2.0, [(title, 32, True, NAVY, 0)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
prs.save(str(OUT))
print("saved:", OUT.name, len(prs.slides), "枚")
