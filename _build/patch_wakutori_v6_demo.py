# -*- coding: utf-8 -*-
"""ワクトリ v5 → v6：具体施策のイメージ（P27〜P33）を「左＝デモの画面／右＝ほかの案3つ」にする（2026-10-30 殿村さん）。

    python -m _build.patch_wakutori_v6_demo <v5.pptx> <v6.pptx>

  最強資料へのFB（左メイン＝LINEでこういうことができるを見せる／右の枠は3つでよい）を受けて。
  ・左：施策名の見出しは残し、流れの図形を外して、ワクトリのデモ（`_data/wakutori/ワクトリ_LINE施策デモ_20261029v2.html`）の
    ①ユーザーが見る画面＋②裏側の処理 を1人分貼る（`_images/wakutori/demo_p27〜33.png`）
  ・右：4つの枠（目的／対象／中身／見る数字）を、3つの枠（入口／LINEの仕組み／配信内容）＝「ほかにもこういう案」にする
  ほかのページは触らない。
"""
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Emu, Inches, Pt

from _build.pptx_parts import style_run

NAVY, INK, GRAY = "1F2A5C", "333333", "7F7F7F"
IMG = "_images/wakutori/demo_p{}.png"
LABELS = ["入口", "LINEの仕組み", "配信内容"]
IDEAS = {
    27: [["SNSのプロフィール・投稿や", "体験談のページにも入口を置く"],
         ["入口ごとにURLを分けて記録し、", "あいさつの文面を出し分ける"],
         ["追加の特典に「寮生活の", "持ち物リスト」などを届ける"]],
    28: [["離れるときだけでなく、", "スクロールや滞在時間でも表示"],
         ["表示したページ（スキー場・海）", "で、追加後の求人を出し分ける"],
         ["「この求人に似た求人」や", "締切が近い求人を届ける"]],
    29: [["リッチメニューや配信から、", "いつでも診断をやり直せる"],
         ["重視したいポイント（立地・", "稼げる・出会い・住環境）も聞く"],
         ["結果に合う働き方のタイプと", "体験談を一緒に届ける"]],
    30: [["診断に答えていない方には、", "翌日に答えるお願いを届ける"],
         ["押した内容（お金・寮など）で", "タグを付け、次の配信を変える"],
         ["寮の写真や1日の流れなど、", "不安ごとの配信を足す"]],
    31: [["特集から、リッチメニューの", "開始月別の求人へつなげる"],
         ["開始時期が近づいたら、", "自動で特集を届ける"],
         ["「LINE友だち限定」の先行公開や", "締切が近い求人の特集"]],
    32: [["応募（ENTRY）の前と後で、", "メニューを切り替える"],
         ["押した項目をタグにして、", "関心に合う配信につなげる"],
         ["季節ごとに特集の枠を入れ替える", "（冬はスキー場、夏は海）"]],
    33: [["お問い合わせなど、ほかの", "フォームの完了画面にも置く"],
         ["ENTRYで伺った内容と同じ質問は", "省いて、診断を短くする"],
         ["面談・お仕事開始までの流れ、", "持ち物・よくある質問を届ける"]],
}


def inch(v):
    return Emu(v).inches


def main(src, out):
    prs = Presentation(src)
    for n in range(27, 34):
        s = prs.slides[n - 1]
        big = next(sh for sh in s.shapes if abs(inch(sh.left) - 0.29) < 0.05 and abs(inch(sh.top) - 2.03) < 0.05
                   and inch(sh.width) > 6)
        head = next(sh for sh in s.shapes if sh.has_text_frame and abs(inch(sh.top) - 2.15) < 0.1 and inch(sh.left) < 1)
        right_boxes = sorted([sh for sh in s.shapes if inch(sh.left) > 6.8 and inch(sh.top) > 1.9 and not sh.has_text_frame
                              or (inch(sh.left) > 6.8 and inch(sh.top) > 1.9 and sh.has_text_frame and not sh.text_frame.text.strip())],
                             key=lambda x: inch(x.top))
        right_texts = [sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip() and inch(sh.left) > 6.8
                       and inch(sh.top) > 1.9]
        flow = [sh for sh in s.shapes if sh is not big and sh is not head and inch(sh.left) < 6.8 and inch(sh.top) > 2.8]
        assert len(right_boxes) == 4 and len(right_texts) == 4 and flow, (n, len(right_boxes), len(right_texts), len(flow))
        tmpl = right_boxes[0]
        for sh in flow + right_texts + right_boxes[1:]:
            sh._element.getparent().remove(sh._element)

        # 左：デモの画面（①＋②）を枠の中央に
        h = 3.62
        w = h * 1744 / 1468
        pic = s.shapes.add_picture(IMG.format(n), Inches(0.29 + (6.44 - w) / 2), Inches(2.9), Inches(w), Inches(h))
        pic.line.color.rgb = RGBColor.from_string("D9D9D9")
        pic.line.width = Pt(0.75)
        cap = s.shapes.add_textbox(Inches(0.45), Inches(6.53), Inches(6.1), Inches(0.24))
        r = cap.text_frame.paragraphs[0].add_run()
        r.text = "画面：ワクトリ様向けに作成したデモ（①ユーザーが見る画面／②裏側の処理）。求人はデモ用のサンプル"
        style_run(r, 8.5, False, GRAY)

        # 右：3つの枠（入口／LINEの仕組み／配信内容）＝ほかの案
        top, bottom, gap = 2.03, 6.80, 0.1
        bh = (bottom - top - gap * 2) / 3
        tmpl.top, tmpl.height = Inches(top), Inches(bh)
        boxes = [tmpl]
        for k in (1, 2):
            el = __import__("copy").deepcopy(tmpl._element)
            s.shapes._spTree.append(el)
            sh = s.shapes[-1]
            sh.top = Inches(top + k * (bh + gap))
            boxes.append(sh)
        for k, (box, label, lines) in enumerate(zip(boxes, LABELS, IDEAS[n])):
            tb = s.shapes.add_textbox(box.left + Inches(0.12), box.top + Inches(0.1), box.width - Inches(0.24), Inches(bh - 0.2))
            tf = tb.text_frame
            tf.word_wrap = True
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            p0 = tf.paragraphs[0]
            r = p0.add_run()
            r.text = f"ほかの案｜{label}"
            style_run(r, 13.5, True, NAVY)
            p0.space_after = Pt(4)
            for line in lines:
                p = tf.add_paragraph()
                r = p.add_run()
                r.text = line
                style_run(r, 12, False, INK)
        # 図形のIDの重複をなくす
        used = set()
        nid = max(int(e.get("id")) for e in s.shapes._spTree.iter("{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr")) + 1
        for e in s.shapes._spTree.iter("{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr"):
            if e.get("id") in used:
                e.set("id", str(nid))
                nid += 1
            used.add(e.get("id"))
    prs.save(out)
    print("saved:", out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
