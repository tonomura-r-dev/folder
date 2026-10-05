# -*- coding: utf-8 -*-
"""チェングロウス19枚 ver1.1 → ver1.2（2026-10-05 殿村さん指示）。
17P・19Pを「カードを並べる形」から図解に変える。17＝流入→サイト→LINE→応募の一本線＋費用の置き場所／19＝11月〜4月以降の実スケジュール図。
文言・数字・時期は ver1.1 のまま。他のページは触らない。
  python3 _build/patch_chengrowth_bridge_v12_layout.py <ver1.1.pptx>
"""
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).resolve().parent))
from pptx_parts import *  # noqa

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "20261005_チェングロウス_サイトリニューアルとLINE同時導入のご提案ver1.2.pptx"
prs = Presentation(sys.argv[1])
S = list(prs.slides)

# ================= 17 導入費用の考え方：集客→サイト→LINE→応募 の一本線に費用を置く =================
s17 = S[16]
keep_header_only(s17)
retitle(s17, "導入費用の考え方｜サイト制作費とは分けて明示",
        "集客 → サイト → LINE → 応募の流れの中で、費用がかかる場所を分けて見る")
X = [2.2, 7.9, 13.6, 19.3]
names = ["広告・流入", "Webサイト", "LINEで接点を残す", "応募"]
fills = [GRAY, GRAY, GREEN, NAVY]
for i in range(4):
    shape(s17, MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON, X[i], 4.8, 6.0, 2.3,
          [(names[i], 15, True, WHITE, 0)], fill=fills[i], adj=0.3, margins=(0.5 if i else 0.3, 0.05, 0.6, 0.05))
rect(s17, 2.2, 7.5, 5.6, 0.12, GRAY)
rect(s17, 7.9, 7.5, 5.6, 0.12, GRAY)
rect(s17, 13.6, 7.5, 11.7, 0.12, GREEN)
blocks = [(2.2, 5.6, "広告費", INK, "担い手：各広告媒体", "サイトへの来訪を増やす"),
          (7.9, 5.6, "サイト改修費", INK, "担い手：サイト制作会社", "見やすさ・検索性・応募フォームを改善"),
          (13.6, 11.7, "LINE独自構築費", GREEN_TX, "担い手：弊社（本提案）", "来訪者を友だち・会員として蓄積し、応募までつなげる")]
for x, w, h, c, who, role in blocks:
    label(s17, x, 7.85, w, 2.9, [(h, 18, True, c, 6), (who, 12, False, GRAY, 4), (role, 12.5, False, INK, 0)])
# 左：この図で一番伝えたいこと／右：LINEの費用の置き場所（1つの説明エリア）
rect(s17, 2.2, 11.5, 0.12, 2.6, NAVY)
label(s17, 2.55, 11.5, 10.4, 2.6, [("LINEの費用は、サイト制作費・広告費とは役割が異なる", 17, True, NAVY, 0)], anchor=MSO_ANCHOR.MIDDLE)
rect(s17, 13.6, 11.5, 11.7, 2.6, CARD, margins=(0.5, 0.2, 0.4, 0.2), align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE,
     paras=[("LINE独自構築費のかかり方", 11, True, GREEN_TX, 4),
            ("初期費用：11月頃〜　｜　月額費用：4月〜", 13.5, True, NAVY, 5),
            ("LINE公式アカウント利用料：別途", 11.5, False, INK, 2),
            ("LINE連携開発：内容に応じて別途お見積もり", 11.5, False, INK, 0)])
hline(s17, 2.2, 25.3, 15.0, NAVY, 1.5)
label(s17, 2.2, 15.25, 23.1, 1.3, [("LINEの費用を分けることで、友だち・会員・応募への効果を単独で把握できる", 16, True, NAVY, 0)],
      align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

# ================= 19 スケジュール：11月〜4月以降の一本の工程図 =================
s19 = S[18]
keep_header_only(s19)
retitle(s19, "スケジュール｜11月頃から初期構築、4月の公開と同時に本格運用へ",
        "11月頃から初期構築に入り、4月のサイト公開後すぐに本格運用へ移行")
LX, LW = 2.2, 4.6          # 工程名の列
TX, MW, GX = 7.0, 2.1, 17.5  # 時間軸の左端／1か月の幅／4月の境目
EX = 25.3
months = ["11月", "12月", "1月", "2月", "3月"]
# 上の帯：設計・構築（11月〜3月）が、そのまま運用・改善（4月〜）へ続く
shape(s19, MSO_SHAPE.PENTAGON, TX, 4.4, GX - TX + 0.3, 0.9, [("11月〜3月｜設計・構築", 12.5, True, WHITE, 0)], fill=NAVY, adj=0.5)
shape(s19, MSO_SHAPE.CHEVRON, GX - 0.1, 4.4, EX - GX + 0.1, 0.9, [("4月〜｜運用・改善", 12.5, True, WHITE, 0)], fill=GREEN, adj=0.5,
      margins=(0.6, 0.05, 0.3, 0.05))
for i, m in enumerate(months):
    label(s19, TX + i * MW, 5.4, MW, 0.6, [(m, 11.5, True, INK, 0)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
label(s19, GX, 5.4, EX - GX, 0.6, [("4月〜", 11.5, True, GREEN_TX, 0)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
rows = [  # (高さ, 工程名, 時期, 開始x, 終了x, 文言, 種類)
    (1.2, "① LINE設計", "11月頃〜", TX, GX, "LINE公式アカウント・リッチメニュー・あいさつメッセージ・ステップ配信・計測の設計", "navy"),
    (1.6, "② サイト・LINE連携", "11月頃〜3月", TX, GX, "LINE Profile+の申請・「LINEで登録」の仕様整理・会員連携・サンクスLINE誘導の仕様整理・条件別検索ページとの連携・サイト制作会社との仕様調整", "navy"),
    (1.15, "③ 実装・検証", "サイト制作の進行に合わせて〜3月", TX, GX, "開発・実装・公開前テスト", "light"),
    (1.3, "④ サイト公開", "4月", None, None, None, "milestone"),
    (1.45, "⑤ 本格運用", "4月〜", GX, EX, "友だち／会員の獲得・ステップ配信・企画配信・リッチメニュー運用・応募への追客", "green"),
    (1.25, "⑥ 分析・改善", "4月〜継続", GX, EX, "KPI計測・月次分析・レポーティング・導線／配信／セグメントの改善", "greenlight"),
]
y0 = 6.1
ys = []
y = y0
for h, *_ in rows:
    ys.append(y)
    y += h + 0.1
bottom = y - 0.1
for i in range(1, 6):  # 月の区切り線（薄く）
    vline(s19, TX + i * MW, 5.5, bottom, "E7E7E7", 0.5)
for (h, nm, when, x1, x2, txt, kind), yy in zip(rows, ys):
    hline(s19, LX, EX, yy - 0.05, LGRAY, 0.5)
    label(s19, LX, yy, LW + 0.1, h, [(nm, 11.5, True, NAVY, 1), (when, 9, False, GRAY, 0)], anchor=MSO_ANCHOR.MIDDLE)
    if kind == "milestone":
        continue
    fill, line, tc = {"navy": (NAVY, None, WHITE), "light": ("E4E8F6", NAVY, NAVY),
                      "green": (GREEN, None, WHITE), "greenlight": (GREEN_BG, GREEN, GREEN_TX)}[kind]
    rect(s19, x1 + 0.05, yy + 0.06, x2 - x1 - 0.1, h - 0.12, fill, line=line, lw=1.0,
         paras=[(txt, 10, True, tc, 0)], align=PP_ALIGN.LEFT, margins=(0.3, 0.03, 0.3, 0.03))
# 4月の境目：マイルストーン
vline(s19, GX, 5.4, bottom, GREEN, 2.5)
my = ys[3]
shape(s19, MSO_SHAPE.DIAMOND, GX - 0.4, my + 0.25, 0.8, 0.8, fill=GREEN)
label(s19, GX + 0.6, my, EX - GX - 0.6, rows[3][0], [("サイトリニューアル公開", 12, True, GREEN_TX, 1), ("LINE本格稼働", 12, True, GREEN_TX, 0)],
      anchor=MSO_ANCHOR.MIDDLE)
label(s19, TX, my, GX - TX - 0.6, rows[3][0], [("初期構築で設計した内容を、そのまま運用へ  ▶", 11.5, True, NAVY, 0)],
      align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
hline(s19, 2.2, 25.3, bottom + 0.3, NAVY, 1.5)
label(s19, 2.2, bottom + 0.4, 23.1, 0.9, [("初期構築から運用・改善まで一貫して支援するため、サイト公開後もスムーズに本格運用へ移行", 14, True, NAVY, 0)],
      align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
label(s19, 2.2, bottom + 1.35, 23.1, 1.4, [
    ("※11月頃〜3月は、定期配信・月次分析・改善運用は原則行わず、設計・構築に集中。必要に応じて一部施策の先行稼働も検討可能（先行稼働分は、シミュレーションに含めない）", 8.5, False, GRAY, 1),
    ("※LINE Profile+の審査期間は、申請内容により変わります／実装の時期は、サイト制作の進行に応じて決定", 8.5, False, GRAY, 1),
    ("※求人ボックス・Indeed経由については、着地先・応募完了先を確認したうえでLINE導線を設計", 8.5, False, GRAY, 0)])
print("bottom", bottom)

AFTER = {qn(x) for x in ("a:sym", "a:hlinkClick", "a:hlinkMouseOver", "a:rtl", "a:extLst")}
for sl in (s17, s19):
    for tag in ("rPr", "endParaRPr"):
        for rpr in sl.shapes._spTree.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}" + tag):
            if rpr.get("sz") and int(rpr.get("sz")) < 800:
                rpr.set("sz", "800")
prs.save(str(OUT))
print("saved:", OUT.name, len(prs.slides), "枚")
