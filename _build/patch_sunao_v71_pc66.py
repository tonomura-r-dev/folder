# -*- coding: utf-8 -*-
"""すなお：殿村さんのPC保存版ver6.6（2026-10-22）に、ver7.0までの直しと、竹の開催月の金額の直しを当てる → ver7.1。

    python -m _build.patch_sunao_v71_pc66 <PC保存版ver6.6.pptx> <ver7.1.pptx>

  ① ver6.7：想定動線の表のフェーズ名を、上の8フェーズの帯に合わせる
  ② ver6.8：「平常期」→「開催のない月」（PC保存版で消えた注記・変わった文言は飛ばす）／「なぜLINEか」の表を左右の比較に
  ③ ver7.0：想定動線を①〜④／⑤〜⑧の2ページに分ける（60枚 → 61枚）
  ④ 竹（プラン②）の開催月のコンサル費 300,000円 → 250,000円 に合わせて、概算お見積り・費用の比較の金額を直す
     （料金プランのページは殿村さんがPCで直し済み）
"""
import sys

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches

from _build.patch_sunao_v67_phase_names import COLS, PHASES
from _build.patch_sunao_v68_kaisai import EDITS, compare, paras, ptext, set_para
from _build.patch_sunao_v70_s17_2p import PAGES, TITLE, build
from _build.patch_tanaka_v5 import copy_slide

MORE = {23: [("平常期 月2回", "開催のない月は月2回")]}   # PC保存版で、プラン③を月3回 → 月2回に直してあるため
PRICE = {  # 竹の開催月：コンサル費 300,000 → 250,000（月額 335,000 → 285,000、サンクスLINE誘導を足すと 365,000 → 315,000）
    20: [("開催月 ¥335,000程度", "開催月 ¥285,000程度"),
         ("（追加した場合の合計：初期 ¥295,000〜／開催月 ¥365,000程度／開催のない月 ¥165,000〜）",
          "（追加した場合の合計：初期 ¥295,000〜／開催月 ¥315,000程度／開催のない月 ¥165,000〜）")],
    21: [("335,000円程度", "285,000円程度"),
         ("365,000円程度（＋30,000円〜）", "315,000円程度（＋30,000円〜）")],
}


def apply(slides, edits, strict):
    skipped = []
    for no, pairs in edits.items():
        ps = paras(slides[no - 1])
        for old, new in pairs:
            hits = [p for p in ps if ptext(p) == old]
            if not hits:
                assert not strict, (no, old)
                skipped.append((no, old[:30]))
            for p in hits:
                set_para(p, new)
    return skipped


def phase_names(s):
    gf = next(sh for sh in s.shapes if sh.has_table)
    tbl = gf.table
    assert tbl.cell(0, 0).text == "フェーズ" and len(tbl.rows) == 9
    for i, w in enumerate(COLS):
        tbl.columns[i].width = Inches(w)
    for r, name in enumerate(PHASES, start=1):
        p = tbl.cell(r, 0)._tc.find(qn("a:txBody")).find(qn("a:p"))
        runs = p.findall(qn("a:r"))
        for x in runs[1:]:
            p.remove(x)
        runs[0].find(qn("a:t")).text = name


def main(src, out):
    prs = Presentation(src)
    slides = list(prs.slides)
    assert len(slides) == 60
    phase_names(slides[16])                                                   # ①
    skipped = apply(slides, EDITS, strict=False) + apply(slides, MORE, strict=True)   # ②
    compare(slides[10])
    apply(slides, PRICE, strict=True)                                         # ④（ページ番号は分ける前の番号）

    first = slides[16]                                                        # ③
    second = copy_slide(prs, first)
    lst = prs.slides._sldIdLst
    ids = list(lst)
    new = ids[-1]
    lst.remove(new)
    ids[16].addnext(new)
    for s, (idx, title, lead) in zip((first, second), PAGES):
        assert build(s, idx, title, lead) <= 7.0

    left = [i + 1 for i, s in enumerate(prs.slides) if any("平常" in ptext(p) for p in paras(s))]
    assert not left, left
    assert len(prs.slides) == 61
    prs.save(out)
    print("saved:", out)
    print("飛ばした直し（PC保存版に元の文が無い）:", skipped)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
