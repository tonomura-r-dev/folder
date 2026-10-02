# -*- coding: utf-8 -*-
"""チェングロウス ver2.7（殿村さんPC保存版）→ ver2.8：章の切り替わりに扉（Chapter）を6枚追加（2026-10-02）
扉の見本は DYM提案FMT の「資料アジェンダ」ページ（タイトル＋中央に章名）。
  python3 _build/patch_chengrowth_v28_chapters.py <ver2.7.pptx> <DYM_LINEOA_提案FMT_202607.pptx>
"""
import copy, sys
from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.util import Pt

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "20260929_チェングロウス_LINE公式アカウント運用のご提案ver2.8.pptx"
prs = Presentation(sys.argv[1])
fmt = Presentation(sys.argv[2]).slides[1]
NAVY = RGBColor(0x1F, 0x28, 0x5A)

# (挿入する直前のスライド番号, 章名)  ※番号はver2.7のもの
CHAPTERS = [
    (3, "LINE活用の事例とメリット"),
    (9, "貴社での活用イメージ"),
    (12, "施策のご提案"),
    (16, "導入の進め方"),
    (19, "成果シミュレーションと費用"),
    (21, "市場と求職者の動き"),
]
layout = prs.slides[2].slide_layout
orig = list(prs.slides._sldIdLst)
for n, (before, name) in enumerate(CHAPTERS, start=1):
    s = prs.slides.add_slide(layout)
    for sh in list(s.shapes):
        sh._element.getparent().remove(sh._element)
    for el in fmt.shapes._spTree:
        if el.tag in (qn("p:nvGrpSpPr"), qn("p:grpSpPr")):
            continue
        s.shapes._spTree.append(copy.deepcopy(el))
    title, box = list(s.shapes)
    title.text_frame.paragraphs[0].runs[0].text = "資料アジェンダ"
    W = prs.slide_width
    box.width = int(W * 0.8)
    box.left = int((W - box.width) / 2)
    tf = box.text_frame
    p0 = tf.paragraphs[0]
    for r in p0.runs[1:]:
        r._r.getparent().remove(r._r)
    p0.runs[0].text = name
    # 「Chapter n」を章名の上に1行足す
    p_new = copy.deepcopy(p0._p)
    p0._p.addprevious(p_new)
    p_new.findall(qn("a:r"))[0].find(qn("a:t")).text = f"Chapter {n}"
    first = tf.paragraphs[0]
    first.runs[0].font.size = Pt(24)
    first.runs[0].font.color.rgb = NAVY
    first.space_after = Pt(10)
    lst = prs.slides._sldIdLst
    item = lst[-1]
    lst.remove(item)
    lst.insert(list(lst).index(orig[before - 1]), item)

for sl in prs.slides:
    for k, c in enumerate(sl.shapes._spTree.iter(qn("p:cNvPr")), start=2):
        c.set("id", str(k))

AFTER = {qn(x) for x in ("a:sym", "a:hlinkClick", "a:hlinkMouseOver", "a:rtl", "a:extLst")}
for sl in prs.slides:
    if not any("資料アジェンダ" in (sh.text_frame.text if sh.has_text_frame else "") for sh in sl.shapes):
        continue
    for r in sl.shapes._spTree.iter(qn("a:r")):
        if r.find(qn("a:rPr")) is None:
            r.insert(0, r.makeelement(qn("a:rPr"), {"lang": "ja-JP"}))
    for tag in ("a:rPr", "a:endParaRPr", "a:defRPr"):
        for rpr in sl.shapes._spTree.iter(qn(tag)):
            fs = []
            for ft in ("a:latin", "a:ea", "a:cs"):
                e = rpr.find(qn(ft))
                if e is None:
                    e = rpr.makeelement(qn(ft), {})
                else:
                    rpr.remove(e)
                e.set("typeface", "メイリオ")
                fs.append(e)
            nxt = next((c for c in rpr if c.tag in AFTER), None)
            for e in fs:
                (nxt.addprevious(e) if nxt is not None else rpr.append(e))
prs.save(str(OUT))
print("saved:", OUT.name, len(prs.slides), "枚")
