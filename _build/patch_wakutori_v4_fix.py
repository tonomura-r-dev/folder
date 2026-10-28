# -*- coding: utf-8 -*-
"""ワクトリ v3 → v4（2026-10-28 殿村さんの確認ポイント）。

    python -m _build.patch_wakutori_v4_fix <v3.pptx> <v4.pptx>

  1. アンケートの設計を15ページ（今のアンケート3問＝エントリー状況・開始時期・エリア）に合わせる：24ページのDay0・29ページ
     （「重視したいポイント」を足す案は外す）
  2. 契約満了の話は載せない：34ページ（契約満了前の次のお仕事案内）を外す・25ページの「契約満了の前」の行を外す
     （殿村さんのPC保存版v3＝55枚は34ページを外し済み。そのときは外さない）
  3. 概算お見積りのページに、SIMの前提を1行足す（殿村さん「今回は入れよう」）
  ほかのページは触らない。
"""
import sys

from pptx import Presentation
from pptx.oxml.ns import qn

from _build.patch_wakutori_v2_shisaku import by_id, set_lines
from pptx.enum.text import MSO_ANCHOR

from _build.pptx_parts import drop_slide, rich

PREMISE = ("※前提：サイトの訪問 月85,887人／友だちの入口＝サイトのLINEボタン・ポップアップ／ブロック率34.1%から／"
           "企画配信 月3本・ステップ配信5通／応募＝LINEからのクリックの0.4〜0.9%")

PARA = {  # (ページ, 今の段落の文字) → 新しい文字
    (29, "今のアンケートと「重視したいポイント」を発展"): "今のアンケート3問の答えで、合う求人を出す",
    (29, "3問タップ（開始時期・エリア・重視したいポイント）"): "3問タップ（エントリー状況・開始時期・エリア）",
    (29, "今のアンケート（開始時期・エリア）に「重視したいポイント（立地・稼げる・出会い多数・住環境）」を足す"):
        "今のアンケート（エントリー状況・開始時期・エリア）の答えで、合う求人と配信を出し分ける",
}


def para_text(p):
    return "".join(t.text or "" for t in p.iter(qn("a:t")))


def set_para(p, text):
    rs = p.findall(qn("a:r"))
    for r in rs[1:]:
        p.remove(r)
    rs[0].find(qn("a:t")).text = text


def main(src, out):
    prs = Presentation(src)
    assert len(prs.slides) in (55, 56)
    for (page, old), new in PARA.items():
        hit = [p for p in prs.slides[page - 1]._element.iter(qn("a:p")) if para_text(p) == old]
        assert len(hit) == 1, (page, old)
        set_para(hit[0], new)
    set_lines(by_id(prs.slides[23], 191), ["3問タップで", "エントリー状況・", "開始時期・エリアを分岐", "・タグ付け", "（今のアンケートを活かす）"])
    t46 = by_id(prs.slides[24], 46)
    set_lines(t46, ["・9月：冬の求人の先行案内", "・4〜5月：夏の求人の先行案内", "・「LINE友だち限定・先行公開」"])
    s34 = prs.slides[33]
    if any(sh.has_text_frame and sh.text_frame.text.startswith("契約満了前") for sh in s34.shapes):
        drop_slide(prs, s34)
    assert len(prs.slides) == 55
    est = next(s for s in prs.slides if any(sh.has_text_frame and sh.text_frame.text.startswith("概算お見積り") for sh in s.shapes))
    rich(est, 0.29 * 2.54, 6.6 * 2.54, 10.2 * 2.54, 0.45 * 2.54, [[(PREMISE, 9, False, "595959")]], anchor=MSO_ANCHOR.TOP)
    prs.save(out)
    print("saved:", out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
