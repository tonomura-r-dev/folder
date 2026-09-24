# -*- coding: utf-8 -*-
"""チェングロウスver1.1：前後検索の4分類をS6・S20・S21に反映（2026-09-24 殿村さんOK済みの案）。
- S6  画像を4色版に、凡例4色、時期の帯を4本（前/当日/直後/後半）、所見3つと締めを差し替え
- S20 14日間ステップを前後検索ベースに組み替え（面談①を3日目に前倒し）＋「根拠（前後検索）」列を追加
- S21 11月に「冬の繁忙期＝現場のリアル」を追加
  python3 _build/patch_chengrowth_0924d.py <入力pptx> <出力pptx>
"""
import copy
import sys
from pathlib import Path
from pptx import Presentation
from pptx.util import Cm, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
SRC, OUT = sys.argv[1], sys.argv[2]
prs = Presentation(SRC)


def by_id(slide, sid):
    for sh in slide.shapes:
        if sh.shape_id == sid:
            return sh
    raise KeyError(sid)


def set_text(tf, lines, tmpl_para=0):
    """段落・runの書式を保ったまま全行差し替え。lines は str か (str, size_pt, bold, color)"""
    txBody = tf._txBody
    ps = txBody.findall(qn("a:p"))
    tp = ps[min(tmpl_para, len(ps) - 1)]
    tr = tp.find(qn("a:r"))
    p_tmpl = copy.deepcopy(tp)
    for r in p_tmpl.findall(qn("a:r")) + p_tmpl.findall(qn("a:br")) + p_tmpl.findall(qn("a:fld")):
        p_tmpl.remove(r)
    r_tmpl = copy.deepcopy(tr) if tr is not None else None
    for p in ps:
        txBody.remove(p)
    for ln in lines:
        text, size, bold, color = (ln, None, None, None) if isinstance(ln, str) else (list(ln) + [None] * 4)[:4]
        p = copy.deepcopy(p_tmpl)
        end = p.find(qn("a:endParaRPr"))
        if r_tmpl is not None:
            r = copy.deepcopy(r_tmpl)
        else:
            r = p.makeelement(qn("a:r"), {})
            r.append(r.makeelement(qn("a:rPr"), {"lang": "ja-JP"}))
            r.append(r.makeelement(qn("a:t"), {}))
        r.find(qn("a:t")).text = text
        rPr = r.find(qn("a:rPr"))
        if size is not None:
            rPr.set("sz", str(int(size * 100)))
        if bold is not None:
            rPr.set("b", "1" if bold else "0")
        if color is not None:
            for f in rPr.findall(qn("a:solidFill")):
                rPr.remove(f)
            sf = rPr.makeelement(qn("a:solidFill"), {})
            sf.append(sf.makeelement(qn("a:srgbClr"), {"val": color}))
            rPr.insert(0, sf)
        if end is not None:
            end.addprevious(r)
        else:
            p.append(r)
        txBody.append(p)


def resize_table(gf, n_rows=None, n_cols=None, col_widths_cm=None):
    tbl = gf.table._tbl
    if n_cols is not None:
        grid = tbl.tblGrid
        while len(grid.findall(qn("a:gridCol"))) < n_cols:
            grid.append(copy.deepcopy(grid.findall(qn("a:gridCol"))[-1]))
            for tr in tbl.findall(qn("a:tr")):
                tr.append(copy.deepcopy(tr.findall(qn("a:tc"))[-1]))
    if n_rows is not None:
        trs = tbl.findall(qn("a:tr"))
        while len(trs) < n_rows:
            tbl.append(copy.deepcopy(trs[-1]))
            trs = tbl.findall(qn("a:tr"))
    if col_widths_cm:
        for gc, w in zip(tbl.tblGrid.findall(qn("a:gridCol")), col_widths_cm):
            gc.set("w", str(int(Cm(w))))
        gf.width = Cm(sum(col_widths_cm))
    gf.height = sum(int(tr.get("h")) for tr in tbl.findall(qn("a:tr")))


# ================= S6 =================
s = prs.slides[5]
pic = by_id(s, 16)
new = s.shapes.add_picture(str(ROOT / "_images/chengrowth_zengo_4color.png"), pic.left, pic.top, width=pic.width)
pic._element.addprevious(new._element)
pic._element.getparent().remove(pic._element)
L, W = pic.left / 360000, pic.width / 360000
set_text(by_id(s, 3).text_frame, [
    "LINEヤフーの前後検索データ（起点KW：カーディーラー）。",
    "職業への関心は検索当日に集中し、3日で消える。"])
leg = by_id(s, 15)
leg.left, leg.top, leg.width, leg.height = Cm(12.3), Cm(4.3), Cm(6.2), Cm(2.3)
set_text(leg.text_frame, [("■ 職業・働き方", 10.5, True, "D9661F"), ("■ 整備・工具・カー用品", 10.5, True, "3467B2"),
                          ("■ お金・手続き", 10.5, True, "8E4EC6"), ("■ クルマ選び（車種・店舗）", 10.5, True, "00897B")])
px = lambda p: L + p / 2000 * W
BARS = [(12, 80, 795, "8C8C8C", "前：ほかの職種と比較"), (13, 795, 1270, "D9661F", "当日：年収・職種"),
        (14, 1270, 1507, "E8A06A", "働き方")]
top = by_id(s, 12).top
for sid, a, b, col, text in BARS:
    sh = by_id(s, sid)
    sh.left, sh.width = Cm(px(a)), Cm(px(b) - px(a))
    sh.fill.fore_color.rgb = RGBColor.from_string(col)
    set_text(sh.text_frame, [(text, 10, True, "FFFFFF")])
last = copy.deepcopy(by_id(s, 14)._element)
by_id(s, 14)._element.addnext(last)
b4 = [sh for sh in s.shapes if sh._element is last][0]
b4.left, b4.width = Cm(px(1507)), Cm(px(1975) - px(1507))
b4.fill.fore_color.rgb = RGBColor.from_string("3467B2")
set_text(b4.text_frame, [("後半：現場・季節作業", 10, True, "FFFFFF")])
last.find(".//" + qn("p:cNvPr")).set("id", "17")
CARDS = {
    6: ("① 検索の65%は「買う人」", ["車種・店舗が129件中84件。", "検索広告では狙えない"]),
    7: ("② 職業の関心は当日〜2日後", ["年収・職種 → 働き方の順に検索。", "最初の3日で出し切る"]),
    8: ("③ 整備・工具の関心は前後ずっと", ["工具・タイヤ交換・オイル交換。", "「クルマ好きを仕事に」が長く効く"]),
}
for sid, (head, body) in CARDS.items():
    tf = by_id(s, sid).text_frame
    set_text(tf, [(head, 12, True, "1F285A")] + [(b, 10.5, False, None) for b in body], tmpl_para=1)
set_text(by_id(s, 9).text_frame, ["関心が消える前の3日間に、LINEで答えを届ける。"])
set_text(by_id(s, 10).text_frame, [
    "出典：LINEヤフー 前後検索データ（起点KW：カーディーラー／2026-09-23取得）。色分けはDYM分類（ラベル129件）"])

# ================= S20 =================
s = prs.slides[19]
set_text(by_id(s, 6).text_frame, [
    "友だち追加後14日間で、前後検索で見えた関心の順番（年収・職種 → 働き方 → 比較 → 現場）に沿って配信する",
    "（関心が消える前の3日目に1回目の面談案内。未予約者には再アプローチ、面談予約者は別シナリオへ分岐）"])
gf = by_id(s, 46)
resize_table(gf, n_rows=10, n_cols=4, col_widths_cm=[2.35, 2.9, 12.3, 7.6])
STEP = [
    ("Day", "対象", "配信内容", "根拠（前後検索）"),
    ("0日後", "全員", "あいさつ＋30秒診断＋「整備士の年収相場」を一言", "当日：カーディーラー 年収"),
    ("1日後", "全員", "職種図鑑（整備・受付・営業の仕事と年収の違い）", "当日：受付嬢・ディーラーとは"),
    ("2日後", "全員", "働き方（休み・残業・「ホワイトな職場」の見分け方）", "直後：ホワイト企業"),
    ("3日後", "未予約者", "面談のご案内①「履歴書、まだ要りません」", "関心が消える前に案内"),
    ("5日後", "未予約者", "ほかの職種との比較（整備士 vs 営業・施工管理）", "前：営業マン・施工管理"),
    ("7日後", "資格別", "資格別の年収相場・キャリアマップ", "当日：年収"),
    ("10日後", "未予約者", "現場のリアル（タイヤ交換の繁忙期・工具・手当）", "後半：スタッドレス・タイヤ交換"),
    ("12日後", "未予約者", "地域別の求人ピックアップ（オートバックスG）", "－"),
    ("14日後", "未予約者", "面談のご案内②＋以後は月数回の定常配信へ切替", "－"),
]
tbl = gf.table
for i, row in enumerate(STEP):
    for j, v in enumerate(row):
        if i > 0 and j == 3:
            set_text(tbl.cell(i, j).text_frame, [(v, 11, False, "7F7F7F")])
        elif i > 0 and j == 2:
            set_text(tbl.cell(i, j).text_frame, [(v, 12.5, None, None)])
        else:
            set_text(tbl.cell(i, j).text_frame, [v])
concl = by_id(s, 47)
concl.top = gf.top + gf.height + Cm(0.25)
set_text(concl.text_frame, [
    "職業への関心は検索当日〜2日後に集中 → 関心が消える前の3日目に1回目の面談案内",
    "面談・応募 95件（6ヶ月累計・SIM①）＝ 累計CPA 22,768円として算出",
    "面談予約で止めず、就業までLINEで追いきります（面談前日リマインド→当日案内→面談後フォロー）"])
set_text(by_id(s, 48).text_frame, ["出典：LINEヤフー 前後検索データ（起点KW：カーディーラー）／コンテンツ案v1（別紙・配信文面の全文）"])

# ================= S21 =================
s = prs.slides[20]
gf = by_id(s, 46)
resize_table(gf, n_rows=11)
PLAN = [
    ("1-2月", "「冬のうちに相談だけ」訴求", "3月の山に向けて前倒し。在職中でも相談OKを打ち出す"),
    ("3月", "年度替わり求人特集", "4月入社求人を地域別に配信（3月＝検索の山）"),
    ("3-4月", "資格取得者向け特集", "2級・検査員の合格者へ資格別の年収相場を配信"),
    ("5月", "連休明けの見直し訴求", "GW明けの「このままでいいか」層を面談へ"),
    ("6月", "LINEリサーチでニーズ収集", "転職理由・希望条件を簡易アンケートで把握し企画に反映"),
    ("7-8月", "賞与後の転職検討特集", "夏のボーナス後に動く層へ年収相場・求人を配信"),
    ("9月", "資格取得・キャリア特集", "無資格・3級向けに整備士育成プロジェクトを案内"),
    ("10月", "年内入社ラストコール", "年内に動きたい層へ面談枠をご案内"),
    ("11月", "冬の繁忙期＝現場のリアル", "スタッドレス・タイヤ交換期の現場と手当を紹介"),
    ("12月", "冬も相談だけOK訴求", "検索の谷でも検討は続く。配信は止めず関係を維持"),
]
tbl = gf.table
for i, row in enumerate(PLAN, start=1):
    for j, v in enumerate(row):
        set_text(tbl.cell(i, j).text_frame, [v])

prs.save(OUT)
print("saved", OUT)
