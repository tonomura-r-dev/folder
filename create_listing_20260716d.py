import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def make_xlsx(filename, data):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "架電リスト"

    headers = ["企業名", "LP URL", "電話番号", "商材", "検索KW", "業界"]

    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="メイリオ", bold=True, color="FFFFFF", size=10)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center
        cell.border = border

    for row_idx, row in enumerate(data, 2):
        fill_color = "F2F7FC" if row_idx % 2 == 0 else "FFFFFF"
        row_fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
        for col_idx, val in enumerate(row, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.fill = row_fill
            cell.font = Font(name="メイリオ", size=9)
            cell.alignment = left
            cell.border = border

    col_widths = [30, 45, 16, 20, 30, 18]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.row_dimensions[1].height = 20
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    wb.save(filename)
    print(f"保存完了: {filename}  行数: {len(data)}")


# 第152弾（2026-07-16d）美容クリニック・エステサロンチェーン 10社・LP URLは予約ページ

data_0716d = [
    ["株式会社ビューティースリー", "https://c-3-esthe.com/introduction/reserve/", "", "月額制プレミアム全身脱毛（C3 シースリー）", "全身脱毛 通い放題, シースリー 予約, 脱毛サロン 月額", "脱毛サロンチェーン"],
    ["医療法人社団あおばクリニック", "https://www.aoba-cg.com/before_appointment/web_reservation_info.html", "", "医療脱毛（全身脱毛）", "医療脱毛 安い, あおばクリニック 予約, 全身脱毛 都度払い", "医療脱毛クリニックチェーン"],
    ["FAVORIX BEAUTY株式会社", "https://www.ginza-blv.jp/reserve/", "092-737-2528", "痩身エステ（GINZA BLV）", "痩身エステ 体験, GINZA BLV 予約, ギンザブルー 口コミ", "痩身エステサロンチェーン"],
    ["株式会社シェイプアップハウス", "https://www.dandy-house.co.jp/trialform/", "06-6311-0706", "メンズエステ（ダンディハウス）", "メンズエステ 体験, ダンディハウス 予約, 男性 エステ 痩身", "メンズエステサロンチェーン"],
    ["株式会社ソシエ・ワールド", "https://webreserve.socie.jp/?landCompTicketCd=hp2504001&nav", "03-5843-5840", "痩身・フェイシャルエステ（ソシエ）", "エステ 体験 予約, ソシエ 痩身, フェイシャルエステ", "エステサロンチェーン"],
    ["株式会社シーズ・ラボ", "https://sa.winboard.biz/contact/reserveNew?companyCode=ci_zlabo", "03-3797-4000", "メディカルエステ（シーズ・ラボ）", "シーズラボ 予約, メディカルエステ 毛穴, フェイシャルエステ 体験", "メディカルエステサロンチェーン"],
    ["医療法人社団光芒会", "https://jennyc.jp/reserve_choice/", "045-290-4112", "女性専門医療脱毛（ジェニークリニック）", "医療脱毛 女性専門, ジェニークリニック 予約, 全身脱毛 安い", "医療脱毛クリニック"],
    ["株式会社ビ・メーク", "https://van-veal.com/lp/renew/", "083-974-0588", "ブライダル・トータルエステ（ヴァン・ベール）", "ブライダルエステ 体験, ヴァンベール 予約, エステ 痩身", "トータルエステサロンチェーン"],
    ["株式会社PMKメディカルラボ", "https://www.pmk-j.com/reservation", "03-5363-4421", "バスト・痩身・小顔エステ（PMK）", "痩身エステ 体験, PMK 予約, 小顔エステ", "エステサロンチェーン"],
    ["株式会社アスクビューティー", "https://www.e-defi.com/inquiry/index.html", "06-6371-0601", "トータルエステ（defi デフィー）", "エステ 体験予約, デフィー 口コミ, フェイシャルエステ", "エステサロンチェーン"],
]

base = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー"
make_xlsx(f"{base}\\架電リスト_Meta_2026-07-16d.xlsx", data_0716d)
