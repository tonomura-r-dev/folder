import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUTPUT_PATH = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー\架電リスト_リスティング_2026-06-03bt.xlsx"
HEADER = ["企業名", "LP URL", "広告種別", "法人番号", "商材", "検索KW", "業界"]
COL_WIDTHS = [32, 56, 16, 16, 22, 32, 22]
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="メイリオ", size=10, bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
ROW_FILLS = [PatternFill("solid", fgColor="F2F7FC"), PatternFill("solid", fgColor="FFFFFF")]
ROW_FONT = Font(name="メイリオ", size=9)
ROW_ALIGN = Alignment(vertical="center", wrap_text=False)
URL_FONT = Font(name="メイリオ", size=9, color="0563C1", underline="single")

DATA = [
    # ── SEOツール ──
    ("Mangools",                   "https://mangools.com/pricing",                            "リスティング広告", "", "Mangools・SEOツール・キーワード調査",          "SEO ツール 安い, Mangools 料金 評判, キーワード 調査 ツール",       "SEO SaaS"),
    # ── SEOツール ──
    ("SE Ranking",                 "https://seranking.com/pricing.html",                      "リスティング広告", "", "SE Ranking・SEO/順位計測・サイト監査",         "SEO ツール 比較, SE Ranking 料金 評判, 検索 順位 計測",            "SEO SaaS"),
    # ── 日程調整 ──
    ("SavvyCal",                   "https://savvycal.com/pricing",                            "リスティング広告", "", "SavvyCal・日程調整・カレンダーオーバーレイ",   "日程調整 ツール, SavvyCal 料金 評判, スケジュール 調整 SaaS",      "業務DXSaaS"),
    # ── 日程調整 ──
    ("YouCanBookMe",               "https://youcanbook.me/pricing",                           "リスティング広告", "", "YouCanBookMe・予約/日程調整・低価格",          "予約 日程調整 SaaS, YouCanBookMe 料金 評判, 予約 自動化 ツール",   "業務DXSaaS"),
    # ── 採用管理ATS ──
    ("Workable",                   "https://www.workable.com/pricing",                        "リスティング広告", "", "Workable・採用管理ATS・AIスクリーニング",      "採用管理 ATS 比較, Workable 料金 評判, リクルーティング SaaS",     "HR SaaS"),
    # ── 採用管理ATS ──
    ("Teamtailor",                 "https://www.teamtailor.com/pricing",                      "リスティング広告", "", "Teamtailor・採用管理ATS・採用ブランディング",  "採用管理 ATS SaaS, Teamtailor 料金 評判, 採用 サイト 作成",        "HR SaaS"),
    # ── 採用管理ATS ──
    ("Ashby",                      "https://www.ashbyhq.com/pricing",                         "リスティング広告", "", "Ashby・採用管理ATS・データ分析重視",           "採用管理 ATS データ, Ashby 料金 評判, 採用 分析 SaaS",            "HR SaaS"),
    # ── 採用管理ATS ──
    ("Recruitee",                  "https://recruitee.com/pricing",                           "リスティング広告", "", "Recruitee・採用管理ATS・コラボ採用",           "採用管理 ATS 中小, Recruitee 料金 評判, コラボ 採用 ツール",       "HR SaaS"),
    # ── 人事評価/タレント ──
    ("Lattice",                    "https://lattice.com/pricing",                             "リスティング広告", "", "Lattice・人事評価/タレントマネジメント・HRIS","タレントマネジメント SaaS, Lattice 料金 評判, 人事評価 ツール",   "HR SaaS"),
    # ── 人事評価/エンゲージ ──
    ("15Five",                     "https://www.15five.com/pricing",                          "リスティング広告", "", "15Five・人事評価/エンゲージメント・1on1",      "人事評価 SaaS, 15Five 料金 評判, エンゲージメント サーベイ",       "HR SaaS"),
    # ── 人事評価/学習 ──
    ("Leapsome",                   "https://www.leapsome.com/pricing",                        "リスティング広告", "", "Leapsome・人事評価/OKR/学習/サーベイ",         "人事評価 学習 SaaS, Leapsome 料金 評判, OKR 管理 ツール",          "HR SaaS"),
    # ── HRIS ──
    ("Personio",                   "https://www.personio.com/pricing/",                       "リスティング広告", "", "Personio・HRIS/給与/採用・欧州主要",           "HRIS SaaS 比較, Personio 料金 評判, 人事 管理 プラットフォーム",   "HR SaaS"),
    # ── HRIS ──
    ("Factorial",                  "https://factorialhr.com/pricing",                         "リスティング広告", "", "Factorial・HRIS/勤怠/人事管理",                "HRIS ツール, Factorial 料金 評判, 人事 勤怠 管理 SaaS",            "HR SaaS"),
    # ── HRIS ──
    ("HiBob",                      "https://www.hibob.com/pricing/",                          "リスティング広告", "", "HiBob・HRIS・モダンな人事プラットフォーム",    "HRIS SaaS モダン, HiBob 料金 評判, 人事 管理 クラウド",            "HR SaaS"),
    # ── 採用管理ATS ──
    ("SmartRecruiters",            "https://www.smartrecruiters.com/pricing/",                "リスティング広告", "", "SmartRecruiters・採用管理ATS・大企業向け",     "採用管理 ATS 大企業, SmartRecruiters 料金 評判, 採用 プラットフォーム","HR SaaS"),
    # ── エンゲージメント ──
    ("Culture Amp",               "https://www.cultureamp.com/pricing",                      "リスティング広告", "", "Culture Amp・従業員エンゲージメント/評価",     "従業員 エンゲージメント SaaS, Culture Amp 料金 評判, 組織 サーベイ","HR SaaS"),
    # ── 称賛/報酬 ──
    ("Bonusly",                    "https://bonus.ly/pricing",                                "リスティング広告", "", "Bonusly・社内称賛/ピアボーナス・エンゲージ",   "ピアボーナス SaaS, Bonusly 料金 評判, 社内 称賛 ツール",           "HR SaaS"),
    # ── HRIS ──
    ("Humaans",                    "https://humaans.io/pricing",                              "リスティング広告", "", "Humaans・モダンHRIS・スタートアップ向け",      "HRIS スタートアップ, Humaans 料金 評判, 人事 管理 SaaS",          "HR SaaS"),
    # ── グローバル雇用 ──
    ("Oyster",                     "https://www.oysterhr.com/pricing",                        "リスティング広告", "", "Oyster・グローバル雇用/EOR・コンプラ",         "グローバル 雇用 EOR, Oyster 料金 評判, 海外 採用 SaaS",            "HR SaaS"),
    # ── エンゲージメント/評価 ──
    ("Workleap",                   "https://www.workleap.com/pricing",                        "リスティング広告", "", "Workleap・従業員体験/エンゲージメント/評価",   "従業員 体験 SaaS, Workleap Officevibe 料金, エンゲージメント ツール","HR SaaS"),
]


def make_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "架電リスト"

    for col_idx, header in enumerate(HEADER, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
    ws.row_dimensions[1].height = 20

    for row_idx, row_data in enumerate(DATA, 2):
        fill = ROW_FILLS[(row_idx - 2) % 2]
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.fill = fill
            cell.alignment = ROW_ALIGN
            if col_idx == 2:
                cell.hyperlink = value
                cell.font = URL_FONT
            else:
                cell.font = ROW_FONT
        ws.row_dimensions[row_idx].height = 16

    for col_idx, width in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADER))}1"

    wb.save(OUTPUT_PATH)
    print(f"Saved: {OUTPUT_PATH}")
    print(f"Rows: {len(DATA)}")


if __name__ == "__main__":
    make_xlsx()
