import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUTPUT_PATH = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー\架電リスト_リスティング_2026-06-03ba.xlsx"
HEADER = ["企業名", "LP URL", "広告種別", "法人番号", "商材", "検索KW", "業界"]
COL_WIDTHS = [32, 56, 16, 16, 22, 32, 22]
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="メイリオ", size=10, bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
ROW_FILLS = [PatternFill("solid", fgColor="F2F7FC"), PatternFill("solid", fgColor="FFFFFF")]
ROW_FONT = Font(name="メイリオ", size=9)
ROW_ALIGN = Alignment(vertical="center", wrap_text=False)
URL_FONT = Font(name="メイリオ", size=9, color="0563C1", underline="single")

# 全社 WebSearch実確認済みLP（企業名はカッコなし正式名称のみ）
DATA = [
    # ── 勤怠・給与SaaS ──
    ("株式会社DONUTS",            "https://all.jobcan.ne.jp/",                               "リスティング広告", "", "勤怠管理・給与計算・労務一体型クラウド",    "ジョブカン 勤怠 料金 評判, 給与計算 クラウド 比較, 勤怠 管理 無料 試し",  "HR SaaS"),
    # ── Google Workspace拡張 ──
    ("rakumo株式会社",            "https://rakumo.com/form/trial/",                          "リスティング広告", "", "Google Workspace拡張・ワークフロー・勤怠",  "Google Workspace 拡張 ツール, ワークフロー システム 比較, rakumo 評判",   "グループウェア"),
    # ── 人事労務クラウド ──
    ("jinjer株式会社",            "https://hcm-jinjer.com/campaign009/",                     "リスティング広告", "", "人事労務クラウド・勤怠管理・無料トライアル", "jinjer 勤怠 料金 評判, 人事 労務 クラウド 比較, HR システム 無料 試し",   "HR SaaS"),
    # ── 労務管理SaaS ──
    ("フリー株式会社",            "https://www.freee.co.jp/lp/hr/yell/",                    "リスティング広告", "", "freee人事労務・労務管理・シェアNo.1",       "freee 人事 労務 料金 評判, 労務管理 ソフト 比較, 社会保険 手続き DX",     "HR SaaS"),
    # ── ECプラットフォーム ──
    ("Shopify Japan株式会社",     "https://www.shopify.com/jp",                              "リスティング広告", "", "ネットショップ・ECサイト構築・無料体験",     "Shopify ネットショップ 開設, EC サイト 作り方 比較, 通販 構築 無料",      "ECプラットフォーム"),
    # ── BtoBマーケSaaS ──
    ("株式会社ベーシック",        "https://ferret-one.com/",                                 "リスティング広告", "", "BtoBマーケティングSaaS・LP制作・MA",         "BtoB マーケ ツール 比較, LP 制作 ツール 簡単, ferret One 評判 料金",      "マーケSaaS"),
    # ── 市場調査 ──
    ("ゼネラルリサーチ株式会社",  "https://general-research.co.jp/lp/marketresearch/",       "リスティング広告", "", "ネットリサーチ・アンケート調査・市場調査",   "アンケート調査 依頼 安い, ネット リサーチ 会社 比較, 市場調査 外注 費用", "マーケティングリサーチ"),
    # ── 請求書・バックオフィスAI ──
    ("株式会社LayerX",            "https://bakuraku.jp/lp/ad_invoice/",                      "リスティング広告", "", "AI請求書処理・経費精算・電子帳簿保存",       "請求書 処理 自動化, バクラク 料金 評判, 電子帳簿 保存 対応 ツール",      "バックオフィスSaaS"),
    # ── ウェビナー集客 ──
    ("株式会社マジセミ",          "https://majisemi.com/service/lp/",                        "リスティング広告", "", "IT企業向けウェビナー集客・独自ハウスリスト", "ウェビナー 集客 代行 IT, BtoB セミナー 集客 方法, マジセミ 評判 料金",   "マーケ支援"),
    # ── WAF セキュリティ ──
    ("株式会社サイバーセキュリティクラウド", "https://lp.cscloud.co.jp/ksk/mail/mailmagazine/lp", "リスティング広告", "", "クラウド型WAF・サイバー攻撃対策",        "WAF クラウド 比較 おすすめ, サイバー攻撃 対策 ツール, 攻撃遮断くん 料金", "セキュリティSaaS"),
    # ── 営業リスト・法人DB ──
    ("株式会社SalesNow",          "https://salesnow.jp/lp_c/",                               "リスティング広告", "", "営業リスト・法人データベース・1,400万件",     "営業 リスト 作成 ツール, 法人 データベース 比較, SalesNow 料金 評判",    "セールスSaaS"),
    # ── BIダッシュボード ──
    ("ウイングアーク1st株式会社", "https://lp.wingarc.com/MB/001",                           "リスティング広告", "", "BIツール・ダッシュボード・データ可視化",      "BI ツール 比較 国産, データ 可視化 ソフト, MotionBoard 料金 評判",       "BI SaaS"),
    # ── SaaS比較メディア ──
    ("スマートキャンプ株式会社",  "https://boxil.jp/campaign/partners/",                     "リスティング広告", "", "SaaS比較サイト掲載・リード獲得・BOXIL",      "SaaS 比較 サイト 掲載, BtoB リード 獲得, BOXIL 掲載 費用 評判",         "デジタルマーケ"),
    # ── CRM SaaS ──
    ("ゾーホージャパン株式会社",  "https://www.zoho.com/jp/crm/lp/crm-software-new.html",   "リスティング広告", "", "CRM・MA・無料プランあり",                    "Zoho CRM 料金 評判, CRM ツール 比較 無料, 顧客 管理 SaaS 安い",         "CRM SaaS"),
    # ── MAツール ──
    ("アドビ株式会社",            "https://www.adobe.com/jp/marketing-cloud.html",           "リスティング広告", "", "マーケティングクラウド・データ分析・広告",     "Adobe マーケティング クラウド 料金, デジタル 広告 分析 ツール, MA 比較", "マーケSaaS"),
    # ── 勤怠SaaS（採用管理LP） ──
    ("株式会社DONUTS",            "https://ats.jobcan.ne.jp/lp/",                            "リスティング広告", "", "採用管理システム・ATS・応募から内定まで",     "採用 管理 システム 比較, ATS ツール 無料, ジョブカン 採用 料金",          "HRSaaS"),
    # ── 人事評価SaaS ──
    ("株式会社カオナビ",          "https://www.kaonavi.jp/",                                 "リスティング広告", "", "タレントマネジメント・人材育成・シェアNo.1",  "カオナビ 料金 評判, 人事評価 クラウド 比較, タレントマネジメント 導入",   "HR SaaS"),
    # ── ワークフロー ──
    ("株式会社コラボスタイル",    "https://lp.collabo-style.co.jp/trial.html",               "リスティング広告", "", "ワークフロー・30日無料トライアル",            "ワークフロー システム 比較 無料, 申請 承認 クラウド, コラボフロー 料金",  "業務DXSaaS"),
    # ── 採用メディア ──
    ("株式会社スタンバイ",        "https://jinji.stanby.co.jp/",                             "リスティング広告", "", "求人検索エンジン・採用広告・クリック課金",    "スタンバイ 求人 掲載 料金, 採用 求人 効果 比較, 求人 検索 エンジン",    "採用メディア"),
    # ── 店舗向け通信 ──
    ("株式会社USEN",              "https://usen.com/service/wifi/uhikari/",                  "リスティング広告", "", "店舗向け光回線・Wi-Fi・AIR UNLIMITED",        "店舗 光回線 おすすめ, USEN AIR 評判 料金, 飲食店 Wi-Fi 業務用",          "通信・回線"),
]


def make_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "架電リスト"

    for col_idx, header in enumerate(HEADER, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
    ws.row_dimensions[1].height = 20

    for row_idx, row_data in enumerate(DATA, 2):
        fill = ROW_FILLS[(row_idx - 2) % 2]
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.fill = fill
            cell.alignment = ROW_ALIGN
            if col_idx == 2:
                cell.hyperlink = value
                cell.font = URL_FONT
            else:
                cell.font = ROW_FONT
        ws.row_dimensions[row_idx].height = 16

    for col_idx, width in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADER))}1"

    wb.save(OUTPUT_PATH)
    print(f"Saved: {OUTPUT_PATH}")
    print(f"Rows: {len(DATA)}")


if __name__ == "__main__":
    make_xlsx()
