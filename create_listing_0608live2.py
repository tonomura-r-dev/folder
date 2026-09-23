import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUTPUT_PATH = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー\架電リスト_リスティング_2026-06-08_live2.xlsx"
HEADER = ["企業名", "LP URL", "広告種別", "法人番号", "商材", "検索KW", "業界"]
COL_WIDTHS = [30, 50, 16, 16, 26, 24, 16]
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="メイリオ", size=10, bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
ROW_FILLS = [PatternFill("solid", fgColor="F2F7FC"), PatternFill("solid", fgColor="FFFFFF")]
ROW_FONT = Font(name="メイリオ", size=9)
ROW_ALIGN = Alignment(vertical="center", wrap_text=False)
URL_FONT = Font(name="メイリオ", size=9, color="0563C1", underline="single")

# Google検索の「スポンサー」枠（今配信中の広告）から抽出。比較/ポータル/代理店/使用済み除外。社名カッコなし。
DATA = [
    ("株式会社大塚シロアリ研究所",        "https://www.otsuka-shiroari.com",                  "リスティング広告", "", "シロアリ駆除・防除",            "シロアリ駆除 徳島",   "害虫駆除"),
    ("株式会社友清白蟻",                  "https://www.tomokiyo.co.jp",                       "リスティング広告", "", "シロアリ駆除・防除",            "シロアリ駆除 四国",   "害虫駆除"),
    ("株式会社Enishi",                    "https://www.enishi-tokushima.co.jp",               "リスティング広告", "", "外構・エクステリア工事",        "外構 エクステリア 徳島","外構"),
    ("四電エナジーサービス株式会社",      "https://www.yes-e-life.jp",                        "リスティング広告", "", "給湯器交換・オール電化",        "給湯器交換 徳島",     "住宅設備"),
    ("株式会社Cools",                     "https://ecocute-shikoku.cools-kyutoki.com",        "リスティング広告", "", "エコキュート・給湯器交換",      "給湯器交換 四国",     "住宅設備"),
    ("ボイラサービスコーポレーション株式会社","https://bsc-ecolife.jp/lp/tokushima/",          "リスティング広告", "", "給湯器・ボイラー修理/交換",     "給湯器交換 徳島",     "住宅設備"),
    ("株式会社加藤自動車相談所",          "https://colorful-tokushima.com/",                  "リスティング広告", "", "軽自動車販売（カラフル）",      "中古車 軽自動車 徳島", "自動車販売"),
]


def make_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "架電リスト"
    for col_idx, header in enumerate(HEADER, 1):
        c = ws.cell(row=1, column=col_idx, value=header)
        c.fill = HEADER_FILL; c.font = HEADER_FONT; c.alignment = HEADER_ALIGN
    ws.row_dimensions[1].height = 20
    for row_idx, row in enumerate(DATA, 2):
        fill = ROW_FILLS[(row_idx - 2) % 2]
        for col_idx, value in enumerate(row, 1):
            c = ws.cell(row=row_idx, column=col_idx, value=value)
            c.fill = fill; c.alignment = ROW_ALIGN
            if col_idx == 2:
                c.hyperlink = value; c.font = URL_FONT
            else:
                c.font = ROW_FONT
        ws.row_dimensions[row_idx].height = 16
    for col_idx, width in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADER))}1"
    wb.save(OUTPUT_PATH)
    print(f"Saved: {OUTPUT_PATH}  Rows: {len(DATA)}")


if __name__ == "__main__":
    make_xlsx()
