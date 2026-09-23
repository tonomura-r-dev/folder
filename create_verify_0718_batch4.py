import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUTPUT_PATH = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー\架電リスト_全社再検証_2026-07-18_batch4.xlsx"
HEADER = ["企業名", "LP URL(広告確認済み)", "電話番号", "検索KW", "媒体"]
COL_WIDTHS = [40, 50, 16, 32, 14]
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="メイリオ", size=10, bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
ROW_FILLS = [PatternFill("solid", fgColor="F2F7FC"), PatternFill("solid", fgColor="FFFFFF")]
ROW_FONT = Font(name="メイリオ", size=9)
ROW_ALIGN = Alignment(vertical="center", wrap_text=False)
URL_FONT = Font(name="メイリオ", size=9, color="0563C1", underline="single")

DATA = [
    ("株式会社ミュージックプラネット", "https://music-planet.jp/", "03-6427-0758", "歌手 オーディション, デビュー 音楽", "Meta"),
    ("株式会社RESTA", "https://www.diy-shop.jp/", "078-976-5323", "フローリング DIY, 床材 通販", "Meta"),
    ("ひまわりネットワーク株式会社", "https://www.himawari.co.jp/", "0565-35-3311", "ケーブルテレビ 豊田市, ネット 開通", "Meta"),
    ("株式会社RESコーポレーション", "https://www.res-corp.jp/", "03-5937-4901", "不動産相続 相談, 借地 底地 トラブル", "Meta"),
    ("株式会社ジムフィールド", "https://www.gym-field.com/", "06-7777-1276", "HYROX 体験, パーソナルトレーニング 大阪", "Meta"),
    ("株式会社ロッピングライフ", "https://ropping.jp/", "03-5772-2105", "テレビ通販, ポータブル電源", "Meta"),
    ("株式会社アイシティ", "https://www.eyecity.jp/", "0120-110-027", "コンタクトレンズ 通販, カラコン 安い", "Meta"),
    ("株式会社ジンズホールディングス", "https://www.jins.com/", "03-5275-7001", "JINS メガネ, 眼鏡 おしゃれ", "Meta"),
    ("株式会社インターメスティック", "https://www.zoff.com/", "03-5468-8650", "Zoff メガネ, 眼鏡 即日", "Meta"),
    ("寺田倉庫株式会社", "https://minikura.com/", "03-5479-1611", "トランクルーム 宅配, 荷物 預ける", "Meta"),
]

NOT_VERIFIED = [
    "株式会社クリーマ", "株式会社おそうじ革命", "セントラル警備保障株式会社", "ピクスタ株式会社",
    "株式会社エンスポーツ", "弁護士法人本田総合法律事務所（第三者記事のみ確認）",
    "株式会社エメトレ", "株式会社RSLジャパン", "株式会社MAJIDE",
]


def make_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "確認済み"
    for col_idx, header in enumerate(HEADER, 1):
        c = ws.cell(row=1, column=col_idx, value=header)
        c.fill = HEADER_FILL; c.font = HEADER_FONT; c.alignment = HEADER_ALIGN
    ws.row_dimensions[1].height = 20
    for row_idx, row in enumerate(DATA, 2):
        fill = ROW_FILLS[(row_idx - 2) % 2]
        for col_idx, value in enumerate(row, 1):
            c = ws.cell(row=row_idx, column=col_idx, value=value)
            c.fill = fill; c.alignment = ROW_ALIGN
            if col_idx == 2:
                c.hyperlink = value; c.font = URL_FONT
            else:
                c.font = ROW_FONT
        ws.row_dimensions[row_idx].height = 16
    for col_idx, width in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADER))}1"

    ws2 = wb.create_sheet("未確認")
    ws2.cell(row=1, column=1, value="広告ライブラリで配信中の広告が見つからなかった企業（要再調査・除外候補）").font = Font(bold=True)
    for i, name in enumerate(NOT_VERIFIED, 2):
        ws2.cell(row=i, column=1, value=name)
    ws2.column_dimensions["A"].width = 55

    wb.save(OUTPUT_PATH)
    print(f"Saved: {OUTPUT_PATH}  Confirmed: {len(DATA)}  Not verified: {len(NOT_VERIFIED)}")


if __name__ == "__main__":
    make_xlsx()
