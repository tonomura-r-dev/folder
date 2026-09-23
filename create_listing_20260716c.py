import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def make_xlsx(filename, data):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "架電リスト"

    headers = ["企業名", "LP URL", "電話番号", "商材", "検索KW", "業界"]

    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="メイリオ", bold=True, color="FFFFFF", size=10)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center
        cell.border = border

    for row_idx, row in enumerate(data, 2):
        fill_color = "F2F7FC" if row_idx % 2 == 0 else "FFFFFF"
        row_fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
        for col_idx, val in enumerate(row, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.fill = row_fill
            cell.font = Font(name="メイリオ", size=9)
            cell.alignment = left
            cell.border = border

    col_widths = [30, 45, 16, 20, 30, 18]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.row_dimensions[1].height = 20
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    wb.save(filename)
    print(f"保存完了: {filename}  行数: {len(data)}")


# 第151弾（2026-07-16c）人材紹介・派遣（業界特化型）10社・LP URLは登録ページ

data_0716c = [
    ["株式会社ネクストビート", "https://hoikushibank.com/form/application/step/register", "", "保育士専門転職支援（保育士バンク！）", "保育士 転職, 保育士 求人, 保育士バンク", "保育士業界人材紹介"],
    ["株式会社ギークリー", "https://www.geekly.co.jp/entry/", "03-6418-9113", "IT/Web/ゲーム業界特化転職エージェント（Geekly）", "エンジニア 転職, IT 転職エージェント, Geekly", "IT業界人材紹介"],
    ["クックビズ株式会社", "https://cookbiz.co.jp/entry", "", "飲食・調理師業界特化転職支援（クックビズ）", "料理人 転職, 飲食 求人 正社員, クックビズ", "飲食・調理師業界人材紹介"],
    ["株式会社ドラEVER", "https://doraever.jp/mypage/login#register", "03-6371-4192", "ドライバー専門求人サイト（ドラEVER）", "ドライバー 転職, トラック運転手 求人, ドラEVER", "運送・ドライバー業界人材紹介"],
    ["株式会社トライトキャリア", "https://kaigoworker.jp/entry/site/", "", "介護職専門転職支援（介護ワーカー）", "介護士 転職, 介護 求人, 介護ワーカー", "介護業界人材紹介"],
    ["株式会社ニッソーネット", "https://hoikubatake.jp/entry.php", "", "保育士・幼稚園教諭人材派遣（ほいく畑）", "保育士 派遣, 保育士 求人 派遣, ほいく畑", "保育士人材派遣"],
    ["株式会社セイファート", "https://www.qjnavi.jp/register", "", "美容師専門転職・求人（QJナビ／リクエストQJ）", "美容師 転職, 美容室 求人, QJナビ", "美容師業界人材紹介"],
    ["株式会社レクリー", "https://sekou-kyujin.com/agent/", "", "施工管理・建設業界特化転職エージェント（ジョブリー建設）", "施工管理 転職, 建設 求人 エージェント, ジョブリー建設", "建設業界人材紹介"],
    ["エニーキャリア株式会社", "https://pharmacareer.jp/form/entry", "", "薬剤師専門転職支援（ファーマキャリア）", "薬剤師 転職, 薬剤師 求人, ファーマキャリア", "薬剤師業界人材紹介"],
    ["ロジHR株式会社", "https://logirec.jp/register.html", "", "物流業界特化求人・転職支援（ロジリク）", "物流 転職, 倉庫 求人, ロジリク", "物流業界人材紹介"],
]

base = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー"
make_xlsx(f"{base}\\架電リスト_Meta_2026-07-16c.xlsx", data_0716c)
