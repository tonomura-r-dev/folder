import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUTPUT_PATH = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー\架電リスト_リスティング_2026-05-25b.xlsx"

HEADER = ["企業名", "LP URL", "広告種別", "電話番号", "商材", "検索KW", "業界"]
COL_WIDTHS = [32, 45, 16, 16, 20, 30, 18]

DATA = [
    ("株式会社アットコンタクト（アットコンタクト）", "https://www.at-contact.jp/", "リスティング広告", "", "コンタクトレンズ通販", "コンタクトレンズ 通販 安い, ワンデー コンタクト 購入, コンタクト 最安値", "コンタクトレンズ通販"),
    ("株式会社レンズアップル（レンズアップル）", "https://www.lens-apple.jp/", "リスティング広告", "", "コンタクトレンズ通販", "2weekコンタクト 通販, コンタクト 即日発送, コンタクトレンズ 安い", "コンタクトレンズ通販"),
    ("株式会社デイトラ（デイトラ）", "https://dailytrial.net/reskilling_lp/", "リスティング広告", "", "オンラインプログラミングスクール", "webデザイン 独学, プログラミング 在宅, リスキリング IT転職", "教育・スキルアップ"),
    ("株式会社キカガク（キカガク）", "https://longterm.kikagaku.ai/", "リスティング広告", "", "AI・データサイエンス教育コース", "データサイエンス 学習, AI 資格 取得, 機械学習 入門 講座", "教育・スキルアップ"),
    ("株式会社クリコム（クリコム）", "https://cricom.jp/", "リスティング広告", "", "宅配クリーニングサービス", "宅配クリーニング 月額, クリーニング 宅配 安い, コート クリーニング 宅配", "家事代行・クリーニング"),
    ("有限会社小柴クリーニング（クリーニングパンダ）", "https://panda-cl.com/", "リスティング広告", "", "宅配クリーニング", "布団 クリーニング 宅配, スーツ クリーニング 宅配, 宅配クリーニング 評判", "家事代行・クリーニング"),
    ("株式会社白洋舎", "https://www.hakuyosha.co.jp/", "リスティング広告", "", "クリーニング・宅配クリーニング", "宅配クリーニング 全国, クリーニング 集配, クリーニング 老舗", "家事代行・クリーニング"),
    ("株式会社ロイブ（ホットヨガスタジオ ロイブ）", "https://www.hotyoga-loive.com/", "リスティング広告", "", "ホットヨガスタジオ", "ホットヨガ おしゃれ, ヨガスタジオ 女性, ホットヨガ 体験 無料", "フィットネス・スポーツ・ゴルフ"),
    ("株式会社ゼンゴルフ（ZEN GOLF RANGE）", "https://zengolf.jp/", "リスティング広告", "", "インドアゴルフスクール・シミュレーター", "インドアゴルフ 月額, ゴルフシミュレーター レッスン, ゴルフ 初心者 練習", "フィットネス・スポーツ・ゴルフ"),
    ("RIZAP株式会社（ライザップイングリッシュ）", "https://www.rizap-english.jp/", "リスティング広告", "", "英語コーチングスクール", "英語 コーチング, 短期間 英語上達, ビジネス英語 プログラム", "英会話・語学・資格スクール"),
    ("出張買取サービス 買いクル（株式会社KAIKU）", "https://kaikuru.com/", "リスティング広告", "", "不用品出張買取サービス", "出張買取 即日, 不用品 高価買取, 家電 家具 買取 出張", "買取・リユース・フリマ"),
    ("株式会社プリンスフラワー（プリフラ）", "https://kaitori-prince.com/", "リスティング広告", "", "不用品出張買取", "出張買取 無料, ブランド品 買取 訪問, 高価買取 出張 無料", "買取・リユース・フリマ"),
    ("SBIペット少額短期保険株式会社", "https://www.sbipet-ssi.co.jp/lp/", "リスティング広告", "", "ペット保険", "ペット保険 安い, 犬 保険 通院, ペット保険 プラン比較", "ペット・ペット保険"),
    ("SEモバイル・アンド・オンライン株式会社（スマリッジ）", "https://s-marriage.jp/", "リスティング広告", "", "オンライン結婚相談所", "オンライン 結婚相談所, 婚活 安い 月額, 婚活 真剣 サポート", "婚活・マッチングアプリ"),
    ("万田発酵株式会社（万田酵素）", "https://www.mandahakko.com/", "リスティング広告", "", "発酵酵素健康食品", "万田酵素 定期購入, 酵素 健康食品, 発酵食品 サプリメント", "化粧品・健康食品通販D2C"),
    ("株式会社新谷酵素（新谷酵素 RAKUFAS）", "https://www.shinyakoso.com/product/rakufas/", "リスティング広告", "", "酵素サプリ・ファスティングプログラム", "酵素 断食, ファスティング サプリ, 酵素 ドリンク 通販", "化粧品・健康食品通販D2C"),
    ("TOKYO CRAFTS株式会社（TOKYO CRAFTS）", "https://tokyocrafts.jp/", "リスティング広告", "", "キャンプ用品直販", "キャンプ用品 直販, 焚き火台 通販, アウトドア ブランド 公式", "フィットネス・スポーツ・ゴルフ"),
    ("株式会社ヤマップ（hinataストア）", "https://store.hinata.me/", "リスティング広告", "", "アウトドア・キャンプ用品セレクト通販", "アウトドア用品 通販, キャンプ テント 購入, キャンプギア おすすめ", "フィットネス・スポーツ・ゴルフ"),
    ("ヴィトゥレ（Vitule）", "https://www.vitule.jp/", "リスティング広告", "", "脱毛・痩身エステサロン", "脱毛 エステ 安い, 痩身 無料体験, エステサロン キャビテーション", "美容クリニック・脱毛・エステ"),
    ("株式会社Wellness（あんしん漢方）", "https://www.kamposupport.com/anshin1.0/lp/", "リスティング広告", "", "AI×専門家 オーダーメイド漢方", "漢方 オンライン 相談, 漢方薬 通販 処方, 漢方 定期便 自動", "化粧品・健康食品通販D2C"),
    ("医療法人社団クリニクフォア（クリニクフォア）", "https://www.clinicfor.life/", "リスティング広告", "", "オンライン診療・処方薬宅配", "オンライン 診療 病院, 処方箋 宅配, オンライン診療 皮膚科 予約", "美容クリニック・脱毛・エステ"),
    ("株式会社ツカモトエイム（ペピイ）", "https://www.peppy.co.jp/", "リスティング広告", "", "ペット用品・フード通販", "ペット用品 通販 安い, ドッグフード 通販, ペット 消耗品 まとめ買い", "ペット・ペット保険"),
    ("Photoback株式会社（Photoback）", "https://www.photoback.jp/", "リスティング広告", "", "フォトブック・アルバム作成", "フォトブック 高品質, 写真アルバム 作成 おしゃれ, フォトブック 送料無料", "その他（EC・通販）"),
    ("株式会社クロスエッジ（Dr.つるかめキッチン）", "https://tsurukame-kitchen.com/", "リスティング広告", "", "制限食・健康宅配弁当", "カロリー制限 食事 宅配, 制限食 弁当 通販, 医師監修 宅配弁当", "食品EC・宅配弁当・ミールキット"),
    ("株式会社ユニマット リタイアメント・コミュニティ（食のそよ風）", "https://shokunosoyokaze.com/", "リスティング広告", "", "高齢者向け宅配弁当", "高齢者 宅配弁当, 介護食 宅配, シニア 食事 宅配 毎日", "食品EC・宅配弁当・ミールキット"),
    ("オリックス自動車株式会社（オリックスカーシェア）", "https://carshare.orix.co.jp/", "リスティング広告", "", "カーシェアリングサービス", "カーシェア 月額, カーシェアリング 個人, 車 シェアリング 登録", "車・カーシェア・バイク"),
    ("株式会社subsclife（subsclife）", "https://subsclife.com/", "リスティング広告", "", "家具・家電サブスクリプション", "家具 サブスク, ソファ 月額 レンタル, 家具 定額 利用 インテリア", "その他（EC・通販）"),
    ("株式会社タイヘイ（タイヘイ）", "https://www.taiheiyakuhin.co.jp/", "リスティング広告", "", "制限食・高齢者向け宅食", "塩分制限 食事 宅配, 糖質制限 弁当 宅配, 高齢者 食事 制限食", "食品EC・宅配弁当・ミールキット"),
    ("株式会社Be-Legend（ビーレジェンド）", "https://be-legend.com/", "リスティング広告", "", "プロテイン・スポーツサプリ通販", "プロテイン 通販 安い, ホエイプロテイン 国産, プロテイン 筋トレ 初心者", "化粧品・健康食品通販D2C"),
    ("オンラインコンタクト株式会社（オンラインコンタクト）", "https://www.online-contact.cc/", "リスティング広告", "", "コンタクトレンズ通販", "コンタクト 最安値 通販, 2weekコンタクト 安い, コンタクト 定期購入", "コンタクトレンズ通販"),
]

def make_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "架電リスト"

    header_font = Font(name="メイリオ", size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    fill_odd  = PatternFill("solid", fgColor="F2F7FC")
    fill_even = PatternFill("solid", fgColor="FFFFFF")
    data_font = Font(name="メイリオ", size=9)
    data_align = Alignment(vertical="center", wrap_text=False)
    url_font  = Font(name="メイリオ", size=9, color="0563C1", underline="single")

    thin_side = Side(style="thin", color="D0D7DE")
    thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    for col_idx, col_name in enumerate(HEADER, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    for row_idx, row_data in enumerate(DATA, start=2):
        fill = fill_odd if (row_idx % 2 == 0) else fill_even
        for col_idx, value in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.fill = fill
            cell.border = thin_border
            cell.alignment = data_align
            if col_idx == 2 and value:
                cell.hyperlink = value
                cell.font = url_font
            else:
                cell.font = data_font

    for col_idx, width in enumerate(COL_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.row_dimensions[1].height = 22
    for row_idx in range(2, len(DATA) + 2):
        ws.row_dimensions[row_idx].height = 18

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADER))}1"

    wb.save(OUTPUT_PATH)
    print(f"保存完了: {OUTPUT_PATH}")
    print(f"件数: {len(DATA)}社")

if __name__ == "__main__":
    make_xlsx()
