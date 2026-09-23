import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUTPUT_PATH = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー\架電リスト_リスティング_2026-05-29i.xlsx"

HEADER = ["企業名", "LP URL", "広告種別", "電話番号", "商材", "検索KW", "業界"]
COL_WIDTHS = [32, 50, 16, 16, 20, 32, 22]

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="メイリオ", size=10, bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)

ROW_FILLS = [PatternFill("solid", fgColor="F2F7FC"), PatternFill("solid", fgColor="FFFFFF")]
ROW_FONT = Font(name="メイリオ", size=9)
ROW_ALIGN = Alignment(vertical="center", wrap_text=False)
URL_FONT = Font(name="メイリオ", size=9, color="0563C1", underline="single")

ROWS = [
    ["株式会社ラグザス・クリエイト",    "https://ninjacode.work/lp/reskilling",                                 "リスティング広告", "",  "プログラミングスクール・リスキリング",   "エンジニア 転職 プログラミング, リスキリング 給付金, IT 未経験 転職",   "IT教育"],
    ["株式会社FURDI",                   "https://furdi.jp/lp/lp06/",                                            "リスティング広告", "",  "女性専用フィットネスジム",               "女性専用 ジム 入会, パーソナルジム 女性, フィットネス レディース",      "フィットネス"],
    ["株式会社サンクチュアリゴルフ",    "https://www.sanctuarygolf.jp/lp/",                                     "リスティング広告", "",  "ゴルフスクール",                         "ゴルフ スクール 体験, 初心者 ゴルフ 習う, ゴルフレッスン 予約",        "ゴルフ"],
    ["株式会社サークフルエステート",    "https://www.surcful.co.jp/lp/",                                        "リスティング広告", "",  "不動産投資コンサル",                     "不動産投資 相談, マンション 投資 始める, 資産形成 不動産",             "不動産投資"],
    ["株式会社HAL",                     "https://hal-tanteisya.com/lp/uwaki",                                   "リスティング広告", "",  "探偵・浮気調査",                         "探偵 浮気調査 費用, 不貞 証拠 つかむ, 素行調査 依頼",                 "探偵"],
    ["株式会社日本エコシステム",        "https://www.shouene.com/lp/solar-power-reduces-electricity-bills/",    "リスティング広告", "",  "太陽光発電システム",                     "太陽光 発電 設置 費用, ソーラーパネル 補助金, 電気代 削減 太陽光",    "省エネ・太陽光"],
    ["株式会社あいプラン",              "https://www.gojyokai.co.jp/lp/index.html",                             "リスティング広告", "",  "冠婚葬祭互助会",                         "互助会 積立 比較, 冠婚葬祭 費用 備え, 葬儀 準備 互助会",              "冠婚葬祭"],
    ["医療法人社団風林会",              "https://www.mens-rize.com/lp/01/",                                     "リスティング広告", "",  "メンズ医療脱毛",                         "メンズ 医療脱毛 安い, 男性 脱毛 クリニック, ひげ脱毛 全身",           "美容クリニック"],
    ["株式会社KIREI produce",           "https://www.osoujikakumei.jp/lp/",                                     "リスティング広告", "",  "ハウスクリーニング",                     "ハウスクリーニング 業者 おすすめ, 部屋 清掃 プロ, エアコン 内部洗浄",  "清掃サービス"],
    ["株式会社アップル",                "https://www.apple-hikkoshi.co.jp/lp/speed-reservation/",               "リスティング広告", "",  "引越しサービス",                         "引越し 業者 安い 口コミ, 見積もり 比較, 単身 引越し 費用",            "引越し"],
    ["みんなのマーケット株式会社",      "https://curama.jp/lp/shop/housekeeping/",                              "リスティング広告", "",  "家事代行マッチング（くらしのマーケット）", "家事代行 おすすめ, 家政婦 派遣 料金, 掃除 代行 単発",                "マッチング・家事代行"],
    ["株式会社グラングレス",            "https://member.rcawaii.com/lp/new",                                    "リスティング広告", "",  "ファッションレンタル（rcawaii）",         "ファッション レンタル 定額, 服 サブスク 月額, コーデ 提案 レンタル",  "ファッションレンタル"],
    ["ウォータースタンド株式会社",      "https://waterstand.jp/lp/biz_cm2025/",                                 "リスティング広告", "",  "水道直結ウォーターサーバー",             "ウォーターサーバー 水道直結, 浄水 サーバー 月額, 給水機 法人",        "浄水サービス"],
    ["株式会社トライグループ",          "https://www.trygroup.co.jp/lp/guarantee/",                             "リスティング広告", "",  "家庭教師（成績保証）",                   "家庭教師 成績保証, 派遣 家庭教師 料金, 受験 家庭教師 個別",           "教育"],
    ["株式会社やる気スイッチグループ",  "https://www.schoolie-net.jp/fs/lp/lp001/taiken/",                      "リスティング広告", "",  "個別指導塾（スクールIE）体験",           "スクールIE 体験授業, 個別指導 塾 無料体験, 学習塾 近く 申込",         "学習塾"],
    ["アニコム損害保険株式会社",        "https://www.anicom-sompo.co.jp/lp/bridge/",                            "リスティング広告", "",  "ペット保険（アニコム）",                 "ペット保険 比較 口コミ, 犬 猫 医療保険, アニコム 申込 資料",          "ペット保険"],
    ["医療法人社団エムズ",              "https://www.clinicfor.life/lp/online-insurance/",                      "リスティング広告", "",  "オンライン診療（clinicfor）",             "オンライン診療 予約, 内科 スマホ 診察, 保険適用 クリニック",          "医療"],
    ["株式会社FPパートナー",            "https://fp-moneydoctor.com/lp/mdmk005/",                               "リスティング広告", "",  "FP相談・保険見直し",                     "FP 無料 相談, 保険 見直し ファイナンシャルプランナー, 老後 資産 相談", "保険代理"],
    ["株式会社ETERNAL",                 "https://hoken-eshop.com/lp/how-to-join_online-consult/",               "リスティング広告", "",  "保険相談オンライン",                     "保険 相談 無料 オンライン, 生命保険 見直し 専門家, 保険 比較 申込",   "保険代理"],
    ["株式会社アラジン",                "https://www.rescue-center.jp/lp/all_media/index.html",                 "リスティング広告", "",  "データ復旧サービス",                     "データ復旧 業者 おすすめ, HDD データ救出 料金, パソコン データ復元",   "データ復旧"],
]

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "架電リスト"

# ヘッダー行
for col, (h, w) in enumerate(zip(HEADER, COL_WIDTHS), 1):
    cell = ws.cell(row=1, column=col, value=h)
    cell.fill = HEADER_FILL
    cell.font = HEADER_FONT
    cell.alignment = HEADER_ALIGN
    ws.column_dimensions[get_column_letter(col)].width = w

ws.row_dimensions[1].height = 20

# データ行
for row_idx, row_data in enumerate(ROWS, 2):
    fill = ROW_FILLS[(row_idx - 2) % 2]
    for col_idx, value in enumerate(row_data, 1):
        cell = ws.cell(row=row_idx, column=col_idx, value=value)
        cell.fill = fill
        cell.alignment = ROW_ALIGN
        if col_idx == 2 and value:
            cell.font = URL_FONT
            cell.hyperlink = value
        else:
            cell.font = ROW_FONT
    ws.row_dimensions[row_idx].height = 16

ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADER))}1"

wb.save(OUTPUT_PATH)
print(f"保存完了: {OUTPUT_PATH}")
print(f"行数: {len(ROWS)}社")
