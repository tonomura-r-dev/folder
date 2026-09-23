import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUTPUT_PATH = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー\架電リスト_リスティング_2026-06-03by.xlsx"
HEADER = ["企業名", "LP URL", "広告種別", "法人番号", "商材", "検索KW", "業界"]
COL_WIDTHS = [32, 56, 16, 16, 22, 32, 22]
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="メイリオ", size=10, bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
ROW_FILLS = [PatternFill("solid", fgColor="F2F7FC"), PatternFill("solid", fgColor="FFFFFF")]
ROW_FONT = Font(name="メイリオ", size=9)
ROW_ALIGN = Alignment(vertical="center", wrap_text=False)
URL_FONT = Font(name="メイリオ", size=9, color="0563C1", underline="single")

DATA = [
    # ── フードデリバリー加盟店募集（日本・実LP確認） ──
    ("株式会社出前館",             "https://service.demae-can.co.jp/restaurant/",            "リスティング広告", "", "出前館・飲食店向け加盟店募集・初期費用0円",     "出前館 出店 費用, フードデリバリー 加盟店 募集, 出前館 加盟",       "フードデリバリー"),
    # ── 期間工求人（日本・実LP確認） ──
    ("株式会社BREXA Next",         "https://kikankou.jp/",                                   "リスティング広告", "", "期間工.jp・大手メーカー期間従業員求人",         "期間工 求人 寮費無料, 期間工.jp 評判, 自動車 工場 期間従業員",     "人材・求人"),
    # ── 治験ボランティア募集（日本・実LP確認） ──
    ("株式会社インクロム",         "https://www.196189.com/lp_healthy/index_A.php",          "リスティング広告", "", "治験ボランティア募集・健康な成人向け",          "治験 ボランティア 募集, 治験 バイト 高額, インクロム 治験",       "治験・CRO"),
    # ── 保育士転職（日本・実LP確認） ──
    ("株式会社ニッソーネット",     "https://hoikubatake.jp/lp-careerup2/",                   "リスティング広告", "", "ほいく畑・保育士求人/転職・未経験可",           "保育士 求人 未経験, ほいく畑 評判, 保育士 転職 サポート",         "人材・求人"),
    # ── カスタマーサポート（OSS） ──
    ("Chatwoot",                   "https://www.chatwoot.com/pricing/",                      "リスティング広告", "", "Chatwoot・OSSカスタマーサポート/ライブチャット","カスタマーサポート OSS, Chatwoot 料金 評判, ライブ チャット SaaS","カスタマーサポートSaaS"),
    # ── B2Bサポート ──
    ("Pylon",                      "https://www.usepylon.com/pricing",                       "リスティング広告", "", "Pylon・B2Bカスタマーサポート・AI",            "B2B カスタマーサポート SaaS, Pylon 料金 評判, AI サポート ツール", "カスタマーサポートSaaS"),
    # ── AIスケジューリング ──
    ("Reclaim",                    "https://reclaim.ai/pricing",                             "リスティング広告", "", "Reclaim・AIカレンダー/自動スケジューリング",   "AI スケジュール SaaS, Reclaim 料金 評判, カレンダー 自動 調整",    "業務DXSaaS"),
    # ── AIカレンダー/タスク ──
    ("Motion",                     "https://www.usemotion.com/pricing",                      "リスティング広告", "", "Motion・AIカレンダー/タスク自動化",            "AI タスク カレンダー SaaS, Motion 料金 評判, AI 予定 管理",        "業務DXSaaS"),
    # ── デイリープランナー ──
    ("Sunsama",                    "https://www.sunsama.com/pricing",                        "リスティング広告", "", "Sunsama・デイリープランナー/タイムボックス",   "デイリー プランナー SaaS, Sunsama 料金 評判, 時間 管理 ツール",    "業務DXSaaS"),
    # ── タスク/時間管理 ──
    ("Akiflow",                    "https://akiflow.com/pricing",                            "リスティング広告", "", "Akiflow・タスク/時間管理・タイムブロッキング", "タスク 時間 管理 SaaS, Akiflow 料金 評判, タイムブロッキング",     "業務DXSaaS"),
    # ── 高速メール ──
    ("Superhuman",                 "https://superhuman.com/plans",                           "リスティング広告", "", "Superhuman・高速メールクライアント・AI",        "メール クライアント SaaS, Superhuman 料金 評判, 高速 メール AI",   "業務DXSaaS"),
    # ── AIメール ──
    ("Shortwave",                  "https://www.shortwave.com/pricing/",                     "リスティング広告", "", "Shortwave・AIメール自動化・受信整理",          "AI メール SaaS, Shortwave 料金 評判, メール 自動化 ツール",        "業務DXSaaS"),
    # ── 共有受信トレイ ──
    ("Missive",                    "https://missiveapp.com/pricing",                         "リスティング広告", "", "Missive・共有受信トレイ/チームメール",         "共有 受信トレイ SaaS, Missive 料金 評判, チーム メール 管理",      "コラボSaaS"),
    # ── 非同期チャット ──
    ("Twist",                      "https://twist.com/pricing",                              "リスティング広告", "", "Twist・非同期チームチャット・スレッド型",       "非同期 チャット SaaS, Twist 料金 評判, スレッド型 チーム チャット","コラボSaaS"),
    # ── チームチャット ──
    ("Pumble",                     "https://pumble.com/pricing",                             "リスティング広告", "", "Pumble・チームチャット・無制限無料",           "チーム チャット 無料 SaaS, Pumble 料金 評判, Slack 代替 ツール",   "コラボSaaS"),
    # ── チームチャット（OSS） ──
    ("Rocket.Chat",                "https://www.rocket.chat/pricing",                        "リスティング広告", "", "Rocket.Chat・OSSセキュアチームチャット",       "OSS チーム チャット, Rocket.Chat 料金 評判, セキュア チャット SaaS","コラボSaaS"),
    # ── チームチャット ──
    ("Zulip",                      "https://zulip.com/plans/",                               "リスティング広告", "", "Zulip・スレッド整理型チームチャット・OSS",     "チーム チャット OSS, Zulip 料金 評判, スレッド型 チャット ツール", "コラボSaaS"),
    # ── セキュアチャット（Matrix） ──
    ("Element",                    "https://element.io/en/pricing",                          "リスティング広告", "", "Element・Matrixベースセキュアチャット",        "セキュア チャット SaaS, Element 料金 評判, Matrix メッセージング",  "コラボSaaS"),
    # ── タスク管理 ──
    ("Todoist",                    "https://www.todoist.com/pricing",                        "リスティング広告", "", "Todoist・タスク管理/ToDoリスト",              "タスク 管理 アプリ, Todoist 料金 評判, ToDo リスト ツール",        "業務DXSaaS"),
    # ── タスク管理 ──
    ("TickTick",                   "https://ticktick.com/upgrade",                           "リスティング広告", "", "TickTick・タスク/カレンダー/習慣管理",         "タスク 管理 アプリ, TickTick 料金 評判, ToDo カレンダー 習慣",     "業務DXSaaS"),
]


def make_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "架電リスト"

    for col_idx, header in enumerate(HEADER, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
    ws.row_dimensions[1].height = 20

    for row_idx, row_data in enumerate(DATA, 2):
        fill = ROW_FILLS[(row_idx - 2) % 2]
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.fill = fill
            cell.alignment = ROW_ALIGN
            if col_idx == 2:
                cell.hyperlink = value
                cell.font = URL_FONT
            else:
                cell.font = ROW_FONT
        ws.row_dimensions[row_idx].height = 16

    for col_idx, width in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADER))}1"

    wb.save(OUTPUT_PATH)
    print(f"Saved: {OUTPUT_PATH}")
    print(f"Rows: {len(DATA)}")


if __name__ == "__main__":
    make_xlsx()
