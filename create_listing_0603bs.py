import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUTPUT_PATH = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー\架電リスト_リスティング_2026-06-03bs.xlsx"
HEADER = ["企業名", "LP URL", "広告種別", "法人番号", "商材", "検索KW", "業界"]
COL_WIDTHS = [32, 56, 16, 16, 22, 32, 22]
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="メイリオ", size=10, bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
ROW_FILLS = [PatternFill("solid", fgColor="F2F7FC"), PatternFill("solid", fgColor="FFFFFF")]
ROW_FONT = Font(name="メイリオ", size=9)
ROW_ALIGN = Alignment(vertical="center", wrap_text=False)
URL_FONT = Font(name="メイリオ", size=9, color="0563C1", underline="single")

DATA = [
    # ── AIコードエディタ ──
    ("Cursor",                     "https://cursor.com/pricing",                              "リスティング広告", "", "Cursor・AIコードエディタ・開発支援",          "AI コード エディタ, Cursor 料金 評判, AI 開発 ツール",             "開発者SaaS"),
    # ── AIアプリ開発 ──
    ("Lovable",                    "https://lovable.dev/pricing",                             "リスティング広告", "", "Lovable・AIアプリ開発・ノーコード生成",        "AI アプリ 開発 ツール, Lovable 料金 評判, AI ノーコード 生成",     "開発者SaaS"),
    # ── オンライン開発環境 ──
    ("Replit",                     "https://replit.com/pricing",                              "リスティング広告", "", "Replit・AI開発環境・ブラウザ開発",            "オンライン 開発 環境, Replit 料金 評判, AI コーディング ツール",   "開発者SaaS"),
    # ── ホスティング ──
    ("Netlify",                    "https://www.netlify.com/pricing/",                        "リスティング広告", "", "Netlify・フロントエンドホスティング/デプロイ", "フロントエンド ホスティング, Netlify 料金 評判, Web デプロイ SaaS","クラウドSaaS"),
    # ── BaaS ──
    ("Supabase",                   "https://supabase.com/pricing",                            "リスティング広告", "", "Supabase・オープンソースBaaS・DB/認証",        "BaaS ツール 比較, Supabase 料金 評判, Firebase 代替 SaaS",         "開発者SaaS"),
    # ── サーバーレスDB ──
    ("Neon",                       "https://neon.tech/pricing",                               "リスティング広告", "", "Neon・サーバーレスPostgres・DB",              "サーバーレス DB SaaS, Neon 料金 評判, Postgres クラウド",          "開発者SaaS"),
    # ── クラウドデプロイ ──
    ("Render",                     "https://render.com/pricing",                              "リスティング広告", "", "Render・クラウドアプリホスティング/デプロイ", "クラウド デプロイ SaaS, Render 料金 評判, アプリ ホスティング",    "クラウドSaaS"),
    # ── クラウドデプロイ ──
    ("Railway",                    "https://railway.com/pricing",                             "リスティング広告", "", "Railway・インフラデプロイ・開発者向け",        "アプリ デプロイ SaaS, Railway 料金 評判, インフラ ホスティング",   "クラウドSaaS"),
    # ── ヘッドレスCMS ──
    ("Sanity",                     "https://www.sanity.io/pricing",                           "リスティング広告", "", "Sanity・ヘッドレスCMS・コンテンツ基盤",        "ヘッドレス CMS 比較, Sanity 料金 評判, コンテンツ 管理 SaaS",     "CMS SaaS"),
    # ── ヘッドレスCMS ──
    ("Storyblok",                  "https://www.storyblok.com/pricing",                       "リスティング広告", "", "Storyblok・ヘッドレスCMS・ビジュアル編集",     "ヘッドレス CMS SaaS, Storyblok 料金 評判, コンテンツ 管理 API",    "CMS SaaS"),
    # ── ビジュアルCMS ──
    ("Builder.io",                 "https://www.builder.io/m/pricing",                        "リスティング広告", "", "Builder.io・ビジュアルCMS/開発・AI",          "ビジュアル CMS SaaS, Builder.io 料金 評判, ノーコード CMS",        "CMS SaaS"),
    # ── 機能フラグ/実験 ──
    ("Statsig",                    "https://statsig.com/pricing",                             "リスティング広告", "", "Statsig・機能フラグ/A/Bテスト・実験基盤",      "機能 フラグ SaaS, Statsig 料金 評判, A/B テスト 実験 ツール",      "プロダクト分析SaaS"),
    # ── 機能フラグ ──
    ("LaunchDarkly",               "https://launchdarkly.com/pricing/",                       "リスティング広告", "", "LaunchDarkly・機能フラグ管理・リリース制御",   "機能 フラグ 管理 SaaS, LaunchDarkly 料金 評判, フィーチャー フラグ","開発者SaaS"),
    # ── 作図/ホワイトボード ──
    ("Whimsical",                  "https://whimsical.com/pricing",                           "リスティング広告", "", "Whimsical・作図/ホワイトボード/ドキュメント", "作図 ツール 比較, Whimsical 料金 評判, ホワイトボード SaaS",       "コラボSaaS"),
    # ── ユーザーリサーチ ──
    ("Maze",                       "https://maze.co/pricing/",                                "リスティング広告", "", "Maze・ユーザーリサーチ/ユーザビリティテスト", "ユーザー リサーチ SaaS, Maze 料金 評判, ユーザビリティ テスト",    "プロダクトSaaS"),
    # ── ユーザーリサーチ ──
    ("Dovetail",                   "https://dovetail.com/pricing/",                           "リスティング広告", "", "Dovetail・ユーザーリサーチ/インサイト管理",   "リサーチ インサイト SaaS, Dovetail 料金 評判, 定性 分析 ツール",   "プロダクトSaaS"),
    # ── メッセージング/MA ──
    ("Customer.io",                "https://customer.io/pricing",                             "リスティング広告", "", "Customer.io・メッセージング/マーケ自動化",     "メッセージング SaaS, Customer.io 料金 評判, マーケ 自動化 ツール", "マーケSaaS"),
    # ── メール配信API ──
    ("Resend",                     "https://resend.com/pricing",                              "リスティング広告", "", "Resend・開発者向けメールAPI・配信",           "メール API 開発者, Resend 料金 評判, トランザクション メール SaaS","コミュニケーションAPI"),
    # ── ニュースレター ──
    ("beehiiv",                    "https://www.beehiiv.com/pricing",                         "リスティング広告", "", "beehiiv・ニュースレター/メディア運営",         "ニュースレター プラットフォーム, beehiiv 料金 評判, メルマガ 配信","メディアSaaS"),
    # ── パブリッシング ──
    ("Ghost",                      "https://ghost.org/pricing/",                              "リスティング広告", "", "Ghost・パブリッシング/ニュースレター・OSS",   "パブリッシング SaaS, Ghost 料金 評判, ブログ メディア 運営 OSS",  "メディアSaaS"),
]


def make_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "架電リスト"

    for col_idx, header in enumerate(HEADER, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
    ws.row_dimensions[1].height = 20

    for row_idx, row_data in enumerate(DATA, 2):
        fill = ROW_FILLS[(row_idx - 2) % 2]
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.fill = fill
            cell.alignment = ROW_ALIGN
            if col_idx == 2:
                cell.hyperlink = value
                cell.font = URL_FONT
            else:
                cell.font = ROW_FONT
        ws.row_dimensions[row_idx].height = 16

    for col_idx, width in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADER))}1"

    wb.save(OUTPUT_PATH)
    print(f"Saved: {OUTPUT_PATH}")
    print(f"Rows: {len(DATA)}")


if __name__ == "__main__":
    make_xlsx()
