import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUTPUT_PATH = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー\架電リスト_全社再検証_2026-07-21_batch6.xlsx"
HEADER = ["企業名", "LP URL(広告確認済み)", "電話番号", "検索KW", "媒体"]
COL_WIDTHS = [40, 50, 16, 32, 14]
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="メイリオ", size=10, bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
ROW_FILLS = [PatternFill("solid", fgColor="F2F7FC"), PatternFill("solid", fgColor="FFFFFF")]
ROW_FONT = Font(name="メイリオ", size=9)
ROW_ALIGN = Alignment(vertical="center", wrap_text=False)
URL_FONT = Font(name="メイリオ", size=9, color="0563C1", underline="single")

# Meta広告ライブラリで配信中の広告を個別確認できた企業のみ。URLは広告に表示されていた遷移先。
DATA = [
    ("株式会社ノバレーゼ", "https://dress.novarese.jp/", "03-5524-3344", "ウェディングドレス, 結婚式場", "Meta"),
    ("株式会社Francfranc", "https://francfranc.com/", "03-4216-4021", "フランフラン, インテリア雑貨", "Meta"),
    ("ベルトラ株式会社", "https://www.veltra.com/", "03-6823-7990", "海外 オプショナルツアー, 体験 予約", "Meta"),
    ("ハウステンボス株式会社", "https://www.huistenbosch.co.jp/", "0570-064-110", "ハウステンボス, テーマパーク 長崎", "Meta"),
    ("学校法人大原学園（資格の大原）", "https://www.o-hara.ac.jp/", "03-3292-6266", "資格の大原 パススル, 資格 スキマ学習", "Meta"),
    ("株式会社ドリームファクトリー（ドクターエア）", "https://www.dr-air.com/", "06-6346-1170", "ドクターエア, マッサージシート", "Meta"),
    ("株式会社DECENCIA（ディセンシア）", "https://www.decencia.co.jp/", "03-3494-1570", "ディセンシア, 敏感肌 シワ改善", "Meta"),
    ("株式会社PETOKOTO（ペトコトフーズ）", "https://foods.petokoto.com/", "03-6555-2833", "ペトコトフーズ, フレッシュペットフード", "Meta"),
]

NOT_VERIFIED = [
    "株式会社タカミ（タカミブライダル）※他社広告内での言及のみ",
    "株式会社ピーチ・ジョン",
    "株式会社デサント",
    "株式会社スクー（schoo）",
    "ユーデミー株式会社（Udemy）※日本向け自社広告なし",
    "株式会社ちふれ化粧品",
]


def make_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "確認済み"
    for col_idx, header in enumerate(HEADER, 1):
        c = ws.cell(row=1, column=col_idx, value=header)
        c.fill = HEADER_FILL; c.font = HEADER_FONT; c.alignment = HEADER_ALIGN
    ws.row_dimensions[1].height = 20
    for row_idx, row in enumerate(DATA, 2):
        fill = ROW_FILLS[(row_idx - 2) % 2]
        for col_idx, value in enumerate(row, 1):
            c = ws.cell(row=row_idx, column=col_idx, value=value)
            c.fill = fill; c.alignment = ROW_ALIGN
            if col_idx == 2:
                c.hyperlink = value; c.font = URL_FONT
            else:
                c.font = ROW_FONT
        ws.row_dimensions[row_idx].height = 16
    for col_idx, width in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADER))}1"

    ws2 = wb.create_sheet("未確認")
    ws2.cell(row=1, column=1, value="広告ライブラリで配信中の広告が見つからなかった企業（要再調査・除外候補）").font = Font(bold=True)
    for i, name in enumerate(NOT_VERIFIED, 2):
        ws2.cell(row=i, column=1, value=name)
    ws2.column_dimensions["A"].width = 55

    wb.save(OUTPUT_PATH)
    print(f"Saved: {OUTPUT_PATH}  Confirmed: {len(DATA)}  Not verified: {len(NOT_VERIFIED)}")


if __name__ == "__main__":
    make_xlsx()
