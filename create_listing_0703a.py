import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUTPUT_PATH = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー\架電リスト_リスティング_2026-07-03a.xlsx"
HEADER = ["企業名", "LP URL", "広告種別", "法人番号", "商材", "検索KW", "業界"]
COL_WIDTHS = [32, 56, 16, 16, 22, 32, 22]
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="メイリオ", size=10, bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
ROW_FILLS = [PatternFill("solid", fgColor="F2F7FC"), PatternFill("solid", fgColor="FFFFFF")]
ROW_FONT = Font(name="メイリオ", size=9)
ROW_ALIGN = Alignment(vertical="center", wrap_text=False)
URL_FONT = Font(name="メイリオ", size=9, color="0563C1", underline="single")

DATA = [
    ("Dual Life Partners株式会社", "https://paytoday.jp/", "リスティング広告", "", "AIファクタリングPAYTODAY", "ファクタリング, 即日 資金調達, 請求書 買取", "ファクタリング"),
    ("株式会社アクティブサポート", "https://ququmo.net/", "リスティング広告", "", "オンラインファクタリングQuQuMo", "ファクタリング 個人事業主, 即日 ファクタリング", "ファクタリング"),
    ("A-LIFE株式会社", "https://endeal.net/", "リスティング広告", "", "遺品整理・生前整理エンディール", "遺品整理 大阪, 生前整理 見積もり", "遺品整理"),
    ("株式会社アートスペース", "https://www.lp.art-sp.com/", "リスティング広告", "", "不動産投資セミナー", "不動産投資 セミナー, マンション投資", "不動産投資"),
    ("株式会社アールケイエンタープライズ", "https://kaitori.rodeodrive.co.jp/", "リスティング広告", "", "ブランド買取ロデオドライブ", "ブランド 買取, ロレックス 買取", "リユース・買取"),
    ("株式会社ブランド買取ブランドハンズ", "https://brand-hands.co.jp/takuhai-kaitori/", "リスティング広告", "", "ブランド品 宅配買取", "ブランド 高価買取, 宅配買取", "リユース・買取"),
    ("First fit株式会社", "https://the-personal-gym.com/", "リスティング広告", "", "THE PERSONAL GYM", "パーソナルジム, パーソナルトレーニング", "パーソナルジム"),
    ("株式会社シスモール", "https://rinshosiken.com/", "リスティング広告", "", "治験モニター募集コーメディカルクラブ", "治験 モニター, 治験 ボランティア", "治験・臨床試験"),
    ("株式会社FiiT", "https://dietpartner.jp/", "リスティング広告", "", "ダイエットパートナー", "パーソナルジム 安い, ダイエット ジム", "パーソナルジム"),
    ("エクササイズコーチ株式会社", "https://reborn-myself.com/", "リスティング広告", "", "女性専用ジム リボーンマイセルフ", "女性専用 ジム, パーソナルジム 女性", "パーソナルジム"),
    ("株式会社ライズ", "https://www.p-gyms.jp/", "リスティング広告", "", "女性専用パーソナルジムGYMS", "女性専用 パーソナルジム, ダイエット 女性", "パーソナルジム"),
    ("株式会社ハッピーカーズ", "https://happycars.jp/", "リスティング広告", "", "出張車買取ハッピーカーズ", "車 買取, 車 出張査定", "車買取"),
    ("株式会社アユザック", "https://whiteningcafe.jp/", "リスティング広告", "", "セルフホワイトニングカフェ", "セルフホワイトニング, ホワイトニング サロン", "ホワイトニング"),
    ("株式会社日本住宅管理", "https://www.njk-reform.jp/", "リスティング広告", "", "水回りリフォーム専門", "水回り リフォーム, トイレ 交換", "リフォーム"),
    ("株式会社シャリオン", "https://charion.co.jp/", "リスティング広告", "", "美歯口 セルフホワイトニング導入", "セルフホワイトニング 開業, ホワイトニング FC", "ホワイトニング・FC"),
    ("株式会社アロン", "https://al-on.com/", "リスティング広告", "", "ハウスクリーニングのオン", "ハウスクリーニング, エアコンクリーニング", "ハウスクリーニング"),
    ("株式会社リライフ", "https://www.relife-r.co.jp/", "リスティング広告", "", "総合リフォーム(水回り・内外装)", "リフォーム 世田谷, 水回り リフォーム", "リフォーム"),
    ("株式会社エクセル・クリーン・サービス", "https://machino-ihinseiri.com/", "リスティング広告", "", "まちの遺品整理屋さん", "遺品整理 大阪, 不用品回収", "遺品整理"),
    ("株式会社カービューティーアイアイシー", "https://www.pro-iic.com/", "リスティング広告", "", "カーコーティング専門IIC", "カーコーティング, ガラスコーティング", "カーコーティング"),
    ("株式会社ポリッシュファクトリー", "https://polishfactory.jp/", "リスティング広告", "", "カーコーティング・車磨き専門", "カーコーティング 東京, セラミックコーティング", "カーコーティング"),
]


def make_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "架電リスト"
    for col_idx, header in enumerate(HEADER, 1):
        c = ws.cell(row=1, column=col_idx, value=header)
        c.fill = HEADER_FILL; c.font = HEADER_FONT; c.alignment = HEADER_ALIGN
    ws.row_dimensions[1].height = 20
    for row_idx, row in enumerate(DATA, 2):
        fill = ROW_FILLS[(row_idx - 2) % 2]
        for col_idx, value in enumerate(row, 1):
            c = ws.cell(row=row_idx, column=col_idx, value=value)
            c.fill = fill; c.alignment = ROW_ALIGN
            if col_idx == 2:
                c.hyperlink = value; c.font = URL_FONT
            else:
                c.font = ROW_FONT
        ws.row_dimensions[row_idx].height = 16
    for col_idx, width in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADER))}1"
    wb.save(OUTPUT_PATH)
    print(f"Saved: {OUTPUT_PATH}  Rows: {len(DATA)}")


if __name__ == "__main__":
    make_xlsx()
