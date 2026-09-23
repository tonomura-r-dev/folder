import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUTPUT_PATH = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー\架電リスト_全社再検証_2026-07-22_batch7_最終.xlsx"
HEADER = ["企業名", "LP URL(広告確認済み)", "電話番号", "検索KW", "媒体"]
COL_WIDTHS = [40, 50, 16, 32, 14]
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="メイリオ", size=10, bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
ROW_FILLS = [PatternFill("solid", fgColor="F2F7FC"), PatternFill("solid", fgColor="FFFFFF")]
ROW_FONT = Font(name="メイリオ", size=9)
ROW_ALIGN = Alignment(vertical="center", wrap_text=False)
URL_FONT = Font(name="メイリオ", size=9, color="0563C1", underline="single")

# Meta広告ライブラリで配信中の広告を個別確認できた企業のみ。
DATA = [
    ("アデコ株式会社", "https://www.adecco.co.jp/", "050-2000-7024", "正社員 就職支援, 人材派遣", "Meta"),
    ("株式会社MIXI（minimo）", "https://minimodel.jp/", "", "ミニモ, 美容室 当日予約", "Meta"),
    ("株式会社リブセンス（マッハバイト）", "https://j-sen.jp/", "03-6683-0082", "マッハバイト, バイト 週1", "Meta"),
    ("株式会社ディー・エヌ・エー（Pococha）", "https://pococha.com/", "", "Pococha, ライブ配信 アプリ", "Meta"),
    ("株式会社ビオ・マーケット（ビオ・マルシェの宅配）", "https://biomarche.jp/", "06-6866-1438", "有機野菜 宅配, オーガニック 定期便", "Meta"),
    ("株式会社はせがわ", "https://www.hasegawa.jp/", "0570-01-7676", "仏壇 位牌, 四十九日 法要", "Meta"),
]

NOT_VERIFIED = [
    "株式会社ドレスガーデン", "株式会社Mrk&Co（Dine）", "株式会社ディー・エヌ・エー（YYC）",
    "株式会社リュミエリーナ", "株式会社ロゼット", "株式会社Dcollection", "株式会社ウニコ",
    "株式会社クラッシュゲート", "医療法人社団慶育会（銀座よしえクリニック）", "医療法人社団エストワン",
    "医療法人社団恵比寿美容外科", "株式会社SOX（バイク館）", "株式会社レッドバロン（求人広告のみ確認）",
    "株式会社cotree", "株式会社ドワンゴ", "アルインコ株式会社", "東武トップツアーズ株式会社",
    "株式会社プレシャスウォーター", "西濃運輸株式会社", "株式会社ペットホームウェブ（ブリーダーナビ）",
    "株式会社みんなのペット（みんなのブリーダー）", "株式会社ロボ団", "株式会社サイバーエージェント（Tech Kids School）",
    "株式会社アクティビティジャパン", "株式会社信濃湧水", "株式会社Kyash", "株式会社スタイルストイック",
    "株式会社アウトソーシング", "トーセイ株式会社", "株式会社キャリアカレッジジャパン",
    "株式会社Global Step Academy", "EF Education Japan株式会社", "株式会社ポニークリーニング",
    "eBay Japan合同会社（Qoo10）", "株式会社ノジマ", "タカラスタンダード株式会社",
    "株式会社ベネッセスタイルケア", "株式会社MOSH", "株式会社ストリートアカデミー（ストアカ）",
    "17Live Japan株式会社", "株式会社タイムチケット", "株式会社エコノバ",
    "株式会社シューマツワーカー", "株式会社エンファクトリー（Workship）", "株式会社NANGA",
    "株式会社タビックスジャパン", "CJ FOODS JAPAN株式会社（美酢）",
    "弁護士法人ひばり法律事務所", "弁護士法人杉山法律事務所",
    "株式会社シロク（N organic）", "株式会社Hyatt",
    "高光製薬株式会社", "森川健康堂株式会社", "サンテミナ株式会社", "株式会社LEFT-U",
    "株式会社For-S", "株式会社アストリション", "株式会社エターナル", "有限会社ルーティ",
    "mog株式会社", "日本ドクターヘルスケア合同会社", "株式会社PROE", "株式会社ヘルシープラス",
]


def make_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "確認済み"
    for col_idx, header in enumerate(HEADER, 1):
        c = ws.cell(row=1, column=col_idx, value=header)
        c.fill = HEADER_FILL; c.font = HEADER_FONT; c.alignment = HEADER_ALIGN
    ws.row_dimensions[1].height = 20
    for row_idx, row in enumerate(DATA, 2):
        fill = ROW_FILLS[(row_idx - 2) % 2]
        for col_idx, value in enumerate(row, 1):
            c = ws.cell(row=row_idx, column=col_idx, value=value)
            c.fill = fill; c.alignment = ROW_ALIGN
            if col_idx == 2:
                c.hyperlink = value; c.font = URL_FONT
            else:
                c.font = ROW_FONT
        ws.row_dimensions[row_idx].height = 16
    for col_idx, width in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADER))}1"

    ws2 = wb.create_sheet("未確認")
    ws2.cell(row=1, column=1, value="広告ライブラリで配信中の広告が見つからなかった企業（要再調査・除外候補）").font = Font(bold=True)
    for i, name in enumerate(NOT_VERIFIED, 2):
        ws2.cell(row=i, column=1, value=name)
    ws2.column_dimensions["A"].width = 55

    wb.save(OUTPUT_PATH)
    print(f"Saved: {OUTPUT_PATH}  Confirmed: {len(DATA)}  Not verified: {len(NOT_VERIFIED)}")


if __name__ == "__main__":
    make_xlsx()
