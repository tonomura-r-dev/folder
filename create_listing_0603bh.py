import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUTPUT_PATH = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー\架電リスト_リスティング_2026-06-03bh.xlsx"
HEADER = ["企業名", "LP URL", "広告種別", "法人番号", "商材", "検索KW", "業界"]
COL_WIDTHS = [32, 56, 16, 16, 22, 32, 22]
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="メイリオ", size=10, bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
ROW_FILLS = [PatternFill("solid", fgColor="F2F7FC"), PatternFill("solid", fgColor="FFFFFF")]
ROW_FONT = Font(name="メイリオ", size=9)
ROW_ALIGN = Alignment(vertical="center", wrap_text=False)
URL_FONT = Font(name="メイリオ", size=9, color="0563C1", underline="single")

DATA = [
    # ── コーディングテスト ──
    ("株式会社ギブリー",            "https://tracks.run/tracktest-no1/",                      "リスティング広告", "", "Track Test・コーディングテスト・導入No.1",    "コーディングテスト 採用 ツール, エンジニア スキル評価, ギブリー 評判", "HR SaaS"),
    # ── 日程調整SaaS ──
    ("株式会社TimeRex",             "https://timerex.net/",                                   "リスティング広告", "", "TimeRex・自動日程調整・無料プランあり",       "日程調整 ツール 無料, スケジュール 自動化 SaaS, TimeRex 評判",       "業務DXSaaS"),
    # ── エンジニア採用テスト ──
    ("株式会社ハイヤールー",        "https://hireroo.io/pricing",                             "リスティング広告", "", "HireRoo・エンジニア採用スキル評価・AI面接",   "エンジニア採用 スキル テスト, コーディング 評価 ツール, HireRoo 料金", "HR SaaS"),
    # ── ITエンジニア派遣 ──
    ("パーソルクロステクノロジー株式会社", "https://persol-tech-s.co.jp/lp/001a/",           "リスティング広告", "", "ITエンジニア派遣・オンライン登録・IT人材",    "ITエンジニア 派遣 登録, システム エンジニア 仕事 探し, パーソル IT", "人材派遣"),
    # ── ノーコード自動化 ──
    ("Yoom株式会社",                "https://lp.yoom.fun/features/app_connect",               "リスティング広告", "", "Yoom・ノーコード業務自動化・100+アプリ連携", "ノーコード 自動化 SaaS, アプリ 連携 RPA, Yoom 料金 評判",            "業務自動化SaaS"),
    # ── 翻訳AI ──
    ("DeepL SE",                    "https://www.deepl.com/ja/for-business",                  "リスティング広告", "", "DeepL Pro・AIビジネス翻訳・法人向け",         "DeepL 法人 翻訳 料金, ビジネス AI 翻訳 SaaS, DeepL Pro 評判",       "AI翻訳SaaS"),
    # ── Web多言語化 ──
    ("WOVN Technologies株式会社",   "https://mx.wovn.io/lp/global-site",                      "リスティング広告", "", "WOVN.io・Webサイト多言語化・ノーコード",       "Web サイト 多言語化 ツール, グローバル 展開 LP 翻訳, WOVN 料金",    "多言語化SaaS"),
    # ── 採用広報 ──
    ("talentbook株式会社",          "https://product.talent-book.jp/",                        "リスティング広告", "", "talentbook・採用広報・AI自動化ソリューション", "採用広報 ツール SaaS, 求職者 インナー ブランディング, タレントブック", "採用マーケSaaS"),
    # ── クラウドカメラ ──
    ("セーフィー株式会社",          "https://safie.co.jp/service/",                           "リスティング広告", "", "クラウドカメラ・録画シェアNo.1・店舗DX",      "クラウドカメラ シェアNo1, 店舗 監視 カメラ DX, セーフィー 料金",    "IoT SaaS"),
    # ── コミュニティSaaS ──
    ("株式会社Asobica",             "https://coorum.jp/lp/community_ver3/",                   "リスティング広告", "", "coorum・コミュニティ構築・CS・UGC創出",        "コミュニティ ツール SaaS 比較, CS 顧客 定着 ツール, coorum 料金",   "カスタマーサクセスSaaS"),
    # ── 株式報酬SaaS ──
    ("Nstock株式会社",              "https://nstock.co.jp/business",                          "リスティング広告", "", "株式報酬SaaS・ストックオプション管理",         "ストックオプション 管理 SaaS, 株式報酬 スタートアップ, Nstock 料金", "フィンテックSaaS"),
    # ── 新卒採用マッチング ──
    ("株式会社ジェイック",          "https://j.futurefinder.net/lp/biz04b/",                  "リスティング広告", "", "FutureFinder・適職診断×採用マッチング",      "新卒 採用 マッチング, 適職診断 就活, FutureFinder 企業 料金",         "採用SaaS"),
    # ── 受付システム ──
    ("株式会社RECEPTIONIST",        "https://receptionist.jp/lp04/",                          "リスティング広告", "", "クラウド受付システム・4,000社導入・iPadで無人", "受付 システム SaaS 比較, クラウド 来客 管理, RECEPTIONIST 料金",    "業務DXSaaS"),
    # ── 現場DX ──
    ("株式会社カミナシ",            "https://lp.kaminashi.jp/factory-audit",                  "リスティング広告", "", "現場DXプラットフォーム・チェックシート電子化", "現場 DX プラットフォーム, チェックシート 電子化 SaaS, カミナシ 料金", "現場DXSaaS"),
    # ── SEOツール ──
    ("Ahrefs Pte.Ltd.",             "https://ahrefs.com/ja/signup?interval=monthly&plan=trial","リスティング広告", "", "Ahrefs・SEOツール・バックリンク分析",          "SEO ツール 比較, Ahrefs 料金 評判, バックリンク 分析 SaaS",          "SEO SaaS"),
    # ── HR評価SaaS ──
    ("株式会社HRBrain",             "https://www.hrbrain.jp/lp/dxplaness",                    "リスティング広告", "", "タレントマネジメント・人事評価・顧客満足度No.1","HRBrain 料金 評判, 人事評価 DX, タレントマネジメント SaaS",          "HR SaaS"),
    # ── 電子契約SaaS ──
    ("GMOグローバルサイン・ホールディングス株式会社", "https://www.gmosign.com/lp/transfer/", "リスティング広告", "", "GMOサイン・電子印鑑・乗り換えキャンペーン",   "GMOサイン 乗り換え, 電子契約 移行, 電子印鑑 SaaS 比較",            "電子契約SaaS"),
    # ── インテントデータ ──
    ("株式会社Sales Marker",        "https://sales-marker.jp/lp/demo/",                       "リスティング広告", "", "インテントセールス・BtoB商談獲得・デモ",       "インテント データ 営業, BtoB 商談 ツール, Sales Marker 料金",        "セールスSaaS"),
    # ── CXM・SNS管理 ──
    ("Sprinklr Japan株式会社",      "https://www.sprinklr.com/jp/lp/modern-care-platform/",   "リスティング広告", "", "Sprinklr・SNS管理・CXMプラットフォーム",       "SNS 管理 ツール 法人, カスタマーケア SaaS, Sprinklr 料金",           "CXMSaaS"),
    # ── 広告計測 ──
    ("株式会社イルグルム",          "https://go.ebis.ne.jp/lp/",                              "リスティング広告", "", "AD EBiS・広告効果測定・CVR分析",               "広告 効果 測定 ツール, アドエビス 料金 評判, CV 計測 SaaS",          "AdTech"),
]


def make_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "架電リスト"

    for col_idx, header in enumerate(HEADER, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
    ws.row_dimensions[1].height = 20

    for row_idx, row_data in enumerate(DATA, 2):
        fill = ROW_FILLS[(row_idx - 2) % 2]
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.fill = fill
            cell.alignment = ROW_ALIGN
            if col_idx == 2:
                cell.hyperlink = value
                cell.font = URL_FONT
            else:
                cell.font = ROW_FONT
        ws.row_dimensions[row_idx].height = 16

    for col_idx, width in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADER))}1"

    wb.save(OUTPUT_PATH)
    print(f"Saved: {OUTPUT_PATH}")
    print(f"Rows: {len(DATA)}")


if __name__ == "__main__":
    make_xlsx()
