import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUTPUT_PATH = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー\架電リスト_リスティング_2026-06-02c.xlsx"
HEADER = ["企業名", "LP URL", "広告種別", "法人番号", "商材", "検索KW", "業界"]
COL_WIDTHS = [32, 50, 16, 16, 20, 32, 22]
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="メイリオ", size=10, bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
ROW_FILLS = [PatternFill("solid", fgColor="F2F7FC"), PatternFill("solid", fgColor="FFFFFF")]
ROW_FONT = Font(name="メイリオ", size=9)
ROW_ALIGN = Alignment(vertical="center", wrap_text=False)
URL_FONT = Font(name="メイリオ", size=9, color="0563C1", underline="single")

DATA = [
    # ── 夏キャンペーン×高客単価 ──────────────────────────────────────────
    # 省エネ・住設
    ("株式会社和上ホールディングス",          "https://wajo-holdings.jp/solar/lp-assessment/",   "リスティング広告", "3120001162667", "太陽光発電システム設置",          "太陽光発電 見積もり 無料, 太陽光パネル 設置 費用, 自家消費 太陽光 補助金",  "太陽光発電"),
    ("株式会社キンライサー",                  "https://www.kyutooki.com/",                       "リスティング広告", "8120901013139", "エコキュート・給湯器交換",          "給湯器 交換 費用, エコキュート 設置 補助金, 給湯器 故障 即日対応",         "住設・省エネ"),
    ("株式会社イキナリデンキ",                "https://ikinari.net/lp/",                         "リスティング広告", "8010001226783", "エアコン修理・交換工事（即日対応）", "エアコン 修理 即日, エアコン 取り付け 費用 東京, エアコン 交換 業者",      "エアコン工事"),
    # 学習・夏期講習
    ("株式会社四谷大塚",                     "https://www.yotsuyaotsuka.com/koushu/summer/",    "リスティング広告", "8011201005655", "中学受験夏期講習",                   "中学受験 夏期講習 2026, 四谷大塚 夏期 申込, 受験 塾 夏休み 費用",          "学習塾"),
    ("株式会社スクールナビ",                  "https://www.navi-school.com/summer/",             "リスティング広告", "8120001146095", "ナビ個別指導学院 夏期講習",          "個別指導 夏期講習 申込, 夏休み 塾 近く おすすめ, 個別 夏期講習 安い",      "学習塾"),
    # スポーツ・体験
    ("株式会社乗馬クラブクレイン",            "https://www.uma-crane.com/",                      "リスティング広告", "5120101032132", "乗馬体験・乗馬クラブ入会",          "乗馬 体験 東京 おすすめ, 乗馬クラブ 入会 費用, 乗馬 初心者 体験",          "乗馬"),
    ("株式会社イトマンスイミングスクール",    "https://www.itoman.com/",                         "リスティング広告", "6120001030177", "スイミングスクール夏期入会",         "スイミングスクール 夏 体験, 水泳教室 子供 夏休み, 水泳 習い事 無料体験",   "スポーツスクール"),
    ("株式会社ナンバメイト",                  "https://menkyolive.net/",                         "リスティング広告", "7120001024665", "合宿免許（夏の繁忙期）",             "合宿免許 夏休み 安い, 免許合宿 申込 最安値, 夏 運転免許 合宿 おすすめ",    "免許取得"),
    ("株式会社パパラギダイビングスクール横浜","https://www.papalagi.co.jp/",                     "リスティング広告", "9020001105465", "ダイビングCカード取得スクール",      "ダイビング ライセンス 取得 費用, スキューバ 体験 東京, PADI 初心者 海",     "マリンスポーツ"),
    # アウトドア・リゾート
    ("株式会社ピカ（PICAリゾート）",          "https://www.pica-resort.jp/glamping.html",        "リスティング広告", "5090001010316", "グランピング・夏アウトドア宿泊",     "グランピング 富士山 予約, 夏 グランピング 関東 おすすめ, キャンプ 高規格",  "グランピング"),
]


def make_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "架電リスト"

    for col_idx, header in enumerate(HEADER, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
    ws.row_dimensions[1].height = 20

    for row_idx, row_data in enumerate(DATA, 2):
        fill = ROW_FILLS[(row_idx - 2) % 2]
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.fill = fill
            cell.alignment = ROW_ALIGN
            if col_idx == 2:
                cell.hyperlink = value
                cell.font = URL_FONT
            else:
                cell.font = ROW_FONT
        ws.row_dimensions[row_idx].height = 16

    for col_idx, width in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADER))}1"

    wb.save(OUTPUT_PATH)
    print(f"Saved: {OUTPUT_PATH}")
    print(f"Rows: {len(DATA)}")


if __name__ == "__main__":
    make_xlsx()
