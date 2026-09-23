import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUTPUT_PATH = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー\架電リスト_リスティング_2026-07-05a.xlsx"
HEADER = ["企業名", "LP URL", "広告種別", "法人番号", "商材", "検索KW", "業界"]
COL_WIDTHS = [32, 56, 16, 16, 22, 32, 22]
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="メイリオ", size=10, bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
ROW_FILLS = [PatternFill("solid", fgColor="F2F7FC"), PatternFill("solid", fgColor="FFFFFF")]
ROW_FONT = Font(name="メイリオ", size=9)
ROW_ALIGN = Alignment(vertical="center", wrap_text=False)
URL_FONT = Font(name="メイリオ", size=9, color="0563C1", underline="single")

DATA = [
    ("ゲンナイ製薬株式会社", "https://gennai-seiyaku.co.jp/", "リスティング広告", "", "葉酸サプリ 笑顔の種", "葉酸サプリ, 妊活 サプリ", "サプリ(妊活)"),
    ("株式会社さくらの森", "https://sakura-forest.com/", "リスティング広告", "", "きなり(DHA)・めなり(ルテイン)", "DHA サプリ, ルテイン サプリ", "サプリ"),
    ("ステラ漢方株式会社", "https://www.stella-s.com/", "リスティング広告", "", "ステラの贅沢青汁", "青汁, 贅沢青汁", "サプリ(青汁)"),
    ("株式会社シードコムス", "https://seedcoms.net/", "リスティング広告", "", "サプリメント専門通販", "サプリ 激安, サプリメント 通販", "サプリ"),
    ("株式会社健康家族", "https://www.kenkoukazoku.co.jp/", "リスティング広告", "", "伝統にんにく卵黄", "にんにく卵黄, 黒酢 にんにく", "サプリ"),
    ("株式会社自然食研", "https://www.sizenshokken.co.jp/", "リスティング広告", "", "しじみ習慣", "しじみ サプリ, オルニチン", "サプリ"),
    ("株式会社サン・クラルテ製薬", "https://www.sain-clarte.com/", "リスティング広告", "", "ヒアルログルコ・KALCALA", "膝 サプリ, ダイエット サプリ", "サプリ"),
    ("株式会社メタボリック", "https://shop.mdc.co.jp/", "リスティング広告", "", "酵素×酵母 生酵素", "生酵素 サプリ, 酵素 ダイエット", "サプリ"),
    ("株式会社やまちや", "https://www.yamachiya.co.jp/", "リスティング広告", "", "きょうの青汁", "青汁 国産, イソフラボン 青汁", "サプリ(青汁)"),
    ("株式会社しまのや", "https://www.simanoya.com/", "リスティング広告", "", "琉球醪酢(もろみ酢)", "もろみ酢, 黒酢 サプリ", "サプリ"),
    ("株式会社ホコニコヘルスケア", "https://hoconico.com/", "リスティング広告", "", "こだわり酵素青汁", "酵素青汁, 青汁 サプリ", "サプリ(青汁)"),
    ("株式会社リフレ", "https://hc-refre.jp/", "リスティング広告", "", "ブルーベリー&ルテイン", "ルテイン サプリ, ヘラシボウ", "サプリ"),
    ("株式会社ヘルスアップ", "https://shizen-labo.jp/", "リスティング広告", "", "和麹づくしの雑穀生酵素", "生酵素, 麹 サプリ", "サプリ"),
    ("株式会社愛しとーと", "https://aishitoto.co.jp/", "リスティング広告", "", "うるおい宣言(コラーゲン)", "コラーゲン, コラーゲンゼリー", "サプリ(美容)"),
    ("株式会社エーエフシー", "https://464981.com/", "リスティング広告", "", "AFC サプリ 30日分", "葉酸 サプリ, ルテイン サプリ", "サプリ"),
    ("株式会社ユーワ", "https://www.yuwa-shop.com/", "リスティング広告", "", "青汁・グルコサミン", "青汁, グルコサミン サプリ", "サプリ"),
    ("リブ・ラボラトリーズ株式会社", "http://www.liv-net.co.jp/", "リスティング広告", "", "ヘラシボウ・ダイエット食品", "ダイエット サプリ, プロテイン", "サプリ(ダイエット)"),
    ("株式会社エポラ", "https://www.epauler.jp/", "リスティング広告", "", "みどりむし(ユーグレナ)サプリ", "ユーグレナ サプリ, みどりむし", "サプリ"),
    ("ツインガーデン株式会社", "https://orkis.jp/", "リスティング広告", "", "B.B.B(トリプルビー)HMB", "HMB 女性, ダイエット サプリ", "サプリ(ダイエット)"),
    ("アスクレピオス製薬株式会社", "https://asklepios.co.jp/", "リスティング広告", "", "ダイエット・健康食品サプリ", "ダイエット サプリ, 健康食品 通販", "サプリ"),
]


def make_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "架電リスト"
    for col_idx, header in enumerate(HEADER, 1):
        c = ws.cell(row=1, column=col_idx, value=header)
        c.fill = HEADER_FILL; c.font = HEADER_FONT; c.alignment = HEADER_ALIGN
    ws.row_dimensions[1].height = 20
    for row_idx, row in enumerate(DATA, 2):
        fill = ROW_FILLS[(row_idx - 2) % 2]
        for col_idx, value in enumerate(row, 1):
            c = ws.cell(row=row_idx, column=col_idx, value=value)
            c.fill = fill; c.alignment = ROW_ALIGN
            if col_idx == 2:
                c.hyperlink = value; c.font = URL_FONT
            else:
                c.font = ROW_FONT
        ws.row_dimensions[row_idx].height = 16
    for col_idx, width in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADER))}1"
    wb.save(OUTPUT_PATH)
    print(f"Saved: {OUTPUT_PATH}  Rows: {len(DATA)}")


if __name__ == "__main__":
    make_xlsx()
