import sys
import io
from pptx import Presentation

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

path = r"C:\Users\tonomura-r\Downloads\LINEOA.pptx"
p = Presentation(path)
print("total slides:", len(p.slides))

targets = set(sys.argv[1:]) if len(sys.argv) > 1 else None

def walk(shapes, depth=0):
    for shape in shapes:
        if shape.shape_type == 6:  # group
            walk(shape.shapes, depth + 1)
            continue
        if shape.has_text_frame:
            text = shape.text_frame.text
            if text.strip():
                print(f"{'  '*depth}--- shape '{shape.name}' ---")
                print(text)
        if shape.has_table:
            print(f"{'  '*depth}--- table '{shape.name}' ---")
            for row in shape.table.rows:
                print(" | ".join(cell.text for cell in row.cells))
        if shape.shape_type == 13:  # picture
            print(f"{'  '*depth}--- picture '{shape.name}' ---")

for i, slide in enumerate(p.slides, start=1):
    if targets and str(i) not in targets:
        continue
    print(f"\n===== Slide {i} =====")
    walk(slide.shapes)
