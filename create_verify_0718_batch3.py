import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUTPUT_PATH = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー\架電リスト_全社再検証_2026-07-18_batch3.xlsx"
HEADER = ["企業名", "LP URL(広告確認済み)", "電話番号", "検索KW", "媒体"]
COL_WIDTHS = [40, 50, 16, 32, 14]
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="メイリオ", size=10, bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
ROW_FILLS = [PatternFill("solid", fgColor="F2F7FC"), PatternFill("solid", fgColor="FFFFFF")]
ROW_FONT = Font(name="メイリオ", size=9)
ROW_ALIGN = Alignment(vertical="center", wrap_text=False)
URL_FONT = Font(name="メイリオ", size=9, color="0563C1", underline="single")

DATA = [
    ("PayPay銀行株式会社", "https://login.paypay-bank.co.jp/", "03-6739-5000", "PayPay銀行 カードローン, ネット銀行", "Meta"),
    ("auじぶん銀行株式会社", "https://lp.jibunbank.co.jp/", "0120-926-111", "auじぶん銀行 カードローン, au限定割", "Meta"),
    ("セコム株式会社", "https://www.secom.co.jp/", "03-5775-8301", "防犯カメラ 設置, ホームセキュリティ", "Meta"),
    ("綜合警備保障株式会社（ALSOK）", "https://lp.alsok.co.jp/", "", "高齢者 みまもり, ALSOK 資料請求", "Meta"),
    ("株式会社クラシコム（北欧、暮らしの道具店）", "https://hokuoh.com/", "0120-096-456", "北欧 雑貨 通販, ライフスタイル EC", "Meta"),
    ("株式会社共立メンテナンス（ドーミーイン）", "https://www.hotespa.net/", "03-5295-7777", "ドーミーイン 予約, 温泉付き ホテル", "Meta"),
    ("Wondertech株式会社（SwitchBot）", "https://www.switchbot.jp/", "03-6416-9946", "SwitchBot, スマートホーム 後付け", "Meta"),
]

NOT_VERIFIED = [
    "株式会社シェアフル", "株式会社アムタス", "NHN株式会社（comico）", "株式会社生活の木",
    "株式会社ルートインジャパン", "株式会社スーパーホテル", "チョイスホテルズジャパン株式会社",
    "Nature株式会社", "アンカー・ジャパン株式会社",
]


def make_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "確認済み"
    for col_idx, header in enumerate(HEADER, 1):
        c = ws.cell(row=1, column=col_idx, value=header)
        c.fill = HEADER_FILL; c.font = HEADER_FONT; c.alignment = HEADER_ALIGN
    ws.row_dimensions[1].height = 20
    for row_idx, row in enumerate(DATA, 2):
        fill = ROW_FILLS[(row_idx - 2) % 2]
        for col_idx, value in enumerate(row, 1):
            c = ws.cell(row=row_idx, column=col_idx, value=value)
            c.fill = fill; c.alignment = ROW_ALIGN
            if col_idx == 2:
                c.hyperlink = value; c.font = URL_FONT
            else:
                c.font = ROW_FONT
        ws.row_dimensions[row_idx].height = 16
    for col_idx, width in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADER))}1"

    ws2 = wb.create_sheet("未確認")
    ws2.cell(row=1, column=1, value="広告ライブラリで配信中の広告が見つからなかった企業（要再調査・除外候補）").font = Font(bold=True)
    for i, name in enumerate(NOT_VERIFIED, 2):
        ws2.cell(row=i, column=1, value=name)
    ws2.column_dimensions["A"].width = 55

    wb.save(OUTPUT_PATH)
    print(f"Saved: {OUTPUT_PATH}  Confirmed: {len(DATA)}  Not verified: {len(NOT_VERIFIED)}")


if __name__ == "__main__":
    make_xlsx()
