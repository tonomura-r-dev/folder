import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def make_xlsx(filename, data):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "架電リスト"

    headers = ["企業名", "LP URL", "電話番号", "商材", "検索KW", "業界"]

    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="メイリオ", bold=True, color="FFFFFF", size=10)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center
        cell.border = border

    for row_idx, row in enumerate(data, 2):
        fill_color = "F2F7FC" if row_idx % 2 == 0 else "FFFFFF"
        row_fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
        for col_idx, val in enumerate(row, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.fill = row_fill
            cell.font = Font(name="メイリオ", size=9)
            cell.alignment = left
            cell.border = border

    col_widths = [30, 45, 16, 20, 30, 18]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.row_dimensions[1].height = 20
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    wb.save(filename)
    print(f"保存完了: {filename}  件数: {len(data)}")


# 列順: 企業名 | LP URL | 電話番号 | 商材 | 検索KW | 業界
# 全社、公式サイト・PR情報等でWebSearch裏取り済み。使用済み企業リストとGrep照合し重複を除外。
# 電話番号はフリーダイヤル(0800等)や未確認のものは空欄。

data_0717 = [
    ["株式会社和とわ", "https://kimono-rentalier.jp/furisode", "0191-26-5291", "宅配着物レンタル（振袖・きものレンタリエ）", "振袖 レンタル 宅配／成人式 着物レンタル", "着物レンタルEC"],
    ["株式会社マッハ", "https://www.mach5.jp/", "093-474-0855", "車検・車買取（マッハ車検、FC本部）", "車検 安い／車検 予約 最短", "車検FC"],
    ["株式会社ホリデー", "https://www.holiday-fc.co.jp/", "06-6633-8500", "車検サービス（ホリデー車検、FC本部）", "車検 立会い／車検 予約 安い", "車検FC"],
    ["株式会社速太郎本部", "https://www.hayataro.com/shop/", "", "車検サービス（車検の速太郎、FC本部）", "車検 45分／車検 安い 早い", "車検FC"],
    ["株式会社ANELLA Group", "https://wan-time.jp/", "03-4500-2644", "保護犬猫ふれあいカフェ（ANELLA CAFE、FC本部）", "保護犬 カフェ／猫カフェ 保護猫", "ペットカフェFC"],
    ["株式会社ラブグラフ", "https://lovegraph.me/", "03-5708-5382", "出張撮影サービス（Lovegraph）", "家族写真 出張撮影／前撮り 出張カメラマン", "出張撮影サービス"],
    ["日本介護システム株式会社", "https://franchise-salon.net/", "06-6271-2725", "訪問理美容FC加盟店募集（KamiBito／髪人）", "訪問理美容 独立／FC 加盟 募集", "訪問美容FC本部"],
    ["株式会社Timers", "https://school-lp.famm.us/", "", "女性向けWebデザインスクール（Fammスクール）", "Webデザイン 講座／在宅ワーク 資格", "オンラインスクール"],
]

base = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー"
make_xlsx(f"{base}\\架電リスト_Meta_2026-07-17.xlsx", data_0717)
