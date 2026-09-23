import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUTPUT_PATH = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー\架電リスト_リスティング_2026-05-29e.xlsx"

HEADER = ["企業名", "LP URL", "広告種別", "電話番号", "商材", "検索KW", "業界"]
COL_WIDTHS = [32, 50, 16, 16, 20, 32, 22]

DATA = [
    [
        "ラクサスマネジメント株式会社",
        "https://www.raxus-management.co.jp/lp/",
        "リスティング広告",
        "0120-046-115",
        "不動産投資セミナー・個別相談",
        "不動産投資 セミナー マンション投資 東京 初心者",
        "不動産投資",
    ],
    [
        "株式会社セオリーファクトリー",
        "https://www.theoryfactory.jp/lp/",
        "リスティング広告",
        "0120-86-3702",
        "区分マンション投資・不動産投資",
        "不動産投資 区分マンション 東京 サラリーマン 節税",
        "不動産投資",
    ],
    [
        "セイハ株式会社",
        "https://english-academy.seiha.com/lp/summer2025/",
        "リスティング広告",
        "0120-815-718",
        "子ども英会話教室・無料体験レッスン",
        "子ども英会話 英語教室 幼児 小学生 無料体験",
        "教育・スキルアップ",
    ],
    [
        "株式会社オーダースーツSADA",
        "https://www.ordersuit.info/lp/ordersuit/",
        "リスティング広告",
        "0120-351-298",
        "フルオーダースーツ・採寸予約",
        "オーダースーツ フルオーダー 安い 採寸 スーツ",
        "ファッション・アパレル",
    ],
    [
        "株式会社やる気スイッチグループ",
        "https://www.schoolie-net.jp/fs/lp/lp001/taiken/",
        "リスティング広告",
        "0120-869-104",
        "個別指導学習塾・無料体験授業",
        "個別指導塾 学習塾 スクールIE 無料体験 中学生",
        "教育・スキルアップ",
    ],
    [
        "株式会社ZWEI",
        "https://www.zwei.com/iryou/lp",
        "リスティング広告",
        "0120-374-281",
        "結婚相談所・婚活サポート",
        "結婚相談所 婚活 マッチング 無料相談 ツヴァイ",
        "婚活・マッチング",
    ],
    [
        "株式会社キタムラ",
        "https://shop.kitamura.jp/lp/brand-01/index.html",
        "リスティング広告",
        "0120-640-055",
        "ブランド品・貴金属・カメラ買取",
        "ブランド品 買取 高価買取 カメラ 貴金属 査定",
        "買取・リユース",
    ],
    [
        "リビン・テクノロジーズ株式会社",
        "https://www.lvnmatch.jp/lp/sell/promo/",
        "リスティング広告",
        "0120-935-565",
        "不動産売却・一括査定サービス",
        "不動産 売却 査定 一括査定 マンション 売りたい",
        "不動産・売買",
    ],
    [
        "株式会社プランドゥ",
        "https://www.my-kaigo-home.com/lp/soudan/input/",
        "リスティング広告",
        "0120-175-155",
        "老人ホーム・介護施設の入居相談",
        "老人ホーム 介護施設 入居 相談 費用 探し方",
        "介護・福祉",
    ],
    [
        "トライズ株式会社",
        "https://toraiz.jp/lp/re/",
        "リスティング広告",
        "0120-961-406",
        "ビジネス英語コーチング・短期集中",
        "英語コーチング ビジネス英語 短期 英会話 話せる",
        "教育・スキルアップ",
    ],
]

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="メイリオ", size=10, bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
ROW_FILLS = [PatternFill("solid", fgColor="F2F7FC"), PatternFill("solid", fgColor="FFFFFF")]
ROW_FONT = Font(name="メイリオ", size=9)
ROW_ALIGN = Alignment(vertical="center", wrap_text=False)
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def make_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "架電リスト"

    # Header row
    for col_idx, h in enumerate(HEADER, start=1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
        cell.border = BORDER
    ws.row_dimensions[1].height = 22

    # Data rows
    for row_idx, row_data in enumerate(DATA, start=2):
        fill = ROW_FILLS[(row_idx - 2) % 2]
        for col_idx, val in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.fill = fill
            cell.border = BORDER
            if col_idx == 2 and val:  # LP URL column: hyperlink
                cell.hyperlink = val
                cell.font = Font(name="メイリオ", size=9, color="0563C1", underline="single")
                cell.alignment = ROW_ALIGN
            else:
                cell.font = ROW_FONT
                cell.alignment = ROW_ALIGN
        ws.row_dimensions[row_idx].height = 18

    # Column widths
    for col_idx, width in enumerate(COL_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # Freeze panes and auto filter
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    wb.save(OUTPUT_PATH)
    print(f"保存完了: {OUTPUT_PATH}")
    print(f"総件数: {len(DATA)}社")


if __name__ == "__main__":
    make_xlsx()
