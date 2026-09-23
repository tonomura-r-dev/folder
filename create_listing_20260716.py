import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def make_xlsx(filename, data):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "架電リスト"

    headers = ["企業名", "LP URL", "電話番号", "商材", "検索KW", "業界"]

    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="メイリオ", bold=True, color="FFFFFF", size=10)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center
        cell.border = border

    for row_idx, row in enumerate(data, 2):
        fill_color = "F2F7FC" if row_idx % 2 == 0 else "FFFFFF"
        row_fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
        for col_idx, val in enumerate(row, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.fill = row_fill
            cell.font = Font(name="メイリオ", size=9)
            cell.alignment = left
            cell.border = border

    col_widths = [30, 45, 16, 20, 30, 18]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.row_dimensions[1].height = 20
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    wb.save(filename)
    print(f"保存完了: {filename}  行数: {len(data)}")


# 第149弾（2026-07-16）買取・リユース業界拡充分（20社）

data_0716 = [
    ["株式会社水野", "https://manekiya.shop/", "03-6908-5688", "金・プラチナ・ダイヤモンド・宝石買取（まねきや）", "貴金属買取, 金 買取 相場, ダイヤ 買取", "貴金属買取"],
    ["株式会社銀蔵", "https://ginzo.biz/", "03-5366-7800", "ロレックス等高級腕時計・ブランド品買取", "ロレックス 買取, 腕時計 買取 専門店, ブランド品 買取", "時計買取"],
    ["グローバルトレード株式会社", "https://tokyo-brand.jp/", "03-6455-7677", "ヴィトン・エルメス・シャネル等ブランドバッグ・財布買取（東京ぶらんど）", "ブランドバッグ 買取, エルメス 財布 売る, ブランド品 買取 東京", "ブランド品買取"],
    ["株式会社京都きもの市場", "https://www.kimonoichiba.com/user_data/kaitori/", "050-8885-0580", "着物・帯・和装小物買取", "着物 買取, 着物 売る 宅配, 振袖 買取", "着物買取"],
    ["株式会社新岐阜商会", "https://www.shingifu.co.jp/kaitori/", "058-264-4777", "切手・古銭・記念硬貨買取（新岐阜切手古銭商会）", "切手 買取, 古銭 買取, 記念硬貨 売る", "切手・古銭買取"],
    ["いわの美術株式会社", "https://iwano.biz/contact/", "", "骨董品・美術品・掛軸・茶道具買取", "骨董品 買取, 美術品 買取 出張, 掛軸 売る", "骨董品買取"],
    ["株式会社楽器堂", "https://www.gakkido-opus.jp/", "088-824-1840", "ギター・管楽器等中古楽器買取（楽器堂OPUS）", "楽器 買取, ギター 買取, 中古楽器 売る", "楽器買取"],
    ["株式会社フジヤカメラ店", "https://www.fujiya-camera.co.jp/shop/kaitori/pc/0c-kaitor/", "03-5318-6320", "中古カメラ・交換レンズ買取", "カメラ 買取, レンズ 買取, 中古カメラ 売る", "カメラ買取"],
    ["株式会社ブックスドリーム", "https://www.mottainaihonpo.com/kaitori/", "06-6210-1675", "古本・コミック・全巻セット買取（もったいない本舗）", "古本 買取, コミック 買取, 本 売る 宅配", "本・コミック買取"],
    ["株式会社RUSH", "https://www.cardrush.jp/", "03-6260-9030", "ポケモンカード・遊戯王等トレーディングカード買取（カードラッシュ）", "トレカ 買取, ポケモンカード 買取, 遊戯王 カード 売る", "トレーディングカード買取"],
    ["株式会社GKファクトリー", "https://golf-king.com/", "052-508-7272", "中古ゴルフクラブ・ゴルフ用品買取（ゴルフキング）", "ゴルフクラブ 買取, ドライバー 買取, ゴルフ用品 売る", "ゴルフ用品買取"],
    ["株式会社リンク", "https://noukiguou.com/sateiform/", "0748-36-3697", "トラクター・コンバイン等中古農機具買取（農機具王）", "農機具 買取, トラクター 買取, 中古農機具 売る", "農機具買取"],
    ["株式会社オフィスバスターズ", "https://www.officebusters.com/caitori/", "03-6262-3123", "事務机・OA機器等オフィス家具買取", "オフィス家具 買取, 中古オフィス用品 売る, 事務椅子 買取", "オフィス家具・什器買取"],
    ["株式会社じゃんぱら", "https://buy.janpara.co.jp/", "03-5298-3682", "PC・スマートフォン・デジタル家電買取", "パソコン 買取, スマホ 買取, 中古PC 売る", "パソコン・スマホ・家電買取"],
    ["OMO株式会社", "https://kougu-up.com/", "03-6874-2693", "電動工具等中古工具買取（工具UP）", "工具 買取, 電動工具 買取, 中古工具 売る", "工具買取"],
    ["株式会社タックルベリー", "https://www.tackleberry.co.jp/kaitori/", "0466-52-0877", "中古釣具買取", "釣具 買取, ロッド リール 買取, 釣り道具 売る", "釣具買取"],
    ["株式会社VOLCA", "https://mountain-c.com/buy/", "", "キャンプ・登山用品買取（マウンテンシティ）", "キャンプ用品 買取, アウトドア用品 買取, 登山用品 売る", "アウトドア・キャンプ用品買取"],
    ["株式会社交通趣味ギャラリー", "http://www.tokyo-collection.co.jp/kaitori.html", "03-5157-5072", "鉄道模型・鉄道部品・切符買取", "鉄道グッズ 買取, 鉄道模型 買取, 切符 買取", "鉄道グッズ・切符買取"],
    ["株式会社ナガツマ", "https://www.bikeboy.jp/form/", "04-7166-7711", "旧車・不動車含むバイク買取（バイクボーイ）", "バイク 買取, 旧車 バイク 売る, バイク 出張買取", "バイク買取"],
    ["JOYLAB株式会社", "https://joylab.jp/online/", "03-6234-0520", "ウイスキー・ワイン・日本酒等お酒買取", "お酒 買取, ウイスキー 買取, ワイン 売る", "酒類買取"],
]

base = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー"
make_xlsx(f"{base}\\架電リスト_Meta_2026-07-16.xlsx", data_0716)
