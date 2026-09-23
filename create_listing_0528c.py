import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUTPUT_PATH = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー\架電リスト_リスティング_2026-05-28c.xlsx"
HEADER = ["企業名", "LP URL", "広告種別", "電話番号", "商材", "検索KW", "業界"]
COL_WIDTHS = [32, 50, 16, 16, 20, 32, 22]

DATA = [
    [
        "株式会社ファーストロジック（楽待）",
        "https://www.rakumachi.jp/lp/apps/rakumachi/",
        "リスティング広告",
        "",
        "不動産投資情報プラットフォーム",
        "不動産投資 収益物件 利回り 物件情報 売却",
        "不動産・賃貸・リノベーション・投資",
    ],
    [
        "株式会社ブランディングエンジニア（tech boost）",
        "https://tech-boost.jp/lp/jobchange-001",
        "リスティング広告",
        "",
        "プログラミングスクール（転職保証付き）",
        "プログラミングスクール エンジニア転職 未経験 副業",
        "教育・スキルアップ",
    ],
    [
        "株式会社ほねごり",
        "https://www.honegori.co.jp/lp/honegori_sango/",
        "リスティング広告",
        "042-703-5548",
        "整骨院・接骨院・はりきゅう院（51店舗）",
        "整骨院 接骨院 腰痛 肩こり 産後骨盤矯正",
        "フィットネス・スポーツ・ゴルフ",
    ],
    [
        "株式会社スプリックス（そら塾）",
        "https://www.sorajuku.jp/lp/05/",
        "リスティング広告",
        "03-6416-5190",
        "オンライン個別指導塾（小中高対応）",
        "個別指導塾 オンライン塾 小学生 中学生 高校生",
        "教育・スキルアップ",
    ],
    [
        "株式会社トラーナ（トイサブ！）",
        "https://toysub.net/lp/lp001.html",
        "リスティング広告",
        "",
        "0〜3歳向け知育玩具サブスクリプション",
        "知育玩具 おもちゃ サブスク レンタル 赤ちゃん",
        "教育・スキルアップ",
    ],
    [
        "株式会社ホワイトウィングスリテール（イーピュア）",
        "https://pr-cleaning.co/lp/",
        "リスティング広告",
        "",
        "宅配クリーニングサービス",
        "宅配クリーニング 洋服クリーニング 宅配 送料無料",
        "家事代行・クリーニング",
    ],
    [
        "株式会社たちばな（きものたちばな）",
        "https://www.tachibana-group.co.jp/lp/furisode/",
        "リスティング広告",
        "026-238-0242",
        "振袖レンタル・成人式着物・フォトスタジオ",
        "振袖レンタル 成人式振袖 着物レンタル 前撮り 購入",
        "ウエディング・葬儀",
    ],
    [
        "iCracked Japan株式会社",
        "https://www.icracked.jp/lp/ic-tokyo001/",
        "リスティング広告",
        "03-6629-8001",
        "スマートフォン修理（iPhone・Android）",
        "iPhone修理 スマホ修理 画面割れ バッテリー交換 即日",
        "格安スマホ・通信",
    ],
    [
        "デジタルデータソリューション株式会社（デジタルデータリカバリー）",
        "https://www.ino-inc.com/lp/pc/",
        "リスティング広告",
        "0120-706-332",
        "データ復旧・データ復元サービス",
        "データ復旧 データ復元 HDD 外付けHDD iPhone",
        "格安スマホ・通信",
    ],
    [
        "株式会社プリシラ",
        "https://www.prisila.jp/lp/kawaii/",
        "リスティング広告",
        "078-671-6722",
        "ウィッグ・かつら・エクステ通販",
        "ウィッグ かつら 女性 ファッション 通販 自然",
        "化粧品・健康食品通販D2C",
    ],
]

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="メイリオ", size=10, bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)

ROW_FILLS = [
    PatternFill("solid", fgColor="F2F7FC"),
    PatternFill("solid", fgColor="FFFFFF"),
]
ROW_FONT = Font(name="メイリオ", size=9)
ROW_ALIGN = Alignment(vertical="center", wrap_text=False)

LINK_FONT_BASE = {"name": "メイリオ", "size": 9, "color": "0563C1", "underline": "single"}

THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def make_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "架電リスト"

    # ヘッダー行
    for col_idx, col_name in enumerate(HEADER, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN
        cell.border = BORDER
    ws.row_dimensions[1].height = 22

    # データ行
    for row_idx, row_data in enumerate(DATA, 2):
        fill = ROW_FILLS[(row_idx) % 2]
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.fill = fill
            cell.border = BORDER
            if col_idx == 2:  # LP URL 列
                cell.hyperlink = value
                cell.font = Font(
                    name=LINK_FONT_BASE["name"],
                    size=LINK_FONT_BASE["size"],
                    color=LINK_FONT_BASE["color"],
                    underline=LINK_FONT_BASE["underline"],
                )
                cell.alignment = ROW_ALIGN
            else:
                cell.font = ROW_FONT
                cell.alignment = ROW_ALIGN
        ws.row_dimensions[row_idx].height = 18

    # 列幅
    for col_idx, width in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # 先頭行固定・オートフィルター
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    wb.save(OUTPUT_PATH)
    print(f"保存完了: {OUTPUT_PATH}")
    print(f"データ行数: {len(DATA)}")


if __name__ == "__main__":
    make_xlsx()
