import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUTPUT_PATH = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー\架電リスト_リスティング_2026-05-26d.xlsx"

HEADER = ["企業名", "LP URL", "広告種別", "電話番号", "商材", "検索KW", "業界"]
COL_WIDTHS = [32, 50, 16, 16, 20, 32, 22]

# LP URLはすべて /lp/ を含む実在確認済みページ
DATA = [
    (
        "株式会社オリエントコーポレーション（オリコ）",
        "https://www.orico.co.jp/cardorder/lp/gw/thepoint/",
        "リスティング広告", "",
        "クレジットカード・カードローン（Orico Card THE POINT等）",
        "オリコカード 年会費無料 申し込み, ポイント 還元率 高い カード, カードローン 即日 審査",
        "保険・金融・証券・カード"
    ),
    (
        "株式会社リノベる。",
        "https://www.renoveru.jp/lp/document_request",
        "リスティング広告", "",
        "中古マンション＋リノベーション（ワンストップサービス）",
        "中古マンション リノベーション 一括, リノベーション 費用 相場, 中古マンション 購入 リノベ 東京",
        "不動産・賃貸・リノベーション・投資"
    ),
    (
        "株式会社シャケンカン（車検館）",
        "https://www.shakenkan.co.jp/lp/",
        "リスティング広告", "",
        "車検専門店（格安・最短・整備済み全国展開）",
        "車検 安い 近く 予約, 格安 車検 専門店, 車検 最短 当日",
        "車・カーシェア・バイク"
    ),
    (
        "株式会社学情（Re就活）",
        "https://re-katsu.jp/career/lp/tenshoku/index02.html",
        "リスティング広告", "",
        "20代・第二新卒専門転職サイト",
        "第二新卒 転職 サイト おすすめ, 20代 未経験 転職 エージェント, 既卒 転職 正社員",
        "転職・就職・人材派遣"
    ),
    (
        "湘南AGAクリニック（SBCメディカルグループ）",
        "https://www.sbc-aga.jp/lp/cm-m/",
        "リスティング広告", "",
        "AGA・薄毛治療クリニック（全国70院以上・オンライン診療）",
        "AGA 治療 クリニック 予約, 薄毛 治療 東京 安い, 男性 ハゲ 対策 薬",
        "美容クリニック・脱毛・エステ"
    ),
    (
        "LDT株式会社（やさしいお葬式）",
        "https://y-osohshiki.com/lp/kazokusoh/",
        "リスティング広告", "",
        "家族葬・一日葬・火葬式（8.9万円〜、全国対応）",
        "家族葬 費用 安い 地域, 葬儀 料金 相場 比較, 格安 葬式 家族のみ",
        "ウエディング・葬儀"
    ),
    (
        "ディップ株式会社（はたらこねっと）",
        "https://www.hatarako.net/lp/jimu/",
        "リスティング広告", "",
        "派遣・アルバイト求人サイト（事務・オフィスワーク特化）",
        "派遣 事務 仕事 近く 未経験, 事務職 パート 時給 高い 主婦, オフィスワーク 派遣 登録 無料",
        "転職・就職・人材派遣"
    ),
    (
        "株式会社GA technologies（RENOSY）",
        "https://pwa.renosy.com/lp/executive-plan",
        "リスティング広告", "",
        "不動産投資プラットフォーム（AI活用・区分マンション投資）",
        "不動産投資 始め方 サラリーマン, マンション 投資 利回り 東京, 資産運用 不動産 少額",
        "不動産・賃貸・リノベーション・投資"
    ),
    (
        "株式会社しちだ・教育研究所（七田式）",
        "https://www.shichida.co.jp/shichipri/lp/",
        "リスティング広告", "",
        "幼児向けプリント学習教材（しちぷり・七田式教材）",
        "幼児 プリント 学習 通信 毎日, 七田式 教材 家庭学習 2歳, 子供 知育 プリント 無料体験",
        "教育・スキルアップ"
    ),
    (
        "株式会社オプテージ（mineo/マイネオ）",
        "https://mineo.jp/lp/lp_10.html",
        "リスティング広告", "",
        "格安SIM・格安スマホ（au/docomo/SoftBank 3キャリア対応）",
        "格安SIM 乗り換え おすすめ 比較, mineo 料金プラン 月額, スマホ 維持費 安くする 方法",
        "格安スマホ・通信"
    ),
]

def make_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "架電リスト"

    header_font   = Font(name="メイリオ", size=10, bold=True, color="FFFFFF")
    header_fill   = PatternFill("solid", fgColor="1F4E79")
    header_align  = Alignment(horizontal="center", vertical="center", wrap_text=True)

    fill_odd  = PatternFill("solid", fgColor="F2F7FC")
    fill_even = PatternFill("solid", fgColor="FFFFFF")
    data_font = Font(name="メイリオ", size=9)
    data_align = Alignment(vertical="center", wrap_text=False)
    url_font  = Font(name="メイリオ", size=9, color="0563C1", underline="single")

    thin = Side(style="thin", color="D0D7DE")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col, name in enumerate(HEADER, start=1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = border

    for row_idx, row in enumerate(DATA, start=2):
        fill = fill_odd if row_idx % 2 == 0 else fill_even
        for col, val in enumerate(row, start=1):
            cell = ws.cell(row=row_idx, column=col, value=val)
            cell.fill = fill
            cell.border = border
            cell.alignment = data_align
            if col == 2 and val:
                cell.hyperlink = val
                cell.font = url_font
            else:
                cell.font = data_font

    for col, width in enumerate(COL_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width

    ws.row_dimensions[1].height = 22
    for r in range(2, len(DATA) + 2):
        ws.row_dimensions[r].height = 18

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADER))}1"

    wb.save(OUTPUT_PATH)
    print(f"保存完了: {OUTPUT_PATH}")
    print(f"件数: {len(DATA)}社")

if __name__ == "__main__":
    make_xlsx()
