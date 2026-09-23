import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def make_xlsx(filename, data):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "架電リスト"

    headers = ["企業名", "LP URL", "電話番号", "商材", "検索KW", "業界"]

    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="メイリオ", bold=True, color="FFFFFF", size=10)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center
        cell.border = border

    for row_idx, row in enumerate(data, 2):
        fill_color = "F2F7FC" if row_idx % 2 == 0 else "FFFFFF"
        row_fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
        for col_idx, val in enumerate(row, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.fill = row_fill
            cell.font = Font(name="メイリオ", size=9)
            cell.alignment = left
            cell.border = border

    col_widths = [30, 45, 16, 20, 30, 18]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.row_dimensions[1].height = 20
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    wb.save(filename)
    print(f"保存完了: {filename}  件数: {len(data)}")


# 列順: 企業名 | LP URL | 電話番号 | 商材 | 検索KW | 業界
# 買取(ホビー/自転車/ブランド品) と 美容系(増毛・発毛) を追加。
# 全社、公式サイト等でWebSearch裏取り済み。使用済み企業リストとGrep照合し重複なしを確認済み。

data_0717b = [
    ["有限会社ヤマト", "https://www.toysking.jp/", "03-5817-8666", "おもちゃ・フィギュア・プラモデル買取（トイズキング）", "フィギュア 買取／プラモデル 買取", "ホビー買取"],
    ["株式会社クレイン", "https://jewel-cafe.jp/kaitori/cosme/perfume/", "03-5684-3821", "ブランド品・香水買取（ジュエルカフェ）", "香水 買取／ブランド品 買取", "ブランド買取FC"],
    ["R-CUBE株式会社", "https://www.secondroad.jp/", "", "ロードバイク・自転車買取（セカンドロード）", "ロードバイク 買取／自転車 売る", "自転車買取"],
    ["株式会社プロピア", "https://propia.co.jp/", "", "増毛・かつら・発毛（Propia）", "増毛 サロン／かつら 育毛", "増毛・発毛サロン"],
    ["株式会社Zoo MASTER", "https://zo-mo.com/", "", "増毛エクステ専門サロン（SENSE by plushair）", "増毛エクステ／薄毛 サロン", "増毛サロンFC"],
]

base = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー"
make_xlsx(f"{base}\\架電リスト_Meta_2026-07-17b.xlsx", data_0717b)
