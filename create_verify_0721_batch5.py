import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUTPUT_PATH = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー\架電リスト_全社再検証_2026-07-21_batch5.xlsx"
HEADER = ["企業名", "LP URL(広告確認済み)", "電話番号", "検索KW", "媒体"]
COL_WIDTHS = [40, 50, 16, 32, 14]
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="メイリオ", size=10, bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
ROW_FILLS = [PatternFill("solid", fgColor="F2F7FC"), PatternFill("solid", fgColor="FFFFFF")]
ROW_FONT = Font(name="メイリオ", size=9)
ROW_ALIGN = Alignment(vertical="center", wrap_text=False)
URL_FONT = Font(name="メイリオ", size=9, color="0563C1", underline="single")

# Meta広告ライブラリで配信中の広告を個別確認できた企業のみ。URLは広告に表示されていた遷移先。
DATA = [
    ("株式会社早稲田アカデミー", "https://www.waseda-ac.co.jp/", "03-3590-4011", "早稲田アカデミー 夏期講習, 中学受験 塾", "Meta"),
    ("株式会社日能研", "https://nichinoken-kansai.com/", "045-473-2311", "日能研 夏期講習, 中学受験", "Meta"),
    ("NTTソルマーレ株式会社（コミックシーモア）", "https://www.cmoa.jp/", "06-6228-8861", "コミックシーモア, 電子書籍 漫画", "Meta"),
    ("株式会社BookLive（ブックライブ）", "https://booklive.jp/", "03-6633-6290", "ブックライブ, 電子書籍 マンガ", "Meta"),
    ("株式会社ボディワーク（Re.Ra.Ku）", "https://reraku.jp/", "", "リラク 肩甲骨, マッサージ 駅近", "Meta"),
    ("株式会社GDOゴルフテック", "https://golftec.golfdigest.co.jp/", "03-5656-2870", "ゴルフレッスン, ゴルフスクール マンツーマン", "Meta"),
    ("WAQ株式会社", "https://waq-online.com/", "06-6195-8288", "WAQ テント, キャンプマット", "Meta"),
    ("有限会社ヤマト（トイズキング）", "https://www.toysking.jp/", "03-5817-8666", "フィギュア 買取, おもちゃ 出張買取", "Meta"),
    ("R-CUBE株式会社（セカンドロード）", "https://www.secondroad.jp/", "", "ロードバイク 買取, E-BIKE 売る", "Meta"),
    ("株式会社プロピア", "https://propia.co.jp/", "", "増毛 ヘアコンタクト, ウィッグ 自然", "Meta"),
    ("株式会社Zoo MASTER（SENSE by plushair）", "https://zo-mo.com/", "", "増毛エクステ, 薄毛 サロン", "Meta"),
    ("株式会社ケイズイノベーション（R Dresser）", "https://www.rina-ogawa.com/", "", "パーソナルカラー診断, 骨格診断", "Meta"),
    ("株式会社ファイヤーワークス（MR.BROTHERS CUT CLUB）", "https://mr-brothers-cutclub.com/", "03-6455-0319", "メンズ バーバー, メンズカット 専門店", "Meta"),
    ("株式会社ブラスト（美.design）", "https://kogaokyosei.com/", "0120-980-439", "小顔矯正, 美容整体", "Meta"),
    ("株式会社さつま骨格矯正", "https://beauty.hotpepper.jp/kr/slnH000319620/", "03-5468-0910", "小顔矯正, 顎顔面矯正", "Meta"),
    ("株式会社Racbaki", "https://racbaki.com/", "", "骨格矯正靴, 脚痩せ サンダル", "Meta"),
    ("株式会社羅針（GINZA RASIN）", "https://www.kaitori-ginza.com/", "03-6252-1515", "腕時計 買取, ロレックス 売る", "Meta"),
    ("日本介護システム株式会社（KamiBito FC加盟募集）", "https://franchise-salon.net/", "06-6271-2725", "訪問理美容 FC, 独立 開業", "Meta"),
    ("株式会社Timers（Fammスクール）", "https://school-lp.famm.us/", "", "Webデザイン 講座, ママ 在宅ワーク", "Meta"),
]

NOT_VERIFIED = [
    "株式会社浜学園", "株式会社Kids Duo International", "株式会社すらら",
    "株式会社りらくる", "株式会社テモミン", "株式会社U-NEXT",
    "スターツコーポレーション株式会社", "株式会社JCOM",
    "SOMPOひまわり生命保険株式会社", "朝日生命保険相互会社",
    "株式会社開祥（e品整理）", "有限会社トライ リコレクションズ（トライホビーズ）",
    "株式会社K-Dream1（美cuol）",
    "株式会社おさるとめぐま", "株式会社エバンス", "株式会社Altomare",
    "株式会社ライフバケーション", "株式会社LJ", "株式会社バーバリーアートスペース",
    "株式会社秋華洞", "株式会社日晃堂", "株式会社貴瞬", "株式会社買取王国",
]


def make_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "確認済み"
    for col_idx, header in enumerate(HEADER, 1):
        c = ws.cell(row=1, column=col_idx, value=header)
        c.fill = HEADER_FILL; c.font = HEADER_FONT; c.alignment = HEADER_ALIGN
    ws.row_dimensions[1].height = 20
    for row_idx, row in enumerate(DATA, 2):
        fill = ROW_FILLS[(row_idx - 2) % 2]
        for col_idx, value in enumerate(row, 1):
            c = ws.cell(row=row_idx, column=col_idx, value=value)
            c.fill = fill; c.alignment = ROW_ALIGN
            if col_idx == 2:
                c.hyperlink = value; c.font = URL_FONT
            else:
                c.font = ROW_FONT
        ws.row_dimensions[row_idx].height = 16
    for col_idx, width in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADER))}1"

    ws2 = wb.create_sheet("未確認")
    ws2.cell(row=1, column=1, value="広告ライブラリで配信中の広告が見つからなかった/確認できなかった企業（要再調査・除外候補）").font = Font(bold=True)
    for i, name in enumerate(NOT_VERIFIED, 2):
        ws2.cell(row=i, column=1, value=name)
    ws2.column_dimensions["A"].width = 55

    wb.save(OUTPUT_PATH)
    print(f"Saved: {OUTPUT_PATH}  Confirmed: {len(DATA)}  Not verified: {len(NOT_VERIFIED)}")


if __name__ == "__main__":
    make_xlsx()
