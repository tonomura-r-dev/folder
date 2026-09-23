import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUTPUT_PATH = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー\架電リスト_リスティング_2026-06-03bx.xlsx"
HEADER = ["企業名", "LP URL", "広告種別", "法人番号", "商材", "検索KW", "業界"]
COL_WIDTHS = [32, 56, 16, 16, 22, 32, 22]
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="メイリオ", size=10, bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
ROW_FILLS = [PatternFill("solid", fgColor="F2F7FC"), PatternFill("solid", fgColor="FFFFFF")]
ROW_FONT = Font(name="メイリオ", size=9)
ROW_ALIGN = Alignment(vertical="center", wrap_text=False)
URL_FONT = Font(name="メイリオ", size=9, color="0563C1", underline="single")

DATA = [
    # ── Apple端末管理 ──
    ("Jamf",                       "https://www.jamf.com/pricing/",                           "リスティング広告", "", "Jamf・Apple端末管理MDM/セキュリティ",          "MDM Apple 端末管理, Jamf 料金 評判, デバイス 管理 SaaS",          "セキュリティSaaS"),
    # ── クロスプラットフォームMDM ──
    ("Hexnode",                    "https://www.hexnode.com/pricing/",                        "リスティング広告", "", "Hexnode・MDM/端末管理・クロスOS",             "MDM ツール 比較, Hexnode 料金 評判, モバイル デバイス 管理",      "セキュリティSaaS"),
    # ── 認証基盤 ──
    ("FusionAuth",                 "https://fusionauth.io/pricing",                           "リスティング広告", "", "FusionAuth・認証/認可基盤・MFA",              "認証 基盤 SaaS, FusionAuth 料金 評判, ユーザー 認証 ツール",      "認証SaaS"),
    # ── 認証基盤 ──
    ("Stytch",                     "https://stytch.com/pricing",                              "リスティング広告", "", "Stytch・パスワードレス認証/MFA",              "パスワードレス 認証 SaaS, Stytch 料金 評判, 認証 API",            "認証SaaS"),
    # ── エンタープライズ認証 ──
    ("WorkOS",                     "https://workos.com/pricing",                              "リスティング広告", "", "WorkOS・SSO/SCIM・B2B SaaS向け認証",          "SSO SCIM 認証 SaaS, WorkOS 料金 評判, エンタープライズ 認証",     "認証SaaS"),
    # ── 認証基盤 ──
    ("Clerk",                      "https://clerk.com/pricing",                               "リスティング広告", "", "Clerk・ユーザー認証/管理・開発者向け",         "ユーザー 認証 SaaS, Clerk 料金 評判, 認証 ライブラリ ツール",     "認証SaaS"),
    # ── 認証基盤 ──
    ("Descope",                    "https://www.descope.com/pricing",                         "リスティング広告", "", "Descope・パスワードレス認証/フロー構築",       "パスワードレス 認証 ツール, Descope 料金 評判, 認証 フロー SaaS",  "認証SaaS"),
    # ── パスワードレス認証 ──
    ("MojoAuth",                   "https://mojoauth.com/pricing",                            "リスティング広告", "", "MojoAuth・パスワードレス認証/パスキー",        "パスワードレス 認証 API, MojoAuth 料金 評判, パスキー 認証",       "認証SaaS"),
    # ── 認証基盤 ──
    ("Kinde",                      "https://kinde.com/pricing/",                              "リスティング広告", "", "Kinde・認証/ユーザー管理・スタートアップ向け", "認証 SaaS スタートアップ, Kinde 料金 評判, ユーザー 管理 認証",   "認証SaaS"),
    # ── 認証基盤（OSS） ──
    ("SuperTokens",                "https://supertokens.com/pricing",                         "リスティング広告", "", "SuperTokens・OSS認証/セッション管理",          "OSS 認証 SaaS, SuperTokens 料金 評判, 認証 セッション 管理",      "認証SaaS"),
    # ── ディレクトリ/IDaaS ──
    ("JumpCloud",                  "https://jumpcloud.com/pricing",                           "リスティング広告", "", "JumpCloud・クラウドディレクトリ/IDaaS/MDM",   "IDaaS ディレクトリ SaaS, JumpCloud 料金 評判, クラウド 認証 管理", "セキュリティSaaS"),
    # ── コンプライアンス自動化 ──
    ("Vanta",                      "https://www.vanta.com/pricing",                           "リスティング広告", "", "Vanta・SOC2/ISO等コンプライアンス自動化",      "SOC2 コンプライアンス SaaS, Vanta 料金 評判, セキュリティ 監査 自動化","セキュリティSaaS"),
    # ── コンプライアンス自動化 ──
    ("Drata",                      "https://drata.com/pricing",                               "リスティング広告", "", "Drata・コンプライアンス自動化・SOC2/ISO",      "コンプライアンス 自動化 SaaS, Drata 料金 評判, セキュリティ 認証 監査","セキュリティSaaS"),
    # ── コンプライアンス自動化 ──
    ("Secureframe",                "https://secureframe.com/pricing",                         "リスティング広告", "", "Secureframe・コンプライアンス自動化・監査",    "セキュリティ コンプライアンス SaaS, Secureframe 料金 評判, SOC2 監査","セキュリティSaaS"),
    # ── コンプライアンス自動化 ──
    ("Sprinto",                    "https://sprinto.com/pricing/",                            "リスティング広告", "", "Sprinto・コンプライアンス自動化・SOC2/ISO",    "コンプライアンス 自動化 ツール, Sprinto 料金 評判, セキュリティ 認証","セキュリティSaaS"),
    # ── ゼロトラストVPN ──
    ("Tailscale",                  "https://tailscale.com/pricing",                           "リスティング広告", "", "Tailscale・ゼロトラストVPN/メッシュネット",    "ゼロトラスト VPN SaaS, Tailscale 料金 評判, メッシュ ネットワーク","セキュリティSaaS"),
    # ── ゼロトラストアクセス ──
    ("Twingate",                   "https://www.twingate.com/pricing",                        "リスティング広告", "", "Twingate・ゼロトラストネットワークアクセス",   "ゼロトラスト ZTNA SaaS, Twingate 料金 評判, リモート アクセス 管理","セキュリティSaaS"),
    # ── エンタープライズ認証 ──
    ("ScaleKit",                   "https://www.scalekit.com/pricing",                        "リスティング広告", "", "ScaleKit・B2B SaaS向けSSO/SCIM認証",          "SSO SCIM SaaS, ScaleKit 料金 評判, エンタープライズ 認証 API",    "認証SaaS"),
    # ── アーティファクト管理 ──
    ("Cloudsmith",                 "https://cloudsmith.com/pricing",                          "リスティング広告", "", "Cloudsmith・成果物管理/パッケージリポジトリ", "アーティファクト 管理 SaaS, Cloudsmith 料金 評判, パッケージ 管理","DevOps SaaS"),
    # ── 認証基盤（OSS） ──
    ("Ory",                        "https://www.ory.sh/pricing",                              "リスティング広告", "", "Ory・OSS認証/アイデンティティ基盤",            "OSS 認証 基盤, Ory 料金 評判, アイデンティティ 管理 SaaS",        "認証SaaS"),
]


def make_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "架電リスト"

    for col_idx, header in enumerate(HEADER, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
    ws.row_dimensions[1].height = 20

    for row_idx, row_data in enumerate(DATA, 2):
        fill = ROW_FILLS[(row_idx - 2) % 2]
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.fill = fill
            cell.alignment = ROW_ALIGN
            if col_idx == 2:
                cell.hyperlink = value
                cell.font = URL_FONT
            else:
                cell.font = ROW_FONT
        ws.row_dimensions[row_idx].height = 16

    for col_idx, width in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADER))}1"

    wb.save(OUTPUT_PATH)
    print(f"Saved: {OUTPUT_PATH}")
    print(f"Rows: {len(DATA)}")


if __name__ == "__main__":
    make_xlsx()
