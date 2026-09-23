from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

SRC = r"C:\Users\tonomura-r\Downloads\LINEOA.pptx"
DST = r"C:\Users\tonomura-r\Downloads\LINEOA.pptx"

NOTE_TEXT = "※記載の金額は全て税別表示です。"

# raw slide indices (1-based) that need the tax notation footnote added
TARGET_SLIDES = [18, 19, 20, 21, 24, 68]

p = Presentation(SRC)

for idx in TARGET_SLIDES:
    slide = p.slides[idx - 1]
    box = slide.shapes.add_textbox(Inches(0.5), Inches(7.12), Inches(9.8), Inches(0.3))
    tf = box.text_frame
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    para = tf.paragraphs[0]
    para.alignment = PP_ALIGN.LEFT
    run = para.add_run()
    run.text = NOTE_TEXT
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
    run.font.name = "Calibri"
    print(f"Added tax notation to slide {idx}")

p.save(DST)
print("Saved.")
