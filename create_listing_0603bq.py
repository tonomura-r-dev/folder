import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUTPUT_PATH = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー\架電リスト_リスティング_2026-06-03bq.xlsx"
HEADER = ["企業名", "LP URL", "広告種別", "法人番号", "商材", "検索KW", "業界"]
COL_WIDTHS = [32, 56, 16, 16, 22, 32, 22]
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="メイリオ", size=10, bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
ROW_FILLS = [PatternFill("solid", fgColor="F2F7FC"), PatternFill("solid", fgColor="FFFFFF")]
ROW_FONT = Font(name="メイリオ", size=9)
ROW_ALIGN = Alignment(vertical="center", wrap_text=False)
URL_FONT = Font(name="メイリオ", size=9, color="0563C1", underline="single")

DATA = [
    # ── AIコピーライティング ──
    ("Jasper AI Inc.",             "https://www.jasper.ai/pricing",                           "リスティング広告", "", "Jasper・AIマーケティング・コピーライティング", "AI ライティング ツール, Jasper 料金 評判, AI コピー 生成",         "AIライティングSaaS"),
    # ── AIライティング ──
    ("Writesonic Inc.",            "https://writesonic.com/pricing",                          "リスティング広告", "", "Writesonic・AI記事/SEO・コンテンツ生成",      "AI 記事 作成 ツール, Writesonic 料金 評判, SEO コンテンツ AI",     "AIライティングSaaS"),
    # ── AI動画生成 ──
    ("Synthesia Ltd.",             "https://www.synthesia.io/ja",                             "リスティング広告", "", "Synthesia・AIアバター動画生成・140言語",       "AI 動画 生成 ツール, Synthesia 料金 評判, アバター 動画 作成",     "AI動画SaaS"),
    # ── AI音声 ──
    ("ElevenLabs Inc.",            "https://elevenlabs.io/ja/pricing",                        "リスティング広告", "", "ElevenLabs・AI音声合成・ナレーション",         "AI 音声 合成 ツール, ElevenLabs 料金 評判, テキスト 音声 読み上げ","AI音声SaaS"),
    # ── 動画/音声編集 ──
    ("Descript Inc.",              "https://www.descript.com/pricing",                        "リスティング広告", "", "Descript・AI動画/ポッドキャスト編集",          "動画 編集 AI ツール, Descript 料金 評判, ポッドキャスト 編集",     "AI動画SaaS"),
    # ── 文字起こしAI ──
    ("Otter.ai (AISense Inc.)",    "https://otter.ai/jp",                                     "リスティング広告", "", "Otter.ai・AI文字起こし・議事録自動作成",       "文字起こし AI ツール, Otter 料金 評判, 議事録 自動 作成 SaaS",     "AI文字起こしSaaS"),
    # ── 会議文字起こしAI ──
    ("Fireflies.ai (Firefly Inc.)","https://fireflies.ai/pricing",                            "リスティング広告", "", "Fireflies.ai・会議文字起こし・AI議事録",       "会議 文字起こし AI, Fireflies 料金 評判, AI 議事録 SaaS",         "AI文字起こしSaaS"),
    # ── AI検索 ──
    ("Perplexity AI Inc.",         "https://www.perplexity.ai/enterprise/pricing",            "リスティング広告", "", "Perplexity・AI検索・企業向けエンタープライズ", "AI 検索 ツール 企業, Perplexity 料金 評判, AI リサーチ SaaS",      "AI検索SaaS"),
    # ── AIライティング/GTM ──
    ("Copy.ai Inc.",               "https://www.copy.ai/pricing",                             "リスティング広告", "", "Copy.ai・GTM AI・セールス/コンテンツ自動化",  "GTM AI プラットフォーム, Copy.ai 料金 評判, AI セールス ツール",   "AIライティングSaaS"),
    # ── AI動画生成 ──
    ("Runway AI Inc.",             "https://runwayml.com/pricing",                            "リスティング広告", "", "Runway・生成AI動画・クリエイティブツール",     "AI 動画 生成 クリエイティブ, Runway 料金 評判, 生成 AI 映像",      "AI動画SaaS"),
    # ── デザイン/プレゼン ──
    ("Visme (Easy WebContent)",    "https://www.visme.co/pricing/",                           "リスティング広告", "", "Visme・AIデザイン/プレゼン/インフォグラフィック","デザイン ツール 比較, Visme 料金 評判, プレゼン 作成 AI",         "デザインSaaS"),
    # ── AI動画作成 ──
    ("Pictory Inc.",               "https://pictory.ai/pricing/",                             "リスティング広告", "", "Pictory・AI動画作成・テキストから動画",        "AI 動画 作成 ツール, Pictory 料金 評判, テキスト 動画 変換",       "AI動画SaaS"),
    # ── AIプレゼン ──
    ("Gamma Tech Inc.",            "https://gamma.app/pricing",                               "リスティング広告", "", "Gamma・AIプレゼン/スライド/サイト作成",        "AI プレゼン 作成 ツール, Gamma 料金 評判, AI スライド 生成",       "デザインSaaS"),
    # ── 文章校正AI ──
    ("QuillBot (Course Hero)",     "https://quillbot.com/premium",                            "リスティング広告", "", "QuillBot・AI文章校正/パラフレーズ",           "文章 校正 AI ツール, QuillBot 料金 評判, パラフレーズ ツール",     "AIライティングSaaS"),
    # ── AIプレゼン ──
    ("Beautiful.ai Inc.",          "https://www.beautiful.ai/pricing",                        "リスティング広告", "", "Beautiful.ai・AIプレゼンテーション・自動デザイン","AI プレゼン デザイン, Beautiful.ai 料金 評判, スライド 自動 作成", "デザインSaaS"),
    # ── AI音声合成 ──
    ("Murf AI (Murf Inc.)",        "https://murf.ai/ja/onsei-gosei",                          "リスティング広告", "", "Murf・AI音声合成・ナレーション・120音声",      "AI 音声 合成 ナレーション, Murf 料金 評判, テキスト 読み上げ AI",  "AI音声SaaS"),
    # ── AIアバター動画 ──
    ("HeyGen (HeyGen Inc.)",       "https://www.heygen.com/pricing",                          "リスティング広告", "", "HeyGen・AIアバター動画生成・翻訳",            "AI アバター 動画 ツール, HeyGen 料金 評判, AI 動画 翻訳 SaaS",     "AI動画SaaS"),
    # ── オンライン動画編集 ──
    ("Veed.io (VEED Ltd.)",        "https://www.veed.io/ja-JP",                               "リスティング広告", "", "Veed.io・オンライン動画編集・AI字幕",          "オンライン 動画 編集 ツール, Veed 料金 評判, AI 字幕 翻訳",        "AI動画SaaS"),
    # ── AI音楽生成 ──
    ("Suno Inc.",                  "https://suno.com/pricing",                                "リスティング広告", "", "Suno・AI音楽生成・楽曲作成",                  "AI 音楽 生成 ツール, Suno 料金 評判, AI 作曲 ツール",             "AI生成SaaS"),
    # ── AIノート ──
    ("Tana Inc.",                  "https://tana.inc/pricing",                                "リスティング広告", "", "Tana・AIメモ/ノート・ナレッジグラフ",          "AI メモ ノート ツール, Tana 料金 評判, Notion 代替 AI",           "コラボSaaS"),
]


def make_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "架電リスト"

    for col_idx, header in enumerate(HEADER, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
    ws.row_dimensions[1].height = 20

    for row_idx, row_data in enumerate(DATA, 2):
        fill = ROW_FILLS[(row_idx - 2) % 2]
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.fill = fill
            cell.alignment = ROW_ALIGN
            if col_idx == 2:
                cell.hyperlink = value
                cell.font = URL_FONT
            else:
                cell.font = ROW_FONT
        ws.row_dimensions[row_idx].height = 16

    for col_idx, width in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADER))}1"

    wb.save(OUTPUT_PATH)
    print(f"Saved: {OUTPUT_PATH}")
    print(f"Rows: {len(DATA)}")


if __name__ == "__main__":
    make_xlsx()
