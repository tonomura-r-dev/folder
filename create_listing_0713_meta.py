import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def make_xlsx(filename, data):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "架電リスト"

    headers = ["企業名", "LP URL", "電話番号", "商材", "検索KW", "業界"]

    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="メイリオ", bold=True, color="FFFFFF", size=10)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center
        cell.border = border

    for row_idx, row in enumerate(data, 2):
        fill_color = "F2F7FC" if row_idx % 2 == 0 else "FFFFFF"
        row_fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
        for col_idx, val in enumerate(row, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.fill = row_fill
            cell.font = Font(name="メイリオ", size=9)
            cell.alignment = left
            cell.border = border

    col_widths = [30, 45, 16, 20, 30, 18]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.row_dimensions[1].height = 20
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    wb.save(filename)
    print(f"保存完了: {filename}")


# 列順: 企業名 | LP URL | 電話番号 | 商材 | 検索KW | 業界
# 2026-07-13 Meta広告ライブラリで配信中広告を確認・全社LP実確認済み・Grep重複確認済み

data_0713 = [
    ["株式会社CHICKEN GYM", "https://chicken-golf.com/", "03-6628-5353", "インドアゴルフレッスンスクール（FC展開）", "ゴルフスクール 初心者 / インドアゴルフ 体験", "ゴルフスクールFC"],
    ["合同会社Forza", "https://beat-pilates.com/", "080-3580-4126", "暗闇系マシンピラティススタジオ（FC展開）", "ビートピラティス 口コミ / 暗闇ピラティス 体験", "ピラティススタジオFC"],
    ["株式会社Wellness X Asia", "https://clubpilates.co.jp/", "", "マシンピラティススタジオ（世界最大級ブランド）", "クラブピラティス 料金 / CLUB PILATES 体験", "ピラティススタジオ"],
    ["株式会社ティーバランス", "https://t-balance-gym.com/", "", "パーソナルトレーニングジム", "パーソナルジム 料金 / ティーバランス 口コミ", "パーソナルジム"],
    ["株式会社ASPIREST", "https://aspirest.com/", "", "パーソナルトレーニングジム", "ASPI パーソナルジム / パーソナルトレーニング 継続", "パーソナルジム"],
    ["株式会社REJUVENATE", "https://revina-personalgym.com/", "090-4243-3895", "パーソナルジム×ピラティス（FC展開）", "ReViNa パーソナルジム / パーソナルジム FC 開業", "パーソナルジムFC"],
    ["株式会社BODY RUN Alive Soul", "https://mittness.jp/", "03-6459-3345", "女性専用暗闇キックボクシングジム", "ミットネス 料金 / 暗闇キックボクシング 体験", "フィットネス（格闘技系）"],
    ["トイカツ株式会社", "https://tkdj.net/", "03-5340-7604", "総合格闘技・キックボクシングジム", "ファイトフィット 料金 / キックボクシング 通い放題", "格闘技フィットネス"],
    ["医療法人社団康英会", "https://united-clinic.jp/", "03-6907-1942", "ED治療・AGA治療オンライン診療", "ED治療 オンライン / ユナイテッドクリニック 評判", "メンズクリニック"],
    ["名古屋植毛fittoクリニック", "https://fitto.clinic/", "052-212-7072", "自毛植毛", "自毛植毛 名古屋 / 植毛 モニター 料金", "植毛クリニック"],
    ["株式会社リリーフ", "https://relief-company.jp/", "", "遺品整理・不用品回収（FC展開）", "遺品整理 業者 / 遺品整理 FC 開業", "遺品整理FC"],
    ["株式会社篠原化学", "https://kaimin-times.com/", "052-841-1505", "高反発マットレス・寝具通販", "マットレス 高反発 通販 / 快眠タイムズ 口コミ", "寝具D2C"],
    ["株式会社こぱんはうすさくら", "https://copain-sakura.com/", "03-3527-3847", "児童発達支援・放課後等デイサービス（FC展開）", "放課後等デイサービス 教室 / 児童発達支援 見学", "児童福祉サービスFC"],
    ["エイトデザイン株式会社", "https://eightdesign.jp/", "052-883-8748", "中古マンション・戸建リノベーション", "リノベーション 名古屋 / 中古マンション リノベ 事例", "リノベーション"],
    ["株式会社リノリビング", "https://www.renoliving.jp/", "092-554-2332", "マンション・戸建リノベーション", "リノベーション 福岡 / 中古戸建 リノベ 費用", "リノベーション"],
    ["クジラ株式会社", "https://kujira.ltd/", "06-6375-7720", "マンション・戸建・古民家リノベーション", "リノベーション 大阪 / 中崎町 リノベ 会社", "リノベーション"],
    ["ファミリア株式会社", "https://familia-group.co.jp/", "029-896-5322", "ローコスト注文住宅・平屋住宅", "平屋 つくば市 / 注文住宅 ローコスト 茨城", "注文住宅"],
    ["ペッツオーライ株式会社", "https://wanpass.me/", "", "愛犬同伴施設検索アプリ", "犬 同伴 お店 アプリ / ペット同伴施設 検索", "ペット関連アプリ"],
    ["株式会社杢目金屋", "https://www.mokumeganeya.com/", "03-3408-7863", "オーダーメイド結婚指輪・婚約指輪", "結婚指輪 オーダーメイド / 婚約指輪 木目金", "ブライダルジュエリー"],
    ["株式会社ファイブスター", "https://tru.salon/", "070-8388-3128", "ネイル・まつげエクステサロン（全国展開）", "ネイルサロン 求人 / まつげエクステ 就職", "ネイル・アイラッシュサロンFC"],
]

base = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー"
make_xlsx(f"{base}\\架電リスト_Meta_2026-07-13.xlsx", data_0713)
