import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUTPUT_PATH = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー\架電リスト_リスティング_2026-06-03ad.xlsx"
HEADER = ["企業名", "LP URL", "広告種別", "法人番号", "商材", "検索KW", "業界"]
COL_WIDTHS = [32, 56, 16, 16, 22, 32, 22]
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="メイリオ", size=10, bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
ROW_FILLS = [PatternFill("solid", fgColor="F2F7FC"), PatternFill("solid", fgColor="FFFFFF")]
ROW_FONT = Font(name="メイリオ", size=9)
ROW_ALIGN = Alignment(vertical="center", wrap_text=False)
URL_FONT = Font(name="メイリオ", size=9, color="0563C1", underline="single")

# ── 全社、実際の広告クリック先LP（CV/キャンペーン/申込ページ）をWeb確認済み ──
DATA = [
    # ── スキンケア（高単価・エイジングケア） ──
    ("株式会社資生堂（エリクシール）",          "https://www.shiseido.co.jp/elixir/promotion/the-serum/", "リスティング広告", "", "エリクシール ザ セラム（美容液）",   "エリクシール 美容液 効果, シワ 改善 美容液 おすすめ, ザ セラム 口コミ",       "スキンケア"),
    # ── 補聴器（高単価・無料試聴LP） ──
    ("スターキージャパン株式会社",              "https://www.starkeyjp.com/try-starkey",                  "リスティング広告", "", "AI充電式補聴器・無料試聴",            "補聴器 おすすめ 充電式, 補聴器 無料 試聴, スターキー 補聴器 評判",            "補聴器"),
    # ── 宅配クリーニング（初回トライアルLP） ──
    ("株式会社ホワイトプラス（リネット）",      "https://www.lenet.jp/campaign/s/trial_organic/",         "リスティング広告", "", "宅配クリーニング・初回トライアル",    "宅配 クリーニング おすすめ, リネット 評判 料金, クリーニング 宅配 比較",      "宅配クリーニング"),
    # ── 健康食品（免疫ケア・定期コースLP） ──
    ("キリンホールディングス株式会社（iMUSE）", "https://kirin-kyowahakko-bio.kirin.co.jp/s/regular/",    "リスティング広告", "", "プラズマ乳酸菌 免疫ケアサプリ・定期", "免疫ケア サプリ おすすめ, プラズマ乳酸菌 効果, iMUSE 定期 価格",            "健康食品"),
    ("サントリーウエルネス株式会社",            "https://www.suntory-kenko.com/service/regular-course/",  "リスティング広告", "", "セサミンEX・DHA EPA サプリ・継続便", "セサミン 効果 おすすめ, サントリー サプリ 評判, DHA EPA 定期 比較",        "健康食品"),
    ("株式会社世田谷自然食品",                  "https://www.shizensyokuhin.jp/course/",                  "リスティング広告", "", "青汁・グルコサミン・おトク定期便",    "青汁 おすすめ 飲みやすい, グルコサミン サプリ 効果, 世田谷自然食品 評判",     "健康食品"),
    ("サンスター株式会社（健康道場）",          "https://www.kenkodojo.com/",                             "リスティング広告", "", "粉末青汁・健康道場・定期購入",        "粉末 青汁 無添加 おすすめ, サンスター 青汁 口コミ, 健康道場 青汁 定期",       "健康食品"),
]


def make_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "架電リスト"

    for col_idx, header in enumerate(HEADER, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
    ws.row_dimensions[1].height = 20

    for row_idx, row_data in enumerate(DATA, 2):
        fill = ROW_FILLS[(row_idx - 2) % 2]
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.fill = fill
            cell.alignment = ROW_ALIGN
            if col_idx == 2:
                cell.hyperlink = value
                cell.font = URL_FONT
            else:
                cell.font = ROW_FONT
        ws.row_dimensions[row_idx].height = 16

    for col_idx, width in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADER))}1"

    wb.save(OUTPUT_PATH)
    print(f"Saved: {OUTPUT_PATH}")
    print(f"Rows: {len(DATA)}")


if __name__ == "__main__":
    make_xlsx()
