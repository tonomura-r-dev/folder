import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def make_xlsx(filename, data):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "架電リスト"

    headers = ["企業名", "LP URL", "電話番号", "商材", "検索KW", "業界"]

    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="メイリオ", bold=True, color="FFFFFF", size=10)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center
        cell.border = border

    for row_idx, row in enumerate(data, 2):
        fill_color = "F2F7FC" if row_idx % 2 == 0 else "FFFFFF"
        row_fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
        for col_idx, val in enumerate(row, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.fill = row_fill
            cell.font = Font(name="メイリオ", size=9)
            cell.alignment = left
            cell.border = border

    col_widths = [30, 45, 16, 20, 30, 18]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.row_dimensions[1].height = 20
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    wb.save(filename)
    print(f"保存完了: {filename}  行数: {len(data)}")


# 第155弾（2026-07-16g）化粧品・スキンケア・美容家電D2C 7社・LP URLは購入/申込ページ

data_0716g = [
    ["Areti株式会社", "https://areti.jp/en-us/collections/beauty_device", "03-6661-6995", "ヘアアイロン・ドライヤー・美顔器等プレミアム美容家電", "美容家電 通販, ヘアアイロン 口コミ, 美顔器 おすすめ", "美容家電D2C"],
    ["株式会社ドクターケイ", "https://doctork.jp/lp?u=otameshi-vitamin", "03-5464-1217", "ビタミンC配合ドクターズコスメ（薬用Cクリアシリーズ）", "ビタミンC 化粧水, ドクターズコスメ 通販, 毛穴 美容液", "スキンケアD2C"],
    ["DINETTE株式会社", "https://phoebebeautyup.com/shop/products/PUELS-01004ZZ52-1-EYELASH-20240201-a", "", "まつ毛美容液・スキンケア（PHOEBE BEAUTY UP）", "まつ毛美容液 効果, フィービー まつげ美容液, D2C コスメ", "コスメD2C"],
    ["ドクターリセラ株式会社", "https://www.recella3d.com/abouts/start_set.php", "06-6990-4700", "無添加化粧品・美容機器（Aqua Venus等）", "ドクターリセラ 化粧品, エステサロン専売化粧品, 深層水コスメ", "スキンケアD2C"],
    ["株式会社SISI", "https://sisi.tokyo/shop/products/194", "", "敏感肌向け機能性スキンケア（ウェルエイジングセット等）", "敏感肌 スキンケア, クリーンビューティ 化粧品, 化粧水 美容液", "スキンケアD2C"],
    ["株式会社トゥヴェール", "https://vioteras.jp/shopping/lp.php?p=856", "072-726-2117", "ビタミンC美容液（ヴィオテラスCセラム等）", "ビタミンC美容液 通販, 美白美容液, 毛穴 美容液", "スキンケアD2C"],
    ["株式会社長寿乃里", "https://www.chojyu.com/lp/sekken001/sp.html", "045-640-3594", "無添加スキンケア（然-しかり-よかせっけん等）", "然 よかせっけん, 無添加 洗顔石鹸, 毛穴 黒ずみ 石鹸", "スキンケアD2C"],
]

base = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー"
make_xlsx(f"{base}\\架電リスト_Meta_2026-07-16g.xlsx", data_0716g)
