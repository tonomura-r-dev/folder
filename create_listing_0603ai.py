import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUTPUT_PATH = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー\架電リスト_リスティング_2026-06-03ai.xlsx"
HEADER = ["企業名", "LP URL", "広告種別", "法人番号", "商材", "検索KW", "業界"]
COL_WIDTHS = [32, 56, 16, 16, 22, 32, 22]
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="メイリオ", size=10, bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
ROW_FILLS = [PatternFill("solid", fgColor="F2F7FC"), PatternFill("solid", fgColor="FFFFFF")]
ROW_FONT = Font(name="メイリオ", size=9)
ROW_ALIGN = Alignment(vertical="center", wrap_text=False)
URL_FONT = Font(name="メイリオ", size=9, color="0563C1", underline="single")

# ── 全社、実際の広告クリック先LP（アートメイク/無料カウンセリングCVページ）をWeb確認済み ──
DATA = [
    # ── 医療アートメイク（無料カウンセリングLP） ──
    ("アイエスクリニック株式会社",              "https://artcell.jp/artmake-lp/",                  "リスティング広告", "", "医療アートメイク・眉毛・無料カウンセリング",   "アートメイク 眉毛 東京 おすすめ, 医療 アートメイク 評判, 眉 施術 無料相談", "医療アートメイク"),
    ("株式会社Gメディカルアートクリニック",    "https://artmake-glow-clinic.com/lp002/",         "リスティング広告", "", "アートメイク・眉・リップ・無料カウンセリング", "グロウクリニック アートメイク 東京 評判, 眉毛 自然 仕上がり 比較, メディカルメイク", "医療アートメイク"),
    ("医療法人社団みずほ台サンクリニック",      "https://sunclinic.or.jp/beauty-lp/artmake/",      "リスティング広告", "", "医療アートメイク・眉・埼玉・無料カウンセリング", "アートメイク 埼玉 安い おすすめ, 医療 アートメイク 眉 費用, 美容皮膚科 施術", "医療アートメイク"),
]


def make_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "架電リスト"

    for col_idx, header in enumerate(HEADER, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
    ws.row_dimensions[1].height = 20

    for row_idx, row_data in enumerate(DATA, 2):
        fill = ROW_FILLS[(row_idx - 2) % 2]
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.fill = fill
            cell.alignment = ROW_ALIGN
            if col_idx == 2:
                cell.hyperlink = value
                cell.font = URL_FONT
            else:
                cell.font = ROW_FONT
        ws.row_dimensions[row_idx].height = 16

    for col_idx, width in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADER))}1"

    wb.save(OUTPUT_PATH)
    print(f"Saved: {OUTPUT_PATH}")
    print(f"Rows: {len(DATA)}")


if __name__ == "__main__":
    make_xlsx()
