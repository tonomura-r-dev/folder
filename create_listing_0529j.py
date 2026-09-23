import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUTPUT_PATH = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー\架電リスト_リスティング_2026-05-29j.xlsx"
HEADER = ["企業名", "LP URL", "広告種別", "電話番号", "商材", "検索KW", "業界"]
COL_WIDTHS = [32, 50, 16, 16, 20, 32, 22]
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="メイリオ", size=10, bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
ROW_FILLS = [PatternFill("solid", fgColor="F2F7FC"), PatternFill("solid", fgColor="FFFFFF")]
ROW_FONT = Font(name="メイリオ", size=9)
ROW_ALIGN = Alignment(vertical="center", wrap_text=False)
URL_FONT = Font(name="メイリオ", size=9, color="0563C1", underline="single")

DATA = [
    ("株式会社ユーポス", "https://www.u-pohs.co.jp/lp/", "リスティング広告", "", "中古車買取", "中古車買取 東京, 中古車 高額査定, 車買取 おすすめ", "中古車買取"),
    ("マルコ株式会社", "https://www.maruko.com/lp/G01/", "リスティング広告", "", "補正下着・ランジェリー", "補正下着 効果, 体型補正 ガードル, ボディスーツ おすすめ", "下着・ランジェリー通販"),
    ("ワールドファミリー株式会社", "https://world-family.co.jp/lp/camp/all/DG/guide.html", "リスティング広告", "", "子ども英語教材（DWE）", "子ども 英語教材, ディズニー英語 口コミ, 幼児英語 おすすめ", "子ども英語教育"),
    ("株式会社リーフラス", "https://liberta.sport-school.com/lp/", "リスティング広告", "", "子どもサッカースクール", "サッカースクール 子ども, サッカー教室 小学生, スポーツスクール 習い事", "子どもスポーツスクール"),
    ("株式会社ノア", "https://www.noaballet.jp/lp/", "リスティング広告", "", "バレエスクール", "バレエ教室 子ども, バレエスクール 習い事, バレエ 大人 初心者", "バレエスクール"),
    ("株式会社チアリー", "https://www.star-programming-school.com/lp/startdash2025/", "リスティング広告", "", "子どもプログラミングスクール", "プログラミング教室 子ども, プログラミング 小学生, コーディング 習い事", "子どもプログラミングスクール"),
    ("株式会社はるやまホールディングス", "https://www.haruyama-co.jp/lp/20240315_zougaku_shitadori.html", "リスティング広告", "", "スーツ・紳士服", "スーツ 購入 安い, ビジネススーツ おすすめ, 紳士服 セール", "スーツ・紳士服"),
    ("株式会社一蔵", "https://www.ondine.jp/lp/festa_tokyo_2504/", "リスティング広告", "", "振袖レンタル・販売", "振袖 レンタル, 成人式 振袖, 振袖 購入 フルセット", "振袖レンタル・販売"),
    ("株式会社スカイロボット", "https://skyrobot.co.jp/droneschoojapan/lp/", "リスティング広告", "", "ドローンスクール", "ドローン スクール, ドローン 資格 取得, ドローン 操縦 講習", "ドローンスクール"),
    ("株式会社コンテンツラボ", "https://www.suit-ya.com/lp/request.php", "リスティング広告", "", "オーダースーツ通販", "オーダースーツ 通販, 格安 オーダーメイドスーツ, スーツ 採寸 自宅", "オーダースーツ通販"),
    ("株式会社ALL&ソリューションズ", "http://rabbit-tantei.com/lp/ad/kmp", "リスティング広告", "", "浮気調査・探偵サービス", "浮気調査 費用, 探偵 依頼, 不倫調査 格安", "探偵・調査"),
    ("株式会社MIRAI", "https://mirai-tantei.jp/lp/uwaki_ls/organic01/", "リスティング広告", "", "浮気調査・探偵サービス", "浮気調査 おすすめ, 探偵事務所 口コミ, 不倫 証拠 集め方", "探偵・調査"),
    ("IdeaLink株式会社", "https://www.stretch-up.jp/lp/shinsaibashi02/", "リスティング広告", "", "ストレッチ専門スタジオ", "ストレッチ 専門店, 体のコリ 解消, ストレッチ 整体 違い", "フィットネス・ストレッチ"),
    ("株式会社All in Motions", "https://allinmotions.co.jp/lp/index.html", "リスティング広告", "", "動画編集スクール", "動画編集 スクール, 動画編集 副業 初心者, YouTube 編集 学ぶ", "動画編集スクール"),
    ("株式会社日本デザイン", "https://japan-design.jp/lp/dssp_1h/hd1_k.html", "リスティング広告", "", "Webデザインスクール", "Webデザイン スクール, デザイン 学校 社会人, ウェブデザイン 資格", "Webデザインスクール"),
    ("株式会社クラスタイル", "http://www.clastyle.com/lp/tsugaku/003/osaka.html", "リスティング広告", "", "ネイルスクール", "ネイルスクール 通学, ネイリスト 資格 取得, ネイル 学校 費用", "ネイルスクール"),
    ("株式会社受験ドクター", "https://www.chugakujuken.com/lp/live/", "リスティング広告", "", "中学受験個別指導", "中学受験 個別指導, 中学受験 塾 おすすめ, 受験 家庭教師 小学生", "中学受験・個別指導"),
    ("株式会社メガネトップ", "https://www.meganeichiba.jp/lp/hearingaid/", "リスティング広告", "", "補聴器・メガネ", "補聴器 おすすめ 価格, 補聴器 安い 種類, 難聴 補聴器 選び方", "補聴器・メガネ"),
    ("SBIプリズム少額短期保険株式会社", "https://jac.app.sbiprism.co.jp/lp/prismpet/smalls.html", "リスティング広告", "", "ペット保険（小動物）", "ペット保険 ハムスター, 小動物 保険, うさぎ ペット保険 比較", "ペット保険"),
    ("株式会社GronG", "https://shop.grong.jp/blogs/lp/trial-pack-recommend/", "リスティング広告", "", "プロテイン・スポーツ栄養", "プロテイン 国産 安い, プロテイン 初心者 おすすめ, ホエイプロテイン 比較", "スポーツ栄養・健康食品EC"),
]


def make_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "架電リスト"

    # Header
    for col_idx, header in enumerate(HEADER, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
    ws.row_dimensions[1].height = 20

    # Data
    for row_idx, row_data in enumerate(DATA, 2):
        fill = ROW_FILLS[(row_idx - 2) % 2]
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.fill = fill
            cell.alignment = ROW_ALIGN
            if col_idx == 2:  # LP URL column
                cell.hyperlink = value
                cell.font = URL_FONT
            else:
                cell.font = ROW_FONT
        ws.row_dimensions[row_idx].height = 16

    # Column widths
    for col_idx, width in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # Freeze pane & auto filter
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADER))}1"

    wb.save(OUTPUT_PATH)
    print(f"Saved: {OUTPUT_PATH}")
    print(f"Rows: {len(DATA)}")


if __name__ == "__main__":
    make_xlsx()
