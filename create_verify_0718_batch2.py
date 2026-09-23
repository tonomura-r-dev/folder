import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUTPUT_PATH = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー\架電リスト_全社再検証_2026-07-18_batch2.xlsx"
HEADER = ["企業名", "LP URL(広告確認済み)", "電話番号", "検索KW", "媒体"]
COL_WIDTHS = [40, 50, 16, 32, 14]
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="メイリオ", size=10, bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
ROW_FILLS = [PatternFill("solid", fgColor="F2F7FC"), PatternFill("solid", fgColor="FFFFFF")]
ROW_FONT = Font(name="メイリオ", size=9)
ROW_ALIGN = Alignment(vertical="center", wrap_text=False)
URL_FONT = Font(name="メイリオ", size=9, color="0563C1", underline="single")

DATA = [
    ("株式会社スタンダード（個別指導塾スタンダード）", "https://std-ie.jp/", "092-283-1188", "個別指導塾, 塾 体験授業", "Meta"),
    ("株式会社アユーラ（AYURA）", "https://www.ayura.co.jp/", "0120-090-030", "アユーラ, スキンケア 敏感肌", "Meta"),
    ("株式会社カネボウ化粧品（リサージ）", "https://www.lissage.jp/", "0120-518-520", "リサージ, スキンケア 大人", "Meta"),
    ("上新電機株式会社（Joshin）", "https://joshinweb.jp/", "06-6631-1221", "ジョーシン, 家電 通販", "Meta"),
    ("株式会社東京個別指導学院", "https://www.kobetsu.co.jp/", "03-6911-3246", "個別指導塾 東京, 高校受験 塾", "Meta"),
    ("グレイル株式会社（GRL）", "https://www.grail.bz/", "06-6532-2211", "GRL グレイル, プチプラ 服 通販", "Meta"),
    ("株式会社ストライプインターナショナル（Maison de FLEUR）", "https://stripe-club.com/", "086-235-8216", "メゾンドフルール, レディース 通販", "Meta"),
    ("株式会社中萬学院", "https://www.chuman.co.jp/", "045-840-1701", "中萬学院, 神奈川 学習塾", "Meta"),
    ("株式会社湘南ゼミナール", "https://www.shozemi.com/", "045-565-5100", "湘南ゼミナール, 神奈川 塾", "Meta"),
    ("バルミューダ株式会社", "https://www.balmuda.com/", "0120-686-717", "バルミューダ, デザイン家電", "Meta"),
    ("株式会社シロカ", "https://www.siroca.co.jp/", "03-3234-5490", "シロカ, 生活家電 通販", "Meta"),
    ("テスコム電機株式会社", "https://www.tescom-japan.co.jp/", "03-5719-2061", "テスコム, 美容家電 ハンドミキサー", "Meta"),
    ("カゴメ株式会社", "https://shop.kagome.co.jp/", "0120-401-831", "カゴメ 通販, スムージー", "Meta"),
    ("株式会社フィールドア（FIELDOOR）", "https://fieldoor.com/", "03-5772-1251", "FIELDOOR, アウトドア用品", "Meta"),
]

NOT_VERIFIED = [
    "パール金属株式会社", "株式会社京進", "ポラス株式会社", "ヤマダホームズ株式会社",
    "株式会社新昭和", "株式会社ニチレイフーズ", "株式会社マークスアンドウェブ", "イプサカンパニー株式会社",
    "株式会社バイクブロス", "株式会社ツクイ", "株式会社日本旅行", "株式会社カルド",
    "クラウドクレジット株式会社", "ファンズ株式会社", "株式会社STAY JAPAN", "株式会社ニッコートラベル",
    "株式会社ITTO個別指導学院", "株式会社キャリアリンク", "フォーラムエンジニアリング株式会社", "UTグループ株式会社",
    "テクノプロ・ホールディングス株式会社", "株式会社ディーホリック", "フランスベッド株式会社", "株式会社ファインラボ",
    "ユーハ味覚糖株式会社", "株式会社リプサ", "コンビ株式会社", "医療法人ヴォーグ",
    "合同会社YAMAtoMICHI", "株式会社ルートート", "株式会社コクーニスト", "株式会社コーセー",
    "株式会社フーミー", "日神不動産株式会社", "オリックス自動車株式会社", "KINTO株式会社",
    "株式会社Mr.チーズケーキ", "ル・クルーゼ ジャポン株式会社", "株式会社タニタ", "オムロンヘルスケア株式会社",
    "株式会社セルヴォーク",
]


def make_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "確認済み"
    for col_idx, header in enumerate(HEADER, 1):
        c = ws.cell(row=1, column=col_idx, value=header)
        c.fill = HEADER_FILL; c.font = HEADER_FONT; c.alignment = HEADER_ALIGN
    ws.row_dimensions[1].height = 20
    for row_idx, row in enumerate(DATA, 2):
        fill = ROW_FILLS[(row_idx - 2) % 2]
        for col_idx, value in enumerate(row, 1):
            c = ws.cell(row=row_idx, column=col_idx, value=value)
            c.fill = fill; c.alignment = ROW_ALIGN
            if col_idx == 2:
                c.hyperlink = value; c.font = URL_FONT
            else:
                c.font = ROW_FONT
        ws.row_dimensions[row_idx].height = 16
    for col_idx, width in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADER))}1"

    ws2 = wb.create_sheet("未確認")
    ws2.cell(row=1, column=1, value="広告ライブラリで配信中の広告が見つからなかった企業（要再調査・除外候補）").font = Font(bold=True)
    for i, name in enumerate(NOT_VERIFIED, 2):
        ws2.cell(row=i, column=1, value=name)
    ws2.column_dimensions["A"].width = 55

    wb.save(OUTPUT_PATH)
    print(f"Saved: {OUTPUT_PATH}  Confirmed: {len(DATA)}  Not verified: {len(NOT_VERIFIED)}")


if __name__ == "__main__":
    make_xlsx()
