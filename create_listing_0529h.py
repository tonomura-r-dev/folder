import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUTPUT_PATH = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー\架電リスト_リスティング_2026-05-29h.xlsx"

HEADER = ["企業名", "LP URL", "広告種別", "電話番号", "商材", "検索KW", "業界"]
COL_WIDTHS = [32, 50, 16, 16, 20, 32, 22]

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="メイリオ", size=10, bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)

ROW_FILLS = [PatternFill("solid", fgColor="F2F7FC"), PatternFill("solid", fgColor="FFFFFF")]
ROW_FONT = Font(name="メイリオ", size=9)
ROW_ALIGN = Alignment(vertical="center", wrap_text=False)
URL_FONT = Font(name="メイリオ", size=9, color="0563C1", underline="single")

ROWS = [
    ["ポケットカード株式会社",            "https://fcard.pocketcard.co.jp/lp/unify-f202603/",           "リスティング広告", "",               "クレジットカード（Fカード）",             "Fカード 発行 審査, クレジットカード 年会費無料 ポイント",    "クレジットカード"],
    ["大樹生命保険株式会社",              "https://www.taiju-life.co.jp/products/lp/taiju_gannohoken.htm", "リスティング広告", "",             "がん保険",                               "がん保険 比較 大樹生命, がん診断 一時金, がん保険 申込",     "生命保険"],
    ["株式会社ニチリョク",                "https://www.nichiryoku.co.jp/lp/public_houjyou.html",         "リスティング広告", "",               "霊園・墓地",                             "霊園 区画 購入, 永代供養 相談, 墓地 申込 受付",              "霊園・葬儀"],
    ["株式会社カナエキモノハーツ",        "https://portal.kimono-hearts.co.jp/lp/mamafuri/",            "リスティング広告", "",               "振袖レンタル（ママ振）",                 "振袖 ママ振 レンタル, 成人式 着物 費用, 振袖 購入 比較",     "振袖・着物レンタル"],
    ["株式会社ベビーリース",              "https://www.nicebaby.co.jp/lp/lp_catalog/index.html",        "リスティング広告", "",               "ベビー用品レンタル",                     "ベビーカー レンタル, ベビー用品 月額 安い, チャイルドシート", "ベビー用品レンタル"],
    ["株式会社コスモライフ",              "https://www.hummingwater.com/lp/general02/",                  "リスティング広告", "",               "ウォーターサーバー",                     "ウォーターサーバー 月額 安い, 天然水 定期配達, 水 宅配",     "ウォーターサーバー"],
    ["株式会社鎌倉新書",                  "https://www.e-sogi.com/lp/safety/",                          "リスティング広告", "",               "葬儀比較・紹介サービス",                 "葬儀社 比較 見積もり, 家族葬 費用 相場, いい葬儀 相談",      "葬儀比較"],
    ["株式会社みんなのウエディング",      "https://myfavoritepart.mwed.jp/lp/resort-wedding-japan/",    "リスティング広告", "",               "結婚式場口コミ比較",                     "結婚式場 口コミ, リゾートウエディング 費用, 式場 比較",       "婚礼情報"],
    ["株式会社チャーム・ケア・コーポレーション", "https://www.charmcc.jp/lp/premiergrand4/",           "リスティング広告", "",               "介護付有料老人ホーム",                   "老人ホーム 入居 相談, 介護施設 費用 目安, 有料老人ホーム",   "介護施設"],
    ["株式会社ホリデー",                  "https://www.holiday-fc.co.jp/lp/tokyo",                      "リスティング広告", "",               "車検（ホリデー車検）",                   "車検 安い 東京, 車検 予約 即日, 格安車検 近く",               "車検"],
    ["株式会社Cotree",                   "https://cotree.jp/lp/video_tel_counseling",                   "リスティング広告", "",               "オンラインカウンセリング",               "カウンセリング オンライン, 心理士 相談 予約, メンタルヘルス", "心理・カウンセリング"],
    ["株式会社日本トリム",                "https://www.nihon-trim.co.jp/lp/cure_details_d_000/",         "リスティング広告", "",               "整水器・電解水素水",                     "電解水素水 整水器, 水素水 効果 家庭用, アルカリイオン水",    "健康機器"],
    ["株式会社メディプラス",              "https://mediplus-orders.jp/lp/pc/1791392/3218/",             "リスティング広告", "",               "スキンケア定期購入",                     "スキンケア 定期, 敏感肌 化粧水 通販, オルゴーラン 効果",     "スキンケア通販"],
    ["株式会社中央コンタクト",            "https://www.chuo-contact.co.jp/lp/",                         "リスティング広告", "",               "コンタクトレンズ通販",                   "コンタクトレンズ 通販, 1日 使い捨て 安い, コンタクト 購入",  "コンタクトレンズ通販"],
    ["学校法人角川ドワンゴ学園",          "https://nnn.ed.jp/lp/hs01_ori_01/",                          "リスティング広告", "",               "通信制高校（N高等学校）",                "通信制高校 N高, 高校 ネット 通信, 不登校 高校 資料請求",     "通信制高校"],
    ["株式会社Kaien",                    "https://mlg.kaien-lab.com/lp/jobassistance",                  "リスティング広告", "",               "発達障害就労移行支援",                   "発達障害 就労移行, ADHD 仕事 支援, 就労支援 体験 相談",      "就労移行支援"],
    ["学校法人佐藤学園",                  "https://www.hchs.ed.jp/lp/sougou/",                          "リスティング広告", "",               "通信制高校（ヒューマンキャンパス）",     "ヒューマンキャンパス 高校, 通信制高校 資料, 不登校 高校",    "通信制高校"],
    ["学校法人KTC学園",                   "https://www.ohzora.ac.jp/lp/20180801/tsushin/a/",            "リスティング広告", "",               "通信制高校（KTCおおぞら）",              "おおぞら高等学院, 通信制高校 自分のペース, 不登校 転校",      "通信制高校"],
    ["株式会社リナビス",                  "https://rinavis.com/s/leather/lp/",                          "リスティング広告", "058-391-6906",   "宅配クリーニング（革製品・ブランド品）", "レザー クリーニング 宅配, ブランドバッグ 洗い, 革製品 洗濯", "宅配クリーニング"],
    ["REXT株式会社",                     "https://www.wonderrex.jp/lp/hikkoshi/",                       "リスティング広告", "",               "リサイクルショップ・引越し宅買便",       "引越し 不用品買取, 宅配買取 引越し, リサイクル 宅配 申込",   "リユース・リサイクル"],
    ["株式会社SODA",                     "https://snkrdunk.com/lp/sneaker-listing-campaign",            "リスティング広告", "",               "スニーカー・ブランド品売買",             "スニーカー 売買, スニダン 出品, 限定スニーカー 購入 正規品",  "フリマ・リセール"],
]

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "架電リスト"

# ヘッダー行
for col, (h, w) in enumerate(zip(HEADER, COL_WIDTHS), 1):
    cell = ws.cell(row=1, column=col, value=h)
    cell.fill = HEADER_FILL
    cell.font = HEADER_FONT
    cell.alignment = HEADER_ALIGN
    ws.column_dimensions[get_column_letter(col)].width = w

ws.row_dimensions[1].height = 20

# データ行
for row_idx, row_data in enumerate(ROWS, 2):
    fill = ROW_FILLS[(row_idx - 2) % 2]
    for col_idx, value in enumerate(row_data, 1):
        cell = ws.cell(row=row_idx, column=col_idx, value=value)
        cell.fill = fill
        cell.alignment = ROW_ALIGN
        if col_idx == 2 and value:
            cell.font = URL_FONT
            cell.hyperlink = value
        else:
            cell.font = ROW_FONT
    ws.row_dimensions[row_idx].height = 16

ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADER))}1"

wb.save(OUTPUT_PATH)
print(f"保存完了: {OUTPUT_PATH}")
print(f"行数: {len(ROWS)}社")
