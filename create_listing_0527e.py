import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUTPUT_PATH = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー\架電リスト_リスティング_2026-05-27e.xlsx"
HEADER = ["企業名", "LP URL", "広告種別", "電話番号", "商材", "検索KW", "業界"]
COL_WIDTHS = [32, 50, 16, 16, 20, 32, 22]

DATA = [
    [
        "株式会社Builds（スタディコーチ）",
        "https://studycoach.co.jp/lp/",
        "リスティング広告",
        "03-6721-1160",
        "難関大受験向けオンライン個別コーチング",
        "受験コーチング 難関大受験 個別指導 オンライン塾",
        "教育・スキルアップ",
    ],
    [
        "株式会社福屋不動産販売",
        "https://www.fukuya-k.co.jp/lp/sellers",
        "リスティング広告",
        "0120-298-195",
        "不動産売却・査定仲介サービス",
        "不動産売却 不動産査定 マンション売却 土地売却",
        "不動産・賃貸・リノベーション・投資",
    ],
    [
        "リビン・テクノロジーズ株式会社（リビンマッチ）",
        "https://www.lvnmatch.jp/bp/lp/",
        "リスティング広告",
        "03-5847-8558",
        "不動産一括査定サービス",
        "不動産査定 一括査定 不動産売却 マンション売却",
        "不動産・賃貸・リノベーション・投資",
    ],
    [
        "株式会社ユニマットライフ",
        "https://www.unimat-shop.jp/water/lp/",
        "リスティング広告",
        "0120-346-634",
        "宅配水ウォーターサーバーサービス",
        "ウォーターサーバー 宅配水 天然水 水宅配",
        "食品EC・宅配弁当・ミールキット",
    ],
    [
        "株式会社FPパートナー（マネードクター）",
        "https://fp-moneydoctor.com/lp/mdmk003/",
        "リスティング広告",
        "0120-366-043",
        "無料ファイナンシャル・保険相談サービス",
        "保険相談 FP相談 無料相談 ファイナンシャルプランナー マネードクター",
        "保険・金融・証券・カード",
    ],
    [
        "株式会社レアジョブ（RareJob英会話）",
        "https://www.rarejob.com/campaign/lp/22my02",
        "リスティング広告",
        "03-4477-3595",
        "オンライン英会話レッスン（フィリピン人講師）",
        "オンライン英会話 英会話 英語学習 英語 レアジョブ",
        "教育・スキルアップ",
    ],
    [
        "株式会社AQ Group（アキュラホーム）",
        "https://www.aqura.co.jp/lp/wagaya1/",
        "リスティング広告",
        "0120-984-351",
        "自由設計注文住宅（高性能・適正価格）",
        "注文住宅 工務店 家づくり 一戸建て マイホーム アキュラホーム",
        "不動産・賃貸・リノベーション・投資",
    ],
    [
        "チューリッヒ保険会社",
        "https://www.zurich.co.jp/lp/adm/auto/kakakuhoujin/",
        "リスティング広告",
        "0120-860-001",
        "ダイレクト型自動車保険（個人・法人対応）",
        "自動車保険 ダイレクト保険 車保険 任意保険 チューリッヒ",
        "保険・金融・証券・カード",
    ],
    [
        "株式会社ユーキャン",
        "https://www.u-can.co.jp/topics/lp/LP_00392/",
        "リスティング広告",
        "0120-552-476",
        "資格・通信講座（宅建・医療事務・介護など140講座）",
        "通信講座 資格取得 通信教育 ユーキャン 資格",
        "教育・スキルアップ",
    ],
    [
        "株式会社オーネット（ONet）",
        "https://www.onet.co.jp/lp/index_area.html",
        "リスティング広告",
        "0120-135-029",
        "結婚相談所・婚活サービス",
        "婚活 結婚相談所 お見合い 婚活サービス 婚活パーティー",
        "婚活・マッチングアプリ",
    ],
]


def make_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "架電リスト"

    header_font = Font(name="メイリオ", size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    fill_odd = PatternFill("solid", fgColor="F2F7FC")
    fill_even = PatternFill("solid", fgColor="FFFFFF")
    data_font = Font(name="メイリオ", size=9)
    data_align = Alignment(vertical="center", wrap_text=False)
    url_font = Font(name="メイリオ", size=9, color="0563C1", underline="single")

    thin = Side(style="thin", color="D0D7DE")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # ヘッダー行
    for col_idx, col_name in enumerate(HEADER, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = border

    ws.row_dimensions[1].height = 22

    # データ行
    for row_idx, row_data in enumerate(DATA, start=2):
        fill = fill_odd if (row_idx % 2 == 0) else fill_even
        for col_idx, value in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.fill = fill
            cell.border = border
            if col_idx == 2:  # LP URL列
                cell.font = url_font
                cell.hyperlink = value
                cell.alignment = data_align
            else:
                cell.font = data_font
                cell.alignment = data_align
        ws.row_dimensions[row_idx].height = 18

    # 列幅
    for col_idx, width in enumerate(COL_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # 先頭行固定・オートフィルター
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADER))}1"

    wb.save(OUTPUT_PATH)
    print(f"保存完了: {OUTPUT_PATH}")
    print(f"件数: {len(DATA)}社")


if __name__ == "__main__":
    make_xlsx()
