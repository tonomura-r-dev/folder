import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUTPUT_PATH = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー\架電リスト_リスティング_2026-06-03br.xlsx"
HEADER = ["企業名", "LP URL", "広告種別", "法人番号", "商材", "検索KW", "業界"]
COL_WIDTHS = [32, 56, 16, 16, 22, 32, 22]
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="メイリオ", size=10, bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
ROW_FILLS = [PatternFill("solid", fgColor="F2F7FC"), PatternFill("solid", fgColor="FFFFFF")]
ROW_FONT = Font(name="メイリオ", size=9)
ROW_ALIGN = Alignment(vertical="center", wrap_text=False)
URL_FONT = Font(name="メイリオ", size=9, color="0563C1", underline="single")

DATA = [
    # ── 住宅設備交換EC（日本・実LP） ──
    ("株式会社交換できるくん",      "https://www.sunrefre.jp/",                                "リスティング広告", "", "給湯器・トイレ交換・住設EC・10年保証",        "給湯器 交換 即日, トイレ 交換 費用, 交換できるくん 評判",          "住宅設備"),
    # ── 生活トラブル110番（日本・実LP） ──
    ("株式会社シェアリングテクノロジー","https://www.sharing-tech.co.jp/gaichu/",                "リスティング広告", "", "害虫駆除110番・生活110番・緊急駆けつけ",       "ゴキブリ 駆除 即日, 害虫駆除 業者, 蜂 駆除 費用",                 "生活サービス"),
    # ── オンライン秘書（日本・実LP） ──
    ("BPOテクノロジー株式会社",     "https://fujiko-san.com/",                                 "リスティング広告", "", "フジ子さん・オンラインアシスタント・秘書代行", "オンライン 秘書 比較, オンラインアシスタント 料金, フジ子さん 評判","BPOサービス"),
    # ── 探偵・浮気調査（日本・実LP） ──
    ("株式会社原一",               "https://www.haraichi.co.jp/",                             "リスティング広告", "", "原一探偵事務所・浮気調査・無料相談",          "浮気 調査 探偵, 浮気調査 費用, 原一探偵事務所 評判",              "調査サービス"),
    # ── 営業データベース ──
    ("Apollo.io",                  "https://www.apollo.io/pricing",                           "リスティング広告", "", "Apollo・セールスインテリジェンス・営業DB",     "営業 リスト ツール, Apollo.io 料金 評判, セールス DB SaaS",        "セールスSaaS"),
    # ── コールドメール ──
    ("Smartlead.ai",              "https://www.smartlead.ai/pricing",                        "リスティング広告", "", "Smartlead・コールドメール・アウトバウンド",   "コールドメール ツール, Smartlead 料金 評判, 営業 メール 自動化",  "セールスSaaS"),
    # ── データエンリッチメント ──
    ("Clay",                       "https://www.clay.com/pricing",                            "リスティング広告", "", "Clay・データエンリッチメント・営業自動化",     "リード エンリッチメント SaaS, Clay 料金 評判, 営業 データ 自動化","セールスSaaS"),
    # ── コールドメール ──
    ("Instantly.ai",              "https://instantly.ai/pricing",                            "リスティング広告", "", "Instantly・コールドメール・送信ウォームアップ","コールドメール 配信 SaaS, Instantly 料金 評判, 営業 メール ツール","セールスSaaS"),
    # ── ノーコードWebアプリ ──
    ("Bubble",                     "https://bubble.io/pricing",                               "リスティング広告", "", "Bubble・ノーコードWebアプリ開発",             "ノーコード Web アプリ, Bubble 料金 評判, ノーコード 開発 ツール",  "ノーコードSaaS"),
    # ── ノーコードモバイル ──
    ("FlutterFlow",                "https://www.flutterflow.io/pricing",                      "リスティング広告", "", "FlutterFlow・ノーコードモバイルアプリ開発",   "ノーコード アプリ 開発, FlutterFlow 料金 評判, モバイル アプリ 作成","ノーコードSaaS"),
    # ── ノーコードアプリ ──
    ("Glide",                      "https://www.glideapps.com/pricing",                       "リスティング広告", "", "Glide・スプレッドシートからアプリ作成",        "ノーコード アプリ Glide, Glide 料金 評判, 業務 アプリ 作成",       "ノーコードSaaS"),
    # ── ノーコードWebサイト/アプリ ──
    ("Softr",                      "https://www.softr.io/pricing",                            "リスティング広告", "", "Softr・ノーコードWebアプリ/ポータル",         "ノーコード ポータル サイト, Softr 料金 評判, 業務 アプリ ノーコード","ノーコードSaaS"),
    # ── 社内ツール開発 ──
    ("Retool",                     "https://retool.com/pricing",                              "リスティング広告", "", "Retool・社内ツール/業務アプリ高速開発",       "社内 ツール 開発 SaaS, Retool 料金 評判, 業務 アプリ ローコード", "ローコードSaaS"),
    # ── 営業インテリジェンス ──
    ("ZoomInfo",                   "https://www.zoominfo.com/pricing",                        "リスティング広告", "", "ZoomInfo・B2B企業データ・営業インテリジェンス","B2B 企業 データ SaaS, ZoomInfo 料金 評判, 営業 リスト 取得",      "セールスSaaS"),
    # ── コールドメール ──
    ("lemlist",                    "https://www.lemlist.com/pricing",                         "リスティング広告", "", "lemlist・コールドメール・パーソナライズ",     "コールドメール SaaS, lemlist 料金 評判, 営業 アウトバウンド ツール","セールスSaaS"),
    # ── Webスクレイピング ──
    ("Apify",                      "https://apify.com/pricing",                               "リスティング広告", "", "Apify・Webスクレイピング/自動化プラットフォーム","Web スクレイピング ツール, Apify 料金 評判, データ 収集 自動化",  "データSaaS"),
    # ── ワークフロー自動化 ──
    ("n8n",                        "https://n8n.io/pricing",                                  "リスティング広告", "", "n8n・OSSワークフロー自動化・AI連携",          "ワークフロー 自動化 OSS, n8n 料金 評判, 業務 自動化 ノーコード",  "業務自動化SaaS"),
    # ── 金融API ──
    ("Plaid",                      "https://plaid.com/pricing/",                              "リスティング広告", "", "Plaid・金融データ連携API・口座連携",          "金融 API データ 連携, Plaid 料金 評判, 口座 連携 フィンテック",    "フィンテックSaaS"),
    # ── スタートアップ向けバンキング ──
    ("Mercury",                    "https://mercury.com/pricing",                             "リスティング広告", "", "Mercury・スタートアップ向けビジネスバンキング","スタートアップ 銀行 サービス, Mercury 評判, ビジネス バンキング",  "フィンテックSaaS"),
    # ── メールアドレス検索 ──
    ("Hunter.io",                  "https://hunter.io/pricing",                               "リスティング広告", "", "Hunter・メールアドレス検索/検証・営業",       "メールアドレス 検索 ツール, Hunter.io 料金 評判, 営業 メール 取得","セールスSaaS"),
]


def make_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "架電リスト"

    for col_idx, header in enumerate(HEADER, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
    ws.row_dimensions[1].height = 20

    for row_idx, row_data in enumerate(DATA, 2):
        fill = ROW_FILLS[(row_idx - 2) % 2]
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.fill = fill
            cell.alignment = ROW_ALIGN
            if col_idx == 2:
                cell.hyperlink = value
                cell.font = URL_FONT
            else:
                cell.font = ROW_FONT
        ws.row_dimensions[row_idx].height = 16

    for col_idx, width in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADER))}1"

    wb.save(OUTPUT_PATH)
    print(f"Saved: {OUTPUT_PATH}")
    print(f"Rows: {len(DATA)}")


if __name__ == "__main__":
    make_xlsx()
