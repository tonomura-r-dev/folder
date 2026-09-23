import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def make_xlsx(filename, data):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "架電リスト"

    headers = ["企業名", "LP URL", "電話番号", "商材", "検索KW", "業界"]

    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="メイリオ", bold=True, color="FFFFFF", size=10)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center
        cell.border = border

    for row_idx, row in enumerate(data, 2):
        fill_color = "F2F7FC" if row_idx % 2 == 0 else "FFFFFF"
        row_fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
        for col_idx, val in enumerate(row, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.fill = row_fill
            cell.font = Font(name="メイリオ", size=9)
            cell.alignment = left
            cell.border = border

    col_widths = [30, 45, 16, 20, 30, 18]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.row_dimensions[1].height = 20
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    wb.save(filename)
    print(f"保存完了: {filename}  行数: {len(data)}")


# 第153弾（2026-07-16e）法律事務所・司法書士法人・税理士法人 9社・LP URLは相談ページ

data_0716e = [
    ["弁護士法人グレイス", "https://www.grace-law.jp/contact/", "03-6432-9783", "離婚問題（親権・養育費・財産分与・慰謝料・モラハラ/DV対応）", "離婚 弁護士, モラハラ 慰謝料, 親権 弁護士", "離婚問題弁護士"],
    ["弁護士法人東京新宿法律事務所", "https://www.shinjuku-law.jp/services/traffic/", "03-5339-0356", "交通事故・労働問題・相続・離婚等の個人向け法律相談", "交通事故 弁護士, 後遺障害 慰謝料, 残業代請求 弁護士", "交通事故弁護士"],
    ["円満相続税理士法人", "https://osd-souzoku.jp/contact/", "", "相続税申告・生前対策専門コンサルティング", "相続税 申告 税理士, 生前対策 相続, 相続税 節税", "相続専門税理士法人"],
    ["弁護士法人エース", "https://ace-law.or.jp/zangyodai/", "03-6625-4140", "残業代請求・不当解雇・交通事故・離婚・相続・企業法務", "残業代請求 弁護士, 未払い残業代 相談, 不当解雇 弁護士", "労働問題弁護士"],
    ["弁護士法人春田法律事務所", "https://haruta-lo.com/case/rikon/", "", "離婚・男女問題、相続、刑事事件等", "離婚 弁護士 相談, 財産分与 弁護士, 離婚 慰謝料", "離婚問題弁護士"],
    ["司法書士法人NCP", "https://www.ncp-law.com/shihoushoshi/", "03-5367-5930", "相続手続き・遺言書作成サポート", "相続手続き 司法書士, 遺言書 作成, 相続登記 代行", "相続専門司法書士法人"],
    ["杠司法書士法人", "https://olao.jp/contact/", "06-6253-7707", "企業法務・相続・遺言・M&A・事業承継・成年後見", "企業法務 司法書士, 事業承継 相談, 相続 遺言 司法書士", "企業法務司法書士法人"],
    ["税理士法人朝日中央綜合事務所", "https://asahichuo-tax.jp/inheritance-tax-reduction-consulting/", "", "相続税申告・相続税軽減対策・事業承継税制コンサルティング", "相続税 軽減対策, 相続税 税理士, 事業承継 税理士", "相続・事業承継専門税理士法人"],
    ["弁護士法人小杉法律事務所", "https://personal-injury.jp/contact", "050-3668-1155", "交通事故・労災事故・学校事故等の被害者側損害賠償請求", "交通事故 弁護士 後遺障害, 労災事故 弁護士, 慰謝料請求 弁護士", "交通事故・労災弁護士"],
]

base = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー"
make_xlsx(f"{base}\\架電リスト_Meta_2026-07-16e.xlsx", data_0716e)
