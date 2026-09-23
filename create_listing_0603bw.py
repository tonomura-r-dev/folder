import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUTPUT_PATH = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー\架電リスト_リスティング_2026-06-03bw.xlsx"
HEADER = ["企業名", "LP URL", "広告種別", "法人番号", "商材", "検索KW", "業界"]
COL_WIDTHS = [32, 56, 16, 16, 22, 32, 22]
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="メイリオ", size=10, bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
ROW_FILLS = [PatternFill("solid", fgColor="F2F7FC"), PatternFill("solid", fgColor="FFFFFF")]
ROW_FONT = Font(name="メイリオ", size=9)
ROW_ALIGN = Alignment(vertical="center", wrap_text=False)
URL_FONT = Font(name="メイリオ", size=9, color="0563C1", underline="single")

DATA = [
    # ── オープンソースBI ──
    ("Metabase",                   "https://www.metabase.com/pricing",                        "リスティング広告", "", "Metabase・オープンソースBI・ダッシュボード",   "BI ツール OSS, Metabase 料金 評判, データ 可視化 SaaS",            "データSaaS"),
    # ── クラウドBI ──
    ("Sigma Computing",            "https://www.sigmacomputing.com/pricing",                  "リスティング広告", "", "Sigma・スプレッドシート型クラウドBI",          "クラウド BI ツール, Sigma 料金 評判, データ 分析 スプレッドシート","データSaaS"),
    # ── BI/分析 ──
    ("Mode Analytics",             "https://mode.com/pricing",                                "リスティング広告", "", "Mode・SQL/BI分析・データチーム向け",          "BI 分析 SaaS, Mode 料金 評判, SQL データ 分析 ツール",             "データSaaS"),
    # ── BI/分析 ──
    ("ThoughtSpot",                "https://www.thoughtspot.com/pricing",                     "リスティング広告", "", "ThoughtSpot・検索型BI・AI分析",                "検索型 BI SaaS, ThoughtSpot 料金 評判, AI データ 分析",            "データSaaS"),
    # ── BI/分析 ──
    ("Holistics",                  "https://www.holistics.io/pricing/",                       "リスティング広告", "", "Holistics・BI/データモデリング・セルフサービス","BI ツール 比較, Holistics 料金 評判, データ モデリング SaaS",      "データSaaS"),
    # ── BI（Superset） ──
    ("Preset",                     "https://preset.io/pricing",                               "リスティング広告", "", "Preset・Apache SupersetベースBI",             "BI OSS SaaS, Preset 料金 評判, Superset データ 可視化",           "データSaaS"),
    # ── 分析/ノートブック ──
    ("Hex",                        "https://hex.tech/pricing/",                               "リスティング広告", "", "Hex・データノートブック/分析・AI",            "データ 分析 ノートブック, Hex 料金 評判, データ チーム ツール",   "データSaaS"),
    # ── BI（OSS/dbt連携） ──
    ("Lightdash",                  "https://www.lightdash.com/pricing",                       "リスティング広告", "", "Lightdash・dbt連携BI・OSS",                   "dbt BI ツール, Lightdash 料金 評判, データ 可視化 OSS",           "データSaaS"),
    # ── リソース管理 ──
    ("Float",                      "https://www.float.com/pricing",                           "リスティング広告", "", "Float・リソース管理/プロジェクト計画",         "リソース 管理 SaaS, Float 料金 評判, プロジェクト 計画 ツール",    "プロジェクト管理SaaS"),
    # ── リソース管理 ──
    ("Resource Guru",              "https://resourceguruapp.com/pricing",                     "リスティング広告", "", "Resource Guru・リソース/スケジュール管理",     "リソース 管理 SaaS, Resource Guru 料金 評判, チーム スケジュール",  "プロジェクト管理SaaS"),
    # ── PSA/プロジェクト ──
    ("Productive",                 "https://productive.io/pricing/",                          "リスティング広告", "", "Productive・PSA/プロジェクト/収益管理",        "PSA ツール SaaS, Productive 料金 評判, プロジェクト 収益 管理",     "プロジェクト管理SaaS"),
    # ── リソース計画 ──
    ("Runn",                       "https://www.runn.io/pricing",                             "リスティング広告", "", "Runn・リソース計画/キャパシティ管理",          "リソース 計画 SaaS, Runn 料金 評判, キャパシティ 管理 ツール",     "プロジェクト管理SaaS"),
    # ── リソース管理 ──
    ("Hub Planner",               "https://hubplanner.com/pricing/",                         "リスティング広告", "", "Hub Planner・リソース管理/スケジューリング",   "リソース 管理 ツール, Hub Planner 料金 評判, チーム 計画 SaaS",    "プロジェクト管理SaaS"),
    # ── 時間管理 ──
    ("Toggl Track",                "https://toggl.com/track/pricing/",                        "リスティング広告", "", "Toggl Track・時間管理/工数トラッキング",       "時間 管理 工数 SaaS, Toggl 料金 評判, タイム トラッキング ツール", "業務DXSaaS"),
    # ── データスタック ──
    ("Mozart Data",               "https://www.mozartdata.com/pricing",                      "リスティング広告", "", "Mozart Data・モダンデータスタック・ELT",      "データ スタック SaaS, Mozart Data 料金 評判, ELT データ 統合",     "データSaaS"),
    # ── ELT/データ統合 ──
    ("Airbyte",                    "https://airbyte.com/pricing",                             "リスティング広告", "", "Airbyte・ELT/データ統合・OSSコネクタ",        "ELT データ 統合 OSS, Airbyte 料金 評判, データ パイプライン SaaS", "データSaaS"),
    # ── コラボ分析 ──
    ("Count",                      "https://count.co/pricing",                                "リスティング広告", "", "Count・コラボ型データ分析キャンバス",          "データ 分析 コラボ SaaS, Count 料金 評判, データ キャンバス ツール","データSaaS"),
    # ── データプラットフォーム ──
    ("Y42",                        "https://www.y42.com/pricing",                             "リスティング広告", "", "Y42・モダンデータスタック/オーケストレーション","データ オーケストレーション SaaS, Y42 料金 評判, データ パイプライン","データSaaS"),
    # ── リバースETL ──
    ("Census",                     "https://www.getcensus.com/pricing",                       "リスティング広告", "", "Census・リバースETL/データ活性化",            "リバース ETL SaaS, Census 料金 評判, データ 連携 ツール",          "データSaaS"),
    # ── ETL/データ統合 ──
    ("Stitch",                     "https://www.stitchdata.com/pricing/",                     "リスティング広告", "", "Stitch・ETL/データパイプライン・シンプル",     "ETL ツール SaaS, Stitch 料金 評判, データ 統合 パイプライン",      "データSaaS"),
]


def make_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "架電リスト"

    for col_idx, header in enumerate(HEADER, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
    ws.row_dimensions[1].height = 20

    for row_idx, row_data in enumerate(DATA, 2):
        fill = ROW_FILLS[(row_idx - 2) % 2]
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.fill = fill
            cell.alignment = ROW_ALIGN
            if col_idx == 2:
                cell.hyperlink = value
                cell.font = URL_FONT
            else:
                cell.font = ROW_FONT
        ws.row_dimensions[row_idx].height = 16

    for col_idx, width in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADER))}1"

    wb.save(OUTPUT_PATH)
    print(f"Saved: {OUTPUT_PATH}")
    print(f"Rows: {len(DATA)}")


if __name__ == "__main__":
    make_xlsx()
