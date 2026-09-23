import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def make_xlsx(filename, data):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "架電リスト"

    headers = ["企業名", "LP URL", "電話番号", "商材", "検索KW", "業界"]

    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="メイリオ", bold=True, color="FFFFFF", size=10)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center
        cell.border = border

    for row_idx, row in enumerate(data, 2):
        fill_color = "F2F7FC" if row_idx % 2 == 0 else "FFFFFF"
        row_fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
        for col_idx, val in enumerate(row, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.fill = row_fill
            cell.font = Font(name="メイリオ", size=9)
            cell.alignment = left
            cell.border = border

    col_widths = [30, 45, 16, 20, 30, 18]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.row_dimensions[1].height = 20
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    wb.save(filename)
    print(f"保存完了: {filename}  行数: {len(data)}")


# 第159弾（2026-07-16k）美容系サロン・クリニック拡充第3弾 7社・LP URLは予約ページ

data_0716k = [
    ["株式会社HIROGINZA", "https://www.hiroginza.com/ladies/bridal/shop/", "03-3538-3317", "ブライダルシェービング・レディースシェービング", "ブライダルシェービング, 花嫁 シェービング, 顔そり サロン", "ブライダルシェービングサロン"],
    ["一般社団法人表参道メディカルクリニックグループ", "https://medicalbrows.jp/booking", "", "医療アートメイク（眉・アイライン・リップ）", "眉 アートメイク, 医療アートメイク 東京, アートメイク 大阪", "医療アートメイククリニック"],
    ["カリスタ株式会社", "https://www.shinq-compass.jp/salon/reserve/28", "03-6721-7416", "女性専用美容鍼灸・リフトアップ美容鍼（CALISTA）", "美容鍼 恵比寿, 女性専用 鍼灸サロン, リフトアップ美容鍼", "美容鍼灸サロン"],
    ["株式会社癒しテック", "https://iyasheep.com/", "", "ドライヘッドスパ専門店", "ドライヘッドスパ 専門店, 頭皮ケア サロン, 頭のマッサージ 睡眠", "ドライヘッドスパ専門店"],
    ["株式会社横浜ヘルシー", "https://www.e-br.co.jp/trial/", "", "EMS痩身・ボディメイクエステ", "EMS 痩身 横浜, ボディメイク エステ 神奈川, 女性専用エステ 湘南", "EMS痩身・ボディメイクサロン"],
    ["株式会社クォーク", "https://www.bidan.co.jp/", "", "メンズ発毛・育毛メディカルサロン", "発毛 メンズ 大阪, 育毛サロン 男性, 薄毛 メディカルサロン", "メンズ発毛・育毛専門サロン"],
    ["株式会社T・S・S", "https://www.mellow-wax.com/reservation/", "022-263-1645", "ブラジリアンワックス脱毛・小さな総合美容サロン", "ブラジリアンワックス脱毛, ワックス脱毛 サロン, 眉毛ワックス", "ワックス脱毛専門サロン"],
]

base = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー"
make_xlsx(f"{base}\\架電リスト_Meta_2026-07-16k.xlsx", data_0716k)
