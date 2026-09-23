import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUTPUT_PATH = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー\架電リスト_リスティング_2026-05-27d.xlsx"
HEADER = ["企業名", "LP URL", "広告種別", "電話番号", "商材", "検索KW", "業界"]
COL_WIDTHS = [32, 50, 16, 16, 20, 32, 22]

DATA = [
    [
        "株式会社ライジングトラスト",
        "https://www.rising-trust.co.jp/lp/request/",
        "リスティング広告",
        "03-5338-3320",
        "不動産投資コンサルティング",
        "不動産投資 収益物件 資産運用 マンション投資",
        "不動産・賃貸・リノベーション・投資",
    ],
    [
        "鈴与商事株式会社（鈴与のソーラー）",
        "https://solar.suzuyoshoji.co.jp/lp/zero/index.html",
        "リスティング広告",
        "0120-224-215",
        "0円ソーラー（初期費用無料の太陽光発電）",
        "ソーラーパネル 0円ソーラー 太陽光発電 蓄電池",
        "不動産・賃貸・リノベーション・投資",
    ],
    [
        "株式会社ビットポイントジャパン（BITPOINT）",
        "https://www.bitpoint.co.jp/lp/cpn/",
        "リスティング広告",
        "0120-004-430",
        "暗号資産（仮想通貨）取引所サービス",
        "仮想通貨 暗号資産 ビットコイン BITPOINT 口座開設",
        "保険・金融・証券・カード",
    ],
    [
        "株式会社アイケンジャパン",
        "https://aikenjapan.jp/lp/",
        "リスティング広告",
        "092-739-6655",
        "新築アパート経営・不動産投資",
        "アパート経営 新築アパート 不動産投資 土地活用",
        "不動産・賃貸・リノベーション・投資",
    ],
    [
        "株式会社fundbook",
        "https://fundbook.co.jp/lp/",
        "リスティング広告",
        "050-1751-2206",
        "M&A仲介・事業承継サービス",
        "M&A 事業売却 事業承継 会社売却 M&A仲介",
        "法律・税務・相談",
    ],
    [
        "住宅情報館株式会社",
        "https://www.jutakujohokan.co.jp/lp/fullsupport/",
        "リスティング広告",
        "042-704-7333",
        "土地探しから注文住宅実現フルサポート",
        "注文住宅 土地探し マイホーム 新築一戸建て",
        "不動産・賃貸・リノベーション・投資",
    ],
    [
        "株式会社パレット（カラーズ・ラボ）",
        "https://crpalette.co.jp/lp/colorsLABO/",
        "リスティング広告",
        "050-5497-9488",
        "IT特化型就労移行支援（障害者向けITスキル習得）",
        "就労移行支援 障害者就労 ITスキル プログラミング",
        "教育・スキルアップ",
    ],
    [
        "株式会社インヴァランス（INVALANCE）",
        "https://www.invalance.co.jp/lp/seminar/basic_c/",
        "リスティング広告",
        "03-5302-7177",
        "不動産投資セミナー（区分マンション投資）",
        "不動産投資セミナー マンション投資 区分マンション 資産形成",
        "不動産・賃貸・リノベーション・投資",
    ],
    [
        "ステップゴルフ株式会社",
        "https://www.stepgolf.co.jp/lp/001/",
        "リスティング広告",
        "03-4577-6786",
        "インドアゴルフスクール（月会費制・24時間営業）",
        "ゴルフスクール インドアゴルフ 初心者ゴルフ ゴルフレッスン",
        "フィットネス・スポーツ・ゴルフ",
    ],
    [
        "セントラル短資ＦＸ株式会社",
        "https://www.central-tanshifx.com/ad/lp/800/",
        "リスティング広告",
        "0120-308-806",
        "FX（外国為替証拠金取引）口座開設サービス",
        "FX 外国為替 FX口座開設 スワップポイント FX取引",
        "保険・金融・証券・カード",
    ],
]


def make_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "架電リスト"

    header_font = Font(name="メイリオ", size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    fill_odd = PatternFill("solid", fgColor="F2F7FC")
    fill_even = PatternFill("solid", fgColor="FFFFFF")
    data_font = Font(name="メイリオ", size=9)
    data_align = Alignment(vertical="center", wrap_text=False)
    url_font = Font(name="メイリオ", size=9, color="0563C1", underline="single")

    thin = Side(style="thin", color="D0D7DE")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # ヘッダー行
    for col_idx, col_name in enumerate(HEADER, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = border

    ws.row_dimensions[1].height = 22

    # データ行
    for row_idx, row_data in enumerate(DATA, start=2):
        fill = fill_odd if (row_idx % 2 == 0) else fill_even
        for col_idx, value in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.fill = fill
            cell.border = border
            if col_idx == 2:  # LP URL列
                cell.font = url_font
                cell.hyperlink = value
                cell.alignment = data_align
            else:
                cell.font = data_font
                cell.alignment = data_align
        ws.row_dimensions[row_idx].height = 18

    # 列幅
    for col_idx, width in enumerate(COL_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # 先頭行固定・オートフィルター
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADER))}1"

    wb.save(OUTPUT_PATH)
    print(f"保存完了: {OUTPUT_PATH}")
    print(f"件数: {len(DATA)}社")


if __name__ == "__main__":
    make_xlsx()
