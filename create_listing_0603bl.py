import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUTPUT_PATH = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー\架電リスト_リスティング_2026-06-03bl.xlsx"
HEADER = ["企業名", "LP URL", "広告種別", "法人番号", "商材", "検索KW", "業界"]
COL_WIDTHS = [32, 56, 16, 16, 22, 32, 22]
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="メイリオ", size=10, bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
ROW_FILLS = [PatternFill("solid", fgColor="F2F7FC"), PatternFill("solid", fgColor="FFFFFF")]
ROW_FONT = Font(name="メイリオ", size=9)
ROW_ALIGN = Alignment(vertical="center", wrap_text=False)
URL_FONT = Font(name="メイリオ", size=9, color="0563C1", underline="single")

DATA = [
    # ── 人事SaaS ──
    ("BambooHR LLC",                "https://www.bamboohr.com/pricing/",                       "リスティング広告", "", "BambooHR・人事・採用・従業員管理",            "HR ソフト 中小企業, BambooHR 料金 評判, 従業員 管理 SaaS",          "HR SaaS"),
    # ── 採用管理ATS ──
    ("Employ Inc.",                 "https://www.lever.co/pricing/",                           "リスティング広告", "", "Lever・採用管理ATS・CRM・AI候補者選考",       "採用管理 ATS 比較, Lever 料金 評判, リクルーティング SaaS",          "HR SaaS"),
    # ── ECメールマーケ ──
    ("Klaviyo Inc.",               "https://www.klaviyo.com/pricing",                         "リスティング広告", "", "Klaviyo・AIメール&SMS・EC向けCRM",            "EC メールマーケ ツール, Klaviyo 料金 評判, SMS マーケ SaaS",        "メールマーケSaaS"),
    # ── 顧客エンゲージメント ──
    ("Braze Inc.",                 "https://www.braze.com/pricing",                           "リスティング広告", "", "Braze・モバイルマーケ・顧客エンゲージメント", "顧客 エンゲージメント SaaS, Braze 料金 評判, アプリ マーケ CRM",   "マーケSaaS"),
    # ── CDP ──
    ("mParticle Inc.",             "https://www.mparticle.com/pricing/",                      "リスティング広告", "", "mParticle・CDP・顧客データ連携・300+連携",    "CDP ツール 比較, mParticle 料金 評判, 顧客 データ 統合 SaaS",      "データSaaS"),
    # ── チャットAPI ──
    ("Sendbird Japan株式会社",     "https://sendbird.com/pricing/chat",                       "リスティング広告", "", "Sendbird・アプリ内チャットAPI・メッセージ",   "チャット API 比較, Sendbird 料金 評判, アプリ メッセージ SaaS",    "コミュニケーションAPI"),
    # ── 音声映像API ──
    ("Agora Inc.",                 "https://www.agora.io/en/pricing/",                        "リスティング広告", "", "Agora・リアルタイム音声/映像API・低遅延",     "音声 映像 API 比較, Agora 料金 評判, ビデオ通話 SDK SaaS",         "コミュニケーションAPI"),
    # ── サブスク請求管理 ──
    ("Chargebee Inc.",             "https://www.chargebee.com/pricing/",                      "リスティング広告", "", "Chargebee・サブスク請求管理・収益運用",        "サブスク 請求 管理 SaaS, Chargebee 料金 評判, 定期課金 ツール",    "フィンテックSaaS"),
    # ── ヘッドレスCMS ──
    ("Contentstack Inc.",          "https://www.contentstack.com/platforms/headless-cms",     "リスティング広告", "", "Contentstack・ヘッドレスCMS・APIファースト",  "ヘッドレス CMS 比較, Contentstack 料金 評判, コンテンツ 管理 API", "CMS SaaS"),
    # ── 検索API ──
    ("Algolia Inc.",               "https://www.algolia.com/lp/request-pricing",              "リスティング広告", "", "Algolia・AI検索/レコメンドAPI",                "サイト内 検索 API, Algolia 料金 評判, 検索 SaaS 比較",             "検索SaaS"),
    # ── 検索/分析エンジン ──
    ("Elasticsearch株式会社",      "https://www.elastic.co/jp/pricing/serverless-search",     "リスティング広告", "", "Elasticsearch・検索/分析・サーバーレス",      "Elasticsearch 料金 評判, 全文 検索 エンジン SaaS, ログ 分析",      "検索SaaS"),
    # ── プロダクト行動分析 ──
    ("Heap Inc.",                  "https://www.heap.io/pricing",                             "リスティング広告", "", "Heap・自動キャプチャ・プロダクト分析",        "プロダクト 分析 ツール, Heap 料金 評判, デジタル 体験 分析 SaaS", "プロダクト分析SaaS"),
    # ── 会話インテリジェンス ──
    ("Gong.io Inc.",               "https://www.gong.io/pricing",                             "リスティング広告", "", "Gong・営業会話分析・Revenue AI",              "営業 会話 分析 SaaS, Gong 料金 評判, セールス AI ツール",          "セールスSaaS"),
    # ── 作図/ホワイトボード ──
    ("Lucid Software Inc.",        "https://lucid.co/pricing",                                "リスティング広告", "", "Lucidchart・作図・ホワイトボード・図解",      "作図 ツール 比較, Lucidchart 料金 評判, フローチャート SaaS",      "コラボSaaS"),
    # ── 英文校正AI ──
    ("Grammarly Inc.",             "https://www.grammarly.com/plans",                         "リスティング広告", "", "Grammarly・AI英文校正・ライティング支援",     "英文 校正 ツール, Grammarly 料金 評判, ビジネス 英語 AI",          "AIライティングSaaS"),
    # ── 日程調整 ──
    ("Calendly LLC",               "https://calendly.com/pricing",                            "リスティング広告", "", "Calendly・日程調整自動化・無料プラン",        "日程調整 ツール 比較, Calendly 料金 評判, 予約 自動化 SaaS",       "業務DXSaaS"),
    # ── メール/SMSマーケ ──
    ("Brevo SAS",                  "https://www.brevo.com/pricing/",                          "リスティング広告", "", "Brevo・メール/SMS/CRM・無料プランあり",       "メール マーケ ツール 無料, Brevo 料金 評判, SMS 配信 SaaS",        "メールマーケSaaS"),
    # ── ヒートマップ分析 ──
    ("Hotjar Ltd.",                "https://www.hotjar.com/ja/pricing/",                      "リスティング広告", "", "Hotjar・ヒートマップ・行動分析・録画",        "ヒートマップ ツール 比較, Hotjar 料金 評判, サイト 行動 分析",     "プロダクト分析SaaS"),
    # ── 会話マーケ ──
    ("Drift.com Inc.",             "https://www.drift.com/pricing/",                          "リスティング広告", "", "Drift・カンバセーションAI・チャットマーケ",   "チャット マーケ SaaS, Drift 料金 評判, 会話 型 マーケ ツール",     "マーケSaaS"),
    # ── デジタル体験分析 ──
    ("Fullstory Inc.",             "https://www.fullstory.com/plans/",                        "リスティング広告", "", "Fullstory・デジタル体験分析・セッション録画", "デジタル 体験 分析 SaaS, Fullstory 料金 評判, セッション リプレイ", "プロダクト分析SaaS"),
]


def make_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "架電リスト"

    for col_idx, header in enumerate(HEADER, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
    ws.row_dimensions[1].height = 20

    for row_idx, row_data in enumerate(DATA, 2):
        fill = ROW_FILLS[(row_idx - 2) % 2]
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.fill = fill
            cell.alignment = ROW_ALIGN
            if col_idx == 2:
                cell.hyperlink = value
                cell.font = URL_FONT
            else:
                cell.font = ROW_FONT
        ws.row_dimensions[row_idx].height = 16

    for col_idx, width in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADER))}1"

    wb.save(OUTPUT_PATH)
    print(f"Saved: {OUTPUT_PATH}")
    print(f"Rows: {len(DATA)}")


if __name__ == "__main__":
    make_xlsx()
