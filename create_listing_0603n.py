import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUTPUT_PATH = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー\架電リスト_リスティング_2026-06-03n.xlsx"
HEADER = ["企業名", "LP URL", "広告種別", "法人番号", "商材", "検索KW", "業界"]
COL_WIDTHS = [32, 50, 16, 16, 20, 32, 22]
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="メイリオ", size=10, bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
ROW_FILLS = [PatternFill("solid", fgColor="F2F7FC"), PatternFill("solid", fgColor="FFFFFF")]
ROW_FONT = Font(name="メイリオ", size=9)
ROW_ALIGN = Alignment(vertical="center", wrap_text=False)
URL_FONT = Font(name="メイリオ", size=9, color="0563C1", underline="single")

DATA = [
    # ── プレミアムオーディオ（高単価） ──────────────────────────────────────────
    ("ボーズ・ジャパン合同会社（BOSE）",                   "https://www.bose.co.jp/",                       "リスティング広告", "", "プレミアムスピーカー・ヘッドフォン・サウンドバー","BOSE ヘッドフォン おすすめ 評判, スピーカー 高音質 ブランド, ボーズ ノイキャン 比較","プレミアムオーディオ"),
    ("株式会社オーディオテクニカ",                         "https://www.audio-technica.co.jp/",             "リスティング広告", "", "ヘッドフォン・マイク・レコードプレーヤー",        "オーディオテクニカ ヘッドフォン 評判, マイク 高音質 配信 おすすめ, ワイヤレス イヤホン 比較","プレミアムオーディオ"),
    ("ゼンハイザージャパン株式会社（Sennheiser）",         "https://www.sennheiser.com/ja-JP/",             "リスティング広告", "", "プレミアムヘッドフォン・イヤフォン・マイク",      "ゼンハイザー ヘッドフォン おすすめ, 高級 イヤホン 音質 比較, Sennheiser 評判 購入","プレミアムオーディオ"),
    # ── PC（高単価） ──────────────────────────────────────────
    ("株式会社VAIO",                                       "https://vaio.com/",                             "リスティング広告", "", "プレミアムノートPC・モバイルPC・法人PC",           "VAIO ノートPC 評判 おすすめ, 軽量 ノートPC 高性能 比較, ビジネス PC 高品質 購入","PC"),
    ("デル・テクノロジーズ株式会社",                       "https://www.dell.com/ja-jp/",                   "リスティング広告", "", "ノートPC・デスクトップ・ゲーミングPC",            "デル PC 評判 おすすめ, ノートPC 高性能 コスパ, XPS 購入 価格 比較",          "PC"),
    ("ASUS JAPAN株式会社（ROG・Zenbook）",                 "https://www.asus.com/jp/",                      "リスティング広告", "", "ゲーミングPC・ノートPC・ディスプレイ",            "ROG ゲーミング PC おすすめ, Zenbook 軽量 ノートPC 評判, ASUS PC 購入 価格",  "PC"),
    # ── PC周辺機器（高単価） ──────────────────────────────────────────
    ("株式会社ロジクール（Logicool）",                     "https://www.logicool.co.jp/",                   "リスティング広告", "", "マウス・キーボード・Webカメラ・ゲーミング機器",   "ロジクール マウス おすすめ ワイヤレス, ゲーミング キーボード 比較 高機能, Logicool ヘッドセット","PC周辺機器"),
    # ── 家電（高単価） ──────────────────────────────────────────
    ("ダイソン株式会社（Dyson）",                          "https://www.dyson.co.jp/",                      "リスティング広告", "", "コードレス掃除機・ドライヤー・空気清浄機",        "ダイソン 掃除機 おすすめ コードレス, ダイソン ドライヤー 評判 効果, 空気清浄 Dyson 価格","家電"),
    ("セイコーエプソン株式会社",                           "https://www.epson.jp/",                         "リスティング広告", "", "プリンター・プロジェクター・インクジェット",      "エプソン プリンター おすすめ 家庭用, インクジェット 写真 印刷 比較, プロジェクター 家庭用 購入","家電"),
    # ── 調理器具（高単価） ──────────────────────────────────────────
    ("WMFジャパン合同会社",                                "https://www.wmf.com/ja-jp/",                    "リスティング広告", "", "高級鍋・圧力鍋・カトラリー・キッチン用品",        "WMF 鍋 おすすめ 評判, 圧力鍋 高品質 ブランド 比較, 調理器具 高級 ギフト",   "調理器具"),
    ("ビタクラフト・ジャパン株式会社",                     "https://vitacraft.co.jp/",                      "リスティング広告", "", "無水鍋・多重構造鍋・ヘルシー調理器具",            "ビタクラフト 鍋 評判 おすすめ, 無水 調理 鍋 栄養 損なわない, 高級 鍋 種類 比較","調理器具"),
    ("フィスラー・ジャパン株式会社",                       "https://www.fissler.com/jp/",                   "リスティング広告", "", "圧力鍋・フライパン・高端調理器具",                "フィスラー 圧力鍋 評判 使い方, 高級 フライパン ブランド 比較, 圧力鍋 時短 料理",  "調理器具"),
    # ── バッグ・旅行用品（高単価） ──────────────────────────────────────────
    ("TUMIジャパン株式会社（TUMI）",                       "https://www.tumi.com/ja/",                      "リスティング広告", "", "ビジネスバッグ・スーツケース・トラベル用品",      "TUMI バッグ おすすめ 評判, ビジネス バッグ 高品質 ブランド, スーツケース TUMI 購入","バッグ"),
    ("サムソナイト・ジャパン株式会社",                     "https://www.samsonite.co.jp/",                  "リスティング広告", "", "スーツケース・旅行かばん・バックパック",           "サムソナイト スーツケース おすすめ, 旅行 かばん ブランド 軽量, スーツケース 耐久性 比較","旅行用品"),
    # ── 自転車（高単価） ──────────────────────────────────────────
    ("株式会社ジャイアント・ジャパン",                     "https://www.giant.co.jp/",                      "リスティング広告", "", "ロードバイク・クロスバイク・電動自転車",           "ロードバイク おすすめ 入門, クロスバイク 選び方 価格, ジャイアント 自転車 評判","自転車"),
    # ── 文具（高単価） ──────────────────────────────────────────
    ("セーラー万年筆株式会社",                             "https://www.sailor.co.jp/",                     "リスティング広告", "", "万年筆・高級ボールペン・インク・文具",             "万年筆 おすすめ 初心者 選び方, セーラー 万年筆 評判, 高級 ペン ギフト プレゼント","文具"),
    ("株式会社パイロットコーポレーション",                 "https://www.pilot.co.jp/",                      "リスティング広告", "", "万年筆・高級ボールペン・ペン・文具",               "パイロット 万年筆 評判 おすすめ, 高級 ボールペン ブランド 比較, 文具 ギフト 選び方","文具"),
    # ── 高級家具（高単価） ──────────────────────────────────────────
    ("ハーマンミラージャパン株式会社（Herman Miller）",    "https://www.hermanmiller.com/ja_jp/",           "リスティング広告", "", "高級オフィスチェア・エルゴノミクス家具",           "ハーマンミラー チェア 評判 おすすめ, アーロンチェア 購入 価格, 高級 オフィスチェア 腰痛","高級家具"),
    # ── カメラ（高単価） ──────────────────────────────────────────
    ("ライカカメラジャパン株式会社（Leica）",              "https://www.leica-camera.com/ja-JP/",           "リスティング広告", "", "プレミアムカメラ・Mシステム・レンズ",              "ライカ カメラ 評判 おすすめ, Leica M 購入 価格, プレミアム カメラ ブランド 選び方","カメラ"),
    # ── 高級文具・ラグジュアリー（高単価） ──────────────────────────────────────────
    ("モンブラン・インターナショナル日本支店（Montblanc）", "https://www.montblanc.com/ja-jp/",             "リスティング広告", "", "高級万年筆・ボールペン・レザーグッズ・腕時計",    "モンブラン 万年筆 評判 価格, 高級 ペン ブランド ギフト, Montblanc ボールペン 購入","ラグジュアリー文具"),
]


def make_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "架電リスト"

    for col_idx, header in enumerate(HEADER, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
    ws.row_dimensions[1].height = 20

    for row_idx, row_data in enumerate(DATA, 2):
        fill = ROW_FILLS[(row_idx - 2) % 2]
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.fill = fill
            cell.alignment = ROW_ALIGN
            if col_idx == 2:
                cell.hyperlink = value
                cell.font = URL_FONT
            else:
                cell.font = ROW_FONT
        ws.row_dimensions[row_idx].height = 16

    for col_idx, width in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADER))}1"

    wb.save(OUTPUT_PATH)
    print(f"Saved: {OUTPUT_PATH}")
    print(f"Rows: {len(DATA)}")


if __name__ == "__main__":
    make_xlsx()
