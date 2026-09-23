import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def make_xlsx(filename, data):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "架電リスト"

    headers = ["企業名", "LP URL", "電話番号", "商材", "検索KW", "業界"]

    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="メイリオ", bold=True, color="FFFFFF", size=10)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center
        cell.border = border

    for row_idx, row in enumerate(data, 2):
        fill_color = "F2F7FC" if row_idx % 2 == 0 else "FFFFFF"
        row_fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
        for col_idx, val in enumerate(row, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.fill = row_fill
            cell.font = Font(name="メイリオ", size=9)
            cell.alignment = left
            cell.border = border

    col_widths = [30, 45, 16, 20, 30, 18]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.row_dimensions[1].height = 20
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    wb.save(filename)
    print(f"保存完了: {filename}  行数: {len(data)}")


# 第157弾（2026-07-16i）買取・リユース業界拡充第2弾 18社・LP URLは買取査定ページ

data_0716i = [
    ["株式会社Real Standard", "https://www.kaitorisenmon.com/toriatukai/perfume/", "0743-74-0154", "香水・ブランド香水買取（コスラボ）", "香水買取, ブランド香水 売る, コスラボ", "香水買取"],
    ["株式会社STAYGOLD", "https://bring-kicks-kaitori.com/", "03-6427-1153", "スニーカー買取（NIKE・adidas等）", "スニーカー買取, NIKE買取, レアスニーカー 売る", "スニーカー買取"],
    ["株式会社山徳", "https://kaitori-beerecords.jp/form", "", "レコード・アナログ盤買取（ビーレコーズ）", "レコード買取, アナログレコード 売る, 宅配買取", "レコード買取"],
    ["株式会社美術刀剣みらい", "https://katana-kaitori.com/contact", "", "日本刀・刀剣買取", "日本刀買取, 刀剣買取, 刀 売る", "刀剣買取"],
    ["株式会社トリアイナ", "https://hareya.jp/sewing/", "03-6912-7207", "ミシン買取（家庭用・工業用）（はれや）", "ミシン買取, 職業用ミシン 売る, 古いミシン処分", "ミシン買取"],
    ["株式会社リサイクルマイスター", "https://audiosound.co.jp/form/", "070-3893-2731", "オーディオ機器買取（アンプ・スピーカー）", "オーディオ買取, アンプ買取, スピーカー 売る", "オーディオ機器買取"],
    ["株式会社トレードランド", "https://tradeshop-drone.com/form/", "", "ドローン買取（DJI等）", "ドローン買取, DJI買取, Mavic 売る", "ドローン買取"],
    ["株式会社フェアーメディカル", "https://www.fair-medical.jp/inquiry", "", "中古医療機器・介護用品買取", "医療機器買取, 介護ベッド買取, 車椅子 売る", "医療機器・介護用品買取"],
    ["株式会社アイエイト", "https://topkenki-juuki.com/contact/", "", "建設機械・重機買取", "重機買取, 油圧ショベル買取, ユンボ 売る", "建設機械買取"],
    ["株式会社カイショー", "https://taiya-kaitori-no1.jp/free-assessment/", "092-600-1086", "タイヤ・ホイール買取", "タイヤ買取, ホイール買取, 中古タイヤ 処分", "タイヤ・ホイール買取"],
    ["株式会社ラフジュ工房", "https://www.rafuju.jp/user_data/user_guide/buy_out/form/", "0294-70-3730", "北欧家具・デザイナーズ家具買取", "北欧家具買取, アンティーク家具 売る, デザイナーズ家具買取", "北欧家具買取"],
    ["株式会社スモビー", "https://ramipass.com/", "050-1721-8202", "制服・学生服買取（ラミパス）", "制服買取, 学生服 売る, 宅配買取", "制服買取"],
    ["株式会社ノースフィールド", "https://nbcpens.com/", "042-689-6992", "万年筆・筆記具買取", "万年筆買取, モンブラン 売る, 筆記具買取", "万年筆買取"],
    ["株式会社ササキコーポレーション", "https://www.sasaki-marine.net/buy", "0836-21-8181", "ボート・船舶買取", "ボート買取, クルーザー 売る, 船舶買取", "船舶買取"],
    ["株式会社山田書店", "https://www.yamada-shoten.com/kaitori.php", "03-3295-0252", "浮世絵・版画買取", "浮世絵買取, 版画 売る, 古美術買取", "浮世絵買取"],
    ["株式会社我楽洞", "https://www.garakudo.co.jp/chadougu/", "045-590-5660", "茶道具買取", "茶道具買取, 茶碗 売る, 骨董品買取", "茶道具買取"],
    ["株式会社コードファクトリー", "https://www.koodofactory.com/kaitoriform/", "", "スーツ・礼服買取", "スーツ買取, 礼服 売る, ブランドスーツ買取", "スーツ買取"],
    ["株式会社まるげん", "https://remove-formen.com/item_category/sunglasses/", "", "サングラス・メガネ買取", "サングラス買取, メガネ買取, ブランドメガネ 売る", "サングラス買取"],
]

base = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー"
make_xlsx(f"{base}\\架電リスト_Meta_2026-07-16i.xlsx", data_0716i)
