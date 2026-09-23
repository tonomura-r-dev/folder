import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUTPUT_PATH = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー\架電リスト_リスティング_2026-06-03ag.xlsx"
HEADER = ["企業名", "LP URL", "広告種別", "法人番号", "商材", "検索KW", "業界"]
COL_WIDTHS = [32, 56, 16, 16, 22, 32, 22]
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="メイリオ", size=10, bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
ROW_FILLS = [PatternFill("solid", fgColor="F2F7FC"), PatternFill("solid", fgColor="FFFFFF")]
ROW_FONT = Font(name="メイリオ", size=9)
ROW_ALIGN = Alignment(vertical="center", wrap_text=False)
URL_FONT = Font(name="メイリオ", size=9, color="0563C1", underline="single")

# ── 全社、実際の広告クリック先LP（定期/トライアル/購入CVページ）をWeb確認済み ──
DATA = [
    # ── 着圧レギンス（定期購入LP） ──
    ("YB-LAB.株式会社（グラマラスパッツ）",      "https://ybl-store.net/lp?u=pnsonebs",                 "リスティング広告", "", "着圧レギンス・定期購入",                "グラマラスパッツ 評判 効果, 着圧 レギンス おすすめ, 脚 痩せ 定期",          "着圧インナー"),
    # ── 女性育毛剤（定期コースLP） ──
    ("株式会社レッドビジョン（マイナチュレ）",    "https://www.my-nature.jp/shop/pages/subscription",    "リスティング広告", "", "女性向け無添加育毛剤・定期コース",    "マイナチュレ 女性育毛剤 評判, 薄毛 抜け毛 女性 育毛, 無添加 育毛 定期", "育毛"),
    # ── スキンケアジェル（定期お試しコースLP） ──
    ("メビウス製薬株式会社（シミウス）",          "https://lp.mebiusseiyaku.co.jp/w01fb026_01",          "リスティング広告", "", "薬用ホワイトニングリフトケアジェル", "シミウス 定期 評判 効果, オールインワンジェル おすすめ, シミ ケア 美白", "スキンケア"),
]


def make_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "架電リスト"

    for col_idx, header in enumerate(HEADER, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
    ws.row_dimensions[1].height = 20

    for row_idx, row_data in enumerate(DATA, 2):
        fill = ROW_FILLS[(row_idx - 2) % 2]
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.fill = fill
            cell.alignment = ROW_ALIGN
            if col_idx == 2:
                cell.hyperlink = value
                cell.font = URL_FONT
            else:
                cell.font = ROW_FONT
        ws.row_dimensions[row_idx].height = 16

    for col_idx, width in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADER))}1"

    wb.save(OUTPUT_PATH)
    print(f"Saved: {OUTPUT_PATH}")
    print(f"Rows: {len(DATA)}")


if __name__ == "__main__":
    make_xlsx()
