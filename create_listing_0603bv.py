import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUTPUT_PATH = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー\架電リスト_リスティング_2026-06-03bv.xlsx"
HEADER = ["企業名", "LP URL", "広告種別", "法人番号", "商材", "検索KW", "業界"]
COL_WIDTHS = [32, 56, 16, 16, 22, 32, 22]
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="メイリオ", size=10, bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
ROW_FILLS = [PatternFill("solid", fgColor="F2F7FC"), PatternFill("solid", fgColor="FFFFFF")]
ROW_FONT = Font(name="メイリオ", size=9)
ROW_ALIGN = Alignment(vertical="center", wrap_text=False)
URL_FONT = Font(name="メイリオ", size=9, color="0563C1", underline="single")

DATA = [
    # ── 企業向けLMS ──
    ("TalentLMS",                  "https://www.talentlms.com/pricing",                       "リスティング広告", "", "TalentLMS・企業研修LMS・eラーニング",          "LMS 比較 企業研修, TalentLMS 料金 評判, eラーニング システム",     "LMS SaaS"),
    # ── 協働型LMS ──
    ("360Learning",                "https://360learning.com/pricing/",                        "リスティング広告", "", "360Learning・協働型LMS・社内ナレッジ研修",     "LMS 協働 研修 SaaS, 360Learning 料金 評判, 社内 教育 ツール",      "LMS SaaS"),
    # ── 企業向けLMS ──
    ("Absorb LMS",                 "https://www.absorblms.com/pricing",                       "リスティング広告", "", "Absorb LMS・企業研修/教育・AI対応",            "LMS 企業 研修 比較, Absorb LMS 料金 評判, 学習 管理 システム",    "LMS SaaS"),
    # ── オンライン講座作成 ──
    ("Thinkific",                  "https://www.thinkific.com/pricing/",                      "リスティング広告", "", "Thinkific・オンライン講座作成/販売",           "オンライン 講座 作成 SaaS, Thinkific 料金 評判, eラーニング 販売", "eラーニングSaaS"),
    # ── オンライン講座作成 ──
    ("Teachable",                  "https://teachable.com/pricing",                           "リスティング広告", "", "Teachable・オンライン講座作成/販売",           "オンライン コース 販売 SaaS, Teachable 料金 評判, 講座 作成 ツール","eラーニングSaaS"),
    # ── クリエイター講座 ──
    ("Kajabi",                     "https://kajabi.com/pricing",                              "リスティング広告", "", "Kajabi・オンライン講座/会員サイト/マーケ",     "オンライン 講座 プラットフォーム, Kajabi 料金 評判, 会員 サイト 作成","eラーニングSaaS"),
    # ── オンライン講座作成 ──
    ("LearnWorlds",                "https://www.learnworlds.com/pricing/",                    "リスティング広告", "", "LearnWorlds・オンライン講座作成/販売",         "オンライン スクール 作成, LearnWorlds 料金 評判, 講座 販売 SaaS",  "eラーニングSaaS"),
    # ── クリエイター販売 ──
    ("Podia",                      "https://www.podia.com/pricing",                           "リスティング広告", "", "Podia・オンライン講座/デジタル商品販売",       "デジタル 商品 販売 SaaS, Podia 料金 評判, オンライン 講座 販売",   "eラーニングSaaS"),
    # ── 契約ライフサイクル管理 ──
    ("Juro",                       "https://juro.com/pricing",                                "リスティング広告", "", "Juro・契約管理CLM・ブラウザ完結",             "契約 管理 CLM SaaS, Juro 料金 評判, 契約 ライフサイクル 管理",     "契約管理SaaS"),
    # ── 契約管理 ──
    ("Concord",                    "https://www.concord.app/pricing",                         "リスティング広告", "", "Concord・契約管理/電子署名・無制限署名",       "契約 管理 SaaS, Concord 料金 評判, 電子署名 契約 ツール",          "契約管理SaaS"),
    # ── 契約ライフサイクル管理（日本） ──
    ("株式会社ContractS",          "https://www.contracts.co.jp/lp/lp-management/",           "リスティング広告", "", "ContractS CLM・契約ライフサイクル管理",        "契約 管理 システム, ContractS CLM 料金 評判, 契約 ライフサイクル", "契約管理SaaS"),
    # ── 契約管理 ──
    ("SpotDraft",                  "https://www.spotdraft.com/pricing",                       "リスティング広告", "", "SpotDraft・契約管理CLM・AI契約レビュー",       "契約 管理 CLM, SpotDraft 料金 評判, AI 契約 レビュー SaaS",        "契約管理SaaS"),
    # ── 契約管理 ──
    ("LinkSquares",               "https://www.linksquares.com/pricing",                     "リスティング広告", "", "LinkSquares・契約管理/AI分析・法務向け",       "契約 管理 法務 SaaS, LinkSquares 料金 評判, AI 契約 分析",         "契約管理SaaS"),
    # ── 電子署名 ──
    ("SignWell",                   "https://www.signwell.com/pricing/",                       "リスティング広告", "", "SignWell・電子署名・低価格/簡単",              "電子署名 SaaS 安い, SignWell 料金 評判, 電子契約 ツール",          "電子署名SaaS"),
    # ── コミュニティ ──
    ("Circle.so",                  "https://circle.so/pricing",                               "リスティング広告", "", "Circle・オンラインコミュニティ/会員サイト",    "オンライン コミュニティ SaaS, Circle.so 料金 評判, 会員 サイト 作成","コミュニティSaaS"),
    # ── コミュニティ ──
    ("Skool",                      "https://www.skool.com/pricing",                           "リスティング広告", "", "Skool・コミュニティ/オンライン講座統合",       "コミュニティ 講座 SaaS, Skool 料金 評判, オンライン サロン 作成",  "コミュニティSaaS"),
    # ── コミュニティ ──
    ("Mighty Networks",            "https://www.mightynetworks.com/pricing",                  "リスティング広告", "", "Mighty Networks・コミュニティ/講座/会員",      "コミュニティ プラットフォーム, Mighty Networks 料金 評判, 会員 サイト","コミュニティSaaS"),
    # ── 企業向けLMS ──
    ("Docebo",                     "https://www.docebo.com/learning-platform-pricing/",       "リスティング広告", "", "Docebo・AI搭載企業向けLMS・大規模研修",        "LMS 大企業 AI, Docebo 料金 評判, 学習 管理 システム 比較",        "LMS SaaS"),
    # ── 契約管理 ──
    ("Ironclad",                   "https://ironcladapp.com/pricing/",                        "リスティング広告", "", "Ironclad・契約管理CLM・ワークフロー自動化",    "契約 管理 CLM 大企業, Ironclad 料金 評判, 法務 ワークフロー SaaS", "契約管理SaaS"),
    # ── ラーニングコミュニティ ──
    ("Disco",                      "https://www.disco.co/pricing",                            "リスティング広告", "", "Disco・AIラーニングコミュニティ/講座運営",     "ラーニング コミュニティ SaaS, Disco 料金 評判, オンライン 講座 運営","eラーニングSaaS"),
]


def make_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "架電リスト"

    for col_idx, header in enumerate(HEADER, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
    ws.row_dimensions[1].height = 20

    for row_idx, row_data in enumerate(DATA, 2):
        fill = ROW_FILLS[(row_idx - 2) % 2]
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.fill = fill
            cell.alignment = ROW_ALIGN
            if col_idx == 2:
                cell.hyperlink = value
                cell.font = URL_FONT
            else:
                cell.font = ROW_FONT
        ws.row_dimensions[row_idx].height = 16

    for col_idx, width in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADER))}1"

    wb.save(OUTPUT_PATH)
    print(f"Saved: {OUTPUT_PATH}")
    print(f"Rows: {len(DATA)}")


if __name__ == "__main__":
    make_xlsx()
