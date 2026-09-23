import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUTPUT_PATH = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー\架電リスト_リスティング_2026-06-03bo.xlsx"
HEADER = ["企業名", "LP URL", "広告種別", "法人番号", "商材", "検索KW", "業界"]
COL_WIDTHS = [32, 56, 16, 16, 22, 32, 22]
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="メイリオ", size=10, bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
ROW_FILLS = [PatternFill("solid", fgColor="F2F7FC"), PatternFill("solid", fgColor="FFFFFF")]
ROW_FONT = Font(name="メイリオ", size=9)
ROW_ALIGN = Alignment(vertical="center", wrap_text=False)
URL_FONT = Font(name="メイリオ", size=9, color="0563C1", underline="single")

DATA = [
    # ── マーケティングオートメーション ──
    ("Act-On Software Inc.",       "https://act-on.com/marketo-vs-act-on/",                   "リスティング広告", "", "Act-On・MA・顧客エンゲージメント",            "マーケ オートメーション 比較, Act-On 料金 評判, MA ツール 中堅",   "マーケSaaS"),
    # ── CRM ──
    ("SugarCRM Inc.",              "https://www.sugarcrm.com/pricing/",                       "リスティング広告", "", "SugarCRM・B2B営業CRM・自動フォロー",          "B2B CRM SaaS, SugarCRM 料金 評判, 営業 支援 ツール",              "CRM SaaS"),
    # ── SNS管理 ──
    ("Sprout Social Inc.",         "https://sproutsocial.com/pricing/",                       "リスティング広告", "", "Sprout Social・SNS管理・分析・エンゲージ",    "SNS 管理 ツール 法人, Sprout Social 料金 評判, ソーシャル 運用",   "SNSマーケSaaS"),
    # ── SNS管理 ──
    ("Hootsuite Inc.",             "https://www.hootsuite.com/plans",                         "リスティング広告", "", "Hootsuite・SNS統合管理・予約投稿・分析",       "SNS 管理 ツール 比較, Hootsuite 料金 評判, ソーシャル メディア 運用","SNSマーケSaaS"),
    # ── SNS投稿予約 ──
    ("Buffer (Buffer Inc.)",       "https://buffer.com/pricing",                              "リスティング広告", "", "Buffer・SNS投稿予約・スケジュール・無料枠",   "SNS 予約 投稿 ツール, Buffer 料金 評判, ソーシャル スケジュール",  "SNSマーケSaaS"),
    # ── SNSスケジュール ──
    ("Later (Later Social)",       "https://later.com/pricing",                               "リスティング広告", "", "Later・Instagram予約・インフルエンサーマーケ","Instagram 予約 ツール, Later 料金 評判, SNS スケジュール SaaS",    "SNSマーケSaaS"),
    # ── SEOツール ──
    ("Semrush Inc.",               "https://ja.semrush.com/pricing/",                         "リスティング広告", "", "Semrush・SEO/競合分析・キーワード調査",        "SEO ツール 比較, Semrush 料金 評判, 競合 分析 SaaS",              "SEO SaaS"),
    # ── メール配信API ──
    ("Mailgun (Sinch)",            "https://www.mailgun.com/pricing/",                        "リスティング広告", "", "Mailgun・トランザクションメールAPI・配信",     "メール 配信 API, Mailgun 料金 評判, トランザクション メール SaaS", "コミュニケーションAPI"),
    # ── メール配信API ──
    ("Postmark (ActiveCampaign)",  "https://postmarkapp.com/pricing",                         "リスティング広告", "", "Postmark・高速トランザクションメール配信",     "トランザクション メール SaaS, Postmark 料金 評判, メール API 高速","コミュニケーションAPI"),
    # ── 中小企業向けCRM ──
    ("Keap (Infusionsoft)",        "https://keap.com/pricing",                                "リスティング広告", "", "Keap・中小企業CRM・マーケ自動化・EC",          "中小 企業 CRM SaaS, Keap 料金 評判, 営業 マーケ 自動化",          "CRM SaaS"),
    # ── オンボーディング ──
    ("Appcues Inc.",              "https://www.appcues.com/pricing",                         "リスティング広告", "", "Appcues・プロダクトオンボーディング・UX改善", "オンボーディング SaaS, Appcues 料金 評判, プロダクト 体験 改善",   "プロダクトSaaS"),
    # ── A/Bテスト ──
    ("Wingify (VWO)",             "https://vwo.com/pricing/",                                "リスティング広告", "", "VWO・A/Bテスト・CRO・ヒートマップ",           "A/B テスト ツール, VWO 料金 評判, CRO 最適化 SaaS",               "プロダクト分析SaaS"),
    # ── 共有受信トレイ ──
    ("Front (FrontApp Inc.)",      "https://front.com/pricing",                               "リスティング広告", "", "Front・共有受信トレイ・チームカスタマー対応", "共有 受信 トレイ SaaS, Front 料金 評判, チーム メール 管理",       "カスタマーサポートSaaS"),
    # ── プロダクトマネジメント ──
    ("Productboard Inc.",          "https://www.productboard.com/pricing/",                   "リスティング広告", "", "Productboard・プロダクトマネジメント・AI",     "プロダクト マネジメント SaaS, Productboard 料金 評判, ロードマップ","プロダクトSaaS"),
    # ── 開発プロジェクト管理 ──
    ("Shortcut (Shortcut Inc.)",   "https://www.shortcut.com/pricing",                        "リスティング広告", "", "Shortcut・エンジニア向けプロジェクト管理",     "開発 プロジェクト 管理 SaaS, Shortcut 料金 評判, Jira 代替",       "DevOps SaaS"),
    # ── ナレッジ管理 ──
    ("Guru (GetGuru)",            "https://www.getguru.com/pricing",                         "リスティング広告", "", "Guru・AIナレッジ管理・エンタープライズ検索",  "AI ナレッジ 管理 SaaS, Guru 料金 評判, 社内 検索 wiki",           "コラボSaaS"),
    # ── 電子署名 ──
    ("Signeasy (Glykka LLC)",      "https://signeasy.com/pricing",                            "リスティング広告", "", "Signeasy・電子署名・契約管理・低価格",         "電子署名 SaaS 比較, Signeasy 料金 評判, 契約 管理 ツール",         "電子署名SaaS"),
    # ── iPaaS ──
    ("Workato Inc.",               "https://www.workato.com/pricing",                         "リスティング広告", "", "Workato・iPaaS・エンタープライズ業務連携",     "iPaaS ツール 比較, Workato 料金 評判, 業務 連携 自動化 SaaS",      "業務自動化SaaS"),
    # ── iPaaS ──
    ("Tray.ai (Tray.io)",         "https://tray.ai/pricing/",                                "リスティング広告", "", "Tray.ai・AI統合自動化・700+連携",             "iPaaS 自動化 SaaS, Tray.io 料金 評判, AI 連携 プラットフォーム",   "業務自動化SaaS"),
    # ── パスワード管理 ──
    ("Bitwarden Inc.",            "https://bitwarden.com/pricing/business/",                 "リスティング広告", "", "Bitwarden・法人パスワード管理・OSS",          "パスワード 管理 法人 OSS, Bitwarden 料金 評判, パスワードマネージャー","セキュリティSaaS"),
]


def make_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "架電リスト"

    for col_idx, header in enumerate(HEADER, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
    ws.row_dimensions[1].height = 20

    for row_idx, row_data in enumerate(DATA, 2):
        fill = ROW_FILLS[(row_idx - 2) % 2]
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.fill = fill
            cell.alignment = ROW_ALIGN
            if col_idx == 2:
                cell.hyperlink = value
                cell.font = URL_FONT
            else:
                cell.font = ROW_FONT
        ws.row_dimensions[row_idx].height = 16

    for col_idx, width in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADER))}1"

    wb.save(OUTPUT_PATH)
    print(f"Saved: {OUTPUT_PATH}")
    print(f"Rows: {len(DATA)}")


if __name__ == "__main__":
    make_xlsx()
