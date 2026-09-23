import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUTPUT_PATH = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー\架電リスト_リスティング_2026-05-29f.xlsx"

HEADER = ["企業名", "LP URL", "広告種別", "電話番号", "商材", "検索KW", "業界"]
COL_WIDTHS = [32, 50, 16, 16, 20, 32, 22]

DATA = [
    [
        "株式会社アガルート",
        "https://www.agaroot.jp/lp/cm202007/",
        "リスティング広告",
        "03-6279-0186",
        "司法書士・宅建・行政書士資格予備校",
        "司法書士 資格 オンライン 予備校 行政書士 宅建",
        "教育・スキルアップ",
    ],
    [
        "株式会社フォーサイト",
        "https://www.foresight.jp/lp/nenkin/request/",
        "リスティング広告",
        "03-5205-2222",
        "宅建・社労士・FP資格取得の通信講座",
        "宅建 通信講座 資格 社労士 FP 合格",
        "教育・スキルアップ",
    ],
    [
        "KIYOラーニング株式会社",
        "https://studying.jp/top/lp/202207/index.html",
        "リスティング広告",
        "03-6259-0742",
        "オンライン資格取得講座（スタディング）",
        "資格 オンライン 宅建 行政書士 スタディング",
        "教育・スキルアップ",
    ],
    [
        "株式会社テイクアンドギヴ・ニーズ",
        "https://www.tgn.co.jp/lp/cp/tg-wedding/",
        "リスティング広告",
        "03-6833-1122",
        "オリジナルウェディングプロデュース",
        "結婚式 ウェディング プロデュース 式場 挙式",
        "ウエディング・葬儀",
    ],
    [
        "株式会社ビズリーチ",
        "https://www.bizreach.jp/lp/official/pc/",
        "リスティング広告",
        "03-4500-6066",
        "ハイクラス転職・スカウト型転職サービス",
        "転職 ハイクラス スカウト 年収アップ キャリア",
        "転職・就職",
    ],
    [
        "パーソルキャリア株式会社",
        "https://doda.jp/promo/lp/004.html",
        "リスティング広告",
        "03-6213-3200",
        "転職エージェント・求人紹介（doda）",
        "転職 エージェント doda 求人 転職サイト",
        "転職・就職",
    ],
    [
        "株式会社ベースフード",
        "https://shop.basefood.co.jp/lp/diet_05/",
        "リスティング広告",
        "03-4500-6633",
        "完全栄養食定期購入（BASE BREAD）",
        "完全栄養食 ダイエット 食事 定期購入 ベースブレッド",
        "食品EC・D2C",
    ],
    [
        "株式会社ティーケーピー",
        "https://www.tkp.jp/lp/officepass/",
        "リスティング広告",
        "03-5322-7600",
        "貸会議室・シェアオフィスサービス",
        "貸会議室 シェアオフィス テレワーク 会議室 東京",
        "オフィス・コワーキング",
    ],
    [
        "綜合警備保障株式会社",
        "https://www.alsok.co.jp/lp/hs10/",
        "リスティング広告",
        "03-5781-1120",
        "ホームセキュリティ・防犯システム",
        "ホームセキュリティ 防犯 ALSOK 警備 一戸建て",
        "セキュリティ・警備",
    ],
    [
        "株式会社グロービス",
        "https://globis.jp/lp/np_22010_v2/",
        "リスティング広告",
        "03-5275-3578",
        "ビジネス動画学習サービス（グロービス学び放題）",
        "ビジネス 学習 MBA 動画 スキルアップ グロービス",
        "教育・スキルアップ",
    ],
]

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="メイリオ", size=10, bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
ROW_FILLS = [PatternFill("solid", fgColor="F2F7FC"), PatternFill("solid", fgColor="FFFFFF")]
ROW_FONT = Font(name="メイリオ", size=9)
ROW_ALIGN = Alignment(vertical="center", wrap_text=False)
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def make_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "架電リスト"

    # Header row
    for col_idx, h in enumerate(HEADER, start=1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
        cell.border = BORDER
    ws.row_dimensions[1].height = 22

    # Data rows
    for row_idx, row_data in enumerate(DATA, start=2):
        fill = ROW_FILLS[(row_idx - 2) % 2]
        for col_idx, val in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.fill = fill
            cell.border = BORDER
            if col_idx == 2 and val:  # LP URL column: hyperlink
                cell.hyperlink = val
                cell.font = Font(name="メイリオ", size=9, color="0563C1", underline="single")
                cell.alignment = ROW_ALIGN
            else:
                cell.font = ROW_FONT
                cell.alignment = ROW_ALIGN
        ws.row_dimensions[row_idx].height = 18

    # Column widths
    for col_idx, width in enumerate(COL_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # Freeze panes and auto filter
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    wb.save(OUTPUT_PATH)
    print(f"保存完了: {OUTPUT_PATH}")
    print(f"総件数: {len(DATA)}社")


if __name__ == "__main__":
    make_xlsx()
