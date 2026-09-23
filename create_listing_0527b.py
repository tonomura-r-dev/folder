import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUTPUT_PATH = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー\架電リスト_リスティング_2026-05-27b.xlsx"

HEADER = ["企業名", "LP URL", "広告種別", "電話番号", "商材", "検索KW", "業界"]
COL_WIDTHS = [32, 50, 16, 16, 20, 32, 22]

DATA = [
    (
        "株式会社メディウェル（薬剤師転職ドットコム）",
        "https://www.ph-10.com/lp/ph_privatly/form.html",
        "リスティング広告", "",
        "薬剤師・薬局専門の転職エージェントサービス",
        "薬剤師 転職 エージェント 高収入, 薬剤師 求人 クリニック おすすめ, 薬局 転職 専門 サイト",
        "転職・就職・人材派遣"
    ),
    (
        "株式会社イクウェル・アンド・カンパニー（EQWELチャイルドアカデミー）",
        "https://www.eqwel.jp/lp/",
        "リスティング広告", "",
        "0歳からの幼児教室・右脳脳育・知育プログラム",
        "幼児教室 右脳 知育 体験, 0歳 1歳 習い事 教室, 子供 脳育 フラッシュカード 教育",
        "教育・スキルアップ"
    ),
    (
        "株式会社グランフィットネス（グランフィットネス24）",
        "https://fitness365.jp/lp/",
        "リスティング広告", "",
        "24時間フィットネスジム（関西15拠点展開）",
        "フィットネスジム 24時間 関西 月額, ジム 安い 通い放題 大阪, 筋トレ ジム 近く 入会",
        "フィットネス・スポーツ・ゴルフ"
    ),
    (
        "株式会社ライフサロン（ほけんの相談ショップ）",
        "https://lifesalon.jp/lp/ls",
        "リスティング広告", "",
        "無料保険相談・見直しサービス（全国44店舗）",
        "保険 無料相談 近く 店舗, 生命保険 見直し FP 相談, 保険相談 専門家 無料 窓口",
        "保険・金融・証券・カード"
    ),
    (
        "はなさく生命保険株式会社",
        "https://www.life8739.co.jp/lp/teiki001",
        "リスティング広告", "",
        "定期保険・引受緩和型生命保険（通販型・ネット申込）",
        "定期保険 通販 ネット申込み 安い, 生命保険 掛け捨て シンプル 比較, がん保険 引受緩和 持病 加入",
        "保険・金融・証券・カード"
    ),
    (
        "株式会社EPARK（EPARKビューティー）",
        "https://beauty.epark.jp/lp/famous-003/",
        "リスティング広告", "",
        "美容院・ヘアサロン・ネイル・まつげ検索予約プラットフォーム",
        "美容院 予約 ポイント 当日, ヘアサロン 近く 口コミ 検索, ネイル まつ毛 予約 割引",
        "美容クリニック・脱毛・エステ"
    ),
    (
        "株式会社サンライフコーポレーション（サンライフプラス）",
        "https://www.sunlife-corporation.jp/sunlife-plus/lp/",
        "リスティング広告", "",
        "家庭用太陽光発電・蓄電池の設置・施工サービス",
        "太陽光発電 設置 費用 補助金 業者, 蓄電池 家庭用 価格 工事 比較, ソーラーパネル 設置 業者 見積もり",
        "不動産・賃貸・リノベーション・投資"
    ),
    (
        "スマートソーラー株式会社",
        "https://www.smartsolar.co.jp/solarchikuden/lp/",
        "リスティング広告", "",
        "太陽光発電＋蓄電池セット導入・販売（無料見積もり）",
        "ソーラー 蓄電池 セット 価格 補助金, 太陽光 蓄電池 一緒に 設置 費用, 蓄電池 見積もり 無料 メーカー",
        "不動産・賃貸・リノベーション・投資"
    ),
    (
        "株式会社レンデックス（LENDEX）",
        "https://lendex.jp/lp/",
        "リスティング広告", "",
        "融資型クラウドファンディング（平均利回り8%超の資産運用）",
        "ソーシャルレンディング 高利回り 登録, 融資型 クラウドファンディング 始め方 少額, 不動産 担保 利回り 8% 投資",
        "保険・金融・証券・カード"
    ),
    (
        "株式会社保険デザイン（保険deあんしん館）",
        "https://www.hoken-anshinkan.jp/lp/",
        "リスティング広告", "",
        "訪問・オンライン対応の無料保険相談・見直しサービス",
        "保険 見直し 無料 相談 訪問, 生命保険 比較 専門家 アドバイス, 保険 相談 自宅 無料 FP",
        "保険・金融・証券・カード"
    ),
]

def make_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "架電リスト"

    header_font   = Font(name="メイリオ", size=10, bold=True, color="FFFFFF")
    header_fill   = PatternFill("solid", fgColor="1F4E79")
    header_align  = Alignment(horizontal="center", vertical="center", wrap_text=True)

    fill_odd  = PatternFill("solid", fgColor="F2F7FC")
    fill_even = PatternFill("solid", fgColor="FFFFFF")
    data_font = Font(name="メイリオ", size=9)
    data_align = Alignment(vertical="center", wrap_text=False)
    url_font  = Font(name="メイリオ", size=9, color="0563C1", underline="single")

    thin = Side(style="thin", color="D0D7DE")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col, name in enumerate(HEADER, start=1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = border

    for row_idx, row in enumerate(DATA, start=2):
        fill = fill_odd if row_idx % 2 == 0 else fill_even
        for col, val in enumerate(row, start=1):
            cell = ws.cell(row=row_idx, column=col, value=val)
            cell.fill = fill
            cell.border = border
            cell.alignment = data_align
            if col == 2 and val:
                cell.hyperlink = val
                cell.font = url_font
            else:
                cell.font = data_font

    for col, width in enumerate(COL_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width

    ws.row_dimensions[1].height = 22
    for r in range(2, len(DATA) + 2):
        ws.row_dimensions[r].height = 18

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADER))}1"

    wb.save(OUTPUT_PATH)
    print(f"保存完了: {OUTPUT_PATH}")
    print(f"件数: {len(DATA)}社")

if __name__ == "__main__":
    make_xlsx()
