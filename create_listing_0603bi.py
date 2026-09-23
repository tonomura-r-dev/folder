import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUTPUT_PATH = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー\架電リスト_リスティング_2026-06-03bi.xlsx"
HEADER = ["企業名", "LP URL", "広告種別", "法人番号", "商材", "検索KW", "業界"]
COL_WIDTHS = [32, 56, 16, 16, 22, 32, 22]
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="メイリオ", size=10, bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
ROW_FILLS = [PatternFill("solid", fgColor="F2F7FC"), PatternFill("solid", fgColor="FFFFFF")]
ROW_FONT = Font(name="メイリオ", size=9)
ROW_ALIGN = Alignment(vertical="center", wrap_text=False)
URL_FONT = Font(name="メイリオ", size=9, color="0563C1", underline="single")

DATA = [
    # ── ビジネスチャット（ByteDance） ──
    ("Lark Technologies株式会社",   "https://www.larksuite.com/paid/chat-alternative-jp",    "リスティング広告", "", "Lark・チャット+ドキュメント+会議統合",        "ビジネスチャット 無料 比較, Lark ツール 評判 料金, チームコミュニケーション", "コラボSaaS"),
    # ── チームコミュニケーション ──
    ("Slack Japan株式会社",          "https://slack.com/intl/ja-jp/productivity-tts",          "リスティング広告", "", "Slack・AIビジネスチャット・プロジェクト管理",  "Slack 料金 評判 法人, ビジネスチャット 比較 SaaS, チャット ツール 無料", "コラボSaaS"),
    # ── ドキュメント管理 ──
    ("Notion Labs Japan合同会社",    "https://www.notion.so/ja-jp/pricing",                    "リスティング広告", "", "Notion・ドキュメント+DB・無料プランあり",     "Notion 料金 評判, ドキュメント 管理 SaaS, チーム wiki ツール",      "コラボSaaS"),
    # ── デザインツール ──
    ("Figma Japan合同会社",          "https://www.figma.com/pricing/",                          "リスティング広告", "", "Figma・UIデザイン・プロトタイピング",          "Figma 料金 評判, デザイン ツール SaaS, UIデザイン 共同作業",         "デザインSaaS"),
    # ── プロジェクト管理SaaS ──
    ("Atlassian株式会社",            "https://ja.atlassian.com/software/confluence/try",        "リスティング広告", "", "Confluence・Jira・ナレッジ管理・無料試用",     "Confluence 料金 評判, Jira プロジェクト管理, Atlassian 無料 試し", "コラボSaaS"),
    # ── ワークマネジメント ──
    ("monday.com株式会社",           "https://monday.com/lang/ja/lp/crm",                       "リスティング広告", "", "monday.com・CRM・プロジェクト管理・245,000社", "monday.com 料金 評判, プロジェクト管理 SaaS 比較, ワーク管理 ツール", "ワークマネジメントSaaS"),
    # ── カスタマーサポートSaaS ──
    ("Intercom Inc.",                "https://www.intercom.com/webinar-series/demo",             "リスティング広告", "", "Intercom・AI顧客サポート・チャット・デモ",    "Intercom 料金 評判, カスタマーサポート SaaS AI, チャット CS ツール", "カスタマーサポートSaaS"),
    # ── ノーコードWebサイト ──
    ("Webflow Inc.",                 "https://webflow.com/pricing",                             "リスティング広告", "", "Webflow・ノーコードWebサイト・エンタープライズ", "Webflow 料金 評判, ノーコード サイト 作成, Web デザイン ツール",    "ノーコードSaaS"),
    # ── オンラインホワイトボード ──
    ("Miro Japan株式会社",           "https://miro.com/ja/pricing/",                            "リスティング広告", "", "Miro・オンラインホワイトボード・コラボ",       "Miro 料金 評判, オンライン ホワイトボード SaaS, チーム 図解 ツール", "コラボSaaS"),
    # ── デザインツール ──
    ("Canva Japan株式会社",          "https://www.canva.com/ja_jp/pricing/",                    "リスティング広告", "", "Canva・グラフィックデザイン・LP制作・無料",   "Canva 料金 評判, デザイン ツール 無料 使い方, 画像 作成 クラウド",  "デザインSaaS"),
    # ── フォームSaaS ──
    ("Typeform S.L.",                "https://www.typeform.com/pricing",                        "リスティング広告", "", "Typeform・AIフォーム・3.5倍のデータ取得",    "フォーム 作成 ツール 比較, Typeform 料金 評判, 調査 申し込み SaaS", "フォームSaaS"),
    # ── メールマーケ ──
    ("The Rocket Science Group LLC", "https://mailchimp.com/pricing/marketing/",                "リスティング広告", "", "Mailchimp・メール配信・14日間無料試用",       "Mailchimp 料金 評判, メール 配信 SaaS 比較, マーケ オートメーション", "メールマーケSaaS"),
    # ── CRM営業管理 ──
    ("Pipedrive株式会社",            "https://www.pipedrive.com/en/pricing",                    "リスティング広告", "", "Pipedrive・CRM・営業パイプライン管理",         "Pipedrive 料金 評判, CRM 営業 管理 SaaS, セールス パイプライン",    "CRM SaaS"),
    # ── プロダクト分析SaaS ──
    ("Mixpanel Inc.",                "https://mixpanel.com/m/product-analytics-ja/",            "リスティング広告", "", "Mixpanel・プロダクト分析・行動データ可視化",  "Mixpanel 料金 評判, プロダクト 分析 SaaS, ユーザー 行動 分析",      "プロダクト分析SaaS"),
    # ── デジタル分析SaaS ──
    ("Amplitude Inc.",               "https://amplitude.com/ja-jp/get-started",                 "リスティング広告", "", "Amplitude・デジタル分析・45,000サービス導入", "Amplitude 料金 評判, デジタル アナリティクス SaaS, A/B テスト 分析", "プロダクト分析SaaS"),
    # ── コーディングテスト ──
    ("株式会社ギブリー",             "https://tracks.run/tracktest-no1/",                       "リスティング広告", "", "Track Test・コーディングテスト・AI評価",      "コーディングテスト 採用 ツール, エンジニア スキル評価 SaaS, Track Test", "HR SaaS"),
    # ── eKYC本人確認 ──
    ("株式会社TRUSTDOCK",            "https://biz.trustdock.io/lp-ekyc-01v1",                   "リスティング広告", "", "eKYC・デジタル本人確認・導入No.1",           "eKYC ツール 比較, 本人確認 API SaaS, TRUSTDOCK 料金 評判",          "IDtech SaaS"),
    # ── 現場DX ──
    ("株式会社カミナシ",             "https://lp.kaminashi.jp/factory-audit",                   "リスティング広告", "", "カミナシ・現場DX・チェックシート電子化",      "現場 DX SaaS, チェックシート 電子化 比較, カミナシ 料金 評判",       "現場DXSaaS"),
    # ── タレントマネジメント ──
    ("株式会社カオナビ",             "https://www.kaonavi.jp/",                                 "リスティング広告", "", "カオナビ・タレントマネジメント・シェア8年No.1", "カオナビ 料金 評判, タレントマネジメント SaaS, 人事評価 クラウド",  "HR SaaS"),
    # ── 受付システム ──
    ("株式会社RECEPTIONIST",         "https://receptionist.jp/lp04/",                           "リスティング広告", "", "RECEPTIONIST・クラウド受付・4,000社導入",     "受付 システム SaaS 比較, クラウド 来客 管理, RECEPTIONIST 料金",    "業務DXSaaS"),
]


def make_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "架電リスト"

    for col_idx, header in enumerate(HEADER, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
    ws.row_dimensions[1].height = 20

    for row_idx, row_data in enumerate(DATA, 2):
        fill = ROW_FILLS[(row_idx - 2) % 2]
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.fill = fill
            cell.alignment = ROW_ALIGN
            if col_idx == 2:
                cell.hyperlink = value
                cell.font = URL_FONT
            else:
                cell.font = ROW_FONT
        ws.row_dimensions[row_idx].height = 16

    for col_idx, width in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADER))}1"

    wb.save(OUTPUT_PATH)
    print(f"Saved: {OUTPUT_PATH}")
    print(f"Rows: {len(DATA)}")


if __name__ == "__main__":
    make_xlsx()
