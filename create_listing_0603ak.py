import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUTPUT_PATH = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー\架電リスト_リスティング_2026-06-03ak.xlsx"
HEADER = ["企業名", "LP URL", "広告種別", "法人番号", "商材", "検索KW", "業界"]
COL_WIDTHS = [32, 56, 16, 16, 22, 32, 22]
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="メイリオ", size=10, bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
ROW_FILLS = [PatternFill("solid", fgColor="F2F7FC"), PatternFill("solid", fgColor="FFFFFF")]
ROW_FONT = Font(name="メイリオ", size=9)
ROW_ALIGN = Alignment(vertical="center", wrap_text=False)
URL_FONT = Font(name="メイリオ", size=9, color="0563C1", underline="single")

# ── 全社、WebSearchで広告出稿を確認しLPを実際に取得済み ──
DATA = [
    # ── 光回線（申し込みLP） ──
    ("スマート株式会社",                   "https://smart.ne.jp/lp/listing/",                  "リスティング広告", "", "光回線・プロバイダ一体型インターネット",   "光回線 おすすめ 比較, プロバイダ 一体型 安い, スマート光 評判",             "光回線"),
    # ── 薬局宅配（申し込みLP） ──
    ("おかぴファーマシーシステム株式会社", "https://todokusuri.com/lp/003/",                   "リスティング広告", "", "処方箋薬 宅配・ポスト受け取り",           "処方箋 宅配 送料無料, お薬 自宅 届く, 薬局 オンライン 配達",               "薬局宅配"),
    # ── ファクタリング（資金調達LP） ──
    ("株式会社ビートレーディング",         "https://betrading.jp/",                             "リスティング広告", "", "ファクタリング・売掛金・即日資金調達",     "ファクタリング 即日 おすすめ, 売掛金 資金調達 最短, ビートレーディング 評判","ファクタリング"),
    # ── クラウドPOSレジ（無料体験LP） ──
    ("株式会社スマレジ",                   "https://smaregi.jp/",                               "リスティング広告", "", "クラウドPOSレジ・無料トライアル30日",     "POSレジ クラウド おすすめ, レジアプリ 無料 比較, スマレジ 評判 料金",       "POSシステム"),
    # ── スキンケア（公式ECサイト） ──
    ("株式会社ニコリオ",                   "https://www.nicorio.co.jp/",                        "リスティング広告", "", "美容家電・スキンケアデバイス・定期便",     "ニコリオ 美顔器 効果, スキンケア デバイス おすすめ, 美容器 定期 比較",      "美容家電"),
    # ── 婦人科・ブライダルチェック（予約LP） ──
    ("クリニックTEN株式会社",              "https://clinicten.jp/lp/bridalcheck/",              "リスティング広告", "", "ブライダルチェック・婦人科・予約",         "ブライダルチェック 東京 予約, 婦人科 渋谷 検査, 結婚前 健康診断 女性",      "婦人科クリニック"),
    # ── ペット保険（申し込みLP） ──
    ("au損害保険株式会社",                 "https://www.au-sonpo.co.jp/pc/pet/",                "リスティング広告", "", "ペット保険・犬・猫・月払い",               "ペット保険 おすすめ 比較, au 犬 猫 保険, ペット 医療費 保険 安い",          "ペット保険"),
    # ── 就活プラットフォーム（会員登録LP） ──
    ("株式会社ワンキャリア",               "https://www.onecareer.jp/users/sign_up",            "リスティング広告", "", "新卒就活・ES対策・選考情報",               "就活 サイト おすすめ 新卒, インターン 情報 無料, ワンキャリア 登録",         "就活サービス"),
    # ── セブンイレブン保険窓口（ペット保険LP） ──
    ("株式会社セブン‐イレブン（au損保提携）", "https://www.7-insurance.jp/lp/sej_aupet/",     "リスティング広告", "", "ペット保険・セブン窓口・かんたん申込",     "ペット保険 コンビニ 申し込み, 犬 猫 保険 簡単, セブンイレブン 保険 比較",  "ペット保険"),
]


def make_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "架電リスト"

    for col_idx, header in enumerate(HEADER, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
    ws.row_dimensions[1].height = 20

    for row_idx, row_data in enumerate(DATA, 2):
        fill = ROW_FILLS[(row_idx - 2) % 2]
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.fill = fill
            cell.alignment = ROW_ALIGN
            if col_idx == 2:
                cell.hyperlink = value
                cell.font = URL_FONT
            else:
                cell.font = ROW_FONT
        ws.row_dimensions[row_idx].height = 16

    for col_idx, width in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADER))}1"

    wb.save(OUTPUT_PATH)
    print(f"Saved: {OUTPUT_PATH}")
    print(f"Rows: {len(DATA)}")


if __name__ == "__main__":
    make_xlsx()
