import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUTPUT_PATH = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー\架電リスト_リスティング_2026-06-03a.xlsx"
HEADER = ["企業名", "LP URL", "広告種別", "法人番号", "商材", "検索KW", "業界"]
COL_WIDTHS = [32, 50, 16, 16, 20, 32, 22]
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="メイリオ", size=10, bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
ROW_FILLS = [PatternFill("solid", fgColor="F2F7FC"), PatternFill("solid", fgColor="FFFFFF")]
ROW_FONT = Font(name="メイリオ", size=9)
ROW_ALIGN = Alignment(vertical="center", wrap_text=False)
URL_FONT = Font(name="メイリオ", size=9, color="0563C1", underline="single")

DATA = [
    # ── 葬儀・家族葬 ──────────────────────────────────────────
    ("株式会社家族葬のファミーユ",        "https://www.famille-kazokusou.com/",              "リスティング広告", "6010401121066", "家族葬・小規模葬儀プラン",           "家族葬 費用 相場, 家族葬 葬儀社 選び方, 小さな葬儀 安い おすすめ",     "葬儀/家族葬"),
    ("株式会社公益社",                   "https://www.koekisha.co.jp/lp/kazokusou/",        "リスティング広告", "6120001108824", "家族葬・一般葬（近畿エリア）",        "公益社 葬儀 費用, 家族葬 大阪 おすすめ, 葬儀社 関西 評判",             "葬儀/家族葬"),
    ("株式会社セレモア",                 "https://www.ceremore.co.jp/",                     "リスティング広告", "8012801009302", "葬儀・家族葬（首都圏）",              "葬儀 費用 東京, 家族葬 東京 比較, 葬儀社 首都圏 安心",                 "葬儀/家族葬"),
    ("株式会社ベルコ",                   "https://www.bellco.co.jp/sougi/",                 "リスティング広告", "9120901021743", "互助会・葬儀全国ネットワーク",        "互助会 葬儀 積立, 全国 葬儀社 おすすめ, 葬式 費用 積立 安い",          "冠婚葬祭"),
    ("イオンライフ株式会社",             "https://www.aeonlife.jp/lp/campaign202412",       "リスティング広告", "5120001105005", "イオンのお葬式・家族葬",              "イオン 葬儀 費用, 家族葬 格安 比較, お葬式 シンプル プラン",            "葬儀/家族葬"),
    ("株式会社ティア",                   "https://www.tear.co.jp/",                         "リスティング広告", "6060001028707", "家族葬・直葬（東海・関東）",          "家族葬 名古屋 費用, ティア 葬儀 評判, 直葬 費用 流れ",                 "葬儀/家族葬"),
    ("株式会社日本セレモニー",           "https://www.tenreikaikan.com/",                   "リスティング広告", "1250001006136", "天礼会館・家族葬（中国エリア）",      "家族葬 広島 費用, 葬儀社 中国地方 おすすめ, 天礼会館 評判",            "葬儀/家族葬"),
    ("株式会社メモリアルアートの大野屋", "https://www.ohnoya.co.jp/",                       "リスティング広告", "9011101056574", "葬儀・お墓・仏壇・終活相談",          "葬儀 費用 仏壇 セット, お墓 購入 相談, 終活 始める 資料請求",          "葬儀/終活"),
    ("株式会社セレマ",                   "https://cerema.co.jp/",                           "リスティング広告", "2130001021294", "冠婚葬祭互助会・葬儀（近畿）",        "互助会 葬儀 近畿, 結婚式 葬儀 積立 セット, 葬儀社 京都 安心",          "冠婚葬祭"),
    ("株式会社さがみ典礼",               "https://www.sagamitenrei.com/",                   "リスティング広告", "",              "家族葬・一般葬（神奈川・首都圏）",    "家族葬 神奈川 費用, 葬儀社 横浜 おすすめ, さがみ典礼 評判 口コミ",     "葬儀/家族葬"),
    # ── フォトスタジオ・記念写真 ──────────────────────────────────────────
    ("株式会社キタムラ（スタジオマリオ）","https://www.studio-mario.jp/lp/mario_renewal/",  "リスティング広告", "3490001000435", "七五三・お宮参り・入学写真",          "スタジオマリオ 七五三 予約, 写真館 七五三 費用, お宮参り 撮影 安い",   "フォトスタジオ"),
    ("株式会社キャラット（スタジオキャラット）","https://www.caratt.jp/",                  "リスティング広告", "1150001015064", "七五三・成人式・マタニティ撮影",      "スタジオキャラット 七五三, 成人式 振袖 撮影 費用, 写真館 予約 安い",   "フォトスタジオ"),
    ("株式会社写真館ピノキオ",           "https://www.pinokio.co.jp/",                      "リスティング広告", "4011401003165", "七五三・お誕生日・家族写真",          "写真館 七五三 東北, ピノキオ 七五三 費用, 子供 写真撮影 プロ",          "フォトスタジオ"),
    ("株式会社デコルテ・ホールディングス（ハピスタ）","https://happy-photo-studio.jp/",   "リスティング広告", "2010001180292", "ハピスタ・子ども写真館",              "ハピスタ 七五三 予約, 子供 写真館 関東 おすすめ, 誕生日 撮影 費用",    "フォトスタジオ"),
    ("株式会社ライフスタジオ",           "https://www.lifestudio.jp/",                      "リスティング広告", "",              "ライフスタジオ・家族・子ども撮影",    "ライフスタジオ 予約, 子供 家族写真 おしゃれ, フォトスタジオ 自然光",   "フォトスタジオ"),
    ("株式会社創寫舘",                   "https://soshakan.co.jp/",                         "リスティング広告", "",              "七五三・成人式・証明写真",            "創寫舘 七五三 予約, 成人式 撮影 費用, 写真館 証明写真 おすすめ",        "フォトスタジオ"),
    ("株式会社タートル（フォトスタジオタートル）","https://www.ps-turtle.com/",            "リスティング広告", "",              "七五三・お宮参り（愛知・名古屋）",    "フォトスタジオタートル 七五三, 名古屋 写真館 七五三 費用, 愛知 撮影",   "フォトスタジオ"),
    ("株式会社ハピリィ（ハピリィフォトスタジオ）","https://www.happilyphoto.jp/",          "リスティング広告", "",              "七五三・お宮参り・バースデー撮影",    "ハピリィ 七五三 予約, フォトスタジオ 関東 おすすめ, お宮参り 撮影 費用","フォトスタジオ"),
    ("株式会社らかんスタジオ",           "https://laquan.com/",                             "リスティング広告", "2012401020086", "七五三・成人式・ブライダル撮影",      "らかんスタジオ 七五三, 成人式 写真 振袖 撮影, 写真館 東京 予約",        "フォトスタジオ"),
    ("プレシュスタジオ（株式会社フォトアトリエ）","https://www.precieux-studio.com/",      "リスティング広告", "",              "貸切型こども写真館（関東）",          "プレシュスタジオ 七五三 予約, 貸切 写真館 東京, お宮参り 撮影 完全予約","フォトスタジオ"),
]


def make_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "架電リスト"

    for col_idx, header in enumerate(HEADER, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
    ws.row_dimensions[1].height = 20

    for row_idx, row_data in enumerate(DATA, 2):
        fill = ROW_FILLS[(row_idx - 2) % 2]
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.fill = fill
            cell.alignment = ROW_ALIGN
            if col_idx == 2:
                cell.hyperlink = value
                cell.font = URL_FONT
            else:
                cell.font = ROW_FONT
        ws.row_dimensions[row_idx].height = 16

    for col_idx, width in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADER))}1"

    wb.save(OUTPUT_PATH)
    print(f"Saved: {OUTPUT_PATH}")
    print(f"Rows: {len(DATA)}")


if __name__ == "__main__":
    make_xlsx()
