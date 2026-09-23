import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUTPUT_PATH = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー\架電リスト_リスティング_2026-05-29d.xlsx"

HEADER = ["企業名", "LP URL", "広告種別", "電話番号", "商材", "検索KW", "業界"]
COL_WIDTHS = [32, 50, 16, 16, 20, 32, 22]

DATA = [
    [
        "株式会社トキハナ",
        "https://tokihana.net/lp/tokihana_line/",
        "リスティング広告",
        "03-6555-2646",
        "結婚式場比較・予約",
        "結婚式場 比較 挙式費用 式場探し ウェディング",
        "ウエディング・葬儀",
    ],
    [
        "株式会社ジョイナス",
        "https://zero-wedding.jp/lp/",
        "リスティング広告",
        "06-4397-3337",
        "格安ウェディングプロデュース",
        "結婚式 少人数 格安 プロデュース 大阪 ウェディング",
        "ウエディング・葬儀",
    ],
    [
        "株式会社フジテック",
        "https://ac.fj-tec.co.jp/lp/matome01/lpm.html",
        "リスティング広告",
        "0120-366-244",
        "エアコン取り付け・交換工事",
        "エアコン取り付け 工事 交換 業務用 東京",
        "住宅設備・工事",
    ],
    [
        "株式会社長沼静きもの学院",
        "https://www.naganuma-kimono.co.jp/lp/kitsuke/",
        "リスティング広告",
        "0120-330-911",
        "着物着付け教室・レッスン",
        "着付け教室 着物 着付け 初心者 無料体験",
        "教育・スキルアップ",
    ],
    [
        "株式会社ライフソムリエール",
        "https://lifesommelier.com/housecleaning/lp/aircon",
        "リスティング広告",
        "0120-953-887",
        "家事代行・ハウスクリーニング",
        "家事代行 ハウスクリーニング エアコン掃除 大阪",
        "家事代行・クリーニング",
    ],
    [
        "株式会社ROZINAS",
        "https://rozinas.co.jp/lp/",
        "リスティング広告",
        "0120-540-644",
        "遺品整理・不用品回収",
        "遺品整理 不用品回収 生前整理 片付け",
        "不用品回収・遺品整理",
    ],
    [
        "株式会社新東亜工業",
        "https://shintoakogyo.co.jp/lp/paint/",
        "リスティング広告",
        "0120-663-642",
        "外壁塗装・大規模修繕・防水工事",
        "外壁塗装 マンション 大規模修繕 防水工事",
        "リフォーム・外壁塗装",
    ],
    [
        "株式会社ライフサービス",
        "https://www.l-service.jp/lp/",
        "リスティング広告",
        "0120-320-370",
        "高品質家事代行（掃除・料理・洗濯）",
        "家事代行 料理代行 お掃除代行 東京",
        "家事代行・クリーニング",
    ],
    [
        "東急リデザイン株式会社",
        "https://reform.tokyu-re-design.co.jp/lp/mansion/sanitary/",
        "リスティング広告",
        "0120-935-109",
        "マンション水まわりリフォーム",
        "マンション リフォーム 水まわり キッチン 浴室",
        "リフォーム・住宅設備",
    ],
    [
        "渋谷DSクリニック",
        "https://dsclinic.jp/ac/lp/a/counseling43_05/",
        "リスティング広告",
        "0120-951-135",
        "医療ダイエット・美容皮膚科",
        "医療ダイエット 肥満外来 GLP-1 美容クリニック 渋谷",
        "美容クリニック・医療",
    ],
]

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="メイリオ", size=10, bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
ROW_FILLS = [PatternFill("solid", fgColor="F2F7FC"), PatternFill("solid", fgColor="FFFFFF")]
ROW_FONT = Font(name="メイリオ", size=9)
ROW_ALIGN = Alignment(vertical="center", wrap_text=False)
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def make_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "架電リスト"

    # Header row
    for col_idx, h in enumerate(HEADER, start=1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
        cell.border = BORDER
    ws.row_dimensions[1].height = 22

    # Data rows
    for row_idx, row_data in enumerate(DATA, start=2):
        fill = ROW_FILLS[(row_idx - 2) % 2]
        for col_idx, val in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.fill = fill
            cell.border = BORDER
            if col_idx == 2 and val:  # LP URL column: hyperlink
                cell.hyperlink = val
                cell.font = Font(name="メイリオ", size=9, color="0563C1", underline="single")
                cell.alignment = ROW_ALIGN
            else:
                cell.font = ROW_FONT
                cell.alignment = ROW_ALIGN
        ws.row_dimensions[row_idx].height = 18

    # Column widths
    for col_idx, width in enumerate(COL_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # Freeze panes and auto filter
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    wb.save(OUTPUT_PATH)
    print(f"保存完了: {OUTPUT_PATH}")
    print(f"総件数: {len(DATA)}社")


if __name__ == "__main__":
    make_xlsx()
