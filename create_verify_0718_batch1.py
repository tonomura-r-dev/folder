import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUTPUT_PATH = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー\架電リスト_全社再検証_2026-07-18_batch1.xlsx"
HEADER = ["企業名", "LP URL(広告確認済み)", "電話番号", "検索KW", "媒体"]
COL_WIDTHS = [40, 50, 16, 32, 14]
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="メイリオ", size=10, bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
ROW_FILLS = [PatternFill("solid", fgColor="F2F7FC"), PatternFill("solid", fgColor="FFFFFF")]
ROW_FONT = Font(name="メイリオ", size=9)
ROW_ALIGN = Alignment(vertical="center", wrap_text=False)
URL_FONT = Font(name="メイリオ", size=9, color="0563C1", underline="single")

# 広告ライブラリ(Meta)で実際に配信中の広告を1社ずつ確認できたものだけ記載。
DATA = [
    ("ホームテック株式会社（entrie/エントリエ）", "https://entrie.net/chikufuru/", "", "リノベーション 築古, 戸建てリノベ", "Meta"),
    ("ジューテックホーム株式会社", "https://jutec-home.jp/", "045-595-3222", "注文住宅 東京, 高気密高断熱住宅", "Meta"),
    ("株式会社アーキテックプランニング（DUALHOME）", "https://dualhome.jp/", "011-215-6241", "注文住宅 デザイン住宅, DUALHOME", "Meta"),
    ("株式会社アワハウス（旧アートホーム）", "https://a-ourhouse.co.jp/", "0123-40-8686", "注文住宅 北海道, ミニプロト", "Meta"),
    ("グッドリビング株式会社（TATTA!）", "https://tatta-goodliving.com/", "053-432-7455", "注文住宅 定額, 平屋 注文住宅", "Meta"),
    ("マイコミュニケーション株式会社（保険ほっとライン）", "https://www.hoholine.com/", "", "保険相談, 保険見直し", "Meta"),
    ("株式会社プレアデス（ほいくジョブ）", "https://hoiku-job.net/", "0120-269-600", "保育士 転職, 学童 求人", "Meta"),
    ("株式会社モニクルフィナンシャル（マネイロ）", "https://moneiro.jp/", "03-6868-8637", "資産運用 相談, 投資診断", "Meta"),
    ("株式会社メディカル・コンシェルジュ（MCナースネット）", "https://mc-nurse.net/", "", "看護師 単発バイト, 看護師 求人", "Meta"),
    ("Dr.AGAクリニック（医療法人社団日昇会）", "https://drskinclinic.jp/", "03-3294-7070", "AGA治療, 発毛メソ治療", "Meta"),
    ("ルシアクリニック（医療法人緑風会）", "https://lucia-c.com/", "", "医療脱毛, 脱毛クリニック", "Meta"),
    ("AND美容外科（一般社団法人AND medical group）", "https://and-bg.com/", "", "美容外科, 美容医療", "Meta"),
    ("富国生命保険相互会社（フコク生命）", "https://www.fukoku-life.co.jp/", "03-3508-1101", "生命保険, 保険 見直し", "Meta"),
    ("株式会社タイミー", "https://timee.co.jp/", "03-6822-3013", "スキマバイト, 単発バイト アプリ", "Meta"),
    ("ラクスル株式会社", "https://raksul.com/", "03-4530-3733", "ネット印刷 安い, チラシ 印刷", "Meta"),
]

# 広告ライブラリで検索したが該当する配信中広告が見つからなかった企業（要再調査・除外候補）
NOT_VERIFIED = [
    "株式会社ティー・エフ・オフィス（保険コンパス）",
    "株式会社エフエムエス中央クリニック（中央クリニック）",
    "パール金属株式会社",
    "株式会社京進",
    "ポラス株式会社",
    "ヤマダホームズ株式会社",
    "株式会社マークスアンドウェブ",
    "イプサカンパニー株式会社",
]


def make_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "確認済み"
    for col_idx, header in enumerate(HEADER, 1):
        c = ws.cell(row=1, column=col_idx, value=header)
        c.fill = HEADER_FILL; c.font = HEADER_FONT; c.alignment = HEADER_ALIGN
    ws.row_dimensions[1].height = 20
    for row_idx, row in enumerate(DATA, 2):
        fill = ROW_FILLS[(row_idx - 2) % 2]
        for col_idx, value in enumerate(row, 1):
            c = ws.cell(row=row_idx, column=col_idx, value=value)
            c.fill = fill; c.alignment = ROW_ALIGN
            if col_idx == 2:
                c.hyperlink = value; c.font = URL_FONT
            else:
                c.font = ROW_FONT
        ws.row_dimensions[row_idx].height = 16
    for col_idx, width in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADER))}1"

    ws2 = wb.create_sheet("未確認")
    ws2.cell(row=1, column=1, value="広告ライブラリで配信中の広告が見つからなかった企業（要再調査・除外候補）").font = Font(bold=True)
    for i, name in enumerate(NOT_VERIFIED, 2):
        ws2.cell(row=i, column=1, value=name)
    ws2.column_dimensions["A"].width = 55

    wb.save(OUTPUT_PATH)
    print(f"Saved: {OUTPUT_PATH}  Confirmed: {len(DATA)}  Not verified: {len(NOT_VERIFIED)}")


if __name__ == "__main__":
    make_xlsx()
