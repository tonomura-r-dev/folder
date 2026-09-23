import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def make_xlsx(filename, data):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "架電リスト"
    headers = ["企業名", "LP URL", "電話番号", "商材", "検索KW", "業界"]
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="メイリオ", bold=True, color="FFFFFF", size=10)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center
        cell.border = border

    for row_idx, row in enumerate(data, 2):
        fill_color = "F2F7FC" if row_idx % 2 == 0 else "FFFFFF"
        row_fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
        for col_idx, val in enumerate(row, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.fill = row_fill
            cell.font = Font(name="メイリオ", size=9)
            cell.alignment = left
            cell.border = border

    col_widths = [30, 45, 16, 20, 30, 18]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.row_dimensions[1].height = 20
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    wb.save(filename)
    print(f"保存完了: {filename} ({len(data)}社)")


# 列順: 企業名 | LP URL | 電話番号 | 商材 | 検索KW | 業界
# 2026-06-29 Meta広告架電リスト（第130弾）
data = [
    # ── フィットネス・パーソナルジムFC ──
    ["株式会社REJUVENATE（ReViNa）",         "https://revina-personalgym.com/",          "050-3706-1018", "女性向けパーソナルジムFC（全国100店舗）",            "パーソナルジム 安い / 体型改善 女性",         "パーソナルジムFC"],
    ["株式会社ハコジム（HACOGYM）",           "https://hacogym.jp/",                      "",              "24時間個室セルフジムFC",                            "個室ジム 無人 / 24時間ジム",                  "フィットネスFC"],
    ["トータル・ワークアウトプレミアムマネジメント株式会社", "https://totalworkout.jp/",  "",              "ハイエンドパーソナルジム",                           "パーソナルジム 六本木 / 本格トレーニング",    "パーソナルジム"],
    ["株式会社N-group（NEXUS GYM）",          "https://www.nexus-gym.com/",               "",              "完全個室パーソナルジムFC（全国50店舗）",             "パーソナルジム 個室 / 完全個室ジム 体験",     "パーソナルジムFC"],
    ["エーイーシー株式会社（ECOFIT24）",      "https://www.ecofit24.com/",                "",              "月額2,980円 24時間無人フィットネスFC（50店舗+）",    "24時間ジム 月額安い / セルフジム 全国",       "フィットネスFC"],
    ["TOKIEL（トキエル）",                    "https://tokiel.jp/",                       "03-6231-6505",  "通い放題パーソナルジムFC（全国30店舗）",             "パーソナルジム 通い放題 / ジム 月額 安い",    "パーソナルジムFC"],
    ["オレンジセオリー・ジャパン株式会社",   "https://www.orangetheoryfitness.co.jp/",   "",              "グループパーソナルトレーニング（海外発FC）",          "グループトレーニング ダイエット / フィットネス 体験", "フィットネスFC"],

    # ── ピラティスFC ──
    ["株式会社Wellness X Asia（CLUB PILATES）", "https://clubpilates.co.jp/",             "",              "マシンピラティスFC（日本80店舗以上・世界最大）",      "マシンピラティス 体験 / ピラティス スタジオ", "ピラティスFC"],
    ["株式会社理学ボディ（luluto）",          "https://luluto.jp/",                       "",              "マシンピラティスFC（全国100店舗以上）",              "マシンピラティス 通い放題 / ピラティス 初心者", "ピラティスFC"],
    ["株式会社VB NEXT（URBAN CLASSIC PILATES）", "https://urbanclassic.jp/",             "",              "マシンピラティスFC（全国69店舗）",                   "マシンピラティス ボディメイク / ピラティス 渋谷", "ピラティスFC"],

    # ── 美容エステ・整体サロンFC ──
    ["株式会社美.design（ビデザイン）",       "https://b-design32.com/",                  "",              "小顔矯正・骨盤矯正美容整体サロンFC（38店舗）",       "小顔矯正 通い放題 / 骨盤矯正 体験",          "美容整体FC"],
    ["株式会社SPEEDspring（BUPURA）",         "https://bupura.jp/",                       "",              "小顔専門エステサロンFC（全国140店舗）",              "小顔 エステ 定額 / 痩身 フェイシャル 体験",  "美容エステFC"],
    ["株式会社Lime",                          "https://lime-fit.com/",                    "",              "エステサロンFC（5ブランド 全国150店舗）",            "エステ 痩身 フェイシャル / エステ サブスク",  "エステFC"],

    # ── 脱毛FC ──
    ["セルフ脱毛ハイジ",                      "https://self-datsumou.com/",               "",              "完全無人24時間セルフ脱毛FC（業界最多店舗数）",       "セルフ脱毛 24時間 / 脱毛 安い サロン",       "脱毛FC"],
    ["Tiana（ティアナ）",                     "https://tiana-beauty.com/",                "",              "無人セルフ脱毛・美容複合型FCサロン",                 "セルフ脱毛 無人 / 脱毛 月額 24時間",         "脱毛FC"],

    # ── コスメ・スキンケアD2C ──
    ["株式会社meeth",                         "https://www.meeth.jp/",                    "03-6450-9759",  "スキンケアD2C（乾燥肌特化 韓国発）",                "スキンケア 乾燥肌 / 美容液 通販",            "コスメD2C"],
    ["株式会社SOLIA（PHOEBE BEAUTY UP）",     "https://phoebebeautyup.com/",              "",              "まつ毛美容液・スキンケアD2C（年商62億円）",          "まつ毛美容液 おすすめ / 目元ケア 通販",      "コスメD2C"],
    ["株式会社SOLIA（ALOBABY）",              "https://www.alo-organic.com/",             "",              "純国産オーガニックベビースキンケアD2C",              "ベビースキンケア 国産 / 赤ちゃん 保湿 オーガニック", "ベビーケアD2C"],
    ["ECH株式会社（KAMIKA）",                 "https://kamikacosmetics.jp/",              "03-5456-6900",  "5in1クリームシャンプーD2C",                         "クリームシャンプー 口コミ / シャンプー トリートメント 不要", "ヘアケアD2C"],
    ["HACCI's JAPAN合同会社（HACCI）",        "https://hacci1912.com/",                   "",              "はちみつコスメ・はちみつギフトD2C",                  "はちみつ スキンケア / 蜂蜜 コスメ 保湿",     "コスメD2C"],
    ["株式会社TSUKURI DESIGN（myberyl）",     "https://myberyl.com/",                     "",              "メンズBBクリーム・CCクリームD2C",                   "メンズコスメ BBクリーム / 男性 肌ケア",      "メンズコスメD2C"],

    # ── サプリメントD2C ──
    ["株式会社スピック（Lypo-C）",            "https://lypo-c.shop/",                     "",              "リポソーム型高濃度ビタミンCサプリD2C",               "ビタミンC サプリ 高濃度 / リポソーム 通販",  "サプリD2C"],
    ["ナチュラルテック株式会社（Mitas）",     "https://brands.naturaltech.jp/media/mitas-series/", "",  "時期別葉酸サプリD2C（妊活・マタニティ）",            "葉酸 サプリ 妊活 / 葉酸 妊婦 おすすめ",      "サプリD2C"],

    # ── ペット ──
    ["ペットメディカルサポート株式会社",      "https://www.petfirst.co.jp/",              "",              "ペット保険（犬・猫向け）",                           "ペット保険 比較 / 犬 猫 保険 医療費",        "ペット保険"],
    ["株式会社バイオフィリア（cocogourmet）", "https://coco-gourmet.com/",                "",              "フレッシュドッグフードD2C（サブスク）",              "ドッグフード 手作り / 国産 フレッシュ ペットフード", "ペットフードD2C"],
    ["犬猫生活株式会社",                      "https://inuneko-seikatsu.co.jp/",          "",              "国産ペットフードD2C（前年比55%増 売上17億円）",       "ドッグフード 国産 / キャットフード 無添加",  "ペットフードD2C"],
    ["オネストフード株式会社（レガリエ）",    "https://honestfood.co.jp/",                "03-4500-1780",  "国産無添加グレインフリーペットフードD2C",            "ペットフード 無添加 グレインフリー / 犬 食事", "ペットフードD2C"],

    # ── 教育 ──
    ["株式会社グノーブリンク（グノーブル）",  "https://gno-jr.com/",                      "",              "難関中学受験塾（御三家・最難関特化）",               "難関中学 受験 塾 / 麻布 開成 受験",          "学習塾"],
    ["株式会社こうゆう（スクールFC）",        "https://www.schoolfc.jp/",                 "048-835-5870",  "難関中学受験進学塾FC（花まる学習会グループ）",       "中学受験 進学塾 / 花まる 受験 塾",           "学習塾FC"],

    # ── 不動産投資 ──
    ["株式会社デュアルタップ",               "https://www.dualtap.co.jp/",               "",              "中古マンション投資・賃貸管理",                       "不動産投資 ワンルーム / マンション 投資 初心者", "不動産投資"],
    ["株式会社和不動産",                     "https://nagomi-fudousan.com/",             "03-5256-3033",  "都心ワンルームマンション投資・管理",                  "不動産投資 東京 / 不動産投資 初心者 セミナー", "不動産投資"],
    ["株式会社アセットジャパン",             "https://assetjapan.co.jp/",                "03-6264-2585",  "ワンルームマンション不動産投資",                      "不動産投資 名古屋 / ワンルーム 投資 東京",   "不動産投資"],
]

base = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー"
make_xlsx(f"{base}\\架電リスト_Meta_2026-06-29.xlsx", data)
