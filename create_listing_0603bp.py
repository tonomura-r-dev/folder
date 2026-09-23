import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUTPUT_PATH = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー\架電リスト_リスティング_2026-06-03bp.xlsx"
HEADER = ["企業名", "LP URL", "広告種別", "法人番号", "商材", "検索KW", "業界"]
COL_WIDTHS = [32, 56, 16, 16, 22, 32, 22]
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="メイリオ", size=10, bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
ROW_FILLS = [PatternFill("solid", fgColor="F2F7FC"), PatternFill("solid", fgColor="FFFFFF")]
ROW_FONT = Font(name="メイリオ", size=9)
ROW_ALIGN = Alignment(vertical="center", wrap_text=False)
URL_FONT = Font(name="メイリオ", size=9, color="0563C1", underline="single")

DATA = [
    # ── クラウド会計 ──
    ("Xero (Xero Limited)",        "https://www.xero.com/us/pricing-plans/",                  "リスティング広告", "", "Xero・クラウド会計・中小企業向け",            "クラウド 会計 SaaS, Xero 料金 評判, 中小企業 会計 ソフト",         "会計SaaS"),
    # ── 給与/HR ──
    ("Gusto Inc.",                 "https://gusto.com/pricing",                               "リスティング広告", "", "Gusto・給与計算・人事・福利厚生",             "給与計算 SaaS, Gusto 料金 評判, 人事 HR プラットフォーム",         "HR SaaS"),
    # ── HR/IT統合 ──
    ("Rippling (People Center)",   "https://www.rippling.com/pricing",                        "リスティング広告", "", "Rippling・HR/IT/財務統合プラットフォーム",     "HR IT 統合 SaaS, Rippling 料金 評判, ワークフォース 管理",         "HR SaaS"),
    # ── グローバル雇用 ──
    ("Deel Inc.",                  "https://www.deel.com/ja/pricing/",                        "リスティング広告", "", "Deel・グローバル雇用・EOR・給与処理",          "グローバル 雇用 EOR, Deel 料金 評判, 海外 人材 雇用 SaaS",        "HR SaaS"),
    # ── グローバル雇用 ──
    ("Remote (Remote Europe BV)",  "https://remote.com/pricing",                              "リスティング広告", "", "Remote・EOR・グローバル給与・コンプラ",        "EOR サービス 比較, Remote 料金 評判, 海外 雇用 給与 SaaS",         "HR SaaS"),
    # ── 経費精算 ──
    ("Expensify Inc.",             "https://www.expensify.com/pricing",                       "リスティング広告", "", "Expensify・経費精算・法人カード・$5",          "経費精算 SaaS, Expensify 料金 評判, 経費 管理 ツール",             "フィンテックSaaS"),
    # ── 法人カード/経費 ──
    ("Ramp Business Corp.",        "https://ramp.com/pricing",                                "リスティング広告", "", "Ramp・法人カード・経費自動化・無料",           "法人カード 経費 SaaS, Ramp 料金 評判, 経費 自動化 ツール",         "フィンテックSaaS"),
    # ── スタートアップ向け法人カード ──
    ("Brex Inc.",                  "https://www.brex.com/pricing",                            "リスティング広告", "", "Brex・法人カード・スタートアップ向け財務",     "スタートアップ 法人カード, Brex 料金 評判, 経費 管理 SaaS",        "フィンテックSaaS"),
    # ── 会計/請求 ──
    ("FreshBooks",                 "https://www.freshbooks.com/pricing",                      "リスティング広告", "", "FreshBooks・会計/請求・小規模事業者向け",      "請求書 会計 SaaS, FreshBooks 料金 評判, 小規模 会計 ソフト",       "会計SaaS"),
    # ── 請求書/支払い ──
    ("BILL (Bill.com)",            "https://www.bill.com/product/pricing",                    "リスティング広告", "", "BILL・請求書/支払い・AP/AR自動化",            "請求書 支払い SaaS, Bill.com 料金 評判, AP AR 自動化 ツール",      "フィンテックSaaS"),
    # ── クラウドERP/会計 ──
    ("Sage Intacct (Sage Group)",  "https://www.sage.com/en-us/sage-business-cloud/intacct/pricing/", "リスティング広告", "", "Sage Intacct・クラウドERP・財務管理", "クラウド ERP 会計, Sage Intacct 料金 評判, 財務 管理 SaaS",      "ERP SaaS"),
    # ── 株式管理 ──
    ("Carta (eShares Inc.)",       "https://carta.com/pricing",                               "リスティング広告", "", "Carta・キャップテーブル/株式管理・ファンド",   "株式 管理 SaaS, Carta 料金 評判, キャップテーブル ツール",         "フィンテックSaaS"),
    # ── アプリセキュリティ ──
    ("Snyk Ltd.",                  "https://go.snyk.io/jp-lp",                                "リスティング広告", "", "Snyk・開発者向けセキュリティ・脆弱性管理",     "アプリ セキュリティ SaaS, Snyk 料金 評判, 脆弱性 管理 ツール",     "セキュリティSaaS"),
    # ── DevOps ──
    ("JFrog Ltd.",                 "https://jfrog.com/pricing/",                              "リスティング広告", "", "JFrog・ソフトウェアサプライチェーン・成果物管理","DevOps SaaS, JFrog 料金 評判, アーティファクト 管理 ツール",      "DevOps SaaS"),
    # ── ノーコードアプリ ──
    ("Quickbase Inc.",             "https://www.quickbase.com/pricing",                       "リスティング広告", "", "Quickbase・ノーコード業務アプリ・複雑PM",      "ノーコード 業務 アプリ, Quickbase 料金 評判, 複雑 プロジェクト管理","ノーコードSaaS"),
    # ── ヘルプデスク ──
    ("HappyFox Inc.",              "https://www.happyfox.com/pricing",                        "リスティング広告", "", "HappyFox・ヘルプデスク・無制限エージェント",   "ヘルプデスク SaaS, HappyFox 料金 評判, カスタマーサポート ツール", "カスタマーサポートSaaS"),
    # ── リバースETL/CDP ──
    ("Hightouch Inc.",             "https://hightouch.com/pricing",                           "リスティング広告", "", "Hightouch・リバースETL・データ活用・CDP",     "リバース ETL SaaS, Hightouch 料金 評判, データ 連携 ツール",       "データSaaS"),
    # ── データパイプライン ──
    ("Fivetran Inc.",              "https://www.fivetran.com/pricing",                        "リスティング広告", "", "Fivetran・データパイプライン・ELT自動化",     "データ パイプライン SaaS, Fivetran 料金 評判, ELT データ 統合",    "データSaaS"),
    # ── データ変換 ──
    ("dbt Labs Inc.",             "https://www.getdbt.com/pricing",                          "リスティング広告", "", "dbt・データ変換・トランスフォーメーション",   "データ 変換 ツール, dbt 料金 評判, データ モデリング SaaS",        "データSaaS"),
    # ── データ/AIプラットフォーム ──
    ("Databricks Inc.",            "https://www.databricks.com/product/pricing",              "リスティング広告", "", "Databricks・レイクハウス・データ/AI基盤",     "データ レイクハウス SaaS, Databricks 料金 評判, AI データ 基盤",   "データSaaS"),
]


def make_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "架電リスト"

    for col_idx, header in enumerate(HEADER, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
    ws.row_dimensions[1].height = 20

    for row_idx, row_data in enumerate(DATA, 2):
        fill = ROW_FILLS[(row_idx - 2) % 2]
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.fill = fill
            cell.alignment = ROW_ALIGN
            if col_idx == 2:
                cell.hyperlink = value
                cell.font = URL_FONT
            else:
                cell.font = ROW_FONT
        ws.row_dimensions[row_idx].height = 16

    for col_idx, width in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADER))}1"

    wb.save(OUTPUT_PATH)
    print(f"Saved: {OUTPUT_PATH}")
    print(f"Rows: {len(DATA)}")


if __name__ == "__main__":
    make_xlsx()
