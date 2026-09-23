import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUTPUT_PATH = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー\架電リスト_Meta_2026-07-18_再検証1.xlsx"
HEADER = ["企業名", "LP URL", "広告種別", "法人番号", "商材", "検索KW", "業界"]
COL_WIDTHS = [36, 56, 16, 16, 22, 32, 22]
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="メイリオ", size=10, bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
ROW_FILLS = [PatternFill("solid", fgColor="F2F7FC"), PatternFill("solid", fgColor="FFFFFF")]
ROW_FONT = Font(name="メイリオ", size=9)
ROW_ALIGN = Alignment(vertical="center", wrap_text=False)
URL_FONT = Font(name="メイリオ", size=9, color="0563C1", underline="single")

# Meta広告ライブラリで個別に再検索し、実際に配信中の広告から確認したURLのみ記載。
DATA = [
    ("ホームテック株式会社", "https://entrie.net/chikufuru/", "Meta広告", "", "築古戸建てリノベーション(entrie/エントリエ)", "リノベーション 築古, 戸建てリノベ", "リノベーション"),
    ("ジューテックホーム株式会社", "https://jutec-home.jp/", "Meta広告", "", "注文住宅(高気密高断熱)", "注文住宅 東京, ゼロエミ住宅", "注文住宅"),
    ("株式会社アーキテックプランニング", "https://dualhome.jp/", "Meta広告", "", "注文住宅(DUALHOME)", "注文住宅 デザイン住宅, DUALHOME", "注文住宅FC"),
    ("株式会社アワハウス", "https://a-ourhouse.co.jp/", "Meta広告", "", "注文住宅(旧アートホーム/mini prot)", "注文住宅 北海道, ミニプロト", "注文住宅"),
    ("グッドリビング株式会社", "https://limini.dxbuilders.jp/yamashitaya/", "Meta広告", "", "ワンプライス注文住宅(TATTA!/Limini)", "注文住宅 定額, 平屋 注文住宅", "注文住宅"),
    ("マイコミュニケーション株式会社", "https://www.hoholine.com/", "Meta広告", "", "保険の相談・見直し(保険ほっとライン)", "保険相談, 保険見直し", "保険代理店"),
    ("株式会社プレアデス", "https://hoiku-job.net/", "Meta広告", "", "保育士・学童指導員向け転職(ほいくジョブ)", "保育士 転職, 学童 求人", "人材紹介"),
    ("株式会社モニクルフィナンシャル", "https://moneiro.jp/", "Meta広告", "", "資産運用相談・投資診断(マネイロ)", "資産運用 相談, 投資診断", "金融・資産運用"),
    ("株式会社メディカル・コンシェルジュ", "https://mc-nurse.net/", "Meta広告", "", "看護師単発バイト紹介(MCナースネット)", "看護師 単発バイト, 看護師 求人", "人材紹介"),
    ("Dr.AGAクリニック（医療法人社団日昇会）", "https://drskinclinic.jp/", "Meta広告", "", "AGA治療(はえループ)", "AGA治療, 発毛メソ治療", "AGA・美容クリニック"),
    ("ルシアクリニック（医療法人緑風会）", "https://lucia-c.com/", "Meta広告", "", "医療脱毛", "医療脱毛, 脱毛クリニック", "美容クリニック"),
    ("AND美容外科（一般社団法人AND medical group）", "https://and-bg.com/", "Meta広告", "", "美容医療・美容外科", "美容外科, 美容医療", "美容クリニック"),
]

# 未確認（広告ライブラリで検索したが該当広告が見つからなかった。要注意）
NOT_VERIFIED = [
    "株式会社ティー・エフ・オフィス（保険コンパス）",
    "株式会社エフエムエス中央クリニック（中央クリニック）",
]


def make_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "架電リスト"
    for col_idx, header in enumerate(HEADER, 1):
        c = ws.cell(row=1, column=col_idx, value=header)
        c.fill = HEADER_FILL; c.font = HEADER_FONT; c.alignment = HEADER_ALIGN
    ws.row_dimensions[1].height = 20
    for row_idx, row in enumerate(DATA, 2):
        fill = ROW_FILLS[(row_idx - 2) % 2]
        for col_idx, value in enumerate(row, 1):
            c = ws.cell(row=row_idx, column=col_idx, value=value)
            c.fill = fill; c.alignment = ROW_ALIGN
            if col_idx == 2:
                c.hyperlink = value; c.font = URL_FONT
            else:
                c.font = ROW_FONT
        ws.row_dimensions[row_idx].height = 16
    for col_idx, width in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADER))}1"

    ws2 = wb.create_sheet("未確認")
    ws2.cell(row=1, column=1, value="広告ライブラリで確認できなかった企業（要再調査 or 除外候補）").font = Font(bold=True)
    for i, name in enumerate(NOT_VERIFIED, 2):
        ws2.cell(row=i, column=1, value=name)
    ws2.column_dimensions["A"].width = 50

    wb.save(OUTPUT_PATH)
    print(f"Saved: {OUTPUT_PATH}  Rows: {len(DATA)}  Not verified: {len(NOT_VERIFIED)}")


if __name__ == "__main__":
    make_xlsx()
