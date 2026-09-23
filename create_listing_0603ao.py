import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUTPUT_PATH = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー\架電リスト_リスティング_2026-06-03ao.xlsx"
HEADER = ["企業名", "LP URL", "広告種別", "法人番号", "商材", "検索KW", "業界"]
COL_WIDTHS = [32, 56, 16, 16, 22, 32, 22]
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="メイリオ", size=10, bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
ROW_FILLS = [PatternFill("solid", fgColor="F2F7FC"), PatternFill("solid", fgColor="FFFFFF")]
ROW_FONT = Font(name="メイリオ", size=9)
ROW_ALIGN = Alignment(vertical="center", wrap_text=False)
URL_FONT = Font(name="メイリオ", size=9, color="0563C1", underline="single")

DATA = [
    # ── 交通事故弁護士（無料相談LP） ──
    ("弁護士法人アトム法律事務所",      "https://atomfirm.com/media/free",                          "リスティング広告", "", "交通事故・無料相談・完全成功報酬",      "交通事故 弁護士 無料相談, 慰謝料 増額 弁護士, 交通事故 示談 依頼",       "法律相談"),
    # ── DM発送代行（申し込みLP） ──
    ("ディーエムソリューションズ株式会社", "https://www.dm-s.co.jp/dm-lp/formonly/",               "リスティング広告", "", "DM印刷・発送代行・ダイレクトメール",    "DM 発送 代行 安い, ダイレクトメール 印刷 料金, DM 代行 会社 比較",     "マーケティング支援"),
    # ── オンライン診療（利用方法LP） ──
    ("株式会社OPTiM（ポケットドクター）", "https://www.pocketdoctor.jp/med/lp/howto/",             "リスティング広告", "", "オンライン診療・スマホで受診・処方箋",  "オンライン診療 アプリ おすすめ, スマホ 医師 診察 自宅, 処方箋 宅配",   "オンライン診療"),
    # ── CRM・SFA（30日無料トライアルLP） ──
    ("セールスフォース・ジャパン株式会社", "https://www.salesforce.com/jp/form/signup/freetrial-sales/", "リスティング広告", "", "Sales Cloud・CRM・30日間無料トライアル", "Salesforce 無料トライアル, CRM ツール おすすめ, 営業管理 SFA 比較",  "SaaS・CRM"),
    # ── ECサイト構築（無料開設申込LP） ──
    ("GMOペパボ株式会社（カラーミーショップ）", "https://shop-pro.jp/?mode=signup",                "リスティング広告", "", "ECサイト構築・無料プランあり・ネットショップ", "ネットショップ 作り方 無料, EC サイト 開設 簡単, カラーミー 評判 料金", "ECサイト構築"),
]


def make_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "架電リスト"

    for col_idx, header in enumerate(HEADER, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
    ws.row_dimensions[1].height = 20

    for row_idx, row_data in enumerate(DATA, 2):
        fill = ROW_FILLS[(row_idx - 2) % 2]
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.fill = fill
            cell.alignment = ROW_ALIGN
            if col_idx == 2:
                cell.hyperlink = value
                cell.font = URL_FONT
            else:
                cell.font = ROW_FONT
        ws.row_dimensions[row_idx].height = 16

    for col_idx, width in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADER))}1"

    wb.save(OUTPUT_PATH)
    print(f"Saved: {OUTPUT_PATH}")
    print(f"Rows: {len(DATA)}")


if __name__ == "__main__":
    make_xlsx()
