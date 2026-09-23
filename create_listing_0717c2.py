import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUTPUT_PATH = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー\架電リスト_リスティング_2026-07-17c.xlsx"
HEADER = ["企業名", "LP URL", "広告種別", "法人番号", "商材", "検索KW", "業界"]
COL_WIDTHS = [32, 56, 16, 16, 22, 32, 22]
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="メイリオ", size=10, bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
ROW_FILLS = [PatternFill("solid", fgColor="F2F7FC"), PatternFill("solid", fgColor="FFFFFF")]
ROW_FONT = Font(name="メイリオ", size=9)
ROW_ALIGN = Alignment(vertical="center", wrap_text=False)
URL_FONT = Font(name="メイリオ", size=9, color="0563C1", underline="single")

# Google検索で実際に「スポンサー」表示された広告のみ採用（比較サイト・ポータルは除外）。
DATA = [
    ("株式会社LIXIL TEPCOスマートパートナーズ", "https://www.lixiltepco-sp.co.jp/tatetoku_reform/index.html", "リスティング広告", "", "太陽光発電・断熱リフォーム(建て得)", "太陽光 実質0円, 断熱リフォーム", "太陽光・リフォーム"),
    ("株式会社ガイアフィールド", "https://akiya.steel-house.co.jp/", "リスティング広告", "", "空き家管理サービス", "空き家管理, 空き家 定期巡回", "空き家管理・不動産"),
    ("株式会社上池解体興業", "https://kamiike-kaitai.com/", "リスティング広告", "", "解体工事(ボッコス)", "解体工事 業者, 家 解体", "解体工事"),
]


def make_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "架電リスト"
    for col_idx, header in enumerate(HEADER, 1):
        c = ws.cell(row=1, column=col_idx, value=header)
        c.fill = HEADER_FILL; c.font = HEADER_FONT; c.alignment = HEADER_ALIGN
    ws.row_dimensions[1].height = 20
    for row_idx, row in enumerate(DATA, 2):
        fill = ROW_FILLS[(row_idx - 2) % 2]
        for col_idx, value in enumerate(row, 1):
            c = ws.cell(row=row_idx, column=col_idx, value=value)
            c.fill = fill; c.alignment = ROW_ALIGN
            if col_idx == 2:
                c.hyperlink = value; c.font = URL_FONT
            else:
                c.font = ROW_FONT
        ws.row_dimensions[row_idx].height = 16
    for col_idx, width in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADER))}1"
    wb.save(OUTPUT_PATH)
    print(f"Saved: {OUTPUT_PATH}  Rows: {len(DATA)}")


if __name__ == "__main__":
    make_xlsx()
