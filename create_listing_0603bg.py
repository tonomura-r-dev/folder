import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUTPUT_PATH = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー\架電リスト_リスティング_2026-06-03bg.xlsx"
HEADER = ["企業名", "LP URL", "広告種別", "法人番号", "商材", "検索KW", "業界"]
COL_WIDTHS = [32, 56, 16, 16, 22, 32, 22]
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="メイリオ", size=10, bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
ROW_FILLS = [PatternFill("solid", fgColor="F2F7FC"), PatternFill("solid", fgColor="FFFFFF")]
ROW_FONT = Font(name="メイリオ", size=9)
ROW_ALIGN = Alignment(vertical="center", wrap_text=False)
URL_FONT = Font(name="メイリオ", size=9, color="0563C1", underline="single")

DATA = [
    # ── 1on1支援ツール ──
    ("株式会社KAKEAI",              "https://kakeai.co.jp/lp/ad-base-lp/",                  "リスティング広告", "", "1on1支援・面談ツール・マネジメント改革",       "1on1 ツール 比較, マネジメント 面談 SaaS, KAKEAI 料金 評判",        "HR SaaS"),
    # ── データ活用コンサル ──
    ("株式会社ブレインパッド",      "https://www.brainpad.co.jp/services/professionals/",    "リスティング広告", "", "AI・データ分析・DX推進・1,400社支援",           "データ 分析 コンサル AI, DX 推進 支援 費用, ブレインパッド 評判",    "AIコンサル"),
    # ── バックオフィスSaaS ──
    ("Bizer株式会社",               "https://bizer.jp/team/lp/",                             "リスティング広告", "", "チームタスク管理・業務可視化・バックオフィス",  "タスク 管理 SaaS 比較, バックオフィス 業務 効率化, Bizer 評判",     "業務DXSaaS"),
    # ── 請求書作成 ──
    ("弥生株式会社",                "https://www.misoca.jp/lp/smartphone/index.html",        "リスティング広告", "", "Misoca・請求書作成・スマホアプリ",              "請求書 作成 アプリ 無料, Misoca 評判 料金, 見積書 納品書 クラウド", "バックオフィスSaaS"),
    # ── IoT在庫管理 ──
    ("株式会社スマートショッピング","https://www.smartmat.io/ai-lp",                          "リスティング広告", "", "スマートマット・AI在庫管理・置くだけDX",        "在庫管理 IoT ツール, スマートマット 料金 評判, 在庫 自動発注",     "IoT SaaS"),
    # ── ゲーム配信広告 ──
    ("株式会社ミラティブ",          "https://www.mirrativ.co.jp/product/ad/",                 "リスティング広告", "", "ミラティブ広告・ゲーム配信・スマホゲーム集客",  "ゲーム 広告 配信 媒体, スマホ ゲーム マーケ, ミラティブ 広告 料金","ゲームマーケ"),
    # ── eKYC・本人確認 ──
    ("株式会社TRUSTDOCK",           "https://biz.trustdock.io/lp-ekyc-01v1",                 "リスティング広告", "", "eKYC・デジタル本人確認・導入社数No.1",         "eKYC 本人確認 SaaS, マイナンバー 確認 API, TRUSTDOCK 料金 評判",  "IDtech SaaS"),
    # ── エンジニア情報共有 ──
    ("Qiita株式会社",               "https://teams.qiita.com/pricing/",                      "リスティング広告", "", "Qiita Team・社内wiki・エンジニア情報共有",      "社内 wiki エンジニア SaaS, Qiita Team 料金 評判, ナレッジ 共有",   "ナレッジ管理SaaS"),
    # ── 理系学生採用 ──
    ("株式会社LabBase",             "https://business.labbase.jp/",                          "リスティング広告", "", "理系学生採用・研究室スカウト・11万人DB",        "理系 採用 サービス 比較, 院生 エンジニア 採用, LabBase 料金 評判", "採用SaaS"),
    # ── チャットコマース ──
    ("株式会社ZEALS",               "https://chatcommerce.zeals.co.jp/en/",                  "リスティング広告", "", "チャットコマース・LINE広告・AI会話CV改善",      "チャットコマース LINE 広告, ZEALS 評判 料金, チャットボット EC",   "マーケSaaS"),
    # ── HR人材業界向けDB ──
    ("株式会社フロッグ",            "https://list.hrog.net/",                                 "リスティング広告", "", "HRogリスト・求人企業DB・人材業界営業リスト",    "人材業界 営業 リスト, HR 企業 DB SaaS, HRog リスト 料金",         "セールスSaaS"),
    # ── 位置情報AI分析 ──
    ("レイ・フロンティア株式会社",  "https://www.rei-frontier.jp/request/",                  "リスティング広告", "", "位置情報AI分析・行動データ・SilentLog",         "位置情報 分析 ツール, 行動 データ プラットフォーム, レイフロ 料金", "Location SaaS"),
    # ── 広告・PRコンサル ──
    ("株式会社ベクトル",            "https://vectorinc.co.jp/pr-menu",                       "リスティング広告", "", "PR・デジタルマーケ・広報代行",                  "PR 会社 費用 比較, 広報 代行 おすすめ, ベクトル PR 評判",         "PRコンサルティング"),
    # ── コマースメディア広告 ──
    ("Criteo株式会社",              "https://www.criteo.com/business/advertisers/",           "リスティング広告", "", "コマースメディア・リターゲティング広告",         "リターゲティング SaaS, Criteo 評判 料金, コマース 広告 比較",      "AdTech"),
    # ── セールスイネーブルメント ──
    ("株式会社ナレッジワーク",      "https://knowledgework.cloud/",                           "リスティング広告", "", "セールスイネーブルメントAI・営業力強化",         "セールスイネーブルメント 比較, 営業 SaaS AI, ナレッジワーク 料金", "セールスSaaS"),
    # ── 採用管理SaaS ──
    ("株式会社HERP",                "https://lp.herp.cloud/",                                "リスティング広告", "", "採用管理ATS・デジタル人材・HR SaaS",            "採用管理 システム ATS, HERP 料金 評判, 採用 DX ツール",           "HR SaaS"),
    # ── 物流DX ──
    ("ハコベル株式会社",            "https://lp.hacobell.com/",                              "リスティング広告", "", "物流DX・配送マッチング・即日手配",              "物流 DX SaaS, ハコベル 評判 料金, 配送 マッチング",               "物流SaaS"),
    # ── ABM・BtoB分析 ──
    ("株式会社ユーザベース",        "https://www.lp.forcas.com/",                             "リスティング広告", "", "FORCAS・ABM・BtoB顧客分析",                    "ABM ツール 比較, BtoB マーケ ターゲット 分析, FORCAS 料金",       "BtoBマーケSaaS"),
    # ── HR評価SaaS ──
    ("株式会社HRBrain",             "https://www.hrbrain.jp/lp/dxplaness",                   "リスティング広告", "", "タレントマネジメント・人事評価・顧客満足度No.1", "HRBrain 料金 評判, 人事評価 DX, タレントマネジメント SaaS",      "HR SaaS"),
    # ── インテントデータ ──
    ("株式会社Sales Marker",        "https://sales-marker.jp/lp/demo/",                      "リスティング広告", "", "インテントセールス・BtoB営業・商談獲得",         "インテント データ 営業, BtoB 商談 ツール, Sales Marker 料金",     "セールスSaaS"),
]


def make_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "架電リスト"

    for col_idx, header in enumerate(HEADER, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
    ws.row_dimensions[1].height = 20

    for row_idx, row_data in enumerate(DATA, 2):
        fill = ROW_FILLS[(row_idx - 2) % 2]
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.fill = fill
            cell.alignment = ROW_ALIGN
            if col_idx == 2:
                cell.hyperlink = value
                cell.font = URL_FONT
            else:
                cell.font = ROW_FONT
        ws.row_dimensions[row_idx].height = 16

    for col_idx, width in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADER))}1"

    wb.save(OUTPUT_PATH)
    print(f"Saved: {OUTPUT_PATH}")
    print(f"Rows: {len(DATA)}")


if __name__ == "__main__":
    make_xlsx()
