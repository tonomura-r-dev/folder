import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def make_xlsx(filename, data):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "架電リスト"

    headers = ["企業名", "LP URL", "電話番号", "商材", "検索KW", "業界"]

    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="メイリオ", bold=True, color="FFFFFF", size=10)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center
        cell.border = border

    for row_idx, row in enumerate(data, 2):
        fill_color = "F2F7FC" if row_idx % 2 == 0 else "FFFFFF"
        row_fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
        for col_idx, val in enumerate(row, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.fill = row_fill
            cell.font = Font(name="メイリオ", size=9)
            cell.alignment = left
            cell.border = border

    col_widths = [30, 45, 16, 20, 30, 18]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.row_dimensions[1].height = 20
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    wb.save(filename)
    print(f"保存完了: {filename}  件数: {len(data)}")


# 列順: 企業名 | LP URL | 電話番号 | 商材 | 検索KW | 業界
# 不動産投資／人材紹介／通販D2C／ブランド買取の4カテゴリをMeta広告ライブラリでキーワード検索し、
# 実配信中の広告を確認できた企業のみ採用（推測なし）。使用済み企業リストとGrep照合し重複なしを確認済み。

data_0717f = [
    ["株式会社ファクター・ナイン", "https://9-fund.com/", "", "不動産クラウドファンディング（NINE FUND）", "不動産クラファン／少額 不動産投資", "不動産投資クラファン"],
    ["小田急不動産株式会社", "https://ls1.odakyu-chukai.com/", "03-3370-1110", "収益物件情報・AI査定サービス（LIFE SCAPE 1）", "収益物件 投資／不動産 AI査定", "不動産投資"],
    ["ST Agency Japan株式会社", "https://www.stagency.me/", "03-6824-5199", "外国人材紹介・特定技能人材紹介", "外国人採用／特定技能 人材紹介", "外国人材紹介"],
    ["株式会社High Link", "https://coloria.jp/", "", "香水サブスクリプション（カラリア）", "香水 サブスク／香水 定期便", "香水D2C"],
    ["UIJオークション株式会社", "https://www.brandadorer.com/watch/", "", "高級時計・ブランド品買取（ブランドアドレ）", "時計 買取／ブランド品 買取", "ブランド品買取"],
    ["株式会社ユーズカンパニー", "https://www.housekihiroba-kaitori.jp/shop/shibuya/", "03-5458-7224", "高級時計・ジュエリー買取（宝石広場）", "時計 買取 渋谷／ジュエリー 買取", "時計・ジュエリー買取"],
]

base = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー"
make_xlsx(f"{base}\\架電リスト_Meta_2026-07-17f.xlsx", data_0717f)
