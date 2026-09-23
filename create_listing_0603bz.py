import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUTPUT_PATH = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー\架電リスト_リスティング_2026-06-03bz.xlsx"
HEADER = ["企業名", "LP URL", "広告種別", "法人番号", "商材", "検索KW", "業界"]
COL_WIDTHS = [32, 56, 16, 16, 22, 32, 22]
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="メイリオ", size=10, bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
ROW_FILLS = [PatternFill("solid", fgColor="F2F7FC"), PatternFill("solid", fgColor="FFFFFF")]
ROW_FONT = Font(name="メイリオ", size=9)
ROW_ALIGN = Alignment(vertical="center", wrap_text=False)
URL_FONT = Font(name="メイリオ", size=9, color="0563C1", underline="single")

DATA = [
    # ── 電話代行（日本・実LP確認） ──
    ("株式会社うるる",             "https://www.fondesk.jp/",                                "リスティング広告", "", "fondesk・電話代行・月1万円から",              "電話代行 サービス, fondesk 料金 評判, 電話 代行 安い",           "BPOサービス"),
    # ── 電話代行/オンライン秘書（日本・実LP確認） ──
    ("株式会社ベルシステム24",     "https://www.tas.bell24.co.jp/",                          "リスティング広告", "", "e秘書・電話代行/秘書代行・士業向け",          "電話代行 e秘書, 秘書代行 サービス, ベルシステム24 料金",         "BPOサービス"),
    # ── 軽貨物ドライバー募集（日本・実LP確認） ──
    ("株式会社ロジクエスト",       "https://logiquest.co.jp/partners/cargo/",                "リスティング広告", "", "軽貨物ドライバー業務委託募集・低リスク",       "軽貨物 ドライバー 募集, 業務委託 配送, ロジクエスト 評判",       "物流・求人"),
    # ── 軽貨物ドライバー募集（日本・実LP確認） ──
    ("株式会社貴順",               "https://butsuryu.asia/recruit/driver/",                  "リスティング広告", "", "委託ドライバー募集・軽貨物運送",              "委託 ドライバー 募集, 軽貨物 求人 東京, 配送 業務委託",         "物流・求人"),
    # ── フォーム作成 ──
    ("Tally",                      "https://tally.so/pricing",                               "リスティング広告", "", "Tally・無料フォーム作成・無制限回答",          "フォーム 作成 無料, Tally 料金 評判, アンケート フォーム SaaS",  "フォームSaaS"),
    # ── フォーム作成 ──
    ("Fillout",                    "https://www.fillout.com/pricing",                        "リスティング広告", "", "Fillout・高機能フォーム作成・無料枠",          "フォーム 作成 SaaS, Fillout 料金 評判, オンライン フォーム",     "フォームSaaS"),
    # ── フォーム作成 ──
    ("Paperform",                  "https://paperform.co/pricing/",                          "リスティング広告", "", "Paperform・LP型フォーム作成・決済対応",        "フォーム 作成 ツール, Paperform 料金 評判, 申込 フォーム SaaS",  "フォームSaaS"),
    # ── フォーム/文書/署名 ──
    ("Formstack",                  "https://www.formstack.com/pricing",                      "リスティング広告", "", "Formstack・フォーム/文書/電子署名",            "フォーム SaaS 比較, Formstack 料金 評判, 業務 フォーム 自動化",  "フォームSaaS"),
    # ── 提案書/見積 ──
    ("Qwilr",                      "https://qwilr.com/pricing/",                             "リスティング広告", "", "Qwilr・Webページ型提案書/見積・決済",          "提案書 作成 SaaS, Qwilr 料金 評判, 見積 提案 ツール",            "提案書SaaS"),
    # ── 提案書 ──
    ("Proposify",                  "https://www.proposify.com/pricing",                      "リスティング広告", "", "Proposify・提案書作成/管理・ブランド",         "提案書 作成 ツール, Proposify 料金 評判, 営業 提案 SaaS",        "提案書SaaS"),
    # ── プレゼン ──
    ("Pitch",                      "https://pitch.com/pricing/us",                           "リスティング広告", "", "Pitch・共同プレゼン作成・AI",                 "プレゼン 作成 SaaS, Pitch 料金 評判, スライド 共同 作成",        "デザインSaaS"),
    # ── プレゼン ──
    ("Prezi",                      "https://prezi.com/pricing/",                             "リスティング広告", "", "Prezi・ズーム型プレゼン作成・AI",             "プレゼン ツール, Prezi 料金 評判, プレゼンテーション 作成 SaaS", "デザインSaaS"),
    # ── インフォグラフィック ──
    ("Piktochart",                 "https://piktochart.com/pricing/",                        "リスティング広告", "", "Piktochart・インフォグラフィック/資料作成",     "インフォグラフィック 作成, Piktochart 料金 評判, 資料 デザイン", "デザインSaaS"),
    # ── インフォグラフィック ──
    ("Venngage",                   "https://venngage.com/pricing",                           "リスティング広告", "", "Venngage・インフォグラフィック/資料作成・AI",   "インフォグラフィック SaaS, Venngage 料金 評判, 資料 作成 ツール", "デザインSaaS"),
    # ── インタラクティブ ──
    ("Genially",                   "https://genially.com/plans/",                            "リスティング広告", "", "Genially・インタラクティブコンテンツ作成",     "インタラクティブ 資料 SaaS, Genially 料金 評判, 動く 資料 作成",  "デザインSaaS"),
    # ── ブランドデザイン ──
    ("Marq",                       "https://www.marq.com/pages/pricing/",                    "リスティング広告", "", "Marq・ブランド統制デザイン・テンプレ運用",     "ブランド デザイン SaaS, Marq 料金 評判, テンプレート デザイン",  "デザインSaaS"),
    # ── 画像デザイン ──
    ("Snappa",                     "https://snappa.com/pricing",                             "リスティング広告", "", "Snappa・かんたん画像デザイン・SNS素材",         "画像 デザイン ツール, Snappa 料金 評判, SNS 画像 作成 SaaS",     "デザインSaaS"),
    # ── 動画ホスティング ──
    ("Wistia",                     "https://wistia.com/pricing",                             "リスティング広告", "", "Wistia・ビジネス動画ホスティング/分析",        "動画 ホスティング 法人, Wistia 料金 評判, 動画 マーケ SaaS",     "動画SaaS"),
    # ── 動画ホスティング ──
    ("Vimeo",                      "https://vimeo.com/upgrade-plan",                         "リスティング広告", "", "Vimeo・動画ホスティング/配信・法人",           "動画 配信 ホスティング, Vimeo 料金 評判, 動画 アップロード SaaS","動画SaaS"),
    # ── 画面録画動画 ──
    ("Tella",                      "https://www.tella.tv/help/introduction/plans",           "リスティング広告", "", "Tella・画面録画/動画作成・オールインワン",      "画面 録画 動画 SaaS, Tella 料金 評判, 動画 作成 ツール",         "動画SaaS"),
]


def make_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "架電リスト"

    for col_idx, header in enumerate(HEADER, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
    ws.row_dimensions[1].height = 20

    for row_idx, row_data in enumerate(DATA, 2):
        fill = ROW_FILLS[(row_idx - 2) % 2]
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.fill = fill
            cell.alignment = ROW_ALIGN
            if col_idx == 2:
                cell.hyperlink = value
                cell.font = URL_FONT
            else:
                cell.font = ROW_FONT
        ws.row_dimensions[row_idx].height = 16

    for col_idx, width in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADER))}1"

    wb.save(OUTPUT_PATH)
    print(f"Saved: {OUTPUT_PATH}")
    print(f"Rows: {len(DATA)}")


if __name__ == "__main__":
    make_xlsx()
