import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUTPUT_PATH = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー\架電リスト_リスティング_2026-05-27c.xlsx"

HEADER = ["企業名", "LP URL", "広告種別", "電話番号", "商材", "検索KW", "業界"]
COL_WIDTHS = [32, 50, 16, 16, 20, 32, 22]

DATA = [
    (
        "株式会社トライトキャリア（介護ワーカー）",
        "https://kaigoworker.jp/lp/hellowork-step/",
        "リスティング広告", "03-5436-7670",
        "介護職専門転職エージェント（全国対応・非公開求人多数）",
        "介護 転職 エージェント 求人, 介護職 仕事 高収入 施設, 介護職員 転職 支援 無料",
        "転職・就職・人材派遣"
    ),
    (
        "株式会社ZEN PLACE（ZEN PLACE）",
        "https://www.zenplace.co.jp/lp/pilates/01",
        "リスティング広告", "03-6409-6500",
        "ピラティス・ヨガ専門スタジオ（全国展開・体験レッスン無料）",
        "ピラティス スタジオ 体験 近く, マシンピラティス 予約 初心者, ヨガ ピラティス 教室 通い放題",
        "フィットネス・スポーツ・ゴルフ"
    ),
    (
        "株式会社プランドゥ（マイ介護ホーム）",
        "https://www.my-kaigo-home.com/lp/soudan/input/",
        "リスティング広告", "",
        "老人ホーム・介護施設の無料入居相談サービス",
        "老人ホーム 入居 無料相談 専門家, 介護施設 選び方 比較 サポート, 特養 有料老人ホーム 費用 見学",
        "医療・健康・介護"
    ),
    (
        "株式会社エスタック",
        "https://estac.co.jp/lp/",
        "リスティング広告", "",
        "マンション投資・不動産投資コンサルティング",
        "不動産投資 マンション 始め方 少額, 投資用マンション 利回り 選び方, 不動産 資産運用 セミナー 無料",
        "不動産・賃貸・リノベーション・投資"
    ),
    (
        "ニッカホーム株式会社",
        "https://www.nikka-home.co.jp/lp/chirashi/chubu/product/mizumawari.html",
        "リスティング広告", "0120-87-7549",
        "水廻りリフォーム（キッチン・風呂・洗面所）自社職人施工",
        "水回り リフォーム 見積もり 費用 相場, キッチン 交換 リフォーム 業者 近く, 浴室 リフォーム 費用 比較",
        "不動産・賃貸・リノベーション・投資"
    ),
    (
        "株式会社銀座山形屋",
        "https://ginzayamagataya.tanmono.com/lp/ginza/",
        "リスティング広告", "03-5297-5181",
        "メンズオーダースーツ・仕立て（創業117年老舗テーラー・銀座本店）",
        "オーダースーツ 銀座 老舗 仕立て, メンズスーツ オーダー 高品質 費用, スーツ 仕立て直し テーラー 東京",
        "ファッション・アパレル"
    ),
    (
        "株式会社カーネクスト",
        "https://carnext.jp/lp/rpg/campaign/",
        "リスティング広告", "0120-301-456",
        "廃車・事故車・不動車の無料引取・買取サービス（年中無休）",
        "廃車 買取 無料 引取り 全国, 事故車 処分 お金 もらえる, 不動車 売る 無料 手続き",
        "車・カーシェア・バイク"
    ),
    (
        "株式会社サクシード（ほいくポータル）",
        "https://www.succeed-jinzai.jp/lp/hoiku/portal/",
        "リスティング広告", "03-5287-7259",
        "保育士・幼稚園教諭専門の無料転職サポートサービス",
        "保育士 転職 エージェント 専門, 幼稚園 教諭 転職 求人, 保育士 新しい職場 探し 無料",
        "転職・就職・人材派遣"
    ),
    (
        "株式会社オートランドリータカノ（タカノの宅配クリーニング）",
        "https://cleaning-takano.com/lp/1",
        "リスティング広告", "0120-79-9029",
        "宅配クリーニングサービス（衣類・布団・コート）全国配送",
        "宅配クリーニング 送料無料 全国, 布団 クリーニング 宅配 安い, コート クリーニング 宅配 料金",
        "家事代行・クリーニング"
    ),
    (
        "パスクリエイト株式会社（税理士紹介エージェント）",
        "https://www.zeirishi-shoukaicenter.com/lp/top",
        "リスティング広告", "03-6380-1145",
        "税理士無料紹介・マッチングサービス（中小企業・個人事業主向け）",
        "税理士 探し方 無料 紹介 比較, 税理士 変更 手順 費用 相談, 確定申告 税理士 依頼 費用 個人",
        "法律・税務・相談"
    ),
]

def make_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "架電リスト"

    header_font   = Font(name="メイリオ", size=10, bold=True, color="FFFFFF")
    header_fill   = PatternFill("solid", fgColor="1F4E79")
    header_align  = Alignment(horizontal="center", vertical="center", wrap_text=True)

    fill_odd  = PatternFill("solid", fgColor="F2F7FC")
    fill_even = PatternFill("solid", fgColor="FFFFFF")
    data_font = Font(name="メイリオ", size=9)
    data_align = Alignment(vertical="center", wrap_text=False)
    url_font  = Font(name="メイリオ", size=9, color="0563C1", underline="single")

    thin = Side(style="thin", color="D0D7DE")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col, name in enumerate(HEADER, start=1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = border

    for row_idx, row in enumerate(DATA, start=2):
        fill = fill_odd if row_idx % 2 == 0 else fill_even
        for col, val in enumerate(row, start=1):
            cell = ws.cell(row=row_idx, column=col, value=val)
            cell.fill = fill
            cell.border = border
            cell.alignment = data_align
            if col == 2 and val:
                cell.hyperlink = val
                cell.font = url_font
            else:
                cell.font = data_font

    for col, width in enumerate(COL_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width

    ws.row_dimensions[1].height = 22
    for r in range(2, len(DATA) + 2):
        ws.row_dimensions[r].height = 18

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADER))}1"

    wb.save(OUTPUT_PATH)
    print(f"保存完了: {OUTPUT_PATH}")
    print(f"件数: {len(DATA)}社")

if __name__ == "__main__":
    make_xlsx()
