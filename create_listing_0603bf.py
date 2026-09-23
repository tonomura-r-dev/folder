import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUTPUT_PATH = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー\架電リスト_リスティング_2026-06-03bf.xlsx"
HEADER = ["企業名", "LP URL", "広告種別", "法人番号", "商材", "検索KW", "業界"]
COL_WIDTHS = [32, 56, 16, 16, 22, 32, 22]
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="メイリオ", size=10, bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
ROW_FILLS = [PatternFill("solid", fgColor="F2F7FC"), PatternFill("solid", fgColor="FFFFFF")]
ROW_FONT = Font(name="メイリオ", size=9)
ROW_ALIGN = Alignment(vertical="center", wrap_text=False)
URL_FONT = Font(name="メイリオ", size=9, color="0563C1", underline="single")

DATA = [
    # ── 法人カード（フィンテック） ──
    ("株式会社UPSIDER",              "https://up-sider.com/",                                  "リスティング広告", "", "法人カード・限度額最大10億・経費管理",       "法人カード 限度額 高い, スタートアップ 法人カード 申し込み, UPSIDER 評判", "フィンテック"),
    # ── ABMプラットフォーム ──
    ("株式会社ユーザベース",         "https://www.lp.forcas.com/",                             "リスティング広告", "", "FORCAS・ABM・BtoB顧客分析・ターゲティング", "ABM ツール 比較, BtoB マーケ ターゲティング SaaS, FORCAS 料金 評判", "BtoBマーケSaaS"),
    # ── HR評価SaaS ──
    ("株式会社HRBrain",              "https://www.hrbrain.jp/lp/dxplaness",                    "リスティング広告", "", "タレントマネジメント・人事評価・顧客満足度No.1", "タレントマネジメント SaaS 比較, HRBrain 料金 評判, 人事評価 DX ツール", "HR SaaS"),
    # ── 地方副業マッチング ──
    ("株式会社JOINS",               "https://lp.joins.co.jp/",                                "リスティング広告", "", "副業プロ人材・地方中小企業・人材シェア",      "副業 専門家 活用, 中小企業 プロ人材 外部 活用, JOINS 評判 料金",     "人材マッチング"),
    # ── 広告計測ツール ──
    ("株式会社イルグルム",           "https://go.ebis.ne.jp/lp/",                              "リスティング広告", "", "AD EBiS・広告効果測定・コンバージョン分析",   "広告 効果 測定 ツール 比較, アドエビス 料金 評判, CV 計測 プラットフォーム", "AdTech"),
    # ── コマースメディア広告 ──
    ("Criteo株式会社",               "https://www.criteo.com/business/advertisers/",            "リスティング広告", "", "コマースメディア・リターゲティング・ダイナミック広告", "リターゲティング 広告 比較, コマースメディア プラットフォーム, Criteo 評判", "AdTech"),
    # ── SEOツール ──
    ("Ahrefs Pte.Ltd.",              "https://ahrefs.com/ja/signup?interval=monthly&plan=trial", "リスティング広告", "", "SEOツール・AI検索時代・バックリンク分析",    "SEO ツール 比較 おすすめ, Ahrefs 料金 評判, バックリンク 分析 SaaS", "SEO SaaS"),
    # ── プロダクト分析SaaS ──
    ("Pendo.io株式会社",             "https://jp.pendo.io/pendo-free/",                        "リスティング広告", "", "プロダクト分析・インアップガイド・無料登録",  "プロダクト 分析 ツール 比較, Pendo 料金 評判, アプリ UX 改善 SaaS",   "プロダクトSaaS"),
    # ── D2C広告支援 ──
    ("株式会社売れるネット広告社",   "https://lp.ureru.co.jp/consul02",                        "リスティング広告", "", "D2C通販広告・LP改善・無料コンサルティング",   "D2C 通販 広告 改善, LP CVR 改善 支援, 売れるネット広告社 評判",       "広告代理店"),
    # ── 法人向けモバイル ──
    ("コネクシオ株式会社",           "https://sol.conexio.co.jp/lp_brand/",                    "リスティング広告", "", "法人携帯・MDM・スマートフォン一括管理",       "法人 携帯 導入 管理, MDM ツール 比較, コネクシオ 法人 サービス",     "ICTソリューション"),
    # ── CXM・SNS管理SaaS ──
    ("Sprinklr Japan株式会社",       "https://www.sprinklr.com/jp/lp/modern-care-platform/",   "リスティング広告", "", "SNS管理・カスタマーケア・CXMプラットフォーム", "SNS 管理 ツール 法人 比較, カスタマーケア SaaS, Sprinklr 料金 評判",  "CXMSaaS"),
    # ── PR会社 ──
    ("株式会社ベクトル",             "https://vectorinc.co.jp/pr-menu",                        "リスティング広告", "", "PR・広報・デジタルマーケティング支援",         "PR 会社 おすすめ, 広報 マーケティング 代行, ベクトル 評判 料金",      "PRコンサルティング"),
    # ── ゼロトラスト ──
    ("Cloudflare Japan株式会社",     "https://www.cloudflare.com/ja-jp/plans/zero-trust-services/", "リスティング広告", "", "ゼロトラスト・SASE・セキュリティ無料プラン", "ゼロトラスト セキュリティ SaaS, SASE ツール 比較, Cloudflare 料金",  "セキュリティSaaS"),
    # ── 経営管理 ──
    ("株式会社ログラス",             "https://lp.loglass.jp/3minutes_loglass",                 "リスティング広告", "", "経営管理クラウド・予実管理・シェアNo.1",       "予実 管理 SaaS 比較, 経営 管理 クラウド, ログラス 料金 評判",          "経営管理SaaS"),
    # ── タレントマネジメント ──
    ("株式会社カオナビ",             "https://www.kaonavi.jp/",                                "リスティング広告", "", "タレントマネジメント・人材育成・シェア8年No.1", "カオナビ 料金 評判, タレント マネジメント 比較, 人事評価 クラウド",    "HR SaaS"),
    # ── 工数管理 ──
    ("株式会社クラウドワークス",     "https://lp.crowdlog.jp/to-c",                            "リスティング広告", "", "工数管理・プロジェクト管理・個人向け",         "工数 管理 クラウド フリーランス, プロジェクト管理 個人 安い, クラウドログ", "プロジェクト管理SaaS"),
    # ── 採用管理 ──
    ("株式会社Thinkings",           "https://sonar-ats.jp/lp/cm-2022-23/",                    "リスティング広告", "", "採用管理システム・ATS・LINEフル連携",          "採用管理 ATS 比較 料金, ソナーATS 評判, 採用 DX ツール おすすめ",     "HR SaaS"),
    # ── SNS運用代行 ──
    ("株式会社コムニコ",             "https://www.comnico.jp/service-doc-form",                "リスティング広告", "", "SNS運用代行・コンサルティング・3,000社実績",   "SNS 運用代行 費用 比較, Instagram 代行 おすすめ, コムニコ 評判",      "SNSマーケ"),
    # ── 物流DX ──
    ("ハコベル株式会社",             "https://lp.hacobell.com/",                               "リスティング広告", "", "物流DX・配送マッチング・ハコベル",             "物流 DX サービス, 配送 マッチング プラットフォーム, ハコベル 評判",    "物流SaaS"),
    # ── ABM・インテントデータ ──
    ("株式会社Sales Marker",         "https://sales-marker.jp/lp/demo/",                       "リスティング広告", "", "インテントセールス・BtoB営業DX・デモ申込",     "インテントデータ 営業 SaaS, BtoB 商談 獲得 ツール, セールスマーカー", "セールスSaaS"),
]


def make_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "架電リスト"

    for col_idx, header in enumerate(HEADER, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
    ws.row_dimensions[1].height = 20

    for row_idx, row_data in enumerate(DATA, 2):
        fill = ROW_FILLS[(row_idx - 2) % 2]
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.fill = fill
            cell.alignment = ROW_ALIGN
            if col_idx == 2:
                cell.hyperlink = value
                cell.font = URL_FONT
            else:
                cell.font = ROW_FONT
        ws.row_dimensions[row_idx].height = 16

    for col_idx, width in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADER))}1"

    wb.save(OUTPUT_PATH)
    print(f"Saved: {OUTPUT_PATH}")
    print(f"Rows: {len(DATA)}")


if __name__ == "__main__":
    make_xlsx()
