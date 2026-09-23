import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUTPUT_PATH = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー\架電リスト_リスティング_2026-05-28a.xlsx"
HEADER = ["企業名", "LP URL", "広告種別", "電話番号", "商材", "検索KW", "業界"]
COL_WIDTHS = [32, 50, 16, 16, 20, 32, 22]

DATA = [
    [
        "株式会社LAVAインターナショナル（ホットヨガLAVA）",
        "https://yoga-lava.com/lp/cp/sp13_kireimo/",
        "リスティング広告",
        "0120-64-9766",
        "ホットヨガ体験・月会員（全国280店舗以上）",
        "ホットヨガ LAVA ホットヨガ体験 ヨガ教室 ホットヨガスタジオ 体験レッスン",
        "フィットネス・スポーツ・ゴルフ",
    ],
    [
        "株式会社クリア（メンズクリア）",
        "https://mensclear.com/lp/a/ft/trial/",
        "リスティング広告",
        "0120-993-676",
        "メンズ全身脱毛（通い放題・全国120店舗以上）",
        "メンズ脱毛 男性脱毛 全身脱毛 脱毛通い放題 メンズクリア",
        "美容クリニック・脱毛・エステ",
    ],
    [
        "イオンモバイル株式会社",
        "https://aeonmobile.jp/column/lp/checklist.html",
        "リスティング広告",
        "0120-025-260",
        "格安SIM・格安スマホ（全国イオン店舗対応・通話かけ放題プランあり）",
        "格安SIM イオンモバイル 格安スマホ 乗り換え SIM スマホ料金節約",
        "格安スマホ・通信",
    ],
    [
        "株式会社Jリスクマネージメント（保険見直しラボ）",
        "https://www.hoken-minaoshi-lab.jp/lp/review1hr/",
        "リスティング広告",
        "0120-222-785",
        "生命保険・医療保険の無料FP相談（全国60拠点・平均経験17年のFP）",
        "保険見直し 保険相談 生命保険 医療保険 FP相談 保険無料相談",
        "保険・金融・証券・カード",
    ],
    [
        "株式会社MS-Japan",
        "https://www.jmsc.co.jp/lp/001/",
        "リスティング広告",
        "",
        "管理部門・士業特化の転職エージェント（経理・人事・法務・会計士・税理士）",
        "経理 転職 管理部門 転職 士業 転職 転職エージェント 会計士 転職",
        "転職・就職・人材派遣",
    ],
    [
        "株式会社ベアーズ",
        "https://www.happy-bears.com/lp/kaji_trial/",
        "リスティング広告",
        "0120-552-445",
        "家事代行サービス（初回お試しプラン・全国対応・スタッフ5,000名以上）",
        "家事代行 家政婦 家事代行サービス 家事サポート ハウスクリーニング 家事",
        "家事代行・クリーニング",
    ],
    [
        "株式会社エコリング",
        "https://www.eco-ring.com/lp/consumerlp",
        "リスティング広告",
        "",
        "ブランド品・貴金属・不用品の高価買取（なんでも買取・全国300店舗以上）",
        "買取 ブランド買取 金買取 不用品買取 貴金属買取 エコリング",
        "買取・リユース・フリマ",
    ],
    [
        "みんなのマーケット株式会社（くらしのマーケット）",
        "https://curama.jp/lp/user/osouji/",
        "リスティング広告",
        "",
        "出張・訪問サービスマッチング（ハウスクリーニング・不用品回収・引越しなど）",
        "くらしのマーケット ハウスクリーニング エアコン掃除 不用品回収 引越し 家事代行",
        "家事代行・クリーニング",
    ],
    [
        "株式会社ジェイック（就職カレッジ）",
        "https://www.jaic-college.jp/lp/m/landing03_02m.html",
        "リスティング広告",
        "",
        "フリーター・既卒・第二新卒向け就職支援（無料・書類不要・集団面接会付き）",
        "就職支援 フリーター 就職 第二新卒 既卒 就職カレッジ ジェイック",
        "転職・就職・人材派遣",
    ],
    [
        "ナッシュ株式会社（nosh）",
        "https://nosh.jp/lp/diet2",
        "リスティング広告",
        "",
        "低糖質・低塩分の冷凍宅配弁当（定期便・管理栄養士監修・125種以上）",
        "宅配弁当 冷凍弁当 低糖質 宅食 ナッシュ ダイエット弁当 食事宅配",
        "食品EC・宅配弁当・ミールキット",
    ],
]


def make_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "架電リスト"

    header_font = Font(name="メイリオ", size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    fill_odd = PatternFill("solid", fgColor="F2F7FC")
    fill_even = PatternFill("solid", fgColor="FFFFFF")
    data_font = Font(name="メイリオ", size=9)
    data_align = Alignment(vertical="center", wrap_text=False)
    url_font = Font(name="メイリオ", size=9, color="0563C1", underline="single")

    thin = Side(style="thin", color="D0D7DE")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # ヘッダー行
    for col_idx, col_name in enumerate(HEADER, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = border

    ws.row_dimensions[1].height = 22

    # データ行
    for row_idx, row_data in enumerate(DATA, start=2):
        fill = fill_odd if (row_idx % 2 == 0) else fill_even
        for col_idx, value in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.fill = fill
            cell.border = border
            if col_idx == 2:  # LP URL列
                cell.font = url_font
                cell.hyperlink = value
                cell.alignment = data_align
            else:
                cell.font = data_font
                cell.alignment = data_align
        ws.row_dimensions[row_idx].height = 18

    # 列幅
    for col_idx, width in enumerate(COL_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # 先頭行固定・オートフィルター
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADER))}1"

    wb.save(OUTPUT_PATH)
    print(f"保存完了: {OUTPUT_PATH}")
    print(f"件数: {len(DATA)}社")


if __name__ == "__main__":
    make_xlsx()
