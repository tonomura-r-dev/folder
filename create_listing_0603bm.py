import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUTPUT_PATH = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー\架電リスト_リスティング_2026-06-03bm.xlsx"
HEADER = ["企業名", "LP URL", "広告種別", "法人番号", "商材", "検索KW", "業界"]
COL_WIDTHS = [32, 56, 16, 16, 22, 32, 22]
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="メイリオ", size=10, bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
ROW_FILLS = [PatternFill("solid", fgColor="F2F7FC"), PatternFill("solid", fgColor="FFFFFF")]
ROW_FONT = Font(name="メイリオ", size=9)
ROW_ALIGN = Alignment(vertical="center", wrap_text=False)
URL_FONT = Font(name="メイリオ", size=9, color="0563C1", underline="single")

DATA = [
    # ── プロジェクト管理 ──
    ("Teamwork.com",               "https://www.teamwork.com/pricing/",                       "リスティング広告", "", "Teamwork・クライアントワーク向けPM",          "プロジェクト管理 SaaS 比較, Teamwork 料金 評判, クライアント 管理", "プロジェクト管理SaaS"),
    # ── 業務自動化 ──
    ("Zapier Inc.",                "https://zapier.com/pricing",                              "リスティング広告", "", "Zapier・業務自動化・9000+アプリ連携",          "業務 自動化 ツール, Zapier 料金 評判, アプリ 連携 ノーコード",      "業務自動化SaaS"),
    # ── 業務自動化 ──
    ("Make (Celonis)",            "https://www.make.com/en/pricing",                         "リスティング広告", "", "Make・ビジュアル業務自動化・AIワークフロー",  "ノーコード 自動化 SaaS, Make Integromat 料金, ワークフロー 自動化", "業務自動化SaaS"),
    # ── ヘルプデスク ──
    ("Help Scout PBC",            "https://www.helpscout.com/pricing/",                      "リスティング広告", "", "Help Scout・カスタマーサポート・低価格",      "ヘルプデスク SaaS 比較, Help Scout 料金 評判, カスタマー サポート", "カスタマーサポートSaaS"),
    # ── デジタルアダプション ──
    ("Whatfix Inc.",              "https://whatfix.com/pricing/",                            "リスティング広告", "", "Whatfix・デジタルアダプション・操作ガイド",   "デジタルアダプション ツール, Whatfix 料金 評判, システム 定着 SaaS","業務DXSaaS"),
    # ── CRM ──
    ("Insightly Inc.",            "https://www.insightly.com/comparison-lp-sfdc/",           "リスティング広告", "", "Insightly・CRM・営業/マーケ/サービス統合",    "CRM SaaS 比較, Insightly 料金 評判, 中小 企業 CRM",                "CRM SaaS"),
    # ── ワークフロー自動化 ──
    ("Pipefy Inc.",               "https://www.pipefy.com/pricing/",                         "リスティング広告", "", "Pipefy・ノーコードワークフロー・AI自動化",    "ワークフロー 自動化 SaaS, Pipefy 料金 評判, 業務 プロセス 管理",   "業務自動化SaaS"),
    # ── カスタマーサービスCRM ──
    ("Kustomer LLC",              "https://www.kustomer.com/pricing/",                       "リスティング広告", "", "Kustomer・CRM型カスタマーサービス・AI",       "カスタマーサービス SaaS, Kustomer 料金 評判, CS プラットフォーム",  "カスタマーサポートSaaS"),
    # ── マーケティングオートメーション ──
    ("ActiveCampaign LLC",        "https://www.activecampaign.com/pricing",                  "リスティング広告", "", "ActiveCampaign・MA・メール・自動化",          "マーケ オートメーション SaaS, ActiveCampaign 料金 評判, MA ツール", "マーケSaaS"),
    # ── クリエイター向けメール ──
    ("Kit (ConvertKit Inc.)",     "https://kit.com/pricing",                                 "リスティング広告", "", "Kit・クリエイター向けメール配信・自動化",      "クリエイター メール 配信, Kit ConvertKit 料金, ニュースレター SaaS","メールマーケSaaS"),
    # ── チャットサポート ──
    ("Crisp IM SAS",              "https://crisp.chat/en/pricing/",                          "リスティング広告", "", "Crisp・AIチャットサポート・マルチチャネル",   "チャット サポート SaaS, Crisp 料金 評判, ライブ チャット ツール",   "カスタマーサポートSaaS"),
    # ── ドキュメント/コラボ ──
    ("Coda Inc.",                 "https://coda.io/pricing",                                 "リスティング広告", "", "Coda・ドキュメント+DB・オールインワン",       "ドキュメント コラボ SaaS, Coda 料金 評判, Notion 代替 ツール",     "コラボSaaS"),
    # ── チームチャット ──
    ("Chanty Inc.",               "https://www.chanty.com/pricing/",                         "リスティング広告", "", "Chanty・チームチャット・低価格Slack代替",     "チーム チャット SaaS 比較, Chanty 料金 評判, ビジネス チャット 安い","コラボSaaS"),
    # ── フォーム作成 ──
    ("Jotform Inc.",              "https://www.jotform.com/pricing/",                        "リスティング広告", "", "Jotform・フォーム作成・10000テンプレ",        "フォーム 作成 SaaS 比較, Jotform 料金 評判, オンライン フォーム",   "フォームSaaS"),
    # ── 提案書・電子署名 ──
    ("PandaDoc Inc.",             "https://www.pandadoc.com/pricing",                        "リスティング広告", "", "PandaDoc・提案書/見積/電子署名・40%短縮",    "提案書 作成 SaaS, PandaDoc 料金 評判, 電子署名 ドキュメント",       "電子署名SaaS"),
    # ── セールスCRM ──
    ("Close (Elastic Inc.)",      "https://www.close.com/pricing",                           "リスティング広告", "", "Close・営業特化CRM・通話/メール統合",         "営業 CRM SaaS, Close 料金 評判, インサイドセールス ツール",         "CRM SaaS"),
    # ── ノーコードサイト ──
    ("Framer B.V.",               "https://www.framer.com/pricing",                          "リスティング広告", "", "Framer・ノーコードWebサイト・デザイナー向け", "ノーコード サイト 作成, Framer 料金 評判, Web デザイン ツール",     "ノーコードSaaS"),
    # ── ノーコードDB ──
    ("Baserow B.V.",              "https://baserow.io/pricing",                              "リスティング広告", "", "Baserow・オープンソースDB・Airtable代替",     "ノーコード DB SaaS, Baserow 料金 評判, データベース 作成 ツール",   "ノーコードSaaS"),
    # ── 日程調整 ──
    ("Doodle AG",                 "https://doodle.com/premium/business/info",                "リスティング広告", "", "Doodle・グループ日程調整・投票",              "日程調整 グループ ツール, Doodle 料金 評判, 予定 調整 SaaS",        "業務DXSaaS"),
    # ── 動画セールス ──
    ("Vidyard Inc.",              "https://www.vidyard.com/pricing/",                        "リスティング広告", "", "Vidyard・動画セールス・AI動画・分析",         "動画 セールス SaaS, Vidyard 料金 評判, 画面録画 ビジネス ツール",   "動画マーケSaaS"),
]


def make_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "架電リスト"

    for col_idx, header in enumerate(HEADER, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
    ws.row_dimensions[1].height = 20

    for row_idx, row_data in enumerate(DATA, 2):
        fill = ROW_FILLS[(row_idx - 2) % 2]
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.fill = fill
            cell.alignment = ROW_ALIGN
            if col_idx == 2:
                cell.hyperlink = value
                cell.font = URL_FONT
            else:
                cell.font = ROW_FONT
        ws.row_dimensions[row_idx].height = 16

    for col_idx, width in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADER))}1"

    wb.save(OUTPUT_PATH)
    print(f"Saved: {OUTPUT_PATH}")
    print(f"Rows: {len(DATA)}")


if __name__ == "__main__":
    make_xlsx()
