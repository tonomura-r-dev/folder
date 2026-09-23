import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUTPUT_PATH = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー\架電リスト_リスティング_2026-06-08_live.xlsx"
HEADER = ["企業名", "LP URL", "広告種別", "法人番号", "商材", "検索KW", "業界"]
COL_WIDTHS = [30, 50, 16, 16, 26, 26, 18]
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="メイリオ", size=10, bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
ROW_FILLS = [PatternFill("solid", fgColor="F2F7FC"), PatternFill("solid", fgColor="FFFFFF")]
ROW_FONT = Font(name="メイリオ", size=9)
ROW_ALIGN = Alignment(vertical="center", wrap_text=False)
URL_FONT = Font(name="メイリオ", size=9, color="0563C1", underline="single")

# 全社：Google検索の「スポンサー」枠（今配信中の広告）から抽出した広告主のみ。
# 比較/ポータル/アグリゲーター・代理店・使用済みは除外。企業名はカッコなし。
DATA = [
    ("株式会社オスカーキャピタル",   "https://oscar-capital.net/3g7/",                 "リスティング広告", "", "不動産投資セミナー",            "不動産投資 セミナー",        "不動産投資"),
    ("株式会社クレアスライフ",       "https://sales.clearthlife.co.jp",                "リスティング広告", "", "投資用マンション",              "不動産投資 マンション",      "不動産投資"),
    ("三井ホーム株式会社",           "https://www.mitsuihome.co.jp/lp/works04_05/",    "リスティング広告", "", "注文住宅・カタログ請求",        "注文住宅 カタログ",          "注文住宅"),
    ("株式会社ヤマジホーム",         "https://zesttoku.com/",                          "リスティング広告", "", "注文住宅（徳島）",              "注文住宅 徳島",              "注文住宅"),
    ("株式会社ラッフルズホーム",     "https://www.raffles-home.jp",                    "リスティング広告", "", "注文住宅（徳島）",              "注文住宅 徳島",              "注文住宅"),
    ("株式会社かみくぼ住宅",         "https://kamikuboreform.com/",                    "リスティング広告", "", "リフォーム（徳島）",            "リフォーム 徳島",            "リフォーム"),
    ("株式会社青木商店",             "https://www.aokishoten.net/",                    "リスティング広告", "", "内窓・窓リフォーム（徳島）",    "内窓 補助金",                "リフォーム"),
    ("株式会社創建",                 "https://www.k-skn.com/sokenpaint/",              "リスティング広告", "", "外壁塗装（創建ペイント）",      "外壁塗装 見積もり",          "外壁塗装"),
    ("株式会社モノクローム",         "https://www.monochrome.so/roof",                 "リスティング広告", "", "屋根一体型太陽光パネル",        "太陽光 蓄電池",              "太陽光"),
    ("株式会社ライフワン",           "https://www.k-chikudenchi.com/ga-lp01/",         "リスティング広告", "", "系統用蓄電池・太陽光",          "太陽光 蓄電池",              "太陽光"),
    ("株式会社YAMABISHI",            "https://www.yamabishi.co.jp/ad/yrw/",            "リスティング広告", "", "産業用蓄電システム",            "蓄電池 産業用",              "蓄電池"),
    ("ニチコン株式会社",             "https://www.nichicon.co.jp",                     "リスティング広告", "", "家庭用蓄電池",                  "蓄電池 家庭用",              "蓄電池"),
    ("株式会社はなまる",             "https://www.hanamaru870.net/lp/car/h/",          "リスティング広告", "", "車買取（ソコカラ）",            "車買取",                    "車買取"),
    ("株式会社アレシア",             "https://bestfactor.jp",                          "リスティング広告", "", "ファクタリング（ベストファクター）","ファクタリング 即日",      "ファクタリング"),
    ("株式会社ラボル",               "https://www.labol.co.jp",                        "リスティング広告", "", "ファクタリング（labol）",       "ファクタリング 即日",        "ファクタリング"),
    ("キャッシュファクター",         "https://cash-factor.com/",                       "リスティング広告", "", "ファクタリング ※運営会社名 確認中","ファクタリング 即日",      "ファクタリング"),
    ("株式会社AGE technologies",     "https://so-zo-ku.com/",                          "リスティング広告", "", "相続手続き（そうぞくドットコム）","相続 相談",                "相続・士業"),
    ("ムスベル株式会社",             "https://www.musbell.com",                        "リスティング広告", "", "結婚相談所",                    "結婚相談所",                "婚活"),
    ("有限会社ベルベ",               "https://bellbe.net/",                            "リスティング広告", "", "葬儀・家族葬（徳島）",          "家族葬 徳島",                "葬儀"),
    ("南無なむ堂",                   "https://mercie.jp/",                             "リスティング広告", "", "家族葬メルシエ（徳島）",        "家族葬 徳島",                "葬儀"),
]


def make_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "架電リスト"
    for col_idx, header in enumerate(HEADER, 1):
        c = ws.cell(row=1, column=col_idx, value=header)
        c.fill = HEADER_FILL; c.font = HEADER_FONT; c.alignment = HEADER_ALIGN
    ws.row_dimensions[1].height = 20
    for row_idx, row in enumerate(DATA, 2):
        fill = ROW_FILLS[(row_idx - 2) % 2]
        for col_idx, value in enumerate(row, 1):
            c = ws.cell(row=row_idx, column=col_idx, value=value)
            c.fill = fill; c.alignment = ROW_ALIGN
            if col_idx == 2:
                c.hyperlink = value; c.font = URL_FONT
            else:
                c.font = ROW_FONT
        ws.row_dimensions[row_idx].height = 16
    for col_idx, width in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADER))}1"
    wb.save(OUTPUT_PATH)
    print(f"Saved: {OUTPUT_PATH}  Rows: {len(DATA)}")


if __name__ == "__main__":
    make_xlsx()
