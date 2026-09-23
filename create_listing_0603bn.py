import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUTPUT_PATH = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー\架電リスト_リスティング_2026-06-03bn.xlsx"
HEADER = ["企業名", "LP URL", "広告種別", "法人番号", "商材", "検索KW", "業界"]
COL_WIDTHS = [32, 56, 16, 16, 22, 32, 22]
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="メイリオ", size=10, bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
ROW_FILLS = [PatternFill("solid", fgColor="F2F7FC"), PatternFill("solid", fgColor="FFFFFF")]
ROW_FONT = Font(name="メイリオ", size=9)
ROW_ALIGN = Alignment(vertical="center", wrap_text=False)
URL_FONT = Font(name="メイリオ", size=9, color="0563C1", underline="single")

DATA = [
    # ── サイト作成 ──
    ("Squarespace Inc.",           "https://www.squarespace.com/pricing",                     "リスティング広告", "", "Squarespace・サイト作成・EC・決済対応",       "サイト作成 ツール 比較, Squarespace 料金 評判, ホームページ 制作", "ノーコードSaaS"),
    # ── EC構築 ──
    ("BigCommerce Inc.",           "https://www.bigcommerce.com/essentials/pricing/",         "リスティング広告", "", "BigCommerce・EC構築・組み込み機能充実",        "EC 構築 SaaS 比較, BigCommerce 料金 評判, ネットショップ 開業",    "EC SaaS"),
    # ── 監視/可視化 ──
    ("Grafana Labs Inc.",          "https://grafana.com/pricing/",                            "リスティング広告", "", "Grafana・監視/可視化・オブザーバビリティ",     "監視 可視化 ツール, Grafana 料金 評判, ダッシュボード SaaS",       "DevOps SaaS"),
    # ── パスワード管理 ──
    ("1Password (AgileBits Inc.)", "https://1password.com/jp/pricing/",                       "リスティング広告", "", "1Password・法人パスワード管理・SSO",           "パスワード 管理 法人, 1Password 料金 評判, パスワードマネージャー", "セキュリティSaaS"),
    # ── パスワード管理 ──
    ("Dashlane Inc.",              "https://www.dashlane.com/pricing",                        "リスティング広告", "", "Dashlane・法人パスワード管理・ダークウェブ監視","パスワード 管理 SaaS 比較, Dashlane 料金 評判, 認証情報 管理",     "セキュリティSaaS"),
    # ── チームチャット ──
    ("Mattermost Inc.",            "https://mattermost.com/pricing/",                         "リスティング広告", "", "Mattermost・OSSチームチャット・自社運用",      "オープンソース チャット, Mattermost 料金 評判, Slack 代替 自社",  "コラボSaaS"),
    # ── クラウドストレージ ──
    ("pCloud AG",                  "https://www.pcloud.com/ja/cloud-storage-pricing-plans.html","リスティング広告", "", "pCloud・クラウドストレージ・買い切り",        "クラウド ストレージ 買い切り, pCloud 料金 評判, ファイル 保存",   "クラウドストレージSaaS"),
    # ── ナレッジ管理 ──
    ("Nuclino GmbH",               "https://www.nuclino.com/pricing",                         "リスティング広告", "", "Nuclino・軽量ナレッジ/ドキュメント・wiki",     "社内 wiki ツール 軽量, Nuclino 料金 評判, ナレッジ 共有 SaaS",    "コラボSaaS"),
    # ── プロジェクト管理 ──
    ("Basecamp (37signals LLC)",   "https://basecamp.com/pricing",                            "リスティング広告", "", "Basecamp・プロジェクト管理・チーム連携",       "プロジェクト管理 SaaS, Basecamp 料金 評判, タスク 共有 ツール",   "プロジェクト管理SaaS"),
    # ── プロダクト分析 ──
    ("PostHog Inc.",               "https://posthog.com/pricing",                             "リスティング広告", "", "PostHog・プロダクト分析・OSS・無料枠",         "プロダクト 分析 OSS, PostHog 料金 評判, セッション リプレイ SaaS","プロダクト分析SaaS"),
    # ── 通信API ──
    ("Vonage (Ericsson)",          "https://www.vonage.com/communications-apis/sms/pricing/", "リスティング広告", "", "Vonage・SMS/音声API・通信プラットフォーム",   "SMS API 比較, Vonage Nexmo 料金, 通信 API SaaS",                  "コミュニケーションAPI"),
    # ── 日程調整 ──
    ("Cal.com Inc.",               "https://cal.com/pricing",                                 "リスティング広告", "", "Cal.com・OSS日程調整・チームスケジュール",     "日程調整 OSS 自社, Cal.com 料金 評判, 予約 自動化 SaaS",          "業務DXSaaS"),
    # ── EC向けサポート ──
    ("Gorgias Inc.",               "https://www.gorgias.com/pricing",                         "リスティング広告", "", "Gorgias・EC向けカスタマーサポート・Shopify",  "EC カスタマーサポート SaaS, Gorgias 料金 評判, Shopify サポート", "カスタマーサポートSaaS"),
    # ── ECメールマーケ ──
    ("Omnisend",                   "https://www.omnisend.com/pricing/",                       "リスティング広告", "", "Omnisend・EC向けメール&SMS・無料プラン",       "EC メールマーケ SaaS, Omnisend 料金 評判, Shopify メール 配信",  "メールマーケSaaS"),
    # ── プロダクト体験 ──
    ("Userpilot Inc.",             "https://userpilot.com/pricing",                           "リスティング広告", "", "Userpilot・プロダクト体験・オンボーディング", "プロダクト オンボーディング SaaS, Userpilot 料金 評判, UX 改善",  "プロダクトSaaS"),
    # ── ECチャットボット ──
    ("Tidio LLC",                  "https://www.tidio.com/pricing/",                          "リスティング広告", "", "Tidio・AIチャットボット・EC向けライブチャット","チャットボット EC SaaS, Tidio 料金 評判, ライブ チャット ツール",  "カスタマーサポートSaaS"),
    # ── AIタスク管理 ──
    ("Taskade Inc.",               "https://www.taskade.com/pricing",                         "リスティング広告", "", "Taskade・AIタスク管理・エージェント",         "AI タスク 管理 SaaS, Taskade 料金 評判, ワークスペース ツール",   "プロジェクト管理SaaS"),
    # ── ワークマネジメント ──
    ("SmartSuite LLC",             "https://www.smartsuite.com/pricing",                      "リスティング広告", "", "SmartSuite・ワークマネジメント・ノーコード",  "ワーク マネジメント SaaS, SmartSuite 料金 評判, Airtable 代替",   "ノーコードSaaS"),
    # ── アンケート ──
    ("SurveyMonkey (Momentive)",   "https://jp.surveymonkey.com/pricing/",                    "リスティング広告", "", "SurveyMonkey・アンケート作成・AI・400テンプレ","アンケート 作成 ツール, SurveyMonkey 料金 評判, 調査 フォーム",   "フォームSaaS"),
    # ── SaaS決済 ──
    ("Paddle.com Market Ltd.",     "https://www.paddle.com/pricing",                          "リスティング広告", "", "Paddle・SaaS決済・サブスク・税務コンプラ",    "SaaS 決済 プラットフォーム, Paddle 料金 評判, サブスク 課金 税務","フィンテックSaaS"),
]


def make_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "架電リスト"

    for col_idx, header in enumerate(HEADER, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
    ws.row_dimensions[1].height = 20

    for row_idx, row_data in enumerate(DATA, 2):
        fill = ROW_FILLS[(row_idx - 2) % 2]
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.fill = fill
            cell.alignment = ROW_ALIGN
            if col_idx == 2:
                cell.hyperlink = value
                cell.font = URL_FONT
            else:
                cell.font = ROW_FONT
        ws.row_dimensions[row_idx].height = 16

    for col_idx, width in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADER))}1"

    wb.save(OUTPUT_PATH)
    print(f"Saved: {OUTPUT_PATH}")
    print(f"Rows: {len(DATA)}")


if __name__ == "__main__":
    make_xlsx()
