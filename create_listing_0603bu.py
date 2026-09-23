import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUTPUT_PATH = r"C:\Users\tonomura-r\Downloads\との\新しいフォルダー\架電リスト_リスティング_2026-06-03bu.xlsx"
HEADER = ["企業名", "LP URL", "広告種別", "法人番号", "商材", "検索KW", "業界"]
COL_WIDTHS = [32, 56, 16, 16, 22, 32, 22]
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="メイリオ", size=10, bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
ROW_FILLS = [PatternFill("solid", fgColor="F2F7FC"), PatternFill("solid", fgColor="FFFFFF")]
ROW_FONT = Font(name="メイリオ", size=9)
ROW_ALIGN = Alignment(vertical="center", wrap_text=False)
URL_FONT = Font(name="メイリオ", size=9, color="0563C1", underline="single")

DATA = [
    # ── レビュー/信頼 ──
    ("Trustpilot",                 "https://business.trustpilot.com/plans-pricing",           "リスティング広告", "", "Trustpilot・レビュー/信頼・口コミ収集",         "レビュー 収集 ツール, Trustpilot 料金 評判, 口コミ プラットフォーム","レビューSaaS"),
    # ── EC向けレビュー/ロイヤルティ ──
    ("Yotpo",                      "https://www.yotpo.com/pricing/",                          "リスティング広告", "", "Yotpo・EC向けレビュー/ロイヤルティ/SMS",       "EC レビュー SaaS, Yotpo 料金 評判, UGC ロイヤルティ ツール",       "EC SaaS"),
    # ── Shopifyレビュー ──
    ("Okendo",                     "https://www.okendo.io/pricing",                           "リスティング広告", "", "Okendo・Shopifyレビュー/サーベイ/顧客마케",    "Shopify レビュー アプリ, Okendo 料金 評判, EC レビュー ツール",     "EC SaaS"),
    # ── ECカスタマーサポート ──
    ("Re:amaze",                   "https://www.reamaze.com/pricing",                         "リスティング広告", "", "Re:amaze・ECカスタマーサポート/チャット",      "EC カスタマーサポート SaaS, Reamaze 料金 評判, ヘルプデスク EC",   "カスタマーサポートSaaS"),
    # ── ECカスタマーサポート ──
    ("Richpanel",                  "https://www.richpanel.com/pricing",                       "リスティング広告", "", "Richpanel・ECカスタマーサポート/自己解決",     "EC サポート SaaS, Richpanel 料金 評判, AI カスタマーサポート",     "カスタマーサポートSaaS"),
    # ── レビュー/評判管理 ──
    ("Birdeye",                    "https://birdeye.com/pricing/",                            "リスティング広告", "", "Birdeye・評判管理/レビュー/口コミ獲得",        "評判 管理 SaaS, Birdeye 料金 評判, 口コミ 獲得 ツール",            "レビューSaaS"),
    # ── レビュー ──
    ("Reviews.io",                 "https://www.reviews.io/pricing",                          "リスティング広告", "", "Reviews.io・レビュー収集/UGC・EC向け",         "レビュー 収集 SaaS, Reviews.io 料金 評判, EC レビュー ツール",     "レビューSaaS"),
    # ── Shopifyレビュー ──
    ("Loox",                       "https://loox.app/pricing",                                "リスティング広告", "", "Loox・Shopify写真レビュー/UGC",               "Shopify 写真 レビュー, Loox 料金 評判, EC レビュー アプリ",        "EC SaaS"),
    # ── レビュー/ロイヤルティ ──
    ("Stamped",                    "https://stamped.io/pricing",                              "リスティング広告", "", "Stamped・レビュー/ロイヤルティ・EC向け",       "EC レビュー ロイヤルティ, Stamped 料金 評判, UGC ツール",          "EC SaaS"),
    # ── EC向けSMSマーケ ──
    ("Postscript",                 "https://postscript.io/pricing",                           "リスティング広告", "", "Postscript・Shopify向けSMSマーケティング",     "EC SMS マーケ SaaS, Postscript 料金 評判, Shopify SMS 配信",       "メールマーケSaaS"),
    # ── SMSマーケ ──
    ("Attentive",                  "https://www.attentive.com/pricing",                       "リスティング広告", "", "Attentive・SMSマーケティング・EC向け",         "SMS マーケ SaaS, Attentive 料金 評判, テキスト マーケ ツール",     "メールマーケSaaS"),
    # ── EC向けメール/SMS ──
    ("Sendlane",                   "https://www.sendlane.com/pricing",                        "リスティング広告", "", "Sendlane・EC向けメール/SMS統合マーケ",         "EC メール SMS SaaS, Sendlane 料金 評判, マーケ オートメーション",  "メールマーケSaaS"),
    # ── メールマーケ ──
    ("Drip",                       "https://www.drip.com/pricing",                            "リスティング広告", "", "Drip・EC向けメールマーケティング自動化",       "EC メールマーケ SaaS, Drip 料金 評判, メール 自動化 ツール",       "メールマーケSaaS"),
    # ── コンバージョン最適化 ──
    ("Privy",                      "https://www.privy.com/pricing",                           "リスティング広告", "", "Privy・ポップアップ/メール・EC集客",           "ポップアップ EC SaaS, Privy 料金 評判, EC 集客 ツール",            "EC SaaS"),
    # ── コンバージョン最適化 ──
    ("Justuno",                    "https://www.justuno.com/pricing/",                        "リスティング広告", "", "Justuno・ポップアップ/CRO・リード獲得",        "CRO ポップアップ SaaS, Justuno 料金 評判, リード 獲得 ツール",     "EC SaaS"),
    # ── リード獲得 ──
    ("OptinMonster",               "https://optinmonster.com/pricing/",                       "リスティング広告", "", "OptinMonster・ポップアップ/リード獲得",        "リード 獲得 ポップアップ, OptinMonster 料金 評判, CV 改善 ツール",  "マーケSaaS"),
    # ── EC向けパーソナライズ ──
    ("Rebuy",                      "https://www.rebuyengine.com/pricing",                     "リスティング広告", "", "Rebuy・EC向けパーソナライズ/アップセル",       "EC パーソナライズ SaaS, Rebuy 料金 評判, アップセル ツール",       "EC SaaS"),
    # ── サブスクリプションEC ──
    ("ReCharge",                   "https://getrecharge.com/pricing/",                        "リスティング広告", "", "ReCharge・EC定期購入/サブスク管理",            "EC サブスク SaaS, ReCharge 料金 評判, 定期購入 管理 ツール",       "EC SaaS"),
    # ── 配送追跡 ──
    ("AfterShip",                  "https://www.aftership.com/pricing",                       "リスティング広告", "", "AfterShip・配送追跡/購入後体験・EC向け",       "配送 追跡 SaaS, AfterShip 料金 評判, EC 配送 通知 ツール",         "EC SaaS"),
    # ── ロイヤルティ ──
    ("Smile.io",                   "https://smile.io/pricing",                                "リスティング広告", "", "Smile.io・EC向けロイヤルティ/リワード",        "EC ロイヤルティ SaaS, Smile.io 料金 評判, ポイント リワード ツール","EC SaaS"),
]


def make_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "架電リスト"

    for col_idx, header in enumerate(HEADER, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
    ws.row_dimensions[1].height = 20

    for row_idx, row_data in enumerate(DATA, 2):
        fill = ROW_FILLS[(row_idx - 2) % 2]
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.fill = fill
            cell.alignment = ROW_ALIGN
            if col_idx == 2:
                cell.hyperlink = value
                cell.font = URL_FONT
            else:
                cell.font = ROW_FONT
        ws.row_dimensions[row_idx].height = 16

    for col_idx, width in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADER))}1"

    wb.save(OUTPUT_PATH)
    print(f"Saved: {OUTPUT_PATH}")
    print(f"Rows: {len(DATA)}")


if __name__ == "__main__":
    make_xlsx()
